# tests/torchcell/datasets/ecoli/test_mohiuddin2022.py
# [[tests.torchcell.datasets.ecoli.test_mohiuddin2022]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_mohiuddin2022.py
"""``torchcell/datasets/ecoli/mohiuddin2022.py``: the sheet readers, the fold-change
reproduction, the promoter resolution, the records, the raw mirror and the build.

Everything here is hermetic. The workbook is written at a SMALL shape (four wells, the
real four arms and the real nine read hours) and the module's grid constants are pinned
to it, the resolution tests read the REAL ``EcoliK12MG1655Genome`` over the synthetic
MG1655 assembly of ``tests/torchcell/sequence/genome/_bacterial_fixtures.py`` with the
network refused, and the end-to-end build runs into a ``tmp_path`` DATA_ROOT. The
``--data`` tests at the bottom read the dev-tree store the real release built.

The two properties worth pinning hardest are the ones the design turns on: that the
released ``Fold Change`` sheet IS the stored readings' own ratio (so it is not stored),
and that a treated arm's pre-dose reads carry no environmental edit while its reads from
hour five carry exactly one.
"""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.ecoli.mohiuddin2022 as m
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    MG1655_GAF,
    MG1655_LOCI,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.datamodels.schema import (
    EXPERIMENT_REFERENCE_TYPE_MAP,
    EXPERIMENT_TYPE_MAP,
    ConcentrationUnit,
    DoseBasis,
    HeterologousPathwayPerturbation,
    PromoterActivityExperiment,
    PromoterActivityExperimentReference,
    ReporterReadout,
    SampleUnit,
    SmallMoleculePerturbation,
)
from torchcell.datasets.dataset_registry import dataset_registry
from torchcell.sequence.genome.ecoli.k12 import MG1655_ASSEMBLY, EcoliK12MG1655Genome
from torchcell.verification.report import Level
from torchcell.verification.sourced import ProvenanceGapReason

# --------------------------------------------------------------------------- #
# The synthetic release: four wells on two plates
# --------------------------------------------------------------------------- #
#: ``(plate, well, promoter)``. ``thrA`` heads two wells, which is what the real sheet
#: does nine times over, and ``U66`` names no gene at all.
SYNTHETIC_WELLS: tuple[tuple[str, str, str], ...] = (
    ("AZ01", "A1", "thrL"),
    ("AZ01", "A2", "thrA"),
    ("AZ02", "A1", "thrA"),
    ("AZ02", "A2", "U66"),
)


#: One reading per (well, arm, hour): a base that differs per well, arm and hour, so
#: every ratio the fold-change check forms is distinct and a transposed column shows up.
def _reading(well: int, arm: int, hour: float) -> float:
    return round(10.0 + well + 3.0 * arm + 0.5 * hour, 4)


def _write_release(path: Path) -> Path:
    """Write a workbook with the real sheet layout at the synthetic shape."""
    book = openpyxl.Workbook()
    raw = book.active
    raw.title = m.RAW_DATA_SHEET
    raw.cell(
        1, 1, "High-throughput screening of the Escherichia coli promoter library."
    )
    raw.cell(2, 1, m.SOURCED_VALUES["raw_sheet_legend"].quote)
    arm_position = {arm: index for index, arm in enumerate(m._RAW_BLOCKS)}
    for arm, (plate_c, well_c, promoter_c, first) in m._RAW_BLOCKS.items():
        raw.cell(4, first + 4, arm)
        raw.cell(m._HEADER_ROW, plate_c + 1, "Plate_Number")
        raw.cell(m._HEADER_ROW, well_c + 1, "Well")
        raw.cell(m._HEADER_ROW, promoter_c + 1, "Promoter_Name")
        for index, hour in enumerate(m.READ_HOURS):
            raw.cell(m._HEADER_ROW, first + index + 1, f"t={int(hour)}")
        for row_index, (plate, well, promoter) in enumerate(SYNTHETIC_WELLS):
            row = m._FIRST_DATA_ROW + row_index
            raw.cell(row, plate_c + 1, plate)
            raw.cell(row, well_c + 1, well)
            raw.cell(row, promoter_c + 1, promoter)
            for index, hour in enumerate(m.READ_HOURS):
                raw.cell(
                    row, first + index + 1, _reading(row_index, arm_position[arm], hour)
                )

    fold = book.create_sheet(m.FOLD_CHANGE_SHEET)
    fold.cell(1, 1, "Fold changes of GFP expression in antibiotic-treated cultures.")
    fold.cell(2, 1, m.SOURCED_VALUES["fold_change_rule"].quote)
    for arm, (promoter_c, first) in m._FC_BLOCKS.items():
        fold.cell(4, first + 4, arm)
        fold.cell(m._HEADER_ROW, promoter_c + 1, "Promoter_Name")
        for index, hour in enumerate(m.READ_HOURS):
            fold.cell(m._HEADER_ROW, first + index + 1, f"t={int(hour)}")
        for row_index, (_, _, promoter) in enumerate(SYNTHETIC_WELLS):
            row = m._FIRST_DATA_ROW + row_index
            fold.cell(row, promoter_c + 1, promoter)
            for index, hour in enumerate(m.READ_HOURS):
                treated = _reading(row_index, arm_position[arm], hour)
                control = _reading(row_index, 0, hour)
                fold.cell(row, first + index + 1, treated / control)
    book.save(path)
    return path


@pytest.fixture
def release(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """The synthetic workbook, with the module's grid constants pinned to its shape."""
    path = _write_release(tmp_path / m.DATA_FILE)
    monkeypatch.setattr(m, "EXPECTED_WELLS", len(SYNTHETIC_WELLS))
    monkeypatch.setattr(
        m, "EXPECTED_RECORDS", len(SYNTHETIC_WELLS) * len(m.ARMS) * len(m.READ_HOURS)
    )
    monkeypatch.setattr(
        m,
        "EXPECTED_FOLD_CHANGE_CELLS",
        len(SYNTHETIC_WELLS) * (len(m.ARMS) - 1) * len(m.READ_HOURS),
    )
    return path


# --------------------------------------------------------------------------- #
# The sheet readers
# --------------------------------------------------------------------------- #
def test_the_raw_sheet_reads_one_row_per_well_with_every_arm_and_hour(
    release: Path,
) -> None:
    frame = m.read_raw_sheet(release)
    assert len(frame) == len(SYNTHETIC_WELLS)
    assert list(frame["promoter"]) == [p for _, _, p in SYNTHETIC_WELLS]
    assert list(frame["plate"]) == [p for p, _, _ in SYNTHETIC_WELLS]
    for position, arm in enumerate(m._RAW_BLOCKS):
        for hour in m.READ_HOURS:
            assert frame[f"{arm}|t={int(hour)}"][0] == _reading(0, position, hour)
    # 3 identity columns + 4 arms x 9 hours.
    assert len(frame.columns) == 3 + len(m.ARMS) * len(m.READ_HOURS)


def test_the_fold_change_sheet_reads_three_arms_keyed_by_position(
    release: Path,
) -> None:
    frame = m.read_fold_change_sheet(release)
    assert len(frame) == len(SYNTHETIC_WELLS)
    assert list(frame["Ampicillin Treatment|promoter"]) == [
        p for _, _, p in SYNTHETIC_WELLS
    ]
    assert frame["Ampicillin Treatment|t=2"][0] == pytest.approx(
        _reading(0, 1, 2.0) / _reading(0, 0, 2.0)
    )


def test_a_sheet_name_that_is_not_the_released_pair_raises(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.active.title = "Sheet1"
    path = tmp_path / "wrong.xlsx"
    book.save(path)
    with pytest.raises(m.SheetLayoutError, match="expected sheets"):
        m.read_raw_sheet(path)


def test_a_missing_hour_column_raises(release: Path) -> None:
    book = openpyxl.load_workbook(release)
    book[m.RAW_DATA_SHEET].cell(m._HEADER_ROW, 4, "t=1")
    book.save(release)
    with pytest.raises(m.SheetLayoutError, match="expected hour headers"):
        m.read_raw_sheet(release)


def test_a_renamed_identity_column_raises(release: Path) -> None:
    book = openpyxl.load_workbook(release)
    book[m.RAW_DATA_SHEET].cell(m._HEADER_ROW, 3, "Gene")
    book.save(release)
    with pytest.raises(m.SheetLayoutError, match="is not 'Promoter_Name'"):
        m.read_raw_sheet(release)


def test_arms_naming_different_wells_on_one_row_raises(release: Path) -> None:
    """The four blocks are joined by position, so a row must name one well."""
    book = openpyxl.load_workbook(release)
    sheet = book[m.RAW_DATA_SHEET]
    sheet.cell(m._FIRST_DATA_ROW, m._RAW_BLOCKS["Ofloxacin Treatment"][2] + 1, "proB")
    book.save(release)
    with pytest.raises(m.SheetLayoutError, match="name different wells"):
        m.read_raw_sheet(release)


def test_a_non_numeric_reading_raises(release: Path) -> None:
    book = openpyxl.load_workbook(release)
    book[m.RAW_DATA_SHEET].cell(m._FIRST_DATA_ROW, 4, "n.d.")
    book.save(release)
    with pytest.raises(m.SheetLayoutError, match="is not a number"):
        m.read_raw_sheet(release)


def test_a_well_count_that_is_not_the_release_raises(release: Path) -> None:
    book = openpyxl.load_workbook(release)
    sheet = book[m.RAW_DATA_SHEET]
    sheet.delete_rows(m._FIRST_DATA_ROW)
    book.save(release)
    with pytest.raises(m.SheetLayoutError, match="expected 4 wells, got 3"):
        m.read_raw_sheet(release)


# --------------------------------------------------------------------------- #
# The fold-change reproduction: why the sheet is not stored
# --------------------------------------------------------------------------- #
def test_every_released_fold_change_is_the_stored_readings_ratio(release: Path) -> None:
    raw = m.read_raw_sheet(release)
    fold = m.read_fold_change_sheet(release)
    check = m.check_fold_change(raw, fold)
    assert check.n_cells == len(SYNTHETIC_WELLS) * 3 * len(m.READ_HOURS)
    assert check.n_within_tolerance == check.n_cells
    assert check.max_relative_error < m.FOLD_CHANGE_TOLERANCE
    assert check.promoter_order_matches is True


def test_a_fold_change_that_is_not_the_ratio_stops_the_build(release: Path) -> None:
    book = openpyxl.load_workbook(release)
    book[m.FOLD_CHANGE_SHEET].cell(m._FIRST_DATA_ROW, 2, 42.0)
    book.save(release)
    raw = m.read_raw_sheet(release)
    fold = m.read_fold_change_sheet(release)
    with pytest.raises(m.SheetLayoutError, match="are not the stored"):
        m.check_fold_change(raw, fold)


def test_a_reordered_fold_change_sheet_stops_the_build(release: Path) -> None:
    """The sheet keys only by promoter, so its row ORDER is the join."""
    book = openpyxl.load_workbook(release)
    sheet = book[m.FOLD_CHANGE_SHEET]
    sheet.cell(m._FIRST_DATA_ROW, 1, "thrA")
    sheet.cell(m._FIRST_DATA_ROW + 1, 1, "thrL")
    book.save(release)
    raw = m.read_raw_sheet(release)
    fold = m.read_fold_change_sheet(release)
    with pytest.raises(m.SheetLayoutError, match="not the Raw Data sheet's promoter"):
        m.check_fold_change(raw, fold)


# --------------------------------------------------------------------------- #
# Promoter identity
# --------------------------------------------------------------------------- #
@pytest.fixture
def mg1655(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12MG1655Genome:
    """The real MG1655 class over the synthetic assembly; network refused."""
    files = write_assembly(
        tmp_path / "tier", MG1655_ASSEMBLY, MG1655_LOCI, gaf_rows=MG1655_GAF
    )
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    root = tmp_path / "mg1655"
    root.mkdir()
    return EcoliK12MG1655Genome(genome_root=str(root), overwrite=False)


def test_a_promoter_label_resolves_to_its_locus_and_a_non_gene_to_none(
    release: Path, mg1655: EcoliK12MG1655Genome
) -> None:
    raw = m.read_raw_sheet(release)
    mapping, ledger = m.resolve_promoters(raw, mg1655, label="synthetic")
    assert mapping == {"thrL": "b0001", "thrA": "b0002", "U66": None}
    assert ledger.n_wells == len(SYNTHETIC_WELLS)
    assert ledger.n_labels == 3
    assert ledger.n_labels_resolved == 2
    assert ledger.unresolved_labels == ["U66"]
    assert ledger.labels_in_more_than_one_well == {"thrA": 2}
    assert ledger.reconciliation.assembly_set == MG1655_ASSEMBLY.assembly_set


# --------------------------------------------------------------------------- #
# The records
# --------------------------------------------------------------------------- #
def test_the_genotype_is_one_episomal_promoter_reporter() -> None:
    genotype = m.reporter_genotype("thrA")
    (perturbation,) = genotype.perturbations
    assert isinstance(perturbation, HeterologousPathwayPerturbation)
    assert perturbation.promoter_name == "thrA"
    assert perturbation.localization == "episomal_plasmid"
    assert perturbation.is_heterologous is True
    assert perturbation.systematic_gene_name == m.REPORTER_GENE == "gfp"
    assert perturbation.construct_name is None
    assert perturbation.source_organism == "unreported"
    assert perturbation.copy_number == 1.0


def test_the_untreated_arm_never_carries_an_environmental_edit() -> None:
    for hour in m.READ_HOURS:
        assert m.environment(m.UNTREATED, hour).perturbations == []


def test_a_treated_arms_predose_reads_carry_no_drug_and_its_later_reads_carry_one() -> (
    None
):
    (ampicillin,) = (a for a in m.ARMS if a.compound == "ampicillin")
    for hour in (2.0, 3.0, 4.0):
        assert m.environment(ampicillin, hour).perturbations == []
    for hour in (5.0, 6.0, 10.0):
        (edit,) = m.environment(ampicillin, hour).perturbations
        assert isinstance(edit, SmallMoleculePerturbation)
        assert edit.compound.name == "ampicillin"
        assert edit.concentration.value == 200.0
        assert edit.concentration.unit is ConcentrationUnit.ug_per_ml
        assert edit.concentration.basis is DoseBasis.fixed


def test_the_environment_carries_lb_miller_at_37_and_the_read_hour() -> None:
    env = m.environment(m.UNTREATED, 7.0)
    assert env.media.base_medium == "LB"
    assert env.temperature is not None
    assert env.temperature.value == 37.0
    assert env.aerobicity == "aerobic"
    assert env.duration_hours == 7.0


def test_the_phenotype_carries_the_reading_its_promoter_and_its_well() -> None:
    phenotype = m.activity_phenotype(12.98, "thrA", "b0002", "Untreated|AZ01|A2")
    assert phenotype.promoter_activity == 12.98
    assert phenotype.promoter_name == "thrA"
    assert phenotype.promoter_gene == "b0002"
    assert phenotype.well_id == "Untreated|AZ01|A2"
    assert phenotype.readout is ReporterReadout.plate_reader_fluorescence
    assert phenotype.reporter_gene == "gfp"
    assert phenotype.n_samples == 1
    assert phenotype.sample_unit is SampleUnit.biological_replicate
    assert phenotype.label_name == "promoter_activity"
    assert phenotype.graph_level == "node"


def test_the_screen_ran_once_so_every_dispersion_is_a_typed_gap() -> None:
    phenotype = m.activity_phenotype(12.98, "thrA", "b0002", "Untreated|AZ01|A2")
    assert phenotype.promoter_activity_uncertainty is None
    assert phenotype.promoter_activity_uncertainty_type is None
    assert phenotype.promoter_activity_se is None
    assert {gap.field for gap in phenotype.provenance_gaps} == {
        "promoter_activity_uncertainty",
        "promoter_activity_uncertainty_type",
    }
    assert {gap.reason for gap in phenotype.provenance_gaps} == {
        ProvenanceGapReason.not_reported_by_primary
    }


def test_the_well_id_names_the_arm_so_two_arms_of_one_well_stay_distinct() -> None:
    ids = {m.well_id(arm, "AZ01", "A1") for arm in m.ARMS}
    assert len(ids) == len(m.ARMS)
    assert "Ampicillin Treatment|AZ01|A1" in ids


def test_the_sourced_n_samples_is_one_and_quotes_the_statistics_section() -> None:
    sourced = m.SOURCED_VALUES["n_samples"]
    assert sourced.value == 1
    assert "performed only once" in sourced.quote
    assert sourced.provenance.sha256 == m.PAPER_TEXT_SHA256
    assert sourced.provenance.source_uri == m.PAPER_TEXT_REL


def test_every_sourced_value_is_anchored_to_a_mirrored_text_artifact() -> None:
    """A quote in the binary workbook is not auditable, so the legends go to text."""
    anchors = {v.provenance.source_uri for v in m.SOURCED_VALUES.values()}
    assert anchors == {m.PAPER_TEXT_REL, m.LEGENDS_REL}
    mirrored = {raw.relpath for raw in m.RAW_FILES}
    assert anchors <= mirrored
    assert all(
        v.provenance.citation_key == m.CITATION_KEY for v in m.SOURCED_VALUES.values()
    )


def test_the_dataset_is_registered_under_its_class_name() -> None:
    assert (
        dataset_registry["PromoterReporterMohiuddin2022Dataset"]
        is m.PromoterReporterMohiuddin2022Dataset
    )


def test_the_mirror_dir_points_at_this_keys_raw_mirror() -> None:
    assert m.raw_mirror_dir("/root") == Path("/root/torchcell-raw") / m.CITATION_KEY


# --------------------------------------------------------------------------- #
# The raw mirror
# --------------------------------------------------------------------------- #
@pytest.fixture
def deposited(
    release: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[dict[str, Path], Path]:
    """A tmp DATA_ROOT whose raw mirror holds the synthetic workbook and three stubs."""
    sources: dict[str, Path] = {m.DATA_FILE: release}
    for name in (m.SI_FILE, m.PAPER_PDF_FILE, m.PAPER_TEXT_FILE):
        path = tmp_path / name
        path.write_bytes(f"synthetic {name}".encode())
        sources[name] = path
    legend_text = m.extract_sheet_legends(release)
    files = tuple(
        m.RawFile(
            name=raw.name,
            relpath=raw.relpath,
            role=raw.role,
            sha256=(
                hashlib.sha256(legend_text.encode()).hexdigest()
                if raw.derived
                else hashlib.sha256(sources[raw.name].read_bytes()).hexdigest()
            ),
            bytes=(
                len(legend_text.encode())
                if raw.derived
                else sources[raw.name].stat().st_size
            ),
            description=f"synthetic {raw.name}",
            derived=raw.derived,
        )
        for raw in m.RAW_FILES
    )
    monkeypatch.setattr(m, "RAW_FILES", files)
    monkeypatch.setattr(m, "RETRIEVED_FILES", tuple(f for f in files if not f.derived))
    monkeypatch.setattr(m, "DATA_SHA256", {files[0].name: files[0].sha256})
    monkeypatch.setattr(m, "RAW_FILES_BY_NAME", {f.name: f for f in files})
    monkeypatch.setattr(
        m, "LEGENDS_SHA256", hashlib.sha256(legend_text.encode()).hexdigest()
    )
    data_root = tmp_path / "data_root"
    m.deposit_raw_mirror(sources=sources, data_root=str(data_root))
    return sources, data_root


def test_the_mirror_holds_the_data_the_si_and_the_article_under_their_roles(
    deposited: tuple[dict[str, Path], Path],
) -> None:
    _, data_root = deposited
    manifest = m.load_manifest(str(data_root))
    assert [(r.path, r.role) for r in manifest.files] == [
        (f"data/{m.DATA_FILE}", "raw_data"),
        (f"si/{m.SI_FILE}", "si_pdf"),
        (f"paper/{m.PAPER_PDF_FILE}", "paper_pdf"),
        (f"paper/{m.PAPER_TEXT_FILE}", "paper_text"),
        (m.LEGENDS_REL, "si_text"),
    ]
    assert manifest.doi == m.PAPER_DOI
    assert manifest.provenance_complete is True


def test_the_derived_legend_text_records_its_reader_and_the_workbooks_hash(
    deposited: tuple[dict[str, Path], Path],
) -> None:
    """The workbook is binary; the legend text is what makes its quotes auditable."""
    sources, data_root = deposited
    manifest = m.load_manifest(str(data_root))
    legends = manifest.files[-1]
    assert legends.retrieval is None
    assert legends.processing is not None
    assert legends.processing.processor.endswith("extract_sheet_legends")
    assert legends.processing.tool == "openpyxl"
    assert legends.processing.input_sha256 == [
        hashlib.sha256(sources[m.DATA_FILE].read_bytes()).hexdigest()
    ]
    text = (m.raw_mirror_dir(str(data_root)) / m.LEGENDS_REL).read_text()
    assert m.SOURCED_VALUES["fold_change_rule"].quote in text
    assert m.SOURCED_VALUES["raw_sheet_legend"].quote in text
    assert text.startswith(f"# sheet: {m.RAW_DATA_SHEET}\n")


def test_a_legend_text_that_changed_is_refused_rather_than_overwritten(
    deposited: tuple[dict[str, Path], Path],
) -> None:
    sources, data_root = deposited
    path = m.raw_mirror_dir(str(data_root)) / m.LEGENDS_REL
    path.write_text("not the released legends\n", encoding="utf-8")
    with pytest.raises(RuntimeError, match="different sha256; refusing"):
        m.deposit_raw_mirror(sources=sources, data_root=str(data_root))


def test_the_manifest_records_the_rerunnable_pmc_retrieval(
    deposited: tuple[dict[str, Path], Path],
) -> None:
    _, data_root = deposited
    manifest = m.load_manifest(str(data_root))
    record = manifest.files[0]
    assert record.retrieval is not None
    assert record.retrieval.method == "pmc_cloud"
    assert record.retrieval.params == {"key": f"PMC8865558.1/{m.DATA_FILE}"}
    assert m.manifest_sha256(manifest, f"data/{m.DATA_FILE}") == record.sha256
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        m.manifest_sha256(manifest, "data/nope.xlsx")


def test_the_deposit_is_idempotent_and_refuses_a_changed_mirror_file(
    deposited: tuple[dict[str, Path], Path],
) -> None:
    sources, data_root = deposited
    assert m.deposit_raw_mirror(sources=sources, data_root=str(data_root)) == (
        m.raw_mirror_dir(str(data_root))
    )
    mirrored = m.raw_mirror_dir(str(data_root)) / f"data/{m.DATA_FILE}"
    mirrored.write_bytes(b"not the released bytes")
    with pytest.raises(RuntimeError, match="different sha256; refusing"):
        m.deposit_raw_mirror(sources=sources, data_root=str(data_root))


def test_the_deposit_refuses_bytes_that_do_not_match_the_pin(
    deposited: tuple[dict[str, Path], Path], tmp_path: Path
) -> None:
    sources, data_root = deposited
    other = tmp_path / "other.xlsx"
    other.write_bytes(b"not the released bytes")
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        m.deposit_raw_mirror(
            sources={**sources, m.DATA_FILE: other}, data_root=str(data_root)
        )


def test_the_deposit_refuses_a_missing_source(
    deposited: tuple[dict[str, Path], Path],
) -> None:
    sources, data_root = deposited
    partial = {k: v for k, v in sources.items() if k != m.SI_FILE}
    with pytest.raises(KeyError, match=m.SI_FILE):
        m.deposit_raw_mirror(sources=partial, data_root=str(data_root))


# --------------------------------------------------------------------------- #
# The end-to-end build, on the synthetic release
# --------------------------------------------------------------------------- #
def _pin_assembly(monkeypatch: pytest.MonkeyPatch) -> None:
    """Pin ``assembly_reference`` so no assembly report has to be served."""
    from torchcell.datamodels.schema import AssemblyReferenceGenome

    reference = AssemblyReferenceGenome(
        species="Escherichia coli",
        strain="MG1655",
        assembly_set="ecoli_K12_MG1655_ASM584v2",
        assembly_accession="GCA_000005845.2",
    )
    assert reference.assembly_set == MG1655_ASSEMBLY.assembly_set
    monkeypatch.setattr(m, "assembly_reference", lambda strain, **_: reference)


@pytest.fixture
def built(
    deposited: tuple[dict[str, Path], Path],
    mg1655: EcoliK12MG1655Genome,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> Any:
    """Build the dataset from the synthetic release into the tmp DATA_ROOT."""
    _, data_root = deposited
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    _pin_assembly(monkeypatch)
    dataset = m.PromoterReporterMohiuddin2022Dataset(
        root=str(tmp_path / "store"), ecoli_genome=mg1655
    )
    yield dataset
    dataset.close_lmdb()


def test_the_build_writes_one_record_per_arm_well_and_read_hour(built: Any) -> None:
    assert len(built) == len(SYNTHETIC_WELLS) * len(m.ARMS) * len(m.READ_HOURS) == 144


def test_the_build_refuses_a_grid_that_is_not_the_release(
    deposited: tuple[dict[str, Path], Path],
    mg1655: EcoliK12MG1655Genome,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    _, data_root = deposited
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    _pin_assembly(monkeypatch)
    monkeypatch.setattr(m, "EXPECTED_RECORDS", 1)
    with pytest.raises(RuntimeError, match="wrote 144 records, the grid is 1"):
        m.PromoterReporterMohiuddin2022Dataset(
            root=str(tmp_path / "store"), ecoli_genome=mg1655
        )


def test_the_built_records_are_the_promoter_activity_pair(built: Any) -> None:
    """The store holds plain dicts; the discriminators are what name the pair."""
    assert built.experiment_class is PromoterActivityExperiment
    assert built.reference_class is PromoterActivityExperimentReference
    item = built[0]
    assert item["experiment"]["experiment_type"] == "promoter_activity"
    assert item["reference"]["experiment_reference_type"] == "promoter_activity"
    assert item["experiment"]["phenotype"]["label_name"] == "promoter_activity"
    assert EXPERIMENT_TYPE_MAP["promoter_activity"] is PromoterActivityExperiment
    assert (
        EXPERIMENT_REFERENCE_TYPE_MAP["promoter_activity"]
        is PromoterActivityExperimentReference
    )
    built.close_lmdb()


def test_every_records_reference_is_its_own_untreated_reading_at_the_same_hour(
    built: Any,
) -> None:
    """This is what makes the released fold change recoverable from one record."""
    checked = 0
    for index in range(len(built)):
        item = built[index]
        phenotype = item["experiment"]["phenotype"]
        control = item["reference"]["phenotype_reference"]
        assert control["promoter_name"] == phenotype["promoter_name"]
        assert control["well_id"].startswith(f"{m.UNTREATED.header}|")
        assert (
            item["reference"]["environment_reference"]["duration_hours"]
            == (item["experiment"]["environment"]["duration_hours"])
        )
        assert item["reference"]["environment_reference"]["perturbations"] == []
        checked += 1
    assert checked == 144
    built.close_lmdb()


def test_the_build_writes_the_fold_change_check_and_the_promoter_ledger(
    built: Any,
) -> None:
    preprocess = Path(built.preprocess_dir)
    extraction = json.loads((preprocess / "extraction.json").read_text())
    assert extraction["grid"] == {
        "n_wells": 4,
        "n_arms": 4,
        "n_read_hours": 9,
        "n_records": 144,
    }
    assert extraction["fold_change"]["n_cells"] == 108
    assert extraction["fold_change"]["n_within_tolerance"] == 108
    ledger = json.loads((preprocess / "identifier_reconciliation.json").read_text())
    assert ledger["unresolved_labels"] == ["U66"]
    assert ledger["n_labels_resolved"] == 2
    built.close_lmdb()


def test_the_verifier_passes_every_row_on_the_synthetic_store(
    built: Any, mg1655: EcoliK12MG1655Genome, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The handle is closed first: a held one fails the verifier's own re-read."""
    from torchcell.verification import runners

    built.close_lmdb()
    monkeypatch.setattr(runners, "_genome_for_reference", lambda *_: mg1655)
    monkeypatch.setattr(
        runners,
        "_gene_set_for_reference",
        lambda *_: {locus.tag for locus in MG1655_LOCI},
    )
    data_root = os.environ["DATA_ROOT"]
    monkeypatch.setattr(m, "DATASET_ROOT_REL", os.path.relpath(built.root, data_root))
    monkeypatch.setattr(m, "EXPECTED_RECORDS", 144)
    monkeypatch.setattr(m, "EXPECTED_FOLD_CHANGE_CELLS", 108)
    report = m.run_verification(data_root)
    names = {row.name for row in report.results}
    assert "fold_change_recoverable_from_one_record" in names
    assert "the_drug_is_on_the_reads_from_hour_five_only" in names
    assert "reading_uniqueness" in names
    assert "promoter_genes_in_the_host_gene_universe" in names
    family = [row for row in report.results if row.name != "provenance_audit"]
    assert [(row.name, row.message) for row in family if not row.passed] == []
    assert {row.level for row in family} == {
        Level.L0,
        Level.L1,
        Level.L2,
        Level.L3,
        Level.L4,
    }
    # The mirror here holds the SYNTHETIC release, so every quote audit must report
    # sha256 drift: the audit re-hashes the file instead of trusting the pin. The
    # quotes themselves are checked against the real pinned bytes by the --data report
    # test, and against the rendered legend text by the deposit test above.
    audits = [row for row in report.results if row.name == "provenance_audit"]
    assert len(audits) == len(m.SOURCED_VALUES)
    assert all(not row.passed for row in audits)
    assert all("sha256 drift" in row.message for row in audits)
    assert {m.PAPER_TEXT_FILE, m.LEGENDS_FILE} == {
        row.message.rsplit(" ", 1)[-1].strip("()") for row in audits
    }


# --------------------------------------------------------------------------- #
# The dev-tree store the real release built (--data)
# --------------------------------------------------------------------------- #
@pytest.mark.data
def test_the_real_store_holds_the_whole_released_grid() -> None:
    dataset = m.PromoterReporterMohiuddin2022Dataset(
        root=os.path.join(os.environ["DATA_ROOT"], m.DATASET_ROOT_REL)
    )
    assert len(dataset) == 69480
    phenotype = dataset[0]["experiment"]["phenotype"]
    assert phenotype["promoter_activity"] > 0.0
    assert phenotype["well_id"].count("|") == 2
    assert phenotype["readout"] == "plate_reader_fluorescence"
    dataset.close_lmdb()


@pytest.mark.data
def test_the_real_store_resolves_1761_of_the_1809_promoter_labels() -> None:
    dataset = m.PromoterReporterMohiuddin2022Dataset(
        root=os.path.join(os.environ["DATA_ROOT"], m.DATASET_ROOT_REL)
    )
    ledger = json.loads(
        (Path(dataset.preprocess_dir) / "identifier_reconciliation.json").read_text()
    )
    dataset.close_lmdb()
    assert ledger["n_wells"] == 1930
    assert ledger["n_labels"] == 1809
    assert ledger["n_labels_resolved"] == 1761
    assert len(ledger["unresolved_labels"]) == 48
    assert {"Empty", "U66", "U139", "rrnA", "spr", "ygaD"} <= set(
        ledger["unresolved_labels"]
    )


@pytest.mark.data
def test_the_real_store_reproduces_every_released_fold_change() -> None:
    dataset = m.PromoterReporterMohiuddin2022Dataset(
        root=os.path.join(os.environ["DATA_ROOT"], m.DATASET_ROOT_REL)
    )
    extraction = json.loads(
        (Path(dataset.preprocess_dir) / "extraction.json").read_text()
    )
    dataset.close_lmdb()
    assert extraction["fold_change"]["n_cells"] == 52110
    assert extraction["fold_change"]["n_within_tolerance"] == 52110
    assert extraction["fold_change"]["max_relative_error"] < 1e-15


@pytest.mark.data
def test_the_real_stores_verification_report_passes_every_row() -> None:
    """The report the real `verify` run wrote, read row by row."""
    report = json.loads(
        Path(
            os.environ["DATA_ROOT"],
            m.DATASET_ROOT_REL,
            "preprocess",
            "verification_report.json",
        ).read_text()
    )
    assert report["dataset_name"] == "PromoterReporterMohiuddin2022Dataset"
    rows = {(row["level"], row["name"]): row for row in report["results"]}
    assert [key for key, row in rows.items() if not row["passed"]] == []
    assert rows[("L1", "count")]["details"] == {"observed": 69480, "expected": 69480}
    assert rows[("L1", "reading_uniqueness")]["details"]["n_keys"] == 69480
    assert rows[("L2", "value_fidelity")]["details"]["n_values"] == 69480
    fold = rows[("L3", "fold_change_recoverable_from_one_record")]["details"]
    assert fold["n_checked"] == fold["n_expected"] == 52110
    assert fold["max_relative_error"] < 1e-15
    dose = rows[("L3", "the_drug_is_on_the_reads_from_hour_five_only")]["details"]
    assert (dose["untreated"], dose["treated_predose"], dose["treated_dosed"]) == (
        17370,
        17370,
        34740,
    )
    gene = rows[("L3", "promoter_gene_is_a_locus")]["details"]
    assert gene["n_genes"] == 1761
    assert gene["n_null_records"] == 48 * 36
    assert len([k for k in rows if k[1] == "provenance_audit"]) == len(m.SOURCED_VALUES)
