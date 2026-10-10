# tests/torchcell/datasets/pputida/test_royet2025.py
# [[tests.torchcell.datasets.pputida.test_royet2025]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/pputida/test_royet2025.py
"""The Royet 2025 KT2440 metal Tn-seq loader (``torchcell.datasets.pputida.royet2025``).

Synthetic tests (run everywhere) write Table S5 and Table S3 workbooks in ``tmp_path``
in the released layout (title row, group-header row, header row, a blank row, then one
TEXT-valued row per gene). The hermetic build uses the real ``PPutidaKT2440Genome`` over
the synthetic KT2440 assembly of ``tests/torchcell/sequence/genome/_bacterial_fixtures.py``
served through a stubbed ``resolve`` with the network refused; the module's build oracles
(gene count, per-metal drops, record count, per-metal hit counts) are lowered to the
synthetic release's, and ``verify_raw_files`` is a presence check (synthetic bytes cannot
carry the real pins).

The four synthetic genes, identical in every metal sheet:

    #Orf     Sites  State  Mean A  Mean B  log2FC   q        fate
    PP_0001  8      NE     10.0    2.5     -2.00    0.00000  kept, the one hit
    PP_0002  20     NE     5.0     5.0     -0.00    1.00000  kept, a measured zero
    PP_0005  12     ES     0.0     0.0     0.00     1.00000  dropped, no reads
    PP_0007  0      N/A    0.0     0.0     0.00     1.00000  dropped, no TA site

so the build stores 2 genes x 4 metals = 8 records and drops 8.

Data-gated tests (``--data``) read the real raw mirror, the literature mirror and the
built dev-tree LMDB under ``$DATA_ROOT`` (they never build it).
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
from collections.abc import Mapping
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.pputida.royet2025 as m
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    KT2440_GAF,
    KT2440_LOCI,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.data import ManifestPinMismatchError
from torchcell.datamodels.media import LB
from torchcell.datamodels.schema import (
    AssayType,
    AssemblyReferenceGenome,
    BacterialEnvironmentResponseExperiment,
    BacterialEnvironmentResponseExperimentReference,
    ConcentrationUnit,
    MeasurementType,
    SampleUnit,
    SmallMoleculePerturbation,
    TransposonInsertionPerturbation,
)
from torchcell.datasets.dataset_registry import dataset_registry
from torchcell.literature.manifest import RetrievalMethod
from torchcell.sequence.genome.pputida.kt2440 import (
    KT2440_ASSEMBLY,
    PPutidaKT2440Genome,
)
from torchcell.verification.sourced import ProvenanceGapReason, audit_sourced_value

REFERENCE = AssemblyReferenceGenome(
    species="Pseudomonas putida",
    strain="KT2440",
    assembly_set="pputida_KT2440_ASM756v2",
    assembly_accession="GCA_000007565.2",
)
#: ``(#Orf, Sites, State, Mean A, Mean B, log2FC, q)`` as the release writes them.
SYNTHETIC_ROWS: tuple[tuple[str, int, str, str, str, str, str], ...] = (
    ("PP_0001", 8, "NE", "10.0", "2.5", "-2.00", "0.00000"),
    ("PP_0002", 20, "NE", "5.0", "5.0", "-0.00", "1.00000"),
    ("PP_0005", 12, "ES", "0.0", "0.0", "0.00", "1.00000"),
    ("PP_0007", 0, "N/A", "0.0", "0.0", "0.00", "1.00000"),
)
KEPT = ("PP_0001", "PP_0002")
SYNTHETIC_RECORDS = len(KEPT) * len(m.METALS)


def _s5_row(
    orf: str, sites: int, state: str, a: str, b: str, fc: str, q: str
) -> list[Any]:
    return [orf, "name", "desc", sites, 0, 0, sites, 0, "1.0", a, state, a, b, "0.0",
            fc, q, q]  # fmt: skip


def _write_s5(
    path: Path,
    rows: tuple[tuple[str, int, str, str, str, str, str], ...] = SYNTHETIC_ROWS,
    *,
    header: tuple[str, ...] = m.S5_HEADER,
    order: dict[str, list[int]] | None = None,
    group: dict[str, str] | None = None,
) -> Path:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for spec in m.METALS:
        sheet = book.create_sheet(spec.sheet)
        sheet.append(["RAW DATA of the HMM and RESAMPLING ANALYSIS by TRANSIT"])
        group_row: list[Any] = [None] * len(header)
        group_row[3] = "HMM - LB"
        group_row[11] = (group or {}).get(spec.code, spec.group_header)
        sheet.append(group_row)
        sheet.append(list(header))
        sheet.append([None] * len(header))
        indices = (order or {}).get(spec.code, list(range(len(rows))))
        for index in indices:
            sheet.append(_s5_row(*rows[index]))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def _write_s3(path: Path, arms: tuple[str, ...] = ("GL", *m.REPLICATE_ARMS)) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Feuil1"
    sheet.append(["TABLE S3 Tn-Seq analysis of P. putida KT2440"])
    sheet.append([None])
    sheet.append(["Mutant pool", "Total no. of reads"])
    for arm in arms:
        sheet.append([f"{arm} #1", "1"])
        sheet.append([f"{arm} #2", "1"])
    sheet.append([None])
    sheet.append(["a The number of reads ... for the sample LB #1 "])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def _synthetic_metals() -> tuple[m.MetalSpec, ...]:
    return tuple(spec.model_copy(update={"paper_hits": 1}) for spec in m.METALS)


@pytest.fixture
def oracles(monkeypatch: pytest.MonkeyPatch) -> None:
    """The module's build oracles, lowered to the synthetic release."""
    monkeypatch.setattr(m, "METALS", _synthetic_metals())
    monkeypatch.setattr(m, "N_GENES", len(SYNTHETIC_ROWS))
    monkeypatch.setattr(m, "EXPECTED_DROPPED", {s.code: 2 for s in m.METALS})
    monkeypatch.setattr(m, "EXPECTED_RECORDS", SYNTHETIC_RECORDS)


@pytest.fixture
def kt2440(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> PPutidaKT2440Genome:
    files = write_assembly(tmp_path / "tier", KT2440_ASSEMBLY, KT2440_LOCI, KT2440_GAF)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    return PPutidaKT2440Genome(genome_root=str(tmp_path / "kt2440"), overwrite=True)


@pytest.fixture
def presence_only_pins(monkeypatch: pytest.MonkeyPatch) -> list[Mapping[str, str]]:
    calls: list[Mapping[str, str]] = []

    def record(raw: str, pins: Mapping[str, str]) -> None:
        missing = [f for f in pins if not osp.exists(osp.join(raw, f))]
        if missing:
            raise FileNotFoundError(f"pinned raw files absent: {missing}")
        calls.append(dict(pins))

    monkeypatch.setattr(m, "verify_raw_files", record)
    return calls


@pytest.fixture
def synthetic(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    oracles: None,
    presence_only_pins: list[Mapping[str, str]],
) -> Path:
    root = tmp_path / m.DATASET_ROOT_REL
    _write_s5(root / "raw" / m.TABLE_S5)
    _write_s3(root / "raw" / m.TABLE_S3)
    monkeypatch.setattr(m, "reference_genome", lambda *a, **k: REFERENCE)
    return root


@pytest.fixture
def built(
    synthetic: Path, kt2440: PPutidaKT2440Genome
) -> m.EnvMetalTnseqRoyet2025Dataset:
    return m.EnvMetalTnseqRoyet2025Dataset(root=str(synthetic), pputida_genome=kt2440)


# --------------------------------------------------------------------------- #
# Reading the release
# --------------------------------------------------------------------------- #
def test_read_metal_sheet_parses_the_text_cells(tmp_path: Path, oracles: None) -> None:
    path = _write_s5(tmp_path / "s5.xlsx")
    rows = m.read_metal_sheet(path, m.METALS[0])
    assert [r.orf for r in rows] == [r[0] for r in SYNTHETIC_ROWS]
    first = rows[0]
    assert (first.mean_a, first.mean_b, first.log2fc, first.q_value) == (
        10.0,
        2.5,
        -2.0,
        0.0,
    )
    assert first.sites == 8 and first.state == "NE" and first.p_value == 0.0
    # '-0.00' is stored as +0.0, not as a signed zero
    assert rows[1].log2fc == 0.0 and math.copysign(1.0, rows[1].log2fc) == 1.0
    assert [r.is_empty for r in rows] == [False, False, True, True]


def test_read_metal_sheet_refuses_a_changed_header(
    tmp_path: Path, oracles: None
) -> None:
    path = _write_s5(tmp_path / "s5.xlsx", header=(*m.S5_HEADER[:-1], "q"))
    with pytest.raises(m.SheetFormatError, match="is headed"):
        m.read_metal_sheet(path, m.METALS[0])


def test_read_metal_sheet_refuses_another_comparison(
    tmp_path: Path, oracles: None
) -> None:
    path = _write_s5(tmp_path / "s5.xlsx", group={"Co": "RESAMPLING LB vs Nickel"})
    with pytest.raises(m.SheetFormatError, match="group header"):
        m.read_metal_sheet(path, m.METALS[0])


def test_read_metal_sheet_refuses_a_missing_gene(tmp_path: Path, oracles: None) -> None:
    path = _write_s5(tmp_path / "s5.xlsx", SYNTHETIC_ROWS[:3])
    with pytest.raises(m.SheetFormatError, match="carries 3 rows"):
        m.read_metal_sheet(path, m.METALS[0])


def test_read_metal_sheet_refuses_a_row_with_no_orf(
    tmp_path: Path, oracles: None
) -> None:
    rows = (*SYNTHETIC_ROWS[:3], (None, 0, "NE", "1.0", "1.0", "0.00", "1.00000"))
    path = _write_s5(tmp_path / "s5.xlsx", rows)  # type: ignore[arg-type]
    with pytest.raises(m.SheetFormatError, match="no #Orf"):
        m.read_metal_sheet(path, m.METALS[0])


def test_read_metal_sheet_refuses_a_hit_count_the_results_do_not_state(
    tmp_path: Path, oracles: None
) -> None:
    path = _write_s5(tmp_path / "s5.xlsx")
    spec = m.METALS[0].model_copy(update={"paper_hits": 2})
    with pytest.raises(m.SheetFormatError, match="the Results state 2"):
        m.read_metal_sheet(path, spec)


def test_read_release_reads_all_four_metals(tmp_path: Path, oracles: None) -> None:
    sheets = m.read_release(_write_s5(tmp_path / "s5.xlsx"))
    assert list(sheets) == ["Co", "Cu", "Zn", "Cd"]
    assert all(len(rows) == len(SYNTHETIC_ROWS) for rows in sheets.values())


def test_read_release_refuses_a_reordered_sheet(tmp_path: Path, oracles: None) -> None:
    path = _write_s5(tmp_path / "s5.xlsx", order={"Zn": [1, 0, 2, 3]})
    with pytest.raises(m.SheetFormatError, match="sheet Zn lists the genes"):
        m.read_release(path)


def test_read_replicate_pools_counts_two_per_arm_and_skips_the_footnote(
    tmp_path: Path,
) -> None:
    pools = m.read_replicate_pools(_write_s3(tmp_path / "s3.xlsx"))
    assert pools == {"LB": 2, "Co": 2, "Cu": 2, "Zn": 2, "Cd": 2}


def test_read_replicate_pools_refuses_a_missing_arm(tmp_path: Path) -> None:
    path = _write_s3(tmp_path / "s3.xlsx", arms=("GL", "LB", "Co", "Cu", "Zn"))
    with pytest.raises(m.SheetFormatError, match="'Cd': 0"):
        m.read_replicate_pools(path)


# --------------------------------------------------------------------------- #
# Retention
# --------------------------------------------------------------------------- #
def test_build_drop_log_accounts_for_every_cell(tmp_path: Path, oracles: None) -> None:
    sheets = m.read_release(_write_s5(tmp_path / "s5.xlsx"))
    log = m.build_drop_log("royet", sheets)
    assert (log.source_cells, log.kept_records, log.dropped_records) == (16, 8, 8)
    assert list(log.rules) == [m.DROP_EMPTY_GENE]
    for ledger in log.per_metal:
        assert ledger.dropped_empty == 2
        assert ledger.dropped_empty_states == {"ES": 1, "N/A": 1}
        assert ledger.dropped_empty_zero_sites == 1
        assert ledger.kept_exact_zero == 1
        assert ledger.kept_q_le_threshold == 1


def test_stored_cells_skips_the_empty_genes(tmp_path: Path, oracles: None) -> None:
    sheets = m.read_release(_write_s5(tmp_path / "s5.xlsx"))
    cells = list(m.stored_cells(sheets, {"PP_0001": "parB", "PP_0002": "PP_0002"}))
    assert [(tag, spec.code, value) for tag, _, spec, value in cells] == [
        (tag, code, value)
        for code in ("Co", "Cu", "Zn", "Cd")
        for tag, value in (("PP_0001", -2.0), ("PP_0002", 0.0))
    ]


# --------------------------------------------------------------------------- #
# Environment, genotype, phenotype
# --------------------------------------------------------------------------- #
def test_each_environment_is_lb_at_28_c_for_12_generations_plus_one_metal() -> None:
    expected = {
        "Co": ("cobalt chloride", 10.0, ConcentrationUnit.micromolar),
        "Cu": ("copper(II) chloride", 2.5, ConcentrationUnit.millimolar),
        "Zn": ("zinc chloride", 125.0, ConcentrationUnit.micromolar),
        "Cd": ("cadmium chloride", 12.5, ConcentrationUnit.micromolar),
    }
    for spec in m.METALS:
        env = m.environment(spec)
        assert env.media == LB
        assert env.temperature is not None and env.temperature.value == 28.0
        assert env.aerobicity == "aerobic"
        assert env.duration_generations == 12.0
        (edit,) = env.perturbations
        assert isinstance(edit, SmallMoleculePerturbation)
        name, dose, unit = expected[spec.code]
        assert edit.compound.name == name
        assert edit.compound.inchikey is not None
        assert edit.concentration is not None
        assert (edit.concentration.value, edit.concentration.unit) == (dose, unit)


def test_the_genotype_is_one_gene_level_mariner_insertion() -> None:
    (perturbation,) = m.insertion_genotype("PP_0001", "parB").perturbations
    assert isinstance(perturbation, TransposonInsertionPerturbation)
    assert (perturbation.systematic_gene_name, perturbation.perturbed_gene_name) == (
        "PP_0001",
        "parB",
    )
    assert perturbation.gene_namespace == "pputida_kt2440_locus_tag"
    assert perturbation.transposon == "mariner (Himar1)"
    assert (
        perturbation.barcode,
        perturbation.insertion_position,
        perturbation.insertion_strand,
        perturbation.library_pool,
    ) == (None, None, None, None)


def test_the_three_absent_perturbation_fields_are_typed_gaps() -> None:
    assert [g.field for g in m.PERTURBATION_FIELD_GAPS] == [
        "barcode",
        "insertion_position",
        "insertion_strand",
    ]
    assert {g.reason for g in m.PERTURBATION_FIELD_GAPS} == {
        ProvenanceGapReason.not_reported_by_primary
    }


def test_the_phenotype_is_a_signed_log2_ratio_over_two_replicates() -> None:
    phenotype = m.phenotype(-8.22)
    assert phenotype.environment_response == -8.22
    assert phenotype.measurement_type is MeasurementType.log2_ratio
    assert phenotype.assay_type is AssayType.other
    assert (phenotype.n_samples, phenotype.sample_unit) == (
        2,
        SampleUnit.biological_replicate,
    )
    assert phenotype.environment_response_se is None
    assert [g.field for g in phenotype.provenance_gaps] == [
        "environment_response_uncertainty",
        "environment_response_se",
    ]
    reference = m.reference_phenotype()
    assert reference.environment_response == 0.0
    assert reference.units == m.UNITS_REFERENCE


def test_build_reference_pairs_the_metal_environment_with_a_zero() -> None:
    env = m.environment(m.METALS[2])
    reference = m.build_reference("royet", REFERENCE, env)
    assert isinstance(reference, BacterialEnvironmentResponseExperimentReference)
    assert reference.environment_reference == env
    assert reference.environment_reference is not env
    assert reference.phenotype_reference.environment_response == 0.0


def test_canonical_symbol_keeps_a_unique_symbol_and_falls_back_on_a_shared_one(
    kt2440: PPutidaKT2440Genome,
) -> None:
    assert m.canonical_symbol(kt2440, "PP_0001") == "parB"
    assert m.canonical_symbol(kt2440, "PP_0002") == "PP_0002"
    # 'asd' names both PP_0005 and PP_0006, so it does not resolve back to either
    assert m.canonical_symbol(kt2440, "PP_0006") == "PP_0006"


def test_the_sourced_values_name_the_paper_and_the_metal_table_agrees() -> None:
    assert {sv.provenance.sha256 for sv in m.SOURCED_VALUES.values()} == {
        m.PAPER_MD_SHA256
    }
    assert m.SOURCED_VALUES["hit_counts"].value == {
        spec.code: spec.paper_hits for spec in m.METALS
    }
    assert m.SOURCED_VALUES["n_samples"].value == m.N_REPLICATES
    assert m.EXPECTED_RECORDS == 22916 - 1333 == 21583


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _fake_raw(monkeypatch: pytest.MonkeyPatch, src: Path) -> dict[str, Path]:
    """Two small files standing in for the PMC objects, with RAW_FILES re-pinned."""
    src.mkdir(parents=True, exist_ok=True)
    files = {}
    raws = []
    for raw, payload in zip(m.RAW_FILES, (b"table s5", b"table s3"), strict=True):
        path = src / raw.name
        path.write_bytes(payload)
        sha = m._sha256(path)
        files[raw.name] = path
        raws.append(
            raw.model_copy(
                update={
                    "sha256": sha,
                    "bytes": len(payload),
                    "retrieval": raw.retrieval.model_copy(update={"sha256": sha}),
                }
            )
        )
    monkeypatch.setattr(m, "RAW_FILES", tuple(raws))
    monkeypatch.setattr(m, "DATA_SHA256", {r.name: r.sha256 for r in raws})
    return files


def test_the_raw_files_are_pmc_bucket_objects() -> None:
    assert [raw.name for raw in m.RAW_FILES] == [m.TABLE_S5, m.TABLE_S3]
    for raw in m.RAW_FILES:
        assert raw.retrieval.method is RetrievalMethod.pmc_cloud
        assert raw.retrieval.params == {"key": f"PMC12041740.1/{raw.name}"}
        assert raw.mirror_relpath == f"data/{raw.name}"
        assert raw.retrieval.sha256 == raw.sha256


def test_deposit_writes_the_mirror_and_its_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = _fake_raw(monkeypatch, tmp_path / "src")
    root = m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path))
    assert root == tmp_path / m.RAW_DIR_REL
    manifest = m.load_manifest(str(tmp_path))
    assert manifest.citation_key == m.CITATION_KEY
    for raw in m.RAW_FILES:
        assert m.manifest_sha256(manifest, raw.mirror_relpath) == raw.sha256
    assert manifest.si_expected == list(m.NOT_MIRRORED)
    # idempotent: a second deposit of the same bytes leaves the files alone
    m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path))
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        m.manifest_sha256(manifest, "data/other.xlsx")


def test_deposit_refuses_a_missing_source_wrong_bytes_and_a_changed_mirror(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = _fake_raw(monkeypatch, tmp_path / "src")
    with pytest.raises(KeyError, match="no source given"):
        m.deposit_raw_mirror(
            sources={m.TABLE_S5: sources[m.TABLE_S5]}, data_root=str(tmp_path)
        )
    wrong = tmp_path / "wrong.xlsx"
    wrong.write_bytes(b"other")
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        m.deposit_raw_mirror(
            sources={**sources, m.TABLE_S5: wrong}, data_root=str(tmp_path)
        )
    dest = tmp_path / m.RAW_DIR_REL / "data" / m.TABLE_S5
    dest.parent.mkdir(parents=True)
    dest.write_bytes(b"edited")
    with pytest.raises(RuntimeError, match="different sha256; refusing"):
        m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path))


def test_retrieve_raw_files_runs_each_recorded_retriever(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = _fake_raw(monkeypatch, tmp_path / "src")
    payload = {raw.retrieval.source_url: sources[raw.name] for raw in m.RAW_FILES}
    monkeypatch.setattr(
        m, "run_retriever", lambda record: payload[record.source_url].read_bytes()
    )
    out = m.retrieve_raw_files(tmp_path / "fetched")
    assert sorted(out) == sorted(sources)
    assert all(out[n].read_bytes() == sources[n].read_bytes() for n in out)


def test_download_links_the_pinned_mirror_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = _fake_raw(monkeypatch, tmp_path / "src")
    m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path))
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    dataset = m.EnvMetalTnseqRoyet2025Dataset.__new__(m.EnvMetalTnseqRoyet2025Dataset)
    dataset.root = str(tmp_path / "ds")
    m.EnvMetalTnseqRoyet2025Dataset.download(dataset)
    for name in sources:
        assert osp.islink(osp.join(dataset.raw_dir, name))


def test_download_refuses_a_manifest_that_disagrees_with_the_pin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = _fake_raw(monkeypatch, tmp_path / "src")
    m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path))
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    monkeypatch.setattr(m, "RAW_FILES", tuple(
        raw.model_copy(update={"sha256": "0" * 64}) for raw in m.RAW_FILES
    ))  # fmt: skip
    dataset = m.EnvMetalTnseqRoyet2025Dataset.__new__(m.EnvMetalTnseqRoyet2025Dataset)
    dataset.root = str(tmp_path / "ds")
    with pytest.raises(ManifestPinMismatchError):
        m.EnvMetalTnseqRoyet2025Dataset.download(dataset)


def test_download_refuses_a_mirror_file_that_is_gone(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sources = _fake_raw(monkeypatch, tmp_path / "src")
    m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path))
    (tmp_path / m.RAW_DIR_REL / "data" / m.TABLE_S3).unlink()
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    dataset = m.EnvMetalTnseqRoyet2025Dataset.__new__(m.EnvMetalTnseqRoyet2025Dataset)
    dataset.root = str(tmp_path / "ds")
    with pytest.raises(RuntimeError, match="missing from mirror"):
        m.EnvMetalTnseqRoyet2025Dataset.download(dataset)


# --------------------------------------------------------------------------- #
# The hermetic build
# --------------------------------------------------------------------------- #
def test_the_build_stores_one_record_per_kept_cell(
    built: m.EnvMetalTnseqRoyet2025Dataset, presence_only_pins: list[Mapping[str, str]]
) -> None:
    assert len(built) == SYNTHETIC_RECORDS
    assert presence_only_pins == [m.DATA_SHA256]


def test_the_built_records_carry_the_released_values(
    built: m.EnvMetalTnseqRoyet2025Dataset,
) -> None:
    records = [built[i] for i in range(len(built))]
    first = records[0]["experiment"]
    assert first["genotype"]["perturbations"][0]["systematic_gene_name"] == "PP_0001"
    assert first["genotype"]["perturbations"][0]["perturbed_gene_name"] == "parB"
    assert first["phenotype"]["environment_response"] == -2.0
    assert first["phenotype"]["measurement_type"] == "log2_ratio"
    values = [r["experiment"]["phenotype"]["environment_response"] for r in records]
    assert values == [-2.0, 0.0] * 4
    references = {
        r["reference"]["phenotype_reference"]["environment_response"] for r in records
    }
    assert references == {0.0}
    compounds = [
        r["experiment"]["environment"]["perturbations"][0]["compound"]["name"]
        for r in records[::2]
    ]
    assert compounds == [
        "cobalt chloride",
        "copper(II) chloride",
        "zinc chloride",
        "cadmium chloride",
    ]
    assert {r["publication"]["doi"] for r in records} == {m.PAPER_DOI}


def test_the_build_writes_every_ledger(built: m.EnvMetalTnseqRoyet2025Dataset) -> None:
    out = Path(built.preprocess_dir)
    log = json.loads((out / "dropped_records.json").read_text())
    assert (log["kept_records"], log["dropped_records"]) == (8, 8)
    reconciliation = json.loads((out / "locus_tag_reconciliation.json").read_text())
    assert reconciliation["status_histogram"]["current"] == 3
    assert json.loads((out / "replicate_structure.json").read_text()) == {
        "table_s3_pools_per_arm": {"LB": 2, "Co": 2, "Cu": 2, "Zn": 2, "Cd": 2}
    }
    gaps = json.loads((out / "perturbation_field_gaps.json").read_text())
    assert [g["field"] for g in gaps] == [
        "barcode",
        "insertion_position",
        "insertion_strand",
    ]
    unstored = json.loads((out / "not_stored.json").read_text())
    assert unstored["columns"] == [
        "Mean A",
        "Mean B",
        "Delta sum",
        "p-value",
        "q-value",
    ]
    assert unstored["values"]["Co"]["PP_0001"] == [10.0, 2.5, 0.0, 0.0, 0.0]
    sourced = json.loads((out / "sourced_values.json").read_text())
    assert sorted(sourced) == sorted(m.SOURCED_VALUES)


def test_the_build_refuses_a_count_the_module_does_not_declare(
    synthetic: Path, kt2440: PPutidaKT2440Genome, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(m, "EXPECTED_RECORDS", SYNTHETIC_RECORDS + 1)
    with pytest.raises(RuntimeError, match="records kept"):
        m.EnvMetalTnseqRoyet2025Dataset(root=str(synthetic), pputida_genome=kt2440)


def test_the_build_refuses_a_tag_that_is_not_a_current_locus_tag(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    oracles: None,
    presence_only_pins: list[Mapping[str, str]],
    kt2440: PPutidaKT2440Genome,
) -> None:
    root = tmp_path / "refseq"
    rows = (("PP_RS00005", *SYNTHETIC_ROWS[0][1:]), *SYNTHETIC_ROWS[1:])
    _write_s5(root / "raw" / m.TABLE_S5, rows)
    _write_s3(root / "raw" / m.TABLE_S3)
    with pytest.raises(m.SheetFormatError, match="not current locus tags"):
        m.EnvMetalTnseqRoyet2025Dataset(root=str(root), pputida_genome=kt2440)


def test_the_built_store_passes_l0_to_l4(
    built: m.EnvMetalTnseqRoyet2025Dataset, kt2440: PPutidaKT2440Genome, tmp_path: Path
) -> None:
    built.close_lmdb()
    report = m.verify_build(
        built.root,
        genome=kt2440,
        data_root=str(tmp_path),
        expected_count=SYNTHETIC_RECORDS,
    )
    assert report.passed, report.summary()
    assert Path(built.root, "preprocess", "verification_report.json").exists()


def test_the_dataset_is_registered_and_declares_its_classes() -> None:
    assert dataset_registry["EnvMetalTnseqRoyet2025Dataset"] is (
        m.EnvMetalTnseqRoyet2025Dataset
    )
    dataset = m.EnvMetalTnseqRoyet2025Dataset.__new__(m.EnvMetalTnseqRoyet2025Dataset)
    assert dataset.experiment_class is BacterialEnvironmentResponseExperiment
    assert dataset.reference_class is BacterialEnvironmentResponseExperimentReference
    assert dataset.raw_file_names == [m.TABLE_S5, m.TABLE_S3]
    assert dataset.preprocess_raw("x") == "x"
    with pytest.raises(NotImplementedError):
        dataset.create_experiment()


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def test_main_deposit_build_and_verify(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    monkeypatch.setattr("dotenv.load_dotenv", lambda *a, **k: True)
    deposited: list[dict[str, Any]] = []

    def deposit(*, sources: Mapping[str, Any], data_root: str) -> Path:
        deposited.append(dict(sources))
        return tmp_path

    monkeypatch.setattr(m, "deposit_raw_mirror", deposit)
    monkeypatch.setattr(
        m, "retrieve_raw_files", lambda d: {m.TABLE_S5: Path(d) / m.TABLE_S5}
    )
    assert m.main(["deposit"]) == 0
    assert deposited[0][m.TABLE_S5] == (
        tmp_path / m.LIBRARY_DIR_REL / "si" / "si9.xlsx"
    )
    assert deposited[0][m.TABLE_S3] == (
        tmp_path / m.LIBRARY_DIR_REL / "si" / "si7.xlsx"
    )
    assert m.main(["deposit", "--retrieve-into", str(tmp_path / "f")]) == 0
    assert deposited[1] == {m.TABLE_S5: tmp_path / "f" / m.TABLE_S5}

    class Built:
        def __init__(self, root: str) -> None:
            self.root = root

        def __len__(self) -> int:
            return 3

    monkeypatch.setattr(m, "EnvMetalTnseqRoyet2025Dataset", Built)
    assert m.main(["build"]) == 0
    assert "len = 3" in capsys.readouterr().out

    class Report:
        passed = False

        def summary(self) -> str:
            return "summary"

    monkeypatch.setattr(m, "verify_build", lambda root, data_root: Report())
    assert m.main(["verify"]) == 1


# --------------------------------------------------------------------------- #
# Data-gated: the real mirror and the built dev-tree store
# --------------------------------------------------------------------------- #
@pytest.mark.data
def test_every_sourced_value_is_a_verbatim_quote_of_the_pinned_paper() -> None:
    library = Path(os.environ["DATA_ROOT"]) / "torchcell-library"
    for value in m.SOURCED_VALUES.values():
        result = audit_sourced_value(value, library)
        assert result.passed, result.message


@pytest.mark.data
def test_the_raw_mirror_carries_the_pinned_bytes() -> None:
    manifest = m.load_manifest()
    for raw in m.RAW_FILES:
        assert m.manifest_sha256(manifest, raw.mirror_relpath) == raw.sha256
        assert m._sha256(m.raw_mirror_dir() / raw.mirror_relpath) == raw.sha256


@pytest.mark.data
def test_the_real_release_matches_the_declared_oracles() -> None:
    path = m.raw_mirror_dir() / "data" / m.TABLE_S5
    sheets = m.read_release(path)
    log = m.build_drop_log("royet", sheets)
    assert log.kept_records == m.EXPECTED_RECORDS == 21583
    assert {x.metal: x.dropped_empty for x in log.per_metal} == m.EXPECTED_DROPPED
    assert {x.metal: x.kept_q_le_threshold for x in log.per_metal} == {
        "Co": 9,
        "Cu": 14,
        "Zn": 3,
        "Cd": 8,
    }
    assert sum(x.dropped_empty_states.get("ES", 0) for x in log.per_metal) == 989
    assert sum(x.kept_exact_zero for x in log.per_metal) == 717
    # hand-checked cells of si9.xlsx: copA1 in copper and PP_1663 in cadmium
    copper = {r.orf: r for r in sheets["Cu"]}
    cadmium = {r.orf: r for r in sheets["Cd"]}
    assert (copper["PP_0586"].log2fc, copper["PP_0586"].mean_a) == (-8.19, 197.3)
    assert (cadmium["PP_1663"].log2fc, cadmium["PP_1663"].mean_b) == (-8.22, 1.2)
    pools = m.read_replicate_pools(m.raw_mirror_dir() / "data" / m.TABLE_S3)
    assert set(pools.values()) == {2}
