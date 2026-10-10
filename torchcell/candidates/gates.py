# torchcell/candidates/gates.py
# [[torchcell.candidates.gates]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/candidates/gates.py
# Test file: tests/torchcell/candidates/test_gates.py
"""Deterministic evaluators for the five candidate gates.

Every function here reads bytes we already hold: a candidate row's fields, the dataset
modules' source, the genomes tier, the literature and raw mirror manifests, a dev store's
LMDB (read-only), the committed compound identity table. Nothing fetches, nothing builds,
and nothing that needs reading a paper is decided here: the agentic halves of G1 (a
borderline aggregation), G3 (which SI file is which table) and G5 (the proposed class) come
in as typed inputs (:class:`SourceKindFinding`, a phenotype class name, an issue number).

What each gate reads:

- **G1** a :class:`SourceKindFinding` when one is recorded for the row
  (:mod:`torchcell.candidates.findings`), else the row's ``status``, ``klass`` and
  ``accession_confirmed``: a row marked as an aggregation without a finding is a borderline
  row and stays ``unmeasured``, as does a row whose accession was never confirmed.
- **G2** host strains named in the row's text against :data:`HOSTS`, then the genomes tier
  through :func:`torchcell.sequence.genome.registry.resolve`. Readable means the set is in
  the schema's ``BacterialAssemblySet`` vocabulary or is one of the two yeast sets a loader
  already reads (:data:`READABLE_ASSEMBLY_SETS`).
- **G3** the ``si/`` and ``data/`` entries of ``torchcell-library/<key>/manifest.json`` and
  ``torchcell-raw/<key>/manifest.json``, kept apart (#495).
- **G4** DOI, accessions and PMID against every module under ``torchcell/datasets`` (G4-key),
  then a released table against a dev store's records on matched samples (G4-value).
- **G5** the phenotype class against ``torchcell.datamodels.schema`` and compound names
  against ``compound_identity_table.json``, offline (#726).
"""

from __future__ import annotations

import ast
import hashlib
import io
import os.path as osp
import pickle
import re
import zipfile
import zlib
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any, Final, Literal, get_args

import lmdb
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from torchcell.candidates.verdict import (
    G2_GATE,
    AggregationRecord,
    CandidateTable,
    CandidateVerdict,
    G2Branch,
    G2Host,
    G2Record,
    G4KeyHit,
    G4KeyRecord,
    G4Relation,
    G4ValueMatch,
    G4ValueRecord,
    G5Record,
    GateName,
    GateResult,
    InventoryFile,
    InventoryRecord,
    SheetCount,
    SourceKind,
    ZoteroItem,
)
from torchcell.datamodels import schema
from torchcell.datamodels.compound_identity import resolve_compound_identity
from torchcell.literature.manifest import MANIFEST_FILENAME, Manifest, RetrievalMethod
from torchcell.sequence.genome.registry import (
    ECOLI_B_REL606,
    ECOLI_K12_BW25113,
    ECOLI_K12_MG1655,
    ECOLI_K12_W3110,
    PETER2018_1011,
    PPUTIDA_KT2440,
    SGD_S288C_R64,
    assembly_set_dir,
    load_genome_manifest,
    resolve,
)
from torchcell.verification.sourced import SourcedValue

REPO: Final = Path(__file__).resolve().parents[2]
DATASETS_DIR: Final = REPO / "torchcell" / "datasets"
LIBRARY_DIR: Final = "torchcell-library"
RAW_DIR: Final = "torchcell-raw"


class _Frozen(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")


# --------------------------------------------------------------------------- #
# The candidate row, table-neutral
# --------------------------------------------------------------------------- #
#: Free-text fields scanned for host strain names, per table.
TEXT_FIELDS: Final[dict[str, tuple[str, ...]]] = {
    "bacteria": (
        "citation",
        "genotypes",
        "env",
        "phenotype",
        "seq_basis",
        "modality",
        "why",
        "accession",
        "schema_need",
    ),
    "yeast": (
        "citation",
        "genotypes",
        "env",
        "phenotype",
        "shape",
        "seq_basis",
        "why",
        "accession",
    ),
}

_DOI_URL = re.compile(r"^https?://(?:dx\.)?doi\.org/(.+?)/?$")


class CandidateRow(_Frozen):
    """The fields of either table's ``Candidate`` the gates read."""

    table: CandidateTable
    name: str
    organism: str = Field(description="'S. cerevisiae' for every yeast row")
    status: str
    klass: str
    seq_basis: str
    url: str
    accession: str
    accession_confirmed: bool | None = Field(
        description="the bacteria table's flag; None for the yeast table, which has none"
    )
    text: str = Field(description="the TEXT_FIELDS of the row joined by newlines")

    @property
    def doi(self) -> str | None:
        """The DOI in the row's ``url``, when the url is a doi.org link."""
        match = _DOI_URL.match(self.url.strip())
        return match.group(1) if match else None


def row_from_table(table: CandidateTable, row: BaseModel) -> CandidateRow:
    """Project a table script's ``Candidate`` onto :class:`CandidateRow`."""
    data = row.model_dump()
    return CandidateRow(
        table=table,
        name=data["name"],
        organism=data["organism"] if table == "bacteria" else "S. cerevisiae",
        status=data["status"],
        klass=data["klass"],
        seq_basis=data["seq_basis"],
        url=data["url"],
        accession=data["accession"],
        accession_confirmed=data["accession_confirmed"]
        if table == "bacteria"
        else None,
        text="\n".join(str(data[f]) for f in TEXT_FIELDS[table]),
    )


def find_row(rows: Iterable[CandidateRow], name: str) -> CandidateRow:
    """The row named ``name`` exactly, else the one row whose name starts with it."""
    pool = list(rows)
    exact = [r for r in pool if r.name == name]
    if exact:
        return exact[0]
    prefixed = [r for r in pool if r.name.startswith(name)]
    if len(prefixed) != 1:
        names = ", ".join(repr(r.name) for r in prefixed) or "none"
        raise LookupError(f"row {name!r}: {len(prefixed)} prefix matches ({names})")
    return prefixed[0]


# --------------------------------------------------------------------------- #
# Gate assembly helpers
# --------------------------------------------------------------------------- #
def unmeasured(gate: GateName, reason: str) -> GateResult:
    """A gate that was not evaluated, and why."""
    return GateResult(gate=gate, outcome="unmeasured", reason=reason)


class GateRun(_Frozen):
    """One gate's result plus the record the verdict carries for it (None if none)."""

    result: GateResult
    record: Any = None


class SourceKindFinding(_Frozen):
    """A recorded G1 judgment for one row: the source kind and the quotes that decide it.

    The deterministic G1 cannot tell an aggregation from a transcription off row fields
    alone; a finding is the typed answer, written once from a settled-row module's sourced
    values (:mod:`torchcell.candidates.findings`) or returned by the agentic G1 stage.
    """

    row_name: str
    citation_key: str
    source_kind: SourceKind
    reason: str = Field(min_length=1)
    evidence: tuple[SourcedValue, ...] = Field(min_length=1)
    aggregation: AggregationRecord | None

    @model_validator(mode="after")
    def _aggregation_iff_kind(self) -> SourceKindFinding:
        if (self.source_kind == "aggregation") != (self.aggregation is not None):
            raise ValueError(
                f"{self.row_name}: aggregation record required for, and only for, "
                "an aggregation"
            )
        return self


class G1Run(_Frozen):
    """G1's result, the source kind it settled on, and the aggregation record."""

    result: GateResult
    source_kind: SourceKind | None
    aggregation: AggregationRecord | None


# --------------------------------------------------------------------------- #
# G1: primary measurement
# --------------------------------------------------------------------------- #
AGGREGATION_KLASS: Final = "Aggregation / support"


def evaluate_g1(row: CandidateRow, finding: SourceKindFinding | None) -> G1Run:
    """Classify the row's source kind; pass primary and recorded aggregations."""
    if finding is not None:
        if finding.row_name != row.name:
            raise ValueError(f"finding for {finding.row_name!r} given for {row.name!r}")
        passes = finding.source_kind in ("primary", "aggregation")
        return G1Run(
            result=GateResult(
                gate="G1",
                outcome="pass" if passes else "fail",
                reason=f"{finding.source_kind}: {finding.reason}",
                evidence=finding.evidence,
            ),
            source_kind=finding.source_kind,
            aggregation=finding.aggregation,
        )
    if row.status == "aggregation" or row.klass == AGGREGATION_KLASS:
        return G1Run(
            result=unmeasured(
                "G1",
                f"status={row.status!r}, klass={row.klass!r} mark an aggregation and no "
                "source-kind finding is recorded: aggregation or transcription is the "
                "agentic G1 judgment, with the verbatim quote as evidence",
            ),
            source_kind=None,
            aggregation=None,
        )
    if row.accession_confirmed is False:
        return G1Run(
            result=unmeasured(
                "G1",
                "accession_confirmed=False: whether the values sit in a deposited "
                "artifact is not established",
            ),
            source_kind=None,
            aggregation=None,
        )
    return G1Run(
        result=GateResult(
            gate="G1",
            outcome="pass",
            reason=f"primary: status={row.status!r}, klass={row.klass!r}, no "
            "aggregation marker",
        ),
        source_kind="primary",
        aggregation=None,
    )


# --------------------------------------------------------------------------- #
# G2: sequence route
# --------------------------------------------------------------------------- #
class HostToken(_Frozen):
    """A host strain name as candidate rows write it, and its assembly set if deposited."""

    name: str
    organism: str
    pattern: str = Field(description="regex matched against the row text")
    assembly_set: str | None


#: Hosts a row can name. ``assembly_set`` None: a public strain with no set in the tier
#: yet, so naming it makes the row ``blocked:provision``, never refused.
HOSTS: Final[tuple[HostToken, ...]] = (
    HostToken(
        name="MG1655",
        organism="E. coli",
        pattern=r"\bMG1655\b",
        assembly_set=ECOLI_K12_MG1655,
    ),
    HostToken(
        name="BW25113",
        organism="E. coli",
        pattern=r"\bBW25113\b|\bKeio\b",
        assembly_set=ECOLI_K12_BW25113,
    ),
    HostToken(
        name="W3110",
        organism="E. coli",
        pattern=r"\bW3110\b",
        assembly_set=ECOLI_K12_W3110,
    ),
    HostToken(
        name="REL606",
        organism="E. coli",
        pattern=r"\bREL606\b",
        assembly_set=ECOLI_B_REL606,
    ),
    HostToken(
        name="KT2440",
        organism="P. putida",
        pattern=r"\bKT2440\b",
        assembly_set=PPUTIDA_KT2440,
    ),
    HostToken(name="BL21", organism="E. coli", pattern=r"\bBL21\b", assembly_set=None),
    HostToken(name="DH1", organism="E. coli", pattern=r"\bDH1\b", assembly_set=None),
    HostToken(
        name="DH5alpha",
        organism="E. coli",
        pattern=r"\bDH5(?:α|a|alpha)\b",
        assembly_set=None,
    ),
    HostToken(
        name="DH10B", organism="E. coli", pattern=r"\bDH10B\b", assembly_set=None
    ),
    HostToken(
        name="JM109", organism="E. coli", pattern=r"\bJM109\b", assembly_set=None
    ),
    HostToken(
        name="Nissle 1917", organism="E. coli", pattern=r"\bNissle\b", assembly_set=None
    ),
    HostToken(
        name="MDS42", organism="E. coli", pattern=r"\bMDS42\b", assembly_set=None
    ),
    HostToken(
        name="EM42", organism="P. putida", pattern=r"\bEM42\b", assembly_set=None
    ),
    HostToken(
        name="DOT-T1E", organism="P. putida", pattern=r"\bDOT-T1E\b", assembly_set=None
    ),
    HostToken(
        name="S288C",
        organism="S. cerevisiae",
        pattern=r"\bS288C\b|\bBY47\d\d\b|\bBY4716\b",
        assembly_set=SGD_S288C_R64,
    ),
    HostToken(
        name="Peter 2018 isolates",
        organism="S. cerevisiae",
        pattern=r"\b1,?011\b",
        assembly_set=PETER2018_1011,
    ),
    HostToken(
        name="CEN.PK", organism="S. cerevisiae", pattern=r"\bCEN\.PK", assembly_set=None
    ),
    HostToken(
        name="W303", organism="S. cerevisiae", pattern=r"\bW303\b", assembly_set=None
    ),
    HostToken(
        name="Sigma1278b",
        organism="S. cerevisiae",
        pattern=r"(?:\bSigma|Σ)1278b\b",
        assembly_set=None,
    ),
    HostToken(
        name="RM11-1a", organism="S. cerevisiae", pattern=r"\bRM11", assembly_set=None
    ),
    HostToken(
        name="SK1", organism="S. cerevisiae", pattern=r"\bSK1\b", assembly_set=None
    ),
)

#: "K-12" names the lineage, not the strain (the bacteria table's SeqBasis note: the
#: exact background is a per-row provenance item recorded when the loader is written).
#: A row naming only the lineage resolves to the K-12 sets the tier can read.
K12_LINEAGE: Final = re.compile(r"\bK-12\b")
K12_READABLE_HOSTS: Final = ("MG1655", "BW25113")

#: Sets a genome class already reads: the schema's bacterial vocabulary (W3110 is
#: deliberately outside it, registry.py) plus the two yeast sets the loaders resolve.
READABLE_ASSEMBLY_SETS: Final[frozenset[str]] = frozenset(
    (*get_args(schema.BacterialAssemblySet), SGD_S288C_R64, PETER2018_1011)
)

_ORGANISM_OF_TABLE_ROW: Final = {"E. coli", "P. putida", "S. cerevisiae"}


def hosts_named(row: CandidateRow) -> tuple[str, ...]:
    """Host names the row's text names, in :data:`HOSTS` order, for the row's organism."""
    if row.organism not in _ORGANISM_OF_TABLE_ROW:
        raise ValueError(f"{row.name}: unknown organism {row.organism!r}")
    named = [
        h.name
        for h in HOSTS
        if h.organism == row.organism and re.search(h.pattern, row.text)
    ]
    if not named and row.organism == "E. coli" and K12_LINEAGE.search(row.text):
        named = list(K12_READABLE_HOSTS)
    return tuple(named)


def g2_branch(row: CandidateRow) -> G2Branch:
    """``haplotype_mosaic`` for a segregant panel, ``assembly_set`` otherwise."""
    return "haplotype_mosaic" if row.seq_basis == "segregant-WGS" else "assembly_set"


def probe_host(name: str, data_root: str, *, verify: bool = True) -> G2Host:
    """What the genomes tier says about one host: resolvable, then readable."""
    token = next(h for h in HOSTS if h.name == name)
    if token.assembly_set is None:
        return G2Host(name=name, assembly_set=None, resolvable=False, readable=False)
    manifest_path = (
        Path(assembly_set_dir(token.assembly_set, data_root)) / MANIFEST_FILENAME
    )
    if not manifest_path.is_file():
        return G2Host(
            name=name, assembly_set=token.assembly_set, resolvable=False, readable=False
        )
    manifest = load_genome_manifest(token.assembly_set, data_root)
    for rec in manifest.files:
        resolve(token.assembly_set, rec.path, verify=verify, data_root=data_root)
    return G2Host(
        name=name,
        assembly_set=token.assembly_set,
        resolvable=True,
        readable=token.assembly_set in READABLE_ASSEMBLY_SETS,
    )


def evaluate_g2(
    row: CandidateRow,
    data_root: str,
    *,
    source_kind: SourceKind | None,
    verify: bool = True,
) -> GateRun:
    """Resolve every host the row names; the worst host decides (plan decision 5)."""
    branch = g2_branch(row)
    names = hosts_named(row)
    if not names and source_kind == "aggregation":
        return GateRun(
            result=unmeasured(
                "G2",
                "an aggregation names no single host: hosts are per source record and "
                "are resolved per source study, not off the row",
            )
        )
    if not names or (branch == "haplotype_mosaic" and len(names) < 2):
        record = G2Record(
            branch=branch,
            outcome="absent_unnamed",
            hosts=tuple(probe_host(n, data_root, verify=verify) for n in names),
        )
        reason = (
            f"{branch}: the row names {len(names)} host(s) "
            f"({', '.join(names) or 'none'}); refused like Excluded.rule='no-sequence'"
        )
    else:
        hosts = tuple(probe_host(n, data_root, verify=verify) for n in names)
        record = G2Record(branch=branch, outcome=_worst(hosts), hosts=hosts)
        reason = f"{branch}: " + "; ".join(_host_phrase(h) for h in hosts)
    outcome, kind = G2_GATE[record.outcome]
    return GateRun(
        result=GateResult(gate="G2", outcome=outcome, reason=reason, blocked_kind=kind),
        record=record,
    )


def _worst(
    hosts: tuple[G2Host, ...],
) -> Literal[
    "resolvable_readable", "resolvable_not_ingestible", "absent_provisionable"
]:
    if any(h.assembly_set is None or not h.resolvable for h in hosts):
        return "absent_provisionable"
    if any(not h.readable for h in hosts):
        return "resolvable_not_ingestible"
    return "resolvable_readable"


def _host_phrase(host: G2Host) -> str:
    if host.assembly_set is None:
        return f"{host.name} has no assembly set (genomes-tier deposit PR first)"
    if not host.resolvable:
        return f"{host.name} -> {host.assembly_set} is not in the tier"
    if not host.readable:
        return (
            f"{host.name} -> {host.assembly_set} resolves but no genome class reads it"
        )
    return f"{host.name} -> {host.assembly_set} resolves and reads"


# --------------------------------------------------------------------------- #
# G3: deposit inventory
# --------------------------------------------------------------------------- #
_EOCD_SIGNATURE: Final = b"PK\x05\x06"
_EOCD_FIXED_BYTES: Final = 22
_XLSX_SUFFIXES: Final = (".xlsx", ".xlsm")
INVENTORY_PREFIXES: Final = ("si/", "data/")


def leading_archive(data: bytes) -> bytes:
    """The bytes up to and including the FIRST end-of-central-directory record.

    The concatenated-archive pattern of ``torchcell/datasets/ecoli/tong2020.py``: a file
    that is a complete zip followed by more bytes reads as the trailing archive, so the
    leading archive is hashed separately.
    """
    eocd = data.find(_EOCD_SIGNATURE)
    if eocd < 0:
        raise ValueError("no zip end-of-central-directory record in the bytes")
    comment_length = int.from_bytes(data[eocd + 20 : eocd + 22], "little")
    return data[: eocd + _EOCD_FIXED_BYTES + comment_length]


def _zip_ok(data: bytes) -> bool:
    """``testzip()`` over the archive; a directory that will not open, or a member
    whose deflate stream is corrupt (``zlib.error``, raised before the CRC check), is bad.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return archive.testzip() is None
    except (zipfile.BadZipFile, zlib.error):
        return False


def _sheet_counts(data: bytes) -> tuple[SheetCount, ...]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    counts = tuple(
        SheetCount(sheet=ws.title, rows=sum(1 for _ in ws.iter_rows(values_only=True)))
        for ws in workbook.worksheets
    )
    workbook.close()
    return counts


def inventory_file(
    tree: Literal["library", "raw"], root: Path, entry: Any
) -> InventoryFile:
    """Measure one manifest entry's bytes (``entry`` is an ``ArtifactRecord``)."""
    path = root / entry.path
    method = entry.retrieval.method.value if entry.retrieval is not None else None
    base: dict[str, Any] = {
        "tree": tree,
        "path": entry.path,
        "role": entry.role,
        "bytes": entry.bytes,
        "sha256_manifest": entry.sha256,
        "retrieval_method": method,
    }
    if not path.is_file():
        return InventoryFile(
            **base,
            sha256_disk=None,
            is_zip=False,
            zip_ok=None,
            leading_archive_sha256=None,
        )
    data = path.read_bytes()
    is_zip = data.startswith(b"PK\x03\x04")
    leading_sha: str | None = None
    zip_ok: bool | None = None
    sheets: tuple[SheetCount, ...] = ()
    if is_zip and _EOCD_SIGNATURE not in data:
        zip_ok = False  # a truncated archive: no end-of-central-directory record
    elif is_zip:
        leading = leading_archive(data)
        if len(leading) < len(data):
            leading_sha = hashlib.sha256(leading).hexdigest()
        zip_ok = _zip_ok(leading)
        if zip_ok and entry.path.lower().endswith(_XLSX_SUFFIXES):
            sheets = _sheet_counts(leading)
    return InventoryFile(
        **base,
        sha256_disk=hashlib.sha256(data).hexdigest(),
        is_zip=is_zip,
        zip_ok=zip_ok,
        leading_archive_sha256=leading_sha,
        sheets=sheets,
    )


def inventory(citation_key: str, data_root: str) -> InventoryRecord | None:
    """Walk both mirrors' manifests for ``si/`` and ``data/``; None when neither exists."""
    trees: list[Literal["library", "raw"]] = []
    files: list[InventoryFile] = []
    walls: list[str] = []
    for tree, top in (("library", LIBRARY_DIR), ("raw", RAW_DIR)):
        root = Path(data_root) / top / citation_key
        manifest_path = root / MANIFEST_FILENAME
        if not manifest_path.is_file():
            continue
        trees.append(tree)  # type: ignore[arg-type]
        manifest = Manifest.model_validate_json(manifest_path.read_text())
        for entry in manifest.files:
            if not entry.path.startswith(INVENTORY_PREFIXES):
                continue
            measured = inventory_file(tree, root, entry)  # type: ignore[arg-type]
            files.append(measured)
            if (
                measured.sha256_disk is None
                and entry.retrieval is not None
                and entry.retrieval.method == RetrievalMethod.manual_browser
            ):
                walls.append(f"{tree}:{entry.path}: {entry.retrieval.params}")
    if not trees:
        return None
    return InventoryRecord(
        citation_key=citation_key,
        trees_read=tuple(trees),
        files=tuple(files),
        walls=tuple(walls),
    )


def evaluate_g3(record: InventoryRecord | None) -> GateRun:
    """Pass an intact deposit; block on a wall or a missing scripted file; fail on drift."""
    if record is None:
        return GateRun(
            result=unmeasured("G3", "no manifest in the literature or the raw mirror")
        )
    if not record.files:
        return GateRun(
            result=unmeasured(
                "G3",
                f"the manifests in {record.trees_read} list nothing under si/ or data/",
            ),
            record=None,
        )
    drifted = [f for f in record.files if f.sha256_disk is not None and not f.intact]
    missing = [f for f in record.files if f.sha256_disk is None]
    counted = sum(s.rows for f in record.files for s in f.sheets)
    summary = (
        f"{len(record.files)} file(s) in {'+'.join(record.trees_read)}, "
        f"{counted} worksheet row(s) counted"
    )
    if drifted:
        names = ", ".join(f"{f.tree}:{f.path}" for f in drifted)
        result = GateResult(
            gate="G3",
            outcome="fail",
            reason=f"{summary}; sha256 drift or a bad zip member: {names}",
        )
    elif record.walls:
        result = GateResult(
            gate="G3",
            outcome="blocked",
            blocked_kind="wall",
            reason=f"{summary}; behind a manual_browser wall: {'; '.join(record.walls)}",
        )
    elif missing:
        names = ", ".join(f"{f.tree}:{f.path}" for f in missing)
        result = GateResult(
            gate="G3",
            outcome="blocked",
            blocked_kind="deposit",
            reason=f"{summary}; listed but not on disk, re-run the retrieval: {names}",
        )
    else:
        result = GateResult(gate="G3", outcome="pass", reason=f"{summary}; all intact")
    return GateRun(result=result, record=record)


# --------------------------------------------------------------------------- #
# G4-key: identifiers against every dataset module
# --------------------------------------------------------------------------- #
#: Accession shapes, from experiments/database/scripts/check_candidate_overlap.py (a
#: bare number is never enough, so the PubChem form keeps its AID prefix).
ACCESSION: Final = re.compile(
    r"(GSE\d+|PXD\d+|PRJ[EN][AB]\d+|E-MTAB-\d+|SRP\d+|AID\s*\d+|MTBLS\d+|"
    r"10\.5061/dryad\.[A-Za-z0-9]+)"
)
#: Module-level names a dataset module declares its own paper's DOI under.
OWN_DOI_NAMES: Final = frozenset({"DOI", "PAPER_DOI"})


class ModuleKeys(_Frozen):
    """What a dataset module declares and says: its key, its DOI, its source text."""

    path: str = Field(description="repo-relative path of the module")
    citation_key: str | None
    own_dois: tuple[str, ...]
    text: str


def _module_constants(tree: ast.Module) -> dict[str, str]:
    out: dict[str, str] = {}
    for node in tree.body:
        target: ast.expr | None = None
        value: ast.expr | None = None
        if isinstance(node, ast.Assign) and len(node.targets) == 1:
            target, value = node.targets[0], node.value
        elif isinstance(node, ast.AnnAssign):
            target, value = node.target, node.value
        if (
            isinstance(target, ast.Name)
            and isinstance(value, ast.Constant)
            and isinstance(value.value, str)
        ):
            out[target.id] = value.value
    return out


def module_index(datasets_dir: Path = DATASETS_DIR) -> tuple[ModuleKeys, ...]:
    """Every module under ``datasets_dir``, parsed with ``ast`` (nothing is imported)."""
    out: list[ModuleKeys] = []
    for path in sorted(datasets_dir.rglob("*.py")):
        text = path.read_text(encoding="utf-8")
        constants = _module_constants(ast.parse(text))
        out.append(
            ModuleKeys(
                path=path.relative_to(datasets_dir.parents[1]).as_posix(),
                citation_key=constants.get("CITATION_KEY"),
                own_dois=tuple(
                    sorted(v for k, v in constants.items() if k in OWN_DOI_NAMES)
                ),
                text=text,
            )
        )
    return tuple(out)


def citation_key_for_doi(
    doi: str, index: tuple[ModuleKeys, ...], data_root: str | None
) -> str | None:
    """The citation key a DOI is filed under: a module's own DOI, else a mirror manifest."""
    keys = {m.citation_key for m in index if doi in m.own_dois and m.citation_key}
    if len(keys) > 1:
        raise LookupError(
            f"{doi}: modules disagree on its citation key: {sorted(keys)}"
        )
    if keys:
        return keys.pop()
    if data_root is None:
        return None
    for top in (LIBRARY_DIR, RAW_DIR):
        base = Path(data_root) / top
        if not base.is_dir():
            continue
        for manifest_path in sorted(base.glob(f"*/{MANIFEST_FILENAME}")):
            manifest = Manifest.model_validate_json(manifest_path.read_text())
            if manifest.doi is not None and manifest.doi.lower() == doi.lower():
                return manifest.citation_key
    return None


def evaluate_g4_key(
    row: CandidateRow,
    citation_key: str,
    index: tuple[ModuleKeys, ...],
    *,
    doi_to_pmid: Mapping[str, str] | None,
    aggregator_pmids: Mapping[str, frozenset[str]] | None,
) -> GateRun:
    """DOI, accession and PMID against every dataset module; fail on another's key."""
    doi = row.doi
    own_modules = {m.path for m in index if m.citation_key == citation_key}
    hits: list[G4KeyHit] = []
    relation: G4Relation
    if doi is not None:
        for m in index:
            if doi in m.own_dois:
                relation = "own" if m.path in own_modules else "primary_of_other"
                hits.append(
                    G4KeyHit(key_kind="doi", key=doi, where=m.path, relation=relation)
                )
            elif doi in m.text:
                relation = "own" if m.path in own_modules else "mentioned"
                hits.append(
                    G4KeyHit(key_kind="doi", key=doi, where=m.path, relation=relation)
                )
    accessions = tuple(sorted(set(ACCESSION.findall(f"{row.accession} {row.url}"))))
    for token in accessions:
        for m in index:
            if token in m.text:
                relation = "own" if m.path in own_modules else "primary_of_other"
                hits.append(
                    G4KeyHit(
                        key_kind="accession", key=token, where=m.path, relation=relation
                    )
                )
    pmid = None
    if doi is not None and doi_to_pmid is not None:
        pmid = doi_to_pmid.get(doi.lower())
    pmid_checked = pmid is not None and aggregator_pmids is not None
    if pmid is not None and aggregator_pmids is not None:
        for name, pmids in sorted(aggregator_pmids.items()):
            if pmid in pmids:
                hits.append(
                    G4KeyHit(
                        key_kind="pmid",
                        key=pmid,
                        where=name,
                        relation="aggregator_source",
                    )
                )
    record = G4KeyRecord(
        doi=doi,
        pmid=pmid,
        accessions=accessions,
        hits=tuple(hits),
        pmid_checked=pmid_checked,
    )
    return GateRun(result=g4_result(record, None), record=record)


def g4_result(key: G4KeyRecord, value: G4ValueRecord | None) -> GateResult:
    """G4 from both halves: fail on a key held elsewhere or on subsumed values."""
    overlapping = key.overlapping
    if overlapping:
        where = "; ".join(f"{h.key_kind} {h.key} in {h.where}" for h in overlapping)
        return GateResult(gate="G4", outcome="fail", reason=f"already held: {where}")
    if value is not None and value.subsumed:
        worst = max(m.best_max_abs_diff for m in value.matches)
        margins = [m.margin for m in value.matches if m.margin is not None]
        margin = f", nearest-wrong margin >= {min(margins):.3g}" if margins else ""
        return GateResult(
            gate="G4",
            outcome="fail",
            reason=f"subsumed: every released column matches a sample of {value.store} "
            f"within {value.tolerance} (max abs diff {worst:.3g}{margin})",
        )
    mentioned = sum(h.relation == "mentioned" for h in key.hits)
    pmid = "PMID checked against aggregators" if key.pmid_checked else "no PMID check"
    values = (
        f"values checked on {len(value.matches)} column(s), not subsumed"
        if value is not None
        else "values not compared (no dev store named)"
    )
    return GateResult(
        gate="G4",
        outcome="pass",
        reason=f"no key held by another module ({mentioned} mention(s)); {pmid}; {values}",
    )


# --------------------------------------------------------------------------- #
# G4-value: released values against a dev store, read-only
# --------------------------------------------------------------------------- #
class DirtyStoreError(RuntimeError):
    """The dev store was built from a tree with uncommitted changes (gotcha 6)."""


def _dotted(record: Any, path: str) -> Any:
    out = record
    for part in path.split("."):
        out = out[int(part)] if isinstance(out, list) else out[part]
    return out


def read_store_values(
    store_dir: str,
    *,
    sample_path: str,
    key_path: str,
    value_path: str,
    samples: frozenset[str] | None = None,
) -> tuple[int, str | None, dict[str, dict[str, float]]]:
    """``(records read, build commit, {sample: {key: value}})`` off a dev store's LMDB.

    Opens ``processed/lmdb`` read-only and never instantiates the dataset class, so no
    ``process()`` or download can run. Refuses a store with no build manifest and one
    built from a dirty tree. Interned ``$ref`` pointers are resolved the way
    ``ExperimentDataset.get_single_item`` resolves them.
    """
    from torchcell.data.experiment_dataset import resolve_interned
    from torchcell.provenance.build_manifest import BuildManifest

    root = Path(store_dir)
    manifest_path = root / "preprocess" / "build_manifest.json"
    if not manifest_path.is_file():
        raise FileNotFoundError(
            f"{manifest_path}: no build manifest; not a built store"
        )
    manifest = BuildManifest.model_validate_json(manifest_path.read_text())
    if manifest.torchcell_dirty:
        raise DirtyStoreError(
            f"{store_dir} was built from a dirty tree at {manifest.torchcell_commit}; "
            "rebuild it from a clean commit before comparing values"
        )
    lmdb_dir = root / "processed" / "lmdb"
    if not lmdb_dir.is_dir():
        raise FileNotFoundError(f"{lmdb_dir}: the store has no records LMDB")
    interned: dict[str, Any] = {}
    interned_dir = root / "processed" / "interned"
    if interned_dir.is_dir():
        ienv = lmdb.open(str(interned_dir), readonly=True, lock=False)
        with ienv.begin() as txn:
            for key, value in txn.cursor():
                interned[key.decode()] = pickle.loads(value)
        ienv.close()
    env = lmdb.open(str(lmdb_dir), readonly=True, lock=False, readahead=False)
    values: dict[str, dict[str, float]] = {}
    n = 0
    with env.begin() as txn:
        for _, raw in txn.cursor():
            n += 1
            record = resolve_interned(pickle.loads(raw), interned)
            sample = str(_dotted(record, sample_path))
            if samples is not None and sample not in samples:
                continue
            values.setdefault(sample, {})[str(_dotted(record, key_path))] = float(
                _dotted(record, value_path)
            )
    env.close()
    return n, manifest.torchcell_commit, values


def match_by_value(
    released: Mapping[str, Mapping[str, float]],
    served: Mapping[str, Mapping[str, float]],
) -> tuple[G4ValueMatch, ...]:
    """Score every released column against every served sample by max abs difference.

    A sample is scored only on the keys it shares with the column; a sample sharing none
    is skipped. The best and the runner-up are kept, so the margin of the identification
    is part of the record (the Thompson 2019 lysine measurement, PR #856).
    """
    out: list[G4ValueMatch] = []
    for column, reported in sorted(released.items()):
        scores: list[tuple[float, str, int]] = []
        for sample, held in sorted(served.items()):
            shared = [k for k in reported if k in held]
            if not shared:
                continue
            diff = max(abs(held[k] - reported[k]) for k in shared)
            scores.append((diff, sample, len(shared)))
        if not scores:
            raise LookupError(
                f"released column {column!r} shares no key with any sample"
            )
        scores.sort()
        best, runner = scores[0], scores[1] if len(scores) > 1 else None
        out.append(
            G4ValueMatch(
                released_column=column,
                best_sample=best[1],
                best_max_abs_diff=best[0],
                runner_up_sample=runner[1] if runner else None,
                runner_up_max_abs_diff=runner[0] if runner else None,
                keys_compared=best[2],
            )
        )
    return tuple(out)


# --------------------------------------------------------------------------- #
# G5: schema fit
# --------------------------------------------------------------------------- #
def phenotype_class_exists(name: str) -> bool:
    """``name`` is a ``Phenotype`` subclass defined in ``torchcell.datamodels.schema``."""
    candidate = getattr(schema, name, None)
    return isinstance(candidate, type) and issubclass(candidate, schema.Phenotype)


def evaluate_g5(
    phenotype_class: str,
    compounds: Iterable[str],
    *,
    issue: int | None,
    proposed_class: str | None = None,
) -> GateRun:
    """Pass when the class exists and every compound resolves; else a gap or unfiled."""
    exists = phenotype_class_exists(phenotype_class)
    unresolved = tuple(
        sorted(c for c in set(compounds) if not resolve_compound_identity(c).identified)
    )
    record = G5Record(
        phenotype_class=phenotype_class,
        class_exists=exists,
        unresolved_compounds=unresolved,
        proposed_class=proposed_class,
    )
    problems = []
    if not exists:
        problems.append(f"no Phenotype subclass {phenotype_class!r} in schema.py")
    if unresolved:
        problems.append(
            f"{len(unresolved)} compound(s) unresolved: {', '.join(unresolved)}"
        )
    if not problems:
        result = GateResult(
            gate="G5",
            outcome="pass",
            reason=f"{phenotype_class} exists; compounds resolve",
        )
    elif issue is not None:
        result = GateResult(
            gate="G5", outcome="gap", reason="; ".join(problems), issue=issue
        )
    else:
        result = GateResult(
            gate="G5",
            outcome="blocked",
            blocked_kind="unfiled_gap",
            reason="; ".join(problems) + "; file the issue and re-run with its number",
        )
    return GateRun(result=result, record=record)


# --------------------------------------------------------------------------- #
# Zotero and the verdict
# --------------------------------------------------------------------------- #
def zotero_item(citation_key: str, data_root: str) -> ZoteroItem:
    """``present`` when the literature mirror's manifest names a Zotero item."""
    path = Path(data_root) / LIBRARY_DIR / citation_key / MANIFEST_FILENAME
    if not path.is_file():
        return "absent"
    manifest = Manifest.model_validate_json(path.read_text())
    return "present" if manifest.zotero_item_key else "absent"


LATER_GATES: Final[tuple[GateName, ...]] = ("G2", "G3", "G4", "G5")


def compose_verdict(
    *,
    row: CandidateRow,
    citation_key: str,
    g1: G1Run,
    runs: Mapping[GateName, GateRun],
    zotero: ZoteroItem,
    decided_at: str,
    torchcell_commit: str,
    g4_value: G4ValueRecord | None = None,
) -> CandidateVerdict:
    """Apply the stop rule and assemble the verdict.

    ``runs`` holds G2..G5 where they were evaluated; a gate missing from it is
    ``unmeasured``. After the first ``fail`` or ``blocked`` every later gate becomes
    ``unmeasured`` and drops its record, whatever was computed for it.
    """
    results: list[GateResult] = [g1.result]
    records: dict[GateName, Any] = {}
    stopper: GateName | None = (
        "G1" if g1.result.outcome in ("fail", "blocked") else None
    )
    for gate in LATER_GATES:
        run = runs.get(gate)
        if stopper is not None:
            results.append(
                unmeasured(gate, f"not evaluated: {stopper} stopped the run")
            )
        elif run is None:
            results.append(unmeasured(gate, "not evaluated on this call"))
        else:
            results.append(run.result)
            records[gate] = run.record
            if run.result.outcome in ("fail", "blocked"):
                stopper = gate
    return CandidateVerdict(
        citation_key=citation_key,
        table=row.table,
        row_name=row.name,
        gates=tuple(results),
        zotero_item=zotero,
        source_kind=g1.source_kind,
        aggregation=g1.aggregation,
        g2=records.get("G2"),
        g3=records.get("G3"),
        g4_key=records.get("G4"),
        g4_value=g4_value if "G4" in records else None,
        g5=records.get("G5"),
        decided_at=decided_at,
        torchcell_commit=torchcell_commit,
    )


def aggregator_pmid_sets(data_root: str) -> dict[str, frozenset[str]] | None:
    """Source PMIDs of the built aggregators (check_candidate_overlap.py's AGGREGATORS).

    None when any aggregator file is not built, so a partial set is never read as
    "checked".
    """
    aggregators = (
        ("SynLethDB SL", "data/torchcell/syn_leth_db_yeast/raw/Yeast_SL.csv", 5),
        ("SynLethDB SR", "data/torchcell/syn_rescue_db_yeast/raw/Yeast_SR.csv", 5),
    )
    out: dict[str, frozenset[str]] = {}
    for name, rel, col in aggregators:
        path = osp.join(data_root, rel)
        if not osp.isfile(path):
            return None
        pmids = set()
        for line in Path(path).read_text(errors="replace").splitlines()[1:]:
            parts = line.split(",")
            if len(parts) > col and parts[col].strip().isdigit():
                pmids.add(parts[col].strip())
        out[name] = frozenset(pmids)
    return out
