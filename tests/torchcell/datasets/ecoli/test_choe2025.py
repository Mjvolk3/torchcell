# tests/torchcell/datasets/ecoli/test_choe2025.py
# [[tests.torchcell.datasets.ecoli.test_choe2025]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_choe2025.py
"""The Choe 2025 CRISPRi loader (``torchcell.datasets.ecoli.choe2025``).

Synthetic tests (run everywhere) write three stand-in workbooks with ``openpyxl`` in the
released sheet layout and build over the real ``EcoliK12MG1655Genome`` on a synthetic
assembly of ``tests/torchcell/sequence/genome/_bacterial_fixtures.py``, served through a
stubbed ``resolve`` with the network refused. The synthetic replicon is the fixtures' own
100 bp ``SEQUENCE``, and every synthetic guide window is one of the nine 23 nt windows of
that sequence whose spacer-plus-NGG-PAM reading is real, so the spacer check runs on
genuine arithmetic rather than on a stub.

Three synthetic loci plus one unresolvable b-number, laid out to trip every retention
rule exactly once:

    gene   b-number  locus span  strand  guide window  what it exercises
    thrL   b0001     1-30        +       15-37         a kept guide (log2 1 under CCCP)
    thrL   b0001     1-30        +       24-46         a second guide of one gene
    thrL   b0001     1-30        +       5-27          absent from Table S6 -> rule 4
    thrA   b0002     41-85       -       43-65         a kept guide
    thrA   b0002     41-85       -       55-77         released strand + -> rule 3
    thrA   b0002     41-85       -       48-70         zero LB control -> rule 5
    thrA   b0002     41-85       -       60-82         zero novobiocin -> rule 6
    yaaP   b0003     86-100      + (ps)  70-92         a kept pseudogene guide, and the
                                                       only zero-library guide
    yaaP   b0003     86-100      + (ps)  32-54         window misses the locus -> rule 2
    zzzA   b9999     (none)              25-47         b-number not on the genome -> 1

Data-gated tests (``@pytest.mark.data``) read the real raw mirror, the literature mirror
and the built dev-tree LMDB under ``$DATA_ROOT`` (they never build it): the pinned
sha256s, every sourced quote, and the released tables' own counts.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import os.path as osp
from collections.abc import Mapping, Sequence
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pandas as pd
import pytest
from openpyxl import Workbook

import torchcell.datasets.ecoli.choe2025 as m
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    MG1655_GAF,
    SEQUENCE,
    SyntheticLocus,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.datamodels.media import LB
from torchcell.datamodels.schema import (
    AssayType,
    AssemblyReferenceGenome,
    BacterialEnvironmentResponseExperiment,
    BacterialEnvironmentResponseExperimentReference,
    ConcentrationUnit,
    DoseBasis,
    MeasurementType,
    MediaComponentRole,
    SmallMoleculePerturbation,
)
from torchcell.datasets.bacteria_common import LocusTagResolutionError
from torchcell.literature.manifest import Manifest, RetrievalMethod
from torchcell.sequence.genome.ecoli.k12 import MG1655_ASSEMBLY, EcoliK12MG1655Genome

# --------------------------------------------------------------------------- #
# The synthetic assembly and its guide windows
# --------------------------------------------------------------------------- #
CHOE_LOCI = [
    SyntheticLocus(
        tag="b0001",
        parts=((1, 30),),
        strand="+",
        symbol="thrL",
        synonyms=("ECK0001",),
        product="thr operon leader peptide",
        protein_id="AAC73112.1",
        protein="MK",
    ),
    SyntheticLocus(
        tag="b0002",
        parts=((41, 85),),
        strand="-",
        symbol="thrA",
        synonyms=("ECK0002",),
        product="aspartokinase I",
        protein_id="AAC73113.1",
        protein="MRVLK",
    ),
    SyntheticLocus(
        tag="b0003",
        parts=((86, 100),),
        strand="+",
        symbol="yaaP",
        synonyms=("ECK0003",),
        pseudo=True,
    ),
]

REFERENCE = AssemblyReferenceGenome(
    species="Escherichia coli",
    strain="MG1655",
    assembly_set="ecoli_K12_MG1655_ASM584v2",
    assembly_accession="GCA_000005845.2",
)


def spacer_of(start: int, end: int) -> str:
    """The 20 nt spacer of a synthetic window, in whichever orientation carries the PAM.

    This mirrors the rule :func:`m.check_spacers` enforces, computed here from the
    fixtures' replicon so the synthetic Table S6 cannot disagree with the genome.
    """
    window = SEQUENCE[start - 1 : end]
    if window[m.SPACER_LENGTH + 1 : m.WINDOW_LENGTH] == "GG":
        return window[: m.SPACER_LENGTH]
    flipped = m.reverse_complement(window)
    if flipped[m.SPACER_LENGTH + 1 : m.WINDOW_LENGTH] == "GG":
        return flipped[: m.SPACER_LENGTH]
    raise AssertionError(f"{start}-{end} is not a PAM-bearing synthetic window")


def oligo(spacer: str) -> str:
    """A Table S6 oligo carrying ``spacer`` between the stated flanks."""
    return f"{m.OLIGO_PREFIX}{spacer}{m.OLIGO_SUFFIX}"


class Row(SimpleNamespace):
    """One synthetic Table S1 row before it is written to a sheet."""


#: ``(gene, start, end, released strand, fragment, LB value, in Table S6, library value,
#: zero in novobiocin)``. The LB value doubles under CCCP, so a kept record is log2 1
#: under CCCP and log2 0 under the other eleven antibiotics.
SYNTHETIC_ROWS: tuple[tuple[str, int, int, str, str, float, bool, float, bool], ...] = (
    ("thrL", 15, 37, "+", "1_1", 50.0, True, 100.0, False),
    ("thrL", 24, 46, "+", "2_1", 100.0, True, 100.0, False),
    ("thrL", 5, 27, "+", "3_1", 150.0, False, 100.0, False),
    ("thrA", 43, 65, "-", "1_2", 50.0, True, 100.0, False),
    ("thrA", 55, 77, "+", "2_2", 100.0, True, 100.0, False),
    ("thrA", 48, 70, "-", "3_2", 0.0, True, 100.0, False),
    ("thrA", 60, 82, "-", "4_2", 200.0, True, 100.0, True),
    ("yaaP", 70, 92, "+", "1_3", 100.0, True, 0.0, False),
    ("yaaP", 32, 54, "+", "2_3", 100.0, True, 100.0, False),
    ("zzzA", 25, 47, "+", "1_4", 100.0, True, 100.0, False),
)
#: ``{released symbol: b-number}``; ``zzzA``'s b-number is on no synthetic locus.
SYNTHETIC_B_NUMBERS = {
    "thrL": "b0001",
    "thrA": "b0002",
    "yaaP": "b0003",
    "zzzA": "b9999",
}
#: Hand-computed medians of (condition mean / library mean) over each gene's rows.
#: thrL: LB [0.5, 1.0, 1.5] -> 1.0; thrA: LB [0.5, 1.0, 0.0, 2.0] -> 0.75, novobiocin
#: [0.5, 1.0, 0.0, 0.0] -> 0.25; zzzA: [1.0] -> 1.0. CCCP doubles every numerator.
#: ``yaaP`` carries the zero-library guide, so the check excludes it and its value is
#: never compared.
SYNTHETIC_ER = {
    "thrL": {"LB": 1.0, "CCCP": 2.0, "other": 1.0},
    "thrA": {"LB": 0.75, "CCCP": 1.5, "other": 0.75, "Novobiocin": 0.25},
    "yaaP": {"LB": 1.0, "CCCP": 2.0, "other": 1.0},
    "zzzA": {"LB": 1.0, "CCCP": 2.0, "other": 1.0},
}
CCCP = "CCCP"
NOVOBIOCIN = "Novobiocin"


def _condition_value(row: Mapping[str, Any], label: str) -> float:
    """The synthetic abundance of one row in one condition."""
    if label == m.CONTROL_LABEL:
        return float(row["lb"])
    if label == CCCP:
        return 2.0 * float(row["lb"])
    if label == NOVOBIOCIN and row["zero_novobiocin"]:
        return 0.0
    return float(row["lb"])


def _rows() -> list[dict[str, Any]]:
    return [
        {
            "gene": gene,
            "start": start,
            "end": end,
            "strand": strand,
            "fragment": fragment,
            "lb": lb,
            "in_s6": in_s6,
            "library": library,
            "zero_novobiocin": zero_novobiocin,
        }
        for gene, start, end, strand, fragment, lb, in_s6, library, zero_novobiocin in (
            SYNTHETIC_ROWS
        )
    ]


def write_table_s1(path: Path, rows: Sequence[Mapping[str, Any]]) -> Path:
    """The released Table S1 layout: title, identifier names, sub-groups, sample labels."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(
        ["Table S1. Normalized sgRNA counts (reads per million mapped reads; RPM)."]
    )
    sheet.append([*m.S1_ID_COLUMNS, m.S1_VALUE_GROUP])
    sheet.append(
        [None] * len(m.S1_ID_COLUMNS) + ["w/o dCas9 (lib)", None, "w/ dCas9 (stress)"]
    )
    samples = [
        f"{label}_{replicate}"
        for label in m.S1_CONDITION_LABELS
        for replicate in (1, 2)
    ]
    sheet.append([None] * len(m.S1_ID_COLUMNS) + [*m.LIBRARY_COLUMNS, *samples])
    for row in rows:
        values: list[Any] = [
            row["gene"],
            row["start"],
            row["end"],
            row["strand"],
            row["fragment"],
            row["library"],
            row["library"],
        ]
        for label in m.S1_CONDITION_LABELS:
            value = _condition_value(row, label)
            values += [value, value]
        sheet.append(values)
    book.save(path)
    return path


def write_table_s2(path: Path, rows: Sequence[Mapping[str, Any]]) -> Path:
    """The released Table S2 layout, with the hand-computed ER of each synthetic gene."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Table S2. Enrichment ratio (ER) values of targeted genes."])
    sheet.append([*m.S2_ID_COLUMNS, "ER"])
    sheet.append([None] * len(m.S2_ID_COLUMNS) + list(m.S2_CONDITION_COLUMNS))
    counts = pd.Series([row["gene"] for row in rows]).value_counts()
    for gene, b_number in SYNTHETIC_B_NUMBERS.items():
        released = SYNTHETIC_ER[gene]
        values: list[Any] = [
            gene,
            b_number,
            "+",
            int(counts[gene]),
            0,
            "-",
            f"synthetic {gene}",
        ]
        for label, column in zip(
            m.S1_CONDITION_LABELS, m.S2_CONDITION_COLUMNS, strict=True
        ):
            values.append(released.get(label, released["other"]))
            assert column
        sheet.append(values)
    book.save(path)
    return path


def write_table_s6(path: Path, rows: Sequence[Mapping[str, Any]]) -> Path:
    """The released Table S6 layout, omitting the row that exercises the spacer rule."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Table S6. A synthetic double stranded sgRNA library"])
    sheet.append(list(m.S6_COLUMNS))
    index = 0
    for row in rows:
        if not row["in_s6"]:
            continue
        index += 1
        spacer = spacer_of(row["start"], row["end"])
        sequence = oligo(spacer)
        name = (
            f"F{index}_{row['start']}_{row['end']}_{row['fragment']}_"
            f"{row['fragment'].split('_')[1]}"
        )
        sheet.append([index, name, sequence, len(sequence)])
    book.save(path)
    return path


@pytest.fixture
def mg1655(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12MG1655Genome:
    """The real MG1655 class over the synthetic assembly; network refused."""
    files = write_assembly(
        tmp_path / "tier", MG1655_ASSEMBLY, CHOE_LOCI, gaf_rows=MG1655_GAF
    )
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    root = tmp_path / "mg1655"
    root.mkdir()
    return EcoliK12MG1655Genome(genome_root=str(root), overwrite=False)


@pytest.fixture
def tables(tmp_path: Path) -> dict[str, Path]:
    """The three synthetic workbooks, in the raw-file names the loader links."""
    rows = _rows()
    directory = tmp_path / "tables"
    directory.mkdir()
    return {
        m.TABLE_S1_FILE: write_table_s1(directory / m.TABLE_S1_FILE, rows),
        m.TABLE_S2_FILE: write_table_s2(directory / m.TABLE_S2_FILE, rows),
        m.TABLE_S6_FILE: write_table_s6(directory / m.TABLE_S6_FILE, rows),
    }


@pytest.fixture
def frames(tables: Mapping[str, Path]) -> dict[str, pd.DataFrame]:
    """The three workbooks parsed by the loader's own readers."""
    return {
        "s1": m.read_table_s1(tables[m.TABLE_S1_FILE]),
        "s2": m.read_table_s2(tables[m.TABLE_S2_FILE]),
        "s6": m.read_table_s6(tables[m.TABLE_S6_FILE]),
    }


@pytest.fixture
def synthetic_constants(monkeypatch: pytest.MonkeyPatch) -> None:
    """The counts and threshold of the synthetic tables, not the release's."""
    monkeypatch.setattr(m, "REPORTED_LIBRARY_COVERAGE", len(SYNTHETIC_ROWS) - 1)
    monkeypatch.setattr(m, "MIN_RESOLVED_FRACTION", 0.5)


@pytest.fixture
def presence_only_pins(monkeypatch: pytest.MonkeyPatch) -> list[Mapping[str, str]]:
    """Replace the build-time byte check with a presence check that records the pins."""
    calls: list[Mapping[str, str]] = []

    def record(raw_dir: str, pins: Mapping[str, str]) -> None:
        missing = [name for name in pins if not osp.exists(osp.join(raw_dir, name))]
        if missing:
            raise FileNotFoundError(f"pinned raw files absent: {missing}")
        calls.append(dict(pins))

    monkeypatch.setattr(m, "verify_raw_files", record)
    return calls


@pytest.fixture
def synthetic_root(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    tables: Mapping[str, Path],
    synthetic_constants: None,
    presence_only_pins: list[Mapping[str, str]],
) -> Path:
    """A dataset root whose ``raw/`` holds the three synthetic workbooks."""
    root = tmp_path / m.DATASET_ROOT_REL
    (root / "raw").mkdir(parents=True)
    for name, path in tables.items():
        (root / "raw" / name).write_bytes(path.read_bytes())
    monkeypatch.setattr(m, "assembly_reference", lambda strain, **_: REFERENCE)
    return root


# --------------------------------------------------------------------------- #
# Sequence helpers
# --------------------------------------------------------------------------- #
def test_reverse_complement_round_trips() -> None:
    assert m.reverse_complement("ACGTACGT") == "ACGTACGT"
    assert m.reverse_complement("AAAC") == "GTTT"
    assert m.reverse_complement(m.reverse_complement(SEQUENCE)) == SEQUENCE


def test_a_window_is_the_spacer_plus_its_pam() -> None:
    assert m.WINDOW_LENGTH == m.SPACER_LENGTH + 3
    assert m.SPACER_LENGTH == 20
    assert len(m.OLIGO_PREFIX) == 26
    assert len(m.OLIGO_SUFFIX) == 33


# --------------------------------------------------------------------------- #
# Reading the workbooks
# --------------------------------------------------------------------------- #
def test_read_table_s1_reads_identifiers_and_every_sample(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s1 = frames["s1"]
    assert len(s1) == len(SYNTHETIC_ROWS)
    assert list(s1.columns)[: len(m.S1_ID_COLUMNS)] == list(m.S1_ID_COLUMNS)
    assert s1["Start"].dtype.kind == "i"
    assert s1.loc[0, "Gene"] == "thrL"
    assert s1.loc[0, f"{m.CONTROL_LABEL}_1"] == 50.0
    assert s1.loc[0, f"{CCCP}_2"] == 100.0
    assert s1.loc[6, f"{NOVOBIOCIN}_1"] == 0.0
    # every library and sample column is present and numeric
    assert set(m.LIBRARY_COLUMNS) <= set(s1.columns)
    for label in m.S1_CONDITION_LABELS:
        assert f"{label}_1" in s1.columns and f"{label}_2" in s1.columns


def test_read_table_s1_refuses_a_changed_header(tmp_path: Path) -> None:
    rows = _rows()
    path = write_table_s1(tmp_path / "s1.xlsx", rows)
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    frame = pd.read_excel(path, sheet_name=0, header=None, engine="openpyxl")
    frame.iat[1, 0] = "Genes"
    for record in frame.itertuples(index=False):
        sheet.append(list(record))
    broken = tmp_path / "broken.xlsx"
    book.save(broken)
    with pytest.raises(m.TableLayoutError, match="Table S1 identifier"):
        m.read_table_s1(broken)


def test_read_table_s1_refuses_a_repeated_window(tmp_path: Path) -> None:
    rows = _rows()
    rows[1]["start"], rows[1]["end"] = rows[0]["start"], rows[0]["end"]
    path = write_table_s1(tmp_path / "repeat.xlsx", rows)
    with pytest.raises(m.TableLayoutError, match="repeats 1"):
        m.read_table_s1(path)


def test_read_table_s2_reads_the_b_number_map_and_the_ers(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s2 = frames["s2"]
    assert dict(zip(s2["Name"], s2["b number"], strict=True)) == SYNTHETIC_B_NUMBERS
    assert dict(zip(s2["Name"], s2["gRNA #"], strict=True)) == {
        "thrL": 3,
        "thrA": 4,
        "yaaP": 2,
        "zzzA": 1,
    }
    released = s2.set_index("Name")
    assert released.loc["thrA", m.CONTROL_LABEL] == 0.75
    assert released.loc["thrA", NOVOBIOCIN] == 0.25


def test_read_table_s2_refuses_a_repeated_symbol(tmp_path: Path) -> None:
    rows = _rows()
    path = write_table_s2(tmp_path / "s2.xlsx", rows)
    frame = pd.read_excel(path, sheet_name=0, header=None, engine="openpyxl")
    frame.iat[4, 0] = "thrL"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for record in frame.itertuples(index=False):
        sheet.append(list(record))
    broken = tmp_path / "broken_s2.xlsx"
    book.save(broken)
    with pytest.raises(m.TableLayoutError, match="repeats gene symbols"):
        m.read_table_s2(broken)


def test_read_table_s6_slices_the_spacer_and_parses_the_name(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s6 = frames["s6"]
    assert len(s6) == sum(1 for row in SYNTHETIC_ROWS if row[6])
    first = s6.iloc[0]
    assert first["start"] == 15 and first["end"] == 37
    assert first["spacer"] == spacer_of(15, 37)
    assert first["orientation"] == "F"
    assert first["fragment"] == "1_1"
    assert (s6["spacer"].str.len() == m.SPACER_LENGTH).all()


def test_read_table_s6_refuses_a_bad_oligo_a_bad_name_and_a_repeat(
    tmp_path: Path,
) -> None:
    def book_of(rows: list[list[Any]]) -> Path:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append(["Table S6."])
        sheet.append(list(m.S6_COLUMNS))
        for row in rows:
            sheet.append(row)
        path = tmp_path / f"s6_{len(list(tmp_path.iterdir()))}.xlsx"
        book.save(path)
        return path

    spacer = spacer_of(15, 37)
    good = oligo(spacer)
    with pytest.raises(m.TableLayoutError, match="do not carry the stated"):
        m.read_table_s6(book_of([[1, "F1_15_37_1_1_1", good[:-1], 78]]))
    with pytest.raises(m.TableLayoutError, match="do not parse"):
        m.read_table_s6(book_of([[1, "X1_15_37_1_1_1", good, 79]]))
    with pytest.raises(m.TableLayoutError, match="repeats 1"):
        m.read_table_s6(
            book_of([[1, "F1_15_37_1_1_1", good, 79], [2, "F2_15_37_1_1_1", good, 79]])
        )


def test_read_table_s6_refuses_a_changed_header(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Table S6."])
    sheet.append(["#", "Primer", "Sequence", "Length"])
    path = tmp_path / "s6_header.xlsx"
    book.save(path)
    with pytest.raises(m.TableLayoutError, match="Table S6 header"):
        m.read_table_s6(path)


def test_spacer_by_window_keys_on_the_released_coordinates(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    mapping = m.spacer_by_window(frames["s6"])
    assert mapping[(15, 37)] == spacer_of(15, 37)
    assert (5, 27) not in mapping


# --------------------------------------------------------------------------- #
# The three fidelity checks
# --------------------------------------------------------------------------- #
def _stub_genome(sequence: str = SEQUENCE) -> Any:
    """The only genome surface :func:`m.check_spacers` reads."""
    return SimpleNamespace(fasta_dna={"synthetic": SimpleNamespace(seq=sequence)})


def test_check_spacers_counts_both_orientations(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s1, s6 = frames["s1"], frames["s6"]
    mapping = m.spacer_by_window(s6)
    joined = s1.assign(
        spacer=pd.Series(
            [
                mapping.get((int(start), int(end)))
                for start, end in zip(s1["Start"], s1["End"], strict=True)
            ],
            index=s1.index,
            dtype="object",
        )
    )
    check = m.check_spacers(joined, _stub_genome())
    assert check.n_joined == len(s6)
    assert check.n_neither == 0
    assert check.n_both == 0
    assert check.n_reverse_read + check.n_forward_read == check.n_joined
    assert check.n_reverse_read > 0 and check.n_forward_read > 0


def test_check_spacers_refuses_a_spacer_that_is_not_in_its_window(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s1 = frames["s1"]
    joined = s1.assign(spacer=["A" * m.SPACER_LENGTH] * len(s1))
    with pytest.raises(m.TableLayoutError, match="are not the 20 nt"):
        m.check_spacers(joined, _stub_genome())


def test_check_spacers_refuses_a_window_that_is_not_23_nt(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s1 = frames["s1"].iloc[:1].copy()
    s1["Start"] = [len(SEQUENCE) - 5]
    s1["End"] = [len(SEQUENCE) + 17]
    joined = s1.assign(spacer=["A" * m.SPACER_LENGTH])
    with pytest.raises(m.TableLayoutError, match="window is"):
        m.check_spacers(joined, _stub_genome())


def test_check_enrichment_ratios_reproduces_the_released_values(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    check = m.check_enrichment_ratios(frames["s1"], frames["s2"])
    assert check.excluded_genes == ("yaaP",)
    assert check.n_genes == 3
    assert check.n_conditions == len(m.S1_CONDITION_LABELS)
    assert check.n_cells == 3 * len(m.S1_CONDITION_LABELS)
    assert check.max_abs_difference == pytest.approx(0.0, abs=1e-12)


def test_check_enrichment_ratios_refuses_a_disagreement(
    frames: Mapping[str, pd.DataFrame],
) -> None:
    s2 = frames["s2"].copy()
    s2.loc[s2["Name"] == "thrA", NOVOBIOCIN] = 0.75
    with pytest.raises(m.TableLayoutError, match="differ from Table S2"):
        m.check_enrichment_ratios(frames["s1"], s2)


def test_check_library_coverage_reproduces_the_reported_count(
    frames: Mapping[str, pd.DataFrame], synthetic_constants: None
) -> None:
    check = m.check_library_coverage(frames["s1"])
    assert check.n_guide_rows == len(SYNTHETIC_ROWS)
    assert check.n_absent_from_library == 1
    assert check.n_present_in_library == check.reported == len(SYNTHETIC_ROWS) - 1


def test_check_library_coverage_refuses_a_mismatch(
    frames: Mapping[str, pd.DataFrame], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(m, "REPORTED_LIBRARY_COVERAGE", 12345)
    with pytest.raises(m.TableLayoutError, match="the paper reports 12345"):
        m.check_library_coverage(frames["s1"])


# --------------------------------------------------------------------------- #
# Identifiers
# --------------------------------------------------------------------------- #
def test_canonical_symbol_prefers_the_annotation_and_falls_back_to_the_tag(
    mg1655: EcoliK12MG1655Genome,
) -> None:
    assert m.canonical_symbol(mg1655, "b0001") == "thrL"
    assert m.canonical_symbol(mg1655, "b0003") == "yaaP"


def test_canonical_symbol_uses_the_tag_when_the_locus_has_no_symbol(
    mg1655: EcoliK12MG1655Genome, monkeypatch: pytest.MonkeyPatch
) -> None:
    locus = mg1655.genbank.loci["b0001"]
    monkeypatch.setitem(
        mg1655.genbank.loci, "b0001", locus.model_copy(update={"symbol": None})
    )
    assert m.canonical_symbol(mg1655, "b0001") == "b0001"


def test_canonical_symbol_uses_the_tag_when_the_symbol_resolves_elsewhere(
    mg1655: EcoliK12MG1655Genome, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        mg1655,
        "resolve_gene_name",
        lambda name: SimpleNamespace(systematic_name="b0002", status=None),
    )
    assert m.canonical_symbol(mg1655, "b0001") == "b0001"


def test_resolve_genes_places_each_gene_and_lists_the_unresolvable(
    frames: Mapping[str, pd.DataFrame],
    mg1655: EcoliK12MG1655Genome,
    synthetic_constants: None,
) -> None:
    identities, unresolved, report = m.resolve_genes(
        frames["s2"], mg1655, label="synthetic"
    )
    assert unresolved == ("zzzA",)
    assert set(identities) == {"thrL", "thrA", "yaaP"}
    thra = identities["thrA"]
    assert thra.locus_tag == "b0002"
    assert thra.symbol == "thrA"
    assert (thra.start, thra.end, thra.strand) == (41, 85, "-")
    assert thra.released_guide_count == 4
    assert report.unique_names == 4
    assert report.resolved == 3


def test_resolve_genes_stops_below_the_resolved_threshold(
    frames: Mapping[str, pd.DataFrame], mg1655: EcoliK12MG1655Genome
) -> None:
    with pytest.raises(LocusTagResolutionError, match="below 0.99"):
        m.resolve_genes(frames["s2"], mg1655, label="synthetic")


# --------------------------------------------------------------------------- #
# Retention
# --------------------------------------------------------------------------- #
@pytest.fixture
def retention(
    frames: Mapping[str, pd.DataFrame],
    mg1655: EcoliK12MG1655Genome,
    synthetic_constants: None,
) -> m.RetentionResult:
    identities, unresolved, _ = m.resolve_genes(frames["s2"], mg1655, label="synthetic")
    return m.select_guides(
        frames["s1"], m.spacer_by_window(frames["s6"]), identities, unresolved
    )


def test_select_guides_applies_every_rule_exactly_once(
    retention: m.RetentionResult,
) -> None:
    assert retention.unresolved_b_number == ["zzzA:25-47"]
    assert retention.outside_locus == ["yaaP:32-54"]
    assert retention.strand_disagreement == ["thrA:55-77"]
    assert retention.missing_spacer == ["thrL:5-27"]
    assert retention.zero_control == 1
    assert retention.zero_antibiotic_cells == 1
    assert [guide.locus_tag for guide in retention.kept] == [
        "b0001",
        "b0001",
        "b0002",
        "b0002",
        "b0003",
    ]


def test_select_guides_carries_the_spacer_and_the_gene_guide_count(
    retention: m.RetentionResult,
) -> None:
    first = retention.kept[0]
    assert first.spacer == spacer_of(15, 37)
    assert first.n_guides == 3  # thrL has three Table S1 rows
    assert first.control_abundance == 50.0
    assert set(first.antibiotic_abundance) == {c.s1_label for c in m.CONDITIONS}
    zero_novobiocin = retention.kept[3]
    assert zero_novobiocin.locus_tag == "b0002"
    assert NOVOBIOCIN not in zero_novobiocin.antibiotic_abundance


def test_select_guides_reports_the_released_guide_count_disagreements(
    retention: m.RetentionResult,
) -> None:
    assert retention.guide_counts == {}


def test_drop_log_accounts_for_every_dropped_record(
    retention: m.RetentionResult,
) -> None:
    ledger = m.drop_log("synthetic", len(SYNTHETIC_ROWS), retention)
    assert ledger.antibiotics == 12
    assert ledger.source_records == len(SYNTHETIC_ROWS) * 12
    assert ledger.kept_guides == 5
    assert ledger.kept_records == 5 * 12 - 1
    assert ledger.dropped_records == ledger.source_records - ledger.kept_records
    assert sum(rule.n_records for rule in ledger.rules) == ledger.dropped_records
    by_rule = {rule.rule: rule for rule in ledger.rules}
    assert by_rule["abundance_is_zero_in_the_antibiotic"].scope == "cell"
    assert by_rule["abundance_is_zero_in_the_antibiotic"].n_records == 1
    assert by_rule["guide_window_does_not_overlap_the_resolved_locus"].items == [
        "yaaP:32-54"
    ]


def test_drop_log_refuses_an_unaccounted_record(retention: m.RetentionResult) -> None:
    broken = retention.model_copy(update={"zero_control": 0})
    with pytest.raises(RuntimeError, match="drop accounting mismatch"):
        m.drop_log("synthetic", len(SYNTHETIC_ROWS), broken)


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
def test_log2_response_is_the_treatment_over_control_ratio(
    retention: m.RetentionResult,
) -> None:
    guide = retention.kept[0]
    assert m.log2_response(guide, CCCP) == pytest.approx(1.0)
    assert m.log2_response(guide, NOVOBIOCIN) == pytest.approx(0.0)
    assert m.log2_response(guide, CCCP) == pytest.approx(
        math.log2(guide.antibiotic_abundance[CCCP] / guide.control_abundance)
    )


def test_response_phenotype_types_the_readout_and_gaps_what_is_unsourced() -> None:
    phenotype = m.response_phenotype(-1.5)
    assert phenotype.measurement_type is MeasurementType.log2_ratio
    assert phenotype.assay_type is AssayType.pooled_competitive_growth_barcode
    assert phenotype.environment_response == -1.5
    assert phenotype.n_samples == 2
    assert phenotype.sample_unit is None
    assert phenotype.environment_response_uncertainty is None
    assert phenotype.environment_response_se is None
    assert phenotype.gapped_fields() == {
        "sample_unit",
        "environment_response_uncertainty",
    }
    assert "untreated" in str(phenotype.units)


def test_the_screening_medium_derives_from_lb_with_two_selection_agents() -> None:
    medium = m.CHOE2025_LB_SELECTION
    assert medium.base_medium == "LB"
    assert medium.state == "liquid"
    assert medium.is_synthetic is False
    assert len(medium.components) == len(LB.components) + 2
    selection = [
        component
        for component in medium.components
        if component.role is MediaComponentRole.selection_agent
    ]
    doses = {
        component.compound.name: component.concentration for component in selection
    }
    assert set(doses) == {"chloramphenicol", "ampicillin"}
    for name, dose in doses.items():
        assert dose is not None
        assert dose.unit is ConcentrationUnit.ug_per_ml
        assert dose.value == (35.0 if name == "chloramphenicol" else 100.0)


def test_environment_doses_the_antibiotic_and_types_the_vehicle() -> None:
    cccp = m.environment(m.CONDITIONS_BY_LABEL[CCCP])
    assert cccp.media is m.CHOE2025_LB_SELECTION
    assert cccp.temperature is not None and cccp.temperature.value == 37.0
    assert cccp.aerobicity == "aerobic"
    assert cccp.duration_hours is None and cccp.duration_generations is None
    assert cccp.gapped_fields() == {"duration_hours", "duration_generations"}
    (perturbation,) = cccp.perturbations
    assert isinstance(perturbation, SmallMoleculePerturbation)
    assert perturbation.compound.name == "CCCP"
    assert perturbation.concentration.value == 5.3
    assert perturbation.concentration.unit is ConcentrationUnit.ug_per_ml
    assert perturbation.concentration.basis is DoseBasis.fixed
    assert perturbation.solvent is not None
    assert perturbation.solvent.name == "DMSO"
    assert perturbation.solvent.percent == 0.4
    assert perturbation.solvent.compound is not None

    verapamil = m.environment(m.CONDITIONS_BY_LABEL["Verapamil"])
    (dose,) = verapamil.perturbations
    assert isinstance(dose, SmallMoleculePerturbation)
    assert dose.solvent is None
    assert dose.concentration.unit is ConcentrationUnit.millimolar
    assert dose.concentration.value == 3.6


def test_control_environment_is_the_same_medium_with_no_drug() -> None:
    control = m.control_environment()
    assert control.perturbations == []
    assert control.media is m.CHOE2025_LB_SELECTION
    assert control.temperature is not None and control.temperature.value == 37.0


def test_knockdown_genotype_is_one_crispri_perturbation_with_its_spacer(
    retention: m.RetentionResult,
) -> None:
    genotype = m.knockdown_genotype(retention.kept[0])
    (perturbation,) = genotype.perturbations
    assert perturbation.perturbation_type == "bacterial_crispr_interference"
    assert perturbation.systematic_gene_name == "b0001"
    assert perturbation.perturbed_gene_name == "thrL"
    assert perturbation.gene_namespace == "ecoli_k12_mg1655_bnumber"
    assert perturbation.expression_direction == "decreased"
    assert perturbation.state == "present"
    assert perturbation.crispr.effector == "dCas9"
    assert perturbation.crispr.guide_sequence == spacer_of(15, 37)
    assert perturbation.crispr.n_guides == 3


def test_two_guides_of_one_gene_are_distinct_genotypes(
    retention: m.RetentionResult,
) -> None:
    first, second = retention.kept[0], retention.kept[1]
    assert first.locus_tag == second.locus_tag
    assert m.knockdown_genotype(first) != m.knockdown_genotype(second)


def test_build_experiment_and_reference_round_trip(
    retention: m.RetentionResult,
) -> None:
    guide = retention.kept[0]
    condition = m.CONDITIONS_BY_LABEL[CCCP]
    experiment = m.build_experiment(
        "choe", guide, condition, m.knockdown_genotype(guide), m.environment(condition)
    )
    assert isinstance(experiment, BacterialEnvironmentResponseExperiment)
    assert experiment.experiment_type == "bacterial_environment_response"
    assert experiment.phenotype.environment_response == pytest.approx(1.0)
    restored = BacterialEnvironmentResponseExperiment.model_validate(
        experiment.model_dump()
    )
    assert restored == experiment

    reference = m.build_reference("choe", REFERENCE)
    assert isinstance(reference, BacterialEnvironmentResponseExperimentReference)
    assert reference.phenotype_reference.environment_response == 0.0
    assert reference.environment_reference.perturbations == []
    assert reference.genome_reference == REFERENCE


# --------------------------------------------------------------------------- #
# The condition table is the paper's
# --------------------------------------------------------------------------- #
def test_the_twelve_conditions_match_table_s3() -> None:
    assert len(m.CONDITIONS) == int(m.SOURCED_VALUES["twelve_antibiotics"].value)
    assert len(m.S1_CONDITION_LABELS) == 13
    assert m.S1_CONDITION_LABELS[0] == m.CONTROL_LABEL
    assert sum(condition.dmso for condition in m.CONDITIONS) == 6
    units = {condition.unit for condition in m.CONDITIONS}
    assert units == {
        ConcentrationUnit.ug_per_ml,
        ConcentrationUnit.millimolar,
        ConcentrationUnit.percent_w_v,
    }
    sulfamethizole = m.CONDITIONS_BY_LABEL["Sulfamethizole"]
    assert sulfamethizole.dose == 200.0
    assert sulfamethizole.unit is ConcentrationUnit.ug_per_ml
    assert "0.2 mg/ml converted to 200 ug/mL" in str(sulfamethizole.dose_sourced.note)
    assert "<td>0.2</td><td>mg/ml</td>" in sulfamethizole.table_s3_row


def test_every_condition_quote_states_its_own_dose_and_sampling_od() -> None:
    for condition in m.CONDITIONS:
        row = condition.table_s3_row
        assert row.startswith("<tr><td>")
        assert f"<td>{condition.sampling_od600:.2f}</td>" in row
        assert condition.mode_of_action in row
        assert ("0.4% DMSO/LB" in row) is condition.dmso
        assert condition.dose_sourced.provenance.source_uri == m.SI1_MD


def test_the_release_names_are_distinct_in_every_table() -> None:
    assert len({c.s1_label for c in m.CONDITIONS}) == 12
    assert len({c.s2_column for c in m.CONDITIONS}) == 12
    assert len({c.compound_label for c in m.CONDITIONS}) == 12
    # Table S1 spells pyocyanin "Pyocyanine" and Table S2 abbreviates mitomycin C
    assert m.CONDITIONS_BY_LABEL["Pyocyanine"].s2_column == "Pyocyanin"
    assert m.CONDITIONS_BY_LABEL["Mitomycin C"].s2_column == "Mitomycin"


# --------------------------------------------------------------------------- #
# Raw-mirror plumbing
# --------------------------------------------------------------------------- #
def test_raw_files_carry_a_pmc_cloud_retrieval() -> None:
    assert [raw.name for raw in m.RAW_FILES] == [
        m.TABLE_S1_FILE,
        m.TABLE_S2_FILE,
        m.TABLE_S6_FILE,
    ]
    for raw in m.RAW_FILES:
        assert raw.mirror_relpath == f"data/{raw.name}"
        assert raw.bucket_key == f"{m.PMC_PREFIX}/{raw.name}"
        assert raw.source_url.endswith(raw.bucket_key)
        record = raw.retrieval
        assert record.method is RetrievalMethod.pmc_cloud
        assert record.params == {"key": raw.bucket_key}
        assert record.sha256 == raw.sha256
        assert len(raw.sha256) == 64
    assert m.DATA_SHA256 == {raw.name: raw.sha256 for raw in m.RAW_FILES}


def test_retrieve_raw_files_runs_the_recorded_retriever(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    payloads = {raw.name: raw.name.encode() for raw in m.RAW_FILES}
    monkeypatch.setattr(
        m, "run_retriever", lambda record: payloads[record.params["key"].split("/")[-1]]
    )
    written: dict[str, bytes] = {}

    def record_write(data: bytes, dest: Any, expected: str, source: str) -> None:
        assert expected and source
        written[Path(dest).name] = data
        Path(dest).write_bytes(data)

    monkeypatch.setattr(m, "write_verified", record_write)
    out = m.retrieve_raw_files(tmp_path / "dl")
    assert set(out) == set(payloads)
    assert written == payloads


def test_deposit_raw_mirror_refuses_a_byte_mismatch(tmp_path: Path) -> None:
    sources = {}
    for raw in m.RAW_FILES:
        path = tmp_path / raw.name
        path.write_bytes(b"not the pinned bytes")
        sources[raw.name] = path
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        m.deposit_raw_mirror(sources=sources, data_root=str(tmp_path / "root"))


def test_deposit_raw_mirror_requires_every_file(tmp_path: Path) -> None:
    with pytest.raises(KeyError, match="no source given"):
        m.deposit_raw_mirror(sources={}, data_root=str(tmp_path))


def test_manifest_sha256_refuses_an_unknown_path() -> None:
    manifest = Manifest(
        citation_key=m.CITATION_KEY,
        doi=m.PAPER_DOI,
        title=m.PAPER_TITLE,
        files=[],
        created_at="2026-10-07T00:00:00+00:00",
    )
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        m.manifest_sha256(manifest, "data/absent.xlsx")


def test_mirror_paths_sit_under_the_given_data_root(tmp_path: Path) -> None:
    assert m.raw_mirror_dir(str(tmp_path)) == tmp_path / m.RAW_DIR_REL
    assert m.library_dir(str(tmp_path)) == tmp_path / m.LIBRARY_DIR_REL


# --------------------------------------------------------------------------- #
# End to end over the synthetic workbooks
# --------------------------------------------------------------------------- #
@pytest.fixture
def built(
    synthetic_root: Path, mg1655: EcoliK12MG1655Genome
) -> m.CrispriChemgenChoe2025Dataset:
    return m.CrispriChemgenChoe2025Dataset(
        root=str(synthetic_root), ecoli_genome=mg1655
    )


def test_process_writes_one_record_per_kept_guide_and_antibiotic(
    built: m.CrispriChemgenChoe2025Dataset,
) -> None:
    assert len(built) == 5 * 12 - 1
    assert built.experiment_class is BacterialEnvironmentResponseExperiment
    assert built.reference_class is BacterialEnvironmentResponseExperimentReference
    assert built.raw_file_names == [m.TABLE_S1_FILE, m.TABLE_S2_FILE, m.TABLE_S6_FILE]
    assert set(built.gene_set) == {"b0001", "b0002", "b0003"}


def test_process_records_carry_the_expected_response(
    built: m.CrispriChemgenChoe2025Dataset,
) -> None:
    responses: dict[str, list[float]] = {}
    for index in range(len(built)):
        record = built[index]
        experiment = record["experiment"]
        (perturbation,) = experiment["genotype"]["perturbations"]
        compound = experiment["environment"]["perturbations"][0]["compound"]["name"]
        responses.setdefault(compound, []).append(
            experiment["phenotype"]["environment_response"]
        )
        assert record["reference"]["phenotype_reference"]["environment_response"] == 0.0
        assert perturbation["crispr"]["effector"] == "dCas9"
    assert len(responses["CCCP"]) == 5
    assert all(value == pytest.approx(1.0) for value in responses["CCCP"])
    assert all(value == pytest.approx(0.0) for value in responses["novobiocin"])
    assert len(responses["novobiocin"]) == 4  # the zero-novobiocin cell is dropped


def test_process_writes_the_three_ledgers(
    built: m.CrispriChemgenChoe2025Dataset,
) -> None:
    out = Path(built.preprocess_dir)
    ledger = m.DropLog.model_validate_json((out / "dropped_records.json").read_text())
    assert ledger.kept_records == len(built)
    assert ledger.guide_rows == len(SYNTHETIC_ROWS)
    identifiers = json.loads((out / "identifier_reconciliation.json").read_text())
    assert identifiers["released_genes"] == 4
    assert identifiers["resolved_genes"] == 3
    assert identifiers["kept_genes"] == 3
    assert identifiers["released_symbol_differs_from_annotation"] == {}
    extraction = json.loads((out / "extraction.json").read_text())
    assert extraction["raw_sha256"] == m.DATA_SHA256
    assert extraction["library_coverage"]["n_absent_from_library"] == 1
    assert extraction["enrichment_ratio"]["excluded_genes"] == ["yaaP"]
    assert extraction["spacer_window"]["n_neither"] == 0


def test_process_refuses_disagreeing_gene_sets(
    synthetic_root: Path, mg1655: EcoliK12MG1655Genome, tmp_path: Path
) -> None:
    rows = _rows()
    rows[0]["gene"] = "notAGene"
    write_table_s1(synthetic_root / "raw" / m.TABLE_S1_FILE, rows)
    with pytest.raises(m.TableLayoutError, match="are not Table S2's"):
        m.CrispriChemgenChoe2025Dataset(root=str(synthetic_root), ecoli_genome=mg1655)


def test_the_dataset_refuses_a_genome_of_another_assembly(
    synthetic_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    wrong = SimpleNamespace(ASSEMBLY_SET="ecoli_K12_BW25113_ASM75055v1")
    dataset = m.CrispriChemgenChoe2025Dataset.__new__(m.CrispriChemgenChoe2025Dataset)
    dataset.ecoli_genome = wrong  # type: ignore[assignment]
    with pytest.raises(ValueError, match="needs the ecoli_K12_MG1655"):
        dataset._genome()


def test_create_experiment_is_not_the_entry_point(
    built: m.CrispriChemgenChoe2025Dataset,
) -> None:
    with pytest.raises(NotImplementedError):
        built.create_experiment()
    frame = pd.DataFrame({"a": [1]})
    assert built.preprocess_raw(frame) is frame


def test_the_reference_strain_is_the_screen_host() -> None:
    assert m.CrispriChemgenChoe2025Dataset.REFERENCE_STRAIN == "MG1655"
    assert m.SOURCED_VALUES["host_strain"].value == "MG1655"
    assert m.GENE_NAMESPACE == "ecoli_k12_mg1655_bnumber"
    assert m.EXPECTED_RECORDS == 466569


def test_main_dispatches_each_subcommand(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    monkeypatch.setattr(m, "load_dotenv", lambda *a, **k: None, raising=False)
    retrieved: list[Path] = []
    monkeypatch.setattr(
        m, "retrieve_raw_files", lambda dest: retrieved.append(Path(dest))
    )
    monkeypatch.setattr(
        m,
        "deposit_raw_mirror",
        lambda *, sources, data_root: Path(data_root) / "mirror",
    )
    assert (
        m.main(["deposit", "--download-dir", str(tmp_path / "dl"), "--retrieve"]) == 0
    )
    assert retrieved == [tmp_path / "dl"]
    assert "mirror" in capsys.readouterr().out

    monkeypatch.setattr(
        m, "CrispriChemgenChoe2025Dataset", lambda root: ["a", "b"], raising=True
    )
    assert m.main(["build"]) == 0
    assert "len = 2" in capsys.readouterr().out

    monkeypatch.setattr(
        m,
        "run_verification",
        lambda data_root: SimpleNamespace(summary=lambda: "SUMMARY", passed=True),
    )
    assert m.main(["verify"]) == 0
    assert "SUMMARY" in capsys.readouterr().out


# --------------------------------------------------------------------------- #
# Data-gated: the real mirror and the released tables
# --------------------------------------------------------------------------- #
def _data_root() -> str | None:
    return os.environ.get("DATA_ROOT")


def _mirror_data() -> Path | None:
    root = _data_root()
    if root is None:
        return None
    directory = m.raw_mirror_dir(root) / "data"
    return directory if directory.is_dir() else None


@pytest.mark.data
def test_real_mirror_matches_every_pin() -> None:
    directory = _mirror_data()
    if directory is None:
        pytest.skip("the Choe 2025 raw mirror is not on this host")
    root = _data_root()
    assert root is not None
    manifest = m.load_manifest(root)
    assert manifest.citation_key == m.CITATION_KEY
    for raw in m.RAW_FILES:
        assert m.manifest_sha256(manifest, raw.mirror_relpath) == raw.sha256
        path = directory / raw.name
        assert path.stat().st_size == raw.bytes
        assert m._sha256(path) == raw.sha256


@pytest.mark.data
def test_every_sourced_value_is_backed_by_its_quote() -> None:
    root = _data_root()
    if root is None or not (Path(root) / "torchcell-library" / m.CITATION_KEY).is_dir():
        pytest.skip("the Choe 2025 literature mirror is not on this host")
    from torchcell.verification.sourced import audit_sourced_value

    library = Path(root) / "torchcell-library"
    for name, value in m.SOURCED_VALUES.items():
        result = audit_sourced_value(value, library)
        assert result.passed, f"{name}: {result.message}"
    for condition in m.CONDITIONS:
        result = audit_sourced_value(condition.dose_sourced, library)
        assert result.passed, f"{condition.s1_label}: {result.message}"
    assert audit_sourced_value(m.CONTROL_ROW, library).passed


@pytest.mark.data
def test_the_released_tables_hold_the_counts_the_paper_reports() -> None:
    directory = _mirror_data()
    if directory is None:
        pytest.skip("the Choe 2025 raw mirror is not on this host")
    s1 = m.read_table_s1(directory / m.TABLE_S1_FILE)
    s2 = m.read_table_s2(directory / m.TABLE_S2_FILE)
    s6 = m.read_table_s6(directory / m.TABLE_S6_FILE)
    n_genes, n_oligos = m.SOURCED_VALUES["library_design"].value
    assert len(s2) == n_genes
    assert len(s6) == n_oligos
    assert s1["Gene"].nunique() == n_genes
    assert set(s1["Gene"]) == set(s2["Name"])
    coverage = m.check_library_coverage(s1)
    assert coverage.n_present_in_library == m.REPORTED_LIBRARY_COVERAGE


@pytest.mark.data
def test_the_built_store_holds_the_expected_records() -> None:
    root = _data_root()
    if root is None:
        pytest.skip("DATA_ROOT is not set on this host")
    preprocess = Path(root) / m.DATASET_ROOT_REL / "preprocess"
    if not (preprocess / "dropped_records.json").is_file():
        pytest.skip("the Choe 2025 dev store is not built on this host")
    ledger = m.DropLog.model_validate_json(
        (preprocess / "dropped_records.json").read_text()
    )
    assert ledger.kept_records == m.EXPECTED_RECORDS
    assert ledger.guide_rows == 39580
    assert ledger.kept_guides == 38966
    gene_set = json.loads((preprocess / "gene_set.json").read_text())
    assert len(gene_set) == 4156
    extraction = json.loads((preprocess / "extraction.json").read_text())
    assert extraction["spacer_window"]["n_neither"] == 0
    assert extraction["spacer_window"]["n_both"] == 0
    assert extraction["enrichment_ratio"]["max_abs_difference"] < m.ER_TOLERANCE


# --------------------------------------------------------------------------- #
# Depositing and linking one stand-in raw file
# --------------------------------------------------------------------------- #
def _synthetic_raw_file(tmp_path: Path) -> tuple[m.RawFile, Path]:
    """A one-file ``RAW_FILES`` stand-in and the bytes it pins."""
    path = tmp_path / "mmcX.xlsx"
    payload = b"synthetic workbook bytes"
    path.write_bytes(payload)
    raw = m.RawFile(
        name=path.name,
        sha256=hashlib.sha256(payload).hexdigest(),
        bytes=len(payload),
        description="synthetic",
    )
    return raw, path


@pytest.fixture
def one_raw_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[m.RawFile, Path]:
    raw, path = _synthetic_raw_file(tmp_path)
    monkeypatch.setattr(m, "RAW_FILES", (raw,))
    monkeypatch.setattr(m, "DATA_SHA256", {raw.name: raw.sha256})
    monkeypatch.setattr(m, "RAW_FILES_BY_NAME", {raw.name: raw})
    return raw, path


def test_deposit_is_idempotent_and_refuses_a_differing_file(
    tmp_path: Path, one_raw_file: tuple[m.RawFile, Path]
) -> None:
    raw, path = one_raw_file
    data_root = str(tmp_path / "root")
    root = m.deposit_raw_mirror(sources={raw.name: path}, data_root=data_root)
    assert root == Path(data_root) / m.RAW_DIR_REL
    m.deposit_raw_mirror(sources={raw.name: path}, data_root=data_root)
    manifest = m.load_manifest(data_root)
    (record,) = manifest.files
    assert (record.path, record.sha256, record.role) == (
        f"data/{raw.name}",
        raw.sha256,
        "raw_data",
    )
    assert record.retrieval is not None
    assert record.retrieval.method is RetrievalMethod.pmc_cloud
    assert manifest.si_expected == list(m.NOT_MIRRORED)
    assert manifest.doi == m.PAPER_DOI

    (root / "data" / raw.name).write_bytes(b"edited")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        m.deposit_raw_mirror(sources={raw.name: path}, data_root=data_root)


def test_download_links_the_mirror_and_checks_the_manifest(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    one_raw_file: tuple[m.RawFile, Path],
) -> None:
    from torchcell.data import ManifestPinMismatchError

    raw, path = one_raw_file
    data_root = tmp_path / "root"
    m.deposit_raw_mirror(sources={raw.name: path}, data_root=str(data_root))
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    dataset = object.__new__(m.CrispriChemgenChoe2025Dataset)
    monkeypatch.setattr(
        m.CrispriChemgenChoe2025Dataset, "raw_dir", str(tmp_path / "raw"), raising=False
    )
    dataset.download()
    linked = tmp_path / "raw" / raw.name
    assert linked.is_symlink()
    assert os.readlink(linked) == str(data_root / m.RAW_DIR_REL / "data" / raw.name)

    manifest_path = data_root / m.RAW_DIR_REL / "manifest.json"
    manifest_path.write_text(manifest_path.read_text().replace(raw.sha256, "0" * 64))
    with pytest.raises(ManifestPinMismatchError, match=f"data/{raw.name}"):
        dataset.download()
    manifest_path.write_text(manifest_path.read_text().replace("0" * 64, raw.sha256))
    (data_root / m.RAW_DIR_REL / "data" / raw.name).unlink()
    linked.unlink()
    with pytest.raises(RuntimeError, match="required raw artifact missing"):
        dataset.download()


# --------------------------------------------------------------------------- #
# The supplementary locus row of the verifier
# --------------------------------------------------------------------------- #
def _tag_record(tag: str) -> dict[str, Any]:
    return {
        "experiment": {"genotype": {"perturbations": [{"systematic_gene_name": tag}]}}
    }


def test_stored_tags_are_loci_accepts_a_gene_and_a_pseudogene(
    mg1655: EcoliK12MG1655Genome,
) -> None:
    result = m.stored_tags_are_loci(
        [_tag_record("b0001"), _tag_record("b0003")], mg1655
    )
    assert result.passed
    assert result.details["statuses"]["current"] == 1
    assert result.details["statuses"]["non_gene_feature"] == 1
    assert result.details["not_a_locus"] == []


def test_stored_tags_are_loci_flags_a_tag_that_is_not_one(
    mg1655: EcoliK12MG1655Genome,
) -> None:
    result = m.stored_tags_are_loci([_tag_record("b9999")], mg1655)
    assert not result.passed
    assert result.details["not_a_locus"] == ["b9999"]
