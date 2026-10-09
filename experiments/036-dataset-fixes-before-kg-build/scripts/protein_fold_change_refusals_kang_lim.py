# experiments/036-dataset-fixes-before-kg-build/scripts/protein_fold_change_refusals_kang_lim.py
# [[experiments.036-dataset-fixes-before-kg-build.scripts.protein_fold_change_refusals_kang_lim]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/036-dataset-fixes-before-kg-build/scripts/protein_fold_change_refusals_kang_lim
"""Measure why Kang 2026 and Lim 2025 stay refused after ``ProteinFoldChangePhenotype``.

Issue #770 names five papers that release a protein-level fold change the schema could
not hold. The class landed; two of the five are still not loadable, and for reasons that
have nothing to do with the phenotype. This script measures both, so the refusals in
``notes/torchcell.datasets.pputida.kang2026.md`` and
``notes/torchcell.datasets.pputida.lim2025.md`` rest on bytes rather than on a reading.

1. **Kang 2026 Table S6.** Its caption, its ten columns, its row count, its key form and
   how many rows clear ``p < 0.05``, read through the loader's own ``si_table`` so the
   rendering is the one the loader would get. Then the decisive number: how many
   ``/db_xref="UniProtKB/..."`` entries the PINNED *P. putida* KT2440 assembly carries,
   beside the same count on the pinned MG1655 assembly that Gupta 2024 resolves 3,225 of
   3,262 proteins through. The new ``uniprot_db_xref`` identifier route reads exactly
   that cross-reference, so a zero here is what refuses the table.
2. **Lim 2025 ``si/si2.xlsx``.** Every proteome sheet's row count, its two
   ``log2_mean_*`` arms and its ``log2_Fold_change_A/B`` column. An evolved isolate on
   one side of every released contrast is what refuses the sheet, because an evolved
   clone's genotype cannot be written with the existing classes (#731).

Run from the repo root:
    python experiments/036-dataset-fixes-before-kg-build/scripts/protein_fold_change_refusals_kang_lim.py
"""

from __future__ import annotations

import gzip
import hashlib
import os
import os.path as osp

import openpyxl
from dotenv import load_dotenv

from torchcell.datasets.pputida.kang2026 import si_table

load_dotenv()
DATA_ROOT = os.environ["DATA_ROOT"]

KANG_SI = osp.join(
    DATA_ROOT,
    "torchcell-library",
    "kangMultilayeredMetabolicRemodeling2026",
    "si",
    "si1.docx",
)
LIM_SI = osp.join(
    DATA_ROOT,
    "torchcell-library",
    "limEvolutionguidedToleranceEngineering2025",
    "si",
    "si2.xlsx",
)
PPUTIDA_GBFF = osp.join(
    DATA_ROOT,
    "torchcell-genomes",
    "pputida_KT2440_ASM756v2",
    "GCA_000007565.2_ASM756v2_genomic.gbff.gz",
)
MG1655_GBFF = osp.join(
    DATA_ROOT,
    "torchcell-genomes",
    "ecoli_K12_MG1655_ASM584v2",
    "GCA_000005845.2_ASM584v2_genomic.gbff.gz",
)
KANG_TABLE_INDEX = 6
UNIPROT_DB_XREF = 'db_xref="UniProtKB'


def sha256_of(path: str) -> str:
    """The pin of a file we quote from."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def uniprot_db_xrefs(path: str) -> int:
    """How many UniProtKB cross-references a GenBank flat file carries."""
    total = 0
    with gzip.open(path, "rt") as handle:
        for line in handle:
            if UNIPROT_DB_XREF in line:
                total += 1
    return total


def report_kang() -> None:
    """Kang 2026 Table S6: what it releases, and why its keys cannot be resolved."""
    print("== Kang 2026 Supplementary Table S6 ==")
    print(f"si/si1.docx sha256 {sha256_of(KANG_SI)}")
    table = si_table(KANG_SI, KANG_TABLE_INDEX)
    header = table.header()
    data = table.rows[1:]
    print(f"caption: {table.caption}")
    print(f"columns ({len(header)}): {header}")
    print(f"rows including the header: {len(table.rows)}; data rows: {len(data)}")
    accessions = [row[table.column("Protein Group")] for row in data]
    print(f"distinct keys in 'Protein Group': {len(set(accessions))}")
    print(f"first three keys: {accessions[:3]}")
    p_values = [float(row[table.column("p-Value (Equal Variance)")]) for row in data]
    print(
        f"rows with p < 0.05: {sum(1 for p in p_values if p < 0.05)} of {len(p_values)}"
    )
    print(
        f"UniProtKB db_xrefs, pinned P. putida KT2440: {uniprot_db_xrefs(PPUTIDA_GBFF)}"
    )
    print(
        f"UniProtKB db_xrefs, pinned E. coli MG1655:   {uniprot_db_xrefs(MG1655_GBFF)}"
    )


def report_lim() -> None:
    """Lim 2025: the two arms of every released proteome fold change."""
    print("\n== Lim 2025 si/si2.xlsx proteome sheets ==")
    print(f"si/si2.xlsx sha256 {sha256_of(LIM_SI)}")
    workbook = openpyxl.load_workbook(LIM_SI, read_only=True, data_only=True)
    print(f"sheets: {workbook.sheetnames}")
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = [cell for cell in header if isinstance(cell, str)]
        arms = [cell for cell in columns if cell.startswith("log2_mean")]
        if not arms:
            continue
        fold_changes = [cell for cell in columns if "Fold_change" in cell]
        print(
            f"[{name}] rows={worksheet.max_row - 1} arms={arms} "
            f"fold_change={fold_changes}"
        )
    workbook.close()


if __name__ == "__main__":
    report_kang()
    report_lim()
