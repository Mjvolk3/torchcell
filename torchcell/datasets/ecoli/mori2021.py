# torchcell/datasets/ecoli/mori2021
# [[torchcell.datasets.ecoli.mori2021]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/mori2021
# Test file: tests/torchcell/datasets/ecoli/test_mori2021.py
"""Mori 2021, the absolute E. coli proteome across diverse growth conditions (row 34).

Mori et al. 2021 (Mol Syst Biol 17:e9536, doi:10.15252/msb.20209536, PMID 34032011)
measured absolute protein abundances for *E. coli* over 66 LC-MS samples spanning
nutrient limitations, non-metabolic stresses and non-planktonic states, using DIA/SWATH
with the xTop protein-inference algorithm. :class:`ProteomeMori2021Dataset` serves one
``BacterialProteinAbundanceExperiment`` per loaded sample.

WHICH PER-PROTEIN QUANTITY IS STORED, AND ITS DEFINITION. The release carries several
per-protein quantities and they are NOT interchangeable: ribosome-profiling synthesis
mass fractions from Li 2014, xTop / TopPep1 / TopPep3 / iBAQ mass fractions for the
seven calibration samples (Dataset EV6), the ribosome-profiling-corrected absolute mass
fractions (Datasets EV8 and EV9), the per-limitation response slopes (Dataset EV11) and
the peptide-precursor intensities (Datasets EV4 and EV5). This loader stores the
**absolute protein mass fraction** of Dataset EV9, whose sheet description is
"Absolute protein mass fractions computed from xTop protein intensities and corrected
with ribosome profiling synthesis rates for the samples described in Dataset EV3
(Samples-2)." The quantity itself is defined in the paper: "The resulting absolute
protein abundances are expressed in "protein mass fractions", i.e., mass of a given
protein over the total mass of all detected proteins, which can readily be converted to
cellular protein concentration (Appendix Note S1)." It is dimensionless and normalized
per sample: measured on the pinned workbook, every one of the 36 EV9 columns sums to 1
to within 3e-8, which ``check_normalization`` asserts at build time.

The unit is NOT copies per cell, and the paper refuses that unit on purpose: "Note that
the frequently used absolute unit "protein copies/cell" is avoided here, as cell size is
highly variable across growth conditions". The ``measurement_type`` therefore names the
quantity and the pipeline that produced it.

SEVEN RECORDS FROM SIXTY-SIX RELEASED SAMPLES, AND EVERY DROP IS A MEDIUM. The released
mass fractions are 30 columns of Dataset EV8 (``Lib-01`` .. ``Lib-30``, the Dataset EV2
samples) plus 36 columns of Dataset EV9 (the Dataset EV3 samples), which is the paper's
own "66 different samples". The Appendix names the base media: "Unless otherwise
indicated, growth media used are based on one of the following base media: modified
Record's MOPS medium (Chen et al., 1990), phosphate-buffered "N-C-" medium (Csonka et
al, 1994; Gutnick et al, 1969), M9 medium (Kochanowski et al, 2013), and Luria-Bertani
(LB) medium." Of those, only one released medium is an object ``MEDIA_LIBRARY`` already
states, and ``media.py`` is a value-surface file this branch does not edit, so the other
59 samples are left out rather than written against a medium that joins nothing:

- ``medium_has_no_media_library_entry`` (52 samples): modified Record's MOPS (19),
  M9 as Kochanowski 2013 states it (23, nitrogen source 11.34 mM (NH4)2SO4, which the
  library's ``M9`` salts do not carry), phosphate-buffered N-C- (7), the anaerobic
  phosphate-buffered medium the Appendix spells out component by component (1) and the
  low-osmolarity MOPS of the biofilm and maltose-pyruvate samples (2).
- ``medium_differs_from_the_library_entry_it_names`` (2 samples): ``Lib-28`` and
  ``Lib-30`` name "MOPS (Neidhardt's)" but their Dataset EV2 nitrogen source is 20 mM
  NH4Cl, where ``MOPS_MINIMAL`` states 9.5 mM; the object is not those two samples'
  medium.
- ``medium_formulation_not_stated_by_the_source`` (5 samples): the LB samples. The
  Appendix names "Luria-Bertani (LB) medium" and states no amounts, and the library
  holds both ``LB`` (Miller, 10 g/L NaCl) and ``LB_LENNOX`` (5 g/L NaCl), two of the
  fifty rows using each. Picking one would be a guess about which broth was weighed.

The seven loaded samples are Dataset EV3's calibration group (``A1-1``, ``A1-2``,
``A1-3``, ``C1``, ``F1-1``, ``F1-2``, ``F1-3``), whose released growth medium is
"MOPS (Neidhardt)" with nitrogen source "9.5 mM NH4Cl" and carbon source "0.2%
glucose". That is :data:`MOPS_MINIMAL` exactly: the library object is Neidhardt 1974's
MOPS minimal without a carbon source at 9.5 mM NH4Cl and 50 mM NaCl, and the paper's own
sentence pins the 50 mM, because it states the biofilm medium as a departure from it
("This medium is Neidhardt's MOPS minimal medium (Neidhardt et al, 1974), but with 1/4 of
the stated MOPS buffer (10 mM final concentration); further, NaCl concentration was
adjusted to a final concentration of 10 mM (instead of 50 mM)."). The glucose is carried
as ``EnvironmentPhysicalPerturbation(factor=carbon_source)``; the ammonium is already a
component of the medium, so it is not a second edit of it.

52 + 2 + 5 dropped + 7 loaded = 66.

ONE RECORD PER RELEASED SAMPLE COLUMN, AND NO VALUE IS AVERAGED. The seven loaded
samples are three biological replicates of one culture condition, two of which were
injected three times: ``A1-1`` .. ``A1-3`` are injections of culture ``A1``, ``F1-1`` ..
``F1-3`` of culture ``F1``, and ``C1`` is a single injection ("Biological replicate of
A1"). The stored value is the released column verbatim, so ``n_replicates`` is 1 for
every protein of every record and ``protein_abundance_se`` is ``None``: one LC-MS
injection has no replicate, and an SE of the group is not an SE of the injection.

THE REFERENCE IS THE CALIBRATION SAMPLE, AND ITS SE IS DERIVED FROM THE RELEASED
COLUMNS. The release designates this condition as the sample the absolute scale was set
on: "Exclusively to the three biological replicates of the calibration sample (E. coli
strain K-12 MG1655 grown in glucose minimal media at exponential growth phase) a set of
29 stable isotope labeled peptides (AQUA peptides (Gerber et al, 2003)) was spiked after
digestion and before C18 purification." ``phenotype_reference`` is therefore that
sample's own profile: per protein, the mean of the three biological replicate means
(``A1``, ``C1``, ``F1``), where a culture's mean is taken over its own quantified
injections. ``n_replicates`` is how many of the three cultures quantified that protein
and the SE is ``stdev(culture means) / sqrt(n)``, which is exact arithmetic on the
released columns rather than a released summary statistic; a protein only one culture
quantified stores ``nan``, which is what no replicate means. Measured on the pinned
workbook: the relative SE is 0.031 at the median and 0.122 at the 95th percentile over
the 1,719 proteins all seven columns quantify.

THE IDENTIFIER ROUTE IS THE RELEASED B-NUMBER, AND IT IS THE RECORDS' OWN NAMESPACE.
Every loaded sample is strain EQ353 ("For the calibration samples A1, C1 and F1, we used
strain EQ353, which is the specific MG1655 strain used in Li et al. (Li et al, 2014)
2014."), so ``REFERENCE_STRAIN`` is MG1655 and the release's ``Gene locus`` column IS an
MG1655 locus tag rather than a foreign identifier. Resolved through
:func:`reconcile_locus_tags` against the pinned MG1655 GenBank annotation, 2,073 of
2,073 released b-numbers resolve (1.0000), 2,070 on the locus-tag layer and 3 on a gene
synonym, with no ambiguity and no collision. The ``Gene name`` symbol is kept in
``preprocess/protein_identifiers.csv`` as the source's cross-reference; it is NOT the
key, because the symbol route resolves 2,072 of 2,077 and leaves four retired symbols
(``gapC_1``, ``ilvG_1``, ``rdoA``, ``yedS_1``) and one ambiguous symbol (``rffT``, which
the annotation carries on both ``b3793`` and ``b4481``) outside the namespace.

THE RETENTION LEDGER, WITH ITS ARITHMETIC. 4,342 released EV9 rows become 2,073 protein
keys: 4,342 - 2,265 - 4 = 2,073.

- 2,265 rows are 0 in all seven loaded columns. A released 0 is "not quantified in this
  sample", not a measured zero: the column sums to 1 over the quantified proteins, and
  Dataset EV11 states the same convention for its own slopes ("Proteins that are not
  detected in all three growth limitations are assigned to the "X"-sector").
- 4 of the 2,077 rows quantified in at least one loaded column carry no released
  ``Gene locus`` (``ygaU``, ``yghZ``, ``yifE``, ``ymfP``), so they name no gene.

Each record carries the subset of the 2,073 its own column quantifies: 1,897 (``A1-1``),
1,879 (``A1-2``), 1,911 (``A1-3``), 1,930 (``C1``), 1,896 (``F1-1``), 1,931 (``F1-2``)
and 1,914 (``F1-3``). Every record's reference is key-matched to that record.

AN IDENTITY FAULT IN THE RELEASED TABLES, FOUND BY MEASUREMENT AND NOT CONSUMED. EV8 and
EV9 carry the same 4,342 rows in the same ``Gene name`` order, but their identity blocks
disagree: EV9 leaves ``Gene locus`` blank for 30 rows, and for 12 of them EV8 supplies a
b-number (``ybaO`` b0447, ``ybbB`` b0503, ``ydfP`` b1553, ``yfaW`` b2247, ``yfhG``
b2555, ``ygaU`` b2665, ``yghZ`` b3001, ``yhfG`` b3362, ``yifE`` b3764, ``ykgN`` b4505,
``ymfO`` b1151, ``ymfP`` b1152). The remaining 18 are the IS-element rows (``insAB-*``,
``insCD-*``, ``insEF-*``, ``insJK``), which neither sheet identifies. This loader keys on
EV9's own identity block and does NOT read EV8's as a substitute, so the four dropped
rows stay dropped; the fault is asserted at build time so the finding stays pinned to the
bytes.

MORI AND SCHMIDT 2016 DO NOT OVERLAP, AND THE CORRELATION SAYS SO. Both are absolute
*E. coli* proteomes and Mori reused Schmidt's sample-preparation protocol, so the two
releases were checked against each other. Mori's loaded records are MG1655 (EQ353) in
Neidhardt MOPS glucose; the landed Schmidt 2016 records are BW25113 in M9, so no loaded
(protein, strain, condition) triple is shared at all. Comparing the nearest pair anyway
-- Mori ``A1-1`` against Schmidt's ``Glucose``, with Schmidt's copies/cell converted to a
mass fraction through its own released molecular weights -- 1,812 UniProt accessions are
shared, the log10 values correlate at Pearson r = 0.796 (Spearman rho = 0.832), and ZERO
of the 1,812 pairs agree even to 1e-6 relative, with a median ratio of 1.61 and a 95th
percentile of 29.9. These are two independent measurements, and the spread is the
direction the paper itself reports ("While the number of detected proteins were
comparable (1,901 vs 1,843), we observed a systematic underestimation of low-abundant
proteins"). Mori ``A1-1`` quantifies exactly 1,901 proteins with a UniProt accession,
which is the paper's own figure. Nothing is deduplicated and neither dataset attributes
its values to the other.

DATA. Six files are consumed, all from the PMC OA Cloud bucket (``PMC8144880.1``) and all
scriptable, deposited under ``$DATA_ROOT/torchcell-raw/<citation key>/data/``:
``si1.docx`` (the Appendix, whose Extended Experimental Methods is the only statement of
the base media and the culture temperature), ``si2.xlsx`` (Dataset EV1, the strain
table), ``si3.xlsx`` (Dataset EV2, the Samples-1 metadata), ``si4.xlsx`` (Dataset EV3,
the Samples-2 metadata), ``si9.xlsx`` (Dataset EV8, whose 30 column headers are the
Samples-1 half of the 66 and whose identity block pins the EV9 fault) and ``si10.xlsx``
(Dataset EV9, the stored mass fractions). The raw mass spectra live in PRIDE/SWATHAtlas
(PASS01421) and in Panorama Public, which the mirror manifest records and which no loader
reads.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import math
import os
import os.path as osp
import pickle
import re
import shutil
import statistics
import zipfile
from collections.abc import Callable, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, ClassVar

import openpyxl
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
    write_verified,
)
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.media import MOPS_MINIMAL
from torchcell.datamodels.schema import (
    BacterialAssemblySet,
    BacterialProteinAbundanceExperiment,
    BacterialProteinAbundanceExperimentReference,
    BacterialStrainBackground,
    Concentration,
    ConcentrationUnit,
    Environment,
    EnvironmentPerturbationType,
    EnvironmentPhysicalPerturbation,
    Experiment,
    ExperimentReference,
    Genotype,
    PhysicalFactor,
    ProteinAbundancePhenotype,
    Publication,
    Temperature,
)
from torchcell.datasets.bacteria_common import (
    BACTERIAL_ASSEMBLY_SETS,
    LocusTagReconciliation,
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ROLE_RAW_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.literature.provenance import run_retriever
from torchcell.sequence import GeneSet
from torchcell.sequence.genome.ecoli.k12 import EcoliK12Genome, EcoliK12StrainName
from torchcell.verification.report import Provenance, VerificationReport
from torchcell.verification.sourced import SourcedValue

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# The pinned artifacts
# --------------------------------------------------------------------------- #
CITATION_KEY = "moriCoarseFineAbsolute2021"
PAPER_DOI = "10.15252/msb.20209536"
PAPER_TITLE = (
    "From coarse to fine: the absolute Escherichia coli proteome under diverse "
    "growth conditions"
)
#: Resolved from the PMC id converter, which returns this PMID and the DOI above for
#: ``PMC8144880`` (the PMC record the mirrored SI was retrieved from).
PUBMED_ID = "34032011"
PMC_ID = "PMC8144880"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
LIBRARY_DIR_REL = f"torchcell-library/{CITATION_KEY}"

PAPER_MD = "paper.md"
PAPER_MD_SHA256 = "c86b36d90657acebac0672d866b398de8b0a115f63024fac9aaf54a2fd11f4f8"

APPENDIX = "si1.docx"
EV1 = "si2.xlsx"
EV2 = "si3.xlsx"
EV3 = "si4.xlsx"
EV8 = "si9.xlsx"
EV9 = "si10.xlsx"

#: ``{raw file name: (pinned sha256, PMC Cloud key)}`` of every consumed file.
CONSUMED: dict[str, tuple[str, str]] = {
    APPENDIX: (
        "3c9490d042c152bd478aa82e9f37759dbc9f8beba521ff717a01aee2dd89aed5",
        f"{PMC_ID}.1/MSB-17-e9536-s013.docx",
    ),
    EV1: (
        "78b8629a8e13229d9ac6fd4c12b54c059a115779215731ef938368d5bb739e8c",
        f"{PMC_ID}.1/MSB-17-e9536-s001.xlsx",
    ),
    EV2: (
        "2bb6c8c3b647740c6fe371fc39bf1bb9e69484d0fff264b67e05f61b86062eb9",
        f"{PMC_ID}.1/MSB-17-e9536-s014.xlsx",
    ),
    EV3: (
        "3c13d83bc11996c042f8856ac40ca3b0f312d826ef5118b3f13e060c357c7b46",
        f"{PMC_ID}.1/MSB-17-e9536-s011.xlsx",
    ),
    EV8: (
        "15795fc3405735fdfb607d2344d14ca84ac0c3320c89fea0336731874b94b854",
        f"{PMC_ID}.1/MSB-17-e9536-s012.xlsx",
    ),
    EV9: (
        "50efbf4a3c7f9c1bb4cb6525053d5f599bd13f1b25b296819315ca5d7d244e74",
        f"{PMC_ID}.1/MSB-17-e9536-s007.xlsx",
    ),
}
#: ``{raw file name: pinned sha256}``, re-checked at the start of ``process()``.
DATA_SHA256: dict[str, str] = {name: sha for name, (sha, _) in CONSUMED.items()}
#: Where each consumed file lives inside the raw mirror.
MIRROR_RELPATH: dict[str, str] = {name: f"data/{name}" for name in CONSUMED}
RETRIEVED_AT = "2026-10-07T11:49:21.802498+00:00"


def _retrieval(name: str) -> RetrievalRecord:
    """The recorded PMC Cloud retrieval of one consumed file."""
    sha, key = CONSUMED[name]
    return RetrievalRecord(
        method=RetrievalMethod.pmc_cloud,
        source_url=f"https://pmc-oa-opendata.s3.amazonaws.com/{key}",
        retriever="torchcell.literature.retrieve.pmc_cloud_object",
        params={"key": key},
        sha256=sha,
        retrieved_at=RETRIEVED_AT,
    )


RETRIEVALS: dict[str, RetrievalRecord] = {name: _retrieval(name) for name in CONSUMED}

#: Released locations this loader does NOT mirror, with why.
NOT_MIRRORED = (
    "SWATHAtlas PASS01421 (the E. coli spectral library and peptide query parameters): "
    "no loader reads a spectral library",
    "Panorama Public (the AQUA-peptide Skyline documents): no loader reads spectra",
    "si/si5.xlsx and si/si6.xlsx (Datasets EV4 and EV5, peptide-precursor intensities): "
    "the stored quantity is a protein mass fraction, not a peptide intensity",
    "si/si7.xlsx, si/si8.xlsx, si/si11.xlsx, si/si12.xlsx, si/si13.xlsx (Datasets EV6, "
    "EV7, EV10, EV11, EV12): other per-protein quantities (TopPep1/3, iBAQ, AQUA, "
    "ribosome profiling, sector slopes, GO enrichment), none of them the stored one",
    "si/si14.pdf (the review process file): no value is read from it",
)


# --------------------------------------------------------------------------- #
# Verbatim quotes. Every one is a substring of the pinned ``paper.md``, of the pinned
# Appendix as :func:`docx_paragraphs` renders it, or of a workbook row rendering (its
# non-empty cells joined by " | "), the source's own typography included.
# --------------------------------------------------------------------------- #
_Q_MASS_FRACTION = (
    "The resulting absolute protein abundances are expressed in “protein mass "
    "fractions”, i.e., mass of a given protein over the total mass of all detected "
    "proteins, which can readily be converted to cellular protein concentration "
    "(Appendix Note S1)."
)
_Q_NOT_COPIES = (
    "Note that the frequently used absolute unit “protein copies/cell” is "
    "avoided here, as cell size is highly variable across growth conditions"
)
_Q_TOTAL_2335 = (
    "A total of 2,335 proteins were detected from 66 samples across these conditions."
)
_Q_66_SAMPLES = (
    "The workflow outlined above allowed us to analyze E. coli proteomes over 66 "
    "different samples representing an array of different treatments, strains, and "
    "growth conditions."
)
_Q_WORKFLOW = (
    "we developed a versatile mass spectrometric workflow based on data-independent "
    "acquisition proteomics (DIA/SWATH) together with a novel protein inference "
    "algorithm (xTop)"
)
_Q_SCHMIDT_COMPARE = (
    "We compared the latter to our results for MG1655 (EQ353) in glucose match culture. "
    "While the number of detected proteins were comparable (1,901 vs 1,843), we observed "
    "a systematic underestimation of low-abundant proteins"
)
_Q_EQ353 = (
    "For the calibration samples A1, C1 and F1, we used strain EQ353, which is the "
    "specific MG1655 strain used in Li et al. (Li et al, 2014) 2014."
)
_Q_BASE_MEDIA = (
    "Unless otherwise indicated, growth media used are based on one of the following "
    "base media: modified Record’s MOPS medium (Chen et al., 1990), "
    "phosphate-buffered “N-C-” medium (Csonka et al, 1994; Gutnick et al, "
    "1969), M9 medium (Kochanowski et al, 2013), and Luria-Bertani (LB) medium."
)
_Q_BATCH_37C = (
    "Batch cultures were grown in a 37°C water bath shaker shaking at 250 rpm for "
    "aeration."
)
_Q_CALIBRATION_IDENTICAL = (
    "For the A1, C1 and F1 samples, both strain, growth media and experimental "
    "procedures are identical to those used in a ribosome profiling study (Li et al., "
    "2014) to measure protein synthesis rates for E. coli in glucose-limited media."
)
_Q_NEIDHARDT_NACL = (
    "This medium is Neidhardt’s MOPS minimal medium (Neidhardt et al, 1974), but "
    "with ¼ of the stated MOPS buffer (10 mM final concentration); further, NaCl "
    "concentration was adjusted to a final concentration of 10 mM (instead of 50 mM)."
)
_Q_THREE_BIOLOGICAL = (
    "Exclusively to the three biological replicates of the calibration sample (E. coli "
    "strain K-12 MG1655 grown in glucose minimal media at exponential growth phase) a "
    "set of 29 stable isotope labeled peptides (AQUA peptides (Gerber et al, 2003)) was "
    "spiked after digestion and before C18 purification."
)
_Q_EV1_EQ353 = (
    "MG1655 (EQ353) | Wild type E. coli strain - same strain used in Li et al. (2014) | "
    "Originarily obtained from Carol Gross Lab"
)
_Q_EV9_DESCRIPTION = (
    "Absolute protein mass fractions computed from xTop protein intensities and "
    "corrected with ribosome profiling synthesis rates for the samples described in "
    "Dataset EV3 (Samples-2)."
)
_Q_EV8_DESCRIPTION = (
    "Absolute protein mass fractions computed from xTop protein intensities and "
    "corrected with ribosome profiling synthesis rates for the samples described in "
    "Dataset EV2 (Samples-1)."
)
_Q_EV3_DESCRIPTION = (
    'Table with informations on the 7 "calibration" samples for E. coli MG1655 '
    "(EQ353), plus the three growth limitation series (C-, A- and R-limitation), "
    "obtained with either E. coli NCM3722 or NCM3722-derived strains."
)
_Q_EV11_NOT_DETECTED = (
    "Proteins that are not detected in all three growth limitations are assigned to the "
    "“X”-sector"
)

PAPER = Provenance(
    source_uri=PAPER_MD, citation_key=CITATION_KEY, sha256=PAPER_MD_SHA256
)
APPENDIX_SOURCE = Provenance(
    source_uri=MIRROR_RELPATH[APPENDIX],
    citation_key=CITATION_KEY,
    sha256=CONSUMED[APPENDIX][0],
)
EV1_SOURCE = Provenance(
    source_uri=MIRROR_RELPATH[EV1], citation_key=CITATION_KEY, sha256=CONSUMED[EV1][0]
)
EV3_SOURCE = Provenance(
    source_uri=MIRROR_RELPATH[EV3], citation_key=CITATION_KEY, sha256=CONSUMED[EV3][0]
)
EV8_SOURCE = Provenance(
    source_uri=MIRROR_RELPATH[EV8], citation_key=CITATION_KEY, sha256=CONSUMED[EV8][0]
)
EV9_SOURCE = Provenance(
    source_uri=MIRROR_RELPATH[EV9], citation_key=CITATION_KEY, sha256=CONSUMED[EV9][0]
)

#: Quotes whose provenance is ``paper.md``, by the constant's own name.
PAPER_QUOTES: dict[str, str] = {
    "mass_fraction": _Q_MASS_FRACTION,
    "not_copies_per_cell": _Q_NOT_COPIES,
    "total_2335": _Q_TOTAL_2335,
    "sixty_six_samples": _Q_66_SAMPLES,
    "workflow": _Q_WORKFLOW,
    "schmidt_comparison": _Q_SCHMIDT_COMPARE,
}
#: Quotes whose provenance is the pinned Appendix (``si1.docx``).
APPENDIX_QUOTES: dict[str, str] = {
    "eq353": _Q_EQ353,
    "base_media": _Q_BASE_MEDIA,
    "batch_37c": _Q_BATCH_37C,
    "calibration_identical": _Q_CALIBRATION_IDENTICAL,
    "neidhardt_nacl": _Q_NEIDHARDT_NACL,
    "three_biological_replicates": _Q_THREE_BIOLOGICAL,
}
#: Quotes whose provenance is a workbook row rendering, by file name.
WORKBOOK_QUOTES: dict[str, dict[str, str]] = {
    EV1: {"eq353_row": _Q_EV1_EQ353},
    EV3: {"description": _Q_EV3_DESCRIPTION},
    EV8: {"description": _Q_EV8_DESCRIPTION},
    EV9: {"description": _Q_EV9_DESCRIPTION},
}


def _paper(value: object, quote: str, *, note: str | None = None) -> SourcedValue:
    """A ``SourcedValue`` quoting the pinned ``paper.md``."""
    return SourcedValue(value=value, provenance=PAPER, quote=quote, note=note)


def _appendix(value: object, quote: str, *, note: str | None = None) -> SourcedValue:
    """A ``SourcedValue`` quoting the pinned Appendix."""
    return SourcedValue(value=value, provenance=APPENDIX_SOURCE, quote=quote, note=note)


def _workbook(
    source: Provenance, value: object, quote: str, *, note: str | None = None
) -> SourcedValue:
    """A ``SourcedValue`` quoting one pinned workbook's row rendering."""
    return SourcedValue(value=value, provenance=source, quote=quote, note=note)


#: The strain every loaded record is written against, as the Appendix states it.
MORI_REFERENCE_STRAIN: EcoliK12StrainName = "MG1655"
#: That strain's deposited assembly set, typed by the schema's Literal so a typo fails
#: type checking; the test module asserts it is what ``BACTERIAL_ASSEMBLY_SETS`` says.
MORI_ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_MG1655_ASM584v2"
#: What one stored abundance IS: the quantity and the pipeline that produced it.
MEASUREMENT_TYPE = "absolute_protein_mass_fraction_xtop_dia_swath_riboprofiling_scaled"

SOURCED_VALUES: dict[str, SourcedValue] = {
    "reference_strain": _appendix(MORI_REFERENCE_STRAIN, _Q_EQ353),
    "strain_label": _workbook(EV1_SOURCE, "MG1655 (EQ353)", _Q_EV1_EQ353),
    "stored_quantity": _paper("absolute protein mass fraction", _Q_MASS_FRACTION),
    "stored_quantity_unit": _paper(
        "dimensionless fraction of total detected protein mass",
        _Q_MASS_FRACTION,
        note="the released columns are normalized per sample and sum to 1, which "
        "check_normalization asserts against the bytes",
    ),
    "unit_is_not_copies_per_cell": _paper(None, _Q_NOT_COPIES),
    "abundance_model": _workbook(
        EV9_SOURCE,
        "xTop protein intensities corrected with the Li 2014 ribosome-profiling "
        "synthesis rates",
        _Q_EV9_DESCRIPTION,
    ),
    "quantification_workflow": _paper("DIA/SWATH with xTop", _Q_WORKFLOW),
    "n_released_samples": _paper(66, _Q_66_SAMPLES),
    "n_detected_proteins": _paper(2335, _Q_TOTAL_2335),
    "base_media": _appendix(
        [
            "modified Record's MOPS medium (Chen et al., 1990)",
            "phosphate-buffered N-C- medium (Csonka et al, 1994; Gutnick et al, 1969)",
            "M9 medium (Kochanowski et al, 2013)",
            "Luria-Bertani (LB) medium",
        ],
        _Q_BASE_MEDIA,
        note="only MOPS minimal (Neidhardt 1974), which the calibration samples use, is "
        "an object MEDIA_LIBRARY states; the other four are the drop rules",
    ),
    "temperature_c": _appendix(37.0, _Q_BATCH_37C),
    "aerobicity": _appendix(
        "aerobic",
        _Q_BATCH_37C,
        note="a water-bath shaker at 250 rpm, which the Appendix states is for aeration",
    ),
    "neidhardt_mops_nacl_mm": _appendix(
        50.0,
        _Q_NEIDHARDT_NACL,
        note="stated as the amount the biofilm medium departs from, which pins the "
        "unmodified Neidhardt MOPS this loader writes the calibration samples against",
    ),
    "calibration_is_the_reference": _appendix(
        ["A1", "C1", "F1"],
        _Q_THREE_BIOLOGICAL,
        note="why the calibration sample's three-culture mean is the phenotype_reference",
    ),
    "calibration_matches_li2014": _appendix(None, _Q_CALIBRATION_IDENTICAL),
    # Dataset EV11 states the same convention for its own slopes, but it is not a
    # consumed file, so the convention is sourced from the paper's own definition.
    "not_detected_is_zero": _paper(
        0.0,
        _Q_MASS_FRACTION,
        note="a released 0 is 'not quantified in this sample': the definition is over "
        "the detected proteins and each released column sums to 1 over its non-zero "
        "rows, so a 0 row contributes nothing and is not a measured zero",
    ),
}


# --------------------------------------------------------------------------- #
# Reading the pinned Appendix
# --------------------------------------------------------------------------- #
_WORD_TEXT = re.compile(r"<w:t[^>]*>(.*?)</w:t>", re.DOTALL)
_WORD_PARAGRAPH = re.compile(r"</w:p>")
_XML_ENTITIES = (("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'))


def docx_paragraphs(path: str | Path) -> list[str]:
    """The Appendix's paragraph texts, in document order.

    A deterministic reader of ``word/document.xml``: every ``<w:t>`` run of a paragraph
    joined as written, paragraphs split on ``</w:p>``, XML entities decoded, empty
    paragraphs dropped. This is the rendering every Appendix quote is checked against,
    so the quotes verify against the pinned bytes rather than against a description of
    them.
    """
    with zipfile.ZipFile(path) as book:
        xml = book.read("word/document.xml").decode("utf-8")
    out: list[str] = []
    for chunk in _WORD_PARAGRAPH.split(xml):
        text = "".join(_WORD_TEXT.findall(chunk))
        for entity, char in _XML_ENTITIES:
            text = text.replace(entity, char)
        if text.strip():
            out.append(text)
    return out


def appendix_text(path: str | Path) -> str:
    """The Appendix's paragraphs joined by newlines, for a substring check."""
    return "\n".join(docx_paragraphs(path))


# --------------------------------------------------------------------------- #
# Workbook geometry
# --------------------------------------------------------------------------- #
SHEET_EV1 = "EV1-Strains"
SHEET_EV2 = "EV2-Samples-1"
SHEET_EV3 = "EV3-Samples-2"
SHEET_EV8 = "EV8-AbsoluteMassFractions-1"
SHEET_EV9 = "EV9-AbsoluteMassFractions-2"
SHEET_DESCRIPTION = "Description"
#: The identity columns of EV8 and EV9, in order, before the sample columns.
IDENTITY_COLUMNS = ("Gene name", "Gene locus", "Protein ID")
N_IDENTITY_COLUMNS = len(IDENTITY_COLUMNS)
#: Dataset EV3's metadata headers.
COL_SAMPLE_ID = "Sample ID"
COL_GROUP = "Group"
COL_GROWTH_RATE = "Growth rate (1/h)"
COL_STRAIN = "Strain"
COL_MEDIUM = "Growth medium"
COL_CARBON = "Carbon source"
COL_NITROGEN = "Nitrogen source"
COL_SUPPLEMENT = "Supplement"
COL_DESCRIPTION = "Description"
COL_SWATH = "SWATH file name"
#: Dataset EV2's metadata headers that this loader reads (its own column spelling).
EV2_COL_SAMPLE_ID = "Sample ID"
EV2_COL_STRAIN = "Strain"
EV2_COL_MEDIUM = "Base medium"
EV2_COL_NITROGEN = "Nitrogen Source"


# --------------------------------------------------------------------------- #
# The released samples and the rules that keep or drop them
# --------------------------------------------------------------------------- #
class DropReason(BaseModel):
    """Why one released sample is not a record."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    rule: str
    description: str
    needed_addition: str | None = None


DROP_NO_MEDIA_ENTRY = DropReason(
    rule="medium_has_no_media_library_entry",
    description=(
        "the sample's released base medium is a formulation no MEDIA_LIBRARY key "
        "states; media.py is a value-surface file this branch does not edit, so the "
        "sample is left out rather than written against a medium that joins nothing"
    ),
    needed_addition=(
        "MOPS_RECORDS_MORI2021, M9_MORI2021, NC_MINUS_MORI2021, "
        "NC_MINUS_ANAEROBIC_MORI2021 and MOPS_LOW_OSMOLARITY_MORI2021"
    ),
)
DROP_MEDIA_ENTRY_DIFFERS = DropReason(
    rule="medium_differs_from_the_library_entry_it_names",
    description=(
        "the sample names MOPS (Neidhardt's) but its released nitrogen source is 20 mM "
        "NH4Cl, where MOPS_MINIMAL states 9.5 mM; the library object is not this "
        "sample's medium"
    ),
    needed_addition="MOPS_MINIMAL_20MM_NH4CL_MORI2021",
)
DROP_FORMULATION_NOT_STATED = DropReason(
    rule="medium_formulation_not_stated_by_the_source",
    description=(
        "the source names 'Luria-Bertani (LB) medium' and states no amounts, and the "
        "library holds both LB (Miller, 10 g/L NaCl) and LB_LENNOX (5 g/L NaCl); "
        "choosing one would be a guess about which broth was weighed"
    ),
    needed_addition=(
        "a sourced decision on which LB formulation Mori 2021 used, before either "
        "existing object can carry these five samples"
    ),
)


class SampleSpec(BaseModel):
    """One released mass-fraction column, and how it becomes a record or a drop."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    column: str = Field(description="The EV8 or EV9 header, verbatim.")
    table: str = Field(description="'EV8' or 'EV9'.")
    strain: str = Field(description="The released Strain cell, verbatim.")
    medium: str = Field(description="The released growth-medium cell, verbatim.")
    culture: str | None = Field(
        default=None,
        description="The biological culture a loaded injection belongs to.",
    )
    drop: DropReason | None = None


def _ev9(
    column: str,
    strain: str,
    medium: str,
    *,
    culture: str | None = None,
    drop: DropReason | None = None,
) -> SampleSpec:
    """One Dataset EV9 sample column."""
    return SampleSpec(
        column=column,
        table="EV9",
        strain=strain,
        medium=medium,
        culture=culture,
        drop=drop,
    )


def _ev8(column: str, strain: str, medium: str, drop: DropReason) -> SampleSpec:
    """One Dataset EV8 sample column; every one of the 30 is dropped."""
    return SampleSpec(
        column=column, table="EV8", strain=strain, medium=medium, drop=drop
    )


_RECORDS = "MOPS (Record's)"
_NEIDHARDT = "MOPS (Neidhardt)"
_NEIDHARDT_S = "MOPS (Neidhardt's)"
_M9 = "M9"
_NC = "Phosphate buffered N-C- medium"
_NC_SPECIAL = "Special N-C- (see Methods)"
_MOPS_SPECIAL = "Special MOPS (see Methods)"
_LB = "LB"

#: Every released sample column, in table order: the 30 of Dataset EV8 then the 36 of
#: Dataset EV9. The strain and medium cells are the released metadata of Datasets EV2
#: and EV3, which ``check_sample_metadata`` re-reads from the bytes.
SAMPLES: tuple[SampleSpec, ...] = (
    _ev8("Lib-01", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-02", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-03", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-04", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-05", "NCM3722", _NC_SPECIAL, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-06", "NCM3722", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-07", "NCM3722", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-08", "NQ393", _M9, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-09", "NQ1431", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-10", "EQ 59", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-11", "Nissle1917", _MOPS_SPECIAL, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-12", "MG1655 (CGSC#6300)", _LB, DROP_FORMULATION_NOT_STATED),
    _ev8("Lib-13", "NQ1527", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-14", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-15", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-16", "EQ 59", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-17", "EQ 59", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-18", "EQ 59", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-19", "NCM3722", _LB, DROP_FORMULATION_NOT_STATED),
    _ev8("Lib-20", "NCM3722", _LB, DROP_FORMULATION_NOT_STATED),
    _ev8("Lib-21", "NCM3722", _LB, DROP_FORMULATION_NOT_STATED),
    _ev8("Lib-22", "NCM3722", _LB, DROP_FORMULATION_NOT_STATED),
    _ev8("Lib-23", "MG1655 (CGSC#6300)", _MOPS_SPECIAL, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-24", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-25", "EQ 59", _NC, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-26", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-27", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-28", "NCM3722", _NEIDHARDT_S, DROP_MEDIA_ENTRY_DIFFERS),
    _ev8("Lib-29", "NCM3722", _RECORDS, DROP_NO_MEDIA_ENTRY),
    _ev8("Lib-30", "NCM3722", _NEIDHARDT_S, DROP_MEDIA_ENTRY_DIFFERS),
    _ev9("A1-1", "EQ353", _NEIDHARDT, culture="A1"),
    _ev9("A1-2", "EQ353", _NEIDHARDT, culture="A1"),
    _ev9("A1-3", "EQ353", _NEIDHARDT, culture="A1"),
    _ev9("C1", "EQ353", _NEIDHARDT, culture="C1"),
    _ev9("F1-1", "EQ353", _NEIDHARDT, culture="F1"),
    _ev9("F1-2", "EQ353", _NEIDHARDT, culture="F1"),
    _ev9("F1-3", "EQ353", _NEIDHARDT, culture="F1"),
    _ev9("C2", "NCM3722", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("C3", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("C4", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("C5", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("C6", "NQ1390", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("C7", "NQ1390", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("C8", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D6", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D7", "NQ1390", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D8", "NQ1390", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F4", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F5", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F6", "NQ1243", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F7", "NQ1390", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F8", "NQ1390", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D1", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D2", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D3", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D4", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("D5", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F2", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("F3", "NQ393", _M9, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("A2", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("E1", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("E2", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("E3", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("E4", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("H1", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
    _ev9("H5", "NCM3722", _RECORDS, drop=DROP_NO_MEDIA_ENTRY),
)
LOADED = tuple(spec for spec in SAMPLES if spec.drop is None)
#: ``{culture: its loaded injection columns}``, in table order.
CULTURES: dict[str, tuple[str, ...]] = {
    culture: tuple(s.column for s in LOADED if s.culture == culture)
    for culture in dict.fromkeys(s.culture for s in LOADED if s.culture is not None)
}
#: 66 released samples - 52 - 2 - 5 = 7.
EXPECTED_RECORDS = 7
#: Released rows, rows quantified in at least one loaded column, and stored keys.
EXPECTED_SOURCE_ROWS = 4342
EXPECTED_QUANTIFIED_ROWS = 2077
EXPECTED_PROTEIN_KEYS = 2073
#: Measured: every released column sums to 1 to within this absolute tolerance.
NORMALIZATION_ATOL = 1e-7
#: Measured: 2,073 of 2,073 released b-numbers resolve, so the floor sits just below 1.
MIN_RESOLVED_FRACTION = 0.999
#: The EV9 rows whose ``Gene locus`` is blank and for which EV8 supplies a b-number.
EV8_SUPPLIES_LOCUS: dict[str, str] = {
    "ybaO": "b0447",
    "ybbB": "b0503",
    "ydfP": "b1553",
    "yfaW": "b2247",
    "yfhG": "b2555",
    "ygaU": "b2665",
    "yghZ": "b3001",
    "yhfG": "b3362",
    "yifE": "b3764",
    "ykgN": "b4505",
    "ymfO": "b1151",
    "ymfP": "b1152",
}
#: The EV9 rows neither sheet identifies: the IS-element rows.
UNIDENTIFIED_ROWS = 18


# --------------------------------------------------------------------------- #
# The strain
# --------------------------------------------------------------------------- #
EQ353_BACKGROUND = BacterialStrainBackground(
    name="MG1655 (EQ353)",
    reference_strain=MORI_REFERENCE_STRAIN,
    assembly_set=MORI_ASSEMBLY_SET,
    parents=["MG1655"],
    construction="the MG1655 laboratory stock of the Li 2014 ribosome-profiling study, "
    "originally obtained from the Carol Gross lab; the release states no lesion relative "
    "to MG1655",
    provenance=[
        SourcedValue(value="MG1655 (EQ353)", provenance=EV1_SOURCE, quote=_Q_EV1_EQ353),
        SourcedValue(value="EQ353", provenance=APPENDIX_SOURCE, quote=_Q_EQ353),
    ],
)
"""Which MG1655 stock the records are, as a background on the MG1655 assembly.

The stock is recorded rather than flattened into bare MG1655 because this paper's own
Figure 9 result is that it matters: "different laboratory strains of E. coli, even those
labeled as common MG1655 strain, exhibit important biological differences". ``alleles``
is empty, which is what the release's "Wild type E. coli strain" states.
"""


# --------------------------------------------------------------------------- #
# Reading the pinned workbooks
# --------------------------------------------------------------------------- #
def _rows(path: str, sheet: str) -> list[tuple[Any, ...]]:
    """Every row of one sheet, values only."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(book[sheet].iter_rows(values_only=True))
    finally:
        book.close()


def _render(row: Sequence[Any]) -> str:
    """One row as its non-empty cells joined by ``" | "``, the quote rendering."""
    return " | ".join(
        str(cell).strip() for cell in row if cell is not None and str(cell).strip()
    )


class MassFractionRow(BaseModel):
    """One Dataset EV9 row: its released identifiers and its per-sample mass fractions."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    row_number: int = Field(description="1-based row in the sheet, for the ledger.")
    gene_name: str
    gene_locus: str | None
    protein_id: str | None
    mass_fraction: dict[str, float] = Field(
        description="loaded sample column -> released mass fraction (0 = not quantified)"
    )

    def quantified(self, column: str) -> bool:
        """True when the release carries a non-zero mass fraction in ``column``."""
        return self.mass_fraction[column] > 0.0


def read_mass_fractions(path: str, columns: Sequence[str]) -> list[MassFractionRow]:
    """Read Dataset EV9's rows over ``columns``.

    The identity block must be the three declared headers in order and every requested
    column must be present, so a re-export that renames or moves a column stops the
    build rather than shifting one sample's values onto another.
    """
    rows = _rows(path, SHEET_EV9)
    header = rows[0]
    found = tuple(str(cell).strip() for cell in header[:N_IDENTITY_COLUMNS])
    if found != IDENTITY_COLUMNS:
        raise RuntimeError(
            f"{SHEET_EV9}'s identity block is {found}, the module declares "
            f"{IDENTITY_COLUMNS}"
        )
    index = {str(cell).strip(): i for i, cell in enumerate(header) if cell is not None}
    missing = [name for name in columns if name not in index]
    if missing:
        raise RuntimeError(f"{SHEET_EV9} carries no sample column {missing}")
    out: list[MassFractionRow] = []
    for offset, row in enumerate(rows[1:]):
        values: dict[str, float] = {}
        for name in columns:
            cell = row[index[name]]
            if not isinstance(cell, (int, float)) or isinstance(cell, bool):
                raise RuntimeError(
                    f"{SHEET_EV9} row {offset + 2}, column {name!r}: {cell!r} is not a "
                    "number"
                )
            values[name] = float(cell)
        locus = row[index["Gene locus"]]
        protein = row[index["Protein ID"]]
        out.append(
            MassFractionRow(
                row_number=offset + 2,
                gene_name=str(row[index["Gene name"]]).strip(),
                gene_locus=None if locus is None else str(locus).strip(),
                protein_id=None if protein is None else str(protein).strip(),
                mass_fraction=values,
            )
        )
    return out


def check_normalization(path: str, table: str) -> dict[str, Any]:
    """Every released column of one mass-fraction sheet sums to 1.

    This is what makes the stored quantity a mass fraction rather than an unnormalized
    intensity, so it is asserted against the bytes instead of read off the description.
    """
    sheet = SHEET_EV8 if table == "EV8" else SHEET_EV9
    rows = _rows(path, sheet)
    header = rows[0]
    worst = 0.0
    sums: dict[str, float] = {}
    for index in range(N_IDENTITY_COLUMNS, len(header)):
        name = str(header[index]).strip()
        total = sum(
            float(row[index])
            for row in rows[1:]
            if isinstance(row[index], (int, float)) and not isinstance(row[index], bool)
        )
        sums[name] = total
        worst = max(worst, abs(total - 1.0))
    if worst > NORMALIZATION_ATOL:
        raise RuntimeError(
            f"{sheet}: a released column sums to {1.0 + worst} rather than 1 "
            f"(tolerance {NORMALIZATION_ATOL}); sums {sums}"
        )
    return {"sheet": sheet, "n_columns": len(sums), "worst_abs_deviation": worst}


def check_sample_columns(path: str, table: str) -> tuple[str, ...]:
    """The declared sample columns of one table are that sheet's headers, in order."""
    sheet = SHEET_EV8 if table == "EV8" else SHEET_EV9
    header = _rows(path, sheet)[0]
    found = tuple(
        str(cell).strip() for cell in header[N_IDENTITY_COLUMNS:] if cell is not None
    )
    declared = tuple(spec.column for spec in SAMPLES if spec.table == table)
    if found != declared:
        raise RuntimeError(
            f"{sheet}'s sample headers are {found}, the module declares {declared}"
        )
    return found


def check_sample_metadata(ev2_path: str, ev3_path: str) -> dict[str, Any]:
    """Every declared strain and medium is the released metadata cell, verbatim.

    Dataset EV2 leaves a repeated cell blank for a run of technical replicates, so a
    blank strain or medium inherits the value above it, which is how the sheet reads.
    """
    released: dict[str, tuple[str, str]] = {}
    for path, sheet, sample, strain, medium in (
        (ev2_path, SHEET_EV2, EV2_COL_SAMPLE_ID, EV2_COL_STRAIN, EV2_COL_MEDIUM),
        (ev3_path, SHEET_EV3, COL_SAMPLE_ID, COL_STRAIN, COL_MEDIUM),
    ):
        rows = _rows(path, sheet)
        index = {
            str(cell).strip(): i for i, cell in enumerate(rows[0]) if cell is not None
        }
        last = ("", "")
        for row in rows[1:]:
            name = row[index[sample]]
            if name is None or not str(name).strip():
                continue
            cells = (row[index[strain]], row[index[medium]])
            values = tuple(
                last[i] if cell is None or not str(cell).strip() else str(cell).strip()
                for i, cell in enumerate(cells)
            )
            last = (values[0], values[1])
            released[str(name).strip()] = last
    for spec in SAMPLES:
        if spec.column not in released:
            raise RuntimeError(f"no released metadata row for sample {spec.column!r}")
        if released[spec.column] != (spec.strain, spec.medium):
            raise RuntimeError(
                f"{spec.column}: the release says {released[spec.column]}, the module "
                f"declares {(spec.strain, spec.medium)}"
            )
    nitrogen = _nitrogen_sources(ev3_path)
    for spec in LOADED:
        if nitrogen[spec.column] != LOADED_NITROGEN_SOURCE:
            raise RuntimeError(
                f"{spec.column}: released nitrogen source {nitrogen[spec.column]!r}, "
                f"not {LOADED_NITROGEN_SOURCE!r}, so MOPS_MINIMAL is not its medium"
            )
    return {"n_samples": len(released), "n_loaded": len(LOADED)}


#: The nitrogen source every loaded sample releases, which is what ``MOPS_MINIMAL``
#: states as its own ammonium component.
LOADED_NITROGEN_SOURCE = "9.5 mM NH4Cl"
#: The carbon source every loaded sample releases.
LOADED_CARBON_SOURCE = "0.2% glucose"
LOADED_CARBON_COMPOUND = "glucose"
LOADED_CARBON_PERCENT = 0.2


def _nitrogen_sources(ev3_path: str) -> dict[str, str]:
    """``{Dataset EV3 sample id: its released nitrogen-source cell}``."""
    rows = _rows(ev3_path, SHEET_EV3)
    index = {str(cell).strip(): i for i, cell in enumerate(rows[0]) if cell is not None}
    out: dict[str, str] = {}
    for row in rows[1:]:
        name = row[index[COL_SAMPLE_ID]]
        if name is None or not str(name).strip():
            continue
        out[str(name).strip()] = str(row[index[COL_NITROGEN]]).strip()
    return out


def read_growth_rates(ev3_path: str) -> dict[str, float]:
    """``{Dataset EV3 sample id: its released growth rate in 1/h}``.

    ``Environment`` has no growth-rate slot, so this goes to the condition ledger rather
    than to a record.
    """
    rows = _rows(ev3_path, SHEET_EV3)
    index = {str(cell).strip(): i for i, cell in enumerate(rows[0]) if cell is not None}
    out: dict[str, float] = {}
    for row in rows[1:]:
        name = row[index[COL_SAMPLE_ID]]
        if name is None or not str(name).strip():
            continue
        cell = row[index[COL_GROWTH_RATE]]
        if isinstance(cell, (int, float)) and not isinstance(cell, bool):
            out[str(name).strip()] = float(cell)
    return out


def check_identity_fault(ev8_path: str, ev9_path: str) -> dict[str, Any]:
    """Dataset EV9 blanks a ``Gene locus`` that Dataset EV8 carries, and this pins it.

    EV8 and EV9 hold the same rows in the same ``Gene name`` order. EV9 leaves the locus
    blank for 30 rows, 12 of which EV8 identifies; the other 18 are the IS-element rows
    neither sheet identifies. No stored key reads EV8's identity block; the check exists
    so the finding stays attached to the bytes rather than to a note.
    """
    ev8 = _rows(ev8_path, SHEET_EV8)
    ev9 = _rows(ev9_path, SHEET_EV9)
    if len(ev8) != len(ev9):
        raise RuntimeError(
            f"{SHEET_EV8} has {len(ev8)} rows and {SHEET_EV9} has {len(ev9)}"
        )
    supplied: dict[str, str] = {}
    blank = 0
    for left, right in zip(ev8[1:], ev9[1:], strict=True):
        gene = str(right[0]).strip()
        if str(left[0]).strip() != gene:
            raise RuntimeError(
                f"{SHEET_EV8} and {SHEET_EV9} no longer share one row order: "
                f"{str(left[0]).strip()!r} against {gene!r}"
            )
        if right[1] is not None and str(right[1]).strip():
            continue
        blank += 1
        if left[1] is not None and str(left[1]).strip():
            supplied[gene] = str(left[1]).strip()
    if supplied != EV8_SUPPLIES_LOCUS:
        raise RuntimeError(
            f"{SHEET_EV8} now supplies {supplied} for the rows {SHEET_EV9} blanks, the "
            f"module declares {EV8_SUPPLIES_LOCUS}"
        )
    if blank - len(supplied) != UNIDENTIFIED_ROWS:
        raise RuntimeError(
            f"{blank - len(supplied)} rows are identified by neither sheet, the module "
            f"declares {UNIDENTIFIED_ROWS}"
        )
    return {
        "n_blank_in_ev9": blank,
        "n_supplied_by_ev8": len(supplied),
        "n_unidentified": blank - len(supplied),
    }


# --------------------------------------------------------------------------- #
# The retention ledger
# --------------------------------------------------------------------------- #
class DropRule(BaseModel):
    """One rule that removed released samples or rows, with the items it removed."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    rule: str
    scope: str = Field(description="'sample' or 'protein_row'")
    description: str
    n_items: int
    items: list[str]
    needed_addition: str | None = None


class DropLog(BaseModel):
    """Every released sample and protein row, and what became of it."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    dataset: str
    source_samples: int
    kept_records: int
    dropped_records: int
    source_protein_rows: int
    kept_protein_keys: int
    dropped_protein_rows: int
    rules: list[DropRule]
    reconciliation: LocusTagReconciliation
    notes: list[str]

    def check(self) -> None:
        """Refuse a ledger whose rules do not account for every drop."""
        samples = sum(r.n_items for r in self.rules if r.scope == "sample")
        if self.kept_records + samples != self.source_samples:
            raise RuntimeError(
                f"{self.dataset}: {self.kept_records} records + {samples} dropped "
                f"samples != {self.source_samples} released samples"
            )
        if self.dropped_records != samples:
            raise RuntimeError(
                f"{self.dataset}: {self.dropped_records} dropped samples stated, "
                f"{samples} accounted by rules"
            )
        rows = sum(r.n_items for r in self.rules if r.scope == "protein_row")
        if self.kept_protein_keys + rows != self.source_protein_rows:
            raise RuntimeError(
                f"{self.dataset}: {self.kept_protein_keys} kept keys + {rows} dropped "
                f"rows != {self.source_protein_rows} released rows"
            )
        if self.dropped_protein_rows != rows:
            raise RuntimeError(
                f"{self.dataset}: {self.dropped_protein_rows} dropped rows stated, "
                f"{rows} accounted by rules"
            )


class ProteinSelection(BaseModel):
    """The kept rows, their stored locus tags, and every drop that got there."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    kept: list[MassFractionRow]
    locus_tag: dict[str, str] = Field(description="released gene name -> locus tag")
    reconciliation: LocusTagReconciliation
    rules: list[DropRule]

    @property
    def dropped_rows(self) -> int:
        """Released rows removed by the row-scoped rules."""
        return sum(rule.n_items for rule in self.rules if rule.scope == "protein_row")


def select_proteins(
    rows: Sequence[MassFractionRow], genome: EcoliK12Genome, *, label: str
) -> ProteinSelection:
    """Keep the rows a loaded sample quantifies and resolve their released b-numbers."""
    columns = [spec.column for spec in LOADED]
    quantified = [row for row in rows if any(row.quantified(c) for c in columns)]
    silent = [row for row in rows if row not in quantified]
    no_locus = [row for row in quantified if row.gene_locus is None]
    with_locus = [row for row in quantified if row.gene_locus is not None]

    stored, report = reconcile_locus_tags(
        genome, pd.Series([str(row.gene_locus) for row in with_locus]), label=label
    )
    report.require_resolved(MIN_RESOLVED_FRACTION)
    outside = set(report.outside_namespace)
    kept: list[MassFractionRow] = []
    locus_tag: dict[str, str] = {}
    unresolved: list[MassFractionRow] = []
    for row, tag in zip(with_locus, stored, strict=True):
        if tag in outside:
            unresolved.append(row)
            continue
        kept.append(row)
        locus_tag[row.gene_name] = str(tag)
    if len({*locus_tag.values()}) != len(kept):
        raise RuntimeError(f"{label}: two kept rows share one locus tag")

    rules = [
        DropRule(
            rule="not_quantified_in_any_loaded_sample",
            scope="protein_row",
            description=(
                "the release writes 0 in every loaded column. A released 0 is 'not "
                "quantified in this sample' rather than a measured zero: the stored "
                "quantity is a fraction of the detected protein mass and each released "
                "column sums to 1 over its non-zero rows, which check_normalization "
                "asserts against the bytes"
            ),
            n_items=len(silent),
            items=[row.gene_name for row in silent],
        ),
        DropRule(
            rule="release_files_no_gene_locus_for_this_row",
            scope="protein_row",
            description=(
                "Dataset EV9 leaves this row's Gene locus blank, so it names no gene of "
                "the pinned assembly. Dataset EV8 does carry a b-number for it "
                "(check_identity_fault pins which), and that is deliberately NOT read "
                "as a substitute: the consumed table's own identity block is the key"
            ),
            n_items=len(no_locus),
            items=[f"{row.gene_name} ({row.protein_id})" for row in no_locus],
        ),
        DropRule(
            rule="b_number_resolves_to_no_locus_of_the_pinned_assembly",
            scope="protein_row",
            description=(
                "the released b-number resolves to no locus of "
                f"{BACTERIAL_ASSEMBLY_SETS[MORI_REFERENCE_STRAIN]} through any resolver "
                "layer and is kept as given by the retain-all policy, so it names no "
                "gene node"
            ),
            n_items=len(unresolved),
            items=[f"{row.gene_name} ({row.gene_locus})" for row in unresolved],
        ),
    ]
    return ProteinSelection(
        kept=kept, locus_tag=locus_tag, reconciliation=report, rules=rules
    )


# --------------------------------------------------------------------------- #
# Environment, phenotype, record
# --------------------------------------------------------------------------- #
def build_environment() -> Environment:
    """The one environment every loaded sample shares.

    The medium is Neidhardt's MOPS minimal without a carbon source, whose stated 9.5 mM
    NH4Cl is the released nitrogen source (``check_sample_metadata`` asserts that), so
    the only edit of the medium is the released carbon source.
    """
    perturbations: list[EnvironmentPerturbationType] = [
        EnvironmentPhysicalPerturbation(
            factor=PhysicalFactor.carbon_source,
            magnitude=Concentration(
                value=LOADED_CARBON_PERCENT, unit=ConcentrationUnit.percent_w_v
            ),
            agent=resolved_compound(LOADED_CARBON_COMPOUND),
        )
    ]
    return Environment(
        media=MOPS_MINIMAL,
        temperature=Temperature(
            value=float(str(SOURCED_VALUES["temperature_c"].value))
        ),
        perturbations=perturbations,
        aerobicity=str(SOURCED_VALUES["aerobicity"].value),
    )


def build_phenotype(
    rows: Sequence[MassFractionRow], locus_tag: Mapping[str, str], column: str
) -> ProteinAbundancePhenotype:
    """One released sample column, verbatim, over the rows it quantifies.

    One LC-MS injection has no replicate, so ``n_replicates`` is 1 for every protein and
    there is no standard error: the spread of the culture's injections is a property of
    the culture, not of this injection.
    """
    abundance = {
        locus_tag[row.gene_name]: row.mass_fraction[column]
        for row in rows
        if row.quantified(column)
    }
    return ProteinAbundancePhenotype(
        protein_abundance=abundance,
        protein_abundance_se=None,
        n_replicates=dict.fromkeys(abundance, 1),
        measurement_type=MEASUREMENT_TYPE,
    )


def build_reference_phenotype(
    rows: Sequence[MassFractionRow], locus_tag: Mapping[str, str]
) -> ProteinAbundancePhenotype:
    """The calibration sample's profile: the mean of its three culture means.

    A culture's mean is over its own quantified injections; a protein no injection of a
    culture quantifies contributes nothing to that protein's mean. ``n_replicates`` is
    how many of the three cultures quantified the protein and the SE is
    ``stdev(culture means) / sqrt(n)``, exact arithmetic on the released columns.
    """
    abundance: dict[str, float] = {}
    standard_error: dict[str, float] = {}
    n_replicates: dict[str, int] = {}
    for row in rows:
        means = [
            statistics.fmean(
                [row.mass_fraction[c] for c in columns if row.quantified(c)]
            )
            for columns in CULTURES.values()
            if any(row.quantified(c) for c in columns)
        ]
        if not means:
            continue
        tag = locus_tag[row.gene_name]
        abundance[tag] = statistics.fmean(means)
        n_replicates[tag] = len(means)
        standard_error[tag] = (
            statistics.stdev(means) / math.sqrt(len(means))
            if len(means) > 1
            else float("nan")
        )
    return ProteinAbundancePhenotype(
        protein_abundance=abundance,
        protein_abundance_se=standard_error,
        n_replicates=n_replicates,
        measurement_type=MEASUREMENT_TYPE,
    )


def restrict(
    phenotype: ProteinAbundancePhenotype, keys: Sequence[str]
) -> ProteinAbundancePhenotype:
    """The same profile over ``keys`` only, for a key-matched reference."""
    wanted = set(keys)
    missing = wanted - set(phenotype.protein_abundance)
    if missing:
        raise RuntimeError(
            f"the calibration reference quantifies none of {sorted(missing)[:5]}"
        )
    return ProteinAbundancePhenotype(
        protein_abundance={
            k: v for k, v in phenotype.protein_abundance.items() if k in wanted
        },
        protein_abundance_se={
            k: v
            for k, v in (phenotype.protein_abundance_se or {}).items()
            if k in wanted
        },
        n_replicates={k: v for k, v in phenotype.n_replicates.items() if k in wanted},
        measurement_type=phenotype.measurement_type,
    )


def publication() -> Publication:
    """The paper, as the PMC id converter resolves it from ``PMC8144880``."""
    return Publication(
        pubmed_id=PUBMED_ID,
        pubmed_url=f"https://pubmed.ncbi.nlm.nih.gov/{PUBMED_ID}/",
        doi=PAPER_DOI,
        doi_url=f"https://doi.org/{PAPER_DOI}",
    )


# --------------------------------------------------------------------------- #
# The raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the environment (both mirrors and the build tree live there)."""
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/moriCoarseFineAbsolute2021``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def library_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-library/moriCoarseFineAbsolute2021``."""
    return Path(data_root or _data_root()) / LIBRARY_DIR_REL


def _sha256(path: str | Path) -> str:
    """Streaming sha256 of a file."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def retrieve_raw_files(dest_dir: str | Path) -> dict[str, Path]:
    """Re-run the recorded retrievals and write the verified bytes into ``dest_dir``.

    Each ``RetrievalRecord`` is what runs, so this IS the recorded retrieval rather than
    a description of it; a byte mismatch raises before anything is written.
    """
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}
    for name, record in RETRIEVALS.items():
        path = dest / name
        write_verified(
            run_retriever(record), path, DATA_SHA256[name], str(record.source_url)
        )
        out[name] = path
    return out


def deposit_raw_mirror(*, source_dir: str | Path, data_root: str | None = None) -> Path:
    """Write the raw mirror from already-retrieved files plus their manifest.

    Idempotent by sha256: a mirror file that already hashes to its pin is left alone,
    and one with any other hash raises rather than being overwritten.
    """
    src_dir = Path(source_dir)
    root = raw_mirror_dir(data_root)
    records: list[ArtifactRecord] = []
    for name, relpath in MIRROR_RELPATH.items():
        src = src_dir / name
        observed = _sha256(src)
        if observed != DATA_SHA256[name]:
            raise RuntimeError(
                f"{src} sha256 mismatch: got {observed}, expected {DATA_SHA256[name]}"
            )
        dest = root / relpath
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            if _sha256(dest) != DATA_SHA256[name]:
                raise RuntimeError(f"{dest} exists with a different sha256; refusing")
        else:
            shutil.copy2(src, dest)
        records.append(
            ArtifactRecord(
                path=relpath,
                role=ROLE_RAW_DATA,
                bytes=dest.stat().st_size,
                sha256=DATA_SHA256[name],
                source=str(RETRIEVALS[name].source_url),
                retrieval=RETRIEVALS[name],
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title=PAPER_TITLE,
        files=records,
        si_data_sources=[
            *(str(record.source_url) for record in RETRIEVALS.values()),
            f"https://pmc.ncbi.nlm.nih.gov/articles/{PMC_ID}/",
            "https://www.swathatlas.org (PASS01421)",
            "https://panoramaweb.org/public (the AQUA-peptide Skyline documents)",
        ],
        si_expected=list(NOT_MIRRORED),
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class ProteomeMori2021Dataset(ExperimentDataset):
    """Absolute MG1655 proteome mass fractions, one record per loaded sample."""

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = MORI_REFERENCE_STRAIN

    def __init__(
        self,
        root: str = "data/torchcell/proteome_mori2021",
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialProteinAbundanceExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialProteinAbundanceExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The six consumed files, linked from the raw mirror."""
        return list(CONSUMED)

    def download(self) -> None:
        """Link the mirrored files into ``raw/`` after checking manifest and sha256."""
        data_root = _data_root()
        manifest = load_manifest(data_root)
        os.makedirs(self.raw_dir, exist_ok=True)
        for name, relpath in MIRROR_RELPATH.items():
            check_manifest_pin(
                relpath, manifest_sha256(manifest, relpath), DATA_SHA256[name]
            )
            src = raw_mirror_dir(data_root) / relpath
            if not src.exists():
                raise RuntimeError(f"required raw artifact missing from mirror: {src}")
            link_verified(src, osp.join(self.raw_dir, name), DATA_SHA256[name])
        log.info("Mori 2021 raw files linked into %s (sha256 verified)", self.raw_dir)

    def compute_gene_set(self) -> GeneSet:
        """The MG1655 loci the stored abundance profiles are keyed by.

        Every record is wild type, so no genotype names a gene and the base class's
        genotype scan would return the empty set it refuses. The dataset's genes are the
        loci it measures, which is how the landed Caglar 2017 and Schmidt 2016 loaders
        state the same situation.
        """
        if self.env is None:
            self._init_db()
        genes = GeneSet()
        with self.env.begin() as txn:
            for _, value in txn.cursor():
                record = pickle.loads(value)
                genes.update(record["experiment"]["phenotype"]["protein_abundance"])
        self.close_lmdb()
        return genes

    def _genome(self) -> EcoliK12Genome:
        """The injected genome, or the reference strain's default cache."""
        if self.ecoli_genome is None:
            self.ecoli_genome = bacterial_genome("ecoli", self.REFERENCE_STRAIN)
        expected = BACTERIAL_ASSEMBLY_SETS[self.REFERENCE_STRAIN]
        if self.ecoli_genome.ASSEMBLY_SET != expected:
            raise ValueError(
                f"{type(self).__name__} needs the {expected} genome, got "
                f"{self.ecoli_genome.ASSEMBLY_SET}"
            )
        return self.ecoli_genome

    @post_process
    def process(self) -> None:
        """Build one abundance record per loaded sample and write the LMDB."""
        verify_raw_files(self.raw_dir, DATA_SHA256)
        paths = {name: osp.join(self.raw_dir, name) for name in CONSUMED}
        check_quotes(paths)
        check_sample_columns(paths[EV8], "EV8")
        check_sample_columns(paths[EV9], "EV9")
        metadata_check = check_sample_metadata(paths[EV2], paths[EV3])
        normalization = [
            check_normalization(paths[EV8], "EV8"),
            check_normalization(paths[EV9], "EV9"),
        ]
        identity_fault = check_identity_fault(paths[EV8], paths[EV9])

        columns = [spec.column for spec in LOADED]
        rows = read_mass_fractions(paths[EV9], columns)
        if len(rows) != EXPECTED_SOURCE_ROWS:
            raise RuntimeError(
                f"{len(rows)} released rows, the module states {EXPECTED_SOURCE_ROWS}"
            )
        genome = self._genome()
        selection = select_proteins(rows, genome, label=f"{self.name} b-numbers")
        if len(selection.kept) != EXPECTED_PROTEIN_KEYS:
            raise RuntimeError(
                f"{len(selection.kept)} protein keys, the module states "
                f"{EXPECTED_PROTEIN_KEYS}"
            )

        environment = build_environment()
        phenotypes = {
            spec.column: build_phenotype(
                selection.kept, selection.locus_tag, spec.column
            )
            for spec in LOADED
        }
        reference_phenotype = build_reference_phenotype(
            selection.kept, selection.locus_tag
        )
        reference_genome = assembly_reference(
            self.REFERENCE_STRAIN, background=EQ353_BACKGROUND
        )
        pub = publication()
        if len(LOADED) != EXPECTED_RECORDS:
            raise RuntimeError(
                f"{len(LOADED)} records, the module states {EXPECTED_RECORDS}"
            )
        growth_rates = read_growth_rates(paths[EV3])

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        sample_rows: list[dict[str, Any]] = []
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for idx, spec in enumerate(tqdm(LOADED, desc="mori2021")):
                phenotype = phenotypes[spec.column]
                experiment = BacterialProteinAbundanceExperiment(
                    dataset_name=self.name,
                    genotype=Genotype(perturbations=[]),
                    environment=environment,
                    phenotype=phenotype,
                )
                reference = BacterialProteinAbundanceExperimentReference(
                    dataset_name=self.name,
                    genome_reference=reference_genome,
                    environment_reference=environment,
                    phenotype_reference=restrict(
                        reference_phenotype, list(phenotype.protein_abundance)
                    ),
                )
                txn.put(
                    f"{idx}".encode(),
                    self._intern_record(experiment, reference, pub, itxn),
                )
                sample_rows.append(
                    {
                        "column": spec.column,
                        "table": spec.table,
                        "culture": spec.culture,
                        "strain": spec.strain,
                        "medium": spec.medium,
                        "library_media": MOPS_MINIMAL.name,
                        "carbon_source": LOADED_CARBON_SOURCE,
                        "nitrogen_source": LOADED_NITROGEN_SOURCE,
                        "temperature_c": SOURCED_VALUES["temperature_c"].value,
                        "growth_rate_per_h": growth_rates.get(spec.column),
                        "n_protein_keys": len(phenotype.protein_abundance),
                    }
                )
        env.close()
        interned_env.close()

        self._write_ledgers(
            rows,
            selection,
            reference_phenotype,
            sample_rows,
            metadata_check,
            normalization,
            identity_fault,
        )
        log.info(
            "Mori 2021: %d records (of %d released samples) x %d-%d protein keys from "
            "%d released rows; %d rows dropped",
            len(LOADED),
            len(SAMPLES),
            min(row["n_protein_keys"] for row in sample_rows),
            max(row["n_protein_keys"] for row in sample_rows),
            len(rows),
            selection.dropped_rows,
        )

    def _write_ledgers(
        self,
        rows: Sequence[MassFractionRow],
        selection: ProteinSelection,
        reference_phenotype: ProteinAbundancePhenotype,
        sample_rows: Sequence[Mapping[str, Any]],
        metadata_check: Mapping[str, Any],
        normalization: Sequence[Mapping[str, Any]],
        identity_fault: Mapping[str, Any],
    ) -> None:
        """The drop log, the sourcing table, the identifier table, the sample table."""
        out = Path(self.preprocess_dir)
        sample_rules = [
            DropRule(
                rule=reason.rule,
                scope="sample",
                description=reason.description,
                n_items=len([s for s in SAMPLES if s.drop == reason]),
                items=[
                    f"{s.column} ({s.strain}, {s.medium})"
                    for s in SAMPLES
                    if s.drop == reason
                ],
                needed_addition=reason.needed_addition,
            )
            for reason in (
                DROP_NO_MEDIA_ENTRY,
                DROP_MEDIA_ENTRY_DIFFERS,
                DROP_FORMULATION_NOT_STATED,
            )
        ]
        log_model = DropLog(
            dataset=self.name,
            source_samples=len(SAMPLES),
            kept_records=len(sample_rows),
            dropped_records=len(SAMPLES) - len(sample_rows),
            source_protein_rows=len(rows),
            kept_protein_keys=len(selection.kept),
            dropped_protein_rows=selection.dropped_rows,
            rules=[*sample_rules, *selection.rules],
            reconciliation=selection.reconciliation,
            notes=[
                f"{len(SAMPLES)} released sample columns = 30 (Dataset EV8) + 36 "
                "(Dataset EV9), which is the paper's own 66 samples",
                f"{len(SAMPLES)} released samples - "
                + " - ".join(f"{r.n_items} ({r.rule})" for r in sample_rules)
                + f" = {len(sample_rows)} records",
                f"{len(rows)} released rows - "
                + " - ".join(f"{r.n_items} ({r.rule})" for r in selection.rules)
                + f" = {len(selection.kept)} protein keys",
                "per-record keys are that column's own quantified subset: "
                + ", ".join(
                    f"{row['column']} {row['n_protein_keys']}" for row in sample_rows
                ),
                f"the calibration reference quantifies "
                f"{len(reference_phenotype.protein_abundance)} of those keys over "
                f"{len(CULTURES)} biological cultures "
                f"({', '.join(f'{k}={len(v)} injection(s)' for k, v in CULTURES.items())})",
                "every loaded record carries one environment and one empty wild-type "
                "genotype: the seven samples are injections of three cultures of ONE "
                "condition, so they are seven measurements rather than seven conditions",
                "Environment has no growth-rate, growth-phase or culture-mode slot, so "
                "the released growth rate is in samples.csv and not on a record",
            ],
        )
        log_model.check()
        (out / "dropped_records.json").write_text(log_model.model_dump_json(indent=2))
        (out / "released_statistics_check.json").write_text(
            json.dumps(
                {
                    "sample_metadata": dict(metadata_check),
                    "normalization": [dict(row) for row in normalization],
                    "ev9_identity_fault": dict(identity_fault),
                    "normalization_atol": NORMALIZATION_ATOL,
                },
                indent=2,
            )
        )
        (out / "sourced_values.json").write_text(
            json.dumps(
                {
                    name: value.model_dump(mode="json")
                    for name, value in SOURCED_VALUES.items()
                },
                indent=2,
            )
        )
        pd.DataFrame(list(sample_rows)).to_csv(out / "samples.csv", index=False)
        pd.DataFrame(
            [
                {
                    "released_gene_name": row.gene_name,
                    "released_gene_locus": row.gene_locus,
                    "released_protein_id": row.protein_id,
                    "stored_locus_tag": selection.locus_tag[row.gene_name],
                    "n_quantified_samples": sum(
                        1 for spec in LOADED if row.quantified(spec.column)
                    ),
                }
                for row in selection.kept
            ]
        ).to_csv(out / "protein_identifiers.csv", index=False)

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process() for this dataset."""
        raise NotImplementedError


def check_quotes(paths: Mapping[str, str]) -> dict[str, int]:
    """Every workbook and Appendix quote is a substring of the pinned bytes.

    ``paper.md`` lives in the literature mirror rather than in ``raw/``, so its quotes
    are checked by the test module against the pinned library artifact; the quotes this
    build can check against its own raw files are checked here.
    """
    text = appendix_text(paths[APPENDIX])
    for name, quote in APPENDIX_QUOTES.items():
        if quote not in text:
            raise RuntimeError(f"Appendix quote {name!r} is not in the pinned bytes")
    checked = len(APPENDIX_QUOTES)
    for file_name, quotes in WORKBOOK_QUOTES.items():
        book = openpyxl.load_workbook(paths[file_name], read_only=True, data_only=True)
        try:
            rendered = {
                _render(row)
                for sheet in book.worksheets
                for row in sheet.iter_rows(values_only=True)
            }
        finally:
            book.close()
        joined = "\n".join(rendered)
        for name, quote in quotes.items():
            if quote not in joined:
                raise RuntimeError(
                    f"{file_name} quote {name!r} is not a row rendering of the pinned "
                    "bytes"
                )
            checked += 1
    return {"n_quotes_checked": checked}


# --------------------------------------------------------------------------- #
# L0-L4 verification
# --------------------------------------------------------------------------- #
def assembly_pin_rule(records: Sequence[Mapping[str, Any]]) -> Any:
    """SUPPLEMENTARY L3: every record pins the MG1655 GenBank assembly."""
    from torchcell.verification.report import Level, LevelResult

    pins = {
        (
            record["reference"]["genome_reference"].get("assembly_set"),
            record["reference"]["genome_reference"].get("assembly_accession"),
        )
        for record in records
    }
    expected = {(BACTERIAL_ASSEMBLY_SETS[MORI_REFERENCE_STRAIN], "GCA_000005845.2")}
    return LevelResult(
        level=Level.L3,
        name="assembly_pin",
        passed=pins == expected,
        message=f"SUPPLEMENTARY: assembly pins {sorted(map(str, pins))}",
        details={"pins": sorted(map(str, pins))},
    )


def strain_label_rule(records: Sequence[Mapping[str, Any]]) -> Any:
    """SUPPLEMENTARY L1: every record names the EQ353 stock as its strain."""
    from torchcell.verification.report import Level, LevelResult

    labels = {
        str(record["reference"]["genome_reference"].get("strain")) for record in records
    }
    expected = {EQ353_BACKGROUND.name}
    return LevelResult(
        level=Level.L1,
        name="strain_label",
        passed=labels == expected,
        message=f"SUPPLEMENTARY: strain labels {sorted(labels)}",
        details={"labels": sorted(labels)},
    )


def single_injection_rule(records: Sequence[Mapping[str, Any]]) -> Any:
    """SUPPLEMENTARY L2: a record is one injection, so it has no replicate and no SE."""
    from torchcell.verification.report import Level, LevelResult

    bad = 0
    n = 0
    for record in records:
        phenotype = record["experiment"]["phenotype"]
        if phenotype.get("protein_abundance_se") is not None:
            bad += 1
        for value in phenotype["n_replicates"].values():
            n += 1
            if int(value) != 1:
                bad += 1
    return LevelResult(
        level=Level.L2,
        name="one_injection_per_record",
        passed=bad == 0,
        message=(
            f"SUPPLEMENTARY: {bad} of {n} stored values disagree with one injection per "
            "record"
        ),
        details={"n_values": n, "n_bad": bad},
    )


def reference_replicate_rule(records: Sequence[Mapping[str, Any]]) -> Any:
    """SUPPLEMENTARY L3: the reference has a finite SE exactly where n > 1."""
    from torchcell.verification.report import Level, LevelResult

    bad = 0
    counts: dict[int, int] = {}
    for record in records:
        reference = record["reference"]["phenotype_reference"]
        errors = reference.get("protein_abundance_se") or {}
        for key, n in reference["n_replicates"].items():
            counts[int(n)] = counts.get(int(n), 0) + 1
            finite = not math.isnan(float(errors[key]))
            if finite != (int(n) > 1):
                bad += 1
    return LevelResult(
        level=Level.L3,
        name="reference_se_matches_culture_count",
        passed=bad == 0,
        message=(
            f"SUPPLEMENTARY: {bad} of {sum(counts.values())} reference values disagree "
            "with their culture count"
        ),
        details={"n_bad": bad, "n_by_cultures": counts},
    )


def gene_containment_rule(
    records: Sequence[Mapping[str, Any]], universe: set[str]
) -> Any:
    """L4: every stored protein key is a locus of the pinned MG1655 assembly."""
    from torchcell.verification.report import Level, LevelResult

    measured: set[str] = set()
    for record in records:
        measured.update(record["experiment"]["phenotype"]["protein_abundance"])
        measured.update(record["reference"]["phenotype_reference"]["protein_abundance"])
    outside = sorted(measured - universe)
    return LevelResult(
        level=Level.L4,
        name="gene_containment_mg1655",
        passed=bool(measured) and not outside,
        message=(
            f"{len(measured)} measured protein keys; {len(outside)} outside the "
            f"{BACTERIAL_ASSEMBLY_SETS[MORI_REFERENCE_STRAIN]} locus universe"
        ),
        details={"outside": outside[:20], "n_universe": len(universe)},
    )


def verify_build(
    dataset_root: str,
    data_root: str | None = None,
    *,
    genome: EcoliK12Genome | None = None,
    expected_count: int = EXPECTED_RECORDS,
) -> VerificationReport:
    """Run this module's L0-L4 gate over a built tree and write the report.

    The shared ``verify_protein_dataset`` supplies L0 to L3 (structure, count, value
    fidelity, a key-matched finite reference, one measurement type); four SUPPLEMENTARY
    rows and the host-aware L4 containment are added here, because the yeast
    deletion-collection overlap that ``run_protein`` adds would say nothing about an
    MG1655 b-number. The report is written to
    ``preprocess/verification_report.json``.
    """
    from torchcell.verification.protein import verify_protein_dataset
    from torchcell.verification.runners import load_records

    records = load_records(dataset_root)
    if genome is None:
        genome = bacterial_genome("ecoli", MORI_REFERENCE_STRAIN, data_root)
    report = verify_protein_dataset(
        records,
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=Provenance(
            source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{MIRROR_RELPATH[EV9]}",
            citation_key=CITATION_KEY,
            sha256=DATA_SHA256[EV9],
            method=(
                "Dataset EV9 'EV9-AbsoluteMassFractions-2': one "
                "BacterialProteinAbundanceExperiment per loaded MG1655 (EQ353) "
                "calibration sample, storing that column's absolute protein mass "
                "fraction verbatim, with the reference the mean of the three culture "
                "means and SE = stdev(culture means) / sqrt(n)"
            ),
            page="si10.xlsx sheet 'EV9-AbsoluteMassFractions-2'",
            retrieved=RETRIEVED_AT,
        ),
        expected_count=expected_count,
    )
    report.add(strain_label_rule(records))
    report.add(single_injection_rule(records))
    report.add(assembly_pin_rule(records))
    report.add(reference_replicate_rule(records))
    report.add(gene_containment_rule(records, set(genome.genbank.loci)))
    out = osp.join(dataset_root, "preprocess", "verification_report.json")
    os.makedirs(osp.dirname(out), exist_ok=True)
    with open(out, "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main() -> None:
    """Deposit the raw mirror, or build and verify the dataset."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--deposit",
        action="store_true",
        help="re-run the recorded retrievals and write the raw mirror, then exit",
    )
    parser.add_argument(
        "--root",
        default="data/torchcell/proteome_mori2021",
        help="build tree, relative to DATA_ROOT",
    )
    args = parser.parse_args()

    from dotenv import load_dotenv

    load_dotenv()
    data_root = _data_root()
    if args.deposit:
        staging = raw_mirror_dir(data_root) / "_staging"
        retrieve_raw_files(staging)
        root = deposit_raw_mirror(source_dir=staging, data_root=data_root)
        shutil.rmtree(staging)
        print(f"deposited {len(CONSUMED)} files into {root}")
        return

    build_root = osp.join(data_root, args.root)
    dataset = ProteomeMori2021Dataset(root=build_root)
    print(f"len = {len(dataset)}")
    ledger = json.loads(
        Path(build_root, "preprocess", "dropped_records.json").read_text()
    )
    print(
        json.dumps(
            {
                k: ledger[k]
                for k in (
                    "source_samples",
                    "kept_records",
                    "dropped_records",
                    "source_protein_rows",
                    "kept_protein_keys",
                    "dropped_protein_rows",
                    "notes",
                )
            },
            indent=2,
        )
    )
    print(verify_build(build_root, data_root).summary())


if __name__ == "__main__":
    main()
