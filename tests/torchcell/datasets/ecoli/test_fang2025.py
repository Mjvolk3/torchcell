# tests/torchcell/datasets/ecoli/test_fang2025.py
# [[tests.torchcell.datasets.ecoli.test_fang2025]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_fang2025.py
"""The Fang 2025 CRISPRi-FACS loader (``torchcell.datasets.ecoli.fang2025``).

Synthetic tests (run everywhere) build every input in ``tmp_path``. The build uses the
real ``EcoliK12MG1655Genome`` over the synthetic MG1655 assembly of
``tests/torchcell/sequence/genome/_bacterial_fixtures.py`` (thrL b0001, thrA b0002,
thrW b0003 tRNA, pseudogene yaaP b0004, proB b0005, proC b0006), served through a
stubbed ``resolve`` with the network refused, and replaces the loader's
``verify_raw_files`` with a presence check (the synthetic workbooks cannot carry the
real pins; those are asserted by the refusal test and the data-gated tests).

The synthetic release is read-count driven, because the ONE retention rule is a read
floor: every expected count below falls out of the counts in ``ROUND_1_READS`` and
``ROUND_2_READS`` rather than being asserted separately.

    guide            cluster members   round 1 (tra/bs/as)   round 2 (bs/as)
    thrLb0001_10     b0001             100/100/50  keep      100/100  keep
    thrLb0001_25     b0001             100/100/200 keep      100/10   floor
    thrAb0002_12     b0002             100/100/5   floor     100/100  keep
    proBb0005_60     b0005, b0006      100/100/100 keep      100/100  keep
    yaaPb0004_45     b0004             100/100/40  keep      100/5    floor
    ybfKb4590_5      b4590 (retired)   100/100/60  retired   100/100  retired
    thrWb0003_31     b0003             10/100/100  floor     10/100   floor

Round 1: 5 rows clear the floor, 1 of them is the retired singleton, so 4 records over
5 perturbations (``proBb0005_60`` carries two, one per cluster member). Round 2: 4 rows
clear the floor, 1 retired, so 3 records, each carrying ONE EXTRA knockdown for the
synthetic pcnBi background (``yaaPb0004_45`` on b0004), giving 7 perturbations. Seven
records, 12 perturbations, 5 distinct genes. ``yaaPb0004_45`` is the background guide
AND is below the round-2 floor on purpose, so no record repeats a gene.

Both Methods equations are written into the synthetic workbook by computing them from
the synthetic read counts, which is what lets the loader's own equation check and its
floor cross-tabulation be exercised rather than bypassed.

Data-gated tests (``@pytest.mark.data``) read the real raw mirror and the built dev-tree
LMDB under ``$DATA_ROOT`` (they never build it): the manifest pins, the provenance audit
of every sourced value, the per-round arithmetic, and the L0-L4 verifier.
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.ecoli.fang2025 as m
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    MG1655_GAF,
    MG1655_LOCI,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.datamodels.media import M9_MODIFIED_FANG2025
from torchcell.datamodels.schema import (
    AssayType,
    AssemblyReferenceGenome,
    BacterialCrisprInterferencePerturbation,
    BacterialEnvironmentResponseExperiment,
    BacterialEnvironmentResponseExperimentReference,
    ConcentrationUnit,
    MeasurementType,
    SampleUnit,
    SmallMoleculePerturbation,
)
from torchcell.datasets.ecoli.wang2018 import CLUSTER_COLUMNS, LIBRARY_COLUMNS
from torchcell.literature.manifest import (
    ROLE_SI_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
)
from torchcell.sequence.genome.ecoli.k12 import MG1655_ASSEMBLY, EcoliK12MG1655Genome
from torchcell.verification.sourced import SourcedValue, audit_sourced_value

# --------------------------------------------------------------------------- #
# The synthetic release
# --------------------------------------------------------------------------- #
PROTEIN_CLUSTERS: list[tuple[str, str]] = [
    ("thrLb0001", "thrLb0001"),
    ("thrAb0002", "thrAb0002"),
    ("proBb0005", "proBb0005,proCb0006"),
    ("yaaPb0004", "yaaPb0004"),
    ("ybfKb4590", "ybfKb4590"),
]
NCRNA_CLUSTERS: list[tuple[str, str]] = [("thrWb0003", "thrWb0003")]
N_SYNTHETIC_CLUSTERS = len(PROTEIN_CLUSTERS) + len(NCRNA_CLUSTERS)
N_SYNTHETIC_MEMBERS = 7

GUIDES: dict[str, str] = {
    "thrLb0001_10": "AAAACGCGCGTCACGCGTCC",
    "thrLb0001_25": "AAAGCGCGCGTCACGCGTCC",
    "thrAb0002_12": "AACGCGCGCGTCACGCGTCC",
    "proBb0005_60": "AAGGCGCGCGTCACGCGTCC",
    "yaaPb0004_45": "AATGCGCGCGTCACGCGTCC",
    "ybfKb4590_5": "ACAGCGCGCGTCACGCGTCC",
    "thrWb0003_31": "ACCGCGCGCGTCACGCGTCC",
    "NC_1": "ACGGCGCGCGTCACGCGTCC",
    "NC_2": "ACTGCGCGCGTCACGCGTCC",
}
N_SYNTHETIC_TARGETING = 7
N_SYNTHETIC_CONTROLS = 2

#: ``guide -> (transformation, before sorting, after sorting)``.
ROUND_1_READS: dict[str, tuple[int, int, int]] = {
    "thrLb0001_10": (100, 100, 50),
    "thrLb0001_25": (100, 100, 200),
    "thrAb0002_12": (100, 100, 5),
    "proBb0005_60": (100, 100, 100),
    "yaaPb0004_45": (100, 100, 40),
    "ybfKb4590_5": (100, 100, 60),
    "thrWb0003_31": (10, 100, 100),
}
#: ``guide -> (before sorting, after sorting)``.
ROUND_2_READS: dict[str, tuple[int, int]] = {
    "thrLb0001_10": (100, 100),
    "thrLb0001_25": (100, 10),
    "thrAb0002_12": (100, 100),
    "proBb0005_60": (100, 100),
    "yaaPb0004_45": (100, 5),
    "ybfKb4590_5": (100, 100),
    "thrWb0003_31": (10, 100),
}
#: The synthetic pcnBi background: a single-gene guide below the round-2 floor.
SYNTHETIC_BACKGROUND_GUIDE = "yaaPb0004_45"

SCREEN_RAW = m.RawFile(
    moesm=8, data_number=6, description="synthetic source data", sha256="0" * 64
)
CLUSTERS_RAW = m.RawFile(
    moesm=5, data_number=2, description="synthetic clusters", sha256="1" * 64
)
LIBRARY_RAW = m.RawFile(
    moesm=6, data_number=3, description="synthetic library", sha256="2" * 64
)

ROUND_1 = m.Round(
    round_id="round_1_cf",
    sheet="Figure 1",
    panel="Figure 1d",
    host_strain="CF",
    libraries=("transformation", "before_sorting", "after_sorting"),
    id_column=10,
    read_columns=(11, 12, 13),
    fitness_column=16,
    excluded_columns=(18, 19, 20),
    fitness_all_column=21,
    carries_background_knockdown=False,
    floor_kept_rows=5,
    kept_records=4,
)
ROUND_2 = m.Round(
    round_id="round_2_pcnbi",
    sheet="Figure 5",
    panel="Figure 5c",
    host_strain="pcnBi",
    libraries=("before_sorting", "after_sorting"),
    id_column=10,
    read_columns=(11, 12),
    fitness_column=15,
    excluded_columns=(17, 18),
    fitness_all_column=19,
    carries_background_knockdown=True,
    floor_kept_rows=4,
    kept_records=3,
)
SYNTHETIC_ROUNDS = (ROUND_1, ROUND_2)
N_SYNTHETIC_RECORDS = ROUND_1.kept_records + ROUND_2.kept_records
N_SYNTHETIC_PERTURBATIONS = 12
N_SYNTHETIC_GENES = 5

REFERENCE = AssemblyReferenceGenome(
    species="Escherichia coli",
    strain="MG1655",
    assembly_set="ecoli_K12_MG1655_ASM584v2",
    assembly_accession="GCA_000005845.2",
)


def _sha256_of(path: Path) -> str:
    """Hex sha256 of a synthetic workbook, so its RawFile pin is its real bytes."""
    import hashlib

    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_sheets(
    path: Path, sheets: Sequence[tuple[str, Sequence[Sequence[Any]]]]
) -> None:
    """Write one workbook, one sheet per ``(name, rows)`` pair."""
    workbook = openpyxl.Workbook()
    default = workbook.active
    assert default is not None
    workbook.remove(default)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    workbook.save(path)


def write_clusters(path: Path) -> None:
    """A synthetic Supplementary Data 2 (Wang 2018's) with its two sheets."""
    _write_sheets(
        path,
        [
            ("protein-coding genes", [CLUSTER_COLUMNS, *PROTEIN_CLUSTERS]),
            ("ncRNA-coding genes", [CLUSTER_COLUMNS, *NCRNA_CLUSTERS]),
        ],
    )


def write_library(path: Path) -> None:
    """A synthetic Supplementary Data 3 (Wang 2018's)."""
    _write_sheets(path, [("sheet1", [LIBRARY_COLUMNS, *GUIDES.items()])])


def _round_rows(
    reads: Mapping[str, tuple[int, ...]], round_: m.Round, floor: int = 20
) -> list[list[Any]]:
    """One round's sheet body, with both Methods equations computed from the reads.

    The ``Fitness`` column is populated exactly where every library clears ``floor``,
    and the value otherwise goes into the per-library excluded column, which is the
    split the released workbook carries and the loader cross-tabulates against.
    """
    width = round_.fitness_all_column + 1
    totals = [
        sum(values[position] for values in reads.values())
        for position in range(len(round_.libraries))
    ]
    before = round_.libraries.index("before_sorting")
    after = round_.libraries.index("after_sorting")
    body: list[list[Any]] = []
    for guide, values in reads.items():
        row: list[Any] = [None] * width
        row[round_.id_column] = guide
        for position, column in enumerate(round_.read_columns):
            row[column] = values[position]
        normalized_before = (values[before] + 1) / totals[before]
        normalized_after = (values[after] + 1) / totals[after]
        row[round_.fitness_column - 2] = normalized_before
        row[round_.fitness_column - 1] = normalized_after
        fitness = math.log2(normalized_after / normalized_before)
        row[round_.fitness_all_column] = fitness
        if all(count >= floor for count in values):
            row[round_.fitness_column] = fitness
        else:
            low = next(
                position for position, count in enumerate(values) if count < floor
            )
            row[round_.excluded_columns[low]] = fitness
        body.append(row)
    return body


def write_screen(path: Path) -> None:
    """A synthetic Supplementary Data 6: two panel rows, a header row, then the guides."""
    sheets: list[tuple[str, Sequence[Sequence[Any]]]] = []
    for round_, reads in ((ROUND_1, ROUND_1_READS), (ROUND_2, ROUND_2_READS)):
        width = round_.fitness_all_column + 1
        panels: list[Any] = [None] * width
        panels[round_.id_column - 9] = round_.panel
        header: list[Any] = [None] * width
        header[round_.id_column] = "sgRNA-ID"
        sheets.append((round_.sheet, [panels, header, *_round_rows(reads, round_)]))
    _write_sheets(path, sheets)


def write_raw(raw: Path) -> None:
    """Every synthetic file the loader reads, under ``raw``."""
    raw.mkdir(parents=True, exist_ok=True)
    write_clusters(raw / CLUSTERS_RAW.filename)
    write_library(raw / LIBRARY_RAW.filename)
    write_screen(raw / SCREEN_RAW.filename)


@pytest.fixture
def synthetic_library_size(monkeypatch: pytest.MonkeyPatch) -> None:
    """Point the released-size check at the synthetic library's seven guides."""
    monkeypatch.setattr(m, "N_TARGETING_GUIDES", N_SYNTHETIC_TARGETING)


@pytest.fixture
def mg1655(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12MG1655Genome:
    """The real MG1655 class over the synthetic assembly; network refused."""
    files = write_assembly(tmp_path / "tier", MG1655_ASSEMBLY, MG1655_LOCI, MG1655_GAF)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    root = tmp_path / "mg1655"
    root.mkdir()
    return EcoliK12MG1655Genome(genome_root=str(root), overwrite=False)


@pytest.fixture
def presence_only_pins(monkeypatch: pytest.MonkeyPatch) -> list[Mapping[str, str]]:
    """Replace the build-time byte check with a presence check that records the pins."""
    calls: list[Mapping[str, str]] = []

    def record(raw_dir: str, pins: Mapping[str, str]) -> None:
        missing = [f for f in pins if not osp.exists(osp.join(raw_dir, f))]
        if missing:
            raise FileNotFoundError(f"pinned raw files absent: {missing}")
        calls.append(dict(pins))

    monkeypatch.setattr(m, "verify_raw_files", record)
    return calls


@pytest.fixture
def synthetic(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    mg1655: EcoliK12MG1655Genome,
    presence_only_pins: list[Mapping[str, str]],
) -> Path:
    """A dataset root whose ``raw/`` holds the synthetic release."""
    root = tmp_path / "data/torchcell/crispri_guide_ffa_enrichment_fang2025"
    write_raw(root / "raw")
    monkeypatch.setattr(m, "ROUNDS", SYNTHETIC_ROUNDS)
    monkeypatch.setattr(m, "RAW_FILES", (SCREEN_RAW,))
    monkeypatch.setattr(m, "SCREEN_FILE", SCREEN_RAW)
    monkeypatch.setattr(m, "WANG_RAW_FILES", (CLUSTERS_RAW, LIBRARY_RAW))
    monkeypatch.setattr(m, "WANG_CLUSTERS_FILE", CLUSTERS_RAW)
    monkeypatch.setattr(m, "WANG_LIBRARY_FILE", LIBRARY_RAW)
    monkeypatch.setattr(m, "N_CLUSTERS", N_SYNTHETIC_CLUSTERS)
    monkeypatch.setattr(m, "N_CLUSTER_MEMBERS", N_SYNTHETIC_MEMBERS)
    monkeypatch.setattr(m, "N_TARGETING_GUIDES", N_SYNTHETIC_TARGETING)
    monkeypatch.setattr(m, "N_CONTROL_GUIDES", N_SYNTHETIC_CONTROLS)
    monkeypatch.setattr(m, "EXPECTED_RECORDS", N_SYNTHETIC_RECORDS)
    monkeypatch.setattr(m, "EXPECTED_PERTURBATIONS", N_SYNTHETIC_PERTURBATIONS)
    monkeypatch.setattr(m, "EXPECTED_GENES", N_SYNTHETIC_GENES)
    monkeypatch.setattr(m, "host_reference", lambda *a, **k: REFERENCE)
    monkeypatch.setattr(
        m,
        "ROUND_TWO_BACKGROUND_GUIDE",
        SourcedValue(
            value=SYNTHETIC_BACKGROUND_GUIDE,
            quote=m.ROUND_TWO_BACKGROUND_GUIDE.quote,
            provenance=m.ROUND_TWO_BACKGROUND_GUIDE.provenance,
        ),
    )
    return root


@pytest.fixture
def built(
    synthetic: Path, monkeypatch: pytest.MonkeyPatch, mg1655: EcoliK12MG1655Genome
) -> m.CrispriGuideFfaEnrichmentFang2025Dataset:
    """The synthetic build, with the resolved-fraction floor lowered for 6 of 7 names."""
    monkeypatch.setattr(
        m.CrispriGuideFfaEnrichmentFang2025Dataset, "MIN_RESOLVED_FRACTION", 0.8
    )
    return m.CrispriGuideFfaEnrichmentFang2025Dataset(
        root=str(synthetic), ecoli_genome=mg1655
    )


def _records(
    dataset: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> list[dict[str, Any]]:
    """Every stored record, with the LMDB handle closed before anything re-reads it."""
    out = [dataset[index] for index in range(len(dataset))]
    dataset.close_lmdb()
    return out


# --------------------------------------------------------------------------- #
# Reading one round
# --------------------------------------------------------------------------- #
def test_read_round_reproduces_both_equations_and_splits_on_the_read_floor(
    tmp_path: Path, synthetic_library_size: None
) -> None:
    """Round 1: 5 of 7 rows clear the floor, and both equations reproduce exactly."""
    path = tmp_path / SCREEN_RAW.filename
    write_screen(path)
    rows = m.read_round(path, ROUND_1)
    assert rows.round_id == "round_1_cf"
    assert rows.source_rows == N_SYNTHETIC_TARGETING
    assert list(rows.kept["guide_id"]) == [
        "thrLb0001_10",
        "thrLb0001_25",
        "proBb0005_60",
        "yaaPb0004_45",
        "ybfKb4590_5",
    ]
    assert rows.excluded_rows == 2
    assert rows.excluded_by_library == {
        "transformation": 1,
        "before_sorting": 0,
        "after_sorting": 1,
    }
    assert rows.total_reads == {
        "transformation": 610,
        "before_sorting": 700,
        "after_sorting": 555,
    }
    assert rows.max_equation_1_error < 1e-12
    assert rows.max_equation_2_error < 1e-12


def test_read_round_two_reads_two_libraries_and_keeps_four(
    tmp_path: Path, synthetic_library_size: None
) -> None:
    """Round 2 has no transformation library, so its floor is a two-way test."""
    path = tmp_path / SCREEN_RAW.filename
    write_screen(path)
    rows = m.read_round(path, ROUND_2)
    assert list(rows.kept["guide_id"]) == [
        "thrLb0001_10",
        "thrAb0002_12",
        "proBb0005_60",
        "ybfKb4590_5",
    ]
    assert rows.excluded_by_library == {"before_sorting": 1, "after_sorting": 2}
    assert set(rows.total_reads) == {"before_sorting", "after_sorting"}


def test_read_round_refuses_a_fitness_that_is_not_the_equation(
    tmp_path: Path, synthetic_library_size: None
) -> None:
    """A released value that is not log2(normalized AS / normalized BS) stops the read."""
    path = tmp_path / SCREEN_RAW.filename
    write_screen(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[ROUND_1.sheet]
    sheet.cell(row=3, column=ROUND_1.fitness_all_column + 1).value = 99.0
    workbook.save(path)
    with pytest.raises(ValueError, match="Methods equations do not reproduce"):
        m.read_round(path, ROUND_1)


def test_read_round_refuses_a_floor_that_disagrees_with_the_fitness_column(
    tmp_path: Path, synthetic_library_size: None
) -> None:
    """A qualified value on a row below the floor means the released split moved."""
    path = tmp_path / SCREEN_RAW.filename
    write_screen(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[ROUND_1.sheet]
    # thrAb0002_12 is row 5 and is below the floor on 'after_sorting'.
    sheet.cell(row=5, column=ROUND_1.fitness_column + 1).value = -1.0
    workbook.save(path)
    with pytest.raises(ValueError, match="floor disagrees with the released Fitness"):
        m.read_round(path, ROUND_1)


def test_read_round_refuses_a_library_that_is_not_the_released_size(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A sheet whose guide count is not the library's is a different release."""
    path = tmp_path / SCREEN_RAW.filename
    write_screen(path)
    monkeypatch.setattr(m, "N_TARGETING_GUIDES", N_SYNTHETIC_TARGETING + 1)
    with pytest.raises(ValueError, match="guide rows, the library is"):
        m.read_round(path, ROUND_1)


def test_read_round_refuses_a_repeated_guide_id(
    tmp_path: Path, synthetic_library_size: None
) -> None:
    """One sheet row per guide; a repeat would make two records of one measurement."""
    path = tmp_path / SCREEN_RAW.filename
    write_screen(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[ROUND_1.sheet]
    sheet.cell(row=4, column=ROUND_1.id_column + 1).value = "thrLb0001_10"
    workbook.save(path)
    with pytest.raises(ValueError, match="repeats a guide id"):
        m.read_round(path, ROUND_1)


# --------------------------------------------------------------------------- #
# The objects a record is made of
# --------------------------------------------------------------------------- #
def test_screen_environment_is_the_one_culture_both_arms_share() -> None:
    """One environment: BS and AS differ by the sort gate, not by the medium."""
    environment = m.screen_environment()
    assert environment.media == M9_MODIFIED_FANG2025
    assert environment.temperature is not None and environment.temperature.value == 30.0
    assert environment.duration_hours == 40.0
    assert environment.aerobicity == "aerobic"
    (inducer,) = environment.perturbations
    assert isinstance(inducer, SmallMoleculePerturbation)
    # IPTG is not in the shared compound table, so it carries a typed InChIKey gap. It
    # is deliberately NOT added: the resolver returns the TABLE's canonical name, so a
    # row for it would rename the IPTG compound node of the already-served Foo 2014
    # records, which is the full-rebuild case rather than an additive change.
    assert inducer.compound.name == "IPTG"
    assert inducer.compound.inchikey is None
    assert [g.field for g in inducer.compound.provenance_gaps] == ["inchikey"]
    assert inducer.concentration.value == 1.0
    assert inducer.concentration.unit is ConcentrationUnit.millimolar


def test_guide_phenotype_is_a_signed_log2_sort_enrichment() -> None:
    """The typed axes: a log2 ratio read off a biosensor sort, not a growth selection."""
    phenotype = m.guide_phenotype(ROUND_1, -6.25)
    assert phenotype.measurement_type is MeasurementType.log2_ratio
    assert phenotype.assay_type is AssayType.biosensor_readout
    assert phenotype.environment_response == -6.25
    assert phenotype.n_samples == 1
    assert phenotype.sample_unit is SampleUnit.pooled
    assert phenotype.screen_id == "round_1_cf"
    assert {gap.field for gap in phenotype.provenance_gaps} == {
        "environment_response_se",
        "environment_response_uncertainty",
        "environment_response_uncertainty_type",
    }


def test_reference_is_the_unsorted_library_at_exactly_zero() -> None:
    """A guide the sort does not move has fitness log2(1) = 0 by equation (2)."""
    reference = m.build_reference("fang", ROUND_2, REFERENCE)
    assert reference.phenotype_reference.environment_response == 0.0
    assert reference.phenotype_reference.screen_id == "round_2_pcnbi"
    assert reference.environment_reference == m.screen_environment()
    assert reference.genome_reference == REFERENCE


def test_host_background_keeps_the_lesion_as_a_verbatim_statement() -> None:
    """``alleles`` is empty because the paper prints no locus tag for ``fadE``."""
    background = m.host_background()
    assert background.name == "CF"
    assert background.alleles == []
    assert background.genotype_statement is not None
    assert "fadE deletion" in background.genotype_statement
    assert background.parents == ["E. coli MG1655(DE3)"]
    assert background.provenance is not None
    assert len(background.provenance) == 2


def test_build_genotype_adds_the_background_knockdown_only_when_asked() -> None:
    """Round 1 genotypes are single, round 2 genotypes carry the pcnBi repression too."""
    targets = (("b0005", "proB"), ("b0006", "proC"))
    plain = m.build_genotype(targets, GUIDES["proBb0005_60"])
    assert [p.systematic_gene_name for p in plain.perturbations] == ["b0005", "b0006"]
    doubled = m.build_genotype(
        targets, GUIDES["proBb0005_60"], (("b0143", "pcnB", "CAGTGGGTACCAGAACATGG"),)
    )
    assert [p.systematic_gene_name for p in doubled.perturbations] == [
        "b0005",
        "b0006",
        "b0143",
    ]
    knockdowns = [
        p
        for p in doubled.perturbations
        if isinstance(p, BacterialCrisprInterferencePerturbation)
    ]
    assert len(knockdowns) == len(doubled.perturbations)
    assert knockdowns[-1].crispr is not None
    assert knockdowns[-1].crispr.guide_sequence == "CAGTGGGTACCAGAACATGG"
    assert all(
        p.crispr is not None and p.crispr.effector == "dCas9" for p in knockdowns
    )


# --------------------------------------------------------------------------- #
# The synthetic build
# --------------------------------------------------------------------------- #
def test_build_writes_one_record_per_kept_guide_round_pair(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    assert len(built) == N_SYNTHETIC_RECORDS == 7


def test_build_validates_every_record_against_the_schema_classes(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    for record in _records(built):
        experiment = BacterialEnvironmentResponseExperiment.model_validate(
            record["experiment"]
        )
        reference = BacterialEnvironmentResponseExperimentReference.model_validate(
            record["reference"]
        )
        assert experiment.phenotype.measurement_type is MeasurementType.log2_ratio
        assert reference.phenotype_reference.environment_response == 0.0


def test_build_keeps_the_two_rounds_apart_by_screen_id(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    screens = [r["experiment"]["phenotype"]["screen_id"] for r in _records(built)]
    assert screens == ["round_1_cf"] * 4 + ["round_2_pcnbi"] * 3


def test_build_drops_the_retired_singleton_from_both_rounds(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    """``ybfKb4590_5`` clears the floor in both rounds and is stored in neither."""
    stored = {
        perturbation["systematic_gene_name"]
        for record in _records(built)
        for perturbation in record["experiment"]["genotype"]["perturbations"]
    }
    assert "b4590" not in stored
    assert stored == {"b0001", "b0002", "b0004", "b0005", "b0006"}
    assert len(stored) == N_SYNTHETIC_GENES


def test_build_gives_every_round_two_record_the_background_knockdown(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    """Round 2 is a double knockdown; the extra leaf is the same gene on every record."""
    records = _records(built)
    round_two = [
        r
        for r in records
        if r["experiment"]["phenotype"]["screen_id"] == "round_2_pcnbi"
    ]
    assert len(round_two) == 3
    for record in round_two:
        perturbations = record["experiment"]["genotype"]["perturbations"]
        names = [p["systematic_gene_name"] for p in perturbations]
        assert "b0004" in names
        assert len(names) == len(set(names))
        (background,) = [
            p for p in perturbations if p["systematic_gene_name"] == "b0004"
        ]
        assert (
            background["crispr"]["guide_sequence"] == GUIDES[SYNTHETIC_BACKGROUND_GUIDE]
        )
    round_one = [
        r for r in records if r["experiment"]["phenotype"]["screen_id"] == "round_1_cf"
    ]
    carries_b0004 = [
        r
        for r in round_one
        if any(
            p["systematic_gene_name"] == "b0004"
            for p in r["experiment"]["genotype"]["perturbations"]
        )
    ]
    # Round 1 keeps yaaPb0004_45 as a LIBRARY guide, and that single record is the only
    # round-1 record on b0004; no round-1 record gains the background leaf.
    assert len(carries_b0004) == 1
    assert len(carries_b0004[0]["experiment"]["genotype"]["perturbations"]) == 1


def test_build_carries_one_knockdown_per_cluster_member(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    """``proBb0005_60`` targets a two-member cluster, so its record carries two leaves."""
    records = _records(built)
    counts = [
        len(r["experiment"]["genotype"]["perturbations"])
        for r in records
        if r["experiment"]["phenotype"]["screen_id"] == "round_1_cf"
    ]
    assert sorted(counts) == [1, 1, 1, 2]
    total = sum(len(r["experiment"]["genotype"]["perturbations"]) for r in records)
    assert total == N_SYNTHETIC_PERTURBATIONS == 12


def test_build_stores_the_released_fitness_verbatim(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
) -> None:
    """The stored value is the workbook's own number, recomputed here from the reads."""
    totals_before = sum(v[1] for v in ROUND_1_READS.values())
    totals_after = sum(v[2] for v in ROUND_1_READS.values())
    reads = ROUND_1_READS["thrLb0001_10"]
    expected = math.log2(
        ((reads[2] + 1) / totals_after) / ((reads[1] + 1) / totals_before)
    )
    first = _records(built)[0]["experiment"]["phenotype"]["environment_response"]
    assert first == pytest.approx(expected, abs=1e-15)


def test_build_writes_the_round_retention_accounting(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset, synthetic: Path
) -> None:
    built.close_lmdb()
    accounting = json.loads(
        (synthetic / "preprocess" / "build_accounting.json").read_text()
    )
    assert accounting["kept_records"] == N_SYNTHETIC_RECORDS
    assert accounting["perturbations"] == N_SYNTHETIC_PERTURBATIONS
    assert accounting["distinct_genes"] == N_SYNTHETIC_GENES
    assert accounting["source_rows"] == 2 * N_SYNTHETIC_TARGETING
    assert accounting["read_floor"] == 20
    assert accounting["library_citation_key"] == "wangPooledCRISPRInterference2018"
    assert accounting["strain_background"] == "CF"
    assert accounting["retired_targets_dropped"] == ["b4590"]
    by_round = {r["round_id"]: r for r in accounting["rounds"]}
    assert by_round["round_1_cf"]["dropped_below_read_floor"] == 2
    assert by_round["round_1_cf"]["dropped_retired_target"] == 1
    assert by_round["round_1_cf"]["kept_records"] == 4
    assert by_round["round_2_pcnbi"]["dropped_below_read_floor"] == 3
    assert by_round["round_2_pcnbi"]["dropped_retired_target"] == 1
    assert by_round["round_2_pcnbi"]["background_knockdowns"] == 3
    assert (synthetic / "preprocess" / "round_retention.csv").exists()


def test_build_verifies_the_pins_of_both_mirrors(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
    presence_only_pins: list[Mapping[str, str]],
) -> None:
    """One ``verify_raw_files`` call covering this paper's file AND Wang 2018's two."""
    built.close_lmdb()
    (pins,) = presence_only_pins
    assert set(pins) == {
        SCREEN_RAW.filename,
        CLUSTERS_RAW.filename,
        LIBRARY_RAW.filename,
    }


def test_round_accounting_refuses_arithmetic_that_does_not_close() -> None:
    """The two drop counts must account for every released row."""
    accounting = m.RoundAccounting(
        round_id="round_1_cf",
        sheet="Figure 1",
        panel="Figure 1d",
        host_strain="CF",
        libraries=["before_sorting", "after_sorting"],
        total_reads={"before_sorting": 10, "after_sorting": 10},
        source_rows=7,
        dropped_below_read_floor=2,
        dropped_by_library={"before_sorting": 1, "after_sorting": 1},
        dropped_retired_target=1,
        kept_records=5,
        perturbations=5,
        background_knockdowns=0,
        max_equation_1_error=0.0,
        max_equation_2_error=0.0,
    )
    with pytest.raises(ValueError, match="not the 5 records written"):
        accounting.check()


def test_background_knockdown_refuses_a_guide_outside_the_library(
    synthetic: Path, monkeypatch: pytest.MonkeyPatch, mg1655: EcoliK12MG1655Genome
) -> None:
    """If the short-form resolution were wrong, the build stops rather than deriving."""
    monkeypatch.setattr(
        m.CrispriGuideFfaEnrichmentFang2025Dataset, "MIN_RESOLVED_FRACTION", 0.8
    )
    monkeypatch.setattr(
        m,
        "ROUND_TWO_BACKGROUND_GUIDE",
        SourcedValue(
            value="pcnBb0143_956",
            quote=m.ROUND_TWO_BACKGROUND_GUIDE.quote,
            provenance=m.ROUND_TWO_BACKGROUND_GUIDE.provenance,
        ),
    )
    with pytest.raises(RuntimeError, match="does not resolve"):
        m.CrispriGuideFfaEnrichmentFang2025Dataset(
            root=str(synthetic), ecoli_genome=mg1655
        )


def test_build_refuses_a_genome_of_another_assembly_set(
    synthetic: Path, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """A BW25113 tag is not a b-number, so resolving this release against it is refused."""
    from tests.torchcell.sequence.genome._bacterial_fixtures import BW25113_LOCI
    from torchcell.sequence.genome.ecoli.k12 import (
        BW25113_ASSEMBLY,
        EcoliK12BW25113Genome,
    )

    files = write_assembly(tmp_path / "bw-tier", BW25113_ASSEMBLY, BW25113_LOCI)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    root = tmp_path / "bw25113"
    root.mkdir()
    bw25113 = EcoliK12BW25113Genome(genome_root=str(root), overwrite=False)
    monkeypatch.setattr(
        m.CrispriGuideFfaEnrichmentFang2025Dataset, "MIN_RESOLVED_FRACTION", 0.8
    )
    with pytest.raises(ValueError, match="needs the ecoli_K12_MG1655_ASM584v2 genome"):
        m.CrispriGuideFfaEnrichmentFang2025Dataset(
            root=str(synthetic), ecoli_genome=bw25113
        )


def test_dataset_exposes_its_schema_classes() -> None:
    dataset = m.CrispriGuideFfaEnrichmentFang2025Dataset.__new__(
        m.CrispriGuideFfaEnrichmentFang2025Dataset
    )
    assert dataset.experiment_class is BacterialEnvironmentResponseExperiment
    assert dataset.reference_class is BacterialEnvironmentResponseExperimentReference
    with pytest.raises(NotImplementedError, match="builds records in process"):
        dataset.create_experiment()
    assert dataset.preprocess_raw("frame") == "frame"


def test_manifest_sha256_refuses_a_path_it_does_not_carry() -> None:
    from torchcell.literature.manifest import Manifest

    manifest = Manifest(citation_key=m.CITATION_KEY, doi=m.DOI, title=m.TITLE, files=[])
    with pytest.raises(KeyError, match="is not in the Fang 2025 raw-mirror manifest"):
        m.manifest_sha256(manifest, "si/si_data/nope.xlsx")


def test_raw_file_names_the_publisher_object_and_its_cloud_key() -> None:
    assert m.SCREEN_FILE.filename == "41467_2025_58368_MOESM8_ESM.xlsx"
    assert m.SCREEN_FILE.relpath == "si/si_data/41467_2025_58368_MOESM8_ESM.xlsx"
    assert m.SCREEN_FILE.cloud_key == ("PMC11954867.1/41467_2025_58368_MOESM8_ESM.xlsx")
    assert m.SCREEN_FILE.url.startswith("https://pmc-oa-opendata.s3.amazonaws.com/")


def test_deposit_refuses_a_source_whose_bytes_are_not_the_pin(tmp_path: Path) -> None:
    """A mismatching file is refused BEFORE anything is written."""
    wrong = tmp_path / "wrong.xlsx"
    wrong.write_bytes(b"not the workbook")
    with pytest.raises(RuntimeError, match="hashes to"):
        m.deposit_raw_mirror(
            {m.SCREEN_FILE.relpath: wrong}, data_root=str(tmp_path / "root")
        )
    assert not (tmp_path / "root").exists()


def test_deposit_writes_the_mirror_and_its_manifest_from_a_staged_source(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The ``sources=`` path copies the file, records its retrieval and is idempotent."""
    workbook = tmp_path / "staged.xlsx"
    write_screen(workbook)
    staged = m.RawFile(
        moesm=8,
        data_number=6,
        description="synthetic source data",
        sha256=_sha256_of(workbook),
    )
    monkeypatch.setattr(m, "RAW_FILES", (staged,))
    root = m.deposit_raw_mirror(
        {staged.relpath: workbook}, data_root=str(tmp_path / "root")
    )
    assert (root / staged.relpath).exists()
    manifest = m.load_manifest(str(tmp_path / "root"))
    assert manifest.citation_key == m.CITATION_KEY
    (record,) = manifest.files
    assert record.path == staged.relpath
    assert record.sha256 == staged.sha256
    assert record.retrieval is not None
    assert record.retrieval.method is RetrievalMethod.pmc_cloud
    assert record.retrieval.params == {"key": staged.cloud_key}
    assert manifest.provenance_complete is True
    assert m.manifest_sha256(manifest, staged.relpath) == staged.sha256
    # Idempotent: a second deposit of identical bytes leaves the mirror alone.
    assert (
        m.deposit_raw_mirror(
            {staged.relpath: workbook}, data_root=str(tmp_path / "root")
        )
        == root
    )


def test_deposit_refuses_to_overwrite_a_mirror_file_that_differs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A deposited file whose bytes changed is a NEW version, never an overwrite."""
    workbook = tmp_path / "staged.xlsx"
    write_screen(workbook)
    staged = m.RawFile(
        moesm=8,
        data_number=6,
        description="synthetic source data",
        sha256=_sha256_of(workbook),
    )
    monkeypatch.setattr(m, "RAW_FILES", (staged,))
    root = m.deposit_raw_mirror(
        {staged.relpath: workbook}, data_root=str(tmp_path / "root")
    )
    (root / staged.relpath).write_bytes(b"different bytes")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        m.deposit_raw_mirror(
            {staged.relpath: workbook}, data_root=str(tmp_path / "root")
        )


def test_download_links_both_mirrors_after_verifying_each_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mg1655: EcoliK12MG1655Genome
) -> None:
    """One ``raw/`` fed by two mirrors, each read under its own citation key."""
    data_root = tmp_path / "root"
    screen = tmp_path / "screen.xlsx"
    write_screen(screen)
    clusters = tmp_path / "clusters.xlsx"
    write_clusters(clusters)
    library = tmp_path / "library.xlsx"
    write_library(library)
    screen_raw = m.RawFile(
        moesm=8, data_number=6, description="screen", sha256=_sha256_of(screen)
    )
    clusters_raw = m.RawFile(
        moesm=5, data_number=2, description="clusters", sha256=_sha256_of(clusters)
    )
    library_raw = m.RawFile(
        moesm=6, data_number=3, description="library", sha256=_sha256_of(library)
    )
    monkeypatch.setattr(m, "RAW_FILES", (screen_raw,))
    monkeypatch.setattr(m, "WANG_RAW_FILES", (clusters_raw, library_raw))
    m.deposit_raw_mirror({screen_raw.relpath: screen}, data_root=str(data_root))
    monkeypatch.setattr(m, "_data_root", lambda: str(data_root))

    wang_manifest = Manifest(
        citation_key="wangPooledCRISPRInterference2018",
        doi="10.1038/s41467-018-04899-x",
        title="Wang 2018",
        files=[
            ArtifactRecord(
                path=raw.relpath,
                role=ROLE_SI_DATA,
                bytes=source.stat().st_size,
                sha256=raw.sha256,
                source=raw.url,
                original_filename=raw.filename,
            )
            for raw, source in ((clusters_raw, clusters), (library_raw, library))
        ],
    )
    wang_root = tmp_path / "wang"
    for raw, source in ((clusters_raw, clusters), (library_raw, library)):
        destination = wang_root / raw.relpath
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(source.read_bytes())
    (wang_root / "manifest.json").write_text(wang_manifest.model_dump_json(indent=2))
    monkeypatch.setattr(m, "wang_raw_mirror_dir", lambda *a, **k: wang_root)
    monkeypatch.setattr(m, "wang_load_manifest", lambda *a, **k: wang_manifest)
    monkeypatch.setattr(
        m,
        "wang_manifest_sha256",
        lambda manifest, relpath: next(
            f.sha256 for f in manifest.files if f.path == relpath
        ),
    )

    dataset = m.CrispriGuideFfaEnrichmentFang2025Dataset.__new__(
        m.CrispriGuideFfaEnrichmentFang2025Dataset
    )
    raw_dir = tmp_path / "dataset" / "raw"
    monkeypatch.setattr(
        type(dataset), "raw_dir", property(lambda self: str(raw_dir)), raising=False
    )
    dataset.download()
    assert sorted(f.name for f in raw_dir.iterdir()) == sorted(
        [screen_raw.filename, clusters_raw.filename, library_raw.filename]
    )
    assert dataset.raw_file_names == [
        screen_raw.filename,
        clusters_raw.filename,
        library_raw.filename,
    ]


def test_download_refuses_a_mirror_missing_the_pinned_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A manifest that pins a file the mirror does not hold stops the build."""
    data_root = tmp_path / "root"
    screen = tmp_path / "screen.xlsx"
    write_screen(screen)
    screen_raw = m.RawFile(
        moesm=8, data_number=6, description="screen", sha256=_sha256_of(screen)
    )
    monkeypatch.setattr(m, "RAW_FILES", (screen_raw,))
    root = m.deposit_raw_mirror({screen_raw.relpath: screen}, data_root=str(data_root))
    (root / screen_raw.relpath).unlink()
    monkeypatch.setattr(m, "_data_root", lambda: str(data_root))
    dataset = m.CrispriGuideFfaEnrichmentFang2025Dataset.__new__(
        m.CrispriGuideFfaEnrichmentFang2025Dataset
    )
    raw_dir = tmp_path / "dataset" / "raw"
    monkeypatch.setattr(
        type(dataset), "raw_dir", property(lambda self: str(raw_dir)), raising=False
    )
    with pytest.raises(RuntimeError, match="missing from mirror"):
        dataset.download()


def test_verify_build_runs_every_level_on_the_synthetic_store(
    built: m.CrispriGuideFfaEnrichmentFang2025Dataset,
    synthetic: Path,
    mg1655: EcoliK12MG1655Genome,
) -> None:
    """The L0-L4 verifier over the synthetic build, with its report written beside it."""
    built.close_lmdb()
    report = m.verify_build(
        str(synthetic), genome=mg1655, expected_count=N_SYNTHETIC_RECORDS
    )
    failed = [r.name for r in report.results if not r.passed]
    assert failed == []
    assert {int(r.level) for r in report.results} >= {0, 1, 2, 3, 4}
    written = json.loads(
        (synthetic / "preprocess" / "verification_report.json").read_text()
    )
    assert written["results"]
    assert report.summary()


def test_publication_carries_the_doi_and_asserts_no_pubmed_id() -> None:
    publication = m.publication()
    assert publication.doi == m.DOI
    assert publication.doi_url == f"https://doi.org/{m.DOI}"
    assert publication.pubmed_id is None


# --------------------------------------------------------------------------- #
# Data-gated: the real mirror and the built dev-tree store
# --------------------------------------------------------------------------- #
def _data_root() -> str | None:
    from dotenv import load_dotenv

    load_dotenv()
    return os.environ.get("DATA_ROOT")


DATA_ROOT = _data_root()
MIRROR = (
    Path(DATA_ROOT) / m.RAW_DIR_REL / "manifest.json" if DATA_ROOT else Path("absent")
)
STORE = (
    Path(DATA_ROOT) / "data/torchcell/crispri_guide_ffa_enrichment_fang2025"
    if DATA_ROOT
    else Path("absent")
)
LIBRARY = Path(DATA_ROOT) / "torchcell-library" if DATA_ROOT else Path("absent")
needs_mirror = pytest.mark.skipif(
    not MIRROR.exists(), reason="the Fang 2025 raw mirror is not deposited here"
)
needs_store = pytest.mark.skipif(
    not (STORE / "processed/lmdb").exists(),
    reason="the Fang 2025 dev-tree LMDB is not built here",
)
needs_library = pytest.mark.skipif(
    not (LIBRARY / m.CITATION_KEY / m.PAPER_MD).exists(),
    reason="the Fang 2025 paper OCR is not mirrored here",
)

#: Every module-level sourced value, which is what the quote audit walks.
SOURCED_VALUES: dict[str, SourcedValue] = {
    name: value for name, value in vars(m).items() if isinstance(value, SourcedValue)
}


@pytest.mark.data
@needs_mirror
def test_the_mirror_pins_the_screen_workbook_to_a_rerunnable_retrieval() -> None:
    """One deposited file, with a scriptable ``pmc_cloud`` retrieval and its sha256."""
    manifest = m.load_manifest(DATA_ROOT)
    assert {f.path for f in manifest.files} == {r.relpath for r in m.RAW_FILES}
    for raw in m.RAW_FILES:
        assert m.manifest_sha256(manifest, raw.relpath) == raw.sha256
        record = next(f for f in manifest.files if f.path == raw.relpath)
        assert record.retrieval is not None
        assert record.retrieval.method is RetrievalMethod.pmc_cloud
        assert record.retrieval.params == {"key": raw.cloud_key}
        assert record.retrieval.source_url == raw.url
    assert manifest.provenance_complete is True
    assert any("wangPooledCRISPRInterference2018" in e for e in manifest.si_expected)


@pytest.mark.data
@needs_library
@pytest.mark.parametrize("name", sorted(SOURCED_VALUES))
def test_every_quote_is_still_verbatim_in_the_pinned_ocr(name: str) -> None:
    """The sha256 still matches and the quote is still a substring of those bytes."""
    result = audit_sourced_value(SOURCED_VALUES[name], LIBRARY)
    assert result.passed, result.message


@pytest.mark.data
@needs_store
def test_the_built_store_holds_the_measured_record_count() -> None:
    """15,708 records over 2,628 genes, with the per-round arithmetic balancing."""
    payload = json.loads((STORE / "preprocess/build_accounting.json").read_text())
    accounting = m.BuildAccounting.model_validate(payload)
    accounting.check()
    assert accounting.source_rows == 111342
    assert accounting.kept_records == m.EXPECTED_RECORDS == 15708
    assert accounting.perturbations == m.EXPECTED_PERTURBATIONS == 16187
    assert accounting.distinct_genes == m.EXPECTED_GENES == 2628
    assert accounting.read_floor == 20
    assert accounting.library_citation_key == "wangPooledCRISPRInterference2018"
    by_round = {r.round_id: r for r in accounting.rounds}
    assert by_round["round_1_cf"].kept_records == 15378
    assert by_round["round_2_pcnbi"].kept_records == 330
    assert by_round["round_1_cf"].dropped_by_library["after_sorting"] == 38379
    assert by_round["round_2_pcnbi"].dropped_by_library["after_sorting"] == 55334
    for round_ in accounting.rounds:
        assert round_.max_equation_1_error < 1e-12
        assert round_.max_equation_2_error < 1e-12


@pytest.mark.data
@needs_store
def test_the_built_store_carries_the_round_two_background_knockdown() -> None:
    """Every round-2 record is a double knockdown whose extra leaf is pcnB on b0143."""
    dataset = m.CrispriGuideFfaEnrichmentFang2025Dataset(root=str(STORE))
    try:
        last = dataset[len(dataset) - 1]["experiment"]
        assert last["phenotype"]["screen_id"] == "round_2_pcnbi"
        names = {p["systematic_gene_name"] for p in last["genotype"]["perturbations"]}
        assert "b0143" in names
        assert len(names) == 2
        first = dataset[0]["experiment"]
        assert first["phenotype"]["screen_id"] == "round_1_cf"
        assert len(first["genotype"]["perturbations"]) == 1
    finally:
        dataset.close_lmdb()


@pytest.mark.data
@needs_store
def test_the_built_store_passes_every_verification_level() -> None:
    report = m.verify_build(str(STORE), data_root=DATA_ROOT)
    failed = [r.name for r in report.results if not r.passed]
    assert failed == []
