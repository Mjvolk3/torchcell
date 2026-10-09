# tests/torchcell/datasets/pputida/test_carruthers2025_fold_change.py
# [[tests.torchcell.datasets.pputida.test_carruthers2025_fold_change]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/pputida/test_carruthers2025_fold_change.py
"""``ProteomeFoldChangeCarruthers2025Dataset``: the released differential proteomics.

Synthetic tests run everywhere: the -log10 p-value inversion and its round trip, the
ragged two-column reader, the contrast-column pairing, the phenotype and its neutral
reference, the three build-time assertions (the Supplementary Fig. 9 pointer cell, the
contrast-to-strain join, the Supplementary Fig. 10 refusal), and the fold-change L0-L3
verifier.

The ``@pytest.mark.data`` tests read the real ``$DATA_ROOT``: they measure both released
sheets on the pinned workbook, and they read the built LMDB and run L0 to L4 over it.

Derived expectations for the pinned Source Data workbook, every one measured on sha256
``1b3a7ab5...``: ``Figure 5b`` releases 1,290 keyed rows over 14 contrast columns and
5,770 fold changes, of which 5,732 survive key reconciliation; ``Figure 6b`` releases
372 keyed rows (beside 181 orphan rows that carry only a ``primary_name``) over 10
contrast columns and 452 fold changes, of which the two stored contrasts contribute 14
(``PP_0812``) and 212 (``PP_0815``); 16 records hold 5,958 fold changes in all; and
``Supplementary Figure 10`` holds 26 keys x 25 strain columns = 650 dense values in
[-7.05757655263963, 3.92161053900316] with no p-value column.
"""

from __future__ import annotations

import math
import os
import os.path as osp
import statistics
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.pputida.carruthers2025 as c25
from torchcell.datamodels.schema import (
    BacterialProteinFoldChangeExperiment,
    FoldChangeScale,
    Genotype,
    ProteinFoldChangePhenotype,
)
from torchcell.verification.protein_fold_change import (
    fold_change_p_value_round_trip,
    protein_fold_change_locus_set,
    verify_protein_fold_change_dataset,
)
from torchcell.verification.report import Level, Provenance

PROV = Provenance(
    source_uri="si/source.xlsx",
    citation_key=c25.CITATION_KEY,
    sha256=c25.SOURCE_DATA_SHA256,
    method="unit test",
    page="unit test",
)


# --------------------------------------------------------------------------- #
# Synthetic: the p-value inversion
# --------------------------------------------------------------------------- #
def test_p_value_from_neg_log10_is_exact_arithmetic() -> None:
    """``p = 10**-x``, and the stored number inverts back to the released one."""
    assert c25.p_value_from_neg_log10(2.0) == 0.01
    assert c25.p_value_from_neg_log10(3.0) == 0.001
    released = 19.3125325679514
    probability = c25.p_value_from_neg_log10(released)
    assert probability == 10.0**-released
    assert -math.log10(probability) == pytest.approx(released, abs=1e-12)


def test_p_value_from_neg_log10_refuses_a_value_at_or_above_the_stated_ceiling() -> (
    None
):
    """The captions state p<0.05, so -log10(p) = 1.0 (p = 0.1) is not from these sheets."""
    with pytest.raises(RuntimeError, match="is not below the 0.05"):
        c25.p_value_from_neg_log10(1.0)


def test_p_value_from_neg_log10_refuses_a_negative_or_non_finite_input() -> None:
    """A -log10 p-value is non-negative and finite; anything else is not that column."""
    with pytest.raises(RuntimeError, match="non-negative finite number"):
        c25.p_value_from_neg_log10(-0.5)
    with pytest.raises(RuntimeError, match="non-negative finite number"):
        c25.p_value_from_neg_log10(float("nan"))


def test_the_released_p_value_ceiling_and_round_trip_tolerance_are_the_measured_ones() -> (
    None
):
    """The two constants the conversion is gated on, pinned."""
    assert c25.FOLD_CHANGE_P_VALUE_CEILING == 0.05
    assert c25.P_VALUE_ROUND_TRIP_TOL == 1e-12
    assert c25.NEG_LOG10_P_VALUE_SUFFIX == "_log10_pval"
    assert c25.FOLD_CHANGE_SUFFIX == "_log2_FC"


# --------------------------------------------------------------------------- #
# Synthetic: the ragged two-column reader
# --------------------------------------------------------------------------- #
def _fold_change_workbook(
    path: Path, sheet: str, header: list[Any], rows: list[list[Any]]
) -> str:
    """Write one fold-change sheet verbatim, header row included."""
    book = openpyxl.Workbook()
    page = book.active
    assert page is not None
    page.title = sheet
    page.append(header)
    for row in rows:
        page.append(row)
    book.save(path)
    return str(path)


SINGLE_HEADER = [
    "Locus Name",
    "Primary Name",
    "PP_0368_log2_FC",
    "PP_0368_log10_pval",
    "PP_0815_log2_FC",
    "PP_0815_log10_pval",
]


def test_read_fold_change_rows_keeps_a_ragged_release_ragged(tmp_path: Path) -> None:
    """An untested protein is no row at all, never a 0 and never a 1."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_FOLD_CHANGE,
        SINGLE_HEADER,
        [
            ["PP_0001", "abcA", 0.5, 2.0, None, None],
            ["PP_0002", None, None, None, -1.5, 3.0],
        ],
    )
    rows = c25.read_fold_change_rows(path, c25.SHEET_FOLD_CHANGE)
    assert [(row.contrast, row.protein, row.log2_fold_change) for row in rows] == [
        ("PP_0368", "PP_0001", 0.5),
        ("PP_0815", "PP_0002", -1.5),
    ]
    assert [row.neg_log10_p_value for row in rows] == [2.0, 3.0]
    assert {row.sheet for row in rows} == {c25.SHEET_FOLD_CHANGE}


def test_read_fold_change_rows_skips_the_orphan_rows_the_ko_sheet_carries(
    tmp_path: Path,
) -> None:
    """A row with a ``primary_name`` and no ``Locus Name`` and no value is not a record."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_KO_FOLD_CHANGE,
        SINGLE_HEADER,
        [
            [None, "aspS", None, None, None, None],
            ["PP_0001", None, 0.5, 2.0, None, None],
        ],
    )
    rows = c25.read_fold_change_rows(path, c25.SHEET_KO_FOLD_CHANGE)
    assert [row.protein for row in rows] == ["PP_0001"]


def test_read_fold_change_rows_refuses_a_keyless_row_that_carries_a_value(
    tmp_path: Path,
) -> None:
    """Dropping the key column can never drop a measurement."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_KO_FOLD_CHANGE,
        SINGLE_HEADER,
        [[None, "aspS", 0.5, 2.0, None, None]],
    )
    with pytest.raises(
        RuntimeError, match="a row with no 'Locus Name' carries a value"
    ):
        c25.read_fold_change_rows(path, c25.SHEET_KO_FOLD_CHANGE)


def test_read_fold_change_rows_refuses_a_half_released_pair(tmp_path: Path) -> None:
    """A fold change and its p-value are released together or not at all."""
    path = _fold_change_workbook(
        tmp_path / "a.xlsx",
        c25.SHEET_FOLD_CHANGE,
        SINGLE_HEADER,
        [["PP_0001", None, 0.5, None, None, None]],
    )
    with pytest.raises(RuntimeError, match="a fold change with no p-value"):
        c25.read_fold_change_rows(path, c25.SHEET_FOLD_CHANGE)
    path = _fold_change_workbook(
        tmp_path / "b.xlsx",
        c25.SHEET_FOLD_CHANGE,
        SINGLE_HEADER,
        [["PP_0001", None, None, 2.0, None, None]],
    )
    with pytest.raises(RuntimeError, match="a p-value with no fold change"):
        c25.read_fold_change_rows(path, c25.SHEET_FOLD_CHANGE)


def test_read_fold_change_rows_refuses_one_protein_twice_in_one_contrast(
    tmp_path: Path,
) -> None:
    """One contrast cannot give one protein two fold changes."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_FOLD_CHANGE,
        SINGLE_HEADER,
        [
            ["PP_0001", None, 0.5, 2.0, None, None],
            ["PP_0001", None, 0.6, 2.1, None, None],
        ],
    )
    with pytest.raises(RuntimeError, match="appears twice"):
        c25.read_fold_change_rows(path, c25.SHEET_FOLD_CHANGE)


def test_read_fold_change_rows_refuses_a_misplaced_p_value_column(
    tmp_path: Path,
) -> None:
    """The p-value column must be its fold change's immediate right neighbor."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_FOLD_CHANGE,
        ["Locus Name", "PP_0368_log2_FC", "PP_0815_log10_pval"],
        [["PP_0001", 0.5, 2.0]],
    )
    with pytest.raises(RuntimeError, match="is not followed by 'PP_0368_log10_pval'"):
        c25.read_fold_change_rows(path, c25.SHEET_FOLD_CHANGE)


def test_read_fold_change_rows_refuses_a_sheet_with_no_fold_change_column(
    tmp_path: Path,
) -> None:
    """A renamed export is not the sheet this loader reads."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_FOLD_CHANGE,
        ["Locus Name", "Primary Name"],
        [["PP_0001", "abcA"]],
    )
    with pytest.raises(RuntimeError, match="no _log2_FC column"):
        c25.read_fold_change_rows(path, c25.SHEET_FOLD_CHANGE)


def test_fold_change_contrasts_are_the_released_column_order(tmp_path: Path) -> None:
    """The contrast order is the sheet's, not sorted."""
    path = _fold_change_workbook(
        tmp_path / "fc.xlsx",
        c25.SHEET_FOLD_CHANGE,
        [
            "Locus Name",
            "PP_0815_log2_FC",
            "PP_0815_log10_pval",
            "PP_0368_log2_FC",
            "PP_0368_log10_pval",
        ],
        [["PP_0001", 0.5, 2.0, 0.6, 2.1]],
    )
    assert c25.fold_change_contrasts(path, c25.SHEET_FOLD_CHANGE) == (
        "PP_0815",
        "PP_0368",
    )


# --------------------------------------------------------------------------- #
# Synthetic: the phenotype and its neutral reference
# --------------------------------------------------------------------------- #
def _cells(*triples: tuple[str, float, float]) -> list[c25.FoldChangeRow]:
    return [
        c25.FoldChangeRow(
            sheet=c25.SHEET_FOLD_CHANGE,
            contrast="PP_0815",
            protein=protein,
            log2_fold_change=value,
            neg_log10_p_value=significance,
        )
        for protein, value, significance in triples
    ]


def test_the_phenotype_stores_the_log2_value_the_probability_and_the_triplicate() -> (
    None
):
    """Scale, basis, p = 10**-x and n = 3 per protein."""
    phenotype = c25.ProteomeFoldChangeCarruthers2025Dataset._phenotype(
        _cells(("PP_0001", -1.25, 2.0), ("PP_0002", 3.5, 4.0)),
        {"PP_0001": "PP_0001", "PP_0002": "PP_0002"},
        (),
        reference_basis="the non-target control strain",
    )
    assert phenotype.protein_fold_change == {"PP_0001": -1.25, "PP_0002": 3.5}
    assert phenotype.protein_fold_change_p_value == {"PP_0001": 0.01, "PP_0002": 0.0001}
    assert phenotype.n_replicates == {"PP_0001": 3, "PP_0002": 3}
    assert phenotype.fold_change_scale is FoldChangeScale.log2
    assert phenotype.reference_basis == "the non-target control strain"
    assert phenotype.measurement_type == (
        "dia_log2_fold_change_paired_two_tailed_t_test"
    )
    assert phenotype.protein_fold_change_se is None
    assert phenotype.protein_fold_change_p_value_adjusted is None
    assert phenotype.p_value_adjustment_method is None


def test_the_phenotype_drops_a_key_that_is_not_a_locus_of_the_pinned_assembly() -> None:
    """A dropped key has no gene node, so it has no fold change in the record."""
    phenotype = c25.ProteomeFoldChangeCarruthers2025Dataset._phenotype(
        _cells(("PP_0001", -1.25, 2.0), ("Q9FD70", 3.5, 4.0)),
        {"PP_0001": "PP_0001", "Q9FD70": "Q9FD70"},
        ("Q9FD70",),
        reference_basis="the non-target control strain",
    )
    assert set(phenotype.protein_fold_change) == {"PP_0001"}


def test_the_phenotype_refuses_two_released_keys_that_reconcile_to_one_locus() -> None:
    """One locus cannot carry two fold changes in one contrast."""
    with pytest.raises(RuntimeError, match="two released keys reconcile to PP_0001"):
        c25.ProteomeFoldChangeCarruthers2025Dataset._phenotype(
            _cells(("alpha", -1.25, 2.0), ("beta", 3.5, 4.0)),
            {"alpha": "PP_0001", "beta": "PP_0001"},
            (),
            reference_basis="the non-target control strain",
        )


def test_the_phenotype_refuses_a_contrast_with_no_resolved_key() -> None:
    """An empty abundance map is not a measurement."""
    with pytest.raises(RuntimeError, match="a contrast with no resolved protein key"):
        c25.ProteomeFoldChangeCarruthers2025Dataset._phenotype(
            _cells(("Q9FD70", 3.5, 4.0)),
            {"Q9FD70": "Q9FD70"},
            ("Q9FD70",),
            reference_basis="the non-target control strain",
        )


def test_the_reference_phenotype_is_the_scales_neutral_value_and_carries_no_p_value() -> (
    None
):
    """A fold change's denominator is 0.0 on the log2 scale by definition."""
    phenotype = c25.ProteomeFoldChangeCarruthers2025Dataset._phenotype(
        _cells(("PP_0002", 3.5, 4.0), ("PP_0001", -1.25, 2.0)),
        {"PP_0001": "PP_0001", "PP_0002": "PP_0002"},
        (),
        reference_basis="the non-target control strain",
    )
    reference = c25.ProteomeFoldChangeCarruthers2025Dataset._reference_phenotype(
        phenotype
    )
    assert reference.protein_fold_change == {"PP_0001": 0.0, "PP_0002": 0.0}
    assert list(reference.protein_fold_change) == ["PP_0001", "PP_0002"]
    assert reference.protein_fold_change_p_value is None
    assert reference.n_replicates == phenotype.n_replicates
    assert reference.reference_basis == phenotype.reference_basis
    assert reference.fold_change_scale is phenotype.fold_change_scale


# --------------------------------------------------------------------------- #
# Synthetic: the three build-time assertions
# --------------------------------------------------------------------------- #
def _pointer_workbook(path: Path, cells: list[Any]) -> str:
    book = openpyxl.Workbook()
    page = book.active
    assert page is not None
    page.title = c25.SHEET_FOLD_CHANGE_POINTER
    for cell in cells:
        page.append([cell])
    book.save(path)
    return str(path)


def test_the_supplementary_figure_9_pointer_is_what_sources_the_denominator(
    tmp_path: Path,
) -> None:
    """One cell, ``See Figure 5b``, is what ties that caption to this sheet."""
    proof = c25.assert_fold_change_pointer(
        _pointer_workbook(tmp_path / "p.xlsx", ["See Figure 5b"])
    )
    assert "See Figure 5b" in proof
    assert c25.SHEET_FOLD_CHANGE in proof
    assert c25.FOLD_CHANGE_POINTER_CELL == "See Figure 5b"


def test_a_changed_pointer_cell_stops_the_build(tmp_path: Path) -> None:
    """If the sheet stops pointing at Figure 5b the sourced basis is gone."""
    with pytest.raises(RuntimeError, match="holds \\['See Figure 7'\\]"):
        c25.assert_fold_change_pointer(
            _pointer_workbook(tmp_path / "p.xlsx", ["See Figure 7"])
        )


def test_a_missing_pointer_sheet_stops_the_build(tmp_path: Path) -> None:
    """A renamed sheet means the pinned workbook is not the one this loader reads."""
    book = openpyxl.Workbook()
    page = book.active
    assert page is not None
    page.title = "other"
    book.save(tmp_path / "p.xlsx")
    with pytest.raises(
        RuntimeError, match=f"has no sheet {c25.SHEET_FOLD_CHANGE_POINTER!r}"
    ):
        c25.assert_fold_change_pointer(str(tmp_path / "p.xlsx"))


def _heatmap_workbook(path: Path, header: list[Any], rows: list[list[Any]]) -> str:
    book = openpyxl.Workbook()
    page = book.active
    assert page is not None
    page.title = c25.SHEET_BEST_ARRAY_HEATMAP
    page.append(header)
    for row in rows:
        page.append(row)
    book.save(path)
    return str(path)


def test_the_best_array_heatmap_refusal_is_a_measurement(tmp_path: Path) -> None:
    """The refusal records the shape, the density, the range and the key form."""
    path = _heatmap_workbook(
        tmp_path / "h.xlsx",
        ["POI", "PP_0368_PP_0815", "PP_0528_PP_0815"],
        [["PP_2793", 1.0, -2.0], ["AcsA1", -1.0, 3.0]],
    )
    proofs = c25.assert_best_array_heatmap_refusal(path)
    assert len(proofs) == 2
    assert "2 keys x 2 strain columns = 4 values with no empty cell" in proofs[0]
    assert "[-2.0, 3.0] with median 0.0" in proofs[0]
    assert "only 1 of 2 keys a PP_ locus tag" in proofs[0]
    assert "no mirrored statement names a scale" in proofs[1]


def test_a_heatmap_that_gains_a_p_value_column_invalidates_the_refusal(
    tmp_path: Path,
) -> None:
    """The refusal rests on there being no p-value; a new one must be re-measured."""
    path = _heatmap_workbook(
        tmp_path / "h.xlsx",
        ["POI", "PP_0368_PP_0815", "PP_0368_PP_0815_log10_pval"],
        [["PP_2793", 1.0, 2.0]],
    )
    with pytest.raises(RuntimeError, match="the recorded refusal is stale"):
        c25.assert_best_array_heatmap_refusal(path)


def test_a_heatmap_that_is_no_longer_dense_invalidates_the_refusal(
    tmp_path: Path,
) -> None:
    """A ragged heatmap is a different release and is not the one that was measured."""
    path = _heatmap_workbook(
        tmp_path / "h.xlsx",
        ["POI", "PP_0368_PP_0815", "PP_0528_PP_0815"],
        [["PP_2793", 1.0, None]],
    )
    with pytest.raises(RuntimeError, match="is no longer dense"):
        c25.assert_best_array_heatmap_refusal(path)


# --------------------------------------------------------------------------- #
# Synthetic: the sourced contrasts and the ledgered refusals
# --------------------------------------------------------------------------- #
def test_the_two_stored_ko_contrasts_are_the_two_the_caption_names() -> None:
    """The Fig. 6 caption names two KO strains, and these are they."""
    assert c25.KO_FOLD_CHANGE_CONTRASTS == ("PP_0812", "PP_0815")
    assert c25.KO_FOLD_CHANGE_SOURCE.value == c25.KO_FOLD_CHANGE_CONTRASTS
    assert "two KO strains" in str(c25.KO_FOLD_CHANGE_SOURCE.quote)


def test_the_ledger_covers_the_other_eight_ko_columns_and_nothing_else() -> None:
    """Every refused column has its own measurement, and no stored one is in the ledger."""
    assert set(c25.KO_FOLD_CHANGE_UNSOURCED_REASONS) == {
        "Control",
        "PP_0368",
        "PP_0751",
        "PP_0812_15",
        "PP_0813",
        "PP_0814",
        "PP_0751_PP_0812",
        "PP_1317_PP_0812",
    }
    assert (
        set(c25.KO_FOLD_CHANGE_UNSOURCED_REASONS) & set(c25.KO_FOLD_CHANGE_CONTRASTS)
        == set()
    )
    assert (
        "BUILT two-sgRNA CRISPRi array"
        in (c25.KO_FOLD_CHANGE_UNSOURCED_REASONS["PP_0751_PP_0812"])
    )
    assert (
        "matches no released strain"
        in (c25.KO_FOLD_CHANGE_UNSOURCED_REASONS["PP_1317_PP_0812"])
    )
    assert (
        "pronoun 'both deletion strains'"
        in (c25.KO_FOLD_CHANGE_UNSOURCED_REASONS["PP_0812_15"])
    )


def test_the_record_counts_are_the_sourced_contrast_counts() -> None:
    """14 single-guide columns plus the 2 sourced KO columns."""
    assert c25.EXPECTED_FOLD_CHANGE_RECORDS_FIG5B == 14
    assert c25.EXPECTED_FOLD_CHANGE_RECORDS_FIG6B == 2
    assert c25.EXPECTED_FOLD_CHANGE_RECORDS == 16


def test_the_sourced_scale_basis_and_test_quote_mirrored_bytes() -> None:
    """Every scalar the phenotype carries is bound to a verbatim statement."""
    assert c25.FOLD_CHANGE_SCALE.value == "log2"
    assert "Log2(Fold-change)" in str(c25.FOLD_CHANGE_SCALE.quote)
    assert c25.FOLD_CHANGE_REFERENCE_BASIS.value == "the non-target control strain"
    assert c25.FOLD_CHANGE_REFERENCE_BASIS.value in str(
        c25.FOLD_CHANGE_REFERENCE_BASIS.quote
    )
    assert c25.KO_FOLD_CHANGE_REFERENCE_BASIS.value == (
        "a non-targeting control sgRNA in a knockout background"
    )
    assert c25.KO_FOLD_CHANGE_REFERENCE_BASIS.value in str(
        c25.KO_FOLD_CHANGE_REFERENCE_BASIS.quote
    )
    assert c25.FOLD_CHANGE_TEST.value == c25.FOLD_CHANGE_MEASUREMENT_TYPE
    assert c25.BEST_ARRAY_HEATMAP_REFUSAL.value is None
    assert "significantly changed" in str(c25.BEST_ARRAY_HEATMAP_REFUSAL.quote)


# --------------------------------------------------------------------------- #
# Synthetic: the fold-change verifier
# --------------------------------------------------------------------------- #
def _record(
    *,
    tags: tuple[str, ...] = ("PP_0001",),
    values: dict[str, float] | None = None,
    p_values: dict[str, float] | None = None,
    basis: str = "the non-target control strain",
    reference: dict[str, float] | None = None,
    scale: str = "log2",
    measurement_type: str = c25.FOLD_CHANGE_MEASUREMENT_TYPE,
) -> dict[str, Any]:
    """One dumped fold-change record: a real experiment plus a dict reference."""
    fold_change = values if values is not None else {"PP_0001": -1.25}
    phenotype = ProteinFoldChangePhenotype(
        protein_fold_change=fold_change,
        fold_change_scale=FoldChangeScale(scale),
        reference_basis=basis,
        protein_fold_change_p_value=(
            p_values if p_values is not None else dict.fromkeys(fold_change, 0.01)
        ),
        n_replicates=dict.fromkeys(fold_change, 3),
        measurement_type=measurement_type,
    )
    experiment = BacterialProteinFoldChangeExperiment(
        dataset_name="proteome_fold_change_carruthers2025",
        genotype=Genotype(
            perturbations=[
                *c25.pathway_perturbations(),
                *(c25.crispri_perturbation(tag, tag) for tag in tags),
            ]
        ),
        environment=c25.production_environment(),
        phenotype=phenotype,
    )
    neutral = reference if reference is not None else phenotype.neutral_reference()
    return {
        "experiment": experiment.model_dump(),
        "reference": {
            "genome_reference": {"species": "Pseudomonas putida"},
            "phenotype_reference": {
                "protein_fold_change": neutral,
                "fold_change_scale": scale,
            },
        },
    }


def _row(report: Any, name: str) -> Any:
    (match,) = [result for result in report.results if result.name == name]
    return match


def test_the_verifier_passes_a_well_formed_pair_of_contrasts() -> None:
    """The eight L0-L3 rows the fold-change family asserts."""
    records = [_record(tags=("PP_0001",)), _record(tags=("PP_0002",))]
    report = verify_protein_fold_change_dataset(
        records, dataset_name="fc", provenance=PROV, expected_count=2
    )
    assert report.passed is True
    assert [result.name for result in report.results] == [
        "structural",
        "count",
        "contrast_uniqueness",
        "value_fidelity",
        "p_values_are_probabilities",
        "reference_is_the_scales_neutral_value",
        "fold_change_scale_consistent",
        "measurement_type_consistent",
    ]
    assert _row(report, "structural").level is Level.L0
    assert _row(report, "contrast_uniqueness").details["n_contrasts"] == 2


def test_the_verifier_fails_one_contrast_stored_twice() -> None:
    """Two records of one contrast are the same measurement stored twice."""
    report = verify_protein_fold_change_dataset(
        [_record(), _record()], dataset_name="fc", provenance=PROV, expected_count=2
    )
    assert _row(report, "contrast_uniqueness").passed is False
    assert _row(report, "contrast_uniqueness").details["n_duplicated"] == 1


def test_the_verifier_separates_two_contrasts_by_their_denominator() -> None:
    """One genotype against two different bases is two contrasts, not a duplicate."""
    report = verify_protein_fold_change_dataset(
        [
            _record(),
            _record(basis="a non-targeting control sgRNA in a knockout background"),
        ],
        dataset_name="fc",
        provenance=PROV,
        expected_count=2,
    )
    assert _row(report, "contrast_uniqueness").passed is True


def test_the_verifier_separates_two_contrasts_by_their_environment() -> None:
    """A wild-type panel varies only the environment, and those are distinct contrasts."""
    first = _record()
    second = _record()
    second["experiment"]["environment"]["temperature"]["value"] = 30.0
    report = verify_protein_fold_change_dataset(
        [first, second], dataset_name="fc", provenance=PROV, expected_count=2
    )
    assert _row(report, "contrast_uniqueness").passed is True
    assert _row(report, "contrast_uniqueness").details["n_contrasts"] == 2


def test_the_verifier_fails_a_p_value_of_exactly_zero() -> None:
    """The schema admits [0, 1]; no test reports a probability of 0."""
    report = verify_protein_fold_change_dataset(
        [_record(p_values={"PP_0001": 0.0})],
        dataset_name="fc",
        provenance=PROV,
        expected_count=1,
    )
    assert _row(report, "p_values_are_probabilities").passed is False
    assert _row(report, "p_values_are_probabilities").details == {
        "n_values": 1,
        "n_bad": 1,
    }


def test_the_verifier_fails_a_reference_that_is_not_the_neutral_value() -> None:
    """A denominator of 1.0 on a log2 scale would not reproduce the released number."""
    report = verify_protein_fold_change_dataset(
        [_record(reference={"PP_0001": 1.0})],
        dataset_name="fc",
        provenance=PROV,
        expected_count=1,
    )
    assert _row(report, "reference_is_the_scales_neutral_value").passed is False


def test_the_verifier_fails_a_reference_whose_keys_do_not_match() -> None:
    """A fold change with no denominator key is not referenced at all."""
    report = verify_protein_fold_change_dataset(
        [_record(reference={"PP_9999": 0.0})],
        dataset_name="fc",
        provenance=PROV,
        expected_count=1,
    )
    assert _row(report, "reference_is_the_scales_neutral_value").passed is False


def test_the_verifier_fails_two_measurement_types_mixed() -> None:
    """One statistic per dataset, or two contrasts pool into nothing."""
    report = verify_protein_fold_change_dataset(
        [
            _record(tags=("PP_0001",)),
            _record(tags=("PP_0002",), measurement_type="other"),
        ],
        dataset_name="fc",
        provenance=PROV,
        expected_count=2,
    )
    assert _row(report, "measurement_type_consistent").passed is False
    assert _row(report, "fold_change_scale_consistent").passed is True


def test_the_locus_set_is_the_tested_proteins_and_the_perturbed_host_genes() -> None:
    """The heterologous pathway tokens are excluded; the host identifiers are not."""
    record = _record(tags=("PP_0001",), values={"PP_0500": 1.0, "PP_0501": -1.0})
    assert protein_fold_change_locus_set([record]) == {"PP_0001", "PP_0500", "PP_0501"}


def test_the_p_value_round_trip_helper_inverts_the_loaders_own_conversion() -> None:
    """The oracle is the inverse of the conversion, not a re-derivation."""
    assert fold_change_p_value_round_trip(2.0, 0.01) is True
    assert fold_change_p_value_round_trip(2.0, 0.011) is False
    assert fold_change_p_value_round_trip(2.0, 0.0) is False
    assert fold_change_p_value_round_trip(float("inf"), 0.01) is False


# --------------------------------------------------------------------------- #
# Data-gated: the real workbook and the built store
# --------------------------------------------------------------------------- #
DATA_ROOT = os.environ.get("DATA_ROOT", "")
FOLD_CHANGE_ROOT = osp.join(
    DATA_ROOT, "data/torchcell/proteome_fold_change_carruthers2025"
)
SOURCE_DATA = osp.join(DATA_ROOT, c25.RAW_DIR_REL, c25.SOURCE_DATA_REL)
MIRROR_PRESENT = bool(DATA_ROOT) and osp.isfile(SOURCE_DATA)
BUILT = osp.isdir(osp.join(FOLD_CHANGE_ROOT, "processed/lmdb"))

requires_mirror = pytest.mark.skipif(
    not MIRROR_PRESENT,
    reason="requires the Carruthers 2025 raw mirror under $DATA_ROOT",
)
requires_built = pytest.mark.skipif(
    not BUILT,
    reason="requires the built Carruthers 2025 fold-change LMDB under $DATA_ROOT",
)


@pytest.mark.data
@requires_mirror
def test_the_released_single_guide_sheet_has_the_measured_shape() -> None:
    """14 contrast columns and 5,770 fold changes over 1,290 keyed rows."""
    rows = c25.read_fold_change_rows(SOURCE_DATA, c25.SHEET_FOLD_CHANGE)
    contrasts = c25.fold_change_contrasts(SOURCE_DATA, c25.SHEET_FOLD_CHANGE)
    assert contrasts == (
        "PP_0368",
        "PP_0437",
        "PP_0528",
        "PP_0812",
        "PP_0813",
        "PP_0814",
        "PP_0815",
        "PP_1317",
        "PP_1506",
        "PP_2136",
        "PP_4120",
        "PP_4189",
        "PP_4191",
        "PP_4192",
    )
    assert len(rows) == 5770
    assert len({row.protein for row in rows}) == 1290
    assert max(row.neg_log10_p_value for row in rows) == 19.3125325679514
    assert min(row.neg_log10_p_value for row in rows) == 1.30103425637741


@pytest.mark.data
@requires_mirror
def test_the_released_ko_sheet_has_the_measured_shape() -> None:
    """10 contrast columns and 452 fold changes over 372 keyed rows."""
    rows = c25.read_fold_change_rows(SOURCE_DATA, c25.SHEET_KO_FOLD_CHANGE)
    contrasts = c25.fold_change_contrasts(SOURCE_DATA, c25.SHEET_KO_FOLD_CHANGE)
    assert len(contrasts) == 10
    assert set(contrasts) == set(c25.KO_FOLD_CHANGE_CONTRASTS) | set(
        c25.KO_FOLD_CHANGE_UNSOURCED_REASONS
    )
    assert len(rows) == 452
    assert len({row.protein for row in rows}) == 372
    per_contrast = {
        contrast: sum(1 for row in rows if row.contrast == contrast)
        for contrast in contrasts
    }
    assert per_contrast["PP_0812"] == 14
    assert per_contrast["PP_0815"] == 220


@pytest.mark.data
@requires_mirror
def test_the_source_data_points_supplementary_figure_9_at_figure_5b() -> None:
    """The one cell that sources the single-guide panel's denominator."""
    assert "See Figure 5b" in c25.assert_fold_change_pointer(SOURCE_DATA)


@pytest.mark.data
@requires_mirror
def test_the_refused_heatmap_measures_as_recorded() -> None:
    """26 keys x 25 strain columns, dense, signed, with no p-value column."""
    keys, strains, values = c25.read_best_array_heatmap(SOURCE_DATA)
    assert len(keys) == 26
    assert len(strains) == 25
    assert len(values) == 650
    assert min(values) == -7.05757655263963
    assert max(values) == 3.92161053900316
    assert statistics.median(values) == -1.0118185429312652
    assert sum(1 for key in keys if c25.LOCUS_TAG_RE.fullmatch(key)) == 2


@pytest.mark.data
@requires_mirror
def test_every_single_guide_contrast_is_a_released_single_guide_construct() -> None:
    """What licenses the genotype each Figure 5b record is written with."""
    proofs = c25.assert_fold_change_contrasts_are_released_strains(SOURCE_DATA)
    assert proofs[0] == (
        "all 14 Figure 5b contrast columns are single PP_ locus tags released as "
        "single-guide Figure 4b constructs"
    )
    joined = {line.split(":")[0]: line for line in proofs[1:]}
    assert (
        "Figure 6a KO background False, Figure 4b line name False"
        in joined["Figure 6b column 'PP_1317_PP_0812'"]
    )
    assert (
        "Figure 6a KO background False, Figure 4b line name True"
        in joined["Figure 6b column 'PP_0751_PP_0812'"]
    )
    assert (
        "Figure 6a KO background True, Figure 4b line name False"
        in joined["Figure 6b column 'PP_0812_15'"]
    )


@pytest.mark.data
@requires_built
def test_the_built_store_holds_one_record_per_sourced_contrast() -> None:
    """16 records: 14 single-guide CRISPRi strains and the 2 sourced KO strains."""
    from torchcell.verification.runners import load_records

    records = load_records(FOLD_CHANGE_ROOT)
    assert len(records) == 16
    by_contrast = {c25._fold_change_contrast(record): record for record in records}
    assert len(by_contrast) == 16
    assert sum(1 for sheet, _ in by_contrast if sheet == c25.SHEET_FOLD_CHANGE) == 14
    assert {
        contrast for sheet, contrast in by_contrast if sheet == c25.SHEET_KO_FOLD_CHANGE
    } == {"PP_0812", "PP_0815"}
    total = sum(
        len(record["experiment"]["phenotype"]["protein_fold_change"])
        for record in records
    )
    assert total == 5958


@pytest.mark.data
@requires_built
def test_a_built_ko_record_carries_the_deletion_beside_its_guide() -> None:
    """The KO panel's strain is the gene deleted AND its own sgRNA present."""
    from torchcell.verification.runners import load_records

    records = load_records(FOLD_CHANGE_ROOT)
    (record,) = [
        record
        for record in records
        if c25._fold_change_contrast(record) == (c25.SHEET_KO_FOLD_CHANGE, "PP_0812")
    ]
    types = {
        (p["perturbation_type"], p["systematic_gene_name"])
        for p in record["experiment"]["genotype"]["perturbations"]
    }
    assert ("bacterial_deletion", "PP_0812") in types
    assert ("bacterial_crispr_interference", "PP_0812") in types
    assert len(record["experiment"]["phenotype"]["protein_fold_change"]) == 14
    assert record["experiment"]["phenotype"]["reference_basis"] == (
        "a non-targeting control sgRNA in a knockout background"
    )


@pytest.mark.data
@requires_built
def test_a_built_single_guide_record_carries_only_its_guide() -> None:
    """A Figure 5b strain has no chromosomal deletion and one CRISPRi leaf."""
    from torchcell.verification.runners import load_records

    records = load_records(FOLD_CHANGE_ROOT)
    (record,) = [
        record
        for record in records
        if c25._fold_change_contrast(record) == (c25.SHEET_FOLD_CHANGE, "PP_2136")
    ]
    types = [
        p["perturbation_type"]
        for p in record["experiment"]["genotype"]["perturbations"]
    ]
    assert types.count("bacterial_crispr_interference") == 1
    assert types.count("bacterial_deletion") == 0
    assert types.count("heterologous_pathway") == 5
    assert len(record["experiment"]["phenotype"]["protein_fold_change"]) == 176
    assert record["experiment"]["phenotype"]["reference_basis"] == (
        "the non-target control strain"
    )
    assert set(
        record["reference"]["phenotype_reference"]["protein_fold_change"].values()
    ) == {0.0}


@pytest.mark.data
@requires_built
def test_the_accounting_names_every_refused_contrast_and_the_refused_sheet() -> None:
    """16 kept + 8 refused KO columns + 25 refused heatmap columns = 49 candidates."""
    import json

    accounting = json.loads(
        Path(osp.join(FOLD_CHANGE_ROOT, "preprocess/build_accounting.json")).read_text()
    )
    assert accounting["kept_records"] == 16
    assert accounting["dropped_records"] == 33
    assert accounting["candidate_records"] == 49
    assert accounting["source_rows"] == 6222
    assert accounting["control_rows"] == 0
    rules = {rule["rule"]: rule for rule in accounting["rules"]}
    assert rules["ko_fold_change_contrast_has_no_sourced_denominator"]["n_records"] == 8
    assert rules["best_array_heatmap_states_no_scale"]["n_records"] == 25
    assert rules["protein_key_is_not_a_locus_of_the_pinned_assembly"]["n_records"] == 0
    assert (
        len(rules["protein_key_is_not_a_locus_of_the_pinned_assembly"]["items"]) == 19
    )


@pytest.mark.data
@requires_built
def test_the_fold_change_report_passes_every_level() -> None:
    """L0 to L4 over the built store, with the rows this family asserts."""
    from torchcell.verification.runners import load_records

    records = load_records(FOLD_CHANGE_ROOT)
    report = c25.fold_change_report(records, DATA_ROOT)
    assert report.passed is True
    assert [result.name for result in report.results] == [
        "structural",
        "count",
        "contrast_uniqueness",
        "value_fidelity",
        "p_values_are_probabilities",
        "reference_is_the_scales_neutral_value",
        "fold_change_scale_consistent",
        "measurement_type_consistent",
        "fold_change_contrast_coverage",
        "fold_change_biological_triplicate",
        "stored_fold_changes_vs_released_sheets",
        "stored_p_values_invert_to_released_neg_log10",
    ]
    assert _row(report, "value_fidelity").details["n_values"] == 5958
    assert _row(report, "stored_fold_changes_vs_released_sheets").level is Level.L4
