# experiments/036-dataset-fixes-before-kg-build/scripts/piotrowski2017_release_inventory.py
# [[experiments.036-dataset-fixes-before-kg-build.scripts.piotrowski2017_release_inventory]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/036-dataset-fixes-before-kg-build/scripts/piotrowski2017_release_inventory.py
"""Measure the Piotrowski 2017 release: the matrix, its metadata, the pool, Data Set 6.

Every release-level number in the 2026.10.11 section of
``[[torchcell.datasets.scerevisiae.piotrowski2017]]`` that is not a build-ledger or
verifier row comes from this script. It reads only sha256-pinned bytes through the
loader's own readers: Data Sets 5, 4 and 2 from the raw mirror and Data Set 6
(``si/si14.xlsx``) from the literature mirror.

Writes ``experiments/036-dataset-fixes-before-kg-build/results/piotrowski2017_release_inventory.json``.

    PYTHONPATH=<worktree> python experiments/036-dataset-fixes-before-kg-build/scripts/piotrowski2017_release_inventory.py
"""

from __future__ import annotations

import os
import os.path as osp
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
from dotenv import load_dotenv
from pydantic import BaseModel, ConfigDict

from torchcell.datasets.scerevisiae import piotrowski2017 as p
from torchcell.literature.manifest import sha256_file

RESULTS = osp.join(
    osp.dirname(osp.dirname(osp.abspath(__file__))),
    "results",
    "piotrowski2017_release_inventory.json",
)
#: Data Set 6, the 1,522 high-confidence profiles (library mirror path and pin).
HCS_REL = "si/si14.xlsx"
HCS_SHA256 = "b07029062cb404a7edaf98e932b94f8b992446672b826ce257a84965aa6a55ae"


class MatrixCensus(BaseModel):
    """Data Set 5 as released."""

    model_config = ConfigDict(extra="forbid")

    n_strains: int
    n_conditions: int
    n_cells: int
    n_numeric: int
    n_na: int
    na_columns_per_strain: dict[str, int]
    columns_by_screen: dict[str, int]
    label_differs_from_data_set_4: dict[str, str]


class ConditionCensus(BaseModel):
    """Data Set 4 and how it covers the matrix."""

    model_config = ConfigDict(extra="forbid")

    n_rows: int
    n_in_final_dataset: int
    n_high_confidence: int
    matrix_columns_not_final: list[str]
    final_not_in_matrix: list[str]
    matrix_columns_by_collection: dict[str, int]
    dose_kinds: dict[str, int]
    range_doses: list[str]
    same_name_and_dose_groups: int
    same_name_dose_and_collection_groups: list[str]
    matrix_columns_without_smiles: int
    matrix_columns_without_inchi: int


class PoolCensus(BaseModel):
    """Data Set 2's pool version 2.0 against the matrix rows."""

    model_config = ConfigDict(extra="forbid")

    n_pool_v2: int
    n_matrix_strains_in_pool: int
    pool_strains_not_in_matrix: list[str]
    sensitized_fitness_sheet_genes: int


class HcsCheck(BaseModel):
    """Data Set 6 compared cell by cell with the same cells of Data Set 5."""

    model_config = ConfigDict(extra="forbid")

    sha256: str
    n_conditions: int
    n_strains: int
    n_cells_compared: int
    n_cells_unequal: int


class Inventory(BaseModel):
    """Everything this script measures."""

    model_config = ConfigDict(extra="forbid")

    matrix: MatrixCensus
    conditions: ConditionCensus
    pool: PoolCensus
    high_confidence_set: HcsCheck


def sensitized_fitness_genes(path: Path) -> int:
    """Data Set 2's second sheet: genes with an original and a sensitized SGA fitness."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["Fitness of WT vs sensitized mut"].iter_rows(values_only=True))
    workbook.close()
    if tuple(rows[2][:4]) != (
        "Gene",
        "Original SMF",
        "Hypersensitive SMF",
        "Difference",
    ):
        raise ValueError(f"unexpected header {rows[2][:4]!r}")
    return sum(1 for row in rows[3:] if row[0] is not None)


def hcs_check(path: Path, matrix: p.ReleasedMatrix) -> HcsCheck:
    """Compare every Data Set 6 cell with the Data Set 5 cell of the same strain and id."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    workbook.close()
    header = rows[2]
    column = {cid: j for j, cid in enumerate(matrix.condition_ids)}
    strain = {label: i for i, label in enumerate(matrix.strains)}
    compared = 0
    unequal = 0
    ids = [str(label).split("_", 1)[0] for label in header[1:]]
    n_strains = 0
    for row in rows[3:]:
        if row[0] is None:
            continue
        n_strains += 1
        i = strain[str(row[0])]
        for k, cid in enumerate(ids):
            cell = row[k + 1]
            released = matrix.values[i, column[cid]]
            compared += 1
            if cell == p.NOT_SCORED:
                unequal += int(not np.isnan(released))
            else:
                unequal += int(float(cell) != released)
    return HcsCheck(
        sha256=sha256_file(path),
        n_conditions=len(ids),
        n_strains=n_strains,
        n_cells_compared=compared,
        n_cells_unequal=unequal,
    )


def main() -> None:
    """Measure the release and write the inventory JSON."""
    load_dotenv()
    data_root = os.environ["DATA_ROOT"]
    mirror = p.raw_mirror_dir(data_root)
    for release in p.RELEASE_FILES:
        if sha256_file(mirror / release.mirror_rel) != release.sha256:
            raise RuntimeError(f"{release.mirror_rel} does not match its pin")
    matrix = p.read_matrix(mirror / p.MATRIX.mirror_rel)
    conditions = p.read_conditions(mirror / p.CONDITIONS.mirror_rel)
    pool = p.read_pool(mirror / p.POOL.mirror_rel)
    hcs_path = Path(data_root) / p.LIBRARY_DIR_REL / HCS_REL
    if sha256_file(hcs_path) != HCS_SHA256:
        raise RuntimeError(f"{hcs_path} does not match its pin")

    rows = [conditions[cid] for cid in matrix.condition_ids]
    na = np.isnan(matrix.values)
    doses = [
        p.parse_dose(r.concentration_cell, r.concentration_number_format) for r in rows
    ]
    name_dose = Counter((r.name, d.source) for r, d in zip(rows, doses, strict=True))
    name_dose_lib = Counter(
        (r.name, d.source, r.collection) for r, d in zip(rows, doses, strict=True)
    )
    in_matrix = set(matrix.condition_ids)
    labels = set(matrix.strains)
    inventory = Inventory(
        matrix=MatrixCensus(
            n_strains=len(matrix.strains),
            n_conditions=len(matrix.condition_ids),
            n_cells=int(na.size),
            n_numeric=int((~na).sum()),
            n_na=int(na.sum()),
            na_columns_per_strain={
                matrix.strains[i]: int(n) for i, n in enumerate(na.sum(axis=1)) if n > 0
            },
            columns_by_screen=dict(Counter(r.screen for r in rows)),
            label_differs_from_data_set_4={
                cid: f"{label} | {conditions[cid].name}"
                for cid, label in zip(
                    matrix.condition_ids, matrix.condition_labels, strict=True
                )
                if label != conditions[cid].name
            },
        ),
        conditions=ConditionCensus(
            n_rows=len(conditions),
            n_in_final_dataset=sum(r.in_final_dataset for r in conditions.values()),
            n_high_confidence=sum(
                r.in_high_confidence_set for r in conditions.values()
            ),
            matrix_columns_not_final=[
                r.condition_id for r in rows if not r.in_final_dataset
            ],
            final_not_in_matrix=sorted(
                cid
                for cid, r in conditions.items()
                if r.in_final_dataset and cid not in in_matrix
            ),
            matrix_columns_by_collection=dict(Counter(r.collection for r in rows)),
            dose_kinds=dict(Counter(d.kind for d in doses)),
            range_doses=[
                f"{r.condition_id} {r.name} ({d.source})"
                for r, d in zip(rows, doses, strict=True)
                if d.kind == "range"
            ],
            same_name_and_dose_groups=sum(1 for n in name_dose.values() if n > 1),
            same_name_dose_and_collection_groups=sorted(
                f"{name} {dose} {lib} x{n}"
                for (name, dose, lib), n in name_dose_lib.items()
                if n > 1
            ),
            matrix_columns_without_smiles=sum(1 for r in rows if r.smiles is None),
            matrix_columns_without_inchi=sum(1 for r in rows if r.inchi is None),
        ),
        pool=PoolCensus(
            n_pool_v2=len(pool),
            n_matrix_strains_in_pool=sum(
                1 for e in pool if e.common_name in labels or e.orf in labels
            ),
            pool_strains_not_in_matrix=sorted(
                f"{e.common_name} ({e.orf})"
                for e in pool
                if e.common_name not in labels and e.orf not in labels
            ),
            sensitized_fitness_sheet_genes=sensitized_fitness_genes(
                mirror / p.POOL.mirror_rel
            ),
        ),
        high_confidence_set=hcs_check(hcs_path, matrix),
    )
    os.makedirs(osp.dirname(RESULTS), exist_ok=True)
    with open(RESULTS, "w") as handle:
        handle.write(inventory.model_dump_json(indent=2))
    print(inventory.model_dump_json(indent=2))


if __name__ == "__main__":
    main()
