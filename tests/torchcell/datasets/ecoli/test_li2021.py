# tests/torchcell/datasets/ecoli/test_li2021.py
# [[tests.torchcell.datasets.ecoli.test_li2021]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_li2021.py
"""Li 2021's three released 13C-MFA maps: the reader, the records and the gapped reference.

The hermetic tests build a workbook with the pinned NETWORK layout in ``tmp_path``; the
``data`` tests read the deposited Additional file 2 and the built dev store.
"""

from __future__ import annotations

import json
import os
import os.path as osp
from pathlib import Path

import pytest
from openpyxl import Workbook

from torchcell.data.experiment_dataset import file_sha256
from torchcell.datamodels.schema import (
    FluxPhenotype,
    PromoterReplacementPerturbation,
    StrainReferenceGenome,
)
from torchcell.datasets.ecoli import li2021 as m
from torchcell.verification.sourced import ProvenanceGap, ProvenanceGapReason

REACTIONS = (("vUPTU", 1.5, 1.5, 1.5), ("vGPI", -0.4, -0.5, -0.3))


def _workbook(path: Path, *, exchange_rows: int = 1) -> Path:
    """A NETWORK sheet with the pinned layout, two reactions and an exchange table."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "NETWORK"
    for offset, label in enumerate(m.DIAGNOSTIC_LABELS):
        ws[f"B{offset + 2}"] = label
    for index, spec in enumerate(m.STRAINS):
        ws[spec.label_cell] = spec.block_label
        label_column = spec.label_cell[0]
        for offset in range(len(m.DIAGNOSTIC_LABELS)):
            ws[f"{label_column}{offset + 2}"] = 10.0 * (index + 1) + offset
        ws[spec.unit_cell] = m.UNIT_HEADER
        for column, header in zip(spec.absolute_columns, m.BOUND_HEADERS, strict=True):
            ws[f"{column}8"] = header
        for row, (name, best, lower, upper) in enumerate(REACTIONS, start=11):
            ws[f"B{row}"] = name
            for column, value in zip(
                spec.absolute_columns, (best, lower, upper), strict=True
            ):
                ws[f"{column}{row}"] = value + index
        for k in range(exchange_rows):
            row = spec.exchange_first_row + k
            ws[f"{spec.exchange_name_column}{row}"] = f"vX{k}.x"
            for column, value in zip(
                spec.exchange_columns, (0.5, 0.1, 0.9), strict=True
            ):
                ws[f"{column}{row}"] = value
    ws["B10"] = "//  Uptake of substrates"
    column, row, text = m.EXCHANGE_HEADER_CELL
    ws[f"{column}{row}"] = text
    out = path / "network.xlsx"
    wb.save(out)
    return out


def test_read_block_keeps_every_number_verbatim_and_signed(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    ws = load_workbook(_workbook(tmp_path), data_only=True)["NETWORK"]
    rows = m.reaction_rows(ws)
    assert rows == [(11, "vUPTU"), (12, "vGPI")]
    block = m.read_block(ws, m.STRAINS[0], rows)
    assert block.net_flux == {"vUPTU": 1.5, "vGPI": -0.4}
    assert block.net_flux_lower == {"vUPTU": 1.5, "vGPI": -0.5}
    assert block.net_flux_upper == {"vUPTU": 1.5, "vGPI": -0.3}
    assert block.exchange_flux == {"vX0.x": 0.5}
    assert block.diagnostics["χ2 90%"] == 12.0


def test_read_workbook_refuses_a_reaction_count_that_is_not_pinned(
    tmp_path: Path,
) -> None:
    with pytest.raises(m.WorkbookFormatError, match="2 reaction rows"):
        m.read_workbook(_workbook(tmp_path))


def test_read_block_refuses_a_moved_unit_header(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = _workbook(tmp_path)
    wb = load_workbook(path)
    wb["NETWORK"][m.STRAINS[1].unit_cell] = "Fluxes normalized"
    wb.save(path)
    ws = load_workbook(path, data_only=True)["NETWORK"]
    with pytest.raises(m.WorkbookFormatError, match="L7"):
        m.read_block(ws, m.STRAINS[1], m.reaction_rows(ws))


def test_read_block_refuses_a_text_cell_rather_than_inventing_a_value(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    path = _workbook(tmp_path)
    wb = load_workbook(path)
    wb["NETWORK"]["E11"] = "n.d."
    wb.save(path)
    ws = load_workbook(path, data_only=True)["NETWORK"]
    with pytest.raises(m.WorkbookFormatError, match="E11"):
        m.read_block(ws, m.STRAINS[0], m.reaction_rows(ws))


def _block(label: str, shift: float = 0.0) -> m.BlockMap:
    return m.BlockMap(
        block_label=label,
        net_flux={"a": 1.0 + shift, "b": -2.0},
        net_flux_lower={"a": 0.5 + shift, "b": -2.0},
        net_flux_upper={"a": 1.5 + shift, "b": -2.0},
        exchange_flux={"a.x": 0.3},
        exchange_flux_lower={"a.x": 0.0},
        exchange_flux_upper={"a.x": 0.9},
        diagnostics={},
    )


def test_identical_rows_counts_rows_equal_in_value_and_both_bounds() -> None:
    assert m.identical_rows(_block("x"), _block("y")) == ["a", "b"]
    assert m.identical_rows(_block("x"), _block("y", shift=0.1)) == ["b"]


def test_flux_phenotype_stores_the_block_at_the_stated_level() -> None:
    phenotype = m.flux_phenotype(_block("BW-P08"))
    assert phenotype.net_flux == {"a": 1.0, "b": -2.0}
    assert phenotype.exchange_flux_upper == {"a.x": 0.9}
    assert phenotype.confidence_level == 0.90
    assert phenotype.measurement_type == "c13_mfa_net_flux_mmol_per_gdw_per_h"
    assert phenotype.gapped_fields() == {"n_samples", "sample_unit"}


def test_reference_map_is_a_typed_gap_not_a_refusal() -> None:
    reference = m.reference_phenotype()
    assert reference.net_flux is None
    assert "net_flux" in reference.gapped_fields()


def test_a_none_map_without_the_gap_is_refused() -> None:
    with pytest.raises(ValueError, match="typed gap"):
        FluxPhenotype(net_flux=None, measurement_type="x")


def test_a_gapped_map_refuses_bounds_and_exchange_values() -> None:
    gap = ProvenanceGap(
        field="net_flux", reason=ProvenanceGapReason.not_reported_by_primary
    )
    with pytest.raises(ValueError, match="exchange_flux is set on a gapped"):
        FluxPhenotype(
            net_flux=None,
            exchange_flux={"a.x": 0.1},
            measurement_type="x",
            provenance_gaps=[gap],
        )


def test_exchange_bounds_must_key_on_a_stored_exchange() -> None:
    with pytest.raises(ValueError, match="not in exchange_flux"):
        FluxPhenotype(
            net_flux={"a": 1.0},
            exchange_flux={"a.x": 0.1},
            exchange_flux_lower={"b.x": 0.0},
            measurement_type="x",
        )
    with pytest.raises(ValueError, match="without exchange_flux"):
        FluxPhenotype(
            net_flux={"a": 1.0}, exchange_flux_upper={"a.x": 0.2}, measurement_type="x"
        )


def test_genotype_is_the_strains_zwf_promoter_swap_under_the_released_symbol() -> None:
    [perturbation] = m.genotype(m.STRAINS[2]).perturbations
    assert isinstance(perturbation, PromoterReplacementPerturbation)
    assert perturbation.systematic_gene_name == "BW25113_1852"
    assert perturbation.perturbed_gene_name == "zwf"
    assert perturbation.promoter_name == "BBa-J 23100"
    assert perturbation.expression_direction == "increased"
    assert perturbation.identifier_mapping is not None
    assert perturbation.identifier_mapping.source_identifier == "zwf"


def test_labeling_environment_states_the_tracer_and_gaps_the_rest() -> None:
    environment = m.labeling_environment()
    assert environment.temperature is None
    assert environment.gapped_fields() == {"temperature"}
    assert environment.media.gapped_fields() == {"name", "state", "is_synthetic"}
    [component] = environment.media.components
    assert component.concentration is not None
    assert component.concentration.value == 10.0


def test_reference_background_is_bw_p_bf_stated_verbatim() -> None:
    background = m.reference_background()
    assert background.name == "BW-P BF"
    assert background.genotype_statement == "BW25113 derivative, ∆pfkA:FRT"
    assert background.alleles == []


def _data_root() -> str:
    return os.environ.get("DATA_ROOT", "")


@pytest.mark.data
def test_released_workbook_reads_three_blocks_of_99_reactions() -> None:
    blocks = m.read_workbook(m.raw_mirror_dir(_data_root()) / m.WORKBOOK_RELPATH)
    assert [b.block_label for b in blocks] == ["BW-P08", "BW-P04", "BW-P10"]
    assert [len(b.net_flux) for b in blocks] == [99, 99, 99]
    assert [len(b.exchange_flux) for b in blocks] == [24, 23, 23]
    by = {b.block_label: b for b in blocks}
    assert by["BW-P08"].net_flux["vGPI"] == -0.445077375
    assert len(m.identical_rows(by["BW-P04"], by["BW-P10"])) == 50
    assert len(m.identical_rows(by["BW-P08"], by["BW-P10"])) == 0


@pytest.mark.data
def test_every_paper_quote_is_verbatim_in_the_mirror() -> None:
    from torchcell.verification.sourced import audit_sourced_value

    library = Path(_data_root()) / "torchcell-library"
    values = [*m.SOURCED_VALUES.values()]
    values += [v for s in m.STRAINS for v in (s.table_1, s.promoter_strain)]
    assert all(audit_sourced_value(v, library).passed for v in values)


@pytest.mark.data
def test_dev_store_verifies() -> None:
    root = osp.join(_data_root(), m.DATASET_ROOT_REL)
    if not osp.isdir(osp.join(root, "processed")):
        pytest.skip("flux_li2021 dev store not built")
    report = m.verify_build(root, _data_root())
    assert report.passed, report.summary()


def test_exchange_values_must_be_finite() -> None:
    with pytest.raises(ValueError, match="must be finite"):
        FluxPhenotype(
            net_flux={"a": 1.0},
            exchange_flux={"a.x": float("nan")},
            measurement_type="x",
        )


# --------------------------------------------------------------------------- #
# The hermetic build: a synthetic library copy, deposit, download, process, verify
# --------------------------------------------------------------------------- #
REFERENCE = StrainReferenceGenome(
    species="Escherichia coli",
    strain="BW-P BF",
    assembly_set="ecoli_K12_BW25113_ASM75055v1",
    assembly_accession="GCA_000750555.1",
    background=m.reference_background(),
)


@pytest.fixture
def mirror(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A data root whose library holds a synthetic workbook, pinned by its own hash."""
    data_root = tmp_path / "data_root"
    library = data_root / m.LIBRARY_DIR_REL / "si"
    library.mkdir(parents=True)
    source = _workbook(tmp_path)
    target = library / "si2.xlsx"
    target.write_bytes(source.read_bytes())
    monkeypatch.setattr(m, "WORKBOOK_SHA256", file_sha256(target))
    monkeypatch.setattr(m, "EXPECTED_REACTIONS", 2)
    monkeypatch.setattr(m, "EXPECTED_EXCHANGES", {s.block_label: 1 for s in m.STRAINS})
    monkeypatch.setattr(
        m, "assembly_reference", lambda strain, background=None: REFERENCE
    )
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    return data_root


def test_deposit_is_idempotent_and_writes_the_manifest(mirror: Path) -> None:
    root = m.deposit_raw_mirror(str(mirror))
    assert (root / m.WORKBOOK_RELPATH).exists()
    assert m.manifest_sha256(str(mirror)) == m.WORKBOOK_SHA256
    assert m.deposit_raw_mirror(str(mirror)) == root


def test_hermetic_build_writes_three_records_and_verifies(
    mirror: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from torchcell.verification import runners
    from torchcell.verification.report import Level, LevelResult

    m.deposit_raw_mirror(str(mirror))
    root = mirror / m.DATASET_ROOT_REL
    dataset = m.FluxLi2021Dataset(root=str(root))
    assert len(dataset) == 3
    not_stored = json.loads((root / "preprocess" / "not_stored.json").read_text())
    assert not_stored["not_stored"][0]["count"] == 3 * 2 * 3
    findings = json.loads((root / "preprocess" / "release_findings.json").read_text())
    assert findings["control_map_released"] is False

    monkeypatch.setattr(
        runners, "_bacterial_gene_set", lambda assembly, base: {"BW25113_1852"}
    )
    monkeypatch.setattr(
        m,
        "audit_sourced_value",
        lambda value, root: LevelResult(
            level=Level.L3, name="provenance_audit", passed=True, message="stub"
        ),
    )
    report = m.verify_build(str(root), str(mirror))
    assert report.passed, report.summary()
    assert (root / "preprocess" / "verification_report.json").exists()
    assert m.main(["verify"]) == 0


def test_verifier_rows_fail_on_a_moved_value_and_a_missing_gap() -> None:
    blocks = [_block(spec.block_label) for spec in m.STRAINS]
    records = [
        {
            "experiment": {
                "genotype": {"perturbations": [{"promoter_name": spec.promoter}]},
                "phenotype": m.flux_phenotype(blocks[i]).model_dump(mode="json"),
            },
            "reference": {
                "phenotype_reference": m.reference_phenotype().model_dump(mode="json")
            },
        }
        for i, spec in enumerate(m.STRAINS)
    ]
    assert m._l1_one_record_per_block(records).passed
    assert m._l2_values_equal_workbook(records, blocks).passed
    assert m._l3_interval_and_reference(records).passed
    moved = [_block(spec.block_label, shift=1.0) for spec in m.STRAINS]
    assert not m._l2_values_equal_workbook(records, moved).passed
    records[0]["reference"]["phenotype_reference"]["provenance_gaps"] = []
    assert not m._l3_interval_and_reference(records).passed
    finding = m._l4_zwf_is_a_bw25113_gene(records[:0], set())
    assert finding.passed
