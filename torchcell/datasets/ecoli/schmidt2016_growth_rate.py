# torchcell/datasets/ecoli/schmidt2016_growth_rate
# [[torchcell.datasets.ecoli.schmidt2016_growth_rate]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/schmidt2016_growth_rate
# Test file: tests/torchcell/datasets/ecoli/test_schmidt2016_growth_rate.py
"""Schmidt 2016 Table S24: the rim-deletion growth rates, as WT-relative fitness.

Schmidt et al. 2016 (Nat Biotechnol 34:104, doi:10.1038/nbt.3418) released one table of
gene-perturbation phenotype: Supplementary Table 24, "Growth rates determined for WT,
ΔrimI, ΔrimJ and ΔrimL E. coli strains grown in glucose and acetate medium". This module
serves it as :class:`GrowthRateSchmidt2016Dataset`, one
``BacterialFitnessExperiment`` per (deletion strain, medium): **six records and two
references**, the wild-type row of each medium being that medium's reference.

It is a SEPARATE MODULE from ``schmidt2016.py`` on purpose. ``build_manifest`` keys a
built store's staleness on the schema closure of the loader MODULE's own
``torchcell.datamodels`` imports (``provenance/schema_deps.py:loader_closure``), so
adding ``FitnessPhenotype`` and ``BacterialDeletionPerturbation`` to ``schmidt2016.py``
would mark the already-served ``proteome_schmidt2016`` store stale and force a full
rebuild for a change that does not touch one of its records. The pinned artifact, the
media objects and the condition table are imported FROM ``schmidt2016`` instead, so
there is still exactly one copy of each.

THE DELETION STRAINS, AND WHERE THEY COME FROM. Online Methods, *Strains and plasmids*:
"Mutant strains with either the rimL, rimJ or rimI gene deleted were taken from the KEIO
collection19. Correctness of the deletions were checked by PCR." So ``collection`` is
``"KEIO collection"`` verbatim and the background is the same BW25113 the proteome map
uses. The replacement cassette is NOT stated by this paper -- it is a property of the
Keio collection that Baba 2006 states and that is not in this mirror -- so ``cassette``
is a typed gap rather than a guess. Measured against the pinned BW25113 annotation, all
three symbols resolve through the gene-symbol layer with no collision and no ambiguity:
``rimI -> BW25113_4373``, ``rimJ -> BW25113_1066``, ``rimL -> BW25113_1427``.

WHY ``FitnessPhenotype`` AND NOT A LOG RATIO. ``FitnessPhenotype`` is documented as
``ko_growth_rate/wt_growth_rate`` and its validator CLAMPS a non-positive value to 0.0,
which would silently destroy a measurement. Measured on the released table, every one of
the six ratios is strictly positive -- 0.454350 (ΔrimJ acetate) to 1.012567 (ΔrimI
acetate) -- so the clamp never fires and no information is lost. The log2 route through
``EnvironmentResponsePhenotype`` is therefore not taken. (Table S23's per-condition
growth rates DO include non-positive values, which is one of the reasons that table is
not loaded here; see the dendron note.)

THE REPLICATE DESIGN IS BACK-SOLVED, NOT ASSUMED. Table S24 releases three ``Replicate``
columns plus ``Average`` and ``Stdev``, and the Methods never name the replicate count
behind the ``Stdev``. Measured on the pinned workbook for all eight (strain, medium)
cells: ``Average`` equals the mean of that cell's non-empty replicates to a maximum
absolute deviation of 1.1e-16, and ``Stdev`` equals their SAMPLE standard deviation
(the n-1 denominator) to 8.3e-17, while the POPULATION standard deviation is off by
1.0e-3 to 8.5e-3. The released statistic therefore identifies its own ``n`` -- the
number of non-empty replicate cells -- and its own kind, ``sample_sd``. That is the
back-solve rule, not the conservative-lower-end fallback. Two cells release two
replicates rather than three (WT in glucose, ΔrimJ in acetate), and ``n_samples`` is 2
for those: the strains were grown in three replicates (Table S21's title, verbatim, for
glucose: "... from biological triplicates"; Table S20 lists Replicates 1, 2 and 3 for
every strain in acetate), but only two growth rates are released, and the released
count is both the back-solved denominator and the conservative one.

WHAT THE STORED UNCERTAINTY IS, EXACTLY. ``fitness`` is a ratio of two independently
measured means, and the released ``Stdev`` is a dispersion of the NUMERATOR in h^-1, so
it cannot be stored verbatim against a dimensionless ratio. Two numbers are stored and
each is named:

- ``fitness_uncertainty`` is ``Stdev_strain / mean(WT)`` with
  ``fitness_uncertainty_type=sample_sd`` and ``n_samples`` the strain's released
  replicate count. This is exact arithmetic: it is the sample standard deviation of the
  n released ratio observations ``{rate_r / mean(WT)}``, so the released statistic's kind
  and n survive the change of units.
- ``fitness_se`` is supplied explicitly as the delta-method standard error of the ratio,
  ``fitness * sqrt((SE_strain/mean_strain)^2 + (SE_WT/mean_WT)^2)`` with
  ``SE_x = Stdev_x / sqrt(n_x)``. The ONE assumption is that the strain's and the
  wild-type's cultures are independent, which the cultivation protocol makes true (each
  is its own shake flask inoculated from a preculture). It is always at least as large as
  the auto-derived ``fitness_uncertainty / sqrt(n)``, which conditions on the wild-type
  mean as a fixed denominator and so understates the spread; the build asserts that
  direction over every record, so the stored SE can never be the optimistic one.

THE MEDIUM HEADERS ARE VERIFIED AGAINST THE VALUES, NOT TRUSTED. This release is already
known to mislabel a block: Table S6 carries LB's coefficient of variation under the
header ``Glucose`` and glucose's under ``LB`` (``schmidt2016.py``,
``check_cv_header_swap``). Table S24's two blocks are headed ``Glucose:`` and
``Acetate:``, so they are checked against an independent released statistic -- Table
S23's own BW25113 growth rate per condition, 0.58 h^-1 in glucose and 0.30 h^-1 in
acetate. Measured: the ``Glucose:`` block's wild-type average is 0.5921 h^-1 (1.9% from
Table S23's glucose rate, 97% from its acetate rate) and the ``Acetate:`` block's is
0.3103 h^-1 (3.4% from acetate, 46% from glucose). The headers are correct, and
``check_medium_headers`` asserts that each block is nearer its own medium than the other
so a future re-export cannot swap them silently.

DATA. The one consumed file is ``si2.xlsx``, the same pinned Supplementary-tables
workbook ``schmidt2016.py`` consumes, read from the same raw mirror under the same
sha256 pin. Nothing is retrieved here that that module does not already record.
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import os
import os.path as osp
import statistics
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any, ClassVar

import openpyxl
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.datamodels.schema import (
    BacterialDeletionPerturbation,
    BacterialFitnessExperiment,
    BacterialFitnessExperimentReference,
    DerivedIdentifierMapping,
    Environment,
    Experiment,
    ExperimentReference,
    FitnessPhenotype,
    Genotype,
    SampleUnit,
    UncertaintyType,
)
from torchcell.datasets.bacteria_common import (
    BACTERIAL_ASSEMBLY_SETS,
    STRAIN_GENE_NAMESPACES,
    LocusTagReconciliation,
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.datasets.ecoli import schmidt2016 as sm
from torchcell.sequence.genome.ecoli.k12 import EcoliK12Genome, EcoliK12StrainName
from torchcell.verification.report import Provenance, VerificationReport
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
)

log = logging.getLogger(__name__)

DATASET_ROOT_REL = "data/torchcell/growth_rate_schmidt2016"

SHEET_S24 = "Table S24"
SHEET_S23 = "Table S23"
#: Table S24's per-block header row reads ``<medium>: | Replicate 1 | ... | Stdev``.
REPLICATE_HEADERS = ("Replicate 1", "Replicate 2", "Replicate 3")
AVERAGE_HEADER = "Average"
STDEV_HEADER = "Stdev"
#: Table S23's columns the medium-header cross-check reads.
S23_CONDITION = "Growth condition"
S23_STRAIN = "Strain"
S23_RATE = "Growth rate (h-1)"
#: Table S24's wild-type row label, and the deletion labels in sheet order.
WILD_TYPE_LABEL = "WT"
DELETION_LABELS: tuple[str, ...] = ("ΔrimI", "ΔrimJ", "ΔrimL")
#: ``Δ`` plus the gene symbol is how the sheet writes a deletion strain.
DELETION_PREFIX = "Δ"
#: Table S24's two medium blocks, and the ``schmidt2016.CONDITIONS`` entry each is.
MEDIUM_BLOCKS: tuple[tuple[str, str], ...] = (
    ("Glucose", "Glucose"),
    ("Acetate", "Acetate"),
)

SCHMIDT_REFERENCE_STRAIN: EcoliK12StrainName = sm.SCHMIDT_REFERENCE_STRAIN
BW25113_NAMESPACE = STRAIN_GENE_NAMESPACES[SCHMIDT_REFERENCE_STRAIN]
KEIO_COLLECTION = "KEIO collection"
#: 3 deletions x 2 media; the two wild-type rows are the references, not records.
EXPECTED_RECORDS = 6
EXPECTED_REFERENCES = 2
#: Relative tolerance of the released-statistics identities against the bytes.
IDENTITY_RTOL = 1e-9
#: Measured: all three symbols resolve, so the floor is 1.0 and a renamed strain label
#: stops the build instead of silently dropping a record.
MIN_RESOLVED_FRACTION = 1.0
#: Table S23's own BW25113 rate is within this relative distance of Table S24's
#: wild-type average for the SAME medium (measured: 0.019 glucose, 0.034 acetate) and
#: far outside it for the other one (0.97 and 0.46).
MEDIUM_HEADER_RTOL = 0.10


# --------------------------------------------------------------------------- #
# Verbatim quotes. Every one is a substring of the pinned ``paper.md`` or of a
# ``si2.xlsx`` row rendering (its non-empty cells joined by " | ").
# --------------------------------------------------------------------------- #
_Q_KEIO_STRAINS = (
    "Mutant strains with either the rimL, rimJ or rimI gene deleted were"
    " taken from the KEIO collection19. Correctness of the deletions were"
    " checked by PCR."
)
_Q_GROWTH_RATE_FIT = (
    "The growth rate of the cultures was determined from the cell counts "
    "over time at cell concentrations from $1 0 ^ { 5 } c e l l s / "
    "\\mathrm { m l }$ to $1 0 ^ { 9 } \\mathrm { c e l l s / m l }$ The "
    "growth rate was calculated from at least four consecutive "
    "measurements."
)
_Q_S23_S24_ARE_THE_CONDITIONS = (
    "An overview about the used growth conditions can be found in "
    "Supplementary Tables 23 and 24."
)
_Q_TABLE_S24 = (
    "Table S24 | Growth rates determined for WT, ΔrimI, ΔrimJ and ΔrimL E."
    " coli strains grown in glucose and acetate medium"
)
_Q_TABLE_S23 = (
    "Table S23 | Experimental details for all E. coli samples analzed in "
    "this study including growth rate, harvesting conditions, OD-values "
    "and number of identified proteins"
)
_Q_TABLE_S21_TRIPLICATES = (
    "Table S21 | List of Nα-acetylated peptides relatively quantified "
    "from WT, ΔrimI, ΔrimJ and ΔrimL strains grown in glucose medium "
    "using label-free quantification from biological triplicates"
)

PAPER = Provenance(
    source_uri=sm.PAPER_MD, citation_key=sm.CITATION_KEY, sha256=sm.PAPER_MD_SHA256
)
SI2_SOURCE = Provenance(
    source_uri=sm.SI2_MIRROR_RELPATH, citation_key=sm.CITATION_KEY, sha256=sm.SI2_SHA256
)

#: Quotes whose provenance is the pinned ``paper.md``, by the constant's own name.
PAPER_QUOTES: dict[str, str] = {
    "keio_strains": _Q_KEIO_STRAINS,
    "growth_rate_fit": _Q_GROWTH_RATE_FIT,
    "s23_s24_are_the_conditions": _Q_S23_S24_ARE_THE_CONDITIONS,
    "batch_37c": sm._Q_BATCH_37C,
    "carbon_sources": sm._Q_CARBON_SOURCES,
}
#: Quotes whose provenance is a ``si2.xlsx`` row rendering.
SI2_QUOTES: dict[str, str] = {
    "table_s24": _Q_TABLE_S24,
    "table_s23": _Q_TABLE_S23,
    "table_s21_biological_triplicates": _Q_TABLE_S21_TRIPLICATES,
}


def _paper(value: object, quote: str, *, note: str | None = None) -> SourcedValue:
    """A ``SourcedValue`` quoting the pinned ``paper.md``."""
    return SourcedValue(value=value, provenance=PAPER, quote=quote, note=note)


def _si2(value: object, quote: str, *, note: str | None = None) -> SourcedValue:
    """A ``SourcedValue`` quoting the pinned ``si2.xlsx``."""
    return SourcedValue(value=value, provenance=SI2_SOURCE, quote=quote, note=note)


SOURCED_VALUES: dict[str, SourcedValue] = {
    "deletion_collection": _paper(KEIO_COLLECTION, _Q_KEIO_STRAINS),
    "deletions_pcr_verified": _paper(True, _Q_KEIO_STRAINS),
    "reference_strain": _paper(SCHMIDT_REFERENCE_STRAIN, sm._Q_STRAIN),
    "temperature_c": _paper(37.0, sm._Q_BATCH_37C),
    "aerobicity": _paper(
        "aerobic",
        sm._Q_BATCH_37C,
        note="orbital shaking at 300 r.p.m. in a sponge-closed Erlenmeyer flask",
    ),
    "growth_rate_estimator": _paper(
        "exponential fit to flow-cytometry cell counts over at least four consecutive "
        "measurements, between 1e5 and 1e9 cells/ml",
        _Q_GROWTH_RATE_FIT,
    ),
    "media_are_table_s23_and_s24": _paper(
        "the Media section's carbon sources and concentrations are the growth "
        "conditions of Supplementary Tables 23 and 24",
        _Q_S23_S24_ARE_THE_CONDITIONS,
        note="this is what ties Table S24's 'Glucose:' and 'Acetate:' blocks to M9 "
        "with glucose 5 g/L and with sodium acetate 3.5 g/L",
    ),
    "sample_unit": _si2(
        SampleUnit.biological_replicate.value,
        _Q_TABLE_S21_TRIPLICATES,
        note="the same four strains in the same glucose medium are quantified 'from "
        "biological triplicates', and Table S20 lists Replicates 1, 2 and 3 for every "
        "strain in acetate, so one Table S24 replicate is one independently grown "
        "culture",
    ),
    "growth_table": _si2("Table S24", _Q_TABLE_S24),
    "condition_details_table": _si2(
        "Table S23",
        _Q_TABLE_S23,
        note="read only to cross-check Table S24's two medium headers against an "
        "independent released growth rate; no record is built from it",
    ),
}

_GAP_CASSETTE = ProvenanceGap(
    field="cassette",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="the paper names the KEIO collection and the PCR verification but never the "
    "replacement cassette; that is a property of the collection, stated by Baba 2006, "
    "which this mirror does not hold, so it is left unset rather than guessed",
)
_GAP_CONSTRUCTION = ProvenanceGap(
    field="construction",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="no plate, well or strain accession of the three KEIO strains is released",
)
#: ``BacterialDeletionPerturbation`` is not a ``ProvenanceGapMixin``, so these two
#: absences are recorded in ``preprocess/provenance_gaps.json`` rather than on the leaf.
#: That is where ``schmidt2016.py`` already files the BW25113 background lesions for the
#: same reason: the field exists, the carrier for its typed absence does not.
PERTURBATION_GAPS: tuple[ProvenanceGap, ...] = (_GAP_CASSETTE, _GAP_CONSTRUCTION)


# --------------------------------------------------------------------------- #
# Reading Table S24
# --------------------------------------------------------------------------- #
class StrainRates(BaseModel):
    """One Table S24 cell: a strain's released growth rates in one medium."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    medium: str = Field(description="Table S24 block label, without its colon.")
    strain: str = Field(description="Row label, verbatim (``WT`` or ``Δ<gene>``).")
    row_number: int = Field(description="1-based row in the sheet, for the ledger.")
    replicates: tuple[float, ...] = Field(
        description="the non-empty Replicate cells, in sheet order"
    )
    average_released: float
    stdev_released: float

    @property
    def n_samples(self) -> int:
        """Released replicate measurements of this strain's growth rate."""
        return len(self.replicates)

    @property
    def mean(self) -> float:
        """Mean of the released replicates."""
        return statistics.fmean(self.replicates)

    @property
    def sample_sd(self) -> float:
        """Sample (n-1) standard deviation of the released replicates."""
        return statistics.stdev(self.replicates)

    @property
    def standard_error(self) -> float:
        """Standard error of this strain's mean growth rate."""
        return self.sample_sd / math.sqrt(self.n_samples)

    @property
    def gene_symbol(self) -> str:
        """The deleted gene symbol; raises for the wild-type row, which deletes none."""
        if not self.strain.startswith(DELETION_PREFIX):
            raise RuntimeError(f"{self.strain!r} is not a deletion strain label")
        return self.strain.removeprefix(DELETION_PREFIX)


class MediumBlock(BaseModel):
    """One Table S24 medium block: its wild-type row and its three deletion rows."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    medium: str
    wild_type: StrainRates
    deletions: tuple[StrainRates, ...]

    @property
    def rows(self) -> tuple[StrainRates, ...]:
        """Every released row of this block, the wild type first."""
        return (self.wild_type, *self.deletions)


def read_table_s24(path: str) -> tuple[MediumBlock, ...]:
    """Read Table S24's two medium blocks, checking each block's own header row.

    A block is located by its ``<medium>:`` label cell, and the five headers beside
    that label must be the declared replicate, average and stdev names, so a re-export
    that adds a replicate column or renames one stops the build rather than reading a
    statistic out of the wrong column.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book[SHEET_S24].iter_rows(values_only=True))
    finally:
        book.close()

    declared = dict(MEDIUM_BLOCKS)
    blocks: list[MediumBlock] = []
    index = 0
    while index < len(rows):
        label = rows[index][0]
        name = "" if label is None else str(label).strip()
        if not name.endswith(":") or name.removesuffix(":") not in declared:
            index += 1
            continue
        medium = name.removesuffix(":")
        header = tuple(
            "" if cell is None else str(cell).strip() for cell in rows[index][1:6]
        )
        expected = (*REPLICATE_HEADERS, AVERAGE_HEADER, STDEV_HEADER)
        if header != expected:
            raise RuntimeError(
                f"Table S24's {medium!r} block is headed {header}, the module declares "
                f"{expected}"
            )
        labels: dict[str, StrainRates] = {}
        for offset, row in enumerate(rows[index + 1 :], start=index + 2):
            strain = "" if row[0] is None else str(row[0]).strip()
            if strain not in (WILD_TYPE_LABEL, *DELETION_LABELS):
                break
            replicates = tuple(
                float(cell)
                for cell in row[1:4]
                if isinstance(cell, (int, float)) and not isinstance(cell, bool)
            )
            if len(replicates) < 2:
                raise RuntimeError(
                    f"Table S24 {medium}/{strain} releases {len(replicates)} replicate"
                    " growth rates; a sample standard deviation needs at least two"
                )
            labels[strain] = StrainRates(
                medium=medium,
                strain=strain,
                row_number=offset,
                replicates=replicates,
                average_released=float(row[4]),
                stdev_released=float(row[5]),
            )
        missing = [
            strain
            for strain in (WILD_TYPE_LABEL, *DELETION_LABELS)
            if strain not in labels
        ]
        if missing:
            raise RuntimeError(
                f"Table S24's {medium!r} block releases no {missing} row"
            )
        blocks.append(
            MediumBlock(
                medium=medium,
                wild_type=labels[WILD_TYPE_LABEL],
                deletions=tuple(labels[strain] for strain in DELETION_LABELS),
            )
        )
        index += 1 + len(labels)
    found = tuple(block.medium for block in blocks)
    if found != tuple(declared):
        raise RuntimeError(
            f"Table S24 carries the medium blocks {found}, the module declares "
            f"{tuple(declared)}"
        )
    return tuple(blocks)


def read_table_s23_wild_type_rates(path: str) -> dict[str, float]:
    """``{growth condition: growth rate}`` of Table S23's BW25113 rows.

    Table S23 is NOT loaded as records (its absolute rate has no ``MeasurementType``
    member and two of its rows are negative). It is read here for one purpose: an
    independent released growth rate per condition, against which Table S24's two
    medium headers are checked.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book[SHEET_S23].iter_rows(values_only=True))
    finally:
        book.close()
    header = rows[sm.HEADER_ROW - 1]
    index = {
        str(cell).strip(): position
        for position, cell in enumerate(header)
        if cell is not None and str(cell).strip()
    }
    out: dict[str, float] = {}
    for row in rows[sm.HEADER_ROW :]:
        condition = row[index[S23_CONDITION]]
        strain = row[index[S23_STRAIN]]
        rate = row[index[S23_RATE]]
        if condition is None or strain is None:
            continue
        if str(strain).strip() != SCHMIDT_REFERENCE_STRAIN:
            continue
        if not isinstance(rate, (int, float)) or isinstance(rate, bool):
            continue
        out[str(condition).strip()] = float(rate)
    return out


# --------------------------------------------------------------------------- #
# The released statistics, and the medium headers
# --------------------------------------------------------------------------- #
def check_released_statistics(blocks: Sequence[MediumBlock]) -> dict[str, Any]:
    """``Average`` is the mean and ``Stdev`` the SAMPLE SD of the released replicates.

    This is what back-solves the replicate design the Methods never state: the released
    ``Stdev`` identifies both its own ``n`` (the number of non-empty replicate cells)
    and its own kind (``sample_sd``, the n-1 denominator). The population SD is checked
    too and must NOT match, because a population SD would mean the released number is
    already scaled by a different n.
    """
    worst_mean = 0.0
    worst_sample = 0.0
    nearest_population = math.inf
    rows = 0
    for block in blocks:
        for row in block.rows:
            rows += 1
            mean_deviation = abs(row.mean - row.average_released)
            sample_deviation = abs(row.sample_sd - row.stdev_released)
            population = statistics.pstdev(row.replicates)
            worst_mean = max(worst_mean, mean_deviation)
            worst_sample = max(worst_sample, sample_deviation)
            nearest_population = min(
                nearest_population, abs(population - row.stdev_released)
            )
            scale = abs(row.average_released)
            if mean_deviation / scale > IDENTITY_RTOL:
                raise RuntimeError(
                    f"Table S24 {row.medium}/{row.strain}: released Average "
                    f"{row.average_released} against {row.mean} computed from "
                    f"{row.n_samples} replicates"
                )
            if sample_deviation / abs(row.stdev_released) > IDENTITY_RTOL:
                raise RuntimeError(
                    f"Table S24 {row.medium}/{row.strain}: released Stdev "
                    f"{row.stdev_released} is not the sample standard deviation "
                    f"{row.sample_sd} of its {row.n_samples} replicates"
                )
    if nearest_population <= worst_sample:
        raise RuntimeError(
            "Table S24's Stdev no longer tells a sample standard deviation from a "
            f"population one (nearest population deviation {nearest_population}, worst "
            f"sample deviation {worst_sample})"
        )
    return {
        "n_rows": rows,
        "worst_average_abs_dev": worst_mean,
        "worst_sample_sd_abs_dev": worst_sample,
        "nearest_population_sd_abs_dev": nearest_population,
        "uncertainty_type": UncertaintyType.sample_sd.value,
        "n_samples_by_row": {
            f"{block.medium}/{row.strain}": row.n_samples
            for block in blocks
            for row in block.rows
        },
    }


def check_medium_headers(
    blocks: Sequence[MediumBlock], s23_rates: Mapping[str, float]
) -> dict[str, Any]:
    """Each Table S24 block's wild-type average matches Table S23's rate for ITS medium.

    The same release mislabels Table S6's first two CV columns, so no header here is
    trusted. The check is positional-proof: for each block the wild-type average must be
    within ``MEDIUM_HEADER_RTOL`` of Table S23's BW25113 rate for that medium AND nearer
    to it than to any other block's medium, so a swapped pair of headers stops the build.
    """
    distances: dict[str, dict[str, float]] = {}
    for block in blocks:
        own = block.wild_type.mean
        distances[block.medium] = {
            medium: abs(own - s23_rates[medium]) / s23_rates[medium]
            for medium, _ in MEDIUM_BLOCKS
        }
    for block in blocks:
        row = distances[block.medium]
        nearest = min(row, key=lambda medium: row[medium])
        if nearest != block.medium or row[block.medium] > MEDIUM_HEADER_RTOL:
            raise RuntimeError(
                f"Table S24's {block.medium!r} block has a wild-type average of "
                f"{block.wild_type.mean} h-1, which is nearest Table S23's "
                f"{nearest!r} rate ({s23_rates[nearest]}); the block headers no longer "
                "agree with the released per-condition growth rates"
            )
    return {
        "table_s23_bw25113_rates": {
            medium: s23_rates[medium] for medium, _ in MEDIUM_BLOCKS
        },
        "table_s24_wild_type_means": {
            block.medium: block.wild_type.mean for block in blocks
        },
        "relative_distances": distances,
        "swapped": False,
    }


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
def fitness_phenotype(strain: StrainRates, wild_type: StrainRates) -> FitnessPhenotype:
    """One deletion strain's growth rate as a ratio to the wild type of its medium.

    ``fitness_uncertainty`` is the released ``Stdev`` divided by the wild-type mean,
    which is the sample standard deviation of the n released ratio observations, and
    ``fitness_se`` is the delta-method standard error of the ratio, which also carries
    the wild type's own spread. The second is never smaller than what the first derives.
    """
    fitness = strain.mean / wild_type.mean
    uncertainty = strain.stdev_released / wild_type.mean
    propagated = fitness * math.sqrt(
        (strain.standard_error / strain.mean) ** 2
        + (wild_type.standard_error / wild_type.mean) ** 2
    )
    conditioned = uncertainty / math.sqrt(strain.n_samples)
    if propagated < conditioned:
        raise RuntimeError(
            f"{strain.medium}/{strain.strain}: the propagated SE {propagated} is "
            f"smaller than the wild-type-conditioned {conditioned}"
        )
    return FitnessPhenotype(
        fitness=fitness,
        fitness_se=propagated,
        fitness_uncertainty=uncertainty,
        fitness_uncertainty_type=UncertaintyType.sample_sd,
        n_samples=strain.n_samples,
        sample_unit=SampleUnit.biological_replicate,
    )


def reference_phenotype(wild_type: StrainRates) -> FitnessPhenotype:
    """The wild type of one medium: the ratio of its own mean to itself, 1.0.

    Its uncertainty is the wild type's own relative spread, ``Stdev / mean``, which is
    the sample standard deviation of the n ratio observations the reference contributes.
    """
    return FitnessPhenotype(
        fitness=1.0,
        fitness_uncertainty=wild_type.stdev_released / wild_type.mean,
        fitness_uncertainty_type=UncertaintyType.sample_sd,
        n_samples=wild_type.n_samples,
        sample_unit=SampleUnit.biological_replicate,
    )


def build_genotype(strain: StrainRates, locus_tag: str) -> Genotype:
    """One KEIO single deletion, written against the pinned BW25113 assembly."""
    return Genotype(
        perturbations=[
            BacterialDeletionPerturbation(
                systematic_gene_name=locus_tag,
                perturbed_gene_name=strain.gene_symbol,
                gene_namespace=BW25113_NAMESPACE,
                identifier_mapping=DerivedIdentifierMapping(
                    source_identifier=strain.strain, route="gene_symbol"
                ),
                collection=KEIO_COLLECTION,
            )
        ]
    )


def build_environments() -> dict[str, Environment]:
    """``{Table S24 medium: environment}``, from ``schmidt2016``'s condition table.

    The media objects, the carbon-source reagents and their g/L are the ones the proteome
    loader already sources from the Methods, which is why nothing is re-stated here: the
    Media section's own closing sentence says those recipes ARE Table S23's and Table
    S24's growth conditions.
    """
    environments = {
        medium: sm.build_environment(sm.CONDITIONS_BY_COLUMN[column])
        for medium, column in MEDIUM_BLOCKS
    }
    sm.check_environments_distinct(environments)
    return environments


def resolve_deletions(
    blocks: Sequence[MediumBlock], genome: EcoliK12Genome, *, label: str
) -> tuple[dict[str, str], LocusTagReconciliation]:
    """``{gene symbol: locus tag}`` for the deletion labels, plus the resolver report."""
    symbols = sorted({row.gene_symbol for block in blocks for row in block.deletions})
    stored, report = reconcile_locus_tags(genome, pd.Series(symbols), label=label)
    report.require_resolved(MIN_RESOLVED_FRACTION)
    outside = set(report.outside_namespace)
    unresolved = [tag for tag in stored if tag in outside]
    if unresolved:
        raise RuntimeError(f"{label}: {unresolved} resolve to no BW25113 locus")
    locus_tag = {symbol: str(tag) for symbol, tag in zip(symbols, stored, strict=True)}
    if len(set(locus_tag.values())) != len(locus_tag):
        raise RuntimeError(f"{label}: two deletion labels share one locus tag")
    return locus_tag, report


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class GrowthRateSchmidt2016Dataset(ExperimentDataset):
    """Schmidt 2016 Table S24: rim-deletion growth rates as WT-relative fitness."""

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = SCHMIDT_REFERENCE_STRAIN

    def __init__(
        self,
        root: str = DATASET_ROOT_REL,
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialFitnessExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialFitnessExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The one consumed file, linked from the raw mirror."""
        return [sm.SI2]

    def download(self) -> None:
        """Link the mirrored workbook into ``raw/`` after checking manifest and sha256."""
        data_root = os.environ["DATA_ROOT"]
        manifest = sm.load_manifest(data_root)
        check_manifest_pin(
            sm.SI2_MIRROR_RELPATH,
            sm.manifest_sha256(manifest, sm.SI2_MIRROR_RELPATH),
            sm.SI2_SHA256,
        )
        src = sm.raw_mirror_dir(data_root) / sm.SI2_MIRROR_RELPATH
        if not src.exists():
            raise RuntimeError(f"required raw artifact missing from mirror: {src}")
        os.makedirs(self.raw_dir, exist_ok=True)
        link_verified(src, osp.join(self.raw_dir, sm.SI2), sm.SI2_SHA256)
        log.info(
            "Schmidt 2016 Table S24 raw file linked into %s (sha256 verified)",
            self.raw_dir,
        )

    def _genome(self) -> EcoliK12Genome:
        """The injected genome, or the reference strain's default cache."""
        if self.ecoli_genome is None:
            self.ecoli_genome = bacterial_genome("ecoli", self.REFERENCE_STRAIN)
        expected = BACTERIAL_ASSEMBLY_SETS[self.REFERENCE_STRAIN]
        if self.ecoli_genome.ASSEMBLY_SET != expected:
            raise ValueError(
                f"{type(self).__name__} needs the {expected} genome, got "
                f"{self.ecoli_genome.ASSEMBLY_SET}"
            )
        return self.ecoli_genome

    @post_process
    def process(self) -> None:
        """Build one fitness record per (deletion strain, medium) and write the LMDB."""
        verify_raw_files(self.raw_dir, sm.DATA_SHA256)
        path = osp.join(self.raw_dir, sm.SI2)
        blocks = read_table_s24(path)
        statistics_check = check_released_statistics(blocks)
        header_check = check_medium_headers(
            blocks, read_table_s23_wild_type_rates(path)
        )
        genome = self._genome()
        locus_tag, report = resolve_deletions(
            blocks, genome, label=f"{self.name} deletion labels"
        )
        environments = build_environments()
        reference_genome = assembly_reference(self.REFERENCE_STRAIN)
        pub = sm.publication()

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        rows: list[dict[str, Any]] = []
        index = 0
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for block in tqdm(blocks, desc="schmidt2016 growth rate"):
                environment = environments[block.medium]
                reference = BacterialFitnessExperimentReference(
                    dataset_name=self.name,
                    genome_reference=reference_genome,
                    environment_reference=environment,
                    phenotype_reference=reference_phenotype(block.wild_type),
                )
                for strain in block.deletions:
                    phenotype = fitness_phenotype(strain, block.wild_type)
                    experiment = BacterialFitnessExperiment(
                        dataset_name=self.name,
                        genotype=build_genotype(strain, locus_tag[strain.gene_symbol]),
                        environment=environment,
                        phenotype=phenotype,
                    )
                    txn.put(
                        f"{index}".encode(),
                        self._intern_record(experiment, reference, pub, itxn),
                    )
                    index += 1
                    rows.append(
                        {
                            "medium": block.medium,
                            "strain": strain.strain,
                            "gene_symbol": strain.gene_symbol,
                            "locus_tag": locus_tag[strain.gene_symbol],
                            "n_samples": strain.n_samples,
                            "growth_rate_h1": strain.mean,
                            "wild_type_growth_rate_h1": block.wild_type.mean,
                            "wild_type_n_samples": block.wild_type.n_samples,
                            "fitness": phenotype.fitness,
                            "fitness_uncertainty": phenotype.fitness_uncertainty,
                            "fitness_se": phenotype.fitness_se,
                        }
                    )
        env.close()
        interned_env.close()

        if index != EXPECTED_RECORDS:
            raise RuntimeError(f"{index} records, the module states {EXPECTED_RECORDS}")
        self._write_ledgers(blocks, rows, report, statistics_check, header_check)
        log.info(
            "Schmidt 2016 Table S24: %d records (+ %d wild-type references) over %d "
            "media; fitness range %.6f to %.6f",
            index,
            len(blocks),
            len(blocks),
            min(row["fitness"] for row in rows),
            max(row["fitness"] for row in rows),
        )

    def _write_ledgers(
        self,
        blocks: Sequence[MediumBlock],
        rows: Sequence[Mapping[str, Any]],
        report: LocusTagReconciliation,
        statistics_check: Mapping[str, Any],
        header_check: Mapping[str, Any],
    ) -> None:
        """The retention ledger, the sourcing table and the released-statistics checks."""
        out = Path(self.preprocess_dir)
        released = sum(len(block.rows) for block in blocks)
        (out / "dropped_records.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "source_rows": released,
                    "kept_records": len(rows),
                    "reference_rows": [
                        f"{block.medium}/{WILD_TYPE_LABEL}" for block in blocks
                    ],
                    "dropped_records": 0,
                    "rules": [],
                    "notes": [
                        f"{released} released Table S24 rows = {len(rows)} records + "
                        f"{len(blocks)} wild-type rows, which are the per-medium "
                        "phenotype_reference rather than records: an unperturbed parent "
                        "is the reference a fitness ratio is taken against",
                        "no row is dropped; both media have a MEDIA_LIBRARY entry and "
                        "all three deletion labels resolve to a BW25113 locus",
                        "Table S23's 26 per-condition growth rates are NOT loaded: an "
                        "absolute rate has no MeasurementType member (gap 1 of the E. "
                        "coli SI audit), two of its rows are negative, and its Stdev's "
                        "replicate design is neither stated nor back-solvable",
                    ],
                    "reconciliation": report.model_dump(mode="json"),
                },
                indent=2,
            )
        )
        (out / "released_statistics_check.json").write_text(
            json.dumps(
                {
                    "replicate_design_back_solved": dict(statistics_check),
                    "table_s24_medium_headers": dict(header_check),
                    "rtol": IDENTITY_RTOL,
                },
                indent=2,
            )
        )
        (out / "sourced_values.json").write_text(
            json.dumps(
                {
                    name: value.model_dump(mode="json")
                    for name, value in SOURCED_VALUES.items()
                },
                indent=2,
            )
        )
        (out / "provenance_gaps.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "carrier": "BacterialDeletionPerturbation",
                    "note": "the leaf is not a ProvenanceGapMixin, so these typed "
                    "absences are recorded here rather than on the record",
                    "gaps": [gap.model_dump(mode="json") for gap in PERTURBATION_GAPS],
                },
                indent=2,
            )
        )
        pd.DataFrame(list(rows)).to_csv(out / "strains.csv", index=False)

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process() for this dataset."""
        raise NotImplementedError(
            "GrowthRateSchmidt2016Dataset builds records in process()"
        )


# --------------------------------------------------------------------------- #
# L0-L4 verification
# --------------------------------------------------------------------------- #
def verify_build(
    dataset_root: str,
    data_root: str | None = None,
    *,
    genome: EcoliK12Genome | None = None,
    expected_count: int = EXPECTED_RECORDS,
) -> VerificationReport:
    """Run the fitness L0-L4 gate over a built tree and write the report.

    The gene universe and the canonical-name resolver are the record's own pinned
    BW25113 assembly, not the yeast default, which is how the landed Campos 2018 loader
    states the same situation for the same collection.
    """
    from torchcell.verification.fitness import verify_fitness_dataset
    from torchcell.verification.runners import load_records

    records = load_records(dataset_root)
    if genome is None:
        genome = bacterial_genome("ecoli", SCHMIDT_REFERENCE_STRAIN, data_root)
    report = verify_fitness_dataset(
        records,
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=Provenance(
            source_uri=f"$DATA_ROOT/{sm.RAW_DIR_REL}/{sm.SI2_MIRROR_RELPATH}",
            citation_key=sm.CITATION_KEY,
            sha256=sm.SI2_SHA256,
            method=(
                "Table S24 'Growth rate (h-1)' replicate columns: one "
                "BacterialFitnessExperiment per (KEIO deletion strain, medium), "
                "fitness = mean(strain replicates) / mean(WT replicates) of the SAME "
                "medium, fitness_se = the delta-method SE of that ratio"
            ),
            page=f"si2.xlsx sheets '{SHEET_S24}' and '{SHEET_S23}'",
            retrieved=sm.SI2_RETRIEVED_AT,
        ),
        expected_count=expected_count,
        sgd_genes=set(genome.genbank.loci),
        gene_universe_label=f"BW25113 ({genome.ASSEMBLY_SET})",
        resolve_gene_name=genome.resolve_gene_name,
    )
    out = osp.join(dataset_root, "preprocess", "verification_report.json")
    os.makedirs(osp.dirname(out), exist_ok=True)
    with open(out, "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main(argv: list[str] | None = None) -> int:
    """CLI: ``build`` the dev-tree LMDB, or ``verify`` an already built one."""
    from dotenv import load_dotenv

    parser = argparse.ArgumentParser(
        prog="python -m torchcell.datasets.ecoli.schmidt2016_growth_rate"
    )
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("build", help="build (or load) the dev-tree LMDB")
    sub.add_parser("verify", help="run L0-L4 on the built dev-tree LMDB")
    args = parser.parse_args(argv)

    load_dotenv()
    data_root = os.environ["DATA_ROOT"]
    root = osp.join(data_root, DATASET_ROOT_REL)
    if args.command == "build":
        dataset = GrowthRateSchmidt2016Dataset(root=root)
        print(f"len = {len(dataset)}")
        print(Path(root, "preprocess", "strains.csv").read_text())
        return 0
    report = verify_build(root, data_root)
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
