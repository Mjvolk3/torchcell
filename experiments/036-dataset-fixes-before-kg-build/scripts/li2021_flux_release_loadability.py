# experiments/036-dataset-fixes-before-kg-build/scripts/li2021_flux_release_loadability.py
# [[experiments.036-dataset-fixes-before-kg-build.scripts.li2021_flux_release_loadability]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/036-dataset-fixes-before-kg-build/scripts/li2021_flux_release_loadability
"""Settle schedule row 50, Li 2021 mevalonate flux, by measuring Additional file 2.

The row's recorded need is "fitted flux as phenotype", so three things decide it, and
each is measured on the pinned bytes rather than read off the abstract:

1. **Is the released quantity a fit, and can an existing class carry it?**
   ``FluxPhenotype`` / ``FluxExperiment`` / ``FluxExperimentReference`` already exist,
   are exported from ``torchcell.datamodels``, and have no consumer; the ``flux
   phenotype`` graph class and three ``CellAdapter`` methods for it are already landed
   too. A real ``FluxPhenotype`` is therefore BUILT from one strain's released column,
   with its fixed bounds and its signed fluxes, to show what the class does and does not
   accept.

2. **What is the shape of the release?** The row states 198 reactions per strain. This
   counts the sheet: reaction rows, atom-transition rows, and ``//`` section comments,
   and reports how many reactions are PINNED (``LB90 == best fit == UB90``) rather than
   carrying a genuine interval, which is what decides how much of a stored map would be
   fitted at all.

3. **Can the released maps be keyed to strains?** The sheets name strains three
   different ways (``NETWORK`` BW-P08/BW-P04/BW-P10, ``LABEL_MEASUREMENTS``
   PB108/PB104/PB100, ``Biomass Composition`` PB10/PB04/PB08). Each ``NETWORK`` block is
   matched to a ``Biomass Composition`` column by converting that column's released
   byproduct yields (g per g CDW) into molar ratios against glucose uptake and comparing
   them with the block's own output fluxes. That is an independent test of the column
   labels, and it is what surfaces the collision the row's triage suspected.

Writes ``results/li2021_flux_release_loadability.json`` and
``results/li2021_flux_release_loadability_reactions.csv``.

Usage (from the repo root)::

    python experiments/036-dataset-fixes-before-kg-build/scripts/li2021_flux_release_loadability.py
"""

from __future__ import annotations

import hashlib
import json
import os
import os.path as osp
from typing import Any

import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook

from torchcell.datamodels import schema as s

CITATION_KEY = "liFineTuningGlycolytic2021"
LIBRARY_DIR_REL = f"torchcell-library/{CITATION_KEY}"
SI1_RELPATH = "si/si1.docx"
SI2_RELPATH = "si/si2.xlsx"
SI1_SHA256 = "14250d74f8e7bb60335635d51d406e3da6756df185228a6b861af896fa2f5094"
SI2_SHA256 = "24c836f8459e5cce4f6a322477a7554e1f7dd16e8bb6b7cdd0ac4890b24c585a"

SHEET_NETWORK = "NETWORK"
SHEET_LABEL_INPUT = "LABEL_INPUT"
SHEET_LABEL_MEASUREMENTS = "LABEL_MEASUREMENTS"
SHEET_BIOMASS = "Biomass Composition"

#: The three strain blocks of ``NETWORK``, measured off rows 1, 7 and 8: the cell that
#: names the strain, the absolute-flux triple and the normalized triple. Columns A, J, Q
#: and X to AB are entirely empty, so there is no fourth block.
STRAIN_BLOCKS: tuple[
    tuple[str, str, tuple[str, str, str], tuple[str, str, str]], ...
] = (
    ("BW-P08", "F1", ("D", "E", "F"), ("G", "H", "I")),
    ("BW-P04", "N1", ("K", "L", "M"), ("N", "O", "P")),
    ("BW-P10", "S1", ("R", "S", "T"), ("U", "V", "W")),
)
EXPECTED_EMPTY_COLUMNS = ("A", "J", Q := "Q", "X", "Y", "Z", "AA", "AB")

#: ``NETWORK`` rows 2 to 5: the per-strain fit diagnostics, by their own labels.
DIAGNOSTIC_ROWS: tuple[tuple[int, str], ...] = (
    (2, "Number of fitted measurements :"),
    (3, "Freedom of flux"),
    (4, "χ2 90%"),
    (5, "SSR :"),
)

#: The net-flux table (rows 9 to 325) and the separate exchange-flux table below it.
NET_FLUX_ROWS = (9, 325)
#: ``Biomass Composition``: the byproduct-yield block, its strain columns and its rows.
BIOMASS_STRAIN_CELLS: tuple[tuple[str, str], ...] = (
    ("PB10", "C"),
    ("PB04", "D"),
    ("PB08", "F"),
)
BIOMASS_YIELD_ROWS: tuple[tuple[str, int], ...] = (
    ("Acetate (g/g.cdw)", 22),
    ("Lactate (g/g.cdw)", 23),
    ("Ethanol (g/g.cdw)", 24),
    ("Citrate (g/g.cdw)", 25),
    ("Glucose (g/g.cdw)", 26),
)
#: The output reaction whose flux each released byproduct yield should reproduce, with
#: the molar mass that converts a mass yield into a molar one.
BYPRODUCT_REACTIONS: tuple[tuple[str, str, float], ...] = (
    ("Acetate (g/g.cdw)", "vACTout", 60.052),
    ("Lactate (g/g.cdw)", "vLACout", 90.078),
    ("Ethanol (g/g.cdw)", "vETHout", 46.068),
    ("Citrate (g/g.cdw)", "vCIT_out", 192.123),
)
GLUCOSE_MW = 180.156
GLUCOSE_YIELD_ROW = "Glucose (g/g.cdw)"
UPTAKE_REACTION = "vUPTU"

#: What the paper says about the released file and the fit, verbatim.
Q_FIT = (
    "Metabolic fluxes were estimated by minimizing the residual sum of squares between "
    "experimentally measured and model predicted 13C-enrichment using 13C-Flux "
    "software obtained from Dr. Wiechert [33]."
)
Q_FILE_NAME = "Additional file 2. MFA simulated result."
Q_TRACER = (
    "13C-MFA was performed using 100% 1-13C1 glucose as the feeding substrate was "
    "added to a concentration of 10 g/L."
)
Q_CORRECTION = (
    "The data obtained from GC-MS were corrected by reduction of the natural abundance "
    "ratio of C, H, O, N, and Si isotopes [30]."
)
Q_INSTRUMENT = (
    "The resulting proteinogenic acids were derivatized with "
    "N-(tert-butyldimethylsilyl)-N-methyl-trifluoroacetamide containing "
    "tert-butyldimethylchlorosilane in acetonitrile at 105 C for 1 h, and then analyzed "
    "by a GC-MS [Agilent 7890 A GC and 5975 C Mass Selective Detector (Agilent "
    "Technologies, Santa Clara, USA)] equipped with a DB-1column (Agilent Technologies)."
)
Q_WHICH_STRAINS = (
    "13C-MFA was performed to detect the metabolic flux distribution in strains "
    "BW-P08 BF and BW-P10 BF (which had high MVA titers) and the control strain BW-P BF."
)
Q_FIG5 = (
    "Fig. 5 Metabolic flux diagram of zwf-strengthened EP-bifido strains. The metabolic "
    "flux shown for strains from top to bottom are BW-P10 BF, BW-P08 BF, BW-P BF"
)
Q_HARVEST = (
    "Cells at the exponential growth phase were harvested by centrifugation at 7000 g "
    "for 5 min at 4 C."
)


def sha256_of(path: str) -> str:
    """Hex sha256 of a file, read in chunks."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_si(data_root: str) -> dict[str, Any]:
    """sha256-verify both SI files against the library manifest and this module's pins."""
    library = osp.join(data_root, LIBRARY_DIR_REL)
    manifest = json.loads(open(osp.join(library, "manifest.json")).read())
    out: dict[str, Any] = {"citation_key": CITATION_KEY, "files": []}
    for relpath, pinned in ((SI1_RELPATH, SI1_SHA256), (SI2_RELPATH, SI2_SHA256)):
        entry = next(f for f in manifest["files"] if f["path"] == relpath)
        path = osp.join(library, relpath)
        observed = sha256_of(path)
        if observed != pinned:
            raise RuntimeError(f"{path}: sha256 {observed}, module pins {pinned}")
        if entry["sha256"] != pinned:
            raise RuntimeError(f"manifest sha256 {entry['sha256']} != {pinned}")
        out["files"].append(
            {
                "path": path,
                "relpath": relpath,
                "sha256": observed,
                "bytes": entry["bytes"],
                "original_filename": entry["original_filename"],
                "retrieval_method": entry["retrieval"]["method"],
                "source_url": entry["retrieval"]["source_url"],
                "retrieval_command": entry["retrieval"]["retriever"],
                "retrieval_params": entry["retrieval"]["params"],
            }
        )
    out["si_data_sources"] = manifest["si_data_sources"]
    return out


# --------------------------------------------------------------------------- #
# The NETWORK sheet
# --------------------------------------------------------------------------- #
def read_diagnostics(ws: Any) -> dict[str, dict[str, Any]]:
    """The four fit-diagnostic rows, per strain, with their own labels checked."""
    out: dict[str, dict[str, Any]] = {}
    for row, label in DIAGNOSTIC_ROWS:
        observed = ws[f"B{row}"].value
        if observed != label:
            raise RuntimeError(f"NETWORK B{row} is {observed!r}, expected {label!r}")
    for name, name_cell, _absolute, _normalized in STRAIN_BLOCKS:
        if ws[name_cell].value != name:
            raise RuntimeError(f"NETWORK {name_cell} is not {name!r}")
        column = name_cell[0]
        out[name] = {
            label: ws[f"{column}{row}"].value for row, label in DIAGNOSTIC_ROWS
        }
    return out


def read_reaction_rows(ws: Any) -> list[dict[str, Any]]:
    """One record per net-flux reaction row, with all three strains' triples.

    A reaction row is one whose ``reaction  name`` cell (column B) holds a name that is
    not a ``//`` section comment. Its atom-transition line is the NEXT row's column C
    with an empty column B, which is why column C has roughly twice as many non-blank
    cells as there are reactions.
    """
    first, last = NET_FLUX_ROWS
    out: list[dict[str, Any]] = []
    for row in range(first, last + 1):
        name = ws[f"B{row}"].value
        if name is None or str(name).strip().startswith("//"):
            continue
        record: dict[str, Any] = {
            "row": row,
            "reaction_name": str(name).strip(),
            "equation": ws[f"C{row}"].value,
            "atom_transitions": ws[f"C{row + 1}"].value
            if ws[f"B{row + 1}"].value is None
            else None,
        }
        for strain, _cell, absolute, normalized in STRAIN_BLOCKS:
            for scale, triple in (("abs", absolute), ("norm", normalized)):
                best, lower, upper = (ws[f"{c}{row}"].value for c in triple)
                record[f"{strain}_{scale}_best"] = best
                record[f"{strain}_{scale}_lb90"] = lower
                record[f"{strain}_{scale}_ub90"] = upper
        out.append(record)
    return out


def count_sheet_shape(ws: Any, reactions: list[dict[str, Any]]) -> dict[str, Any]:
    """Where the row's "198 reactions" came from, and what the real count is."""
    first, last = NET_FLUX_ROWS
    comments = sum(
        1
        for row in range(first, last + 1)
        if (value := ws[f"B{row}"].value) is not None
        and str(value).strip().startswith("//")
    )
    column_c_nonblank = sum(
        1 for row in range(1, ws.max_row + 1) if ws[f"C{row}"].value is not None
    )
    atom_rows = sum(1 for r in reactions if r["atom_transitions"] is not None)
    empty = [
        column
        for column in EXPECTED_EMPTY_COLUMNS
        if all(cell.value is None for cell in ws[column])
    ]
    return {
        "n_reaction_rows": len(reactions),
        "n_atom_transition_rows": atom_rows,
        "n_section_comment_rows": comments,
        "column_c_nonblank_cells": column_c_nonblank,
        "schedule_row_claim": 198,
        "where_198_comes_from": (
            f"column C holds {column_c_nonblank} non-blank cells = 1 header "
            f"('Reactions') + {len(reactions)} reaction equations + {atom_rows} "
            "atom-transition lines. 198 is the reaction rows PLUS their atom-transition "
            f"rows, which double-counts every reaction. The reaction count is "
            f"{len(reactions)}."
        ),
        "columns_expected_empty_that_are_empty": empty,
        "n_strain_blocks": len(STRAIN_BLOCKS),
        "fourth_strain_block_present": len(empty) != len(EXPECTED_EMPTY_COLUMNS),
    }


def count_pinned_fluxes(reactions: list[dict[str, Any]]) -> dict[str, Any]:
    """How many reactions are PINNED (LB90 == best fit == UB90) rather than fitted.

    This is what decides how much of a stored map is a fitted quantity at all: a row
    whose two bounds equal its best fit was held at a measured rate or a biomass drain,
    not identified from the labeling data. The test runs on the ABSOLUTE columns, because
    dividing by the glucose uptake re-introduces floating-point spread into bounds that
    were bit-identical in absolute units, which understates how much is pinned.
    """
    out: dict[str, Any] = {}
    for strain, _cell, _absolute, _normalized in STRAIN_BLOCKS:
        pinned = [
            r["reaction_name"]
            for r in reactions
            if r[f"{strain}_abs_best"]
            == r[f"{strain}_abs_lb90"]
            == r[f"{strain}_abs_ub90"]
        ]
        pinned_norm = sum(
            1
            for r in reactions
            if r[f"{strain}_norm_best"]
            == r[f"{strain}_norm_lb90"]
            == r[f"{strain}_norm_ub90"]
        )
        out[strain] = {
            "n_pinned_absolute": len(pinned),
            "n_interval_absolute": len(reactions) - len(pinned),
            "n_pinned_normalized": pinned_norm,
            "pinned_reactions": pinned,
        }
    out["note"] = (
        "a pinned row is a constraint, not a fitted flux: the pinned sets are the "
        "measured uptake and secretion rates plus the biomass-stoichiometry drains. The "
        "three fits did not pin the SAME reactions, which is itself a finding."
    )
    return out


def find_duplicate_blocks(reactions: list[dict[str, Any]]) -> dict[str, Any]:
    """Pairwise: how many reaction rows are bit-identical between two strain blocks.

    Two independent fits of two different strains share only their trivially fixed rows.
    A pair sharing most of its rows is one constraint vector fitted twice, which is what
    the row's triage suspected when it said a strain's map "cannot be matched to the
    published figure".
    """
    names = [name for name, *_ in STRAIN_BLOCKS]
    out: dict[str, Any] = {}
    for i, left in enumerate(names):
        for right in names[i + 1 :]:
            identical = [
                r["reaction_name"]
                for r in reactions
                if r[f"{left}_abs_best"] == r[f"{right}_abs_best"]
                and r[f"{left}_abs_lb90"] == r[f"{right}_abs_lb90"]
                and r[f"{left}_abs_ub90"] == r[f"{right}_abs_ub90"]
            ]
            out[f"{left} vs {right}"] = {
                "n_identical_rows": len(identical),
                "of_n_reactions": len(reactions),
                "uptake_flux_equal": next(
                    r[f"{left}_abs_best"] == r[f"{right}_abs_best"]
                    for r in reactions
                    if r["reaction_name"] == UPTAKE_REACTION
                ),
            }
    return out


# --------------------------------------------------------------------------- #
# Keying the blocks to strains through the released byproduct yields
# --------------------------------------------------------------------------- #
def read_biomass_yields(ws: Any) -> dict[str, dict[str, float | None]]:
    """The released byproduct yields (g per g CDW) per ``Biomass Composition`` column."""
    for label, row in BIOMASS_YIELD_ROWS:
        observed = ws[f"B{row}"].value
        if observed != label:
            raise RuntimeError(
                f"'{SHEET_BIOMASS}' B{row} is {observed!r}, expected {label!r}"
            )
    out: dict[str, dict[str, float | None]] = {}
    for name, column in BIOMASS_STRAIN_CELLS:
        out[name] = {
            label: ws[f"{column}{row}"].value for label, row in BIOMASS_YIELD_ROWS
        }
    return out


def match_blocks_to_biomass_columns(
    reactions: list[dict[str, Any]], yields: dict[str, dict[str, float | None]]
) -> dict[str, Any]:
    """Match each NETWORK block to a Biomass column by its own output fluxes.

    A released mass yield (g byproduct per g CDW) and a fitted output flux are the same
    quantity in different units, so converting both to a molar ratio against glucose
    gives an independent test of which column belongs to which strain. The comparison is
    ``(flux_out / flux_uptake)`` against ``(yield_byproduct / MW_byproduct) / (yield_glucose
    / MW_glucose)``; the block's own match is the column minimizing the summed absolute
    difference over the four byproducts.
    """
    by_name = {r["reaction_name"]: r for r in reactions}
    observed: dict[str, dict[str, float]] = {}
    for strain, *_rest in STRAIN_BLOCKS:
        uptake = by_name[UPTAKE_REACTION][f"{strain}_abs_best"]
        observed[strain] = {
            label: float(by_name[reaction][f"{strain}_abs_best"]) / float(uptake)
            for label, reaction, _mw in BYPRODUCT_REACTIONS
        }
    expected: dict[str, dict[str, float]] = {}
    for column, row in yields.items():

        def released(label: str, row: dict[str, float | None] = row) -> float:
            """The released yield, which every column states; a blank one is a fault."""
            value = row[label]
            if value is None:
                raise RuntimeError(
                    f"'{SHEET_BIOMASS}' column {column}: {label} is blank"
                )
            return float(value)

        glucose_mol = released(GLUCOSE_YIELD_ROW) / GLUCOSE_MW
        expected[column] = {
            label: (released(label) / mw) / glucose_mol
            for label, _reaction, mw in BYPRODUCT_REACTIONS
        }
    out: dict[str, Any] = {
        "flux_molar_ratios": observed,
        "yield_molar_ratios": expected,
    }
    matches: dict[str, Any] = {}
    for strain, ratios in observed.items():
        scores = {
            column: sum(abs(ratios[label] - values[label]) for label in ratios)
            for column, values in expected.items()
        }
        best = min(scores, key=lambda column: scores[column])
        ordered = sorted(scores.items(), key=lambda pair: pair[1])
        matches[strain] = {
            "best_match": best,
            "scores": scores,
            "margin_over_runner_up": ordered[1][1] - ordered[0][1],
            "label_implies": {"BW-P08": "PB08", "BW-P04": "PB04", "BW-P10": "PB10"}[
                strain
            ],
            "constraints_agree_with_the_label": best
            == {"BW-P08": "PB08", "BW-P04": "PB04", "BW-P10": "PB10"}[strain],
        }
    out["matches"] = matches
    return out


# --------------------------------------------------------------------------- #
# The measured inputs
# --------------------------------------------------------------------------- #
def read_label_measurements(ws: Any) -> dict[str, Any]:
    """The labeling data the fit was run against, and whether it carries uncertainty."""
    strain_columns = {ws[f"{column}4"].value: column for column in ("D", "E", "F")}
    rows = 0
    metabolites: list[str] = []
    groups = 0
    current = None
    for row in range(5, ws.max_row + 1):
        if ws[f"C{row}"].value is None:
            continue
        rows += 1
        name = ws[f"B{row}"].value
        if name is not None:
            metabolites.append(str(name).strip())
        constraint = str(ws[f"C{row}"].value)
        prefix = (name, "x" in constraint)
        if name is not None or prefix[1] != (current[1] if current else None):
            groups += 1
        current = prefix
    empty_columns = [
        column
        for column in ("G", "H", "I", "J", "K")
        if all(cell.value is None for cell in ws[column])
    ]
    return {
        "strain_columns": strain_columns,
        "n_constraint_rows_per_strain": rows,
        "n_metabolites": len(metabolites),
        "metabolites": metabolites,
        "n_mdv_groups_approx": groups,
        "correction_note_cell": ws["D3"].value,
        "columns_g_to_k_empty": empty_columns,
        "carries_any_uncertainty_column": len(empty_columns) < 5,
        "note": (
            "columns G to K are entirely empty: the fitting inputs carry NO standard "
            "deviation, standard error or weight, so the least-squares objective was "
            "unweighted or weighted by an undisclosed constant. The paper states no "
            "replicate count for the labeling experiment."
        ),
    }


# --------------------------------------------------------------------------- #
# Can an existing class carry it?
# --------------------------------------------------------------------------- #
def build_real_flux_phenotype(
    reactions: list[dict[str, Any]], diagnostics: dict[str, dict[str, Any]], strain: str
) -> dict[str, Any]:
    """Build a real ``FluxPhenotype`` from one strain's released absolute column.

    This is the measurement behind "can an existing class carry it". Nothing is rounded
    or clipped: the signed fluxes, the pinned bounds where the two equal the best fit,
    and the stated 90 percent level all go in as released.
    """
    best = {
        r["reaction_name"]: float(r[f"{strain}_abs_best"])
        for r in reactions
        if r[f"{strain}_abs_best"] is not None
    }
    lower = {
        r["reaction_name"]: float(r[f"{strain}_abs_lb90"])
        for r in reactions
        if r[f"{strain}_abs_lb90"] is not None
    }
    upper = {
        r["reaction_name"]: float(r[f"{strain}_abs_ub90"])
        for r in reactions
        if r[f"{strain}_abs_ub90"] is not None
    }
    phenotype = s.FluxPhenotype(
        net_flux=best,
        net_flux_lower=lower,
        net_flux_upper=upper,
        confidence_level=0.90,
        measurement_type="c13_mfa_net_flux_mmol_per_gdcw_per_h",
        n_samples=None,
        sample_unit=None,
        provenance_gaps=[],
    )
    negative = {name: value for name, value in best.items() if value < 0}
    return {
        "strain": strain,
        "accepted": True,
        "n_reactions_stored": len(phenotype.net_flux),
        "n_signed_negative": len(negative),
        "example_negative": dict(sorted(negative.items())[:4]),
        "label_statistic_name": phenotype.label_statistic_name,
        "graph_level": phenotype.graph_level,
        "fit_diagnostics_that_have_no_field": diagnostics[strain],
        "note": (
            "FluxPhenotype takes the released column WHOLE: signed net fluxes, bounds "
            "that equal the best fit where the row was pinned, and the stated 90 percent "
            "level, with label_statistic_name None by the class's own design. The four "
            "fit diagnostics have no field on it and would live in the loader's ledger."
        ),
    }


def probe_the_reference(reactions: list[dict[str, Any]]) -> dict[str, Any]:
    """Can a ``FluxExperimentReference`` be built? This is the blocking question.

    Every ``Experiment`` in this schema is stored with an ``ExperimentReference`` whose
    ``phenotype_reference`` is a phenotype of the same family, and
    ``FluxPhenotype.net_flux`` must be non-empty. For an ABSOLUTE readout the reference is
    the parent strain's own measured value, as the landed Fuhrer 2017 and Mori 2021
    loaders both do; there is no definitional zero to fall back on the way a log fold
    change has one. So the question is whether the control strain's map was released.
    """
    released = [name for name, *_ in STRAIN_BLOCKS]
    try:
        s.FluxPhenotype(
            net_flux={}, measurement_type="c13_mfa_net_flux_mmol_per_gdcw_per_h"
        )
        empty_accepted = True
        error = None
    except Exception as exc:
        empty_accepted = False
        error = " ".join(str(exc).split())[:300]
    return {
        "maps_released": released,
        "control_strain_named_by_the_text": "BW-P BF",
        "control_map_released": False,
        "quote_naming_the_control_as_analyzed": Q_WHICH_STRAINS,
        "quote_fig5_rows": Q_FIG5,
        "empty_net_flux_accepted": empty_accepted,
        "empty_net_flux_error": error,
        "conclusion": (
            "no reference can be built. The text says 13C-MFA covered BW-P08 BF, "
            "BW-P10 BF and the control BW-P BF, and Fig. 5 prints three rows naming "
            "those three strains, but Additional file 2 holds maps for BW-P08, BW-P04 "
            "and BW-P10 -- the control's map is absent and BW-P04's is present without "
            "ever being mentioned. FluxExperimentReference.phenotype_reference is "
            "required and FluxPhenotype.net_flux must be non-empty, so there is nothing "
            "to put in it, and an absolute flux has no definitional zero the way a log "
            "ratio does. This blocks every record, including the one strain whose "
            "column is cleanly identified."
        ),
    }


def check_chi_square_typo(diagnostics: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Is the ``χ2 90%`` cell of one strain a digit transposition of the others?"""
    label = DIAGNOSTIC_ROWS[2][1]
    values = {strain: row[label] for strain, row in diagnostics.items()}
    distinct = sorted({float(v) for v in values.values()})
    return {
        "chi_square_90_by_strain": values,
        "distinct_values": distinct,
        "agrees_across_strains": len(distinct) == 1,
        "note": (
            "the three fits share a degrees-of-freedom row (85) and so must share one "
            "chi-square quantile, yet one cell reads 102.8 against 102.08 in the other "
            "two. The sheet contradicts itself; which value is intended is not stated, "
            "and no sentence in the paper interprets SSR or this quantile at all."
        ),
    }


# --------------------------------------------------------------------------- #
def main() -> int:
    """Measure Additional file 2, probe the flux classes, write the results."""
    load_dotenv()
    data_root = os.environ["DATA_ROOT"]
    here = osp.dirname(osp.dirname(osp.abspath(__file__)))
    results = osp.join(here, "results")
    os.makedirs(results, exist_ok=True)

    pinned = verify_si(data_root)
    si2 = next(f for f in pinned["files"] if f["relpath"] == SI2_RELPATH)["path"]
    workbook = load_workbook(si2, data_only=True, read_only=False)
    if workbook.sheetnames != [
        SHEET_NETWORK,
        SHEET_LABEL_INPUT,
        SHEET_LABEL_MEASUREMENTS,
        SHEET_BIOMASS,
    ]:
        raise RuntimeError(f"unexpected sheets: {workbook.sheetnames}")

    network = workbook[SHEET_NETWORK]
    diagnostics = read_diagnostics(network)
    reactions = read_reaction_rows(network)
    shape = count_sheet_shape(network, reactions)
    pinned_counts = count_pinned_fluxes(reactions)
    duplicates = find_duplicate_blocks(reactions)
    yields = read_biomass_yields(workbook[SHEET_BIOMASS])
    keying = match_blocks_to_biomass_columns(reactions, yields)
    measurements = read_label_measurements(workbook[SHEET_LABEL_MEASUREMENTS])

    report: dict[str, Any] = {
        "schedule_row": 50,
        "dataset": "Li 2021 mevalonate flux",
        "citation_key": CITATION_KEY,
        "pinned_si": pinned,
        "q1_is_it_a_fit": {
            "quote_methods": Q_FIT,
            "quote_the_file_name_the_paper_gives_it": Q_FILE_NAME,
            "quote_tracer": Q_TRACER,
            "quote_instrument": Q_INSTRUMENT,
            "quote_correction": Q_CORRECTION,
            "quote_harvest": Q_HARVEST,
            "software": "13CFLUX2, named only through reference [33] (Weitzel 2013); no "
            "version, no solver settings, no weighting scheme",
            "objective": "least squares on measured minus model-predicted 13C "
            "enrichment",
            "interval_procedure_stated": False,
            "column_headers_that_state_the_level": ["LB90", "UB90", "χ2 90%"],
            "fit_diagnostics_by_strain": diagnostics,
            "chi_square_check": check_chi_square_typo(diagnostics),
            "verdict": (
                "a MODEL FIT, not a measurement: a least-squares best fit of a whole "
                "reaction network to GC-MS labeling data from a single [1-13C]glucose "
                "experiment. The paper's own one-line name for the file is 'MFA "
                "simulated result', and the procedure behind LB90/UB90 is never stated "
                "anywhere, so only the LEVEL is sourced, not the method."
            ),
        },
        "q2_what_was_measured": {
            "label_measurements": measurements,
            "released_byproduct_yields": yields,
            "fitted_measurements_claimed_by_the_sheet": {
                strain: row["Number of fitted measurements :"]
                for strain, row in diagnostics.items()
            },
            "note": (
                "the fitting inputs are the labeling data of 13 proteinogenic amino "
                "acids plus the released byproduct yields and the biomass composition, "
                "all in this same file, which is what makes the fit reproducible. The "
                "measured extracellular rates appear ONLY as unitless yields in a sheet "
                "the paper never references, and no growth rate is reported anywhere."
            ),
        },
        "q3_shape_of_the_release": {
            "sheet_shape": shape,
            "pinned_vs_fitted": pinned_counts,
            "pairwise_identical_blocks": duplicates,
            "strain_keying": keying,
        },
        "q4_can_an_existing_class_carry_it": {
            "classes_that_already_exist_and_have_no_consumer": [
                "FluxPhenotype",
                "FluxExperiment",
                "FluxExperimentReference",
            ],
            "graph_class_already_declared": "flux phenotype",
            "adapter_methods_already_landed": [
                "_flux_properties",
                "_flux_phenotype_node",
                "_get_flux_phenotype_reference_nodes",
            ],
            "built_phenotype": build_real_flux_phenotype(
                reactions, diagnostics, "BW-P08"
            ),
            "reference": probe_the_reference(reactions),
        },
    }

    out_json = osp.join(results, "li2021_flux_release_loadability.json")
    with open(out_json, "w") as handle:
        json.dump(report, handle, indent=2, default=str)
    out_csv = osp.join(results, "li2021_flux_release_loadability_reactions.csv")
    pd.DataFrame(reactions).to_csv(out_csv, index=False)

    print(json.dumps(report["q1_is_it_a_fit"], indent=2, default=str))
    print(json.dumps(report["q3_shape_of_the_release"]["sheet_shape"], indent=2))
    print(
        json.dumps(
            {
                strain: {k: v for k, v in row.items() if k != "pinned_reactions"}
                for strain, row in pinned_counts.items()
                if strain != "note"
            },
            indent=2,
        )
    )
    print(json.dumps(duplicates, indent=2, default=str))
    print(json.dumps(keying["matches"], indent=2))
    print(json.dumps(measurements, indent=2, default=str))
    print(
        json.dumps(report["q4_can_an_existing_class_carry_it"], indent=2, default=str)
    )
    print(f"\nwrote {out_json}\nwrote {out_csv}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
