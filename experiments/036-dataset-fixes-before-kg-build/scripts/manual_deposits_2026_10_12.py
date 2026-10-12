# experiments/036-dataset-fixes-before-kg-build/scripts/manual_deposits_2026_10_12.py
# [[experiments.036-dataset-fixes-before-kg-build.scripts.manual_deposits_2026_10_12]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/036-dataset-fixes-before-kg-build/scripts/manual_deposits_2026_10_12.py
r"""Record and measure the six bacterial manual deposits of 2026-10-12.

On 2026-10-12 the owner downloaded, in a browser, the files scripts could not fetch for
six bacterial papers, and they were copied into ``$DATA_ROOT/torchcell-raw/<key>/``.
This script does two things, each idempotent:

``--record KEY ...`` writes each deposit's ``DEPOSIT.md`` and ``SHA256SUMS.txt`` beside
its files and adds the files to the key's ``manifest.json``
(:mod:`torchcell.literature.manual_deposit`). Three keys have a loader that owns its
deposit and its manifest, so the script calls the loader's own recorder:
``tong2020.record_carpe_deposit`` (the CarPE print, its layout text and the parsed CSV),
``balakrishnan2022.record_si_deposit`` (the Science supplement and the SM layout text)
and ``ishii2007.record_som_deposit`` (the SOM and its MinerU OCR). The other three have
no loader, so their deposits are declared here (Thompson 2020, Nichols 2011) or in their
gate module (``teteneva2024.SI_DEPOSIT``), and their manifests are extended here; the
Nichols and Teteneva manifests are CREATED, since neither key had a raw-mirror manifest
(the Mohiuddin 2022 precedent: a raw-mirror key may exist before the paper is in the
literature mirror).

``--measure`` reads the deposited bytes and writes
``experiments/036-dataset-fixes-before-kg-build/results/manual_deposits_2026_10_12.json``:

- Balakrishnan 2022: the Science Table S3 against the GEO copy the loader reads, cell for
  cell, and the shape and content of Tables S4 to S7 (Table S6 is a candidate dataset).
- Nichols 2011: every table (format, sheets, dimensions, header), the chemical-genomic
  fitness-score matrix of Table S2 (strains, conditions, values, identifier patterns),
  and what ``mmc8.pdf`` holds.
- Teteneva 2024: Table S4 (rows, water samples, fitness values, identifiers) and how its
  gene symbols land on the two W3110 annotation members, Table S3 (sequencing samples,
  the QC-failed one) and Table S2 (the water samples).
- Thompson 2020: the correction's text, the sentences that change anything, and the 46
  Borchert 2024 samples attributed to Thompson 2020, checked for any trace of what the
  correction rescopes.

Run from the repo root::

    python experiments/036-dataset-fixes-before-kg-build/scripts/manual_deposits_2026_10_12.py --record thompson nichols teteneva
    python experiments/036-dataset-fixes-before-kg-build/scripts/manual_deposits_2026_10_12.py --measure
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import subprocess
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Final

import openpyxl
import xlrd
from dotenv import load_dotenv

from torchcell.datasets.ecoli import balakrishnan2022 as bk
from torchcell.datasets.ecoli import ishii2007, mori2021, teteneva2024, tong2020
from torchcell.datasets.pputida import borchert2024
from torchcell.literature.manifest import (
    ROLE_SI_DATA,
    ROLE_SI_PDF,
    ROLE_SI_TEXT,
    ArtifactRecord,
    ProcessingRecord,
)
from torchcell.literature.manual_deposit import (
    DepositedFile,
    ManualDeposit,
    ScriptedEquivalent,
    artifact_records,
    extend_manifest,
    verify_deposit,
    write_deposit_records,
)
from torchcell.literature.retrieve import elsevier_mmc_url

REPO: Final = Path(__file__).resolve().parents[3]
RESULTS: Final = (
    REPO
    / "experiments/036-dataset-fixes-before-kg-build/results"
    / "manual_deposits_2026_10_12.json"
)
PDFTOTEXT_ARGS: Final = ("-layout", "-enc", "UTF-8")
PDFTOTEXT_VERSION: Final = "pdftotext version 21.01.0"

# --------------------------------------------------------------------------- #
# Thompson 2020: the SI figure PDF and the 2021 correction (#986)
# --------------------------------------------------------------------------- #
THOMPSON_KEY: Final = "thompsonFattyAcidAlcohol2020"
THOMPSON_DOI: Final = "10.1128/AEM.01665-20"
THOMPSON_TITLE: Final = (
    "Fatty Acid and Alcohol Metabolism in Pseudomonas putida: Functional Analysis "
    "Using Random Barcode Transposon Sequencing"
)
THOMPSON_SI_PAGE: Final = "https://journals.asm.org/doi/suppl/10.1128/AEM.01665-20"
#: The correction's own DOI, printed on it ("Appl Environ Microbiol 87:e00177-21").
CORRECTION_DOI: Final = "10.1128/AEM.00177-21"
CORRECTION_PDF: Final = "si/thompson2021_correction_PMC8091123.pdf"
CORRECTION_SHA256: Final = (
    "2c7e387f4b9a47bfa4cb8633f437fedaf471409b487acdf269c54e781b998319"
)
CORRECTION_TEXT: Final = "si/thompson2021_correction_PMC8091123.layout.txt"
THOMPSON_DEPOSIT: Final = ManualDeposit(
    citation_key=THOMPSON_KEY,
    doi=THOMPSON_DOI,
    title=THOMPSON_TITLE,
    source_url=THOMPSON_SI_PAGE,
    retrieval_command=(
        f"open {THOMPSON_SI_PAGE} in a browser and download Supplemental File 1; open "
        "the correction (PMC8091123, https://pmc.ncbi.nlm.nih.gov/articles/PMC8091123/, "
        f"publisher DOI {CORRECTION_DOI}) and download its PDF; copy both unchanged to "
        "si/ (the correction renamed to thompson2021_correction_PMC8091123.pdf), then "
        "verify every sha256"
    ),
    retrieved_by="the owner (mjvolk3), browser download",
    deposited_at="2026-10-12",
    issues=(699, 986),
    note="the SI is one figure PDF (Figs. S1 to S8 with legends) and carries no "
    "per-gene data. The correction was reached from its PMC record PMC8091123, but its "
    "bytes are the publisher's PDF: the page carries ASM's stamp 'Downloaded from "
    "https://journals.asm.org/journal/aem on 12 October 2026' and prints its own DOI "
    f"{CORRECTION_DOI}. Its pdftotext layout text is written beside it",
    files=(
        DepositedFile(
            path="si/aem.01665-20-s0001.pdf",
            role=ROLE_SI_PDF,
            sha256="8b3a2ce6399e3aeb644d183232e8bd421581015c26011926b02625c683f52430",
            bytes=3220545,
            original_filename="aem.01665-20-s0001.pdf",
            arrived_at="2026-10-12T03:41:11-05:00",
            description="Supplemental File 1, 8 pages: Figures S1 to S8 with legends",
        ),
        DepositedFile(
            path=CORRECTION_PDF,
            role=ROLE_SI_PDF,
            sha256=CORRECTION_SHA256,
            bytes=179342,
            original_filename=(
                "thompson-et-al-2021-correction-for-thompson-et-al-fatty-acid-and-"
                "alcohol-metabolism-in-pseudomonas-putida-functional.pdf"
            ),
            arrived_at="2026-10-12T03:41:12-05:00",
            source_url=f"https://doi.org/{CORRECTION_DOI}",
            description="the 2021 author correction, 1 page (Appl Environ Microbiol "
            "87:e00177-21)",
        ),
    ),
)
THOMPSON_SI_EXPECTED: Final = (
    "Thompson 2020 released no per-gene data file: its fitness data are published only "
    "on the Fitness Browser, which is probed rather than mirrored, and the values are "
    "served through the Borchert 2024 compendium. Its Supplemental File 1 (Figs. S1 to "
    "S8) and its 2021 correction (PMC8091123, DOI 10.1128/AEM.00177-21) were DEPOSITED "
    "2026-10-12 by the owner (manual_browser, issues #699 and #986). The correction "
    "rescopes the PP_2675 deletion mutant (it also removed the first 14 codons of "
    "PP_2676) and the JBEI registry link; it touches no RB-TnSeq fitness value",
)
#: The correction's substantive statements, verbatim; a layout line break folds to one
#: space and the text layer's "ﬁ" ligature to "fi" (NFKC) when they are checked against
#: the layout text. The registry URL is printed across a line break after "jbei", so it
#: is two fragments.
CORRECTION_STATEMENTS: Final = {
    "pp_2675_deletion": (
        "Readers should note that when the authors created the deletion mutant of "
        "PP_2675, they deleted the entire region between the gene’s start and stop "
        "codons. However, this also deleted the DNA that codes for the first 14 amino "
        "acids of the overlapping ORF PP_2676."
    ),
    "registry_link_wrong": (
        "Furthermore, the originally published URL for the JBEI public strain "
        "repository was incor-"
    ),
    "registry_link": (
        "The correct link to all strains and plasmids in this study is "
        "https://public-registry.jbei"
    ),
    "registry_link_end": ".org/folders/641.",
}

# --------------------------------------------------------------------------- #
# Nichols 2011: the Cell supplemental information (#691)
# --------------------------------------------------------------------------- #
NICHOLS_KEY: Final = "nicholsPhenotypicLandscapeBacterial2011"
NICHOLS_DOI: Final = "10.1016/j.cell.2010.11.052"
NICHOLS_TITLE: Final = "Phenotypic Landscape of a Bacterial Cell"
#: The article's PII; Elsevier's CDN serves its supplements to scripts.
NICHOLS_PII: Final = "S0092867410013747"
_NICHOLS_FILES: Final = (
    ("mmc1.xls", ROLE_SI_DATA, 34304, "07.743", "349ae759627153dbb6921aee6cc6fc4ef2d6a3e3304da5c68afeb16f5eb02e30", "Table S1, the screen list"),
    ("mmc2.xls", ROLE_SI_DATA, 13938077, "02.837", "54be9cbc384bdbc7b1aa5a1da872d0e10ca513526f1f446ed068d27675f10fa8", "Table S2, the complete data set (an OOXML workbook under an .xls name)"),
    ("mmc3.xls", ROLE_SI_DATA, 18432, "07.843", "bd1de75036367825e51e70387e2818ce10d952ea5195dacc24dbf21bd0df5f5f", "Table S3, multi-stress responders"),
    ("mmc4.xls", ROLE_SI_DATA, 25088, "07.611", "20cbe16a3e17fbc955a623c2eb6eb293c8846c2524aabf19d14f255b45c05527", "Table S4, conditionally essential genes"),
    ("mmc5.xls", ROLE_SI_DATA, 62976, "07.244", "c5407840e72ba9a673d9d08fa202beb0e684f39e696b19b769070b221adc099e", "Table S5, rich-media conditionally essential lethal calls"),
    ("mmc6.xls", ROLE_SI_DATA, 124416, "07.125", "0cb5f840b8c017cfcb667d4de54e6b2a51f9fa7df240e4fc6e3991026fde73ea", "Table S6, orphan genes correlated with annotated genes"),
    ("mmc7.xls", ROLE_SI_DATA, 64512, "07.351", "0b0b1248c8101ef8d7148cd67bb6bfbd19197938cfd629586c2c38c5fd1417cd", "Table S7, drug-GO interactions"),
    ("mmc8.pdf", ROLE_SI_PDF, 4341546, "10.319", "e0c4e40678d83bc47252d8e47cad64bdd8178ca4713cad2496e0a72b77909e2c", "the article plus Supplemental Information, 23 pages"),
)  # fmt: skip
NICHOLS_DEPOSIT: Final = ManualDeposit(
    citation_key=NICHOLS_KEY,
    doi=NICHOLS_DOI,
    title=NICHOLS_TITLE,
    source_url=f"https://doi.org/{NICHOLS_DOI}",
    retrieval_command=(
        f"open https://doi.org/{NICHOLS_DOI} in a browser, open 'Supplemental "
        "Information', download mmc1 to mmc8 and copy them unchanged to si/; or, "
        "measured equivalent, fetch each with torchcell.literature.retrieve.elsevier_mmc"
        f"(pii='{NICHOLS_PII}', filename=...)"
    ),
    retrieved_by="the owner (mjvolk3), browser download",
    deposited_at="2026-10-12",
    issues=(691,),
    note="the Cell supplemental information, all eight files; Elsevier's CDN has no "
    "mmc9 (HEAD over pdf, xls, xlsx, docx, doc and zip answered 404 on 2026-10-12). "
    "Every file is byte-identical to the CDN object of the same name, fetched and "
    "compared on 2026-10-12. mmc8.pdf holds the article itself (pages 1 to 14) ahead "
    "of the Supplemental Information, so the article text is in hand even though the "
    "Zotero item REEL3RAD still has no PDF (#691)",
    files=tuple(
        DepositedFile(
            path=f"si/{name}",
            role=role,
            sha256=sha,
            bytes=size,
            original_filename=name,
            arrived_at=f"2026-10-12T03:41:{arrived[:2]}-05:00",
            description=holds,
            scripted_equivalent=ScriptedEquivalent(
                retriever="torchcell.literature.retrieve.elsevier_mmc",
                params={"pii": NICHOLS_PII, "filename": name},
                url=elsevier_mmc_url(NICHOLS_PII, name),
                checked_at="2026-10-12",
            ),
        )
        for name, role, size, arrived, sha, holds in _NICHOLS_FILES
    ),
)
NICHOLS_SI_EXPECTED: Final = (
    "Tables S1 to S7 (mmc1-mmc7) and the article plus Supplemental Information "
    "(mmc8.pdf): DEPOSITED 2026-10-12 by the owner (manual_browser, #691), each "
    "byte-identical to Elsevier's CDN object. No loader reads them yet; Table S2 "
    "(mmc2) is the chemical-genomic fitness-score matrix a loader would read",
    "the article PDF in Zotero: item REEL3RAD has no attachment (#691); the owner is "
    "attaching it, after which lit_sync mirrors it into torchcell-library",
    "the data set web site the article names, http://ecoliwiki.net/tools/chemgen/: not "
    "probed",
)
TETENEVA_SI_EXPECTED: Final = (
    "Supplementary Figures S1, S2, Tables S2, S3 and S4: DEPOSITED 2026-10-12 by the "
    "owner (manual_browser, #910), each byte-identical to the PMC Article Datasets "
    "object under PMC11188689.1. Table S4 is the loader's input once the W3110 "
    "annotation route is decided (#910); no loader yet",
    "Supplementary Table S1 and Figures S3 and S4: listed in the PMC bucket under "
    "PMC11188689.1 and not deposited; nothing reads them",
    "the article: filed in Zotero on 2026-10-12 (PUKZFNFZ), mirrored into "
    "torchcell-library by the next lit_sync run",
)


# --------------------------------------------------------------------------- #
# Recording
# --------------------------------------------------------------------------- #
def raw_root(data_root: str) -> Path:
    """``$DATA_ROOT/torchcell-raw``."""
    return Path(data_root) / "torchcell-raw"


def layout_text(pdf: Path) -> str:
    """``pdftotext -layout -enc UTF-8`` of one PDF, refusing another poppler release."""
    version = subprocess.run(
        ["pdftotext", "-v"], capture_output=True, text=True, check=True
    )
    line = (version.stderr or version.stdout).strip().split("\n")[0]
    if line != PDFTOTEXT_VERSION:
        raise RuntimeError(f"{line!r} is not the pinned {PDFTOTEXT_VERSION!r}")
    return subprocess.run(
        ["pdftotext", *PDFTOTEXT_ARGS, str(pdf), "-"],
        capture_output=True,
        text=True,
        check=True,
    ).stdout


def record_thompson(data_root: str) -> Path:
    """The SI and the correction, plus the correction's layout text, into the manifest."""
    root = raw_root(data_root) / THOMPSON_KEY
    write_deposit_records(THOMPSON_DEPOSIT, root)
    text = layout_text(root / CORRECTION_PDF).encode("utf-8")
    dest = root / CORRECTION_TEXT
    if dest.exists() and dest.read_bytes() != text:
        raise RuntimeError(f"{dest} exists with other content; refusing")
    dest.write_bytes(text)
    derived = ArtifactRecord(
        path=CORRECTION_TEXT,
        role=ROLE_SI_TEXT,
        bytes=len(text),
        sha256=hashlib.sha256(text).hexdigest(),
        source=f"derived from {CORRECTION_PDF}",
        processing=ProcessingRecord(
            processor=(
                "experiments/036-dataset-fixes-before-kg-build/scripts/"
                "manual_deposits_2026_10_12.py:layout_text"
            ),
            tool="pdftotext",
            version=PDFTOTEXT_VERSION,
            params={"args": list(PDFTOTEXT_ARGS)},
            input_sha256=[CORRECTION_SHA256],
        ),
    )
    return extend_manifest(
        root,
        [*artifact_records(THOMPSON_DEPOSIT), derived],
        citation_key=THOMPSON_KEY,
        doi=THOMPSON_DOI,
        title=THOMPSON_TITLE,
        si_expected=list(THOMPSON_SI_EXPECTED),
    )


def record_nichols(data_root: str) -> Path:
    """The eight Cell supplements; the key's manifest is created."""
    root = raw_root(data_root) / NICHOLS_KEY
    write_deposit_records(NICHOLS_DEPOSIT, root)
    return extend_manifest(
        root,
        artifact_records(NICHOLS_DEPOSIT),
        citation_key=NICHOLS_KEY,
        doi=NICHOLS_DOI,
        title=NICHOLS_TITLE,
        si_expected=list(NICHOLS_SI_EXPECTED),
    )


def record_teteneva(data_root: str) -> Path:
    """The five ISME supplements; the key's manifest is created."""
    deposit = teteneva2024.SI_DEPOSIT
    root = raw_root(data_root) / deposit.citation_key
    write_deposit_records(deposit, root)
    return extend_manifest(
        root,
        artifact_records(deposit),
        citation_key=deposit.citation_key,
        doi=deposit.doi,
        title=deposit.title,
        si_expected=list(TETENEVA_SI_EXPECTED),
        si_data_sources=[
            f"https://pmc-oa-opendata.s3.amazonaws.com/{teteneva2024.PAPER_PMCID}.1/"
        ],
    )


RECORDERS: Final = {
    "tong": lambda root: tong2020.record_carpe_deposit(root),
    "balakrishnan": lambda root: bk.record_si_deposit(root),
    "ishii": lambda root: ishii2007.record_som_deposit(root),
    "thompson": record_thompson,
    "nichols": record_nichols,
    "teteneva": record_teteneva,
}


# --------------------------------------------------------------------------- #
# Measuring
# --------------------------------------------------------------------------- #
def _sheets(path: Path) -> dict[str, list[tuple[Any, ...]]]:
    """Every sheet of an OOXML workbook as row tuples (read from bytes)."""
    book = openpyxl.load_workbook(
        io.BytesIO(path.read_bytes()), read_only=True, data_only=True
    )
    out = {ws.title: list(ws.iter_rows(values_only=True)) for ws in book.worksheets}
    book.close()
    return out


def _nonempty(rows: list[tuple[Any, ...]]) -> list[tuple[Any, ...]]:
    return [r for r in rows if any(c is not None and str(c).strip() for c in r)]


def _width(rows: list[tuple[Any, ...]]) -> int:
    """The last column any row fills."""
    return max(
        (max((i + 1 for i, c in enumerate(r) if c is not None), default=0))
        for r in rows
    )


def measure_balakrishnan(data_root: str) -> dict[str, Any]:
    """Table S3 (Science) against the GEO copy; Tables S4 to S7 inventoried."""
    root = bk.raw_mirror_dir(data_root)
    verify_deposit(bk.SI_DEPOSIT, root)
    geo = _sheets(root / bk.MIRROR_RELPATH[bk.TABLE_S3])
    science = _sheets(root / "data/science.abk2066_table_s3.xlsx")
    s3: dict[str, Any] = {"same_sheet_names": list(geo) == list(science)}
    for name in geo:
        a, b = geo[name], science[name]
        diffs: list[tuple[int, int, Any, Any]] = []
        for i in range(max(len(a), len(b))):
            ra = a[i] if i < len(a) else ()
            rb = b[i] if i < len(b) else ()
            for j in range(max(len(ra), len(rb))):
                x = ra[j] if j < len(ra) else None
                y = rb[j] if j < len(rb) else None
                if x != y:
                    diffs.append((i + 1, j + 1, x, y))
        floats = [
            (abs(x - y) / max(abs(x), abs(y)))
            for _, _, x, y in diffs
            if isinstance(x, float) and isinstance(y, float)
        ]
        s3[name] = {
            "rows": (len(a), len(b)),
            "cells_differing": len(diffs),
            "float_differences": len(floats),
            "max_relative_float_difference": max(floats, default=0.0),
            "other_differences": [
                {"row": r, "column": c, "geo": x, "science": y}
                for r, c, x, y in diffs
                if not (isinstance(x, float) and isinstance(y, float))
            ],
        }
    tables: dict[str, Any] = {}
    for n in (4, 5, 6, 7):
        sheets = _sheets(root / f"data/science.abk2066_table_s{n}.xlsx")
        tables[f"S{n}"] = {
            name: {
                "nonempty_rows": len(_nonempty(rows)),
                "columns": _width(rows),
                "header": [c for c in _nonempty(rows)[0] if c is not None],
            }
            for name, rows in sheets.items()
        }
    s6 = _sheets(root / "data/science.abk2066_table_s6.xlsx")
    s6_parameters: dict[str, Any] = {}
    for name, rows in s6.items():
        header, body = rows[0], _nonempty(rows[1:])
        s6_parameters[name] = {
            "genes": len(body),
            "distinct_loci": len({r[1] for r in body}),
            "reported_per_column": {
                str(header[j]): sum(
                    1
                    for r in body
                    if r[j] is not None and str(r[j]).strip().lower() != "nan"
                )
                for j in range(2, len(header))
            },
        }
    s4 = _sheets(root / "data/science.abk2066_table_s4.xlsx")
    description = _nonempty(s4["1 - Proteomics (description)"])
    ev9 = {s.column: s.strain for s in mori2021.SAMPLES if s.table == "EV9"}
    s4_strain = {str(r[0]): str(r[3]) for r in description[1:]}
    return {
        "table_s3_science_vs_geo": s3,
        "tables_s4_to_s7": tables,
        "table_s6_parameters": s6_parameters,
        "table_s4_samples": [r[0] for r in description[1:]],
        "table_s4_strains": dict(Counter(str(r[3]) for r in description[1:])),
        "table_s4_samples_named_in_mori2021_ev9": sorted(set(s4_strain) & set(ev9)),
        "table_s4_samples_whose_mori2021_ev9_strain_differs": sorted(
            k for k in set(s4_strain) & set(ev9) if s4_strain[k] != ev9[k]
        ),
    }


def _xls_rows(path: Path) -> dict[str, list[list[Any]]]:
    """Every sheet of a BIFF (.xls) or OOXML workbook, by its bytes' real format."""
    data = path.read_bytes()
    if data[:2] == b"PK":
        return {k: [list(r) for r in v] for k, v in _sheets(path).items()}
    book = xlrd.open_workbook(file_contents=data)
    return {
        sh.name: [sh.row_values(r) for r in range(sh.nrows)] for sh in book.sheets()
    }


_ECK_ROW = re.compile(r"ECK\d{4}-[A-Z0-9]+")


def _strain_kind(identifier: str) -> str:
    """A Table S2 row label's shape, by the pattern the release uses."""
    if identifier.endswith(" - SPA"):
        return "SPA-tagged essential gene (' - SPA')"
    if "/" in identifier.split("-", 1)[0]:
        return "several ECK numbers (a multi-gene deletion)"
    if re.fullmatch(r"ECK\d{4}", identifier):
        return "ECK number only"
    if identifier.endswith("'"):
        return "ECK number and symbol, primed (')"
    if _ECK_ROW.fullmatch(identifier):
        return "ECK number and symbol"
    return "other (alleles, truncations, insertions)"


def measure_nichols(data_root: str) -> dict[str, Any]:
    """Every Cell supplement, and the Table S2 fitness-score matrix in detail."""
    root = raw_root(data_root) / NICHOLS_KEY
    verify_deposit(NICHOLS_DEPOSIT, root)
    tables: dict[str, Any] = {}
    for f in NICHOLS_DEPOSIT.files:
        if not f.name.endswith(".xls"):
            continue
        data = (root / f.path).read_bytes()
        sheets = _xls_rows(root / f.path)
        tables[f.name] = {
            "format": "OOXML (zip)" if data[:2] == b"PK" else "BIFF (Excel 97-2004)",
            "sheets": {
                name: {
                    "rows": len(rows),
                    "nonempty_rows": len(_nonempty([tuple(r) for r in rows])),
                    "columns": _width([tuple(r) for r in rows]),
                    "title": str(rows[0][0]) if rows else None,
                    "header": [c for c in rows[1] if str(c).strip()]
                    if len(rows) > 1
                    else [],
                }
                for name, rows in sheets.items()
            },
        }
    (matrix,) = _xls_rows(root / "si/mmc2.xls").values()
    header, body = matrix[1], matrix[2:]
    conditions = [str(c) for c in header[1:]]
    values = [c for r in body for c in r[1:] if isinstance(c, (int, float))]
    identifiers = [str(r[0]) for r in body]
    repeated = Counter(identifiers)
    (screens,) = _xls_rows(root / "si/mmc1.xls").values()
    screen_rows = _nonempty([tuple(r) for r in screens[2:]])
    mmc8 = layout_text(root / "si/mmc8.pdf").split("\f")
    supplemental_start = next(
        i + 1 for i, page in enumerate(mmc8) if page.lstrip().startswith("Supplemental")
    )
    return {
        "tables": tables,
        "table_s2_matrix": {
            "file": "si/mmc2.xls",
            "title": str(matrix[0][0]),
            "strain_rows": len(body),
            "distinct_strain_labels": len(repeated),
            "labels_on_two_rows": sorted(k for k, v in repeated.items() if v > 1),
            "conditions": len(conditions),
            "distinct_conditions": len(set(conditions)),
            "condition_label_suffixes": dict(
                Counter(c.rsplit(" - ", 1)[-1] for c in conditions)
            ),
            "numeric_values": len(values),
            "empty_cells": sum(1 for r in body for c in r[1:] if c is None or c == ""),
            "value_min": min(values),
            "value_max": max(values),
            "strain_label_kinds": dict(Counter(_strain_kind(i) for i in identifiers)),
            "first_conditions": conditions[:5],
        },
        "table_s1_screens": {
            "condition_rows": len(screen_rows),
            "screens_column_sum": sum(
                float(r[5]) for r in screen_rows if str(r[5]).strip()
            ),
        },
        "mmc8_pdf": {
            "pages": len([p for p in mmc8 if p.strip()]),
            "supplemental_information_starts_on_page": supplemental_start,
            "first_line": mmc8[0].strip().splitlines()[0],
        },
    }


def measure_teteneva(data_root: str) -> dict[str, Any]:
    """Table S4 and its symbol routes, Table S3's samples, Table S2's water samples."""
    deposit = teteneva2024.SI_DEPOSIT
    root = raw_root(data_root) / deposit.citation_key
    verify_deposit(deposit, root)
    inventory = teteneva2024.table_s4_inventory(root / teteneva2024.TABLE_S4_RELPATH)
    book = openpyxl.load_workbook(
        io.BytesIO((root / teteneva2024.TABLE_S3_RELPATH).read_bytes())
    )
    sheet = book["DATA"]
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    samples = [[c.value for c in r] for r in sheet.iter_rows(min_row=2)]
    highlighted = [
        r[0].value
        for r in sheet.iter_rows(min_row=2)
        if any(c.fill.fill_type and c.fill.fgColor.rgb == "FFFFC7CE" for c in r)
    ]
    legend = [c.value for r in book["LEGEND"].iter_rows() for c in r if c.value]
    water = layout_text(root / "si/supplementary_table_s2_wrae096.pdf")
    return {
        "table_s4": inventory.model_dump(exclude={"distinct_symbols"})
        | {"distinct_symbols": len(inventory.distinct_symbols)},
        "table_s4_symbol_routes": teteneva2024.symbol_routes(
            inventory.distinct_symbols, data_root
        ),
        "table_s3": {
            "header": header,
            "samples": len(samples),
            "by_sample_type": dict(Counter(str(r[3]) for r in samples)),
            "by_sampling_day": dict(Counter(str(r[2]) for r in samples)),
            "by_water_sample": dict(Counter(str(r[1]) for r in samples)),
            "highlighted_red": highlighted,
            "legend": legend,
        },
        "table_s2_water_samples": [
            " ".join(line.split()) for line in water.splitlines() if "#" in line
        ],
    }


def _fold(text: str) -> str:
    """Layout line breaks to one space, and NFKC (the text layer's "ﬁ" ligature to "fi")."""
    return " ".join(unicodedata.normalize("NFKC", text).split())


def measure_thompson(data_root: str) -> dict[str, Any]:
    """The correction's statements, and the 46 Borchert samples they could touch."""
    root = raw_root(data_root) / THOMPSON_KEY
    verify_deposit(THOMPSON_DEPOSIT, root)
    text = (root / CORRECTION_TEXT).read_text(encoding="utf-8")
    folded = _fold(text)
    release = borchert2024.read_release(
        borchert2024.raw_mirror_dir(data_root) / borchert2024.DATA_RELPATH
    )
    thompson = [
        s
        for s in release.samples
        if borchert2024.attribute_sample(s).study == "thompson2020"
    ]
    pattern = re.compile(r"PP_?2675|PP_?2676|pedF|delet|knock|Δ", re.IGNORECASE)
    return {
        "correction_sha256": CORRECTION_SHA256,
        "statements_in_layout_text": {
            name: _fold(statement) in folded
            for name, statement in CORRECTION_STATEMENTS.items()
        },
        "statements": CORRECTION_STATEMENTS,
        "borchert2024_thompson_samples": {
            "samples": len(thompson),
            "records_at_4732_loci": len(thompson) * borchert2024.N_GENES,
            "sets": dict(Counter(s.set_name for s in thompson)),
            "mutant_libraries": dict(Counter(s.mutant_library for s in thompson)),
            "exp_groups": dict(Counter(s.exp_group for s in thompson)),
            "samples_naming_pp2675_pp2676_pedf_or_a_deletion": [
                s.exp_name
                for s in thompson
                if pattern.search(
                    " ".join(
                        str(v)
                        for v in (s.exp_desc, s.condition_1, s.condition_2, s.media)
                        if v is not None
                    )
                )
            ],
        },
    }


def main() -> None:
    """Record the named deposits, or measure all of them."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--record", nargs="*", choices=sorted(RECORDERS), default=[])
    parser.add_argument("--measure", action="store_true")
    args = parser.parse_args()
    load_dotenv(REPO / ".env")
    data_root = os.environ["DATA_ROOT"]
    for key in args.record:
        print(f"recorded {key}: {RECORDERS[key](data_root)}")
    if args.measure:
        results = {
            "script": "experiments/036-dataset-fixes-before-kg-build/scripts/"
            "manual_deposits_2026_10_12.py",
            "balakrishnan2022": measure_balakrishnan(data_root),
            "nichols2011": measure_nichols(data_root),
            "teteneva2024": measure_teteneva(data_root),
            "thompson2020": measure_thompson(data_root),
        }
        RESULTS.write_text(json.dumps(results, indent=2, default=str) + "\n")
        print(f"wrote {RESULTS}")


if __name__ == "__main__":
    main()
