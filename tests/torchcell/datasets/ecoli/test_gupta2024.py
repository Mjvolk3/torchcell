# tests/torchcell/datasets/ecoli/test_gupta2024.py
# [[tests.torchcell.datasets.ecoli.test_gupta2024]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_gupta2024.py
"""Tests for the Gupta 2024 protein-turnover loader.

Everything outside the last block runs with NO network and NO ``$DATA_ROOT``: the two
supplementary workbooks are synthesized to the exact column shape the pinned ones have,
the raw mirror is deposited under ``tmp_path`` by the module's own
``deposit_raw_mirror`` with its pins monkeypatched to the synthetic digests, and the
MG1655 annotation is the real ``EcoliK12MG1655Genome`` class over a synthetic assembly
whose CDS features carry UniProt ``db_xref`` qualifiers. The last block, skipped without
the mirror, pins the numbers measured on the REAL bytes.
"""

from __future__ import annotations

import hashlib
import json
import math
import os.path as osp
from pathlib import Path
from typing import Any, cast

import openpyxl
import pytest
from pydantic import ValidationError

import torchcell.datasets.ecoli.gupta2024 as gp
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    SyntheticLocus,
    forbid_network,
    gaf_row,
    serve_tier,
    write_assembly,
)
from torchcell.datamodels.schema import (
    AssemblyReferenceGenome,
    BacterialStrainBackground,
    Censoring,
    DerivedIdentifierMapping,
    EnvironmentPhysicalPerturbation,
    ProteinTurnoverExperiment,
    ProteinTurnoverExperimentReference,
)
from torchcell.sequence.genome.ecoli.k12 import MG1655_ASSEMBLY, EcoliK12MG1655Genome

# --------------------------------------------------------------------------- #
# A synthetic MG1655 assembly: four protease / SsrA loci plus three others.
# Only some loci carry a UniProt db_xref, so both identifier layers are exercised.
# --------------------------------------------------------------------------- #
LOCI = [
    SyntheticLocus(
        tag="b0001",
        parts=((1, 9),),
        strand="+",
        symbol="thrL",
        synonyms=("ECK0001",),
        xref="UniProtKB/Swiss-Prot:P00001",
        product="thr operon leader peptide",
        protein_id="AAC73112.1",
        protein="MK",
    ),
    SyntheticLocus(
        tag="b0002",
        parts=((12, 26),),
        strand="-",
        symbol="thrA",
        synonyms=("ECK0002",),
        product="aspartokinase I",
        protein_id="AAC73113.1",
        protein="MRVLK",
    ),
    SyntheticLocus(
        tag="b0003",
        parts=((30, 41),),
        strand="+",
        symbol="thrW",
        synonyms=("ECK0003",),
        xref="UniProtKB/Swiss-Prot:P00003",
        product="a third protein",
        protein_id="AAC73114.1",
        protein="MQQ",
    ),
    SyntheticLocus(
        tag="b0437",
        parts=((45, 56),),
        strand="+",
        symbol="clpP",
        synonyms=("ECK0437",),
        xref="UniProtKB/Swiss-Prot:P00437",
        product="ATP-dependent Clp protease proteolytic subunit",
        protein_id="AAC73115.1",
        protein="MSY",
    ),
    SyntheticLocus(
        tag="b0439",
        parts=((60, 71),),
        strand="-",
        symbol="lon",
        synonyms=("ECK0439",),
        product="DNA-binding ATP-dependent protease La",
        protein_id="AAC73116.1",
        protein="MNP",
    ),
    SyntheticLocus(
        tag="b3932",
        parts=((75, 86),),
        strand="+",
        symbol="hslV",
        synonyms=("ECK3932",),
        xref="UniProtKB/Swiss-Prot:P03932",
        product="ATP-dependent protease subunit HslV",
        protein_id="AAC73117.1",
        protein="MTT",
    ),
    SyntheticLocus(
        tag="b2621",
        parts=((88, 99),),
        strand="-",
        symbol="smpB",
        synonyms=("ECK2621",),
        product="SsrA-binding protein",
        protein_id="AAC73118.1",
        protein="MAK",
    ),
]

#: One GO row per locus symbol: the MG1655 assembly's GO route is the GAF's synonym
#: column, so the synthetic release needs one row per symbol it annotates.
GAF = [
    gaf_row(cast(str, locus.symbol), f"{locus.symbol}|{locus.tag}", "GO:0000001")
    for locus in LOCI
]

#: One released row per tuple: ``(protein id, released gene name, expected stored key)``.
#: ``b0002`` has no xref so it goes through the gene name; the ``THRW`` row carries the
#: WRONG gene name on purpose, so the accession layer overrides a disagreeing symbol;
#: the isoform row resolves through neither layer and is the only dropped key.
ROWS: tuple[tuple[str, str, str], ...] = (
    ("sp|P00001|THRL_ECOLI", "thrL", "b0001"),
    ("sp|P00002|THRA_ECOLI", "thrA", "b0002"),
    ("sp|P00003|THRW_ECOLI", "thrL", "b0003"),
    ("sp|P00437|CLPP_ECOLI", "clpP", "b0437"),
    ("sp|P00439|LON_ECOLI", "lon", "b0439"),
    ("sp|P03932|HSLV_ECOLI", "hslV", "b3932"),
    ("sp|P02621|SMPB_ECOLI", "smpB", "b2621"),
    ("sp|P00002-2|THRA_ECOLI", "", "sp|P00002-2|THRA_ECOLI"),
)
DROPPED = "sp|P00002-2|THRA_ECOLI"
#: Row 4 (``lon``) is absent from the first condition; row 1 (``thrA``) has one
#: replicate there; row 3 (``clpP``) is ceiling-flagged in replicate 1 everywhere.
ABSENT_ROW = 4
SINGLE_REPLICATE_ROW = 1
CEILING_ROW = 3


def _cells(condition_index: int, row: int) -> tuple[str, str]:
    """The two replicate cell texts of one released row in one condition."""
    if condition_index == 0 and row == ABSENT_ROW:
        return "", ""
    first = f"{1.0 + row:.6f}"
    second = f"{1.5 + row:.6f}"
    if row == CEILING_ROW:
        first = f"{gp.CONDITIONS[condition_index].ceiling_hours:.3f}*"
    if condition_index == 0 and row == SINGLE_REPLICATE_ROW:
        second = ""
    return first, second


def _mean(first: str, second: str) -> str:
    """The authors' mean cell: the arithmetic mean of the available replicates."""
    values = [float(text.rstrip("*")) for text in (first, second) if text]
    if not values:
        return ""
    return f"{sum(values) / len(values):.12f}"


def _interval(text: str) -> str:
    """One Supplementary Data 7 cell for a replicate cell of ``text``."""
    if not text:
        return ""
    if text.endswith("*"):
        return "Undetermined"
    rate = math.log(2.0) / float(text)
    return f"[{math.log(2.0) / (rate * 1.2):.3f} {math.log(2.0) / (rate * 0.8):.3f}]"


def _write_sheet(
    path: Path, sheet: str, note: str, headers: list[str], body: list[list[str]]
) -> None:
    """Write one workbook with the four-row preamble the released ones have."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = sheet
    worksheet.cell(row=1, column=2, value=f"{sheet} - synthetic")
    worksheet.cell(row=3, column=2, value=note)
    worksheet.cell(row=gp.HEADER_ROW, column=2, value="Uniprot Protein IDs")
    worksheet.cell(row=gp.HEADER_ROW, column=3, value="-")
    for offset, header in enumerate(headers):
        worksheet.cell(row=gp.HEADER_ROW, column=4 + offset, value=header)
    worksheet.cell(row=gp.HEADER_ROW + 1, column=2, value="Protein ID")
    worksheet.cell(row=gp.HEADER_ROW + 1, column=3, value="Gene names ")
    for index, row in enumerate(body):
        worksheet.cell(row=gp.HEADER_ROW + 2 + index, column=2, value=row[0])
        worksheet.cell(row=gp.HEADER_ROW + 2 + index, column=3, value=row[1] or None)
        for offset, value in enumerate(row[2:]):
            worksheet.cell(
                row=gp.HEADER_ROW + 2 + index, column=4 + offset, value=value or None
            )
    workbook.save(path)


def write_workbooks(directory: Path) -> tuple[Path, Path]:
    """Both synthetic workbooks, column-for-column as the released ones are shaped."""
    half_headers: list[str] = []
    ci_headers: list[str] = []
    for condition in gp.CONDITIONS:
        half_headers += [
            condition.header_replicate_1,
            condition.header_replicate_2,
            condition.header_mean,
        ]
        ci_headers += [condition.header_ci_1, condition.header_ci_2]
    half_body: list[list[str]] = []
    ci_body: list[list[str]] = []
    for row, (protein_id, gene_name, _) in enumerate(ROWS):
        half_row = [protein_id, gene_name]
        ci_row = [protein_id, gene_name]
        for index in range(len(gp.CONDITIONS)):
            first, second = _cells(index, row)
            half_row += [first, second, _mean(first, second)]
            ci_row += [_interval(first), _interval(second)]
        half_body.append(half_row)
        ci_body.append(ci_row)
    half_path = directory / gp.HALF_LIVES_FILENAME
    ci_path = directory / gp.CONFIDENCE_FILENAME
    _write_sheet(
        half_path, gp.HALF_LIVES_SHEET, gp.CEILING_NOTE.quote, half_headers, half_body
    )
    _write_sheet(
        ci_path, gp.CONFIDENCE_SHEET, gp.UNDETERMINED_NOTE.quote, ci_headers, ci_body
    )
    return half_path, ci_path


# --------------------------------------------------------------------------- #
# Synthetic: the released-cell parser and the rate transform
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("", None),
        ("   ", None),
        ("3.5", gp.HalfLifeCell(hours=3.5, ceiling=False)),
        (3.5, gp.HalfLifeCell(hours=3.5, ceiling=False)),
        ("8.000*", gp.HalfLifeCell(hours=8.0, ceiling=True)),
    ],
)
def test_a_released_cell_parses_with_its_ceiling_flag(
    value: Any, expected: gp.HalfLifeCell | None
) -> None:
    """A trailing ``*`` is kept as a flag; an empty cell is an absent value."""
    assert gp.parse_half_life(value) == expected


def test_the_total_turnover_rate_is_ln2_over_the_half_life() -> None:
    """The Supplementary Note's identity, applied as written."""
    assert gp.total_turnover_rate(1.0) == pytest.approx(math.log(2.0))
    assert gp.total_turnover_rate(6.0) == pytest.approx(math.log(2.0) / 6.0)


@pytest.mark.parametrize("hours", [0.0, -1.0])
def test_a_non_positive_half_life_has_no_rate(hours: float) -> None:
    """``ln 2 / 0`` is not a rate, so the build stops rather than storing an infinity."""
    with pytest.raises(RuntimeError, match="has no turnover rate"):
        gp.total_turnover_rate(hours)


@pytest.mark.parametrize(
    ("protein_id", "accession"),
    [
        ("sp|P00001|THRL_ECOLI", "P00001"),
        ("tr|A5A614|YCIZ_ECOLI", "A5A614"),
        ("sp|P07363-2|CHEA_ECOLI", "P07363"),
    ],
)
def test_the_canonical_accession_drops_the_isoform_suffix(
    protein_id: str, accession: str
) -> None:
    """A ``-2`` isoform key still names its canonical accession."""
    assert gp.base_accession(protein_id) == accession


def test_a_protein_key_of_another_shape_is_refused() -> None:
    """A key that is not ``db|accession|entry`` means the release changed shape."""
    with pytest.raises(RuntimeError, match="is not a 'db|accession|entry' key"):
        gp.base_accession("P00001")


# --------------------------------------------------------------------------- #
# Synthetic: the condition table
# --------------------------------------------------------------------------- #
def test_every_condition_is_distinct_in_what_identifies_a_record() -> None:
    """13 conditions, 13 identities: no two records could collide."""
    assert len(gp.CONDITIONS) == gp.EXPECTED_RECORDS
    identities = {gp._condition_identity(c) for c in gp.CONDITIONS}
    assert len(identities) == gp.EXPECTED_RECORDS


def test_every_condition_pins_its_censored_and_interval_key_counts() -> None:
    """Both #753 maps are pinned per condition, so neither can drift unnoticed."""
    assert sum(c.censored_keys for c in gp.CONDITIONS) == 1989
    assert sum(c.interval_keys for c in gp.CONDITIONS) == 4270
    for condition in gp.CONDITIONS:
        assert 0 < condition.censored_keys < condition.stored_proteins
        assert 0 < condition.interval_keys < condition.stored_proteins


def test_duration_hours_is_the_last_sampling_time_of_the_doubling_time() -> None:
    """Methods Table 2's series, in hours: what separates three dilution rates."""
    by_key = {c.key: c for c in gp.CONDITIONS}
    assert by_key["wt_minimal_batch_42min"].duration_hours == pytest.approx(175 / 60)
    assert by_key["wt_clim_3h"].duration_hours == pytest.approx(729 / 60)
    assert by_key["wt_clim_6h"].duration_hours == pytest.approx(1612 / 60)
    assert by_key["wt_clim_12h"].duration_hours == pytest.approx(2166 / 60)


def test_the_measurement_type_names_the_assay_and_its_unit_only() -> None:
    """The dilution regime moved onto the environment, so it is out of this string."""
    assert gp.MEASUREMENT_TYPE == "n15_ammonium_tmtproc_total_turnover_rate_per_hour"
    assert "doubling" not in gp.MEASUREMENT_TYPE
    assert "chemostat" not in gp.MEASUREMENT_TYPE
    values = gp.condition_values(
        _condition("wt_nlim_6h"), ["b0001"], _triple(2.0, 4.0), _pair(None, None)
    )
    for condition in gp.CONDITIONS:
        assert gp.phenotype(condition, values).measurement_type == gp.MEASUREMENT_TYPE


def test_the_dilution_rate_is_ln2_over_each_doubling_time() -> None:
    """The dilution-limited half-life IS the doubling time, through ln 2 / T."""
    by_key = {c.key: c for c in gp.CONDITIONS}
    assert by_key["wt_clim_3h"].dilution_rate_per_hour == pytest.approx(
        math.log(2.0) / 3.0
    )
    assert by_key["wt_clim_6h"].dilution_rate_per_hour == pytest.approx(
        math.log(2.0) / 6.0
    )
    assert by_key["wt_clim_12h"].dilution_rate_per_hour == pytest.approx(
        math.log(2.0) / 12.0
    )
    assert by_key["wt_minimal_batch_42min"].dilution_rate_per_hour is None
    assert by_key["wt_minimal_batch_42min"].doubling_hours == pytest.approx(0.7)
    continuous = [c for c in gp.CONDITIONS if c.dilution_rate_per_hour is not None]
    assert len(continuous) == 12


def test_the_13_identities_survive_the_dilution_rate_moving_off_measurement_type() -> (
    None
):
    """One constant measurement_type for all 13, and 13 distinct records regardless."""
    identities = {gp._condition_identity(c) for c in gp.CONDITIONS}
    assert len(identities) == 13
    windows = {c.duration_hours for c in gp.CONDITIONS}
    assert len(windows) == 4
    assert gp._condition_identity(_condition("wt_clim_3h")) != gp._condition_identity(
        _condition("wt_clim_6h")
    )


def test_the_stored_count_is_table_1_less_the_isoform_rows() -> None:
    """Every condition pins both counts, and the stored one is never the larger."""
    for condition in gp.CONDITIONS:
        assert condition.stored_proteins < condition.table1_proteins
        assert condition.table1_proteins - condition.stored_proteins <= 2


# --------------------------------------------------------------------------- #
# Synthetic: media, environment, genotype
# --------------------------------------------------------------------------- #
def _amount(media: Any, name: str) -> float:
    return next(
        float(component.concentration.value)
        for component in media.components
        if component.compound.name == name
    )


def test_the_unlimited_medium_is_mops_with_the_stated_glucose() -> None:
    """No limitation: the base's own ammonium and phosphate, plus 0.4 percent glucose."""
    media = gp.medium("none")
    assert media.base_medium == "MOPS_MINIMAL"
    assert media.is_synthetic is True
    assert media.state == "liquid"
    assert _amount(media, gp.GLUCOSE_COMPONENT) == pytest.approx(0.4)
    assert _amount(media, gp.AMMONIUM_COMPONENT) == pytest.approx(9.5)
    assert _amount(media, gp.PHOSPHATE_COMPONENT) == pytest.approx(1.32)


@pytest.mark.parametrize(
    ("limitation", "glucose", "ammonium", "phosphate"),
    [("C", 0.08, 9.5, 1.32), ("N", 0.4, 1.9, 1.32), ("P", 0.4, 9.5, 0.132)],
)
def test_each_limitation_reduces_exactly_its_own_component(
    limitation: str, glucose: float, ammonium: float, phosphate: float
) -> None:
    """The fivefold carbon and nitrogen cuts and the tenfold phosphate cut, as stated."""
    media = gp.medium(cast(Any, limitation))
    assert _amount(media, gp.GLUCOSE_COMPONENT) == pytest.approx(glucose)
    assert _amount(media, gp.AMMONIUM_COMPONENT) == pytest.approx(ammonium)
    assert _amount(media, gp.PHOSPHATE_COMPONENT) == pytest.approx(phosphate)


def test_the_four_media_are_four_distinct_entities() -> None:
    """A medium's name is what a cross-dataset aggregate groups on."""
    names = {gp.medium(cast(Any, lim)).name for lim in ("none", "C", "N", "P")}
    assert len(names) == 4


def test_a_renamed_base_component_stops_the_build(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The overrides are matched by compound name, so a rename must not pass silently."""
    monkeypatch.setattr(gp, "AMMONIUM_COMPONENT", "ammonium chloride hexahydrate")
    with pytest.raises(RuntimeError, match="is not a component of the built medium"):
        gp.medium("none")


def test_a_chemostat_environment_carries_the_controlled_ph_and_a_batch_one_does_not() -> (
    None
):
    """PH was maintained in the chemostat only, so only those records state it."""
    by_key = {c.key: c for c in gp.CONDITIONS}
    chemostat = gp.environment(by_key["wt_nlim_6h"])
    batch = gp.environment(by_key["wt_minimal_batch_42min"])
    assert len(chemostat.perturbations) == 1
    factor = chemostat.perturbations[0]
    assert isinstance(factor, EnvironmentPhysicalPerturbation)
    assert factor.factor.value == "pH"
    assert factor.magnitude is not None
    assert factor.magnitude.value == pytest.approx(7.2)
    assert batch.perturbations == []
    assert chemostat.temperature is not None
    assert chemostat.temperature.value == pytest.approx(37.0)
    assert batch.duration_hours == pytest.approx(175 / 60)
    assert chemostat.aerobicity == "aerobic"


def test_a_chemostat_environment_carries_the_dilution_rate_and_a_batch_one_does_not() -> (
    None
):
    """The dilution rate is the chemostat's controlled variable; a batch has none."""
    by_key = {c.key: c for c in gp.CONDITIONS}
    assert gp.environment(by_key["wt_nlim_6h"]).dilution_rate_per_hour == pytest.approx(
        math.log(2.0) / 6.0
    )
    assert gp.environment(
        by_key["wt_clim_12h"]
    ).dilution_rate_per_hour == pytest.approx(math.log(2.0) / 12.0)
    assert gp.environment(by_key["wt_minimal_batch_42min"]).dilution_rate_per_hour is (
        None
    )


def test_the_three_c_limited_environments_stay_three_distinct_environments() -> None:
    """Same medium, three dilution rates: the identity has to keep them apart."""
    from torchcell.datamodels.identity import environment_identity

    identities = {
        json.dumps(environment_identity(gp.environment(condition)), sort_keys=True)
        for condition in gp.CONDITIONS
        if condition.key in ("wt_clim_3h", "wt_clim_6h", "wt_clim_12h")
    }
    assert len(identities) == 3
    assert (
        len(
            {
                json.dumps(environment_identity(gp.environment(c)), sort_keys=True)
                for c in gp.CONDITIONS
            }
        )
        == 8
    )


def test_the_ncm3722_background_declares_what_the_paper_does_not_say() -> None:
    """NCM3722's lesions and construction are typed absences, not guesses."""
    background = gp.ncm3722_background()
    assert isinstance(background, BacterialStrainBackground)
    assert background.name == "NCM3722"
    assert background.reference_strain == "MG1655"
    assert background.alleles == []
    assert background.gapped_fields() == {"genotype_statement", "construction"}
    assert background.provenance is not None


def test_a_deletion_records_the_symbol_it_was_derived_from() -> None:
    """The paper releases symbols, so the b-number is a DERIVED identifier."""
    perturbation = gp.deletion_perturbation("clpP", "b0437")
    assert perturbation.systematic_gene_name == "b0437"
    assert perturbation.perturbed_gene_name == "clpP"
    assert perturbation.gene_namespace == "ecoli_k12_mg1655_bnumber"
    assert perturbation.identifier_mapping is not None
    assert perturbation.identifier_mapping.source_identifier == "clpP"
    assert perturbation.identifier_mapping.route == "gene_symbol"
    assert perturbation.collection == gp.KEIO_COLLECTION


def test_the_gift_triple_knockout_carries_no_keio_attribution() -> None:
    """``smpB`` is not a Keio-derived allele in this paper, so it claims no collection."""
    assert gp.deletion_perturbation("smpB", "b2621").collection is None


# --------------------------------------------------------------------------- #
# Synthetic: the per-condition label maps
# --------------------------------------------------------------------------- #
def _condition(key: str) -> gp.Condition:
    return next(c for c in gp.CONDITIONS if c.key == key)


def _triple(
    first: float | None, second: float | None, *, ceiling: bool = False
) -> list[gp.HalfLifeCell | None]:
    cells: list[gp.HalfLifeCell | None] = []
    values = [first, second]
    for index, value in enumerate(values):
        cells.append(
            None
            if value is None
            else gp.HalfLifeCell(hours=value, ceiling=ceiling and index == 0)
        )
    available = [v for v in values if v is not None]
    cells.append(
        None
        if not available
        else gp.HalfLifeCell(hours=sum(available) / len(available), ceiling=False)
    )
    return cells


def _pair(
    first: tuple[float, float] | str | None, second: tuple[float, float] | str | None
) -> list[gp.ConfidenceCell | None]:
    """Two Supplementary Data 7 cells: an interval, ``Undetermined``, or an empty cell."""
    cells: list[gp.ConfidenceCell | None] = []
    for value in (first, second):
        if value is None:
            cells.append(None)
        elif isinstance(value, str):
            assert value == "Undetermined"
            cells.append(gp.ConfidenceCell(undetermined=True))
        else:
            cells.append(
                gp.ConfidenceCell(
                    undetermined=False, lower_hours=value[0], upper_hours=value[1]
                )
            )
    return cells


def test_two_clean_replicates_get_the_papers_own_standard_error() -> None:
    """sd/sqrt(2) on the rate scale, which for n = 2 is half the rate difference."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"), ["b0001"], _triple(2.0, 4.0), _pair(None, None)
    )
    assert values.half_life == {"b0001": pytest.approx(3.0)}
    assert values.degradation_rate["b0001"] == pytest.approx(math.log(2.0) / 3.0)
    assert values.n_replicates == {"b0001": 2}
    rates = [math.log(2.0) / 2.0, math.log(2.0) / 4.0]
    assert values.degradation_rate_se["b0001"] == pytest.approx(
        abs(rates[0] - rates[1]) / 2.0
    )
    assert values.two_replicate_keys == 1
    assert values.ceiling_cells == ()
    assert values.censoring == {"b0001": Censoring.uncensored}
    assert values.degradation_rate_lower == {}
    assert values.degradation_rate_upper == {}
    assert values.refused_intervals == ()


def test_one_replicate_has_no_replicate_standard_error() -> None:
    """A single measurement carries no dispersion, so the key stores NaN."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"), ["b0001"], _triple(2.0, None), _pair((1.6, 2.5), None)
    )
    assert values.n_replicates == {"b0001": 1}
    assert math.isnan(values.degradation_rate_se["b0001"])
    assert values.two_replicate_keys == 0
    assert values.censoring == {"b0001": Censoring.uncensored}
    assert values.degradation_rate_lower == {
        "b0001": pytest.approx(math.log(2.0) / 2.5)
    }
    assert values.degradation_rate_upper == {
        "b0001": pytest.approx(math.log(2.0) / 1.6)
    }


def test_a_ceiling_replicate_is_kept_but_never_enters_a_standard_error() -> None:
    """A censored value is part of the released mean and not part of a dispersion."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"),
        ["b0001"],
        _triple(8.0, 4.0, ceiling=True),
        _pair("Undetermined", (3.5, 4.6)),
    )
    assert values.half_life["b0001"] == pytest.approx(6.0)
    assert values.n_replicates == {"b0001": 2}
    assert math.isnan(values.degradation_rate_se["b0001"])
    assert values.ceiling_cells == (("b0001", 1),)
    assert values.censoring == {"b0001": Censoring.right}
    assert values.degradation_rate_lower == {}


def test_a_protein_the_condition_did_not_quantify_is_not_a_key() -> None:
    """A missing value is an absent key, never a zero."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"), ["b0001"], _triple(None, None), _pair(None, None)
    )
    assert values.half_life == {}
    assert values.degradation_rate == {}
    assert values.n_replicates == {}
    assert values.censoring == {}


def test_a_mean_that_is_not_the_mean_of_its_replicates_is_refused() -> None:
    """The consumed column must be the authors' own average, or it is not understood."""
    cells = _triple(2.0, 4.0)
    cells[2] = gp.HalfLifeCell(hours=9.0, ceiling=False)
    with pytest.raises(RuntimeError, match="is not the arithmetic mean"):
        gp.condition_values(
            _condition("wt_nlim_6h"), ["b0001"], cells, _pair(None, None)
        )


def test_replicate_values_with_no_mean_cell_are_refused() -> None:
    """The mean column is the consumed column, so its absence is not recoverable."""
    cells = _triple(2.0, 4.0)
    cells[2] = None
    with pytest.raises(RuntimeError, match="with no mean cell"):
        gp.condition_values(
            _condition("wt_nlim_6h"), ["b0001"], cells, _pair(None, None)
        )


def test_a_mean_cell_with_no_replicate_value_is_refused() -> None:
    """A mean with nothing behind it is not a measurement."""
    cells: list[gp.HalfLifeCell | None] = [
        None,
        None,
        gp.HalfLifeCell(hours=3.0, ceiling=False),
    ]
    with pytest.raises(RuntimeError, match="a mean cell with no replicate value"):
        gp.condition_values(
            _condition("wt_nlim_6h"), ["b0001"], cells, _pair(None, None)
        )


def test_the_phenotype_declares_the_unreleased_synthesis_rate() -> None:
    """One fitted parameter, so a synthesis rate is a typed absence."""
    condition = _condition("wt_nlim_6h")
    values = gp.condition_values(
        condition, ["b0001"], _triple(2.0, 4.0), _pair(None, None)
    )
    phenotype = gp.phenotype(condition, values)
    assert phenotype.label_name == "degradation_rate"
    assert phenotype.label_statistic_name == "degradation_rate_se"
    assert phenotype.synthesis_rate is None
    assert phenotype.gapped_fields() == {"synthesis_rate"}
    assert phenotype.measurement_type == gp.MEASUREMENT_TYPE


# --------------------------------------------------------------------------- #
# Synthetic: the censoring flag and the published interval (#753)
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("", None),
        ("   ", None),
        ("Undetermined", gp.ConfidenceCell(undetermined=True)),
        (
            "[2.996 5.721]",
            gp.ConfidenceCell(undetermined=False, lower_hours=2.996, upper_hours=5.721),
        ),
        (
            "[0.169 -0.523]",
            gp.ConfidenceCell(
                undetermined=False, lower_hours=0.169, upper_hours=-0.523
            ),
        ),
    ],
)
def test_a_released_interval_cell_parses_to_its_two_endpoints(
    value: Any, expected: gp.ConfidenceCell | None
) -> None:
    """``[lower upper]`` in hours, or the ``Undetermined`` that IS the ceiling flag."""
    assert gp.parse_confidence_interval(value) == expected


@pytest.mark.parametrize("value", ["2.996 5.721", "[1.0]", "[1.0 2.0 3.0]"])
def test_an_interval_cell_of_another_shape_is_refused(value: str) -> None:
    """An unparsed cell is never skipped: the record's bounds come out of it."""
    with pytest.raises(RuntimeError):
        gp.parse_confidence_interval(value)


def test_an_undetermined_cell_has_no_endpoints_to_read() -> None:
    """It is the censoring flag, so asking it for bounds is a build error."""
    with pytest.raises(RuntimeError, match="has no endpoints"):
        gp.ConfidenceCell(undetermined=True).bounds_hours


def test_the_rate_bounds_swap_the_released_endpoints() -> None:
    """``r = ln 2 / T`` is decreasing, so the upper half-life gives the lower rate."""
    lower, upper = gp.rate_bounds(2.0, 8.0)
    assert lower == pytest.approx(math.log(2.0) / 8.0)
    assert upper == pytest.approx(math.log(2.0) / 2.0)


@pytest.mark.parametrize(("low", "high"), [(0.0, 0.0), (0.169, -0.523), (-1.0, 2.0)])
def test_a_non_positive_released_endpoint_has_no_rate_bound(
    low: float, high: float
) -> None:
    """A rate bound that crossed zero cannot be clamped into existence."""
    with pytest.raises(RuntimeError, match="non-positive endpoint"):
        gp.rate_bounds(low, high)


def test_a_one_replicate_key_whose_upper_endpoint_is_negative_is_refused_not_clamped() -> (
    None
):
    """The key carries no interval and is listed, with the released endpoints."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"),
        ["b0001"],
        _triple(0.5, None),
        _pair((0.169, -0.523), None),
    )
    assert values.degradation_rate_lower == {}
    assert values.degradation_rate_upper == {}
    assert len(values.refused_intervals) == 1
    refused = values.refused_intervals[0]
    assert refused.locus_tag == "b0001"
    assert refused.condition == "wt_nlim_6h"
    assert refused.released_lower_hours == pytest.approx(0.169)
    assert refused.released_upper_hours == pytest.approx(-0.523)
    assert refused.reason == "non_positive_released_half_life_endpoint"
    assert refused.half_life_hours == pytest.approx(0.5)


def test_a_released_interval_that_does_not_bracket_the_stored_rate_stops_the_build() -> (
    None
):
    """A stored bound that excludes the stored value would misreport both."""
    with pytest.raises(RuntimeError, match="does not bracket the stored rate"):
        gp.condition_values(
            _condition("wt_nlim_6h"),
            ["b0001"],
            _triple(2.0, None),
            _pair((3.0, 4.0), None),
        )


def test_a_censored_one_replicate_key_must_carry_an_undetermined_interval() -> None:
    """The two workbooks agree on every cell, so a determined one here is a conflict."""
    with pytest.raises(RuntimeError, match="the two workbooks disagree"):
        gp.condition_values(
            _condition("wt_nlim_6h"),
            ["b0001"],
            _triple(8.0, None, ceiling=True),
            _pair((7.0, 9.0), None),
        )


def test_a_censored_one_replicate_key_stores_no_interval() -> None:
    """A capped fit has no interval to store, and the key says it is right-censored."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"),
        ["b0001"],
        _triple(8.0, None, ceiling=True),
        _pair("Undetermined", None),
    )
    assert values.censoring == {"b0001": Censoring.right}
    assert values.degradation_rate_lower == {}
    assert values.n_replicates == {"b0001": 1}


def test_a_one_replicate_estimate_with_no_interval_cell_is_refused() -> None:
    """The release publishes an interval for every determined fit; 0 exceptions."""
    with pytest.raises(RuntimeError, match="publishes an interval for every"):
        gp.condition_values(
            _condition("wt_nlim_6h"), ["b0001"], _triple(2.0, None), _pair(None, None)
        )


def test_a_two_replicate_key_gets_no_interval_even_when_both_are_published() -> None:
    """The interval of the MEAN of two fits is not published, so none is stored."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"),
        ["b0001"],
        _triple(2.0, 4.0),
        _pair((1.6, 2.5), (3.5, 4.6)),
    )
    assert values.n_replicates == {"b0001": 2}
    assert values.degradation_rate_lower == {}
    assert values.degradation_rate_upper == {}


def test_the_second_replicate_is_the_one_read_when_the_first_is_absent() -> None:
    """The interval comes from the replicate the key actually has."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"),
        ["b0001"],
        _triple(None, 4.0),
        _pair((1.0, 2.0), (3.5, 4.6)),
    )
    assert values.degradation_rate_lower == {
        "b0001": pytest.approx(math.log(2.0) / 4.6)
    }
    assert values.degradation_rate_upper == {
        "b0001": pytest.approx(math.log(2.0) / 3.5)
    }


def test_the_phenotype_states_the_level_and_the_method_of_its_bounds() -> None:
    """Bounds without a level and a construction are two studies silently compared."""
    condition = _condition("wt_nlim_6h")
    values = gp.condition_values(
        condition, ["b0001"], _triple(2.0, None), _pair((1.6, 2.5), None)
    )
    phenotype = gp.phenotype(condition, values)
    assert phenotype.confidence_level == 0.95
    assert phenotype.interval_method == gp.INTERVAL_METHOD
    assert phenotype.interval_method.startswith("curve_fit_parameter_variance_t_ppf")
    assert phenotype.censoring == {"b0001": Censoring.uncensored}
    assert phenotype.degradation_rate_lower is not None
    assert set(phenotype.degradation_rate_lower) == {"b0001"}
    assert phenotype.label_statistic_name == "degradation_rate_se"


def test_the_level_is_the_one_the_supplementary_note_states() -> None:
    """0.95, quoted; the workbook's own column headers say it too."""
    assert gp.CONFIDENCE_LEVEL == 0.95
    assert gp.CONFIDENCE_LEVEL_SOURCE.value == 0.95
    assert "$9 5 \\%$ CI" in gp.CONFIDENCE_LEVEL_SOURCE.quote
    for condition in gp.CONDITIONS:
        assert condition.header_ci_1.startswith(
            "Total half-life 95% confidence interval"
        )


def test_every_stored_key_states_whether_it_is_censored() -> None:
    """The oracle is complete, so uncensored is stated rather than spelled by absence."""
    values = gp.condition_values(
        _condition("wt_nlim_6h"),
        ["b0001", "b0002"],
        _triple(2.0, 4.0) + _triple(8.0, 4.0, ceiling=True),
        _pair((1.6, 2.5), (3.5, 4.6)) + _pair("Undetermined", (3.5, 4.6)),
    )
    assert values.censoring == {"b0001": Censoring.uncensored, "b0002": Censoring.right}
    assert set(values.censoring) == set(values.degradation_rate)


# --------------------------------------------------------------------------- #
# Synthetic: the workbook readers
# --------------------------------------------------------------------------- #
def test_the_readers_return_every_released_row_in_order(tmp_path: Path) -> None:
    """Both workbooks list the same proteins in the same order, column-matched."""
    half_path, ci_path = write_workbooks(tmp_path)
    protein_ids, gene_names, cells = gp.read_half_lives(str(half_path))
    assert protein_ids == [row[0] for row in ROWS]
    assert gene_names == [row[1] for row in ROWS]
    assert set(cells) == {condition.key for condition in gp.CONDITIONS}
    assert len(cells["wt_nlim_6h"]) == 3 * len(ROWS)
    ci_ids, intervals = gp.read_confidence_intervals(str(ci_path))
    assert ci_ids == protein_ids
    assert len(intervals["wt_nlim_6h"]) == 2 * len(ROWS)
    flagged = intervals["wt_nlim_6h"][2 * CEILING_ROW]
    assert flagged == gp.ConfidenceCell(undetermined=True)
    determined = intervals["wt_nlim_6h"][2 * SINGLE_REPLICATE_ROW]
    assert determined is not None
    assert determined.undetermined is False
    lower, upper = determined.bounds_hours
    assert lower < upper
    second = intervals["wt_nlim_6h"][2 * CEILING_ROW + 1]
    assert second is not None
    assert second.undetermined is False


def test_a_column_the_loader_pins_and_the_workbook_lacks_is_refused(
    tmp_path: Path,
) -> None:
    """A reshaped workbook must not shift a condition's values onto another."""
    half_path, _ = write_workbooks(tmp_path)
    workbook = openpyxl.load_workbook(half_path)
    worksheet = workbook[gp.HALF_LIVES_SHEET]
    worksheet.cell(row=gp.HEADER_ROW, column=4, value="something else")
    workbook.save(half_path)
    with pytest.raises(RuntimeError, match="0 columns are headed"):
        gp.read_half_lives(str(half_path))


def test_a_duplicated_pinned_column_is_refused(tmp_path: Path) -> None:
    """Two columns under one header leave no single column to read."""
    half_path, _ = write_workbooks(tmp_path)
    workbook = openpyxl.load_workbook(half_path)
    worksheet = workbook[gp.HALF_LIVES_SHEET]
    worksheet.cell(
        row=gp.HEADER_ROW, column=5, value=gp.CONDITIONS[0].header_replicate_1
    )
    workbook.save(half_path)
    with pytest.raises(RuntimeError, match="2 columns are headed"):
        gp.read_half_lives(str(half_path))


# --------------------------------------------------------------------------- #
# Synthetic: the identifier route over the real genome class
# --------------------------------------------------------------------------- #
@pytest.fixture
def genome(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12MG1655Genome:
    """The real MG1655 genome class over the synthetic assembly; network refuses."""
    files = write_assembly(tmp_path / "tier", MG1655_ASSEMBLY, LOCI, GAF)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    return EcoliK12MG1655Genome(genome_root=str(tmp_path / "mg1655"), overwrite=True)


def test_the_uniprot_map_comes_from_the_pinned_assemblys_own_xrefs(
    genome: EcoliK12MG1655Genome,
) -> None:
    """Only the loci that carry a UniProt db_xref are in the map."""
    assert gp.uniprot_locus_tags(genome) == {
        "P00001": "b0001",
        "P00003": "b0003",
        "P00437": "b0437",
        "P03932": "b3932",
    }


def test_an_accession_claimed_by_two_loci_is_excluded(
    genome: EcoliK12MG1655Genome, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A one-to-many accession cannot key a gene, so it is not in the map at all."""
    table = genome.locus_table.copy()
    table.loc[table["locus_tag"] == "b0002", "db_xrefs"] = table.loc[
        table["locus_tag"] == "b0002", "db_xrefs"
    ].apply(lambda _: ("UniProtKB/Swiss-Prot:P00001",))
    monkeypatch.setattr(type(genome), "locus_table", property(lambda _: table))
    assert "P00001" not in gp.uniprot_locus_tags(genome)


def test_the_accession_layer_takes_precedence_over_a_disagreeing_symbol(
    genome: EcoliK12MG1655Genome,
) -> None:
    """The ``THRW`` row's gene name points at another locus; the accession wins."""
    names, mappings, route = gp.identifier_names(
        genome, [row[0] for row in ROWS], [row[1] for row in ROWS]
    )
    assert names == [
        "b0001",
        "b0002",
        "b0003",
        "b0437",
        "b0439",
        "b3932",
        "b2621",
        DROPPED,
    ]
    assert route.resolved_by_uniprot_xref == 4
    assert route.resolved_by_gene_name == 3
    assert route.resolved == len(ROWS) - 1
    assert route.unresolved == (DROPPED,)
    assert route.disagreements == (("sp|P00003|THRW_ECOLI", "thrL", "b0003", "b0001"),)
    assert route.assembly_uniprot_xrefs == 4
    assert route.unique_accessions == len(ROWS) - 1
    assert route.rows_per_route == {"uniprot_db_xref": 4, "gene_symbol": 3}
    assert [None if m is None else m.route for m in mappings] == [
        "uniprot_db_xref",
        "gene_symbol",
        "uniprot_db_xref",
        "uniprot_db_xref",
        "gene_symbol",
        "uniprot_db_xref",
        "gene_symbol",
        None,
    ]
    assert [None if m is None else m.source_identifier for m in mappings] == [
        "sp|P00001|THRL_ECOLI",
        "thrA",
        "sp|P00003|THRW_ECOLI",
        "sp|P00437|CLPP_ECOLI",
        "lon",
        "sp|P03932|HSLV_ECOLI",
        "smpB",
        None,
    ]


def test_a_db_xref_mapping_validates_the_released_accession_through_the_schema(
    genome: EcoliK12MG1655Genome,
) -> None:
    """The typed route reads the accession out of the ``sp|ACC|ENTRY`` header itself."""
    _, mappings, _ = gp.identifier_names(
        genome, [row[0] for row in ROWS], [row[1] for row in ROWS]
    )
    first = mappings[0]
    assert first is not None
    assert first.uniprot_accession() == "P00001"
    with pytest.raises(ValidationError):
        DerivedIdentifierMapping(
            source_identifier="not-an-accession", route="uniprot_db_xref"
        )


def test_a_row_neither_layer_resolves_carries_no_mapping(
    genome: EcoliK12MG1655Genome,
) -> None:
    """An unresolved row keys nothing, so there is no derivation to record for it."""
    names, mappings, route = gp.identifier_names(genome, [DROPPED], [""])
    assert names == [DROPPED]
    assert mappings == [None]
    assert route.unresolved == (DROPPED,)
    assert route.rows_per_route == {"uniprot_db_xref": 0, "gene_symbol": 0}


def test_a_deleted_symbol_that_resolves_nowhere_stops_the_build(
    genome: EcoliK12MG1655Genome, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A genotype cannot be keyed on a symbol the pinned assembly does not carry."""
    dataset = gp.ProteinTurnoverGupta2024Dataset.__new__(
        gp.ProteinTurnoverGupta2024Dataset
    )
    monkeypatch.setattr(
        gp,
        "CONDITIONS",
        tuple(
            condition.model_copy(update={"deleted_symbols": ("notAGene",)})
            if condition.key == "lon_nlim_6h"
            else condition
            for condition in gp.CONDITIONS
        ),
    )
    with pytest.raises(RuntimeError, match="does not resolve to one locus"):
        dataset._deleted_loci(genome)


# --------------------------------------------------------------------------- #
# Synthetic: the raw mirror
# --------------------------------------------------------------------------- #
@pytest.fixture
def pinned(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Path]:
    """The synthetic workbooks, with the module's pins set to their digests."""
    source = tmp_path / "source"
    source.mkdir()
    half_path, ci_path = write_workbooks(source)
    monkeypatch.setattr(
        gp, "HALF_LIVES_SHA256", hashlib.sha256(half_path.read_bytes()).hexdigest()
    )
    monkeypatch.setattr(
        gp, "CONFIDENCE_SHA256", hashlib.sha256(ci_path.read_bytes()).hexdigest()
    )
    return half_path, ci_path


def test_the_mirror_records_a_rerunnable_retrieval_for_every_file(
    tmp_path: Path, pinned: tuple[Path, Path]
) -> None:
    """Each deposited file carries its bucket key, its url and its sha256."""
    half_path, ci_path = pinned
    data_root = tmp_path / "data_root"
    root = gp.deposit_raw_mirror(
        half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
    )
    manifest = gp.load_manifest(str(data_root))
    assert {record.path for record in manifest.files} == {
        gp.HALF_LIVES_REL,
        gp.CONFIDENCE_REL,
    }
    for record in manifest.files:
        assert record.role == "raw_data"
        assert record.retrieval is not None
        assert record.retrieval.method == "pmc_cloud"
        source_url = record.retrieval.source_url
        assert source_url is not None
        assert source_url.endswith(osp.basename(record.path))
        assert record.sha256 == gp.manifest_sha256(manifest, record.path)
    assert (root / gp.HALF_LIVES_REL).is_file()
    assert manifest.doi == gp.DOI
    assert any(gp.PRIDE_ACCESSION in line for line in manifest.si_expected)


def test_depositing_twice_is_a_no_op(tmp_path: Path, pinned: tuple[Path, Path]) -> None:
    """Idempotent by sha256: a matching mirror file is left exactly as it is."""
    half_path, ci_path = pinned
    data_root = tmp_path / "data_root"
    gp.deposit_raw_mirror(
        half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
    )
    before = (data_root / gp.RAW_DIR_REL / gp.HALF_LIVES_REL).stat().st_mtime_ns
    gp.deposit_raw_mirror(
        half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
    )
    after = (data_root / gp.RAW_DIR_REL / gp.HALF_LIVES_REL).stat().st_mtime_ns
    assert before == after


def test_a_differing_mirror_file_is_refused_rather_than_overwritten(
    tmp_path: Path, pinned: tuple[Path, Path]
) -> None:
    """Upstream drift is detected, never followed."""
    half_path, ci_path = pinned
    data_root = tmp_path / "data_root"
    gp.deposit_raw_mirror(
        half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
    )
    target = data_root / gp.RAW_DIR_REL / gp.HALF_LIVES_REL
    target.write_bytes(b"different bytes")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        gp.deposit_raw_mirror(
            half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
        )


def test_a_source_file_off_its_pin_leaves_no_partial_deposit(
    tmp_path: Path, pinned: tuple[Path, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Both files are verified before anything is written."""
    half_path, ci_path = pinned
    monkeypatch.setattr(gp, "CONFIDENCE_SHA256", "0" * 64)
    data_root = tmp_path / "data_root"
    with pytest.raises(RuntimeError, match="hashes to"):
        gp.deposit_raw_mirror(
            half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
        )
    assert not (data_root / gp.RAW_DIR_REL).exists()


def test_a_file_outside_the_manifest_has_no_recorded_digest(
    tmp_path: Path, pinned: tuple[Path, Path]
) -> None:
    """``manifest_sha256`` names the file it cannot find instead of returning None."""
    half_path, ci_path = pinned
    data_root = tmp_path / "data_root"
    gp.deposit_raw_mirror(
        half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
    )
    manifest = gp.load_manifest(str(data_root))
    with pytest.raises(KeyError, match="si/nothing.xlsx"):
        gp.manifest_sha256(manifest, "si/nothing.xlsx")


# --------------------------------------------------------------------------- #
# Synthetic: the accounting ledger and the censoring oracle
# --------------------------------------------------------------------------- #
def _accounting(**overrides: Any) -> gp.BuildAccounting:
    from torchcell.datasets.bacteria_common import LocusTagReconciliation

    fields: dict[str, Any] = {
        "dataset": "x",
        "source_rows": 10,
        "released_cells": 100,
        "ceiling_cells": 1,
        "censored_keys": 1,
        "uncensored_keys": 8,
        "interval_keys": 2,
        "one_replicate_keys": 2,
        "refused_intervals": [],
        "confidence_level": gp.CONFIDENCE_LEVEL,
        "interval_method": gp.INTERVAL_METHOD,
        "measurement_type": gp.MEASUREMENT_TYPE,
        "dilution_rate_per_hour": {},
        "candidate_records": 13,
        "kept_records": 13,
        "dropped_records": 0,
        "kept_protein_keys": 9,
        "dropped_protein_keys": 1,
        "per_condition_keys": {},
        "per_condition_censored_keys": {},
        "per_condition_interval_keys": {},
        "rules": [],
        "identifier_route": gp.IdentifierRoute(
            source_rows=10,
            unique_accessions=10,
            assembly_uniprot_xrefs=4,
            resolved_by_uniprot_xref=5,
            resolved_by_gene_name=4,
            rows_per_route={"uniprot_db_xref": 5, "gene_symbol": 4},
            unresolved=(),
        ),
        "reconciliation": LocusTagReconciliation(
            label="x",
            assembly_set=gp.MG1655_ASSEMBLY_SET,
            gene_namespace=gp.MG1655_NAMESPACE,
            unique_names=10,
            status_histogram={},
            layer_histogram={},
            remapped=0,
            kept_on_collision=(),
            retired_kept=(),
            ambiguous_kept={},
            case_insensitive=(),
            outside_namespace=(),
        ),
        "notes": [],
    }
    fields.update(overrides)
    return gp.BuildAccounting(**fields)


def test_a_balanced_ledger_passes() -> None:
    """Kept plus dropped is what was read, on both axes."""
    accounting = _accounting()
    accounting.check()
    assert accounting.kept_records + accounting.dropped_records == (
        accounting.candidate_records
    )
    assert accounting.kept_protein_keys + accounting.dropped_protein_keys == (
        accounting.source_rows
    )


def test_an_unbalanced_record_ledger_is_refused() -> None:
    """A record that is neither kept nor accounted for is a silent loss."""
    with pytest.raises(RuntimeError, match="!= 13 candidates"):
        _accounting(kept_records=12).check()


def test_an_unbalanced_key_ledger_is_refused() -> None:
    """Every released row is either a key of some record or a named drop."""
    with pytest.raises(RuntimeError, match="released rows"):
        _accounting(kept_protein_keys=8).check()


def test_the_two_workbooks_must_agree_on_which_cells_are_censored() -> None:
    """``*`` in Supplementary Data 1 and ``Undetermined`` in 7 are one fact, twice."""
    keys = ["b0001"]
    cells = {c.key: _triple(8.0, 4.0, ceiling=True) for c in gp.CONDITIONS}
    good = {c.key: _pair("Undetermined", (3.5, 4.6)) for c in gp.CONDITIONS}
    gp.ProteinTurnoverGupta2024Dataset._check_undetermined_oracle(keys, cells, good)
    bad = dict(good)
    bad[gp.CONDITIONS[0].key] = _pair((7.0, 9.0), (3.5, 4.6))
    with pytest.raises(RuntimeError, match="disagree between Supplementary"):
        gp.ProteinTurnoverGupta2024Dataset._check_undetermined_oracle(keys, cells, bad)


# --------------------------------------------------------------------------- #
# Synthetic: the whole loader, end to end
# --------------------------------------------------------------------------- #
#: The synthetic release's stored key counts: 7 rows resolve, the isoform row is
#: dropped, and the first condition drops ``lon`` (absent) as well.
SYNTHETIC_STORED = {condition.key: 7 for condition in gp.CONDITIONS}
SYNTHETIC_STORED[gp.CONDITIONS[0].key] = 6
#: ``clpP`` is ceiling-flagged in replicate 1 of every condition, so exactly one stored
#: key per record is right-censored.
SYNTHETIC_CENSORED = {condition.key: 1 for condition in gp.CONDITIONS}
#: Only the first condition has a one-replicate key (``thrA``), so only it stores a
#: published interval.
SYNTHETIC_INTERVALS = {condition.key: 0 for condition in gp.CONDITIONS}
SYNTHETIC_INTERVALS[gp.CONDITIONS[0].key] = 1


def _pin(
    strain: Any, *, background: BacterialStrainBackground | None = None, **_: Any
) -> AssemblyReferenceGenome:
    """The assembly pin, without reading the genomes tier's assembly report."""
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain=strain if background is None else background.name,
        assembly_set=cast(Any, gp.MG1655_ASSEMBLY_SET),
        assembly_accession="GCA_000005845.2",
        background=background,
    )


@pytest.fixture
def mirrored(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    pinned: tuple[Path, Path],
    genome: EcoliK12MG1655Genome,
) -> Path:
    """A tmp ``DATA_ROOT`` holding the synthetic mirror, with the synthetic pins."""
    half_path, ci_path = pinned
    data_root = tmp_path / "data_root"
    gp.deposit_raw_mirror(
        half_lives_path=half_path, confidence_path=ci_path, data_root=str(data_root)
    )
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.delenv("TC_DATA_URL", raising=False)
    monkeypatch.setattr(gp, "SOURCE_ROWS", len(ROWS))
    monkeypatch.setattr(gp, "DROPPED_PROTEIN_KEYS", (DROPPED,))
    monkeypatch.setattr(gp, "MIN_RESOLVED_FRACTION", 0.8)
    monkeypatch.setattr(
        gp,
        "CONDITIONS",
        tuple(
            condition.model_copy(
                update={
                    "table1_proteins": SYNTHETIC_STORED[condition.key] + 1,
                    "stored_proteins": SYNTHETIC_STORED[condition.key],
                    "censored_keys": SYNTHETIC_CENSORED[condition.key],
                    "interval_keys": SYNTHETIC_INTERVALS[condition.key],
                }
            )
            for condition in gp.CONDITIONS
        ),
    )
    monkeypatch.setattr(gp, "assembly_reference", _pin)
    monkeypatch.setattr(
        gp, "bacterial_genome", lambda host, strain, data_root=None: genome
    )
    return data_root


def test_the_loader_builds_the_synthetic_release_end_to_end(
    tmp_path: Path, mirrored: Path, genome: EcoliK12MG1655Genome
) -> None:
    """One record per condition, with the counts, keys and gene set the ledger states."""
    root = tmp_path / "dataset"
    dataset = gp.ProteinTurnoverGupta2024Dataset(root=str(root), ecoli_genome=genome)
    assert len(dataset) == gp.EXPECTED_RECORDS
    assert sorted(dataset.gene_set) == ["b0437", "b0439", "b2621", "b3932"]
    references = dataset.experiment_reference_index
    assert references is not None

    records = [dataset[index] for index in range(len(dataset))]
    by_identity = {
        gp._record_identity({"experiment": record["experiment"]}): record
        for record in records
    }
    assert len(by_identity) == gp.EXPECTED_RECORDS
    wild_type = by_identity[
        gp._condition_identity(next(c for c in gp.CONDITIONS if c.key == "wt_nlim_6h"))
    ]
    assert wild_type["experiment"]["genotype"]["perturbations"] == []
    phenotype = wild_type["experiment"]["phenotype"]
    assert set(phenotype["degradation_rate"]) == {
        "b0001",
        "b0002",
        "b0003",
        "b0437",
        "b0439",
        "b2621",
        "b3932",
    }
    assert DROPPED not in phenotype["degradation_rate"]
    for key, rate in phenotype["degradation_rate"].items():
        assert rate == pytest.approx(math.log(2.0) / phenotype["half_life"][key])
    assert math.isnan(phenotype["degradation_rate_se"]["b0437"])
    assert phenotype["measurement_type"] == gp.MEASUREMENT_TYPE
    assert phenotype["confidence_level"] == gp.CONFIDENCE_LEVEL
    assert phenotype["interval_method"] == gp.INTERVAL_METHOD
    assert set(phenotype["censoring"]) == set(phenotype["degradation_rate"])
    assert phenotype["censoring"]["b0437"] == Censoring.right.value
    assert phenotype["censoring"]["b0001"] == Censoring.uncensored.value
    assert phenotype["degradation_rate_lower"] == {}
    assert phenotype["degradation_rate_upper"] == {}
    assert wild_type["experiment"]["environment"][
        "dilution_rate_per_hour"
    ] == pytest.approx(math.log(2.0) / 6.0)

    batch = by_identity[
        gp._condition_identity(
            next(c for c in gp.CONDITIONS if c.key == "wt_minimal_batch_42min")
        )
    ]
    assert batch["experiment"]["environment"]["dilution_rate_per_hour"] is None
    batch_phenotype = batch["experiment"]["phenotype"]
    assert set(batch_phenotype["degradation_rate_lower"]) == {"b0002"}
    assert set(batch_phenotype["degradation_rate_upper"]) == {"b0002"}
    low = batch_phenotype["degradation_rate_lower"]["b0002"]
    high = batch_phenotype["degradation_rate_upper"]["b0002"]
    assert low < batch_phenotype["degradation_rate"]["b0002"] < high
    assert batch_phenotype["n_replicates"]["b0002"] == 1

    triple = by_identity[
        gp._condition_identity(
            next(c for c in gp.CONDITIONS if c.key == "clpp_lon_hslv_nlim_6h")
        )
    ]
    assert sorted(
        p["systematic_gene_name"]
        for p in triple["experiment"]["genotype"]["perturbations"]
    ) == ["b0437", "b0439", "b3932"]

    accounting = json.loads((root / "preprocess" / "build_accounting.json").read_text())
    assert accounting["source_rows"] == len(ROWS)
    assert accounting["kept_records"] == gp.EXPECTED_RECORDS
    assert accounting["dropped_protein_keys"] == 1
    assert accounting["rules"][0]["items"] == [DROPPED]
    assert accounting["per_condition_keys"] == SYNTHETIC_STORED
    assert accounting["ceiling_cells"] == gp.EXPECTED_RECORDS
    assert accounting["censored_keys"] == gp.EXPECTED_RECORDS
    assert accounting["uncensored_keys"] == sum(SYNTHETIC_STORED.values()) - (
        gp.EXPECTED_RECORDS
    )
    assert accounting["interval_keys"] == 1
    assert accounting["one_replicate_keys"] == 1
    assert accounting["refused_intervals"] == []
    assert accounting["confidence_level"] == gp.CONFIDENCE_LEVEL
    assert accounting["interval_method"] == gp.INTERVAL_METHOD
    assert accounting["measurement_type"] == gp.MEASUREMENT_TYPE
    assert accounting["per_condition_censored_keys"] == SYNTHETIC_CENSORED
    assert accounting["per_condition_interval_keys"] == SYNTHETIC_INTERVALS
    assert accounting["dilution_rate_per_hour"]["wt_nlim_6h"] == pytest.approx(
        math.log(2.0) / 6.0
    )
    assert accounting["dilution_rate_per_hour"]["wt_minimal_batch_42min"] is None
    assert accounting["identifier_route"]["rows_per_route"] == {
        "uniprot_db_xref": 4,
        "gene_symbol": 3,
    }
    route = (root / "preprocess" / "identifier_route.csv").read_text()
    assert route.splitlines()[0] == (
        "protein_id,accession,released_gene_name,route,source_identifier,stored_key,kept"
    )
    assert (
        "sp|P00003|THRW_ECOLI,P00003,thrL,uniprot_db_xref,"
        "sp|P00003|THRW_ECOLI,b0003,yes" in route
    )
    assert "sp|P00002|THRA_ECOLI,P00002,thrA,gene_symbol,thrA,b0002,yes" in route
    assert f"{DROPPED},P00002,,,,{DROPPED},no" in route
    ceilings = (root / "preprocess" / "ceiling_cells.csv").read_text()
    assert ceilings.count("b0437,1,") == gp.EXPECTED_RECORDS
    assert (root / "preprocess" / "build_manifest.json").is_file()
    assert (root / "raw" / gp.HALF_LIVES_FILENAME).is_file()


def test_the_verifier_passes_on_the_synthetic_build(
    tmp_path: Path, mirrored: Path, genome: EcoliK12MG1655Genome
) -> None:
    """The loader-level L0-L4 gate, over the build the fixture above makes."""
    root = tmp_path / "dataset"
    gp.ProteinTurnoverGupta2024Dataset(root=str(root), ecoli_genome=genome)
    report = gp.verify_build(str(root), genome=genome)
    assert report.passed, report.summary()
    names = {result.name for result in report.results}
    assert "degradation_rate_is_ln2_over_half_life" in names
    assert "stored_protein_keys_are_loci_of_the_pinned_assembly" in names
    assert "dilution_rate_per_hour_is_ln2_over_the_doubling_time" in names
    assert "stored_interval_brackets_the_stored_rate" in names
    assert "per_condition_censored_and_interval_key_counts_are_the_pinned_ones" in names
    assert (root / "preprocess" / "verification_report.json").is_file()


def test_the_verifier_fails_a_record_count_that_is_not_the_pinned_one(
    tmp_path: Path, mirrored: Path, genome: EcoliK12MG1655Genome
) -> None:
    """L1 is an oracle, so a wrong expectation must FAIL rather than be tolerated."""
    root = tmp_path / "dataset"
    gp.ProteinTurnoverGupta2024Dataset(root=str(root), ecoli_genome=genome)
    report = gp.verify_build(str(root), genome=genome, expected_count=12)
    assert not report.passed


def test_a_direct_run_opens_the_mg1655_genome_itself(
    tmp_path: Path, mirrored: Path, genome: EcoliK12MG1655Genome
) -> None:
    """A loader built without an injected genome resolves one from the tier."""
    dataset = gp.ProteinTurnoverGupta2024Dataset.__new__(
        gp.ProteinTurnoverGupta2024Dataset
    )
    dataset.ecoli_genome = None
    assert dataset._genome() is genome
    assert dataset.ecoli_genome is genome


def test_the_loader_declares_its_schema_pair_and_its_raw_files() -> None:
    """The classes ``transform_item`` validates against, and the two workbooks."""
    dataset = gp.ProteinTurnoverGupta2024Dataset.__new__(
        gp.ProteinTurnoverGupta2024Dataset
    )
    assert dataset.experiment_class is ProteinTurnoverExperiment
    assert dataset.reference_class is ProteinTurnoverExperimentReference
    assert dataset.raw_file_names == [gp.HALF_LIVES_FILENAME, gp.CONFIDENCE_FILENAME]
    assert gp.ProteinTurnoverGupta2024Dataset.REFERENCE_STRAIN == "MG1655"
    with pytest.raises(NotImplementedError):
        dataset.create_experiment()
    assert dataset.preprocess_raw("x") == "x"


def test_the_publication_is_this_paper() -> None:
    """DOI only: the mirror's manifest records no PubMed id for this key."""
    publication = gp.publication()
    assert publication.doi == gp.DOI
    doi_url = publication.doi_url
    assert doi_url is not None
    assert doi_url.endswith(gp.DOI)


# --------------------------------------------------------------------------- #
# The REAL released bytes: skipped unless the raw mirror is on this machine
# --------------------------------------------------------------------------- #
def _mirror_path(relpath: str) -> Path | None:
    import os

    data_root = os.environ.get("DATA_ROOT")
    if not data_root:
        return None
    path = Path(data_root) / gp.RAW_DIR_REL / relpath
    return path if path.is_file() else None


requires_mirror = pytest.mark.skipif(
    _mirror_path(gp.HALF_LIVES_REL) is None or _mirror_path(gp.CONFIDENCE_REL) is None,
    reason="the Gupta 2024 raw mirror is not on this machine",
)


@requires_mirror
def test_the_released_workbooks_are_on_their_pins() -> None:
    """The deposited bytes still hash to what the loader pins."""
    for relpath, pin in (
        (gp.HALF_LIVES_REL, gp.HALF_LIVES_SHA256),
        (gp.CONFIDENCE_REL, gp.CONFIDENCE_SHA256),
    ):
        path = _mirror_path(relpath)
        assert path is not None
        assert hashlib.sha256(path.read_bytes()).hexdigest() == pin


@requires_mirror
def test_the_released_per_condition_counts_are_table_1s() -> None:
    """The cross-source oracle on the real bytes: 13 of 13 conditions agree."""
    path = _mirror_path(gp.HALF_LIVES_REL)
    assert path is not None
    protein_ids, _, cells = gp.read_half_lives(str(path))
    assert len(protein_ids) == gp.SOURCE_ROWS
    for condition in gp.CONDITIONS:
        triples = cells[condition.key]
        released = sum(
            1 for row in range(len(protein_ids)) if triples[3 * row + 2] is not None
        )
        assert released == condition.table1_proteins, condition.key


@requires_mirror
def test_the_two_released_workbooks_agree_on_every_censored_cell() -> None:
    """61,811 released cells, and the ``*`` and ``Undetermined`` flags never disagree."""
    half_path = _mirror_path(gp.HALF_LIVES_REL)
    ci_path = _mirror_path(gp.CONFIDENCE_REL)
    assert half_path is not None and ci_path is not None
    protein_ids, _, cells = gp.read_half_lives(str(half_path))
    ci_ids, intervals = gp.read_confidence_intervals(str(ci_path))
    assert protein_ids == ci_ids
    gp.ProteinTurnoverGupta2024Dataset._check_undetermined_oracle(
        protein_ids, cells, intervals
    )
    # The flat triples hold two replicate cells then the authors' mean cell.
    replicates = [
        cell
        for condition in gp.CONDITIONS
        for index, cell in enumerate(cells[condition.key])
        if index % 3 != 2 and cell is not None
    ]
    means = [
        cell
        for condition in gp.CONDITIONS
        for index, cell in enumerate(cells[condition.key])
        if index % 3 == 2 and cell is not None
    ]
    assert len(replicates) == 61811
    assert len(means) == sum(c.table1_proteins for c in gp.CONDITIONS) == 33201
    assert sum(1 for cell in replicates if cell.ceiling) == 2084
    # No mean cell is itself flagged: the ceiling lives on the replicate it capped.
    assert not any(cell.ceiling for cell in means)


@requires_mirror
def test_the_released_intervals_reproduce_the_pinned_censoring_and_bound_counts() -> (
    None
):
    """The #753 maps, measured on the real bytes: 2,082 cells, 1,989 keys, 4,270 bounds.

    Keyed on the released protein ids rather than locus tags, which needs no genome: the
    two isoform rows are then popped exactly as ``process`` pops them, so the per-condition
    numbers are the ones each ``Condition`` pins.
    """
    half_path = _mirror_path(gp.HALF_LIVES_REL)
    ci_path = _mirror_path(gp.CONFIDENCE_REL)
    assert half_path is not None and ci_path is not None
    protein_ids, _, cells = gp.read_half_lives(str(half_path))
    _, intervals = gp.read_confidence_intervals(str(ci_path))
    censored_cells = 0
    censored_keys = 0
    uncensored_keys = 0
    interval_keys = 0
    one_replicate_keys = 0
    refused = 0
    for condition in gp.CONDITIONS:
        values = gp.condition_values(
            condition, protein_ids, cells[condition.key], intervals[condition.key]
        )
        for key in gp.DROPPED_PROTEIN_KEYS:
            values.half_life.pop(key, None)
            values.degradation_rate.pop(key, None)
            values.degradation_rate_lower.pop(key, None)
            values.degradation_rate_upper.pop(key, None)
            values.censoring.pop(key, None)
            values.n_replicates.pop(key, None)
        kept = set(values.half_life)
        this_censored = sum(
            1 for state in values.censoring.values() if state is Censoring.right
        )
        assert this_censored == condition.censored_keys, condition.key
        assert len(values.degradation_rate_lower) == condition.interval_keys, (
            condition.key
        )
        censored_cells += sum(1 for key, _ in values.ceiling_cells if key in kept)
        censored_keys += this_censored
        uncensored_keys += sum(
            1 for state in values.censoring.values() if state is Censoring.uncensored
        )
        interval_keys += len(values.degradation_rate_lower)
        one_replicate_keys += sum(1 for n in values.n_replicates.values() if n == 1)
        refused += sum(1 for r in values.refused_intervals if r.locus_tag in kept)
    assert censored_cells == 2082
    assert censored_keys == 1989
    assert uncensored_keys == 31198
    assert censored_keys + uncensored_keys == 33187
    assert one_replicate_keys == 4587
    assert interval_keys == 4270
    assert refused == 25
    assert interval_keys + refused + 292 == one_replicate_keys


@requires_mirror
def test_every_sourced_quote_is_still_verbatim_in_its_pinned_mirror_file() -> None:
    """A quote that drifted is a sourced value that no longer says what it claims."""
    import os

    data_root = os.environ.get("DATA_ROOT")
    assert data_root
    library = Path(data_root) / "torchcell-library" / gp.CITATION_KEY
    if not (library / gp.PAPER_MD).is_file():
        pytest.skip("the Gupta 2024 OCR mirror is not on this machine")
    texts = {
        gp.PAPER_MD: (library / gp.PAPER_MD).read_text(),
        gp.SI1_MD: (library / gp.SI1_MD).read_text(),
    }
    for sourced in (
        gp.CONFIDENCE_LEVEL_SOURCE,
        gp.INTERVAL_FROM_CURVE_FIT_VARIANCE,
        gp.INTERVAL_T_QUANTILE,
        gp.INTERVAL_DOF,
        gp.DILUTION_IS_THE_CHEMOSTAT_VARIABLE,
        gp.DILUTION_LIMIT_IS_THE_DOUBLING_TIME,
        gp.DOUBLING_TIMES_MEASURED,
        gp.DOUBLING_TIMES_CHEMOSTAT,
    ):
        text = texts[sourced.provenance.source_uri]
        assert sourced.quote in text, sourced.quote
        assert (
            hashlib.sha256(
                (library / sourced.provenance.source_uri).read_bytes()
            ).hexdigest()
            == sourced.provenance.sha256
        )


@requires_mirror
def test_every_released_half_life_is_strictly_positive() -> None:
    """Why ``ln 2 / T`` is always finite on this release."""
    path = _mirror_path(gp.HALF_LIVES_REL)
    assert path is not None
    _, _, cells = gp.read_half_lives(str(path))
    values = [
        cell.hours
        for condition in gp.CONDITIONS
        for cell in cells[condition.key]
        if cell is not None
    ]
    assert min(values) > 0.0
