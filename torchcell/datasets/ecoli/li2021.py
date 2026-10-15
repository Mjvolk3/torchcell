# torchcell/datasets/ecoli/li2021
# [[torchcell.datasets.ecoli.li2021]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/li2021
# Test file: tests/torchcell/datasets/ecoli/test_li2021.py
"""Li 2021: the three released 13C-MFA flux maps of the zwf-promoter EP-bifido strains.

Li, ..., Xian 2021 (Microb Cell Fact 20:32, doi:10.1186/s12934-021-01526-1, citation key
``liFineTuningGlycolytic2021``) replaced the native promoter of ``zwf`` in the
mevalonate-producing EP-bifido strain BW-P (BW25113 ``∆pfkA:FRT`` carrying pBSA and pFF)
with synthetic promoters, and ran 13C metabolic flux analysis on three of them.
Additional file 2 ("MFA simulated result.", ``12934_2021_1526_MOESM2_ESM.xlsx``) releases
the fits.

WHAT IS STORED. One ``FluxExperiment`` per ``NETWORK`` block the workbook releases,
under the block's own strain label: ``BW-P08``, ``BW-P04`` and ``BW-P10``. Each record is
the block's 99 net fluxes in mmol/gdw/h (the absolute columns, keyed by the release's
reaction names, signed and unclamped), their ``LB90`` and ``UB90`` bounds at
``confidence_level=0.90``, and the block's exchange-flux table (``vAco.x`` ...) with its
bounds, verbatim. Record count: 3.

THE REFERENCE IS A TYPED GAP, NOT A REFUSAL. The paper names the control strain BW-P BF as
the third strain of the analysis and draws its map in Fig. 5, but the workbook releases no
block for it. The reference is BW-P BF on the same labeling medium with its ``net_flux`` a
``ProvenanceGap`` (``CONTROL_MAP_GAP``); the experiment records are the released maps.

WHAT THE RELEASE SAYS ABOUT ITSELF, recorded as findings rather than acted on (the data
is the data, so none of them changes what is stored):

* The workbook releases a ``BW-P04`` block the paper never mentions.
* 50 of the 99 reaction rows are bit-identical between the ``BW-P04`` and ``BW-P10``
  blocks (best fit and both bounds), with equal glucose uptake, so the ``BW-P10`` block
  carries the extracellular constraints of the ``BW-P04`` block (issue #800 measured this
  by keying each block's byproduct fluxes to the ``Biomass Composition`` yields). The
  block is stored under the label the release gives it.
* The ``χ2 90%`` cell is 102.8 for ``BW-P08`` and 102.08 for the other two although all
  three state 85 degrees of freedom.
* Neither the interval procedure nor a replicate count is stated anywhere in the paper,
  so ``n_samples`` and ``sample_unit`` are typed gaps; the bounds are stored as released.
* The labeling culture's medium name, physical state, temperature and duration are not
  stated; the tracer and its concentration are, and they are the medium's one component.

NOT STORED, each with its count, in ``preprocess/not_stored.json``: the normalized
columns (3 blocks x 99 reactions x 3 numbers), which are the absolute fit divided by the
block's own released ``vUPTU`` and so restate the stored map on a second scale; the four
fit-diagnostic rows (3 x 4 cells), which describe the fit and have no field; and the
``LABEL_INPUT``, ``LABEL_MEASUREMENTS`` and ``Biomass Composition`` sheets, which are the
fit's inputs rather than its result.

GENOTYPE. Each record's edit relative to the reference BW-P BF is the ``zwf`` promoter
replacement Table 1 states for its strain (BBa-J 23108, 23104 or 23100). ``zwf`` is keyed
to BW25113's locus through the ``gene_symbol`` route and the released symbol is kept as
``identifier_mapping.source_identifier``. The constant content (``∆pfkA:FRT``, pBSA, pFF)
is the reference's strain background, stated verbatim.

DATA. ``$DATA_ROOT/torchcell-raw/liFineTuningGlycolytic2021/data/`` holds the workbook,
copied sha256-verified from the literature mirror's ``si/si2.xlsx`` with the PMC Article
Datasets retrieval recorded (:func:`deposit_raw_mirror`).
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import os.path as osp
import shutil
from collections.abc import Callable, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, ClassVar, Final

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict

from torchcell.data import (
    ExperimentDataset,
    Organism,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.data.experiment_dataset import verify_sha256
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.schema import (
    BacterialAssemblySet,
    BacterialStrainBackground,
    Concentration,
    ConcentrationUnit,
    DerivedIdentifierMapping,
    Environment,
    Experiment,
    ExperimentReference,
    FluxExperiment,
    FluxExperimentReference,
    FluxPhenotype,
    Genotype,
    Media,
    MediaComponent,
    MediaComponentRole,
    PromoterReplacementPerturbation,
    Publication,
    ReferenceGenome,
)
from torchcell.datasets.bacteria_common import (
    STRAIN_GENE_NAMESPACES,
    assembly_reference,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.sequence.genome.ecoli.k12 import EcoliK12StrainName
from torchcell.verification.report import (
    Level,
    LevelResult,
    Provenance,
    VerificationReport,
)
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
    audit_sourced_value,
)

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
CITATION_KEY: Final = "liFineTuningGlycolytic2021"
PAPER_DOI: Final = "10.1186/s12934-021-01526-1"
PAPER_TITLE: Final = (
    "Fine tuning the glycolytic flux ratio of EP-bifido pathway for mevalonate "
    "production by enhancing glucose-6-phosphate dehydrogenase (Zwf) and CRISPRi "
    "suppressing 6-phosphofructose kinase (PfkA) in Escherichia coli"
)
PAPER_MD: Final = "paper.md"
PAPER_MD_SHA256: Final = (
    "7d174c8ef7af0e0c275f6194271d39346c22411c27904803ee16b7caf95b7773"
)
RAW_DIR_REL: Final = f"torchcell-raw/{CITATION_KEY}"
LIBRARY_DIR_REL: Final = f"torchcell-library/{CITATION_KEY}"
DATASET_ROOT_REL: Final = "data/torchcell/flux_li2021"

WORKBOOK_NAME: Final = "12934_2021_1526_MOESM2_ESM.xlsx"
WORKBOOK_RELPATH: Final = f"data/{WORKBOOK_NAME}"
WORKBOOK_LIBRARY_RELPATH: Final = "si/si2.xlsx"
WORKBOOK_SHA256: Final = (
    "24c836f8459e5cce4f6a322477a7554e1f7dd16e8bb6b7cdd0ac4890b24c585a"
)
WORKBOOK_URL: Final = (
    "https://pmc-oa-opendata.s3.amazonaws.com/PMC7852082.1/"
    "12934_2021_1526_MOESM2_ESM.xlsx"
)
WORKBOOK_RETRIEVED_AT: Final = "2026-10-07T11:41:55.853631+00:00"

PAPER: Final = Provenance(
    source_uri=PAPER_MD, citation_key=CITATION_KEY, sha256=PAPER_MD_SHA256
)
PUBLICATION: Final = Publication(doi=PAPER_DOI, doi_url=f"https://doi.org/{PAPER_DOI}")

REFERENCE_STRAIN: Final[EcoliK12StrainName] = "BW25113"
ASSEMBLY_SET: Final[BacterialAssemblySet] = "ecoli_K12_BW25113_ASM75055v1"
GENE_NAMESPACE: Final = STRAIN_GENE_NAMESPACES[REFERENCE_STRAIN]
#: ``zwf`` on the pinned BW25113 GenBank annotation, measured with
#: ``resolve_gene_name("zwf")`` (status ``renamed``: the gene symbol of current gene
#: ``BW25113_1852``); the L4 rule re-checks it is a gene row of the assembly.
ZWF: Final = ("zwf", "BW25113_1852")

MEASUREMENT_TYPE: Final = "c13_mfa_net_flux_mmol_per_gdw_per_h"


def _paper(
    value: Any, quote: str, *, page: str, note: str | None = None
) -> SourcedValue:
    """A value backed by a verbatim quote of the OCR'd paper."""
    return SourcedValue(
        value=value,
        provenance=PAPER.model_copy(update={"page": page}),
        quote=quote,
        note=note,
    )


SOURCED_VALUES: Final[dict[str, SourcedValue]] = {
    "analysed_strains": _paper(
        ["BW-P08 BF", "BW-P10 BF", "BW-P BF"],
        "MFA was performed to detect the metabolic fux distribution in strains BW-P08 "
        "BF and BW-P10 BF (which had high MVA titers) and the control strain BW-P BF.",
        page="Discussion",
        note="the control BW-P BF is named as analysed; its map is not in the release",
    ),
    "figure_5_strains": _paper(
        ["BW-P10 BF", "BW-P08 BF", "BW-P BF"],
        "Fig. 5 Metabolic fux diagram of zwf-strengthened EP-bifdo strains. The "
        "metabolic fux shown for strains from top to bottom are BW-P10 BF, BW-P08 BF, "
        "BW-P BF",
        page="Fig. 5 caption",
    ),
    "released_file": _paper(
        "Additional file 2",
        # the OCR text carries a no-break space (U+00A0) after "fle", kept verbatim
        "Additional fle\u00a02. MFA simulated result.",
        page="Supplementary Information",
        note="the paper's only name for the workbook; the Methods describe a "
        "least-squares fit, so 'simulated' is the authors' word and is kept",
    ),
    "fit": _paper(
        "13C-Flux least-squares fit",
        "Metabolic fuxes were estimated by minimizing the residual sum of squares "
        "between experimentally measured and model predicted",
        page="Methods, 13C-MFA",
    ),
    "tracer": _paper(
        {"tracer": "100% [1-13C1]glucose", "g_per_l": 10.0},
        "13C-MFA was performed using $1 0 0 \\% 1 \\mathrm { { - } ^ { 1 3 } C _ { 1 } }$ "
        "glucose as the feeding substrate was added to a concentration of $1 0 ~ "
        "\\mathrm { g / L }$",
        page="Methods, 13C-MFA",
    ),
    "harvest": _paper(
        "exponential phase",
        "Cells at the exponential growth phase were harvested by centrifugation",
        page="Methods, 13C-MFA",
        note="a growth phase, not a time: duration_hours stays None",
    ),
    "control_strain": _paper(
        "BW-P BF",
        "transformed into the fve zwf-enhanced strains and cultivated with the control "
        "strain BW-P/pBSA pFF (abbreviated to BW-P BF).",
        page="Results",
        note="'fve zwf-enhanced strains' is also the direction of every promoter swap",
    ),
    "bw_p": _paper(
        "BW25113 derivative, ∆pfkA:FRT",
        "<td>BW-P</td><td>BW25113 derivative, ∆pfkA:FRT</td>",
        page="Table 1",
    ),
    "pff": _paper(
        "pFF",
        "<td>pFF</td><td>pCDFduet, with fxpk from B. Adolescentis, and fbp gene from "
        "E. coli, AmpR</td>",
        page="Table 1",
    ),
    "pbsa": _paper(
        "pBSA",
        "<td>pBSA</td><td>pTr9a with atoB from E. coli, mvaS and mvaA from L. casei. tac "
        "promoter, Amp</td>",
        page="Table 1",
    ),
    "promoter_strategy": _paper(
        "increased",
        "we fnely tuned the glycolytic fux ratio in two ways. First, we enhanced PPP fux "
        "for NADPH supply by replacing the promoter of zwf on the genome with a set of "
        "diferent strength promoters",
        page="Abstract",
    ),
}


class StrainSpec(BaseModel):
    """One NETWORK block: its released label and the Table 1 strain it names."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    block_label: str
    label_cell: str
    unit_cell: str
    absolute_columns: tuple[str, str, str]
    exchange_name_column: str
    exchange_columns: tuple[str, str, str]
    exchange_first_row: int
    strain: str
    promoter: str
    table_1: SourcedValue
    promoter_strain: SourcedValue


def _table_1(strain: str, statement: str) -> SourcedValue:
    return _paper(strain, f"<td>{strain}</td><td>{statement}</td>", page="Table 1")


#: The three blocks of ``NETWORK``, read off rows 1, 7, 8 and 327 to 328, with the Table 1
#: strain each label names. ``BW-P08`` is Table 1's ``BW-P 08 BF``: Table 1 writes that
#: label with a space, the NETWORK sheet and the Discussion without one.
STRAINS: Final[tuple[StrainSpec, ...]] = (
    StrainSpec(
        block_label="BW-P08",
        label_cell="F1",
        unit_cell="D7",
        absolute_columns=("D", "E", "F"),
        exchange_name_column="B",
        exchange_columns=("E", "F", "G"),
        exchange_first_row=329,
        strain="BW-P08 BF",
        promoter="BBa-J 23108",
        table_1=_table_1("BW-P 08 BF", "BW-P 23108-zwf carrying pBSA pFF"),
        promoter_strain=_table_1(
            "BW-P 23108-zwf", "BW-P derived, zwf promoter: BBa-J 23108, CmR"
        ),
    ),
    StrainSpec(
        block_label="BW-P04",
        label_cell="N1",
        unit_cell="L7",
        absolute_columns=("K", "L", "M"),
        exchange_name_column="K",
        exchange_columns=("L", "M", "N"),
        exchange_first_row=328,
        strain="BW-P04 BF",
        promoter="BBa-J 23104",
        table_1=_table_1("BW-P04 BF", "BW-P 23104-zwf carrying pBSA pFF"),
        promoter_strain=_table_1(
            "BW-P 23104-zwf", "BW-P derived, zwf promoter: BBa-J 23104, CmR"
        ),
    ),
    StrainSpec(
        block_label="BW-P10",
        label_cell="S1",
        unit_cell="R7",
        absolute_columns=("R", "S", "T"),
        exchange_name_column="S",
        exchange_columns=("T", "U", "V"),
        exchange_first_row=328,
        strain="BW-P10 BF",
        promoter="BBa-J 23100",
        table_1=_table_1("BW-P10 BF", "BW-P 23100-zwf carrying pBSA pFF"),
        promoter_strain=_table_1(
            "BW-P 23100-zwf", "BW-P derived, zwf promoter: BBa-J 23100, CmR"
        ),
    ),
)

#: The workbook cells whose text is the only statement of the unit and the level.
UNIT_HEADER: Final = "Absolute  Flux  value (mmol/gdw/h)"
BOUND_HEADERS: Final = ("best fit", "LB90", "UB90")
CONFIDENCE_LEVEL: Final = 0.90
NET_FLUX_ROWS: Final = (9, 325)
EXCHANGE_HEADER_CELL: Final = ("B", 327, "Exchange flux ")
DIAGNOSTIC_LABELS: Final = (
    "Number of fitted measurements :",
    "Freedom of flux",
    "χ2 90%",
    "SSR :",
)

#: Build oracles, measured on the pinned workbook before being written here.
EXPECTED_RECORDS: Final = 3
EXPECTED_REACTIONS: Final = 99
EXPECTED_EXCHANGES: Final = {"BW-P08": 24, "BW-P04": 23, "BW-P10": 23}

_WORKBOOK_LOOKED_IN = Provenance(
    source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{WORKBOOK_RELPATH}",
    citation_key=CITATION_KEY,
    sha256=WORKBOOK_SHA256,
    page="NETWORK, rows 1 to 8",
)
_PAPER_LOOKED_IN = PAPER.model_copy(update={"page": "Methods and Results, full text"})

CONTROL_MAP_GAP: Final = ProvenanceGap(
    field="net_flux",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_WORKBOOK_LOOKED_IN,
    note="the paper names the control BW-P BF as analysed and draws its map in Fig. 5, "
    "but Additional file 2's NETWORK sheet releases blocks for BW-P08, BW-P04 and "
    "BW-P10 only; columns A, J, Q and X to AB are empty, so no fourth block exists",
)
N_SAMPLES_GAP: Final = ProvenanceGap(
    field="n_samples",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_PAPER_LOOKED_IN,
    note="no replicate count for the labeling experiment is stated, and the "
    "LABEL_MEASUREMENTS sheet's uncertainty columns G to K are empty",
)
SAMPLE_UNIT_GAP: Final = ProvenanceGap(
    field="sample_unit",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_PAPER_LOOKED_IN,
    note="what one labeling experiment is, is not stated",
)
RECORD_GAPS: Final = (N_SAMPLES_GAP, SAMPLE_UNIT_GAP)
MEDIUM_GAPS: Final = tuple(
    ProvenanceGap(
        field=field,
        reason=ProvenanceGapReason.not_reported_by_primary,
        looked_in=_PAPER_LOOKED_IN,
        note="the 13C-MFA paragraph states only the tracer and its concentration; the "
        "one medium the paper does describe (M9 with 0.2% yeast extract and 20 g/L "
        "glucose, for MVA production) carries twice the labeling run's glucose, so it is "
        "a different condition and is not borrowed",
    )
    for field in ("name", "state", "is_synthetic")
)
TEMPERATURE_GAP: Final = ProvenanceGap(
    field="temperature",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_PAPER_LOOKED_IN,
    note="no temperature is stated for the labeling culture",
)


# --------------------------------------------------------------------------- #
# Reading the workbook
# --------------------------------------------------------------------------- #
class WorkbookFormatError(ValueError):
    """The workbook's cells are not where the pinned layout puts them."""


class BlockMap(BaseModel):
    """One released NETWORK block, verbatim."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    block_label: str
    net_flux: dict[str, float]
    net_flux_lower: dict[str, float]
    net_flux_upper: dict[str, float]
    exchange_flux: dict[str, float]
    exchange_flux_lower: dict[str, float]
    exchange_flux_upper: dict[str, float]
    diagnostics: dict[str, float]


def _number(value: Any, where: str) -> float:
    """A numeric cell, refusing text or an empty cell (no value is invented)."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise WorkbookFormatError(f"{where} is {value!r}, not a number")
    return float(value)


def _require(ws: Any, cell: str, expected: str) -> None:
    observed = ws[cell].value
    if observed != expected:
        raise WorkbookFormatError(
            f"NETWORK {cell} is {observed!r}, expected {expected!r}"
        )


def reaction_rows(ws: Any) -> list[tuple[int, str]]:
    """``(row, reaction name)`` of every net-flux reaction row, comments excluded."""
    first, last = NET_FLUX_ROWS
    rows: list[tuple[int, str]] = []
    for row in range(first, last + 1):
        name = ws[f"B{row}"].value
        if name is None or str(name).strip().startswith("//"):
            continue
        rows.append((row, str(name).strip()))
    return rows


def read_block(ws: Any, spec: StrainSpec, rows: Sequence[tuple[int, str]]) -> BlockMap:
    """One block's net fluxes, bounds, exchange table and fit diagnostics, verbatim."""
    _require(ws, spec.label_cell, spec.block_label)
    best, lower, upper = spec.absolute_columns
    _require(ws, spec.unit_cell, UNIT_HEADER)
    for column, header in zip(spec.absolute_columns, BOUND_HEADERS, strict=True):
        _require(ws, f"{column}8", header)
    label_column = spec.label_cell[0]
    diagnostics: dict[str, float] = {}
    for offset, label in enumerate(DIAGNOSTIC_LABELS):
        _require(ws, f"B{offset + 2}", label)
        diagnostics[label] = _number(
            ws[f"{label_column}{offset + 2}"].value, f"{label_column}{offset + 2}"
        )
    net: dict[str, float] = {}
    net_lower: dict[str, float] = {}
    net_upper: dict[str, float] = {}
    for row, name in rows:
        if name in net:
            raise WorkbookFormatError(f"reaction {name} appears twice")
        net[name] = _number(ws[f"{best}{row}"].value, f"{best}{row}")
        net_lower[name] = _number(ws[f"{lower}{row}"].value, f"{lower}{row}")
        net_upper[name] = _number(ws[f"{upper}{row}"].value, f"{upper}{row}")
    exchange: dict[str, float] = {}
    exchange_lower: dict[str, float] = {}
    exchange_upper: dict[str, float] = {}
    e_best, e_lower, e_upper = spec.exchange_columns
    row = spec.exchange_first_row
    while (name := ws[f"{spec.exchange_name_column}{row}"].value) is not None:
        key = str(name).strip()
        if key in exchange:
            raise WorkbookFormatError(f"exchange {key} appears twice")
        exchange[key] = _number(ws[f"{e_best}{row}"].value, f"{e_best}{row}")
        exchange_lower[key] = _number(ws[f"{e_lower}{row}"].value, f"{e_lower}{row}")
        exchange_upper[key] = _number(ws[f"{e_upper}{row}"].value, f"{e_upper}{row}")
        row += 1
    return BlockMap(
        block_label=spec.block_label,
        net_flux=net,
        net_flux_lower=net_lower,
        net_flux_upper=net_upper,
        exchange_flux=exchange,
        exchange_flux_lower=exchange_lower,
        exchange_flux_upper=exchange_upper,
        diagnostics=diagnostics,
    )


def read_workbook(path: str | Path) -> list[BlockMap]:
    """The three released blocks, refusing a layout or count that is not the pinned one."""
    workbook = load_workbook(str(path), data_only=True)
    ws = workbook["NETWORK"]
    column, row, text = EXCHANGE_HEADER_CELL
    _require(ws, f"{column}{row}", text)
    rows = reaction_rows(ws)
    if len(rows) != EXPECTED_REACTIONS:
        raise WorkbookFormatError(
            f"{len(rows)} reaction rows, the module declares {EXPECTED_REACTIONS}"
        )
    blocks = [read_block(ws, spec, rows) for spec in STRAINS]
    for block in blocks:
        n_exchange = len(block.exchange_flux)
        if n_exchange != EXPECTED_EXCHANGES[block.block_label]:
            raise WorkbookFormatError(
                f"{block.block_label}: {n_exchange} exchange rows, the module declares "
                f"{EXPECTED_EXCHANGES[block.block_label]}"
            )
    return blocks


def identical_rows(left: BlockMap, right: BlockMap) -> list[str]:
    """Reaction rows whose best fit AND both bounds are equal between two blocks."""
    return sorted(
        name
        for name in left.net_flux
        if left.net_flux[name] == right.net_flux[name]
        and left.net_flux_lower[name] == right.net_flux_lower[name]
        and left.net_flux_upper[name] == right.net_flux_upper[name]
    )


# --------------------------------------------------------------------------- #
# Records (pure)
# --------------------------------------------------------------------------- #
def labeling_environment() -> Environment:
    """The labeling culture: its tracer at 10 g/L; everything else a typed gap.

    ``aerobicity`` is the field's default. The schema has no gap for it (a non-optional
    ``str``), and the labeling run's aeration is not stated; the dataset note says so.
    """
    tracer = SOURCED_VALUES["tracer"]
    return Environment(
        media=Media(
            name=None,
            state=None,
            is_synthetic=None,
            components=[
                MediaComponent(
                    compound=resolved_compound("D-glucose"),
                    role=MediaComponentRole.carbon_source,
                    concentration=Concentration(
                        value=float(tracer.value["g_per_l"]),
                        unit=ConcentrationUnit.g_per_l,
                    ),
                    provenance=[tracer],
                    note="100% [1-13C1]glucose, the labeling tracer; the compound "
                    "table has no isotopologue entry, so the label is stated here",
                )
            ],
            provenance=[tracer],
            provenance_gaps=list(MEDIUM_GAPS),
        ),
        temperature=None,
        provenance_gaps=[TEMPERATURE_GAP],
    )


def reference_background() -> BacterialStrainBackground:
    """BW-P BF: BW25113 ``∆pfkA:FRT`` carrying pBSA and pFF, verbatim, no typed allele."""
    return BacterialStrainBackground(
        name="BW-P BF",
        reference_strain=REFERENCE_STRAIN,
        assembly_set=ASSEMBLY_SET,
        parents=["BW-P"],
        construction="BW-P/pBSA pFF (abbreviated to BW-P BF)",
        genotype_statement="BW25113 derivative, ∆pfkA:FRT",
        provenance=[
            SOURCED_VALUES["control_strain"],
            SOURCED_VALUES["bw_p"],
            SOURCED_VALUES["pbsa"],
            SOURCED_VALUES["pff"],
        ],
    )


def reference_phenotype() -> FluxPhenotype:
    """The control BW-P BF's map: named by the paper, not released (a typed gap)."""
    return FluxPhenotype(
        net_flux=None,
        confidence_level=CONFIDENCE_LEVEL,
        measurement_type=MEASUREMENT_TYPE,
        provenance_gaps=[CONTROL_MAP_GAP, *RECORD_GAPS],
    )


def flux_phenotype(block: BlockMap) -> FluxPhenotype:
    """One released block, verbatim, at the workbook's stated 90% level."""
    return FluxPhenotype(
        net_flux=dict(block.net_flux),
        net_flux_lower=dict(block.net_flux_lower),
        net_flux_upper=dict(block.net_flux_upper),
        exchange_flux=dict(block.exchange_flux),
        exchange_flux_lower=dict(block.exchange_flux_lower),
        exchange_flux_upper=dict(block.exchange_flux_upper),
        confidence_level=CONFIDENCE_LEVEL,
        measurement_type=MEASUREMENT_TYPE,
        provenance_gaps=list(RECORD_GAPS),
    )


def genotype(spec: StrainSpec) -> Genotype:
    """The strain's edit relative to BW-P BF: its ``zwf`` promoter replacement."""
    symbol, tag = ZWF
    return Genotype(
        perturbations=[
            PromoterReplacementPerturbation(
                systematic_gene_name=tag,
                perturbed_gene_name=symbol,
                gene_namespace=GENE_NAMESPACE,
                identifier_mapping=DerivedIdentifierMapping(
                    source_identifier=symbol, route="gene_symbol"
                ),
                promoter_name=spec.promoter,
                expression_direction=str(SOURCED_VALUES["promoter_strategy"].value),
                is_inducible=False,
            )
        ]
    )


def build_reference(
    dataset_name: str, genome_reference: ReferenceGenome
) -> FluxExperimentReference:
    """BW-P BF on the labeling medium, its map a typed gap."""
    return FluxExperimentReference(
        dataset_name=dataset_name,
        genome_reference=genome_reference,
        environment_reference=labeling_environment(),
        phenotype_reference=reference_phenotype(),
    )


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/liFineTuningGlycolytic2021``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def library_dir(data_root: str | None = None) -> Path:
    """This key's directory in the OCR library mirror."""
    return Path(data_root or _data_root()) / LIBRARY_DIR_REL


def workbook_retrieval() -> RetrievalRecord:
    """The PMC Article Datasets retrieval the library manifest records for the file."""
    return RetrievalRecord(
        method=RetrievalMethod.pmc_cloud,
        source_url=WORKBOOK_URL,
        retriever="torchcell.literature.retrieve.pmc_cloud_object",
        params={"key": "PMC7852082.1/12934_2021_1526_MOESM2_ESM.xlsx"},
        sha256=WORKBOOK_SHA256,
        retrieved_at=WORKBOOK_RETRIEVED_AT,
    )


def deposit_raw_mirror(data_root: str | None = None) -> Path:
    """Copy the verified workbook from the library mirror into the raw mirror.

    Idempotent by sha256: an existing mirror file with the pinned hash is left alone and
    any other bytes raise. The manifest records the PMC retrieval and the library copy.
    """
    src = library_dir(data_root) / WORKBOOK_LIBRARY_RELPATH
    verify_sha256(src, WORKBOOK_SHA256)
    root = raw_mirror_dir(data_root)
    dest = root / WORKBOOK_RELPATH
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        verify_sha256(dest, WORKBOOK_SHA256)
    else:
        shutil.copy2(src, dest)
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title=PAPER_TITLE,
        files=[
            ArtifactRecord(
                path=WORKBOOK_RELPATH,
                role="raw_data",
                bytes=dest.stat().st_size,
                sha256=WORKBOOK_SHA256,
                source=WORKBOOK_URL,
                original_filename=WORKBOOK_NAME,
                retrieval=workbook_retrieval(),
            )
        ],
        si_data_sources=[WORKBOOK_URL],
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def manifest_sha256(data_root: str | None = None) -> str:
    """The sha256 the raw mirror's manifest records for the workbook."""
    manifest = Manifest.model_validate_json(
        (raw_mirror_dir(data_root) / "manifest.json").read_text()
    )
    return next(f.sha256 for f in manifest.files if f.path == WORKBOOK_RELPATH)


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class FluxLi2021Dataset(ExperimentDataset):
    """Li 2021's three released 13C-MFA maps, against BW-P BF's unreleased map."""

    organism = Organism.ecoli

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = REFERENCE_STRAIN

    def __init__(
        self,
        root: str = DATASET_ROOT_REL,
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize at the dev-tree root."""
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return FluxExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return FluxExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The one consumed file, Additional file 2."""
        return [WORKBOOK_NAME]

    def download(self) -> None:
        """Link the mirrored workbook after checking the manifest pin and sha256."""
        data_root = _data_root()
        check_manifest_pin(
            WORKBOOK_RELPATH, manifest_sha256(data_root), WORKBOOK_SHA256
        )
        src = raw_mirror_dir(data_root) / WORKBOOK_RELPATH
        os.makedirs(self.raw_dir, exist_ok=True)
        link_verified(src, osp.join(self.raw_dir, WORKBOOK_NAME), WORKBOOK_SHA256)

    @post_process
    def process(self) -> None:
        """One record per released NETWORK block, 3 in all."""
        verify_raw_files(self.raw_dir, {WORKBOOK_NAME: WORKBOOK_SHA256})
        blocks = read_workbook(osp.join(self.raw_dir, WORKBOOK_NAME))
        if len(blocks) != EXPECTED_RECORDS:
            raise RuntimeError(f"{len(blocks)} blocks, expected {EXPECTED_RECORDS}")
        by_label = {block.block_label: block for block in blocks}
        os.makedirs(self.processed_dir, exist_ok=True)
        reference = build_reference(
            self.name,
            assembly_reference(REFERENCE_STRAIN, background=reference_background()),
        )
        environment = labeling_environment()
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for index, spec in enumerate(STRAINS):
                experiment = FluxExperiment(
                    dataset_name=self.name,
                    genotype=genotype(spec),
                    environment=environment,
                    phenotype=flux_phenotype(by_label[spec.block_label]),
                )
                txn.put(
                    f"{index}".encode(),
                    self._intern_record(experiment, reference, PUBLICATION, itxn),
                )
        env.close()
        interned_env.close()
        self._write_ledgers(blocks)
        log.info("Li 2021: %d released flux maps written", len(blocks))

    def _write_ledgers(self, blocks: Sequence[BlockMap]) -> None:
        """What is not stored (with counts), the release findings, the sourcing."""
        out = Path(self.preprocess_dir)
        out.mkdir(parents=True, exist_ok=True)
        by_label = {block.block_label: block for block in blocks}
        pairs = {
            f"{a} vs {b}": len(identical_rows(by_label[a], by_label[b]))
            for a, b in (
                ("BW-P08", "BW-P04"),
                ("BW-P08", "BW-P10"),
                ("BW-P04", "BW-P10"),
            )
        }
        (out / "not_stored.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "stored_records": len(blocks),
                    "not_stored": [
                        {
                            "what": "normalized columns (fluxes per 100 units of "
                            "glucose uptake), best fit and both bounds",
                            "count": len(blocks) * EXPECTED_REACTIONS * 3,
                            "why": "the absolute fit divided by the block's own released "
                            "vUPTU; storing it would store each map twice",
                        },
                        {
                            "what": "fit diagnostics (fitted measurements, freedom, "
                            "chi2 90%, SSR) per block",
                            "count": len(blocks) * len(DIAGNOSTIC_LABELS),
                            "values": {b.block_label: b.diagnostics for b in blocks},
                            "why": "properties of the fit with no field on FluxPhenotype",
                        },
                        {
                            "what": "LABEL_INPUT, LABEL_MEASUREMENTS and Biomass "
                            "Composition sheets",
                            "why": "the fit's inputs, not its result",
                        },
                    ],
                },
                indent=2,
            )
        )
        (out / "release_findings.json").write_text(
            json.dumps(
                {
                    "identical_reaction_rows_between_blocks": pairs,
                    "chi2_90_cells": {
                        b.block_label: b.diagnostics["χ2 90%"] for b in blocks
                    },
                    "control_map_released": False,
                    "unmentioned_block": "BW-P04",
                },
                indent=2,
            )
        )
        (out / "sourced_values.json").write_text(
            json.dumps(
                {k: v.model_dump(mode="json") for k, v in SOURCED_VALUES.items()},
                indent=2,
            )
        )

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process()."""
        raise NotImplementedError("FluxLi2021Dataset builds records in process()")


# --------------------------------------------------------------------------- #
# L0-L4 verification
# --------------------------------------------------------------------------- #
Record = Mapping[str, Any]

VERIFIER_PROVENANCE: Final = Provenance(
    source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{WORKBOOK_RELPATH}",
    citation_key=CITATION_KEY,
    sha256=WORKBOOK_SHA256,
    method="one FluxExperiment per released NETWORK block of Additional file 2: the "
    "absolute best fit, LB90 and UB90 of 99 reactions and the exchange table, verbatim; "
    "the control BW-P BF's map is a typed gap on the reference",
    page="NETWORK",
)


def _l1_one_record_per_block(records: Sequence[Record]) -> LevelResult:
    promoters = [
        r["experiment"]["genotype"]["perturbations"][0]["promoter_name"]
        for r in records
    ]
    passed = sorted(promoters) == sorted(spec.promoter for spec in STRAINS)
    return LevelResult(
        level=Level.L1,
        name="one_record_per_released_block",
        passed=passed,
        message=f"{len(records)} records carry promoters {sorted(promoters)}",
        details={"promoters": promoters},
    )


def _l2_values_equal_workbook(
    records: Sequence[Record], blocks: Sequence[BlockMap]
) -> LevelResult:
    """L2: every stored number equals the re-read workbook cell, bit for bit."""
    by_promoter = {spec.promoter: spec.block_label for spec in STRAINS}
    by_label = {block.block_label: block for block in blocks}
    mismatched: list[str] = []
    n_values = 0
    for record in records:
        phenotype = record["experiment"]["phenotype"]
        promoter = record["experiment"]["genotype"]["perturbations"][0]["promoter_name"]
        block = by_label[by_promoter[promoter]]
        for field in (
            "net_flux",
            "net_flux_lower",
            "net_flux_upper",
            "exchange_flux",
            "exchange_flux_lower",
            "exchange_flux_upper",
        ):
            stored = phenotype[field]
            released = getattr(block, field)
            n_values += len(released)
            if stored != released:
                mismatched.append(f"{block.block_label}.{field}")
    return LevelResult(
        level=Level.L2,
        name="value_fidelity_to_workbook",
        passed=not mismatched,
        message=f"{n_values} stored values equal the workbook; mismatched {mismatched}",
        details={"n_values": n_values, "mismatched": mismatched},
    )


def _l3_interval_and_reference(records: Sequence[Record]) -> LevelResult:
    """L3: one measurement type, a complete 90% interval, and a gapped reference map."""
    types = {r["experiment"]["phenotype"]["measurement_type"] for r in records}
    incomplete = [
        i
        for i, r in enumerate(records)
        if r["experiment"]["phenotype"]["net_flux_lower"] is None
        or r["experiment"]["phenotype"]["net_flux_upper"] is None
        or r["experiment"]["phenotype"]["confidence_level"] != CONFIDENCE_LEVEL
    ]
    reference_gapped = [
        r["reference"]["phenotype_reference"]["net_flux"] is None
        and any(
            gap["field"] == "net_flux"
            for gap in r["reference"]["phenotype_reference"]["provenance_gaps"]
        )
        for r in records
    ]
    passed = types == {MEASUREMENT_TYPE} and not incomplete and all(reference_gapped)
    return LevelResult(
        level=Level.L3,
        name="flux_interval_and_gapped_reference",
        passed=passed,
        message=f"types {sorted(types)}; {len(incomplete)} incomplete intervals; "
        f"{sum(reference_gapped)} of {len(records)} references carry the control-map gap",
        details={"incomplete": incomplete},
    )


def _l4_zwf_is_a_bw25113_gene(
    records: Sequence[Record], universe: set[str]
) -> LevelResult:
    """L4 (a finding, never a gate): the keyed locus is a gene row of BW25113."""
    loci = {
        p["systematic_gene_name"]
        for r in records
        for p in r["experiment"]["genotype"]["perturbations"]
    }
    missing = sorted(loci - universe)
    return LevelResult(
        level=Level.L4,
        name="perturbed_locus_in_bw25113_annotation",
        passed=True,
        message=f"{len(loci) - len(missing)} of {len(loci)} keyed loci are BW25113 gene "
        f"rows; missing {missing} (reported as a finding: the genome never gates)",
        details={"loci": sorted(loci), "missing": missing},
    )


def verify_build(dataset_root: str, data_root: str | None = None) -> VerificationReport:
    """L0-L4 over the built store; writes ``preprocess/verification_report.json``."""
    from pydantic import TypeAdapter

    from torchcell.datamodels.schema import ExperimentType
    from torchcell.verification.levels import l0_structural, l1_count
    from torchcell.verification.runners import (
        BACTERIAL_GENE_ASSEMBLIES,
        _bacterial_gene_set,
        load_records,
    )

    base = data_root or _data_root()
    records = load_records(dataset_root)
    validate: Callable[[Any], object] = TypeAdapter(ExperimentType).validate_python
    report = VerificationReport(
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=VERIFIER_PROVENANCE,
    )
    report.add(l0_structural((r["experiment"] for r in records), validate))
    report.add(l1_count(len(records), EXPECTED_RECORDS))
    report.add(_l1_one_record_per_block(records))
    blocks = read_workbook(raw_mirror_dir(base) / WORKBOOK_RELPATH)
    report.add(_l2_values_equal_workbook(records, blocks))
    report.add(_l3_interval_and_reference(records))
    for value in SOURCED_VALUES.values():
        report.add(audit_sourced_value(value, Path(base) / "torchcell-library"))
    universe = _bacterial_gene_set(BACTERIAL_GENE_ASSEMBLIES[ASSEMBLY_SET], base)
    report.add(_l4_zwf_is_a_bw25113_gene(records, universe))
    out = osp.join(dataset_root, "preprocess", "verification_report.json")
    os.makedirs(osp.dirname(out), exist_ok=True)
    with open(out, "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main(argv: list[str] | None = None) -> int:
    """CLI: ``deposit`` the raw mirror, ``build`` the dev store, or ``verify`` it."""
    from dotenv import load_dotenv

    parser = argparse.ArgumentParser(prog="python -m torchcell.datasets.ecoli.li2021")
    parser.add_argument("command", choices=("deposit", "build", "verify"))
    args = parser.parse_args(argv)
    load_dotenv()
    data_root = _data_root()
    root = osp.join(data_root, DATASET_ROOT_REL)
    if args.command == "deposit":
        print(deposit_raw_mirror(data_root))
        return 0
    if args.command == "build":
        print(f"len = {len(FluxLi2021Dataset(root=root))}")
        return 0
    report = verify_build(root, data_root)
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
