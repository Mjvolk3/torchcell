# tests/torchcell/datasets/ecoli/test_fuhrer2017.py
# [[tests.torchcell.datasets.ecoli.test_fuhrer2017]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_fuhrer2017.py
"""The Fuhrer 2017 Keio metabolome loader (``torchcell.datasets.ecoli.fuhrer2017``).

Synthetic tests (run everywhere) build every input in ``tmp_path``. The hermetic build
uses the real ``EcoliK12BW25113Genome`` over the synthetic BW25113 assembly of
``tests/torchcell/sequence/genome/_bacterial_fixtures.py`` (JW4367 on BW25113_0001 thrL,
JW0001 on BW25113_0002 thrA, hokC on BW25113_4412), served through a stubbed ``resolve``
with the network refused, and replaces the loader's ``verify_raw_files`` with a
presence check (the synthetic files cannot carry the real pins; the pins are asserted by
the refusal test and by the data-gated tests). The deposit of the synthetic frame, in
the deposit's alphabetical column order:

    column  name   Table EV1A entries                      raw columns
    1       hokC   JW4367 b0001 (ready)                    4   -> BW25113_0001 by JW id;
                                                               the NAME resolves to 4412
    2       ilvG   JW3740 b3767, JW3741 b3768 (ready)      8   -> pooled, dropped
    3       thrA   JW0001 b0002 (ready)                    4   -> BW25113_0002
    4       thrL   JW8888 b0099 (ready)                    4   -> JW id retired, dropped;
                                                               the NAME is hokC's locus
    5       wt     (none)                                  8   -> reference, n = 4
    6       yaaQ   JW9999 b9999 (ready)                    4   -> JW id retired, dropped

Each z-score row is a permutation of the standardized 1..6 scores (median 0, SD 1 over
the six columns), so the build's standardization check passes. Two of the four JW ids
resolve, below the 0.95 default, so the success path lowers ``MIN_RESOLVED_FRACTION``
to 0.5 and a separate test shows the default stops the build.

Data-gated tests (``@pytest.mark.data``) read the real raw mirror and the built dev-tree
LMDB under ``$DATA_ROOT`` (they never build it): the manifest pins, the provenance audit
of every sourced value, the record count, one hand-checked record (``thrA``, values read
off ``zscore_neg.tsv`` / ``zscore_pos.tsv`` column 1,915 with ``cut``), and the L0-L4
verifier.
"""

from __future__ import annotations

import json
import os
import os.path as osp
from collections.abc import Mapping
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pytest

import torchcell.datasets.ecoli.fuhrer2017 as m
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    BW25113_LOCI,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.data import ManifestPinMismatchError, RawSha256MismatchError
from torchcell.datamodels.media import M9_GLUCOSE_CASEIN_FUHRER2017
from torchcell.datamodels.schema import (
    AssemblyReferenceGenome,
    Environment,
    Temperature,
)
from torchcell.datasets.bacteria_common import LocusTagResolutionError
from torchcell.literature.manifest import RetrievalMethod, RetrievalRecord
from torchcell.sequence.genome.ecoli.k12 import BW25113_ASSEMBLY, EcoliK12BW25113Genome
from torchcell.verification.sourced import audit_sourced_value

Z6 = [-1.3363, -0.8018, -0.2673, 0.2673, 0.8018, 1.3363]
NEG_ROWS = [
    [Z6[0], Z6[1], Z6[2], Z6[3], Z6[4], Z6[5]],
    [Z6[5], Z6[4], Z6[3], Z6[2], Z6[1], Z6[0]],
    [Z6[2], Z6[0], Z6[4], Z6[1], Z6[5], Z6[3]],
]
POS_ROWS = [
    [Z6[1], Z6[2], Z6[0], Z6[5], Z6[3], Z6[4]],
    [Z6[3], Z6[5], Z6[1], Z6[0], Z6[2], Z6[4]],
]
COLUMNS = ["hokC", "ilvG", "thrA", "thrL", "wt", "yaaQ"]
RAW_COUNTS = {"hokC": 4, "ilvG": 8, "thrA": 4, "thrL": 4, "wt": 8, "yaaQ": 4}
EV1A_ROWS: list[tuple[str, str, str, str]] = [
    ("JW4367", "b0001", "hokC", "ready to distribute"),
    ("JW3740", "b3767", "ilvG", "ready to distribute"),
    ("JW3741", "b3768", "ilvG", "ready to distribute"),
    ("JW0001", "b0002", "thrA", "ready to distribute"),
    ("JW8888", "b0099", "thrL", "ready to distribute"),
    ("JW9999", "b9999", "yaaQ", "ready to distribute"),
    ("JW5765", " ", " ", "not current_JW ORF"),
]
NEG_MZ = [100.0, 200.0, 300.0]
POS_MZ = [150.0, 250.0]
EV1B_NEG_MZ = [100.001, 200.001, 300.001]
EV1B_POS_MZ = [150.0005, 250.0005]
REFERENCE = AssemblyReferenceGenome(
    species="Escherichia coli",
    strain="BW25113",
    assembly_set="ecoli_K12_BW25113_ASM75055v1",
    assembly_accession="GCA_000750555.1",
)


def _entries() -> list[m.KeioEntry]:
    return [
        m.KeioEntry(jw_id=jw, blattner_id=b, gene_name=g.strip(), delivery_status=s)
        for jw, b, g, s in EV1A_ROWS
    ]


def _write_column(path: Path, values: list[Any]) -> None:
    """A one-column, headerless sheet (xlsx bytes; pandas sniffs the container)."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for value in values:
        sheet.append([value])
    workbook.save(path)


def _write_table_ev1(path: Path) -> None:
    workbook = openpyxl.Workbook()
    ev1a = workbook.active
    assert ev1a is not None
    ev1a.title = m.TABLE_EV1A_SHEET
    ev1a.append(["title"])
    ev1a.append([None, None, None, None, None, None, "Growth-ratesc"])
    ev1a.append(["blank"])
    ev1a.append(
        [
            "Sample Number",
            "JW ID",
            "Blattner ID",
            "Gene Name",
            "Annotationa",
            "Keio Delivery Statusb",
            "Mean",
            "Standard Deviation",
        ]
    )
    for i, (jw, b, g, s) in enumerate(EV1A_ROWS):
        ev1a.append([i + 1, jw, b, g, "-", s, 0.8, 0.05])
    ev1a.append(["cMean and standard deviations are calculated from duplicate"])
    ev1b = workbook.create_sheet(m.TABLE_EV1B_SHEET)
    ev1b.append(["Ionization Mode", "Ion Index", "m/z", "Annotations"])
    ev1b.append([None, None, None, None])
    for mode, mzs in (("neg", EV1B_NEG_MZ), ("pos", EV1B_POS_MZ)):
        for i, mz in enumerate(mzs):
            ev1b.append([mode, i + 1, mz, "no annotation"])
    workbook.save(path)


def _write_raw(raw: Path) -> None:
    raw.mkdir(parents=True, exist_ok=True)
    (raw / m.STUDY_JSON).write_text(
        json.dumps(
            {
                "accno": "S-BSST5",
                "section": {
                    "attributes": [
                        {"name": "Organism", "value": "Escherichia coli"},
                        {
                            "name": "Experimental design",
                            "value": "2 Biological and 2 technical replicates",
                        },
                    ]
                },
            }
        )
    )
    for name, rows in ((m.ZSCORE_NEG, NEG_ROWS), (m.ZSCORE_POS, POS_ROWS)):
        (raw / name).write_text(
            "".join("\t".join(f"{v:.4f}" for v in row) + "\n" for row in rows)
        )
    _write_column(raw / m.SAMPLE_ID_ZSCORE, COLUMNS)
    _write_column(
        raw / m.SAMPLE_ID_ALL,
        [name for name in COLUMNS for _ in range(RAW_COUNTS[name])],
    )
    _write_column(raw / m.NEG_ION_MZ, NEG_MZ)
    _write_column(raw / m.POS_ION_MZ, POS_MZ)
    _write_table_ev1(raw / m.TABLE_EV1)


@pytest.fixture
def bw25113(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12BW25113Genome:
    """The real BW25113 class over the synthetic assembly; network refused."""
    files = write_assembly(tmp_path / "tier", BW25113_ASSEMBLY, BW25113_LOCI)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    root = tmp_path / "bw25113"
    root.mkdir()
    return EcoliK12BW25113Genome(genome_root=str(root), overwrite=False)


@pytest.fixture
def presence_only_pins(monkeypatch: pytest.MonkeyPatch) -> list[Mapping[str, str]]:
    """Replace the build-time byte check with a presence check that records the pins."""
    calls: list[Mapping[str, str]] = []

    def record(raw_dir: str, pins: Mapping[str, str]) -> None:
        missing = [f for f in pins if not osp.exists(osp.join(raw_dir, f))]
        if missing:
            raise FileNotFoundError(f"pinned raw files absent: {missing}")
        calls.append(dict(pins))

    monkeypatch.setattr(m, "verify_raw_files", record)
    return calls


@pytest.fixture
def synthetic(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    bw25113: EcoliK12BW25113Genome,
    presence_only_pins: list[Mapping[str, str]],
) -> Path:
    """A dataset root whose raw/ holds the synthetic deposit, with small ion counts."""
    root = tmp_path / "metabolome_fuhrer2017"
    _write_raw(root / "raw")
    monkeypatch.setattr(m, "EXPECTED_ION_COUNTS", {"neg": 3, "pos": 2})
    monkeypatch.setattr(m, "assembly_reference", lambda strain, **_: REFERENCE)
    return root


def _expected_environment() -> dict[str, Any]:
    return Environment(
        media=M9_GLUCOSE_CASEIN_FUHRER2017, temperature=Temperature(value=37.0)
    ).model_dump()


# --------------------------------------------------------------------------- #
# Pure parsing and selection
# --------------------------------------------------------------------------- #
def test_ion_keys_are_mode_and_one_based_row() -> None:
    assert m.ion_keys(3, 2) == [
        "neg_0001",
        "neg_0002",
        "neg_0003",
        "pos_0001",
        "pos_0002",
    ]
    assert m.ion_key("pos", 4365) == "pos_4365"


def test_select_strains_keeps_single_entries_and_drops_pools() -> None:
    """One record per single-entry name; the shared name is a pooled drop; wt is the
    reference with 8 raw columns = 4 cultures.
    """
    selection = m.select_strains(COLUMNS, _entries(), RAW_COUNTS)
    assert [(s.column, s.deposit_name, s.entry.jw_id) for s in selection.strains] == [
        (0, "hokC", "JW4367"),
        (2, "thrA", "JW0001"),
        (3, "thrL", "JW8888"),
        (5, "yaaQ", "JW9999"),
    ]
    assert [s.n_biological for s in selection.strains] == [2, 2, 2, 2]
    assert (selection.wt_column, selection.wt_raw_columns) == (4, 8)
    assert selection.wt_n_biological == 4
    assert selection.pooled.rule == "pooled_keio_entries"
    assert selection.pooled.items == [
        "ilvG: JW3740 b3767 (ready to distribute), JW3741 b3768 (ready to "
        "distribute); 2 loci"
    ]


def test_select_strains_names_a_same_locus_pool() -> None:
    entries = [
        m.KeioEntry(
            jw_id="JW5182", blattner_id="b1185", gene_name="dsbB", delivery_status=s
        )
        for s in ("ready to distribute", "Eliminated; wrong primers")
    ]
    selection = m.select_strains(["dsbB", "wt"], entries, {"dsbB": 8, "wt": 4})
    assert selection.strains == []
    assert selection.pooled.items == [
        "dsbB: JW5182 b1185 (ready to distribute), JW5182 b1185 (Eliminated; wrong "
        "primers); same locus"
    ]


@pytest.mark.parametrize(
    ("columns", "counts", "error", "message"),
    [
        (
            ["nope", "wt"],
            {"nope": 4, "wt": 4},
            m.UnmappedDepositColumnError,
            "deposit column 1 'nope' matches no Table EV1A entry",
        ),
        (
            ["thrA", "wt"],
            {"thrA": 6, "wt": 4},
            m.RawColumnCountError,
            "'thrA': 6 raw columns for 1 Table EV1A entries (expected 4 each)",
        ),
        (["thrA"], {"thrA": 4}, m.UnmappedDepositColumnError, "no 'wt' column"),
        (["thrA", "wt"], {"thrA": 4, "wt": 3}, m.RawColumnCountError, "'wt' has 3"),
    ],
)
def test_select_strains_refuses_unexpected_deposits(
    columns: list[str], counts: dict[str, int], error: type[Exception], message: str
) -> None:
    with pytest.raises(error, match=message.replace("(", r"\(").replace(")", r"\)")):
        m.select_strains(columns, _entries(), counts)


def test_select_strains_refuses_an_undistributed_single_entry() -> None:
    entries = [
        m.KeioEntry(
            jw_id="JW0001",
            blattner_id="b0002",
            gene_name="thrA",
            delivery_status="Eliminated; weak growth",
        )
    ]
    with pytest.raises(m.UndistributedKeioEntryError) as err:
        m.select_strains(["thrA", "wt"], entries, {"thrA": 4, "wt": 4})
    assert str(err.value) == (
        "'thrA' (JW0001) has Keio status 'Eliminated; weak growth'"
    )


def test_ion_alignment_counts_rows_nearest_their_own_index() -> None:
    result = m.ion_alignment("neg", NEG_MZ, EV1B_NEG_MZ)
    assert result.nearest_row_agrees == 3
    assert result.agreement == 1.0
    assert result.max_abs_mz_difference == pytest.approx(0.001)
    with pytest.raises(ValueError, match="neg: 3 deposit ions but 2 Table EV1B ions"):
        m.ion_alignment("neg", NEG_MZ, [100.0, 200.0])
    with pytest.raises(ValueError, match="only 1 of 3 Table EV1B ions"):
        m.ion_alignment("neg", NEG_MZ, [300.0, 200.0, 100.0])


def test_zscore_scale_accepts_standardized_rows_and_refuses_others() -> None:
    scale = m.zscore_scale("neg", np.asarray(NEG_ROWS))
    assert scale.max_abs_column_median == 0.0
    assert scale.min_column_sd == pytest.approx(1.0, abs=1e-4)
    with pytest.raises(ValueError, match="neg: released z-scores are not standardized"):
        m.zscore_scale("neg", np.asarray(NEG_ROWS) * 2.0)


def test_read_zscores_refuses_a_wrong_shape_and_a_nan(tmp_path: Path) -> None:
    path = tmp_path / "z.tsv"
    path.write_text("0.1\t0.2\n0.3\tnan\n")
    with pytest.raises(ValueError, match=r"shape \(2, 2\), expected \(2, 3\)"):
        m.read_zscores(path, 2, 3)
    with pytest.raises(m.NonFiniteZScoreError) as err:
        m.read_zscores(path, 2, 2)
    assert "1 non-finite cells, first at row 2, column 2" in str(err.value)


def test_read_study_design(tmp_path: Path) -> None:
    _write_raw(tmp_path)
    assert m.read_study_design(tmp_path / m.STUDY_JSON) == (
        "2 Biological and 2 technical replicates"
    )


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
def test_metabolite_phenotype_carries_the_clone_count_and_the_se_gap() -> None:
    phenotype = m.metabolite_phenotype(["neg_0001", "pos_0001"], [0.5, -1.25], 2)
    assert phenotype.metabolite_level == {"neg_0001": 0.5, "pos_0001": -1.25}
    assert phenotype.n_replicates == {"neg_0001": 2, "pos_0001": 2}
    assert phenotype.metabolite_level_se is None
    assert phenotype.measurement_type == "fia_tof_ms_ion_modified_z_score"
    assert phenotype.provenance_gaps == [m.METABOLITE_LEVEL_SE_GAP]
    assert phenotype.gapped_fields() == {"metabolite_level_se"}


def test_build_experiment_is_one_keio_deletion_on_the_bw25113_namespace() -> None:
    entry = m.KeioEntry(
        jw_id="JW0001",
        blattner_id="b0002",
        gene_name="thrA",
        delivery_status="ready to distribute",
    )
    resolved = m.ResolvedStrain(
        strain=m.KeioStrain(column=2, deposit_name="thrA", entry=entry, raw_columns=4),
        locus_tag="BW25113_0002",
        symbol="thrA",
    )
    experiment = m.build_experiment(
        "MetabolomeFuhrer2017Dataset", resolved, ["neg_0001"], [1.5], m.environment()
    )
    dumped = experiment.model_dump()
    assert dumped["experiment_type"] == "bacterial_metabolite"
    assert dumped["environment"] == _expected_environment()
    (perturbation,) = dumped["genotype"]["perturbations"]
    assert {
        k: perturbation[k]
        for k in (
            "systematic_gene_name",
            "perturbed_gene_name",
            "perturbation_type",
            "gene_namespace",
            "collection",
            "cassette",
            "construction",
            "state",
        )
    } == {
        "systematic_gene_name": "BW25113_0002",
        "perturbed_gene_name": "thrA",
        "perturbation_type": "bacterial_deletion",
        "gene_namespace": "ecoli_k12_bw25113_locus_tag",
        "collection": "KEIO knockout collection",
        "cassette": "kanamycin cassette flanked by FLP recognition target sites",
        "construction": {
            "strain_accession": "JW0001",
            "lab": None,
            "batch": None,
            "plate": None,
            "well": None,
        },
        "state": "absent",
    }
    assert dumped["phenotype"]["metabolite_level"] == {"neg_0001": 1.5}
    assert dumped["phenotype"]["n_replicates"] == {"neg_0001": 2}


def test_build_reference_keeps_the_assembly_pin_through_a_dump() -> None:
    reference = m.build_reference(
        "MetabolomeFuhrer2017Dataset",
        REFERENCE,
        ["neg_0001"],
        [0.25],
        96,
        m.environment(),
    )
    dumped = reference.model_dump()
    assert dumped["genome_reference"]["assembly_set"] == "ecoli_K12_BW25113_ASM75055v1"
    assert dumped["genome_reference"]["assembly_accession"] == "GCA_000750555.1"
    assert dumped["phenotype_reference"]["metabolite_level"] == {"neg_0001": 0.25}
    assert dumped["phenotype_reference"]["n_replicates"] == {"neg_0001": 96}


# --------------------------------------------------------------------------- #
# Hermetic build
# --------------------------------------------------------------------------- #
def test_build_two_records_with_the_wt_reference_and_ledgers(
    synthetic: Path,
    monkeypatch: pytest.MonkeyPatch,
    bw25113: EcoliK12BW25113Genome,
    presence_only_pins: list[Mapping[str, str]],
) -> None:
    """Two records, hokC (JW4367) and thrA (JW0001); ilvG is pooled and the JW ids of
    thrL and yaaQ are retired, all ledgered. hokC's record is BW25113_0001 (thrL), the JW
    id's locus, while its deposit NAME resolves to BW25113_4412; that disagreement is
    ledgered, and the dropped thrL's name is marked as resolving to that kept locus.
    """
    monkeypatch.setattr(m, "MIN_RESOLVED_FRACTION", 0.5)
    dataset = m.MetabolomeFuhrer2017Dataset(root=str(synthetic), ecoli_genome=bw25113)
    assert presence_only_pins == [m.DATA_SHA256]
    assert len(dataset) == 2
    keys = ["neg_0001", "neg_0002", "neg_0003", "pos_0001", "pos_0002"]

    first = dataset[0]["experiment"]
    assert (
        first["genotype"]["perturbations"][0]["systematic_gene_name"] == "BW25113_0001"
    )
    assert first["genotype"]["perturbations"][0]["perturbed_gene_name"] == "thrL"
    assert first["genotype"]["perturbations"][0]["construction"][
        "strain_accession"
    ] == ("JW4367")
    assert first["phenotype"]["metabolite_level"] == dict(
        zip(keys, [Z6[0], Z6[5], Z6[2], Z6[1], Z6[3]], strict=True)
    )
    assert first["phenotype"]["n_replicates"] == dict.fromkeys(keys, 2)
    assert first["environment"] == _expected_environment()

    second = dataset[1]["experiment"]
    assert (
        second["genotype"]["perturbations"][0]["systematic_gene_name"] == "BW25113_0002"
    )
    assert second["phenotype"]["metabolite_level"] == dict(
        zip(keys, [Z6[2], Z6[3], Z6[4], Z6[0], Z6[1]], strict=True)
    )

    reference = dataset[0]["reference"]
    assert reference["genome_reference"] == REFERENCE.model_dump()
    assert reference["phenotype_reference"]["metabolite_level"] == dict(
        zip(keys, [Z6[4], Z6[1], Z6[5], Z6[3], Z6[2]], strict=True)
    )
    assert reference["phenotype_reference"]["n_replicates"] == dict.fromkeys(keys, 4)
    assert dataset[0]["publication"] == m.PUBLICATION.model_dump()

    preprocess = synthetic / "preprocess"
    drops = json.loads((preprocess / "dropped_records.json").read_text())
    assert (
        drops["source_records"],
        drops["kept_records"],
        drops["dropped_records"],
    ) == (5, 2, 3)
    assert [(r["rule"], r["items"]) for r in drops["rules"]] == [
        (
            "pooled_keio_entries",
            [
                "ilvG: JW3740 b3767 (ready to distribute), JW3741 b3768 (ready to "
                "distribute); 2 loci"
            ],
        ),
        (
            "jw_id_not_on_a_bw25113_locus",
            [
                "JW8888 (thrL, b0099): retired; the name resolves renamed "
                "BW25113_0001, a kept record's locus",
                "JW9999 (yaaQ, b9999): retired; the name resolves retired yaaQ",
            ],
        ),
    ]
    ledger = json.loads((preprocess / "identifier_reconciliation.json").read_text())
    assert ledger["symbol_disagreements"] == [
        {
            "deposit_name": "hokC",
            "jw_id": "JW4367",
            "jw_locus": "BW25113_0001",
            "name_resolution": "renamed BW25113_4412",
        }
    ]
    assert ledger["reconciliation"]["status_histogram"]["renamed"] == 2
    assert ledger["reconciliation"]["status_histogram"]["retired"] == 2
    ions = (preprocess / "ions.csv").read_text().splitlines()
    assert ions[0] == "key,mode,ion_index,deposit_mz,table_ev1b_mz"
    assert ions[1] == "neg_0001,neg,1,100.0,100.001"
    strains = (preprocess / "strains.csv").read_text().splitlines()
    assert strains[1:] == [
        "0,1,hokC,JW4367,b0001,BW25113_0001,thrL,2",
        "1,3,thrA,JW0001,b0002,BW25113_0002,thrA,2",
    ]
    scale = json.loads((preprocess / "zscore_scale.json").read_text())
    assert [s["mode"] for s in scale] == ["neg", "pos"]


def test_build_stops_below_the_resolution_threshold(
    synthetic: Path, bw25113: EcoliK12BW25113Genome
) -> None:
    """Two of four JW ids resolve (0.500 < 0.95): the build stops, nothing dropped."""
    with pytest.raises(LocusTagResolutionError, match=r"2 of 4 names \(0\.500\)"):
        m.MetabolomeFuhrer2017Dataset(root=str(synthetic), ecoli_genome=bw25113)
    assert not (synthetic / "processed" / "lmdb").exists()


def test_build_refuses_a_genome_of_another_strain(
    synthetic: Path, monkeypatch: pytest.MonkeyPatch, bw25113: EcoliK12BW25113Genome
) -> None:
    monkeypatch.setattr(
        m, "BACTERIAL_ASSEMBLY_SETS", {"BW25113": "ecoli_K12_MG1655_ASM584v2"}
    )
    with pytest.raises(ValueError, match="needs the ecoli_K12_MG1655_ASM584v2 genome"):
        m.MetabolomeFuhrer2017Dataset(root=str(synthetic), ecoli_genome=bw25113)


def test_build_refuses_a_design_other_than_two_by_two(
    synthetic: Path, bw25113: EcoliK12BW25113Genome
) -> None:
    study = synthetic / "raw" / m.STUDY_JSON
    study.write_text(study.read_text().replace("2 Biological", "3 Biological"))
    with pytest.raises(ValueError, match="S-BSST5 design '3 Biological"):
        m.MetabolomeFuhrer2017Dataset(root=str(synthetic), ecoli_genome=bw25113)


def test_process_verifies_the_real_pins(
    synthetic: Path, monkeypatch: pytest.MonkeyPatch, bw25113: EcoliK12BW25113Genome
) -> None:
    """With the real check back, the synthetic study JSON is refused by its pin."""
    from torchcell.data import verify_raw_files

    monkeypatch.setattr(m, "verify_raw_files", verify_raw_files)
    with pytest.raises(
        RawSha256MismatchError, match=f"expected {m.DATA_SHA256[m.STUDY_JSON]}"
    ):
        m.MetabolomeFuhrer2017Dataset(root=str(synthetic), ecoli_genome=bw25113)


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _synthetic_raw_files(
    tmp_path: Path,
) -> tuple[tuple[m.RawFile, ...], dict[str, Path]]:
    import hashlib

    sources: dict[str, Path] = {}
    files: list[m.RawFile] = []
    for name, payload in (("a.tsv", b"1\t2\n"), ("b.xls", b"bytes")):
        path = tmp_path / "download" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
        sha = hashlib.sha256(payload).hexdigest()
        url = f"https://example.org/{name}"
        files.append(
            m.RawFile(
                name=name,
                sha256=sha,
                bytes=len(payload),
                description=name,
                retrieval=RetrievalRecord(
                    method=RetrievalMethod.direct_url,
                    source_url=url,
                    retriever="torchcell.literature.retrieve.direct_url",
                    params={"url": url},
                    sha256=sha,
                    retrieved_at="2026-10-07",
                ),
            )
        )
        sources[name] = path
    return tuple(files), sources


def test_deposit_is_idempotent_and_refuses_a_differing_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    files, sources = _synthetic_raw_files(tmp_path)
    monkeypatch.setattr(m, "RAW_FILES", files)
    monkeypatch.setattr(m, "DATA_SHA256", {f.name: f.sha256 for f in files})
    data_root = str(tmp_path / "root")
    root = m.deposit_raw_mirror(sources=sources, data_root=data_root)
    assert root == Path(data_root) / m.RAW_DIR_REL
    m.deposit_raw_mirror(sources=sources, data_root=data_root)
    manifest = m.load_manifest(data_root)
    assert [
        (f.path, f.sha256, f.retrieval.method if f.retrieval else None)
        for f in manifest.files
    ] == [
        ("data/a.tsv", files[0].sha256, RetrievalMethod.direct_url),
        ("data/b.xls", files[1].sha256, RetrievalMethod.direct_url),
    ]
    assert manifest.si_expected == list(m.NOT_MIRRORED)
    assert m.manifest_sha256(manifest, "data/b.xls") == files[1].sha256
    with pytest.raises(KeyError, match="data/c.xls is not in the raw-mirror manifest"):
        m.manifest_sha256(manifest, "data/c.xls")

    (root / "data" / "a.tsv").write_bytes(b"edited")
    with pytest.raises(RuntimeError, match="exists with a different sha256; refusing"):
        m.deposit_raw_mirror(sources=sources, data_root=data_root)
    (root / "data" / "a.tsv").write_bytes(b"1\t2\n")
    sources["b.xls"].write_bytes(b"other")
    with pytest.raises(RuntimeError, match="b.xls sha256 mismatch"):
        m.deposit_raw_mirror(sources=sources, data_root=data_root)
    with pytest.raises(KeyError, match="no source given for"):
        m.deposit_raw_mirror(sources={}, data_root=data_root)


def test_download_links_the_mirror_and_checks_the_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    files, sources = _synthetic_raw_files(tmp_path)
    monkeypatch.setattr(m, "RAW_FILES", files)
    monkeypatch.setattr(m, "DATA_SHA256", {f.name: f.sha256 for f in files})
    data_root = tmp_path / "root"
    m.deposit_raw_mirror(sources=sources, data_root=str(data_root))
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    dataset = object.__new__(m.MetabolomeFuhrer2017Dataset)
    monkeypatch.setattr(
        m.MetabolomeFuhrer2017Dataset, "raw_dir", str(tmp_path / "raw"), raising=False
    )
    dataset.download()
    linked = tmp_path / "raw" / "a.tsv"
    assert linked.is_symlink()
    assert os.readlink(linked) == str(data_root / m.RAW_DIR_REL / "data" / "a.tsv")

    manifest_path = data_root / m.RAW_DIR_REL / "manifest.json"
    manifest_path.write_text(
        manifest_path.read_text().replace(files[0].sha256, "0" * 64)
    )
    with pytest.raises(ManifestPinMismatchError, match="data/a.tsv"):
        dataset.download()


def test_retrieve_runs_the_recorded_retriever_and_verifies(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    files, _ = _synthetic_raw_files(tmp_path)
    monkeypatch.setattr(m, "RAW_FILES", files)
    calls: list[RetrievalRecord] = []

    def fake(record: RetrievalRecord) -> bytes:
        calls.append(record)
        return b"1\t2\n" if record.params["url"].endswith("a.tsv") else b"wrong"

    monkeypatch.setattr(m, "run_retriever", fake)
    out = m.retrieve_raw_files(tmp_path / "fetched", ["a.tsv"])
    assert out == {"a.tsv": tmp_path / "fetched" / "a.tsv"}
    assert [c.source_url for c in calls] == ["https://example.org/a.tsv"]
    with pytest.raises(RawSha256MismatchError):
        m.retrieve_raw_files(tmp_path / "fetched", ["b.xls"])
    assert not (tmp_path / "fetched" / "b.xls").exists()


def test_every_consumed_file_is_pinned_once() -> None:
    assert [f.name for f in m.RAW_FILES] == [
        "S-BSST5.json",
        "zscore_neg.tsv",
        "zscore_pos.tsv",
        "sample_id_zscore.xls",
        "sample_id_all.xls",
        "neg_ionMz.xls",
        "pos_ionMz.xls",
        "si2.xlsx",
    ]
    assert all(
        len(f.sha256) == 64 and f.sha256 == f.retrieval.sha256 for f in m.RAW_FILES
    )
    assert m.RAW_FILES_BY_NAME["si2.xlsx"].retrieval.method == RetrievalMethod.pmc_cloud


def test_sourced_value_roots_split_the_two_mirrors(tmp_path: Path) -> None:
    root = str(tmp_path)
    assert m.sourced_value_root(m.SOURCED_VALUES["design"], root) == (
        tmp_path / "torchcell-raw"
    )
    assert m.sourced_value_root(m.SOURCED_VALUES["temperature_c"], root) == (
        tmp_path / "torchcell-library"
    )


# --------------------------------------------------------------------------- #
# Real data (``--data``): the mirror and the built dev-tree LMDB, never rebuilt
# --------------------------------------------------------------------------- #
def _real_data_root() -> str:
    return os.environ["DATA_ROOT"]


def _built_root() -> str:
    return osp.join(_real_data_root(), "data/torchcell/metabolome_fuhrer2017")


@pytest.mark.data
def test_real_mirror_matches_the_module_pins() -> None:
    from torchcell.data.experiment_dataset import file_sha256

    manifest = m.load_manifest(_real_data_root())
    for raw in m.RAW_FILES:
        assert m.manifest_sha256(manifest, raw.mirror_relpath) == raw.sha256
        path = m.raw_mirror_dir(_real_data_root()) / raw.mirror_relpath
        assert path.stat().st_size == raw.bytes
        assert file_sha256(path) == raw.sha256


@pytest.mark.data
@pytest.mark.parametrize("key", sorted(m.SOURCED_VALUES))
def test_real_sourced_values_are_verbatim(key: str) -> None:
    value = m.SOURCED_VALUES[key]
    result = audit_sourced_value(value, m.sourced_value_root(value, _real_data_root()))
    assert result.passed, result.message


@pytest.mark.data
def test_real_build_counts_and_the_thra_record() -> None:
    """3,735 records (3,806 mutant columns - 25 pooled - 46 JW ids off BW25113); thrA
    (JW0001) is BW25113_0002 with values read off the deposit by hand.
    """
    from torchcell.verification.runners import load_records

    drops = json.loads(
        Path(_built_root(), "preprocess", "dropped_records.json").read_text()
    )
    assert (drops["source_records"], drops["kept_records"]) == (3806, 3735)
    assert [r["n_records"] for r in drops["rules"]] == [25, 46]
    records = load_records(_built_root())
    assert len(records) == 3735
    (thra,) = [
        r
        for r in records
        if r["experiment"]["genotype"]["perturbations"][0]["construction"][
            "strain_accession"
        ]
        == "JW0001"
    ]
    perturbation = thra["experiment"]["genotype"]["perturbations"][0]
    assert perturbation["systematic_gene_name"] == "BW25113_0002"
    assert perturbation["perturbed_gene_name"] == "thrA"
    levels = thra["experiment"]["phenotype"]["metabolite_level"]
    assert len(levels) == 3169 + 4365
    assert {k: levels[k] for k in HAND_CHECKED_THRA} == HAND_CHECKED_THRA
    assert set(thra["experiment"]["phenotype"]["n_replicates"].values()) == {2}
    reference = thra["reference"]["phenotype_reference"]
    assert {k: reference["metabolite_level"][k] for k in HAND_CHECKED_WT} == (
        HAND_CHECKED_WT
    )
    assert set(reference["n_replicates"].values()) == {96}


@pytest.mark.data
def test_real_build_passes_l0_to_l4() -> None:
    report = m.run_verification(_real_data_root())
    assert report.passed, report.summary()


#: Read off the deposit by hand (thrA is entry 1,915 and wt entry 2,060 of
#: ``sample_id_zscore.xls``): ``sed -n <row>p zscore_neg.tsv | cut -f 1915`` for rows 1
#: and 3169, the same on ``zscore_pos.tsv`` for rows 1 and 4365, and ``cut -f 2060``.
HAND_CHECKED_THRA: dict[str, float] = {
    "neg_0001": 1.3524,
    "neg_3169": 0.4154,
    "pos_0001": 0.6878,
    "pos_4365": 1.3324,
}
HAND_CHECKED_WT: dict[str, float] = {"neg_0001": -0.1128, "pos_4365": -0.0166}
