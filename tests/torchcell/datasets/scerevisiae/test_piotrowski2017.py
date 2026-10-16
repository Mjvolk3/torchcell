# tests/torchcell/datasets/scerevisiae/test_piotrowski2017.py
# [[tests.torchcell.datasets.scerevisiae.test_piotrowski2017]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/scerevisiae/test_piotrowski2017.py
"""Piotrowski 2017 loader: release readers, dose parsing, strain matching, the build.

Everything but the two mirror tests runs on synthetic workbooks written with openpyxl in
the three released layouts (Data Set 5 matrix, Data Set 4 conditions, Data Set 2 pool),
so no build tree and no network are needed. The loader consults no genome. The mirror tests bind
every module-level ``SourcedValue`` to its verbatim quote and the release pins to the
literature mirror's manifest, and skip when the mirror is not mounted.

The fixture (``_write_release``): seven strain rows x nine condition columns, value
``v(i, j) = -(10 i + j) / 4`` for row ``i`` and column ``j`` except four ``NA`` cells
(TFC3 in NPD0002's column, VPS8 in the three non-RIKEN columns).

- Strains: TFC3 (YAL001C), VPS8 (YAL002W), OLD1 (YAL003W, an R64 alias of YAL002W),
  GONE1 (YPL999C, not an R64 gene), SNQ2 (YDR011W, a locus the SGA selections fix),
  YAL058C-A (an R64 alias of YAL056C-A) and YBR001C (labelled by its ORF, whose pool
  common name is the ORF). Every one is stored under its released ORF and common name.
- Conditions: Benomyl 10-ug/ml (a curated table row), NPD0001 100-uM (an InChI and no
  table row), NPD0002 50-ng/ml (a library code with no structure: held out as an
  unidentified agent), Ouabain 10-100-ug/ml (a range: dropped), NPE0003
  (In_final_dataset False: dropped), MMS 4e-05 under a percent format (POSCON, so the
  NCI/NIH/GSK screen), NSC0001 at 100-uM twice (two NCI libraries, one compound, one
  dose), and Kundrymycin 50-ng/ml (a released name with no structure: stored, gapped).

Counts: 63 cells, 4 NA, so 59 source records; the two dropped columns hold 7 each and
NPD0002 holds 6, so 59 - 20 = 39 kept.
"""

from __future__ import annotations

import os
import os.path as osp
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import numpy as np
import openpyxl
import pytest

from torchcell.datamodels.compound_identity import resolve_compound_identity
from torchcell.datamodels.schema import ConcentrationUnit, DoseBasis
from torchcell.datasets.scerevisiae import piotrowski2017 as p
from torchcell.datasets.scerevisiae.hoepfner2014 import ReleasedIdentity
from torchcell.datasets.scerevisiae.vanacloig2022 import LIBRARY_STRAIN
from torchcell.verification.report import Provenance, StatDerivation, VerificationReport
from torchcell.verification.sourced import SourcedValue

# --------------------------------------------------------------------------- #
# Synthetic release
# --------------------------------------------------------------------------- #
_ETHANOL_INCHI = "InChI=1S/C2H6O/c1-2-3/h3H,2H2,1H3"
_ETHANOL_KEY = "LFQSCWFLJHTTHZ-UHFFFAOYSA-N"
_BENOMYL_SMILES = "CCCCNC(=O)n1c(NC(=O)OC)nc2ccccc21"

#: (label, ORF, barcode) of the Data Set 5 rows, in row order.
_STRAINS = [
    ("TFC3", "YAL001C", "AAAAAAAAAAAAAAAAAAAA"),
    ("VPS8", "YAL002W", "CCCCCCCCCCCCCCCCCCCC"),
    ("OLD1", "YAL003W", "GGGGGGGGGGGGGGGGGGGG"),
    ("GONE1", "YPL999C", "TTTTTTTTTTTTTTTTTTTT"),
    ("SNQ2", "YDR011W", "ACACACACACACACACACAC"),
    ("YAL058C-A", "YAL058C-A", "AGAGAGAGAGAGAGAGAGAG"),
    ("YBR001C", "YBR001C", "ATATATATATATATATATAT"),
]

#: (id, name, dose cell, number format, collection, final, HCS, SMILES, InChI)
_CONDITIONS: list[tuple[Any, ...]] = [
    (
        "Cond-000001",
        "Benomyl",
        "10-ug/ml",
        "General",
        "RIKEN NPDepo",
        True,
        True,
        _BENOMYL_SMILES,
        "NA",
    ),
    (
        "Cond-000002",
        "NPD0001",
        "100-uM",
        "General",
        "RIKEN NPDepo",
        True,
        False,
        "CCO",
        _ETHANOL_INCHI,
    ),
    (
        "Cond-000003",
        "NPD0002",
        "50-ng/ml",
        "General",
        "RIKEN NPDepo",
        True,
        False,
        None,
        None,
    ),
    (
        "Cond-000004",
        "Ouabain",
        "10-100-ug/ml",
        "General",
        "RIKEN NPDepo",
        True,
        False,
        "CCO",
        _ETHANOL_INCHI,
    ),
    (
        "Cond-000005",
        "NPE0003",
        "10-ug/ml",
        "General",
        "RIKEN NPDepo",
        False,
        False,
        "CCO",
        _ETHANOL_INCHI,
    ),
    (
        "Cond-000006",
        "MMS",
        4e-05,
        "0.000%",
        "POSCON-CLINICAL-2013",
        True,
        True,
        "COS(C)(=O)=O",
        "NA",
    ),
    (
        "Cond-000007",
        "NSC0001",
        "100-uM",
        "General",
        "NCI-STRUCT-DIV",
        True,
        False,
        "CCO",
        _ETHANOL_INCHI,
    ),
    (
        "Cond-000008",
        "NSC0001",
        "100-uM",
        "General",
        "NCI-MECH-DIV",
        True,
        False,
        "CCO",
        _ETHANOL_INCHI,
    ),
    (
        "Cond-000010",
        "Kundrymycin",
        "50-ng/ml",
        "General",
        "RIKEN NPDepo",
        True,
        False,
        "NA",
        "NA",
    ),
]
#: A Data Set 4 row the matrix does not carry (Data Set 4 lists every condition).
_UNSCREENED = (
    "Cond-000009",
    "NPD0009",
    "10-ug/ml",
    "General",
    "RIKEN NPDepo",
    False,
    False,
    "CCO",
    _ETHANOL_INCHI,
)
_NA_CELLS = {(0, 2), (1, 5), (1, 6), (1, 7)}


def _value(i: int, j: int) -> float:
    return -(10 * i + j) / 4


def _write_release(raw: Path) -> None:
    """The three released workbooks, in the publisher's sheet layouts."""
    raw.mkdir(parents=True, exist_ok=True)
    matrix = openpyxl.Workbook()
    sheet = matrix.active
    assert sheet is not None
    sheet.title = "All_conditions"
    sheet.append([p.MATRIX_TITLE.quote])
    sheet.append([None])
    sheet.append([None, *(f"{c[0]}_{c[1]}" for c in _CONDITIONS)])
    for i, (label, _, _) in enumerate(_STRAINS):
        sheet.append(
            [
                label,
                *(
                    "NA" if (i, j) in _NA_CELLS else _value(i, j)
                    for j in range(len(_CONDITIONS))
                ),
            ]
        )
    matrix.save(raw / p.MATRIX.filename)

    conditions = openpyxl.Workbook()
    sheet = conditions.active
    assert sheet is not None
    sheet.title = "Suppl. dataset 4"
    for line in ("title", "note 1", "note 2", None, None):
        sheet.append([line])
    sheet.append(list(p._CONDITION_HEADER))
    for cid, name, dose, fmt, collection, final, hcs, smiles, inchi in [
        *_CONDITIONS,
        _UNSCREENED,
    ]:
        sheet.append([cid, name, dose, collection, final, hcs, 50.0, smiles, inchi])
        sheet.cell(row=sheet.max_row, column=3).number_format = fmt
    conditions.save(raw / p.CONDITIONS.filename)

    pool = openpyxl.Workbook()
    sheet = pool.active
    assert sheet is not None
    sheet.title = "Suppl. dataset 2"
    sheet.append(["Supplementary dataset 2"])
    sheet.append([None])
    sheet.append([None])
    sheet.append([p.POOL_COLLECTION.value])
    sheet.append(list(p._POOL_HEADER))
    for label, orf, barcode in _STRAINS:
        sheet.append([orf, label, barcode])
    for k in range(p.POOL_V2.value - len(_STRAINS)):
        sheet.append([f"YZZ{k:03d}W", f"FIL{k}", "G" * 20])
    sheet.append(["total strains", 310])
    pool.save(raw / p.POOL.filename)


@pytest.fixture()
def built(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """Build the synthetic dataset end to end into a temporary root."""
    monkeypatch.setattr(
        p.EnvChemgenPiotrowski2017Dataset, "download", lambda self: None
    )
    root = tmp_path / "env_chemgen_piotrowski2017"
    _write_release(root / "raw")
    return p.EnvChemgenPiotrowski2017Dataset(root=str(root))


def _drop_log(built: Any) -> p.DropLog:
    path = Path(built.root, "preprocess", "dropped_records.json")
    return p.DropLog.model_validate_json(path.read_text())


def _condition_ledger(built: Any) -> p.ConditionLedger:
    path = Path(built.root, "preprocess", "conditions.json")
    return p.ConditionLedger.model_validate_json(path.read_text())


# --------------------------------------------------------------------------- #
# Dose parsing
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("cell", "fmt", "value", "unit"),
    [
        ("10-ug/ml", "General", 10.0, ConcentrationUnit.ug_per_ml),
        ("100-uM", "General", 100.0, ConcentrationUnit.micromolar),
        ("12.5-nM", "General", 12.5, ConcentrationUnit.nanomolar),
        ("26-mM", "General", 26.0, ConcentrationUnit.millimolar),
        ("26 mM", "General", 26.0, ConcentrationUnit.millimolar),
        ("50-ng/ml", "General", 0.05, ConcentrationUnit.ug_per_ml),
        ("7-ng/ml", "General", 0.007, ConcentrationUnit.ug_per_ml),
        (4e-05, "0.000%", 0.004, ConcentrationUnit.percent),
        (4e-05, "0.00%", 0.004, ConcentrationUnit.percent),
    ],
)
def test_parse_dose_reads_every_single_dose_form(
    cell: str | float, fmt: str, value: float, unit: ConcentrationUnit
) -> None:
    dose = p.parse_dose(cell, fmt)
    assert dose.kind == "single"
    assert dose.concentration is not None
    assert (dose.concentration.value, dose.concentration.unit) == (value, unit)
    assert dose.concentration.basis is DoseBasis.fixed


@pytest.mark.parametrize(
    ("cell", "low", "high", "unit"),
    [
        ("10-100-ug/ml", 10.0, 100.0, ConcentrationUnit.ug_per_ml),
        ("0.5-0.7-mM", 0.5, 0.7, ConcentrationUnit.millimolar),
        ("0.2-10ug/ml", 0.2, 10.0, ConcentrationUnit.ug_per_ml),
    ],
)
def test_parse_dose_reads_a_range_as_no_single_dose(
    cell: str, low: float, high: float, unit: ConcentrationUnit
) -> None:
    dose = p.parse_dose(cell, "General")
    assert (dose.kind, dose.concentration) == ("range", None)
    assert (dose.range_low, dose.range_high, dose.range_unit) == (low, high, unit)


@pytest.mark.parametrize(
    ("cell", "fmt"), [("10 ug", "General"), ("ten-uM", "General"), (4e-05, "General")]
)
def test_parse_dose_refuses_a_cell_it_was_not_taught(
    cell: str | float, fmt: str
) -> None:
    with pytest.raises(ValueError):
        p.parse_dose(cell, fmt)


# --------------------------------------------------------------------------- #
# Strains, as released
# --------------------------------------------------------------------------- #
def test_resolve_strains_stores_every_pool_row_as_released() -> None:
    """No genome: the ORF and common name are Data Set 2's, YAL058C-A included."""
    pool = [
        p.PoolStrain(orf=orf, common_name=label, barcode=bc)
        for label, orf, bc in _STRAINS
    ]
    strains = p.resolve_strains([s[0] for s in _STRAINS], pool)
    assert [(s.label, s.systematic_name, s.perturbed_gene_name) for s in strains] == [
        ("TFC3", "YAL001C", "TFC3"),
        ("VPS8", "YAL002W", "VPS8"),
        ("OLD1", "YAL003W", "OLD1"),
        ("GONE1", "YPL999C", "GONE1"),
        ("SNQ2", "YDR011W", "SNQ2"),
        ("YAL058C-A", "YAL058C-A", "YAL058C-A"),
        ("YBR001C", "YBR001C", "YBR001C"),
    ]
    assert "drop_rule" not in p.ResolvedStrain.model_fields
    assert not hasattr(p, "constructed_orf")
    assert not hasattr(p, "_canonical_common_names")
    assert not hasattr(p, "default_genome")


def _twins(
    values: list[list[float]],
) -> tuple[p.ReleasedMatrix, list[p.ResolvedStrain]]:
    entry = p.PoolStrain(orf="YAL001C", common_name="TFC3", barcode="A" * 20)
    other = p.PoolStrain(orf="YAL002W", common_name="VPS8", barcode="C" * 20)
    matrix = p.ReleasedMatrix(
        strains=["YAL001C", "TFC3", "VPS8"],
        condition_ids=["Cond-000001", "Cond-000002"],
        condition_labels=["A", "B"],
        values=np.asarray(values, dtype=np.float64),
    )
    strains = p.resolve_strains(matrix.strains, [entry, other])
    return matrix, strains


def test_one_pool_strain_under_two_labels_in_disjoint_columns_is_recorded() -> None:
    """The TCD1/YHR003C shape: one strain, labelled by ORF in one screen's columns."""
    matrix, strains = _twins([[1.0, np.nan], [np.nan, 2.0], [3.0, 4.0]])
    assert [s.systematic_name for s in strains] == ["YAL001C", "YAL001C", "YAL002W"]
    assert p.strains_labelled_twice(matrix, strains) == {"YAL001C": ["YAL001C", "TFC3"]}


def test_two_labels_scored_in_one_condition_are_refused_as_a_duplicate() -> None:
    matrix, strains = _twins([[1.0, np.nan], [5.0, 2.0], [3.0, 4.0]])
    with pytest.raises(ValueError, match="scored twice in one condition"):
        p.strains_labelled_twice(matrix, strains)


def test_resolve_strains_refuses_a_label_outside_the_pool() -> None:
    pool = [p.PoolStrain(orf="YAL001C", common_name="TFC3", barcode="A" * 20)]
    with pytest.raises(ValueError, match="not a pool v2.0 member"):
        p.resolve_strains(["NOPE"], pool)


def test_read_pool_refuses_a_pool_that_is_not_310_strains(tmp_path: Path) -> None:
    _write_release(tmp_path)
    workbook = openpyxl.load_workbook(tmp_path / p.POOL.filename)
    workbook["Suppl. dataset 2"].delete_rows(6)
    workbook.save(tmp_path / p.POOL.filename)
    with pytest.raises(ValueError, match="holds 309 strains"):
        p.read_pool(tmp_path / p.POOL.filename)


def test_read_matrix_maps_na_to_nan_and_keeps_the_ids(tmp_path: Path) -> None:
    _write_release(tmp_path)
    matrix = p.read_matrix(tmp_path / p.MATRIX.filename)
    assert matrix.strains == [s[0] for s in _STRAINS]
    assert matrix.condition_ids == [c[0] for c in _CONDITIONS]
    assert matrix.condition_labels[0] == "Benomyl"
    assert int(np.isnan(matrix.values).sum()) == len(_NA_CELLS)
    assert matrix.values[3, 4] == _value(3, 4)


# --------------------------------------------------------------------------- #
# The build
# --------------------------------------------------------------------------- #
def test_drop_log_accounts_for_every_numeric_cell(built: Any) -> None:
    log = _drop_log(built)
    rules = {rule.rule: rule for rule in log.rules}
    assert (log.released_cells, log.not_scored_cells) == (63, 4)
    assert (log.source_records, log.kept_records, log.dropped_records) == (59, 39, 20)
    assert [rule.rule for rule in log.rules] == [
        "condition_not_in_the_final_dataset",
        "dose_is_a_released_range",
        p.UNIDENTIFIED_COMPOUND_RULE,
    ]
    assert log.source_records == log.kept_records + sum(r.n_records for r in log.rules)
    assert rules["condition_not_in_the_final_dataset"].n_records == 7
    assert rules["condition_not_in_the_final_dataset"].items == [
        "Cond-000005 NPE0003 (10-ug/ml)"
    ]
    assert rules["dose_is_a_released_range"].n_records == 7
    assert rules["dose_is_a_released_range"].items == [
        "Cond-000004 Ouabain (10-100-ug/ml)"
    ]
    assert rules[p.UNIDENTIFIED_COMPOUND_RULE].n_records == 6
    assert rules[p.UNIDENTIFIED_COMPOUND_RULE].items == [
        "Cond-000003 NPD0002 (50-ng/ml)"
    ]
    assert log.not_scored_strains_by_screen == {
        "RIKEN": {"TFC3": 1},
        "NCI/NIH/GSK": {"VPS8": 3},
    }
    assert len(built) == log.kept_records


def test_records_are_condition_major_with_the_released_scores(built: Any) -> None:
    kept_rows = list(range(len(_STRAINS)))
    kept_columns = [0, 1, 5, 6, 7, 8]
    expected = [
        (j, i) for j in kept_columns for i in kept_rows if (i, j) not in _NA_CELLS
    ]
    assert len(expected) == len(built) == 39
    for idx, (j, i) in enumerate(expected):
        phenotype = built[idx]["experiment"]["phenotype"]
        assert phenotype["environment_response"] == _value(i, j)
        assert phenotype["replicate_id"] == _CONDITIONS[j][0]


def test_genotype_is_the_one_barcoded_pool_deletion(built: Any) -> None:
    (deletion,) = built[0]["experiment"]["genotype"]["perturbations"]
    assert deletion["perturbation_type"] == "barcoded_kanmx_deletion"
    assert (deletion["systematic_gene_name"], deletion["perturbed_gene_name"]) == (
        "YAL001C",
        "TFC3",
    )
    assert deletion["barcode"] == "AAAAAAAAAAAAAAAAAAAA"
    assert deletion["collection"] == p.POOL_COLLECTION.value
    assert deletion["cassette"] == "kanMX"
    assert deletion["constructed_orf"] is None
    # Column 0 holds every strain row, so records 0..6 are the seven strains in order.
    stored = [
        (
            perturbation["systematic_gene_name"],
            perturbation["perturbed_gene_name"],
            perturbation["constructed_orf"],
        )
        for idx in range(len(_STRAINS))
        for perturbation in built[idx]["experiment"]["genotype"]["perturbations"]
    ]
    assert stored == [(orf, label, None) for label, orf, _ in _STRAINS]
    (released,) = built[5]["experiment"]["genotype"]["perturbations"]
    assert released["systematic_gene_name"] == "YAL058C-A"  # stored as released
    assert released["perturbed_gene_name"] == "YAL058C-A"


def test_environment_is_ypgal_liquid_at_30c_for_48h_in_dmso(built: Any) -> None:
    environment = built[0]["experiment"]["environment"]
    assert environment["media"] == p.YP_GALACTOSE_LIQUID.model_dump()
    assert environment["media"]["state"] == "liquid"
    assert environment["media"]["base_medium"] == "YP"
    assert environment["temperature"]["value"] == 30.0
    assert environment["duration_hours"] == 48.0
    assert environment["duration_generations"] is None
    assert {g["field"] for g in environment["provenance_gaps"]} == {
        "duration_generations",
        "auxotroph_supplements",
    }
    culture = environment["culture_format"]
    assert (culture["vessel"], culture["working_volume_ul"]) == (
        "96-well plates",
        200.0,
    )
    assert culture["inoculum_cells_per_strain"] == 250.0
    assert culture["endpoint"] == "fixed_duration"
    assert [g["field"] for g in culture["provenance_gaps"]] == ["shaking_rpm"]
    assert environment["pre_culture"]["source"] == "frozen_stock"
    (treatment,) = environment["perturbations"]
    assert treatment["solvent"]["name"] == "DMSO"
    assert treatment["solvent"]["percent"] == 1.0


def test_doses_keep_the_released_number_and_unit(built: Any) -> None:
    def dose(idx: int) -> tuple[Any, ...]:
        (treatment,) = built[idx]["experiment"]["environment"]["perturbations"]
        c = treatment["concentration"]
        return (c["value"], c["unit"], c["basis"])

    assert dose(0) == (10.0, "ug/mL", "fixed")  # Benomyl 10-ug/ml
    assert dose(7) == (100.0, "uM", "fixed")  # NPD0001 100-uM
    assert dose(14) == (0.004, "percent", "fixed")  # MMS 4e-05 under 0.000%
    assert dose(32) == (0.05, "ug/mL", "fixed")  # Kundrymycin 50-ng/ml


def test_compound_takes_the_table_the_released_inchi_or_a_typed_gap(built: Any) -> None:
    def compound(idx: int) -> dict[str, Any]:
        (treatment,) = built[idx]["experiment"]["environment"]["perturbations"]
        result: dict[str, Any] = treatment["compound"]
        return result

    benomyl = compound(0)
    assert benomyl["name"] == resolve_compound_identity(name="Benomyl").name
    assert benomyl["inchikey"] == resolve_compound_identity(name="Benomyl").inchikey
    assert benomyl["smiles"] == _BENOMYL_SMILES  # the released SMILES rides along
    released = compound(7)
    assert (released["name"], released["inchikey"]) == ("NPD0001", _ETHANOL_KEY)
    assert released["inchi"] == _ETHANOL_INCHI
    gapped = compound(32)  # a released name with no structure: identified, gapped
    assert (gapped["name"], gapped["inchikey"], gapped["smiles"]) == (
        "Kundrymycin",
        None,
        None,
    )
    assert [(g["field"], g["reason"]) for g in gapped["provenance_gaps"]] == [
        ("inchikey", "deferred_pending_source_review")
    ]


def test_screen_and_replicate_ids_and_the_per_screen_reference(built: Any) -> None:
    riken, nci = built[0], built[14]
    assert riken["experiment"]["phenotype"]["screen_id"] == "RIKEN"
    assert nci["experiment"]["phenotype"]["screen_id"] == "NCI/NIH/GSK"
    assert riken["reference"]["phenotype_reference"]["screen_id"] == "RIKEN"
    assert nci["reference"]["phenotype_reference"]["screen_id"] == "NCI/NIH/GSK"
    # the same compound at the same dose in two NCI libraries stays two conditions
    first, second = built[20], built[26]
    assert first["experiment"]["environment"] == second["experiment"]["environment"]
    assert first["experiment"]["phenotype"]["replicate_id"] == "Cond-000007"
    assert second["experiment"]["phenotype"]["replicate_id"] == "Cond-000008"


def test_reference_is_the_screens_dmso_control_at_zero(built: Any) -> None:
    reference = built[0]["reference"]
    (dmso,) = reference["environment_reference"]["perturbations"]
    assert dmso["compound"]["name"] == resolve_compound_identity(name="DMSO").name
    concentration = dmso["concentration"]
    assert (concentration["value"], concentration["unit"]) == (1.0, "percent")
    phenotype = reference["phenotype_reference"]
    assert phenotype["environment_response"] == 0.0
    assert phenotype["measurement_type"] == "z_score"
    assert {g["field"] for g in phenotype["provenance_gaps"]} == {
        "environment_response_se",
        "environment_response_uncertainty",
        "environment_response_uncertainty_type",
        "n_samples",
        "sample_unit",
    }
    genome = reference["genome_reference"]
    assert genome["strain"] == LIBRARY_STRAIN
    assert genome["background"]["mating_type"] == "a"


def test_phenotype_states_the_conservative_replicate_count(built: Any) -> None:
    phenotype = built[0]["experiment"]["phenotype"]
    assert phenotype["measurement_type"] == "z_score"
    assert phenotype["assay_type"] == "pooled_competitive_growth_barcode"
    assert (phenotype["n_samples"], phenotype["sample_unit"]) == (
        1,
        "technical_replicate",
    )
    assert "positive = resistant, negative = sensitive" in phenotype["units"]
    derivation = StatDerivation.model_validate_json(
        Path(built.root, "preprocess", "n_samples_derivation.json").read_text()
    )
    assert (derivation.value, derivation.range_low, derivation.range_high) == (
        1.0,
        1.0,
        3.0,
    )
    assert derivation.method.value == "conservative_low"


def test_condition_ledger_lists_routes_and_the_pool_census(built: Any) -> None:
    ledger = _condition_ledger(built)
    assert [row.condition_id for row in ledger.conditions] == [
        c[0] for c in _CONDITIONS
    ]
    assert ledger.route_counts == {"curated_table": 2, "gapped": 1, "released_inchi": 3}
    by_id = {row.condition_id: row for row in ledger.conditions}
    assert by_id["Cond-000003"].drop_rule == p.UNIDENTIFIED_COMPOUND_RULE
    assert by_id["Cond-000003"].compound is None
    assert by_id["Cond-000010"].compound is not None
    assert by_id["Cond-000010"].compound.route == "gapped"
    assert by_id["Cond-000004"].drop_rule == "dose_is_a_released_range"
    assert by_id["Cond-000004"].compound is None
    assert by_id["Cond-000006"].screen == "NCI/NIH/GSK"
    assert by_id["Cond-000001"].in_high_confidence_set is True
    assert by_id["Cond-000002"].n_scored_strains == 7
    assert "renamed_strains" not in p.ConditionLedger.model_fields
    assert ledger.strains_labelled_twice == {}
    assert len(ledger.absent_pool_strains) == p.POOL_V2.value - len(_STRAINS)
    assert ledger.released_key_comparisons == 3
    assert ledger.released_key_disagreements == 0


def test_unidentified_agent_is_held_out_with_its_count_and_quotes(built: Any) -> None:
    log = _drop_log(built)
    (held,) = log.unidentified_compounds_held_out
    assert held.model_dump() == {
        "condition_id": "Cond-000003",
        "source_name": "NPD0002",
        "collection": "RIKEN NPDepo",
        "screen": "RIKEN",
        "source_dose": "50-ng/ml",
        "n_records": 6,
    }
    assert log.unidentified_compound_rule == p.UNIDENTIFIED_COMPOUND_RULE
    assert [sv.provenance.sha256 for sv in log.unidentified_compound_sources] == [
        p.CONDITIONS.sha256,
        p.CONDITIONS.sha256,
        p.PAPER_MD_SHA256,
    ]
    stored = {
        treatment["compound"]["name"]
        for idx in range(len(built))
        for treatment in built[idx]["experiment"]["environment"]["perturbations"]
    }
    assert "NPD0002" not in stored
    assert "Kundrymycin" in stored
    assert not any(
        built[idx]["experiment"]["phenotype"]["replicate_id"] == "Cond-000003"
        for idx in range(len(built))
    )
    assert log.source_records == log.kept_records + sum(r.n_records for r in log.rules)


@pytest.mark.parametrize(
    ("name", "smiles", "inchi", "identified"),
    [
        ("NPD403", None, None, False),  # a library code, no structure
        ("NSC76455", None, None, False),  # a registry number, no structure
        ("IB01212", None, None, False),
        ("GI261520A", "Cl.COc1ccc2ncnc(Nc3ccc(OCc4ccccc4)cc3)c2c1", None, True),
        ("NSC305780", "c1c(ccc2c1cccc2)Sc1cc2c(nc(nc2cc1)N)N.*.CCO", None, True),
        ("Kundrymycin", None, None, True),  # a released name
        ("Mucopeptin B", None, None, True),
        ("Lactoferrin, from Bovin Milk", None, None, True),
    ],
)
def test_condition_is_identified_by_a_released_name_or_structure(
    name: str, smiles: str | None, inchi: str | None, identified: bool
) -> None:
    row = p.ConditionRow(
        condition_id="Cond-000001",
        name=name,
        concentration_cell="10-ug/ml",
        concentration_number_format="General",
        collection="RIKEN NPDepo",
        in_final_dataset=True,
        in_high_confidence_set=False,
        percent_growth=None,
        smiles=smiles,
        inchi=inchi,
    )
    assert p.condition_is_identified(row) is identified


def test_a_released_identity_admits_a_held_out_label(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = SourcedValue(
        value="ethanol",
        quote="synthetic release",
        provenance=Provenance(
            source_uri="synthetic.txt", citation_key="synthetic", sha256="0" * 64
        ),
    )
    monkeypatch.setitem(
        p.RELEASED_IDENTITIES,
        "NPD0002",
        ReleasedIdentity(name="NPD0002", smiles="CCO", source=source),
    )
    monkeypatch.setattr(
        p.EnvChemgenPiotrowski2017Dataset, "download", lambda self: None
    )
    root = tmp_path / "env_chemgen_piotrowski2017"
    _write_release(root / "raw")
    dataset = p.EnvChemgenPiotrowski2017Dataset(root=str(root))
    log = _drop_log(dataset)
    assert log.unidentified_compounds_held_out == []
    assert (log.source_records, log.kept_records) == (59, 45)
    ledger = _condition_ledger(dataset)
    entry = {row.condition_id: row for row in ledger.conditions}["Cond-000003"]
    assert entry.drop_rule is None
    assert entry.compound is not None
    assert (entry.compound.route, entry.compound.stored_inchikey) == (
        "released_identity",
        _ETHANOL_KEY,
    )


def test_sourced_values_are_written_beside_the_build(built: Any) -> None:
    written = p.SourcedValues.model_validate_json(
        Path(built.root, "preprocess", "sourced_values.json").read_text()
    )
    assert set(written.values) == set(p.sourced_values())
    assert "Z_SCORE_POLARITY" in written.values


# --------------------------------------------------------------------------- #
# Verification entry point
# --------------------------------------------------------------------------- #
def test_verify_build_refuses_a_root_the_registry_does_not_name(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="is not the registered store"):
        p.verify_build(str(tmp_path / "elsewhere"), data_root=str(tmp_path))


def test_verify_build_delegates_to_the_family_entry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from torchcell.verification import runners

    report = VerificationReport(
        dataset_name=p.DATASET_SLUG, provenance=Provenance(source_uri="stand-in")
    )
    calls: list[tuple[str, str]] = []

    def entry(name: str, data_root: str) -> VerificationReport:
        calls.append((name, data_root))
        return report

    monkeypatch.setattr(runners, "verify_environment_response_entry", entry)
    root = osp.join(str(tmp_path), "data/torchcell", p.DATASET_SLUG)
    assert p.verify_build(root, data_root=str(tmp_path)) is report
    assert calls == [(p.DATASET_SLUG, str(tmp_path))]


# --------------------------------------------------------------------------- #
# Raw mirror and download
# --------------------------------------------------------------------------- #
def _stage_library(data_root: Path) -> None:
    """Three stand-in release files in the literature mirror, pins repointed at them."""
    for release in p.RELEASE_FILES:
        path = data_root / p.LIBRARY_DIR_REL / release.library_rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(release.filename.encode())


def _repin(monkeypatch: pytest.MonkeyPatch) -> None:
    import hashlib

    for name in ("MATRIX", "CONDITIONS", "POOL"):
        release = getattr(p, name)
        digest = hashlib.sha256(release.filename.encode()).hexdigest()
        monkeypatch.setattr(p, name, release.model_copy(update={"sha256": digest}))
    monkeypatch.setattr(p, "RELEASE_FILES", (p.MATRIX, p.CONDITIONS, p.POOL))


def test_deposit_raw_mirror_copies_and_records_the_pmc_retrieval(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _repin(monkeypatch)
    _stage_library(tmp_path)
    mirror = p.deposit_raw_mirror(str(tmp_path))
    manifest = p.load_manifest(str(tmp_path))
    assert [record.path for record in manifest.files] == [
        f"si/{release.filename}" for release in p.RELEASE_FILES
    ]
    for record, release in zip(manifest.files, p.RELEASE_FILES, strict=True):
        assert record.sha256 == release.sha256
        assert record.retrieval is not None
        assert record.retrieval.method.value == "pmc_cloud"
        assert record.retrieval.params == {"key": release.pmc_key}
        assert record.retrieval.source_url == (
            f"https://pmc-oa-opendata.s3.amazonaws.com/PMC6056180.1/{release.filename}"
        )
        assert (mirror / record.path).read_bytes() == release.filename.encode()
    # idempotent: a second deposit over the same bytes is a no-op
    assert p.deposit_raw_mirror(str(tmp_path)) == mirror
    assert p.manifest_sha256(manifest, p.MATRIX.mirror_rel) == p.MATRIX.sha256
    with pytest.raises(KeyError):
        p.manifest_sha256(manifest, "si/absent.xlsx")


def test_deposit_raw_mirror_refuses_drifted_bytes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _repin(monkeypatch)
    _stage_library(tmp_path)
    p.deposit_raw_mirror(str(tmp_path))
    (p.raw_mirror_dir(str(tmp_path)) / p.POOL.mirror_rel).write_bytes(b"edited")
    with pytest.raises(RuntimeError, match="exists with another sha256"):
        p.deposit_raw_mirror(str(tmp_path))
    (tmp_path / p.LIBRARY_DIR_REL / p.MATRIX.library_rel).write_bytes(b"drift")
    with pytest.raises(RuntimeError, match="!= pinned"):
        p.deposit_raw_mirror(str(tmp_path))


def test_download_links_the_mirror_and_refuses_an_off_pin_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from torchcell.data import ManifestPinMismatchError

    _repin(monkeypatch)
    _stage_library(tmp_path)
    p.deposit_raw_mirror(str(tmp_path))
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    dataset = p.EnvChemgenPiotrowski2017Dataset.__new__(
        p.EnvChemgenPiotrowski2017Dataset
    )
    dataset.root = str(tmp_path / "store")
    dataset.download()
    linked = sorted(os.listdir(dataset.raw_dir))
    assert linked == sorted(release.filename for release in p.RELEASE_FILES)
    manifest_path = p.raw_mirror_dir(str(tmp_path)) / "manifest.json"
    manifest = p.load_manifest(str(tmp_path))
    manifest.files[0].sha256 = "0" * 64
    manifest_path.write_text(manifest.model_dump_json())
    with pytest.raises(ManifestPinMismatchError):
        dataset.download()


# --------------------------------------------------------------------------- #
# Mirror audits (skip when the literature mirror is not mounted)
# --------------------------------------------------------------------------- #
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

#: The worksheet and cell each sheet-sourced value quotes.
_SHEET_CELLS = {
    "MATRIX_TITLE": ("All_conditions", "A1"),
    "POOL_COLLECTION": ("Suppl. dataset 2", "A4"),
    "FINAL_DATASET_FLAG": ("Suppl. dataset 4", "E6"),
    "NPDEPO_NOMENCLATURE": ("Suppl. dataset 4", "A2"),
    "STRUCTURE_STRINGS": ("Suppl. dataset 4", "A3"),
}


def _library() -> Path:
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT not set")
    library = Path(data_root) / p.LIBRARY_DIR_REL
    if not (library / p.PAPER_MD).is_file():
        pytest.skip("paper mirror not mounted")
    return library


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    return "\n".join(
        "".join(node.text or "" for node in paragraph.iter(_W + "t"))
        for paragraph in root.iter(_W + "p")
    )


@pytest.mark.data
def test_every_sourced_value_is_a_verbatim_quote_in_the_mirror() -> None:
    library = _library()
    values: dict[str, SourcedValue] = p.sourced_values()
    # 32 of this module's own, plus ARRAY_KANMX (vanacloig2022's quote of this paper)
    assert len(values) == 33
    paper = (library / p.PAPER_MD).read_text(encoding="utf-8")
    note = _docx_text(library / p.NOTE_DOCX)
    for name, value in values.items():
        uri = value.provenance.source_uri
        assert value.provenance.citation_key == p.CITATION_KEY, name
        assert p._sha256(library / uri) == value.provenance.sha256, name
        if uri == p.PAPER_MD:
            assert value.quote in paper, name
        elif uri == p.NOTE_DOCX:
            assert value.quote in note, name
        else:
            sheet, cell = _SHEET_CELLS[name]
            workbook = openpyxl.load_workbook(library / uri, read_only=True)
            assert workbook[sheet][cell].value == value.quote, name
            workbook.close()


@pytest.mark.data
def test_release_pins_match_the_literature_mirror_manifest() -> None:
    from torchcell.literature.manifest import Manifest

    library = _library()
    manifest = Manifest.model_validate_json((library / "manifest.json").read_text())
    records = {record.path: record for record in manifest.files}
    for release in p.RELEASE_FILES:
        record = records[release.library_rel]
        assert record.sha256 == release.sha256
        assert record.original_filename == release.filename
        assert record.retrieval is not None
        assert record.retrieval.source_url == release.source_url
        assert record.retrieval.retrieved_at == p.RETRIEVED_AT
    assert records[p.NOTE_DOCX].sha256 == p.NOTE_DOCX_SHA256
    assert records[p.PAPER_MD].sha256 == p.PAPER_MD_SHA256
