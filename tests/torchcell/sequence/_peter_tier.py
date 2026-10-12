# tests/torchcell/sequence/_peter_tier.py
"""Synthetic Peter 2018 genomes-tier files for the natural-isolate tests.

``torchcell.sequence.natural_isolate`` reads six files from the Peter tier (the
reference-gene tarball, the presence and copy-number matrices, the two loss-of-function
matrices and the Supplementary Tables workbook) and the R64 chromosomes and GFF from the
SGD tier. The loader tests already write the first three; this module writes the other
three in the release's own formats and the GFF, so every sha256 check on the build path
runs against real manifests:

- ``1011LossOfFunction.xls.gz``: gzip of tab-separated text with CR line ends, columns
  ``CODE NAME Gene SGDorder CH START STOP`` then one 0/1 column per isolate.
- ``genesMatrix_Frameshift.tab.gz``: gzip of tab-separated text, columns
  ``CODE Gene SGDorder CH START STOP`` then the isolates; ``CODE`` may be empty.
- ``41586_2018_30_MOESM3_ESM.xls``: an xlsx written under the publisher's ``.xls`` name
  (pandas picks the engine from the bytes). Sheet ``Table S1`` has a title row, the
  aneuploidy legend, a blank row, the header on sheet row 4, one row per isolate from
  sheet row 5, then a blank row and a reference list whose rows have no standardized name.

Tests with no loss-of-function rows get two matrices that index the panel and carry no
row, and a Table S1 row per isolate with ``HO deletion`` no, ``Ploidy`` 2, ``euploid``
and ``homozygous``.
"""

from __future__ import annotations

import gzip
import hashlib
import io
import re
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path

import openpyxl
from pydantic import BaseModel, ConfigDict

from torchcell.literature.manifest import ArtifactRecord
from torchcell.sequence import natural_isolate as ni
from torchcell.sequence.genome.registry import GenomeManifest


class LofRow(BaseModel):
    """One row of the loss-of-function matrix; ``calls`` are the isolates set to 1."""

    model_config = ConfigDict(frozen=True)

    code: int
    name: str
    gene: str
    ch: int
    start: int
    stop: int
    calls: frozenset[str] = frozenset()


class FrameshiftRow(BaseModel):
    """One row of the frameshift matrix; ``code`` None is a row released without one."""

    model_config = ConfigDict(frozen=True)

    code: int | None
    gene: str
    ch: int
    start: int
    stop: int
    calls: frozenset[str] = frozenset()


class TableS1Row(BaseModel):
    """One isolate row of Supplementary Table 1 (``None`` writes an empty cell)."""

    model_config = ConfigDict(frozen=True)

    isolate: str
    ho_deletion: str = "no"
    ploidy: int | str = 2
    aneuploidies: str | None = "euploid"
    zygosity: str = "homozygous"


class GffGene(BaseModel):
    """One gene-level GFF row."""

    model_config = ConfigDict(frozen=True)

    id: str
    chrom: str
    start: int
    end: int
    standard: str | None = None
    aliases: tuple[str, ...] = ()


def lof_matrix(isolates: Sequence[str], rows: Iterable[LofRow]) -> bytes:
    """The loss-of-function matrix, gzip, CR line ends, byte-stable."""
    lines = ["\t".join([*ni.LOF_META, *isolates])]
    for i, r in enumerate(rows):
        cells = [str(r.code), r.name, r.gene, str(i + 1), str(r.ch), str(r.start)]
        cells += [str(r.stop), *("1" if iso in r.calls else "0" for iso in isolates)]
        lines.append("\t".join(cells))
    return gzip.compress("\r".join(lines).encode("latin-1"), mtime=0)


def frameshift_matrix(isolates: Sequence[str], rows: Iterable[FrameshiftRow]) -> bytes:
    """The frameshift matrix, gzip, LF line ends, byte-stable."""
    lines = ["\t".join([*ni.FRAMESHIFT_META, *isolates])]
    for i, r in enumerate(rows):
        code = "" if r.code is None else str(r.code)
        cells = [code, r.gene, str(i + 1), str(r.ch), str(r.start), str(r.stop)]
        cells += ["1" if iso in r.calls else "0" for iso in isolates]
        lines.append("\t".join(cells))
    return gzip.compress(("\n".join(lines) + "\n").encode(), mtime=0)


def tables_workbook(
    rows: Iterable[TableS1Row], table_s2: Mapping[tuple[int, int], str] | None = None
) -> bytes:
    """The Supplementary Tables workbook with a ``Table S1`` sheet (xlsx bytes).

    ``table_s2`` writes ``{(sheet row, column): cell}`` into the ``Table S2`` sheet, for
    the Peter 2018 verifier, which reads both sheets of the one workbook.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = ni.TABLE_S1_SHEET
    sheet.append(["Table S1: Overview of the S. cerevisiae collection"])
    sheet.append([ni.ANEUPLOIDY_LEGEND])
    sheet.append([])
    sheet.append(
        [
            "Isolate name",
            "Standardized name",
            "HO deletion",
            "Ploidy",
            "Aneuploidies",
            "Zygosity",
        ]
    )
    for r in rows:
        sheet.append(
            [f"strain {r.isolate}", r.isolate, r.ho_deletion, r.ploidy]
            + [r.aneuploidies, r.zygosity]
        )
    sheet.append([])
    sheet.append(["References"])
    sheet.append(["1. A reference row with no standardized name."])
    s2 = book.create_sheet("Table S2")
    for (row, column), value in (table_s2 or {}).items():
        s2.cell(row=row, column=column, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return _byte_stable(buffer.getvalue())


def _byte_stable(xlsx: bytes) -> bytes:
    """``xlsx`` with its save-time stamps fixed, so its sha256 is a constant.

    openpyxl writes the save time into ``docProps/core.xml`` and into every zip entry's
    date, so two saves of one workbook differ and a tier pin written by a fixture would
    not equal the one an expectation computes.
    """
    stamp = b"2026-10-12T00:00:00Z"
    source = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(
                    rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                    rb"\g<1>" + stamp + rb"\g<3>",
                    body,
                )
            entry = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 12, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, body)
    return out.getvalue()


def gff(genes: Iterable[GffGene]) -> bytes:
    """A gene-level R64 GFF (attribute values written as SGD writes them)."""
    lines = ["##gff-version 3"]
    for g in genes:
        attrs = [f"ID={g.id}", f"Name={g.id}"]
        if g.standard is not None:
            attrs.append(f"gene={g.standard}")
        if g.aliases:
            attrs.append("Alias=" + ",".join(g.aliases))
        fields = [g.chrom, "SGD", "gene", str(g.start), str(g.end), ".", "+", "."]
        lines.append("\t".join([*fields, ";".join(attrs)]))
    lines.append("##FASTA")
    return ("\n".join(lines) + "\n").encode()


def genotype_files(
    isolates: Sequence[str],
    *,
    lof: Iterable[LofRow] = (),
    frameshift: Iterable[FrameshiftRow] = (),
    table_s1: Iterable[TableS1Row] | None = None,
) -> dict[str, bytes]:
    """The three Peter-tier files the shared module adds to a loader's raw set."""
    rows = (
        list(table_s1)
        if table_s1 is not None
        else [TableS1Row(isolate=i) for i in isolates]
    )
    return {
        ni.LOF_NAME: lof_matrix(isolates, lof),
        ni.FRAMESHIFT_NAME: frameshift_matrix(isolates, frameshift),
        ni.SUPPLEMENTARY_TABLES_NAME: tables_workbook(rows),
    }


def write_tier(data_root: Path, assembly_set: str, files: dict[str, bytes]) -> None:
    """A genomes-tier set holding ``files`` with a real ``GenomeManifest``."""
    tier = data_root / "torchcell-genomes" / assembly_set
    tier.mkdir(parents=True)
    records = []
    for name, data in files.items():
        (tier / name).write_bytes(data)
        records.append(
            ArtifactRecord(
                path=name,
                role="sequence",
                bytes=len(data),
                sha256=hashlib.sha256(data).hexdigest(),
            )
        )
    manifest = GenomeManifest(
        assembly_set=assembly_set,
        organism="Saccharomyces cerevisiae",
        strain_or_population="synthetic",
        source="test",
        release="test",
        files=records,
        provenance_complete=True,
        created_at="2026-10-12T00:00:00+00:00",
    )
    (tier / "manifest.json").write_text(manifest.model_dump_json(indent=2))
