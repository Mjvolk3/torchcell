# torchcell/datasets/ecoli/schastnaya2021
# [[torchcell.datasets.ecoli.schastnaya2021]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/schastnaya2021
# Test file: tests/torchcell/datasets/ecoli/test_schastnaya2021.py
"""Schastnaya 2021 metabolome of the KEIO deletion arm of an E. coli phosphosite screen.

Schastnaya, Raguz Nakic, Gruber, Doubleday, Krishnan, Johns, Park, Wang and Sauer 2021
(Nat Commun 12:5650, doi:10.1038/s41467-021-25988-4) mutated 52 phosphosites on 23
enzymes by MAGE, grew each strain plus the matching gene-deletion strain on M9 with one
of five carbon sources, and profiled the intracellular metabolome by flow-injection
TOF mass spectrometry in negative mode. Supplementary Data 3 releases, per (strain,
carbon source), the log2 fold change of each annotated ion against the wild type with a
Benjamini-Hochberg adjusted p-value.

WHAT THE STORED QUANTITY IS. A RELATIVE metabolite level: the log2 fold change of one
annotated deprotonated ion's abundance in the mutant over its abundance in the wild
type, pooled over the strain's replicates (``SOURCED_VALUES["log2_fold_change"]``). It
is not a concentration, not a phosphosite occupancy, not an enzyme activity and not a
flux: the paper's flux arm covers four strains in one condition and its activity arm
four purified enzymes, neither of which is this release.

DATA. Two publisher workbooks from the PMC Article Datasets bucket: Supplementary Data
3 (``si6.xlsx``, the values; 464 ion rows x 200 (strain, carbon source) columns in a
LOG2(FC) block and the same 200 in an adjusted-p-value block) and Supplementary Data 1
(``si4.xlsx``, the 89 phosphomutant strains with the b-number of their enzyme and
whether the substitution abolishes or mimics phosphorylation), read to name every
dropped record in the retention ledger. The MassIVE deposit MSV000087795 holds the raw
spectra, not this matrix, and is NOT mirrored.

STRAIN. ``REFERENCE_STRAIN = "MG1655"``: the wild type the fold change is taken against
is "E. coli MG1655 dmutS with kanamycin resistance (referred to as wild-type)", carried
as a ``BacterialStrainBackground`` on the pinned MG1655 assembly, and the b-numbers
Supplementary Data 1 releases for the screened enzymes are MG1655 tags. The deletion
strains themselves "were retrieved from the KEIO collection", which Baba 2006 (mirrored)
builds in BW25113, so each released fold change compares a BW25113 deletion against an
MG1655 wild type. The schema has no field for a perturbed strain whose background
differs from its control's, so that is recorded here, in ``SOURCED_VALUES`` and in
``preprocess/strains.csv``, and the deletion carries ``collection`` and ``cassette``
from Baba.

RECORDS DROPPED (rules + items in ``preprocess/dropped_records.json``). 171 of the 200
released columns are phosphomutants, and the schema has NO bacterial codon-substitution
perturbation leaf: ``AllelePerturbation`` is the amino-acid-substitution leaf but
inherits the R64 systematic-name validator, so ``b2029`` is refused, and the five
bacterial leaves (deletion, transposon insertion, CRISPRi, promoter replacement,
heterologous pathway) all assert something the MAGE point mutation is not. Writing one
of them would state a genotype the paper did not make, so those records are dropped
under ``phosphosite_substitution_has_no_bacterial_leaf`` and the gap is the finding.
The 29 kept columns are the 16 KEIO deletion strains over their carbon sources.

PHENOTYPE. ``MetabolitePhenotype`` keyed by the released ``Formula``, the neutral
formula of the deprotonated ion, which is unique over the 464 rows (measured at build
time). A column's ``nd`` cells are not stored: detection varies by extraction protocol,
so a record carries 332, 309 or 284 keys, and the build asserts that split against the
paper's own hot/cold extraction lists. ``target_metabolite_ids`` maps a key to its KEGG
id wherever the row names exactly ONE; a row whose mass cannot separate several KEGG
compounds is ABSENT from the map rather than assigned one of its candidates, with every
candidate set in ``preprocess/metabolite_identity.json`` (Rapp 2026's rule; a typed gap
cannot sit beside a partially populated map, issue #753). ``n_replicates`` is 3 for
every key, the conservative floor of a count the release does not carry per record.
``metabolite_level_se`` is a typed gap. The reference is the wild type, whose level is 0
for every key BY THE DEFINITION of the released statistic.

NOT STORED. The adjusted p-value of every stored value is released and has no field on
``MetabolitePhenotype``; it is written to ``preprocess/adjusted_p_values.csv`` rather
than forced into ``metabolite_level_se``, which is a standard error, not a p-value.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import os.path as osp
import shutil
from collections.abc import Callable, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, ClassVar, Final

import openpyxl
import pandas as pd
from pydantic import BaseModel, ConfigDict
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
    write_verified,
)
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.schema import (
    BACTERIAL_ASSEMBLY_SETS,
    AssemblyReferenceGenome,
    BacterialAssemblySet,
    BacterialDeletionPerturbation,
    BacterialMetaboliteExperiment,
    BacterialMetaboliteExperimentReference,
    BacterialStrainBackground,
    Concentration,
    ConcentrationUnit,
    Environment,
    Experiment,
    ExperimentReference,
    Genotype,
    Media,
    MediaComponent,
    MediaComponentRole,
    MetabolitePhenotype,
    Publication,
    Temperature,
)
from torchcell.datasets.bacteria_common import (
    LOCUS_TAG_PATTERNS,
    STRAIN_GENE_NAMESPACES,
    LocusTagReconciliation,
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ROLE_RAW_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.literature.provenance import run_retriever
from torchcell.sequence.genome.ecoli.k12 import EcoliK12Genome, EcoliK12StrainName
from torchcell.verification.report import (
    Level,
    LevelResult,
    Provenance,
    VerificationReport,
)
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
    audit_sourced_value,
)

logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
CITATION_KEY = "schastnayaExtensiveRegulationEnzyme2021"
PAPER_DOI = "10.1038/s41467-021-25988-4"
PMCID = "PMC8463566"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
LIBRARY_DIR_REL = f"torchcell-library/{CITATION_KEY}"

PAPER_MD = "paper.md"
PAPER_MD_SHA256 = "f299eee75a1f022e08ba7f1bc444922db357f220be94f19466a36881214546c3"
#: The publisher's own description of each Supplementary Data file, MinerU OCR.
SI3_MD = "si/si3.md"
SI3_MD_SHA256 = "32bb2072d4dc7b8021ce217423bc764ac2071d75e3b42a948e308e7cfb25ab5e"
#: The Nature Research Reporting Summary. Its form is a scan with no text layer, so the
#: sample-size and replication cells live only in the MinerU table IMAGE below; they are
#: quoted in notes, never as an audited quote (the audit reads text).
REPORTING_SUMMARY_PDF = "si/si12.pdf"
REPORTING_SUMMARY_SHA256 = (
    "0102c0961ce1044544f6147ab36e516db08f2dfd3bcc0def988b0ea28f8c13c2"
)
REPORTING_SUMMARY_TABLE_IMAGE = (
    "si/images/si12/"
    "c9ce0df8de37533f33875de9d8b4f61debd4b1ac0bb6755506bab5db4f7495d1.jpg"
)
REPORTING_SUMMARY_TABLE_SHA256 = (
    "530ff753d501b747552fd55b25432c104546f90856e44ef8281d9944c6ed1092"
)
#: Baba 2006, the collection paper Schastnaya defers its deletion strains to.
BABA_KEY = "babaConstructionEscherichiaColi2006"
BABA_SHA256 = "ca71475baf25d070562c70296d5aad8ec17a59d5b843f264b1e226362b6daf0d"

SD1_FILE = "si4.xlsx"
SD1_SHEET = "SD1"
SD3_FILE = "si6.xlsx"
SD3_SHEET = "SD3"
MASSIVE_ACCESSION = "MSV000087795"
PRIDE_ACCESSION = "PXD027243"


class RawFile(BaseModel):
    """One file the loader consumes: its pinned bytes and how they were retrieved."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    name: str
    sha256: str
    bytes: int
    description: str
    retrieval: RetrievalRecord

    @property
    def mirror_relpath(self) -> str:
        """Path inside the raw mirror (``data/<name>``)."""
        return f"data/{self.name}"


def _pmc_file(
    name: str, moesm: int, sha256: str, size: int, description: str
) -> RawFile:
    """A publisher supplementary workbook from the PMC Article Datasets bucket.

    The retrieval record is copied from the literature mirror's ``manifest.json``, which
    is where these bytes were first captured.
    """
    key = f"{PMCID}.1/41467_2021_25988_MOESM{moesm}_ESM.xlsx"
    return RawFile(
        name=name,
        sha256=sha256,
        bytes=size,
        description=description,
        retrieval=RetrievalRecord(
            method=RetrievalMethod.pmc_cloud,
            source_url=f"https://pmc-oa-opendata.s3.amazonaws.com/{key}",
            retriever="torchcell.literature.retrieve.pmc_cloud_object",
            params={"key": key},
            sha256=sha256,
            retrieved_at="2026-10-07T11:42:57.784551+00:00",
        ),
    )


RAW_FILES: tuple[RawFile, ...] = (
    _pmc_file(
        SD1_FILE,
        4,
        "8acdfa9fcc3a09f73145e9188f11779b75da8c30cea36ebb0a120d4921a083ff",
        13644,
        "Supplementary Data 1: the 89 phosphomutant strains, each with its enzyme's "
        "b-number and whether the substitution abolishes or mimics phosphorylation",
    ),
    _pmc_file(
        SD3_FILE,
        6,
        "ad9fb64f29b4d7cac7c9575217f7fa8fadb401d4910fb8d330685d42b43d49d4",
        1624605,
        "Supplementary Data 3: per (strain, carbon source) log2 fold change of every "
        "annotated ion against the wild type, plus the adjusted p-value",
    ),
)
#: ``{raw file name: pinned sha256}``, the build-time check of every consumed file.
DATA_SHA256: dict[str, str] = {f.name: f.sha256 for f in RAW_FILES}
RAW_FILES_BY_NAME: dict[str, RawFile] = {f.name: f for f in RAW_FILES}

#: Released files deliberately not mirrored (the loader does not read them).
NOT_MIRRORED = (
    "MassIVE MSV000087795 (the metabolomics raw spectra): the released matrix the "
    "loader consumes is Supplementary Data 3, not the spectra",
    f"PRIDE {PRIDE_ACCESSION} (intact-protein mass spectrometry of purified enzymes): "
    "a different assay, not this dataset",
    "Supplementary Data 2, 4, 5, 6, 7, 8 (growth rates, local metabolic changes, "
    "Spearman correlations, melting temperatures, in vitro activities, oligos): not "
    "consumed",
)


def _paper(value: Any, quote: str, *, note: str | None = None) -> SourcedValue:
    """A value bound to a verbatim quote in the sha256-pinned Schastnaya ``paper.md``."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=PAPER_MD,
            citation_key=CITATION_KEY,
            sha256=PAPER_MD_SHA256,
            method="MinerU OCR of the publisher PDF (torchcell-library mirror)",
        ),
    )


def _si3(value: Any, quote: str, *, note: str | None = None) -> SourcedValue:
    """A value bound to the publisher's description of the Supplementary Data files."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=SI3_MD,
            citation_key=CITATION_KEY,
            sha256=SI3_MD_SHA256,
            method="MinerU OCR of the publisher PDF (torchcell-library mirror)",
            page="Description of Additional Supplementary Information File",
        ),
    )


def _baba(value: Any, quote: str, *, note: str | None = None) -> SourcedValue:
    """A value bound to Baba 2006, the collection paper Schastnaya defers to."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=PAPER_MD,
            citation_key=BABA_KEY,
            sha256=BABA_SHA256,
            method="MinerU OCR of the publisher PDF (torchcell-library mirror)",
        ),
    )


# --------------------------------------------------------------------------- #
# Sourced values (verbatim quotes, pinned sha256)
# --------------------------------------------------------------------------- #
_CULTURE_QUOTE = (
    "E. coli strains were grown in $1 \\mathrm { m L }$ M9 minimal medium supplemented "
    "with the same carbon sources used for growth analysis. Cultures were grown in deep "
    "96-well plates at $3 7 ^ { \\circ } \\mathrm { C }$ and $2 5 0 \\mathrm { r p m }$ "
    "until the $\\mathrm { O D } _ { 6 0 0 }$ of 0.4–1.5 was reached "
    "(mid-exponential phase)."
)
_MEDIA_QUOTE = (
    "E. coli cells were grown in M9 minimal medium supplemented with either $5 \\mathrm "
    "{ g / L }$ glucose, $5 \\mathrm { g / L }$ fructose, $6 . 8 \\mathrm { g / L }$ "
    "sodium acetate, $6 . 1 \\mathrm { g / L }$ sodium pyruvate or ${ 5 . 1 \\mathrm { g "
    "} } / { \\mathrm"
)
_LOG2FC_QUOTE = (
    "For every ion, the abundances from all replicates of a given mutant were pooled and "
    "compared to the pooled abundances of the wild-type sample. The $\\log _ { 2 }$ fold "
    "change of an ion abundance in the mutant compared to the wild-type was determined, "
    "and a two-sided 2-sample t-test with unequal variance was performed. The obtained "
    "$\\boldsymbol { p }$ -values were corrected for multiple testing using the "
    "Benjamini–Hoch"
)
_REPLICATE_QUOTE = (
    "$n = 1 0 ,$ , 3, 4, 5 biological replicates with two technical replicates were "
    "measured for the wild-type, knockout, S115E, S115A mutants, respectively."
)
_REPORTING_SUMMARY_NOTE = (
    "the Reporting Summary's 'Sample size' cell reads 'All experiments were performed "
    "at least in triplicates, and the sample sizes were increased whenever possible' "
    "and its 'Replication' cell 'Biological replicates were performed independently at "
    f"least three times' ({REPORTING_SUMMARY_PDF}, sha256 {REPORTING_SUMMARY_SHA256}; "
    "that form is a scan with no text layer, so the cells are only in the MinerU table "
    f"image {REPORTING_SUMMARY_TABLE_IMAGE}, sha256 {REPORTING_SUMMARY_TABLE_SHA256}, "
    "and are quoted here rather than audited)"
)

SOURCED_VALUES: dict[str, SourcedValue] = {
    "released_table": _si3(
        "Changes in metabolites for phosphomutant and knockout strains",
        "File name: Supplementary data 3\n\nDescription: Changes in metabolites for "
        "phosphomutant and knockout strains",
        note="what Supplementary Data 3 is, in the publisher's own words; the loader "
        "consumes it and nothing else for the values",
    ),
    "log2_fold_change": _paper(
        "log2 fold change of an ion abundance in the mutant over the wild type",
        _LOG2FC_QUOTE,
        note="the stored metabolite_level. Relative, unitless, and centered on the "
        "wild type by construction, which is why the reference level is 0",
    ),
    "ion_annotation": _paper(
        "deprotonated ions annotated by mass within 0.001 Da against an E. coli "
        "genome-scale model",
        "Deprotonated ions were annotated based on mass using 0.001 Da tolerance using a "
        "genome-wide reconstruction model of $E$ . coli metabolism41.",
        note="why one row can name several KEGG compounds of equal mass, and why the "
        "stored key is the neutral formula of the deprotonated ion",
    ),
    "ionization_mode": _paper(
        "negative",
        "Mass spectra were recorded in negative ionization mode within a mass/charge "
        "ratio range of $5 0 { - } 1 0 0 0 \\mathrm { m / z }$ using the highest "
        "resolving power (4 GHz HiRes) with an acquisition rate of 1.4 spectra per "
        "second.",
    ),
    "significance_cutoff": _paper(
        {"abs_log2_fold_change": 0.38, "adjusted_p_value": 0.05},
        "Ions with a $\\log _ { 2 }$ fold change $> \\pm 0 . 3 8 $ and a corrected "
        "$\\boldsymbol { p }$ - value $< 0 . 0 5$ were considered si",
        note="the paper's own call threshold. Every released value is stored, called "
        "or not; the threshold is recorded, never applied",
    ),
    "wild_type_strain": _paper(
        "MG1655 dmutS kan",
        "E. coli MG1655 ΔmutS with kanamycin resistance (referred to as wild-type) "
        "harboring the temperature-sensitive $\\lambda$ -Red recombineering plasmid "
        "$\\mathsf { p S I M } 5 ^ { 6 7 }$ with chloramphenicol resistance",
        note="the strain every fold change is taken against. pSIM5 was cured after "
        "MAGE, so the background carries the dmutS lesion and the kanamycin marker",
    ),
    "deletion_collection": _paper(
        "KEIO collection",
        "E. coli gene deletion mutants were retrieved from the KEIO collection38.",
        note="the stored collection string is this quote's own wording",
    ),
    "deletion_background": _baba(
        "BW25113",
        "The Keio collection is comprised of 3985 deletions in duplicate (7970 total) of "
        "E. coli K-12 strain BW25113 (Datsenko and Wanner, 2000)",
        note="the deferral target of Schastnaya's '(KEIO collection38)'. Schastnaya "
        "never names the deletion strains' background, so each released fold change "
        "compares a BW25113 deletion against an MG1655 wild type. The record pins "
        "MG1655, the strain the control and the released b-numbers are written "
        "against; the deletion leaf has no field for a perturbed strain's own "
        "background, so this crossing is recorded here and in preprocess/strains.csv",
    ),
    "cassette": _baba(
        "kanamycin cassette flanked by FLP recognition target sites",
        "Open-reading frame coding regions were replaced with a kanamycin cassette "
        "flanked by FLP recognition target sites",
        note="Schastnaya does not say the cassette was excised, so the collection "
        "strain (kanamycin resistant, cassette in place) is what is recorded",
    ),
    "media_recipe": _paper(
        "M9 minimal medium plus one carbon source",
        _MEDIA_QUOTE,
        note="the quote is cut at the OCR's glycerol markup. The salts of the M9 base "
        "are not stated anywhere in the paper, so each medium carries base_medium='M9' "
        "and exactly the one carbon source the paper weighs",
    ),
    "culture": _paper(
        "deep 96-well plates, 37 C, 250 rpm, harvested mid-exponentially",
        _CULTURE_QUOTE,
        note="the metabolome cultures; harvest is by growth phase, not by time, so "
        "duration_hours stays None",
    ),
    "temperature_c": _paper(37.0, _CULTURE_QUOTE),
    "aerobic_shaking": _paper(
        "aerobic",
        _CULTURE_QUOTE,
        note="shaken 1 mL cultures in deep-well plates; the paper names no other "
        "oxygen regime",
    ),
    "extraction_split": _paper(
        (
            "AcnB",
            "Adk",
            "AhpC",
            "GpmM",
            "KdsD",
            "ManX",
            "MetK",
            "Pck",
            "Pgi",
            "Pta",
            "TalB",
        ),
        "Metabolic extracts of AcnB, Adk, AhpC, GpmM, KdsD, ManX, MetK, Pck, Pgi, Pta, "
        "and TalB mutants were prepared using the hot extraction procedure, and for "
        "other mutants cold extraction was used.",
        note="measured on the release: every hot-extraction deletion column carries "
        "332 detected ions on glucose and 309 on acetate, every cold-extraction one "
        "284, and the build refuses any other split (detected_ion_counts.json)",
    ),
    "n_biological_replicates": _paper(
        3,
        _REPLICATE_QUOTE,
        note="Fig. 3b. The per-record count is not a released column and varies (10, "
        "3, 4, 5 for the wild type, the sucB knockout and two phosphomutants on "
        "fructose), so the CONSERVATIVE LOWER END is stored, never the optimistic one. "
        + _REPORTING_SUMMARY_NOTE,
    ),
    "technical_replicates": _paper(
        2,
        _REPLICATE_QUOTE,
        note="two injections of one extract are not independent replicates, so "
        "n_replicates counts biological replicates only",
    ),
    "screen_scope": _paper(
        {"phosphosites": 52, "enzymes": 23},
        "we selected 52 single phosphosites or multiple phosphosites in close proximity "
        "located on 23 enzymes, including transferases, isomerases, oxidoreductases, "
        "and lyases (Fig. 1a and Supplementa",
        note="the phosphomutant arm, which this loader cannot store",
    ),
    "mutation_method": _paper(
        "multiplex automated genome engineering (MAGE)",
        "Genomic point mutations of $E _ { * }$ coli phosphosites were constructed "
        "using MAGE34.",
        note="the dropped arm's perturbation is a genomic codon substitution; the "
        "schema has no bacterial leaf for one",
    ),
    "mutation_scheme": _paper(
        "S/T -> A and Y -> F abolish; S/T -> E mimics",
        "To abolish phosphorylation, phosphorylatable hydroxy groups were removed by "
        "mutating S and $\\mathrm { T }$ to alanine (A), and $\\mathrm { Y }$ to "
        "phenylalanine (F). To mimic phosphorylation, the negative charge of "
        "phosphorylation was imitated by subs",
    ),
    "data_availability": _paper(
        MASSIVE_ACCESSION,
        "The metabolomics data generated in this study have been deposited in the "
        "MassIVE database under accession code MSV000087795.",
        note="the raw spectra. The per-strain matrix the loader consumes is "
        "Supplementary Data 3, so MassIVE is recorded and not mirrored",
    ),
}

TEMPERATURE_C: float = SOURCED_VALUES["temperature_c"].value
N_BIOLOGICAL_REPLICATES: int = SOURCED_VALUES["n_biological_replicates"].value
COLLECTION: str = SOURCED_VALUES["deletion_collection"].value
CASSETTE: str = SOURCED_VALUES["cassette"].value
HOT_EXTRACTION_ENZYMES: tuple[str, ...] = SOURCED_VALUES["extraction_split"].value

MEASUREMENT_TYPE = "fia_tof_ms_ion_log2_fold_change_vs_wild_type"
REFERENCE_STRAIN_NAME: Final[EcoliK12StrainName] = "MG1655"
MG1655_ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_MG1655_ASM584v2"
if BACTERIAL_ASSEMBLY_SETS[REFERENCE_STRAIN_NAME] != MG1655_ASSEMBLY_SET:
    raise RuntimeError(
        f"MG1655's assembly set is "
        f"{BACTERIAL_ASSEMBLY_SETS[REFERENCE_STRAIN_NAME]!r}, not "
        f"{MG1655_ASSEMBLY_SET!r}"
    )
WILD_TYPE_BACKGROUND_NAME = "MG1655 dmutS kan"
KNOCKOUT_SUFFIX = " knockout"
NOT_DETECTED = "nd"
#: Checklist item 4: below this fraction of deletion gene symbols resolving to MG1655
#: b-numbers the build stops. Every one of the 16 resolves on the pinned annotation.
MIN_RESOLVED_FRACTION = 0.95

PUBLICATION = Publication(doi=PAPER_DOI, doi_url=f"https://doi.org/{PAPER_DOI}")
"""The mirrored article prints no PubMed id, so the DOI is the only identifier."""

METABOLITE_LEVEL_SE_GAP = ProvenanceGap(
    field="metabolite_level_se",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=Provenance(
        source_uri=PAPER_MD,
        citation_key=CITATION_KEY,
        sha256=PAPER_MD_SHA256,
        method="full Methods read, plus every column of Supplementary Data 3 "
        f"(sha256 {DATA_SHA256[SD3_FILE]}) and the file descriptions in {SI3_MD}",
        page="Methods, 'FIA TOF-MS measurement and untargeted metabolomics data "
        "processing'",
    ),
    note="the release carries one pooled log2 fold change per (strain, carbon source) "
    "and ion plus its Benjamini-Hochberg adjusted p-value, and no spread. The "
    "per-replicate abundances are on MassIVE as spectra, not as the normalized values "
    "this matrix was pooled from, so no standard error can be recovered. The adjusted "
    "p-value is NOT a standard error and is written to preprocess/adjusted_p_values.csv "
    "instead of this field",
)


# --------------------------------------------------------------------------- #
# The five media (one carbon source each)
#
# MEDIA_LIBRARY has no entry for Schastnaya's M9 and ``media.py`` is a value surface, so
# each medium is built here with ``base_medium="M9"`` (which resolves in the library, so
# the record still joins the M9 family). The paper states no salt recipe, so the one
# stated component is the carbon source.
# --------------------------------------------------------------------------- #
CARBON_SOURCES: dict[str, tuple[str, float]] = {
    "GLUCOSE": ("D-glucose", 5.0),
    "FRUCTOSE": ("D-fructose", 5.0),
    "ACETATE": ("sodium acetate", 6.8),
    "PYRUVATE": ("sodium pyruvate", 6.1),
    "GLYCEROL": ("glycerol", 5.1),
}
"""Each released carbon-source token, its compound and the paper's g/L."""


def carbon_source_media(token: str) -> Media:
    """The M9 medium of one released carbon-source token."""
    compound_name, grams_per_litre = CARBON_SOURCES[token]
    return Media(
        name=f"M9 minimal medium (Schastnaya 2021) + {grams_per_litre:g} g/L "
        f"{compound_name}",
        state="liquid",
        is_synthetic=True,
        base_medium="M9",
        components=[
            MediaComponent(
                compound=resolved_compound(compound_name),
                role=MediaComponentRole.carbon_source,
                concentration=Concentration(
                    value=grams_per_litre, unit=ConcentrationUnit.g_per_l
                ),
                provenance=[SOURCED_VALUES["media_recipe"]],
                note="the five carbon sources are weighed 'to keep C-atoms constant "
                "between the conditions'",
            )
        ],
        provenance=[SOURCED_VALUES["media_recipe"], SOURCED_VALUES["culture"]],
    )


MEDIA: dict[str, Media] = {
    token: carbon_source_media(token) for token in CARBON_SOURCES
}
"""One medium per released carbon-source token, built once."""


def environment(token: str) -> Environment:
    """The metabolome culture of one carbon source: M9 + that source, 37 C, shaken."""
    return Environment(
        media=MEDIA[token],
        temperature=Temperature(value=TEMPERATURE_C),
        aerobicity=SOURCED_VALUES["aerobic_shaking"].value,
    )


# --------------------------------------------------------------------------- #
# Raw mirror (the loader reads the mirror, never a live URL)
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the environment (the mirrors and build tree live under it)."""
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/schastnayaExtensiveRegulationEnzyme2021``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def library_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-library/schastnayaExtensiveRegulationEnzyme2021``."""
    return Path(data_root or _data_root()) / LIBRARY_DIR_REL


def _sha256(path: str | Path) -> str:
    """Streaming sha256 of a file."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def retrieve_raw_files(
    dest_dir: str | Path, names: Sequence[str] | None = None
) -> dict[str, Path]:
    """Run each file's recorded retriever and write the verified bytes to ``dest_dir``.

    The recorded ``RetrievalRecord`` is what runs (``run_retriever``), so this is the
    re-runnable retrieval itself; a byte mismatch raises before anything is written.
    """
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}
    for raw in RAW_FILES:
        if names is not None and raw.name not in names:
            continue
        path = dest / raw.name
        write_verified(
            run_retriever(raw.retrieval),
            path,
            raw.sha256,
            raw.retrieval.source_url or raw.name,
        )
        out[raw.name] = path
    return out


def deposit_raw_mirror(
    *, sources: Mapping[str, str | Path], data_root: str | None = None
) -> Path:
    """Write the raw mirror from already-retrieved files plus its ``manifest.json``.

    ``sources`` maps every name in ``RAW_FILES`` to a local file. Idempotent by sha256:
    a mirror file with the pinned hash is left alone, and one with any other hash
    raises rather than being overwritten.
    """
    missing = sorted(set(DATA_SHA256) - set(sources))
    if missing:
        raise KeyError(f"no source given for {missing}")
    root = raw_mirror_dir(data_root)
    records: list[ArtifactRecord] = []
    for raw in RAW_FILES:
        src = Path(sources[raw.name])
        got = _sha256(src)
        if got != raw.sha256:
            raise RuntimeError(
                f"{src} sha256 mismatch: got {got}, expected {raw.sha256}"
            )
        dest = root / raw.mirror_relpath
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            if _sha256(dest) != raw.sha256:
                raise RuntimeError(f"{dest} exists with a different sha256; refusing")
        else:
            shutil.copy2(src, dest)
        records.append(
            ArtifactRecord(
                path=raw.mirror_relpath,
                role=ROLE_RAW_DATA,
                bytes=dest.stat().st_size,
                sha256=raw.sha256,
                source=raw.retrieval.source_url,
                retrieval=raw.retrieval,
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title="Extensive regulation of enzyme activity by phosphorylation in "
        "Escherichia coli",
        files=records,
        si_data_sources=[
            f"https://pmc-oa-opendata.s3.amazonaws.com/{PMCID}.1/",
            f"https://doi.org/{PAPER_DOI}",
        ],
        si_expected=list(NOT_MIRRORED),
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


def sourced_value_root(value: SourcedValue, data_root: str | None = None) -> Path:
    """The mirror root a ``SourcedValue``'s ``source_uri`` is relative to.

    Every quote of this loader comes from a paper in the literature mirror.
    """
    return Path(data_root or _data_root()) / "torchcell-library"


# --------------------------------------------------------------------------- #
# Parsing the two released workbooks
# --------------------------------------------------------------------------- #
#: Supplementary Data 3's fixed geometry: a title row, a header row, two label rows and
#: then the ions; four label columns and then two equal blocks of strain columns.
SD3_TITLE = "Supplementary Data 3 | Changes in metabolites for phosphomutant and knockout strains"
SD3_LABEL_COLUMNS = ("Kegg ID", "Annotation", "Formula", "m/z")
SD3_FOLD_CHANGE_HEADER = "LOG2(FC)"
SD3_P_VALUE_HEADER = "Adjusted p-value"
SD3_N_LABEL_COLUMNS = 4
SD3_FIRST_DATA_ROW = 5
#: Supplementary Data 1's columns, of which the loader reads three.
SD1_TITLE = "Supplementary Data 1 | Phosphomutant strains"
SD1_STRAIN_COLUMN = 0
SD1_STATE_COLUMN = 1
SD1_LOCUS_COLUMN = 4


class TableFormatError(ValueError):
    """A released workbook does not have the geometry the loader pins."""


class IonRow(BaseModel):
    """One annotated ion of Supplementary Data 3, with its candidate identities."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    row: int
    formula: str
    mz: float
    annotation: str
    kegg_ids: tuple[str, ...]

    @property
    def n_candidates(self) -> int:
        """KEGG compounds this ion's mass cannot separate."""
        return len(self.kegg_ids)

    @property
    def target_metabolite_id(self) -> str | None:
        """The one KEGG id, or ``None`` when the row is a merged isobaric set."""
        return self.kegg_ids[0] if self.n_candidates == 1 else None


class StrainColumn(BaseModel):
    """One released (strain, carbon source) column: its values and their p-values."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    column: int
    strain: str
    carbon_source: str
    fold_change: tuple[float | None, ...]
    adjusted_p_value: tuple[float | None, ...]

    @property
    def is_knockout(self) -> bool:
        """True when the released token names a gene deletion, not a phosphomutant."""
        return self.strain.endswith(KNOCKOUT_SUFFIX)

    @property
    def gene_symbol(self) -> str:
        """The deleted gene's symbol; only a knockout token has one."""
        if not self.is_knockout:
            raise ValueError(f"{self.strain!r} is not a knockout column")
        return self.strain[: -len(KNOCKOUT_SUFFIX)].strip()

    @property
    def detected_rows(self) -> tuple[int, ...]:
        """Indices of the ions this column detected (the rest are ``nd``)."""
        return tuple(i for i, v in enumerate(self.fold_change) if v is not None)

    @property
    def n_detected(self) -> int:
        """Ions this column detected."""
        return len(self.detected_rows)

    @property
    def extraction(self) -> str:
        """``hot`` or ``cold``, from the paper's own list of hot-extraction enzymes."""
        hot = {name.lower() for name in HOT_EXTRACTION_ENZYMES}
        return "hot" if self.gene_symbol.lower() in hot else "cold"


class Sd3Table(BaseModel):
    """Supplementary Data 3 as parsed: its ions and its strain columns."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    ions: tuple[IonRow, ...]
    columns: tuple[StrainColumn, ...]


class PhosphomutantEntry(BaseModel):
    """One Supplementary Data 1 row: a phosphomutant strain and what it changes."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    strain: str
    mutant_state: str
    enzyme_locus: str


def _cell_float(value: Any, where: str) -> float | None:
    """A numeric cell, or ``None`` for the release's ``nd``; anything else raises."""
    if isinstance(value, str):
        if value.strip() == NOT_DETECTED:
            return None
        raise TableFormatError(f"{where}: {value!r} is neither a number nor 'nd'")
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise TableFormatError(f"{where}: {value!r} is not a number")
    return float(value)


def read_sd3(path: str | Path) -> Sd3Table:
    """Parse Supplementary Data 3, refusing any geometry other than the pinned one.

    The sheet is a title row, a header row naming the two blocks, two label rows
    (carbon source, strain) and then one row per ion. The fold-change block and the
    adjusted-p-value block must carry the SAME (strain, carbon source) pairs in the
    SAME order, which is what lets one column object hold both.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [SD3_SHEET]:
            raise TableFormatError(f"{path}: sheets {workbook.sheetnames}")
        rows = [tuple(row) for row in workbook[SD3_SHEET].iter_rows(values_only=True)]
    finally:
        workbook.close()
    if rows[0][0] != SD3_TITLE:
        raise TableFormatError(f"{path}: title {rows[0][0]!r}")
    header = rows[1]
    if tuple(header[:SD3_N_LABEL_COLUMNS]) != SD3_LABEL_COLUMNS:
        raise TableFormatError(f"{path}: label columns {header[:4]!r}")
    if header[SD3_N_LABEL_COLUMNS] != SD3_FOLD_CHANGE_HEADER:
        raise TableFormatError(f"{path}: block 1 header {header[4]!r}")
    block_starts = [i for i, value in enumerate(header) if value == SD3_P_VALUE_HEADER]
    if len(block_starts) != 1:
        raise TableFormatError(f"{path}: {len(block_starts)} p-value block headers")
    p_start = block_starts[0]
    n_columns = p_start - SD3_N_LABEL_COLUMNS
    carbon = rows[2]
    strains = rows[3]
    fold_pairs = [(strains[j], carbon[j]) for j in range(SD3_N_LABEL_COLUMNS, p_start)]
    p_pairs = [(strains[j], carbon[j]) for j in range(p_start, p_start + n_columns)]
    if fold_pairs != p_pairs:
        raise TableFormatError(f"{path}: the two blocks name different columns")
    unknown = sorted({c for _, c in fold_pairs} - set(CARBON_SOURCES))
    if unknown:
        raise TableFormatError(f"{path}: carbon sources {unknown} are not the five")

    ions: list[IonRow] = []
    data: list[tuple[Any, ...]] = [
        row for row in rows[SD3_FIRST_DATA_ROW - 1 :] if any(v is not None for v in row)
    ]
    for index, row in enumerate(data, start=1):
        kegg, annotation, formula, mz = row[:SD3_N_LABEL_COLUMNS]
        if not isinstance(formula, str) or not formula.strip():
            raise TableFormatError(f"SD3 data row {index}: formula {formula!r}")
        if not isinstance(kegg, str) or not isinstance(annotation, str):
            raise TableFormatError(f"SD3 data row {index}: kegg/annotation {kegg!r}")
        if isinstance(mz, bool) or not isinstance(mz, (int, float)):
            raise TableFormatError(f"SD3 data row {index}: m/z {mz!r}")
        ions.append(
            IonRow(
                row=index,
                formula=formula.strip(),
                mz=float(mz),
                annotation=annotation.strip(),
                kegg_ids=tuple(part.strip() for part in kegg.split(";")),
            )
        )
    formulas = [ion.formula for ion in ions]
    if len(set(formulas)) != len(formulas):
        raise TableFormatError(f"{path}: the Formula column is not unique")

    columns: list[StrainColumn] = []
    for offset in range(n_columns):
        fold = SD3_N_LABEL_COLUMNS + offset
        pvalue = p_start + offset
        strain, carbon_source = fold_pairs[offset]
        if not isinstance(strain, str) or not strain.strip():
            raise TableFormatError(f"SD3 column {fold + 1}: strain {strain!r}")
        values = tuple(
            _cell_float(row[fold], f"SD3 row {i} column {fold + 1}")
            for i, row in enumerate(data, start=1)
        )
        pvalues = tuple(
            _cell_float(row[pvalue], f"SD3 row {i} column {pvalue + 1}")
            for i, row in enumerate(data, start=1)
        )
        missing = [
            i
            for i, (v, p) in enumerate(zip(values, pvalues, strict=True))
            if (v is None) != (p is None)
        ]
        if missing:
            raise TableFormatError(
                f"SD3 column {fold + 1} ({strain}): {len(missing)} rows carry a value "
                "without its p-value, or the other way round"
            )
        columns.append(
            StrainColumn(
                column=fold + 1,
                strain=strain.strip(),
                carbon_source=str(carbon_source),
                fold_change=values,
                adjusted_p_value=pvalues,
            )
        )
    return Sd3Table(ions=tuple(ions), columns=tuple(columns))


def read_sd1(path: str | Path) -> dict[str, PhosphomutantEntry]:
    """Parse Supplementary Data 1 into ``{strain token: entry}``."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [SD1_SHEET]:
            raise TableFormatError(f"{path}: sheets {workbook.sheetnames}")
        rows = [tuple(row) for row in workbook[SD1_SHEET].iter_rows(values_only=True)]
    finally:
        workbook.close()
    if rows[0][0] != SD1_TITLE:
        raise TableFormatError(f"{path}: title {rows[0][0]!r}")
    entries: dict[str, PhosphomutantEntry] = {}
    for index, row in enumerate(rows[2:], start=1):
        if row[SD1_STRAIN_COLUMN] is None:
            continue
        strain, state, locus = (
            row[SD1_STRAIN_COLUMN],
            row[SD1_STATE_COLUMN],
            row[SD1_LOCUS_COLUMN],
        )
        if not all(isinstance(v, str) and v.strip() for v in (strain, state, locus)):
            raise TableFormatError(f"SD1 row {index}: {strain!r} {state!r} {locus!r}")
        entry = PhosphomutantEntry(
            strain=str(strain).strip(),
            mutant_state=str(state).strip(),
            enzyme_locus=str(locus).strip(),
        )
        if entry.strain in entries:
            raise TableFormatError(f"SD1 names {entry.strain!r} twice")
        entries[entry.strain] = entry
    return entries


# --------------------------------------------------------------------------- #
# Retention: the deletion columns are kept, the phosphomutant columns are not
# --------------------------------------------------------------------------- #
PHOSPHOMUTANT_DROP_RULE = "phosphosite_substitution_has_no_bacterial_leaf"
PHOSPHOMUTANT_DROP_DESCRIPTION = (
    "the record's perturbation is a genomic codon substitution at a phosphosite, made "
    "by MAGE, and the schema has no bacterial sequence-level perturbation leaf: "
    "AllelePerturbation is the amino-acid-substitution leaf but inherits "
    "GenePerturbation's R64 systematic-name validator, which refuses a b-number, and "
    "the five bacterial leaves (bacterial_deletion, transposon_insertion, "
    "bacterial_crispr_interference, promoter_replacement, heterologous_pathway) each "
    "assert a genotype this strain does not have. Storing one would state a "
    "perturbation the paper did not make, so the record is dropped and the missing "
    "leaf is the finding"
)


class DropRule(BaseModel):
    """One retention rule, the released columns it removed, and why."""

    rule: str
    description: str
    n_records: int
    items: list[str] = []


class DropLog(BaseModel):
    """Every retention rule applied to a build, in the order they were applied."""

    dataset: str
    released_columns: int
    source_records: int
    kept_records: int
    dropped_records: int
    rules: list[DropRule]


def phosphomutant_drop_rule(
    dropped: Sequence[StrainColumn], entries: Mapping[str, PhosphomutantEntry]
) -> DropRule:
    """The dropped phosphomutant columns, one ledger item per distinct strain.

    Each item names the strain's enzyme locus and whether the substitution abolishes or
    mimics phosphorylation, so the blocked arm is legible without the workbook.
    """
    by_strain: dict[str, list[str]] = {}
    for column in dropped:
        by_strain.setdefault(column.strain, []).append(column.carbon_source)
    items: list[str] = []
    for strain in sorted(by_strain):
        entry = entries.get(strain)
        if entry is None:
            raise TableFormatError(
                f"Supplementary Data 3 column {strain!r} is in no Supplementary Data 1 "
                "row, so its enzyme and mutant state cannot be named"
            )
        items.append(
            f"{strain} ({entry.enzyme_locus}, {entry.mutant_state}): "
            + ", ".join(sorted(by_strain[strain]))
        )
    return DropRule(
        rule=PHOSPHOMUTANT_DROP_RULE,
        description=PHOSPHOMUTANT_DROP_DESCRIPTION,
        n_records=len(dropped),
        items=items,
    )


class DetectedIonLedger(BaseModel):
    """Which ions each kept group detected, and that it depends only on the group.

    Measured on the release: among the deletion columns the detected set is a property
    of the (extraction protocol, carbon source) pair, not of the strain, which is what
    the paper's hot/cold extraction split predicts.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    n_ions: int
    groups: dict[str, int]
    n_strains: dict[str, int]
    cold_set_shared_across_carbon_sources: bool


def detected_ion_ledger(kept: Sequence[StrainColumn]) -> DetectedIonLedger:
    """Group the kept columns by (extraction, carbon source), refusing a split by strain."""
    masks: dict[str, set[tuple[int, ...]]] = {}
    counts: dict[str, int] = {}
    for column in kept:
        key = f"{column.extraction}|{column.carbon_source}"
        masks.setdefault(key, set()).add(column.detected_rows)
        counts[key] = counts.get(key, 0) + 1
    mixed = sorted(key for key, found in masks.items() if len(found) > 1)
    if mixed:
        raise TableFormatError(
            f"the detected-ion set differs between strains of {mixed}, so it is not a "
            "property of the extraction protocol and the carbon source"
        )
    cold = {
        next(iter(found)) for key, found in masks.items() if key.startswith("cold|")
    }
    n_ions = len(kept[0].fold_change) if kept else 0
    return DetectedIonLedger(
        n_ions=n_ions,
        groups={key: len(next(iter(found))) for key, found in sorted(masks.items())},
        n_strains=dict(sorted(counts.items())),
        cold_set_shared_across_carbon_sources=len(cold) == 1,
    )


class IdentityLedger(BaseModel):
    """The metabolite-identity status of the ions the loader stores."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    n_ions: int
    n_single_identity: int
    n_merged_isobaric: int
    candidate_size_histogram: dict[int, int]
    merged_candidates: dict[str, tuple[str, ...]]

    @property
    def target_metabolite_ids_covered(self) -> float:
        """Share of ions carrying exactly one KEGG id."""
        return self.n_single_identity / self.n_ions


def identity_ledger(ions: Sequence[IonRow]) -> IdentityLedger:
    """Count how many ions name one KEGG compound and list every merged candidate set."""
    sizes: dict[int, int] = {}
    merged: dict[str, tuple[str, ...]] = {}
    for ion in ions:
        sizes[ion.n_candidates] = sizes.get(ion.n_candidates, 0) + 1
        if ion.target_metabolite_id is None:
            merged[ion.formula] = ion.kegg_ids
    return IdentityLedger(
        n_ions=len(ions),
        n_single_identity=len(ions) - len(merged),
        n_merged_isobaric=len(merged),
        candidate_size_histogram=dict(sorted(sizes.items())),
        merged_candidates=dict(sorted(merged.items())),
    )


class ResolvedDeletion(BaseModel):
    """One kept column with the MG1655 locus tag of its deleted gene."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    column: StrainColumn
    locus_tag: str


def resolve_deletions(
    genome: EcoliK12Genome, kept: Sequence[StrainColumn], *, label: str
) -> tuple[list[ResolvedDeletion], LocusTagReconciliation]:
    """Reconcile each deletion's gene symbol on ``genome``, refusing an unplaced one.

    Every one of the released symbols is a current MG1655 gene, so a symbol the
    reconciler keeps as given is a changed annotation or the wrong genome, not a
    record to drop.
    """
    symbols = pd.Series([column.gene_symbol for column in kept], dtype=object)
    stored, report = reconcile_locus_tags(genome, symbols, label=label)
    report.require_resolved(MIN_RESOLVED_FRACTION)
    pattern = LOCUS_TAG_PATTERNS[report.gene_namespace]
    resolved: list[ResolvedDeletion] = []
    unplaced: list[str] = []
    for column, tag in zip(kept, stored.tolist(), strict=True):
        if pattern.match(str(tag)) is None:
            unplaced.append(f"{column.gene_symbol} -> {tag}")
            continue
        resolved.append(ResolvedDeletion(column=column, locus_tag=str(tag)))
    if unplaced:
        raise TableFormatError(
            f"{len(unplaced)} deletion symbols are not on an MG1655 locus: {unplaced}"
        )
    return resolved, report


# --------------------------------------------------------------------------- #
# Record construction (pure, no files)
# --------------------------------------------------------------------------- #
def wild_type_background() -> BacterialStrainBackground:
    """The strain every released fold change is taken against."""
    return BacterialStrainBackground(
        name=WILD_TYPE_BACKGROUND_NAME,
        reference_strain=REFERENCE_STRAIN_NAME,
        assembly_set=MG1655_ASSEMBLY_SET,
        parents=["MG1655"],
        construction=None,
        genotype_statement="MG1655 ΔmutS with kanamycin resistance",
        alleles=[],
        provenance=[SOURCED_VALUES["wild_type_strain"]],
    )


def metabolite_phenotype(
    ions: Sequence[IonRow], values: Sequence[float | None]
) -> MetabolitePhenotype:
    """One column's detected ions: their log2 fold changes and their KEGG identities.

    An ``nd`` ion is absent from the record, because the release states no value for
    it. ``target_metabolite_ids`` carries only the ions naming exactly one KEGG
    compound; a merged isobaric set is left out rather than assigned a candidate.
    """
    if len(ions) != len(values):
        raise ValueError(f"{len(ions)} ions for {len(values)} values")
    levels: dict[str, float] = {}
    targets: dict[str, str] = {}
    for ion, value in zip(ions, values, strict=True):
        if value is None:
            continue
        levels[ion.formula] = value
        target = ion.target_metabolite_id
        if target is not None:
            targets[ion.formula] = target
    if not levels:
        raise ValueError("a stored column must detect at least one ion")
    return MetabolitePhenotype(
        metabolite_level=levels,
        metabolite_level_se=None,
        n_replicates=dict.fromkeys(levels, N_BIOLOGICAL_REPLICATES),
        measurement_type=MEASUREMENT_TYPE,
        target_metabolite_ids=targets,
        provenance_gaps=[METABOLITE_LEVEL_SE_GAP],
    )


def wild_type_phenotype(phenotype: MetabolitePhenotype) -> MetabolitePhenotype:
    """The reference profile of one record: 0 for every key the record measured.

    Not a measured profile. The released statistic is the log2 fold change AGAINST this
    wild type, so its own value is 0 by construction, which is also what the metabolite
    family's ``reference_zero`` check asserts.
    """
    return MetabolitePhenotype(
        metabolite_level=dict.fromkeys(phenotype.metabolite_level, 0.0),
        metabolite_level_se=None,
        n_replicates=dict.fromkeys(phenotype.metabolite_level, N_BIOLOGICAL_REPLICATES),
        measurement_type=MEASUREMENT_TYPE,
        target_metabolite_ids=dict(phenotype.target_metabolite_ids or {}),
        provenance_gaps=[METABOLITE_LEVEL_SE_GAP],
    )


def deletion_genotype(resolved: ResolvedDeletion) -> Genotype:
    """The one KEIO deletion, named by its MG1655 b-number."""
    return Genotype(
        perturbations=[
            BacterialDeletionPerturbation(
                systematic_gene_name=resolved.locus_tag,
                perturbed_gene_name=resolved.column.gene_symbol,
                gene_namespace=STRAIN_GENE_NAMESPACES[REFERENCE_STRAIN_NAME],
                collection=COLLECTION,
                cassette=CASSETTE,
            )
        ]
    )


def build_experiment(
    dataset_name: str, resolved: ResolvedDeletion, ions: Sequence[IonRow]
) -> BacterialMetaboliteExperiment:
    """The record of one deletion strain in one carbon source."""
    return BacterialMetaboliteExperiment(
        dataset_name=dataset_name,
        genotype=deletion_genotype(resolved),
        environment=environment(resolved.column.carbon_source),
        phenotype=metabolite_phenotype(ions, resolved.column.fold_change),
    )


def build_reference(
    dataset_name: str,
    genome_reference: AssemblyReferenceGenome,
    experiment: BacterialMetaboliteExperiment,
) -> BacterialMetaboliteExperimentReference:
    """The wild type in the same medium, on the same key set, at 0."""
    return BacterialMetaboliteExperimentReference(
        dataset_name=dataset_name,
        genome_reference=genome_reference,
        environment_reference=experiment.environment.model_copy(),
        phenotype_reference=wild_type_phenotype(experiment.phenotype),
    )


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class MetabolomeSchastnaya2021Dataset(ExperimentDataset):
    """FIA-TOF-MS metabolome of the KEIO deletion arm of an E. coli phosphosite screen."""

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = REFERENCE_STRAIN_NAME

    def __init__(
        self,
        root: str = "data/torchcell/metabolome_schastnaya2021",
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialMetaboliteExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialMetaboliteExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """Every consumed file, linked from the raw mirror."""
        return [raw.name for raw in RAW_FILES]

    def download(self) -> None:
        """Link each mirror file into ``raw/`` after checking the manifest and sha256.

        The mirror plus ``DATA_SHA256`` is canonical; the PMC URLs are retrieval
        metadata that ``retrieve_raw_files`` re-runs, never a build input.
        """
        data_root = _data_root()
        manifest = load_manifest(data_root)
        os.makedirs(self.raw_dir, exist_ok=True)
        for raw in RAW_FILES:
            check_manifest_pin(
                raw.mirror_relpath,
                manifest_sha256(manifest, raw.mirror_relpath),
                raw.sha256,
            )
            src = raw_mirror_dir(data_root) / raw.mirror_relpath
            if not src.exists():
                raise RuntimeError(f"required raw artifact missing from mirror: {src}")
            link_verified(src, osp.join(self.raw_dir, raw.name), raw.sha256)
        log.info(
            "Schastnaya 2021 raw files linked into %s (sha256 verified)", self.raw_dir
        )

    def _raw(self, name: str) -> str:
        return osp.join(self.raw_dir, name)

    def _genome(self) -> EcoliK12Genome:
        """The injected genome, or the reference strain's default cache (a direct run);
        a genome of another assembly set is refused.
        """
        if self.ecoli_genome is None:
            self.ecoli_genome = bacterial_genome("ecoli", self.REFERENCE_STRAIN)
        expected = BACTERIAL_ASSEMBLY_SETS[self.REFERENCE_STRAIN]
        if self.ecoli_genome.ASSEMBLY_SET != expected:
            raise ValueError(
                f"{type(self).__name__} needs the {expected} genome, got "
                f"{self.ecoli_genome.ASSEMBLY_SET}"
            )
        return self.ecoli_genome

    @post_process
    def process(self) -> None:
        """Parse Supplementary Data 3 into one record per kept deletion column + LMDB."""
        verify_raw_files(self.raw_dir, DATA_SHA256)
        table = read_sd3(self._raw(SD3_FILE))
        entries = read_sd1(self._raw(SD1_FILE))

        kept_columns = [c for c in table.columns if c.is_knockout]
        dropped_columns = [c for c in table.columns if not c.is_knockout]
        drop_rule = phosphomutant_drop_rule(dropped_columns, entries)
        detected = detected_ion_ledger(kept_columns)
        identities = identity_ledger(table.ions)

        genome = self._genome()
        resolved, reconciliation = resolve_deletions(
            genome, kept_columns, label=f"{self.name} deletion gene symbols"
        )
        drop_log = DropLog(
            dataset=self.name,
            released_columns=len(table.columns),
            source_records=len(table.columns),
            kept_records=len(resolved),
            dropped_records=len(table.columns) - len(resolved),
            rules=[drop_rule],
        )
        if sum(rule.n_records for rule in drop_log.rules) != drop_log.dropped_records:
            raise RuntimeError("drop rules do not account for every dropped column")
        log.info(
            "Schastnaya 2021: %d released columns -> %d records; %d phosphomutant "
            "columns dropped for the missing bacterial substitution leaf; %d of %d "
            "ions name one KEGG compound; detected ions per group %s",
            len(table.columns),
            len(resolved),
            drop_rule.n_records,
            identities.n_single_identity,
            identities.n_ions,
            detected.groups,
        )

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        self._write_ledgers(
            drop_log, reconciliation, detected, identities, table.ions, resolved
        )

        genome_reference = assembly_reference(
            self.REFERENCE_STRAIN, background=wild_type_background()
        )
        env_out, interned_env = self._open_write_lmdb(
            osp.join(self.processed_dir, "lmdb")
        )
        with env_out.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for idx, item in enumerate(tqdm(resolved, desc="schastnaya2021")):
                experiment = build_experiment(self.name, item, table.ions)
                reference = build_reference(self.name, genome_reference, experiment)
                txn.put(
                    f"{idx}".encode(),
                    self._intern_record(experiment, reference, PUBLICATION, itxn),
                )
        env_out.close()
        interned_env.close()
        log.info(
            "Wrote %d Schastnaya 2021 metabolome experiments to LMDB", len(resolved)
        )

    def _write_ledgers(
        self,
        drop_log: DropLog,
        reconciliation: LocusTagReconciliation,
        detected: DetectedIonLedger,
        identities: IdentityLedger,
        ions: Sequence[IonRow],
        resolved: Sequence[ResolvedDeletion],
    ) -> None:
        """The drop log, the reconciliation, the ion tables and the p-value matrix."""
        out = Path(self.preprocess_dir)
        (out / "dropped_records.json").write_text(drop_log.model_dump_json(indent=2))
        (out / "identifier_reconciliation.json").write_text(
            reconciliation.model_dump_json(indent=2)
        )
        (out / "detected_ion_counts.json").write_text(
            detected.model_dump_json(indent=2)
        )
        (out / "metabolite_identity.json").write_text(
            json.dumps(
                identities.model_dump()
                | {
                    "target_metabolite_ids_covered": identities.target_metabolite_ids_covered
                },
                indent=2,
            )
        )
        pd.DataFrame(
            [
                {
                    "row": ion.row,
                    "formula": ion.formula,
                    "mz": ion.mz,
                    "n_candidates": ion.n_candidates,
                    "kegg_ids": "; ".join(ion.kegg_ids),
                    "target_metabolite_id": ion.target_metabolite_id or "",
                    "annotation": ion.annotation,
                }
                for ion in ions
            ]
        ).to_csv(out / "ions.csv", index=False)
        pd.DataFrame(
            [
                {
                    "record": idx,
                    "released_column": item.column.column,
                    "released_strain": item.column.strain,
                    "gene_symbol": item.column.gene_symbol,
                    "locus_tag": item.locus_tag,
                    "carbon_source": item.column.carbon_source,
                    "extraction": item.column.extraction,
                    "n_detected": item.column.n_detected,
                    "deletion_background": SOURCED_VALUES["deletion_background"].value,
                    "reference_background": WILD_TYPE_BACKGROUND_NAME,
                }
                for idx, item in enumerate(resolved)
            ]
        ).to_csv(out / "strains.csv", index=False)
        p_values = pd.DataFrame(
            {
                f"{item.column.strain}|{item.column.carbon_source}": (
                    item.column.adjusted_p_value
                )
                for item in resolved
            },
            index=[ion.formula for ion in ions],
        )
        p_values.index.name = "formula"
        p_values.to_csv(out / "adjusted_p_values.csv")

    def preprocess_raw(
        self, df: pd.DataFrame, preprocess: dict[str, Any] | None = None
    ) -> pd.DataFrame:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled by ``build_experiment``."""
        raise NotImplementedError


# --------------------------------------------------------------------------- #
# Verification (L0-L4) of a built LMDB
#
# The metabolite family verifier keys L1 uniqueness on the GENOTYPE alone, and this
# dataset is one record per (genotype, carbon source), so its rows are composed here
# instead: the same L0/L1-count/L2/L3 checks with a (deletion set, carbon source) key.
# --------------------------------------------------------------------------- #
VERIFIER_PROVENANCE = Provenance(
    source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/data/{SD3_FILE}",
    citation_key=CITATION_KEY,
    sha256=DATA_SHA256[SD3_FILE],
    method="Supplementary Data 3: one BacterialMetaboliteExperiment per (KEIO deletion "
    "strain, carbon source), log2 fold change of each annotated ion against the "
    "MG1655 dmutS wild type; the 171 phosphomutant columns are not stored",
    page="Nat Commun 12:5650, Supplementary Data 3",
)


def _genotype_environment_key(record: Mapping[str, Any]) -> tuple[Any, ...]:
    """The strain and its carbon source: this dataset's record identity."""
    experiment = record["experiment"]
    perturbations = tuple(
        sorted(
            (p["systematic_gene_name"], p["perturbation_type"])
            for p in experiment["genotype"]["perturbations"]
        )
    )
    return (perturbations, experiment["environment"]["media"]["name"])


def _l1_strain_condition_uniqueness(
    records: Sequence[Mapping[str, Any]],
) -> LevelResult:
    """L1: one record per (deletion set, medium)."""
    seen: dict[tuple[Any, ...], int] = {}
    for record in records:
        key = _genotype_environment_key(record)
        seen[key] = seen.get(key, 0) + 1
    duplicated = sorted(str(k) for k, n in seen.items() if n > 1)
    return LevelResult(
        level=Level.L1,
        name="genotype_environment_uniqueness",
        passed=not duplicated,
        message=(
            f"{len(seen)} unique (deletion set, medium) pairs, one record each"
            if not duplicated
            else f"{len(duplicated)} pairs appear in more than one record"
        ),
        details={"n_pairs": len(seen), "n_duplicated": len(duplicated)},
    )


def _l3_reference_zero_on_the_same_keys(
    records: Sequence[Mapping[str, Any]],
) -> LevelResult:
    """L3: the reference level is 0 on exactly the keys the record measured."""
    worst = 0.0
    n = 0
    mismatched = 0
    for record in records:
        measured = record["experiment"]["phenotype"]["metabolite_level"]
        reference = record["reference"]["phenotype_reference"]["metabolite_level"]
        if set(reference) != set(measured):
            mismatched += 1
        for value in reference.values():
            n += 1
            worst = max(worst, abs(float(value)))
    passed = worst == 0.0 and mismatched == 0
    return LevelResult(
        level=Level.L3,
        name="reference_zero",
        passed=passed,
        message=(
            f"reference level == 0 on the record's own keys for all {n} values"
            if passed
            else f"max|reference|={worst:.3g}, {mismatched} records key-mismatched"
        ),
        details={"n_values": n, "worst_abs": worst, "n_key_mismatched": mismatched},
    )


def _l3_measurement_type(records: Sequence[Mapping[str, Any]]) -> LevelResult:
    """L3: one measurement_type over the dataset, and the reference shares it."""
    types = {r["experiment"]["phenotype"]["measurement_type"] for r in records}
    types |= {
        r["reference"]["phenotype_reference"]["measurement_type"] for r in records
    }
    passed = types == {MEASUREMENT_TYPE}
    return LevelResult(
        level=Level.L3,
        name="measurement_type_consistent",
        passed=passed,
        message=(
            f"single measurement_type: {MEASUREMENT_TYPE!r}"
            if passed
            else f"measurement types {sorted(types)}"
        ),
        details={"measurement_types": sorted(types)},
    )


def verify_build(
    dataset_root: str, *, data_root: str | None = None
) -> VerificationReport:
    """Run L0-L4 on a built tree and write ``preprocess/verification_report.json``.

    L4 is containment of every deleted b-number in the MG1655 GenBank gene rows the
    records' own ``genome_reference`` pins, read through the verification runners'
    assembly resolver, plus the provenance audit of every ``SOURCED_VALUES`` entry.
    """
    from pydantic import TypeAdapter

    from torchcell.datamodels.schema import ExperimentType
    from torchcell.verification.levels import l0_structural, l1_count, l2_value_fidelity
    from torchcell.verification.runners import _gene_set_for_reference, load_records

    base = data_root or _data_root()
    records = load_records(dataset_root)
    drops = DropLog.model_validate_json(
        Path(dataset_root, "preprocess", "dropped_records.json").read_text()
    )
    validate: Callable[[Any], object] = TypeAdapter(ExperimentType).validate_python
    report = VerificationReport(
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=VERIFIER_PROVENANCE,
    )
    report.add(l0_structural((r["experiment"] for r in records), validate))
    report.add(l1_count(len(records), drops.kept_records))
    report.add(_l1_strain_condition_uniqueness(records))
    report.add(
        l2_value_fidelity(
            [
                float(v)
                for r in records
                for v in r["experiment"]["phenotype"]["metabolite_level"].values()
            ],
            allow_nan=False,
        )
    )
    report.add(_l3_reference_zero_on_the_same_keys(records))
    report.add(_l3_measurement_type(records))

    universe: set[str] = set()
    for reference in {
        json.dumps(r["reference"]["genome_reference"], sort_keys=True) for r in records
    }:
        universe |= _gene_set_for_reference(json.loads(reference), base)
    deleted = {
        p["systematic_gene_name"]
        for r in records
        for p in r["experiment"]["genotype"]["perturbations"]
    }
    missing = sorted(deleted - universe)
    report.add(
        LevelResult(
            level=Level.L4,
            name="gene_containment_mg1655_b_numbers",
            passed=not missing,
            message=f"{len(deleted) - len(missing)} of {len(deleted)} deleted loci are "
            "MG1655 GenBank gene rows",
            details={
                "n_deleted": len(deleted),
                "n_universe": len(universe),
                "missing": missing[:20],
            },
        )
    )
    for value in SOURCED_VALUES.values():
        report.add(audit_sourced_value(value, sourced_value_root(value, base)))

    preprocess = osp.join(dataset_root, "preprocess")
    os.makedirs(preprocess, exist_ok=True)
    with open(osp.join(preprocess, "verification_report.json"), "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main(argv: list[str] | None = None) -> int:
    """CLI: ``deposit`` the raw mirror, ``build`` the dev LMDB, or ``verify`` it."""
    from dotenv import load_dotenv

    parser = argparse.ArgumentParser(
        prog="python -m torchcell.datasets.ecoli.schastnaya2021"
    )
    sub = parser.add_subparsers(dest="command", required=True)
    deposit = sub.add_parser("deposit", help="deposit the raw mirror")
    deposit.add_argument(
        "--download-dir",
        default=None,
        help="directory holding the retrieved workbooks; the literature mirror's si/ "
        "by default",
    )
    deposit.add_argument(
        "--retrieve",
        action="store_true",
        help="run every recorded retriever into --download-dir first",
    )
    sub.add_parser("build", help="build (or load) the dev-tree LMDB")
    sub.add_parser("verify", help="run L0-L4 on the built dev-tree LMDB")
    args = parser.parse_args(argv)

    load_dotenv()
    data_root = _data_root()
    dataset_root = osp.join(data_root, "data/torchcell/metabolome_schastnaya2021")
    if args.command == "deposit":
        download = (
            Path(args.download_dir)
            if args.download_dir
            else library_dir(data_root) / "si"
        )
        if args.retrieve:
            retrieve_raw_files(download)
        sources: dict[str, str | Path] = {name: download / name for name in DATA_SHA256}
        print(deposit_raw_mirror(sources=sources, data_root=data_root))
        return 0
    if args.command == "build":
        dataset = MetabolomeSchastnaya2021Dataset(root=dataset_root)
        print(f"len = {len(dataset)}")
        return 0
    report = verify_build(dataset_root, data_root=data_root)
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
