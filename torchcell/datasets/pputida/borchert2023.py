# torchcell/datasets/pputida/borchert2023
# [[torchcell.datasets.pputida.borchert2023]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/pputida/borchert2023
# Test file: tests/torchcell/datasets/pputida/test_borchert2023.py
"""Borchert 2023 lignin tolerance: the KT2440 RB-TnSeq fitness the compendium DROPPED.

Borchert, Bleem and Beckham 2023 (Metab Eng 77:208-218, doi:10.1016/j.ymben.2023.04.007,
citation key ``borchertRBTnSeqIdentifiesGenetic2023``) ran 14 RB-TnSeq experiments in
biological triplicate on the KT2440 ``ML-5`` library: M9 + 20 mM glucose against eleven
lignin-relevant stressors, plus a second glucose reference and a protocatechuate
enrichment run on a later day. Row 29 of the fifty bacterial rows.

MOST OF IT IS ALREADY SERVED, AND THAT IS MEASURED HERE, NOT ASSUMED. The 42 cultures are
42 of the 332 samples of the Borchert 2024 compendium that ``RbTnseqBorchert2024Dataset``
already serves. ``si1.xlsx``'s 13 pairwise comparison sheets carry 6 per-replicate fitness
columns each; every one of the 42 distinct (experiment, replicate) arms behind them equals
exactly one compendium sample column on all 4,732 shared loci, to a maximum absolute
difference of 0.0005, which is the half-unit of the compendium's three-decimal rounding.
The runner-up compendium column is never closer than 1.29. The build re-measures this over
all 13 sheets and refuses to write if it drifts, so none of those 198,744 values is stored
a second time.

WHAT IS GENUINELY NEW, and why this dataset exists. The compendium eliminated 832 of the
5,564 protein-coding genes ("In instances where gene fitness data for a particular gene did
not exist across all 332 data sets, the gene was eliminated from analysis"). MEASURED:
**271 of those eliminated loci carry full triplicate fitness in at least one Borchert 2023
comparison sheet**, which is 10,824 (locus, arm) values the compendium does not have. Those
are the records this dataset stores, and they are disjoint from the served store by
construction and by a build-time check of the served LMDB's own gene set.

RECORDS. One ``BacterialEnvironmentResponseExperiment`` per (locus, arm): the genotype is a
``TransposonInsertionPerturbation`` of the locus in ``Putida_ML5_JBEI``, the environment is
M9 + 20 mM glucose plus the arm's stressor at the Methods dose, and the phenotype is the
replicate's gene fitness. Replicates are NOT averaged, matching the compendium; the
``_mean`` columns are MEASURED to be the arithmetic mean of their three replicates to
7.3e-14, so they are a derived column and are not stored.

PHENOTYPE CLASS. RB-TnSeq gene fitness is a signed log2 ratio centered on zero, and
``FitnessPhenotype.validate_fitness`` clamps every non-positive value to 0.0, so the record
family is ``EnvironmentResponsePhenotype`` with ``measurement_type=log2_ratio`` and
``assay_type=pooled_competitive_growth_barcode``, exactly as the compendium loader stores
the same measurement.

DURATION. Unlike the compendium, Borchert 2023 releases a per-replicate growth duration
(Table S4), so ``Environment.duration_hours`` is filled here where the compendium loader
records it as a gap.

NOT STORED, each for a stated reason (``preprocess/not_stored.json``):

* the 198,744 (compendium locus, arm) fitness values: measured identical to the served
  Borchert 2024 records, to 0.0005;
* the 26 ``_mean`` columns: measured derived from their replicates;
* the ``t-statistic``, ``p-value``, ``q-value`` and ``adjusted_q-value`` of all 64,853
  (gene, comparison) rows: **no phenotype class in the schema carries a significance
  triple.** ``GeneInteractionPhenotype.gene_interaction_p_value`` is the only p-value field
  and it is not applicable to a fitness contrast. Recorded on issue #776, which already
  asks for fields on this same class; storing the fitness and dropping its significance
  call silently is what this refusal avoids;
* ``Exp1_All_Poolcount`` and ``Exp2_All_Poolcount``: 186,957 barcodes x 42 raw read counts.
  The perturbation leaf for that grain EXISTS (``TransposonInsertionPerturbation.barcode``,
  ``insertion_position``, ``insertion_strand``), so the blocker is the phenotype side: a
  per-barcode sequencing read count is a raw read tally, not a phenotype, and no class or
  ``MeasurementType`` member holds one;
* ``Figure_2..6_growth_data``: back-scattered 620 nm light in arbitrary units over time.
  ``MeasurementType`` has no optical-density member (issue #776 item 3) and no phenotype
  carries a time series.

Design, the measurements and the partition: [[torchcell.datasets.pputida.borchert2023]].
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import os.path as osp
import pickle
import shutil
from collections.abc import Callable, Iterable, Mapping, Sequence
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, ClassVar, Literal

import numpy as np
import numpy.typing as npt
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field
from tqdm import tqdm

import torchcell.datasets.pputida.borchert2024 as b24
from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    file_sha256,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.datamodels.schema import (
    AssayType,
    AssemblyReferenceGenome,
    BacterialEnvironmentResponseExperiment,
    BacterialEnvironmentResponseExperimentReference,
    Concentration,
    ConcentrationUnit,
    Environment,
    EnvironmentPerturbationType,
    EnvironmentPhysicalPerturbation,
    EnvironmentResponsePhenotype,
    Experiment,
    ExperimentReference,
    Genotype,
    MeasurementType,
    PhysicalFactor,
    Publication,
    SmallMoleculePerturbation,
    Temperature,
)
from torchcell.datasets.bacteria_common import (
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ROLE_RAW_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
    SourceCheck,
)
from torchcell.sequence.genome.pputida.kt2440 import PPutidaKT2440Genome
from torchcell.verification.report import Provenance, VerificationReport
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
)

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
CITATION_KEY = "borchertRBTnSeqIdentifiesGenetic2023"
DOI = "10.1016/j.ymben.2023.04.007"
TITLE = (
    "RB-TnSeq identifies genetic targets for improved tolerance of Pseudomonas putida "
    "towards compounds relevant to lignin conversion"
)
#: sha256 of the mirrored OCR every paper quote is a literal substring of.
PAPER_MD_SHA256 = "e156d3b2139add2feb8b3ffb1b319f9a7d369d6ee9b6454f67c38d33c924265a"
#: sha256 of the mirrored SI OCR every Table S4 quote is a literal substring of.
SI2_MD_SHA256 = "568d7edfa237314d9ac775b1cb4bd69645dbf4e3ba4ba03e106f1f5760c23891"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"

#: Elsevier's supplementary-file slot for this article.
ARTICLE_PII = "S1096717623000599"
DATA_FILE = "si1.xlsx"
DATA_RELPATH = f"data/{DATA_FILE}"
DATA_URL = f"https://ars.els-cdn.com/content/image/1-s2.0-{ARTICLE_PII}-mmc1.xlsx"
DATA_SHA256 = "b80c6866c6fbb95696067532f4890951fcbbf2dc72f4351df0bdec86956a0784"
DATA_BYTES = 70872248
RAW_RETRIEVED_AT = "2026-10-07"

_PAPER = Provenance(
    source_uri="paper.md", citation_key=CITATION_KEY, sha256=PAPER_MD_SHA256
)
_SI2 = Provenance(
    source_uri="si/si2.md", citation_key=CITATION_KEY, sha256=SI2_MD_SHA256
)
_RELEASE = Provenance(
    source_uri=DATA_RELPATH,
    citation_key=CITATION_KEY,
    sha256=DATA_SHA256,
    method=f"raw mirror $DATA_ROOT/{RAW_DIR_REL}/{DATA_RELPATH}",
)

# --------------------------------------------------------------------------- #
# Verbatim quotes (literal substrings of the pinned mirrors above)
# --------------------------------------------------------------------------- #
Q_DOSES = (
    "M9 minimal medium with $2 0 \\mathrm { m M }$ glucose and supplemented with either "
    "nothing, 60 mM 4-coumarate, $6 0 ~ \\mathrm { m M }$ ferulate, $6 0 ~ \\mathrm { m M "
    "}$ 4-hydroxybenzoate, $6 0 ~ \\mathrm { m M }$ vanillate, $1 0 \\ \\mathrm { m M }$ "
    "4-hydroxybenzaldehyde, $1 0 \\ \\mathrm { \\ m M }$ vanillin, $3 0 \\mathrm { \\ m M "
    "}$ protocatechuate, $7 5 ~ \\mathrm { m M }$ acetate, $5 0 0 ~ \\mathrm { m M }$ "
    "lactate, $1 2 5 ~ \\mathrm { m M }$ glycolate, $5 0 0 \\mathrm { m M N a C l }$ , or "
    "$5 0 0 \\mathrm { m M N a } _ { 2 } S \\mathrm { O } _ { 4 }$"
)
Q_CULTURE = (
    "Cultures were grown at $3 0 ~ ^ { \\circ } \\mathbf { C }$ ,shaking at $2 2 5 ~ "
    "\\mathrm { r p m }$ , until reaching an $\\mathrm { O D } _ { 6 0 0 \\mathrm { n m } "
    "}$ of 1.0, when 1 mL aliquots were taken"
)
Q_STRAIN_FITNESS = (
    "Transposon insertion counts were used to determine strain fitness calculated as a "
    "normalized $\\log _ { 2 }$ ratio of barcode reads in the enrichment sample vs. the "
    "baseline sample."
)
Q_INSERTION_FILTER = (
    "If a gene did not contain at least three transposon insertions in the enrichment "
    "condition or ${ > } 3 0$ transposon insertions in the baseline condition, it was "
    "excluded from analysis."
)
Q_SIGNIFICANCE = (
    "Comparison of mean fitness values between enrichment and medium reference cultures "
    "$( \\mathbf { M 9 } + 2 0 \\mathbf { m M }$ glucose alone) was performed using a "
    "two-sample $t$ test, where the $p$ value was corrected for multiple testing via the "
    "positive false discovery rate (pFDR) method"
)
Q_ADJUSTED_Q = (
    "pFDR $q$ values were then adjusted for monotonicity (Yekutieli and Benjamini, 1999), "
    "and both unadjusted and adjusted $q$ values are reported."
)
Q_POOLCOUNT = (
    "reads were tabulated according to the number of times each barcode was seen in each "
    "sample and the table of barcode counts was then processed according to a table of "
    "previously defined genomic barcode locations in the KT2440 library to generate a "
    "table (all.poolcount) of tabulated strain counts for each transposon insertion "
    "across all samples"
)
Q_BACKSCATTER = (
    "Growth was assessed as the change in measured in arbitrary units [a.u.] of "
    "back-scatted $6 2 0 \\mathrm { n m }$ light with gain set to 3."
)
Q_TABLE_S4 = (
    "Table S4. Experimental layout description with corresponding $\\mathrm { O D } _ { 6 "
    "0 0 }$ readings and total growth duration taken at time of sampling for each of the "
    "three biological replicates. \\*These experiments were performed later than the "
    "preceding experiments."
)


def _sv(
    source: Provenance, value: Any, quote: str, note: str | None = None
) -> SourcedValue:
    """A value quoted from one pinned artifact."""
    return SourcedValue(value=value, provenance=source, quote=quote, note=note)


# --------------------------------------------------------------------------- #
# Table S4: the experimental layout, with the per-replicate growth duration
# --------------------------------------------------------------------------- #
class ExperimentLayout(BaseModel):
    """One Table S4 row: an experiment, its condition and its three replicates.

    ``od600`` and ``duration`` are the released values for replicates A, B and C, in
    that column order; ``duration`` keeps the released ``h:mm`` text verbatim and
    :func:`duration_hours` converts it. ``quote`` is the row as the pinned OCR renders
    it, so the audit can find it.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    experiment: int
    condition: str
    od600: tuple[float, float, float]
    duration: tuple[str, str, str]
    quote: str


TABLE_S4: tuple[ExperimentLayout, ...] = (
    ExperimentLayout(
        experiment=1,
        condition="M9 + 20 mM glucose",
        od600=(1.2, 0.99, 1.21),
        duration=("10:40", "10:05", "10:05"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose</td><td rowspan=1 colspan=1>1"
            "</td><td rowspan=1 colspan=1>1.20</td><td rowspan=1 colspan=1>10:40</td><td "
            "rowspan=1 colspan=1>0.99</td><td rowspan=1 colspan=1>10:05</td><td rowspan=1"
            " colspan=1>1.21</td><td rowspan=1 colspan=1>10:05</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=2,
        condition="M9 + 20 mM glucose,60 mM 4-coumarate",
        od600=(1.11, 1.08, 0.99),
        duration=("12:45", "12:25", "11:55"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,60 mM 4-coumarate</td><td row"
            "span=1 colspan=1>2</td><td rowspan=1 colspan=1>1.11</td><td rowspan=1 colspa"
            "n=1>12:45</td><td rowspan=1 colspan=1>1.08</td><td rowspan=1 colspan=1>12:25"
            "</td><td rowspan=1 colspan=1>0.99</td><td rowspan=1 colspan=1>11:55</td></tr"
            ">"
        ),
    ),
    ExperimentLayout(
        experiment=3,
        condition="M9 + 20 mM glucose,60 mM ferulate",
        od600=(1.08, 1.63, 1.21),
        duration=("14:25", "15:10", "15:20"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,60 mM ferulate</td><td rowspa"
            "n=1 colspan=1>3</td><td rowspan=1 colspan=1>1.08</td><td rowspan=1 colspan=1"
            ">14:25</td><td rowspan=1 colspan=1>1.63</td><td rowspan=1 colspan=1>15:10</t"
            "d><td rowspan=1 colspan=1>1.21</td><td rowspan=1 colspan=1>15:20</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=4,
        condition="M9 + 20 mM glucose,10 mM 4-hydroxybenzaldehyde",
        od600=(1.13, 2.09, 2.18),
        duration=("19:30", "19:30", "17:50"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,10 mM 4-hydroxybenzaldehyde</"
            "td><td rowspan=1 colspan=1>4</td><td rowspan=1 colspan=1>1.13</td><td rowspa"
            "n=1 colspan=1>19:30</td><td rowspan=1 colspan=1>2.09</td><td rowspan=1 colsp"
            "an=1>19:30</td><td rowspan=1 colspan=1>2.18</td><td rowspan=1 colspan=1>17:5"
            "0</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=5,
        condition="M9 + 20 mM glucose,60 mM 4-hydroxybenzoate",
        od600=(1.12, 1.01, 0.97),
        duration=("9:30", "10:04", "9:30"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,60 mM 4-hydroxybenzoate</td><"
            "td rowspan=1 colspan=1>5</td><td rowspan=1 colspan=1>1.12</td><td rowspan=1 "
            "colspan=1>9:30</td><td rowspan=1 colspan=1>1.01</td><td rowspan=1 colspan=1>"
            "10:04</td><td rowspan=1 colspan=1>0.97</td><td rowspan=1 colspan=1>9:30</td>"
            "</tr>"
        ),
    ),
    ExperimentLayout(
        experiment=6,
        condition="M9 + 20 mM glucose,10 mM vanillin",
        od600=(1.08, 1.47, 1.15),
        duration=("13:15", "15:10", "13:15"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,10 mM vanillin</td><td rowspa"
            "n=1 colspan=1>6</td><td rowspan=1 colspan=1>1.08</td><td rowspan=1 colspan=1"
            ">13:15</td><td rowspan=1 colspan=1>1.47</td><td rowspan=1 colspan=1>15:10</t"
            "d><td rowspan=1 colspan=1>1.15</td><td rowspan=1 colspan=1>13:15</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=7,
        condition="M9 + 20 mM glucose,60 mM vanillate",
        od600=(1.05, 1.64, 1.2),
        duration=("18:40", "18:40", "20:00"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,60 mM vanillate</td><td rowsp"
            "an=1 colspan=1>7</td><td rowspan=1 colspan=1>1.05</td><td rowspan=1 colspan="
            "1>18:40</td><td rowspan=1 colspan=1>1.64</td><td rowspan=1 colspan=1>18:40</"
            "td><td rowspan=1 colspan=1>1.20</td><td rowspan=1 colspan=1>20:00</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=8,
        condition="M9 + 20 mM glucose,500 mM NaCl",
        od600=(0.95, 1.04, 1.01),
        duration=("19:30", "20:30", "20:00"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,500 mM NaCl</td><td rowspan=1"
            " colspan=1>8</td><td rowspan=1 colspan=1>0.95</td><td rowspan=1 colspan=1>19"
            ":30</td><td rowspan=1 colspan=1>1.04</td><td rowspan=1 colspan=1>20:30</td><"
            "td rowspan=1 colspan=1>1.01</td><td rowspan=1 colspan=1>20:00</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=9,
        condition="M9 + 20 mM glucose,500 mM Na2SO4",
        od600=(0.94, 1.09, 1.02),
        duration=("24:40", "24:40", "24:00"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,500 mM Na2SO4</td><td rowspan"
            "=1 colspan=1>9</td><td rowspan=1 colspan=1>0.94</td><td rowspan=1 colspan=1>"
            "24:40</td><td rowspan=1 colspan=1>1.09</td><td rowspan=1 colspan=1>24:40</td"
            "><td rowspan=1 colspan=1>1.02</td><td rowspan=1 colspan=1>24:00</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=10,
        condition="M9 + 20 mM glucose,75 mM acetate",
        od600=(1.23, 1.05, 0.95),
        duration=("30:25", "30:25", "31:00"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,75 mM acetate</td><td rowspan"
            "=1 colspan=1>10</td><td rowspan=1 colspan=1>1.23</td><td rowspan=1 colspan=1"
            ">30:25</td><td rowspan=1 colspan=1>1.05</td><td rowspan=1 colspan=1>30:25</t"
            "d><td rowspan=1 colspan=1>0.95</td><td rowspan=1 colspan=1>31:00</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=11,
        condition="M9 + 20 mM glucose,500 mM lactate",
        od600=(1.0, 1.05, 1.05),
        duration=("25:10", "26:20", "24:40"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,500 mM lactate</td><td rowspa"
            "n=1 colspan=1>11</td><td rowspan=1 colspan=1>1.00</td><td rowspan=1 colspan="
            "1>25:10</td><td rowspan=1 colspan=1>1.05</td><td rowspan=1 colspan=1>26:20</"
            "td><td rowspan=1 colspan=1>1.05</td><td rowspan=1 colspan=1>24:40</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=12,
        condition="M9 + 20 mM glucose,125 mM glycolate",
        od600=(1.03, 1.05, 1.31),
        duration=("15:20", "16:40", "16:40"),
        quote=(
            "<tr><td rowspan=1 colspan=1>M9 + 20 mM glucose,125 mM glycolate</td><td rows"
            "pan=1 colspan=1>12</td><td rowspan=1 colspan=1>1.03</td><td rowspan=1 colspa"
            "n=1>15:20</td><td rowspan=1 colspan=1>1.05</td><td rowspan=1 colspan=1>16:40"
            "</td><td rowspan=1 colspan=1>1.31</td><td rowspan=1 colspan=1>16:40</td></tr"
            ">"
        ),
    ),
    ExperimentLayout(
        experiment=13,
        condition="*M9 + 20 mM glucose",
        od600=(1.0, 1.09, 1.25),
        duration=("10:10", "10:50", "10:10"),
        quote=(
            "<tr><td rowspan=1 colspan=1>*M9 + 20 mM glucose</td><td rowspan=1 colspan=1>"
            "13</td><td rowspan=1 colspan=1>1.00</td><td rowspan=1 colspan=1>10:10</td><t"
            "d rowspan=1 colspan=1>1.09</td><td rowspan=1 colspan=1>10:50</td><td rowspan"
            "=1 colspan=1>1.25</td><td rowspan=1 colspan=1>10:10</td></tr>"
        ),
    ),
    ExperimentLayout(
        experiment=14,
        condition="*M9 + 20 mM glucose,30 mM protecatechuate",
        od600=(1.13, 0.9, 0.93),
        duration=("20:20", "20:20", "18:10"),
        quote=(
            "<tr><td rowspan=1 colspan=1>*M9 + 20 mM glucose,30 mM protecatechuate</td><t"
            "d rowspan=1 colspan=1>14</td><td rowspan=1 colspan=1>1.13</td><td rowspan=1 "
            "colspan=1>20:20</td><td rowspan=1 colspan=1>0.9</td><td rowspan=1 colspan=1>"
            "20:20</td><td rowspan=1 colspan=1>0.93</td><td rowspan=1 colspan=1>18:10</td"
            "></tr>"
        ),
    ),
)
"""Table S4, one row per experiment, transcribed from the pinned ``si/si2.md`` and
re-read off the rendered page of ``si/si2.pdf`` (page S5) to catch an OCR digit."""

LAYOUT_BY_EXPERIMENT: dict[int, ExperimentLayout] = {
    layout.experiment: layout for layout in TABLE_S4
}

REPLICATES: tuple[str, str, str] = ("A", "B", "C")


def duration_hours(released: str) -> float:
    """The released ``h:mm`` growth duration as hours."""
    hours, minutes = released.split(":")
    return int(hours) + int(minutes) / 60.0


# --------------------------------------------------------------------------- #
# The 14 experiments: their sheet columns, dose and compendium samples
# --------------------------------------------------------------------------- #
#: The 11 comparison sheets whose reference arm is experiment 1's glucose culture.
DAY1_SHEETS: tuple[str, ...] = (
    "Glu_v_M9_Glu_4CA",
    "Glu_v_Glu_4HBA",
    "Glu_v_Glu_4HBald",
    "Glu_v_Glu_FA",
    "Glu_v_Glu_VA",
    "Glu_v_Glu_Van",
    "Glu_v_Glu_AA",
    "Glu_v_Glu_GA",
    "Glu_v_Glu_LA",
    "Glu_v_Glu_Na2SO4",
    "Glu_v_Glu_NaCl",
)


class ExperimentColumns(BaseModel):
    """One experiment: where its three replicate columns are, and what it was.

    ``prefix`` plus ``_Rep{A,B,C}`` is the column name in every sheet of ``sheets``.
    ``stressor`` is the added compound, named as the shared compound-identity table
    resolves it (see :data:`STRESSOR_NAMING`), and ``dose_mm`` is the Methods dose.
    ``compendium`` is the DECLARED Borchert 2024 sample of each replicate; the build
    re-derives it by value and refuses any disagreement.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    experiment: int
    prefix: str
    sheets: tuple[str, ...]
    stressor: str | None
    dose_mm: float | None
    compendium: tuple[str, str, str]

    def column(self, replicate: str) -> str:
        """The sheet column holding this experiment's replicate."""
        return f"{self.prefix}_Rep{replicate}"

    def arm(self, replicate: str) -> str:
        """Table S3's own ``Exp./Rep.`` key for this culture, e.g. ``Exp1A``."""
        return f"Exp{self.experiment}{replicate}"


EXPERIMENTS: tuple[ExperimentColumns, ...] = (
    ExperimentColumns(
        experiment=1,
        prefix="M9_Glucose",
        sheets=DAY1_SHEETS,
        stressor=None,
        dose_mm=None,
        compendium=("set100IT008", "set100IT009", "set100IT010"),
    ),
    ExperimentColumns(
        experiment=2,
        prefix="M9_Glucose_4-Coumarate",
        sheets=("Glu_v_M9_Glu_4CA",),
        stressor="p-Coumaric acid",
        dose_mm=60.0,
        compendium=("set100IT011", "set100IT012", "set100IT013"),
    ),
    ExperimentColumns(
        experiment=3,
        prefix="M9_Glucose_Ferulate",
        sheets=("Glu_v_Glu_FA", "Glu_FA_v_Glu_PCA"),
        stressor="Ferulic acid",
        dose_mm=60.0,
        compendium=("set100IT014", "set100IT015", "set100IT016"),
    ),
    ExperimentColumns(
        experiment=4,
        prefix="M9_Glucose_4-HBaldehyde",
        sheets=("Glu_v_Glu_4HBald",),
        stressor="4-Hydroxybenzaldehyde",
        dose_mm=10.0,
        compendium=("set100IT017", "set100IT018", "set100IT019"),
    ),
    ExperimentColumns(
        experiment=5,
        prefix="M9_Glucose_4-HBA",
        sheets=("Glu_v_Glu_4HBA",),
        stressor="4-Hydroxybenzoic acid",
        dose_mm=60.0,
        compendium=("set100IT020", "set100IT021", "set100IT022"),
    ),
    ExperimentColumns(
        experiment=6,
        prefix="M9_Glucose_Vanillin",
        sheets=("Glu_v_Glu_Van",),
        stressor="Vanillin",
        dose_mm=10.0,
        compendium=("set100IT023", "set100IT024", "set100IT025"),
    ),
    ExperimentColumns(
        experiment=7,
        prefix="M9_Glucose_Vanillate",
        sheets=("Glu_v_Glu_VA",),
        stressor="Vanillic acid",
        dose_mm=60.0,
        compendium=("set100IT026", "set100IT027", "set100IT028"),
    ),
    ExperimentColumns(
        experiment=8,
        prefix="M9_Glucose_NaCl",
        sheets=("Glu_v_Glu_NaCl",),
        stressor="Sodium chloride",
        dose_mm=500.0,
        compendium=("set100IT032", "set100IT033", "set100IT034"),
    ),
    ExperimentColumns(
        experiment=9,
        prefix="M9_Glucose_Na2SO4",
        sheets=("Glu_v_Glu_Na2SO4",),
        stressor="Sodium sulfate",
        dose_mm=500.0,
        compendium=("set100IT035", "set100IT036", "set100IT037"),
    ),
    ExperimentColumns(
        experiment=10,
        prefix="M9_Glucose_Acetate",
        sheets=("Glu_v_Glu_AA",),
        stressor="Acetic acid",
        dose_mm=75.0,
        compendium=("set100IT038", "set100IT039", "set100IT040"),
    ),
    ExperimentColumns(
        experiment=11,
        prefix="M9_Glucose_Lactate",
        sheets=("Glu_v_Glu_LA",),
        stressor="Lactic acid",
        dose_mm=500.0,
        compendium=("set100IT041", "set100IT042", "set100IT043"),
    ),
    ExperimentColumns(
        experiment=12,
        prefix="M9_Glucose_Glycolate",
        sheets=("Glu_v_Glu_GA",),
        stressor="Glycolic acid",
        dose_mm=125.0,
        compendium=("set100IT044", "set100IT045", "set100IT046"),
    ),
    ExperimentColumns(
        experiment=13,
        prefix="M9_Glucose",
        sheets=("Glu_v_Glu_PCA",),
        stressor=None,
        dose_mm=None,
        compendium=("set101IT007", "set101IT008", "set101IT009"),
    ),
    ExperimentColumns(
        experiment=14,
        prefix="M9_Glucose_PCA",
        sheets=("Glu_v_Glu_PCA", "Glu_FA_v_Glu_PCA"),
        stressor="Protecatechuic acid",
        dose_mm=30.0,
        compendium=("set101IT013", "set101IT014", "set101IT015"),
    ),
)
N_EXPERIMENTS = 14
N_ARMS = N_EXPERIMENTS * len(REPLICATES)

STRESSOR_NAMING = (
    "Borchert 2023's Methods name the conjugate bases ('4-coumarate', 'ferulate', "
    "'vanillate', ...), which the shared compound-identity table resolves to name-only "
    "compounds with no InChIKey; the compendium's acid spellings resolve to a structure. "
    "MEASURED over the 12 stressors: 4 of 12 agree on both spellings, 8 of 12 resolve "
    "only under the compendium's. So the agent is named as the served Borchert 2024 "
    "records name it, which keeps ONE compound node per chemical in the graph, and the "
    "build asserts the dose against the Methods"
)

#: The ``Glu_v_Glu_4HBald`` Read_Me says 20 mM; three other released statements say 10.
READ_ME_4HBALD_CONFLICT = (
    "the Read_Me description of sheet Glu_v_Glu_4HBald says '20 mM "
    "4-hydroxybenzaldehyde', while the Methods, Table S4 and the Borchert 2024 "
    "compendium metadata all say 10 mM. 10 mM is stored, 3 released statements to 1"
)

#: Every value this loader hardcodes, each with its verbatim quote and pinned sha256.
SOURCED_VALUES: dict[str, SourcedValue] = {
    "glucose_mm": _sv(_PAPER, 20.0, Q_DOSES),
    "stressor_doses_mm": _sv(
        _PAPER,
        {exp.stressor: exp.dose_mm for exp in EXPERIMENTS if exp.stressor is not None},
        Q_DOSES,
        note=STRESSOR_NAMING,
    ),
    "temperature_c": _sv(_PAPER, 30.0, Q_CULTURE),
    "sampling_od600": _sv(
        _PAPER,
        1.0,
        Q_CULTURE,
        note="the harvest rule; Table S4 reports the OD600 actually reached per replicate",
    ),
    "fitness_definition": _sv(
        _PAPER, "normalized log2(enrichment/baseline) barcode ratio", Q_STRAIN_FITNESS
    ),
    "gene_fitness_normalization": _sv(
        _PAPER,
        "weighted mean over untrimmed insertions, 251-gene sliding-window median = 0",
        b24.Q_BORCHERT23_FITNESS,
        note="so the reference phenotype of every record is 0.0",
    ),
    "insertion_filter": _sv(
        _PAPER,
        {"enrichment_min_insertions": 3, "baseline_min_insertions": 30},
        Q_INSERTION_FILTER,
    ),
    "triplicate_filter": _sv(_PAPER, 3, b24.Q_BORCHERT23_TRIPLICATE),
    "library": _sv(
        _PAPER,
        "ML-5",
        b24.Q_BORCHERT23_LIBRARY,
        note="the compendium labels these same 42 cultures 'Putida_ML5_JBEI', which is "
        "the library_pool stored, so a record joins the served ones on the strain",
    ),
    "protocatechuate_separate_day": _sv(
        _PAPER,
        {"experiments": [13, 14]},
        b24.Q_BORCHERT23_PCA_DAY,
        note="MEASURED: the Glu_v_Glu_PCA sheet's M9_Glucose columns differ from the "
        "other 11 sheets' by up to 6.93, so they are a second glucose culture, not a "
        "re-export; the 11 day-1 sheets agree to exactly 0.0",
    ),
    "significance_test": _sv(
        _PAPER, "two-sample t test, pFDR q values", Q_SIGNIFICANCE
    ),
    "adjusted_q": _sv(_PAPER, "monotonicity-adjusted pFDR q", Q_ADJUSTED_Q),
    "poolcount_is_raw_counts": _sv(_PAPER, "all.poolcount", Q_POOLCOUNT),
    "growth_curve_readout": _sv(
        _PAPER, "back-scattered 620 nm light [a.u.], gain 3", Q_BACKSCATTER
    ),
    "sra": _sv(_PAPER, "SRP385031", b24.Q_BORCHERT23_SRA),
    "growth_duration_source": _sv(
        _SI2,
        "per-replicate growth duration at the time of sampling",
        Q_TABLE_S4,
        note="the compendium releases no growth time, so Environment.duration_hours is a "
        "gap on the served records and a value here",
    ),
    **{
        f"layout_experiment_{layout.experiment}": _sv(
            _SI2,
            {
                "condition": layout.condition,
                "od600": list(layout.od600),
                "duration": list(layout.duration),
            },
            layout.quote,
        )
        for layout in TABLE_S4
    },
}

# --------------------------------------------------------------------------- #
# The comparison sheets
# --------------------------------------------------------------------------- #
ID_COLUMNS: tuple[str, ...] = (
    "old_locus_tag",
    "new_locus_tag",
    "gene_name",
    "description",
)
STAT_COLUMNS: tuple[str, ...] = (
    "t-statistic",
    "p-value",
    "q-value",
    "adjusted_q-value",
)
READ_ME_SHEET = "Read_Me"
POOLCOUNT_SHEETS: tuple[str, ...] = ("Exp1_All_Poolcount", "Exp2_All_Poolcount")
GROWTH_SHEETS: tuple[str, ...] = tuple(f"Figure_{n}_growth_data" for n in range(2, 7))


class ComparisonPlan(BaseModel):
    """One comparison sheet: which experiment is the reference and which the condition."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    sheet: str
    reference_experiment: int
    condition_experiment: int


COMPARISONS: tuple[ComparisonPlan, ...] = (
    ComparisonPlan(
        sheet="Glu_v_M9_Glu_4CA", reference_experiment=1, condition_experiment=2
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_4HBA", reference_experiment=1, condition_experiment=5
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_4HBald", reference_experiment=1, condition_experiment=4
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_FA", reference_experiment=1, condition_experiment=3
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_VA", reference_experiment=1, condition_experiment=7
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_Van", reference_experiment=1, condition_experiment=6
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_PCA", reference_experiment=13, condition_experiment=14
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_AA", reference_experiment=1, condition_experiment=10
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_GA", reference_experiment=1, condition_experiment=12
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_LA", reference_experiment=1, condition_experiment=11
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_Na2SO4", reference_experiment=1, condition_experiment=9
    ),
    ComparisonPlan(
        sheet="Glu_v_Glu_NaCl", reference_experiment=1, condition_experiment=8
    ),
    ComparisonPlan(
        sheet="Glu_FA_v_Glu_PCA", reference_experiment=3, condition_experiment=14
    ),
)
N_COMPARISONS = 13
#: The workbook's sheets, in released order; any other layout stops the build.
SHEET_ORDER: tuple[str, ...] = (
    (READ_ME_SHEET,)
    + POOLCOUNT_SHEETS
    + GROWTH_SHEETS
    + tuple(c.sheet for c in COMPARISONS)
)


class Comparison(BaseModel):
    """One parsed comparison sheet."""

    model_config = ConfigDict(extra="forbid", frozen=True, arbitrary_types_allowed=True)

    sheet: str
    tags: tuple[str, ...] = Field(description="old_locus_tag, in sheet order")
    columns: tuple[str, ...] = Field(description="the 8 value columns, in sheet order")
    values: npt.NDArray[np.float64] = Field(description="rows x 8")
    stats: npt.NDArray[np.float64] = Field(description="rows x 4 (t, p, q, adjusted q)")

    def column(self, name: str) -> npt.NDArray[np.float64]:
        """One value column of this sheet."""
        return self.values[:, self.columns.index(name)]


COLUMNS_BY_EXPERIMENT: dict[int, ExperimentColumns] = {
    experiment.experiment: experiment for experiment in EXPERIMENTS
}


def _expected_columns(plan: ComparisonPlan) -> tuple[str, ...]:
    """The 8 value-column headers a comparison sheet must carry, in released order."""
    names: list[str] = []
    for number in (plan.reference_experiment, plan.condition_experiment):
        prefix = COLUMNS_BY_EXPERIMENT[number].prefix
        names.extend(f"{prefix}_Rep{r}" for r in REPLICATES)
        names.append(f"{prefix}_mean")
    return tuple(names)


def read_comparisons(path: str | Path) -> tuple[Comparison, ...]:
    """Parse the 13 comparison sheets of ``si1.xlsx``, refusing any other layout.

    The workbook's sheet list, every sheet's 16 headers and the absence of an empty
    value cell are all asserted; the two poolcount sheets and the five growth sheets
    are not read, because nothing in them is stored (see the module docstring).
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if tuple(workbook.sheetnames) != SHEET_ORDER:
        raise ValueError(f"unexpected sheets {workbook.sheetnames}")
    parsed: list[Comparison] = []
    for plan in COMPARISONS:
        rows = list(workbook[plan.sheet].iter_rows(values_only=True))
        header = tuple(str(c) for c in rows[0])
        expected = ID_COLUMNS + _expected_columns(plan) + STAT_COLUMNS
        if header != expected:
            raise ValueError(f"{plan.sheet}: header {header} is not {expected}")
        body = rows[1:]
        tags = tuple(str(r[0]) for r in body)
        if len(set(tags)) != len(tags):
            raise ValueError(f"{plan.sheet}: old_locus_tag repeats")
        values = np.array([r[4:12] for r in body], dtype=np.float64)
        stats = np.array([r[12:16] for r in body], dtype=np.float64)
        if np.isnan(values).any() or np.isnan(stats).any():
            raise ValueError(f"{plan.sheet} has an empty cell")
        parsed.append(
            Comparison(
                sheet=plan.sheet,
                tags=tags,
                columns=_expected_columns(plan),
                values=values,
                stats=stats,
            )
        )
    workbook.close()
    return tuple(parsed)


def read_read_me(path: str | Path) -> dict[str, str]:
    """The ``Read_Me`` sheet's per-tab description, which names each sheet's dose."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook[READ_ME_SHEET].iter_rows(values_only=True))
    workbook.close()
    described = {
        str(row[0]): str(row[1])
        for row in rows
        if row[0] is not None and row[1] is not None
    }
    del described["Tab"]
    return described


# --------------------------------------------------------------------------- #
# The partition against Borchert 2024
# --------------------------------------------------------------------------- #
#: Half a unit in the last place of the compendium's three-decimal fitness.
COMPENDIUM_HALF_ULP = 0.0005
#: A matched column must beat every other compendium column by at least this much.
RUNNER_UP_FLOOR = 0.5


class ArmMatch(BaseModel):
    """One arm's measured identity with a Borchert 2024 sample column."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    arm: str
    experiment: int
    replicate: str
    sheet: str
    column: str
    compendium_sample: str
    n_shared_loci: int
    max_abs_diff: float
    runner_up_sample: str
    runner_up_max_abs_diff: float


def match_arms(
    comparisons: Sequence[Comparison], release: b24.Release
) -> tuple[ArmMatch, ...]:
    """Prove, column by column, that every arm IS a served Borchert 2024 sample.

    For each of the 42 arms the released values are joined to the compendium on the
    locus tag and compared against all 332 sample columns. The closest column must be
    the one :data:`EXPERIMENTS` declares, must agree on every shared locus to
    :data:`COMPENDIUM_HALF_ULP`, and the runner-up must be at least
    :data:`RUNNER_UP_FLOOR` away, so the identification is not a coincidence.
    """
    index = {tag: row for row, tag in enumerate(release.genes)}
    samples = [sample.column for sample in release.samples]
    by_name = {sample.exp_name: sample.column for sample in release.samples}
    matches: list[ArmMatch] = []
    for experiment in EXPERIMENTS:
        sheet = experiment.sheets[0]
        comparison = next(c for c in comparisons if c.sheet == sheet)
        rows = [index[tag] for tag in comparison.tags if tag in index]
        keep = [i for i, tag in enumerate(comparison.tags) if tag in index]
        block = release.fitness[rows, :]
        for position, replicate in enumerate(REPLICATES):
            released = comparison.column(experiment.column(replicate))[keep]
            spread = np.abs(block - released[:, None]).max(axis=0)
            order = np.argsort(spread)
            declared = experiment.compendium[position]
            best = samples[int(order[0])]
            if best != by_name[declared]:
                raise RuntimeError(
                    f"{experiment.arm(replicate)}: the closest compendium column is "
                    f"{best!r}, not the declared {declared!r}"
                )
            if spread[order[0]] > COMPENDIUM_HALF_ULP:
                raise RuntimeError(
                    f"{experiment.arm(replicate)} against {declared}: max abs diff "
                    f"{spread[order[0]]:.6g} exceeds {COMPENDIUM_HALF_ULP}, so the "
                    "subsumption measurement has drifted"
                )
            if spread[order[1]] < RUNNER_UP_FLOOR:
                raise RuntimeError(
                    f"{experiment.arm(replicate)}: the runner-up column "
                    f"{samples[int(order[1])]!r} is only {spread[order[1]]:.6g} away"
                )
            matches.append(
                ArmMatch(
                    arm=experiment.arm(replicate),
                    experiment=experiment.experiment,
                    replicate=replicate,
                    sheet=sheet,
                    column=experiment.column(replicate),
                    compendium_sample=declared,
                    n_shared_loci=len(rows),
                    max_abs_diff=float(spread[order[0]]),
                    runner_up_sample=samples[int(order[1])],
                    runner_up_max_abs_diff=float(spread[order[1]]),
                )
            )
    return tuple(matches)


def assert_arms_agree_across_sheets(comparisons: Sequence[Comparison]) -> list[str]:
    """An arm exported by more than one sheet carries the same values in each.

    Also the reason experiments 1 and 13 are separate arms despite sharing the column
    name ``M9_Glucose_Rep*``: the 11 day-1 sheets agree exactly and the protocatechuate
    sheet, run on a later day, does not.
    """
    by_sheet = {c.sheet: c for c in comparisons}
    proofs: list[str] = []
    for experiment in EXPERIMENTS:
        if len(experiment.sheets) == 1:
            continue
        head, *rest = experiment.sheets
        first = by_sheet[head]
        for other_name in rest:
            other = by_sheet[other_name]
            shared = sorted(set(first.tags) & set(other.tags))
            left = {tag: i for i, tag in enumerate(first.tags)}
            right = {tag: i for i, tag in enumerate(other.tags)}
            worst = 0.0
            for replicate in REPLICATES:
                name = experiment.column(replicate)
                a = first.column(name)[[left[t] for t in shared]]
                b = other.column(name)[[right[t] for t in shared]]
                worst = max(worst, float(np.abs(a - b).max()))
            if worst != 0.0:
                raise RuntimeError(
                    f"experiment {experiment.experiment}: {head} and {other_name} "
                    f"disagree by {worst:.6g} on the same column"
                )
            proofs.append(
                f"experiment {experiment.experiment}: {head} and {other_name} agree "
                f"exactly on {len(shared)} loci x 3 replicates"
            )
    day1, pca = EXPERIMENTS[0], EXPERIMENTS[12]
    reference = by_sheet[day1.sheets[0]]
    later = by_sheet[pca.sheets[0]]
    shared = sorted(set(reference.tags) & set(later.tags))
    left = {tag: i for i, tag in enumerate(reference.tags)}
    right = {tag: i for i, tag in enumerate(later.tags)}
    separation = max(
        float(
            np.abs(
                reference.column(day1.column(r))[[left[t] for t in shared]]
                - later.column(pca.column(r))[[right[t] for t in shared]]
            ).max()
        )
        for r in REPLICATES
    )
    if separation <= COMPENDIUM_HALF_ULP:
        raise RuntimeError(
            "experiments 1 and 13 share the column name M9_Glucose_Rep* and now agree "
            f"to {separation:.6g}; they were measured to be two different cultures"
        )
    proofs.append(
        f"experiments 1 and 13 share the column name M9_Glucose_Rep* but differ by up "
        f"to {separation:.6g} over {len(shared)} loci, so they are two cultures"
    )
    return proofs


#: Largest absolute error allowed when re-deriving a ``_mean`` column.
MEAN_TOLERANCE = 1e-12


def assert_means_are_derived(comparisons: Sequence[Comparison]) -> list[str]:
    """Every ``_mean`` column is the arithmetic mean of its three replicates."""
    proofs: list[str] = []
    for comparison in comparisons:
        worst = 0.0
        for name in comparison.columns:
            if not name.endswith("_mean"):
                continue
            prefix = name.removesuffix("_mean")
            replicates = np.stack(
                [comparison.column(f"{prefix}_Rep{r}") for r in REPLICATES]
            )
            worst = max(
                worst,
                float(np.abs(replicates.mean(axis=0) - comparison.column(name)).max()),
            )
        if worst > MEAN_TOLERANCE:
            raise RuntimeError(
                f"{comparison.sheet}: a _mean column is not the mean of its replicates "
                f"(max abs error {worst:.3g})"
            )
        proofs.append(
            f"{comparison.sheet}: both _mean columns are the mean of their three "
            f"replicates to {worst:.3g}"
        )
    return proofs


def new_loci_by_arm(
    comparisons: Sequence[Comparison], served_loci: Iterable[str]
) -> dict[str, tuple[str, ...]]:
    """Each arm's loci that the Borchert 2024 compendium does not carry, sorted.

    An arm exported by several sheets takes the union of their loci; the values agree
    (:func:`assert_arms_agree_across_sheets`), so the union stores each one once.
    """
    served = set(served_loci)
    by_sheet = {c.sheet: c for c in comparisons}
    out: dict[str, tuple[str, ...]] = {}
    for experiment in EXPERIMENTS:
        extra: set[str] = set()
        for sheet in experiment.sheets:
            extra |= {tag for tag in by_sheet[sheet].tags if tag not in served}
        ordered = tuple(sorted(extra))
        for replicate in REPLICATES:
            out[experiment.arm(replicate)] = ordered
    return out


class ServedPartition(BaseModel):
    """What the already-served Borchert 2024 store holds, and the disjointness proof."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    served_root: str
    served_records: int
    served_loci: int
    served_screens: int
    stored_loci: int
    shared_loci: int
    matched_samples_present: int


def read_served_loci(served_root: str) -> tuple[frozenset[str], frozenset[str], int]:
    """Stream the served Borchert 2024 LMDB once: its loci, its screen ids, its count.

    The generator is fully consumed, so ``stream_records`` closes the environment
    before this returns; a held handle makes the next open of the same path fail.
    """
    from torchcell.verification.runners import stream_records

    loci: set[str] = set()
    screens: set[str] = set()
    records = 0
    for record in stream_records(served_root):
        experiment = record["experiment"]
        for perturbation in experiment["genotype"]["perturbations"]:
            loci.add(perturbation["systematic_gene_name"])
        screens.add(experiment["phenotype"]["screen_id"])
        records += 1
    return frozenset(loci), frozenset(screens), records


def assert_served_partition(
    served_root: str,
    stored_loci: Iterable[str],
    matches: Sequence[ArmMatch],
    release: b24.Release,
) -> ServedPartition:
    """Prove the partition against the served store, in both directions.

    Forward: every arm this build does NOT store is a screen the served store holds, so
    those 198,744 values are already there. Reverse: not one locus this build stores is
    a locus the served store holds, so nothing is written twice.
    """
    served_loci, screens, records = read_served_loci(served_root)
    stored = frozenset(stored_loci)
    if served_loci != frozenset(release.genes):
        raise RuntimeError(
            f"{served_root} holds {len(served_loci)} loci, the pinned compendium "
            f"release has {len(release.genes)}; the served store is stale"
        )
    shared = stored & served_loci
    if shared:
        raise RuntimeError(
            f"{len(shared)} loci are already served by Borchert 2024, so storing them "
            f"would duplicate a record: {sorted(shared)[:5]}"
        )
    missing = [
        m.compendium_sample for m in matches if m.compendium_sample not in screens
    ]
    if missing:
        raise RuntimeError(
            f"the served store has no screen for {missing}, so the fitness this build "
            "declines to store is not in fact served"
        )
    return ServedPartition(
        served_root=served_root,
        served_records=records,
        served_loci=len(served_loci),
        served_screens=len(screens),
        stored_loci=len(stored),
        shared_loci=0,
        matched_samples_present=len(matches),
    )


# --------------------------------------------------------------------------- #
# Environment
# --------------------------------------------------------------------------- #
GLUCOSE_MM = 20.0
TEMPERATURE_C = 30.0
MUTANT_LIBRARY = "Putida_ML5_JBEI"
MEDIA_LABEL = "M9_medium"

_GAP_GENERATIONS = ProvenanceGap(
    field="duration_generations",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_SI2,
    note="Table S4 reports the growth duration in hours, not in doublings",
)


def assert_dose_agrees(
    experiment: ExperimentColumns, sample: b24.SampleMetadata
) -> None:
    """Borchert 2023's own Methods dose equals the compendium's for the same culture."""
    if (sample.condition_1, sample.units_1, sample.concentration_1) != (
        "D-Glucose",
        "mM",
        GLUCOSE_MM,
    ):
        raise RuntimeError(
            f"{sample.exp_name}: carbon source is "
            f"{sample.condition_1!r} {sample.concentration_1} {sample.units_1!r}"
        )
    if sample.media != MEDIA_LABEL or sample.mutant_library != MUTANT_LIBRARY:
        raise RuntimeError(
            f"{sample.exp_name}: {sample.media!r} / {sample.mutant_library!r}"
        )
    if sample.temperature != TEMPERATURE_C or sample.aerobic != "Aerobic":
        raise RuntimeError(
            f"{sample.exp_name}: {sample.temperature} C, {sample.aerobic!r}"
        )
    if experiment.stressor is None:
        if sample.condition_2 is not None:
            raise RuntimeError(
                f"{sample.exp_name} carries {sample.condition_2!r}, experiment "
                f"{experiment.experiment} has no stressor"
            )
        return
    if (sample.condition_2, sample.units_2, sample.concentration_2) != (
        experiment.stressor,
        "mM",
        experiment.dose_mm,
    ):
        raise RuntimeError(
            f"{sample.exp_name}: the compendium has {sample.condition_2!r} at "
            f"{sample.concentration_2} {sample.units_2!r}, Borchert 2023's Methods give "
            f"{experiment.stressor!r} at {experiment.dose_mm} mM"
        )


def build_environment(experiment: ExperimentColumns, replicate: str) -> Environment:
    """The arm's medium, temperature, stressor and its own growth duration."""
    perturbations: list[EnvironmentPerturbationType] = [
        EnvironmentPhysicalPerturbation(
            factor=PhysicalFactor.carbon_source,
            agent=b24.condition_compound("D-Glucose"),
            magnitude=Concentration(
                value=GLUCOSE_MM, unit=ConcentrationUnit.millimolar
            ),
        )
    ]
    if experiment.stressor is not None:
        perturbations.append(
            SmallMoleculePerturbation(
                compound=b24.condition_compound(experiment.stressor),
                concentration=Concentration(
                    value=experiment.dose_mm, unit=ConcentrationUnit.millimolar
                ),
            )
        )
    layout = LAYOUT_BY_EXPERIMENT[experiment.experiment]
    return Environment(
        media=b24.BORCHERT2023_M9,
        temperature=Temperature(value=TEMPERATURE_C),
        perturbations=perturbations,
        aerobicity="aerobic",
        duration_hours=duration_hours(layout.duration[REPLICATES.index(replicate)]),
        provenance_gaps=[_GAP_GENERATIONS],
    )


# --------------------------------------------------------------------------- #
# Phenotype
# --------------------------------------------------------------------------- #
UNITS_FITNESS = (
    "RB-TnSeq gene fitness (Borchert 2023): normalized log2(enrichment/baseline) barcode "
    "ratio, weighted mean over all untrimmed insertions of the locus, normalized by the "
    "median unnormalized fitness of a 251-gene sliding window so the typical gene is 0"
)
UNITS_REFERENCE = "the typical gene of this culture (gene fitness is normalized to 0)"

_GAP_UNCERTAINTY = ProvenanceGap(
    field="environment_response_uncertainty",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_RELEASE,
    note="the sheets release no per-replicate uncertainty; their t, p, q and adjusted-q "
    "are a two-sample contrast between two arms, not a dispersion of one value",
)
_GAP_SE = ProvenanceGap(
    field="environment_response_se",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_RELEASE,
)
_GAP_N = ProvenanceGap(
    field="n_samples",
    reason=ProvenanceGapReason.not_carried_by_curation,
    looked_in=_RELEASE,
    note="a record is one biological replicate; the strains averaged into its gene "
    "fitness are counted upstream and the sheets do not carry the count",
)
_GAP_UNIT = ProvenanceGap(
    field="sample_unit", reason=ProvenanceGapReason.not_carried_by_curation
)


def build_phenotype(fitness: float, arm: str) -> EnvironmentResponsePhenotype:
    """One locus's fitness in one culture."""
    return EnvironmentResponsePhenotype(
        measurement_type=MeasurementType.log2_ratio,
        assay_type=AssayType.pooled_competitive_growth_barcode,
        environment_response=fitness,
        screen_id=arm,
        units=UNITS_FITNESS,
        provenance_gaps=[_GAP_UNCERTAINTY, _GAP_SE, _GAP_N, _GAP_UNIT],
    )


def build_reference(
    dataset_name: str,
    arm: str,
    environment: Environment,
    genome_reference: AssemblyReferenceGenome,
) -> BacterialEnvironmentResponseExperimentReference:
    """The typical gene of the same culture: fitness 0 by the normalization."""
    return BacterialEnvironmentResponseExperimentReference(
        dataset_name=dataset_name,
        genome_reference=genome_reference,
        environment_reference=environment,
        phenotype_reference=EnvironmentResponsePhenotype(
            measurement_type=MeasurementType.log2_ratio,
            assay_type=AssayType.pooled_competitive_growth_barcode,
            environment_response=0.0,
            screen_id=arm,
            units=UNITS_REFERENCE,
        ),
    )


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the repo-root ``.env``."""
    from dotenv import load_dotenv

    load_dotenv()
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/borchertRBTnSeqIdentifiesGenetic2023``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def library_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-library/borchertRBTnSeqIdentifiesGenetic2023``."""
    return Path(data_root or _data_root()) / "torchcell-library" / CITATION_KEY


def _retrieval(
    sha256: str, retrieved_at: str, check: SourceCheck | None
) -> RetrievalRecord:
    return RetrievalRecord(
        method=RetrievalMethod.direct_url,
        source_url=DATA_URL,
        retriever="torchcell.literature.retrieve.elsevier_mmc",
        params={"pii": ARTICLE_PII, "filename": "mmc1.xlsx"},
        sha256=sha256,
        retrieved_at=retrieved_at,
        last_check=check,
    )


def deposit_raw_mirror(
    source: str | Path | None = None,
    *,
    data_root: str | None = None,
    verify_source: bool = False,
    retrieved_at: str = RAW_RETRIEVED_AT,
) -> Path:
    """Deposit ``si1.xlsx`` into the raw mirror with its provenance record.

    ``source`` is an already-retrieved copy (the library mirror's ``si/si1.xlsx`` is
    one); with ``source=None`` the recorded retrieval runs against Elsevier's
    supplementary-file URL. The bytes must hash to ``DATA_SHA256``. Idempotent by
    sha256: an existing mirror file is kept when it matches and refused when it differs.
    ``verify_source`` re-runs the retrieval and records the comparison as ``last_check``.
    """
    from torchcell.literature.retrieve import elsevier_mmc

    root = raw_mirror_dir(data_root)
    dest = root / DATA_RELPATH
    dest.parent.mkdir(parents=True, exist_ok=True)
    staged = root / f".{DATA_FILE}.incoming"
    if source is None:
        staged.write_bytes(elsevier_mmc(ARTICLE_PII, "mmc1.xlsx"))
        source = staged
    sha = file_sha256(source)
    if sha != DATA_SHA256:
        raise RuntimeError(f"{source} hashes to {sha}, the pin is {DATA_SHA256}")
    check: SourceCheck | None = None
    if verify_source:
        produced = hashlib.sha256(elsevier_mmc(ARTICLE_PII, "mmc1.xlsx")).hexdigest()
        if produced != sha:
            raise RuntimeError(
                f"{DATA_URL} now yields sha256 {produced}; the deposited bytes are {sha}"
            )
        check = SourceCheck(
            checked_at=date.today().isoformat(), produced_sha256=produced, matches=True
        )
    elif (root / "manifest.json").exists():
        previous = load_manifest(data_root).files[0].retrieval
        if previous is not None and previous.sha256 == sha:
            check = previous.last_check
    if dest.exists():
        if file_sha256(dest) != sha:
            raise RuntimeError(f"{dest} exists with a different sha256; refusing")
    else:
        shutil.copy2(source, dest)
    if staged.exists():
        staged.unlink()
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=DOI,
        title=TITLE,
        files=[
            ArtifactRecord(
                path=DATA_RELPATH,
                role=ROLE_RAW_DATA,
                bytes=dest.stat().st_size,
                sha256=sha,
                source=DATA_URL,
                original_filename="mmc1.xlsx",
                retrieval=_retrieval(sha, retrieved_at, check),
            )
        ],
        si_data_sources=[f"https://doi.org/{DOI}"],
        si_expected=[
            "si1.xlsx (mmc1.xlsx), Supplementary File 1: Read_Me, Exp1/Exp2 poolcount, "
            "five Figure_*_growth_data sheets and the 13 pairwise comparison sheets",
            "si2.pdf (mmc2.pdf), Supplementary File 2: Tables S1 to S4 and Figures S1 "
            "to S10. NOT deposited here: only its Table S4 is read, as verbatim quotes "
            "against the library mirror's pinned OCR si/si2.md",
        ],
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    return Manifest.model_validate_json(
        (raw_mirror_dir(data_root) / "manifest.json").read_text()
    )


def manifest_sha256(manifest: Manifest, relpath: str = DATA_RELPATH) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the Borchert 2023 raw-mirror manifest")


# --------------------------------------------------------------------------- #
# What is released and NOT stored
# --------------------------------------------------------------------------- #
class NotStored(BaseModel):
    """One released quantity this loader deliberately does not store."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    quantity: str
    where: str
    values: int
    reason: str
    issue: int | None = None


N_BARCODES = 186957
N_POOLCOUNT_COLUMNS = 42


def not_stored(
    matches: Sequence[ArmMatch], comparisons: Sequence[Comparison]
) -> tuple[NotStored, ...]:
    """The released quantities left out, each with its measured size and its reason."""
    subsumed = sum(m.n_shared_loci for m in matches)
    test_rows = sum(len(c.tags) for c in comparisons)
    return (
        NotStored(
            quantity="per-replicate gene fitness of the 4,732 compendium loci",
            where=f"{DATA_FILE} 13 comparison sheets, 42 arms",
            values=subsumed,
            reason="MEASURED identical to the served RbTnseqBorchert2024Dataset records, "
            f"to {COMPENDIUM_HALF_ULP} on every shared locus of every arm",
        ),
        NotStored(
            quantity="the _mean of each arm triple",
            where=f"{DATA_FILE} 13 comparison sheets, 26 columns",
            values=test_rows * 2,
            reason="MEASURED to be the arithmetic mean of the three released replicates "
            f"to {MEAN_TOLERANCE}, so it is a derived column, not a measurement",
        ),
        NotStored(
            quantity="t-statistic, p-value, q-value and adjusted_q-value",
            where=f"{DATA_FILE} 13 comparison sheets, 4 columns each",
            values=test_rows * len(STAT_COLUMNS),
            reason="no phenotype class carries a significance triple; "
            "GeneInteractionPhenotype.gene_interaction_p_value is the only p-value field "
            "in the schema and it is not applicable to a fitness contrast. Storing the "
            "fitness while dropping its significance call is what this refusal avoids",
            issue=776,
        ),
        NotStored(
            quantity="per-barcode read counts",
            where=f"{DATA_FILE} {' and '.join(POOLCOUNT_SHEETS)}",
            values=N_BARCODES * N_POOLCOUNT_COLUMNS,
            reason="raw sequencing read tallies per barcode, which no loader ingests; "
            "the perturbation leaf for the barcode grain EXISTS "
            "(TransposonInsertionPerturbation.barcode / insertion_position / "
            "insertion_strand), so the gap is on the phenotype side: no class and no "
            "MeasurementType member holds a read count",
        ),
        NotStored(
            quantity="OD600 growth curves of the deletion and overexpression strains",
            where=f"{DATA_FILE} {', '.join(GROWTH_SHEETS)}",
            values=0,
            reason="back-scattered 620 nm light in arbitrary units over time: "
            "MeasurementType has no optical-density member and no phenotype carries a "
            "time series",
            issue=776,
        ),
        NotStored(
            quantity="OD600 at the time of sampling, per replicate",
            where="si/si2.md Table S4",
            values=N_ARMS,
            reason="the same missing MeasurementType member; the table's growth duration "
            "IS stored, as Environment.duration_hours",
            issue=776,
        ),
    )


# --------------------------------------------------------------------------- #
# Dataset
# --------------------------------------------------------------------------- #
#: Records of the full build: the 271 compendium-absent loci over the arms carrying them.
EXPECTED_RECORDS = 10824
#: Distinct loci the compendium eliminated that this dataset recovers.
EXPECTED_NEW_LOCI = 271
#: The dev-tree root of the already-served compendium this build is partitioned against.
SERVED_ROOT_REL = "data/torchcell/rbtnseq_borchert2024"


@register_dataset
class RbTnseqBorchert2023Dataset(ExperimentDataset):
    """Borchert 2023 KT2440 RB-TnSeq fitness for the 271 loci the compendium dropped."""

    REFERENCE_STRAIN: ClassVar[Literal["KT2440"]] = "KT2440"

    def __init__(
        self,
        root: str = "data/torchcell/rbtnseq_borchert2023",
        io_workers: int = 0,
        pputida_genome: PPutidaKT2440Genome | None = None,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``pputida_genome`` is injected by the build entry points."""
        self.pputida_genome = pputida_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialEnvironmentResponseExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialEnvironmentResponseExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The one released file this dataset reads."""
        return [DATA_FILE]

    def download(self) -> None:
        """Link the mirror's Supplementary File 1 into ``raw/`` and verify its sha256."""
        data_root = _data_root()
        manifest = load_manifest(data_root)
        recorded = manifest_sha256(manifest)
        check_manifest_pin(DATA_RELPATH, recorded, DATA_SHA256)
        src = raw_mirror_dir(data_root) / DATA_RELPATH
        if not src.exists():
            raise RuntimeError(f"required raw artifact missing from mirror: {src}")
        os.makedirs(self.raw_dir, exist_ok=True)
        link_verified(src, osp.join(self.raw_dir, DATA_FILE), recorded)

    def _genome(self) -> PPutidaKT2440Genome:
        if self.pputida_genome is None:  # a direct run; build entry points inject it
            self.pputida_genome = bacterial_genome("pputida", self.REFERENCE_STRAIN)
        return self.pputida_genome

    def _pointer(self, obj: Any, hint: str, itxn: Any) -> Any:
        """``obj``'s dump, interned exactly as ``_intern_record`` interns it."""
        holder = {"value": obj.model_dump()}
        self._maybe_intern(holder, "value", obj, hint, itxn)
        return holder["value"]

    def _compendium(self) -> b24.Release:
        """The Borchert 2024 release the served records were built from, sha256-pinned."""
        data_root = _data_root()
        recorded = b24.manifest_sha256(b24.load_manifest(data_root))
        check_manifest_pin(b24.DATA_RELPATH, recorded, b24.DATA_SHA256)
        path = b24.raw_mirror_dir(data_root) / b24.DATA_RELPATH
        if file_sha256(path) != b24.DATA_SHA256:
            raise RuntimeError(f"{path} does not match the Borchert 2024 pin")
        return b24.read_release(path)

    @post_process
    def process(self) -> None:
        """Prove the partition against Borchert 2024, then write what is new."""
        verify_raw_files(self.raw_dir, {DATA_FILE: DATA_SHA256})
        comparisons = read_comparisons(osp.join(self.raw_dir, DATA_FILE))
        release = self._compendium()
        matches = match_arms(comparisons, release)
        proofs = assert_arms_agree_across_sheets(comparisons)
        proofs.extend(assert_means_are_derived(comparisons))
        by_arm = new_loci_by_arm(comparisons, release.genes)
        tags = sorted({tag for loci in by_arm.values() for tag in loci})
        if len(tags) != EXPECTED_NEW_LOCI:
            raise RuntimeError(
                f"{len(tags)} loci are absent from the compendium, {EXPECTED_NEW_LOCI} "
                "were measured"
            )
        partition = assert_served_partition(
            osp.join(_data_root(), SERVED_ROOT_REL), tags, matches, release
        )
        genome = self._genome()
        stored, reconciliation = reconcile_locus_tags(
            genome, pd.Series(tags), label=self.name
        )
        reconciliation.require_resolved(1.0)
        names = b24.perturbed_gene_names(genome, list(stored))
        genome_reference = assembly_reference(self.REFERENCE_STRAIN)
        by_name = {sample.exp_name: sample for sample in release.samples}
        publication = Publication(doi=DOI, doi_url=f"https://doi.org/{DOI}")

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        genotypes: dict[str, dict[str, Any]] = {}
        idx = 0
        for experiment in tqdm(EXPERIMENTS, desc=self.name):
            comparison = next(c for c in comparisons if c.sheet == experiment.sheets[0])
            row_of = {tag: i for i, tag in enumerate(comparison.tags)}
            for position, replicate in enumerate(REPLICATES):
                arm = experiment.arm(replicate)
                assert_dose_agrees(experiment, by_name[experiment.compendium[position]])
                environment = build_environment(experiment, replicate)
                reference = build_reference(
                    self.name, arm, environment, genome_reference
                )
                values = comparison.column(experiment.column(replicate))
                with (
                    env.begin(write=True) as txn,
                    interned_env.begin(write=True) as itxn,
                ):
                    env_ptr = self._pointer(environment, environment.media.name, itxn)
                    ref_ptr = self._pointer(reference, reference.dataset_name, itxn)
                    pub_ptr = self._pointer(publication, "publication", itxn)
                    for tag in by_arm[arm]:
                        if tag not in genotypes:
                            genotypes[tag] = b24.build_genotype(
                                tag, names[tag], MUTANT_LIBRARY
                            ).model_dump()
                        phenotype = build_phenotype(float(values[row_of[tag]]), arm)
                        record = {
                            "experiment": {
                                "experiment_type": "bacterial_environment_response",
                                "dataset_name": self.name,
                                "genotype": genotypes[tag],
                                "environment": env_ptr,
                                "phenotype": phenotype.model_dump(),
                            },
                            "reference": ref_ptr,
                            "publication": pub_ptr,
                        }
                        if idx == 0:
                            self._check_record(
                                record,
                                genotypes[tag],
                                environment,
                                phenotype,
                                reference,
                                publication,
                                itxn,
                            )
                        txn.put(f"{idx}".encode(), pickle.dumps(record))
                        idx += 1
        env.close()
        interned_env.close()
        if idx != EXPECTED_RECORDS:
            raise RuntimeError(f"wrote {idx} records, {EXPECTED_RECORDS} were measured")
        self._write_reports(
            matches, by_arm, comparisons, proofs, partition, reconciliation
        )
        log.info("Wrote %d %s records to LMDB", idx, self.name)

    def _check_record(
        self,
        record: dict[str, Any],
        genotype: dict[str, Any],
        environment: Environment,
        phenotype: EnvironmentResponsePhenotype,
        reference: BacterialEnvironmentResponseExperimentReference,
        publication: Publication,
        itxn: Any,
    ) -> None:
        """The fast-path record equals what ``_intern_record`` writes for the same objects."""
        experiment = BacterialEnvironmentResponseExperiment(
            dataset_name=self.name,
            genotype=Genotype.model_validate(genotype),
            environment=environment,
            phenotype=phenotype,
        )
        expected = pickle.loads(
            self._intern_record(experiment, reference, publication, itxn)
        )
        if expected != record:
            raise AssertionError(
                f"{self.name}: assembled record differs from _intern_record's"
            )

    def _write_reports(
        self,
        matches: Sequence[ArmMatch],
        by_arm: Mapping[str, Sequence[str]],
        comparisons: Sequence[Comparison],
        proofs: Sequence[str],
        partition: ServedPartition,
        reconciliation: Any,
    ) -> None:
        """Write the subsumption, partition, not-stored and identifier reports."""
        reports: dict[str, Any] = {
            "subsumption.json": {
                "claim": "every Borchert 2023 arm IS a served Borchert 2024 sample",
                "half_ulp": COMPENDIUM_HALF_ULP,
                "runner_up_floor": RUNNER_UP_FLOOR,
                "arms": [m.model_dump() for m in matches],
                "internal_proofs": list(proofs),
            },
            "served_partition.json": partition.model_dump(),
            "new_loci.json": {
                "n_loci": len({tag for loci in by_arm.values() for tag in loci}),
                "n_records": sum(len(loci) for loci in by_arm.values()),
                "by_arm": {arm: list(loci) for arm, loci in by_arm.items()},
            },
            "not_stored.json": {
                "released_rows_per_sheet": {c.sheet: len(c.tags) for c in comparisons},
                "quantities": [
                    n.model_dump() for n in not_stored(matches, comparisons)
                ],
                "read_me_conflict": READ_ME_4HBALD_CONFLICT,
            },
            "locus_tag_reconciliation.json": reconciliation.model_dump(mode="json"),
        }
        for name, payload in reports.items():
            with open(osp.join(self.preprocess_dir, name), "w") as handle:
                json.dump(payload, handle, indent=2)

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process() for this dataset."""
        raise NotImplementedError(
            "RbTnseqBorchert2023Dataset builds records in process()"
        )


def verify_build(
    dataset_root: str,
    *,
    genome: PPutidaKT2440Genome | None = None,
    data_root: str | None = None,
    expected_count: int = EXPECTED_RECORDS,
) -> VerificationReport:
    """Run the environment-response L0-L4 verifier on a built tree and write its report.

    Every record is checked against the KT2440 genome its references pin, with the L4
    universe every GenBank locus of the assembly. The report is written to
    ``preprocess/verification_report.json``.
    """
    from torchcell.verification.environment_response import (
        verify_environment_response_dataset_streaming,
    )
    from torchcell.verification.runners import stream_records

    if genome is None:
        genome = bacterial_genome("pputida", "KT2440", data_root)
    report = verify_environment_response_dataset_streaming(
        stream_records(dataset_root),
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=Provenance(
            source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{DATA_RELPATH}",
            citation_key=CITATION_KEY,
            sha256=DATA_SHA256,
            method="si1.xlsx (Elsevier mmc1.xlsx): one "
            "BacterialEnvironmentResponseExperiment per (locus, culture) for the 271 "
            "loci the Borchert 2024 compendium eliminated, log2_ratio gene fitness",
            page=f"{N_COMPARISONS} pairwise comparison sheets",
            retrieved=RAW_RETRIEVED_AT,
        ),
        expected_count=expected_count,
        sgd_genes=set(genome.genbank.loci),
        resolve_gene_name=genome.resolve_gene_name,
    )
    preprocess = osp.join(dataset_root, "preprocess")
    os.makedirs(preprocess, exist_ok=True)
    with open(osp.join(preprocess, "verification_report.json"), "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main() -> None:
    """Build the dataset under ``DATA_ROOT`` and print its length and first record."""
    dataset = RbTnseqBorchert2023Dataset(
        root=osp.join(_data_root(), "data/torchcell/rbtnseq_borchert2023")
    )
    print(f"len = {len(dataset)}")
    print(dataset[0])


if __name__ == "__main__":
    main()
