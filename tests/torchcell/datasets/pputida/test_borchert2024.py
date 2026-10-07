# tests/torchcell/datasets/pputida/test_borchert2024.py
# [[tests.torchcell.datasets.pputida.test_borchert2024]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/pputida/test_borchert2024.py
"""Borchert 2024 fModules loader (``torchcell.datasets.pputida.borchert2024``).

Synthetic tests write a small workbook with the release's exact sheet and column layout
(three genes, one sample per kind of condition the loader handles) and run the reader,
the attribution and drop rules, the environment, phenotype and genotype builders, the
raw-mirror deposit and a full ``process()`` build into ``tmp_path``. The genome, the
locus-tag reconciliation and the assembly pin are replaced by in-test objects, so nothing
reads ``$DATA_ROOT``.

The ``@pytest.mark.data`` tests read the real raw mirror, the paper mirrors and (when it
has been built) the dev-tree LMDB, and pin the measured numbers of the dendron note: 332
samples, 4,732 genes, 254 Fitness Browser plus 78 Beckham samples, the per-study coverage
(Thompson 2020 46 samples over 23 conditions, Schmidt 2022 123 over 71, Borchert 2023 42),
the 42 dropped samples and the 1,372,280 records.
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
import pickle
from pathlib import Path
from typing import Any

import lmdb
import numpy as np
import pandas as pd
import pytest

import torchcell.datasets.pputida.borchert2024 as bt
from torchcell.data import file_sha256
from torchcell.datamodels.media import MOPS_MINIMAL
from torchcell.datamodels.schema import (
    AssemblyReferenceGenome,
    EnvironmentPhysicalPerturbation,
    PhysicalFactor,
    SmallMoleculePerturbation,
)
from torchcell.datasets.bacteria_common import LocusTagReconciliation
from torchcell.sequence.genome.base import GeneNameResolution, GeneNameStatus

GENES = ("PP_0001", "PP_0002", "PP_0003")


def _sample_row(
    exp_name: str,
    set_name: str,
    exp_desc: str,
    *,
    group: str = "carbon source",
    library: str = "Putida_ML5_JBEI",
    media: str = "MOPS minimal media_noCarbon",
    c1: tuple[str, float, str] | None = ("D-Glucose", 10.0, "mM"),
    c2: tuple[str, float, str] | None = None,
) -> list[Any]:
    """A metadata row in METADATA_COLUMNS order."""
    cell: dict[str, Any] = {name: None for name in bt.METADATA_COLUMNS}
    cell.update(
        orgId="Putida",
        expName=exp_name,
        set=set_name,
        expDesc=exp_desc,
        expGroup=group,
        total_rep=2,
        rep=1,
        mutantLibrary=library,
        person="test",
        media=media,
        temperature=30,
        aerobic="Aerobic",
        liquid="Liquid",
    )
    if c1 is not None:
        cell.update(condition_1=c1[0], concentration_1=c1[1], units_1=c1[2])
    if c2 is not None:
        cell.update(condition_2=c2[0], concentration_2=c2[1], units_2=c2[2])
    return [cell[name] for name in bt.METADATA_COLUMNS]


#: One sample per kind of condition: a Thompson 2020 alcohol, a Schmidt nitrogen source
#: and amino-acid drop-out, a Borchert 2023 stressor, the release's mojibake label, a DMSO
#: vehicle, and the two drop rules.
SAMPLES: list[list[Any]] = [
    _sample_row("set16IT001", "set16", "Butanol (C)", c1=("Butanol", 10.0, "mM")),
    _sample_row(
        "set27IT001",
        "set27",
        "L-Serine (N)",
        group="nitrogen source",
        media="MOPS minimal media_Glucose_noNitrogen",
        c1=("L-Serine", 10.0, "mM"),
    ),
    _sample_row(
        "set10IT001",
        "set10",
        "D-Glucose (C) with 20AA_mix_minus_Ala",
        c2=("20AA_mix_minus_Ala", 0.5, "X"),
    ),
    _sample_row(
        "set100IT001",
        "set100",
        "D-Glucose with Vanillin (C)",
        media="M9_medium",
        c1=("D-Glucose", 20.0, "mM"),
        c2=("Vanillin", 10.0, "mM"),
    ),
    _sample_row(
        "set101IT001",
        "set101",
        "D-Glucose with ketoadipate (C)",
        media="M9_medium",
        c1=("D-Glucose", 20.0, "mM"),
        c2=("ÃŸ-ketoadipate", 125.0, "mM"),
    ),
    _sample_row(
        "set12IT001",
        "set12",
        "Vanillin (C)",
        c1=("Vanillin", 10.0, "mM"),
        c2=("Dimethyl Sulfoxide", 1.0, "vol%"),
    ),
    _sample_row(
        "set5IT001",
        "set5",
        "D-Glucose (C)",
        library="Putida_ML5",
        media="RCH2_defined_noCarbon",
        c1=("D-Glucose", 20.0, "mM"),
    ),
    _sample_row(
        "set8IT001",
        "set8",
        "24hr timepoint from 1L M9/1% dextrose BATCH fermentation",
        group="reactor",
        media="M9_1percentGlucose",
        c1=None,
    ),
]
#: fitness and t per (gene, sample): column 3 (set100, Beckham) carries an inverted t.
FITNESS = np.array(
    [
        [-1.5, 0.2, -3.0, -2.0, 0.4, 0.0, 1.0, 1.0],
        [0.3, -0.004, 0.5, 0.6, -0.7, 2.0, 1.0, 1.0],
        [0.0, 1.25, -0.8, -0.9, 0.1, -0.5, 1.0, 1.0],
    ]
)
T_STAT = np.array(
    [
        [-7.5, 1.0, -10.0, 4.0, 1.5, 0.0, 2.0, 2.0],
        [1.5, -0.02, 2.5, -3.0, -3.5, 10.0, 2.0, 2.0],
        [0.0, 6.25, -4.0, 4.5, 0.5, -2.5, 2.0, 2.0],
    ]
)


def write_workbook(path: Path) -> Path:
    """The synthetic release, in the pinned sheet and column layout."""
    import openpyxl

    workbook = openpyxl.Workbook()
    meta = workbook.active
    meta.title = bt.METADATA_SHEET
    meta.append(list(bt.METADATA_COLUMNS))
    for row in SAMPLES:
        meta.append(row)
    headers = [f"{row[1]} {row[3]}" for row in SAMPLES]
    for name, values in ((bt.FITNESS_SHEET, FITNESS), (bt.T_SHEET, T_STAT)):
        sheet = workbook.create_sheet(name)
        sheet.append([*bt.GENE_COLUMNS, *headers])
        for gene, row_values in zip(GENES, values, strict=True):
            sheet.append(["Putida", gene, gene, None, "desc", *row_values.tolist()])
    workbook.save(path)
    return path


@pytest.fixture
def release(tmp_path: Path) -> bt.Release:
    return bt.read_release(write_workbook(tmp_path / bt.DATA_FILE))


def _sample(release: bt.Release, exp_name: str) -> bt.SampleMetadata:
    return next(s for s in release.samples if s.exp_name == exp_name)


# --------------------------------------------------------------------------- #
# Reader
# --------------------------------------------------------------------------- #
def test_read_release_parses_the_pinned_layout(release: bt.Release) -> None:
    assert release.genes == GENES
    assert [s.exp_name for s in release.samples] == [row[1] for row in SAMPLES]
    np.testing.assert_array_equal(release.fitness, FITNESS)
    np.testing.assert_array_equal(release.t, T_STAT)
    vanillin = _sample(release, "set100IT001")
    assert vanillin.condition_2 == "Vanillin"
    assert vanillin.concentration_2 == 10.0
    assert vanillin.is_beckham
    assert not _sample(release, "set16IT001").is_beckham


def test_read_release_refuses_a_value_column_out_of_metadata_order(
    tmp_path: Path,
) -> None:
    import openpyxl

    path = write_workbook(tmp_path / bt.DATA_FILE)
    workbook = openpyxl.load_workbook(path)
    header = workbook[bt.T_SHEET]["F1"]
    header.value = "set999IT001 Wrong (C)"
    workbook.save(path)
    with pytest.raises(ValueError, match="value columns are not the metadata samples"):
        bt.read_release(path)


def test_read_release_refuses_an_empty_cell(tmp_path: Path) -> None:
    import openpyxl

    path = write_workbook(tmp_path / bt.DATA_FILE)
    workbook = openpyxl.load_workbook(path)
    workbook[bt.FITNESS_SHEET]["F2"].value = None
    workbook.save(path)
    with pytest.raises(ValueError, match="has an empty cell"):
        bt.read_release(path)


# --------------------------------------------------------------------------- #
# Attribution and drops
# --------------------------------------------------------------------------- #
def test_attribution_follows_the_quoted_source_studies(release: bt.Release) -> None:
    studies = {s.exp_name: bt.attribute_sample(s).study for s in release.samples}
    assert studies == {
        "set16IT001": "thompson2020",
        "set27IT001": "schmidt2022",
        "set10IT001": "schmidt2022",
        "set100IT001": "borchert2023",
        "set101IT001": "borchert2024",
        "set12IT001": "borchert2024",
        "set5IT001": "borchert2024",
        "set8IT001": "borchert2024",
    }
    compendium_only = bt.attribute_sample(_sample(release, "set101IT001"))
    assert compendium_only.basis == "compendium_release"
    assert compendium_only.hypothesis == bt.HYPOTHESES["set101"]
    for attribution in map(bt.attribute_sample, release.samples):
        for key in attribution.evidence:
            assert key in bt.SOURCED_VALUES


def test_set12_thompson_conditions_and_the_set101_protocatechuate_pair() -> None:
    def sample(set_name: str, c1: str, c2: str | None) -> bt.SampleMetadata:
        return bt.SampleMetadata(
            exp_name=f"{set_name}IT099",
            set_name=set_name,
            exp_desc="x",
            exp_group="carbon source",
            mutant_library="Putida_ML5_JBEI",
            person="p",
            media="M9_medium",
            temperature=30.0,
            aerobic="Aerobic",
            total_rep=3,
            rep=1,
            condition_1=c1,
            units_1="mM",
            concentration_1=10.0,
            condition_2=c2,
            units_2="mM" if c2 else None,
            concentration_2=30.0 if c2 else None,
        )

    assert bt.attribute_sample(sample("set12", "1,4-Butanediol", None)).study == (
        "thompson2020"
    )
    assert bt.attribute_sample(sample("set12", "Ferulic Acid", None)).study == (
        "borchert2024"
    )
    assert bt.attribute_sample(
        sample("set101", "D-Glucose", "Protecatechuic acid")
    ).study == ("borchert2023")
    assert bt.attribute_sample(sample("set101", "D-Glucose", None)).study == (
        "borchert2023"
    )
    assert bt.attribute_sample(
        sample("set100", "D-Glucose", "cis,cis-muconate")
    ).study == ("borchert2024")


def test_drop_rules(release: bt.Release) -> None:
    reasons = {s.exp_name: bt.drop_reason(s) for s in release.samples}
    assert reasons.pop("set5IT001") == bt.DROP_MEDIUM_NOT_IN_LIBRARY
    assert reasons.pop("set8IT001") == bt.DROP_REACTOR_PROCESS
    assert set(reasons.values()) == {None}
    assert set(bt.DROP_RULES) == {
        bt.DROP_MEDIUM_NOT_IN_LIBRARY,
        bt.DROP_REACTOR_PROCESS,
    }


def test_condition_key_ignores_the_description_spelling(release: bt.Release) -> None:
    a = _sample(release, "set100IT001")
    b = a.model_copy(update={"exp_desc": a.exp_desc.replace(" (C)", "  (C)")})
    assert bt.condition_key(a) == bt.condition_key(b)


# --------------------------------------------------------------------------- #
# Environment
# --------------------------------------------------------------------------- #
def test_carbon_source_sample(release: bt.Release) -> None:
    environment = bt.build_environment(_sample(release, "set16IT001"))
    assert environment.media is MOPS_MINIMAL
    assert environment.temperature is not None
    assert environment.temperature.value == 30.0
    (perturbation,) = environment.perturbations
    assert isinstance(perturbation, EnvironmentPhysicalPerturbation)
    assert perturbation.factor is PhysicalFactor.carbon_source
    assert perturbation.magnitude is not None
    assert perturbation.magnitude.value == 10.0
    assert environment.gapped_fields() == {"duration_hours", "duration_generations"}


def test_nitrogen_source_sample_uses_the_derived_mops(release: bt.Release) -> None:
    environment = bt.build_environment(_sample(release, "set27IT001"))
    assert environment.media is bt.FB_MOPS_GLUCOSE_NO_NITROGEN
    assert environment.media.base_medium == "MOPS_MINIMAL"
    (perturbation,) = environment.perturbations
    assert isinstance(perturbation, EnvironmentPhysicalPerturbation)
    assert perturbation.factor is PhysicalFactor.nitrogen_source
    assert perturbation.agent is not None
    assert perturbation.agent.name == "L-serine"
    component_names = [c.compound.name for c in environment.media.components]
    assert "ammonium chloride" not in component_names
    assert "D-glucose" in component_names
    assert [c.name for c in environment.media.dropouts] == ["ammonium chloride"]


def test_amino_acid_dropout_sample(release: bt.Release) -> None:
    perturbations = bt.build_environment(_sample(release, "set10IT001")).perturbations
    assert [p.perturbation_type for p in perturbations] == [
        "environment_physical",
        "small_molecule",
        "environment_physical",
    ]
    carbon, mix, dropout = perturbations
    assert isinstance(mix, SmallMoleculePerturbation)
    assert mix.compound.name == bt.AMINO_ACID_MIX
    assert mix.compound.gapped_fields() == {"inchikey"}
    assert isinstance(dropout, EnvironmentPhysicalPerturbation)
    assert dropout.factor is PhysicalFactor.nutrient_dropout
    assert dropout.agent is not None
    assert dropout.agent.name == "L-alanine"


def test_stressor_and_label_fix(release: bt.Release) -> None:
    stressed = bt.build_environment(_sample(release, "set100IT001"))
    assert stressed.media is bt.BORCHERT2023_M9
    glucose, vanillin = stressed.perturbations
    assert isinstance(vanillin, SmallMoleculePerturbation)
    assert vanillin.compound.name == "vanillin"
    assert vanillin.concentration.value == 10.0
    fixed = bt.build_environment(_sample(release, "set101IT001")).perturbations[1]
    assert isinstance(fixed, SmallMoleculePerturbation)
    assert fixed.compound.name == "beta-ketoadipate"
    assert bt.condition_compound("Protecatechuic acid").name == (
        bt.condition_compound("Protocatechuic Acid").name
    )


def test_dmso_vehicle_is_a_percent_small_molecule(release: bt.Release) -> None:
    vehicle = bt.build_environment(_sample(release, "set12IT001")).perturbations[1]
    assert isinstance(vehicle, SmallMoleculePerturbation)
    assert vehicle.compound.name == "dimethyl sulfoxide"
    assert vehicle.concentration.unit == "percent_v/v"


def test_borchert2023_m9_replaces_the_base_phosphate() -> None:
    media = bt.BORCHERT2023_M9
    assert media.base_medium == "M9"
    amounts = {
        c.compound.name: c.concentration.value
        for c in media.components
        if c.concentration
    }
    assert amounts["dipotassium hydrogen phosphate"] == 3.0
    assert amounts["ammonium chloride"] == 1.66
    assert [c.name for c in media.dropouts] == ["potassium dihydrogen phosphate"]


# --------------------------------------------------------------------------- #
# Phenotype and genotype
# --------------------------------------------------------------------------- #
def test_derived_se_removes_the_normalization_floor() -> None:
    # q = 0.2, se = sqrt(0.04 - 0.01)
    assert bt.derived_se(-1.5, -7.5) == pytest.approx(math.sqrt(0.03))
    # q = 0.5, se = sqrt(0.25 - 0.01)
    assert bt.derived_se(1.0, 2.0) == pytest.approx(math.sqrt(0.24))
    assert bt.derived_se(0.0, 0.0) is None
    # q = 0.09 sits below the floor: only rounding puts it there
    assert bt.derived_se(0.009, 0.1) is None
    # bound (0.0005/0.004 + 0.0005/0.02) * 0.04 / 0.03 = 0.2 > 0.05
    assert bt.derived_se(-0.004, -0.02) is None
    # near the floor the subtraction amplifies rounding: q = 0.105,
    # (0.0005/0.105 + 0.0005) * 0.011025 / 0.001025 = 0.057 > 0.05
    assert bt.derived_se(0.105, 1.0) is None


def test_floor_violations_and_the_identity_check(release: bt.Release) -> None:
    f = np.array([[0.5, 0.009, 0.0], [0.05, 0.3, 1.0]])
    t = np.array([[2.5, 0.2, 0.0], [0.4, 1.5, 0.0]])
    # 0.009 / 0.2 = 0.045 and 0.05 / 0.4 = 0.125: one breach; zeros are skipped
    assert bt.floor_violations(f, t) == 1
    # kept Fitness Browser columns set16, set27, set10, set12: 12 pairs, 2 of them zero
    assert bt.check_identity_floor(release) == 10
    broken = release.model_copy(update={"t": release.t * 100})
    with pytest.raises(ValueError, match="below"):
        bt.check_identity_floor(broken)


def test_fitness_browser_phenotype(release: bt.Release) -> None:
    sample = _sample(release, "set16IT001")
    phenotype = bt.build_phenotype(-1.5, -7.5, sample)
    assert phenotype.environment_response == -1.5
    assert phenotype.environment_response_se == pytest.approx(math.sqrt(0.03))
    assert phenotype.environment_response_uncertainty is None
    assert phenotype.screen_id == "set16IT001"
    assert phenotype.units == bt.UNITS_FITNESS_BROWSER
    assert phenotype.gapped_fields() == {
        "environment_response_uncertainty",
        "n_samples",
        "sample_unit",
    }
    rounded = bt.build_phenotype(-0.004, -0.02, sample)
    assert rounded.environment_response_se is None
    assert "environment_response_se" in rounded.gapped_fields()


def test_beckham_phenotype_derives_no_se(release: bt.Release) -> None:
    phenotype = bt.build_phenotype(-2.0, 4.0, _sample(release, "set100IT001"))
    assert phenotype.environment_response == -2.0
    assert phenotype.environment_response_se is None
    assert phenotype.units == bt.UNITS_BECKHAM
    assert "environment_response_se" in phenotype.gapped_fields()


def test_negative_fitness_is_not_clamped(release: bt.Release) -> None:
    phenotype = bt.build_phenotype(-12.837, -6.137, _sample(release, "set16IT001"))
    assert phenotype.environment_response == -12.837


def test_genotype_is_a_kt2440_transposon_insertion() -> None:
    (perturbation,) = bt.build_genotype(
        "PP_0292", "hisA", "Putida_ML5_JBEI"
    ).perturbations
    assert perturbation.perturbation_type == "transposon_insertion"
    assert perturbation.systematic_gene_name == "PP_0292"
    assert perturbation.perturbed_gene_name == "hisA"
    assert perturbation.gene_namespace == "pputida_kt2440_locus_tag"
    assert perturbation.library_pool == "Putida_ML5_JBEI"
    assert perturbation.transposon == bt.TRANSPOSON
    assert perturbation.barcode is None


class _Locus:
    def __init__(self, symbol: str | None) -> None:
        self.symbol = symbol


class _Annotation:
    def __init__(self, loci: dict[str, _Locus]) -> None:
        self.loci = loci


class FakeGenome:
    """``genbank.loci[tag].symbol`` and ``resolve_gene_name`` over GENES.

    PP_0001 carries a unique symbol, PP_0002 a symbol shared with another locus, PP_0003
    none.
    """

    def __init__(self) -> None:
        """Three loci: a unique symbol, a shared symbol, no symbol."""
        self.genbank = _Annotation(
            {
                "PP_0001": _Locus("parB"),
                "PP_0002": _Locus("dup"),
                "PP_0003": _Locus(None),
            }
        )

    def resolve_gene_name(self, name: str) -> GeneNameResolution:
        """``parB`` renames PP_0001, ``dup`` is ambiguous, a tag is itself."""
        if name == "parB":
            return GeneNameResolution(
                input_name=name,
                status=GeneNameStatus.RENAMED,
                systematic_name="PP_0001",
                note="gene symbol of PP_0001",
            )
        if name == "dup":
            return GeneNameResolution(
                input_name=name,
                status=GeneNameStatus.AMBIGUOUS,
                systematic_name=None,
                candidates=["PP_0002", "PP_0999"],
            )
        return GeneNameResolution(
            input_name=name, status=GeneNameStatus.CURRENT, systematic_name=name
        )


def test_perturbed_gene_names_use_a_symbol_only_when_it_resolves_back() -> None:
    names = bt.perturbed_gene_names(FakeGenome(), list(GENES))  # type: ignore[arg-type]
    assert names == {"PP_0001": "parB", "PP_0002": "PP_0002", "PP_0003": "PP_0003"}


def test_t_sign_agreement_flags_an_inverted_set(release: bt.Release) -> None:
    agreement = bt.t_sign_agreement(release)
    assert agreement["set16"]["sign_disagreement"] == 0.0
    # set100: f (-2.0, 0.6, -0.9) against t (4.0, -3.0, 4.5) disagree on all three
    assert agreement["set100"]["sign_disagreement"] == 1.0
    assert agreement["set100"]["median_corr_f_t"] < 0


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def test_deposit_is_idempotent_and_refuses_other_bytes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = write_workbook(tmp_path / "source.xlsx")
    sha = file_sha256(source)
    monkeypatch.setattr(bt, "DATA_SHA256", sha)
    data_root = str(tmp_path / "root")
    root = bt.deposit_raw_mirror(source, data_root=data_root)
    manifest = bt.load_manifest(data_root)
    assert bt.manifest_sha256(manifest) == sha
    (record,) = manifest.files
    assert record.retrieval is not None
    assert record.retrieval.retriever == "torchcell.literature.retrieve.direct_url"
    assert record.retrieval.params == {"url": bt.DATA_URL}
    assert bt.DATA_COMMIT in record.retrieval.source_url  # type: ignore[operator]
    bt.deposit_raw_mirror(source, data_root=data_root)
    assert file_sha256(root / bt.DATA_RELPATH) == sha
    other = tmp_path / "other.xlsx"
    other.write_bytes(b"not the release")
    with pytest.raises(RuntimeError, match="the pin is"):
        bt.deposit_raw_mirror(other, data_root=data_root)
    (root / bt.DATA_RELPATH).write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="different sha256"):
        bt.deposit_raw_mirror(source, data_root=data_root)


# --------------------------------------------------------------------------- #
# End-to-end build on the synthetic release
# --------------------------------------------------------------------------- #
def _reference() -> AssemblyReferenceGenome:
    return AssemblyReferenceGenome(
        species="Pseudomonas putida",
        strain="KT2440",
        ploidy="haploid",
        assembly_set="pputida_KT2440_ASM756v2",
        assembly_accession="GCA_000007565.2",
    )


def _identity_reconciliation(
    genome: Any, names: pd.Series, *, label: str
) -> tuple[pd.Series, LocusTagReconciliation]:
    return names, LocusTagReconciliation(
        label=label,
        assembly_set="pputida_KT2440_ASM756v2",
        gene_namespace="pputida_kt2440_locus_tag",
        unique_names=len(names),
        status_histogram={
            s: (len(names) if s is GeneNameStatus.CURRENT else 0)
            for s in GeneNameStatus
        },
        layer_histogram={"locus tag": len(names)},
        remapped=0,
        kept_on_collision=(),
        retired_kept=(),
        ambiguous_kept={},
        case_insensitive=(),
        outside_namespace=(),
    )


def test_process_builds_one_record_per_gene_and_kept_sample(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "rbtnseq_borchert2024"
    raw = root / "raw"
    raw.mkdir(parents=True)
    path = write_workbook(raw / bt.DATA_FILE)
    monkeypatch.setattr(bt, "DATA_SHA256", file_sha256(path))
    monkeypatch.setattr(bt, "N_GENES", len(GENES))
    monkeypatch.setattr(bt, "N_SAMPLES", len(SAMPLES))
    monkeypatch.setattr(bt, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(bt, "assembly_reference", lambda strain: _reference())

    dataset = bt.RbTnseqBorchert2024Dataset(
        root=str(root),
        pputida_genome=FakeGenome(),  # type: ignore[arg-type]
    )
    kept = [
        i for i, row in enumerate(SAMPLES) if row[1] not in ("set5IT001", "set8IT001")
    ]
    assert len(dataset) == len(GENES) * len(kept)

    first = dataset.transform_item(dataset[0])
    (perturbation,) = first["experiment"].genotype.perturbations
    assert perturbation.systematic_gene_name == "PP_0001"
    assert perturbation.perturbed_gene_name == "parB"
    assert first["experiment"].phenotype.environment_response == FITNESS[0, 0]
    assert first["experiment"].phenotype.screen_id == "set16IT001"
    assert first["reference"].phenotype_reference.environment_response == 0.0
    assert first["publication"].doi == bt.SOURCE_STUDIES["thompson2020"].doi
    index = dataset.experiment_reference_index
    assert index is not None
    assert len(index) == len(kept)

    values = sorted(
        dataset[i]["experiment"]["phenotype"]["environment_response"]
        for i in range(len(dataset))
    )
    expected = sorted(FITNESS[:, kept].ravel().tolist())
    assert values == expected

    preprocess = Path(dataset.preprocess_dir)
    dropped = json.loads((preprocess / "dropped_records.json").read_text())
    assert dropped["kept_records"] == len(dataset)
    assert {k: v["n_records_not_written"] for k, v in dropped["by_rule"].items()} == {
        bt.DROP_MEDIUM_NOT_IN_LIBRARY: len(GENES),
        bt.DROP_REACTOR_PROCESS: len(GENES),
    }
    studies = json.loads((preprocess / "source_studies.json").read_text())
    assert {c["study"]: c["n_samples"] for c in studies["coverage"]} == {
        "thompson2020": 1,
        "schmidt2022": 2,
        "borchert2023": 1,
        "borchert2024": 4,
    }
    uncertainty = json.loads((preprocess / "uncertainty.json").read_text())
    assert sum(uncertainty["record_counts"].values()) == len(dataset)
    assert uncertainty["record_counts"]["beckham"] == 2 * len(GENES)
    assert (preprocess / "build_manifest.json").exists()
    dataset.close_lmdb()

    report = bt.verify_build(
        str(root),
        genome=FakeGenome(),  # type: ignore[arg-type]
        expected_count=len(GENES) * len(kept),
    )
    failed = [(r.level, r.name, r.message) for r in report.results if not r.passed]
    assert failed == []
    assert (preprocess / "verification_report.json").exists()


# --------------------------------------------------------------------------- #
# Real data
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    from dotenv import load_dotenv

    load_dotenv()
    return os.environ["DATA_ROOT"]


def _real_release() -> bt.Release:
    path = bt.raw_mirror_dir(_data_root()) / bt.DATA_RELPATH
    if not path.exists():
        pytest.skip(f"raw mirror not deposited: {path}")
    return bt.read_release(path)


@pytest.mark.data
def test_raw_mirror_matches_the_pin() -> None:
    data_root = _data_root()
    manifest = bt.load_manifest(data_root)
    assert bt.manifest_sha256(manifest) == bt.DATA_SHA256
    assert file_sha256(bt.raw_mirror_dir(data_root) / bt.DATA_RELPATH) == bt.DATA_SHA256
    (record,) = manifest.files
    assert record.bytes == bt.DATA_BYTES
    assert record.retrieval is not None
    assert record.retrieval.last_check is not None
    assert record.retrieval.last_check.matches


@pytest.mark.data
def test_every_quote_is_verbatim_in_its_pinned_mirror() -> None:
    from torchcell.verification.sourced import audit_sourced_value

    library = osp.join(_data_root(), "torchcell-library")
    quoted = list(bt.SOURCED_VALUES.values())
    # Components inherited from MOPS_MINIMAL quote Price 2018's xlsx as row renderings
    # and are audited by the media library's own tests; this module's additions are not.
    inherited = list(MOPS_MINIMAL.components)
    for media in (bt.BORCHERT2023_M9, bt.FB_MOPS_GLUCOSE_NO_NITROGEN):
        quoted.extend(media.provenance)
        for component in media.components:
            if component not in inherited:
                quoted.extend(component.provenance)
    failures = [
        (sv.provenance.citation_key, sv.quote[:60])
        for sv in quoted
        if not audit_sourced_value(sv, library).passed
    ]
    assert failures == []


@pytest.mark.data
def test_real_release_shape_and_superset_coverage() -> None:
    release = _real_release()
    assert (len(release.genes), len(release.samples)) == (bt.N_GENES, bt.N_SAMPLES)
    assert sum(s.is_beckham for s in release.samples) == 78
    attributions = {s.exp_name: bt.attribute_sample(s) for s in release.samples}
    coverage = {c.study: c for c in bt.coverage(release.samples, attributions)}
    assert (
        coverage["thompson2020"].n_samples,
        coverage["thompson2020"].n_conditions,
    ) == (46, 23)
    assert (
        coverage["schmidt2022"].n_samples,
        coverage["schmidt2022"].n_conditions,
    ) == (123, 71)
    assert coverage["borchert2023"].n_samples == 42
    assert coverage["borchert2024"].n_samples == 121
    assert coverage["borchert2024"].n_kept_samples == 79
    drops = [bt.drop_reason(s) for s in release.samples]
    assert drops.count(bt.DROP_MEDIUM_NOT_IN_LIBRARY) == 20
    assert drops.count(bt.DROP_REACTOR_PROCESS) == 22
    for _, sample in ((None, s) for s in release.samples if bt.drop_reason(s) is None):
        bt.build_environment(sample)


@pytest.mark.data
def test_real_t_columns_agree_in_sign_only_for_the_fitness_browser() -> None:
    agreement = bt.t_sign_agreement(_real_release())
    for set_name, values in agreement.items():
        assert values["floor_violations"] == 0
        if set_name in bt.BECKHAM_SETS:
            assert 0.45 < values["sign_disagreement"] < 0.55
            assert values["median_corr_f_t"] < -0.4
        else:
            assert values["sign_disagreement"] == 0.0
            assert values["median_corr_f_t"] > 0.75


@pytest.mark.data
def test_every_locus_id_is_a_current_kt2440_locus_tag() -> None:
    from torchcell.datasets.bacteria_common import (
        bacterial_genome,
        reconcile_locus_tags,
    )

    release = _real_release()
    genome = bacterial_genome("pputida", "KT2440", data_root=_data_root())
    stored, report = reconcile_locus_tags(genome, pd.Series(release.genes), label="t")
    assert report.status_histogram[GeneNameStatus.CURRENT] == bt.N_GENES
    assert report.remapped == 0
    assert list(stored) == list(release.genes)


@pytest.mark.data
def test_price2018_carries_no_kt2440_experiment() -> None:
    """The plan listed Price 2018's KT2440 experiments as subsumed; it has none."""
    import openpyxl

    path = osp.join(
        _data_root(), "torchcell-library/priceMutantPhenotypesThousands2018/si/si3.xlsx"
    )
    workbook = openpyxl.load_workbook(path, read_only=True)
    org_ids = {
        row[0] for row in workbook["TableS5_Experiments"].iter_rows(values_only=True)
    }
    bacteria = [
        str(row[0])
        for row in workbook["TableS14_RB_TnSeq_Bacteria"].iter_rows(values_only=True)
        if row and row[0]
    ]
    assert "Putida" not in org_ids
    assert "Keio" in org_ids
    assert not any("putida" in name.lower() for name in bacteria)


@pytest.mark.data
def test_built_lmdb_record_count() -> None:
    root = osp.join(_data_root(), "data/torchcell/rbtnseq_borchert2024")
    lmdb_dir = osp.join(root, "processed", "lmdb")
    if not osp.isdir(lmdb_dir):
        pytest.skip(f"not built: {lmdb_dir}")
    env = lmdb.open(lmdb_dir, readonly=True, lock=False)
    with env.begin() as txn:
        assert txn.stat()["entries"] == 290 * bt.N_GENES
        first = pickle.loads(txn.get(b"0"))
    env.close()
    assert first["experiment"]["experiment_type"] == "bacterial_environment_response"
    assert not math.isnan(first["experiment"]["phenotype"]["environment_response"])
