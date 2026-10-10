# torchcell/datasets/ecoli/wang2024
# [[torchcell.datasets.ecoli.wang2024]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/wang2024
# Test file: tests/torchcell/datasets/ecoli/test_wang2024.py
"""Wang 2024: a rifampicin dose-by-time Tn-seq screen of E. coli K-12 MG1655.

Wang, Fu, Shi, Zhao and Lyu 2024 (Microbiology Spectrum 12(1):e02895-23,
doi:10.1128/spectrum.02895-23, PMC10782999, citation key
``wangGenomeWideScreenRevealsCellular2024``) built a Tn5 insertion library in MG1655
(~310,000 insertions, 71.4 per kb), grew it to mid-exponential phase in LB, split it
into three flasks dosed with 2, 32 and 160 mg/L rifampicin (0.25x, 4x and 20x the
8 mg/L MIC), and recovered the survivors 1 and 3 h after treatment by plating. Tn-seq
of the input and of each of the six post-treatment pools, two biological replicates
each, was compared gene by gene with the TRANSIT ``Resampling`` method. Table S2
releases the WHOLE comparison, not the hits: six sheets, one per (dose, time), each
carrying all 4,419 genes.

RECORD = one (gene x condition) ``BacterialEnvironmentResponseExperiment``.
4,419 genes x 6 conditions = 26,514 released cells; the stored count and every drop
are in RETENTION below.

PHENOTYPE. ``EnvironmentResponsePhenotype``, ``measurement_type=log2_ratio``,
``assay_type=other``. The number is the sheet's ``log2FC`` column verbatim: TRANSIT's
log2 fold change of the gene's insertion read counts, summed over every insertion site
and both replicates, between the post-treatment pool and the input pool. Negative =
mutants of the gene are depleted after rifampicin, so the gene counteracts the drug.
``assay_type`` is ``other`` because ``AssayType`` has no member for insertion-junction
sequencing: it is pooled survival read by sequencing the transposon-genome junction
itself, so there is no molecular barcode (not ``pooled_competitive_growth_barcode``),
and the readout is spelled out in ``units``, as Girgis 2009 does for its footprinting.
``n_samples = 2`` biological replicates, ``sample_unit=biological_replicate``; the
statistic is ONE log2FC over the two replicates' summed counts, not a mean of two.

THE TEST OF THE RESPONSE (issue #863, ``SCHEMA_FINDING_ISSUE``). Each row also
releases ``p-value`` and ``Adj. p-value``: the resampling permutation test of the
log2FC and its FDR correction. Both are stored verbatim on every record, as
``environment_response_p_value`` and ``environment_response_p_value_adjusted``, with
``p_value_adjustment_method="benjamini_hochberg"``. The paper names the correction only
as "the method of FDR", so the procedure is BACK-SOLVED (:func:`adjustment_back_solve`):
within each sheet, the Benjamini-Hochberg adjustment of all 4,419 released p-values
reproduces the released ``Adj. p-value`` to at most 5.0e-6 (measured on the mirror; the
p-values are printed to 4 decimals and the adjusted values to 5), and a build that
misses by more than ``ADJUSTMENT_TOLERANCE`` refuses. The family is the whole sheet,
the rows this loader drops included, so the adjusted value of a kept record is the
released number, never one recomputed over the kept subset. A test is not a dispersion,
so neither value feeds ``environment_response_uncertainty`` or the SE.

WHAT IS NOT STORED. The count columns the log2FC is computed from (``Sites``,
``Mean Ctrl``, ``Mean Exp``, ``Sum Ctrl``, ``Sum Exp``, ``Delta Mean``): they are the
inputs of the stored statistic.

CONDITIONS. The six sheets, each a ``SmallMoleculePerturbation`` of rifampicin at its
ABSOLUTE dose in ug/mL (= mg/L, from the Methods), ``basis=DoseBasis.MIC`` because the
three doses were SET as multiples of the measured MIC, and the multiple carried in the
perturbation's ``description`` (``"rifampicin at 0.25x MIC (MIC 8 mg/L)"``). The
exposure time is ``Environment.duration_hours`` (1 or 3). ``screen_id`` is the sheet
name verbatim (``0.25xMIC-1hour`` ...), which is what keeps the two time points of one
dose L1-distinct. One medium, LB Miller, 37 C, aerobic (shaken). Kanamycin, which the
library was grown in, was washed out before the split and is NOT on the medium; the
kanamycin of the recovery plates is outside the treatment environment.

STRAIN AND GENOTYPE. The library was transformed directly into MG1655, so the records
are written against MG1655 GCA_000005845.2 with no strain background, and the TRANSIT
mapping was to NC_000913.3, the RefSeq copy of the same U00096.3 sequence. One
gene-level ``TransposonInsertionPerturbation`` per record, ``transposon="Tn5"``:
a row aggregates every insertion in the gene's central 90 percent (the Methods
discard the 5 percent at each end), so ``barcode``, ``insertion_position`` and
``insertion_strand`` are None; their typed absences are in
:data:`PERTURBATION_FIELD_GAPS` (the leaf has no ``provenance_gaps`` slot).

IDENTIFIERS. ``#Orf`` holds MG1655 b-numbers. A b-number is stored when it IS a locus
tag of the pinned annotation; any other is dropped with its resolution note, the
Girgis 2009 rule (issue #753 for the missing retired-tag route). The measured split is
written to ``preprocess/identifier_reconciliation.json``.

RETENTION (``preprocess/dropped_records.json``).

1. ``no_insertion_reads_in_either_pool`` -- ``Sum Ctrl`` and ``Sum Exp`` both 0. No
   mutant of the gene was read in the input OR the treated pool, so there is nothing
   whose abundance changed; TRANSIT still prints log2FC 0 and p 1 for such a gene, and
   storing that 0 would assert "no fitness change" for a mutant that was never observed.
   Measured on the release: the input has no reads for the same 303 genes in all six
   sheets, and 277 to 294 of them per sheet also have none after treatment, 1,721 cells
   in all. Every one of them carries log2FC 0. (Hypothesis, untested: most of the 303
   are genes whose disruption the library does not tolerate.)
2. ``b_number_is_not_a_locus_tag_of_the_pinned_annotation`` -- 5 genes x 6 conditions
   = 30 cells: ``b3036``, ``b4223``, ``b4590`` and ``b4700`` are not in the annotation
   and ``b3681`` is a synonym of the pseudogene ``b4556``. The other 4,414 are locus
   tags as released, so no record needs a ``DerivedIdentifierMapping``.

26,514 - 1,721 - 30 = 24,763 records. A gene with no input reads but reads after
treatment (9 to 26 per sheet) is kept: those reads are measured.

DUPLICATION. Independent of the two rifampicin screens already stored. Spearman r of
each condition against Choe 2025's CRISPRi rifampicin screen (3.2 ug/mL, about 3,780
shared genes) is -0.028 to 0.027, and against Shiver 2016's Keio rifampicin screen
(4 ug/mL, about 3,420 shared genes by symbol) -0.001 to 0.121
(``experiments/036-dataset-fixes-before-kg-build/results/
wangGenomeWideScreenRevealsCellular2024_release_inventory.json``).

UNCERTAINTY. No dispersion is released per cell; ``environment_response_uncertainty``
and ``environment_response_se`` carry ``not_reported_by_primary`` gaps. The p-value pair
above is the released test, carried on its own fields.

REFERENCE. One per condition: the unperturbed MG1655 library parent in the same
condition, scoring 0 on the log2 ratio axis (no change in representation).

BUILD-TIME CHECKS (no fallbacks; each raises). The workbook must have exactly the six
sheets in the released order; the first sheet's title and legend rows and every
sheet's header must read as expected; every sheet must carry the same 4,419 b-numbers
in the same order with the same ``Sites``, ``Mean Ctrl`` and ``Sum Ctrl`` (the input
is ONE shared sample, so its columns cannot differ between sheets); and every value
cell must be a finite number.

DATA SOURCE. Table S2 (``spectrum.02895-23-s0002.xlsx``) and the article's own PDF and
plain text, from the PMC Article Datasets bucket prefix ``PMC10782999.1``
(``pmc_cloud``, scriptable), deposited in
``$DATA_ROOT/torchcell-raw/wangGenomeWideScreenRevealsCellular2024/`` with a
``manifest.json``. The paper is NOT in the literature mirror (no
``torchcell-library`` directory carries this DOI, measured 2026-10-10), so every quote
is anchored to ``paper/PMC10782999.1.txt``, the publisher's own text, sha256-pinned in
the raw mirror, and the workbook's legend rows to a deterministic text rendering of
them with a ``ProcessingRecord``.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import math
import os
import os.path as osp
import re
import shutil
from collections import Counter
from collections.abc import Callable, Hashable, Iterable, Iterator, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, ClassVar, Final

import numpy as np
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
    write_verified,
)
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.media import LB
from torchcell.datamodels.schema import (
    BACTERIAL_ASSEMBLY_SETS,
    AssayType,
    AssemblyReferenceGenome,
    BacterialAssemblySet,
    BacterialEnvironmentResponseExperiment,
    BacterialEnvironmentResponseExperimentReference,
    Concentration,
    ConcentrationUnit,
    DoseBasis,
    Environment,
    EnvironmentResponsePhenotype,
    Experiment,
    ExperimentReference,
    Genotype,
    MeasurementType,
    Publication,
    SampleUnit,
    SmallMoleculePerturbation,
    Temperature,
    TransposonInsertionPerturbation,
)
from torchcell.datasets.bacteria_common import (
    LAYER_LOCUS_TAG,
    LOCUS_TAG_PATTERNS,
    STRAIN_GENE_NAMESPACES,
    LocusTagReconciliation,
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
    resolution_layer,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ROLE_PAPER_PDF,
    ROLE_PAPER_TEXT,
    ROLE_RAW_DATA,
    ROLE_SI_TEXT,
    ArtifactRecord,
    Manifest,
    ProcessingRecord,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.literature.provenance import run_retriever
from torchcell.literature.retrieve import pmc_cloud_url
from torchcell.sequence.genome.ecoli.k12 import EcoliK12Genome, EcoliK12StrainName
from torchcell.verification.report import (
    Level,
    LevelResult,
    Provenance,
    VerificationReport,
)
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
    audit_sourced_value,
)

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
CITATION_KEY = "wangGenomeWideScreenRevealsCellular2024"
PAPER_DOI = "10.1128/spectrum.02895-23"
PAPER_TITLE = (
    "Genome-wide screen reveals cellular functions that counteract rifampicin "
    "lethality in Escherichia coli"
)
PMCID = "PMC10782999"
#: The PubMed id of the DOI, read from the PMC id converter on 2026-10-10.
PUBMED_ID = "38054714"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
DATASET_ROOT_REL = "data/torchcell/ecoli_env_chemgen_wang2024"

#: The PMC Article Datasets bucket prefix of this article (version 1).
PMC_PREFIX = f"{PMCID}.1"
#: When the files were retrieved and hashed before deposit.
RETRIEVED_AT = "2026-10-10"

DATA_FILE = "spectrum.02895-23-s0002.xlsx"
PAPER_PDF_FILE = f"{PMC_PREFIX}.pdf"
PAPER_TEXT_FILE = f"{PMC_PREFIX}.txt"

#: Quote anchor: the publisher's plain text, inside the raw mirror.
PAPER_TEXT_REL = f"paper/{PAPER_TEXT_FILE}"
PAPER_TEXT_SHA256 = "0e499878b45bd741ed81bfe557013980faca39ea72006b4b9a30bd9312d89b0d"

#: Quote anchor for Table S2's own legend rows: a deterministic text rendering of
#: them, deposited beside the binary workbook with a ``ProcessingRecord``.
LEGENDS_FILE = "spectrum.02895-23-s0002.legends.txt"
LEGENDS_REL = f"si/{LEGENDS_FILE}"
LEGENDS_SHA256 = "c78aab7d2a945f6153bcb1a2f47a615ce7d8b58ad51f8915dcc05ca13434e94c"

#: The GitHub issue that added the p-value carrier to ``EnvironmentResponsePhenotype``
#: (the schema finding this loader raised), so Table S2's two test columns are stored.
SCHEMA_FINDING_ISSUE = 863


class RawFile(BaseModel):
    """One file the mirror holds: its pinned bytes, role and how it was retrieved."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    name: str
    relpath: str
    role: str
    sha256: str
    bytes: int
    description: str
    derived: bool = False

    @property
    def bucket_key(self) -> str:
        """The PMC Article Datasets bucket key the bytes came from."""
        return f"{PMC_PREFIX}/{self.name}"

    @property
    def source_url(self) -> str:
        """HTTPS URL of the bucket object."""
        return pmc_cloud_url(self.bucket_key)

    @property
    def retrieval(self) -> RetrievalRecord:
        """The re-runnable retrieval of these bytes."""
        return RetrievalRecord(
            method=RetrievalMethod.pmc_cloud,
            source_url=self.source_url,
            retriever="torchcell.literature.retrieve.pmc_cloud_object",
            params={"key": self.bucket_key},
            sha256=self.sha256,
            retrieved_at=RETRIEVED_AT,
        )


RAW_FILES: tuple[RawFile, ...] = (
    RawFile(
        name=DATA_FILE,
        relpath=f"data/{DATA_FILE}",
        role=ROLE_RAW_DATA,
        sha256="143c77c80093adbead1293f79100a5c9a0bf6cbe8d6e7208d0e29710773f9f98",
        bytes=2123965,
        description="Table S2, 'Raw data of resampling': six sheets, one per "
        "(rifampicin dose, time), each carrying all 4,419 MG1655 genes with Sites, "
        "Mean Ctrl, Mean Exp, log2FC, Sum Ctrl, Sum Exp, Delta Mean, p-value and "
        "Adj. p-value (26,514 gene x condition rows)",
    ),
    RawFile(
        name=PAPER_PDF_FILE,
        relpath=f"paper/{PAPER_PDF_FILE}",
        role=ROLE_PAPER_PDF,
        sha256="8c24d5fbf0bad2b56dc2c8669ba2c6eade9a11cd2a981888b27e46d9a2f153ae",
        bytes=4007931,
        description="The article PDF as PMC serves it. Mirrored because the paper is "
        "not in the literature mirror, so no paper.pdf exists for this citation key",
    ),
    RawFile(
        name=PAPER_TEXT_FILE,
        relpath=PAPER_TEXT_REL,
        role=ROLE_PAPER_TEXT,
        sha256=PAPER_TEXT_SHA256,
        bytes=59236,
        description="The article's full text as PMC renders it, the anchor of every "
        "quote in this loader that comes from the paper. Not OCR: the publisher's own "
        "text, so a quote matches character for character",
    ),
    RawFile(
        name=LEGENDS_FILE,
        relpath=LEGENDS_REL,
        role=ROLE_SI_TEXT,
        sha256=LEGENDS_SHA256,
        bytes=299,
        description="Table S2's title and three legend rows (first sheet, rows 1 to "
        "4), rendered to text by extract_sheet_legends so a quote in them is "
        "auditable; the workbook itself is binary",
        derived=True,
    ),
)
#: ``{raw file name: pinned sha256}``, the build-time check of every consumed file.
DATA_SHA256: dict[str, str] = {RAW_FILES[0].name: RAW_FILES[0].sha256}
RAW_FILES_BY_NAME: dict[str, RawFile] = {f.name: f for f in RAW_FILES}
#: The files that come off the PMC bucket; the legend text is derived from the workbook.
RETRIEVED_FILES: tuple[RawFile, ...] = tuple(f for f in RAW_FILES if not f.derived)

#: What the article releases that the loader deliberately does not consume.
NOT_MIRRORED = (
    "Supplemental material spectrum.02895-23-s0001.pdf (Fig. S1, Tables S1 and S5): "
    "the MIC curves, the input-library read counts and the adaptor and primer list. "
    "The MIC and the doses the loader uses are stated in the article text",
    "Table S3 spectrum.02895-23-s0003.xlsx, 'Identified genes': the rows of Table S2 "
    "that pass |log2FC| > 1 and adjusted p < 0.05, plus an 'All' sheet of their log2FC. "
    "A hit call over Table S2, not a measurement",
    "Table S4 spectrum.02895-23-s0004.xlsx, 'Pathway enrichment analysis': DAVID "
    "enrichment of the Table S3 gene lists. A derived analysis",
    "The six article figures (spectrum.02895-23.f001.jpg to f006.jpg) in the same "
    "bucket prefix: no record is built from an image",
)

# --------------------------------------------------------------------------- #
# Sourced values (verbatim quotes, pinned sha256)
# --------------------------------------------------------------------------- #
_STRAINS = "MATERIALS AND METHODS, 'Strains, media, and growth conditions'"
_LIBRARY = "MATERIALS AND METHODS, 'Generation of transposon mutant library'"
_SCREEN = "MATERIALS AND METHODS, 'Library screening'"
_SEQ = "MATERIALS AND METHODS, 'Sequencing analysis'"
_RESULTS = "RESULTS, 'Selection of transposon mutants altering rifampicin efficacy'"


def _paper(
    value: Any, quote: str, *, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to a verbatim quote in the sha256-pinned article text."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=PAPER_TEXT_REL,
            citation_key=CITATION_KEY,
            sha256=PAPER_TEXT_SHA256,
            method="the publisher's own plain text of the article, from the PMC "
            "Article Datasets bucket (raw mirror; the paper is not in the literature "
            "mirror, so there is no OCR'd paper.md to quote)",
            page=page,
        ),
    )


def _legend(value: Any, quote: str, *, note: str | None = None) -> SourcedValue:
    """Bind a value to a verbatim quote in the rendered Table S2 legend rows."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=LEGENDS_REL,
            citation_key=CITATION_KEY,
            sha256=LEGENDS_SHA256,
            method="openpyxl rendering of Table S2's title and legend rows "
            "(extract_sheet_legends), pinned to the workbook's sha256",
            page="Table S2, sheet '0.25xMIC-1hour', rows 1 to 4",
        ),
    )


SOURCED_VALUES: dict[str, SourcedValue] = {
    "reference_strain": _paper(
        "MG1655",
        "E. coli K12 strain MG1655 was used in this study.",
        page=_STRAINS,
        note="pinned to GCA_000005845.2, the assembly of the ecoli_K12_MG1655_ASM584v2 "
        "set, whose locus tags are the b-numbers Table S2 keys on",
    ),
    "temperature_c": _paper(
        37.0, "Cells were grown at 37°C on a rotating shaker at 220 rpm.", page=_STRAINS
    ),
    "aerobicity": _paper(
        "aerobic",
        "The pooled library was grown aerobically in the middle exponential phase",
        page=_RESULTS,
        note="the treatment flasks were incubated 'at 37°C with rotation' (Library "
        "screening), the shaken regime of the strains section",
    ),
    "medium": _paper(
        "LB",
        "Approximately 107 colony-forming units were inoculated into a 100-mL flask "
        "containing 10-mL LB and grown to an OD600 of 0.5.",
        page=_SCREEN,
        note="the treatment culture is LB; '107' is the PMC text's rendering of 10^7. "
        "The strains section names 'Luria–Bertani (LB)' with no recipe, and the LB "
        "library object's Miller recipe is the shared join key",
    ),
    "kanamycin_washed_out": _paper(
        False,
        "Cells were washed with an equal volume of LB to remove kanamycin.",
        page=_SCREEN,
        note="kanamycin selected the library during its growth and on the recovery "
        "plates, but it was removed before the rifampicin split, so it is not on the "
        "treatment medium",
    ),
    "transposon": _paper(
        "Tn5",
        "a modified Tn5 transposon (Epicentre, TSM99K2) was transformed into E. coli "
        "by electroporation",
        page=_LIBRARY,
    ),
    "doses_mg_per_l": _paper(
        (2.0, 32.0, 160.0),
        "treated with different concentrations of rifampicin (2, 32, and 160 mg/L, "
        "Sigma Aldrich, R3501) at 37°C with rotation",
        page=_SCREEN,
        note="mg/L = ug/mL, so the stored values are the printed numbers",
    ),
    "mic_multiples": _paper(
        {2.0: "0.25x", 32.0: "4x", 160.0: "20x"},
        "exposed to 2 (0.25× MIC), 32 (4× MIC), and 160 (20× MIC) mg/L rifampicin.",
        page=_RESULTS,
    ),
    "mic_mg_per_l": _paper(
        8.0,
        "The E. coli K12 strain MG1655 had a rifampicin MIC of 8 mg/L (Fig. S1A).",
        page=_RESULTS,
        note="the Methods define the MIC measured as MIC90 ('Ninety percent "
        "inhibitory concentration (MIC90)')",
    ),
    "time_points_h": _paper(
        (1.0, 3.0),
        "At 1 and 3 h post-treatment, approximately 3 million CFUs were recovered by "
        "plating onto five LB agar plates (200 mm diameter) supplemented with "
        "kanamycin (50 mg/L) and incubated overnight at 37°C.",
        page=_SCREEN,
        note="duration_hours is the rifampicin exposure; the overnight plate outgrowth "
        "that follows is a recovery step, not part of the treatment",
    ),
    "n_biological_replicates": _paper(
        2,
        "Colonies were pooled in LB and stored at −80°C. Two biological replicates "
        "were performed.",
        page=_SCREEN,
    ),
    "statistic": _paper(
        "log2 fold change of summed read counts, treated vs input",
        "For each gene, the normalized read counts at all the insertion sites and all "
        "replicates under each condition were summed. The difference of the summed "
        "read counts between the input and post-treatment samples was calculated for "
        "each mutant and expressed as a log2 fold change (log2FC).",
        page=_RESULTS,
        note="ONE statistic over both replicates, so n_samples=2 is the replicate "
        "count behind it, not a count of averaged values",
    ),
    "p_value_test": _paper(
        "permutation test",
        "The significance of this difference was calculated using a permutation test.",
        page=_RESULTS,
        note="what Table S2's p-value column is: the TRANSIT resampling test of the "
        "log2FC, stored as environment_response_p_value",
    ),
    "p_value_adjustment": _paper(
        "FDR",
        "Read counts, P-value (adjusted by using the method of FDR), and log2FC between "
        "the input and post-treatment were calculated using default parameters.",
        page=_SEQ,
        note="the correction is named only as FDR; which FDR procedure produced "
        "Adj. p-value is back-solved from the released p-values "
        "(adjustment_back_solve) as Benjamini-Hochberg over each whole sheet",
    ),
    "resampling_method": _paper(
        "TRANSIT Resampling v3.2.0",
        "The “Resampling” method of TRANSIT software (v3.2.0) was used to identify "
        "mutants that were differentially represented between input and "
        "post-treatment samples.",
        page=_SEQ,
    ),
    "gene_trimming": _paper(
        0.05,
        "Reads in the 5% N-terminal and 5% C-terminal of the gene sequence were "
        "discarded (61).",
        page=_SEQ,
        note="why a record is gene-level with no insertion position",
    ),
    "mapping_reference": _paper(
        "NC_000913.3",
        "mapped to a unique site on the E. coli MG1655 genome (NCBI GenBank, "
        "NC_000913.3)",
        page=_SEQ,
        note="the RefSeq record of the U00096.3 sequence that GCA_000005845.2 holds, "
        "so the b-numbers are that annotation's locus tags",
    ),
    "table_s2_is_resampling_output": _legend(
        "TRANSIT resampling output, all genes", "Table S2-Raw data of resampling."
    ),
    "mean_ctrl_definition": _legend(
        "input sample",
        "Mean Ctrl:  average number of unique sequenced templates of the input sample",
        note="the control of every sheet is the one input sample, which is why the "
        "build requires Mean Ctrl and Sum Ctrl to be identical across the six sheets",
    ),
    "mean_exp_definition": _legend(
        "post-treatment sample",
        "Mean Exp: average number of unique sequenced templates of the post-treatment "
        "sample",
    ),
}

REFERENCE_STRAIN_NAME: Final[EcoliK12StrainName] = "MG1655"
MG1655_ASSEMBLY_SET: Final[BacterialAssemblySet] = "ecoli_K12_MG1655_ASM584v2"
MG1655_NAMESPACE = STRAIN_GENE_NAMESPACES[REFERENCE_STRAIN_NAME]
TEMPERATURE_C: Final[float] = SOURCED_VALUES["temperature_c"].value
AEROBICITY: Final[str] = SOURCED_VALUES["aerobicity"].value
TRANSPOSON: Final[str] = SOURCED_VALUES["transposon"].value
MIC_MG_PER_L: Final[float] = SOURCED_VALUES["mic_mg_per_l"].value
N_REPLICATES: Final[int] = SOURCED_VALUES["n_biological_replicates"].value
#: Every released ``#Orf`` matches this; a value that does not fails the format check.
B_NUMBER_PATTERN = re.compile(r"^b\d{4}$")
#: Below this fraction of distinct ``#Orf`` b-numbers that are locus tags of the pinned
#: MG1655 annotation the build stops: the release was mapped to the same sequence, so a
#: lower fraction would mean the wrong annotation.
MIN_RESOLVED_FRACTION = 0.98

_PAPER_LOOKED_IN = Provenance(
    source_uri=PAPER_TEXT_REL,
    citation_key=CITATION_KEY,
    sha256=PAPER_TEXT_SHA256,
    method="full Results and Materials and Methods read, and every column of Table S2",
)

PHENOTYPE_GAPS: tuple[ProvenanceGap, ...] = (
    ProvenanceGap(
        field="environment_response_uncertainty",
        reason=ProvenanceGapReason.not_reported_by_primary,
        looked_in=_PAPER_LOOKED_IN,
        note="Table S2 releases one log2FC per cell computed on the two replicates' "
        "SUMMED counts, so no per-replicate value and no dispersion exists. Its "
        "p-value and Adj. p-value columns are a permutation test of the log2FC, not "
        "an uncertainty on it, and are stored on environment_response_p_value and "
        "environment_response_p_value_adjusted instead (#863)",
    ),
    ProvenanceGap(
        field="environment_response_se",
        reason=ProvenanceGapReason.not_reported_by_primary,
        looked_in=_PAPER_LOOKED_IN,
        note="no dispersion is released, so no standard error can be derived",
    ),
)

_PERTURBATION_GAP_NOTES: dict[str, str] = {
    "barcode": "Tn-seq reads the transposon-genome junction itself, so the library "
    "carries no molecular barcode: the insertion's own position is its read-out "
    "identity",
    "insertion_position": "a row of Table S2 sums every insertion site in the gene "
    "(after trimming 5 percent at each end), so no single insertion site exists for it; "
    "the per-site wig files are not released",
    "insertion_strand": "the gene-level resampling output carries no orientation",
}
#: Typed absences of the transposon leaf's own fields. The leaf carries no
#: ``provenance_gaps`` slot, so they live here and in
#: ``preprocess/perturbation_field_gaps.json``, and every record's value is None.
PERTURBATION_FIELD_GAPS: tuple[ProvenanceGap, ...] = tuple(
    ProvenanceGap(
        field=field,
        reason=ProvenanceGapReason.not_reported_by_primary,
        looked_in=_PAPER_LOOKED_IN,
        note=note,
    )
    for field, note in _PERTURBATION_GAP_NOTES.items()
)

PUBLICATION = Publication(
    pubmed_id=PUBMED_ID,
    pubmed_url=f"https://pubmed.ncbi.nlm.nih.gov/{PUBMED_ID}/",
    doi=PAPER_DOI,
    doi_url=f"https://doi.org/{PAPER_DOI}",
)


# --------------------------------------------------------------------------- #
# The six conditions of Table S2
# --------------------------------------------------------------------------- #
class ConditionSpec(BaseModel):
    """One sheet of Table S2: its rifampicin dose and its exposure time."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    sheet: str
    dose_mg_per_l: float
    hours: float

    @property
    def mic_multiple(self) -> str:
        """The dose as the multiple of the MIC the paper names it by."""
        multiples: dict[float, str] = SOURCED_VALUES["mic_multiples"].value
        return multiples[self.dose_mg_per_l]

    @property
    def screen_id(self) -> str:
        """The sheet name verbatim."""
        return self.sheet

    @property
    def dose_description(self) -> str:
        """The perturbation's description: the MIC multiple and the MIC it multiplies."""
        mic = f"{MIC_MG_PER_L:g}"
        return f"rifampicin at {self.mic_multiple} MIC (MIC {mic} mg/L)"


def _conditions() -> tuple[ConditionSpec, ...]:
    """The six sheets in the released order, each checked against the sourced doses."""
    doses: tuple[float, ...] = SOURCED_VALUES["doses_mg_per_l"].value
    hours: tuple[float, ...] = SOURCED_VALUES["time_points_h"].value
    specs = (
        ConditionSpec(sheet="0.25xMIC-1hour", dose_mg_per_l=2.0, hours=1.0),
        ConditionSpec(sheet="0.25xMIC-3hours", dose_mg_per_l=2.0, hours=3.0),
        ConditionSpec(sheet="4xMIC-1hour", dose_mg_per_l=32.0, hours=1.0),
        ConditionSpec(sheet="4xMIC-3hours", dose_mg_per_l=32.0, hours=3.0),
        ConditionSpec(sheet="20xMIC-1hour", dose_mg_per_l=160.0, hours=1.0),
        ConditionSpec(sheet="20xMIC-3hours", dose_mg_per_l=160.0, hours=3.0),
    )
    for spec in specs:
        if spec.dose_mg_per_l not in doses or spec.hours not in hours:
            raise ValueError(f"{spec.sheet} is not a sourced (dose, time)")
        if not spec.sheet.startswith(spec.mic_multiple.removesuffix("x") + "xMIC-"):
            raise ValueError(f"{spec.sheet} does not name {spec.mic_multiple} MIC")
    return specs


CONDITIONS: tuple[ConditionSpec, ...] = _conditions()
CONDITIONS_BY_SHEET: dict[str, ConditionSpec] = {c.sheet: c for c in CONDITIONS}


# --------------------------------------------------------------------------- #
# Raw mirror (the loader reads the mirror, never a live URL)
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the environment (the mirrors and build tree live under it)."""
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/wangGenomeWideScreenRevealsCellular2024``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def _sha256(path: str | Path) -> str:
    """Streaming sha256 of a file."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


#: Rows 1 to 4 of the first sheet: the table title and its three column definitions.
LEGEND_ROWS: Final[int] = 4


def extract_sheet_legends(workbook: str | Path) -> str:
    """Table S2's title and legend rows as deterministic text.

    Rows 1 to 4 of the first sheet are the release's own statement of what the table
    is and what its control and experimental columns mean. The workbook is binary, so
    a quote in it cannot be audited; this rendering can, and its
    :class:`ProcessingRecord` names this function and pins the workbook's sha256.
    """
    import openpyxl

    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    sheet = CONDITIONS[0].sheet
    lines = [f"# sheet: {sheet}"]
    for row in book[sheet].iter_rows(min_row=1, max_row=LEGEND_ROWS, values_only=True):
        lines.append("" if row[0] is None else str(row[0]))
    book.close()
    return "\n".join(lines) + "\n"


def _openpyxl_version() -> str:
    """The openpyxl version the legend text was rendered with."""
    import openpyxl

    return str(openpyxl.__version__)


def legends_processing(workbook_sha256: str) -> ProcessingRecord:
    """How the legend text was produced, and from which bytes."""
    return ProcessingRecord(
        processor="torchcell.datasets.ecoli.wang2024.extract_sheet_legends",
        tool="openpyxl",
        version=_openpyxl_version(),
        params={"sheets": [CONDITIONS[0].sheet], "rows": list(range(1, 5))},
        input_sha256=[workbook_sha256],
    )


def retrieve_raw_files(dest_dir: str | Path) -> dict[str, Path]:
    """Run each file's recorded retriever and write the verified bytes to ``dest_dir``.

    The recorded ``RetrievalRecord`` is what runs (``run_retriever``), so this IS the
    re-runnable retrieval; a byte mismatch raises before anything is written.
    """
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}
    for raw in RETRIEVED_FILES:
        path = dest / raw.name
        write_verified(run_retriever(raw.retrieval), path, raw.sha256, raw.source_url)
        out[raw.name] = path
    return out


def deposit_raw_mirror(
    *, sources: Mapping[str, str | Path], data_root: str | None = None
) -> Path:
    """Write the raw mirror from already-retrieved files plus its ``manifest.json``.

    ``sources`` maps every retrieved name in ``RAW_FILES`` to a local file. Idempotent
    by sha256: a mirror file with the pinned hash is left alone, and one with any other
    hash raises rather than being overwritten. The legend text is rendered from the
    workbook and must hash to :data:`LEGENDS_SHA256`.
    """
    missing = sorted({f.name for f in RETRIEVED_FILES} - set(sources))
    if missing:
        raise KeyError(f"no source given for {missing}")
    root = raw_mirror_dir(data_root)
    workbook_sha256 = _sha256(sources[DATA_FILE])
    records: list[ArtifactRecord] = []
    for raw in RAW_FILES:
        dest = root / raw.relpath
        dest.parent.mkdir(parents=True, exist_ok=True)
        if raw.derived:
            text = extract_sheet_legends(sources[DATA_FILE])
            got = hashlib.sha256(text.encode("utf-8")).hexdigest()
            if got != raw.sha256:
                raise RuntimeError(
                    f"rendered legends sha256 {got}, expected {raw.sha256}"
                )
            if dest.exists() and _sha256(dest) != raw.sha256:
                raise RuntimeError(f"{dest} exists with a different sha256; refusing")
            dest.write_text(text, encoding="utf-8")
            records.append(
                ArtifactRecord(
                    path=raw.relpath,
                    role=raw.role,
                    bytes=dest.stat().st_size,
                    sha256=raw.sha256,
                    source=f"derived from data/{DATA_FILE}",
                    original_filename=raw.name,
                    processing=legends_processing(workbook_sha256),
                )
            )
            continue
        src = Path(sources[raw.name])
        got = _sha256(src)
        if got != raw.sha256:
            raise RuntimeError(
                f"{src} sha256 mismatch: got {got}, expected {raw.sha256}"
            )
        if dest.exists():
            if _sha256(dest) != raw.sha256:
                raise RuntimeError(f"{dest} exists with a different sha256; refusing")
        else:
            shutil.copy2(src, dest)
        records.append(
            ArtifactRecord(
                path=raw.relpath,
                role=raw.role,
                bytes=dest.stat().st_size,
                sha256=raw.sha256,
                source=raw.source_url,
                original_filename=raw.name,
                retrieval=raw.retrieval,
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title=PAPER_TITLE,
        files=records,
        si_data_sources=[raw.source_url for raw in RETRIEVED_FILES],
        si_expected=list(NOT_MIRRORED),
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


# --------------------------------------------------------------------------- #
# Reading Table S2
# --------------------------------------------------------------------------- #
class SheetFormatError(ValueError):
    """A sheet whose title, header, ids or a cell is not what the loader reads."""


#: Every sheet's header row, verbatim.
HEADER: Final[tuple[str, ...]] = (
    "#Orf",
    "Name",
    "Sites",
    "Mean Ctrl",
    "Mean Exp",
    "log2FC",
    "Sum Ctrl",
    "Sum Exp",
    "Delta Mean",
    "p-value",
    "Adj. p-value",
)
#: Columns that describe the shared input sample, identical in every sheet.
INPUT_COLUMNS: Final[tuple[str, ...]] = (
    "#Orf",
    "Name",
    "Sites",
    "Mean Ctrl",
    "Sum Ctrl",
)
NUMERIC_COLUMNS: Final[tuple[str, ...]] = HEADER[2:]
#: The first sheet's rows above its header, verbatim (row 5 is blank).
FIRST_SHEET_PREAMBLE: Final[tuple[str | None, ...]] = (
    "Table S2-Raw data of resampling. ",
    "Sites:number of Tn5 transposon insertion sites sequenced across all comparisons",
    "Mean Ctrl:  average number of unique sequenced templates of the input sample",
    "Mean Exp: average number of unique sequenced templates of the post-treatment "
    "sample",
    None,
)


def _numeric(value: Any, where: str) -> float:
    """A value cell as a finite float; anything else raises."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise SheetFormatError(f"{where}: non-numeric cell {value!r}")
    if not math.isfinite(value):
        raise SheetFormatError(f"{where}: non-finite cell {value!r}")
    return float(value)


def read_table_s2(path: str | Path) -> dict[str, pd.DataFrame]:
    """Every sheet of Table S2 as a frame keyed by sheet name, format-checked.

    Refuses a workbook whose sheet list, first-sheet preamble, header, ids or value
    cells differ from what the loader reads.
    """
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    expected = [c.sheet for c in CONDITIONS]
    if book.sheetnames != expected:
        raise SheetFormatError(f"sheets {book.sheetnames}, expected {expected}")
    frames: dict[str, pd.DataFrame] = {}
    for index, spec in enumerate(CONDITIONS):
        rows = list(book[spec.sheet].iter_rows(values_only=True))
        preamble = FIRST_SHEET_PREAMBLE if index == 0 else ()
        got = tuple(row[0] for row in rows[: len(preamble)])
        if got != preamble:
            raise SheetFormatError(f"{spec.sheet}: preamble {got!r}")
        header = tuple(rows[len(preamble)])
        if header != HEADER:
            raise SheetFormatError(f"{spec.sheet}: header {header!r}")
        data = rows[len(preamble) + 1 :]
        records: list[dict[str, Any]] = []
        for number, row in enumerate(data, start=len(preamble) + 2):
            where = f"{spec.sheet} row {number}"
            orf = row[0]
            if not isinstance(orf, str) or B_NUMBER_PATTERN.match(orf) is None:
                raise SheetFormatError(f"{where}: #Orf {orf!r} is not a b-number")
            record: dict[str, Any] = {"#Orf": orf, "Name": str(row[1])}
            for column, value in zip(HEADER[2:], row[2:], strict=True):
                record[column] = _numeric(value, f"{where} {column}")
            records.append(record)
        frames[spec.sheet] = pd.DataFrame.from_records(records, columns=list(HEADER))
    book.close()
    return frames


def check_sheets_align(frames: Mapping[str, pd.DataFrame]) -> list[str]:
    """The shared b-number list, after checking every sheet carries the same input.

    The control of every comparison is the one input sample, so the input columns
    (ids, names, sites, mean and sum of the control) must be identical across sheets.
    """
    first = frames[CONDITIONS[0].sheet]
    for spec in CONDITIONS[1:]:
        frame = frames[spec.sheet]
        for column in INPUT_COLUMNS:
            if frame[column].tolist() != first[column].tolist():
                raise SheetFormatError(
                    f"{spec.sheet}: column {column!r} differs from {CONDITIONS[0].sheet}"
                )
    b_numbers = [str(v) for v in first["#Orf"].tolist()]
    if len(set(b_numbers)) != len(b_numbers):
        raise SheetFormatError("#Orf repeats a b-number")
    return b_numbers


# --------------------------------------------------------------------------- #
# The released test: which FDR procedure produced ``Adj. p-value`` (#863)
# --------------------------------------------------------------------------- #
#: The correction behind ``Adj. p-value``, back-solved from the released p-values
#: (:func:`adjustment_back_solve`); the paper names only "the method of FDR"
#: (``SOURCED_VALUES["p_value_adjustment"]``).
P_VALUE_ADJUSTMENT_METHOD: Final[str] = "benjamini_hochberg"
#: A released ``Adj. p-value`` further than this from the Benjamini-Hochberg value of
#: its sheet's own p-values refuses the build. The release prints p-values to 4 decimals
#: and adjusted values to 5, so a faithful BH lands within half a unit of the fifth
#: decimal; measured on the mirror over all six sheets: 5.0e-6 at worst.
ADJUSTMENT_TOLERANCE: Final[float] = 1e-5


class AdjustmentBackSolve(BaseModel):
    """The evidence that ``Adj. p-value`` is the Benjamini-Hochberg value of its own
    sheet's ``p-value`` column.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    method: str
    sheets: int
    rows_per_sheet: int = Field(
        description="the size of each sheet's testing family: every released gene, the "
        "rows this loader drops included"
    )
    max_abs_deviation: float
    worst_sheet: str
    tolerance: float


#: A float64 vector, the type :func:`benjamini_hochberg` reads and returns.
FloatArray = np.ndarray[Any, np.dtype[np.float64]]


def benjamini_hochberg(p_values: FloatArray) -> FloatArray:
    """Benjamini-Hochberg adjusted p-values of ``p_values`` (step-up, capped at 1)."""
    n = p_values.size
    order = np.argsort(p_values, kind="stable")
    scaled = p_values[order] * n / np.arange(1, n + 1, dtype=np.float64)
    monotone = np.minimum.accumulate(scaled[::-1])[::-1]
    out = np.empty(n, dtype=np.float64)
    out[order] = np.minimum(monotone, 1.0)
    return out


def adjustment_back_solve(frames: Mapping[str, pd.DataFrame]) -> AdjustmentBackSolve:
    """Back-solve the multiple-testing correction from the released columns.

    The paper names the correction only as "the method of FDR" and defers the procedure
    to TRANSIT's default parameters, so the method is MEASURED: within each sheet, the
    Benjamini-Hochberg adjustment of every released ``p-value`` must reproduce
    ``Adj. p-value``. A sheet that misses by more than ``ADJUSTMENT_TOLERANCE`` refuses
    the build.
    """
    worst = -1.0
    worst_sheet = ""
    sizes: set[int] = set()
    for spec in CONDITIONS:
        frame = frames[spec.sheet]
        sizes.add(len(frame))
        deviation = float(
            np.abs(
                benjamini_hochberg(frame["p-value"].to_numpy(dtype=np.float64))
                - frame["Adj. p-value"].to_numpy(dtype=np.float64)
            ).max()
        )
        if deviation > worst:
            worst = deviation
            worst_sheet = spec.sheet
    if worst > ADJUSTMENT_TOLERANCE:
        raise SheetFormatError(
            f"{worst_sheet}: Adj. p-value differs from the Benjamini-Hochberg value of "
            f"its sheet's p-values by {worst} (tolerance {ADJUSTMENT_TOLERANCE}); the "
            "released correction is not Benjamini-Hochberg"
        )
    (rows,) = sizes
    return AdjustmentBackSolve(
        method=P_VALUE_ADJUSTMENT_METHOD,
        sheets=len(CONDITIONS),
        rows_per_sheet=rows,
        max_abs_deviation=worst,
        worst_sheet=worst_sheet,
        tolerance=ADJUSTMENT_TOLERANCE,
    )


# --------------------------------------------------------------------------- #
# Identifiers and the retention ledger
# --------------------------------------------------------------------------- #
RULE_NO_READS = "no_insertion_reads_in_either_pool"
RULE_IDENTIFIER = "b_number_is_not_a_locus_tag_of_the_pinned_annotation"

DROP_RULE_DESCRIPTIONS: dict[str, str] = {
    RULE_NO_READS: "Sum Ctrl and Sum Exp are both 0: no mutant of the gene was read in "
    "the input or the treated pool, so no abundance changed. TRANSIT prints log2FC 0 "
    "and p-value 1 for such a gene, and storing that 0 would assert 'no fitness "
    "change' for a mutant that was never observed",
    RULE_IDENTIFIER: "the released b-number is not a locus tag of GCA_000005845.2. "
    "Storing a merged locus needs a DerivedIdentifierMapping and DerivedIdentifierRoute "
    "has no member for a retired tag of the pinned strain's own namespace (issue #753), "
    "so every condition's cell for that gene is dropped; each item gives the "
    "annotation's own resolution note",
}


class DropRule(BaseModel):
    """One retention rule, the cells it removed, and why."""

    rule: str
    description: str
    n_records: int
    items: list[str] = []


class ConditionLedger(BaseModel):
    """Per-condition counts: released cells, kept records and their sign split."""

    sheet: str
    source_cells: int
    no_reads: int
    dropped_identifier: int
    no_input_reads_kept: int
    kept_records: int
    n_positive: int
    n_negative: int
    n_zero: int


class DropLog(BaseModel):
    """The retention ledger of one build, in the order the rules were applied."""

    dataset: str
    source_loci: int
    kept_loci: int
    source_records: int
    kept_records: int
    dropped_records: int
    conditions: list[ConditionLedger]
    rules: list[DropRule]


class IdentifierLedger(BaseModel):
    """The reconciliation of every released b-number, plus the stop threshold."""

    reconciliation: LocusTagReconciliation
    min_resolved_fraction: float
    n_locus_tags: int
    not_a_locus_tag: dict[str, str]


def canonical_symbol(genome: EcoliK12Genome, tag: str) -> str:
    """The annotation's own gene symbol of ``tag`` when it resolves back to ``tag``,
    else the tag itself, so the stored pair always round-trips through the resolver.
    """
    symbol = genome.genbank.loci[tag].symbol
    if symbol is None:
        return tag
    resolution = genome.resolve_gene_name(symbol)
    return symbol if resolution.systematic_name == tag else tag


def resolve_b_numbers(
    genome: EcoliK12Genome, b_numbers: Sequence[str], *, label: str
) -> tuple[dict[str, str], IdentifierLedger]:
    """The kept b-numbers mapped to the annotation's symbol, plus the ledger.

    A b-number is kept when it IS a locus tag of the pinned annotation, the only case
    in which the stored tag is the one the source released. Stops below
    :data:`MIN_RESOLVED_FRACTION`.
    """
    _, report = reconcile_locus_tags(
        genome, pd.Series(list(b_numbers), dtype=object), label=label
    )
    report.require_resolved(MIN_RESOLVED_FRACTION)
    pattern = LOCUS_TAG_PATTERNS[report.gene_namespace]
    kept: dict[str, str] = {}
    unplaced: dict[str, str] = {}
    for b_number in b_numbers:
        resolution = genome.resolve_gene_name(b_number)
        if resolution_layer(genome, resolution) == LAYER_LOCUS_TAG:
            if pattern.match(b_number) is None:
                raise SheetFormatError(
                    f"{b_number} is a locus of {report.assembly_set} but not a "
                    f"{report.gene_namespace} tag"
                )
            kept[b_number] = canonical_symbol(genome, b_number)
        else:
            unplaced[b_number] = (
                f"{resolution.status.value}: {resolution.note or 'no note'}"
            )
    ledger = IdentifierLedger(
        reconciliation=report,
        min_resolved_fraction=MIN_RESOLVED_FRACTION,
        n_locus_tags=len(kept),
        not_a_locus_tag=unplaced,
    )
    return kept, ledger


def _no_reads(row: Mapping[Hashable, Any]) -> bool:
    """True when neither pool has a single read of the gene."""
    return bool(row["Sum Ctrl"] == 0 and row["Sum Exp"] == 0)


class StoredCell(BaseModel):
    """One kept (gene, condition) cell of Table S2: the values a record carries."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    locus_tag: str
    symbol: str
    spec: ConditionSpec
    log2fc: float
    p_value: float
    p_value_adjusted: float


def stored_cells(
    frames: Mapping[str, pd.DataFrame], kept: Mapping[str, str]
) -> Iterator[StoredCell]:
    """Every stored cell, condition-major then gene in sheet order (LMDB order)."""
    for spec in CONDITIONS:
        for row in frames[spec.sheet].to_dict("records"):
            symbol = kept.get(row["#Orf"])
            if symbol is None or _no_reads(row):
                continue
            yield StoredCell(
                locus_tag=row["#Orf"],
                symbol=symbol,
                spec=spec,
                log2fc=float(row["log2FC"]),
                p_value=float(row["p-value"]),
                p_value_adjusted=float(row["Adj. p-value"]),
            )


def build_drop_log(
    dataset_name: str,
    frames: Mapping[str, pd.DataFrame],
    b_numbers: Sequence[str],
    kept: Mapping[str, str],
) -> DropLog:
    """The retention ledger of one build, refusing rules that miss a dropped cell."""
    dropped_identifier = sorted(set(b_numbers) - set(kept))
    no_reads_items: list[str] = []
    ledgers: list[ConditionLedger] = []
    for spec in CONDITIONS:
        signs: Counter[str] = Counter()
        identifier = no_reads = no_input = 0
        for row in frames[spec.sheet].to_dict("records"):
            if row["#Orf"] not in kept:
                identifier += 1
                continue
            if _no_reads(row):
                no_reads += 1
                no_reads_items.append(f"{row['#Orf']} {spec.sheet}")
                continue
            if row["Sum Ctrl"] == 0:
                no_input += 1
            value = row["log2FC"]
            signs["positive" if value > 0 else "negative" if value < 0 else "zero"] += 1
        ledgers.append(
            ConditionLedger(
                sheet=spec.sheet,
                source_cells=len(frames[spec.sheet]),
                no_reads=no_reads,
                dropped_identifier=identifier,
                no_input_reads_kept=no_input,
                kept_records=sum(signs.values()),
                n_positive=signs["positive"],
                n_negative=signs["negative"],
                n_zero=signs["zero"],
            )
        )
    rules = [
        DropRule(
            rule=RULE_NO_READS,
            description=DROP_RULE_DESCRIPTIONS[RULE_NO_READS],
            n_records=len(no_reads_items),
            items=no_reads_items,
        ),
        DropRule(
            rule=RULE_IDENTIFIER,
            description=DROP_RULE_DESCRIPTIONS[RULE_IDENTIFIER],
            n_records=len(dropped_identifier) * len(CONDITIONS),
            items=dropped_identifier,
        ),
    ]
    source_records = sum(ledger.source_cells for ledger in ledgers)
    kept_records = sum(ledger.kept_records for ledger in ledgers)
    drop_log = DropLog(
        dataset=dataset_name,
        source_loci=len(b_numbers),
        kept_loci=len(kept),
        source_records=source_records,
        kept_records=kept_records,
        dropped_records=source_records - kept_records,
        conditions=ledgers,
        rules=rules,
    )
    if sum(rule.n_records for rule in rules) != drop_log.dropped_records:
        raise RuntimeError("drop rules do not account for every dropped cell")
    return drop_log


# --------------------------------------------------------------------------- #
# Environment, genotype, phenotype (pure, no files)
# --------------------------------------------------------------------------- #
def rifampicin(spec: ConditionSpec) -> SmallMoleculePerturbation:
    """The condition's rifampicin at its absolute dose, set as a multiple of the MIC."""
    return SmallMoleculePerturbation(
        description=spec.dose_description,
        compound=resolved_compound("rifampicin"),
        concentration=Concentration(
            value=spec.dose_mg_per_l,
            unit=ConcentrationUnit.ug_per_ml,
            basis=DoseBasis.MIC,
        ),
    )


def environment(spec: ConditionSpec) -> Environment:
    """The treatment culture of one condition: LB plus rifampicin for 1 or 3 h."""
    return Environment(
        media=LB,
        temperature=Temperature(value=TEMPERATURE_C),
        perturbations=[rifampicin(spec)],
        aerobicity=AEROBICITY,
        duration_hours=spec.hours,
    )


def insertion_genotype(locus_tag: str, symbol: str) -> Genotype:
    """The gene-level Tn5 disruption of one MG1655 locus."""
    return Genotype(
        perturbations=[
            TransposonInsertionPerturbation(
                systematic_gene_name=locus_tag,
                perturbed_gene_name=symbol,
                gene_namespace=MG1655_NAMESPACE,
                transposon=TRANSPOSON,
            )
        ]
    )


UNITS = (
    "log2 fold change (TRANSIT Resampling v3.2.0) of the gene's Tn5 insertion read "
    "counts, summed over its insertion sites and both biological replicates, in the "
    "pool surviving rifampicin over the input pool; read by Tn-seq of the "
    "transposon-genome junctions of survivors recovered by plating. Negative = the "
    "gene's mutants are depleted, so the gene counteracts rifampicin"
)
UNITS_REFERENCE = (
    "the unperturbed MG1655 library parent in the same condition: 0 on the log2 "
    "ratio axis, no change in representation relative to the input"
)


def phenotype(cell: StoredCell) -> EnvironmentResponsePhenotype:
    """One released log2FC of one gene in one condition, with its released test."""
    return EnvironmentResponsePhenotype(
        measurement_type=MeasurementType.log2_ratio,
        assay_type=AssayType.other,
        environment_response=cell.log2fc,
        n_samples=N_REPLICATES,
        sample_unit=SampleUnit.biological_replicate,
        units=UNITS,
        screen_id=cell.spec.screen_id,
        environment_response_p_value=cell.p_value,
        environment_response_p_value_adjusted=cell.p_value_adjusted,
        p_value_adjustment_method=P_VALUE_ADJUSTMENT_METHOD,
        provenance_gaps=list(PHENOTYPE_GAPS),
    )


def reference_phenotype(spec: ConditionSpec) -> EnvironmentResponsePhenotype:
    """The parent of the same condition: a log2 ratio of 0."""
    return EnvironmentResponsePhenotype(
        measurement_type=MeasurementType.log2_ratio,
        assay_type=AssayType.other,
        environment_response=0.0,
        n_samples=N_REPLICATES,
        sample_unit=SampleUnit.biological_replicate,
        units=UNITS_REFERENCE,
        screen_id=spec.screen_id,
    )


def reference_genome(data_root: str | None = None) -> AssemblyReferenceGenome:
    """MG1655 pinned to its GenBank assembly; the library was built in MG1655 itself."""
    return assembly_reference(REFERENCE_STRAIN_NAME, data_root=data_root)


def build_experiment(
    dataset_name: str, cell: StoredCell, env: Environment
) -> BacterialEnvironmentResponseExperiment:
    """The record of one (gene, condition) cell of Table S2."""
    return BacterialEnvironmentResponseExperiment(
        dataset_name=dataset_name,
        genotype=insertion_genotype(cell.locus_tag, cell.symbol),
        environment=env,
        phenotype=phenotype(cell),
    )


def build_reference(
    dataset_name: str,
    genome: AssemblyReferenceGenome,
    spec: ConditionSpec,
    env: Environment,
) -> BacterialEnvironmentResponseExperimentReference:
    """The unperturbed parent of one condition, scoring 0."""
    return BacterialEnvironmentResponseExperimentReference(
        dataset_name=dataset_name,
        genome_reference=genome,
        environment_reference=env.model_copy(),
        phenotype_reference=reference_phenotype(spec),
    )


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class EnvChemgenWang2024Dataset(ExperimentDataset):
    """TRANSIT log2FC of MG1655 Tn5 mutants at three rifampicin doses x two times."""

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = REFERENCE_STRAIN_NAME

    def __init__(
        self,
        root: str = DATASET_ROOT_REL,
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialEnvironmentResponseExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialEnvironmentResponseExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The one consumed file, linked from the raw mirror."""
        return [DATA_FILE]

    def download(self) -> None:
        """Link Table S2 into ``raw/`` after checking the manifest and sha256.

        The mirror plus ``DATA_SHA256`` is canonical; the PMC bucket URL is retrieval
        metadata that ``retrieve_raw_files`` re-runs, never a build input.
        """
        data_root = _data_root()
        manifest = load_manifest(data_root)
        raw = RAW_FILES_BY_NAME[DATA_FILE]
        check_manifest_pin(
            raw.relpath, manifest_sha256(manifest, raw.relpath), raw.sha256
        )
        src = raw_mirror_dir(data_root) / raw.relpath
        if not src.exists():
            raise RuntimeError(f"required raw artifact missing from mirror: {src}")
        os.makedirs(self.raw_dir, exist_ok=True)
        link_verified(src, osp.join(self.raw_dir, raw.name), raw.sha256)
        log.info("Wang 2024 Table S2 linked into %s (sha256 verified)", self.raw_dir)

    def _genome(self) -> EcoliK12Genome:
        """The injected genome, or the reference strain's default cache (a direct run);
        a genome of another assembly set is refused.
        """
        if self.ecoli_genome is None:
            self.ecoli_genome = bacterial_genome("ecoli", self.REFERENCE_STRAIN)
        expected = BACTERIAL_ASSEMBLY_SETS[self.REFERENCE_STRAIN]
        if self.ecoli_genome.ASSEMBLY_SET != expected:
            raise ValueError(
                f"{type(self).__name__} needs the {expected} genome, got "
                f"{self.ecoli_genome.ASSEMBLY_SET}"
            )
        return self.ecoli_genome

    @post_process
    def process(self) -> None:
        """Parse Table S2 into per-gene, per-condition records + LMDB, checking it."""
        verify_raw_files(self.raw_dir, DATA_SHA256)
        frames = read_table_s2(osp.join(self.raw_dir, DATA_FILE))
        b_numbers = check_sheets_align(frames)
        kept, identifiers = resolve_b_numbers(
            self._genome(), b_numbers, label=f"{self.name} Table S2 #Orf"
        )
        drop_log = build_drop_log(self.name, frames, b_numbers, kept)
        adjustment = adjustment_back_solve(frames)
        log.info(
            "Wang 2024: %d genes x %d conditions = %d cells -> %d records; dropped %s",
            drop_log.source_loci,
            len(CONDITIONS),
            drop_log.source_records,
            drop_log.kept_records,
            {rule.rule: rule.n_records for rule in drop_log.rules},
        )

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        out = Path(self.preprocess_dir)
        (out / "dropped_records.json").write_text(drop_log.model_dump_json(indent=2))
        (out / "identifier_reconciliation.json").write_text(
            identifiers.model_dump_json(indent=2)
        )
        (out / "adjustment_back_solve.json").write_text(
            adjustment.model_dump_json(indent=2)
        )
        (out / "perturbation_field_gaps.json").write_text(
            json.dumps(
                [gap.model_dump(mode="json") for gap in PERTURBATION_FIELD_GAPS],
                indent=2,
            )
        )

        environments = {spec.sheet: environment(spec) for spec in CONDITIONS}
        genome_reference = reference_genome()
        references = {
            spec.sheet: build_reference(
                self.name, genome_reference, spec, environments[spec.sheet]
            )
            for spec in CONDITIONS
        }
        env_out, interned_env = self._open_write_lmdb(
            osp.join(self.processed_dir, "lmdb")
        )
        index = 0
        with env_out.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for cell in tqdm(
                stored_cells(frames, kept), total=drop_log.kept_records, desc="wang2024"
            ):
                sheet = cell.spec.sheet
                experiment = build_experiment(self.name, cell, environments[sheet])
                txn.put(
                    f"{index}".encode(),
                    self._intern_record(
                        experiment, references[sheet], PUBLICATION, itxn
                    ),
                )
                index += 1
        env_out.close()
        interned_env.close()
        if index != drop_log.kept_records:
            raise RuntimeError(
                f"wrote {index} records, the ledger counted {drop_log.kept_records}"
            )
        log.info("Wrote %d Wang 2024 environment-response experiments to LMDB", index)

    def preprocess_raw(
        self, df: pd.DataFrame, preprocess: dict[str, Any] | None = None
    ) -> pd.DataFrame:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled by ``build_experiment``."""
        raise NotImplementedError


# --------------------------------------------------------------------------- #
# Verification (L0-L4) of a built tree
# --------------------------------------------------------------------------- #
#: Records of the full build: 26,514 cells minus 1,721 no-read and 30 identifier drops.
EXPECTED_RECORDS = 24763


def _l2_released_test(
    records: Iterable[Mapping[str, Any]], frames: Mapping[str, pd.DataFrame]
) -> LevelResult:
    """L2: every record carries its cell's released p-value pair verbatim (#863).

    Each record is looked up in the workbook by (sheet = ``screen_id``, ``#Orf``), and
    both stored p-values must equal the released cells exactly, with the back-solved
    correction named.
    """
    released = {
        spec.sheet: frames[spec.sheet].set_index("#Orf")[["p-value", "Adj. p-value"]]
        for spec in CONDITIONS
    }
    n_records = 0
    missing: list[str] = []
    mismatched: list[str] = []
    for record in records:
        n_records += 1
        experiment = record["experiment"]
        ph = experiment["phenotype"]
        tag = experiment["genotype"]["perturbations"][0]["systematic_gene_name"]
        key = f"{ph['screen_id']}/{tag}"
        stored = (
            ph["environment_response_p_value"],
            ph["environment_response_p_value_adjusted"],
            ph["p_value_adjustment_method"],
        )
        if None in stored:
            missing.append(key)
            continue
        row = released[ph["screen_id"]].loc[tag]
        if stored != (
            float(row["p-value"]),
            float(row["Adj. p-value"]),
            P_VALUE_ADJUSTMENT_METHOD,
        ):
            mismatched.append(key)
    return LevelResult(
        level=Level.L2,
        name="released_test_fidelity",
        passed=n_records > 0 and not missing and not mismatched,
        message=f"{n_records - len(missing) - len(mismatched)} of {n_records} records "
        "carry their cell's released p-value and Adj. p-value verbatim, corrected by "
        f"{P_VALUE_ADJUSTMENT_METHOD}",
        details={
            "n_records": n_records,
            "n_missing": len(missing),
            "n_mismatched": len(mismatched),
            "examples": (missing + mismatched)[:10],
        },
    )


def _l3_adjustment_back_solve(evidence: AdjustmentBackSolve) -> LevelResult:
    """L3: the released ``Adj. p-value`` is its sheet's own Benjamini-Hochberg value."""
    return LevelResult(
        level=Level.L3,
        name="p_value_adjustment_back_solve",
        passed=evidence.max_abs_deviation <= evidence.tolerance,
        message=f"{evidence.method} reproduces Adj. p-value over {evidence.sheets} "
        f"sheets x {evidence.rows_per_sheet} rows to {evidence.max_abs_deviation:.3g}",
        details=evidence.model_dump(),
    )


def verify_build(
    dataset_root: str,
    *,
    genome: EcoliK12Genome | None = None,
    data_root: str | None = None,
    expected_count: int | None = None,
) -> VerificationReport:
    """Run the environment-response L0-L4 gate on a built tree and write its report.

    The LMDB is streamed once and checked against the MG1655 genome the references
    pin (its resolver, and every GenBank locus as the L4 universe). A second stream
    checks every record's released p-value pair against the tree's own linked Table S2
    (L2 ``released_test_fidelity``), and the correction is back-solved from that
    workbook (L3 ``p_value_adjustment_back_solve``). Every module-level
    ``SourcedValue`` is additionally audited against the raw mirror, where this
    paper's quote anchors live. The report is written to
    ``preprocess/verification_report.json``.
    """
    from torchcell.verification.environment_response import (
        verify_environment_response_dataset_streaming,
    )
    from torchcell.verification.runners import stream_records

    base = data_root or _data_root()
    if expected_count is None:
        expected_count = EXPECTED_RECORDS
    if genome is None:
        genome = bacterial_genome("ecoli", REFERENCE_STRAIN_NAME, base)
    report = verify_environment_response_dataset_streaming(
        stream_records(dataset_root),
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=Provenance(
            source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/data/{DATA_FILE}",
            citation_key=CITATION_KEY,
            sha256=DATA_SHA256[DATA_FILE],
            method="Table S2, TRANSIT resampling of every MG1655 gene in six "
            "rifampicin conditions (2, 32, 160 mg/L = 0.25x, 4x, 20x MIC, at 1 and "
            "3 h); one BacterialEnvironmentResponseExperiment per (gene, condition), "
            "gene-level Tn5 TransposonInsertionPerturbation against MG1655, reference "
            "= the unperturbed parent at log2FC 0; cells with no reads in either pool "
            "dropped",
            page="Table S2 (spectrum.02895-23-s0002.xlsx), all six sheets",
            retrieved=RETRIEVED_AT,
        ),
        expected_count=expected_count,
        sgd_genes=set(genome.genbank.loci),
        resolve_gene_name=genome.resolve_gene_name,
    )
    frames = read_table_s2(osp.join(dataset_root, "raw", DATA_FILE))
    report.add(_l2_released_test(stream_records(dataset_root), frames))
    report.add(_l3_adjustment_back_solve(adjustment_back_solve(frames)))
    mirror = Path(base) / "torchcell-raw"
    for value in SOURCED_VALUES.values():
        report.add(audit_sourced_value(value, mirror))
    preprocess = osp.join(dataset_root, "preprocess")
    os.makedirs(preprocess, exist_ok=True)
    with open(osp.join(preprocess, "verification_report.json"), "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main(argv: list[str] | None = None) -> int:
    """CLI: ``deposit`` the raw mirror, ``build`` the dev LMDB, or ``verify`` it."""
    from dotenv import load_dotenv

    parser = argparse.ArgumentParser(prog="python -m torchcell.datasets.ecoli.wang2024")
    sub = parser.add_subparsers(dest="command", required=True)
    deposit = sub.add_parser(
        "deposit", help="retrieve (optional) and deposit the raw mirror"
    )
    deposit.add_argument("--download-dir", required=True)
    deposit.add_argument(
        "--retrieve",
        action="store_true",
        help="run the recorded PMC bucket retrievers into --download-dir first",
    )
    sub.add_parser("build", help="build (or load) the dev-tree LMDB")
    sub.add_parser("verify", help="run L0-L4 on the built dev-tree LMDB")
    args = parser.parse_args(argv)

    load_dotenv()
    data_root = _data_root()
    if args.command == "deposit":
        download = Path(args.download_dir)
        if args.retrieve:
            retrieve_raw_files(download)
        sources: dict[str, str | Path] = {
            raw.name: download / raw.name for raw in RETRIEVED_FILES
        }
        print(deposit_raw_mirror(sources=sources, data_root=data_root))
        return 0
    root = osp.join(data_root, DATASET_ROOT_REL)
    if args.command == "build":
        dataset = EnvChemgenWang2024Dataset(root=root)
        print(f"len = {len(dataset)}")
        return 0
    report = verify_build(root, data_root=data_root)
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
