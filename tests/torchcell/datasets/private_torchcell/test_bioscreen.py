# tests/torchcell/datasets/private_torchcell/test_bioscreen.py
# [[tests.torchcell.datasets.private_torchcell.test_bioscreen]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/private_torchcell/test_bioscreen.py
"""Bioscreen C parsing and derivation for the 2021 inhibitor runs.

Synthetic curves pin :func:`generation_time` (the oracle is the doubling time the curve
was built with); synthetic dose tables pin the layout counts. The archive-backed tests
skip when ``/bulk/thesis/thesis_archive`` is not mounted and pin the numbers measured on
2026-10-08 against the archive copies.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import pytest
from scipy.stats import spearmanr

from torchcell.datasets.private_torchcell import bioscreen as b
from torchcell.literature.manifest import Manifest, RetrievalMethod
from torchcell.literature.provenance import check_source, verify_artifact

ARCHIVE = b.ARCHIVE_ROOT
needs_archive = pytest.mark.skipif(
    not (ARCHIVE / b.ARCHIVE_MANIFEST).is_file(),
    reason="the thesis archive on /bulk is not mounted",
)


def _f(values: object) -> b.FloatArray:
    """A float64 array, the type ``generation_time`` takes."""
    return np.asarray(values, dtype=np.float64)


HOURS = _f(np.arange(0.0, 72.0, 0.25))


def _lag_then_exponential(
    onset_h: float, doubling_h: float, jump: float = 0.06, base: float = 0.1
) -> b.FloatArray:
    """Flat at ``base`` until ``onset_h``, then ``base + jump * 2**((t - onset)/d)``."""
    rise = np.where(
        HOURS >= onset_h, jump * 2.0 ** ((HOURS - onset_h) / doubling_h), 0.0
    )
    return _f(np.minimum(base + rise, 2.0))


# --------------------------------------------------------------------------- #
# generation_time on synthetic curves
# --------------------------------------------------------------------------- #
def test_a_clean_exponential_gives_its_doubling_time() -> None:
    gt = b.generation_time(HOURS, _lag_then_exponential(onset_h=4.0, doubling_h=2.0))
    assert gt == pytest.approx(2.0, rel=1e-9)


def test_a_flat_curve_did_not_grow() -> None:
    assert b.generation_time(HOURS, _f(np.full_like(HOURS, 0.12))) is None


def test_a_rise_below_the_growth_threshold_did_not_grow() -> None:
    od = 0.1 + np.where(HOURS > 10, 0.25, 0.0)
    assert b.generation_time(HOURS, _f(od)) is None


def test_a_late_riser_is_found() -> None:
    gt = b.generation_time(HOURS, _lag_then_exponential(onset_h=50.0, doubling_h=3.0))
    assert gt == pytest.approx(3.0, rel=1e-9)


def test_the_steepest_window_wins() -> None:
    """Doubling every 6 h from 2 h, every 1.5 h from 20 h: 1.5 h is reported."""
    slow = (np.clip(HOURS, 2.0, 20.0) - 2.0) / 6.0
    fast = np.clip(HOURS - 20.0, 0.0, None) / 1.5
    od = 0.1 + np.where(HOURS >= 2.0, 0.06 * 2.0 ** (slow + fast), 0.0)
    gt = b.generation_time(HOURS, _f(np.minimum(od, 50.0)))
    assert gt == pytest.approx(1.5, rel=1e-9)


# --------------------------------------------------------------------------- #
# Layouts from synthetic inputs
# --------------------------------------------------------------------------- #
def _synthetic_ex21_doses() -> dict[b.Inhibitor, list[float]]:
    return {i: [float(9 - k) for k in range(9)] for i in b.INHIBITORS}


def test_ex21_layout_counts() -> None:
    plate = b.ex21_layout_from(_synthetic_ex21_doses())
    assert len(plate.wells) == 6 * 9 * 3 + 18
    assert len(plate.wild_type()) == 6 * 3
    for inhibitor in b.INHIBITORS:
        treated = [w for w in plate.wells if w.present() == [inhibitor]]
        assert len(treated) == 27
        assert {w.biological_replicate_id for w in treated} == {1, 2, 3}
    assert not {w.well for w in plate.wells} & {*range(91, 101), *range(191, 201)}


def test_ex21_blocks_open_with_the_uninhibited_well() -> None:
    plate = {w.well: w for w in b.ex21_layout_from(_synthetic_ex21_doses()).wells}
    assert plate[31].is_wild_type and plate[41].is_wild_type
    assert plate[32].doses_g_per_l[b.Inhibitor.AA] == 9.0
    assert plate[40].doses_g_per_l[b.Inhibitor.AA] == 1.0
    assert plate[101].plate == 2 and plate[101].is_wild_type


def test_ex23_layout_counts_and_doses() -> None:
    rows = [b.Ex23Row(name="WT", well=1, biological_replicate_id=1)]
    rows += [b.Ex23Row(name="blank", well=2, biological_replicate_id=1)]
    rows += [b.Ex23Row(name="FF_AA_HMF", well=3, biological_replicate_id=2)]
    plate = b.ex23_layout_from(rows)
    assert [w.well for w in plate.wells] == [1, 3]
    combo = plate.wells[1]
    assert combo.doses_g_per_l == {
        b.Inhibitor.FF: 1.5,
        b.Inhibitor.AA: 2.0,
        b.Inhibitor.HMF: 2.522,
        b.Inhibitor.FA: 0.0,
        b.Inhibitor.LVA: 0.0,
        b.Inhibitor.LA: 0.0,
    }
    assert combo.condition == "FF_AA_HMF"


def test_ex23_rejects_a_repeated_inhibitor() -> None:
    with pytest.raises(ValueError, match="named twice"):
        b.ex23_doses("FF_FF")


@pytest.mark.parametrize("run", b.ISOBOLE_RUNS)
def test_isobole_layout_is_two_ten_by_ten_plates(run: b.Run) -> None:
    plate = b.isobole_layout(run)
    assert len(plate.wells) == 200
    assert [w.well for w in plate.wild_type()] == [1, 101]
    inhibitor, step = b.ISOBOLES[run]
    corner = {w.well: w for w in plate.wells}[100]
    assert corner.doses_g_per_l[inhibitor] == round(9 * step, 3)
    assert corner.doses_g_per_l[b.Inhibitor.AA] == 3.6
    well_11 = {w.well: w for w in plate.wells}[11]
    assert well_11.present() == [b.Inhibitor.AA]
    assert well_11.doses_g_per_l[b.Inhibitor.AA] == 0.4


def test_a_well_named_wt_must_be_uninhibited() -> None:
    with pytest.raises(ValueError, match="only the uninhibited well"):
        b.WellLayout(
            run=b.Run.ex23,
            plate=1,
            well=5,
            condition="WT",
            doses_g_per_l=b.no_doses() | {b.Inhibitor.FF: 1.5},
            biological_replicate_id=1,
        )


def test_grew_must_match_the_generation_time() -> None:
    with pytest.raises(ValueError, match="grew must be True exactly"):
        b.Well(
            run=b.Run.ex26,
            plate=1,
            well=1,
            condition="WT",
            doses_g_per_l=b.no_doses(),
            biological_replicate_id=1,
            generation_time_h=None,
            grew=True,
            trait_source=b.TraitSource.raw_curve,
        )


def test_relative_growth_rate_and_wild_type_mean() -> None:
    def well(n: int, cond: str, gt: float | None) -> b.Well:
        d = b.no_doses() if cond == "WT" else b.no_doses() | {b.Inhibitor.AA: 2.0}
        return b.Well(
            run=b.Run.ex23,
            plate=1,
            well=n,
            condition=cond,
            doses_g_per_l=d,
            biological_replicate_id=1,
            generation_time_h=gt,
            grew=gt is not None,
            trait_source=b.TraitSource.raw_curve,
        )

    wells = [well(1, "WT", 2.0), well(2, "WT", 3.0), well(3, "WT", None)]
    wells += [well(4, "AA2", 5.0), well(5, "AA2", None)]
    wt = b.wild_type_generation_time(wells)
    assert wt == 2.5
    assert b.relative_growth_rate(wells[3], wt) == 0.5
    assert b.relative_growth_rate(wells[4], wt) is None
    with pytest.raises(ValueError, match="no uninhibited well grew"):
        b.wild_type_generation_time([wells[2], wells[3]])


# --------------------------------------------------------------------------- #
# Archive-backed
# --------------------------------------------------------------------------- #
@needs_archive
def test_ex21_doses_come_from_the_dispensed_volumes() -> None:
    doses = b.ex21_doses(b.titration_volumes_ul(ARCHIVE), b.stock_g_per_l(ARCHIVE))
    assert doses[b.Inhibitor.FF] == [30, 24, 18, 12, 6, 3, 1.5, 0.75, 0.375]
    assert doses[b.Inhibitor.FA] == [5, 4, 3, 2, 1, 0.5, 0.25, 0.125, 0.063]
    assert doses[b.Inhibitor.HMF][4] == 2.522
    # The notebook labels agree with the volumes except its two typos.
    for inhibitor, labels in b.EX21_LABELS.items():
        mismatched = [
            (label, dose)
            for label, dose in zip(labels, doses[inhibitor], strict=True)
            if abs(float(label) - dose) > 0.006
        ]
        expected = {b.Inhibitor.FF: [("28", 18.0)], b.Inhibitor.FA: [("1.25", 0.125)]}
        assert mismatched == expected.get(inhibitor, [])


@needs_archive
def test_ex23_layout_from_the_archive() -> None:
    plate = b.ex23_layout(ARCHIVE)
    assert len(plate.wells) == 63 * 3 + 8
    assert len(plate.wild_type()) == 8
    combos = {w.condition for w in plate.wells if not w.is_wild_type}
    assert len(combos) == 63


@needs_archive
def test_ex23_doses_are_the_well_map_volumes_over_1700_ul() -> None:
    import pandas as pd

    stocks = b.stock_g_per_l(ARCHIVE)
    well_map = pd.read_excel(ARCHIVE / b.EX23_WELL_MAP)
    for inhibitor in b.INHIBITORS:
        volumes = set(well_map[f"{inhibitor.value}_v_ul"]) - {0}
        assert len(volumes) == 1
        dose = volumes.pop() * stocks[inhibitor] / b.EX23_TUBE_UL
        assert dose == pytest.approx(b.EX23_G_PER_L[inhibitor], abs=1e-9)
    total = well_map[[f"{i.value}_v_ul" for i in b.INHIBITORS] + ["ypd_v_ul"]].sum(1)
    assert set(total) == {b.EX23_TUBE_UL}


@needs_archive
def test_isobole_steps_match_the_design_sheet() -> None:
    import openpyxl

    book = openpyxl.load_workbook(ARCHIVE / b.ISOBOLE_DESIGN_XLSX, data_only=True)
    sheets = {b.Run.ex26: "FF_AA_new ", b.Run.ex27: "FA_AA", b.Run.ex28: "HMF_AA"}
    for run, sheet in sheets.items():
        ws = book[sheet]
        per_10_ul = 10 / ws["B12"].value
        assert ws["D3"].value * per_10_ul == pytest.approx(b.ISOBOLES[run][1], abs=1e-3)
        assert ws["D7"].value * per_10_ul == pytest.approx(b.AA_STEP, abs=1e-12)


@needs_archive
def test_ex26_raw_curve_agrees_with_the_software_traits() -> None:
    raw = {w.well: w for w in b.derive_wells(b.Run.ex26, ARCHIVE)}
    sw = {
        w.well: w
        for w in b.derive_wells(b.Run.ex26, ARCHIVE, b.TraitSource.bioscreen_software)
    }
    agree = np.mean([raw[k].grew == sw[k].grew for k in raw])
    both = [k for k in raw if raw[k].grew and sw[k].grew]
    rho = spearmanr(
        [sw[k].generation_time_h for k in both],
        [raw[k].generation_time_h for k in both],
    ).statistic
    assert len(raw) == 200
    assert agree == pytest.approx(0.985, abs=1e-12)
    assert len(both) == 45
    assert rho == pytest.approx(0.8043, abs=1e-4)


@needs_archive
def test_wells_grown_per_run_from_the_raw_curves() -> None:
    grown = {run: sum(w.grew for w in b.derive_wells(run, ARCHIVE)) for run in b.Run}
    assert grown == {
        b.Run.ex21: 117,
        b.Run.ex23: 69,
        b.Run.ex26: 45,
        b.Run.ex27: 76,
        b.Run.ex28: 53,
    }


def _ex23_single_means(source: b.TraitSource) -> dict[b.Inhibitor, float]:
    wells = b.derive_wells(b.Run.ex23, ARCHIVE, source)
    wt = b.wild_type_generation_time(wells)
    out: dict[b.Inhibitor, float] = {}
    for inhibitor in b.INHIBITORS:
        rates = [
            b.relative_growth_rate(w, wt)
            for w in wells
            if w.condition == inhibitor.value
        ]
        assert len(rates) == 3
        assert all(r is not None for r in rates)
        out[inhibitor] = float(np.mean([r for r in rates if r is not None]))
    return out


@needs_archive
def test_ex23_single_inhibitor_rates_from_the_software_traits() -> None:
    """The thesis's own numbers (039 note, Fig 11b bars): WT mean 1.8146 h over 8 wells."""
    wells = b.derive_wells(b.Run.ex23, ARCHIVE, b.TraitSource.bioscreen_software)
    assert b.wild_type_generation_time(wells) == pytest.approx(1.8146, abs=1e-4)
    expected = {"AA": 0.947, "FA": 0.937, "LVA": 0.873, "FF": 0.769, "LA": 0.767}
    expected |= {"HMF": 0.437}
    means = _ex23_single_means(b.TraitSource.bioscreen_software)
    for inhibitor, value in expected.items():
        assert means[b.Inhibitor(inhibitor)] == pytest.approx(value, abs=1e-3)


@needs_archive
def test_ex23_single_inhibitor_rates_from_the_raw_curves() -> None:
    """Measured 2026-10-08; differs from the software on FF and LA (module docstring)."""
    wells = b.derive_wells(b.Run.ex23, ARCHIVE)
    assert b.wild_type_generation_time(wells) == pytest.approx(1.345, abs=1e-3)
    expected = {"AA": 0.998, "FA": 0.996, "LVA": 0.869, "FF": 0.507, "LA": 0.982}
    expected |= {"HMF": 0.480}
    means = _ex23_single_means(b.TraitSource.raw_curve)
    for inhibitor, value in expected.items():
        assert means[b.Inhibitor(inhibitor)] == pytest.approx(value, abs=1e-3)


@needs_archive
def test_software_traits_do_not_exist_for_ex27_and_ex28() -> None:
    for run in (b.Run.ex27, b.Run.ex28):
        with pytest.raises(ValueError, match="never processed"):
            b.software_generation_times(run, ARCHIVE)


@needs_archive
@pytest.mark.parametrize("run", list(b.Run))
def test_tray_temperature_is_logged_at_30_c(run: b.Run) -> None:
    log = b.read_temperature_log(ARCHIVE / b.RAW_BSM[run.value])
    summary = b.summarize_tray_temperature(run, log)
    assert summary.tray_median_c == 30.01
    assert 29.98 <= summary.tray_p01_c <= summary.tray_p99_c <= 30.02
    assert summary.warmup_readings_dropped == 4


@needs_archive
def test_bsm_temperature_fields_match_the_labeled_ex1_log() -> None:
    log = b.read_temperature_log(ARCHIVE / b.EX1_BSM)
    text = (ARCHIVE / b.EX1_TEMPERATURE_TXT).read_text(encoding="utf-16")
    rows = [line.split() for line in text.splitlines()[1:6]]
    for k, (_, tray, cover) in enumerate(rows):
        assert round(log.tray_c[k] + 1e-9, 1) == float(tray)
        assert round(log.cover_c[k] + 1e-9, 1) == float(cover)


@needs_archive
def test_run_settings_and_durations() -> None:
    settings = b.read_run_settings(ARCHIVE / b.RAW_BSM["ex23"])
    assert settings["4"] == "900"
    assert settings["3"] == "345600"
    durations = {run: round(b.run_duration_h(ARCHIVE, run), 2) for run in b.Run}
    assert durations == {
        b.Run.ex21: 71.97,
        b.Run.ex23: 84.97,
        b.Run.ex26: 95.99,
        b.Run.ex27: 95.98,
        b.Run.ex28: 95.98,
    }


@needs_archive
def test_deposit_raw_mirror_round_trips(tmp_path: Path) -> None:
    root = b.deposit_raw_mirror(tmp_path)
    manifest = Manifest.model_validate_json((root / "manifest.json").read_text())
    listed = b.read_archive_manifest(ARCHIVE)
    assert [f.path for f in manifest.files] == list(b.RAW_FILES)
    for record in manifest.files:
        assert verify_artifact(record, root)
        assert record.retrieval is not None
        assert record.retrieval.method == RetrievalMethod.local_archive
        assert record.retrieval.source_url == listed[record.path].source_path
        assert check_source(record.retrieval, now="2026-10-08").matches
    # A second deposit leaves the mirror as it is.
    assert b.deposit_raw_mirror(tmp_path) == root


@needs_archive
def test_deposit_refuses_a_mirror_file_with_other_bytes(tmp_path: Path) -> None:
    dest = b.raw_mirror_dir(tmp_path) / b.RAW_FILES[0]
    dest.parent.mkdir(parents=True)
    dest.write_bytes(b"not the export")
    with pytest.raises(RuntimeError, match="different sha256"):
        b.deposit_raw_mirror(tmp_path)


@needs_archive
def test_deposit_report_pins_the_pdf(tmp_path: Path) -> None:
    key_dir = b.deposit_report(tmp_path)
    manifest = Manifest.model_validate_json((key_dir / "manifest.json").read_text())
    (pdf,) = manifest.files
    assert pdf.path == "paper.pdf" and pdf.role == "paper_pdf"
    assert pdf.sha256 == b.REPORT_PDF_SHA256
    assert hashlib.sha256((key_dir / "paper.pdf").read_bytes()).hexdigest() == (
        b.REPORT_PDF_SHA256
    )
    assert pdf.retrieval is not None
    assert pdf.retrieval.method == RetrievalMethod.local_archive
    assert manifest.title == b.REPORT_TITLE
    assert manifest.doi is None and manifest.zotero_item_key is None
