# tests/torchcell/datasets/scerevisiae/test_hoepfner2014_synthetic.py
# [[tests.torchcell.datasets.scerevisiae.test_hoepfner2014_synthetic]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/scerevisiae/test_hoepfner2014_synthetic.py
"""Hoepfner 2014 HIP-HOP loader built end to end on synthetic score matrices.

``process()`` reads two things besides ``raw/``: the committed Table S5 strain CSV of
experiment 036 (sha256-pinned, read from the worktree) and ``DATA_ROOT`` (pointed into
``tmp_path``). No genome is read (2026.10.12, the data is the data): every row is stored
under its released spelling. ``Table_S1.xls`` is written with openpyxl under the
release's ``.xls`` name; pandas sniffs the zip signature, so ``_load_compound_meta`` runs
unchanged.

Table S1: CMB 3 Amitriptyline (SMILES released, curated: InChIKey
KRMDCWKBEZIMAB-UHFFFAOYSA-N, PubChem 2160, ChEBI:2666), CMB 409 Boromycin (released
SMILES RDKit cannot parse, no curated identifier: STORED with the inchikey gap and its
SMILES), CMB 777 (structure only, no name: resolves from its SMILES as ``CMB777``), CMB
888 Unreleased (name only, no SMILES: STORED with the inchikey gap), plus a ``244, 1818``
multi-id structure row, an empty-id row, a nameless CMB 1234 row and a NaN-SMILES row
that exercise the cell parser.

HIP header columns (index: header): 1 ``Ad. 3_50_HIP_0077`` (n = 2), 2 its z-score
companion (skipped), 3 ``MADL 3_50_HIP_0091`` (n = 1), 4 ``Ad. 409_1.2_HIP_0077``, 5
``Ad. 777_10_HIP_0077``, 6 ``Ad. 888_10_HIP_0077``, 7 ``Ad. 3_50_HOP_0077`` (other assay,
skipped). HOP header: 1 ``Ad. 3_50_HOP_0077``, 2 ``Ad. 409_1.2_HOP_0077``. Every
sensitivity column is stored.

HIP rows (cells at columns 1..7) and what they become:

    YAL001C    0.1 0.11 -1.5 0.9 2.25 0.3 9.9   records 0-4
    YAL035C-A  0.2 0.21 ""   ""  0.4  0.3 9.9   R64 alias of YAL034C-B, stored as
                                                released: records 5-7
    YCL074W    0.5 0.5  0.5  0.5 ""   0.5 9.9   pseudogene, stored: records 8-11
    YBR271W   -3.0 -3.1 -2.0 -1.0 -0.5 0.3 9.9  Table S5 positional: records 12-16
    YHR999W    0.7 (all seven)                  no such R64 gene, stored: records 17-21
    YLR074C    0.6 "" 0.6 "" 0.6 (line ends)    Table S5 CL1+ MUT: records 22-24
    YAL034C-B  0.8 (line ends)                  scored in 1 of 5 columns: record 25
    R0010W     0.9 (line ends)                  not ORF-shaped: held out, 1 cell

HOP rows: YAL001C 1.0 0.7 (records 26, 27), YBR271W -0.25 "" (record 28), YAL034C-B 0.33
with no second cell (record 29; the short line exercises the column-bound guard), and
YAL001C again 0.99 "" (record 30; a repeated ORF row is stored, see the Finding test).

Ledger arithmetic: stored HIP 26 + HOP 5 = 31; HIP scored cells 27 = 26 stored + 1 held
out (R0010W); HOP 5 = 5 stored. Boromycin's identity gap covers HIP 4 records (YAL001C,
YCL074W, YBR271W, YHR999W) and HOP 1; CMB 888's covers HIP 5. References are one per
(assay, study): HIP_0091 members [1, 9, 13, 18, 23], HIP_0077 the other HIP records,
HOP_0077 [26-30]. Table S5 flags: YBR271W 5 HIP records, YLR074C 3. Table S5 lists five
of the seven stored HIP rows once (YAL035C-A as ``YAL035C-a``); YHR999W and YAL034C-B
are not listed and carry the construction gap. The records are
``StrainEnvironmentResponseExperiment``s on the BY4743 ``StrainReferenceGenome`` (#506).
"""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from torchcell.data import RawSha256MismatchError
from torchcell.data.experiment_dataset import verify_raw_files
from torchcell.datamodels.media import YPD_LIQUID
from torchcell.datamodels.schema import (
    AssayType,
    BarcodedKanMxDeletionPerturbation,
    Compound,
    Concentration,
    ConcentrationUnit,
    CultureEnvironment,
    CultureFormat,
    DoseBasis,
    EndpointRule,
    EnvironmentResponsePhenotype,
    Genotype,
    HeterozygousDeletionPerturbation,
    MeasurementType,
    PreCulture,
    PreCultureSource,
    Publication,
    SampleUnit,
    SmallMoleculePerturbation,
    Solvent,
    StrainConstruction,
    StrainEnvironmentResponseExperiment,
    StrainEnvironmentResponseExperimentReference,
    Temperature,
)
from torchcell.datasets.scerevisiae import hoepfner2014 as m
from torchcell.verification.sourced import ProvenanceGap, ProvenanceGapReason

_CITATION_KEY = "hoepfnerHighresolutionChemicalDissection2014"
_DOI = "10.1016/j.micres.2013.11.004"
_DATASET = "EnvChemgenHoepfner2014Dataset"
_AMITRIPTYLINE = "CN(C)CCC=C2c1ccccc1CCc3ccccc23"
_BOROMYCIN = (
    "CC(C)C(N)C(=O)OC(C)C1C/C=C\\CCC(O)C(C)(C)C7CCC(C)C4(OB25OC(C(=O)O1)C3(O2)"
    "OC(CCC3C)C(C)(C)C(O)CCCC6CC(OC(=O)C4O5)C(C)O6)O7"
)

_HIP_HEADER = [
    "Systematic Name",
    "Ad. scores for Exp. 3_50_HIP_0077",
    "Ad. scores for Exp. 3_50_HIP_0077 z-score",
    "MADL scores for Exp. 3_50_HIP_0091",
    "Ad. scores for Exp. 409_1.2_HIP_0077",
    "Ad. scores for Exp. 777_10_HIP_0077",
    "Ad. scores for Exp. 888_10_HIP_0077",
    "Ad. scores for Exp. 3_50_HOP_0077",
]
_HIP_ROWS: list[list[str]] = [
    ["YAL001C", "0.1", "0.11", "-1.5", "0.9", "2.25", "0.3", "9.9"],
    ["YAL035C-A", "0.2", "0.21", "", "", "0.4", "0.3", "9.9"],
    ["YCL074W", "0.5", "0.5", "0.5", "0.5", "", "0.5", "9.9"],
    ["YBR271W", "-3.0", "-3.1", "-2.0", "-1.0", "-0.5", "0.3", "9.9"],
    ["YHR999W"] + ["0.7"] * 7,
    ["YLR074C", "0.6", "", "0.6", "", "0.6"],
    ["YAL034C-B", "0.8"],
    ["R0010W", "0.9"],
]
_HOP_HEADER = [
    "Systematic Name",
    "Ad. scores for Exp. 3_50_HOP_0077",
    "Ad. scores for Exp. 409_1.2_HOP_0077",
]
_HOP_ROWS: list[list[str]] = [
    ["YAL001C", "1.0", "0.7"],
    ["YBR271W", "-0.25", ""],
    ["YAL034C-B", "0.33"],
    ["YAL001C", "0.99", ""],
]


def _quoted_matrix(header: list[str], rows: list[list[str]]) -> str:
    """The deposited format: every cell double-quoted, tab-separated."""
    lines = ["\t".join(f'"{cell}"' for cell in header)]
    lines += ["\t".join(f'"{cell}"' for cell in row) for row in rows]
    return "\n".join(lines) + "\n"


def _write_table_s1(path: Path) -> None:
    workbook = openpyxl.Workbook()
    known = workbook.active
    known.title = "Reference Substances known MoA"
    known.append(["CMB ID", "Common Name", "IC30 (uM)"])
    known.append([3, "Amitriptyline", 50])
    known.append([None, "Nameless", 5])
    known.append([409, "Boromycin", 1.2])
    known.append([888, "Unreleased", 10])
    novel = workbook.create_sheet("Substances novel MoA")
    novel.append(["CMB ID", "Common Name"])
    novel.append(["1234", None])
    novel.append([1235, "Novel compound"])
    structures = workbook.create_sheet("All Structures")
    structures.append(["CMB ID", "SMILE string"])
    structures.append([3, _AMITRIPTYLINE])
    structures.append(["409", _BOROMYCIN])
    structures.append([777, _AMITRIPTYLINE])
    structures.append(["244, 1818", "CC(=O)O"])
    structures.append([555, None])
    workbook.save(path)


def _write_raw(raw: Path) -> None:
    raw.mkdir(parents=True, exist_ok=True)
    (raw / "HIP_scores.txt").write_text(_quoted_matrix(_HIP_HEADER, _HIP_ROWS))
    (raw / "HOP_scores.txt").write_text(_quoted_matrix(_HOP_HEADER, _HOP_ROWS))
    _write_table_s1(raw / "Table_S1.xls")


@pytest.fixture
def dataset(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> m.EnvChemgenHoepfner2014Dataset:
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    root = tmp_path / "env_chemgen_hoepfner2014"
    _write_raw(root / "raw")
    return m.EnvChemgenHoepfner2014Dataset(root=str(root))


_AMITRIPTYLINE_COMPOUND = Compound(
    name="amitriptyline",
    inchikey="KRMDCWKBEZIMAB-UHFFFAOYSA-N",
    smiles=_AMITRIPTYLINE,
    pubchem_cid=2160,
    chebi_id="CHEBI:2666",
)
_CMB777_COMPOUND = Compound(
    name="CMB777", inchikey="KRMDCWKBEZIMAB-UHFFFAOYSA-N", smiles=_AMITRIPTYLINE
)
_BOROMYCIN_COMPOUND = Compound(
    name="Boromycin",
    smiles=_BOROMYCIN,
    provenance_gaps=[
        ProvenanceGap(
            field="inchikey", reason=ProvenanceGapReason.deferred_pending_source_review
        )
    ],
)
_DMSO = Compound(
    name="dimethyl sulfoxide",
    inchikey="IAZDPXIOMUYVGZ-UHFFFAOYSA-N",
    smiles="CS(=O)C",
    pubchem_cid=679,
    chebi_id="CHEBI:28262",
)
_SE_NOTE = (
    "no per-cell uncertainty is released: the replicate t-test p-value is folded into the "
    "adjusted score a_L = min(0.05/p, 1) * s_L"
)
_MEASUREMENT_UNITS = (
    "adjusted MADL sensitivity score = (r_L - med(r_L)) / MAD(r_L) over all pool strains, "
    "r_L = log ratio of treated vs control strain abundance (Hoepfner 2014); negative = "
    "hypersensitive, positive = resistant; 0 = growth equal to control. A companion "
    "gene-wise z-score column is deposited per experiment but not stored here."
)
_REFERENCE_UNITS = (
    "the screen's own no-drug DMSO control samples, the denominator of the MADL log "
    "ratio; a strain growing exactly as its controls scores 0"
)


def _uncertainty_gaps() -> list[ProvenanceGap]:
    return [
        ProvenanceGap(
            field=field,
            reason=ProvenanceGapReason.not_reported_by_primary,
            note=_SE_NOTE,
        )
        for field in (
            "environment_response_se",
            "environment_response_uncertainty",
            "environment_response_uncertainty_type",
        )
    ]


def _culture_format(assay: str) -> CultureFormat:
    """24-well Greiner 662102, 1600 ul, 550 rpm; HIP 250 cells/strain with a gapped
    endpoint, HOP 320 cells/strain read after a fixed 16 h.
    """
    sv = m.SOURCED_VALUES
    if assay == "HIP":
        return CultureFormat(
            vessel="24-well plate (Greiner 662102)",
            working_volume_ul=1600.0,
            shaking_rpm=550.0,
            inoculum_cells_per_strain=250.0,
            provenance=[sv["culture_vessel"], sv["shaking_rpm"], sv["hip_inoculum"]],
            provenance_gaps=[
                ProvenanceGap(
                    field="endpoint",
                    reason=ProvenanceGapReason.deferred_pending_source_review,
                    resolve_with=m.HOEPFNER_SUPPLEMENT,
                    note=m._HIP_PASSAGE_NOTE + " (Fig. S2)",
                )
            ],
        )
    return CultureFormat(
        vessel="24-well plate (Greiner 662102)",
        working_volume_ul=1600.0,
        shaking_rpm=550.0,
        inoculum_cells_per_strain=320.0,
        endpoint=EndpointRule.fixed_duration,
        provenance=[
            sv["culture_vessel"],
            sv["shaking_rpm"],
            sv["hop_inoculum"],
            sv["hop_duration_hours"],
        ],
    )


def _environment(
    assay: str, perturbation: SmallMoleculePerturbation
) -> CultureEnvironment:
    sv = m.SOURCED_VALUES
    if assay == "HIP":
        return CultureEnvironment(
            media=YPD_LIQUID,
            temperature=Temperature(value=30.0),
            perturbations=[perturbation],
            aerobicity="aerobic",
            culture_format=_culture_format("HIP"),
            pre_culture=PreCulture(
                source=PreCultureSource.log_phase_culture,
                medium=YPD_LIQUID,
                source_label="overnight log phase pre-culture",
                provenance=[sv["hip_inoculum"]],
            ),
            provenance_gaps=[*m._HIP_DURATION_GAPS, m._AUXOTROPH_GAP],
        )
    return CultureEnvironment(
        media=YPD_LIQUID,
        temperature=Temperature(value=30.0),
        perturbations=[perturbation],
        aerobicity="aerobic",
        duration_hours=16.0,
        duration_generations=5.0,
        culture_format=_culture_format("HOP"),
        pre_culture=PreCulture(
            source=PreCultureSource.thaw_recovery,
            medium=YPD_LIQUID,
            duration_hours=3.0,
            source_label="thawed and recovered for 3 h in YPD",
            provenance=[sv["hop_thaw_recovery"]],
        ),
        provenance_gaps=[m._AUXOTROPH_GAP],
    )


def _treated(compound: Compound, micromolar: float) -> SmallMoleculePerturbation:
    return SmallMoleculePerturbation(
        compound=compound,
        concentration=Concentration(
            value=micromolar, unit=ConcentrationUnit.micromolar, basis=DoseBasis.IC30
        ),
        solvent=Solvent(name="DMSO", percent=2.0, compound=_DMSO),
    )


def _phenotype(
    score: float, n_samples: int, study: str
) -> EnvironmentResponsePhenotype:
    return EnvironmentResponsePhenotype(
        measurement_type=MeasurementType.sensitivity_score,
        assay_type=AssayType.pooled_competitive_growth_barcode,
        environment_response=score,
        n_samples=n_samples,
        sample_unit=SampleUnit.technical_replicate,
        units=_MEASUREMENT_UNITS,
        screen_id=study,
        provenance_gaps=_uncertainty_gaps(),
    )


def _hip_genotype(orf: str, construction: StrainConstruction | None) -> Genotype:
    """The released ORF as systematic and perturbed name; a strain Table S5 does not
    list carries the construction gap.
    """
    gaps = (
        [m._BARCODE_GAP]
        if construction is not None
        else [m._BARCODE_GAP, m._CONSTRUCTION_GAP]
    )
    return Genotype(
        perturbations=[
            HeterozygousDeletionPerturbation(
                systematic_gene_name=orf,
                perturbed_gene_name=orf,
                cassette="kanMX4",
                collection="YSC1055",
                construction=construction,
                provenance_gaps=gaps,
            )
        ]
    )


def _hop_genotype(orf: str) -> Genotype:
    return Genotype(
        perturbations=[
            BarcodedKanMxDeletionPerturbation(
                systematic_gene_name=orf,
                perturbed_gene_name=orf,
                cassette="kanMX4",
                collection="YSC1056",
                provenance_gaps=[m._BARCODE_GAP],
            )
        ]
    )


def _reference(assay: str, study: str) -> dict[str, Any]:
    vehicle = SmallMoleculePerturbation(
        compound=_DMSO,
        concentration=Concentration(
            value=2.0, unit=ConcentrationUnit.percent_v_v, basis=DoseBasis.fixed
        ),
    )
    return StrainEnvironmentResponseExperimentReference(
        dataset_name=_DATASET,
        genome_reference=m.hoepfner_reference_genome(),
        environment_reference=_environment(assay, vehicle),
        phenotype_reference=EnvironmentResponsePhenotype(
            measurement_type=MeasurementType.sensitivity_score,
            assay_type=AssayType.pooled_competitive_growth_barcode,
            environment_response=0.0,
            n_samples=4,
            sample_unit=SampleUnit.technical_replicate,
            units=_REFERENCE_UNITS,
            screen_id=study,
            provenance_gaps=_uncertainty_gaps(),
        ),
    ).model_dump()


def _summary(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> list[tuple[str, str, str, str, float, float, int, str]]:
    """(leaf, ORF, perturbed name, compound, dose, score, n, study) per record."""
    out = []
    for i in range(len(dataset)):
        experiment = dataset[i]["experiment"]
        (perturbation,) = experiment["genotype"]["perturbations"]
        (treatment,) = experiment["environment"]["perturbations"]
        phenotype = experiment["phenotype"]
        out.append(
            (
                perturbation["perturbation_type"],
                perturbation["systematic_gene_name"],
                perturbation["perturbed_gene_name"],
                treatment["compound"]["name"],
                treatment["concentration"]["value"],
                phenotype["environment_response"],
                phenotype["n_samples"],
                phenotype["screen_id"],
            )
        )
    return out


def test_thirty_one_records_stream_hip_then_hop_in_row_then_column_order(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Stored HIP columns are header indices 1, 3, 4, 5, 6, so each HIP row yields up to
    five records (fewer where a cell is empty or the line ends), then the HOP rows follow
    with their two columns; the ``MADL`` column stores n = 1, the ``Ad.`` columns n = 2.
    Every row is stored under its released spelling: the renamed ORF YAL035C-A, the
    pseudogene YCL074W, the gene R64 does not list (YHR999W) and the sparse row
    YAL034C-B (1 of 5 columns). Only R0010W, not ORF-shaped, is held out.
    """
    assert len(dataset) == 31
    hip, hop = "heterozygous_deletion", "barcoded_kanmx_deletion"
    ami, boro, c777, c888 = "amitriptyline", "Boromycin", "CMB777", "Unreleased"

    full = [
        (ami, 50.0, 2, "0077"),
        (ami, 50.0, 1, "0091"),
        (boro, 1.2, 2, "0077"),
        (c777, 10.0, 2, "0077"),
        (c888, 10.0, 2, "0077"),
    ]

    def scored(orf: str, scores: list[float | None]) -> list[Any]:
        return [
            (hip, orf, orf, name, dose, score, n, study)
            for (name, dose, n, study), score in zip(full, scores, strict=True)
            if score is not None
        ]

    assert _summary(dataset) == [
        *scored("YAL001C", [0.1, -1.5, 0.9, 2.25, 0.3]),
        *scored("YAL035C-A", [0.2, None, None, 0.4, 0.3]),
        *scored("YCL074W", [0.5, 0.5, 0.5, None, 0.5]),
        *scored("YBR271W", [-3.0, -2.0, -1.0, -0.5, 0.3]),
        *scored("YHR999W", [0.7, 0.7, 0.7, 0.7, 0.7]),
        *scored("YLR074C", [0.6, 0.6, None, 0.6, None]),
        *scored("YAL034C-B", [0.8, None, None, None, None]),
        (hop, "YAL001C", "YAL001C", ami, 50.0, 1.0, 2, "0077"),
        (hop, "YAL001C", "YAL001C", boro, 1.2, 0.7, 2, "0077"),
        (hop, "YBR271W", "YBR271W", ami, 50.0, -0.25, 2, "0077"),
        (hop, "YAL034C-B", "YAL034C-B", ami, 50.0, 0.33, 2, "0077"),
        (hop, "YAL001C", "YAL001C", ami, 50.0, 0.99, 2, "0077"),
    ]


def test_hip_record_is_a_heterozygous_kanmx4_deletion_in_ypd_with_a_dmso_vehicle(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Record 0 field by field: YAL001C heterozygous kanMX4 deletion from YSC1055 with
    its Table S5 construction (Lab 4, chr1_1, plate 201, well E12) and a barcode gap;
    YPD liquid, 30 C, the 24-well HIP culture with every duration gapped; amitriptyline
    50 uM set at IC30 in 2 percent DMSO; score 0.1 over 2 technical replicates in study
    0077 with the three uncertainty gaps.
    """
    assert (
        dataset[0]["experiment"]
        == StrainEnvironmentResponseExperiment(
            dataset_name=_DATASET,
            genotype=_hip_genotype(
                "YAL001C",
                StrainConstruction(
                    lab="Lab 4", batch="chr1_1", plate="201", well="E12"
                ),
            ),
            environment=_environment("HIP", _treated(_AMITRIPTYLINE_COMPOUND, 50.0)),
            phenotype=_phenotype(0.1, 2, "0077"),
        ).model_dump()
    )
    assert dataset[0]["experiment"]["experiment_type"] == "strain_environment_response"
    assert dataset[0]["reference"] == _reference("HIP", "0077")
    assert (
        dataset[0]["publication"]
        == Publication(doi=_DOI, doi_url=f"https://doi.org/{_DOI}").model_dump()
    )


def test_renamed_row_is_stored_as_released_and_a_structure_only_compound_is_cmb_named(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Record 6: YAL035C-A is stored under YAL035C-A (no rename onto YAL034C-B, no
    ConstructedOrf), with its construction from Table S5's ``YAL035C-a`` entry; CMB 777
    has no released name, so its compound is named ``CMB777`` with the InChIKey derived
    from its SMILES.
    """
    assert (
        dataset[6]["experiment"]
        == StrainEnvironmentResponseExperiment(
            dataset_name=_DATASET,
            genotype=_hip_genotype(
                "YAL035C-A",
                StrainConstruction(
                    lab="Lab 4", batch="chr1_1", plate="262", well="E10"
                ),
            ),
            environment=_environment("HIP", _treated(_CMB777_COMPOUND, 10.0)),
            phenotype=_phenotype(0.4, 2, "0077"),
        ).model_dump()
    )


def test_boromycin_and_a_gene_r64_does_not_list_are_stored(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Record 19: YHR999W (no R64 gene, absent from Table S5, so the construction gap)
    under Boromycin, whose released SMILES RDKit cannot parse: the compound is stored
    under its released name with that SMILES verbatim and the typed inchikey gap.
    """
    assert (
        dataset[19]["experiment"]
        == StrainEnvironmentResponseExperiment(
            dataset_name=_DATASET,
            genotype=_hip_genotype("YHR999W", None),
            environment=_environment("HIP", _treated(_BOROMYCIN_COMPOUND, 1.2)),
            phenotype=_phenotype(0.7, 2, "0077"),
        ).model_dump()
    )


def test_hop_record_is_a_barcoded_kanmx4_deletion_over_a_16_hour_exposure(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Record 26: the HOP arm stores ``BarcodedKanMxDeletionPerturbation`` from YSC1056
    and an environment with both durations (16 h, 5 generations); its reference is the
    HOP control of study 0077.
    """
    assert (
        dataset[26]["experiment"]
        == StrainEnvironmentResponseExperiment(
            dataset_name=_DATASET,
            genotype=_hop_genotype("YAL001C"),
            environment=_environment("HOP", _treated(_AMITRIPTYLINE_COMPOUND, 50.0)),
            phenotype=_phenotype(1.0, 2, "0077"),
        ).model_dump()
    )
    assert dataset[26]["reference"] == _reference("HOP", "0077")


def test_reference_index_has_one_entry_per_assay_and_study(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    index = json.loads(
        (
            Path(dataset.root) / "preprocess" / "experiment_reference_index.json"
        ).read_text()
    )
    hip_0091 = [1, 9, 13, 18, 23]
    assert [entry["member_indices"] for entry in index] == [
        [i for i in range(26) if i not in hip_0091],
        hip_0091,
        [26, 27, 28, 29, 30],
    ]
    assert [entry["reference"] for entry in index] == [
        _reference("HIP", "0077"),
        _reference("HIP", "0091"),
        _reference("HOP", "0077"),
    ]


def test_ledger_accounts_for_every_scored_cell(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """``<root>/dropped_records.json``: HIP 27 scored cells = 26 stored + 1 held out
    (R0010W, not ORF-shaped); HOP 5 = 5 stored. Boromycin (released SMILES, RDKit parse
    failure as the unresolved reason) and CMB 888 (no SMILES) are listed in
    ``compound_identity_gaps`` with their stored columns and records per arm; the
    CMB4019 exclusion has no column here; amitriptyline's 50 uM, Boromycin's 1.2 uM and
    CMB 888's 10 uM equal their Table S1 IC30 (in range).
    """
    report = json.loads((Path(dataset.root) / "dropped_records.json").read_text())
    report.pop("created_at")
    gaps = report.pop("compound_identity_gaps")
    assert report == {
        "dataset": _DATASET,
        "not_orf_shaped_rule": (
            "systematic_name_is_not_orf_shaped_needs_released_name_leaf"
        ),
        "not_orf_shaped_policy": m.NOT_ORF_SHAPED_POLICY,
        "exclusion_rule": m.EXCLUSION_RULE,
        "n_kept": 31,
        "n_dropped": 1,
        "kept_by_assay": {"HIP": 26, "HOP": 5},
        "dropped_by_assay": {"HIP": 1, "HOP": 0},
        "accounting": {
            "HIP": {
                "n_sensitivity_columns": 5,
                "n_rows": 8,
                "n_rows_without_any_score": 0,
                "n_source_cells": 27,
                "n_stored": 26,
                "n_not_orf_shaped": 1,
                "n_excluded_condition": 0,
            },
            "HOP": {
                "n_sensitivity_columns": 2,
                "n_rows": 4,
                "n_rows_without_any_score": 0,
                "n_source_cells": 5,
                "n_stored": 5,
                "n_not_orf_shaped": 0,
                "n_excluded_condition": 0,
            },
        },
        "not_orf_shaped": {
            "HIP": [{"source_name": "R0010W", "n_records": 1}],
            "HOP": [],
        },
        "excluded_conditions": [
            {
                "cmb_id": "4019",
                "table_s1_name": "D-Glucose (starvation)",
                "n_columns_by_assay": {"HIP": 0, "HOP": 0},
                "n_records_by_assay": {"HIP": 0, "HOP": 0},
            }
        ],
        "n_cmb_ids_stored": 4,
        "n_compounds_stored": 4,
        "columns_by_assay": {"HIP": 5, "HOP": 2},
        "positive_control": {"HIP": 0, "HOP": 0},
        "solvent_gap_columns": {"HIP": 0, "HOP": 0},
        "ph_gap_columns": {"HIP": 0, "HOP": 0},
        "dose_outside_ic30_range": [],
        "n_kept_columns_with_ic30": {"HIP": 4, "HOP": 2},
        "hip_construction": {"one_entry": 5, "not_listed": 2},
    }
    assert [
        (
            g["cmb_id"],
            g["source_name"],
            g["smiles"],
            g["n_columns_by_assay"],
            g["n_records_by_assay"],
        )
        for g in gaps
    ] == [
        ("409", "Boromycin", _BOROMYCIN, {"HIP": 1, "HOP": 1}, {"HIP": 4, "HOP": 1}),
        ("888", "Unreleased", None, {"HIP": 1, "HOP": 0}, {"HIP": 5, "HOP": 0}),
    ]
    assert gaps[0]["resolution_status"] == "UNRESOLVED_PUBLIC"
    assert gaps[0]["unresolved_reason"] == (
        "the released Table S1 SMILES does not parse in RDKit 2026.03.6 (boron "
        "chemistry), so no InChIKey can be derived from the primary structure; PubChem "
        "carries a Boromycin record (CID 76962270) but adopting it would substitute a "
        "different structure for the one the screen released"
    )


def test_table_s5_flag_file_counts_kept_hip_records_per_listed_strain(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Two flagged Table S5 strains appear in the HIP matrix: YBR271W (CL1, positional,
    MUT, 5 stored records) and YLR074C (CL1+, not positional, MUT, 3); the HOP rows of
    YBR271W are not flagged. The file pins the regenerated 036 CSV and Table S5 by
    sha256 and carries the paper's cluster paragraph verbatim.
    """
    flags = json.loads(
        (Path(dataset.root) / "table_s5_affected_strains.json").read_text()
    )
    flags.pop("created_at")
    assert flags == {
        "dataset": _DATASET,
        "citation_key": _CITATION_KEY,
        "source_csv": (
            "experiments/036-dataset-fixes-before-kg-build/results/"
            "hoepfner2014_table_s5_strains.csv"
        ),
        "source_csv_sha256": m.TABLE_S5_STRAINS_CSV_SHA256,
        "source_script": m.TABLE_S5_STRAINS_SCRIPT,
        "table_s5_path": "si/Table_S5.xls",
        "table_s5_sha256": (
            "b123dc3e87fc10d3b4256f449fcd2eb38c91d1779000278af5a1a788356624a2"
        ),
        "paper_quote": str(m.SOURCED_VALUES["table_s5_mutations"].quote),
        "policy": m.TABLE_S5_POLICY,
        "n_strains": 2,
        "n_positional_strains": 1,
        "n_records_flagged": 8,
        "n_positional_records_flagged": 5,
        "strains": [
            {
                "systematic_gene_name": "YBR271W",
                "cluster": "CL1",
                "is_positional": True,
                "mutation": "Chromosome XI Aneuploidy",
                "validation_result": "MUT",
                "construction_lab": "Lab 14",
                "construction_batch": "chr00_1",
                "n_table_s5_entries": 1,
                "n_records": 5,
            },
            {
                "systematic_gene_name": "YLR074C",
                "cluster": "CL1+",
                "is_positional": False,
                "mutation": "Chromosome XI Aneuploidy",
                "validation_result": "MUT",
                "construction_lab": "Lab 3",
                "construction_batch": "chr12_2",
                "n_table_s5_entries": 1,
                "n_records": 3,
            },
        ],
    }


def test_sourced_values_file_and_gene_set(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """``preprocess/sourced_values.json`` holds the sourced constants in key order;
    ``temperature_c`` is checked literally and the whole file equals the JSON dump of
    ``SOURCED_VALUES``. The gene set is every released ORF a stored record carries.
    """
    preprocess = Path(dataset.root) / "preprocess"
    payload = json.loads((preprocess / "sourced_values.json").read_text())
    assert list(payload) == sorted(m.SOURCED_VALUES)
    assert "hip_duration_generations" not in payload
    assert payload["temperature_c"] == {
        "value": 30.0,
        "provenance": {
            "source_uri": "paper.md",
            "citation_key": _CITATION_KEY,
            "sha256": "a9877549eff2fe1aaf8aa403d9fea1c381284de030326f4475e869c102af0aeb",
            "method": "MinerU OCR of the publisher PDF (mirrored artifact)",
            "page": None,
            "retrieved": None,
        },
        "quote": (
            "Plates were incubated for $1 6 \\mathrm { h }$ in a robotic shaking "
            "incubator at $3 0 ^ { \\circ } C / 5 5 0$ RPM allowing for ${ \\sim } 5$ "
            "doublings."
        ),
        "note": None,
    }
    assert payload == {
        key: value.model_dump(mode="json") for key, value in m.SOURCED_VALUES.items()
    }
    assert json.loads((preprocess / "gene_set.json").read_text()) == [
        "YAL001C",
        "YAL034C-B",
        "YAL035C-A",
        "YBR271W",
        "YCL074W",
        "YHR999W",
        "YLR074C",
    ]
    manifest = json.loads((preprocess / "build_manifest.json").read_text())
    assert manifest["dataset_name"] == "env_chemgen_hoepfner2014"
    assert manifest["loader_class"] == _DATASET
    assert {
        "StrainEnvironmentResponseExperiment",
        "HeterozygousDeletionPerturbation",
        "BarcodedKanMxDeletionPerturbation",
    } <= set(manifest["closure"])
    assert dataset.experiment_class is StrainEnvironmentResponseExperiment
    assert dataset.reference_class is StrainEnvironmentResponseExperimentReference
    assert dataset.raw_file_names == [
        "HIP_scores.txt",
        "HOP_scores.txt",
        "Table_S1.xls",
    ]
    with pytest.raises(NotImplementedError):
        dataset.create_experiment()
    frame = {"untouched": True}
    assert dataset.preprocess_raw(frame) is frame


def test_compound_meta_reads_names_and_structures_across_the_three_sheets(
    tmp_path: Path,
) -> None:
    """Names come from the two MoA sheets and SMILES from ``All Structures``; an id cell
    may be an int, a digit string or a comma list (``244, 1818`` gives both ids the same
    SMILES); an empty id cell contributes nothing (the ``Nameless`` row), a NaN name
    leaves CMB 1234 nameless, and a NaN SMILES leaves CMB 555 structureless.
    """
    _write_table_s1(tmp_path / "Table_S1.xls")
    assert m._load_compound_meta(str(tmp_path / "Table_S1.xls")) == {
        "3": {"common_name": "Amitriptyline", "smiles": _AMITRIPTYLINE},
        "409": {"common_name": "Boromycin", "smiles": _BOROMYCIN},
        "888": {"common_name": "Unreleased", "smiles": None},
        "1234": {"common_name": None, "smiles": None},
        "1235": {"common_name": "Novel compound", "smiles": None},
        "777": {"common_name": None, "smiles": _AMITRIPTYLINE},
        "244": {"common_name": None, "smiles": "CC(=O)O"},
        "1818": {"common_name": None, "smiles": "CC(=O)O"},
        "555": {"common_name": None, "smiles": None},
    }


def _digests(raw: Path) -> dict[str, dict[str, str]]:
    return {
        name: {**spec, "sha256": hashlib.sha256((raw / name).read_bytes()).hexdigest()}
        for name, spec in m._DRYAD_FILES.items()
    }


def test_download_links_the_mirror_only_when_its_bytes_match_the_pinned_sha256(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """With a mirror at ``$DATA_ROOT/torchcell-raw/<key>/``, a mirror file whose digest
    is not the pinned Dryad digest raises ``RawSha256MismatchError`` with both digests
    and links nothing; with the pins repointed at the fixture's digests each absent raw
    file is symlinked, a raw file already present (``Table_S1.xls``, a plain copy) is
    left as it is, and the build runs to its 31 records under the real build-time
    check.
    """
    monkeypatch.setattr(m, "verify_raw_files", verify_raw_files)
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    mirror = data_root / "torchcell-raw" / _CITATION_KEY
    _write_raw(mirror)
    digest = hashlib.sha256((mirror / "HIP_scores.txt").read_bytes()).hexdigest()
    with pytest.raises(RawSha256MismatchError) as err:
        m.EnvChemgenHoepfner2014Dataset(root=str(tmp_path / "a"))
    assert str(err.value) == (
        f"sha256 mismatch for {mirror / 'HIP_scores.txt'}: expected "
        f"{m._DRYAD_FILES['HIP_scores.txt']['sha256']}, observed {digest}"
    )
    assert list((tmp_path / "a" / "raw").iterdir()) == []
    monkeypatch.setattr(m, "_DRYAD_FILES", _digests(mirror))
    raw = tmp_path / "b" / "raw"
    raw.mkdir(parents=True)
    (raw / "Table_S1.xls").write_bytes((mirror / "Table_S1.xls").read_bytes())
    dataset = m.EnvChemgenHoepfner2014Dataset(root=str(tmp_path / "b"))
    for name in ("HIP_scores.txt", "HOP_scores.txt"):
        assert (raw / name).is_symlink() and os.readlink(raw / name) == str(
            mirror / name
        )
    assert not (raw / "Table_S1.xls").is_symlink()
    assert len(dataset) == 31


class _FakeResponse:
    """The two members of a streamed ``requests.Response`` that ``_fetch_from_dryad`` uses."""

    def __init__(self, body: bytes) -> None:
        self.body = body
        self.closed = False

    def iter_content(self, chunk_size: int) -> list[bytes]:
        return [
            self.body[i : i + chunk_size] for i in range(0, len(self.body), chunk_size)
        ]

    def close(self) -> None:
        self.closed = True


def test_download_fetches_a_file_the_mirror_lacks_and_verifies_the_pin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """When the mirror has no ``HOP_scores.txt``, ``download()`` runs the recorded Dryad
    retrieval: ``_dryad_get`` (stubbed here, no network) streams the body, which is
    written to a ``.partial`` sibling and hashed chunk by chunk; a body whose digest is
    not the pin raises ``RawSha256MismatchError`` naming the URL and both digests and
    leaves no HOP file in ``raw/``; the right body is renamed into place as a plain file
    beside the two symlinks.
    """
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    mirror = data_root / "torchcell-raw" / _CITATION_KEY
    _write_raw(mirror)
    pins = _digests(mirror)
    monkeypatch.setattr(m, "_DRYAD_FILES", pins)
    hop_bytes = (mirror / "HOP_scores.txt").read_bytes()
    (mirror / "HOP_scores.txt").unlink()
    responses: list[_FakeResponse] = []
    calls: list[str] = []

    def fake_dryad_get(session: Any, url: str) -> _FakeResponse:
        calls.append(url)
        responses.append(_FakeResponse(b"wrong body" if len(calls) == 1 else hop_bytes))
        return responses[-1]

    monkeypatch.setattr(m, "_dryad_get", fake_dryad_get)
    wrong = hashlib.sha256(b"wrong body").hexdigest()
    with pytest.raises(RawSha256MismatchError) as err:
        m.EnvChemgenHoepfner2014Dataset(root=str(tmp_path / "a"))
    assert str(err.value) == (
        "sha256 mismatch for https://datadryad.org/downloads/file_stream/4834609: "
        f"expected {pins['HOP_scores.txt']['sha256']}, observed {wrong}"
    )
    assert sorted(p.name for p in (tmp_path / "a" / "raw").iterdir()) == [
        "HIP_scores.txt"
    ]
    dataset = m.EnvChemgenHoepfner2014Dataset(root=str(tmp_path / "b"))
    assert calls == ["https://datadryad.org/downloads/file_stream/4834609"] * 2
    assert [r.closed for r in responses] == [True, True]
    hop = tmp_path / "b" / "raw" / "HOP_scores.txt"
    assert not hop.is_symlink() and hop.read_bytes() == hop_bytes
    assert (tmp_path / "b" / "raw" / "HIP_scores.txt").is_symlink()
    assert len(dataset) == 31


def test_deposit_raw_mirror_records_the_dryad_retrieval_per_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A source file whose digest is not the pin is refused with
    ``RawSha256MismatchError`` and both digests BEFORE anything is copied, so no mirror
    directory exists afterwards; with the pins repointed the manifest lists the three
    files with the ``_dryad_get`` retriever and the Dryad DOI as the SI source.
    """
    source = tmp_path / "source"
    _write_raw(source)
    data_root = tmp_path / "data_root"
    digest = hashlib.sha256((source / "HIP_scores.txt").read_bytes()).hexdigest()
    with pytest.raises(RawSha256MismatchError) as err:
        m.deposit_raw_mirror(
            source_dir=source, retrieved_at="2026-09-27", data_root=str(data_root)
        )
    assert str(err.value) == (
        f"sha256 mismatch for {source / 'HIP_scores.txt'}: expected "
        f"{m._DRYAD_FILES['HIP_scores.txt']['sha256']}, observed {digest}"
    )
    assert not data_root.exists()
    digests = _digests(source)
    monkeypatch.setattr(m, "_DRYAD_FILES", digests)
    root = m.deposit_raw_mirror(
        source_dir=source, retrieved_at="2026-09-27", data_root=str(data_root)
    )
    assert root == data_root / "torchcell-raw" / _CITATION_KEY
    manifest = m.load_manifest(str(data_root))
    assert manifest.citation_key == _CITATION_KEY
    assert manifest.doi == _DOI
    assert manifest.si_data_sources == ["https://doi.org/10.5061/dryad.v5m8v"]
    assert manifest.si_expected == ["HIP_scores.txt", "HOP_scores.txt", "Table_S1.xls"]
    assert [(f.path, f.role, f.sha256, f.source) for f in manifest.files] == [
        (name, "raw_data", spec["sha256"], spec["url"])
        for name, spec in digests.items()
    ]
    for record in manifest.files:
        assert record.retrieval is not None
        assert record.retrieval.method.value == "direct_url"
        assert (
            record.retrieval.retriever
            == "torchcell.datasets.scerevisiae.hoepfner2014._dryad_get"
        )
        assert record.retrieval.retrieved_at == "2026-09-27"
        assert record.bytes == (root / record.path).stat().st_size


def test_a_short_matrix_line_stops_at_its_last_cell(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Record 29 is the HOP row ``YAL034C-B`` whose line carries only the first cell
    (0.33); the Boromycin column lies beyond the line's end and adds no record and no
    scored cell (HOP stores 5 of 5 scored cells).
    """
    (perturbation,) = dataset[29]["experiment"]["genotype"]["perturbations"]
    assert perturbation == _hop_genotype("YAL034C-B").perturbations[0].model_dump()
    assert dataset[29]["experiment"]["phenotype"]["environment_response"] == 0.33
    report = json.loads((Path(dataset.root) / "dropped_records.json").read_text())
    assert report["accounting"]["HOP"]["n_source_cells"] == 5


def test_solve_anubis_returns_the_smallest_nonce_meeting_the_difficulty() -> None:
    """The Dryad frontend's proof of work, solved offline: ``difficulty`` counts leading
    zero hex nibbles of sha256(random_data + nonce). For difficulty 2 (one zero byte)
    and 3 (a zero byte plus a zero high nibble) the returned digest is that hash, meets
    the bound, and no smaller nonce does.
    """
    for difficulty, nibbles in ((2, "00"), (3, "000")):
        digest, nonce = m._solve_anubis("challenge-seed", difficulty)
        assert digest == hashlib.sha256(f"challenge-seed{nonce}".encode()).hexdigest()
        assert digest.startswith(nibbles)
        assert not any(
            hashlib.sha256(f"challenge-seed{n}".encode())
            .hexdigest()
            .startswith(nibbles)
            for n in range(nonce)
        )


def test_a_repeated_orf_row_is_stored_twice_under_one_genotype(
    dataset: m.EnvChemgenHoepfner2014Dataset,
) -> None:
    """Finding: ``_iter_records`` (hoepfner2014.py lines 1266-1269) caches the genotype
    per released ORF and never checks that an ORF row is unique, so the second HOP
    ``YAL001C`` line becomes record 30 with the same genotype and reference as record 26
    and a second (strain, condition) measurement of 0.99; nothing in the ledger reports
    the repeat. (The deposited matrices carry no repeated token: measured 2026.10.12.)
    """
    assert (
        dataset[30]["experiment"]["genotype"] == dataset[26]["experiment"]["genotype"]
    )
    assert dataset[30]["reference"] == dataset[26]["reference"]
    assert (
        dataset[30]["experiment"]["environment"]
        == dataset[26]["experiment"]["environment"]
    )
    assert dataset[30]["experiment"]["phenotype"]["environment_response"] == 0.99
    assert dataset[26]["experiment"]["phenotype"]["environment_response"] == 1.0


def test_a_raw_file_off_the_pin_is_refused_at_build_time(
    off_pin_raw: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Contract (issue #518's sweep): with the three Dryad files already in ``raw/`` PyG
    skips ``download()``, so ``process()`` verifies each against ``_DRYAD_FILES`` first.
    ``HIP_scores.txt`` off its pin raises ``RawSha256MismatchError`` naming it and both
    digests before a row is read; no store is written and the file is left as found.
    """
    monkeypatch.delenv("TC_DATA_URL", raising=False)
    staged = off_pin_raw(m, ["HIP_scores.txt", "HOP_scores.txt", "Table_S1.xls"])
    raw = staged.root / "raw" / "HIP_scores.txt"
    with pytest.raises(RawSha256MismatchError) as err:
        m.EnvChemgenHoepfner2014Dataset(root=str(staged.root))
    assert str(err.value) == (
        f"sha256 mismatch for {raw}: expected "
        "dbc5041defea9c046da0890d5e569f97d5f7afbf50ea0885f539ea8e5980cd24, "
        f"observed {staged.observed}"
    )
    assert list((staged.root / "processed").iterdir()) == []
    assert not (staged.root / "preprocess").exists()
    assert hashlib.sha256(raw.read_bytes()).hexdigest() == staged.observed


# Phase 24: a stored dose outside the IC30 band is listed in the ledger


def test_a_stored_dose_outside_the_ic30_band_is_listed_with_its_ratio(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``IC30_RATIO_RANGE`` is narrowed to (2.0, 4.0), so amitriptyline's 50 uM against
    its Table S1 IC30 of 50 uM (ratio 1.0, in range under the real (0.5, 2.0)) falls
    outside, and so do Boromycin's 1.2 uM and CMB 888's 10 uM (each ratio 1.0). Every
    stored column with an IC30 is listed, in assay then column order. CMB 777 has no
    IC30 row (it is only in the structures sheet), so it is never checked.
    """
    monkeypatch.setattr(m, "IC30_RATIO_RANGE", (2.0, 4.0))
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    root = tmp_path / "env_chemgen_hoepfner2014"
    _write_raw(root / "raw")
    m.EnvChemgenHoepfner2014Dataset(root=str(root))
    report = json.loads((root / "dropped_records.json").read_text())
    assert report["dose_outside_ic30_range"] == [
        {
            "assay": assay,
            "cmb_id": cmb,
            "compound": compound,
            "study": study,
            "concentration_um": dose,
            "ic30_um": dose,
            "ratio": 1.0,
        }
        for assay, cmb, compound, study, dose in (
            ("HIP", "3", "amitriptyline", "0077", 50.0),
            ("HIP", "3", "amitriptyline", "0091", 50.0),
            ("HIP", "409", "Boromycin", "0077", 1.2),
            ("HIP", "888", "Unreleased", "0077", 10.0),
            ("HOP", "3", "amitriptyline", "0077", 50.0),
            ("HOP", "409", "Boromycin", "0077", 1.2),
        )
    ]
