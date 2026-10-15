# torchcell/datasets/ecoli/potel2018
# [[torchcell.datasets.ecoli.potel2018]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/potel2018
# Test file: tests/torchcell/datasets/ecoli/test_potel2018.py
r"""Potel 2018 E. coli phosphoproteome: per-sample phosphosite intensities, as released.

Potel CM, Lin MH, Heck AJR, Lemeer S (2018) "Widespread bacterial protein histidine
phosphorylation revealed by mass spectrometry-based proteomics", Nature Methods
15:187-190, doi:10.1038/nmeth.4580, PMID 29377012; ProteomeXchange PXD008369 (jPOST
JPST000369). The first consumer of ``PhosphositeIntensityPhenotype`` (#921).

WHAT IS STORED. Supplementary Table 2 is the release's one per-record artifact: 2,129
phosphosites (S 1,220, T 501, Y 162, H 246) with a MaxQuant summed-XIC precursor
intensity in each of 12 samples, glucose or glycerol x exponential or stationary x three
replicates. One ``PhosphositeIntensityExperiment`` per sample column, 12
records. Every record carries all 2,129 sites exactly as the table gives them: the
accession string, the gene names, the residue and the positions, verbatim, including the
42 sites listed on more than one gene (10 of them on more than one accession), each as
ONE key. No accession, gene name or position is looked up, mapped or resolved against
any annotation: linking a site to a gene node is the adapter's and the model's concern,
not the record's. A released ``0`` (12,095 of 25,548 cells) is the typed status
``not_quantified`` and never a stored zero; the 245 sites with no intensity in any of
the twelve samples are stored on every record with that status.

WHAT IS NOT STORED, AND WHY.

- Table 2's total ``Intensity`` column. Its legend defines it over "the identified
  phosphopeptide sequence" with no sample, and it is not the sum of the twelve sample
  columns: equal (within 1e-6 relative) for 278 sites, larger for 1,130, smaller for
  477, and 0 for 244 (:class:`BuildAccounting`, re-counted at build time).
- Supplementary Table 5. It is derived from Table 2 (every joinable cell whose Table 2
  intensity is non-zero is its log2, and the rest are Perseus imputations), so it is
  the L4 cross-table check here, not a record source.
- The peptide-level columns (``Number of Phospho (HSTY)``, ``Phospho (HSTY)
  Probabilities``, ``Charge``, ``Mass error [ppm]``) and ``Protein names``: they
  describe the best peptide's spectrum, not the site's quantity.

THE SAMPLE. Each record's ``sample_label`` is its released column header, and the
column's legend line sources its carbon source, growth phase and replicate number
(:data:`SAMPLES`). The harvest optical density is the dissertation's ("mid-log phase
(OD600= 0.6/0.5 for glucose / glycerol-fed cells) and stationary phase (OD600=
1.2/1.0 ...)"). The replicate UNIT is a typed gap: two SI legends call the twelve
samples "n=12 biological replicates", while the reporting summary ("we made technical
triplicates (after digestion, before enrichment) of each" condition) and the
dissertation ("Fe3+-IMAC was performed in technical triplicates for each biological
groups") call each condition's three technical. The published Online Methods would
decide it and are walled (#922), so the gap is ``deferred_pending_source_review``.

THE SOURCE OF THE METHODS. The paper is in neither mirror and in no PMC set (#922). The
first author's open-access dissertation (Utrecht handle 1874/370166), whose Chapter V
reprints this paper with its Materials and Methods, is deposited in the raw mirror with
a pinned ``pdftotext -layout`` rendering, and every Methods quote here is read from
Chapter V of that rendering. ``pdftotext`` sets the subscripted formula of the first
M9 salt out of line in Chapter V, so its 6 g/L is corroborated by the same sentence in
Chapter VI, which renders clean (:data:`SALT_NA2HPO4_CLEAN`).

THE HOST. "E. coli strain W3110" (Chapter V), so every reference pins the W3110
assembly set ``ecoli_K12_W3110_ASM1024v1``.

REFERENCE (a decision, stated for review). The release names no denominator. Each
record's reference is the exponential-phase sample of its own carbon source with the
same replicate number: the one released contrast (Supplementary Table 5) is exponential
against stationary within a carbon source, and the replicate number only picks one real
released sample deterministically. The exponential and stationary samples were injected
in different batches (jPOST ``OR13_A2_20161014`` and ``OR13_20160922``), so the numbers
do not pair the cultures.

Release inventory: ``experiments/036-dataset-fixes-before-kg-build/scripts/
potel2018_release_inventory.py``. Note and L0-L4 table:
[[torchcell.datasets.ecoli.potel2018]].
"""

from __future__ import annotations

import json
import logging
import math
import os
import os.path as osp
from collections import Counter
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any, ClassVar, Literal

import openpyxl
from pydantic import BaseModel, ConfigDict
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    Organism,
    check_manifest_pin,
    file_sha256,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.schema import (
    BacterialAssemblySet,
    Concentration,
    ConcentrationUnit,
    CultureGrowthPhase,
    Environment,
    Experiment,
    ExperimentReference,
    Genotype,
    Media,
    MediaComponent,
    MediaComponentRole,
    PhosphoResidue,
    Phosphosite,
    PhosphositeIntensityExperiment,
    PhosphositeIntensityExperimentReference,
    PhosphositeIntensityPhenotype,
    PhosphositeQuantification,
    ProteinIdNamespace,
    Publication,
    ReferenceGenome,
    Temperature,
)
from torchcell.datasets.bacteria_common import assembly_reference
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import Manifest
from torchcell.verification.levels import l3_convention
from torchcell.verification.phosphosite import verify_phosphosite_dataset
from torchcell.verification.report import (
    Level,
    LevelResult,
    Provenance,
    VerificationReport,
)
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
    audit_sourced_value,
)

log = logging.getLogger(__name__)

__all__ = [
    "CITATION_KEY",
    "SAMPLES",
    "PhosphositeIntensityPotel2018Dataset",
    "load_manifest",
    "raw_mirror_dir",
    "read_table_2",
    "verify_build",
]

# --------------------------------------------------------------------------- #
# Paper, mirror and raw-file pins
# --------------------------------------------------------------------------- #
CITATION_KEY = "potelWidespreadBacterialProtein2018"
DOI = "10.1038/nmeth.4580"
PUBMED_ID = "29377012"
PROTEOMEXCHANGE = "PXD008369"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
RAW_RETRIEVED_AT = "2026-10-10"

#: Supplementary Table 2, the one workbook every record is built from.
TABLE_2_FILENAME = "41592_2018_BFnmeth4580_MOESM4_ESM.xlsx"
TABLE_2_REL = f"data/{TABLE_2_FILENAME}"
TABLE_2_SHA256 = "f69ca61a0f0dba3987b26bcdf4bd9b0c8fcf8744495233bad678a23cf5682481"
TABLE_2_SHEET = "Supplementary Table 2."
#: Supplementary Table 5, derived from Table 2: the L4 cross-table check reads it.
TABLE_5_REL = "si/41592_2018_BFnmeth4580_MOESM6_ESM.xlsx"
TABLE_5_SHA256 = "31a986e2b9a1cc582b1265a0bc124693f4bcdbf941abb3398f27217d3f1d58bb"
TABLE_5_SHEET = "Supplementary Table 5."
#: Text renderings in the raw mirror, each pinned by the release inventory.
TABLE_2_LEGEND_REL = "data/41592_2018_BFnmeth4580_MOESM4_ESM.legends.txt"
TABLE_2_LEGEND_SHA256 = (
    "a85a3a155971d28456c6b3661b96bc662dfc46310c1c0d487283cb4e8dcfb57d"
)
THESIS_TEXT_REL = "paper/thesis/Potel.pdftotext.txt"
THESIS_TEXT_SHA256 = "05a506579dcbc869d056af2bc81ecabf4f82debb4a84dda73b7b35075fdbd81f"
SI_TEXT_REL = "si/41592_2018_BFnmeth4580_MOESM1_ESM.pdftotext.txt"
SI_TEXT_SHA256 = "4dee5c2a098fdaf95cfbe899c773348fe01281d535f9147708ffc9be79bd40de"
REPORTING_SUMMARY_TEXT_REL = "si/41592_2018_BFnmeth4580_MOESM2_ESM.pdftotext.txt"
REPORTING_SUMMARY_TEXT_SHA256 = (
    "9951a1ac02a5090d343134b97440acf618757c369e12310cc3b4ba23ee37ee43"
)

W3110_STRAIN: Literal["W3110"] = "W3110"
W3110_ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_W3110_ASM1024v1"

#: Measured on the pinned workbook by the release inventory; asserted at build time.
SOURCE_SITES = 2129
SOURCE_RESIDUES: dict[str, int] = {"S": 1220, "T": 501, "Y": 162, "H": 246}
MULTI_GENE_SITES = 42
MULTI_ACCESSION_SITES = 10
ZERO_CELLS = 12095
SITES_ZERO_IN_ALL_SAMPLES = 245
EXPECTED_RECORDS = 12

MEASUREMENT_TYPE = "fe3_imac_lc_msms_maxquant_summed_xic"

#: The Table 2 header, verbatim and in order; a release that differs is refused.
TABLE_2_HEADER: tuple[str, ...] = (
    "Acc. No.",
    "Protein names",
    "Gene names",
    "Positions",
    "Localization prob",
    "PEP",
    "Score",
    "Delta score",
    "Diagnostic peak",
    "Number of Phospho (HSTY)",
    "Amino acid",
    "Sequence window",
    "Phospho (HSTY) Probabilities",
    "Charge",
    "Mass error [ppm]",
    "Intensity",
    "Intensity glucose_exp1",
    "Intensity glucose_exp2",
    "Intensity glucose_exp3",
    "Intensity glucose_stat1",
    "Intensity glucose_stat2",
    "Intensity glucose_stat3",
    "Intensity glycerol_exp1",
    "Intensity glycerol_exp2",
    "Intensity glycerol_exp3",
    "Intensity glycerol_stat1",
    "Intensity glycerol_stat2",
    "Intensity glycerol_stat3",
)
LIST_SEPARATOR = ";"


# --------------------------------------------------------------------------- #
# Sourced values: each quote is a verbatim substring of the pinned mirror bytes
# --------------------------------------------------------------------------- #
def _provenance(uri: str, sha256: str, method: str, page: str) -> Provenance:
    return Provenance(
        source_uri=uri,
        citation_key=CITATION_KEY,
        sha256=sha256,
        method=method,
        page=page,
    )


def _thesis(
    value: Any, quote: str, *, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to the pinned rendering of the dissertation (Chapter V is the paper)."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=_provenance(
            THESIS_TEXT_REL,
            THESIS_TEXT_SHA256,
            "pdftotext -layout of the first author's dissertation, whose Chapter V "
            "reprints the paper with its Methods (raw mirror)",
            page,
        ),
    )


def _legend(value: Any, quote: str, *, item: str) -> SourcedValue:
    """Bind a value to one line of Table 2's Description sheet (raw-mirror rendering)."""
    return SourcedValue(
        value=value,
        quote=quote,
        provenance=_provenance(
            TABLE_2_LEGEND_REL,
            TABLE_2_LEGEND_SHA256,
            "text rendering of the Description sheet of Supplementary Table 2 "
            "(raw mirror)",
            f"Supplementary Table 2, Description, column {item}",
        ),
    )


def _si(value: Any, quote: str, *, page: str) -> SourcedValue:
    """Bind a value to the pinned rendering of the Supplementary Text and Figures."""
    return SourcedValue(
        value=value,
        quote=quote,
        provenance=_provenance(
            SI_TEXT_REL,
            SI_TEXT_SHA256,
            "pdftotext -layout of the Supplementary Text and Figures (raw mirror)",
            page,
        ),
    )


def _reporting_summary(value: Any, quote: str, *, page: str) -> SourcedValue:
    """Bind a value to the pinned rendering of the Life Sciences Reporting Summary."""
    return SourcedValue(
        value=value,
        quote=quote,
        provenance=_provenance(
            REPORTING_SUMMARY_TEXT_REL,
            REPORTING_SUMMARY_TEXT_SHA256,
            "pdftotext -layout of the Life Sciences Reporting Summary (raw mirror)",
            page,
        ),
    )


_CH5_GROWTH = "Chapter V, Materials and Methods, 'E. coli growth conditions'"
_CH6_GROWTH = "Chapter VI, Materials and Methods, 'E. coli growth conditions'"
_CH5_ANALYSIS = "Chapter V, Materials and Methods, data analysis"

STRAIN = _thesis(
    W3110_STRAIN,
    "E. coli strain W3110 was grown in M9 minimal medium, consisting of M9 salts (6    2",
    page=_CH5_GROWTH,
    note="the line ends where pdftotext set the subscripted Na2HPO4 out of line",
)
SALT_NA2HPO4_CLEAN = _thesis(
    6.0,
    "M9 salts (6 g/L\nNa2HPO4, 3 g/L KH2PO4, 0.5 g/L NaCl, 1 g/L NH4Cl) supplemented "
    "with either additional",
    page=_CH6_GROWTH,
    note="Chapter VI repeats Chapter V's growth sentence and renders the formula "
    "clean; Chapter V's own rendering of the first salt is '(6    2 / HPO / g/L4Na'",
)
SALTS = _thesis(
    {"KH2PO4": 3.0, "NaCl": 0.5, "NH4Cl": 1.0},
    "3 g/L KH2PO4, 0.5 g/L NaCl, 1 g/L 4NH Cl) supplemented with either additional "
    "0.5% (w/v)",
    page=_CH5_GROWTH,
    note="'4NH Cl' is pdftotext's rendering of NH4Cl, kept verbatim",
)
CARBON_AND_MAGNESIUM = _thesis(
    {"glucose_percent_w_v": 0.5, "glycerol_percent": 0.5, "MgSO4_mM": 1.0},
    "glucose or 0.5% glycerol, 1 mM MgSO",
    page=_CH5_GROWTH,
    note="the sentence reads 'either additional 0.5% (w/v) glucose or 0.5% glycerol', "
    "so the w/v basis is stated for glucose only and glycerol is stored as a percent "
    "with the basis unstated (ConcentrationUnit.percent)",
)
CALCIUM = _thesis(0.1, ", 0.1 mM CaCl2", page=_CH5_GROWTH)
TEMPERATURE_C = _thesis(
    37.0,
    "with vigorous shaking at 37°C. Cells",
    page=_CH5_GROWTH,
    note="vigorous shaking is why every record is aerobic",
)
HARVEST_EXPONENTIAL = _thesis(
    {"glucose": 0.6, "glycerol": 0.5},
    "were collected by centrifugation at mid-log phase (OD600= 0.6/0.5 for glucose / "
    "glycerol-fed",
    page=_CH5_GROWTH,
)
HARVEST_STATIONARY = _thesis(
    {"glucose": 1.2, "glycerol": 1.0},
    "cells) and stationary phase (OD600= 1.2/1.0 for glucose/ glycerol-fed cells) "
    "(Figure S12) and",
    page=_CH5_GROWTH,
)
TECHNICAL_TRIPLICATES_THESIS = _thesis(
    "technical",
    "Fe3+-IMAC was performed in technical triplicates for each biological groups",
    page=_CH5_ANALYSIS,
)
SEARCH_SOFTWARE = _thesis(
    "MaxQuant 1.5.5.1",
    "Raw data files were processed using the MaxQuant software (ver. 1.5.5.1)",
    page=_CH5_ANALYSIS,
)
SEARCH_DATABASE = _thesis(
    "UniProtKB accessions",
    "E. coli K12 database (Uniprot, March 2016)",
    page=_CH5_ANALYSIS,
)
BIOLOGICAL_REPLICATES_SI = _si(
    "biological",
    "The same UV profiles were observed for the n=12 biological replicates.",
    page="Supplementary Figure 3 legend",
)
TECHNICAL_TRIPLICATES_RS = _reporting_summary(
    "technical",
    "we made technical triplicates (after digestion, before enrichment) of each",
    page="1. Sample size",
)
INTENSITY_COLUMN = _legend(
    MEASUREMENT_TYPE,
    "17. Intensity glucose_exp1, summed up eXtracted Ion Current (XIC) of all isotopic "
    "clusters associated with the identified phosphopeptdie sequence in the 1st "
    "replicate of glucose-fed exponential E. coli cells.",
    item="17",
)
TOTAL_INTENSITY_COLUMN = _legend(
    "not stored",
    "16. Intensity, summed up eXtracted Ion Current (XIC) of all isotopic clusters "
    "associated with the identified  phosphopeptide sequence.",
    item="16",
)
LOCALIZATION_COLUMN = _legend(
    "localization_probability",
    "5. Localization prob, the probability of phosphosite assignment.",
    item="5",
)
POSITIONS_COLUMN = _legend(
    "positions",
    "4. Positions, the position of identified phosphosite within the protein "
    "associated with.",
    item="4",
)
DIAGNOSTIC_PEAK_COLUMN = _legend(
    "diagnostic_peak",
    "9. Diagnostic peak, whether the immonium present in the spectra involving this "
    "phosphosite identification.",
    item="9",
)


class SampleSpec(BaseModel):
    """One released sample column: its carbon source, phase, replicate and legend."""

    model_config = ConfigDict(frozen=True)

    column: str
    item: str
    carbon: Literal["glucose", "glycerol"]
    phase: CultureGrowthPhase
    replicate: int
    legend_quote: str

    @property
    def legend(self) -> SourcedValue:
        """The column's legend line, which sources its carbon, phase and replicate."""
        return _legend(
            {
                "carbon": self.carbon,
                "phase": str(self.phase),
                "replicate": self.replicate,
            },
            self.legend_quote,
            item=self.item,
        )

    @property
    def harvest_optical_density(self) -> float:
        """The dissertation's harvest OD600 of this carbon source and phase."""
        source = (
            HARVEST_EXPONENTIAL
            if self.phase is CultureGrowthPhase.exponential
            else HARVEST_STATIONARY
        )
        return float(source.value[self.carbon])


_ORDINAL = {1: "1st", 2: "2nd", 3: "3rd"}
_PHASE_WORD = {
    CultureGrowthPhase.exponential: ("exp", "exponential"),
    CultureGrowthPhase.stationary: ("stat", "stationary"),
}


def _sample(
    item: int,
    carbon: Literal["glucose", "glycerol"],
    phase: CultureGrowthPhase,
    rep: int,
) -> SampleSpec:
    short, word = _PHASE_WORD[phase]
    column = f"Intensity {carbon}_{short}{rep}"
    quote = (
        f"{item}. {column}, summed up eXtracted Ion Current (XIC) of all isotopic "
        "clusters associated with the identified phosphopeptdie sequence in the "
        f"{_ORDINAL[rep]} replicate of {carbon}-fed {word} E. coli cells."
    )
    return SampleSpec(
        column=column,
        item=str(item),
        carbon=carbon,
        phase=phase,
        replicate=rep,
        legend_quote=quote,
    )


Carbon = Literal["glucose", "glycerol"]
_CARBONS: tuple[Carbon, ...] = ("glucose", "glycerol")
#: The twelve sample columns in release order (Table 2 legend items 17 to 28).
SAMPLES: tuple[SampleSpec, ...] = tuple(
    _sample(17 + index, carbon, phase, rep)
    for index, (carbon, phase, rep) in enumerate(
        (carbon, phase, rep)
        for carbon in _CARBONS
        for phase in (CultureGrowthPhase.exponential, CultureGrowthPhase.stationary)
        for rep in (1, 2, 3)
    )
)
#: Sites each sample quantifies (a non-zero cell), measured on the pinned workbook by
#: the release inventory and asserted at build time.
QUANTIFIED_PER_SAMPLE: dict[str, int] = {
    "Intensity glucose_exp1": 989,
    "Intensity glucose_exp2": 1002,
    "Intensity glucose_exp3": 897,
    "Intensity glucose_stat1": 1200,
    "Intensity glucose_stat2": 1241,
    "Intensity glucose_stat3": 1136,
    "Intensity glycerol_exp1": 1064,
    "Intensity glycerol_exp2": 1087,
    "Intensity glycerol_exp3": 1054,
    "Intensity glycerol_stat1": 1263,
    "Intensity glycerol_stat2": 1201,
    "Intensity glycerol_stat3": 1319,
}

_LOOKED_IN = _provenance(
    SI_TEXT_REL,
    SI_TEXT_SHA256,
    "pdftotext -layout of the Supplementary Text and Figures (raw mirror)",
    "Supplementary Figures 3 and 6 legends; si/...MOESM2 reporting summary "
    "'1. Sample size'; paper/thesis Chapter V Materials and Methods",
)
_PUBLISHED_METHODS = Provenance(
    source_uri=f"https://www.nature.com/articles/{DOI.split('/')[1]}",
    citation_key=CITATION_KEY,
    method="the published Online Methods, behind nature.com's subscription (#922)",
    page="Online Methods",
)
#: The release states the replicate unit both ways; the walled Methods would decide.
REPLICATE_UNIT_GAP = ProvenanceGap(
    field="replicate_unit",
    reason=ProvenanceGapReason.deferred_pending_source_review,
    looked_in=_LOOKED_IN,
    resolve_with=_PUBLISHED_METHODS,
    note=(
        "two SI legends call the twelve samples 'n=12 biological replicates' while "
        "the reporting summary says 'we made technical triplicates (after digestion, "
        "before enrichment) of each' condition and the dissertation's Chapter V says "
        "'Fe3+-IMAC was performed in technical triplicates for each biological groups'"
    ),
)


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the environment (the mirror and build tree live under it)."""
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/potelWidespreadBacterialProtein2018``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json`` (written by the release inventory)."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


# --------------------------------------------------------------------------- #
# Reading Table 2
# --------------------------------------------------------------------------- #
class SiteRow(BaseModel):
    """One Table 2 row: the site as released and its twelve sample cells."""

    model_config = ConfigDict(frozen=True)

    site: Phosphosite
    total_intensity: int
    intensities: dict[str, int]


def _split(cell: object) -> list[str]:
    """A released ``;`` list, in its order (a single value is a one-element list)."""
    return str(cell).split(LIST_SEPARATOR)


def _diagnostic_peak(cell: object) -> bool:
    """``+`` is a diagnostic immonium ion seen; an empty cell is none seen."""
    if cell is None:
        return False
    if cell == "+":
        return True
    raise ValueError(f"Diagnostic peak cell {cell!r} is neither '+' nor empty")


def _count(cell: object, column: str) -> int:
    """An intensity cell: a non-negative integer, as every one of them is released."""
    if not isinstance(cell, int) or isinstance(cell, bool) or cell < 0:
        raise ValueError(f"{column} cell {cell!r} is not a non-negative integer")
    return cell


def site_from_row(row: Mapping[str, object]) -> Phosphosite:
    """The site of one row, every identifier as the table gives it."""
    return Phosphosite(
        protein_ids=_split(row["Acc. No."]),
        protein_id_namespace=ProteinIdNamespace.uniprotkb_accession,
        gene_names=_split(row["Gene names"]),
        positions=[int(p) for p in _split(row["Positions"])],
        residue=PhosphoResidue(str(row["Amino acid"])),
        sequence_windows=_split(row["Sequence window"]),
        localization_probability=float(row["Localization prob"]),  # type: ignore[arg-type]  # openpyxl cell: int or float, checked by the model
        posterior_error_probability=float(row["PEP"]),  # type: ignore[arg-type]  # openpyxl cell: int or float, checked by the model
        identification_score=float(row["Score"]),  # type: ignore[arg-type]  # openpyxl cell: int or float, checked by the model
        delta_score=float(row["Delta score"]),  # type: ignore[arg-type]  # openpyxl cell: int or float, checked by the model
        diagnostic_peak=_diagnostic_peak(row["Diagnostic peak"]),
    )


def read_table_2(path: str | Path) -> list[SiteRow]:
    """Every Table 2 row as a site plus its sample cells; the header must be released."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[TABLE_2_SHEET].iter_rows(values_only=True))
    book.close()
    header = tuple(str(cell) for cell in rows[0])
    if header != TABLE_2_HEADER:
        raise RuntimeError(f"Table 2 header {header} != {TABLE_2_HEADER}")
    out: list[SiteRow] = []
    for values in rows[1:]:
        row = dict(zip(header, values, strict=True))
        out.append(
            SiteRow(
                site=site_from_row(row),
                total_intensity=_count(row["Intensity"], "Intensity"),
                intensities={
                    spec.column: _count(row[spec.column], spec.column)
                    for spec in SAMPLES
                },
            )
        )
    return out


# --------------------------------------------------------------------------- #
# Environment, reference and the phenotype
# --------------------------------------------------------------------------- #
def _component(
    name: str,
    role: MediaComponentRole,
    value: float,
    unit: ConcentrationUnit,
    *provenance: SourcedValue,
) -> MediaComponent:
    return MediaComponent(
        compound=resolved_compound(name),
        role=role,
        concentration=Concentration(value=value, unit=unit),
        provenance=list(provenance),
    )


def medium(carbon: Literal["glucose", "glycerol"]) -> Media:
    """M9 salts with Mg and Ca and the record's carbon source, as Chapter V states."""
    salt, g_l, mm = (
        MediaComponentRole.bulk_salt,
        ConcentrationUnit.g_per_l,
        (ConcentrationUnit.millimolar),
    )
    carbon_component = (
        _component(
            "D-glucose",
            MediaComponentRole.carbon_source,
            0.5,
            ConcentrationUnit.percent_w_v,
            CARBON_AND_MAGNESIUM,
        )
        if carbon == "glucose"
        else _component(
            "glycerol",
            MediaComponentRole.carbon_source,
            0.5,
            ConcentrationUnit.percent,
            CARBON_AND_MAGNESIUM,
        )
    )
    basis = "0.5% (w/v) glucose" if carbon == "glucose" else "0.5% glycerol"
    return Media(
        name=f"M9 minimal medium, {basis}, 1 mM MgSO4, 0.1 mM CaCl2 (Potel 2018)",
        state="liquid",
        is_synthetic=True,
        base_medium="M9",
        components=[
            _component(
                "disodium hydrogen phosphate",
                salt,
                6.0,
                g_l,
                STRAIN,
                SALT_NA2HPO4_CLEAN,
            ),
            _component("potassium dihydrogen phosphate", salt, 3.0, g_l, SALTS),
            _component("sodium chloride", salt, 0.5, g_l, SALTS),
            _component(
                "ammonium chloride", MediaComponentRole.nitrogen_source, 1.0, g_l, SALTS
            ),
            _component("magnesium sulfate", salt, 1.0, mm, CARBON_AND_MAGNESIUM),
            _component("calcium chloride", salt, 0.1, mm, CALCIUM),
            carbon_component,
        ],
        provenance=[STRAIN, SALTS, CARBON_AND_MAGNESIUM, CALCIUM],
    )


def environment(carbon: Literal["glucose", "glycerol"]) -> Environment:
    """Aerobic, vigorously shaken 37 C batch culture in the carbon source's M9."""
    return Environment(
        media=medium(carbon),
        temperature=Temperature(value=float(TEMPERATURE_C.value)),
        aerobicity="aerobic",
    )


def strain_reference(data_root: str | None = None) -> ReferenceGenome:
    """The assembly-pinned W3110 reference every record is written against."""
    return assembly_reference(W3110_STRAIN, data_root=data_root)


def phenotype(
    spec: SampleSpec, rows: Sequence[SiteRow]
) -> PhosphositeIntensityPhenotype:
    """One sample's record: every site, its status, and the released intensities."""
    sites = {row.site.key(): row.site for row in rows}
    if len(sites) != len(rows):
        raise RuntimeError(f"{len(rows) - len(sites)} Table 2 rows share a site key")
    quantification = {
        row.site.key(): (
            PhosphositeQuantification.quantified
            if row.intensities[spec.column] > 0
            else PhosphositeQuantification.not_quantified
        )
        for row in rows
    }
    return PhosphositeIntensityPhenotype(
        phosphosite_intensity={
            row.site.key(): float(row.intensities[spec.column])
            for row in rows
            if row.intensities[spec.column] > 0
        },
        quantification=quantification,
        sites=sites,
        measurement_type=MEASUREMENT_TYPE,
        sample_label=spec.column,
        replicate=spec.replicate,
        harvest_growth_phase=spec.phase,
        harvest_optical_density=spec.harvest_optical_density,
        provenance_gaps=[REPLICATE_UNIT_GAP],
    )


def reference_sample(spec: SampleSpec) -> SampleSpec:
    """The exponential sample of the record's carbon source with its replicate number."""
    return next(
        s
        for s in SAMPLES
        if s.carbon == spec.carbon
        and s.phase is CultureGrowthPhase.exponential
        and s.replicate == spec.replicate
    )


def publication() -> Publication:
    """This paper, by DOI and PubMed id."""
    return Publication(
        pubmed_id=PUBMED_ID,
        pubmed_url=f"https://pubmed.ncbi.nlm.nih.gov/{PUBMED_ID}/",
        doi=DOI,
        doi_url=f"https://doi.org/{DOI}",
    )


class TotalIntensityComparison(BaseModel):
    """Table 2's total ``Intensity`` against the sum of the twelve sample columns."""

    equal_within_1e_6_relative: int
    total_greater: int
    total_smaller: int
    total_zero: int


class BuildAccounting(BaseModel):
    """What the build read, what it stored, and the counts that tie the two."""

    dataset: str
    source_sites: int
    residues: dict[str, int]
    multi_gene_sites: int
    multi_accession_sites: int
    diagnostic_peak_by_residue: dict[str, int]
    cells: int
    zero_cells: int
    sites_zero_in_all_samples: int
    quantified_per_sample: dict[str, int]
    records: int
    total_intensity_not_stored: TotalIntensityComparison

    def check(self) -> None:
        """Every count equals the release inventory's measurement of the same bytes."""
        observed = (
            self.source_sites,
            self.residues,
            self.multi_gene_sites,
            self.multi_accession_sites,
            self.zero_cells,
            self.sites_zero_in_all_samples,
            self.quantified_per_sample,
            self.records,
        )
        expected = (
            SOURCE_SITES,
            SOURCE_RESIDUES,
            MULTI_GENE_SITES,
            MULTI_ACCESSION_SITES,
            ZERO_CELLS,
            SITES_ZERO_IN_ALL_SAMPLES,
            QUANTIFIED_PER_SAMPLE,
            EXPECTED_RECORDS,
        )
        if observed != expected:
            raise RuntimeError(f"build counts {observed} != pinned {expected}")
        if self.cells != self.zero_cells + sum(self.quantified_per_sample.values()):
            raise RuntimeError("zero and quantified cells do not add up to the grid")


def build_accounting(name: str, rows: Sequence[SiteRow]) -> BuildAccounting:
    """Count the release the way the release inventory does."""
    columns = [spec.column for spec in SAMPLES]
    comparison = Counter[str]()
    for row in rows:
        total = row.total_intensity
        summed = sum(row.intensities.values())
        if total == 0:
            comparison["zero"] += 1
            continue
        relative = (total - summed) / total
        if abs(relative) <= 1e-6:
            comparison["equal"] += 1
        elif relative > 0:
            comparison["greater"] += 1
        else:
            comparison["smaller"] += 1
    return BuildAccounting(
        dataset=name,
        source_sites=len(rows),
        residues=dict(
            sorted(
                Counter(row.site.residue.value for row in rows).items(),
                key=lambda item: "STYH".index(item[0]),
            )
        ),
        multi_gene_sites=sum(1 for row in rows if len(row.site.gene_names) > 1),
        multi_accession_sites=sum(1 for row in rows if len(row.site.protein_ids) > 1),
        diagnostic_peak_by_residue=dict(
            sorted(
                Counter(
                    row.site.residue.value for row in rows if row.site.diagnostic_peak
                ).items()
            )
        ),
        cells=len(rows) * len(columns),
        zero_cells=sum(1 for row in rows for c in columns if row.intensities[c] == 0),
        sites_zero_in_all_samples=sum(
            1 for row in rows if not any(row.intensities.values())
        ),
        quantified_per_sample={
            c: sum(1 for row in rows if row.intensities[c] > 0) for c in columns
        },
        records=len(SAMPLES),
        total_intensity_not_stored=TotalIntensityComparison(
            equal_within_1e_6_relative=comparison["equal"],
            total_greater=comparison["greater"],
            total_smaller=comparison["smaller"],
            total_zero=comparison["zero"],
        ),
    )


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class PhosphositeIntensityPotel2018Dataset(ExperimentDataset):
    """Potel 2018 phosphosite intensities of W3110 in 12 released samples."""

    organism = Organism.ecoli

    REFERENCE_STRAIN: ClassVar[Literal["W3110"]] = W3110_STRAIN
    #: Every record is the wild type: the twelve records differ by sample only.
    has_gene_perturbations = False

    def __init__(
        self,
        root: str = "data/torchcell/phosphosite_intensity_potel2018",
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; no genome is read, since no identifier is resolved."""
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return PhosphositeIntensityExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return PhosphositeIntensityExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """Supplementary Table 2, the one released per-record workbook."""
        return [TABLE_2_FILENAME]

    def download(self) -> None:
        """Link the deposited Table 2 into ``raw/`` after verifying its pins."""
        manifest = load_manifest()
        check_manifest_pin(
            TABLE_2_REL, manifest_sha256(manifest, TABLE_2_REL), TABLE_2_SHA256
        )
        os.makedirs(self.raw_dir, exist_ok=True)
        link_verified(
            raw_mirror_dir() / TABLE_2_REL,
            osp.join(self.raw_dir, TABLE_2_FILENAME),
            TABLE_2_SHA256,
        )

    @post_process
    def process(self) -> None:
        """Build one phosphosite-intensity record per released sample; write LMDB."""
        verify_raw_files(self.raw_dir, {TABLE_2_FILENAME: TABLE_2_SHA256})
        rows = read_table_2(osp.join(self.raw_dir, TABLE_2_FILENAME))
        accounting = build_accounting(self.name, rows)
        accounting.check()

        genome_reference = strain_reference()
        references = {
            spec.column: PhosphositeIntensityExperimentReference(
                dataset_name=self.name,
                genome_reference=genome_reference,
                environment_reference=environment(spec.carbon),
                phenotype_reference=phenotype(reference_sample(spec), rows),
            )
            for spec in SAMPLES
        }
        pub = publication()

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        index = 0
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for spec in tqdm(SAMPLES, desc="potel2018-phosphosite"):
                experiment = PhosphositeIntensityExperiment(
                    dataset_name=self.name,
                    genotype=Genotype(perturbations=[]),
                    environment=environment(spec.carbon),
                    phenotype=phenotype(spec, rows),
                )
                txn.put(
                    f"{index}".encode(),
                    self._intern_record(experiment, references[spec.column], pub, itxn),
                )
                index += 1
        env.close()
        interned_env.close()
        if index != EXPECTED_RECORDS:
            raise RuntimeError(f"{index} records, not {EXPECTED_RECORDS}")
        with open(
            osp.join(self.preprocess_dir, "build_accounting.json"), "w"
        ) as handle:
            handle.write(accounting.model_dump_json(indent=2))
        log.info("Potel 2018 phosphosites: %d records of %d sites", index, len(rows))

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process() for this dataset."""
        raise NotImplementedError(
            "PhosphositeIntensityPotel2018Dataset builds records in process()"
        )


# --------------------------------------------------------------------------- #
# Verification: the phosphosite family gate plus this release's own L1-L4
# --------------------------------------------------------------------------- #
def _records(dataset_root: str) -> list[dict[str, Any]]:
    """The built records of one tree."""
    from torchcell.verification.runners import load_records

    return load_records(dataset_root)


def _by_label(records: Sequence[Mapping[str, Any]]) -> dict[str, Mapping[str, Any]]:
    return {
        str(record["experiment"]["phenotype"]["sample_label"]): record
        for record in records
    }


def _l1_quantified_per_sample(records: Sequence[Mapping[str, Any]]) -> LevelResult:
    """L1: each sample quantifies the pinned number of sites, out of all 2,129."""
    observed = {
        label: len(record["experiment"]["phenotype"]["phosphosite_intensity"])
        for label, record in _by_label(records).items()
    }
    sites = {len(record["experiment"]["phenotype"]["sites"]) for record in records}
    passed = observed == QUANTIFIED_PER_SAMPLE and sites == {SOURCE_SITES}
    return LevelResult(
        level=Level.L1,
        name="quantified_sites_per_sample_are_the_pinned_ones",
        passed=passed,
        message=f"{len(observed)} samples; site counts per record {sorted(sites)}",
        details={"observed": observed, "expected": QUANTIFIED_PER_SAMPLE},
    )


def _l2_cells_are_the_released_cells(
    records: Sequence[Mapping[str, Any]], workbook: Path
) -> LevelResult:
    """L2: every stored intensity and status is the released cell, re-read.

    25,548 cells: a non-zero cell is the stored intensity exactly, and a zero cell is a
    ``not_quantified`` status with no intensity.
    """
    rows = read_table_2(workbook)
    by_label = _by_label(records)
    mismatched: list[str] = []
    checked = 0
    for spec in SAMPLES:
        stored = by_label[spec.column]["experiment"]["phenotype"]
        for row in rows:
            key = row.site.key()
            cell = row.intensities[spec.column]
            checked += 1
            status = stored["quantification"][key]
            value = stored["phosphosite_intensity"].get(key)
            ok = (
                status == "quantified" and value == float(cell)
                if cell > 0
                else status == "not_quantified" and value is None
            )
            if not ok:
                mismatched.append(f"{spec.column}:{key}")
            if stored["sites"][key] != row.site.model_dump(mode="json"):
                mismatched.append(f"{spec.column}:{key}:site")
    return LevelResult(
        level=Level.L2,
        name="every_stored_cell_and_site_is_the_released_one",
        passed=checked == SOURCE_SITES * EXPECTED_RECORDS and not mismatched,
        message=f"{checked} cells re-read; {len(mismatched)} differ",
        details={"n_checked": checked, "mismatched": mismatched[:20]},
    )


def _l3_sample_metadata(records: Sequence[Mapping[str, Any]]) -> LevelResult:
    """L3: each record's carbon, phase, replicate, OD and reference follow its legend."""
    wrong: list[str] = []
    for spec in SAMPLES:
        record = _by_label(records)[spec.column]
        stored = record["experiment"]["phenotype"]
        media = str(record["experiment"]["environment"]["media"]["name"])
        reference = record["reference"]["phenotype_reference"]["sample_label"]
        if (
            stored["harvest_growth_phase"] != str(spec.phase)
            or stored["replicate"] != spec.replicate
            or stored["harvest_optical_density"] != spec.harvest_optical_density
            or media != medium(spec.carbon).name
            or reference != reference_sample(spec).column
            or stored["replicate_unit"] is not None
        ):
            wrong.append(spec.column)
    return l3_convention(
        "sample_fields_follow_the_column_legend",
        not wrong and len(records) == len(SAMPLES),
        detail=f"{len(records)} records; {len(wrong)} disagree {wrong}",
    )


def _l3_assembly_pin(records: Sequence[Mapping[str, Any]]) -> LevelResult:
    """L3: every reference pins the W3110 assembly set."""
    pins = {
        (
            str(record["reference"]["genome_reference"].get("assembly_set")),
            str(record["reference"]["genome_reference"].get("strain")),
        )
        for record in records
    }
    return l3_convention(
        "assembly_pin_is_w3110",
        pins == {(W3110_ASSEMBLY_SET, W3110_STRAIN)},
        detail=f"{len(pins)} distinct pin(s): {sorted(pins)}",
    )


#: Every text-quoted value; each is re-read against the raw-mirror bytes at L3.
TEXT_SOURCED_VALUES: tuple[SourcedValue, ...] = (
    STRAIN,
    SALT_NA2HPO4_CLEAN,
    SALTS,
    CARBON_AND_MAGNESIUM,
    CALCIUM,
    TEMPERATURE_C,
    HARVEST_EXPONENTIAL,
    HARVEST_STATIONARY,
    TECHNICAL_TRIPLICATES_THESIS,
    SEARCH_SOFTWARE,
    SEARCH_DATABASE,
    BIOLOGICAL_REPLICATES_SI,
    TECHNICAL_TRIPLICATES_RS,
    INTENSITY_COLUMN,
    TOTAL_INTENSITY_COLUMN,
    LOCALIZATION_COLUMN,
    POSITIONS_COLUMN,
    DIAGNOSTIC_PEAK_COLUMN,
    *(spec.legend for spec in SAMPLES),
)


def _l3_quotes_are_verbatim(data_root: str | None) -> LevelResult:
    """L3: every quote is a verbatim substring of its sha256-pinned mirror file."""
    raw_root = Path(data_root or _data_root()) / "torchcell-raw"
    failed = [
        sv.quote[:60]
        for sv in TEXT_SOURCED_VALUES
        if not audit_sourced_value(sv, raw_root).passed
    ]
    return l3_convention(
        "sourced_quotes_are_verbatim_in_the_pinned_mirror",
        not failed,
        detail=f"{len(TEXT_SOURCED_VALUES)} quotes re-read; {len(failed)} failed {failed}",
    )


#: Cells of Table 5 that are log2 of a quantified Table 2 cell, measured by the release
#: inventory (995 + 1,011 + 991 + 1,187 + 1,122 + 1,207 over the six glycerol samples).
TABLE_5_AGREEING_CELLS = 6513
TABLE_5_LOG2_TOLERANCE = 1e-3


def read_table_5(path: str | Path) -> dict[str, dict[str, float]]:
    """Table 5 as site key -> released column -> log2 value (keys as Table 2 writes)."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book[TABLE_5_SHEET].iter_rows(values_only=True))
    book.close()
    header = [str(cell) for cell in rows[0]]
    log2_columns = [h for h in header if h.startswith("log2 (")]
    out: dict[str, dict[str, float]] = {}
    for values in rows[1:]:
        row = dict(zip(header, values, strict=True))
        key = "|".join(
            (str(row["Acc. No."]), str(row["Amino acid"]), str(row["Position"]))
        )
        out[key] = {
            column.removeprefix("log2 (").removesuffix(")"): float(row[column])
            for column in log2_columns
        }
    return out


def _l4_table_5_agrees(
    records: Sequence[Mapping[str, Any]], workbook: Path
) -> LevelResult:
    """L4: Table 5's log2 values are log2 of the stored intensities, cell by cell.

    Table 5 is the release's other per-site table. Its key is the released accession,
    residue and one position, so a site Table 2 lists on several genes joins only when
    Table 5 writes it the same way. Every joined cell this store quantified must equal
    log2 of the stored intensity; a joined cell the store calls ``not_quantified`` is a
    Perseus imputation in Table 5 and is counted, not compared.
    """
    table_5 = read_table_5(workbook)
    by_label = _by_label(records)
    agree = 0
    imputed = 0
    disagree: list[str] = []
    joined = {
        key for key in table_5 if key in records[0]["experiment"]["phenotype"]["sites"]
    }
    for key in sorted(joined):
        for column, value in table_5[key].items():
            intensity = by_label[column]["experiment"]["phenotype"][
                "phosphosite_intensity"
            ].get(key)
            if intensity is None:
                imputed += 1
            elif abs(math.log2(intensity) - value) <= TABLE_5_LOG2_TOLERANCE:
                agree += 1
            else:
                disagree.append(f"{column}:{key}")
    return LevelResult(
        level=Level.L4,
        name="table_5_log2_values_are_log2_of_the_stored_intensities",
        passed=not disagree and agree == TABLE_5_AGREEING_CELLS,
        message=(
            f"{len(joined)} of {len(table_5)} Table 5 sites join; {agree} cells agree, "
            f"{imputed} are imputations of not_quantified cells, {len(disagree)} differ"
        ),
        details={
            "joined_sites": len(joined),
            "table_5_sites": len(table_5),
            "agree": agree,
            "imputed": imputed,
            "disagree": disagree[:20],
        },
    )


def verify_build(
    dataset_root: str,
    *,
    data_root: str | None = None,
    expected_count: int = EXPECTED_RECORDS,
) -> VerificationReport:
    """Run the L0-L4 gate on a built tree and write its report.

    The phosphosite family gate (``torchcell.verification.phosphosite``) runs L0, the
    count, and the shape rules every phosphosite store shares; this release adds its
    per-sample quantified counts (L1), a re-read of all 25,548 cells (L2), the sample
    legend, assembly pin and verbatim quotes (L3) and Table 5's agreement (L4).
    """
    records = _records(dataset_root)
    mirror = raw_mirror_dir(data_root)
    for rel, sha in ((TABLE_2_REL, TABLE_2_SHA256), (TABLE_5_REL, TABLE_5_SHA256)):
        if file_sha256(mirror / rel) != sha:
            raise RuntimeError(f"{rel} no longer has its pinned sha256")
    report = verify_phosphosite_dataset(
        records,
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=Provenance(
            source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{TABLE_2_REL}",
            citation_key=CITATION_KEY,
            sha256=TABLE_2_SHA256,
            method=(
                "Supplementary Table 2, one PhosphositeIntensityExperiment "
                "per sample column, every site as released, 0 -> not_quantified"
            ),
            page=f"sheet {TABLE_2_SHEET}",
            retrieved=RAW_RETRIEVED_AT,
        ),
        expected_count=expected_count,
    )
    report.add(_l1_quantified_per_sample(records))
    report.add(_l2_cells_are_the_released_cells(records, mirror / TABLE_2_REL))
    report.add(_l3_sample_metadata(records))
    report.add(_l3_assembly_pin(records))
    report.add(_l3_quotes_are_verbatim(data_root))
    report.add(_l4_table_5_agrees(records, mirror / TABLE_5_REL))
    preprocess = osp.join(dataset_root, "preprocess")
    os.makedirs(preprocess, exist_ok=True)
    with open(osp.join(preprocess, "verification_report.json"), "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main() -> None:
    """Build the dataset under ``DATA_ROOT`` and verify it, for interactive debugging."""
    from dotenv import load_dotenv

    load_dotenv()
    data_root = _data_root()
    root = osp.join(data_root, "data/torchcell/phosphosite_intensity_potel2018")
    dataset = PhosphositeIntensityPotel2018Dataset(root=root)
    print(f"len = {len(dataset)}")
    print(Path(osp.join(root, "preprocess/build_accounting.json")).read_text())
    print(json.dumps(verify_build(root, data_root=data_root).summary()))


if __name__ == "__main__":
    main()
