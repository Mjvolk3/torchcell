# torchcell/datasets/ecoli/schmidt2016_srm
# [[torchcell.datasets.ecoli.schmidt2016_srm]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/schmidt2016_srm
# Test file: tests/torchcell/datasets/ecoli/test_schmidt2016_srm.py
"""Schmidt 2016 Tables S2 and S3: the SRM + SID absolute abundances, a second assay.

Schmidt et al. 2016 (Nat Biotechnol 34:104, doi:10.1038/nbt.3418) anchored its
proteome-wide label-free estimates on 41 proteins measured ABSOLUTELY by selected
reaction monitoring against synthetic heavy reference peptides: "41 proteins covering
key enzymes and iso-enzymes of carbohydrate metabolic pathways were selected for
absolute quantification by SRM and SID (Supplementary Table 1)." Those anchor
measurements are released per (protein, peptide, growth condition) in Supplementary
Tables 2 and 3, one table per LC-MS data set, and this module serves them:

- :class:`ProteomeSrmSet1Schmidt2016Dataset` -- Table S2, data set 1 (the
  OGE-fractionated arm), **779 released rows = 41 proteins x 19 conditions**, with a
  standard deviation "from Technical Replicates".
- :class:`ProteomeSrmSet2Schmidt2016Dataset` -- Table S3, data set 2 (the unfractionated
  arm), **682 released rows = 31 proteins x 22 conditions**, with a standard deviation
  "from Biological Triplicates".

TWO DATASET CLASSES, NOT ONE, AND THE REASON IS A VERIFIER INVARIANT.
``verify_protein_dataset``'s L3 ``measurement_type_consistent`` requires every record of
one dataset to share a single ``measurement_type``, which exists so heterogeneous
proteomics assays are never silently mixed. The two arms differ in what their released
dispersion IS (two SRM injections of one sample against three independently grown
cultures), so they are two measurement types and therefore two datasets. Merging them
under one type would also put two records on the same (genotype, environment) for the
twelve conditions both arms measured, with nothing on
``ProteinAbundancePhenotype`` to tell them apart -- there is no ``screen_id`` there.
The arm is the DATASET identity instead, which is what the knowledge graph already
distinguishes by.

THIS IS NOT THE STORED TABLE S6 BLOCK, AND THE BUILD PROVES IT RATHER THAN SAYING SO.
``ProteomeSchmidt2016Dataset`` stores Table S6's label-free copies/cell under
``absolute_protein_copies_per_cell_label_free_sid_anchored``, and every one of these 41
proteins is in that block, so the two could be confused. ``check_not_the_stored_block``
asserts at build time that (1) this dataset's ``measurement_type`` differs from the
stored block's and from the other arm's, and (2) **zero** of the shared (accession,
condition) cells agree to ``IDENTITY_RTOL``. Measured on the pinned workbook: 0 of 738
Table S2 cells and 0 of 682 Table S3 cells agree with Table S6, at Pearson r 0.664 and
0.928 on log10 copies/cell and a median absolute log2 ratio of 1.353 and 0.619. The two
arms disagree with each other too: 0 of the 558 cells they share agree, median ratio
0.394. So this is a second measurement of the same proteins, not a re-release of the
stored one, and nothing is deduplicated away.

WHICH TABLE IS WHICH DATA SET IS VERIFIED, NOT READ OFF THE TITLE. This release already
mislabels a block (Table S6's first two coefficient-of-variation headers are swapped,
asserted in ``schmidt2016.py``), so the "dataset 1" / "dataset 2" in the two titles is
checked against Table S6's own coverage pattern. Table S6 marks a cell ``NA`` where a
data set did not measure it, and measured over its 301 dataset-1 rows, every one is
``NA`` in exactly four conditions: Glycerol + AA, Xylose, Mannose and Fructose. Table
S2's 19 condition labels are exactly the remaining 18 plus ``anaerobic``, and Table S3's
22 are exactly all 22. The Methods independently put data set 1 at 19 samples ("Then,
all 19 peptide mixtures were separated on a 12-cm pH 3-10 immobilized PH gradient
strip"). ``check_dataset_arm`` asserts that pattern, so a swapped pair of titles stops
the build.

THE DISPERSION LABELS ARE CHECKED THE SAME WAY. Table S2 says its standard deviation
comes from technical replicates and Table S3 from biological triplicates. Measured:
Table S2's median relative spread is 1.771% and Table S3's is 6.489%, a 3.7-fold
difference in the direction the labels predict, and no Table S2 row has a standard
deviation at or above its abundance (0 of 779) against two of Table S3's 682.
``check_dispersion_labels`` reads BOTH sheets in either build and asserts the ordering,
which is what a swap between the two sheets would break.

A RELEASED CONDITION THE PAPER NEVER DESCRIBES. Table S2's 19th condition is labelled
``anaerobic``, and the string "anaerob" appears NOWHERE in the pinned ``paper.md`` or in
the Supplementary ``si1.md`` -- no medium, no cultivation protocol, no oxygen regime.
It is in no other released table either: not Table S6's 22 columns, not Table S23's
experimental details, not Table S25's sample-name map. There is nothing to build an
``Environment`` from and ``aerobicity="aerobic"``, which every other condition of this
paper carries, would be flatly wrong for it, so the condition is dropped under
``condition_not_described_by_the_source`` and ledgered. It is a real 19th data-set-1
sample, which is why it is recorded rather than ignored.

A RELEASED STANDARD DEVIATION OF EXACTLY ZERO, KEPT VERBATIM. 98 of Table S2's 779 rows
carry a standard deviation of exactly 0, across 13 of its 41 proteins; four of them
(``ytjC``, ``acnA``, ``pgk``, ``aceA``) carry it in all 19 conditions. Table S3 has
none. The released value is stored as given -- an SE of 0.0 -- because reading it as
"not measured" would be a reinterpretation, not a reading; the count is asserted at
build time so a corrected re-export is visible instead of silent, and the finding is
recorded in the dendron note.

THE RECORD GRAIN, AND THE DROPS. One ``BacterialProteinAbundanceExperiment`` per loaded
growth condition, keyed by BW25113 locus tag, which is the grain the stored Table S6
block uses. Each arm's glucose condition is the ``phenotype_reference`` rather than a
record, as it is there. The same three structural condition rules apply, for the same
measured reasons (``schmidt2016.py``'s ``DropReason`` objects are reused verbatim, so
there is one statement of each): the glycerol + amino acid medium has no
``MEDIA_LIBRARY`` entry, the four chemostat arms differ only in a dilution rate
``Environment`` cannot state, and the two stationary-phase arms collapse onto the
exponential glucose environment. Set 1: 19 - 1 anaerobic - 4 chemostat - 2 stationary -
1 reference = **11 records**. Set 2: 22 - 1 glycerol + AA - 4 chemostat - 2 stationary -
1 reference = **14 records**.

IDENTIFIERS. Each released ``Gene`` symbol is resolved against the pinned BW25113
GenBank annotation through ``reconcile_locus_tags``, the same derived route the stored
block takes, because the records are BW25113 and these tables release no b-number at
all. Measured: 41 of 41 and 31 of 31 symbols resolve (39 and 29 through the gene-symbol
layer, 2 and 2 through a gene synonym) with no collision and no ambiguity. The one
identifier inconsistency inside the release is pinned rather than smoothed over:
``P0ACP1`` is ``cra`` in Table S1 and in its own UniProt description's ``GN=`` token but
``fruR`` in Tables S2 and S3. They are synonyms of one gene and both resolve to the same
locus; ``check_identifier_columns`` asserts that this is the ONLY accession whose gene
name disagrees between the panel table and the measurement tables, and that every
released peptide sequence matches Table S1's selected proteotypic peptide exactly.

DATA. The one consumed file is ``si2.xlsx``, the same pinned Supplementary-tables
workbook ``schmidt2016.py`` consumes, read from the same raw mirror under the same
sha256 pin. This is a separate MODULE from ``schmidt2016.py`` so the served
``proteome_schmidt2016`` store's schema closure -- which ``build_manifest`` keys its
staleness on -- is left exactly as it was built.
"""

from __future__ import annotations

import argparse
import inspect
import json
import logging
import math
import os
import os.path as osp
import pickle
import statistics
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any, ClassVar

import openpyxl
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.datamodels.schema import (
    BacterialProteinAbundanceExperiment,
    BacterialProteinAbundanceExperimentReference,
    Experiment,
    ExperimentReference,
    Genotype,
    ProteinAbundancePhenotype,
    Publication,
)
from torchcell.datasets.bacteria_common import (
    BACTERIAL_ASSEMBLY_SETS,
    LocusTagReconciliation,
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.datasets.ecoli import schmidt2016 as sm
from torchcell.sequence import GeneSet
from torchcell.sequence.genome.ecoli.k12 import EcoliK12Genome, EcoliK12StrainName
from torchcell.verification.report import Provenance, VerificationReport
from torchcell.verification.sourced import SourcedValue

log = logging.getLogger(__name__)

SHEET_S1 = "Table S1"
SHEET_S2 = "Table S2"
SHEET_S3 = "Table S3"
#: Table S1's columns: the 41 selected proteins and their proteotypic peptides.
S1_ACCESSION = "Uniprot Accession"
S1_GENE = "Gene"
S1_DESCRIPTION = "Description"
S1_PEPTIDE = "Selected Peptide Sequence*"
S1_SPIKE = (
    "Concentration of Heavy Reference Peptides Spiked into Sample (fmol/ug of total "
    "protein)"
)
#: Tables S2 and S3 share these five leading columns; only the sixth header differs.
SRM_ACCESSION = "Uniprot Accession"
SRM_GENE = "Gene"
SRM_PEPTIDE = "Peptide Sequence"
SRM_CONDITION = "Growth Condition"
SRM_ABUNDANCE = "Protein Abundance (copies/cell)"
SRM_SD_TECHNICAL = "Standard Deviation (copies/cell) from Technical Replicates"
SRM_SD_BIOLOGICAL = "Standard Deviation (copies/cell) from Biological Triplicates"

#: The one accession the release names differently in Table S1 than in Tables S2 and S3.
GENE_NAME_DISAGREEMENT = {"P0ACP1": ("cra", "fruR")}
#: Relative tolerance below which two released numbers count as the same number.
IDENTITY_RTOL = 1e-6
#: Measured: every released symbol resolves in both panels, so a renamed one stops the
#: build instead of silently shrinking the abundance map.
MIN_RESOLVED_FRACTION = 1.0

SCHMIDT_REFERENCE_STRAIN: EcoliK12StrainName = sm.SCHMIDT_REFERENCE_STRAIN
#: Table S6's own condition-coverage pattern: the four columns every dataset-1 row
#: marks ``NA``, measured over its 301 dataset-1 rows.
DATASET1_UNCOVERED: tuple[str, ...] = ("Glycerol + AA", "Xylose", "Mannose", "Fructose")

#: ``{SRM growth-condition label: Table S6 condition column}``. The two sheets spell
#: three labels differently from each other and from Table S6, which is exactly why the
#: mapping is declared and checked rather than derived by string munging.
CONDITION_LABELS: dict[str, str] = {
    "glucose": "Glucose",
    "LB": "LB",
    "glycerol + AA": "Glycerol + AA",
    "acetate": "Acetate",
    "fumarate": "Fumarate",
    "glucosamine": "Glucosamine",
    "glycerol": "Glycerol",
    "pyruvate": "Pyruvate",
    "chemostat µ=0.5": "Chemostat µ=0.5",
    "chemostat µ=0.35": "Chemostat µ=0.35",
    "chemostat µ=0.2": "Chemostat µ=0.20",
    "chemostat µ=0.20": "Chemostat µ=0.20",
    "chemostat µ=0.12": "Chemostat µ=0.12",
    "stationary 1 day": "Stationary phase 1 day",
    "stationary 3 day": "Stationary phase 3 days",
    "stationary 3 days": "Stationary phase 3 days",
    "50 mM NaCl": "Osmotic-stress glucose",
    "42°C": "42°C glucose",
    "pH 6": "pH6 glucose",
    "xylose": "Xylose",
    "mannose": "Mannose",
    "galactose": "Galactose",
    "succinate": "Succinate",
    "fructose": "Fructose",
}
#: The one released condition label the paper never describes, so no environment exists.
UNDESCRIBED_CONDITION = "anaerobic"


# --------------------------------------------------------------------------- #
# Verbatim quotes
# --------------------------------------------------------------------------- #
_Q_SRM_PANEL = (
    "41 proteins covering key enzymes and iso-enzymes of carbohydrate "
    "metabolic pathways were selected for absolute quantification by SRM "
    "and SID (Supplementary Table 1)."
)
_Q_SRM_DUPLICATE = "Each sample was analyzed in duplicate."
_Q_OGE_NINETEEN = (
    "Then, all 19 peptide mixtures were separated on a $1 2 \\mathrm { - c"
    "m p H } 3 \\mathrm { - } 1 0$ immobilized PH gradient strip"
)
_Q_DATASET1_QUALITATIVE = (
    "Data generated in data set 1 were included only in the qualitative "
    "analysis of identified protein modifications illustrated in Table 1 "
    "and Figure ${ \\bf 5 a , b }$ ."
)
_Q_TABLE_S1 = (
    "Table S1 | Proteins selected for absolute quantification and their "
    "selected proteotypic peptides for which heavy reference peptides "
    "were synthesized and employed for quantification by stable isotope "
    "dilution"
)
_Q_TABLE_S2 = (
    "Table S2 | Absolute quantification of selected proteins for dataset "
    "1 (see supplemental Figure 4 for dataset details)"
)
_Q_TABLE_S3 = (
    "Table S3 | Absolute quantification of selected proteins for dataset "
    "2 (see supplemental Figure 4 for dataset details)"
)

PAPER = Provenance(
    source_uri=sm.PAPER_MD, citation_key=sm.CITATION_KEY, sha256=sm.PAPER_MD_SHA256
)
SI2_SOURCE = Provenance(
    source_uri=sm.SI2_MIRROR_RELPATH, citation_key=sm.CITATION_KEY, sha256=sm.SI2_SHA256
)

#: Quotes whose provenance is the pinned ``paper.md``, by the constant's own name.
PAPER_QUOTES: dict[str, str] = {
    "srm_panel": _Q_SRM_PANEL,
    "srm_duplicate": _Q_SRM_DUPLICATE,
    "oge_nineteen": _Q_OGE_NINETEEN,
    "dataset1_qualitative": _Q_DATASET1_QUALITATIVE,
    "copies_per_cell": sm._Q_COPIES_PER_CELL,
    "triplicates": sm._Q_TRIPLICATES,
    "strain": sm._Q_STRAIN,
    "batch_37c": sm._Q_BATCH_37C,
}
#: Quotes whose provenance is a ``si2.xlsx`` row rendering.
SI2_QUOTES: dict[str, str] = {
    "table_s1": _Q_TABLE_S1,
    "table_s2": _Q_TABLE_S2,
    "table_s3": _Q_TABLE_S3,
}


def _paper(value: object, quote: str, *, note: str | None = None) -> SourcedValue:
    """A ``SourcedValue`` quoting the pinned ``paper.md``."""
    return SourcedValue(value=value, provenance=PAPER, quote=quote, note=note)


def _si2(value: object, quote: str, *, note: str | None = None) -> SourcedValue:
    """A ``SourcedValue`` quoting the pinned ``si2.xlsx``."""
    return SourcedValue(value=value, provenance=SI2_SOURCE, quote=quote, note=note)


SOURCED_VALUES: dict[str, SourcedValue] = {
    "reference_strain": _paper(SCHMIDT_REFERENCE_STRAIN, sm._Q_STRAIN),
    "assay": _paper(
        "selected reaction monitoring against one synthetic heavy reference peptide "
        "per protein (stable isotope dilution)",
        _Q_SRM_PANEL,
    ),
    "stored_quantity": _paper("protein copies per cell", sm._Q_COPIES_PER_CELL),
    "n_replicates_set1": _paper(
        2,
        _Q_SRM_DUPLICATE,
        note="two SRM injections of one sample: the dispersion Table S2 releases is "
        "'from Technical Replicates', and data set 1 has no biological replicate",
    ),
    "n_replicates_set2": _paper(
        3,
        sm._Q_TRIPLICATES,
        note="three independently grown cultures: the dispersion Table S3 releases is "
        "'from Biological Triplicates'",
    ),
    "dataset1_sample_count": _paper(
        19,
        _Q_OGE_NINETEEN,
        note="the Methods' own count of data-set-1 samples, which is what makes Table "
        "S2's 19 condition labels (18 Table S6 columns plus the undescribed anaerobic "
        "one) a complete arm rather than a short read",
    ),
    "temperature_c": _paper(37.0, sm._Q_BATCH_37C),
    "panel_table": _si2("Table S1", _Q_TABLE_S1),
    "set1_table": _si2("Table S2", _Q_TABLE_S2),
    "set2_table": _si2("Table S3", _Q_TABLE_S3),
}


# --------------------------------------------------------------------------- #
# Reading the workbook
# --------------------------------------------------------------------------- #
class PanelProtein(BaseModel):
    """One Table S1 row: a selected protein, its peptide and its spike concentration."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    uniprot: str
    gene: str
    description: str
    peptide: str
    spike_fmol_per_ug: float


class SrmCell(BaseModel):
    """One released (protein, peptide, condition) SRM measurement."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    row_number: int = Field(description="1-based row in the sheet, for the ledger.")
    uniprot: str
    gene: str
    peptide: str
    condition: str = Field(description="Growth-condition label, verbatim.")
    copies_per_cell: float
    standard_deviation: float


def _text(value: Any, *, where: str) -> str:
    """A released cell as a non-empty stripped string."""
    if value is None or not str(value).strip():
        raise RuntimeError(f"{where}: the released cell is empty")
    return str(value).strip()


def _float(value: Any, *, where: str) -> float:
    """A released cell as a float; anything else raises."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise RuntimeError(f"{where}: {value!r} is not a number")
    return float(value)


def _header_index(header: Sequence[Any]) -> dict[str, int]:
    """``{stripped header: column index}``; a repeat is refused."""
    out: dict[str, int] = {}
    for position, cell in enumerate(header):
        if cell is None or not str(cell).strip():
            continue
        name = str(cell).strip()
        if name in out:
            raise RuntimeError(f"header {name!r} appears twice in one row")
        out[name] = position
    return out


def read_table_s1(path: str) -> dict[str, PanelProtein]:
    """``{accession: panel protein}`` of the 41 proteins selected for SRM."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book[SHEET_S1].iter_rows(values_only=True))
    finally:
        book.close()
    index = _header_index(rows[sm.HEADER_ROW - 1])
    out: dict[str, PanelProtein] = {}
    # The data block is contiguous and is followed by a blank row and a footnote row
    # ("* Peptides were selected according to ..."), so reading stops at the first blank
    # accession rather than trying to parse prose as a protein.
    for offset, row in enumerate(rows[sm.HEADER_ROW :], start=sm.HEADER_ROW + 1):
        cell = row[index[S1_ACCESSION]]
        if cell is None or not str(cell).strip():
            break
        where = f"{SHEET_S1} row {offset}"
        accession = _text(cell, where=where)
        if accession in out:
            raise RuntimeError(f"{where}: {accession} is filed twice in Table S1")
        out[accession] = PanelProtein(
            uniprot=accession,
            gene=_text(row[index[S1_GENE]], where=where),
            description=_text(row[index[S1_DESCRIPTION]], where=where),
            peptide=_text(row[index[S1_PEPTIDE]], where=where),
            spike_fmol_per_ug=_float(row[index[S1_SPIKE]], where=where),
        )
    return out


def read_srm_table(path: str, sheet: str, sd_header: str) -> list[SrmCell]:
    """Read one SRM sheet; its sixth header must be the declared dispersion name.

    The dispersion header is what says whether the released standard deviation is of
    technical or of biological replicates, so it is matched exactly rather than read
    positionally: a re-export that swaps the two sheets' dispersion columns stops the
    build here.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book[sheet].iter_rows(values_only=True))
    finally:
        book.close()
    index = _header_index(rows[sm.HEADER_ROW - 1])
    declared = (
        SRM_ACCESSION,
        SRM_GENE,
        SRM_PEPTIDE,
        SRM_CONDITION,
        SRM_ABUNDANCE,
        sd_header,
    )
    missing = [name for name in declared if name not in index]
    if missing:
        raise RuntimeError(f"{sheet} carries no column {missing}")
    out: list[SrmCell] = []
    for offset, row in enumerate(rows[sm.HEADER_ROW :], start=sm.HEADER_ROW + 1):
        cell = row[index[SRM_ACCESSION]]
        if cell is None or not str(cell).strip():
            break
        where = f"{sheet} row {offset}"
        out.append(
            SrmCell(
                row_number=offset,
                uniprot=_text(cell, where=where),
                gene=_text(row[index[SRM_GENE]], where=where),
                peptide=_text(row[index[SRM_PEPTIDE]], where=where),
                condition=_text(row[index[SRM_CONDITION]], where=where),
                copies_per_cell=_float(row[index[SRM_ABUNDANCE]], where=where),
                standard_deviation=_float(row[index[sd_header]], where=where),
            )
        )
    keys = {(cell.uniprot, cell.peptide, cell.condition) for cell in out}
    if len(keys) != len(out):
        raise RuntimeError(
            f"{sheet}: {len(out)} rows carry only {len(keys)} distinct "
            "(protein, peptide, condition) keys"
        )
    return out


def read_table_s6_copies(path: str) -> dict[tuple[str, str], float]:
    """``{(accession, Table S6 condition): copies/cell}`` of the STORED block.

    Read only to prove that this dataset is not that one; nothing from it is stored.
    """
    rows, _ = sm.read_table_s6(path)
    return {
        (row.uniprot, column): value
        for row in rows
        for column, value in row.copies.items()
        if value is not None
    }


def dataset1_uncovered_conditions(path: str) -> tuple[str, ...]:
    """Table S6 columns in which EVERY dataset-1 row is ``NA``, from the bytes.

    This is Table S6's own statement of which conditions data set 1 did not measure, and
    it is what ``check_dataset_arm`` holds the two sheet titles against.
    """
    rows, _ = sm.read_table_s6(path)
    dataset1 = [row for row in rows if row.dataset != sm.DATASET_WITH_REPLICATES]
    if not dataset1:
        raise RuntimeError("Table S6 releases no dataset-1 row")
    return tuple(
        spec.s6_column
        for spec in sm.CONDITIONS
        if all(row.copies[spec.s6_column] is None for row in dataset1)
    )


# --------------------------------------------------------------------------- #
# Header-versus-values checks
# --------------------------------------------------------------------------- #
def check_identifier_columns(
    cells: Sequence[SrmCell], panel: Mapping[str, PanelProtein], *, sheet: str
) -> dict[str, Any]:
    """Every released accession, peptide and gene name agrees with Table S1.

    The peptide must match exactly, because Table S1 names the ONE proteotypic peptide
    whose heavy reference was synthesized, so a mismatch would mean the measurement is
    of a peptide no standard was spiked for. The gene name is allowed to disagree for
    exactly the accessions ``GENE_NAME_DISAGREEMENT`` names, with the two released
    spellings pinned, and for no others.
    """
    outside = sorted({cell.uniprot for cell in cells} - set(panel))
    if outside:
        raise RuntimeError(f"{sheet}: {outside} are not Table S1 panel proteins")
    peptide_faults = sorted(
        {
            (cell.uniprot, cell.peptide, panel[cell.uniprot].peptide)
            for cell in cells
            if cell.peptide != panel[cell.uniprot].peptide
        }
    )
    if peptide_faults:
        raise RuntimeError(
            f"{sheet}: released peptide disagrees with Table S1's selected peptide for "
            f"{peptide_faults}"
        )
    disagreements = {
        cell.uniprot: (panel[cell.uniprot].gene, cell.gene)
        for cell in cells
        if cell.gene != panel[cell.uniprot].gene
    }
    if disagreements != GENE_NAME_DISAGREEMENT:
        raise RuntimeError(
            f"{sheet}: the gene names disagreeing with Table S1 are {disagreements}, "
            f"the module declares {GENE_NAME_DISAGREEMENT}"
        )
    return {
        "n_cells": len(cells),
        "n_panel_proteins": len({cell.uniprot for cell in cells}),
        "peptides_match_table_s1": True,
        "gene_name_disagreements": {
            key: list(value) for key, value in disagreements.items()
        },
    }


def check_dataset_arm(
    cells: Sequence[SrmCell], uncovered: Sequence[str], *, sheet: str, is_set1: bool
) -> dict[str, Any]:
    """The sheet's condition labels are its data set's, by Table S6's coverage pattern.

    Data set 1 measured the Table S6 conditions that are NOT in ``uncovered``, plus the
    undescribed ``anaerobic`` sample; data set 2 measured all 22. Holding each title
    against the other sheet's expected set is what a swapped pair of titles fails.
    """
    labels = sorted({cell.condition for cell in cells})
    unmapped = sorted(
        label
        for label in labels
        if label != UNDESCRIBED_CONDITION and label not in CONDITION_LABELS
    )
    if unmapped:
        raise RuntimeError(f"{sheet}: no Table S6 column is declared for {unmapped}")
    mapped = sorted(
        {CONDITION_LABELS[label] for label in labels if label != UNDESCRIBED_CONDITION}
    )
    declared = [spec.s6_column for spec in sm.CONDITIONS]
    expected = (
        sorted(column for column in declared if column not in uncovered)
        if is_set1
        else sorted(declared)
    )
    if mapped != expected:
        raise RuntimeError(
            f"{sheet}: its conditions map onto {mapped}, but the Table S6 coverage "
            f"pattern puts this data set on {expected}"
        )
    extra = [label for label in labels if label == UNDESCRIBED_CONDITION]
    if bool(extra) is not is_set1:
        raise RuntimeError(
            f"{sheet}: {UNDESCRIBED_CONDITION!r} is {'absent' if is_set1 else 'present'}"
            ", which contradicts the declared data-set arm"
        )
    counts = {
        label: sum(1 for cell in cells if cell.condition == label) for label in labels
    }
    if len(set(counts.values())) != 1:
        raise RuntimeError(
            f"{sheet}: the released rows per condition are not uniform: {counts}"
        )
    return {
        "n_conditions": len(labels),
        "rows_per_condition": next(iter(counts.values())),
        "table_s6_columns": mapped,
        "dataset1_uncovered_by_table_s6": list(uncovered),
        "undescribed_condition_present": bool(extra),
    }


def check_dispersion_labels(
    set1: Sequence[SrmCell], set2: Sequence[SrmCell]
) -> dict[str, Any]:
    """The technical-replicate sheet is tighter than the biological-triplicate one.

    Both sheets are read in either build, because the claim this checks is a comparison
    between them: a dispersion from two injections of one sample must be smaller than
    one from three independently grown cultures, and that ordering is what a swap of the
    two sheets' dispersion columns would invert.
    """

    def spread(cells: Sequence[SrmCell]) -> list[float]:
        return [
            100.0 * cell.standard_deviation / cell.copies_per_cell
            for cell in cells
            if cell.copies_per_cell > 0
        ]

    technical, biological = spread(set1), spread(set2)
    technical_median = statistics.median(technical)
    biological_median = statistics.median(biological)
    if technical_median >= biological_median:
        raise RuntimeError(
            f"the Table S2 dispersion labelled 'Technical Replicates' has median "
            f"{technical_median:.4f}% against Table S3's 'Biological Triplicates' "
            f"{biological_median:.4f}%; the two dispersion columns no longer read the "
            "way their headers say"
        )
    return {
        "set1_technical_median_cv_percent": technical_median,
        "set2_biological_median_cv_percent": biological_median,
        "set1_rows_with_sd_at_or_above_abundance": sum(
            1 for cell in set1 if cell.standard_deviation >= cell.copies_per_cell
        ),
        "set2_rows_with_sd_at_or_above_abundance": sum(
            1 for cell in set2 if cell.standard_deviation >= cell.copies_per_cell
        ),
        "set1_rows_with_zero_sd": sum(
            1 for cell in set1 if cell.standard_deviation == 0.0
        ),
        "set2_rows_with_zero_sd": sum(
            1 for cell in set2 if cell.standard_deviation == 0.0
        ),
    }


def check_not_the_stored_block(
    cells: Sequence[SrmCell],
    stored: Mapping[tuple[str, str], float],
    *,
    measurement_type: str,
    sheet: str,
) -> dict[str, Any]:
    """This arm is a second measurement, not a re-release of the stored Table S6 block.

    Two things are asserted. The ``measurement_type`` must differ from the stored
    block's and from the other arm's, so no consumer can join the two as one quantity.
    And NO shared (accession, condition) cell may agree to ``IDENTITY_RTOL``: one that
    did would mean the released SRM number IS the stored label-free number for that
    cell, which is the only way the two blocks could be mixed without anyone noticing.
    """
    declared = (sm.MEASUREMENT_TYPE, *MEASUREMENT_TYPES.values())
    if len(set(declared)) != len(declared):
        raise RuntimeError(
            f"the stored block's measurement_type and the two SRM arms' are not "
            f"pairwise distinct: {declared}"
        )
    if measurement_type not in MEASUREMENT_TYPES.values():
        raise RuntimeError(
            f"{measurement_type!r} is not one of the declared SRM measurement types "
            f"{sorted(MEASUREMENT_TYPES.values())}"
        )
    if measurement_type == sm.MEASUREMENT_TYPE:
        raise RuntimeError(
            f"{measurement_type!r} is the stored Table S6 block's measurement_type"
        )
    pairs = [
        (cell, stored[(cell.uniprot, CONDITION_LABELS[cell.condition])])
        for cell in cells
        if cell.condition != UNDESCRIBED_CONDITION
        and (cell.uniprot, CONDITION_LABELS[cell.condition]) in stored
    ]
    if not pairs:
        raise RuntimeError(f"{sheet}: no cell could be compared against Table S6")
    agreeing = [
        (cell.uniprot, cell.condition)
        for cell, other in pairs
        if other > 0 and abs(cell.copies_per_cell - other) / other < IDENTITY_RTOL
    ]
    if agreeing:
        raise RuntimeError(
            f"{sheet}: {len(agreeing)} cells agree with the stored Table S6 block to "
            f"{IDENTITY_RTOL} ({agreeing[:5]}); the two blocks are no longer separable"
        )
    ratios = sorted(cell.copies_per_cell / other for cell, other in pairs if other > 0)
    left = [math.log10(cell.copies_per_cell) for cell, other in pairs if other > 0]
    right = [math.log10(other) for _cell, other in pairs if other > 0]
    mean_left, mean_right = statistics.fmean(left), statistics.fmean(right)
    covariance = sum(
        (a - mean_left) * (b - mean_right) for a, b in zip(left, right, strict=True)
    )
    spread = math.sqrt(
        sum((a - mean_left) ** 2 for a in left)
        * sum((b - mean_right) ** 2 for b in right)
    )
    return {
        "measurement_type": measurement_type,
        "stored_measurement_type": sm.MEASUREMENT_TYPE,
        "n_shared_cells": len(pairs),
        "n_agreeing_to_rtol": 0,
        "rtol": IDENTITY_RTOL,
        "median_srm_over_stored": statistics.median(ratios),
        "pearson_r_log10": covariance / spread,
        "median_abs_log2_ratio": statistics.median(
            [abs(math.log2(ratio)) for ratio in ratios]
        ),
    }


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
#: The one condition Table S2 releases that the paper never describes.
DROP_CONDITION_UNDESCRIBED = sm.DropReason(
    rule="condition_not_described_by_the_source",
    description=(
        "Table S2 releases a 19th growth condition labelled 'anaerobic', and the string "
        "'anaerob' appears nowhere in the pinned paper.md or si1.md: no medium, no "
        "cultivation protocol and no oxygen regime is stated for it, and it is absent "
        "from Table S6's 22 columns, from Table S23's experimental details and from "
        "Table S25's sample map. There is nothing to build an Environment from, and the "
        "aerobic regime every described condition carries would be wrong for it"
    ),
    needed_addition=None,
)

#: ``{dataset class name: measurement_type}``, so the separability check can see both.
MEASUREMENT_TYPES: dict[str, str] = {
    "ProteomeSrmSet1Schmidt2016Dataset": (
        "absolute_protein_copies_per_cell_srm_sid_dataset1_technical_sd"
    ),
    "ProteomeSrmSet2Schmidt2016Dataset": (
        "absolute_protein_copies_per_cell_srm_sid_dataset2_biological_sd"
    ),
}


class ArmSelection(BaseModel):
    """What one arm keeps: its conditions, its protein keys and every drop."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    locus_tag: dict[str, str] = Field(description="uniprot accession -> locus tag")
    reconciliation: LocusTagReconciliation
    loaded_columns: list[str] = Field(
        description="Table S6 condition columns this arm loads, the reference included"
    )
    dropped: dict[str, list[str]] = Field(
        description="drop rule -> the released condition labels it removed"
    )


def select_arm(
    cells: Sequence[SrmCell], genome: EcoliK12Genome, *, label: str
) -> ArmSelection:
    """Resolve the arm's gene symbols, and partition its conditions by the drop rules."""
    accessions = sorted({cell.uniprot for cell in cells})
    genes = {cell.uniprot: cell.gene for cell in cells}
    stored, report = reconcile_locus_tags(
        genome, pd.Series([genes[key] for key in accessions]), label=label
    )
    report.require_resolved(MIN_RESOLVED_FRACTION)
    outside = set(report.outside_namespace)
    unresolved = [tag for tag in stored if tag in outside]
    if unresolved:
        raise RuntimeError(f"{label}: {unresolved} resolve to no BW25113 locus")
    locus_tag = {key: str(tag) for key, tag in zip(accessions, stored, strict=True)}
    if len(set(locus_tag.values())) != len(locus_tag):
        raise RuntimeError(f"{label}: two released symbols share one locus tag")

    labels = sorted({cell.condition for cell in cells})
    dropped: dict[str, list[str]] = {}
    loaded: list[str] = []
    for released in labels:
        if released == UNDESCRIBED_CONDITION:
            dropped.setdefault(DROP_CONDITION_UNDESCRIBED.rule, []).append(released)
            continue
        column = CONDITION_LABELS[released]
        reason = sm.CONDITIONS_BY_COLUMN[column].drop
        if reason is None:
            loaded.append(column)
        else:
            dropped.setdefault(reason.rule, []).append(released)
    order = [spec.s6_column for spec in sm.CONDITIONS]
    return ArmSelection(
        locus_tag=locus_tag,
        reconciliation=report,
        loaded_columns=sorted(loaded, key=order.index),
        dropped={rule: sorted(items) for rule, items in dropped.items()},
    )


def build_phenotype(
    cells: Sequence[SrmCell],
    locus_tag: Mapping[str, str],
    column: str,
    *,
    measurement_type: str,
    n_replicates: int,
) -> ProteinAbundancePhenotype:
    """The SRM abundance profile of one condition, keyed by BW25113 locus tag.

    The released standard deviation is a dispersion of ``n_replicates`` observations, so
    the stored standard error divides it by ``sqrt(n)``. A released dispersion of exactly
    0 is kept as given: 98 of Table S2's rows carry one, and reading it as "not measured"
    would be a reinterpretation rather than a reading.
    """
    abundance: dict[str, float] = {}
    standard_error: dict[str, float] = {}
    replicates: dict[str, int] = {}
    for cell in cells:
        # The undescribed condition has no Table S6 column and no environment, so it is
        # never a record's column; ``check_dataset_arm`` has already refused any OTHER
        # unmapped label, so a KeyError here would be a real defect.
        if cell.condition == UNDESCRIBED_CONDITION:
            continue
        if CONDITION_LABELS[cell.condition] != column:
            continue
        tag = locus_tag[cell.uniprot]
        abundance[tag] = cell.copies_per_cell
        standard_error[tag] = cell.standard_deviation / math.sqrt(n_replicates)
        replicates[tag] = n_replicates
    return ProteinAbundancePhenotype(
        protein_abundance=abundance,
        protein_abundance_se=standard_error,
        n_replicates=replicates,
        measurement_type=measurement_type,
    )


def publication() -> Publication:
    """The paper, as ``schmidt2016`` resolves it from ``PMC4888949``."""
    return sm.publication()


# --------------------------------------------------------------------------- #
# The datasets
# --------------------------------------------------------------------------- #
class _SrmSchmidt2016Dataset(ExperimentDataset):
    """Shared build for one SRM arm; the subclasses declare which arm they are."""

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = SCHMIDT_REFERENCE_STRAIN
    #: The sheet this arm reads, and the dispersion header it must carry.
    SHEET: ClassVar[str]
    SD_HEADER: ClassVar[str]
    #: What one stored abundance IS, and how many observations its dispersion is of.
    MEASUREMENT_TYPE: ClassVar[str]
    N_REPLICATES: ClassVar[int]
    #: ``True`` for the arm Table S6's coverage pattern puts on 18 of its 22 conditions.
    IS_SET1: ClassVar[bool]
    #: Measured on the pinned workbook: released rows, protein keys, records.
    EXPECTED_ROWS: ClassVar[int]
    EXPECTED_PROTEIN_KEYS: ClassVar[int]
    EXPECTED_RECORDS: ClassVar[int]

    def __init__(
        self,
        root: str,
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialProteinAbundanceExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialProteinAbundanceExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The one consumed file, linked from the raw mirror."""
        return [sm.SI2]

    def download(self) -> None:
        """Link the mirrored workbook into ``raw/`` after checking manifest and sha256."""
        data_root = os.environ["DATA_ROOT"]
        manifest = sm.load_manifest(data_root)
        check_manifest_pin(
            sm.SI2_MIRROR_RELPATH,
            sm.manifest_sha256(manifest, sm.SI2_MIRROR_RELPATH),
            sm.SI2_SHA256,
        )
        src = sm.raw_mirror_dir(data_root) / sm.SI2_MIRROR_RELPATH
        if not src.exists():
            raise RuntimeError(f"required raw artifact missing from mirror: {src}")
        os.makedirs(self.raw_dir, exist_ok=True)
        link_verified(src, osp.join(self.raw_dir, sm.SI2), sm.SI2_SHA256)
        log.info(
            "Schmidt 2016 %s raw file linked into %s (sha256 verified)",
            self.SHEET,
            self.raw_dir,
        )

    def compute_gene_set(self) -> GeneSet:
        """The BW25113 loci this arm's abundance profiles are keyed by.

        Every record is wild type, so no genotype names a gene and the base class's
        genotype scan would return the empty set it refuses. This is the same override
        ``ProteomeSchmidt2016Dataset`` carries, for the same reason.
        """
        if self.env is None:
            self._init_db()
        genes = GeneSet()
        with self.env.begin() as txn:
            for _, value in txn.cursor():
                record = pickle.loads(value)
                genes.update(record["experiment"]["phenotype"]["protein_abundance"])
        self.close_lmdb()
        return genes

    def _genome(self) -> EcoliK12Genome:
        """The injected genome, or the reference strain's default cache."""
        if self.ecoli_genome is None:
            self.ecoli_genome = bacterial_genome("ecoli", self.REFERENCE_STRAIN)
        expected = BACTERIAL_ASSEMBLY_SETS[self.REFERENCE_STRAIN]
        if self.ecoli_genome.ASSEMBLY_SET != expected:
            raise ValueError(
                f"{type(self).__name__} needs the {expected} genome, got "
                f"{self.ecoli_genome.ASSEMBLY_SET}"
            )
        return self.ecoli_genome

    @post_process
    def process(self) -> None:
        """Build one abundance record per loaded condition and write the LMDB."""
        verify_raw_files(self.raw_dir, sm.DATA_SHA256)
        path = osp.join(self.raw_dir, sm.SI2)
        panel = read_table_s1(path)
        cells = read_srm_table(path, self.SHEET, self.SD_HEADER)
        if len(cells) != self.EXPECTED_ROWS:
            raise RuntimeError(
                f"{self.SHEET} releases {len(cells)} data rows, the module states "
                f"{self.EXPECTED_ROWS}"
            )
        identifier_check = check_identifier_columns(cells, panel, sheet=self.SHEET)
        arm_check = check_dataset_arm(
            cells,
            dataset1_uncovered_conditions(path),
            sheet=self.SHEET,
            is_set1=self.IS_SET1,
        )
        dispersion_check = check_dispersion_labels(
            read_srm_table(path, SHEET_S2, SRM_SD_TECHNICAL),
            read_srm_table(path, SHEET_S3, SRM_SD_BIOLOGICAL),
        )
        separability_check = check_not_the_stored_block(
            cells,
            read_table_s6_copies(path),
            measurement_type=self.MEASUREMENT_TYPE,
            sheet=self.SHEET,
        )

        genome = self._genome()
        selection = select_arm(cells, genome, label=f"{self.name} gene symbols")
        environments = {
            column: sm.build_environment(sm.CONDITIONS_BY_COLUMN[column])
            for column in selection.loaded_columns
        }
        sm.check_environments_distinct(environments)
        phenotypes = {
            column: build_phenotype(
                cells,
                selection.locus_tag,
                column,
                measurement_type=self.MEASUREMENT_TYPE,
                n_replicates=self.N_REPLICATES,
            )
            for column in selection.loaded_columns
        }
        keys = {len(phenotype.protein_abundance) for phenotype in phenotypes.values()}
        if keys != {self.EXPECTED_PROTEIN_KEYS}:
            raise RuntimeError(
                f"{self.SHEET}: protein-key counts {sorted(keys)}, the module states "
                f"{self.EXPECTED_PROTEIN_KEYS} on every record"
            )
        reference = BacterialProteinAbundanceExperimentReference(
            dataset_name=self.name,
            genome_reference=assembly_reference(self.REFERENCE_STRAIN),
            environment_reference=environments[sm.REFERENCE_CONDITION],
            phenotype_reference=phenotypes[sm.REFERENCE_CONDITION],
        )
        pub = publication()
        records = [
            column
            for column in selection.loaded_columns
            if column != sm.REFERENCE_CONDITION
        ]
        if len(records) != self.EXPECTED_RECORDS:
            raise RuntimeError(
                f"{len(records)} records, the module states {self.EXPECTED_RECORDS}"
            )

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        rows: list[dict[str, Any]] = []
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for index, column in enumerate(tqdm(records, desc=self.name)):
                phenotype = phenotypes[column]
                experiment = BacterialProteinAbundanceExperiment(
                    dataset_name=self.name,
                    genotype=Genotype(perturbations=[]),
                    environment=environments[column],
                    phenotype=phenotype,
                )
                txn.put(
                    f"{index}".encode(),
                    self._intern_record(experiment, reference, pub, itxn),
                )
                spec = sm.CONDITIONS_BY_COLUMN[column]
                rows.append(
                    {
                        "s6_column": column,
                        "media": spec.media.name,
                        "carbon_source": spec.carbon_source,
                        "carbon_g_per_l": spec.carbon_g_per_l,
                        "temperature_c": spec.temperature_c,
                        "n_protein_keys": len(phenotype.protein_abundance),
                        "n_with_zero_se": sum(
                            1
                            for value in (phenotype.protein_abundance_se or {}).values()
                            if value == 0.0
                        ),
                    }
                )
        env.close()
        interned_env.close()

        self._write_ledgers(
            cells,
            panel,
            selection,
            rows,
            {
                "identifier_columns": identifier_check,
                "dataset_arm": arm_check,
                "dispersion_labels": dispersion_check,
                "not_the_stored_block": separability_check,
            },
        )
        log.info(
            "Schmidt 2016 %s: %d records (+ the %s reference) x %d protein keys from "
            "%d released rows; dropped conditions %s",
            self.SHEET,
            len(records),
            sm.REFERENCE_CONDITION,
            self.EXPECTED_PROTEIN_KEYS,
            len(cells),
            selection.dropped,
        )

    def _write_ledgers(
        self,
        cells: Sequence[SrmCell],
        panel: Mapping[str, PanelProtein],
        selection: ArmSelection,
        rows: Sequence[Mapping[str, Any]],
        checks: Mapping[str, Mapping[str, Any]],
    ) -> None:
        """The drop log, the sourcing table, the identifier table and the checks."""
        out = Path(self.preprocess_dir)
        labels = sorted({cell.condition for cell in cells})
        reasons = {
            DROP_CONDITION_UNDESCRIBED.rule: DROP_CONDITION_UNDESCRIBED,
            sm.DROP_MEDIUM_NOT_IN_LIBRARY.rule: sm.DROP_MEDIUM_NOT_IN_LIBRARY,
            sm.DROP_CULTURE_NOT_BATCH.rule: sm.DROP_CULTURE_NOT_BATCH,
            sm.DROP_GROWTH_PHASE.rule: sm.DROP_GROWTH_PHASE,
        }
        dropped = sum(len(items) for items in selection.dropped.values())
        if len(rows) + dropped + 1 != len(labels):
            raise RuntimeError(
                f"{self.SHEET}: {len(rows)} records + {dropped} dropped conditions + 1 "
                f"reference != {len(labels)} released conditions"
            )
        (out / "dropped_records.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "sheet": self.SHEET,
                    "source_conditions": len(labels),
                    "source_rows": len(cells),
                    "reference_conditions": [sm.REFERENCE_CONDITION],
                    "kept_records": len(rows),
                    "dropped_records": dropped,
                    "kept_protein_keys": self.EXPECTED_PROTEIN_KEYS,
                    "rules": [
                        {
                            "rule": rule,
                            "scope": "condition",
                            "description": reasons[rule].description,
                            "n_items": len(items),
                            "items": items,
                            "needed_addition": reasons[rule].needed_addition,
                        }
                        for rule, items in sorted(selection.dropped.items())
                    ],
                    "reconciliation": selection.reconciliation.model_dump(mode="json"),
                    "notes": [
                        f"{len(cells)} released rows = "
                        f"{len({cell.uniprot for cell in cells})} panel proteins x "
                        f"{len(labels)} growth conditions, one row per (protein, "
                        "peptide, condition); every panel protein carries exactly one "
                        "proteotypic peptide, so the peptide axis has length one",
                        f"the {sm.REFERENCE_CONDITION} condition is the "
                        "phenotype_reference rather than a record, which is the "
                        "convention the stored Table S6 block uses: the release defines "
                        "every fold change against BW25113 in glucose minimal medium",
                        "the three structural condition rules are schmidt2016.py's own "
                        "DropReason objects, reused rather than restated, so there is "
                        "one statement of each",
                        "this block is NOT the stored Table S6 block: zero of the "
                        f"{checks['not_the_stored_block']['n_shared_cells']} shared "
                        "(accession, condition) cells agree to "
                        f"{IDENTITY_RTOL}, at Pearson r "
                        f"{checks['not_the_stored_block']['pearson_r_log10']:.4f} on "
                        "log10 copies/cell",
                    ],
                },
                indent=2,
            )
        )
        (out / "released_statistics_check.json").write_text(
            json.dumps({**{k: dict(v) for k, v in checks.items()}}, indent=2)
        )
        (out / "sourced_values.json").write_text(
            json.dumps(
                {
                    name: value.model_dump(mode="json")
                    for name, value in SOURCED_VALUES.items()
                },
                indent=2,
            )
        )
        pd.DataFrame(list(rows)).to_csv(out / "conditions.csv", index=False)
        pd.DataFrame(
            [
                {
                    "uniprot": accession,
                    "released_gene": next(
                        cell.gene for cell in cells if cell.uniprot == accession
                    ),
                    "table_s1_gene": panel[accession].gene,
                    "peptide": panel[accession].peptide,
                    "spike_fmol_per_ug": panel[accession].spike_fmol_per_ug,
                    "stored_locus_tag": tag,
                    "n_replicates": self.N_REPLICATES,
                }
                for accession, tag in sorted(selection.locus_tag.items())
            ]
        ).to_csv(out / "protein_identifiers.csv", index=False)

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process() for this dataset."""
        raise NotImplementedError(f"{type(self).__name__} builds records in process()")


@register_dataset
class ProteomeSrmSet1Schmidt2016Dataset(_SrmSchmidt2016Dataset):
    """Table S2: the OGE-fractionated arm's SRM absolute abundances, 11 records."""

    SHEET: ClassVar[str] = SHEET_S2
    SD_HEADER: ClassVar[str] = SRM_SD_TECHNICAL
    MEASUREMENT_TYPE: ClassVar[str] = MEASUREMENT_TYPES[
        "ProteomeSrmSet1Schmidt2016Dataset"
    ]
    N_REPLICATES: ClassVar[int] = 2
    IS_SET1: ClassVar[bool] = True
    EXPECTED_ROWS: ClassVar[int] = 779
    EXPECTED_PROTEIN_KEYS: ClassVar[int] = 41
    EXPECTED_RECORDS: ClassVar[int] = 11

    def __init__(
        self,
        root: str = "data/torchcell/proteome_srm_set1_schmidt2016",
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize with this arm's own build root."""
        super().__init__(
            root,
            io_workers,
            transform,
            pre_transform,
            ecoli_genome=ecoli_genome,
            **kwargs,
        )


@register_dataset
class ProteomeSrmSet2Schmidt2016Dataset(_SrmSchmidt2016Dataset):
    """Table S3: the unfractionated arm's SRM absolute abundances, 14 records."""

    SHEET: ClassVar[str] = SHEET_S3
    SD_HEADER: ClassVar[str] = SRM_SD_BIOLOGICAL
    MEASUREMENT_TYPE: ClassVar[str] = MEASUREMENT_TYPES[
        "ProteomeSrmSet2Schmidt2016Dataset"
    ]
    N_REPLICATES: ClassVar[int] = 3
    IS_SET1: ClassVar[bool] = False
    EXPECTED_ROWS: ClassVar[int] = 682
    EXPECTED_PROTEIN_KEYS: ClassVar[int] = 31
    EXPECTED_RECORDS: ClassVar[int] = 14

    def __init__(
        self,
        root: str = "data/torchcell/proteome_srm_set2_schmidt2016",
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize with this arm's own build root."""
        super().__init__(
            root,
            io_workers,
            transform,
            pre_transform,
            ecoli_genome=ecoli_genome,
            **kwargs,
        )


ARMS: tuple[type[_SrmSchmidt2016Dataset], ...] = (
    ProteomeSrmSet1Schmidt2016Dataset,
    ProteomeSrmSet2Schmidt2016Dataset,
)


# --------------------------------------------------------------------------- #
# L0-L4 verification
# --------------------------------------------------------------------------- #
def verify_build(
    dataset_root: str,
    data_root: str | None = None,
    *,
    arm: type[_SrmSchmidt2016Dataset],
    genome: EcoliK12Genome | None = None,
) -> VerificationReport:
    """Run the protein L0-L4 gate over one arm's built tree and write the report.

    The shared ``verify_protein_dataset`` supplies L0 to L3; the three SUPPLEMENTARY
    rows and the host-aware L4 containment are ``schmidt2016``'s own, reused here
    because this arm's records have the same shape as that block's -- a wild-type
    genotype, one environment per record, and BW25113 locus keys.
    """
    from torchcell.verification.protein import verify_protein_dataset
    from torchcell.verification.runners import load_records

    records = load_records(dataset_root)
    if genome is None:
        genome = bacterial_genome("ecoli", SCHMIDT_REFERENCE_STRAIN, data_root)
    report = verify_protein_dataset(
        records,
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=Provenance(
            source_uri=f"$DATA_ROOT/{sm.RAW_DIR_REL}/{sm.SI2_MIRROR_RELPATH}",
            citation_key=sm.CITATION_KEY,
            sha256=sm.SI2_SHA256,
            method=(
                f"{arm.SHEET} 'Protein Abundance (copies/cell)': one "
                "BacterialProteinAbundanceExperiment per loaded BW25113 growth "
                f"condition, SE = the released standard deviation / sqrt("
                f"{arm.N_REPLICATES})"
            ),
            page=f"si2.xlsx sheets '{SHEET_S1}' and '{arm.SHEET}'",
            retrieved=sm.SI2_RETRIEVED_AT,
        ),
        expected_count=arm.EXPECTED_RECORDS,
    )
    report.add(sm.environment_uniqueness_rule(records))
    report.add(sm.assembly_pin_rule(records))
    report.add(sm.gene_containment_rule(records, set(genome.genbank.loci)))
    out = osp.join(dataset_root, "preprocess", "verification_report.json")
    os.makedirs(osp.dirname(out), exist_ok=True)
    with open(out, "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def arm_root(arm: type[_SrmSchmidt2016Dataset]) -> str:
    """The arm's own default build root, relative to ``DATA_ROOT``."""
    return str(inspect.signature(arm.__init__).parameters["root"].default)


def main(argv: list[str] | None = None) -> int:
    """CLI: ``build`` or ``verify`` one SRM arm's dev-tree LMDB."""
    from dotenv import load_dotenv

    by_name = {cls.__name__: cls for cls in ARMS}
    parser = argparse.ArgumentParser(
        prog="python -m torchcell.datasets.ecoli.schmidt2016_srm"
    )
    parser.add_argument("command", choices=("build", "verify"))
    parser.add_argument("--arm", choices=sorted(by_name), required=True)
    args = parser.parse_args(argv)

    load_dotenv()
    data_root = os.environ["DATA_ROOT"]
    arm = by_name[args.arm]
    root = osp.join(data_root, arm_root(arm))
    if args.command == "build":
        dataset = arm(root=root)
        print(f"len = {len(dataset)}")
        return 0
    report = verify_build(root, data_root, arm=arm)
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
