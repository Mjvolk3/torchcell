# experiments/036-dataset-fixes-before-kg-build/scripts/potel2018_release_inventory.py
# [[experiments.036-dataset-fixes-before-kg-build.scripts.potel2018_release_inventory]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/036-dataset-fixes-before-kg-build/scripts/potel2018_release_inventory

r"""What Potel 2018 released, where its bytes now live, and what the schema can hold.

Potel CM, Lin MH, Heck AJR, Lemeer S (2018) "Widespread bacterial protein histidine
phosphorylation revealed by mass spectrometry-based proteomics", Nature Methods
15:187-190, doi:10.1038/nmeth.4580, PMID 29377012; ProteomeXchange PXD008369, which is
hosted by jPOST as JPST000369 (PRIDE answers "not in the database"). Candidate row
"Potel 2018 phosphoproteome" of the bacterial table.

The paper is in neither mirror and in no PMC or Europe PMC full-text set. nature.com
serves a subscription landing page, so the published Online Methods are not scriptable.
Two scriptable sources carry what is needed, and both are deposited in the RAW mirror
(``$DATA_ROOT/torchcell-raw/potelWidespreadBacterialProtein2018/``) on the Mohiuddin
2022 / Li 2014 pattern, never in Zotero:

- the seven Springer ESM files (``retrieve.springer_esm``): the Supplementary Text and
  Figures PDF, the Life Sciences Reporting Summary, Supplementary Tables 1, 2, 3 and 5,
  and the Fig. 1 source data. Table 2 is the per-record artifact (``data/``);
- the first author's open-access dissertation (Utrecht University repository, handle
  1874/370166, "Publisher version", "Open Access (free)"), whose Chapter V reprints this
  paper with its Materials and Methods. The repository wraps every download in a
  generated cover, so two downloads differ in 134 bytes while their ``pdftotext
  -layout`` renderings are identical (measured 2026-10-10). The deposited PDF is
  therefore canonical as held, and the TEXT rendering is what this script pins and what
  every quote is read from. The repository's other copy of the paper (handle
  1874/362159) is the October 2017 submitted version and carries no Methods.

``pdftotext -layout`` renderings of the two SI PDFs and of the dissertation, and a
deterministic text rendering of the Description sheets of Tables 2 and 5, are deposited
beside their sources with a ``ProcessingRecord`` pinning the input sha256, so every
quote below is auditable as text.

Measurements, all off the deposited bytes (``--network`` adds the jPOST FTP listing):

1. **Table 2**: sites, residues, localization classes, proteins, the 12 per-sample
   intensity columns, zero cells, and whether the total ``Intensity`` column is the sum
   of the twelve.
2. **Table 5 against Table 2**: whether its log2 values are log2 of Table 2's
   intensities, and how many are Perseus imputations.
3. **Host, medium, replicate structure**: from the dissertation's Chapter V Methods,
   against the SI figure legends and the reporting summary, which disagree.
4. **Identifier route**: Table 2's gene names through ``reconcile_locus_tags`` on the
   MG1655 assembly (the search database was the UniProt K-12 reference); the host's own
   W3110 assembly is not readable by any genome class yet (#910).
5. **Schema fit**: which ``Phenotype`` subclasses carry a site or residue; none does.
6. **Duplication by content** against the six served proteome dev stores (Schmidt
   2016 x3, Ishii 2007, Brunk 2016, Mori 2021): key namespace and site-shaped keys,
   read-only, with each store's build-manifest dirty flag recorded.

Run from the repo root::

    python experiments/036-dataset-fixes-before-kg-build/scripts/potel2018_release_inventory.py --deposit --network
"""

import argparse
import hashlib
import json
import os
import os.path as osp
import pickle
import re
import shutil
import subprocess
import tempfile
import urllib.request
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal

import lmdb
import numpy as np
import openpyxl
import pandas as pd
from dotenv import load_dotenv
from pydantic import BaseModel, ConfigDict

load_dotenv()

from torchcell.data.experiment_dataset import resolve_interned  # noqa: E402
from torchcell.datamodels import schema  # noqa: E402
from torchcell.datasets.bacteria_common import (  # noqa: E402
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.literature.manifest import (  # noqa: E402
    ROLE_PAPER_PDF,
    ROLE_PAPER_TEXT,
    ROLE_RAW_DATA,
    ROLE_SI_DATA,
    ROLE_SI_PDF,
    ROLE_SI_TEXT,
    ArtifactRecord,
    Manifest,
    ProcessingRecord,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.literature.retrieve import RETRIEVERS  # noqa: E402
from torchcell.provenance.build_manifest import BuildManifest  # noqa: E402

CITATION_KEY = "potelWidespreadBacterialProtein2018"
PAPER_DOI = "10.1038/nmeth.4580"
PAPER_TITLE = (
    "Widespread bacterial protein histidine phosphorylation revealed by mass "
    "spectrometry-based proteomics"
)
PMID = "29377012"
PROTEOMEXCHANGE = "PXD008369"
JPOST = "JPST000369"
JPOST_FTP = f"ftp://ftp.jpostdb.org/{JPOST}/"
RETRIEVED_AT = "2026-10-10"

ESM_URL = (
    "https://static-content.springer.com/esm/art%3A10.1038%2Fnmeth.4580/"
    "MediaObjects/{name}"
)
#: The dissertation PDF, served by the Utrecht DSpace REST API (handle 1874/370166).
THESIS_URL = (
    "https://dspace.library.uu.nl/server/api/core/bitstreams/"
    "2364f282-67ea-482a-a1b4-989c61986639/content"
)
THESIS_PDF = "Potel.pdf"
THESIS_TEXT = "Potel.pdftotext.txt"
#: sha256 of ``pdftotext -layout`` (poppler 21.01.0) of the dissertation; identical for
#: two downloads whose PDF bytes differ (the repository's generated cover).
THESIS_TEXT_SHA256 = "05a506579dcbc869d056af2bc81ecabf4f82debb4a84dda73b7b35075fdbd81f"

TABLE_2 = "41592_2018_BFnmeth4580_MOESM4_ESM.xlsx"
TABLE_5 = "41592_2018_BFnmeth4580_MOESM6_ESM.xlsx"
SI_PDF = "41592_2018_BFnmeth4580_MOESM1_ESM.pdf"
REPORTING_SUMMARY = "41592_2018_BFnmeth4580_MOESM2_ESM.pdf"
SOURCE_DATA_FIG1 = "41592_2018_BFnmeth4580_MOESM7_ESM.xlsx"
TABLE_2_SHEET = "Supplementary Table 2."
TABLE_5_SHEET = "Supplementary Table 5."

RESULTS = osp.join(
    osp.dirname(osp.dirname(osp.abspath(__file__))),
    "results",
    "potel2018_release_inventory.json",
)

#: The six served proteome dev stores the duplication check reads (read-only).
PROTEOME_STORES: tuple[str, ...] = (
    "proteome_schmidt2016",
    "proteome_srm_set1_schmidt2016",
    "proteome_srm_set2_schmidt2016",
    "proteome_ishii2007",
    "proteome_brunk2016",
    "proteome_mori2021",
)

#: A key that names a residue and a position, the shape a phosphosite key would take.
SITE_KEY = re.compile(r"[STYH]\d+$|phospho", re.IGNORECASE)
#: Field names that would carry a site or a modified residue on a phenotype.
SITE_FIELD = re.compile(r"site|residue|phospho|modif", re.IGNORECASE)


class RawFile(BaseModel):
    """One deposited file: pinned bytes, role, and the retrieval that reproduces it."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    name: str
    relpath: str
    role: str
    sha256: str | None  # None: the source is not byte-stable, its text is pinned
    source: Literal["springer_esm", "uu_dspace"]
    description: str

    @property
    def source_url(self) -> str:
        """The URL the retriever GETs."""
        if self.source == "springer_esm":
            return ESM_URL.format(name=self.name)
        return THESIS_URL

    @property
    def retriever(self) -> str:
        """Registry key of the retriever in ``torchcell.literature.retrieve``."""
        if self.source == "springer_esm":
            return "torchcell.literature.retrieve.springer_esm"
        return "torchcell.literature.retrieve.direct_url"

    def fetch(self) -> bytes:
        """Run the recorded retriever on the recorded URL."""
        return RETRIEVERS[self.retriever](url=self.source_url)

    def retrieval(self, sha256: str) -> RetrievalRecord:
        """The re-runnable retrieval of these bytes, anchored to ``sha256``."""
        method = (
            RetrievalMethod.springer_esm
            if self.source == "springer_esm"
            else RetrievalMethod.direct_url
        )
        return RetrievalRecord(
            method=method,
            source_url=self.source_url,
            retriever=self.retriever,
            params={"url": self.source_url},
            sha256=sha256,
            retrieved_at=RETRIEVED_AT,
        )


RAW_FILES: tuple[RawFile, ...] = (
    RawFile(
        name=TABLE_2,
        relpath=f"data/{TABLE_2}",
        role=ROLE_RAW_DATA,
        sha256="f69ca61a0f0dba3987b26bcdf4bd9b0c8fcf8744495233bad678a23cf5682481",
        source="springer_esm",
        description="Supplementary Table 2: 2,129 phosphosites (STYH), score, "
        "localization probability and MaxQuant precursor intensity in 12 samples",
    ),
    RawFile(
        name=SI_PDF,
        relpath=f"si/{SI_PDF}",
        role=ROLE_SI_PDF,
        sha256="f4cc85a6bf9aca3084641dd0340b006840484c00deef5a87aca3c868583edb3d",
        source="springer_esm",
        description="Supplementary Text and Figures: Supplementary Figures 1-12 and "
        "Supplementary Tables 4 and 6",
    ),
    RawFile(
        name=REPORTING_SUMMARY,
        relpath=f"si/{REPORTING_SUMMARY}",
        role=ROLE_SI_PDF,
        sha256="7ab7cabe3d1eb8479f8a3bcd98b3ac263c1429141ba02fda3f43d8d6b0401981",
        source="springer_esm",
        description="Life Sciences Reporting Summary",
    ),
    RawFile(
        name="41592_2018_BFnmeth4580_MOESM3_ESM.xlsx",
        relpath="si/41592_2018_BFnmeth4580_MOESM3_ESM.xlsx",
        role=ROLE_SI_DATA,
        sha256="e6e4ef28ce90357fd497c3616269721557c44f0e3c76ae9c11cf506f3997c2ce",
        source="springer_esm",
        description="Supplementary Table 1: phosphohistidine sites from chemically "
        "phosphorylated protein mixtures (standards, not E. coli)",
    ),
    RawFile(
        name="41592_2018_BFnmeth4580_MOESM5_ESM.xlsx",
        relpath="si/41592_2018_BFnmeth4580_MOESM5_ESM.xlsx",
        role=ROLE_SI_DATA,
        sha256="c2ebaef8910fb2ec5109557943ce36079069c9894b072435c32addac44ab95ef",
        source="springer_esm",
        description="Supplementary Table 3: phosphohistidine sites, antibody-based "
        "(Oslund et al.) against this study",
    ),
    RawFile(
        name=TABLE_5,
        relpath=f"si/{TABLE_5}",
        role=ROLE_SI_DATA,
        sha256="31a986e2b9a1cc582b1265a0bc124693f4bcdbf941abb3398f27217d3f1d58bb",
        source="springer_esm",
        description="Supplementary Table 5: glycerol exponential vs stationary, log2 "
        "intensities (with imputation) and a t-test significance flag",
    ),
    RawFile(
        name=SOURCE_DATA_FIG1,
        relpath=f"si/{SOURCE_DATA_FIG1}",
        role=ROLE_SI_DATA,
        sha256="61ff4b381c5f53f24ee1609ad80c9fd6c42e7c738579bf781dd94e1b83e89589",
        source="springer_esm",
        description="Source data to Fig. 1: phosphopeptides per run by protocol",
    ),
    RawFile(
        name=THESIS_PDF,
        relpath=f"paper/thesis/{THESIS_PDF}",
        role=ROLE_PAPER_PDF,
        sha256=None,
        source="uu_dspace",
        description="C.M. Potel, 'Exploring new avenues in phosphoproteomics in pursuit "
        "of the elusive histidine phosphorylation', Utrecht University 2018, ISBN "
        "978-94-93019-74-4; Chapter V reprints this paper with its Methods",
    ),
)

#: Text renderings of PDFs: (source PDF relpath, rendered text relpath, role).
PDF_TEXTS: tuple[tuple[str, str, str], ...] = (
    (f"si/{SI_PDF}", f"si/{SI_PDF.removesuffix('.pdf')}.pdftotext.txt", ROLE_SI_TEXT),
    (
        f"si/{REPORTING_SUMMARY}",
        f"si/{REPORTING_SUMMARY.removesuffix('.pdf')}.pdftotext.txt",
        ROLE_SI_TEXT,
    ),
    (f"paper/thesis/{THESIS_PDF}", f"paper/thesis/{THESIS_TEXT}", ROLE_PAPER_TEXT),
)

#: Workbook legend renderings: (workbook relpath, legend text relpath).
LEGEND_TEXTS: tuple[tuple[str, str], ...] = (
    (f"data/{TABLE_2}", f"data/{TABLE_2.removesuffix('.xlsx')}.legends.txt"),
    (f"si/{TABLE_5}", f"si/{TABLE_5.removesuffix('.xlsx')}.legends.txt"),
)

#: Dissertation quotes that must have an occurrence in Chapter V (this paper's reprint);
#: the growth paragraph recurs verbatim in Chapter VI, so presence alone does not say
#: which chapter a quote was read from.
CHAPTER_V_QUOTES: frozenset[str] = frozenset(
    {
        "strain_and_medium",
        "carbon_sources",
        "temperature_and_harvest",
        "technical_triplicates_thesis",
        "search_database",
        "imputation",
        "thesis_t_test",
        "thesis_fdr_threshold",
        "phis_counts",
        "class_i_phis",
    }
)
#: The dissertation's running page header, which names the chapter a page belongs to.
CHAPTER_HEADER = re.compile(r"CHAPTER (VI|IV|V|III|II|I)\b")

SI_TEXT_REL = PDF_TEXTS[0][1]
RS_TEXT_REL = PDF_TEXTS[1][1]
THESIS_TEXT_REL = PDF_TEXTS[2][1]
T2_LEGEND_REL = LEGEND_TEXTS[0][1]
T5_LEGEND_REL = LEGEND_TEXTS[1][1]

#: Quotes that ground the inventory, each checked to occur verbatim (whitespace
#: collapsed) in the named text file.
QUOTES: dict[str, tuple[str, str]] = {
    "thesis_chapter_v_is_this_paper": (
        THESIS_TEXT_REL,
        "Chapter V: Widespread bacterial protein histidine phosphorylation revealed "
        "by mass spectrometry-based proteomics",
    ),
    "thesis_publication_entry": (
        THESIS_TEXT_REL,
        "Potel, C. M., Lin, M.-H., Heck, A. J. R. & Lemeer, S. Widespread bacterial "
        "protein histidine phosphorylation revealed by mass spectrometry-based "
        "proteomics. Nature Methods 15, 187 (2018).",
    ),
    "strain_and_medium": (
        THESIS_TEXT_REL,
        "E. coli strain W3110 was grown in M9 minimal medium, consisting of M9 salts",
    ),
    "carbon_sources": (
        THESIS_TEXT_REL,
        "supplemented with either additional 0.5% (w/v) glucose or 0.5% glycerol",
    ),
    "temperature_and_harvest": (
        THESIS_TEXT_REL,
        "with vigorous shaking at 37°C. Cells were collected by centrifugation at "
        "mid-log phase (OD600= 0.6/0.5 for glucose / glycerol-fed cells) and stationary "
        "phase (OD600= 1.2/1.0 for glucose/ glycerol-fed cells) (Figure S12)",
    ),
    "technical_triplicates_thesis": (
        THESIS_TEXT_REL,
        "Fe3+-IMAC was performed in technical triplicates for each biological groups",
    ),
    "search_database": (
        THESIS_TEXT_REL,
        "Raw data files were processed using the MaxQuant software (ver. 1.5.5.1) with "
        "the following search parameters: trypsin digestion with a maximum of 3 missed "
        "cleavages, reviewed E. coli K12 database (Uniprot, March 2016)",
    ),
    "imputation": (
        THESIS_TEXT_REL,
        "missing value imputation was performed with default parameters (width is 0.3 "
        "and down shift is 1.8 for each replicate)",
    ),
    "thesis_t_test": (
        THESIS_TEXT_REL,
        "The log2 scale of phosphosite intensity from technical Fe3+-IMAC triplicates "
        "were subjected into a two tailed t-test with a false discovery rate (FDR)",
    ),
    "thesis_fdr_threshold": (
        THESIS_TEXT_REL,
        "of less than 0.05 and a S0 (the minimal fold change) value of",
    ),
    "phis_counts": (
        THESIS_TEXT_REL,
        "With a total number of 246 pHis sites on 173 histidine phosphoproteins",
    ),
    "class_i_phis": (
        THESIS_TEXT_REL,
        "of the 246 pHis sites, 135 phosphosites could be retained as high- confidence "
        "Class I pHis sites (localization confidence ≥ 75%) (Table S2)",
    ),
    "biological_replicates_si_fig3": (
        SI_TEXT_REL,
        "The same UV profiles were observed for the n=12 biological replicates.",
    ),
    "biological_replicates_si_fig6": (
        SI_TEXT_REL,
        "The different intensities profiles derive from the totality of the n=12 "
        "biological replicates.",
    ),
    "growth_phases_si_fig12": (
        SI_TEXT_REL,
        "E. coli cells were collected at either mid-log phase or stationary phase "
        "according to OD",
    ),
    "technical_triplicates_after_digestion": (
        RS_TEXT_REL,
        "we made technical triplicates (after digestion, before enrichment) of each",
    ),
    "technical_triplicates_reporting_summary": (
        RS_TEXT_REL,
        "Results were reproduced under 4 biological conditions, each consisting of "
        "technical",
    ),
    "t2_intensity_column": (
        T2_LEGEND_REL,
        "17. Intensity glucose_exp1, summed up eXtracted Ion Current (XIC) of all "
        "isotopic clusters associated with the identified phosphopeptdie sequence in "
        "the 1st replicate of glucose-fed exponential E. coli cells.",
    ),
    "t2_localization_column": (
        T2_LEGEND_REL,
        "5. Localization prob, the probability of phosphosite assignment.",
    ),
    "t5_test": (
        T5_LEGEND_REL,
        "statistical analysis of detected changes between exponential and stationary "
        "phases of glycerol-fed E. coli. (t-test, with FDR ≤ 0.02 and S0=0.1)",
    ),
}


def _sha256(path: str | Path) -> str:
    """Streaming sha256 of a file."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def raw_mirror_dir() -> Path:
    """``$DATA_ROOT/torchcell-raw/potelWidespreadBacterialProtein2018``."""
    return Path(os.environ["DATA_ROOT"]) / "torchcell-raw" / CITATION_KEY


def _pdftotext_version() -> str:
    """The poppler ``pdftotext`` version the texts were rendered with."""
    out = subprocess.run(
        ["pdftotext", "-v"], capture_output=True, text=True, check=True
    )
    return (out.stderr or out.stdout).splitlines()[0].strip()


def render_pdf_text(pdf: Path, dest: Path) -> None:
    """``pdftotext -layout`` of a PDF, the auditable text a quote is read from."""
    subprocess.run(["pdftotext", "-layout", str(pdf), str(dest)], check=True)


def extract_sheet_legends(workbook: Path) -> str:
    """Every non-empty cell of the workbook's ``Description`` sheet, one per line.

    The workbook is binary, so a quote in it cannot be audited; this rendering can, and
    its :class:`ProcessingRecord` names this function and pins the workbook's sha256.
    """
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    lines = [f"# sheet: Description of {workbook.name}"]
    for row in book["Description"].iter_rows(values_only=True):
        lines.extend(str(cell) for cell in row if cell is not None)
    book.close()
    return "\n".join(lines) + "\n"


def _fetch(raw: RawFile, staging: Path) -> Path:
    """Run the recorded retriever and verify what arrives; return the staged path."""
    data = raw.fetch()
    staged = staging / raw.name
    staged.write_bytes(data)
    if raw.sha256 is not None:
        got = hashlib.sha256(data).hexdigest()
        if got != raw.sha256:
            raise RuntimeError(f"{raw.source_url}: sha256 {got}, expected {raw.sha256}")
        return staged
    text = staging / f"{raw.name}.txt"
    render_pdf_text(staged, text)
    got = _sha256(text)
    if got != THESIS_TEXT_SHA256:
        raise RuntimeError(
            f"{raw.source_url}: text rendering sha256 {got}, expected "
            f"{THESIS_TEXT_SHA256}"
        )
    return staged


def deposit() -> Path:
    """Run every recorded retriever, verify the bytes, and write mirror + manifest.

    Idempotent: a file already in the mirror is kept when its sha256 is the pinned one
    (byte-stable sources) or when its text rendering is (the dissertation); anything
    else raises instead of being overwritten.
    """
    root = raw_mirror_dir()
    records: list[ArtifactRecord] = []
    version = _pdftotext_version()
    with tempfile.TemporaryDirectory() as tmp:
        for raw in RAW_FILES:
            dest = root / raw.relpath
            dest.parent.mkdir(parents=True, exist_ok=True)
            if not dest.exists():
                shutil.move(str(_fetch(raw, Path(tmp))), dest)
            held = _sha256(dest)
            if raw.sha256 is not None and held != raw.sha256:
                raise RuntimeError(f"{dest} exists with another sha256; refusing")
            records.append(
                ArtifactRecord(
                    path=raw.relpath,
                    role=raw.role,
                    bytes=dest.stat().st_size,
                    sha256=held,
                    source=raw.source_url,
                    original_filename=raw.name,
                    retrieval=raw.retrieval(held),
                )
            )
    for pdf_rel, text_rel, role in PDF_TEXTS:
        pdf, text = root / pdf_rel, root / text_rel
        render_pdf_text(pdf, text)
        if text_rel == THESIS_TEXT_REL and _sha256(text) != THESIS_TEXT_SHA256:
            raise RuntimeError(f"{text}: rendering differs from the pinned text")
        records.append(
            ArtifactRecord(
                path=text_rel,
                role=role,
                bytes=text.stat().st_size,
                sha256=_sha256(text),
                source=f"derived from {pdf_rel}",
                original_filename=text.name,
                processing=ProcessingRecord(
                    processor="experiments/036-dataset-fixes-before-kg-build/scripts/"
                    "potel2018_release_inventory.py:render_pdf_text",
                    tool="pdftotext",
                    version=version,
                    params={"flags": ["-layout"]},
                    input_sha256=[_sha256(pdf)],
                ),
            )
        )
    for book_rel, text_rel in LEGEND_TEXTS:
        book, text = root / book_rel, root / text_rel
        text.write_text(extract_sheet_legends(book), encoding="utf-8")
        records.append(
            ArtifactRecord(
                path=text_rel,
                role=ROLE_SI_TEXT,
                bytes=text.stat().st_size,
                sha256=_sha256(text),
                source=f"derived from {book_rel}",
                original_filename=text.name,
                processing=ProcessingRecord(
                    processor="experiments/036-dataset-fixes-before-kg-build/scripts/"
                    "potel2018_release_inventory.py:extract_sheet_legends",
                    tool="openpyxl",
                    version=str(openpyxl.__version__),
                    params={"sheet": "Description", "cells": "non-empty, row-major"},
                    input_sha256=[_sha256(book)],
                ),
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title=PAPER_TITLE,
        files=records,
        si_data_sources=[raw.source_url for raw in RAW_FILES],
        si_expected=[
            "Supplementary Text and Figures",
            "Life Sciences Reporting Summary",
            "Supplementary Table 1",
            "Supplementary Table 2",
            "Supplementary Table 3",
            "Supplementary Table 5",
            "Source data to Fig. 1",
        ],
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def verify_mirror(root: Path) -> dict[str, str]:
    """Every manifest file re-hashed against its record; raise on any mismatch."""
    manifest = Manifest.model_validate_json((root / "manifest.json").read_text())
    out: dict[str, str] = {}
    for record in manifest.files:
        got = _sha256(root / record.path)
        if got != record.sha256:
            raise RuntimeError(f"{record.path}: sha256 {got} != {record.sha256}")
        out[record.path] = got
    if out[THESIS_TEXT_REL] != THESIS_TEXT_SHA256:
        raise RuntimeError("the dissertation text is not the pinned rendering")
    return out


def chapters_of(flat: str, quote: str) -> list[str]:
    """The dissertation chapter of each occurrence of ``quote`` ("front" before any)."""
    headers = [(m.start(), m.group(1)) for m in CHAPTER_HEADER.finditer(flat)]
    out = []
    for match in re.finditer(re.escape(quote), flat):
        before = [name for start, name in headers if start < match.start()]
        out.append(before[-1] if before else "front")
    return out


def check_quotes(
    root: Path, hashes: dict[str, str]
) -> dict[str, dict[str, str | int | list[str]]]:
    """Each grounding quote, located verbatim in its file (whitespace collapsed)."""
    out: dict[str, dict[str, str | int | list[str]]] = {}
    flat: dict[str, str] = {}
    for name, (relpath, quote) in QUOTES.items():
        if relpath not in flat:
            text = (root / relpath).read_text(encoding="utf-8")
            flat[relpath] = re.sub(r"\s+", " ", text)
        occurrences = flat[relpath].count(quote)
        if occurrences == 0:
            raise AssertionError(f"quote {name!r} not found in {relpath}")
        out[name] = {
            "file": relpath,
            "sha256": hashes[relpath],
            "quote": quote,
            "occurrences": occurrences,
        }
        if relpath == THESIS_TEXT_REL:
            chapters = chapters_of(flat[relpath], quote)
            if name in CHAPTER_V_QUOTES and "V" not in chapters:
                raise AssertionError(f"quote {name!r} has no Chapter V occurrence")
            out[name]["chapters"] = chapters
    return out


def token_counts(root: Path) -> dict[str, dict[str, int]]:
    """How often each deposited text names a K-12 strain (the row named one by recall)."""
    tokens = ("W3110", "MG1655", "BW25113")
    out: dict[str, dict[str, int]] = {}
    for _, text_rel, _ in PDF_TEXTS:
        text = (root / text_rel).read_text(encoding="utf-8")
        out[text_rel] = {t: len(re.findall(rf"\b{t}\b", text)) for t in tokens}
    return out


def read_table_2(root: Path) -> pd.DataFrame:
    """Supplementary Table 2 as strings, exactly as openpyxl reads the cells."""
    return pd.read_excel(root / f"data/{TABLE_2}", sheet_name=TABLE_2_SHEET, dtype=str)


def sample_columns(table: pd.DataFrame) -> list[str]:
    """The twelve per-sample intensity columns (``Intensity <carbon>_<phase><rep>``)."""
    return [c for c in table.columns if re.fullmatch(r"Intensity \w+_\w+\d", c)]


def table_2_inventory(table: pd.DataFrame) -> dict[str, Any]:
    """Sites, residues, localization, proteins, and the per-sample intensity grid."""
    samples = sample_columns(table)
    grid = table[samples].astype(float)
    total = table["Intensity"].astype(float)
    summed = grid.sum(axis=1)
    rel = (total - summed) / total.where(total > 0)
    loc = table["Localization prob"].astype(float)
    residue = table["Amino acid"]
    key = table["Acc. No."] + "|" + residue + "|" + table["Positions"]
    nonzero = grid.where(grid > 0).stack()
    multi_protein = table["Acc. No."].str.contains(";")
    return {
        "rows": len(table),
        "columns": list(table.columns),
        "sample_columns": samples,
        "duplicate_accession_residue_position": int(key.duplicated().sum()),
        "residues": dict(Counter(residue)),
        "localization_prob_at_least_0_75": dict(Counter(residue[loc >= 0.75])),
        "localization_prob_min": float(loc.min()),
        "distinct_accession_strings": int(table["Acc. No."].nunique()),
        "phis_sites": int((residue == "H").sum()),
        "phis_distinct_accession_strings": int(
            table.loc[residue == "H", "Acc. No."].nunique()
        ),
        "rows_naming_more_than_one_accession": int(multi_protein.sum()),
        "rows_naming_more_than_one_gene": int(
            table["Gene names"].str.contains(";").sum()
        ),
        "rows_naming_more_than_one_position": int(
            table["Positions"].str.contains(";").sum()
        ),
        "cells": int(grid.size),
        "zero_cells": int((grid == 0).sum().sum()),
        "negative_or_non_numeric_cells": int(
            (grid < 0).sum().sum() + grid.isna().sum().sum()
        ),
        "nonzero_per_sample": {c: int((grid[c] > 0).sum()) for c in samples},
        "nonzero_min": float(nonzero.min()),
        "nonzero_median": float(nonzero.median()),
        "nonzero_max": float(nonzero.max()),
        "sites_zero_in_all_twelve": int((summed == 0).sum()),
        "sites_total_intensity_zero": int((total == 0).sum()),
        "total_intensity_vs_sum_of_twelve": {
            "equal_within_1e-6_relative": int((rel.abs() <= 1e-6).sum()),
            "total_greater": int((rel > 1e-6).sum()),
            "total_smaller": int((rel < -1e-6).sum()),
        },
        "diagnostic_peak_plus": int((table["Diagnostic peak"] == "+").sum()),
    }


def _multi_position_row_kind(t2: pd.DataFrame, key: str) -> str:
    """Which multi-position Table 2 row lists a site Table 5 gives one position for."""
    accession, residue, position = key.split("|")
    multi = t2[t2["Positions"].str.contains(";") & (t2["Amino acid"] == residue)]
    for _, row in multi.iterrows():
        if accession in row["Acc. No."].split(";") and position in str(
            row["Positions"]
        ).split(";"):
            if ";" in row["Acc. No."]:
                return "multi_accession_row"
            return "one_accession_row_naming_several_genes"
    return "no_table_2_row"


def table_5_relation(root: Path, t2: pd.DataFrame) -> dict[str, Any]:
    """Table 5 against Table 2: log2 of the same intensity, or an imputation."""
    t5 = pd.read_excel(root / f"si/{TABLE_5}", sheet_name=TABLE_5_SHEET, dtype=str)
    k2 = t2["Acc. No."] + "|" + t2["Amino acid"] + "|" + t2["Positions"]
    k5 = t5["Acc. No."] + "|" + t5["Amino acid"] + "|" + t5["Position"]
    joined = t5.assign(k=k5).merge(t2.assign(k=k2), on="k", suffixes=("_5", "_2"))
    per_rep: dict[str, dict[str, int]] = {}
    max_diff = 0.0
    for rep in ("exp1", "exp2", "exp3", "stat1", "stat2", "stat3"):
        log2 = joined[f"log2 (Intensity glycerol_{rep})"].astype(float)
        raw = joined[f"Intensity glycerol_{rep}"].astype(float)
        measured = raw > 0
        diff = (log2[measured] - np.log2(raw[measured])).abs()
        max_diff = max(max_diff, float(diff.max()))
        per_rep[rep] = {
            "table_2_nonzero": int(measured.sum()),
            "equal_log2_within_1e-3": int((diff < 1e-3).sum()),
            "table_2_zero_so_table_5_value_is_imputed": int((~measured).sum()),
        }
    unmatched = sorted(set(k5) - set(k2))
    return {
        "rows": len(t5),
        "significant_plus": int((t5["Significance"] == "+").sum()),
        "keys_found_in_table_2": int(k5.isin(set(k2)).sum()),
        "keys_not_found_in_table_2": unmatched,
        "unmatched_keys_by_table_2_row": dict(
            Counter(_multi_position_row_kind(t2, key) for key in unmatched)
        ),
        "per_replicate": per_rep,
        "max_abs_diff_log2": max_diff,
    }


def resolve_genes(
    t2: pd.DataFrame, strain: Literal["MG1655", "BW25113"]
) -> dict[str, Any]:
    """Single-gene rows' names through ``reconcile_locus_tags`` on ``strain``."""
    genome = bacterial_genome("ecoli", strain)
    names = t2.loc[~t2["Gene names"].str.contains(";"), "Gene names"]
    stored, report = reconcile_locus_tags(
        genome, pd.Series(names.unique()), label=f"potel2018_{strain}"
    )
    offered = list(names.unique())
    tag = re.compile(r"^b\d{4}$" if strain == "MG1655" else rf"^{strain}_\d+$")
    resolved = {n: str(t) for n, t in zip(offered, stored, strict=True) if tag.match(t)}
    return {
        "assembly_set": report.assembly_set,
        "names_offered": len(offered),
        "resolved_to_one_locus": len(resolved),
        "resolved_fraction": round(len(resolved) / len(offered), 4),
        "unresolved": sorted(set(offered) - set(resolved)),
        "status_histogram": {str(k): v for k, v in report.status_histogram.items()},
        "loci": sorted(set(resolved.values())),
    }


def schema_fit() -> dict[str, Any]:
    """Phenotype subclasses whose fields name a site, residue or modification."""
    classes = sorted(
        name
        for name, obj in vars(schema).items()
        if isinstance(obj, type)
        and issubclass(obj, schema.Phenotype)
        and obj is not schema.Phenotype
    )
    with_site = {
        name: [f for f in getattr(schema, name).model_fields if SITE_FIELD.search(f)]
        for name in classes
    }
    return {
        "phenotype_classes": len(classes),
        "classes_with_a_site_or_residue_field": {
            k: v for k, v in with_site.items() if v
        },
    }


def store_key_shapes(store: str, mg1655: set[str], bw25113: set[str]) -> dict[str, Any]:
    """One proteome dev store, read-only: label keys, site-shaped keys, overlap."""
    root = Path(os.environ["DATA_ROOT"]) / "data" / "torchcell" / store
    build = BuildManifest.model_validate_json(
        (root / "preprocess" / "build_manifest.json").read_text()
    )
    interned: dict[str, Any] = {}
    interned_dir = root / "processed" / "interned"
    if interned_dir.is_dir():
        ienv = lmdb.open(str(interned_dir), readonly=True, lock=False)
        with ienv.begin() as txn:
            for key, value in txn.cursor():
                interned[key.decode()] = pickle.loads(value)
        ienv.close()
    env = lmdb.open(
        str(root / "processed" / "lmdb"), readonly=True, lock=False, readahead=False
    )
    types: Counter[str] = Counter()
    keys: set[str] = set()
    with env.begin() as txn:
        for _, raw in txn.cursor():
            record = resolve_interned(pickle.loads(raw), interned)
            experiment = record["experiment"]
            types[str(experiment["experiment_type"])] += 1
            phenotype = experiment["phenotype"]
            keys |= set(phenotype[phenotype["label_name"]])
    env.close()
    return {
        "records": sum(types.values()),
        "experiment_types": dict(types),
        "build_commit": build.torchcell_commit,
        "build_dirty": build.torchcell_dirty,
        "distinct_label_keys": len(keys),
        "site_shaped_keys": sorted(k for k in keys if SITE_KEY.search(k)),
        "shared_with_potel_mg1655_loci": len(keys & mg1655),
        "shared_with_potel_bw25113_loci": len(keys & bw25113),
    }


def workbook_rows(root: Path) -> dict[str, dict[str, dict[str, int]]]:
    """Per deposited workbook and sheet: rows the sheet declares, rows with content.

    ``candidate-gate``'s G3 counts declared rows; Supplementary Table 1's sheet declares
    over a million, almost all empty, so the content count is recorded beside it.
    """
    out: dict[str, dict[str, dict[str, int]]] = {}
    for raw in RAW_FILES:
        if not raw.name.endswith(".xlsx"):
            continue
        book = openpyxl.load_workbook(
            root / raw.relpath, read_only=True, data_only=True
        )
        sheets: dict[str, dict[str, int]] = {}
        for sheet in book.worksheets:
            declared = with_content = 0
            for row in sheet.iter_rows(values_only=True):
                declared += 1
                if any(c is not None and str(c).strip() for c in row):
                    with_content += 1
            sheets[sheet.title] = {"declared": declared, "with_content": with_content}
        book.close()
        out[raw.relpath] = sheets
    return out


def source_data_fig1(root: Path) -> list[list[str]]:
    """The Fig. 1 source-data sheet, every non-empty row as strings."""
    book = openpyxl.load_workbook(
        root / f"si/{SOURCE_DATA_FIG1}", read_only=True, data_only=True
    )
    rows = [
        [str(c) for c in row if c is not None]
        for row in book["Sheet1"].iter_rows(values_only=True)
    ]
    book.close()
    return [r for r in rows if r]


def jpost_listing() -> list[dict[str, Any]]:
    """Name and size of every file in the jPOST FTP directory (network)."""
    with urllib.request.urlopen(JPOST_FTP, timeout=120) as response:
        lines = response.read().decode().splitlines()
    out = []
    for line in lines:
        parts = line.split()
        out.append({"name": parts[-1], "bytes": int(parts[4])})
    return sorted(out, key=lambda r: str(r["name"]))


def jpost_sample_map(
    listing: list[dict[str, Any]], samples: list[str]
) -> dict[str, Any]:
    """Which raw files are the twelve Table 2 columns, by carbon, phase and replicate."""
    raws = [str(r["name"]) for r in listing if str(r["name"]).endswith(".raw")]
    pattern = re.compile(
        r"OR13_.*_Ecoli_(glucose|glycerol)_(exp|stat)_2mg_IMAC_R(\d)\.raw"
    )
    mapped: dict[str, str] = {}
    for raw in raws:
        match = pattern.fullmatch(raw)
        if match:
            carbon, phase, rep = match.groups()
            mapped[f"Intensity {carbon}_{phase}{rep}"] = raw
    return {
        "raw_files": len(raws),
        "raw_bytes": sum(
            int(r["bytes"]) for r in listing if str(r["name"]).endswith(".raw")
        ),
        "table_2_column_to_raw_file": mapped,
        "every_table_2_column_has_one_raw_file": sorted(mapped) == sorted(samples),
        "raw_files_not_in_table_2_columns": sorted(set(raws) - set(mapped.values())),
    }


def main() -> None:
    """Deposit (optional), then measure everything off the mirror bytes."""
    parser = argparse.ArgumentParser(description=(__doc__ or "").splitlines()[0])
    parser.add_argument("--deposit", action="store_true")
    parser.add_argument("--network", action="store_true")
    args = parser.parse_args()
    root = deposit() if args.deposit else raw_mirror_dir()
    hashes = verify_mirror(root)
    t2 = read_table_2(root)
    mg1655 = resolve_genes(t2, "MG1655")
    bw25113 = resolve_genes(t2, "BW25113")
    mg_loci, bw_loci = set(mg1655.pop("loci")), set(bw25113.pop("loci"))
    result: dict[str, Any] = {
        "citation_key": CITATION_KEY,
        "doi": PAPER_DOI,
        "pmid": PMID,
        "proteomexchange": PROTEOMEXCHANGE,
        "jpost": JPOST,
        "mirror": str(root),
        "mirror_sha256": hashes,
        "quotes": check_quotes(root, hashes),
        "strain_tokens_in_deposited_texts": token_counts(root),
        "workbook_rows": workbook_rows(root),
        "table_2": table_2_inventory(t2),
        "table_5_vs_table_2": table_5_relation(root, t2),
        "source_data_fig1": source_data_fig1(root),
        "identifier_resolution": {
            "MG1655": {**mg1655, "distinct_loci": len(mg_loci)},
            "BW25113": {**bw25113, "distinct_loci": len(bw_loci)},
        },
        "schema_fit": schema_fit(),
        "proteome_store_key_shapes": {
            s: store_key_shapes(s, mg_loci, bw_loci) for s in PROTEOME_STORES
        },
    }
    if args.network:
        listing = jpost_listing()
        result["jpost_listing"] = listing
        result["jpost_sample_map"] = jpost_sample_map(
            listing, result["table_2"]["sample_columns"]
        )
    elif osp.exists(RESULTS):
        with open(RESULTS) as handle:
            previous = json.load(handle)
        for key in ("jpost_listing", "jpost_sample_map"):
            if key in previous:
                result[key] = previous[key]
    os.makedirs(osp.dirname(RESULTS), exist_ok=True)
    with open(RESULTS, "w") as handle:
        json.dump(result, handle, indent=2, sort_keys=True, allow_nan=False)
        handle.write("\n")
    summary = {k: v for k, v in result.items() if k not in ("quotes", "jpost_listing")}
    print(json.dumps(summary, indent=2, default=str)[:8000])


if __name__ == "__main__":
    main()
