# torchcell/datasets/scerevisiae/piotrowski2017
# [[torchcell.datasets.scerevisiae.piotrowski2017]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/scerevisiae/piotrowski2017
# Test file: tests/torchcell/datasets/scerevisiae/test_piotrowski2017.py
"""Piotrowski 2017 MOSAIC chemical-genetic screen: env x geno -> CG z-score.

Piotrowski et al. 2017 (Nat Chem Biol 13:982-993, doi:10.1038/nchembio.2436) profiled
13,524 conditions from seven compound libraries against a 310-strain DIAGNOSTIC POOL of
kanMX deletion mutants built in the pdr1 pdr3 snq2 drug-sensitized SGA background, grown
together in 200 uL YP + 2% galactose microcultures for 48 h at 30 C and read out by
multiplexed UPTAG barcode sequencing. The released readout is the chemical-genetic (CG)
interaction z-score of each strain in each condition against the screen's DMSO control
profiles (Supplementary Note, ``si/si1.docx``). The sign is the source's own and is
recorded verbatim (``Z_SCORE_POLARITY``): "Positive z-scores suggest resistance to the
condition, and negative z-scores suggest sensitivity", the same negative-is-sick
orientation as Vanacloig 2022, Hoepfner 2014 and Wildenhain 2015.

RELEASE (PMC Article Datasets bucket, PMC6056180.1, NIHMS884855; mirrored under
``$DATA_ROOT/torchcell-raw/piotrowskiFunctionalAnnotationChemical2017/si/``):

- Supplementary Data Set 5 (``Suppl_Dataset_5.xlsx``, sheet ``All_conditions``): the
  matrix, 291 strain rows x 11,984 conditions (columns ``Cond-NNNNNN_<name>``),
  3,438,421 numeric z-scores and 48,923 ``NA`` cells. Its title reads "Average chemical
  genomic profile of 11983 compounds". The 291 rows name 289 pool strains: TCD1
  (YHR003C) and TCD2 (YKL027W) each appear twice, under the ORF in the RIKEN columns
  and under the standard name in the NCI/NIH/GSK columns, and each row is ``NA`` wherever
  its twin is scored (``ConditionLedger.strains_labelled_twice``).
- Supplementary Data Set 4 (``Suppl_Dataset_4.xlsx``): one row per condition (13,966),
  its name, dose, library, ``In_final_dataset`` and ``In_high_confidence_set`` flags,
  percent growth, and the OpenBabel SMILES and InChI.
- Supplementary Data Set 2 (``Suppl_Dataset_2.xlsx``): the pool members, ORF + common
  name + 20-mer barcode for pool version 2.0 (310 strains, the main screen).

Data Set 6 (the 1,522 high-confidence profiles) is an exact subset of Data Set 5 (442,902
cells, 0 unequal) and is not consumed; ``In_high_confidence_set`` rides in the condition
ledger instead.

RECORDS. One ``StrainEnvironmentResponseExperiment`` per numeric cell of a kept strain
and a kept condition. The genotype is the ONE screened ``BarcodedKanMxDeletionPerturbation``
(barcode, ORF and common name from Data Set 2, verbatim); the constant drug-sensitized background is the same
typed ``StrainBackground`` Vanacloig 2022 serves (``vanacloig2022.library_background``,
whose construction quotes are this paper's own Online Methods), so the two datasets share
one background entity. The environment is liquid YP + 2% galactose
(``YP_GALACTOSE_LIQUID``, base ``YP``) at 30 C for 48 h in 200 uL 96-well cultures
inoculated at 250 cells per strain from the thawed frozen pool, plus one
``SmallMoleculePerturbation``: the condition's compound at its Data Set 4 dose, delivered in
DMSO. The reference is the screen's DMSO control profile (z = 0), one per screen, because
each screen was normalized against its own DMSO profiles with its own filters.

THE TWO SCREENS. The Note splits the work into a "RIKEN" screen (every RIKEN NPDepo
compound) and an "NCI/NIH/GSK" screen ("all remaining compounds"), each with its own
count filters, LDA components and replicate-correlation threshold, so ``screen_id`` is the
screen and the reference is per screen. The release itself confirms the split: the 7
strains ``NA`` in 3,565 columns are ``NA`` in exactly the non-RIKEN columns (the
NCI/NIH/GSK screen ran a strain count filter, the RIKEN screen did not), and the two
twice-labelled strains change label at exactly that boundary. ``replicate_id`` is the condition id
(``Cond-NNNNNN``): the same compound at the same dose is released as more than one
condition (three RIKEN pairs, and NSC compounds held by two NCI libraries), and each is an
independent well, not a duplicate.

COMPOUND IDENTITY. A label with a curated ``compound_identity_table`` row takes that row
(the policy every loader follows, so a compound is one node across datasets); any other
label takes the InChIKey HASHED FROM THE RELEASED STANDARD InChI
(``resolved_compound(..., derive_from_inchi=True)``), the source's own identifier, rather
than one regenerated from its SMILES. The released SMILES and InChI are carried on the
``Compound`` either way. A condition whose InChI is absent, or that RDKit cannot read, is
STORED with the typed ``inchikey`` gap ``resolved_compound`` attaches
(``deferred_pending_source_review``: these labels have never been through the curator,
#726), never dropped and never guessed; ``preprocess/conditions.json`` lists each one.

DOSE. Each condition's dose is its Data Set 4 cell (``10-ug/ml``, ``100-uM``, ...), under
``basis=fixed``: the libraries were screened at fixed concentrations and re-screened at
lower ones when growth collapsed. ``ng/ml`` cells are stored in ug/mL (x 1e-3, exact); the
two MMS cells are numbers carrying an Excel percent format (``4e-05`` shown as ``0.004%``)
and are stored as 0.004 in the basis-free ``ConcentrationUnit.percent`` (#764). A cell that
states a RANGE (``10-100-ug/ml``) names no single dose, so its condition is dropped under
``dose_is_a_released_range`` (8 conditions, 2,312 records). That drop needs a schema
change, not a loader one: ``SmallMoleculePerturbation.concentration`` is a required
``Concentration``, so it cannot be ``None`` and therefore cannot carry a typed gap, and
``Concentration`` itself is not a gap carrier; a ``Concentration`` with ``basis=fixed`` and
no value would build, but it would hold an unstated dose silently with the released range
nowhere on the record. The schema has no interval dose.

STRAINS, AS RELEASED. Each Data Set 5 row label names one Data Set 2 pool row (by its common
name, or by its ORF where the row's common name is the ORF itself). The record stores that
row's ORF verbatim as ``systematic_gene_name`` and its common name verbatim as
``perturbed_gene_name``; no genome is consulted, so ``YAL058C-A`` is stored as
``YAL058C-A`` and every one of the 289 strains Data Set 5 carries is stored.

RECORDS NOT STORED (rules + counts in ``preprocess/dropped_records.json``; every numeric
cell is counted once, in rule order): ``condition_not_in_the_final_dataset``, the one
released column Data Set 4 marks ``In_final_dataset = False`` (Cond-008510, 289 cells; the
Data Set 5 title also counts 11,983, not 11,984). That exclusion is the release's own
metadata, not a filter of ours. Then the range-dose conditions above. ``NA`` cells are not
values and are counted beside the ledger (``not_scored_cells``), not as drops.

REPLICATES. Each released profile is the mean of the condition's technical (PCR)
replicates that survived the Note's batch-effect filter and correlation-graph collapsing:
three by design, a number the release does not give per profile and no released statistic
lets one back-solve. ``n_samples`` is therefore the conservative lower end, 1
(``preprocess/n_samples_derivation.json``, a ``StatDerivation``), with
``sample_unit=technical_replicate``; no uncertainty is released, so the three uncertainty
fields are typed gaps.
"""

from __future__ import annotations

import hashlib
import logging
import math
import os
import os.path as osp
import pickle
import re
import shutil
from collections import Counter
from collections.abc import Callable
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

import numpy as np
import numpy.typing as npt
import openpyxl
from pydantic import BaseModel, ConfigDict
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    Organism,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.datamodels.compound_identity import (
    inchikey_from_inchi,
    inchikey_from_smiles,
    resolve_compound_identity,
    resolved_compound,
)
from torchcell.datamodels.schema import (
    AssayType,
    BarcodedKanMxDeletionPerturbation,
    ComponentDefinition,
    Compound,
    Concentration,
    ConcentrationUnit,
    CultureEnvironment,
    CultureFormat,
    DoseBasis,
    EndpointRule,
    EnvironmentResponsePhenotype,
    Experiment,
    ExperimentReference,
    Genotype,
    MeasurementType,
    Media,
    MediaComponent,
    MediaComponentRole,
    PreCulture,
    PreCultureSource,
    Publication,
    SampleUnit,
    SmallMoleculePerturbation,
    Solvent,
    StrainEnvironmentResponseExperiment,
    StrainEnvironmentResponseExperimentReference,
    StrainReferenceGenome,
    Temperature,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.datasets.scerevisiae.vanacloig2022 import (
    ARRAY_KANMX,
    LIBRARY_STRAIN,
    library_background,
)
from torchcell.literature.manifest import (
    ROLE_SI_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.sequence.genome.registry import (
    SGD_S288C_R64,
    SGD_S288C_R64_GENBANK_ACCESSION,
)
from torchcell.verification.report import (
    DerivationMethod,
    Provenance,
    StatDerivation,
    VerificationReport,
)
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
)

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
CITATION_KEY = "piotrowskiFunctionalAnnotationChemical2017"
PAPER_DOI = "10.1038/nchembio.2436"
PMC_ARTICLE = "PMC6056180.1"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
LIBRARY_DIR_REL = f"torchcell-library/{CITATION_KEY}"
DATASET_SLUG = "env_chemgen_piotrowski2017"

PAPER_MD = "paper.md"
PAPER_MD_SHA256 = "9314a0dd932c1b20b4dc4297de4ce9452b09f314bf100f05f2fe718fb3415fc1"

#: The Supplementary Note (NIHMS884855-supplement-1.docx): the scoring method, the two
#: screens and their filters. Quotes are the text runs of ``word/document.xml``.
NOTE_DOCX = "si/si1.docx"
NOTE_DOCX_SHA256 = "38cddea8f1a497107304599271bcfcabfba7f65a4f7b3b616f26923be119959a"

#: When ``capture_si`` retrieved the three files (the library manifest's record).
RETRIEVED_AT = "2026-10-11T10:42:21.441520+00:00"
_PMC_BUCKET = "https://pmc-oa-opendata.s3.amazonaws.com"


class ReleaseFile(BaseModel):
    """One released file the loader consumes, where it sits, and its pinned bytes."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    library_rel: str
    filename: str
    sha256: str

    @property
    def mirror_rel(self) -> str:
        """The file's path inside the raw mirror, under its publisher filename."""
        return f"si/{self.filename}"

    @property
    def pmc_key(self) -> str:
        """The object key in the PMC Article Datasets bucket."""
        return f"{PMC_ARTICLE}/{self.filename}"

    @property
    def source_url(self) -> str:
        """The bucket URL the file was retrieved from (retrieval metadata)."""
        return f"{_PMC_BUCKET}/{self.pmc_key}"


MATRIX = ReleaseFile(
    library_rel="si/si13.xlsx",
    filename="NIHMS884855-supplement-Suppl_Dataset_5.xlsx",
    sha256="96e07f956f7edeb07359fc8b1fa5c72a167bf9bb18d1e8a680c76362e97aaefd",
)
CONDITIONS = ReleaseFile(
    library_rel="si/si11.xlsx",
    filename="NIHMS884855-supplement-Suppl_Dataset_4.xlsx",
    sha256="eba460fe354bafd4e57c2a8438201cdc9399cf97582e3d86adf3e9362f821f87",
)
POOL = ReleaseFile(
    library_rel="si/si6.xlsx",
    filename="NIHMS884855-supplement-Suppl_Dataset_2.xlsx",
    sha256="394505047acaea5e7e7d1afa2753e76be5356bb43abd4e01ec8eb6e1acd145a2",
)
RELEASE_FILES: tuple[ReleaseFile, ...] = (MATRIX, CONDITIONS, POOL)


def _paper(
    value: Any, quote: str, *, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to a verbatim quote in the sha256-pinned paper OCR mirror."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=PAPER_MD,
            citation_key=CITATION_KEY,
            sha256=PAPER_MD_SHA256,
            method="MinerU OCR of the publisher PDF (torchcell-library mirror)",
            page=page,
        ),
    )


def _note(
    value: Any, quote: str, *, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to a verbatim paragraph run of the Supplementary Note docx."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=NOTE_DOCX,
            citation_key=CITATION_KEY,
            sha256=NOTE_DOCX_SHA256,
            method="paragraph text of word/document.xml (the w:t runs of each w:p "
            "concatenated); the NIHMS884855 Supplementary Note",
            page=page,
        ),
    )


def _sheet(
    value: Any, quote: str, *, release: ReleaseFile, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to a verbatim cell of a released workbook (raw-mirror copy)."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=release.library_rel,
            citation_key=CITATION_KEY,
            sha256=release.sha256,
            method="cell value read with openpyxl (read_only, data_only)",
            page=page,
        ),
    )


_METHODS = "Online Methods"
_SCREEN_PAGE = f"{_METHODS}, 'Screening the NPDepo, NCI, NIH, and GSK collections' (paper.md line 226)"
_OPTIMIZATION_PAGE = (
    f"{_METHODS}, 'Optimization of signal detection and sequencing parameters' "
    "(paper.md line 218)"
)
_NOTE_SCORING = "Supplementary Note, 'Chemical-genetic interaction scoring method'"
_NOTE_SCREENS = (
    "Supplementary Note, 'Applying the chemical genetic interaction scoring method to "
    "chemical genomic screens in this study'"
)

# --------------------------------------------------------------------------- #
# Sourced culture, readout and design constants (verbatim quotes, pinned sha256)
# --------------------------------------------------------------------------- #
SCREEN_DEFERS_TO_OPTIMIZATION = _paper(
    "optimized conditions",
    "We performed our pooled growth assay with the diagnostic mutant collection under "
    "optimized conditions as described above.",
    page=_SCREEN_PAGE,
    note="the main screen restates no medium, temperature, inoculum or vehicle fraction; "
    "it defers them to the optimization section (paper.md line 218) and to the selected "
    "screening conditions (line 33), which is where those values are quoted from",
)
SELECTED_CONDITIONS = _paper(
    {
        "working_volume_ul": 200.0,
        "duration_hours": 48.0,
        "inoculum_cells_per_strain": 250.0,
    },
    "Ultimately, the screening conditions we selected were $2 0 0 \\mathrm { - } \\mu "
    "\\mathrm { L }$ microcultures grown for $4 8 \\mathrm { { h } }$ at an inoculum of "
    "250 cells/strain and 30 PCR cycles for barcode amplification.",
    page="Results, 'Optimizing signal detection and high-throughput screening' (paper.md "
    "line 33)",
)
SCREEN_VESSEL = _paper(
    "96-well plates",
    "The diagnostic mutant pool was grown in $2 0 0 \\mu \\mathrm { L }$ cultures in "
    "96-well plates.",
    page=_SCREEN_PAGE,
)
SCREEN_MEDIUM = _paper(
    "YP + 2% galactose",
    "To assay the mutant pool for drug hypersensitivity, cells were thawed, counted "
    "using a hemocytometer, and diluted to seven different final inoculum densities "
    "(3,727–58 cells/strain) in $\\mathrm { Y P } + 2 \\%$ galactose in a 96-well "
    "flat-bottom plate.",
    page=_OPTIMIZATION_PAGE,
    note="also the pre-culture: the frozen pool is thawed, counted and diluted straight "
    "into the assay, with no recovery culture (PreCultureSource.frozen_stock)",
)
YPGAL_RECIPE = _paper(
    {"yeast extract": 1.0, "peptone": 2.0, "galactose": 2.0},
    "Cultures of the control strain (y13206) were grown to early log phase (0.4 OD) in "
    "YPGal $1 \\%$ yeast extract, $2 \\%$ peptone, $2 \\%$ galactose).",
    page=f"{_METHODS}, flow cytometry validations (paper.md line 280)",
    note="the paper's percent recipe of YPGal; '%' carries no w/v or v/v, recorded as w/v "
    "like every percentage of the media library",
)
YP_GRAMS = _paper(
    {"yeast extract g/L": 10.0, "peptone g/L": 20.0},
    "and $2 \\mathrm { m L }$ was added to a $5 0 \\mathrm { - m L }$ stock of $2 \\%$ "
    "YP $\\mathrm { 1 0 ~ g / L }$ yeast extract, $2 0 \\mathrm { g / L }$ peptone) "
    "$+ 2 \\%$ galactose $+ ~ 1 \\%$ agar (YPGal).",
    page=f"{_METHODS}, halo assay (paper.md line 206)",
    note="the same YP in grams per liter (10 g/L = 1% w/v, 20 g/L = 2% w/v); this "
    "sentence is the solid halo-assay plate, so only its YP amounts are used",
)
TEMPERATURE_C = _paper(
    30.0,
    "After growing for 18, 24, or $4 8 \\mathrm { { h } }$ at $3 0 ~ ^ { \\circ } "
    "\\mathrm { C } ,$ , cells from each well were harvested by centrifugation.",
    page=_OPTIMIZATION_PAGE,
    note="the selected 48 h is the screen's duration (SELECTED_CONDITIONS)",
)
FROZEN_POOL = _paper(
    "frozen pool at -80 C",
    "Finally, the pool was adjusted to a final concentration of $5 0 ~ \\mathrm { { O D } "
    "_ { 6 0 0 } / m L }$ by dilution or centrifugation and stored at $- 8 0 ~ ^ { \\circ "
    "} \\mathrm { C }$ until required.",
    page=_OPTIMIZATION_PAGE,
)
DMSO_CONTROL_PERCENT = _paper(
    1.0,
    "Cultures were then spiked with either $3 4 . 4 \\mu \\mathrm { M }$ benomyl, $2 5 ~ "
    "\\mathrm { n M }$ micafungin, or a $1 \\%$ DMSO control.",
    page=_OPTIMIZATION_PAGE,
    note="the DMSO fraction of the control, written '1%' with no v/v or w/v, so the "
    "reference's DMSO dose is the basis-free ConcentrationUnit.percent (#764); the "
    "screen restates no fraction and defers to this section "
    "(SCREEN_DEFERS_TO_OPTIMIZATION). Consistency check, not a source: the NPDepo's 1 "
    "mg/mL stocks screened at 10 ug/mL are a 1:100 dilution (SCREEN_DOSES)",
)
COMPOUND_SOLVENT = _note(
    "DMSO",
    "The 5724 control profiles were generated by growing the yeast diagnostic mutant pool "
    "in the absence of a compound but in the presence of the compound’s solvent, DMSO (a "
    "negative control)",
    page="Supplementary Note, 'Computing a combined z-score and p-value for compounds and "
    "their process targets'",
    note="names DMSO as every compound's solvent; its fraction is DMSO_CONTROL_PERCENT",
)
SCREEN_DOSES = _paper(
    {
        "RIKEN NPDepo": "10 ug/mL",
        "other collections": "100 uM",
        "NCI-MECH-DIV": "10 uM",
    },
    "The NPDepo is maintained as $1 ~ \\mathrm { m g / m L }$ stocks, and we screened it "
    "at a final concentration of $1 0 ~ \\mu \\mathrm { g / m L }$ , with the exception of "
    "a number of compounds that received additional lower dosing in a pilot experiment "
    "(Supplementary Data Set 4). All remaining collections were screened at $1 0 0 \\ "
    "\\mu \\mathrm { M } ,$ except for the NCI Mechanistic Diversity set $1 0 \\mu "
    "\\mathrm { M }$ ; Supplementary Data Set 4). Selected compounds were re-screened at "
    "lower concentrations if the initial concentration resulted in severe growth "
    "inhibition.",
    page=_SCREEN_PAGE,
    note="each condition's own dose is its Data Set 4 'Concentration' cell, under "
    "DoseBasis.fixed; the 'additional lower dosing in a pilot experiment' is what a "
    "range cell records, which is why a range names no single dose",
)
DESIGN_TECHNICAL_REPLICATES = _paper(
    3,
    "The genomic extractions for each plate were amplified in triplicate using three "
    "unique multiplex primer plates (three technical replicates).",
    page=f"{_METHODS}, screening (paper.md line 228)",
)
FIG2_REPLICATES = _paper(
    3,
    "Mean negative chemical-genetic interactions are represented in red $\\left( n = 3 , "
    "\\right.$ technical replicates).",
    page="Figure 2 caption (paper.md line 41)",
    note="the design count, stated for the high-confidence heat map",
)
REPLICATE_COLLAPSING = _note(
    "mean of the technical replicates that correlate",
    "In the “replicate collapsing” step, profiles from technical replicate conditions "
    "were collapsed into single, mean profiles. To determine which replicates would be "
    "collapsed into each final profile, a graph was constructed in which nodes were the "
    "technical replicates and edges indicated a correlation between two replicates above "
    "a certain threshold (“replicate correlation threshold”).",
    page=_NOTE_SCORING,
    note="a released profile is the mean of the replicates the correlation graph "
    "joined, and the batch-effect filter removed whole multiplex tags before that, so "
    "the count behind a profile is between 1 and the design's 3 and is not released",
)
REPLICATE_THRESHOLDS = _note(
    {"RIKEN": 0.7, "NCI/NIH/GSK": 0.5},
    "Replicate collapsing (replicate correlation threshold: 0.7)",
    page=f"{_NOTE_SCREENS}, RIKEN screen",
    note="the NCI/NIH/GSK screen's line reads 'Replicate collapsing (replicate "
    "correlation threshold: 0.5)'",
)
Z_SCORE_DEFINITION = _note(
    MeasurementType.z_score,
    "Finally, chemical-genetic interactions were computed as z-scores representing the "
    "standardized deviation of each strain in the condition profile with respect to its "
    "counterpart strain in the reference profile.",
    page=_NOTE_SCORING,
)
Z_SCORE_REFERENCE = _note(
    "DMSO control profiles",
    "First, the reference profile was calculated as the log of the strain-wise average "
    "profile across all DMSO control read count profiles (excluding read counts less "
    "than 20, which were set to NaN and excluded from the calculation).",
    page=_NOTE_SCORING,
)
Z_SCORE_POLARITY = _note(
    "positive = resistant, negative = sensitive",
    "Positive z-scores suggest resistance to the condition, and negative z-scores "
    "suggest sensitivity to the condition relative to the fitness of each mutant in the "
    "control condition.",
    page=_NOTE_SCORING,
    note="the orientation recorded before any pooling: negative is the sick strain, as "
    "in Vanacloig 2022, Hoepfner 2014 and Wildenhain 2015 and opposite to both "
    "Hillenmeyer 2008 arms",
)
CG_SCORE_POLARITY_RESULTS = _paper(
    "negative = hypersensitive",
    "Chemical-genetic interactions were identified and scored by comparing the "
    "individual mutant barcode read counts to those from a set of solvent control "
    "conditions. A negative chemical-genetic (CG) interaction score represents "
    "hypersensitivity to a compound, whereas a positive CG score represents resistance "
    "(see Online Methods).",
    page="Results, 'Chemical-genetic profiling of diverse compound libraries' (paper.md "
    "line 45)",
)
ASSAY = _paper(
    AssayType.pooled_competitive_growth_barcode,
    "Primers containing the Illumina sequencing adaptor, common priming site for the "
    "UPTAG barcode, and 1,000 selected 10-bp multiplex tags were synthesized (Sigma), "
    "arrayed in 96-well plates.",
    page=f"{_METHODS}, multiplex tag design (paper.md line 222)",
    note="a pooled culture read out by amplifying each strain's UPTAG barcode; the Data "
    "Set 2 'barcode' column is that UPTAG",
)
BARCODE_READ = _note(
    "barcode = bases 28-47",
    "Each retained sequence was then split into a “multiplex tag” (bases 1–10) and a "
    "“barcode” (bases 28–47).",
    page=_NOTE_SCORING,
)
SCREEN_SPLIT = _note(
    {"RIKEN": "RIKEN NPDepo", "NCI/NIH/GSK": "every other collection"},
    "Data in this study were collected from two large screening efforts. The “RIKEN” "
    "screen encompasses all compounds originating from the RIKEN Natural Product "
    "Depository, and the “NCI/NIH/GSK” screen encompasses all remaining compounds (the "
    "NCI collections, NIH Clinical Collection, and GSK Kinase Inhibitor collection).",
    page=_NOTE_SCREENS,
    note="the POSCON-CLINICAL-2013 positive controls are 'remaining compounds', and the "
    "release agrees: the strains the NCI/NIH/GSK strain-count filter removed are NA in "
    "exactly those columns too",
)
NCI_STRAIN_FILTER = _note(
    "50% of conditions > 20 reads",
    "Strain count filter: For each strain to pass this filter, 50% of the conditions in "
    "the strain vector were required to possess more than 20 read counts.",
    page=f"{_NOTE_SCREENS}, NCI/NIH/GSK screen",
    note="why a strain is NA across one screen's columns; the RIKEN screen ran no strain "
    "filter ('Strain count filter: No filter was applied')",
)
GTR1_AVT5_REMOVED = _note(
    ["GTR1", "AVT5"],
    "First, data from the gtr1∆ and avt5∆ strains were removed, because these strains "
    "were over represented following growth in a specific set of conditions and thus "
    "accounted for a large majority of the read counts (Supplementary Fig. 7).",
    page=_NOTE_SCORING,
    note="two of the 21 pool-v2.0 strains Data Set 5 does not carry; the other 19 are "
    "absent with no stated reason (listed in the condition ledger's absent_pool_strains)",
)
POOL_V2 = _paper(
    310,
    "These were removed to generate a collection of 310 strains for the final version of "
    "the diagnostic strain set (Supplementary Data Set 2, diagnostic pool version 2).",
    page=f"{_METHODS}, diagnostic gene set (paper.md line 216)",
)
POOL_COLLECTION = _sheet(
    "Diagnostic pool version 2.0 (310 strains, used for main screen)",
    "Diagnostic pool version 2.0 (310 strains, used for main screen)",
    release=POOL,
    page="sheet 'Suppl. dataset 2', cell A4",
    note="the perturbation's collection, verbatim",
)
MATRIX_TITLE = _sheet(
    11983,
    "Supplementary dataset 5. Average chemical genomic profile of 11983 compounds",
    release=MATRIX,
    page="sheet 'All_conditions', cell A1",
    note="the release carries 11,984 columns; the extra one is Cond-008510 (NPE1123), "
    "which Data Set 4 marks In_final_dataset = False",
)
FINAL_DATASET_FLAG = _sheet(
    "In_final_dataset",
    "In_final_dataset",
    release=CONDITIONS,
    page="sheet 'Suppl. dataset 4', cell E6",
)
CHEMICAL_IDENTIFIERS = _paper(
    "OpenBabel SMILES and InChI",
    "SMILES and InChI string representations of all molecules were generated using the "
    "OpenBabel cheminformatics toolkit53 (http://openbabel.org) and its python wrapper, "
    "pybel54 .",
    page=f"{_METHODS}, molecular descriptors (paper.md line 234)",
    note="the released InChI is the source's own standard identifier, so its hash is the "
    "InChIKey stored for a label with no curated table row",
)

#: The two screens, their Note names, and the collection that maps to the first.
SCREEN_RIKEN = "RIKEN"
SCREEN_NCI_NIH_GSK = "NCI/NIH/GSK"
RIKEN_COLLECTION = "RIKEN NPDepo"

#: The cell value the release writes for a strain not scored in a condition.
NOT_SCORED = "NA"

#: Hoepfner and Wildenhain's convention for the three uncertainty fields.
_UNCERTAINTY_FIELDS = (
    "environment_response_se",
    "environment_response_uncertainty",
    "environment_response_uncertainty_type",
)

_PERCENT_BASIS_NOTE = (
    "the source writes '%' with no w/v or v/v basis; recorded as w/v, which is the "
    "basis for every other percentage in the media library"
)

# --------------------------------------------------------------------------- #
# The medium: liquid YP + 2% galactose (loader-local, base_medium "YP")
# --------------------------------------------------------------------------- #
YP_GALACTOSE_LIQUID = Media(
    name="YP + 2% galactose (YPGal), liquid",
    state="liquid",
    is_synthetic=False,
    base_medium="YP",
    components=[
        MediaComponent(
            compound=Compound(name="yeast extract"),
            role=MediaComponentRole.complex_ingredient,
            definition=ComponentDefinition.intrinsically_undefined,
            concentration=Concentration(value=1.0, unit=ConcentrationUnit.percent_w_v),
            provenance=[YPGAL_RECIPE, YP_GRAMS],
        ),
        MediaComponent(
            compound=Compound(name="peptone"),
            role=MediaComponentRole.complex_ingredient,
            definition=ComponentDefinition.intrinsically_undefined,
            concentration=Concentration(value=2.0, unit=ConcentrationUnit.percent_w_v),
            provenance=[YPGAL_RECIPE, YP_GRAMS],
        ),
        MediaComponent(
            compound=resolved_compound("galactose"),
            role=MediaComponentRole.carbon_source,
            concentration=Concentration(value=2.0, unit=ConcentrationUnit.percent_w_v),
            provenance=[SCREEN_MEDIUM, YPGAL_RECIPE],
            note=_PERCENT_BASIS_NOTE,
        ),
    ],
    provenance=[SCREEN_MEDIUM, SCREEN_DEFERS_TO_OPTIMIZATION],
)
"""The screen's medium: the library's ``YP`` base plus 2% galactose, in liquid culture.

Loader-local because the library's ``YP_GALACTOSE`` is Bloom 2019's SOLID plate with
Bloom's own bench quote; this one joins it at ``base_medium == "YP"`` and at the
galactose compound without borrowing that quote."""


# --------------------------------------------------------------------------- #
# Typed release rows
# --------------------------------------------------------------------------- #
class PoolStrain(BaseModel):
    """One Data Set 2 pool-version-2.0 row: the strain's ORF, common name and UPTAG."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    orf: str
    common_name: str
    barcode: str


class ConditionRow(BaseModel):
    """The Data Set 4 cells of one condition the loader reads, verbatim."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    condition_id: str
    name: str
    concentration_cell: str | float
    concentration_number_format: str
    collection: str
    in_final_dataset: bool
    in_high_confidence_set: bool
    percent_growth: float | None
    smiles: str | None
    inchi: str | None

    @property
    def screen(self) -> str:
        """The Note's screen this condition belongs to (``SCREEN_SPLIT``)."""
        return (
            SCREEN_RIKEN if self.collection == RIKEN_COLLECTION else SCREEN_NCI_NIH_GSK
        )


class ReleasedMatrix(BaseModel):
    """Data Set 5 as read: strain labels, condition ids and names, and the z-scores.

    ``values`` is strains x conditions with ``NaN`` where the release writes ``NA``.
    """

    model_config = ConfigDict(extra="forbid", arbitrary_types_allowed=True)

    strains: list[str]
    condition_ids: list[str]
    condition_labels: list[str]
    values: npt.NDArray[np.float64]


# --------------------------------------------------------------------------- #
# Dose parsing (Data Set 4 'Concentration')
# --------------------------------------------------------------------------- #
_UNIT_FACTORS: dict[str, tuple[ConcentrationUnit, Decimal]] = {
    "ug/ml": (ConcentrationUnit.ug_per_ml, Decimal(1)),
    "ng/ml": (ConcentrationUnit.ug_per_ml, Decimal("0.001")),
    "uM": (ConcentrationUnit.micromolar, Decimal(1)),
    "nM": (ConcentrationUnit.nanomolar, Decimal(1)),
    "mM": (ConcentrationUnit.millimolar, Decimal(1)),
}
_UNITS = "|".join(re.escape(unit) for unit in _UNIT_FACTORS)
_NUMBER = r"\d+(?:\.\d+)?"
_SINGLE_DOSE = re.compile(rf"^(?P<value>{_NUMBER})[- ](?P<unit>{_UNITS})$")
_RANGE_DOSE = re.compile(
    rf"^(?P<low>{_NUMBER})-(?P<high>{_NUMBER})-?(?P<unit>{_UNITS})$"
)


class ParsedDose(BaseModel):
    """One Data Set 4 dose cell, parsed: a single concentration, or a released range."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    source: str
    kind: Literal["single", "range"]
    concentration: Concentration | None = None
    range_low: float | None = None
    range_high: float | None = None
    range_unit: ConcentrationUnit | None = None


def parse_dose(cell: str | float, number_format: str) -> ParsedDose:
    """Read one Data Set 4 dose cell; anything the loader was not taught raises.

    ``<n>-<unit>`` (or ``<n> <unit>``, one Bortezomib cell) is a single dose; ``ng/ml`` is
    rescaled to ug/mL (x 1e-3, an exact decimal step). A NUMBER cell must carry an Excel
    percent format: its value is a fraction, stored x 100 in the basis-free
    ``ConcentrationUnit.percent`` (the MMS cells: ``4e-05`` -> 0.004). ``<a>-<b>-<unit>``
    (and ``0.2-10ug/ml``) is a range, which names no single dose.
    """
    if isinstance(cell, float):
        if "%" not in number_format:
            raise ValueError(f"numeric dose {cell!r} without a percent format")
        value = float(Decimal(repr(cell)) * 100)
        return ParsedDose(
            source=f"{cell!r} ({number_format})",
            kind="single",
            concentration=Concentration(
                value=value, unit=ConcentrationUnit.percent, basis=DoseBasis.fixed
            ),
        )
    single = _SINGLE_DOSE.fullmatch(cell)
    if single is not None:
        unit, factor = _UNIT_FACTORS[single["unit"]]
        return ParsedDose(
            source=cell,
            kind="single",
            concentration=Concentration(
                value=float(Decimal(single["value"]) * factor),
                unit=unit,
                basis=DoseBasis.fixed,
            ),
        )
    ranged = _RANGE_DOSE.fullmatch(cell)
    if ranged is not None:
        unit, factor = _UNIT_FACTORS[ranged["unit"]]
        return ParsedDose(
            source=cell,
            kind="range",
            range_low=float(Decimal(ranged["low"]) * factor),
            range_high=float(Decimal(ranged["high"]) * factor),
            range_unit=unit,
        )
    raise ValueError(f"unreadable Data Set 4 dose cell {cell!r}")


# --------------------------------------------------------------------------- #
# Reading the release
# --------------------------------------------------------------------------- #
_COLUMN_RE = re.compile(r"^(?P<id>Cond-\d{6})_(?P<label>.+)$")
_CONDITION_HEADER = (
    "Condition_ID",
    "Name",
    "Concentration",
    "Collection",
    "In_final_dataset",
    "In_high_confidence_set",
    "Percent Growth",
    "SMILES",
    "InChI",
)
_POOL_HEADER = ("ORF", "Common name", "barcode")


def _structure_string(value: Any) -> str | None:
    """A SMILES / InChI cell: the string, or ``None`` for an empty or ``NA`` cell."""
    if value is None or value == NOT_SCORED:
        return None
    if not isinstance(value, str) or value.strip() == "":
        raise ValueError(f"unexpected structure cell {value!r}")
    return value


def read_conditions(path: str | Path) -> dict[str, ConditionRow]:
    """Every Data Set 4 condition row, keyed by its ``Condition_ID``.

    The header is row 6 (rows 1-3 are the title and two notes); the dose cell is read
    as a cell, not a value, because the two MMS doses are numbers whose meaning is their
    percent number format.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Suppl. dataset 4"]
    rows = sheet.iter_rows()
    for _ in range(5):
        next(rows)
    header = [cell.value for cell in next(rows)]
    if tuple(header[: len(_CONDITION_HEADER)]) != _CONDITION_HEADER:
        raise ValueError(f"Data Set 4 header changed: {header[:9]!r}")
    conditions: dict[str, ConditionRow] = {}
    for cells in rows:
        cid = cells[0].value
        if cid is None:
            continue
        dose = cells[2]
        growth = cells[6].value
        row = ConditionRow(
            condition_id=str(cid),
            name=str(cells[1].value),
            concentration_cell=dose.value,
            concentration_number_format=str(dose.number_format),
            collection=str(cells[3].value),
            in_final_dataset=bool(cells[4].value),
            in_high_confidence_set=bool(cells[5].value),
            percent_growth=None if growth in (None, NOT_SCORED) else float(growth),
            smiles=_structure_string(cells[7].value),
            inchi=_structure_string(cells[8].value),
        )
        if row.condition_id in conditions:
            raise ValueError(f"Data Set 4 repeats {row.condition_id}")
        conditions[row.condition_id] = row
    workbook.close()
    return conditions


def read_pool(path: str | Path) -> list[PoolStrain]:
    """The Data Set 2 pool-version-2.0 members (columns A-C below the row-5 header).

    The block ends at the first row whose barcode cell is empty (the summary rows
    'total strains', 'no fitness score on SGA' and '%').
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Suppl. dataset 2"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if rows[3][0] != POOL_COLLECTION.value:
        raise ValueError(f"Data Set 2 block title changed: {rows[3][0]!r}")
    if tuple(rows[4][:3]) != _POOL_HEADER:
        raise ValueError(f"Data Set 2 header changed: {rows[4][:3]!r}")
    pool: list[PoolStrain] = []
    for row in rows[5:]:
        if row[0] is None or not isinstance(row[2], str):
            break
        pool.append(PoolStrain(orf=row[0], common_name=row[1], barcode=row[2]))
    if len(pool) != POOL_V2.value:
        raise ValueError(f"Data Set 2 pool v2.0 holds {len(pool)} strains, not 310")
    return pool


def read_matrix(path: str | Path) -> ReleasedMatrix:
    """Data Set 5: strain rows x ``Cond-NNNNNN_<name>`` columns (header row 3)."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["All_conditions"]
    rows = sheet.iter_rows(values_only=True)
    title = next(rows)[0]
    if title != MATRIX_TITLE.quote:
        raise ValueError(f"Data Set 5 title changed: {title!r}")
    next(rows)
    header = next(rows)
    ids: list[str] = []
    labels: list[str] = []
    for column in header[1:]:
        match = _COLUMN_RE.fullmatch(str(column))
        if match is None:
            raise ValueError(f"unparseable Data Set 5 column {column!r}")
        ids.append(match["id"])
        labels.append(match["label"])
    strains: list[str] = []
    values: list[list[float]] = []
    for row in rows:
        if row[0] is None:
            continue
        strains.append(str(row[0]))
        parsed: list[float] = []
        for cell in row[1:]:
            if cell == NOT_SCORED:
                parsed.append(math.nan)
            elif isinstance(cell, int | float) and not isinstance(cell, bool):
                if not math.isfinite(cell):
                    raise ValueError(f"non-finite Data Set 5 cell {cell!r}")
                parsed.append(float(cell))
            else:
                raise ValueError(f"unexpected Data Set 5 cell {cell!r}")
        values.append(parsed)
    workbook.close()
    return ReleasedMatrix(
        strains=strains,
        condition_ids=ids,
        condition_labels=labels,
        values=np.asarray(values, dtype=np.float64),
    )


# --------------------------------------------------------------------------- #
# Strains: the pool row each Data Set 5 label names, stored as released
# --------------------------------------------------------------------------- #
class ResolvedStrain(BaseModel):
    """One Data Set 5 strain row matched to its Data Set 2 pool row, as released."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    label: str
    pool: PoolStrain
    systematic_name: str
    perturbed_gene_name: str


def resolve_strains(labels: list[str], pool: list[PoolStrain]) -> list[ResolvedStrain]:
    """Match each Data Set 5 label to its pool row; names are the release's, verbatim.

    The label is the pool row's common name (or its ORF, for the rows whose common name
    is the ORF). ``systematic_name`` is the pool row's ORF and ``perturbed_gene_name`` its
    common name (the ORF when the row gives none). No genome is consulted.
    """
    by_name = {entry.common_name: entry for entry in pool}
    by_orf = {entry.orf: entry for entry in pool}
    if len(by_name) != len(pool) or len(by_orf) != len(pool):
        raise ValueError("Data Set 2 repeats a common name or an ORF")
    resolved: list[ResolvedStrain] = []
    for label in labels:
        entry = by_name.get(label) or by_orf.get(label)
        if entry is None:
            raise ValueError(f"Data Set 5 strain {label!r} is not a pool v2.0 member")
        resolved.append(
            ResolvedStrain(
                label=label,
                pool=entry,
                systematic_name=entry.orf,
                perturbed_gene_name=entry.common_name or entry.orf,
            )
        )
    return resolved


def strains_labelled_twice(
    matrix: ReleasedMatrix, strains: list[ResolvedStrain]
) -> dict[str, list[str]]:
    """Pool strains that more than one Data Set 5 row names, with those row labels.

    Two rows naming one pool strain are one strain under two labels only when no
    condition scores both; a condition that did would be a duplicate measurement, so it
    raises. Measured on the release: TCD1/YHR003C and TCD2/YKL027W, each pair split
    exactly at the RIKEN vs NCI/NIH/GSK column boundary.
    """
    rows_by_orf: dict[str, list[int]] = {}
    for i, strain in enumerate(strains):
        rows_by_orf.setdefault(strain.pool.orf, []).append(i)
    twice: dict[str, list[str]] = {}
    for orf, rows in sorted(rows_by_orf.items()):
        if len(rows) < 2:
            continue
        scored = (~np.isnan(matrix.values[rows])).sum(axis=0)
        if int((scored > 1).sum()) > 0:
            raise ValueError(
                f"pool strain {orf} is scored twice in one condition by rows "
                f"{[strains[i].label for i in rows]}"
            )
        twice[orf] = [strains[i].label for i in rows]
    return twice


# --------------------------------------------------------------------------- #
# Compounds
# --------------------------------------------------------------------------- #
class CompoundLedgerEntry(BaseModel):
    """How one condition's compound was identified, and what the release says of it."""

    model_config = ConfigDict(extra="forbid")

    route: Literal["curated_table", "released_inchi", "gapped"]
    stored_name: str
    stored_inchikey: str | None
    released_inchi_key: str | None
    released_smiles_key: str | None


def compound_for(row: ConditionRow) -> Compound:
    """The condition's compound: curated row, else the released InChI's hash, else a gap."""
    return resolved_compound(
        row.name, smiles=row.smiles, inchi=row.inchi, derive_from_inchi=True
    )


def compound_ledger_entry(row: ConditionRow, compound: Compound) -> CompoundLedgerEntry:
    """The identity route of ``compound`` plus both released-structure keys, measured.

    The two released keys are the hash of the released InChI and RDKit's key from the
    released SMILES; where a curated row won, a stored key that disagrees with them is
    what the ledger exists to show.
    """
    curated = resolve_compound_identity(name=row.name).curated
    route: Literal["curated_table", "released_inchi", "gapped"] = (
        "curated_table"
        if curated
        else "released_inchi"
        if compound.inchikey is not None
        else "gapped"
    )
    return CompoundLedgerEntry(
        route=route,
        stored_name=compound.name,
        stored_inchikey=compound.inchikey,
        released_inchi_key=None
        if row.inchi is None
        else inchikey_from_inchi(row.inchi),
        released_smiles_key=(
            None if row.smiles is None else inchikey_from_smiles(row.smiles)
        ),
    )


# --------------------------------------------------------------------------- #
# Ledger models
# --------------------------------------------------------------------------- #
class DropRule(BaseModel):
    """One retention rule, the numeric cells it removed, and the items it names."""

    model_config = ConfigDict(extra="forbid")

    rule: str
    scope: Literal["condition"]
    description: str
    n_records: int
    items: list[str] = []


class DropLog(BaseModel):
    """Every retention rule of a build, in the order applied, plus the NA cell count."""

    model_config = ConfigDict(extra="forbid")

    dataset: str
    source_records: int
    kept_records: int
    dropped_records: int
    rules: list[DropRule]
    released_cells: int
    not_scored_cells: int
    not_scored_strains_by_screen: dict[str, dict[str, int]]


class ConditionLedgerRow(BaseModel):
    """One released condition: its Data Set 4 metadata, its fate, and its compound."""

    model_config = ConfigDict(extra="forbid")

    condition_id: str
    matrix_label: str
    name: str
    collection: str
    screen: str
    source_dose: str
    in_high_confidence_set: bool
    percent_growth: float | None
    n_scored_strains: int
    drop_rule: str | None
    compound: CompoundLedgerEntry | None


class ConditionLedger(BaseModel):
    """``preprocess/conditions.json``: every released condition and the pool census."""

    model_config = ConfigDict(extra="forbid")

    dataset: str
    conditions: list[ConditionLedgerRow]
    route_counts: dict[str, int]
    curated_key_disagreements: list[str]
    released_key_disagreements: int
    released_key_comparisons: int
    absent_pool_strains: list[str]
    strains_labelled_twice: dict[str, list[str]]


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the environment (the mirrors and the build tree live there)."""
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/piotrowskiFunctionalAnnotationChemical2017``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def _sha256(path: str | Path) -> str:
    """Streaming sha256 of a file."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def deposit_raw_mirror(data_root: str | None = None) -> Path:
    """Copy the three consumed files from the literature mirror into the raw mirror.

    Each file is sha256-checked before it is copied and an existing mirror file with
    another hash raises. The retrieval record is the one ``capture_si`` wrote when it
    fetched the file from the PMC Article Datasets bucket (``pmc_cloud``), so the
    recorded retrieval re-runs as-is.
    """
    root = data_root or _data_root()
    mirror = raw_mirror_dir(root)
    records: list[ArtifactRecord] = []
    for release in RELEASE_FILES:
        src = Path(root) / LIBRARY_DIR_REL / release.library_rel
        got = _sha256(src)
        if got != release.sha256:
            raise RuntimeError(f"{src} sha256 {got} != pinned {release.sha256}")
        dest = mirror / release.mirror_rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            if _sha256(dest) != release.sha256:
                raise RuntimeError(f"{dest} exists with another sha256; refusing")
        else:
            shutil.copy2(src, dest)
        records.append(
            ArtifactRecord(
                path=release.mirror_rel,
                role=ROLE_SI_DATA,
                bytes=dest.stat().st_size,
                sha256=release.sha256,
                source=release.source_url,
                original_filename=release.filename,
                retrieval=RetrievalRecord(
                    method=RetrievalMethod.pmc_cloud,
                    source_url=release.source_url,
                    retriever="torchcell.literature.retrieve.pmc_cloud_object",
                    params={"key": release.pmc_key},
                    sha256=release.sha256,
                    retrieved_at=RETRIEVED_AT,
                ),
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title="Functional annotation of chemical libraries across diverse biological "
        "processes",
        files=records,
        si_data_sources=[f"{_PMC_BUCKET}/{PMC_ARTICLE}/"],
        si_expected=[
            "Supplementary Data Sets 1-22, the Supplementary Note and the Supplementary "
            "Information PDF are in the literature mirror (si/si1-si24); only the three "
            "files the loader reads are deposited here"
        ],
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (mirror / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return mirror


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


# --------------------------------------------------------------------------- #
# Dataset
# --------------------------------------------------------------------------- #
_UNITS_TEXT = (
    "chemical-genetic interaction z-score (Supplementary Note): the strain's "
    "LOWESS-normalized log read count in the condition minus its log mean read count "
    "over the screen's DMSO control profiles, divided by the larger of the condition's "
    "middle-75% SD and the sign-specific reference SD; multiplex-tag LDA components "
    "(6 RIKEN, 5 NCI/NIH/GSK) and the first SVD component removed; technical replicates "
    "joined by the correlation graph averaged; positive = resistant, negative = "
    "sensitive (Supplementary Data Set 5)"
)
_SE_GAP_NOTE = (
    "Data Set 5 releases one averaged z-score per (strain, condition) and no dispersion "
    "of the technical replicates behind it"
)
BATCH_SIZE = 500_000


@register_dataset
class EnvChemgenPiotrowski2017Dataset(ExperimentDataset):
    """MOSAIC diagnostic-pool chemical genetics: env x geno -> CG interaction z-score."""

    organism = Organism.scerevisiae

    def __init__(
        self,
        root: str = "data/torchcell/env_chemgen_piotrowski2017",
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize the dataset; no genome is consulted."""
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return StrainEnvironmentResponseExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return StrainEnvironmentResponseExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The three released workbooks the build reads."""
        return [release.filename for release in RELEASE_FILES]

    def download(self) -> None:
        """Link the mirror files into ``raw/`` after verifying each against its pin.

        The manifest is the retrieval record and must carry each pin, else
        ``ManifestPinMismatchError`` names both digests.
        """
        data_root = _data_root()
        manifest = load_manifest(data_root)
        os.makedirs(self.raw_dir, exist_ok=True)
        for release in RELEASE_FILES:
            check_manifest_pin(
                release.mirror_rel,
                manifest_sha256(manifest, release.mirror_rel),
                release.sha256,
            )
            src = raw_mirror_dir(data_root) / release.mirror_rel
            link_verified(src, osp.join(self.raw_dir, release.filename), release.sha256)
        log.info(
            "Piotrowski 2017 release linked into %s (sha256 verified)", self.raw_dir
        )

    def _raw(self, release: ReleaseFile) -> str:
        return osp.join(self.raw_dir, release.filename)

    # ---- record builders ---------------------------------------------------- #
    @staticmethod
    def _vehicle() -> Solvent:
        """DMSO at the control's 1% (``COMPOUND_SOLVENT`` + ``DMSO_CONTROL_PERCENT``)."""
        return Solvent(
            name="DMSO",
            percent=DMSO_CONTROL_PERCENT.value,
            compound=resolved_compound("DMSO"),
        )

    @staticmethod
    def _culture_format() -> CultureFormat:
        """200 uL 96-well cultures from 250 cells per strain, read at a fixed 48 h."""
        return CultureFormat(
            vessel=SCREEN_VESSEL.value,
            working_volume_ul=SELECTED_CONDITIONS.value["working_volume_ul"],
            shaking_rpm=None,
            inoculum_cells_per_strain=SELECTED_CONDITIONS.value[
                "inoculum_cells_per_strain"
            ],
            endpoint=EndpointRule.fixed_duration,
            provenance=[SCREEN_VESSEL, SELECTED_CONDITIONS],
            provenance_gaps=[
                ProvenanceGap(
                    field="shaking_rpm",
                    reason=ProvenanceGapReason.not_reported_by_primary,
                    note="neither the screen nor the optimization section says whether "
                    "the 96-well cultures were shaken",
                )
            ],
        )

    @staticmethod
    def _pre_culture() -> PreCulture:
        """The frozen pool thawed, counted and diluted straight into the assay."""
        return PreCulture(
            source=PreCultureSource.frozen_stock,
            source_label="thawed, counted using a hemocytometer, and diluted",
            provenance=[SCREEN_MEDIUM, FROZEN_POOL],
        )

    def _environment(self, perturbations: list[Any]) -> CultureEnvironment:
        """YPGal at 30 C for 48 h in the screen's culture format, with ``perturbations``."""
        return CultureEnvironment(
            media=YP_GALACTOSE_LIQUID,
            temperature=Temperature(value=TEMPERATURE_C.value),
            perturbations=perturbations,
            duration_hours=SELECTED_CONDITIONS.value["duration_hours"],
            duration_generations=None,
            culture_format=self._culture_format(),
            pre_culture=self._pre_culture(),
            auxotroph_supplements=None,
            provenance_gaps=[
                ProvenanceGap(
                    field="duration_generations",
                    reason=ProvenanceGapReason.not_reported_by_primary,
                    note="the paper states the 48 h but no number of doublings",
                ),
                ProvenanceGap(
                    field="auxotroph_supplements",
                    reason=ProvenanceGapReason.not_reported_by_primary,
                    note="the medium is YP, a complex medium, and the paper names no "
                    "supplement for the background's auxotrophies",
                ),
            ],
        )

    def _treated(self, compound: Compound, dose: Concentration) -> CultureEnvironment:
        """The condition: the compound at its Data Set 4 dose, in DMSO."""
        return self._environment(
            [
                SmallMoleculePerturbation(
                    compound=compound, concentration=dose, solvent=self._vehicle()
                )
            ]
        )

    def _control(self) -> CultureEnvironment:
        """The DMSO control the z-score is taken against: DMSO alone at 1%."""
        return self._environment(
            [
                SmallMoleculePerturbation(
                    compound=resolved_compound("DMSO"),
                    concentration=Concentration(
                        value=DMSO_CONTROL_PERCENT.value,
                        unit=ConcentrationUnit.percent,
                        basis=DoseBasis.fixed,
                    ),
                )
            ]
        )

    @staticmethod
    def _genome_reference() -> StrainReferenceGenome:
        """The pool's typed SGA-progeny background (shared with Vanacloig 2022)."""
        return StrainReferenceGenome(
            species="Saccharomyces cerevisiae",
            strain=LIBRARY_STRAIN,
            ploidy="haploid",
            background=library_background(),
            assembly_set=SGD_S288C_R64,
            assembly_accession=SGD_S288C_R64_GENBANK_ACCESSION,
        )

    @staticmethod
    def _uncertainty_gaps() -> list[ProvenanceGap]:
        return [
            ProvenanceGap(
                field=field,
                reason=ProvenanceGapReason.not_reported_by_primary,
                note=_SE_GAP_NOTE,
            )
            for field in _UNCERTAINTY_FIELDS
        ]

    def _reference(self, screen: str) -> StrainEnvironmentResponseExperimentReference:
        """The screen's DMSO reference profile, z = 0 by construction."""
        return StrainEnvironmentResponseExperimentReference(
            dataset_name=self.name,
            genome_reference=self._genome_reference(),
            environment_reference=self._control(),
            phenotype_reference=EnvironmentResponsePhenotype(
                measurement_type=Z_SCORE_DEFINITION.value,
                assay_type=ASSAY.value,
                environment_response=0.0,
                n_samples=None,
                sample_unit=None,
                units=_UNITS_TEXT,
                screen_id=screen,
                provenance_gaps=[
                    *self._uncertainty_gaps(),
                    ProvenanceGap(
                        field="n_samples",
                        reason=ProvenanceGapReason.not_reported_by_primary,
                        note="the reference is the mean over the screen's DMSO control "
                        "profiles; the Note gives 5724 control profiles across both "
                        "screens and no per-screen count",
                    ),
                    ProvenanceGap(
                        field="sample_unit",
                        reason=ProvenanceGapReason.not_reported_by_primary,
                        note="no per-screen DMSO profile count, so no sample unit",
                    ),
                ],
            ),
        )

    def _phenotype_template(self, row: ConditionRow) -> EnvironmentResponsePhenotype:
        """The per-condition phenotype; only ``environment_response`` varies per strain."""
        return EnvironmentResponsePhenotype(
            measurement_type=Z_SCORE_DEFINITION.value,
            assay_type=ASSAY.value,
            environment_response=0.0,
            n_samples=N_SAMPLES_DERIVATION.value_int,
            sample_unit=SampleUnit.technical_replicate,
            units=_UNITS_TEXT,
            screen_id=row.screen,
            replicate_id=row.condition_id,
            provenance_gaps=self._uncertainty_gaps(),
        )

    @staticmethod
    def _genotype(strain: ResolvedStrain) -> Genotype:
        """The ONE screened deletion; the constant background is on the reference."""
        return Genotype(
            perturbations=[
                BarcodedKanMxDeletionPerturbation(
                    systematic_gene_name=strain.systematic_name,
                    perturbed_gene_name=strain.perturbed_gene_name,
                    barcode=strain.pool.barcode,
                    collection=POOL_COLLECTION.value,
                    cassette=ARRAY_KANMX.value,
                )
            ]
        )

    def _intern_ref(self, obj: Any, hint: str, itxn: Any) -> Any:
        """Intern one constant sub-object and return what belongs in the record."""
        container: dict[str, Any] = {"value": obj.model_dump()}
        self._maybe_intern(container, "value", obj, hint, itxn)
        return container["value"]

    # ---- build ---------------------------------------------------------------- #
    @post_process
    def process(self) -> None:
        """Write one record per numeric cell of a kept strain and a kept condition."""
        from pydantic import TypeAdapter

        from torchcell.datamodels.schema import ExperimentType

        verify_raw_files(
            self.raw_dir,
            {release.filename: release.sha256 for release in RELEASE_FILES},
        )
        matrix = read_matrix(self._raw(MATRIX))
        conditions = read_conditions(self._raw(CONDITIONS))
        pool = read_pool(self._raw(POOL))
        strains = resolve_strains(matrix.strains, pool)
        labelled_twice = strains_labelled_twice(matrix, strains)
        validate: Callable[[Any], object] = TypeAdapter(ExperimentType).validate_python

        n_strains, n_conditions = matrix.values.shape
        scored = ~np.isnan(matrix.values)
        source_records = int(scored.sum())

        # --- condition rules, in order ---------------------------------------- #
        rows = [conditions[cid] for cid in matrix.condition_ids]
        doses = [
            parse_dose(row.concentration_cell, row.concentration_number_format)
            for row in rows
        ]
        condition_rule: list[str | None] = []
        for row, dose in zip(rows, doses, strict=True):
            if not row.in_final_dataset:
                condition_rule.append("condition_not_in_the_final_dataset")
            elif dose.kind == "range":
                condition_rule.append("dose_is_a_released_range")
            else:
                condition_rule.append(None)
        kept_columns = [j for j, rule in enumerate(condition_rule) if rule is None]
        kept_rows = list(range(len(strains)))
        column_scored = scored.sum(axis=0)

        rules: list[DropRule] = []
        for rule, description in (
            (
                "condition_not_in_the_final_dataset",
                "Data Set 4 marks the condition In_final_dataset = False and the Data "
                "Set 5 title counts 11,983 profiles, so the release's own metadata "
                "(not a filter of ours) excludes the column Data Set 5 still carries",
            ),
            (
                "dose_is_a_released_range",
                "the Data Set 4 dose cell is a range (e.g. '10-100-ug/ml'), the pilot "
                "dosing the Methods describe; a range names no single dose, the "
                "schema has no interval dose, and the required concentration cannot "
                "carry a typed gap (needs schema)",
            ),
        ):
            columns = [j for j, r in enumerate(condition_rule) if r == rule]
            rules.append(
                DropRule(
                    rule=rule,
                    scope="condition",
                    description=description,
                    n_records=int(column_scored[columns].sum()) if columns else 0,
                    items=[
                        f"{matrix.condition_ids[j]} {rows[j].name} ({doses[j].source})"
                        for j in columns
                    ],
                )
            )
        # --- interned constants ------------------------------------------------ #
        publication = Publication(doi=PAPER_DOI, doi_url=f"https://doi.org/{PAPER_DOI}")
        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        genotype_models = {i: self._genotype(strains[i]) for i in kept_rows}
        genotypes = {i: model.model_dump() for i, model in genotype_models.items()}
        template_genotype = genotype_models[kept_rows[0]]
        templates: dict[int, tuple[dict[str, Any], dict[str, Any]]] = {}
        ledger_compounds: dict[int, CompoundLedgerEntry] = {}
        with interned_env.begin(write=True) as itxn:
            reference_refs = {
                screen: self._intern_ref(self._reference(screen), screen, itxn)
                for screen in (SCREEN_RIKEN, SCREEN_NCI_NIH_GSK)
            }
            pub_ref = self._intern_ref(publication, PAPER_DOI, itxn)
            for j in tqdm(kept_columns, desc="Piotrowski 2017 conditions"):
                row = rows[j]
                compound = compound_for(row)
                ledger_compounds[j] = compound_ledger_entry(row, compound)
                concentration = doses[j].concentration
                if concentration is None:
                    raise RuntimeError(f"{row.condition_id}: a kept dose with no value")
                environment = self._treated(compound, concentration)
                # The template is a real model dump; only the genotype and the score
                # differ per record, and the environment is an interned pointer.
                base = StrainEnvironmentResponseExperiment(
                    dataset_name=self.name,
                    genotype=template_genotype,
                    environment=environment,
                    phenotype=self._phenotype_template(row),
                ).model_dump()
                env_dump = base["environment"]
                base["environment"] = self._intern_ref(environment, compound.name, itxn)
                templates[j] = (base, env_dump)

        # --- records ------------------------------------------------------------ #
        idx = 0
        txn = env.begin(write=True)
        for j in tqdm(kept_columns, desc="Piotrowski 2017 records"):
            base, env_dump = templates[j]
            phenotype = base["phenotype"]
            row = rows[j]
            validated = False
            for i in kept_rows:
                value = matrix.values[i, j]
                if math.isnan(value):
                    continue
                experiment = {
                    **base,
                    "genotype": genotypes[i],
                    "phenotype": {**phenotype, "environment_response": float(value)},
                }
                if not validated:
                    validate({**experiment, "environment": env_dump})
                    validated = True
                txn.put(
                    f"{idx}".encode(),
                    pickle.dumps(
                        {
                            "experiment": experiment,
                            "reference": reference_refs[row.screen],
                            "publication": pub_ref,
                        }
                    ),
                )
                idx += 1
                if idx % BATCH_SIZE == 0:
                    txn.commit()
                    txn = env.begin(write=True)
        txn.commit()
        env.close()
        interned_env.close()

        drop_log = DropLog(
            dataset=self.name,
            source_records=source_records,
            kept_records=idx,
            dropped_records=source_records - idx,
            rules=rules,
            released_cells=n_strains * n_conditions,
            not_scored_cells=int((~scored).sum()),
            not_scored_strains_by_screen=_not_scored_by_screen(matrix, rows),
        )
        accounted = sum(rule.n_records for rule in rules)
        if accounted != drop_log.dropped_records:
            raise RuntimeError(
                f"drop accounting mismatch: rules total {accounted}, "
                f"{drop_log.dropped_records} records missing from the build"
            )
        self._write_json("dropped_records.json", drop_log)
        self._write_json(
            "conditions.json",
            _condition_ledger(
                self.name,
                matrix,
                rows,
                doses,
                condition_rule,
                ledger_compounds,
                pool,
                labelled_twice,
            ),
        )
        self._write_json("n_samples_derivation.json", N_SAMPLES_DERIVATION.derivation)
        self._write_json("sourced_values.json", SourcedValues(values=sourced_values()))
        log.info("Wrote %d Piotrowski 2017 environment-response experiments", idx)

    def _write_json(self, filename: str, model: BaseModel) -> None:
        with open(osp.join(self.preprocess_dir, filename), "w") as handle:
            handle.write(model.model_dump_json(indent=2))

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside ``process()`` for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in ``process()``."""
        raise NotImplementedError


def _not_scored_by_screen(
    matrix: ReleasedMatrix, rows: list[ConditionRow]
) -> dict[str, dict[str, int]]:
    """Per screen, each strain with ``NA`` cells and how many of the screen's columns."""
    out: dict[str, dict[str, int]] = {}
    for screen in (SCREEN_RIKEN, SCREEN_NCI_NIH_GSK):
        columns = [j for j, row in enumerate(rows) if row.screen == screen]
        missing = np.isnan(matrix.values[:, columns]).sum(axis=1)
        out[screen] = {
            matrix.strains[i]: int(n) for i, n in enumerate(missing) if n > 0
        }
    return out


def _condition_ledger(
    dataset: str,
    matrix: ReleasedMatrix,
    rows: list[ConditionRow],
    doses: list[ParsedDose],
    condition_rule: list[str | None],
    compounds: dict[int, CompoundLedgerEntry],
    pool: list[PoolStrain],
    labelled_twice: dict[str, list[str]],
) -> ConditionLedger:
    """Every released condition with its fate and identity route, plus the pool census."""
    scored = (~np.isnan(matrix.values)).sum(axis=0)
    entries = [
        ConditionLedgerRow(
            condition_id=row.condition_id,
            matrix_label=matrix.condition_labels[j],
            name=row.name,
            collection=row.collection,
            screen=row.screen,
            source_dose=doses[j].source,
            in_high_confidence_set=row.in_high_confidence_set,
            percent_growth=row.percent_growth,
            n_scored_strains=int(scored[j]),
            drop_rule=condition_rule[j],
            compound=compounds.get(j),
        )
        for j, row in enumerate(rows)
    ]
    kept = list(compounds.values())
    routes = Counter(entry.route for entry in kept)
    curated_disagree = sorted(
        f"{rows[j].condition_id} {rows[j].name}: table {entry.stored_inchikey} vs "
        f"released InChI {entry.released_inchi_key}"
        for j, entry in compounds.items()
        if entry.route == "curated_table"
        and entry.stored_inchikey is not None
        and entry.released_inchi_key is not None
        and entry.stored_inchikey.split("-")[0]
        != entry.released_inchi_key.split("-")[0]
    )
    both = [
        entry
        for entry in kept
        if entry.released_inchi_key is not None
        and entry.released_smiles_key is not None
    ]
    labels = set(matrix.strains)
    return ConditionLedger(
        dataset=dataset,
        conditions=entries,
        route_counts=dict(sorted(routes.items())),
        curated_key_disagreements=curated_disagree,
        released_key_disagreements=sum(
            1 for entry in both if entry.released_inchi_key != entry.released_smiles_key
        ),
        released_key_comparisons=len(both),
        absent_pool_strains=sorted(
            f"{entry.common_name} ({entry.orf})"
            for entry in pool
            if entry.common_name not in labels and entry.orf not in labels
        ),
        strains_labelled_twice=labelled_twice,
    )


# --------------------------------------------------------------------------- #
# n_samples: the conservative lower end of 1..3 (no per-profile count is released)
# --------------------------------------------------------------------------- #
class NSamplesDerivation(BaseModel):
    """The derived ``n_samples`` and the typed record that justifies it."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    derivation: StatDerivation

    @property
    def value_int(self) -> int:
        """The chosen count as the integer ``n_samples`` takes."""
        return int(self.derivation.value)


N_SAMPLES_DERIVATION = NSamplesDerivation(
    derivation=StatDerivation(
        field="n_samples",
        method=DerivationMethod.conservative_low,
        value=1.0,
        range_low=1.0,
        range_high=float(DESIGN_TECHNICAL_REPLICATES.value),
        provenance=REPLICATE_COLLAPSING.provenance,
        rationale=(
            "Each condition was amplified as three technical (PCR) replicates "
            "(DESIGN_TECHNICAL_REPLICATES, FIG2_REPLICATES), but a released profile is "
            "the mean of only the replicates the batch-effect filter kept and the "
            "correlation graph joined (REPLICATE_COLLAPSING, thresholds 0.7 RIKEN and "
            "0.5 NCI/NIH/GSK), so the count is between 1 and 3 per profile and is not "
            "released. No dispersion is released either, so no back-solve exists; the "
            "CLAUDE.md rule takes the conservative lower end"
        ),
    )
)


class SourcedValues(BaseModel):
    """``preprocess/sourced_values.json``: every module-level ``SourcedValue``."""

    model_config = ConfigDict(extra="forbid")

    values: dict[str, SourcedValue]


def sourced_values() -> dict[str, SourcedValue]:
    """Every module-level ``SourcedValue`` by its constant name."""
    return {
        name: value
        for name, value in sorted(globals().items())
        if isinstance(value, SourcedValue)
    }


# --------------------------------------------------------------------------- #
# Verification
# --------------------------------------------------------------------------- #
def verify_build(
    dataset_root: str, *, data_root: str | None = None
) -> VerificationReport:
    """L0-L4 over the built store, as the environment-response family run verifies it.

    The registry entry (``runners.ENVIRONMENT_RESPONSE_DATASETS[DATASET_SLUG]``) is the
    one specification; ``dataset_root`` must be the store that entry names.
    """
    from torchcell.verification import runners

    root = data_root or _data_root()
    expected = osp.join(
        root, runners.ENVIRONMENT_RESPONSE_DATASETS[DATASET_SLUG]["root"]
    )
    if osp.abspath(dataset_root) != osp.abspath(expected):
        raise ValueError(f"{dataset_root} is not the registered store {expected}")
    return runners.verify_environment_response_entry(DATASET_SLUG, root)


def main() -> None:
    """Build or load the dev store (``deposit_raw_mirror`` is a one-time step before it)."""
    from dotenv import load_dotenv

    load_dotenv()
    data_root = os.environ["DATA_ROOT"]
    root = osp.join(data_root, "data/torchcell", DATASET_SLUG)
    dataset = EnvChemgenPiotrowski2017Dataset(root=root)
    print(f"len = {len(dataset)}")
    print(dataset[0])


if __name__ == "__main__":
    main()
