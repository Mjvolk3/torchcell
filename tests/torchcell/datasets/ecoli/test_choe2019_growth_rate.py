# tests/torchcell/datasets/ecoli/test_choe2019_growth_rate.py
# [[tests.torchcell.datasets.ecoli.test_choe2019_growth_rate]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_choe2019_growth_rate.py
"""Choe 2019 loader (``torchcell.datasets.ecoli.choe2019_growth_rate``).

The synthetic tests write a Source Data workbook with the Fig. 2c and Fig. 2f sheets'
exact layout and three variant tables with each of the three column layouts the release
uses, then drive the readers, all three cross-checks, the phenotype and genotype
builders and a full ``process()`` for both dataset classes into ``tmp_path``. The
verifier's genome is an in-test object and the loaders read no annotation (#1045 part
2), so nothing reads ``$DATA_ROOT``.

One assertion here is a MEASUREMENT of the schema rather than of the paper:
``BacterialStrainBackground.alleles`` defaults to ``[]``, so a ``ProvenanceGap`` on it is
refused, which is why MS56's 55 unenumerated regions live in a ledger instead of a typed
gap. That refusal is pinned so a future schema change that makes the gap expressible
fails this test and sends someone back to the background.

The ``@pytest.mark.data`` tests read the real mirror and pin what the note states: the
Fig. 2c strain means, that the two deletion strains reproduce the paper's own "80% of
eMS57" claim, that the two panels' eMS57/MS56 ratios disagree by a factor near 2, the
exact 117 / 101 / 145 call counts of the three variant tables, and that the 21 named
genes are the contiguous ``b2721``-``b2741`` run on the real MG1655 annotation.
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
import statistics
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest
from pydantic import ValidationError

import torchcell.datasets.ecoli.choe2019_growth_rate as ch
from torchcell.data import file_sha256
from torchcell.datamodels.schema import (
    AlleleEdit,
    BacterialAssemblySet,
    BacterialBackgroundAllele,
    BacterialGeneNamespace,
    BacterialReleasedStrainPerturbation,
    BacterialStrainBackground,
    ReferenceGenome,
    ReleasedStrainKind,
    SampleUnit,
    SequenceVariantPerturbation,
    StrainReferenceGenome,
    UncertaintyType,
)
from torchcell.sequence.genome.base import GeneNameResolution, GeneNameStatus
from torchcell.verification.report import FindingClass, finding_class_of
from torchcell.verification.sourced import ProvenanceGap, ProvenanceGapReason

MG1655_ASSEMBLY: BacterialAssemblySet = "ecoli_K12_MG1655_ASM584v2"
BW25113_ASSEMBLY: BacterialAssemblySet = "ecoli_K12_BW25113_ASM75055v1"
MG1655_NAMESPACE: BacterialGeneNamespace = "ecoli_k12_mg1655_bnumber"
BW25113_NAMESPACE: BacterialGeneNamespace = "ecoli_k12_bw25113_locus_tag"
#: The shared verifier's genome-reference rules. Under the owner rule (#1057) a gene
#: symbol stored as released need not be a locus of the annotation, so these may read it
#: as a failure until the verifier changes (a separate branch); nothing else may fail.
GENOME_REFERENCE_RULES = frozenset(
    {"canonical_gene_names", "gene_containment_sgd", "current_genome_genes"}
)

#: The Fig. 2c rows, verbatim from the real sheet, in sheet order.
FIG2C: dict[str, tuple[float, float, float]] = {
    "MG1655": (0.5828, 0.597, 0.58785),
    "MS56": (0.1834, 0.1876, 0.176),
    "eMS57": (0.5148, 0.5152, 0.5103),
    "Δ21 kb": (0.4073, 0.3891, 0.3993),
    "ΔrpoS": (0.4158, 0.4215, 0.3897),
}
#: The Fig. 2f rows, verbatim from the real sheet, in sheet order.
FIG2F: dict[str, tuple[float, float, float]] = {
    "eMS57/MS56": (1.421071401277824, 1.4733489228066399, 1.3025292733124838),
    "ilvN*/ilvNWT": (1.0073143271396994, 1.0033412933448602, 0.9465918963112657),
    "cspC*/cspCWT": (1.193092236328973, 1.1627038768390179, 1.0976646038935256),
    "yifB*/yifBWT": (0.6154782959906288, 0.7074886092907978, 0.7420058326965318),
}


# --------------------------------------------------------------------------- #
# Synthetic workbooks
# --------------------------------------------------------------------------- #
def write_source_data(path: Path) -> Path:
    """Write a Source Data workbook with the Fig. 2c and Fig. 2f sheets."""
    book = openpyxl.Workbook()
    fig2c = book.active
    assert fig2c is not None
    fig2c.title = ch.GROWTH_SHEET
    fig2c.cell(row=1, column=2, value=ch.GROWTH_SHEET_TITLE)
    for column, header in enumerate(ch.GROWTH_HEADERS, start=2):
        fig2c.cell(row=3, column=column, value=header)
    for offset, (strain, rates) in enumerate(FIG2C.items()):
        fig2c.cell(row=4 + offset, column=2, value=strain)
        for column, value in enumerate(rates, start=3):
            fig2c.cell(row=4 + offset, column=column, value=value)

    fig2f = book.create_sheet(ch.RATIO_SHEET)
    fig2f.cell(row=1, column=2, value=ch.RATIO_SHEET_TITLE)
    fig2f.cell(row=3, column=2, value="Mutated/Parental")
    for column, header in enumerate(("Set1", "Set2", "Set3"), start=3):
        fig2f.cell(row=3, column=column, value=header)
    for offset, (label, ratios) in enumerate(FIG2F.items()):
        fig2f.cell(row=4 + offset, column=2, value=label)
        for column, value in enumerate(ratios, start=3):
            fig2f.cell(row=4 + offset, column=column, value=value)
    book.save(path)
    book.close()
    return path


#: ``(gene, position, type, ref, allele, aa)`` rows for the Supplementary Data 2 layout,
#: whose gene cell comes FIRST. Two rows share ``yafF`` and one is intergenic.
_MS56_ROWS = (
    ("intergenic1", 12086, "SNV", "C", "T", "-"),
    ("ilvH", 82242, "SNV", "G", "T", "Gly21Cys"),
    ("yafF", 240000, "SNV", "A", "G", "Thr5Ala"),
    ("yafF", 240100, "SNV", "C", "T", "Pro9Leu"),
)
#: Rows for the Supplementary Data 3 / 4 layout, whose POSITION cell comes first and
#: whose intergenic calls are written as a bare dash or as the word.
_OTHER_ROWS = (
    (6919, "yaaJ", "C", "T", "SNV", "Trp347*"),
    (131453, "-", "C", "A", "SNV", None),
    (208960, "intergenic", "A", "G", "SNV", None),
)


def write_variant_tables(directory: Path) -> None:
    """Write the three variant tables, each with a stray single-cell row below it."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Sheet1"
    sheet.cell(row=1, column=1, value="Supplementary Data 2.")
    for column, header in enumerate(
        ("Gene", "Position", "Type", "Ref", "Allele", "AA change"), start=1
    ):
        sheet.cell(row=2, column=column, value=header)
    sheet.cell(row=2, column=7, value="Allelic frequency (%)")
    for column, day in enumerate((0, 62), start=7):
        sheet.cell(row=3, column=column, value=day)
    for offset, row in enumerate(_MS56_ROWS):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=4 + offset, column=column, value=value)
        sheet.cell(row=4 + offset, column=7, value=0)
        sheet.cell(row=4 + offset, column=8, value=100)
    # the stray cell the real workbook carries twelve rows past its last call
    sheet.cell(row=20, column=10, value=3)
    book.save(directory / ch.VARIANTS_MS56.name)
    book.close()

    for artifact, header_rows in ((ch.VARIANTS_MG1655, 3), (ch.VARIANTS_EXTRA, 4)):
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Sheet1"
        sheet.cell(row=1, column=1, value=artifact.si_label)
        for column, header in enumerate(
            ("Position", "Gene", "Ref", "Allele", "Type", "AA change"), start=1
        ):
            sheet.cell(row=2, column=column, value=header)
        sheet.cell(row=header_rows, column=7, value="Replicate 1")
        for offset, other in enumerate(_OTHER_ROWS):
            for column, value in enumerate(other, start=1):
                sheet.cell(row=header_rows + 1 + offset, column=column, value=value)
            sheet.cell(row=header_rows + 1 + offset, column=7, value=11.1)
        sheet.cell(row=header_rows + 12, column=9, value=7)
        book.save(directory / artifact.name)
        book.close()


@pytest.fixture
def source_data(tmp_path: Path) -> Path:
    """A Source Data workbook in ``tmp_path``."""
    return write_source_data(tmp_path / ch.SOURCE_DATA.name)


# --------------------------------------------------------------------------- #
# The declared module constants
# --------------------------------------------------------------------------- #
def test_the_large_deletion_genotype_string_expands_to_the_declared_symbols() -> None:
    """The 21 symbols are exactly what Supplementary Table 3's genotype string names."""
    inner = ch.LARGE_DELETION_GENOTYPE_STATEMENT.removeprefix("MS56, (").removesuffix(
        ")::kan"
    )
    runs = inner.split("-")
    assert runs == [
        "hycEDCBA",
        "hypABCDE",
        "fhlA",
        "ygbA",
        "mutS",
        "pphB",
        "ygbIJKLMN",
        "rpoS",
    ]
    expanded: list[str] = []
    for run in runs:
        stem = run[:3] if run[:3] in {"hyc", "hyp", "ygb"} else run
        if stem == run:
            expanded.append(run)
            continue
        expanded += [f"{stem}{letter}" for letter in run[3:]]
    assert set(expanded) == set(ch.LARGE_DELETION_SYMBOLS)
    assert len(ch.LARGE_DELETION_SYMBOLS) == 21
    assert ch.SOURCED_VALUES["large_deletion_gene_count"].value == 21
    assert ch.RPOS_SYMBOL in ch.LARGE_DELETION_SYMBOLS


def test_every_declared_count_and_quote_is_present() -> None:
    assert ch.GROWTH_EXPECTED_RECORDS == 8
    assert ch.GROWTH_EXPECTED_REFERENCES == 7
    assert ch.TF_EXPECTED_RECORDS == 2
    assert ch.KEIO_PANEL_STRAINS == 62
    assert [fitness for _, fitness in ch.KEIO_FITNESS] == [0.724, 0.698]
    # both percentages come from one legend sentence, so one quote justifies both
    assert (
        ch.SOURCED_VALUES["keio_ydas_fitness"].quote
        == ch.SOURCED_VALUES["keio_abgr_fitness"].quote
    )
    assert "7 2 . 4" in ch.SOURCED_VALUES["keio_ydas_fitness"].quote
    assert "6 9 . 8" in ch.SOURCED_VALUES["keio_ydas_fitness"].quote
    assert "did not result from ydaS deletion" in (
        ch.SOURCED_VALUES["keio_ydas_fitness"].quote
    )
    assert ch.SOURCED_VALUES["parent_regions_deleted"].value == 55
    assert ch.MS56_GENOTYPE_STATEMENT in (
        ch.SOURCED_VALUES["parent_genotype_statement"].quote
    )
    assert {name for name in ch.SOURCED_VALUES} >= {
        "cassette",
        "construction",
        "media_recipe",
        "recovery_fraction_of_evolved",
    }


def test_the_raw_mirror_pins_every_file_either_class_reads() -> None:
    assert set(ch.DATA_SHA256) == {raw.name for raw in ch.RAW_FILES}
    assert len(ch.RAW_FILES) == 5
    for raw in ch.RAW_FILES:
        assert len(raw.sha256) == 64
        assert raw.mirror_relpath == f"data/{raw.name}"
        assert raw.source_url.endswith(raw.pmc_object)
        assert raw.retrieval.sha256 == raw.sha256
        assert raw.retrieval.params["key"].endswith(raw.pmc_object)
    assert ch.SUPPLEMENTARY_PDF.name.endswith(".pdf")
    # the four released files no class in this schema can carry are declared, not hidden
    assert len(ch.NOT_MIRRORED) == 4


def test_the_assembly_set_literal_matches_the_registry() -> None:
    """The annotated constant and the registry are one value, so they cannot drift."""
    from torchcell.datasets.bacteria_common import BACTERIAL_ASSEMBLY_SETS

    assert ch.MG1655_ASSEMBLY_SET == BACTERIAL_ASSEMBLY_SETS["MG1655"]
    assert ch.parent_background().assembly_set == BACTERIAL_ASSEMBLY_SETS["MG1655"]


def test_the_medium_is_the_methods_recipe() -> None:
    media = ch.M9_GLUCOSE_CHOE2019
    assert media.base_medium == "M9"
    assert media.is_synthetic is True
    values: dict[str, float] = {}
    for component in media.components:
        assert component.concentration is not None
        value = component.concentration.value
        assert value is not None
        values[component.compound.name] = value
    assert values == {
        "disodium hydrogen phosphate": 47.75,
        "potassium dihydrogen phosphate": 22.04,
        "sodium chloride": 8.56,
        "ammonium chloride": 18.70,
        "magnesium sulfate": 2.0,
        "calcium chloride": 0.1,
        "D-glucose": 2.0,
    }
    for component in media.components:
        assert component.provenance == [ch.SOURCED_VALUES["media_recipe"]]


def test_the_two_environments_differ_only_in_what_the_paper_states() -> None:
    growth = ch.growth_environment()
    keio = ch.keio_environment()
    assert growth.temperature is None
    assert growth.gapped_fields() == {"temperature"}
    assert keio.temperature is not None
    assert keio.temperature.value == 37.0
    assert keio.gapped_fields() == set()
    assert growth.perturbations == []
    assert keio.perturbations == []
    assert growth.media == keio.media


# --------------------------------------------------------------------------- #
# Reading the Source Data workbook
# --------------------------------------------------------------------------- #
def test_read_growth_panel_reads_every_strain_in_sheet_order(source_data: Path) -> None:
    panel = ch.read_growth_panel(str(source_data))
    assert [row.strain for row in panel.rows] == list(ch.GROWTH_STRAIN_ORDER)
    for strain, rates in FIG2C.items():
        assert panel.rates(strain).replicates == rates
        assert panel.rates(strain).n_samples == 3
    assert panel.rates("MS56").mean == pytest.approx(statistics.fmean(FIG2C["MS56"]))


def test_read_growth_panel_refuses_a_changed_title(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    book[ch.GROWTH_SHEET].cell(row=1, column=2, value="Something else")
    path = tmp_path / "retitled.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="title is"):
        ch.read_growth_panel(str(path))


def test_read_growth_panel_refuses_changed_headers(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    book[ch.GROWTH_SHEET].cell(row=3, column=3, value="Replicate 1")
    path = tmp_path / "reheadered.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="headers are"):
        ch.read_growth_panel(str(path))


def test_read_growth_panel_refuses_a_reordered_strain_block(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    sheet = book[ch.GROWTH_SHEET]
    sheet.cell(row=4, column=2, value="MS56")
    sheet.cell(row=5, column=2, value="MG1655")
    path = tmp_path / "reordered.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="strain order is"):
        ch.read_growth_panel(str(path))


def test_read_growth_panel_refuses_a_missing_sheet(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    path = tmp_path / "empty.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="no sheet"):
        ch.read_growth_panel(str(path))


def test_read_reconstruction_ratios_reads_the_refused_panel(source_data: Path) -> None:
    ratios = ch.read_reconstruction_ratios(str(source_data))
    assert list(ratios) == list(ch.RATIO_ROW_ORDER)
    assert ratios["cspC*/cspCWT"] == pytest.approx(FIG2F["cspC*/cspCWT"])


def test_read_reconstruction_ratios_refuses_a_reordered_sheet(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    book[ch.RATIO_SHEET].cell(row=4, column=2, value="yifB*/yifBWT")
    path = tmp_path / "reordered.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="row order is"):
        ch.read_reconstruction_ratios(str(path))


def test_read_variant_table_drops_the_stray_cell_and_counts_intergenic_calls(
    tmp_path: Path,
) -> None:
    write_variant_tables(tmp_path)
    ms56 = ch.read_variant_table(
        str(tmp_path / ch.VARIANTS_MS56.name),
        ch.VARIANTS_MS56,
        header_rows=3,
        gene_column=0,
    )
    # four calls, NOT five: the stray single-cell row names no variant
    assert ms56.rows == len(_MS56_ROWS)
    assert ms56.intergenic_rows == 1
    assert ms56.loci_with_several_calls == {"yafF": 2}
    assert ms56.frequency_columns == ("0", "62")

    other = ch.read_variant_table(
        str(tmp_path / ch.VARIANTS_MG1655.name),
        ch.VARIANTS_MG1655,
        header_rows=3,
        gene_column=1,
    )
    assert other.rows == len(_OTHER_ROWS)
    # a bare dash and the bare word are both intergenic in these two layouts
    assert other.intergenic_rows == 2
    assert other.loci_with_several_calls == {}


# --------------------------------------------------------------------------- #
# The cross-checks on the released numbers
# --------------------------------------------------------------------------- #
def test_check_recovery_fraction_reproduces_the_papers_own_claim(
    source_data: Path,
) -> None:
    panel = ch.read_growth_panel(str(source_data))
    check = ch.check_recovery_fraction(panel)
    measured = check["measured_recovery_fraction_of_evolved"]
    assert measured["Δ21 kb"] == pytest.approx(0.7763, abs=1e-4)
    assert measured["ΔrpoS"] == pytest.approx(0.7966, abs=1e-4)
    assert check["parent_over_wild_type"] == pytest.approx(0.3094, abs=1e-4)


def test_check_recovery_fraction_refuses_a_rescaled_deletion_row(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    sheet = book[ch.GROWTH_SHEET]
    for column in (3, 4, 5):
        sheet.cell(row=7, column=column, value=0.2)
    path = tmp_path / "rescaled.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="more than .* from the stated"):
        ch.check_recovery_fraction(ch.read_growth_panel(str(path)))


def test_check_recovery_fraction_refuses_a_parent_that_is_not_impaired(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    sheet = book[ch.GROWTH_SHEET]
    # make MS56 match eMS57, which keeps the recovery fractions but breaks the reduction
    for column, value in enumerate(FIG2C["eMS57"], start=3):
        sheet.cell(row=5, column=column, value=value)
    for column, value in enumerate(FIG2C["eMS57"], start=3):
        sheet.cell(row=7, column=column, value=value * 0.8)
        sheet.cell(row=8, column=column, value=value * 0.8)
    path = tmp_path / "healthy_parent.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="not the severe reduction"):
        ch.check_recovery_fraction(ch.read_growth_panel(str(path)))


def test_check_panels_disagree_records_the_factor(source_data: Path) -> None:
    panel = ch.read_growth_panel(str(source_data))
    ratios = ch.read_reconstruction_ratios(str(source_data))
    check = ch.check_panels_disagree(panel, ratios)
    assert check["fig2c_evolved_over_parent_from_rates"] == pytest.approx(
        2.8159, abs=1e-4
    )
    assert check["fig2f_evolved_over_parent_released"] == pytest.approx(
        1.3990, abs=1e-4
    )
    assert check["factor"] == pytest.approx(2.013, abs=1e-3)


def test_check_panels_disagree_refuses_agreement(source_data: Path) -> None:
    """If the sheets ever agree, the refusal of Fig. 2f needs rereading, not reusing."""
    panel = ch.read_growth_panel(str(source_data))
    ratios = dict(ch.read_reconstruction_ratios(str(source_data)))
    ratios[ch.RATIO_ROW_ORDER[0]] = (2.8, 2.8, 2.8)
    with pytest.raises(ch.TableFormatError, match="no longer disagree"):
        ch.check_panels_disagree(panel, ratios)


# --------------------------------------------------------------------------- #
# Identifier resolution and the deleted region
# --------------------------------------------------------------------------- #
class _Locus:
    def __init__(self, symbol: str | None) -> None:
        self.symbol = symbol


class _Annotation:
    def __init__(self, loci: dict[str, _Locus]) -> None:
        self.loci = loci


class _Feature:
    def __init__(self, tag: str, start: int, end: int) -> None:
        self.attributes = {"locus_tag": [tag]}
        self.start = start
        self.end = end


class _Db:
    def __init__(self, features: list[_Feature]) -> None:
        self._features = features

    def features_of_type(self, kind: str) -> list[_Feature]:
        assert kind == "gene"
        return self._features


#: ``b2721``-``b2741`` in the order the real annotation carries them, 1 kb apart.
MG1655_TAGS = {
    symbol: f"b{2721 + offset}"
    for offset, symbol in enumerate(ch.LARGE_DELETION_SYMBOLS)
}
BW25113_TAGS = {"ydaS": "BW25113_1357", "abgR": "BW25113_1339"}


class FakeGenome:
    """Carries the deleted gene run, and optionally one extra gene inside it."""

    def __init__(
        self,
        tags: dict[str, str],
        assembly_set: str,
        *,
        extra_gene: tuple[str, int, int] | None = None,
    ) -> None:
        """One locus per symbol, laid out 1 kb apart from 1,000,000."""
        self.ASSEMBLY_SET = assembly_set
        self.genbank = _Annotation(
            {tag: _Locus(symbol) for symbol, tag in tags.items()}
        )
        features = [
            _Feature(tag, 1_000_000 + 1_000 * i, 1_000_500 + 1_000 * i)
            for i, tag in enumerate(tags.values())
        ]
        if extra_gene is not None:
            features.append(_Feature(*extra_gene))
        self.db = _Db(features)

    def resolve_gene_name(self, name: str) -> GeneNameResolution:
        """A locus tag or a carried symbol resolves to its tag; anything else retires."""
        by_symbol = {locus.symbol: tag for tag, locus in self.genbank.loci.items()}
        tag = name if name in self.genbank.loci else by_symbol.get(name)
        return GeneNameResolution(
            input_name=name,
            status=GeneNameStatus.CURRENT
            if tag is not None
            else GeneNameStatus.RETIRED,
            systematic_name=tag,
        )


def test_check_deleted_region_accepts_the_contiguous_run() -> None:
    genome = FakeGenome(MG1655_TAGS, MG1655_ASSEMBLY)
    region = ch.check_deleted_region(genome)  # type: ignore[arg-type]
    assert region["locus_tags"] == [MG1655_TAGS[s] for s in ch.LARGE_DELETION_SYMBOLS]
    assert region["genes_inside_window"] == 21
    assert region["unlisted_genes_inside_window"] == 0
    assert region["released_ms56_window_bp"] == 20965
    assert region["problems"] == []
    rule = ch.deleted_region_rule(genome)  # type: ignore[arg-type]
    assert (rule.name, rule.passed) == (ch.DELETED_REGION_RULE, True)


def test_check_deleted_region_reports_an_unlisted_gene_inside_the_window() -> None:
    """A gene inside the window the source does not list is a finding, never a gate."""
    genome = FakeGenome(
        MG1655_TAGS, MG1655_ASSEMBLY, extra_gene=("b9999", 1_005_100, 1_005_200)
    )
    region = ch.check_deleted_region(genome)  # type: ignore[arg-type]
    assert region["unlisted_genes_inside_window"] == 1
    assert "does not list as deleted" in region["problems"][0]
    rule = ch.deleted_region_rule(genome)  # type: ignore[arg-type]
    assert rule.passed is False
    assert finding_class_of(rule.name) is FindingClass.report


def test_check_deleted_region_reports_a_missing_gene_feature() -> None:
    genome = FakeGenome({"rpoS": "b2741"}, MG1655_ASSEMBLY)
    region = ch.check_deleted_region(genome)  # type: ignore[arg-type]
    assert region["problems"][0].startswith("no gene feature of the pinned assembly")
    assert "hycE" in region["problems"][0]


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
def test_fitness_phenotype_is_the_ratio_to_the_isogenic_parent(
    source_data: Path,
) -> None:
    panel = ch.read_growth_panel(str(source_data))
    parent = panel.rates("MS56")
    phenotype = ch.fitness_phenotype(panel.rates("ΔrpoS"), parent)
    assert phenotype.fitness == pytest.approx(
        statistics.fmean(FIG2C["ΔrpoS"]) / statistics.fmean(FIG2C["MS56"])
    )
    assert phenotype.fitness_uncertainty == pytest.approx(
        statistics.stdev(FIG2C["ΔrpoS"]) / statistics.fmean(FIG2C["MS56"])
    )
    assert phenotype.fitness_uncertainty_type is UncertaintyType.sample_sd
    assert phenotype.n_samples == 3
    assert phenotype.sample_unit is SampleUnit.biological_replicate
    # the stored SE carries the parent's spread too, so it is never the optimistic one
    assert phenotype.fitness_se is not None
    assert phenotype.fitness_uncertainty is not None
    assert phenotype.fitness_se > phenotype.fitness_uncertainty / math.sqrt(3)


def test_parent_phenotype_is_identically_one(source_data: Path) -> None:
    panel = ch.read_growth_panel(str(source_data))
    phenotype = ch.parent_phenotype(panel.rates("MS56"))
    assert phenotype.fitness == 1.0
    assert phenotype.fitness_uncertainty == pytest.approx(
        statistics.stdev(FIG2C["MS56"]) / statistics.fmean(FIG2C["MS56"])
    )


def test_keio_phenotype_carries_the_ratio_and_no_invented_statistics() -> None:
    phenotype = ch.keio_phenotype(0.724)
    assert phenotype.fitness == 0.724
    assert phenotype.n_samples is None
    assert phenotype.sample_unit is None
    assert phenotype.fitness_uncertainty is None
    assert phenotype.fitness_se is None


def test_parent_background_states_the_genotype_and_types_no_allele() -> None:
    background = ch.parent_background()
    assert background.name == "MS56"
    assert background.reference_strain == "MG1655"
    assert background.parents == ["MG1655"]
    assert background.genotype_statement == ch.MS56_GENOTYPE_STATEMENT
    assert background.alleles == []
    assert background.is_fully_sourced is True


def test_a_provenance_gap_on_alleles_is_refused_by_the_schema() -> None:
    """MEASURED: the untyped absence has no typed home, which is why a ledger holds it.

    ``ProvenanceGapMixin`` requires a gapped field to be ``None`` and ``alleles``
    defaults to ``[]``, so MS56's 55 unenumerated regions cannot be declared on the
    background. If a schema change ever makes this expressible, this test fails and the
    background should carry the gap instead of ``genotype_gaps.json``.
    """
    with pytest.raises(ValidationError, match="is not None"):
        BacterialStrainBackground(
            name="MS56",
            reference_strain="MG1655",
            assembly_set=MG1655_ASSEMBLY,
            genotype_statement=ch.MS56_GENOTYPE_STATEMENT,
            provenance=[ch.SOURCED_VALUES["parent_strain"]],
            provenance_gaps=[
                ProvenanceGap(
                    field="alleles",
                    reason=ProvenanceGapReason.not_reported_by_primary,
                    note="the 55 regions are never enumerated",
                )
            ],
        )
    assert ch.BACKGROUND_ABSENCE["field"] == "alleles"
    assert ch.BACKGROUND_ABSENCE["model"] == "BacterialStrainBackground"


def test_no_perturbation_leaf_holds_a_called_bacterial_variant() -> None:
    """MEASURED: the two refusals the module's finding rests on, on real identifiers.

    The reconstructed Fig. 2f strains are DESIGNED single-allele strains with exactly
    specified alleles, so neither refusal is an evolved-population artefact; since #1057
    they are stored on the released-strain leaf instead (see the reconstruction tests).
    """
    # Since P1 (ontology-unify) the refusal is the namespace agreement check: the leaf
    # defaults gene_namespace to S288C, and a b-number is an MG1655 identifier.
    with pytest.raises(
        ValidationError,
        match="'b0078' is a ecoli_k12_mg1655_bnumber tag but gene_namespace is",
    ):
        SequenceVariantPerturbation(
            systematic_gene_name="b0078", perturbed_gene_name="ilvH", strain_id="eMS57"
        )
    with pytest.raises(ValidationError, match="valid boolean"):
        BacterialBackgroundAllele(
            systematic_gene_name="b1823",
            gene_namespace=MG1655_NAMESPACE,
            gene_name="cspC",
            allele_name="cspC(G37A)",
            edit=AlleleEdit.sequence_variant,
            functional=None,  # type: ignore[arg-type]
            provenance=[ch.SOURCED_VALUES["parent_strain"]],
        )


def test_reconstructed_alleles_are_the_fig2f_rows_and_quote_supplementary_table_3() -> (
    None
):
    assert [a.ratio_label for a in ch.RECONSTRUCTED_ALLELES] == list(
        ch.RATIO_ROW_ORDER[1:]
    )
    assert [a.strain for a in ch.RECONSTRUCTED_ALLELES] == [
        "ilvNmut",
        "cspCmut",
        "yifBmut",
    ]
    assert ch.SOURCED_VALUES["cspCmut_genotype"].quote == (
        "cspCmut</td><td rowspan=1 colspan=1>MS56, cspC::cspC(G37A)-kanR</td>"
    )
    assert ch.SOURCED_VALUES["yifBWT_genotype"].value == "MS56, yifB::yifB-kanR"
    assert ch.SOURCED_VALUES["reconstruction_replicates"].value == 3


def test_reconstructed_genotype_stores_the_allele_as_released() -> None:
    allele = ch.RECONSTRUCTED_ALLELES[1]
    (perturbation,) = ch.reconstructed_genotype(allele).perturbations
    assert isinstance(perturbation, BacterialReleasedStrainPerturbation)
    assert (
        perturbation.systematic_gene_name,
        perturbation.perturbed_gene_name,
        perturbation.strain_label,
        perturbation.identifier_kind,
        perturbation.released_kind,
        perturbation.cassette,
        perturbation.gene_namespace,
    ) == (
        "cspC",
        "cspC(G37A)",
        "cspCmut",
        "gene symbol",
        ReleasedStrainKind.sequence_allele,
        "kanR",
        MG1655_NAMESPACE,
    )


def test_control_background_names_the_kan_marked_wild_type_allele_strain() -> None:
    background = ch.control_background(ch.RECONSTRUCTED_ALLELES[0])
    assert (background.name, background.parents, background.genotype_statement) == (
        "ilvNWT",
        ["MS56"],
        "MS56, ilvN::ilvN-kan",
    )
    assert background.alleles == []
    assert background.provenance is not None
    assert background.provenance[0].quote == (
        "ilvNWT</td><td rowspan=1 colspan=1>MS56, ilvN::ilvN-kan</td>"
    )


def test_reconstruction_phenotype_is_the_mean_and_sample_sd_of_the_released_ratios() -> (
    None
):
    ratios = FIG2F["yifB*/yifBWT"]
    phenotype = ch.reconstruction_phenotype(ratios)
    assert phenotype.fitness == pytest.approx(statistics.fmean(ratios))
    assert phenotype.fitness == pytest.approx(0.688324, abs=1e-6)
    assert phenotype.fitness_uncertainty == pytest.approx(statistics.stdev(ratios))
    assert phenotype.fitness_uncertainty_type is UncertaintyType.sample_sd
    assert phenotype.n_samples == 3
    assert phenotype.sample_unit is SampleUnit.biological_replicate
    control = ch.control_phenotype()
    assert (control.fitness, control.n_samples, control.fitness_uncertainty) == (
        1.0,
        None,
        None,
    )


def test_build_growth_genotype_writes_one_deletion_per_named_gene() -> None:
    """#1045 part 2: each deletion keys the symbol the release names, as released."""
    genotype = ch.build_growth_genotype(
        ch.LARGE_DELETION_SYMBOLS, ch.LARGE_DELETION_LABEL
    )
    assert len(genotype.perturbations) == 21
    assert set(genotype.systematic_gene_names) == set(ch.LARGE_DELETION_SYMBOLS)
    for perturbation in genotype.perturbations:
        assert isinstance(perturbation, BacterialReleasedStrainPerturbation)
        assert perturbation.released_kind is ReleasedStrainKind.deletion
        assert perturbation.strain_label == ch.LARGE_DELETION_LABEL
        assert perturbation.cassette == ch.KAN_CASSETTE
        assert perturbation.gene_namespace == MG1655_NAMESPACE
    single = ch.build_growth_genotype((ch.RPOS_SYMBOL,), ch.RPOS_DELETION_LABEL)
    assert len(single.perturbations) == 1
    assert single.perturbations[0].perturbed_gene_name == "rpoS"


def test_build_keio_genotype_names_the_collection() -> None:
    genotype = ch.build_keio_genotype("ydaS")
    assert len(genotype.perturbations) == 1
    perturbation = genotype.perturbations[0]
    assert isinstance(perturbation, BacterialReleasedStrainPerturbation)
    assert perturbation.systematic_gene_name == "ydaS"
    assert perturbation.strain_label == "ΔydaS"
    assert perturbation.collection == ch.KEIO_COLLECTION
    assert perturbation.cassette is None
    assert perturbation.gene_namespace == BW25113_NAMESPACE


# --------------------------------------------------------------------------- #
# End-to-end builds on the synthetic workbooks
# --------------------------------------------------------------------------- #
def _reference_genome(strain: str, assembly_set: str, accession: str) -> Any:
    def build(
        reference_strain: str, *, background: Any = None, data_root: Any = None
    ) -> ReferenceGenome:
        if background is None:
            return ReferenceGenome(
                species="Escherichia coli",
                strain=strain,
                ploidy="haploid",
                assembly_set=assembly_set,  # type: ignore[arg-type]
                assembly_accession=accession,
            )
        return StrainReferenceGenome(
            species="Escherichia coli",
            strain=background.name,
            ploidy="haploid",
            assembly_set=assembly_set,  # type: ignore[arg-type]
            assembly_accession=accession,
            background=background,
        )

    return build


def test_deposit_raw_mirror_copies_from_the_library_and_refuses_a_bad_hash(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    library = tmp_path / "torchcell-library" / ch.CITATION_KEY
    (library / "si").mkdir(parents=True)
    write_source_data(library / ch.SOURCE_DATA.library_relpath)
    write_variant_tables(library / "si")
    for artifact in (ch.VARIANTS_MS56, ch.VARIANTS_MG1655, ch.VARIANTS_EXTRA):
        (library / "si" / artifact.name).replace(library / artifact.library_relpath)
    (library / ch.SUPPLEMENTARY_PDF.library_relpath).write_bytes(b"%PDF-1.4 stub")
    pins = {
        raw.name: file_sha256(library / raw.library_relpath) for raw in ch.RAW_FILES
    }
    patched = tuple(
        raw.model_copy(update={"sha256": pins[raw.name], "bytes": 1})
        for raw in ch.RAW_FILES
    )
    monkeypatch.setattr(ch, "RAW_FILES", patched)
    monkeypatch.setattr(ch, "RAW_BY_NAME", {raw.name: raw for raw in patched})
    monkeypatch.setattr(ch, "DATA_SHA256", pins)

    root = ch.deposit_raw_mirror(str(tmp_path))
    manifest = ch.load_manifest(str(tmp_path))
    assert {record.path for record in manifest.files} == {
        raw.mirror_relpath for raw in patched
    }
    assert (
        ch.manifest_sha256(manifest, ch.SOURCE_DATA.mirror_relpath)
        == pins[ch.SOURCE_DATA.name]
    )
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        ch.manifest_sha256(manifest, "data/nope.xlsx")
    # idempotent: a second deposit over identical bytes is a no-op
    assert ch.deposit_raw_mirror(str(tmp_path)) == root

    (library / ch.SUPPLEMENTARY_PDF.library_relpath).write_bytes(b"different")
    with pytest.raises(RuntimeError, match="sha256"):
        ch.deposit_raw_mirror(str(tmp_path))


def _deposit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, str]:
    """Build a synthetic raw mirror under ``tmp_path`` and pin it; returns the pins."""
    library = tmp_path / "torchcell-library" / ch.CITATION_KEY
    (library / "si").mkdir(parents=True)
    write_source_data(library / ch.SOURCE_DATA.library_relpath)
    write_variant_tables(library / "si")
    for artifact in (ch.VARIANTS_MS56, ch.VARIANTS_MG1655, ch.VARIANTS_EXTRA):
        (library / "si" / artifact.name).replace(library / artifact.library_relpath)
    (library / ch.SUPPLEMENTARY_PDF.library_relpath).write_bytes(b"%PDF-1.4 stub")
    pins = {
        raw.name: file_sha256(library / raw.library_relpath) for raw in ch.RAW_FILES
    }
    patched = tuple(
        raw.model_copy(update={"sha256": pins[raw.name], "bytes": 1})
        for raw in ch.RAW_FILES
    )
    monkeypatch.setattr(ch, "RAW_FILES", patched)
    monkeypatch.setattr(ch, "RAW_BY_NAME", {raw.name: raw for raw in patched})
    monkeypatch.setattr(ch, "DATA_SHA256", pins)
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    ch.deposit_raw_mirror(str(tmp_path))
    return pins


def test_download_links_the_mirror_and_refuses_an_absent_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pins = _deposit(tmp_path, monkeypatch)
    raw_dir = tmp_path / "build" / "raw"
    raw_dir.mkdir(parents=True)
    dataset = ch.GrowthRateChoe2019Dataset.__new__(ch.GrowthRateChoe2019Dataset)
    monkeypatch.setattr(type(dataset), "raw_dir", property(lambda self: str(raw_dir)))
    dataset.download()
    assert file_sha256(raw_dir / ch.SOURCE_DATA.name) == pins[ch.SOURCE_DATA.name]

    (tmp_path / ch.RAW_DIR_REL / ch.SOURCE_DATA.mirror_relpath).unlink()
    with pytest.raises(RuntimeError, match="missing from mirror"):
        dataset.download()


def test_process_builds_the_deletion_and_reconstructed_allele_records(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _deposit(tmp_path, monkeypatch)
    monkeypatch.setattr(
        ch,
        "assembly_reference",
        _reference_genome("MS56", MG1655_ASSEMBLY, "GCA_000005845.2"),
    )
    root = tmp_path / "growth_rate_choe2019"
    dataset = ch.GrowthRateChoe2019Dataset(root=str(root))
    assert len(dataset) == ch.GROWTH_EXPECTED_RECORDS
    # every gene is stored under its released symbol (#1045 part 2)
    assert dataset.gene_set == set(ch.LARGE_DELETION_SYMBOLS) | {"ilvN", "cspC", "yifB"}

    items = [dataset.transform_item(dataset[i]) for i in range(len(dataset))]
    assert sorted(len(i["experiment"].genotype.perturbations) for i in items) == [
        0,
        0,
        0,
        1,
        1,
        1,
        1,
        21,
    ]
    assert {i["publication"].doi for i in items} == {ch.PAPER_DOI}
    assert {i["reference"].phenotype_reference.fitness for i in items} == {1.0}
    by_reference = {
        i["reference"].genome_reference.strain: i["reference"].genome_reference
        for i in items
    }
    assert sorted(by_reference) == ["MS56", "cspCWT", "eMS57", "ilvNWT", "yifBWT"]
    assert {
        strain: genome.background.genotype_statement
        for strain, genome in by_reference.items()
    } == {
        "MS56": ch.MS56_GENOTYPE_STATEMENT,
        "eMS57": "E. coli MS56 adaptively evolved in M9 glucosemedium",
        "ilvNWT": "MS56, ilvN::ilvN-kan",
        "cspCWT": "MS56, cspC::cspC-kan",
        "yifBWT": "MS56, yifB::yifB-kanR",
    }
    assert {len(g.background.alleles) for g in by_reference.values()} == {0}
    assert by_reference["eMS57"].background.parents == ["MS56"]
    edited = [i for i in items if i["experiment"].genotype.perturbations]
    deletions = [i for i in edited if i["reference"].genome_reference.strain == "MS56"]
    # both designed deletions grow FASTER than the reduced parent, so fitness > 1
    assert all(i["experiment"].phenotype.fitness > 1.0 for i in deletions)
    alleles = {
        i["experiment"].genotype.perturbations[0].strain_label: i[
            "experiment"
        ].phenotype
        for i in edited
        if i["reference"].genome_reference.strain != "MS56"
    }
    assert {k: round(v.fitness, 6) for k, v in alleles.items()} == {
        "ilvNmut": round(statistics.fmean(FIG2F["ilvN*/ilvNWT"]), 6),
        "cspCmut": round(statistics.fmean(FIG2F["cspC*/cspCWT"]), 6),
        "yifBmut": round(statistics.fmean(FIG2F["yifB*/yifBWT"]), 6),
    }
    assert {v.n_samples for v in alleles.values()} == {3}
    # #1045 part 3: MS56 over MG1655, and eMS57 over MS56 once per panel, each on its
    # named background with an empty genotype
    unedited = {
        (
            i["reference"].genome_reference.strain,
            i["experiment"].phenotype.screen_id,
        ): i["experiment"].phenotype.fitness
        for i in items
        if not i["experiment"].genotype.perturbations
    }
    mean = statistics.fmean
    assert unedited == pytest.approx(
        {
            ("MS56", None): mean(FIG2C["MS56"]) / mean(FIG2C["MG1655"]),
            ("eMS57", "Fig2c"): mean(FIG2C["eMS57"]) / mean(FIG2C["MS56"]),
            ("eMS57", "Fig2f"): mean(FIG2F["eMS57/MS56"]),
        }
    )
    assert {i["experiment"].environment.temperature for i in items} == {None}
    dataset.close_lmdb()

    preprocess = Path(dataset.preprocess_dir)
    ledger = json.loads((preprocess / "dropped_records.json").read_text())
    assert ledger["source_rows"] == 9
    assert ledger["kept_records"] == ch.GROWTH_EXPECTED_RECORDS
    # nothing is dropped since #1045 part 3; MG1655 is MS56's reference
    assert ledger["dropped_records"] == 0
    assert ledger["rules"] == []
    assert ledger["reference_rows"] == ["MG1655", "MS56", "ilvNWT", "cspCWT", "yifBWT"]
    assert {rule["reason"] for rule in ledger["retired_rules"]} == {
        "evolved_clone_has_no_writable_genotype",
        "parent_genotype_not_enumerated_by_this_paper",
    }
    assert any("ilvN*/ilvNWT" in note for note in ledger["notes"])

    variants = json.loads((preprocess / "called_variants.json").read_text())
    assert variants["carrier"] is None
    assert variants["total_calls"] == len(_MS56_ROWS) + 2 * len(_OTHER_ROWS)
    assert len(variants["tables"]) == 3
    assert "intergenic_call_has_no_locus_to_key_to" in variants["blocking_reasons"]

    gaps = json.loads((preprocess / "genotype_gaps.json").read_text())
    assert [(a["strain"], a["field"]) for a in gaps["untyped_absences"]] == [
        ("MS56", "alleles"),
        ("eMS57", "alleles"),
    ]
    assert gaps["typed_gaps"]["environment"][0]["field"] == "temperature"

    checks = json.loads((preprocess / "released_statistics_check.json").read_text())
    assert checks["recovery_fraction"]["stated_recovery_fraction_of_evolved"] == 0.80
    assert checks["panel_disagreement"]["factor"] == pytest.approx(2.013, abs=1e-3)
    # the region cross-check is a verification finding now, not a build check
    assert "deleted_region" not in checks

    sourced = json.loads((preprocess / "sourced_values.json").read_text())
    assert sourced["cassette"]["value"] == ch.KAN_CASSETTE
    strains = pd.read_csv(preprocess / "strains.csv")
    assert sorted(strains["deleted_genes"].dropna()) == [1, 21]
    assert list(strains["strain"].iloc[2:5]) == ["ilvNmut", "cspCmut", "yifBmut"]
    assert list(strains["strain"].iloc[5:]) == ["MS56", "eMS57", "eMS57/MS56"]
    assert (preprocess / "build_manifest.json").exists()

    report = ch.verify_growth_build(
        str(root),
        genome=FakeGenome(MG1655_TAGS, MG1655_ASSEMBLY),  # type: ignore[arg-type]
    )
    # a released gene symbol is not a locus of the pinned annotation; the rules that
    # say so report and never gate (#1057)
    failed = {r.name for r in report.results if not r.passed}
    assert failed <= GENOME_REFERENCE_RULES, failed
    assert report.passed
    (region,) = [r for r in report.results if r.name == ch.DELETED_REGION_RULE]
    assert region.passed and region.details["genes_inside_window"] == 21
    assert (preprocess / "verification_report.json").exists()


def test_process_builds_the_two_keio_records(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _deposit(tmp_path, monkeypatch)
    monkeypatch.setattr(
        ch,
        "assembly_reference",
        _reference_genome("BW25113", BW25113_ASSEMBLY, "GCA_000750555.1"),
    )
    root = tmp_path / "tf_knockout_growth_choe2019"
    dataset = ch.TranscriptionFactorKnockoutChoe2019Dataset(root=str(root))
    assert len(dataset) == ch.TF_EXPECTED_RECORDS
    assert dataset.gene_set == set(BW25113_TAGS)

    items = [dataset.transform_item(dataset[i]) for i in range(len(dataset))]
    assert {len(i["experiment"].genotype.perturbations) for i in items} == {1}
    assert {i["experiment"].phenotype.fitness for i in items} == {0.724, 0.698}
    assert {i["reference"].phenotype_reference.fitness for i in items} == {1.0}
    assert [
        isinstance(i["reference"].genome_reference, StrainReferenceGenome)
        for i in items
    ] == [False, False]
    assert {i["experiment"].environment.temperature.value for i in items} == {37.0}
    dataset.close_lmdb()

    preprocess = Path(dataset.preprocess_dir)
    ledger = json.loads((preprocess / "dropped_records.json").read_text())
    assert ledger["source_rows"] == ch.KEIO_PANEL_STRAINS
    assert ledger["kept_records"] == ch.TF_EXPECTED_RECORDS
    assert ledger["dropped_records"] == ch.KEIO_PANEL_STRAINS - 2
    assert any("dicA" in rule["detail"] for rule in ledger["rules"])
    assert any("ydaS deletion" in note for note in ledger["notes"])
    gaps = json.loads((preprocess / "provenance_gaps.json").read_text())
    assert {gap["field"] for gap in gaps["gaps"]} == {
        "n_samples",
        "fitness_uncertainty",
        "cassette",
        "construction",
    }
    strains = pd.read_csv(preprocess / "strains.csv")
    assert list(strains["gene_symbol"]) == ["ydaS", "abgR"]

    report = ch.verify_tf_build(
        str(root),
        genome=FakeGenome(BW25113_TAGS, BW25113_ASSEMBLY),  # type: ignore[arg-type]
    )
    # the released symbols are not the annotation's tags; only the genome-reference
    # rows may say so, and they never gate (#1045 part 2)
    assert report.passed
    assert {r.name for r in report.results if not r.passed} <= GENOME_REFERENCE_RULES


def test_process_refuses_a_record_count_the_module_does_not_state(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _deposit(tmp_path, monkeypatch)
    monkeypatch.setattr(
        ch,
        "assembly_reference",
        _reference_genome("BW25113", BW25113_ASSEMBLY, "GCA_000750555.1"),
    )
    monkeypatch.setattr(ch, "KEIO_FITNESS", (("ydaS", 0.724),))
    with pytest.raises(RuntimeError, match="the module states"):
        ch.TranscriptionFactorKnockoutChoe2019Dataset(root=str(tmp_path / "tf_short"))


def test_read_reconstruction_ratios_refuses_a_changed_title(
    tmp_path: Path, source_data: Path
) -> None:
    book = openpyxl.load_workbook(source_data)
    book[ch.RATIO_SHEET].cell(row=1, column=2, value="Something else")
    path = tmp_path / "retitled.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(ch.TableFormatError, match="title is"):
        ch.read_reconstruction_ratios(str(path))


def test_deposit_raw_mirror_refuses_a_mirror_file_that_changed_underneath(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A deposited file with another hash is never overwritten; the deposit stops."""
    _deposit(tmp_path, monkeypatch)
    mirror = tmp_path / ch.RAW_DIR_REL / ch.SUPPLEMENTARY_PDF.mirror_relpath
    mirror.write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        ch.deposit_raw_mirror(str(tmp_path))


def test_growth_process_refuses_a_record_count_the_module_does_not_state(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _deposit(tmp_path, monkeypatch)
    monkeypatch.setattr(
        ch,
        "assembly_reference",
        _reference_genome("MS56", MG1655_ASSEMBLY, "GCA_000005845.2"),
    )
    monkeypatch.setattr(ch, "GROWTH_EXPECTED_RECORDS", 3)
    with pytest.raises(RuntimeError, match="the module states 3"):
        ch.GrowthRateChoe2019Dataset(root=str(tmp_path / "growth_short"))


def test_create_experiment_is_not_the_entry_point() -> None:
    for klass in (
        ch.GrowthRateChoe2019Dataset,
        ch.TranscriptionFactorKnockoutChoe2019Dataset,
    ):
        dataset = klass.__new__(klass)
        with pytest.raises(NotImplementedError):
            dataset.create_experiment()
        frame = pd.DataFrame({"a": [1]})
        assert dataset.preprocess_raw(frame) is frame


# --------------------------------------------------------------------------- #
# Data-gated: the real mirror
# --------------------------------------------------------------------------- #
def _real(name: str) -> str:
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT is not set")
    path = osp.join(data_root, ch.RAW_DIR_REL, "data", name)
    if not osp.exists(path):
        pytest.skip(f"raw mirror not deposited: {path}")
    return path


@pytest.mark.data
def test_real_fig2c_releases_the_five_strain_means() -> None:
    panel = ch.read_growth_panel(_real(ch.SOURCE_DATA.name))
    means = {row.strain: round(row.mean, 6) for row in panel.rows}
    assert means == {
        "MG1655": 0.589217,
        "MS56": 0.182333,
        "eMS57": 0.513433,
        "Δ21 kb": 0.398567,
        "ΔrpoS": 0.409,
    }
    assert {row.n_samples for row in panel.rows} == {3}


@pytest.mark.data
def test_real_panels_reproduce_the_prose_and_disagree_with_each_other() -> None:
    path = _real(ch.SOURCE_DATA.name)
    panel = ch.read_growth_panel(path)
    ratios = ch.read_reconstruction_ratios(path)
    recovery = ch.check_recovery_fraction(panel)
    assert recovery["measured_recovery_fraction_of_evolved"]["Δ21 kb"] == pytest.approx(
        0.7763, abs=1e-4
    )
    assert recovery["parent_over_wild_type"] == pytest.approx(0.3094, abs=1e-4)
    disagreement = ch.check_panels_disagree(panel, ratios)
    assert disagreement["factor"] == pytest.approx(2.013, abs=1e-3)


@pytest.mark.data
def test_real_variant_tables_release_the_counted_population_calls() -> None:
    counts = {}
    for artifact, header_rows, gene_column, _t, _f in ch.VARIANT_LAYOUTS:
        table = ch.read_variant_table(
            _real(artifact.name),
            artifact,
            header_rows=header_rows,
            gene_column=gene_column,
        )
        counts[artifact.name] = table.rows
        assert table.frequency_columns
    assert counts == {"si5.xlsx": 117, "si6.xlsx": 101, "si7.xlsx": 145}


@pytest.mark.data
def test_real_supplementary_data_2_frequencies_span_twenty_timepoints() -> None:
    table = ch.read_variant_table(
        _real(ch.VARIANTS_MS56.name), ch.VARIANTS_MS56, header_rows=3, gene_column=0
    )
    assert table.frequency_columns == tuple(
        str(day)
        for day in (
            0,
            5,
            10,
            15,
            20,
            25,
            27,
            28,
            29,
            30,
            33,
            35,
            37,
            40,
            43,
            45,
            50,
            55,
            60,
            62,
        )
    )
    assert table.intergenic_rows == 22
    assert table.loci_with_several_calls == {
        "iscR": 3,
        "rrsH": 2,
        "tufA": 2,
        "yafF": 5,
        "ydjN": 4,
        "yeaM": 2,
    }


@pytest.mark.data
def test_real_mg1655_annotation_puts_the_named_genes_in_one_contiguous_run() -> None:
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT is not set")
    from torchcell.datasets.bacteria_common import bacterial_genome

    genome = bacterial_genome("ecoli", "MG1655", data_root)
    region = ch.check_deleted_region(genome)
    assert region["problems"] == []
    assert region["locus_tags"] == [f"b{2721 + i}" for i in range(21)]
    assert region["mg1655_window"] == [2844762, 2867551]
    assert region["mg1655_window_bp"] == 22790
    # the released interval is MS56's, 806,266 bp upstream of the MG1655 one
    assert region["ms56_to_mg1655_offset_bp"] == 806266


@pytest.mark.data
def test_real_every_sourced_quote_is_verbatim_in_its_pinned_file() -> None:
    """Each quote, the #1057 Supplementary Table 3 rows included, is in its file."""
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT is not set")
    library = Path(data_root, ch.LIBRARY_DIR_REL)
    if not library.exists():
        pytest.skip(f"library mirror not present: {library}")
    texts: dict[str, str] = {}
    for prov in (ch.PAPER, ch.SI1):
        assert file_sha256(library / prov.source_uri) == prov.sha256
        texts[prov.source_uri] = (library / prov.source_uri).read_text(encoding="utf-8")
    missing = [
        name
        for name, value in ch.SOURCED_VALUES.items()
        if value.quote not in texts[value.provenance.source_uri]
    ]
    assert missing == []
