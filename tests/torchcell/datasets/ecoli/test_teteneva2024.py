# tests/torchcell/datasets/ecoli/test_teteneva2024.py
# [[tests.torchcell.datasets.ecoli.test_teteneva2024]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_teteneva2024.py

"""Teteneva 2024 lake-water RB-TnSeq: the deposits, the readers and the loader.

The hermetic tests run the real readers on synthetic Tables S2, S3 and S4 (the red QC
fill included) and build the dataset end to end from them, then verify it; the records
are keyed by the released symbols verbatim, so no genome is involved. The ``--data``
tests re-hash the W3110 tier set, re-read the deposited tables, and audit every quote
against the raw mirror's pinned bytes.
"""

import hashlib
import json
import math
import os
import os.path as osp
import re
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pytest
from openpyxl.styles import PatternFill

from torchcell.datamodels.media import LAKE_WATER
from torchcell.datamodels.schema import (
    ASSEMBLY_SET_ACCESSIONS,
    BACTERIAL_ASSEMBLY_SETS,
    ComponentDefinition,
    MeasurementType,
    ReferenceGenome,
    SampleUnit,
    StrainBackground,
    StrainReferenceGenome,
)
from torchcell.datasets.ecoli import teteneva2024 as te
from torchcell.literature.manual_deposit import verify_deposit
from torchcell.sequence.genome.registry import (
    ECOLI_K12_W3110,
    resolve,
    verify_assembly_set,
)
from torchcell.verification.report import Level
from torchcell.verification.sourced import (
    ProvenanceGapReason,
    SourcedValue,
    audit_sourced_value,
)

DATA_ROOT = os.environ.get("DATA_ROOT", "")
TIER_MANIFEST = osp.join(
    DATA_ROOT, "torchcell-genomes", ECOLI_K12_W3110, "manifest.json"
)
requires_tier = pytest.mark.skipif(
    not osp.isfile(TIER_MANIFEST),
    reason=f"requires {ECOLI_K12_W3110} under $DATA_ROOT/torchcell-genomes",
)
RAW_KEY_ROOT = Path(DATA_ROOT) / "torchcell-raw" / te.CITATION_KEY
requires_raw = pytest.mark.skipif(
    not (RAW_KEY_ROOT / te.PAPER_TEXT_RELPATH).is_file(),
    reason=f"requires the raw mirror key {te.CITATION_KEY} with its paper text",
)

SHA256 = re.compile(r"^[0-9a-f]{64}$")
MD5 = re.compile(r"^[0-9a-f]{32}$")
RED = PatternFill(
    start_color=te.QC_HIGHLIGHT_FILL, end_color=te.QC_HIGHLIGHT_FILL, fill_type="solid"
)


# --------------------------------------------------------------------------------------
# The deposits, hermetically
# --------------------------------------------------------------------------------------
def test_the_deposit_record_names_the_registry_set_and_the_accession_pair() -> None:
    """The record's set id is the registry constant and its pair is GCA .1 / GCF .2."""
    deposit = te.W3110_TIER_DEPOSIT
    assert deposit.assembly_set == ECOLI_K12_W3110 == "ecoli_K12_W3110_ASM1024v1"
    assert deposit.genbank_accession == "GCA_000010245.1"
    assert deposit.refseq_accession == "GCF_000010245.2"
    assert deposit.genbank_replicon == "AP009048.1"
    assert deposit.refseq_replicon == "NC_007779.1"
    assert deposit.replicon_length_bp == 4_646_332
    assert deposit.assembly_name == "ASM1024v1"
    assert deposit.deposited_at == "2026-10-09"


def test_the_deposit_lists_ten_distinct_members_with_both_digests() -> None:
    """Six GCA members, four GCF members, every md5 and sha256 well formed."""
    members = te.W3110_TIER_DEPOSIT.members
    assert len(members) == 10
    assert len({m.path for m in members}) == 10
    assert [m.path.split("_")[0] for m in members] == ["GCA"] * 6 + ["GCF"] * 4
    for member in members:
        assert MD5.match(member.md5), member.path
        assert SHA256.match(member.sha256), member.path
        assert member.bytes > 0


def test_the_gcf_assembly_report_is_a_member_because_the_gca_one_is_stale() -> None:
    """Only this set deposits the GCF report, and the record says why."""
    paths = {m.path for m in te.W3110_TIER_DEPOSIT.members}
    assert "GCF_000010245.2_ASM1024v1_assembly_report.txt" in paths
    assert te.STALE_GCA_REFSEQ_ACCESSION == "GCF_000010245.1"
    assert te.STALE_GCA_REFSEQ_REPLICON == "AC_000091.1"
    assert te.STALE_GCA_REFSEQ_ACCESSION != te.W3110_TIER_DEPOSIT.refseq_accession


def test_the_set_is_in_the_schema_vocabulary_with_the_gcf_report_s_pair() -> None:
    """W3110 is a reference strain (#910) whose set pairs GCA .1 with GCF .2."""
    assert BACTERIAL_ASSEMBLY_SETS["W3110"] == ECOLI_K12_W3110
    assert ASSEMBLY_SET_ACCESSIONS[ECOLI_K12_W3110] == (
        te.W3110_TIER_DEPOSIT.genbank_accession,
        te.W3110_TIER_DEPOSIT.refseq_accession,
    )
    assert te.GENE_NAMESPACE == "ecoli_k12_w3110_refseq_locus_tag"


def test_the_pmc_deposit_holds_the_workbook_and_the_paper_text() -> None:
    """Sixteen objects were listed; Table S4 and the paper text are two of them."""
    keys = [f.key for f in te.PMC_OA_DEPOSIT]
    assert len(keys) == len(set(keys)) == 16
    assert te.TABLE_S4_OA_KEY in keys
    assert te.PAPER_TEXT_KEY == "PMC11188689.1/PMC11188689.1.txt"
    assert te.PAPER_TEXT_KEY in keys
    assert all(key.startswith(f"{te.PAPER_PMCID}.1/") for key in keys)


def test_the_schedule_counts_are_the_record_count() -> None:
    """The schedule's counts, now the loader's oracle: every value is a record."""
    assert te.SCHEDULE_ROW == 41
    assert te.SCHEDULE_FITNESS_VALUES == te.TABLE_S4_FITNESS_VALUES == 66_162
    assert te.EXPECTED_RECORDS == te.TABLE_S4_FITNESS_VALUES
    assert te.TABLE_S4_FITNESS_VALUES == 6 * te.TABLE_S4_GENE_BY_SAMPLE_ROWS
    assert sum(te.TABLE_S4_ROWS_PER_WATER_SAMPLE.values()) == (
        te.TABLE_S4_GENE_BY_SAMPLE_ROWS
    )
    assert te.EXPECTED_UNPERTURBED == 3 * te.TABLE_S4_ROWS_PER_WATER_SAMPLE[2]


def test_the_si_deposit_is_five_files_each_equal_to_a_listed_bucket_object() -> None:
    """Every deposited file names the PMC object it was measured equal to."""
    deposit = te.SI_DEPOSIT
    assert (deposit.citation_key, deposit.doi, deposit.issues) == (
        te.CITATION_KEY,
        te.PAPER_DOI,
        (910,),
    )
    listed = {f.key for f in te.PMC_OA_DEPOSIT}
    assert len(deposit.files) == 5
    for f in deposit.files:
        assert f.scripted_equivalent is not None
        assert f.scripted_equivalent.params["key"] in listed
    pinned = {f.path: f.sha256 for f in deposit.files}
    assert pinned[te.TABLE_S4_RELPATH] == te.TABLE_S4_SHA256
    assert pinned[te.TABLE_S3_RELPATH] == te.TABLE_S3_SHA256
    assert pinned[te.TABLE_S2_PDF_RELPATH] == te.TABLE_S2_PDF_SHA256
    assert {name: pin for name, (_, pin) in te.LOADER_FILES.items()} == {
        "table_s4.xlsx": te.TABLE_S4_SHA256,
        "table_s3.xlsx": te.TABLE_S3_SHA256,
        "table_s2.layout.txt": te.TABLE_S2_TEXT_SHA256,
    }


def test_the_paper_text_retrieval_is_the_pmc_bucket_object() -> None:
    retrieval = te.paper_text_retrieval()
    assert retrieval.method.value == "pmc_cloud"
    assert retrieval.params == {"key": te.PAPER_TEXT_KEY}
    assert retrieval.sha256 == te.PAPER_TEXT_SHA256
    assert retrieval.source_url == (
        "https://pmc-oa-opendata.s3.amazonaws.com/PMC11188689.1/PMC11188689.1.txt"
    )


def test_every_paper_quote_cites_the_pinned_paper_text() -> None:
    """Each sourced value carries the citation key and the sha256 of its file."""
    values = [v for v in vars(te).values() if isinstance(v, SourcedValue)]
    assert len(values) == 33
    pins = {
        te.PAPER_TEXT_RELPATH: te.PAPER_TEXT_SHA256,
        te.WORKBOOK_TEXT_RELPATH: te.WORKBOOK_TEXT_SHA256,
    }
    for value in values:
        assert value.provenance.citation_key == te.CITATION_KEY
        assert pins[value.provenance.source_uri] == value.provenance.sha256
    assert te.REFERENCE_SEQUENCE.value == "NC_007779.1"
    assert te.TECHNICAL_REPLICATE_UNIT.value is SampleUnit.technical_replicate
    assert te.BARCODE_FITNESS.value is MeasurementType.log2_ratio


# --------------------------------------------------------------------------------------
# Synthetic tables
# --------------------------------------------------------------------------------------
#: Genes every water sample carries, with fitness chosen so the paper's statements hold
#: (Tol-Pal negative; rhoL/rpoS/rbsR up in filtered water; glpK/glpF up in non-filtered).
GENES: dict[str, tuple[float, ...]] = {
    "tolA": (-2.0, -2.1, -2.2, -1.9, -2.3, -1.8),
    "pal": (-3.0, -3.1, -3.2, -2.9, -3.3, -2.8),
    "lpp": (-3.5, -3.4, -3.6, -2.7, -3.0, -2.5),
    "rhoL": (3.0, 6.9, 7.0, 2.2, 0.6, 1.0),
    "rpoS": (0.1, 2.9, 6.1, 0.7, -0.1, -0.3),
    "rbsR": (0.3, 2.0, 5.1, 0.3, -0.3, 0.0),
    "glpK": (0.1, 0.3, 0.1, 0.1, 0.2, 1.2),
    "glpF": (0.3, 0.4, 0.3, 0.3, 0.3, 1.3),
    "thrL": (0.01, -0.02, 0.03, 0.0, 0.05, -0.04),
}


def _pcs(rows: list[tuple[float, ...]]) -> list[tuple[float, float, float]]:
    """PC1, PC2 and their distance by the release's own (centred, scaled) PCA."""
    x = np.array(rows)
    y = (x - x.mean(axis=0)) / x.std(axis=0, ddof=1)
    _, _, vt = np.linalg.svd(y, full_matrices=False)
    scores = y @ vt.T[:, :2]
    return [(float(a), float(b), float(math.hypot(a, b))) for a, b in scores]


def _table_s4(path: Path, *, symbol: str | None = "thrL", empty: bool = False) -> None:
    """A Table S4 with the release's header rows, every gene in water samples 1-3."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "DATA"
    sheet.append(
        ["Lake water sample", "gene fitness values", None, None, None, None, None]
        + ["Gene name", "Principal component", None, "Euclidean distance"]
        + ["Protein description"]
    )
    sheet.append([None, "filtered", None, None, "non-filtered"])
    sheet.append(
        [None, "day 2", "day 4", "day 8", "day 2", "day 4", "day 8", None, 1, 2]
    )
    for k in (1, 2, 3):
        values = [tuple(v + 0.01 * k for v in f) for f in GENES.values()]
        for (gene, fitness), pcs in zip(
            zip(GENES, values, strict=True), _pcs(values), strict=True
        ):
            cells: list[Any] = list(fitness)
            if empty and gene == "thrL" and k == 1:
                cells[5] = None
            name = symbol if gene == "thrL" else gene
            sheet.append([k, *cells, name, *pcs, f"protein {gene}"])
    legend = book.create_sheet("LEGEND")
    legend.append(["Supplementary Table S4.", None, None, " "])
    legend.append([None])
    legend.append(["Gene fitness values, principal component values, and Euclidean."])
    book.save(path)


def _table_s3(
    path: Path, *, failed: str | None = "sample_04", drop: bool = False
) -> None:
    """A Table S3: per water sample a control and one replicate per condition, plus a
    second replicate of #1 filtered day 4; ``failed`` (default: #1 filtered day 4's
    first replicate, as sample_10 is in the release) carries the red fill.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "DATA"
    sheet.append(
        [
            "sample id",
            "biological replicate",
            "sampling day",
            "sample type",
            "technical replicate",
            "duplicate reads",
        ]
    )
    n = 0
    for k in (1, 2, 3):
        rows = [(0, "control", 1)]
        for day in (2, 4, 8):
            rows += [(day, "filtered water", 1), (day, "non-filtered water", 1)]
            if k == 1 and day == 4:
                rows.append((day, "filtered water", 2))
        for day, kind, rep in rows:
            if drop and (k, day, kind) == (3, 8, "non-filtered water"):
                continue
            n += 1
            sheet.append([f"sample_{n:02d}", k, day, kind, rep, 0.05])
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == failed:
            for cell in row:
                cell.fill = RED
    legend = book.create_sheet("LEGEND")
    legend.append(["Supplementary Table S3"])
    legend.append([None])
    legend.append(["One sample that failed the QC is omitted."])
    book.save(path)


TABLE_S2 = (
    "Supplementary Table S2. Parameters of taken lake water samples.\n\n"
    " Date          Water sample    pH    Air temperature, °C   Water temperature, °C\n"
    "\n 9.05.2022     #1              7     13                    12.4\n"
    "\n 19.06.2022    #2              6.5   29                    18.4\n"
    "\n 6.07.2022     #3              6.5   17                    16.1\n"
    "\n 1.06.2023     #4              6.5   17                    14.6\n"
    "\n 8.10.2023     #5              6.5   16                    13.6\n"
)


# --------------------------------------------------------------------------------------
# The readers
# --------------------------------------------------------------------------------------
def test_table_s4_inventory_names_the_columns_and_counts_values_and_symbols(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s4.xlsx"
    _table_s4(path)
    inventory = te.table_s4_inventory(path)
    assert inventory.gene_by_sample_rows == 27
    assert inventory.rows_per_water_sample == {1: 9, 2: 9, 3: 9}
    assert inventory.fitness_columns == te.TABLE_S4_FITNESS_COLUMNS
    assert inventory.fitness_values == 162
    assert inventory.identifier_column == "Gene name"
    assert len(inventory.distinct_symbols) == 9
    assert inventory.other_columns == (
        "Principal component",
        "Euclidean distance",
        "Protein description",
    )


def test_table_s4_inventory_refuses_another_layout_and_a_row_without_a_symbol(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s4.xlsx"
    _table_s4(path, symbol=None)
    with pytest.raises(ValueError, match="no gene symbol"):
        te.table_s4_inventory(path)
    book = openpyxl.load_workbook(path)
    book["DATA"]["H1"] = "Locus tag"
    book.save(path)
    with pytest.raises(ValueError, match="Table S4 header"):
        te.table_s4_inventory(path)


def test_read_table_s4_keeps_each_row_verbatim_in_column_order(tmp_path: Path) -> None:
    path = tmp_path / "s4.xlsx"
    _table_s4(path)
    rows = te.read_table_s4(path)
    assert len(rows) == 27
    first = rows[0]
    assert (first.water_sample, first.symbol, first.description) == (
        1,
        "tolA",
        "protein tolA",
    )
    assert first.fitness == pytest.approx((-1.99, -2.09, -2.19, -1.89, -2.29, -1.79))
    assert [c.label for c in te.FITNESS_COLUMNS] == list(te.TABLE_S4_FITNESS_COLUMNS)
    assert [c.index for c in te.FITNESS_COLUMNS] == [1, 2, 3, 4, 5, 6]


def test_read_table_s4_refuses_an_empty_value_and_a_renamed_column(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s4.xlsx"
    _table_s4(path, empty=True)
    with pytest.raises(ValueError, match="thrL has an empty fitness cell"):
        te.read_table_s4(path)
    _table_s4(path)
    book = openpyxl.load_workbook(path)
    book["DATA"]["C3"] = "day 3"
    book.save(path)
    with pytest.raises(ValueError, match="Table S4 fitness columns are"):
        te.read_table_s4(path)


def test_the_pca_of_the_release_is_reproduced_and_a_changed_value_breaks_it(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s4.xlsx"
    _table_s4(path)
    rows = te.read_table_s4(path)
    reproduction = te.reproduce_pca(rows)
    assert [r.water_sample for r in reproduction] == [1, 2, 3]
    assert [r.n_genes for r in reproduction] == [9, 9, 9]
    result = te.pca_result(reproduction)
    assert (result.level, result.passed) == (Level.L4, True)
    changed = [rows[0].model_copy(update={"pc1": rows[0].pc1 + 0.5}), *rows[1:]]
    assert te.pca_result(te.reproduce_pca(changed)).passed is False


def test_parse_table_s2_reads_day_month_year_and_every_column() -> None:
    samples = te.parse_table_s2(TABLE_S2)
    assert [s.number for s in samples] == [1, 2, 3, 4, 5]
    second = samples[1]
    assert (second.sampled_on.isoformat(), second.ph) == ("2022-06-19", 6.5)
    assert (second.air_temperature_c, second.water_temperature_c) == (29.0, 18.4)
    assert (
        second.row == "19.06.2022    #2              6.5   29                    18.4"
    )
    value = te.water_sample_value(second)
    assert value.quote == second.row
    assert value.value["sampled_on"] == "2022-06-19"
    assert value.provenance.sha256 == te.TABLE_S2_TEXT_SHA256


def test_parse_table_s2_refuses_another_header_or_a_missing_row() -> None:
    with pytest.raises(ValueError, match="no header row"):
        te.parse_table_s2(TABLE_S2.replace("Water sample", "Sample"))
    with pytest.raises(ValueError, match=r"not #1-#5"):
        te.parse_table_s2(TABLE_S2.replace(" 1.06.2023     #4", " x"))


def test_read_table_s3_finds_the_red_sample_and_counts_replicates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s3.xlsx"
    _table_s3(path)
    samples = te.read_table_s3(path)
    assert len(samples) == 22
    assert [s.sample_id for s in samples if s.qc_failed] == ["sample_04"]
    assert te.highlighted_rows(path) == [("sample_04", 1, 4, "filtered water", 1, 0.05)]
    counts = te.replicate_counts(samples)
    # the red replicate is omitted: #1 filtered day 4 keeps its second replicate only
    assert len(counts) == 18
    assert counts[(1, "filtered", 4)] == 1
    assert set(counts.values()) == {1}
    unflagged = te.replicate_counts(te.read_table_s3(_rewrite(path, failed=None)))
    assert unflagged[(1, "filtered", 4)] == 2


def _rewrite(path: Path, *, failed: str | None) -> Path:
    other = path.with_name(f"other_{path.name}")
    _table_s3(other, failed=failed)
    return other


def test_conditions_refuse_one_with_no_kept_replicate(tmp_path: Path) -> None:
    path = tmp_path / "s3.xlsx"
    _table_s3(path, failed=None, drop=True)
    with pytest.raises(ValueError, match=r"no QC-passed technical replicate for \(3"):
        te.conditions(te.replicate_counts(te.read_table_s3(path)))
    _table_s3(path, failed=None)
    conds = te.conditions(te.replicate_counts(te.read_table_s3(path)))
    assert len(conds) == 18
    assert conds[1].screen_id == "lake water #1, filtered, day 4"
    assert conds[1].n_technical_replicates == 2
    assert [c.column for c in conds[:6]] == [0, 1, 2, 3, 4, 5]


def test_read_table_s3_refuses_another_header(tmp_path: Path) -> None:
    path = tmp_path / "s3.xlsx"
    _table_s3(path)
    book = openpyxl.load_workbook(path)
    book["DATA"]["B1"] = "water"
    book.save(path)
    with pytest.raises(ValueError, match="Table S3 header"):
        te.read_table_s3(path)


def test_render_workbook_texts_carries_both_legends_and_the_red_row(
    tmp_path: Path,
) -> None:
    s3, s4 = tmp_path / "s3.xlsx", tmp_path / "s4.xlsx"
    _table_s3(s3)
    _table_s4(s4)
    text = te.render_workbook_texts(s3, s4)
    assert text.splitlines() == [
        "# workbook: s3.xlsx, sheet: LEGEND",
        "Supplementary Table S3",
        "One sample that failed the QC is omitted.",
        "# workbook: s4.xlsx, sheet: LEGEND",
        "Supplementary Table S4.",
        "Gene fitness values, principal component values, and Euclidean.",
        f"# workbook: s3.xlsx, sheet: DATA, rows filled {te.QC_HIGHLIGHT_FILL}",
        "sample_04\t1\t4\tfiltered water\t1\t0.05",
    ]


# --------------------------------------------------------------------------------------
# The records
# --------------------------------------------------------------------------------------
def _reference(*_: Any, **kwargs: Any) -> ReferenceGenome:
    background = kwargs["background"]
    assert isinstance(background, StrainBackground)
    return StrainReferenceGenome(
        species="Escherichia coli",
        strain=background.name,
        assembly_set="ecoli_K12_W3110_ASM1024v1",
        assembly_accession="GCF_000010245.2",
        background=background,
    )


def test_the_host_is_w3110_rpos_with_no_asserted_allele(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    background = te.host_background()
    assert (background.name, background.reference_strain) == ("W3110 RpoS+", "W3110")
    assert (background.genotype_statement, background.alleles) == ("RpoS+", [])
    assert background.provenance == [te.STRAIN, te.BACKGROUND_CONSTRUCTION]
    calls: list[tuple[Any, ...]] = []

    def pin(strain: str, **kwargs: Any) -> ReferenceGenome:
        calls.append((strain, kwargs["data_root"]))
        return _reference(**kwargs)

    monkeypatch.setattr(te, "assembly_reference", pin)
    reference = te.host_reference("/dr")
    assert calls == [("W3110", "/dr")]
    assert reference.strain == "W3110 RpoS+"


def test_each_water_and_filtration_is_its_own_undefined_natural_medium() -> None:
    sample = te.parse_table_s2(TABLE_S2)[0]
    filtered = te.lake_water_medium(sample, "filtered")
    raw = te.lake_water_medium(sample, "non-filtered")
    assert filtered.name == (
        "lake water #1 (sampled 2022-05-09), filter-sterilized 0.22 um (Teteneva 2024)"
    )
    assert (
        raw.name == "lake water #1 (sampled 2022-05-09), non-filtered (Teteneva 2024)"
    )
    assert (filtered.is_synthetic, filtered.state) == (False, "liquid")
    (component,) = filtered.components
    assert component.definition is ComponentDefinition.intrinsically_undefined
    assert filtered.components == raw.components == LAKE_WATER.components
    assert filtered.base_medium == raw.base_medium == "LAKE_WATER"
    assert filtered.provenance[2] == te.WATER_FILTRATION
    assert raw.provenance[2] == te.NON_FILTERED_PURPOSE
    assert filtered.provenance[0].quote == sample.row


def test_the_environment_carries_the_day_in_hours_and_the_temperature_gap(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s3.xlsx"
    _table_s3(path, failed=None)
    conds = te.conditions(te.replicate_counts(te.read_table_s3(path)))
    water = te.parse_table_s2(TABLE_S2)
    environment = te.build_environment(conds[2], water[0])
    assert environment.duration_hours == 192.0
    assert environment.temperature is None
    (gap,) = environment.provenance_gaps
    assert (gap.field, gap.reason) == (
        "temperature",
        ProvenanceGapReason.not_reported_by_primary,
    )
    assert environment.perturbations == []
    with pytest.raises(ValueError, match="given water sample #2"):
        te.build_environment(conds[2], water[1])


def test_a_record_is_keyed_by_the_released_symbol_with_no_lookup() -> None:
    genotype = te.build_genotype("yeiE")
    (leaf,) = genotype.perturbations
    assert leaf.model_dump(
        include={
            "systematic_gene_name",
            "perturbed_gene_name",
            "perturbation_type",
            "gene_namespace",
            "strain_label",
            "identifier_kind",
            "released_kind",
            "cassette",
            "collection",
        }
    ) == {
        "systematic_gene_name": "yeiE",
        "perturbed_gene_name": "yeiE",
        "perturbation_type": "released_identifier",
        "gene_namespace": "ecoli_k12_w3110_refseq_locus_tag",
        "strain_label": "yeiE",
        "identifier_kind": "Gene name",
        "released_kind": "transposon_insertion",
        "cassette": "kanamycin",
        "collection": None,
    }


def _built(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Path]:
    """Synthetic raw files wired into the loader; returns (data root, dataset root)."""
    source = tmp_path / "source"
    source.mkdir()
    _table_s4(source / "table_s4.xlsx")
    _table_s3(source / "table_s3.xlsx")
    (source / "table_s2.layout.txt").write_text(TABLE_S2)
    by_relpath = {relpath: name for name, (relpath, _) in te.LOADER_FILES.items()}
    monkeypatch.setattr(
        te,
        "LOADER_FILES",
        {
            name: (relpath, hashlib.sha256((source / name).read_bytes()).hexdigest())
            for name, (relpath, _) in te.LOADER_FILES.items()
        },
    )
    monkeypatch.setattr(
        te, "raw_path", lambda relpath, data_root=None: source / by_relpath[relpath]
    )
    monkeypatch.setattr(te, "assembly_reference", lambda strain, **kw: _reference(**kw))
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.delenv("TC_DATA_URL", raising=False)
    return data_root, data_root / te.DATASET_ROOT_REL


def test_the_loader_stores_every_released_value_end_to_end(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, root = _built(tmp_path, monkeypatch)
    dataset = te.RbTnseqTeteneva2024Dataset(root=str(root))
    # 27 rows x 6 conditions: every value, no row refused
    assert len(dataset) == 162
    assert sorted(dataset.gene_set) == sorted(GENES)
    assert dataset.organism.value == "ecoli"
    first = dataset[0]
    phenotype = first["experiment"]["phenotype"]
    assert phenotype["screen_id"] == "lake water #1, filtered, day 2"
    assert phenotype["environment_response"] == pytest.approx(-1.99)
    assert (phenotype["n_samples"], phenotype["sample_unit"]) == (
        1,
        "technical_replicate",
    )
    leaf = first["experiment"]["genotype"]["perturbations"][0]
    assert (leaf["systematic_gene_name"], leaf["released_kind"]) == (
        "tolA",
        "transposon_insertion",
    )
    assert first["reference"]["genome_reference"]["assembly_accession"] == (
        "GCF_000010245.2"
    )
    assert first["publication"]["pubmed_id"] == "38874171"
    inventory = json.loads((root / "preprocess" / "inventory.json").read_text())
    assert (inventory["released_values"], inventory["stored_records"]) == (162, 162)
    assert set(inventory["replicates_per_condition"].values()) == {1}
    conditions = json.loads((root / "preprocess" / "conditions.json").read_text())
    assert len(conditions) == 18
    assert sorted(os.listdir(root / "raw")) == sorted(te.LOADER_FILES)
    with pytest.raises(NotImplementedError, match="builds its records in process"):
        dataset.create_experiment()
    assert dataset.preprocess_raw("frame") == "frame"
    assert dataset.experiment_class.__name__ == "EnvironmentResponseExperiment"


def test_verify_build_passes_on_the_synthetic_store(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    data_root, root = _built(tmp_path, monkeypatch)
    te.RbTnseqTeteneva2024Dataset(root=str(root))
    # the modal medium is water #1's filtered medium in key order: 9 genes x 3 days
    monkeypatch.setattr(te, "EXPECTED_UNPERTURBED", 27)
    report = te.verify_build(str(root), data_root=str(data_root), expected_count=162)
    results = {r.name: r for r in report.results}
    assert {n: r.message for n, r in results.items() if not r.passed} == {}
    assert results["stored_gene_identifiers_are_the_released_symbols"].level == Level.L1
    assert results["stored_values_are_table_s4_cells"].details["wrong_value"] == []
    claims = results["paper_statements_hold_on_the_stored_values"].details["checks"]
    assert [c["claim"] for c in claims] == [
        "tol_pal_negative_everywhere",
        "rpos_group_filtered_advantage",
        "glp_non_filtered_advantage",
    ]
    written = json.loads((root / "preprocess" / "verification_report.json").read_text())
    assert written["dataset_name"] == "RbTnseqTeteneva2024Dataset"


def test_the_supplementary_rows_fail_on_a_changed_value_or_symbol(
    tmp_path: Path,
) -> None:
    path = tmp_path / "s4.xlsx"
    _table_s4(path)
    rows = te.read_table_s4(path)
    s3 = tmp_path / "s3.xlsx"
    _table_s3(s3, failed=None)
    conds = te.conditions(te.replicate_counts(te.read_table_s3(s3)))
    values = [
        te.StoredValue(
            screen_id=c.screen_id,
            symbol=r.symbol,
            perturbed_gene_name=r.symbol,
            perturbation_type="released_identifier",
            identifier_kind=te.IDENTIFIER_KIND,
            strain_label=r.symbol,
            gene_namespace=te.GENE_NAMESPACE,
            value=r.fitness[c.column],
            n_samples=c.n_technical_replicates,
        )
        for c in conds
        for r in rows
        if r.water_sample == c.water_sample
    ]
    assert te.stored_values_equal_release(values, rows, conds, 162).passed is True
    assert te.stored_symbols_are_released(values, rows).passed is True
    assert te.claims_result(te.paper_claims(values)).passed is True
    wrong = [values[0].model_copy(update={"value": 9.0, "n_samples": 7}), *values[1:]]
    result = te.stored_values_equal_release(wrong, rows, conds, 162)
    assert (result.passed, len(result.details["wrong_value"])) == (False, 1)
    assert len(result.details["wrong_n_samples"]) == 1
    renamed = [
        values[0].model_copy(update={"identifier_kind": "locus_tag"}),
        *values[1:],
    ]
    assert te.stored_symbols_are_released(renamed, rows).passed is False
    flipped = [
        v.model_copy(update={"value": -v.value}) if v.symbol == "glpK" else v
        for v in values
    ]
    checks = te.paper_claims(flipped)
    assert [c.holds for c in checks] == [True, True, False]
    assert checks[2].failures == ("glpK #1", "glpK #2", "glpK #3")


def test_main_dispatches_record_report_and_verify(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(
        te,
        "record_paper_and_texts",
        lambda paper_text_source=None: f"rec {paper_text_source}",
    )
    monkeypatch.setattr(te, "loader_report", lambda: {"n": 1})

    class _Report:
        def summary(self) -> str:
            return "verified"

    monkeypatch.setattr(te, "verify_build", lambda root, data_root: _Report())
    monkeypatch.setenv("DATA_ROOT", "/dr")
    te.main(["record", "--paper-text", "p.txt"])
    te.main(["report"])
    te.main(["verify"])
    assert capsys.readouterr().out.splitlines() == [
        "rec p.txt",
        "{",
        '  "n": 1',
        "}",
        "verified",
    ]


# --------------------------------------------------------------------------------------
# The deposited bytes (--data)
# --------------------------------------------------------------------------------------
@pytest.mark.data
@requires_tier
def test_every_deposited_member_hashes_to_the_record() -> None:
    """``verify_assembly_set`` re-hashes the tier; the ten digests are the record's."""
    assert verify_assembly_set(ECOLI_K12_W3110) == {
        m.path: m.sha256 for m in te.W3110_TIER_DEPOSIT.members
    }


@pytest.mark.data
@requires_tier
def test_the_two_assembly_reports_disagree_about_the_refseq_release() -> None:
    """The GCA report names the 2006 release; only the GCF report names the pair."""
    gca = Path(
        resolve(ECOLI_K12_W3110, "GCA_000010245.1_ASM1024v1_assembly_report.txt")
    ).read_text()
    gcf = Path(
        resolve(ECOLI_K12_W3110, "GCF_000010245.2_ASM1024v1_assembly_report.txt")
    ).read_text()
    assert "# RefSeq assembly accession: GCF_000010245.1\n" in gca
    assert "# RefSeq assembly accession: GCF_000010245.2\n" in gcf
    assert "\tNC_007779.1\t" in gcf


@pytest.mark.data
@requires_raw
def test_every_quote_is_in_its_pinned_file_and_the_texts_rerender() -> None:
    """Every sourced value audits against the raw mirror, the Table S2 rows included,
    and the two derived texts re-render to their pins off the deposited bytes.
    """
    root = Path(DATA_ROOT) / "torchcell-raw"
    verify_deposit(te.SI_DEPOSIT, RAW_KEY_ROOT)
    water = te.parse_table_s2(
        (RAW_KEY_ROOT / te.TABLE_S2_TEXT_RELPATH).read_text(encoding="utf-8")
    )
    values = [v for v in vars(te).values() if isinstance(v, SourcedValue)]
    values += [te.water_sample_value(s) for s in water]
    for value in values:
        assert audit_sourced_value(value, root).passed, value.quote[:60]
    rendered = te.render_workbook_texts(
        RAW_KEY_ROOT / te.TABLE_S3_RELPATH, RAW_KEY_ROOT / te.TABLE_S4_RELPATH
    )
    assert hashlib.sha256(rendered.encode()).hexdigest() == te.WORKBOOK_TEXT_SHA256
    s2 = te.render_table_s2(RAW_KEY_ROOT / te.TABLE_S2_PDF_RELPATH)
    assert hashlib.sha256(s2.encode()).hexdigest() == te.TABLE_S2_TEXT_SHA256
    assert [(s.number, s.sampled_on.isoformat(), s.ph) for s in water[:3]] == [
        (1, "2022-05-09", 7.0),
        (2, "2022-06-19", 6.5),
        (3, "2022-07-06", 6.5),
    ]


@pytest.mark.data
@requires_raw
def test_the_deposited_tables_give_the_measured_counts_and_pca() -> None:
    """66,162 values over 11,027 rows, sample_10 the one QC failure, the PCA exact."""
    report = te.loader_report(DATA_ROOT)
    inventory = report["inventory"]
    assert inventory["stored_records"] == inventory["released_values"] == 66_162
    assert inventory["rows_per_water_sample"] == {"1": 3_675, "2": 3_679, "3": 3_673}
    assert inventory["symbols"] == te.TABLE_S4_DISTINCT_SYMBOLS
    replicates = inventory["replicates_per_condition"]
    assert replicates.pop("lake water #1, filtered, day 4") == 2
    assert set(replicates.values()) == {3}
    assert report["qc_failed"] == list(te.QC_FAILED_SAMPLES)
    assert all(
        max(p["max_abs_pc1"], p["max_abs_pc2"], p["max_abs_distance"])
        <= te.PCA_TOLERANCE
        for p in report["pca"]
    )
