# tests/torchcell/datasets/ecoli/test_mori2021.py
# [[tests.torchcell.datasets.ecoli.test_mori2021]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_mori2021.py
"""Mori 2021 proteome loader (``torchcell.datasets.ecoli.mori2021``).

The synthetic tests write the six-file release layout the loader consumes: a minimal
Appendix ``.docx`` carrying the Appendix quotes, the Dataset EV1 strain sheet, the
Dataset EV2 and EV3 metadata sheets for all 66 declared samples, and the Dataset EV8 and
EV9 mass-fraction sheets with 34 rows chosen to exercise one row rule each. They drive
the readers, the three released-bytes checks, the row rules, the phenotype and reference
builders, the raw-mirror deposit and a full ``process()`` build into ``tmp_path``. The
genome, the locus-tag reconciliation and the assembly pin are in-test objects, so
nothing reads ``$DATA_ROOT``.

The ``@pytest.mark.data`` tests read the real mirrors and the dev-tree LMDB and pin the
numbers the dendron note states: 4,342 released rows, 2,077 quantified in a loaded
sample, 2,073 protein keys, 7 records from 66 released samples, the per-column key
counts, the Dataset EV9 identity fault, the per-column normalization, and the Schmidt
2016 non-overlap (1,812 shared accessions, Pearson r 0.796 on log10 mass fraction, zero
pairs agreeing to 1e-6).
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
import statistics
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

import torchcell.datasets.ecoli.mori2021 as mo
from torchcell.data import file_sha256
from torchcell.datamodels.schema import (
    AssemblyReferenceGenome,
    BacterialAssemblySet,
    BacterialGeneNamespace,
    EnvironmentPhysicalPerturbation,
)
from torchcell.datasets.bacteria_common import LocusTagReconciliation
from torchcell.sequence.genome.base import GeneNameStatus

ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_MG1655_ASM584v2"
NAMESPACE: BacterialGeneNamespace = "ecoli_k12_mg1655_bnumber"

#: The 18 rows neither released sheet identifies, as the release names them.
UNIDENTIFIED = (
    *(f"insAB-{i}" for i in range(1, 7)),
    *(f"insCD-{i}" for i in range(1, 7)),
    *(f"insEF-{i}" for i in range(1, 6)),
    "insJK",
)
#: One quantified blank-locus row, so the no-locus rule has an item.
QUANTIFIED_BLANK = "ygaU"
#: ``(gene name, b-number)`` of the rows EV9 does identify.
IDENTIFIED: tuple[tuple[str, str], ...] = (
    ("keepA", "b0001"),
    ("keepB", "b0002"),
    ("ghost", "b0003"),
    ("silent", "b0004"),
)
KEPT_GENES = ("keepA", "keepB")
UNRESOLVED = ("b0003",)
#: Rows with a non-zero mass fraction in the loaded columns.
QUANTIFIED = (QUANTIFIED_BLANK, "keepA", "keepB", "ghost")
GENE_ORDER = (*UNIDENTIFIED, *mo.EV8_SUPPLIES_LOCUS, *(g for g, _ in IDENTIFIED))
N_SOURCE_ROWS = len(GENE_ORDER)

EV2_HEADERS = (
    "Sample ID",
    "Short Description",
    "Description",
    "Strain",
    "Base medium",
    "Carbon Source",
    "Nitrogen Source",
    "Supplement",
    "Special condition/procedure",
    "Doubling time (min)",
    "DDA file name",
    "SWATH file name",
)
EV3_HEADERS = (
    "Sample ID",
    "Group",
    "Growth rate (1/h)",
    "Strain",
    "Growth medium",
    "Carbon source",
    "Nitrogen source",
    "Supplement",
    "Description",
    "SWATH file name",
)


def _mass_fraction(gene: str, column: str) -> float:
    """A deterministic mass fraction; every column sums to 1 over the quantified rows.

    The two kept rows trade mass across the seven loaded columns, so the three cultures'
    means differ and the reference standard error is non-zero and checkable.
    """
    if gene not in QUANTIFIED:
        return 0.0
    loaded = [spec.column for spec in mo.LOADED]
    step = 0.01 * loaded.index(column) if column in loaded else 0.0
    return {
        QUANTIFIED_BLANK: 0.10,
        "keepA": 0.30 + step,
        "keepB": 0.40 - step,
        "ghost": 0.20,
    }[gene]


# --------------------------------------------------------------------------- #
# The synthetic release
# --------------------------------------------------------------------------- #
_DOCX_XML_HEAD = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:body>"
)


def write_appendix(path: Path, extra: str = "") -> Path:
    """Write a ``.docx`` whose ``word/document.xml`` carries the Appendix quotes."""
    paragraphs = [*mo.APPENDIX_QUOTES.values()]
    if extra:
        paragraphs.append(extra)
    body = "".join(
        f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'
        for text in paragraphs
    )
    with zipfile.ZipFile(path, "w") as book:
        book.writestr(
            "word/document.xml", _DOCX_XML_HEAD + body + "</w:body></w:document>"
        )
    return path


def write_ev1(path: Path) -> Path:
    """Dataset EV1: the strain table, carrying the quoted EQ353 row."""
    book = openpyxl.Workbook()
    description = book.active
    description.title = mo.SHEET_DESCRIPTION
    description.append(["Mori et al. 2021"])
    sheet = book.create_sheet(mo.SHEET_EV1)
    sheet.append(["Strain Name", "Description", "Source"])
    sheet.append(["NCM3722", "Wild type E. coli strain", None])
    sheet.append(
        [
            "MG1655 (EQ353)",
            "Wild type E. coli strain - same strain used in Li et al. (2014)",
            "Originarily obtained from Carol Gross Lab",
        ]
    )
    book.save(path)
    return path


def write_ev2(path: Path) -> Path:
    """Dataset EV2: the Samples-1 metadata of the 30 Dataset EV8 columns."""
    book = openpyxl.Workbook()
    description = book.active
    description.title = mo.SHEET_DESCRIPTION
    description.append([mo._Q_EV8_DESCRIPTION])
    sheet = book.create_sheet(mo.SHEET_EV2)
    sheet.append(list(EV2_HEADERS))
    for spec in mo.SAMPLES:
        if spec.table != "EV8":
            continue
        row: list[Any] = [None] * len(EV2_HEADERS)
        row[0] = spec.column
        row[3] = spec.strain
        row[4] = spec.medium
        row[6] = "10 mM NH4Cl"
        sheet.append(row)
    book.save(path)
    return path


def write_ev3(path: Path) -> Path:
    """Dataset EV3: the Samples-2 metadata of the 36 Dataset EV9 columns."""
    book = openpyxl.Workbook()
    description = book.active
    description.title = mo.SHEET_DESCRIPTION
    description.append([mo._Q_EV3_DESCRIPTION])
    sheet = book.create_sheet(mo.SHEET_EV3)
    sheet.append(list(EV3_HEADERS))
    for index, spec in enumerate(s for s in mo.SAMPLES if s.table == "EV9"):
        row: list[Any] = [None] * len(EV3_HEADERS)
        row[0] = spec.column
        row[2] = 0.69 + 0.01 * index
        row[3] = spec.strain
        row[4] = spec.medium
        row[5] = mo.LOADED_CARBON_SOURCE if spec.drop is None else "0.2% glucose"
        row[6] = (
            mo.LOADED_NITROGEN_SOURCE if spec.drop is None else "11.34 mM (NH4)2SO4"
        )
        sheet.append(row)
    book.save(path)
    return path


def _write_mass_fractions(path: Path, table: str, description: str) -> Path:
    """Dataset EV8 or EV9: the identity block then that table's sample columns."""
    book = openpyxl.Workbook()
    head = book.active
    head.title = mo.SHEET_DESCRIPTION
    head.append([description])
    sheet = book.create_sheet(mo.SHEET_EV8 if table == "EV8" else mo.SHEET_EV9)
    columns = [spec.column for spec in mo.SAMPLES if spec.table == table]
    sheet.append([*mo.IDENTITY_COLUMNS, *columns])
    locus_of = dict(IDENTIFIED)
    for gene in GENE_ORDER:
        if gene in UNIDENTIFIED:
            locus = None
        elif gene in mo.EV8_SUPPLIES_LOCUS:
            locus = mo.EV8_SUPPLIES_LOCUS[gene] if table == "EV8" else None
        else:
            locus = locus_of[gene]
        protein = None if gene in UNIDENTIFIED else f"P{abs(hash(gene)) % 90000:05d}"
        sheet.append(
            [
                gene,
                locus,
                protein,
                *(_mass_fraction(gene, column) for column in columns),
            ]
        )
    book.save(path)
    return path


def write_release(directory: Path) -> dict[str, str]:
    """Write all six consumed files into ``directory`` and return their paths."""
    directory.mkdir(parents=True, exist_ok=True)
    write_appendix(directory / mo.APPENDIX)
    write_ev1(directory / mo.EV1)
    write_ev2(directory / mo.EV2)
    write_ev3(directory / mo.EV3)
    _write_mass_fractions(directory / mo.EV8, "EV8", mo._Q_EV8_DESCRIPTION)
    _write_mass_fractions(directory / mo.EV9, "EV9", mo._Q_EV9_DESCRIPTION)
    return {name: str(directory / name) for name in mo.CONSUMED}


@pytest.fixture
def release(tmp_path: Path) -> dict[str, str]:
    """The six synthetic release files."""
    return write_release(tmp_path / "release")


# --------------------------------------------------------------------------- #
# The declared sample set
# --------------------------------------------------------------------------- #
def test_the_sixty_six_released_samples_split_into_seven_records_and_three_rules() -> (
    None
):
    """The ledger arithmetic is a property of the declaration, not of the build."""
    assert len(mo.SAMPLES) == 66
    assert len([s for s in mo.SAMPLES if s.table == "EV8"]) == 30
    assert len([s for s in mo.SAMPLES if s.table == "EV9"]) == 36
    assert len(mo.LOADED) == mo.EXPECTED_RECORDS == 7
    counts = {
        reason.rule: len([s for s in mo.SAMPLES if s.drop == reason])
        for reason in (
            mo.DROP_NO_MEDIA_ENTRY,
            mo.DROP_MEDIA_ENTRY_DIFFERS,
            mo.DROP_FORMULATION_NOT_STATED,
        )
    }
    assert counts == {
        "medium_has_no_media_library_entry": 52,
        "medium_differs_from_the_library_entry_it_names": 2,
        "medium_formulation_not_stated_by_the_source": 5,
    }
    assert sum(counts.values()) + len(mo.LOADED) == len(mo.SAMPLES)
    assert {s.strain for s in mo.LOADED} == {"EQ353"}
    assert {s.medium for s in mo.LOADED} == {"MOPS (Neidhardt)"}
    assert mo.CULTURES == {
        "A1": ("A1-1", "A1-2", "A1-3"),
        "C1": ("C1",),
        "F1": ("F1-1", "F1-2", "F1-3"),
    }


def test_the_two_mops_neidhardt_samples_with_the_wrong_ammonium_are_their_own_rule() -> (
    None
):
    """``Lib-28`` and ``Lib-30`` name the library object but are not its medium."""
    named = [s.column for s in mo.SAMPLES if s.medium == "MOPS (Neidhardt's)"]
    assert named == ["Lib-28", "Lib-30"]
    for column in named:
        spec = next(s for s in mo.SAMPLES if s.column == column)
        assert spec.drop == mo.DROP_MEDIA_ENTRY_DIFFERS
        assert spec.drop.needed_addition == "MOPS_MINIMAL_20MM_NH4CL_MORI2021"


def test_the_lb_samples_are_dropped_for_an_unstated_formulation() -> None:
    """Both LB objects exist, so the drop is about which one the source weighed."""
    columns = [s.column for s in mo.SAMPLES if s.drop == mo.DROP_FORMULATION_NOT_STATED]
    assert columns == ["Lib-12", "Lib-19", "Lib-20", "Lib-21", "Lib-22"]
    assert {s.medium for s in mo.SAMPLES if s.column in columns} == {"LB"}


# --------------------------------------------------------------------------- #
# Reading the pinned Appendix
# --------------------------------------------------------------------------- #
def test_docx_reader_joins_runs_decodes_entities_and_drops_blanks(
    tmp_path: Path,
) -> None:
    path = tmp_path / "a.docx"
    body = (
        "<w:p><w:r><w:t>A &amp; B</w:t></w:r><w:r><w:t> &lt;c&gt;</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>   </w:t></w:r></w:p>"
        "<w:p><w:r><w:t>second</w:t></w:r></w:p>"
    )
    with zipfile.ZipFile(path, "w") as book:
        book.writestr(
            "word/document.xml", _DOCX_XML_HEAD + body + "</w:body></w:document>"
        )
    assert mo.docx_paragraphs(path) == ["A & B <c>", "second"]
    assert mo.appendix_text(path) == "A & B <c>\nsecond"


def test_check_quotes_reads_the_appendix_and_every_workbook(
    release: dict[str, str],
) -> None:
    summary = mo.check_quotes(release)
    assert summary["n_quotes_checked"] == len(mo.APPENDIX_QUOTES) + sum(
        len(q) for q in mo.WORKBOOK_QUOTES.values()
    )


def test_check_quotes_refuses_an_appendix_missing_a_quote(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "short.docx"
    with zipfile.ZipFile(path, "w") as book:
        book.writestr("word/document.xml", _DOCX_XML_HEAD + "</w:body></w:document>")
    broken = {**release, mo.APPENDIX: str(path)}
    with pytest.raises(RuntimeError, match="Appendix quote"):
        mo.check_quotes(broken)


def test_check_quotes_refuses_a_workbook_missing_its_description(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "ev9.xlsx"
    _write_mass_fractions(path, "EV9", "a different description")
    with pytest.raises(RuntimeError, match="quote 'description'"):
        mo.check_quotes({**release, mo.EV9: str(path)})


# --------------------------------------------------------------------------- #
# The released-bytes checks
# --------------------------------------------------------------------------- #
def test_sample_columns_must_be_the_declared_headers_in_order(
    release: dict[str, str],
) -> None:
    assert mo.check_sample_columns(release[mo.EV8], "EV8") == tuple(
        s.column for s in mo.SAMPLES if s.table == "EV8"
    )
    assert mo.check_sample_columns(release[mo.EV9], "EV9") == tuple(
        s.column for s in mo.SAMPLES if s.table == "EV9"
    )


def test_sample_columns_refuse_a_renamed_column(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "renamed.xlsx"
    book = openpyxl.load_workbook(release[mo.EV9])
    book[mo.SHEET_EV9].cell(row=1, column=4).value = "A1-9"
    book.save(path)
    with pytest.raises(RuntimeError, match="sample headers are"):
        mo.check_sample_columns(str(path), "EV9")


def test_normalization_holds_and_refuses_a_column_that_stops_summing_to_one(
    release: dict[str, str], tmp_path: Path
) -> None:
    for table, name in (("EV8", mo.EV8), ("EV9", mo.EV9)):
        summary = mo.check_normalization(release[name], table)
        assert summary["worst_abs_deviation"] < mo.NORMALIZATION_ATOL
        assert summary["n_columns"] == len([s for s in mo.SAMPLES if s.table == table])
    path = tmp_path / "unnormalized.xlsx"
    book = openpyxl.load_workbook(release[mo.EV9])
    sheet = book[mo.SHEET_EV9]
    sheet.cell(row=2 + GENE_ORDER.index("keepA"), column=4).value = 0.5
    book.save(path)
    with pytest.raises(RuntimeError, match="rather than 1"):
        mo.check_normalization(str(path), "EV9")


def test_identity_fault_pins_which_rows_ev8_identifies_and_ev9_does_not(
    release: dict[str, str],
) -> None:
    summary = mo.check_identity_fault(release[mo.EV8], release[mo.EV9])
    assert summary == {
        "n_blank_in_ev9": len(UNIDENTIFIED) + len(mo.EV8_SUPPLIES_LOCUS),
        "n_supplied_by_ev8": len(mo.EV8_SUPPLIES_LOCUS),
        "n_unidentified": mo.UNIDENTIFIED_ROWS,
    }


def test_identity_fault_refuses_a_changed_row_order(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "reordered.xlsx"
    book = openpyxl.load_workbook(release[mo.EV8])
    book[mo.SHEET_EV8].cell(row=2, column=1).value = "elsewhere"
    book.save(path)
    with pytest.raises(RuntimeError, match="one row order"):
        mo.check_identity_fault(str(path), release[mo.EV9])


def test_identity_fault_refuses_a_changed_supplied_set(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "nolocus.xlsx"
    book = openpyxl.load_workbook(release[mo.EV8])
    row = 2 + GENE_ORDER.index("ybaO")
    book[mo.SHEET_EV8].cell(row=row, column=2).value = None
    book.save(path)
    with pytest.raises(RuntimeError, match="now supplies"):
        mo.check_identity_fault(str(path), release[mo.EV9])


def test_sample_metadata_must_match_the_declaration(release: dict[str, str]) -> None:
    summary = mo.check_sample_metadata(release[mo.EV2], release[mo.EV3])
    assert summary == {"n_samples": len(mo.SAMPLES), "n_loaded": len(mo.LOADED)}


def test_sample_metadata_refuses_a_changed_medium(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "ev3.xlsx"
    book = openpyxl.load_workbook(release[mo.EV3])
    book[mo.SHEET_EV3].cell(row=2, column=5).value = "M9"
    book.save(path)
    with pytest.raises(RuntimeError, match="the module declares"):
        mo.check_sample_metadata(release[mo.EV2], str(path))


def test_sample_metadata_refuses_a_loaded_sample_whose_ammonium_changed(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "ev3.xlsx"
    book = openpyxl.load_workbook(release[mo.EV3])
    book[mo.SHEET_EV3].cell(row=2, column=7).value = "20 mM NH4Cl"
    book.save(path)
    with pytest.raises(RuntimeError, match="MOPS_MINIMAL is not its medium"):
        mo.check_sample_metadata(release[mo.EV2], str(path))


def test_growth_rates_are_read_for_every_ev3_sample(release: dict[str, str]) -> None:
    rates = mo.read_growth_rates(release[mo.EV3])
    assert set(rates) == {s.column for s in mo.SAMPLES if s.table == "EV9"}
    assert rates["A1-1"] == pytest.approx(0.69)


# --------------------------------------------------------------------------- #
# Reading the mass fractions
# --------------------------------------------------------------------------- #
def test_read_mass_fractions_keeps_the_released_cells_verbatim(
    release: dict[str, str],
) -> None:
    columns = [spec.column for spec in mo.LOADED]
    rows = mo.read_mass_fractions(release[mo.EV9], columns)
    assert [row.gene_name for row in rows] == list(GENE_ORDER)
    assert len(rows) == N_SOURCE_ROWS
    keep_a = next(row for row in rows if row.gene_name == "keepA")
    assert keep_a.gene_locus == "b0001"
    assert keep_a.mass_fraction["A1-1"] == pytest.approx(0.30)
    assert keep_a.mass_fraction["F1-3"] == pytest.approx(0.36)
    assert keep_a.quantified("A1-1")
    silent = next(row for row in rows if row.gene_name == "silent")
    assert not silent.quantified("A1-1")
    assert (
        next(row for row in rows if row.gene_name == QUANTIFIED_BLANK).gene_locus
        is None
    )


def test_read_mass_fractions_refuses_a_renamed_identity_block(
    release: dict[str, str], tmp_path: Path
) -> None:
    path = tmp_path / "ev9.xlsx"
    book = openpyxl.load_workbook(release[mo.EV9])
    book[mo.SHEET_EV9].cell(row=1, column=2).value = "Locus ID"
    book.save(path)
    with pytest.raises(RuntimeError, match="identity block is"):
        mo.read_mass_fractions(str(path), [spec.column for spec in mo.LOADED])


def test_read_mass_fractions_refuses_a_missing_column_and_a_non_number(
    release: dict[str, str], tmp_path: Path
) -> None:
    with pytest.raises(RuntimeError, match="no sample column"):
        mo.read_mass_fractions(release[mo.EV9], ["not-a-sample"])
    path = tmp_path / "ev9.xlsx"
    book = openpyxl.load_workbook(release[mo.EV9])
    book[mo.SHEET_EV9].cell(row=2 + GENE_ORDER.index("keepA"), column=4).value = "NA"
    book.save(path)
    with pytest.raises(RuntimeError, match="is not a"):
        mo.read_mass_fractions(str(path), [spec.column for spec in mo.LOADED])


# --------------------------------------------------------------------------- #
# The row rules
# --------------------------------------------------------------------------- #
class _Locus:
    def __init__(self, symbol: str | None) -> None:
        self.symbol = symbol


class _Annotation:
    def __init__(self, loci: dict[str, _Locus]) -> None:
        self.loci = loci


class FakeGenome:
    """Carries the kept synthetic b-numbers as its loci, and nothing else."""

    ASSEMBLY_SET = ASSEMBLY_SET

    def __init__(self) -> None:
        """One locus per kept b-number."""
        self.genbank = _Annotation(
            {
                locus: _Locus(gene)
                for gene, locus in IDENTIFIED
                if locus not in UNRESOLVED
            }
        )


def _reconciliation(names: pd.Series, label: str) -> LocusTagReconciliation:
    unresolved = [n for n in names if n in UNRESOLVED]
    resolved = len(set(names)) - len(unresolved)
    return LocusTagReconciliation(
        label=label,
        assembly_set=ASSEMBLY_SET,
        gene_namespace=NAMESPACE,
        unique_names=len(set(names)),
        status_histogram={
            GeneNameStatus.CURRENT: resolved,
            GeneNameStatus.RENAMED: 0,
            GeneNameStatus.NON_GENE_FEATURE: 0,
            GeneNameStatus.RETIRED: len(unresolved),
            GeneNameStatus.AMBIGUOUS: 0,
        },
        layer_histogram={"locus tag": resolved, "not found": len(unresolved)},
        remapped=0,
        kept_on_collision=(),
        retired_kept=tuple(unresolved),
        ambiguous_kept={},
        case_insensitive=(),
        outside_namespace=tuple(unresolved),
    )


def _identity_reconciliation(
    genome: Any, names: pd.Series, *, label: str
) -> tuple[pd.Series, LocusTagReconciliation]:
    """Store each b-number as itself, and mark ``UNRESOLVED`` outside the namespace."""
    return names, _reconciliation(names, label)


def _reference_genome() -> AssemblyReferenceGenome:
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain=mo.EQ353_BACKGROUND.name,
        ploidy="haploid",
        assembly_set=ASSEMBLY_SET,
        assembly_accession="GCA_000005845.2",
    )


@pytest.fixture
def patched_resolution(monkeypatch: pytest.MonkeyPatch) -> None:
    """Resolve b-numbers to themselves and pin the synthetic release's counts."""
    monkeypatch.setattr(mo, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(mo, "MIN_RESOLVED_FRACTION", 0.5)
    monkeypatch.setattr(mo, "EXPECTED_SOURCE_ROWS", N_SOURCE_ROWS)
    monkeypatch.setattr(mo, "EXPECTED_PROTEIN_KEYS", len(KEPT_GENES))


def test_select_proteins_applies_the_three_row_rules(
    release: dict[str, str], patched_resolution: None
) -> None:
    rows = mo.read_mass_fractions(release[mo.EV9], [spec.column for spec in mo.LOADED])
    selection = mo.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    assert [row.gene_name for row in selection.kept] == list(KEPT_GENES)
    assert selection.locus_tag == {"keepA": "b0001", "keepB": "b0002"}
    by_rule = {rule.rule: rule for rule in selection.rules}
    assert by_rule["not_quantified_in_any_loaded_sample"].n_items == (
        N_SOURCE_ROWS - len(QUANTIFIED)
    )
    assert by_rule["release_files_no_gene_locus_for_this_row"].n_items == 1
    assert (
        by_rule["release_files_no_gene_locus_for_this_row"]
        .items[0]
        .startswith(QUANTIFIED_BLANK)
    )
    assert by_rule["b_number_resolves_to_no_locus_of_the_pinned_assembly"].items == [
        "ghost (b0003)"
    ]
    assert selection.dropped_rows == N_SOURCE_ROWS - len(KEPT_GENES)


def test_select_proteins_refuses_a_release_below_the_resolution_floor(
    release: dict[str, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(mo, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(mo, "MIN_RESOLVED_FRACTION", 0.999)
    rows = mo.read_mass_fractions(release[mo.EV9], [spec.column for spec in mo.LOADED])
    from torchcell.datasets.bacteria_common import LocusTagResolutionError

    with pytest.raises(LocusTagResolutionError):
        mo.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]


# --------------------------------------------------------------------------- #
# Environment, phenotype, reference
# --------------------------------------------------------------------------- #
def test_environment_is_neidhardt_mops_plus_the_released_carbon_source() -> None:
    environment = mo.build_environment()
    assert environment.media.base_medium == "MOPS_MINIMAL"
    assert environment.temperature is not None
    assert environment.temperature.value == 37.0
    assert environment.aerobicity == "aerobic"
    (perturbation,) = environment.perturbations
    assert isinstance(perturbation, EnvironmentPhysicalPerturbation)
    assert perturbation.factor == "carbon_source"
    magnitude = perturbation.magnitude
    assert magnitude is not None
    assert magnitude.value == pytest.approx(0.2)
    assert magnitude.unit == "percent_w/v"
    assert perturbation.agent is not None
    assert perturbation.agent.name == "D-glucose"
    # The ammonium is a component of the medium, not a second edit of it.
    ammonium = [
        c
        for c in environment.media.components
        if c.compound.name == "ammonium chloride"
    ]
    assert len(ammonium) == 1
    assert ammonium[0].concentration is not None
    assert ammonium[0].concentration.value == pytest.approx(9.5)


def test_phenotype_is_one_injection_with_no_replicate_and_no_error(
    release: dict[str, str], patched_resolution: None
) -> None:
    rows = mo.read_mass_fractions(release[mo.EV9], [spec.column for spec in mo.LOADED])
    selection = mo.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    phenotype = mo.build_phenotype(selection.kept, selection.locus_tag, "A1-2")
    assert phenotype.protein_abundance == {
        "b0001": pytest.approx(0.31),
        "b0002": pytest.approx(0.39),
    }
    assert phenotype.protein_abundance_se is None
    assert set(phenotype.n_replicates.values()) == {1}
    assert phenotype.measurement_type == mo.MEASUREMENT_TYPE


def test_reference_is_the_mean_of_three_culture_means_with_an_exact_error(
    release: dict[str, str], patched_resolution: None
) -> None:
    rows = mo.read_mass_fractions(release[mo.EV9], [spec.column for spec in mo.LOADED])
    selection = mo.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    reference = mo.build_reference_phenotype(selection.kept, selection.locus_tag)
    means = [
        statistics.fmean([_mass_fraction("keepA", c) for c in columns])
        for columns in mo.CULTURES.values()
    ]
    assert reference.protein_abundance["b0001"] == pytest.approx(
        statistics.fmean(means)
    )
    assert reference.n_replicates["b0001"] == 3
    assert reference.protein_abundance_se is not None
    assert reference.protein_abundance_se["b0001"] == pytest.approx(
        statistics.stdev(means) / math.sqrt(3)
    )


def test_reference_stores_nan_where_only_one_culture_quantified_the_protein(
    release: dict[str, str], tmp_path: Path, patched_resolution: None
) -> None:
    """Only the C1 culture quantifies ``keepB``, so its SE is ``nan`` and n is 1."""
    path = tmp_path / "ev9.xlsx"
    book = openpyxl.load_workbook(release[mo.EV9])
    sheet = book[mo.SHEET_EV9]
    row = 2 + GENE_ORDER.index("keepB")
    for spec in mo.LOADED:
        if spec.culture == "C1":
            continue
        column = (
            1
            + mo.N_IDENTITY_COLUMNS
            + [s.column for s in mo.SAMPLES if s.table == "EV9"].index(spec.column)
        )
        sheet.cell(row=row, column=column).value = 0.0
    book.save(path)
    rows = mo.read_mass_fractions(str(path), [spec.column for spec in mo.LOADED])
    selection = mo.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    reference = mo.build_reference_phenotype(selection.kept, selection.locus_tag)
    assert reference.n_replicates["b0002"] == 1
    assert reference.protein_abundance_se is not None
    assert math.isnan(reference.protein_abundance_se["b0002"])


def test_restrict_refuses_a_key_the_reference_does_not_quantify(
    release: dict[str, str], patched_resolution: None
) -> None:
    rows = mo.read_mass_fractions(release[mo.EV9], [spec.column for spec in mo.LOADED])
    selection = mo.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    reference = mo.build_reference_phenotype(selection.kept, selection.locus_tag)
    assert set(mo.restrict(reference, ["b0001"]).protein_abundance) == {"b0001"}
    with pytest.raises(RuntimeError, match="quantifies none of"):
        mo.restrict(reference, ["b9999"])


def test_the_eq353_background_is_an_edit_of_the_mg1655_assembly() -> None:
    from torchcell.datasets.bacteria_common import BACTERIAL_ASSEMBLY_SETS

    assert mo.EQ353_BACKGROUND.reference_strain == mo.MORI_REFERENCE_STRAIN == "MG1655"
    assert mo.MORI_ASSEMBLY_SET == BACTERIAL_ASSEMBLY_SETS[mo.MORI_REFERENCE_STRAIN]
    assert mo.EQ353_BACKGROUND.assembly_set == ASSEMBLY_SET
    assert mo.EQ353_BACKGROUND.alleles == []
    assert mo.EQ353_BACKGROUND.provenance is not None
    assert len(mo.EQ353_BACKGROUND.provenance) == 2


def test_publication_carries_the_resolved_identifiers() -> None:
    pub = mo.publication()
    assert pub.pubmed_id == "34032011"
    assert pub.doi == "10.15252/msb.20209536"


# --------------------------------------------------------------------------- #
# The raw mirror
# --------------------------------------------------------------------------- #
def test_mirror_paths_follow_the_citation_key(tmp_path: Path) -> None:
    assert mo.raw_mirror_dir(str(tmp_path)).name == mo.CITATION_KEY
    assert mo.raw_mirror_dir(str(tmp_path)).parent.name == "torchcell-raw"
    assert mo.library_dir(str(tmp_path)).parent.name == "torchcell-library"


def test_every_consumed_file_records_a_scriptable_pmc_cloud_retrieval() -> None:
    assert set(mo.CONSUMED) == set(mo.RETRIEVALS) == set(mo.MIRROR_RELPATH)
    for name, record in mo.RETRIEVALS.items():
        assert record.method == "pmc_cloud"
        assert record.sha256 == mo.DATA_SHA256[name]
        assert str(record.source_url).startswith(
            "https://pmc-oa-opendata.s3.amazonaws.com/PMC8144880.1/"
        )
        assert record.retriever == "torchcell.literature.retrieve.pmc_cloud_object"


def test_deposit_writes_every_file_with_its_manifest_and_is_idempotent(
    release: dict[str, str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = Path(release[mo.EV9]).parent
    shas = {name: file_sha256(path) for name, path in release.items()}
    monkeypatch.setattr(mo, "DATA_SHA256", shas)
    data_root = tmp_path / "root"
    first = mo.deposit_raw_mirror(source_dir=source, data_root=str(data_root))
    manifest = mo.load_manifest(str(data_root))
    assert {record.path for record in manifest.files} == set(mo.MIRROR_RELPATH.values())
    for name, relpath in mo.MIRROR_RELPATH.items():
        assert mo.manifest_sha256(manifest, relpath) == shas[name]
        assert (first / relpath).exists()
    assert mo.deposit_raw_mirror(source_dir=source, data_root=str(data_root)) == first


def test_deposit_refuses_a_file_whose_bytes_are_not_the_pin(
    release: dict[str, str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = Path(release[mo.EV9]).parent
    monkeypatch.setattr(mo, "DATA_SHA256", {name: "0" * 64 for name in mo.CONSUMED})
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        mo.deposit_raw_mirror(source_dir=source, data_root=str(tmp_path / "root"))


def test_manifest_sha256_refuses_an_unrecorded_path(
    release: dict[str, str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = Path(release[mo.EV9]).parent
    monkeypatch.setattr(
        mo, "DATA_SHA256", {n: file_sha256(p) for n, p in release.items()}
    )
    data_root = tmp_path / "root"
    mo.deposit_raw_mirror(source_dir=source, data_root=str(data_root))
    with pytest.raises(KeyError):
        mo.manifest_sha256(mo.load_manifest(str(data_root)), "data/nope.xlsx")


def test_download_links_every_mirror_file_and_refuses_a_missing_one(
    release: dict[str, str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = Path(release[mo.EV9]).parent
    shas = {name: file_sha256(path) for name, path in release.items()}
    monkeypatch.setattr(mo, "DATA_SHA256", shas)
    data_root = tmp_path / "root"
    mo.deposit_raw_mirror(source_dir=source, data_root=str(data_root))
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    root = tmp_path / "build"
    (root / "raw").mkdir(parents=True)
    dataset = mo.ProteomeMori2021Dataset.__new__(mo.ProteomeMori2021Dataset)
    monkeypatch.setattr(
        type(dataset), "raw_dir", property(lambda self: str(root / "raw"))
    )
    dataset.download()
    for name in mo.CONSUMED:
        assert file_sha256(str(root / "raw" / name)) == shas[name]
    (data_root / mo.RAW_DIR_REL / mo.MIRROR_RELPATH[mo.EV9]).unlink()
    with pytest.raises(RuntimeError, match="missing from mirror"):
        dataset.download()


# --------------------------------------------------------------------------- #
# End-to-end build on the synthetic release
# --------------------------------------------------------------------------- #
def test_process_builds_one_record_per_loaded_sample(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proteome_mori2021"
    paths = write_release(root / "raw")
    monkeypatch.setattr(
        mo, "DATA_SHA256", {n: file_sha256(p) for n, p in paths.items()}
    )
    monkeypatch.setattr(mo, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(mo, "MIN_RESOLVED_FRACTION", 0.5)
    monkeypatch.setattr(mo, "EXPECTED_SOURCE_ROWS", N_SOURCE_ROWS)
    monkeypatch.setattr(mo, "EXPECTED_PROTEIN_KEYS", len(KEPT_GENES))
    monkeypatch.setattr(
        mo, "assembly_reference", lambda strain, *, background=None: _reference_genome()
    )

    dataset = mo.ProteomeMori2021Dataset(
        root=str(root),
        ecoli_genome=FakeGenome(),  # type: ignore[arg-type]
    )
    assert len(dataset) == mo.EXPECTED_RECORDS
    assert dataset.gene_set == {"b0001", "b0002"}

    items = [dataset.transform_item(dataset[i]) for i in range(len(dataset))]
    assert {len(i["experiment"].genotype.perturbations) for i in items} == {0}
    assert {i["experiment"].phenotype.measurement_type for i in items} == {
        mo.MEASUREMENT_TYPE
    }
    assert {i["experiment"].phenotype.protein_abundance_se for i in items} == {None}
    assert {i["publication"].doi for i in items} == {mo.PAPER_DOI}
    assert {i["reference"].genome_reference.strain for i in items} == {
        mo.EQ353_BACKGROUND.name
    }
    # The seven records are seven measurements of ONE condition.
    assert len({i["experiment"].environment.model_dump_json() for i in items}) == 1
    stored = sorted(i["experiment"].phenotype.protein_abundance["b0001"] for i in items)
    assert stored == [
        pytest.approx(0.30 + 0.01 * index) for index in range(mo.EXPECTED_RECORDS)
    ]
    for item in items:
        reference = item["reference"].phenotype_reference
        assert set(reference.protein_abundance) == set(
            item["experiment"].phenotype.protein_abundance
        )
        assert reference.n_replicates["b0001"] == len(mo.CULTURES)
    dataset.close_lmdb()

    preprocess = Path(dataset.preprocess_dir)
    ledger = json.loads((preprocess / "dropped_records.json").read_text())
    assert ledger["source_samples"] == 66
    assert ledger["kept_records"] == mo.EXPECTED_RECORDS
    assert ledger["dropped_records"] == 59
    assert ledger["source_protein_rows"] == N_SOURCE_ROWS
    assert ledger["kept_protein_keys"] == len(KEPT_GENES)
    assert ledger["dropped_protein_rows"] == N_SOURCE_ROWS - len(KEPT_GENES)
    checks = json.loads((preprocess / "released_statistics_check.json").read_text())
    assert checks["ev9_identity_fault"]["n_supplied_by_ev8"] == len(
        mo.EV8_SUPPLIES_LOCUS
    )
    assert [row["sheet"] for row in checks["normalization"]] == [
        mo.SHEET_EV8,
        mo.SHEET_EV9,
    ]
    sourced = json.loads((preprocess / "sourced_values.json").read_text())
    assert sourced["reference_strain"]["value"] == "MG1655"
    assert sourced["stored_quantity"]["value"] == "absolute protein mass fraction"
    identifiers = pd.read_csv(preprocess / "protein_identifiers.csv")
    assert list(identifiers["released_gene_name"]) == list(KEPT_GENES)
    samples = pd.read_csv(preprocess / "samples.csv")
    assert list(samples["column"]) == [spec.column for spec in mo.LOADED]
    assert set(samples["culture"]) == set(mo.CULTURES)
    assert (preprocess / "build_manifest.json").exists()

    report = mo.verify_build(
        str(root),
        genome=FakeGenome(),  # type: ignore[arg-type]
        expected_count=mo.EXPECTED_RECORDS,
    )
    assert [(r.level, r.name) for r in report.results if not r.passed] == []
    assert (preprocess / "verification_report.json").exists()


def test_drop_log_check_refuses_an_unaccounted_drop() -> None:
    reconciliation = _reconciliation(pd.Series(["b0001"]), "test")
    log_model = mo.DropLog(
        dataset="test",
        source_samples=66,
        kept_records=7,
        dropped_records=59,
        source_protein_rows=10,
        kept_protein_keys=2,
        dropped_protein_rows=8,
        rules=[
            mo.DropRule(
                rule="medium_has_no_media_library_entry",
                scope="sample",
                description="x",
                n_items=58,
                items=[],
            ),
            mo.DropRule(
                rule="not_quantified_in_any_loaded_sample",
                scope="protein_row",
                description="x",
                n_items=8,
                items=[],
            ),
        ],
        reconciliation=reconciliation,
        notes=[],
    )
    with pytest.raises(RuntimeError, match="dropped samples !="):
        log_model.check()
    fixed = log_model.model_copy(update={"dropped_records": 58, "source_samples": 65})
    fixed.check()
    mismatched = fixed.model_copy(update={"dropped_protein_rows": 7})
    with pytest.raises(RuntimeError, match="dropped rows stated"):
        mismatched.check()
    short = fixed.model_copy(update={"kept_protein_keys": 5})
    with pytest.raises(RuntimeError, match="dropped rows !="):
        short.check()


# --------------------------------------------------------------------------- #
# Real data
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    from dotenv import load_dotenv

    load_dotenv()
    return os.environ["DATA_ROOT"]


def _real_paths() -> dict[str, str]:
    mirror = mo.raw_mirror_dir(_data_root())
    paths = {name: str(mirror / relpath) for name, relpath in mo.MIRROR_RELPATH.items()}
    missing = [path for path in paths.values() if not osp.exists(path)]
    if missing:
        pytest.skip(f"raw mirror not deposited: {missing[0]}")
    return paths


@pytest.mark.data
def test_raw_mirror_matches_every_pin() -> None:
    paths = _real_paths()
    manifest = mo.load_manifest(_data_root())
    for name, path in paths.items():
        assert file_sha256(path) == mo.DATA_SHA256[name]
        assert (
            mo.manifest_sha256(manifest, mo.MIRROR_RELPATH[name])
            == mo.DATA_SHA256[name]
        )


@pytest.mark.data
def test_every_quote_is_verbatim_in_its_pinned_mirror() -> None:
    paths = _real_paths()
    paper = mo.library_dir(_data_root()) / mo.PAPER_MD
    if not paper.exists():
        pytest.skip(f"paper mirror not present: {paper}")
    assert file_sha256(paper) == mo.PAPER_MD_SHA256
    text = paper.read_text(encoding="utf-8")
    assert [name for name, q in mo.PAPER_QUOTES.items() if q not in text] == []
    assert mo.check_quotes(paths)["n_quotes_checked"] == len(mo.APPENDIX_QUOTES) + sum(
        len(q) for q in mo.WORKBOOK_QUOTES.values()
    )


@pytest.mark.data
def test_real_release_shape_and_the_row_rules() -> None:
    paths = _real_paths()
    mo.check_sample_columns(paths[mo.EV8], "EV8")
    mo.check_sample_columns(paths[mo.EV9], "EV9")
    # Datasets EV2 and EV3 file 89 metadata rows between them: the 66 with a released
    # mass-fraction column, plus the Lib-00, Lib-31, Lib-32 and OffGel rows, which are
    # the DDA library-generation runs and carry no mass fraction anywhere.
    assert mo.check_sample_metadata(paths[mo.EV2], paths[mo.EV3]) == {
        "n_samples": 89,
        "n_loaded": 7,
    }
    for table, name in (("EV8", mo.EV8), ("EV9", mo.EV9)):
        summary = mo.check_normalization(paths[name], table)
        assert summary["worst_abs_deviation"] < 3e-8
    assert mo.check_identity_fault(paths[mo.EV8], paths[mo.EV9]) == {
        "n_blank_in_ev9": 30,
        "n_supplied_by_ev8": 12,
        "n_unidentified": 18,
    }
    columns = [spec.column for spec in mo.LOADED]
    rows = mo.read_mass_fractions(paths[mo.EV9], columns)
    assert len(rows) == 4342
    quantified = [row for row in rows if any(row.quantified(c) for c in columns)]
    assert len(quantified) == 2077
    assert [row.gene_name for row in quantified if row.gene_locus is None] == [
        "ygaU",
        "yghZ",
        "yifE",
        "ymfP",
    ]
    per_column = {c: sum(1 for row in rows if row.quantified(c)) for c in columns}
    assert per_column == {
        "A1-1": 1901,
        "A1-2": 1883,
        "A1-3": 1915,
        "C1": 1934,
        "F1-1": 1900,
        "F1-2": 1935,
        "F1-3": 1918,
    }


@pytest.mark.data
def test_real_b_number_route_resolves_every_released_locus() -> None:
    from torchcell.datasets.bacteria_common import bacterial_genome

    paths = _real_paths()
    genome = bacterial_genome("ecoli", mo.MORI_REFERENCE_STRAIN, _data_root())
    rows = mo.read_mass_fractions(paths[mo.EV9], [spec.column for spec in mo.LOADED])
    selection = mo.select_proteins(rows, genome, label="test")
    assert len(selection.kept) == 2073
    report = selection.reconciliation
    assert report.unique_names == 2073
    assert report.resolved == 2073
    assert report.resolved_fraction == 1.0
    assert report.layer_histogram["locus tag"] == 2070
    assert report.layer_histogram["gene synonym"] == 3
    assert report.outside_namespace == ()
    assert report.ambiguous_kept == {}


@pytest.mark.data
def test_built_lmdb_numbers() -> None:
    root = osp.join(_data_root(), "data/torchcell/proteome_mori2021")
    if not osp.exists(osp.join(root, "processed", "lmdb")):
        pytest.skip(f"dev store not built: {root}")
    from torchcell.verification.runners import load_records

    records = load_records(root)
    assert len(records) == mo.EXPECTED_RECORDS
    keys = sorted(
        len(r["experiment"]["phenotype"]["protein_abundance"]) for r in records
    )
    assert keys == [1879, 1896, 1897, 1911, 1914, 1930, 1931]
    ledger = json.loads(Path(root, "preprocess", "dropped_records.json").read_text())
    assert ledger["source_samples"] == 66
    assert ledger["kept_records"] == 7
    assert ledger["dropped_records"] == 59
    assert ledger["source_protein_rows"] == 4342
    assert ledger["kept_protein_keys"] == 2073
    assert ledger["dropped_protein_rows"] == 2269
    assert ledger["reconciliation"]["unique_names"] == 2073
    by_rule = {rule["rule"]: rule["n_items"] for rule in ledger["rules"]}
    assert by_rule["medium_has_no_media_library_entry"] == 52
    assert by_rule["medium_differs_from_the_library_entry_it_names"] == 2
    assert by_rule["medium_formulation_not_stated_by_the_source"] == 5
    assert by_rule["not_quantified_in_any_loaded_sample"] == 2265
    assert by_rule["release_files_no_gene_locus_for_this_row"] == 4
    assert by_rule["b_number_resolves_to_no_locus_of_the_pinned_assembly"] == 0


@pytest.mark.data
def test_mori_does_not_republish_the_schmidt_2016_proteome() -> None:
    """The nearest conditions correlate but no pair agrees, so neither subsumes the other.

    Mori's loaded records are MG1655 (EQ353) in Neidhardt MOPS glucose and the landed
    Schmidt records are BW25113 in M9, so the two loaded sets share no (protein, strain,
    condition) triple at all. This pins the comparison anyway, on UniProt accession and
    with Schmidt's copies/cell converted to a mass fraction through its own released
    molecular weights.
    """
    import torchcell.datasets.ecoli.schmidt2016 as sm

    paths = _real_paths()
    schmidt = sm.raw_mirror_dir(_data_root()) / sm.SI2_MIRROR_RELPATH
    if not schmidt.exists():
        pytest.skip(f"Schmidt raw mirror not deposited: {schmidt}")

    rows = mo.read_mass_fractions(paths[mo.EV9], ["A1-1"])
    mori = {
        row.protein_id: row.mass_fraction["A1-1"]
        for row in rows
        if row.protein_id is not None and row.quantified("A1-1")
    }
    assert len(mori) == 1901

    schmidt_rows, _ = sm.read_table_s6(str(schmidt))
    copies = {
        row.uniprot: (row.copies["Glucose"], row.molecular_weight_da)
        for row in schmidt_rows
        if row.copies["Glucose"] is not None and row.copies["Glucose"] > 0
    }
    total = sum(value * weight for value, weight in copies.values())
    schmidt_fraction = {
        key: value * weight / total for key, (value, weight) in copies.items()
    }

    shared = sorted(set(mori) & set(schmidt_fraction))
    assert len(shared) == 1812
    agreeing = [
        key
        for key in shared
        if abs(mori[key] - schmidt_fraction[key]) / schmidt_fraction[key] < 1e-6
    ]
    assert agreeing == []
    left = [math.log10(mori[key]) for key in shared]
    right = [math.log10(schmidt_fraction[key]) for key in shared]
    mean_left, mean_right = statistics.fmean(left), statistics.fmean(right)
    covariance = sum(
        (a - mean_left) * (b - mean_right) for a, b in zip(left, right, strict=True)
    )
    spread = math.sqrt(
        sum((a - mean_left) ** 2 for a in left)
        * sum((b - mean_right) ** 2 for b in right)
    )
    assert covariance / spread == pytest.approx(0.796, abs=0.005)
