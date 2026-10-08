# tests/torchcell/datasets/ecoli/test_schmidt2016_srm.py
# [[tests.torchcell.datasets.ecoli.test_schmidt2016_srm]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_schmidt2016_srm.py
"""Schmidt 2016 SRM loaders (``torchcell.datasets.ecoli.schmidt2016_srm``).

The synthetic tests extend the Table S6 / S8 / S25 workbook
``test_schmidt2016.write_workbook`` already builds with the three sheets these loaders
read -- Table S1's panel and Tables S2 and S3's long-form measurements -- so the
separability check runs against a real Table S6 block rather than a stub. They drive the
readers, all four header-versus-values checks, the phenotype builder and a full
``process()`` for BOTH arms into ``tmp_path``; the genome and the locus-tag
reconciliation are in-test objects, so nothing reads ``$DATA_ROOT``.

The ``@pytest.mark.data`` tests read the real mirror and pin the numbers the dendron note
states: 779 and 682 released rows over 41 and 31 panel proteins, 19 and 22 conditions,
the undescribed ``anaerobic`` condition, the 98 rows whose released standard deviation is
exactly 0, and that zero of the 738 and 682 cells shared with the stored Table S6 block
agree with it.
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
import statistics
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

import torchcell.datasets.ecoli.schmidt2016 as sm
import torchcell.datasets.ecoli.schmidt2016_srm as srm
from tests.torchcell.datasets.ecoli.test_schmidt2016 import DATASET1_MISSING
from tests.torchcell.datasets.ecoli.test_schmidt2016 import (
    write_workbook as write_stored_workbook,
)
from torchcell.data import file_sha256
from torchcell.datamodels.schema import (
    AssemblyReferenceGenome,
    BacterialAssemblySet,
    BacterialGeneNamespace,
)
from torchcell.datasets.bacteria_common import LocusTagReconciliation
from torchcell.sequence.genome.base import GeneNameStatus

ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_BW25113_ASM75055v1"
NAMESPACE: BacterialGeneNamespace = "ecoli_k12_bw25113_locus_tag"

#: ``{accession: (Table S1 gene, SRM gene, peptide, spike)}``. ``P00002`` is named
#: differently in Table S1 than in the measurement tables, which is the one
#: identifier inconsistency the release itself carries (``P0ACP1``: ``cra`` / ``fruR``).
PANEL: dict[str, tuple[str, str, str, float]] = {
    "P00001": ("aceA", "aceA", "WEGITRPYSAEDVVK", 200.0),
    "P00002": ("aceB", "aceBalias", "DAVNGTISYTNEAGK", 20.0),
    "P00003": ("dupA", "dupA", "LVPIIADEAR", 200.0),
}
LOCUS_TAGS = {
    "aceA": "BW25113_4015",
    "aceBalias": "BW25113_4014",
    "dupA": "BW25113_4013",
}
DISAGREEMENT = {"P00002": ("aceB", "aceBalias")}
#: One SRM label per Table S6 column, in the spelling each real sheet uses.
SET1_LABELS: dict[str, str] = {
    "Glucose": "glucose",
    "LB": "LB",
    "Acetate": "acetate",
    "Fumarate": "fumarate",
    "Glucosamine": "glucosamine",
    "Glycerol": "glycerol",
    "Pyruvate": "pyruvate",
    "Chemostat µ=0.5": "chemostat µ=0.5",
    "Chemostat µ=0.35": "chemostat µ=0.35",
    "Chemostat µ=0.20": "chemostat µ=0.20",
    "Chemostat µ=0.12": "chemostat µ=0.12",
    "Stationary phase 1 day": "stationary 1 day",
    "Stationary phase 3 days": "stationary 3 days",
    "Osmotic-stress glucose": "50 mM NaCl",
    "42°C glucose": "42°C",
    "pH6 glucose": "pH 6",
    "Galactose": "galactose",
    "Succinate": "succinate",
}
SET2_EXTRA: dict[str, str] = {
    "Glycerol + AA": "glycerol + AA",
    "Xylose": "xylose",
    "Mannose": "mannose",
    "Fructose": "fructose",
}
#: Table S3's own spellings for the two labels the two sheets write differently.
SET2_RESPELLED = {
    "Chemostat µ=0.20": "chemostat µ=0.2",
    "Stationary phase 3 days": "stationary 3 day",
}
#: Set 2 measures a SUBSET of set 1's panel, as the real 31-of-41 nesting does, and
#: both include the accession whose gene name disagrees between Table S1 and the
#: measurement tables.
SET1_ACCESSIONS = ("P00001", "P00002", "P00003")
SET2_ACCESSIONS = ("P00001", "P00002")
#: Relative dispersions: the technical arm must stay tighter than the biological one.
SET1_CV = 0.01
SET2_CV = 0.08


def _set2_labels() -> dict[str, str]:
    """``{Table S6 column: Table S3 label}`` over all 22 released conditions."""
    labels = {**SET1_LABELS, **SET2_EXTRA}
    return {
        column: SET2_RESPELLED.get(column, label) for column, label in labels.items()
    }


def _srm_value(accession: str, column: str) -> float:
    """An SRM copies/cell that is never the Table S6 value for the same cell."""
    index = sorted(PANEL).index(accession)
    offset = [c.s6_column for c in sm.CONDITIONS].index(column)
    return 137.0 * (index + 1) + 1.37 * offset + 0.5


def write_workbook(path: Path) -> Path:
    """The stored block's workbook plus Table S1, Table S2 and Table S3."""
    write_stored_workbook(path)
    book = openpyxl.load_workbook(path)
    _write_s1(book.create_sheet(srm.SHEET_S1))
    _write_srm(
        book.create_sheet(srm.SHEET_S2),
        srm._Q_TABLE_S2,
        srm.SRM_SD_TECHNICAL,
        SET1_ACCESSIONS,
        {**SET1_LABELS, "_anaerobic": srm.UNDESCRIBED_CONDITION},
        SET1_CV,
    )
    _write_srm(
        book.create_sheet(srm.SHEET_S3),
        srm._Q_TABLE_S3,
        srm.SRM_SD_BIOLOGICAL,
        SET2_ACCESSIONS,
        _set2_labels(),
        SET2_CV,
    )
    book.save(path)
    return path


def _write_s1(sheet: Any) -> None:
    sheet.append([srm._Q_TABLE_S1])
    sheet.append([None] * 6)
    sheet.append(
        [
            srm.S1_ACCESSION,
            srm.S1_GENE,
            srm.S1_DESCRIPTION,
            srm.S1_PEPTIDE,
            "Pathway/Function According to www.uniprot.org",
            srm.S1_SPIKE,
        ]
    )
    for accession, (gene, _srm_gene, peptide, spike) in PANEL.items():
        sheet.append(
            [
                accession,
                gene,
                f"Test protein {gene} OS=Escherichia coli (strain K12) GN={gene}",
                peptide,
                "glycolysis",
                spike,
            ]
        )
    sheet.append([None] * 6)
    sheet.append(["* Peptides were selected according to the following requirements"])


def _write_srm(
    sheet: Any,
    title: str,
    sd_header: str,
    accessions: tuple[str, ...],
    labels: dict[str, str],
    relative_sd: float,
) -> None:
    sheet.append([title])
    sheet.append([None] * 6)
    sheet.append(
        [
            srm.SRM_ACCESSION,
            srm.SRM_GENE,
            srm.SRM_PEPTIDE,
            srm.SRM_CONDITION,
            srm.SRM_ABUNDANCE,
            sd_header,
        ]
    )
    for accession in accessions:
        _gene, srm_gene, peptide, _spike = PANEL[accession]
        for column, label in labels.items():
            value = (
                _srm_value(accession, column)
                if not column.startswith("_")
                else 11.0 * (sorted(PANEL).index(accession) + 1)
            )
            sheet.append(
                [accession, srm_gene, peptide, label, value, relative_sd * value]
            )


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    """The synthetic workbook, written once per test."""
    return write_workbook(tmp_path / sm.SI2)


@pytest.fixture
def declared_disagreement(monkeypatch: pytest.MonkeyPatch) -> None:
    """The synthetic panel's own gene-name inconsistency stands in for ``P0ACP1``.

    Requested only by the tests that reach ``check_identifier_columns`` on the synthetic
    workbook; the data-gated tests read the real ``P0ACP1`` disagreement instead, so
    this must not be autouse.
    """
    monkeypatch.setattr(srm, "GENE_NAME_DISAGREEMENT", DISAGREEMENT)


# --------------------------------------------------------------------------- #
# The declared module constants
# --------------------------------------------------------------------------- #
def test_the_two_arms_declare_distinct_measurement_types() -> None:
    types = [arm.MEASUREMENT_TYPE for arm in srm.ARMS]
    assert len(set(types)) == 2
    assert sm.MEASUREMENT_TYPE not in types
    assert all("srm_sid" in value for value in types)
    assert srm.MEASUREMENT_TYPES == {
        arm.__name__: arm.MEASUREMENT_TYPE for arm in srm.ARMS
    }
    assert [arm.N_REPLICATES for arm in srm.ARMS] == [2, 3]
    assert [arm.SHEET for arm in srm.ARMS] == [srm.SHEET_S2, srm.SHEET_S3]
    assert [arm.IS_SET1 for arm in srm.ARMS] == [True, False]


def test_every_declared_condition_label_names_a_released_column() -> None:
    columns = {c.s6_column for c in sm.CONDITIONS}
    assert set(srm.CONDITION_LABELS.values()) == columns
    assert srm.UNDESCRIBED_CONDITION not in srm.CONDITION_LABELS
    assert (
        srm.DROP_CONDITION_UNDESCRIBED.rule == "condition_not_described_by_the_source"
    )
    assert srm.DROP_CONDITION_UNDESCRIBED.needed_addition is None


def test_sourced_values_quote_the_pinned_artifacts() -> None:
    assert srm.SOURCED_VALUES["n_replicates_set1"].value == 2
    assert srm.SOURCED_VALUES["n_replicates_set2"].value == 3
    assert srm.SOURCED_VALUES["dataset1_sample_count"].value == 19
    assert {v.provenance.sha256 for v in srm.SOURCED_VALUES.values()} == {
        sm.PAPER_MD_SHA256,
        sm.SI2_SHA256,
    }
    assert srm.PAPER_QUOTES["srm_duplicate"] == "Each sample was analyzed in duplicate."


# --------------------------------------------------------------------------- #
# Reading the workbook
# --------------------------------------------------------------------------- #
def test_read_table_s1_stops_before_the_footnote_row(workbook: Path) -> None:
    panel = srm.read_table_s1(str(workbook))
    assert set(panel) == set(PANEL)
    assert panel["P00001"].gene == "aceA"
    assert panel["P00001"].peptide == PANEL["P00001"][2]
    assert panel["P00002"].spike_fmol_per_ug == 20.0


def test_read_table_s1_refuses_a_twice_filed_accession(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    book[srm.SHEET_S1].cell(row=5, column=1, value="P00001")
    book.save(path)
    with pytest.raises(RuntimeError, match="filed twice"):
        srm.read_table_s1(str(path))


def test_read_srm_table_parses_the_long_form_and_refuses_a_renamed_column(
    workbook: Path, tmp_path: Path
) -> None:
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    assert len(cells) == len(SET1_ACCESSIONS) * (len(SET1_LABELS) + 1)
    assert {cell.uniprot for cell in cells} == set(SET1_ACCESSIONS)
    assert srm.UNDESCRIBED_CONDITION in {cell.condition for cell in cells}
    glucose = next(
        cell
        for cell in cells
        if cell.uniprot == "P00001" and cell.condition == "glucose"
    )
    assert glucose.copies_per_cell == pytest.approx(_srm_value("P00001", "Glucose"))
    assert glucose.standard_deviation == pytest.approx(
        SET1_CV * _srm_value("P00001", "Glucose")
    )
    assert glucose.row_number == 4

    # Table S3's dispersion header is NOT Table S2's, which is the swap this refuses.
    with pytest.raises(RuntimeError, match="carries no column"):
        srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_BIOLOGICAL)


def test_read_srm_table_refuses_a_repeated_key(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    sheet = book[srm.SHEET_S2]
    sheet.cell(row=5, column=4, value=sheet.cell(row=4, column=4).value)
    book.save(path)
    with pytest.raises(RuntimeError, match="distinct"):
        srm.read_srm_table(str(path), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)


def test_read_srm_table_refuses_a_non_numeric_abundance(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    book[srm.SHEET_S2].cell(row=4, column=5, value="below LOQ")
    book.save(path)
    with pytest.raises(RuntimeError, match="is not a number"):
        srm.read_srm_table(str(path), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)


def test_read_srm_table_refuses_an_empty_identifier(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    book[srm.SHEET_S2].cell(row=4, column=2).value = None
    book.save(path)
    with pytest.raises(RuntimeError, match="the released cell is empty"):
        srm.read_srm_table(str(path), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)


def test_header_index_refuses_a_repeated_name() -> None:
    assert srm._header_index(["a", None, "b"]) == {"a": 0, "b": 2}
    with pytest.raises(RuntimeError, match="appears twice"):
        srm._header_index(["a", "a"])


def test_dataset1_uncovered_conditions_reads_table_s6s_own_na_pattern(
    workbook: Path,
) -> None:
    assert srm.dataset1_uncovered_conditions(str(workbook)) == DATASET1_MISSING


def test_read_table_s6_copies_keys_on_accession_and_condition(workbook: Path) -> None:
    stored = srm.read_table_s6_copies(str(workbook))
    assert ("P00001", "Glucose") in stored
    # the dataset-1 row has no value in the four conditions it did not cover
    assert ("P00002", "Xylose") not in stored


# --------------------------------------------------------------------------- #
# Header-versus-values checks
# --------------------------------------------------------------------------- #
def test_identifier_columns_agree_with_table_s1_and_pin_the_one_disagreement(
    workbook: Path, declared_disagreement: None
) -> None:
    panel = srm.read_table_s1(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    summary = srm.check_identifier_columns(cells, panel, sheet=srm.SHEET_S2)
    assert summary["peptides_match_table_s1"] is True
    assert summary["gene_name_disagreements"] == {"P00002": ["aceB", "aceBalias"]}
    assert summary["n_panel_proteins"] == len(SET1_ACCESSIONS)


def test_identifier_columns_refuse_a_peptide_that_is_not_the_spiked_one(
    workbook: Path, declared_disagreement: None
) -> None:
    panel = srm.read_table_s1(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    broken = [cells[0].model_copy(update={"peptide": "OTHERPEPTIDE"}), *cells[1:]]
    with pytest.raises(RuntimeError, match="selected peptide"):
        srm.check_identifier_columns(broken, panel, sheet=srm.SHEET_S2)


def test_identifier_columns_refuse_an_undeclared_gene_name_disagreement(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    panel = srm.read_table_s1(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    monkeypatch.setattr(srm, "GENE_NAME_DISAGREEMENT", {})
    with pytest.raises(RuntimeError, match="disagreeing with Table S1"):
        srm.check_identifier_columns(cells, panel, sheet=srm.SHEET_S2)


def test_identifier_columns_refuse_an_accession_outside_the_panel(
    workbook: Path, declared_disagreement: None
) -> None:
    panel = srm.read_table_s1(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    stranger = [cells[0].model_copy(update={"uniprot": "P99999"}), *cells[1:]]
    with pytest.raises(RuntimeError, match="not Table S1 panel proteins"):
        srm.check_identifier_columns(stranger, panel, sheet=srm.SHEET_S2)


def test_dataset_arm_is_checked_against_table_s6s_coverage_pattern(
    workbook: Path,
) -> None:
    uncovered = srm.dataset1_uncovered_conditions(str(workbook))
    set1 = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    set2 = srm.read_srm_table(str(workbook), srm.SHEET_S3, srm.SRM_SD_BIOLOGICAL)
    first = srm.check_dataset_arm(set1, uncovered, sheet=srm.SHEET_S2, is_set1=True)
    assert first["n_conditions"] == len(SET1_LABELS) + 1
    assert first["undescribed_condition_present"] is True
    assert first["rows_per_condition"] == len(SET1_ACCESSIONS)
    assert set(first["table_s6_columns"]) == set(SET1_LABELS)

    second = srm.check_dataset_arm(set2, uncovered, sheet=srm.SHEET_S3, is_set1=False)
    assert second["n_conditions"] == len(sm.CONDITIONS)
    assert second["undescribed_condition_present"] is False

    # Swapping the two arms' declared identity is what this refuses: each arm's
    # condition set is the other's with four columns added or removed.
    with pytest.raises(RuntimeError, match="coverage pattern"):
        srm.check_dataset_arm(set1, uncovered, sheet=srm.SHEET_S2, is_set1=False)
    with pytest.raises(RuntimeError, match="coverage pattern"):
        srm.check_dataset_arm(set2, uncovered, sheet=srm.SHEET_S3, is_set1=True)
    # An arm whose conditions DO match but whose undescribed-condition flag does not.
    without = [cell for cell in set1 if cell.condition != srm.UNDESCRIBED_CONDITION]
    with pytest.raises(RuntimeError, match="contradicts the declared data-set arm"):
        srm.check_dataset_arm(without, uncovered, sheet=srm.SHEET_S2, is_set1=True)


def test_dataset_arm_refuses_an_undeclared_condition_label(workbook: Path) -> None:
    uncovered = srm.dataset1_uncovered_conditions(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    renamed = [cells[0].model_copy(update={"condition": "microaerobic"}), *cells[1:]]
    with pytest.raises(RuntimeError, match="no Table S6 column is declared"):
        srm.check_dataset_arm(renamed, uncovered, sheet=srm.SHEET_S2, is_set1=True)


def test_dataset_arm_refuses_a_ragged_condition_block(workbook: Path) -> None:
    uncovered = srm.dataset1_uncovered_conditions(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    thinned = [
        cell
        for cell in cells
        if not (cell.condition == "glucose" and cell.uniprot == "P00002")
    ]
    with pytest.raises(RuntimeError, match="not uniform"):
        srm.check_dataset_arm(thinned, uncovered, sheet=srm.SHEET_S2, is_set1=True)


def test_dispersion_labels_order_the_two_sheets_and_a_swap_is_refused(
    workbook: Path,
) -> None:
    set1 = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    set2 = srm.read_srm_table(str(workbook), srm.SHEET_S3, srm.SRM_SD_BIOLOGICAL)
    summary = srm.check_dispersion_labels(set1, set2)
    assert summary["set1_technical_median_cv_percent"] == pytest.approx(100 * SET1_CV)
    assert summary["set2_biological_median_cv_percent"] == pytest.approx(100 * SET2_CV)
    assert summary["set1_rows_with_sd_at_or_above_abundance"] == 0
    assert summary["set1_rows_with_zero_sd"] == 0
    with pytest.raises(RuntimeError, match="the way their headers say"):
        srm.check_dispersion_labels(set2, set1)


def test_separability_from_the_stored_block_is_asserted(workbook: Path) -> None:
    stored = srm.read_table_s6_copies(str(workbook))
    for arm, sheet, header in (
        (srm.ProteomeSrmSet1Schmidt2016Dataset, srm.SHEET_S2, srm.SRM_SD_TECHNICAL),
        (srm.ProteomeSrmSet2Schmidt2016Dataset, srm.SHEET_S3, srm.SRM_SD_BIOLOGICAL),
    ):
        cells = srm.read_srm_table(str(workbook), sheet, header)
        summary = srm.check_not_the_stored_block(
            cells, stored, measurement_type=arm.MEASUREMENT_TYPE, sheet=sheet
        )
        assert summary["n_agreeing_to_rtol"] == 0
        assert summary["n_shared_cells"] > 0
        assert summary["stored_measurement_type"] == sm.MEASUREMENT_TYPE
        assert summary["measurement_type"] != sm.MEASUREMENT_TYPE


def test_separability_refuses_a_cell_that_equals_the_stored_block(
    workbook: Path,
) -> None:
    stored = srm.read_table_s6_copies(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    cloned = [
        cells[0].model_copy(
            update={"copies_per_cell": stored[(cells[0].uniprot, "Glucose")]}
        ),
        *cells[1:],
    ]
    first = [cell for cell in cloned if cell.condition == "glucose"]
    assert first
    with pytest.raises(RuntimeError, match="no longer separable"):
        srm.check_not_the_stored_block(
            cloned,
            stored,
            measurement_type=srm.ProteomeSrmSet1Schmidt2016Dataset.MEASUREMENT_TYPE,
            sheet=srm.SHEET_S2,
        )


def test_separability_refuses_the_stored_blocks_own_measurement_type(
    workbook: Path,
) -> None:
    stored = srm.read_table_s6_copies(str(workbook))
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    with pytest.raises(RuntimeError, match="not one of the declared SRM"):
        srm.check_not_the_stored_block(
            cells, stored, measurement_type=sm.MEASUREMENT_TYPE, sheet=srm.SHEET_S2
        )


# --------------------------------------------------------------------------- #
# Identifier resolution and the condition partition
# --------------------------------------------------------------------------- #
class _Locus:
    def __init__(self, symbol: str | None) -> None:
        self.symbol = symbol


class _Annotation:
    def __init__(self, loci: dict[str, _Locus]) -> None:
        self.loci = loci


class FakeGenome:
    """Carries the panel's locus tags and nothing else."""

    ASSEMBLY_SET = ASSEMBLY_SET

    def __init__(self) -> None:
        """One locus per panel protein."""
        self.genbank = _Annotation(
            {tag: _Locus(symbol) for symbol, tag in LOCUS_TAGS.items()}
        )


def _fake_reconcile(
    genome: Any, names: pd.Series, *, label: str
) -> tuple[pd.Series, LocusTagReconciliation]:
    """Map each panel symbol to its locus tag; anything else is outside the namespace."""
    stored = pd.Series([LOCUS_TAGS.get(name, name) for name in names])
    resolved = sum(1 for name in set(names) if name in LOCUS_TAGS)
    report = LocusTagReconciliation(
        label=label,
        assembly_set=ASSEMBLY_SET,
        gene_namespace=NAMESPACE,
        unique_names=len(set(names)),
        status_histogram={
            GeneNameStatus.RENAMED: resolved,
            GeneNameStatus.CURRENT: 0,
            GeneNameStatus.NON_GENE_FEATURE: 0,
            GeneNameStatus.RETIRED: len(set(names)) - resolved,
            GeneNameStatus.AMBIGUOUS: 0,
        },
        layer_histogram={"gene symbol": resolved},
        remapped=resolved,
        kept_on_collision=(),
        retired_kept=(),
        ambiguous_kept={},
        case_insensitive=(),
        outside_namespace=tuple(n for n in names if n not in LOCUS_TAGS),
    )
    return stored, report


def _reference_genome() -> AssemblyReferenceGenome:
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain=sm.SCHMIDT_REFERENCE_STRAIN,
        ploidy="haploid",
        assembly_set=ASSEMBLY_SET,
        assembly_accession="GCA_000750555.1",
    )


def test_select_arm_partitions_the_conditions_by_the_shared_drop_rules(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(srm, "reconcile_locus_tags", _fake_reconcile)
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    selection = srm.select_arm(cells, FakeGenome(), label="test")  # type: ignore[arg-type]
    assert selection.locus_tag == {
        "P00001": "BW25113_4015",
        "P00002": "BW25113_4014",
        "P00003": "BW25113_4013",
    }
    assert selection.dropped == {
        "condition_not_described_by_the_source": [srm.UNDESCRIBED_CONDITION],
        "culture_not_batch": [
            "chemostat µ=0.12",
            "chemostat µ=0.20",
            "chemostat µ=0.35",
            "chemostat µ=0.5",
        ],
        "growth_phase_not_representable": ["stationary 1 day", "stationary 3 days"],
    }
    # the loaded columns keep the release's own column order, reference included
    assert selection.loaded_columns[0] == sm.REFERENCE_CONDITION
    assert (
        len(selection.loaded_columns)
        == srm.ProteomeSrmSet1Schmidt2016Dataset.EXPECTED_RECORDS + 1
    )


def test_select_arm_refuses_a_symbol_outside_the_namespace(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(srm, "reconcile_locus_tags", _fake_reconcile)
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    renamed = [cell.model_copy(update={"gene": "ghost"}) for cell in cells]
    from torchcell.datasets.bacteria_common import LocusTagResolutionError

    with pytest.raises(LocusTagResolutionError):
        srm.select_arm(renamed, FakeGenome(), label="test")  # type: ignore[arg-type]


def test_phenotype_divides_the_released_sd_by_the_root_of_the_replicate_count(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(srm, "reconcile_locus_tags", _fake_reconcile)
    cells = srm.read_srm_table(str(workbook), srm.SHEET_S2, srm.SRM_SD_TECHNICAL)
    selection = srm.select_arm(cells, FakeGenome(), label="test")  # type: ignore[arg-type]
    phenotype = srm.build_phenotype(
        cells,
        selection.locus_tag,
        "Glucose",
        measurement_type="test_type",
        n_replicates=2,
    )
    assert set(phenotype.protein_abundance) == set(LOCUS_TAGS.values())
    assert phenotype.measurement_type == "test_type"
    assert set(phenotype.n_replicates.values()) == {2}
    assert phenotype.protein_abundance_se is not None
    for accession, tag in selection.locus_tag.items():
        value = _srm_value(accession, "Glucose")
        assert phenotype.protein_abundance[tag] == pytest.approx(value)
        assert phenotype.protein_abundance_se[tag] == pytest.approx(
            SET1_CV * value / math.sqrt(2)
        )


# --------------------------------------------------------------------------- #
# End-to-end builds on the synthetic workbook
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("arm", "accessions", "records"),
    [
        (srm.ProteomeSrmSet1Schmidt2016Dataset, SET1_ACCESSIONS, 11),
        (srm.ProteomeSrmSet2Schmidt2016Dataset, SET2_ACCESSIONS, 14),
    ],
    ids=["set1", "set2"],
)
def test_process_builds_one_record_per_loaded_condition(
    arm: type[Any],
    accessions: tuple[str, ...],
    records: int,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    declared_disagreement: None,
) -> None:
    root = tmp_path / srm.arm_root(arm).rsplit("/", 1)[-1]
    raw = root / "raw"
    raw.mkdir(parents=True)
    path = write_workbook(raw / sm.SI2)
    conditions = len(SET1_LABELS) + 1 if arm.IS_SET1 else len(sm.CONDITIONS)
    monkeypatch.setattr(sm, "DATA_SHA256", {sm.SI2: file_sha256(path)})
    monkeypatch.setattr(srm, "reconcile_locus_tags", _fake_reconcile)
    monkeypatch.setattr(srm, "assembly_reference", lambda strain: _reference_genome())
    monkeypatch.setattr(arm, "EXPECTED_ROWS", len(accessions) * conditions)
    monkeypatch.setattr(arm, "EXPECTED_PROTEIN_KEYS", len(accessions))

    dataset = arm(root=str(root), ecoli_genome=FakeGenome())
    assert len(dataset) == records == arm.EXPECTED_RECORDS
    assert dataset.gene_set == {LOCUS_TAGS[PANEL[a][1]] for a in accessions}

    items = [dataset.transform_item(dataset[i]) for i in range(len(dataset))]
    assert {len(i["experiment"].genotype.perturbations) for i in items} == {0}
    assert {i["experiment"].phenotype.measurement_type for i in items} == {
        arm.MEASUREMENT_TYPE
    }
    assert {i["publication"].doi for i in items} == {sm.PAPER_DOI}
    assert (
        len({i["experiment"].environment.model_dump_json() for i in items}) == records
    )
    for item in items:
        reference = item["reference"].phenotype_reference
        assert set(reference.protein_abundance) == set(
            item["experiment"].phenotype.protein_abundance
        )
        assert reference.measurement_type == arm.MEASUREMENT_TYPE
    dataset.close_lmdb()

    preprocess = Path(dataset.preprocess_dir)
    ledger = json.loads((preprocess / "dropped_records.json").read_text())
    assert ledger["sheet"] == arm.SHEET
    assert ledger["source_conditions"] == conditions
    assert ledger["kept_records"] == records
    assert ledger["kept_protein_keys"] == len(accessions)
    rules = {rule["rule"]: rule["n_items"] for rule in ledger["rules"]}
    assert rules["culture_not_batch"] == 4
    assert rules["growth_phase_not_representable"] == 2
    if arm.IS_SET1:
        assert rules["condition_not_described_by_the_source"] == 1
    else:
        assert rules["medium_has_no_media_library_entry"] == 1
    checks = json.loads((preprocess / "released_statistics_check.json").read_text())
    assert checks["not_the_stored_block"]["n_agreeing_to_rtol"] == 0
    assert checks["dataset_arm"]["undescribed_condition_present"] is arm.IS_SET1
    assert checks["identifier_columns"]["peptides_match_table_s1"] is True
    sourced = json.loads((preprocess / "sourced_values.json").read_text())
    assert sourced["reference_strain"]["value"] == "BW25113"
    identifiers = pd.read_csv(preprocess / "protein_identifiers.csv")
    assert len(identifiers) == len(accessions)
    assert set(identifiers["n_replicates"]) == {arm.N_REPLICATES}
    built = pd.read_csv(preprocess / "conditions.csv")
    assert len(built) == records
    assert sm.REFERENCE_CONDITION not in set(built["s6_column"])
    assert (preprocess / "build_manifest.json").exists()

    report = srm.verify_build(str(root), arm=arm, genome=FakeGenome())  # type: ignore[arg-type]
    assert [(r.level, r.name) for r in report.results if not r.passed] == []
    assert (preprocess / "verification_report.json").exists()


def test_process_refuses_a_row_count_that_is_not_the_declared_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proteome_srm_set1_schmidt2016"
    raw = root / "raw"
    raw.mkdir(parents=True)
    path = write_workbook(raw / sm.SI2)
    monkeypatch.setattr(sm, "DATA_SHA256", {sm.SI2: file_sha256(path)})
    monkeypatch.setattr(srm, "reconcile_locus_tags", _fake_reconcile)
    with pytest.raises(RuntimeError, match="data rows, the module states"):
        srm.ProteomeSrmSet1Schmidt2016Dataset(
            root=str(root),
            ecoli_genome=FakeGenome(),  # type: ignore[arg-type]
        )


def test_download_links_the_mirror(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = write_workbook(tmp_path / "source.xlsx")
    sha = file_sha256(source)
    monkeypatch.setattr(sm, "SI2_SHA256", sha)
    data_root = tmp_path / "root"
    sm.deposit_raw_mirror(source=source, data_root=str(data_root))
    monkeypatch.setenv("DATA_ROOT", str(data_root))

    root = tmp_path / "build"
    (root / "raw").mkdir(parents=True)
    cls = srm.ProteomeSrmSet2Schmidt2016Dataset
    dataset = cls.__new__(cls)
    monkeypatch.setattr(
        type(dataset), "raw_dir", property(lambda self: str(root / "raw"))
    )
    dataset.download()
    assert file_sha256(root / "raw" / sm.SI2) == sha

    (data_root / sm.RAW_DIR_REL / sm.SI2_MIRROR_RELPATH).unlink()
    with pytest.raises(RuntimeError, match="missing from mirror"):
        dataset.download()


def test_arm_root_is_each_arms_own_slug() -> None:
    roots = [srm.arm_root(arm) for arm in srm.ARMS]
    assert roots == [
        "data/torchcell/proteome_srm_set1_schmidt2016",
        "data/torchcell/proteome_srm_set2_schmidt2016",
    ]
    assert len(set(roots)) == 2


# --------------------------------------------------------------------------- #
# Data-gated: the real mirror
# --------------------------------------------------------------------------- #
def _real_workbook() -> str:
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT is not set")
    path = osp.join(data_root, sm.RAW_DIR_REL, sm.SI2_MIRROR_RELPATH)
    if not osp.exists(path):
        pytest.skip(f"raw mirror not deposited: {path}")
    return path


@pytest.fixture
def real_cells() -> tuple[list[srm.SrmCell], list[srm.SrmCell]]:
    """Both real SRM sheets, read from the pinned mirror."""
    path = _real_workbook()
    return (
        srm.read_srm_table(path, srm.SHEET_S2, srm.SRM_SD_TECHNICAL),
        srm.read_srm_table(path, srm.SHEET_S3, srm.SRM_SD_BIOLOGICAL),
    )


@pytest.mark.data
def test_real_sheets_release_the_pinned_rows_and_panels(
    real_cells: tuple[list[srm.SrmCell], list[srm.SrmCell]],
) -> None:
    set1, set2 = real_cells
    assert len(set1) == 779
    assert len(set2) == 682
    assert len({cell.uniprot for cell in set1}) == 41
    assert len({cell.uniprot for cell in set2}) == 31
    assert {cell.uniprot for cell in set2} <= {cell.uniprot for cell in set1}
    assert len({cell.condition for cell in set1}) == 19
    assert len({cell.condition for cell in set2}) == 22
    # one proteotypic peptide per protein, so the peptide axis has length one
    assert len({(cell.uniprot, cell.peptide) for cell in set1}) == 41


@pytest.mark.data
def test_real_panel_is_the_forty_one_selected_proteins(
    real_cells: tuple[list[srm.SrmCell], list[srm.SrmCell]],
) -> None:
    panel = srm.read_table_s1(_real_workbook())
    assert len(panel) == 41
    assert {record.spike_fmol_per_ug for record in panel.values()} == {20.0, 200.0}
    set1, set2 = real_cells
    for cells, sheet in ((set1, srm.SHEET_S2), (set2, srm.SHEET_S3)):
        summary = srm.check_identifier_columns(cells, panel, sheet=sheet)
        assert summary["gene_name_disagreements"] == {"P0ACP1": ["cra", "fruR"]}
        assert summary["peptides_match_table_s1"] is True


@pytest.mark.data
def test_real_undescribed_anaerobic_condition_is_table_s2s_nineteenth(
    real_cells: tuple[list[srm.SrmCell], list[srm.SrmCell]],
) -> None:
    set1, set2 = real_cells
    labels1 = {cell.condition for cell in set1}
    labels2 = {cell.condition for cell in set2}
    assert srm.UNDESCRIBED_CONDITION in labels1
    assert srm.UNDESCRIBED_CONDITION not in labels2
    uncovered = srm.dataset1_uncovered_conditions(_real_workbook())
    assert uncovered == ("Glycerol + AA", "Xylose", "Mannose", "Fructose")
    mapped = {
        srm.CONDITION_LABELS[label] for label in labels1 - {srm.UNDESCRIBED_CONDITION}
    }
    assert len(mapped) == len(sm.CONDITIONS) - len(uncovered) == 18
    assert (
        srm.check_dataset_arm(set1, uncovered, sheet=srm.SHEET_S2, is_set1=True)[
            "rows_per_condition"
        ]
        == 41
    )


@pytest.mark.data
def test_real_dispersion_labels_order_as_their_headers_say(
    real_cells: tuple[list[srm.SrmCell], list[srm.SrmCell]],
) -> None:
    set1, set2 = real_cells
    summary = srm.check_dispersion_labels(set1, set2)
    assert summary["set1_technical_median_cv_percent"] == pytest.approx(1.771, abs=1e-3)
    assert summary["set2_biological_median_cv_percent"] == pytest.approx(
        6.489, abs=1e-3
    )
    # the released zero-dispersion rows, kept verbatim and pinned so a correction shows
    assert summary["set1_rows_with_zero_sd"] == 98
    assert summary["set2_rows_with_zero_sd"] == 0
    assert summary["set1_rows_with_sd_at_or_above_abundance"] == 0
    assert summary["set2_rows_with_sd_at_or_above_abundance"] == 2
    zero_proteins = {cell.gene for cell in set1 if cell.standard_deviation == 0.0}
    assert len(zero_proteins) == 13
    always = {
        gene
        for gene in zero_proteins
        if sum(
            1 for cell in set1 if cell.gene == gene and cell.standard_deviation == 0.0
        )
        == 19
    }
    assert always == {"ytjC", "acnA", "pgk", "aceA"}


@pytest.mark.data
def test_real_arms_are_separable_from_the_stored_table_s6_block(
    real_cells: tuple[list[srm.SrmCell], list[srm.SrmCell]],
) -> None:
    stored = srm.read_table_s6_copies(_real_workbook())
    set1, set2 = real_cells
    first = srm.check_not_the_stored_block(
        set1,
        stored,
        measurement_type=srm.ProteomeSrmSet1Schmidt2016Dataset.MEASUREMENT_TYPE,
        sheet=srm.SHEET_S2,
    )
    assert first["n_shared_cells"] == 738
    assert first["n_agreeing_to_rtol"] == 0
    assert first["pearson_r_log10"] == pytest.approx(0.664, abs=5e-4)
    assert first["median_abs_log2_ratio"] == pytest.approx(1.352, abs=5e-3)

    second = srm.check_not_the_stored_block(
        set2,
        stored,
        measurement_type=srm.ProteomeSrmSet2Schmidt2016Dataset.MEASUREMENT_TYPE,
        sheet=srm.SHEET_S3,
    )
    assert second["n_shared_cells"] == 682
    assert second["n_agreeing_to_rtol"] == 0
    assert second["pearson_r_log10"] == pytest.approx(0.928, abs=5e-4)
    assert second["median_abs_log2_ratio"] == pytest.approx(0.616, abs=5e-3)


@pytest.mark.data
def test_real_arms_disagree_with_each_other_on_every_shared_cell(
    real_cells: tuple[list[srm.SrmCell], list[srm.SrmCell]],
) -> None:
    """Two arms, two measurements: one dataset each, not one dataset with two values."""
    set1, set2 = real_cells
    first = {
        (cell.uniprot, srm.CONDITION_LABELS[cell.condition]): cell.copies_per_cell
        for cell in set1
        if cell.condition != srm.UNDESCRIBED_CONDITION
    }
    second = {
        (cell.uniprot, srm.CONDITION_LABELS[cell.condition]): cell.copies_per_cell
        for cell in set2
    }
    shared = sorted(set(first) & set(second))
    assert len(shared) == 558
    agreeing = [
        key
        for key in shared
        if abs(first[key] - second[key]) / second[key] < srm.IDENTITY_RTOL
    ]
    assert agreeing == []
    ratios = [first[key] / second[key] for key in shared]
    assert statistics.median(ratios) == pytest.approx(0.394, abs=5e-4)
