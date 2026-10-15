# torchcell/datasets/scerevisiae/jackson2020
# [[torchcell.datasets.scerevisiae.jackson2020]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/scerevisiae/jackson2020
# Test file: tests/torchcell/datasets/scerevisiae/test_jackson2020.py
"""Jackson 2020: single-cell RNA-seq of 72 barcoded TF-deletion strains in 11 conditions.

Jackson et al. 2020 (eLife 9:e51254, doi:10.7554/eLife.51254; citation key
``jacksonGeneRegulatoryNetwork2020``) built a prototrophic diploid array of eleven
nitrogen-regulatory transcription-factor deletions plus an HO-deletion wild-type control,
six independently constructed strains of each genotype, each strain carrying its own
transcribed barcode in the 3' UTR of its KanMX cassette. The 72 strains were grown as one
pool in eleven shake-flask conditions and profiled by 10x Chromium 3' v2 single-cell
RNA-seq; the barcode transcript assigns every cell to its strain.

RELEASE. Source code 2, ``elife-51254-code2.tsv.gz``: "The raw count matrix ... 38,225
observations (cells)", "unmodified transcript counts after deartifacting using UMIs".
38,225 rows x 6,829 count features + five metadata columns (``Genotype`` = the strain,
``Genotype_Group``, ``Replicate`` = the array well, ``Condition``, ``tenXBarcode``). Every
count is a non-negative integer and they sum to 83,703,440, the paper's "83,703,440
transcript counts from a total of 38,225 individual cells". The feature header is R's
``make.names`` form of the Ensembl R64-1-1.93 gene ids (``YDL247W.A``, ``X15S_rRNA``);
undoing it gives exactly the header of the same release's bulk matrix
(``TRIZOL_BULK.tsv``, Source code 1), which the loader checks. Two features are the drug
cassettes the authors added to the reference ("The antibiotic resistance marker cassettes
was added to both the FASTA and GTF files"): ``KANMX`` and ``NATMX``.

RECORD. One ``PseudobulkUmiCountExperiment`` per (strain, condition), the paper's own
pseudobulk: "The raw, unmodified UMI counts of all cells from each biological replicate
(with the same strain barcode) within a specific environmental growth condition were
summed". The six constructions of a genotype stay six records (they are the replicate
structure). 784 records: 72 strains x 11 conditions is 792, and eight (strain, condition)
pairs hold no cell in the release although the Results say "Each of the 72 strains is
found in each of the 11 conditions"; the eight are listed in the release ledger. The
phenotype is ``SingleCellExpressionPhenotype``, the one single-cell UMI-count leaf (#1016):
it keeps the per-cell barcodes and UMI totals, and the two cassette transcripts apart from
the genes. Each cell's library is its release ``Condition`` (the paper's "11 independent
single-cell sequencing batches, each corresponding to a single shake flask"), and the
release assigns no cell-cycle stage, so ``cell_cycle_stage_counts`` is a typed gap.
``measurement_type`` is the leaf's one tag, ``scrnaseq_umi_count``: the chemistry is the
10x Chromium 3' kit (``CHEMISTRY_QUOTE``), the same physical measurement as Boocock 2025
and N'Guessan 2025, and the per-strain sum is analysis that ``n_cells`` records.

REFERENCE. One per condition: the six HO-deletion control strains of that flask, their
cells summed ("six biological replicates of the wild-type control in which we deleted
the neutral HO locus"). The six control strains are also records of their own.

GENOTYPE. A homozygous whole-ORF deletion, ``Δgene::KanMX/Δgene::NatMX`` (Table S2), as a
``BarcodedKanMxDeletionPerturbation``: the 19-nt barcode of the KanMX allele (Table S2,
checked against Source code 4's ``bcdel_1_barcodes.fasta``), the array plate and well,
and the strain label the release uses. Systematic names are the authors' own
(``figure_common_settings.R`` lists the eleven TF ORFs); HO is resolved through the
genome. Seven diploid rows of Table S2 name the other STP gene than their strain label;
the haploid MATa parent row with the same barcode, the strain label and the release's
``Genotype_Group`` all name the label's gene, so the label is followed and the seven are
recorded. The background (FY4 x FY5, MATa/MATalpha, prototrophic) rides on a
``StrainReferenceGenome``.

ENVIRONMENT. ``CultureEnvironment`` per Table 1 condition, media from Supplementary Table
4 through the media library (``MINIMAL_BASE_JACKSON2020`` and its derivatives, liquid YPD,
YP + 2% v/v ethanol). The release's condition labels (``AmmoniumSulfate``) are mapped to
Table 1's abbreviations (``NLIM-NH4``) by the authors' ``figure_common_settings.R``.
Harvest at 4 h, DIAUXY at 10 h; RAPA is YPD with 200 ng/mL rapamycin for the last 30 min.
The flask temperature is not stated (every other growth step of the Methods is at 30 C),
so it is a typed gap.

RAW MIRROR. ``$DATA_ROOT/torchcell-raw/jacksonGeneRegulatoryNetwork2020/``: the matrix,
Supplementary file 1, ``figure_common_settings.R`` and ``TRIZOL_BULK.tsv`` (members of
Source code 1) and ``bcdel_1_barcodes.fasta`` (member of Source code 4), each re-fetched
from the PMC Article Datasets bucket on 2026-10-11 with bytes equal to the library
mirror's copies.
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import os.path as osp
import re
import shutil
import tarfile
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from pydantic import BaseModel, ConfigDict, model_validator
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    Organism,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.media import (
    CSTARVE_JACKSON2020,
    MMD_JACKSON2020,
    MMETOH_JACKSON2020,
    NLIM_GLN_JACKSON2020,
    NLIM_NH4_JACKSON2020,
    NLIM_PRO_JACKSON2020,
    NLIM_UREA_JACKSON2020,
    YP_ETHANOL_LIQUID_JACKSON2020,
    YPD_LIQUID,
    restated,
)
from torchcell.datamodels.schema import (
    BarcodedKanMxDeletionPerturbation,
    Concentration,
    ConcentrationUnit,
    CultureEnvironment,
    CultureFormat,
    EndpointRule,
    EukaryoticStrainBackground,
    Experiment,
    ExperimentReference,
    Genotype,
    MatingType,
    Media,
    PreCulture,
    PreCultureSource,
    PseudobulkUmiCountExperiment,
    PseudobulkUmiCountExperimentReference,
    Publication,
    SingleCellExpressionPhenotype,
    SmallMoleculePerturbation,
    Solvent,
    StrainConstruction,
    StrainReferenceGenome,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.datasets.scerevisiae.gene_name_reconcile import (
    reconcile_systematic_names,
)
from torchcell.literature.manifest import (
    ROLE_RAW_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
    sha256_file,
)
from torchcell.sequence.genome.registry import (
    SGD_S288C_R64,
    SGD_S288C_R64_GENBANK_ACCESSION,
)
from torchcell.sequence.genome.scerevisiae import SCerevisiaeGenome
from torchcell.verification.report import Provenance
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
)

logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
CITATION_KEY = "jacksonGeneRegulatoryNetwork2020"
PAPER_DOI = "10.7554/eLife.51254"
PAPER_TITLE = (
    "Gene regulatory network reconstruction using single-cell RNA sequencing of "
    "barcoded genotypes in diverse environments"
)
GEO_ACCESSION = "GSE125162"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
#: The PMC Article Datasets bucket object prefix of the eLife supplementary files.
PMC_KEY_PREFIX = "PMC7004572.1"
PMC_BUCKET_URL = "https://pmc-oa-opendata.s3.amazonaws.com/"

#: Library-mirror artifacts every quote below is read from (sha256 of the bytes).
PAPER_MD = "paper.md"
PAPER_MD_SHA256 = "453924f3d5ed998e99695768cf5b4067f2dbac3fb0af53dcbf8fd9017f199697"
SI_XLSX = "si/si6.xlsx"
SI_XLSX_SHA256 = "6c96fdf425f9d07c92d6bad8f44e54e1bebca6b83d06a7b7b3fbb06ea1923737"
TABLE_S2_SHEET = "Supplemental Table 2 - Strains"

#: Raw-mirror files (path relative to the mirror) and their sha256 pins.
MATRIX_NAME = "elife-51254-code2.tsv.gz"
MATRIX_REL = f"data/{MATRIX_NAME}"
MATRIX_SHA256 = "4bb210b7a1e2e9724e475ab516de646fbc36f2638ac6ade3bf15bf66ffb8c011"
SUPP1_NAME = "elife-51254-supp1.xlsx"
SUPP1_REL = f"si/{SUPP1_NAME}"
SUPP1_SHA256 = SI_XLSX_SHA256
SETTINGS_NAME = "figure_common_settings.R"
SETTINGS_REL = f"code/{SETTINGS_NAME}"
SETTINGS_SHA256 = "bf6b384c4f341d1a2c4a13024829967593b697bdfa109fa728f686d2c78acf9c"
BULK_NAME = "TRIZOL_BULK.tsv"
BULK_REL = f"data/{BULK_NAME}"
BULK_SHA256 = "95b5823d5da02d8c762399f5f537d8a92a16a8f8267bcad291a4c47d29944b1b"
BARCODES_NAME = "bcdel_1_barcodes.fasta"
BARCODES_REL = f"data/{BARCODES_NAME}"
BARCODES_SHA256 = "f16e245dde905c79f487b32c56a8eaeb07a5b530b71df8033afafa7c24904359"
#: The two tarballs three of the files are members of (sha256 of the containers).
CODE1_NAME = "elife-51254-code1.tar.gz"
CODE1_SHA256 = "ee6d8c92552669dbd11915f35d131d0d2d1de9e2fc403bc8a0600a8995e24d25"
CODE4_NAME = "elife-51254-code4.tar.gz"
CODE4_SHA256 = "371d9e791f4722d0bfd4df987680f029dace1d38a5a950256d35995bb41dd8b0"

RAW_FILES: dict[str, str] = {
    MATRIX_NAME: MATRIX_SHA256,
    SUPP1_NAME: SUPP1_SHA256,
    SETTINGS_NAME: SETTINGS_SHA256,
    BULK_NAME: BULK_SHA256,
    BARCODES_NAME: BARCODES_SHA256,
}
RAW_RELPATHS: dict[str, str] = {
    MATRIX_NAME: MATRIX_REL,
    SUPP1_NAME: SUPP1_REL,
    SETTINGS_NAME: SETTINGS_REL,
    BULK_NAME: BULK_REL,
    BARCODES_NAME: BARCODES_REL,
}
#: Which bucket object each mirror file is, and the tar member when it is one.
RAW_SOURCES: dict[str, tuple[str, str | None, str | None]] = {
    MATRIX_REL: ("elife-51254-code2.tsv.gz", None, None),
    SUPP1_REL: ("elife-51254-supp1.xlsx", None, None),
    SETTINGS_REL: (CODE1_NAME, SETTINGS_NAME, CODE1_SHA256),
    BULK_REL: (CODE1_NAME, "data/TRIZOL_BULK.tsv", CODE1_SHA256),
    BARCODES_REL: (CODE4_NAME, BARCODES_NAME, CODE4_SHA256),
}

#: The release as the paper counts it.
PAPER_CELLS = 38225
PAPER_TRANSCRIPT_COUNTS = 83703440
N_STRAINS = 72
N_CONDITIONS = 11
N_RELEASED_FEATURES = 6829
#: The release's metadata columns, in header order.
META_COLUMNS = ("Genotype", "Genotype_Group", "Replicate", "Condition", "tenXBarcode")
#: The drug-cassette transcripts the authors added to the reference.
MARKER_FEATURES = ("KANMX", "NATMX")
#: The wild-type control's strain-label prefix (``WT(ho)_A1``) and its deleted gene.
WT_LABEL = "WT(ho)"
WT_GENE_TOKEN = "ho"
#: Seven Table S2 diploid rows whose ``Genotype`` cell names the other STP gene than
#: the strain label; measured on the pinned workbook (sha256 6c96fdf4...). A change to
#: this set means the workbook changed and the rule must be re-read, so it is pinned.
TABLE_S2_GENOTYPE_DISAGREEMENTS = (
    "stp1_H11",
    "stp2_A12",
    "stp2_D12",
    "stp2_E12",
    "stp2_F12",
    "stp2_G12",
    "stp2_H12",
)


def _sv(
    value: object,
    quote: str,
    *,
    uri: str = PAPER_MD,
    sha: str = PAPER_MD_SHA256,
    page: str | None = None,
    note: str | None = None,
) -> SourcedValue:
    """A SourcedValue pinned to a mirrored Jackson 2020 artifact (quote + sha256)."""
    return SourcedValue(
        value=value,
        provenance=Provenance(
            source_uri=uri, citation_key=CITATION_KEY, sha256=sha, page=page
        ),
        quote=quote,
        note=note,
    )


def _table_s2_sv(
    value: object, quote: str, row: int, note: str | None = None
) -> SourcedValue:
    """A SourcedValue on one row of Supplemental Table 2 (sheet row is 1-based)."""
    return _sv(
        value,
        quote,
        uri=SI_XLSX,
        sha=SI_XLSX_SHA256,
        page=f"sheet '{TABLE_S2_SHEET}' row {row}",
        note=note,
    )


# --------------------------------------------------------------------------- #
# Verbatim quotes (paper.md is MinerU OCR, so LaTeX markup is kept as written)
# --------------------------------------------------------------------------- #
BACKGROUND_QUOTE = (
    "All yeast strains were generated from the prototrophic FY4 (MATa) or FY5 (MATa) "
    "background strains."
)
DESIGN_QUOTE = (
    "We engineered eleven different TF knockout genotypes, using six independently "
    "constructed biological replicates for each genotype."
)
WT_CONTROL_QUOTE = (
    "In addition, we constructed six biological replicates of the wild-type control in "
    "which we deleted the neutral HO locus."
)
WHOLE_ORF_QUOTE = (
    "We engineered TF gene deletions by precisely excising the entire TF open reading "
    "frame and introducing a unique transcriptional barcode that enables multiplexed "
    "analysis of genotypes using scRNAseq."
)
BARCODE_QUOTE = (
    "Genes were deleted using a modified kanMX cassette such that each of the 72 strains "
    "contains a unique transcriptional barcode in the ${ \\boldsymbol { 3 ^ { \\prime } } }$ "
    "untranslated region (UTR) of the G418 resistance gene"
)
HOMOZYGOUS_QUOTE = (
    "Homozygous diploids were constructed by mating to a strain containing the same TF "
    "knockout marked with a nourseothricin drug resistance cassette."
)
RELEASE_QUOTE = "Source code 2. The raw count matrix as a gzipped TSV file."
RELEASE_COUNT_QUOTE = (
    "This file contains 38,225 observations (cells). Doublets and low-count cells have "
    "already been removed; gene expression values are unmodified transcript counts after "
    "deartifacting using UMIs"
)
TOTAL_QUOTE = (
    "we recovered 83,703,440 transcript counts from a total of 38,225 individual cells."
)
PSEUDOBULK_QUOTE = (
    "The raw, unmodified UMI counts of all cells from each biological replicate (with "
    "the same strain barcode) within a specific environmental growth condition were "
    "summed, resulting in 72 samples per condition (six biological replicates for each "
    "of the 12 transcription factor deletions)."
)
EVERY_STRAIN_QUOTE = "Each of the 72 strains is found in each of the 11 conditions"
BATCH_QUOTE = (
    "This final data matrix is assembled from 11 independent single-cell sequencing "
    "batches, each corresponding to a single shake flask with a different growth "
    "condition."
)
REFERENCE_GENOME_QUOTE = (
    "The reference genome was obtained from Ensembl (Version R64-1-1) as a FASTA file, "
    "and the reference annotations were obtained from Ensembl (Version R64-1-1.93) in GTF "
    "format."
)
MARKER_QUOTE = (
    "The antibiotic resistance marker cassettes was added to both the FASTA and GTF files "
    "using command line tools."
)
CHEMISTRY_QUOTE = (
    "Single cell library preparation was done using the $ 1 0 \\times$ Genomics Chromium "
    "${ \\boldsymbol { 3 ^ { \\prime } } }$ v2 Single Cell Gene Expression Kit"
)
GEO_QUOTE = "accession number GEO: GSE125162."
PRECULTURE_QUOTE = (
    "The second recovery plate was pinned to a 96 well roundbottom plate containing "
    "$2 0 0 ~ \\mu \\ L$ YPD in each well and grown overnight at "
    "$3 0 ^ { \\circ } \\mathsf { C }$ ."
)
INOCULATION_QUOTE = (
    "of the washed cells were used to inoculate $5 0 ~ \\mathsf { m L }$ of the relevant "
    "media for the specific experimental condition in a shake flask."
)
HARVEST_QUOTE = (
    "All cultures were harvested $4 \\ : \\mathrm { h r }$ after inoculation, except for "
    "the culture harvested after the diauxic shift [DIAUXY], which was harvested "
    "$1 0 \\ : \\mathrm { h r }$ after inoculation."
)
RAPAMYCIN_QUOTE = (
    "The experiment treated with rapamycin was grown for 3 hr and "
    "$3 0 ~ \\mathsf { m i n }$ in YPD, and then $1 0 ~ \\mu \\ L$ of rapamycin stock "
    "$1 \\ m g / \\mathsf { m L }$ Millipore #553210 in ethanol) was added to a final "
    "concentration of $2 0 0 ~ \\mathrm { { n g / m L } }$ ."
)
HARVEST_DENSITY_QUOTE = (
    "Cell density (cells/mL) for each condition at harvest was as follows: (YPD 1.4e7; "
    "RAPA"
)
MEDIA_TABLE_QUOTE = "Specific media formulations are listed in Supplementary file 1-Supplemental Table 4."

#: Every paper.md quote the loader states, for the verifier's provenance audit.
PAPER_QUOTES: tuple[str, ...] = (
    BACKGROUND_QUOTE,
    DESIGN_QUOTE,
    WT_CONTROL_QUOTE,
    WHOLE_ORF_QUOTE,
    BARCODE_QUOTE,
    HOMOZYGOUS_QUOTE,
    RELEASE_QUOTE,
    RELEASE_COUNT_QUOTE,
    TOTAL_QUOTE,
    PSEUDOBULK_QUOTE,
    EVERY_STRAIN_QUOTE,
    BATCH_QUOTE,
    REFERENCE_GENOME_QUOTE,
    MARKER_QUOTE,
    CHEMISTRY_QUOTE,
    GEO_QUOTE,
    PRECULTURE_QUOTE,
    INOCULATION_QUOTE,
    HARVEST_QUOTE,
    RAPAMYCIN_QUOTE,
    HARVEST_DENSITY_QUOTE,
    MEDIA_TABLE_QUOTE,
)

#: Harvest cell densities (cells/mL) the Methods state per condition; DIAUXY is not
#: listed. NLIM-PRO's value is ``0 . 4 \phantom { - } \Theta ^ { 7 }`` in the OCR, read
#: as 0.4e7 by the sentence's pattern; kept in the ledger and the note, not in records.
HARVEST_CELLS_PER_ML: dict[str, float] = {
    "YPD": 1.4e7,
    "RAPA": 1.2e7,
    "YPEtOH": 1.0e7,
    "NLIM-GLN": 0.5e7,
    "NLIM-NH4": 0.8e7,
    "NLIM-PRO": 0.4e7,
    "NLIM-UREA": 0.5e7,
    "MMD": 1.1e7,
    "MMEtOH": 0.7e7,
    "CSTARVE": 0.1e7,
}

# --------------------------------------------------------------------------- #
# Strain background
# --------------------------------------------------------------------------- #
BACKGROUND_NAME = "FY4/FY5"
BACKGROUND_PROVENANCE: list[SourcedValue] = [
    _sv(
        "FY4 (MATa) and FY5 (MATalpha), prototrophic",
        BACKGROUND_QUOTE,
        page="Materials and methods, 'Yeast strain construction and growth'",
        note="the OCR prints both mating types as 'MATa'; Table S2 (rows 2 and 3) "
        "gives FY4 MATa and FY5 MATα. The paper names no background allele: the FY "
        "strains' S288C ancestry (Winston et al. 1995) is not mirrored, so the allele "
        "list is the paper's 'prototrophic' and nothing more",
    ),
    _table_s2_sv("FY4, MATa", "FY4 | MATa", 2),
    _table_s2_sv("FY5, MATalpha", "FY5 | MATα", 3),
    _sv(
        "homozygous diploids by mating the MATa KanMX deletant to the MATalpha NatMX "
        "deletant",
        HOMOZYGOUS_QUOTE,
        page="Results, 'Engineering a library of Prototrophic, "
        "Transcriptionally-Barcoded Gene Deletion Strains'",
    ),
]


def fy_background() -> EukaryoticStrainBackground:
    """The FY4 x FY5 diploid every strain of the array was built in."""
    return EukaryoticStrainBackground(
        reference_strain="S288C",
        assembly_set=SGD_S288C_R64,
        name=BACKGROUND_NAME,
        parents=["FY4", "FY5"],
        construction="MATa FY4 and MATalpha FY5 deletants of the same gene pinned onto "
        "one YPD plate to mate, then three rounds of diploid selection on YPD + G418 + "
        "nourseothricin",
        mating_type=MatingType.a_alpha,
        ploidy="diploid",
        alleles=[],
        provenance=list(BACKGROUND_PROVENANCE),
    )


def reference_genome() -> StrainReferenceGenome:
    """The diploid FY4/FY5 background on the S288C R64 coordinate system."""
    return StrainReferenceGenome(
        species="Saccharomyces cerevisiae",
        strain=BACKGROUND_NAME,
        ploidy="diploid",
        background=fy_background(),
        assembly_set=SGD_S288C_R64,
        assembly_accession=SGD_S288C_R64_GENBANK_ACCESSION,
    )


# --------------------------------------------------------------------------- #
# Strains (Supplemental Table 2)
# --------------------------------------------------------------------------- #
class StrainTableError(ValueError):
    """Supplemental Table 2 does not have the rows this loader reads."""


class StrainSpec(BaseModel):
    """One diploid strain of the array: its Table S2 row and its MATa parent's row."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    strain: str
    gene_token: str
    well: str
    plate_location: str
    barcode: str
    diploid_genotype: str
    diploid_row: int
    diploid_row_quote: str
    parent_dgy: int
    parent_genotype: str
    parent_row: int
    natmx_parent_row: int

    @property
    def is_wild_type(self) -> bool:
        """True for the six HO-deletion control strains."""
        return self.gene_token == WT_GENE_TOKEN

    @property
    def diploid_cell_disagrees(self) -> bool:
        """True when the diploid row's Genotype cell names another gene than the label."""
        return f"Δ{self.gene_token}::KanMX/Δ{self.gene_token}::NatMX" not in (
            self.diploid_genotype
        )


def render_row(cells: tuple[Any, ...]) -> str:
    """A sheet row as its quote form: the non-empty cells joined by ``' | '``.

    The rendering the media library's xlsx quotes are cut from (``str`` of each cell
    openpyxl reads, so a numeric cell keeps openpyxl's ``1.0``); quotes are chosen from
    string cells.
    """
    return " | ".join(str(cell) for cell in cells if cell is not None)


def sheet_text(xlsx_path: str | Path) -> str:
    """Every sheet's non-empty rows rendered by ``render_row``, joined by ``' / '``."""
    import openpyxl

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
    try:
        sheets = [
            " / ".join(
                render_row(tuple(row))
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None for cell in row)
            )
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()
    return "\n".join(sheets)


def read_sheet_rows(xlsx_path: str | Path, sheet: str) -> list[tuple[Any, ...]]:
    """Every row of one sheet as cell tuples (row 1 first)."""
    import openpyxl

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
    try:
        return [tuple(r) for r in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _label_token(strain: str) -> tuple[str, str]:
    """``gat1_H5`` -> (``gat1``, ``H5``); ``WT(ho)_A1`` -> (``ho``, ``A1``)."""
    prefix, well = strain.rsplit("_", 1)
    return (WT_GENE_TOKEN if prefix == WT_LABEL else prefix), well


def read_strain_table(xlsx_path: str | Path) -> list[StrainSpec]:
    """The 72 diploid strains of Table S2, each joined to its two haploid parents.

    A diploid row has a ``MATa/MATα`` genotype and a ``BC_DEL_1_<well>`` plate location;
    its MATa parent is the FY4 row with the same strain label and barcode, its MATalpha
    parent the FY5 NatMX row of the label's gene. The deleted gene is the label's: the
    MATa parent's genotype must name it (``MATa Δ<gene>::KanMX``). A diploid ``Genotype``
    cell that names another gene is kept verbatim and must be one of the pinned
    ``TABLE_S2_GENOTYPE_DISAGREEMENTS``.
    """
    rows = read_sheet_rows(xlsx_path, TABLE_S2_SHEET)
    header = rows[0]
    if header[:6] != (
        "Strain ID (DGY)",
        "Strain",
        "Genotype",
        "Parent",
        "Barcode",
        "Plate Location",
    ):
        raise StrainTableError(f"unexpected Table S2 header {header!r}")
    diploid: dict[str, tuple[int, tuple[Any, ...]]] = {}
    mata: dict[str, tuple[int, tuple[Any, ...]]] = {}
    natmx: dict[str, int] = {}
    for number, row in enumerate(rows[1:], start=2):
        dgy, strain, genotype, parent, barcode, location = row[:6]
        if strain is None:
            continue
        strain = str(strain)
        if genotype is not None and str(genotype).startswith("MATa/MATα"):
            diploid[strain] = (number, row)
        elif parent == "FY4":
            mata[strain] = (number, row)
        elif parent == "FY5":
            natmx[strain] = number
    if len(diploid) != N_STRAINS or len(mata) != N_STRAINS:
        raise StrainTableError(
            f"Table S2 lists {len(diploid)} diploid and {len(mata)} MATa strains, "
            f"expected {N_STRAINS} of each"
        )
    specs: list[StrainSpec] = []
    disagreeing: list[str] = []
    for strain, (number, row) in diploid.items():
        token, well = _label_token(strain)
        if strain not in mata:
            raise StrainTableError(f"{strain}: no MATa parent row")
        parent_number, parent = mata[strain]
        if parent[4] != row[4]:
            raise StrainTableError(f"{strain}: diploid and MATa barcodes differ")
        if parent[2] != f"MATa Δ{token}::KanMX":
            raise StrainTableError(
                f"{strain}: MATa parent genotype {parent[2]!r} does not delete {token}"
            )
        if token not in natmx:
            raise StrainTableError(f"{strain}: no FY5 NatMX parent for {token}")
        if row[5] != f"BC_DEL_1_{well}":
            raise StrainTableError(f"{strain}: plate location {row[5]!r} is not {well}")
        spec = StrainSpec(
            strain=strain,
            gene_token=token,
            well=well,
            plate_location=str(row[5]),
            barcode=str(row[4]),
            diploid_genotype=str(row[2]),
            diploid_row=number,
            diploid_row_quote=render_row(row),
            parent_dgy=int(parent[0]),
            parent_genotype=str(parent[2]),
            parent_row=parent_number,
            natmx_parent_row=natmx[token],
        )
        if spec.diploid_cell_disagrees:
            disagreeing.append(strain)
        specs.append(spec)
    if tuple(sorted(disagreeing)) != TABLE_S2_GENOTYPE_DISAGREEMENTS:
        raise StrainTableError(
            f"diploid Genotype cells naming another gene than their label: "
            f"{sorted(disagreeing)}, pinned {list(TABLE_S2_GENOTYPE_DISAGREEMENTS)}"
        )
    return specs


def read_barcode_fasta(path: str | Path) -> dict[str, str]:
    """``{strain label: barcode}`` from Source code 4's ``bcdel_1_barcodes.fasta``."""
    out: dict[str, str] = {}
    label: str | None = None
    for line in Path(path).read_text().splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith(">"):
            label = line[1:].strip()
            continue
        if label is None or label in out:
            raise ValueError(f"{path}: sequence line without a fresh header: {line!r}")
        out[label] = line
        label = None
    return out


# --------------------------------------------------------------------------- #
# The authors' condition labels and TF systematic names (figure_common_settings.R)
# --------------------------------------------------------------------------- #
_R_VECTOR = re.compile(r"^(?P<name>[A-Z.]+) <<- c\((?P<body>[^)]*)\)\s*$", re.M)
_R_TF = re.compile(r'"(?P<orf>Y[A-P][LR]\d{3}[WC])"\s*,?\s*# (?P<gene>[a-z0-9]+)')


def _r_strings(body: str) -> list[str]:
    return re.findall(r'"([^"]*)"', body)


class AuthorSettings(BaseModel):
    """What ``figure_common_settings.R`` states: label maps and the TF ORFs."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    condition_label: dict[str, str]
    condition_levels_line: str
    condition_labels_line: str
    tf_systematic_name: dict[str, str]
    tf_lines: dict[str, str]


def read_author_settings(path: str | Path) -> AuthorSettings:
    """Parse the release-label -> Table 1 abbreviation map and the eleven TF ORFs.

    ``CONDITION.LEVELS`` (the release's ``Condition`` values) and ``CONDITION.LABELS``
    (Table 1's abbreviations) are parallel vectors; ``nitrogen.tfs`` lists each TF ORF
    with its gene in a trailing comment (``"YFL021W", # gat1``).
    """
    text = Path(path).read_text()
    vectors = {m.group("name"): m for m in _R_VECTOR.finditer(text)}
    levels = _r_strings(vectors["CONDITION.LEVELS"].group("body"))
    labels = _r_strings(vectors["CONDITION.LABELS"].group("body"))
    if len(levels) != N_CONDITIONS or len(labels) != N_CONDITIONS:
        raise ValueError(f"{path}: {len(levels)} levels and {len(labels)} labels")
    tfs: dict[str, str] = {}
    tf_lines: dict[str, str] = {}
    for line in text.splitlines():
        match = _R_TF.search(line)
        if match is None:
            continue
        gene = match.group("gene")
        if gene in tfs:
            raise ValueError(f"{path}: {gene} is listed twice")
        tfs[gene] = match.group("orf")
        tf_lines[gene] = line.strip()
    return AuthorSettings(
        condition_label=dict(zip(levels, labels, strict=True)),
        condition_levels_line=vectors["CONDITION.LEVELS"].group(0).strip(),
        condition_labels_line=vectors["CONDITION.LABELS"].group(0).strip(),
        tf_systematic_name=tfs,
        tf_lines=tf_lines,
    )


# --------------------------------------------------------------------------- #
# Conditions (Table 1 + Supplementary Table 4)
# --------------------------------------------------------------------------- #
JACKSON_YPD: Media = restated(
    YPD_LIQUID,
    _sv(
        "2% peptone, 2% glucose, 1% yeast extract",
        "YPD | 2% (w/v) peptone | BD Difco #DF0118-17-0",
        uri=SI_XLSX,
        sha=SI_XLSX_SHA256,
        page="sheet 'Supplemental Table 4 - Media' rows 2-4",
        note="the block's other rows are '2% (w/v) glucose | BD Difco #DF0155-08-5' and "
        "'1% (w/v) yeast extract | BD Difco #DF0127-17-9': the same three rows and "
        "amounts as the library's YPD_LIQUID",
    ),
)
"""Liquid YPD of Table 4: the library ``YPD_LIQUID`` restated with Table 4's rows."""

RAPAMYCIN_UG_PER_ML = 0.2  # 200 ng/mL
#: 10 uL of an ethanol stock into the 50 mL culture.
RAPAMYCIN_SOLVENT_PERCENT = 10.0 / 50_000.0 * 100.0
GROWTH_HOURS = 4.0
DIAUXY_HOURS = 10.0


class ConditionSpec(BaseModel):
    """One Table 1 condition: its release label, Table 1 row, medium and timing."""

    model_config = ConfigDict(extra="forbid", arbitrary_types_allowed=True)

    abbreviation: str
    release_label: str
    table1_row: str
    media: Media
    duration_hours: float
    rapamycin: bool = False


def build_conditions(settings: AuthorSettings) -> dict[str, ConditionSpec]:
    """The eleven conditions keyed by Table 1 abbreviation, in Table 1's order.

    The release label is the authors' own map (``figure_common_settings.R``); the
    medium is the Table 4 block Table 1's carbon and nitrogen sources name.
    """
    release_of = {abbrev: level for level, abbrev in settings.condition_label.items()}

    def spec(
        abbreviation: str,
        row: str,
        media: Media,
        *,
        hours: float = GROWTH_HOURS,
        rapamycin: bool = False,
    ) -> ConditionSpec:
        return ConditionSpec(
            abbreviation=abbreviation,
            release_label=release_of[abbreviation],
            table1_row=row,
            media=media,
            duration_hours=hours,
            rapamycin=rapamycin,
        )

    specs = [
        spec("YPD", "<td>Yeast Extract, Peptone, Glucose</td><td>YPD</td><td>YP</td><td>D-Glucose</td>", JACKSON_YPD),
        spec("DIAUXY", "<td>YPD (Harvested after Post-Diauxic Shift)</td><td>DIAUXY</td><td>YP</td><td>D-Glucose</td>", JACKSON_YPD, hours=DIAUXY_HOURS),
        spec("RAPA", "<td>YPD + 200 ng/mL Rapamycin</td><td>RAPA</td><td>YP</td><td>D-Glucose</td>", JACKSON_YPD, rapamycin=True),
        spec("YPEtOH", "<td>Yeast Extract, Peptone, Ehanol</td><td>YPEtOH</td><td>YP</td><td>Ethanol</td>", YP_ETHANOL_LIQUID_JACKSON2020),
        spec("MMD", "<td>Minimal Media (Glucose)</td><td>MMD</td><td>20 mM (NH4)2SO4</td><td>D-Glucose</td>", MMD_JACKSON2020),
        spec("MMEtOH", "<td>Minimal Media (Ethanol)</td><td>MMEtOH</td><td>20 mM (NH4)2SO4</td><td>Ethanol</td>", MMETOH_JACKSON2020),
        spec("NLIM-GLN", "<td>Nitrogen Limited Minimal Media (with Glutamine)</td><td>NLIM-GLN</td><td>0.8 mM L-Glutamine</td><td>D-Glucose</td>", NLIM_GLN_JACKSON2020),
        spec("NLIM-PRO", "<td>Nitrogen Limited Minimal Media (with Proline)</td><td>NLIM-PRO</td><td>0.8 mM L-Proline</td><td>D-Glucose</td>", NLIM_PRO_JACKSON2020),
        spec("NLIM-NH4", "<td>Nitrogen Limited Minimal Media (with NH4)</td><td>NLIM-NH4</td><td>0.8 mM (NH4)2SO4</td><td>D-Glucose</td>", NLIM_NH4_JACKSON2020),
        spec("NLIM-UREA", "<td>Nitrogen Limited Minimal Media (with Urea)</td><td>NLIM-UREA</td><td>0.8 mM Urea</td><td>D-Glucose</td>", NLIM_UREA_JACKSON2020),
        spec("CSTARVE", "<td>Carbon Starvation</td><td>CSTARVE</td><td>1 mM (NH4)2SO4</td><td>None</td>", CSTARVE_JACKSON2020),
    ]  # fmt: skip
    conditions = {s.abbreviation: s for s in specs}
    if set(conditions) != set(release_of):
        raise ValueError(
            f"Table 1 abbreviations {sorted(set(conditions) ^ set(release_of))} differ "
            "from figure_common_settings.R's labels"
        )
    return conditions


_TEMPERATURE_GAP_NOTE = (
    "the shake-flask growth sentence ('These flasks were grown for 4 hr.') states no "
    "temperature; the recovery plates, the 96-well YPD pre-culture, the array steps and "
    "the separate bulk-RNA flasks are each stated at 30 C, so 30 C is likely and "
    "unconfirmed"
)


def rapamycin_perturbation() -> SmallMoleculePerturbation:
    """200 ng/mL rapamycin from a 1 mg/mL ethanol stock, 30 min before harvest."""
    return SmallMoleculePerturbation(
        description="rapamycin added 3.5 h after inoculation, 30 min before the 4 h "
        "harvest",
        compound=resolved_compound("rapamycin"),
        concentration=Concentration(
            value=RAPAMYCIN_UG_PER_ML, unit=ConcentrationUnit.ug_per_ml
        ),
        solvent=Solvent(
            name="ethanol",
            percent=RAPAMYCIN_SOLVENT_PERCENT,
            compound=resolved_compound("ethanol"),
        ),
    )


def culture_format() -> CultureFormat:
    """The 50 mL shake flask every condition was grown in, read at a fixed time."""
    return CultureFormat(
        vessel="shake flask",
        working_volume_ul=50_000.0,
        endpoint=EndpointRule.fixed_duration,
        provenance=[
            _sv("50 mL shake flask", INOCULATION_QUOTE, page="'Culturing and harvest'"),
            _sv(
                "harvest 4 h after inoculation, DIAUXY 10 h",
                HARVEST_QUOTE,
                page="Table 1 legend",
            ),
        ],
    )


def pre_culture() -> PreCulture:
    """The pooled overnight 96-well YPD culture each flask was inoculated from."""
    return PreCulture(
        source=PreCultureSource.overnight_culture,
        medium=JACKSON_YPD,
        provenance=[
            _sv(
                "overnight in 200 uL YPD per well at 30 C, then pooled and washed twice "
                "in PBS",
                PRECULTURE_QUOTE,
                page="'Culturing and harvest'",
            )
        ],
    )


def condition_environment(spec: ConditionSpec) -> CultureEnvironment:
    """A condition's culture: medium, timing, rapamycin for RAPA, temperature gapped."""
    return CultureEnvironment(
        media=spec.media,
        temperature=None,
        perturbations=[rapamycin_perturbation()] if spec.rapamycin else [],
        aerobicity="aerobic",
        duration_hours=spec.duration_hours,
        culture_format=culture_format(),
        pre_culture=pre_culture(),
        auxotroph_supplements=[],
        provenance_gaps=[
            ProvenanceGap(
                field="temperature",
                reason=ProvenanceGapReason.not_reported_by_primary,
                note=_TEMPERATURE_GAP_NOTE,
            )
        ],
    )


# --------------------------------------------------------------------------- #
# Features: the release header -> R64 identifiers
# --------------------------------------------------------------------------- #
class FeatureMapError(ValueError):
    """The release header does not undo to the bulk matrix's header."""


def unmangle(name: str) -> str:
    """Undo R's ``make.names`` on an Ensembl R64-1-1 gene id.

    ``make.names`` prefixes an ``X`` to a name starting with a digit (``15S_rRNA`` ->
    ``X15S_rRNA``) and turns ``-`` into ``.`` (``YDL247W-A`` -> ``YDL247W.A``).
    """
    if len(name) > 1 and name[0] == "X" and name[1].isdigit():
        name = name[1:]
    return name.replace(".", "-")


class FeatureMap(BaseModel):
    """How every released count column is keyed in the records."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    released: list[str]
    source_ids: list[str]
    keys: list[str]
    gene_columns: list[int]
    marker_columns: list[int]
    status: dict[str, str]
    resolved: dict[str, str | None]

    @property
    def gene_keys(self) -> list[str]:
        """The record keys of the gene columns, in column order."""
        return [self.keys[i] for i in self.gene_columns]

    def legacy(self) -> dict[str, str]:
        """Gene ids kept verbatim although they are not an R64-4-1 feature, and why.

        A retired id is not in the annotation; a merged id names a feature that now
        carries another id, which a second released column already maps to, so both
        keep their source ids and stay two keys.
        """
        out: dict[str, str] = {}
        for i in self.gene_columns:
            source = self.source_ids[i]
            if self.keys[i] != source:
                continue
            if self.status[source] == "retired":
                out[source] = "retired: not in R64-4-1"
            elif self.resolved[source] != source:
                out[source] = f"merged into {self.resolved[source]}; kept distinct"
        return out

    @property
    def universe_sha256(self) -> str:
        """sha256 of the sorted gene keys joined by newlines."""
        text = "\n".join(sorted(self.gene_keys))
        return hashlib.sha256(text.encode("utf-8")).hexdigest()


def build_feature_map(
    released: list[str], bulk_header: list[str], genome: SCerevisiaeGenome
) -> FeatureMap:
    """Key every released count column: an R64 identifier for a gene, the marker name
    for a cassette transcript.

    The un-mangled header must equal the bulk matrix's header exactly (same pipeline,
    same annotation). Gene ids are reconciled to R64-4-1 by the shared retain-all policy
    (``reconcile_systematic_names``): a renamed or non-``gene`` locus takes its current
    id, two source ids merged into one current feature both keep their source ids, a
    retired id is kept verbatim. Every key is distinct.
    """
    source_ids = [unmangle(name) for name in released]
    if source_ids != bulk_header:
        diff = sorted(set(source_ids) ^ set(bulk_header))
        raise FeatureMapError(
            f"the un-mangled release header differs from TRIZOL_BULK.tsv's: {diff[:10]}"
        )
    marker_columns = [i for i, s in enumerate(source_ids) if s in MARKER_FEATURES]
    if sorted(source_ids[i] for i in marker_columns) != sorted(MARKER_FEATURES):
        raise FeatureMapError(
            f"marker columns {marker_columns} are not {MARKER_FEATURES}"
        )
    gene_columns = [i for i in range(len(source_ids)) if i not in set(marker_columns)]
    gene_ids = pd.Series([source_ids[i] for i in gene_columns])
    reconciled = reconcile_systematic_names(genome, gene_ids, label="Jackson 2020")
    keys = list(source_ids)
    status: dict[str, str] = {}
    resolved: dict[str, str | None] = {}
    for column, source, key in zip(gene_columns, gene_ids, reconciled, strict=True):
        keys[column] = str(key)
        resolution = genome.resolve_gene_name(source)
        if resolution.status.value == "ambiguous":
            raise FeatureMapError(f"{source} resolves to several features")
        status[source] = resolution.status.value
        resolved[source] = resolution.systematic_name
    if len(set(keys)) != len(keys):
        raise FeatureMapError("two released columns map to one key")
    return FeatureMap(
        released=list(released),
        source_ids=source_ids,
        keys=keys,
        gene_columns=gene_columns,
        marker_columns=marker_columns,
        status=status,
        resolved=resolved,
    )


# --------------------------------------------------------------------------- #
# Release reading
# --------------------------------------------------------------------------- #
class ReleaseShapeError(ValueError):
    """The release does not have the shape the paper and this loader state."""


class StrainConditionSum(BaseModel):
    """The summed counts of one (strain, condition) and the cells they came from."""

    model_config = ConfigDict(extra="forbid", arbitrary_types_allowed=True)

    counts: np.ndarray[Any, Any]
    cell_barcodes: list[str]
    cell_libraries: list[str]
    cell_total_umis: list[int]


def read_release_header(path: str | Path) -> list[str]:
    """The count-column names of the matrix (the five metadata columns excluded)."""
    header = pd.read_csv(path, sep="\t", nrows=0, compression="gzip").columns.tolist()
    if tuple(header[-len(META_COLUMNS) :]) != META_COLUMNS:
        raise ReleaseShapeError(
            f"the last columns are {header[-5:]}, not {META_COLUMNS}"
        )
    return header[: -len(META_COLUMNS)]


def read_release(
    path: str | Path,
    strains: dict[str, StrainSpec],
    release_labels: set[str],
    chunksize: int = 4000,
) -> tuple[dict[tuple[str, str], StrainConditionSum], int]:
    """Sum the count matrix per (strain, condition), keeping every cell's barcode/total.

    Refuses a non-integer or negative count, a strain outside Table S2, a condition
    outside the authors' label map, a ``Genotype_Group`` or ``Replicate`` that disagrees
    with the strain label, a droplet barcode seen twice in one condition, and a cell
    count other than the paper's.
    """
    features = read_release_header(path)
    sums: dict[tuple[str, str], StrainConditionSum] = {}
    seen_barcodes: dict[str, set[str]] = {}
    n_cells = 0
    reader = pd.read_csv(
        path, sep="\t", index_col=0, compression="gzip", chunksize=chunksize
    )
    for chunk in tqdm(reader, desc="reading the count matrix"):
        counts = chunk[features].to_numpy()
        if counts.dtype.kind != "i":
            raise ReleaseShapeError(f"non-integer counts (dtype {counts.dtype})")
        if (counts < 0).any():
            raise ReleaseShapeError("a negative count")
        totals = counts.sum(axis=1)
        for row, (strain, group, well, condition, barcode) in enumerate(
            chunk[list(META_COLUMNS)].itertuples(index=False, name=None)
        ):
            if strain not in strains:
                raise ReleaseShapeError(f"strain {strain!r} is not in Table S2")
            if condition not in release_labels:
                raise ReleaseShapeError(f"condition {condition!r} is not a known label")
            spec = strains[strain]
            expected_group = WT_LABEL if spec.is_wild_type else spec.gene_token
            if group != expected_group or well != spec.well:
                raise ReleaseShapeError(
                    f"{strain}: Genotype_Group {group!r} / Replicate {well!r} disagree "
                    "with the label"
                )
            barcodes = seen_barcodes.setdefault(condition, set())
            if barcode in barcodes:
                raise ReleaseShapeError(f"{condition}: barcode {barcode} seen twice")
            barcodes.add(barcode)
            key = (strain, condition)
            if key not in sums:
                sums[key] = StrainConditionSum(
                    counts=np.zeros(len(features), dtype=np.int64),
                    cell_barcodes=[],
                    cell_libraries=[],
                    cell_total_umis=[],
                )
            entry = sums[key]
            entry.counts += counts[row]
            entry.cell_barcodes.append(str(barcode))
            # One 10x library per condition (BATCH_QUOTE): the release's Condition
            # column is the cell's library.
            entry.cell_libraries.append(str(condition))
            entry.cell_total_umis.append(int(totals[row]))
        n_cells += len(chunk)
    if n_cells != PAPER_CELLS:
        raise ReleaseShapeError(f"{n_cells} cells, the paper states {PAPER_CELLS}")
    return sums, n_cells


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
#: The release labels no cell with a cell-cycle stage: its five metadata columns are
#: ``META_COLUMNS``, and the paper's cell-cycle reading (Figure 3) clusters cells by
#: marker-gene expression without releasing a per-cell label.
CELL_CYCLE_GAP = ProvenanceGap(
    field="cell_cycle_stage_counts",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=Provenance(
        source_uri=MATRIX_REL, citation_key=CITATION_KEY, sha256=MATRIX_SHA256
    ),
    note="the release's metadata columns are Genotype, Genotype_Group, Replicate, "
    "Condition and tenXBarcode; no cell carries a cell-cycle stage. Figure 3 clusters "
    "cells within a condition by cell-cycle marker expression, which is not released "
    "as a per-cell label",
)


def pseudobulk_phenotype(
    entry: StrainConditionSum, features: FeatureMap
) -> SingleCellExpressionPhenotype:
    """One (strain, condition) sum as the phenotype the records store."""
    counts = entry.counts
    umi = {
        features.keys[i]: int(counts[i]) for i in features.gene_columns if counts[i] > 0
    }
    markers = {features.keys[i]: int(counts[i]) for i in features.marker_columns}
    return SingleCellExpressionPhenotype(
        umi_count=umi,
        marker_umi_count=markers,
        total_umis=int(counts.sum()),
        n_features=len(features.gene_columns),
        feature_universe_sha256=features.universe_sha256,
        n_cells=len(entry.cell_barcodes),
        cell_cycle_stage_counts=None,
        library_cell_counts=dict(Counter(entry.cell_libraries)),
        cell_barcodes=list(entry.cell_barcodes),
        cell_libraries=list(entry.cell_libraries),
        cell_total_umis=list(entry.cell_total_umis),
        provenance_gaps=[CELL_CYCLE_GAP],
    )


def pooled(entries: list[StrainConditionSum]) -> StrainConditionSum:
    """Several strains' cells of one condition as one sum (the reference's)."""
    return StrainConditionSum(
        counts=np.sum([e.counts for e in entries], axis=0),
        cell_barcodes=[b for e in entries for b in e.cell_barcodes],
        cell_libraries=[b for e in entries for b in e.cell_libraries],
        cell_total_umis=[n for e in entries for n in e.cell_total_umis],
    )


def strain_genotype(spec: StrainSpec, systematic: str, standard: str) -> Genotype:
    """The homozygous barcoded deletion of one array strain."""
    return Genotype(
        perturbations=[
            BarcodedKanMxDeletionPerturbation(
                systematic_gene_name=systematic,
                perturbed_gene_name=standard,
                barcode=spec.barcode,
                collection="BC_DEL_1 (Jackson 2020 prototrophic barcoded transcription "
                "factor deletion array)",
                cassette="KanMX/NatMX",
                construction=StrainConstruction(
                    strain_accession=spec.strain, plate="BC_DEL_1", well=spec.well
                ),
            )
        ]
    )


def publication() -> Publication:
    """The eLife paper."""
    return Publication(doi=PAPER_DOI, doi_url=f"https://doi.org/{PAPER_DOI}")


class ReleaseLedger(BaseModel):
    """Every released cell and count, accounted for (``preprocess/release_ledger.json``)."""

    model_config = ConfigDict(extra="forbid")

    matrix_sha256: str
    n_cells_released: int
    n_cells_served: int
    n_wild_type_cells: int
    n_transcript_counts_released: int
    n_transcript_counts_served: int
    paper_cells: int
    paper_transcript_counts: int
    n_features_released: int
    n_gene_features: int
    marker_features: list[str]
    feature_status: dict[str, int]
    remapped_features: dict[str, str]
    legacy_features: dict[str, str]
    n_strains: int
    n_conditions: int
    n_records: int
    n_references: int
    strain_conditions_without_cells: list[str]
    cells_per_condition: dict[str, int]
    records_per_condition: dict[str, int]
    table_s2_genotype_disagreements: list[str]
    harvest_cells_per_ml: dict[str, float]
    refused: dict[str, int]

    @model_validator(mode="after")
    def _accounted(self) -> ReleaseLedger:
        if self.n_cells_served + sum(self.refused.values()) != self.n_cells_released:
            raise ValueError("served + refused cells must equal released cells")
        if self.n_cells_released != self.paper_cells:
            raise ValueError("the release's cell count is not the paper's")
        if self.n_transcript_counts_released != self.paper_transcript_counts:
            raise ValueError("the release's count total is not the paper's")
        if self.n_records + len(self.strain_conditions_without_cells) != (
            self.n_strains * self.n_conditions
        ):
            raise ValueError("records + empty pairs must equal strains x conditions")
        return self


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    from dotenv import load_dotenv

    load_dotenv()
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/jacksonGeneRegulatoryNetwork2020``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def _bucket_url(name: str) -> str:
    return f"{PMC_BUCKET_URL}{PMC_KEY_PREFIX}/{name}"


def _retrieval(relpath: str, sha: str, retrieved_at: str) -> RetrievalRecord:
    """The bucket object (or the tar member of one) a mirror file was cut from."""
    container, member, container_sha = RAW_SOURCES[relpath]
    url = _bucket_url(container)
    if member is None:
        return RetrievalRecord(
            method=RetrievalMethod.pmc_cloud,
            source_url=url,
            retriever="torchcell.literature.retrieve.pmc_cloud_object",
            params={"key": f"{PMC_KEY_PREFIX}/{container}"},
            sha256=sha,
            retrieved_at=retrieved_at,
        )
    return RetrievalRecord(
        method=RetrievalMethod.pmc_cloud,
        source_url=url,
        retriever="torchcell.literature.retrieve.tar_member",
        params={"url": url, "member": member, "container_sha256": container_sha},
        sha256=sha,
        retrieved_at=retrieved_at,
    )


def deposit_raw_mirror(
    *, fetch_dir: str | Path, retrieved_at: str, data_root: str | None = None
) -> Path:
    """Write the raw mirror from the bucket objects already fetched into ``fetch_dir``.

    Run once by hand after ``curl -sSf -O <PMC_BUCKET_URL>PMC7004572.1/<name>`` for
    ``elife-51254-code1.tar.gz``, ``-code2.tsv.gz``, ``-code4.tar.gz`` and
    ``-supp1.xlsx``. Each container must hash to its pin, each member is cut out of its
    container, and every deposited file must hash to the loader's pin. Idempotent by
    sha256: an existing file with the recorded hash is left alone, a differing one
    raises.
    """
    fetch = Path(fetch_dir)
    containers = {
        CODE1_NAME: CODE1_SHA256,
        CODE4_NAME: CODE4_SHA256,
        MATRIX_NAME: MATRIX_SHA256,
        SUPP1_NAME: SUPP1_SHA256,
    }
    for name, pin in containers.items():
        if sha256_file(fetch / name) != pin:
            raise ValueError(f"{fetch / name} does not hash to its pin {pin}")
    root = raw_mirror_dir(data_root)
    files: list[ArtifactRecord] = []
    for name, relpath in RAW_RELPATHS.items():
        container, member, _ = RAW_SOURCES[relpath]
        if member is None:
            data = (fetch / container).read_bytes()
        else:
            with tarfile.open(fetch / container, "r:gz") as archive:
                extracted = archive.extractfile(member)
                if extracted is None:
                    raise ValueError(f"{container}: member {member!r} did not extract")
                data = extracted.read()
        sha = hashlib.sha256(data).hexdigest()
        if sha != RAW_FILES[name]:
            raise ValueError(
                f"{relpath} hashes to {sha}, the loader pins {RAW_FILES[name]}"
            )
        dest = root / relpath
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            if sha256_file(dest) != sha:
                raise RuntimeError(f"{dest} exists with a different sha256; refusing")
        else:
            with open(dest, "wb") as handle:
                shutil.copyfileobj(io.BytesIO(data), handle)
        files.append(
            ArtifactRecord(
                path=relpath,
                role=ROLE_RAW_DATA,
                bytes=dest.stat().st_size,
                sha256=sha,
                source=_bucket_url(container),
                original_filename=container if member is None else member,
                retrieval=_retrieval(relpath, sha, retrieved_at),
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title=PAPER_TITLE,
        files=files,
        si_data_sources=[_bucket_url(c) for c in sorted(containers)],
        si_expected=sorted(RAW_RELPATHS.values()),
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


# --------------------------------------------------------------------------- #
# Dataset
# --------------------------------------------------------------------------- #
@register_dataset
class TfDeletionPseudobulkJackson2020Dataset(ExperimentDataset):
    """Jackson 2020: one record per (barcoded strain, condition), 784 records."""

    organism = Organism.scerevisiae

    def __init__(
        self,
        root: str = "data/torchcell/tf_deletion_pseudobulk_jackson2020",
        io_workers: int = 0,
        genome: SCerevisiaeGenome | None = None,
        transform: Any | None = None,
        pre_transform: Any | None = None,
        **kwargs: Any,
    ) -> None:
        """Build or load the store; a genome is required to key the gene features."""
        self.genome = genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return PseudobulkUmiCountExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return PseudobulkUmiCountExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The five raw-mirror files the build reads."""
        return list(RAW_FILES)

    def download(self) -> None:
        """Link the five mirror files, each checked against its pin and the manifest."""
        data_root = _data_root()
        manifest = load_manifest(data_root)
        os.makedirs(self.raw_dir, exist_ok=True)
        for name, relpath in RAW_RELPATHS.items():
            check_manifest_pin(
                relpath, manifest_sha256(manifest, relpath), RAW_FILES[name]
            )
            link_verified(
                raw_mirror_dir(data_root) / relpath,
                osp.join(self.raw_dir, name),
                RAW_FILES[name],
            )
        log.info(
            "Jackson 2020 raw files linked into %s (sha256 verified)", self.raw_dir
        )

    def _require_genome(self) -> SCerevisiaeGenome:
        if self.genome is None:
            raise RuntimeError(
                "TfDeletionPseudobulkJackson2020Dataset needs a genome to key the gene "
                "features; inject SCerevisiaeGenome(...)"
            )
        return self.genome

    @post_process
    def process(self) -> None:
        """Build the 784 (strain, condition) records and the 11 references."""
        verify_raw_files(self.raw_dir, RAW_FILES)
        genome = self._require_genome()
        os.makedirs(self.preprocess_dir, exist_ok=True)
        raw = Path(self.raw_dir)

        specs = read_strain_table(raw / SUPP1_NAME)
        fasta = read_barcode_fasta(raw / BARCODES_NAME)
        if fasta != {s.strain: s.barcode for s in specs}:
            raise StrainTableError(
                "Table S2 barcodes differ from bcdel_1_barcodes.fasta"
            )
        strains = {s.strain: s for s in specs}
        settings = read_author_settings(raw / SETTINGS_NAME)
        conditions = build_conditions(settings)
        gene_ids = self._deleted_gene_ids(specs, settings, genome)

        bulk_header = pd.read_csv(raw / BULK_NAME, sep="\t", nrows=0).columns.tolist()
        released = read_release_header(raw / MATRIX_NAME)
        if len(released) != N_RELEASED_FEATURES:
            raise ReleaseShapeError(
                f"{len(released)} count features, expected {N_RELEASED_FEATURES}"
            )
        features = build_feature_map(released, bulk_header, genome)
        sums, n_cells = read_release(
            raw / MATRIX_NAME, strains, set(settings.condition_label)
        )

        environments = {a: condition_environment(c) for a, c in conditions.items()}
        references: dict[str, PseudobulkUmiCountExperimentReference] = {}
        wild_type = [s.strain for s in specs if s.is_wild_type]
        for abbreviation, condition in conditions.items():
            entries = [
                sums[(strain, condition.release_label)]
                for strain in wild_type
                if (strain, condition.release_label) in sums
            ]
            if not entries:
                raise ReleaseShapeError(
                    f"{abbreviation}: no HO-deletion control cell, so the condition has "
                    "no reference"
                )
            references[abbreviation] = PseudobulkUmiCountExperimentReference(
                dataset_name=self.name,
                genome_reference=reference_genome(),
                environment_reference=environments[abbreviation],
                phenotype_reference=pseudobulk_phenotype(pooled(entries), features),
            )

        pub = publication()
        rows: list[dict[str, str]] = []
        empty: list[str] = []
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        idx = 0
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for abbreviation, condition in conditions.items():
                for spec in specs:
                    key = (spec.strain, condition.release_label)
                    if key not in sums:
                        empty.append(f"{spec.strain} x {abbreviation}")
                        continue
                    systematic, standard = gene_ids[spec.gene_token]
                    experiment = PseudobulkUmiCountExperiment(
                        dataset_name=self.name,
                        genotype=strain_genotype(spec, systematic, standard),
                        environment=environments[abbreviation],
                        phenotype=pseudobulk_phenotype(sums[key], features),
                    )
                    txn.put(
                        f"{idx}".encode(),
                        self._intern_record(
                            experiment, references[abbreviation], pub, itxn
                        ),
                    )
                    rows.append(
                        {
                            "strain": spec.strain,
                            "condition": abbreviation,
                            "release_label": condition.release_label,
                            "n_cells": str(len(sums[key].cell_barcodes)),
                        }
                    )
                    idx += 1
        env.close()
        interned_env.close()
        pd.DataFrame(rows).to_csv(
            osp.join(self.preprocess_dir, "data.csv"), index=False
        )
        self._write_ledger(
            sums, n_cells, features, specs, conditions, empty, idx, len(references)
        )
        log.info("Wrote %d Jackson 2020 records (%d empty pairs)", idx, len(empty))

    @staticmethod
    def _deleted_gene_ids(
        specs: list[StrainSpec], settings: AuthorSettings, genome: SCerevisiaeGenome
    ) -> dict[str, tuple[str, str]]:
        """``{gene token: (systematic, standard name)}`` for the twelve deleted genes.

        The eleven TF ORFs are the authors' (``figure_common_settings.R``); each must be
        a current R64 gene whose standard name is the token. HO is resolved through the
        genome by name.
        """
        table = genome.gene_attribute_table
        standard = dict(zip(table["ID"].astype(str), table["gene"].astype(str)))
        out: dict[str, tuple[str, str]] = {}
        for token in sorted({s.gene_token for s in specs}):
            if token == WT_GENE_TOKEN:
                resolution = genome.resolve_gene_name("HO")
                if (
                    resolution.status.value != "renamed"
                    or resolution.systematic_name is None
                ):
                    raise ValueError(
                        f"HO did not resolve to one current gene: {resolution}"
                    )
                systematic = resolution.systematic_name
            else:
                systematic = settings.tf_systematic_name[token]
                if genome.resolve_gene_name(systematic).status.value != "current":
                    raise ValueError(f"{token}: {systematic} is not a current R64 gene")
            name = standard[systematic]
            if name.upper() != token.upper():
                raise ValueError(f"{systematic}'s standard name {name} is not {token}")
            out[token] = (systematic, name)
        if set(out) - {WT_GENE_TOKEN} != set(settings.tf_systematic_name):
            raise ValueError("the strain table's genes are not the settings file's TFs")
        return out

    def _write_ledger(
        self,
        sums: dict[tuple[str, str], StrainConditionSum],
        n_cells: int,
        features: FeatureMap,
        specs: list[StrainSpec],
        conditions: dict[str, ConditionSpec],
        empty: list[str],
        n_records: int,
        n_references: int,
    ) -> None:
        label_to_abbrev = {c.release_label: a for a, c in conditions.items()}
        cells = Counter[str]()
        records = Counter[str]()
        for (strain, label), entry in sums.items():
            cells[label_to_abbrev[label]] += len(entry.cell_barcodes)
            records[label_to_abbrev[label]] += 1
        wild_type = {s.strain for s in specs if s.is_wild_type}
        total = int(sum(int(e.counts.sum()) for e in sums.values()))
        ledger = ReleaseLedger(
            matrix_sha256=MATRIX_SHA256,
            n_cells_released=n_cells,
            n_cells_served=sum(len(e.cell_barcodes) for e in sums.values()),
            n_wild_type_cells=sum(
                len(e.cell_barcodes) for (s, _), e in sums.items() if s in wild_type
            ),
            n_transcript_counts_released=total,
            n_transcript_counts_served=total,
            paper_cells=PAPER_CELLS,
            paper_transcript_counts=PAPER_TRANSCRIPT_COUNTS,
            n_features_released=len(features.released),
            n_gene_features=len(features.gene_columns),
            marker_features=list(MARKER_FEATURES),
            feature_status=dict(Counter(features.status.values())),
            remapped_features={
                features.source_ids[i]: features.keys[i]
                for i in features.gene_columns
                if features.keys[i] != features.source_ids[i]
            },
            legacy_features=features.legacy(),
            n_strains=len(specs),
            n_conditions=len(conditions),
            n_records=n_records,
            n_references=n_references,
            strain_conditions_without_cells=empty,
            cells_per_condition=dict(cells),
            records_per_condition=dict(records),
            table_s2_genotype_disagreements=list(TABLE_S2_GENOTYPE_DISAGREEMENTS),
            harvest_cells_per_ml=dict(HARVEST_CELLS_PER_ML),
            refused={},
        )
        with open(osp.join(self.preprocess_dir, "release_ledger.json"), "w") as handle:
            handle.write(ledger.model_dump_json(indent=2))

    def preprocess_raw(
        self, df: pd.DataFrame, preprocess: dict[str, Any] | None = None
    ) -> pd.DataFrame:
        """Preprocessing happens inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Records are assembled inside process()."""
        raise NotImplementedError(
            "TfDeletionPseudobulkJackson2020Dataset builds its records in process()"
        )


def verify_build(dataset_root: str, data_root: str | None = None) -> Any:
    """L0-L4 report for a built store (``build_dataset_lmdb --verify`` calls this)."""
    from torchcell.verification import runners

    return runners.verify_pseudobulk_umi_count_dataset(
        "tf_deletion_pseudobulk_jackson2020",
        data_root or _data_root(),
        abs_root=dataset_root,
    )


def main() -> None:
    """Build/load the dataset and print one record's shape."""
    from torchcell.datasets.scerevisiae.gene_name_reconcile import default_genome

    data_root = _data_root()
    dataset = TfDeletionPseudobulkJackson2020Dataset(
        root=osp.join(data_root, "data/torchcell/tf_deletion_pseudobulk_jackson2020"),
        genome=default_genome(),
    )
    print(f"len = {len(dataset)}")
    record = dataset[0]
    print(record["experiment"].phenotype)
    print(json.dumps(record["reference"].model_dump(mode="json"))[:400])


if __name__ == "__main__":
    main()
