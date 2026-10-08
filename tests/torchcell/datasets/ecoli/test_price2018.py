# tests/torchcell/datasets/ecoli/test_price2018.py
# [[tests.torchcell.datasets.ecoli.test_price2018]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_price2018.py
"""Price 2018 compendium, E. coli BW25113: sample rules, the SE, the ECK route, records.

The synthetic tests pin each builder with frames, crosswalks and files made here. The
``data`` tests (``--data``) re-derive the build's numbers from the sha256-pinned mirrors
under ``$DATA_ROOT``: every quote, the sample inventory by source paper, the
standard-error identity with the released t, the ECK route on the deposited K-12
annotations, the raw mirror, and (when the dev LMDB exists) its record count.
"""

from __future__ import annotations

import hashlib
import inspect
import json
import math
import os
import os.path as osp
import pickle
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import lmdb
import numpy as np
import openpyxl
import pandas as pd
import pytest

import torchcell.datasets.ecoli.price2018 as p
import torchcell.datasets.ecoli.wetmore2015 as w
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    BW25113_LOCI,
    MG1655_GAF,
    MG1655_LOCI,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.data import ManifestPinMismatchError, RawSha256MismatchError
from torchcell.datamodels.media import (
    LB_LENNOX,
    M9_NOCARBON_PRICE2018,
    M9_NONITROGEN_PRICE2018,
    MOPS_MINIMAL,
)
from torchcell.datamodels.schema import (
    AssayType,
    AssemblyReferenceGenome,
    BacterialGeneEssentialityExperiment,
    BacterialGeneEssentialityExperimentReference,
    ConcentrationUnit,
    EnvironmentPhysicalPerturbation,
    Genotype,
    MeasurementType,
    MediaComponentRole,
    PhysicalFactor,
    SmallMoleculePerturbation,
    TransposonInsertionPerturbation,
    UncertaintyType,
)
from torchcell.datasets.bacteria_common import (
    LocusTagReconciliation,
    LocusTagResolutionError,
    bacterial_genome,
    declared_reference_strain,
)
from torchcell.datasets.dataset_registry import dataset_registry
from torchcell.literature.manifest import Manifest
from torchcell.sequence.genome.base import GeneNameResolution, GeneNameStatus
from torchcell.sequence.genome.ecoli.k12 import (
    BW25113_ASSEMBLY,
    MG1655_ASSEMBLY,
    EckCrosswalk,
    EckPair,
    EcoliK12BW25113Genome,
    EcoliK12MG1655Genome,
)
from torchcell.verification.report import Level
from torchcell.verification.sourced import audit_sourced_value

DATA_ROOT = os.environ.get("DATA_ROOT", "")
MIRRORS_PRESENT = all(
    osp.isfile(osp.join(DATA_ROOT, rel, "manifest.json"))
    for rel in (
        p.RAW_DIR_REL,
        w.RAW_DIR_REL,
        f"{p.LIBRARY_DIR_REL}/{p.CITATION_KEY}",
        f"{w.LIBRARY_DIR_REL}/{w.CITATION_KEY}",
    )
)
needs_mirrors = pytest.mark.skipif(not MIRRORS_PRESENT, reason="requires the mirrors")


# --------------------------------------------------------------------------- #
# Synthetic Table S5 rows
# --------------------------------------------------------------------------- #
def _row(
    name: str,
    group: str,
    media: str,
    condition: str | None,
    concentration: float | None,
    units: str | None,
    temperature: float = 37.0,
    aerobic: str = "Aerobic",
) -> dict[str, Any]:
    return {
        "orgId": "Keio",
        "organism": "Escherichia coli BW25113",
        "name": name,
        "Group": group,
        "short": f"{condition} ({group})",
        "Media": media,
        "Condition_1": condition,
        "Concentration_1": concentration,
        "Units_1": units,
        "Temperature": temperature,
        "pH": None,
        "Aerobic_v_Anaerobic": aerobic,
        "Shaking": "200 rpm",
        "Growth.Method": "tube",
        "has_exact_replicate": "TRUE",
        "has_similar_replicate": "TRUE",
    }


ROWS = [
    _row(
        "set1IT003",
        "carbon source",
        "M9 minimal media_noCarbon",
        "D-Glucose",
        20.0,
        "mM",
    ),
    _row(
        "set1IT007", "carbon source", "M9 minimal media_noCarbon", "Sucrose", 20.0, "mM"
    ),
    _row(
        "set1IT067",
        "carbon source",
        "MOPS Rich Defined media_noCarbon",
        "D-Glucose",
        22.0,
        "mM",
    ),
    _row("set6IT068", "motility", "LB", "Agar", 3.0, "g/L", temperature=30.0),
    _row(
        "set2IT026",
        "stress",
        "LB",
        "Cephalothin sodium salt",
        0.01,
        "mg/ml",
        temperature=28.0,
    ),
    _row("set2IT045", "lb", "LB", None, None, None, temperature=28.0),
    _row(
        "set1IT071",
        "nitrogen source",
        "M9 minimal media_noNitrogen",
        "L-Arginine",
        10.0,
        "mM",
    ),
    _row(
        "set2IT053", "stress", "LB", "Dimethyl Sulfoxide", 7.5, "vol%", temperature=28.0
    ),
    _row(
        "set2IT094",
        "carbon source",
        "MOPS minimal media_noCarbon",
        "D-Glucose",
        22.0,
        "mM",
    ),
]
CARRIED = frozenset({"set1IT003", "set1IT007", "set1IT067", "set1IT071"})

#: Supplementary Table 4's ``Solvent`` keyed as :func:`read_solvents` keys it, over the
#: compounds these rows name plus the two non-water vehicles of the real sheet.
SOLVENTS = {
    "cephalothin sodium salt": "water",
    "dimethyl sulfoxide": "water",
    "chloramphenicol": "Ethanol",
    "vanillin": "Dimethyl Sulfoxide",
}


def _specs(rows: list[dict[str, Any]] | None = None) -> dict[str, p.SampleSpec]:
    table = pd.DataFrame(ROWS if rows is None else rows)
    return {s.name: s for s in p.classify_samples(table, CARRIED)}


# --------------------------------------------------------------------------- #
# Class contract and pins
# --------------------------------------------------------------------------- #
def test_the_loader_is_registered_against_bw25113() -> None:
    cls = p.RbTnseqPrice2018EcoliDataset
    assert dataset_registry["RbTnseqPrice2018EcoliDataset"] is cls
    assert declared_reference_strain(cls) == "BW25113"
    params = inspect.signature(cls.__init__).parameters
    assert "ecoli_genome" in params
    assert "genome" not in params
    assert params["root"].default == "data/torchcell/rbtnseq_price2018_ecoli"
    shell = cls.__new__(cls)
    assert shell.raw_file_names == [
        "fit_logratios_good.tab",
        "fit_standard_error_obs.tab",
        "fit_standard_error_naive.tab",
        "fit_t.tab",
        "si3.xlsx",
    ]


def test_the_pins_are_the_ones_the_wetmore_record_uses() -> None:
    assert p.PAPER_MD_SHA256 == w.SUPERSET_PAPER_MD_SHA256
    assert p.SUPPLEMENTARY_TABLES_SHA256 == w.SUPERSET_TABLE_S5_SHA256
    assert p.SUPPLEMENTARY_TABLES == w.SUPERSET_TABLE_S5
    assert p.TABLE_S5_SHEET == w.SUPERSET_TABLE_S5_SHEET
    assert p.ORG_ID == w.SUPERSET_ORG_ID
    assert p.REFERENCE_STRAIN == w.REFERENCE_STRAIN
    for relpath, ref in p.REFERENCED_RAW_FILES.items():
        assert ref.citation_key == w.CITATION_KEY
        assert ref.sha256 == w.RAW_FILES[relpath].sha256
    assert p.RAW_FILE_NAMES["fit_logratios_good.tab"] == (
        "wetmore",
        w.SUPERSET_FITNESS,
        "b37f038702ef0b792bdc02af9f42adf3a5838bc08513249c89983f6a7647e8fd",
    )


def test_expected_records_is_kept_samples_times_mapped_genes() -> None:
    assert (p.EXPECTED_SAMPLES, p.EXPECTED_GENES) == (147, 3768)
    assert p.EXPECTED_RECORDS == 553896


# --------------------------------------------------------------------------- #
# Samples
# --------------------------------------------------------------------------- #
def test_each_sample_gets_one_rule_and_its_source_paper() -> None:
    specs = _specs()
    assert {n: s.drop_rule for n, s in specs.items()} == {
        "set1IT003": None,
        "set1IT007": p.DROP_DISREGARDED,
        "set1IT067": p.DROP_MEDIUM_NOT_IN_LIBRARY,
        "set6IT068": p.DROP_MOTILITY,
        "set2IT026": None,
        "set2IT045": None,
        "set1IT071": None,
        "set2IT053": None,
        "set2IT094": None,
    }
    assert {n for n, s in specs.items() if s.source is p.SampleSource.wetmore2015} == (
        set(CARRIED)
    )
    assert specs["set2IT026"].source is p.SampleSource.price2018
    assert specs["set1IT003"].screen_id == "Keio:set1IT003"
    assert specs["set2IT045"].condition is None
    assert specs["set2IT045"].temperature_c == 28.0


def test_the_withdrawal_outranks_the_medium_rule() -> None:
    rows = [
        _row(
            "set1IT043",
            "carbon source",
            "MOPS Rich Defined media_noCarbon",
            "D-Mannitol",
            20.0,
            "mM",
        )
    ]
    table = pd.DataFrame(rows)
    (spec,) = p.classify_samples(table, frozenset())
    assert spec.drop_rule == p.DROP_DISREGARDED


@pytest.mark.parametrize(
    ("row", "message"),
    [
        (
            _row("s1", "carbon source", "R2A", "D-Glucose", 20.0, "mM"),
            "unknown medium 'R2A'",
        ),
        (_row("s2", "pH", "LB", None, None, None), "unknown group 'pH'"),
        (
            _row("s3", "stress", "M9 minimal media_noCarbon", "NaCl", 1.0, "mM"),
            "group 'stress' on medium 'M9 minimal media_noCarbon'",
        ),
        (
            _row("s4", "lb", "LB", None, None, None, aerobic="Anaerobic"),
            "not aerobic (Anaerobic)",
        ),
    ],
)
def test_a_row_outside_the_rules_is_refused(row: dict[str, Any], message: str) -> None:
    with pytest.raises(
        ValueError, match=message.replace("(", r"\(").replace(")", r"\)")
    ):
        p.classify_samples(pd.DataFrame([row]), frozenset())


def test_repeated_names_and_unknown_carried_samples_are_refused() -> None:
    with pytest.raises(ValueError, match="repeats a Keio sample name"):
        p.classify_samples(pd.DataFrame([ROWS[0], ROWS[0]]), frozenset())
    with pytest.raises(ValueError, match=r"absent from Table S5: \['set9IT001'\]"):
        p.classify_samples(pd.DataFrame(ROWS), frozenset({"set9IT001"}))


def test_a_carbon_source_is_the_varied_factor_of_a_carbon_free_medium() -> None:
    environment = p.build_environment(_specs()["set1IT003"], SOLVENTS)
    assert environment.media == M9_NOCARBON_PRICE2018
    assert environment.temperature is not None
    assert environment.temperature.value == 37.0
    assert environment.aerobicity == "aerobic"
    (perturbation,) = environment.perturbations
    assert isinstance(perturbation, EnvironmentPhysicalPerturbation)
    assert perturbation.factor is PhysicalFactor.carbon_source
    assert perturbation.magnitude is not None
    assert (perturbation.magnitude.value, perturbation.magnitude.unit) == (
        20.0,
        ConcentrationUnit.millimolar,
    )
    assert perturbation.agent is not None
    assert perturbation.agent.name == "D-glucose"
    assert perturbation.agent.inchikey == "WQZGKKKJIJFFOK-GASJEMHNSA-N"


def test_a_nitrogen_source_and_the_mops_medium_map_to_their_library_entries() -> None:
    specs = _specs()
    nitrogen = p.build_environment(specs["set1IT071"], SOLVENTS)
    assert nitrogen.media == M9_NONITROGEN_PRICE2018
    (perturbation,) = nitrogen.perturbations
    assert isinstance(perturbation, EnvironmentPhysicalPerturbation)
    assert perturbation.factor is PhysicalFactor.nitrogen_source
    assert p.build_environment(specs["set2IT094"], SOLVENTS).media == MOPS_MINIMAL


def test_a_stress_compound_keeps_its_printed_dose_and_table_s4s_vehicle() -> None:
    specs = _specs()
    environment = p.build_environment(specs["set2IT026"], SOLVENTS)
    assert environment.media == LB_LENNOX
    (perturbation,) = environment.perturbations
    assert isinstance(perturbation, SmallMoleculePerturbation)
    assert (perturbation.concentration.value, perturbation.concentration.unit) == (
        0.01,
        ConcentrationUnit.g_per_l,
    )
    assert perturbation.solvent is not None
    # water is not in the pinned identity table, so the vehicle carries the typed gap
    assert perturbation.solvent.name == "water"
    assert perturbation.solvent.percent is None
    assert perturbation.solvent.compound is not None
    assert [g.field for g in perturbation.solvent.compound.provenance_gaps] == [
        "inchikey"
    ]
    assert perturbation.provenance_gaps == []
    assert perturbation.description == p.STRESS_DESCRIPTION
    assert "wild-type IC50 prescreen" in perturbation.description
    assert perturbation.compound.name == "Cephalothin sodium salt"
    assert [g.field for g in perturbation.compound.provenance_gaps] == ["inchikey"]
    (dmso,) = p.build_environment(specs["set2IT053"], SOLVENTS).perturbations
    assert isinstance(dmso, SmallMoleculePerturbation)
    assert (dmso.concentration.value, dmso.concentration.unit) == (
        7.5,
        ConcentrationUnit.percent_v_v,
    )
    assert dmso.solvent is not None and dmso.solvent.name == "water"


def test_an_identified_vehicle_carries_its_structure() -> None:
    spec = _specs()["set2IT026"].model_copy(update={"condition": "Chloramphenicol"})
    solvent = p.stress_solvent(spec, SOLVENTS)
    assert solvent.name == "Ethanol"
    assert solvent.compound is not None
    assert solvent.compound.name == "ethanol"
    assert solvent.compound.inchikey == "LFQSCWFLJHTTHZ-UHFFFAOYSA-N"
    assert solvent.compound.provenance_gaps == []


def test_a_stress_compound_outside_table_s4_is_refused() -> None:
    spec = _specs()["set2IT026"].model_copy(update={"condition": "Unobtainium"})
    with pytest.raises(ValueError, match="is not a TableS4_Stress compound"):
        p.stress_solvent(spec, SOLVENTS)
    plain = _specs()["set2IT045"]
    with pytest.raises(ValueError, match="a stress sample names no condition"):
        p.stress_solvent(plain, SOLVENTS)


def test_plain_lb_carries_no_perturbation_and_drops_build_nothing() -> None:
    specs = _specs()
    plain = p.build_environment(specs["set2IT045"], SOLVENTS)
    assert plain.perturbations == []
    assert plain.media == LB_LENNOX
    with pytest.raises(ValueError, match=r"set1IT007 is dropped"):
        p.build_environment(specs["set1IT007"], SOLVENTS)
    with pytest.raises(ValueError, match="a plain sample names 'Agar'"):
        p.build_environment(
            specs["set2IT045"].model_copy(update={"condition": "Agar"}), SOLVENTS
        )
    with pytest.raises(ValueError, match="unknown unit 'uL'"):
        p.build_environment(
            specs["set2IT026"].model_copy(update={"units": "uL"}), SOLVENTS
        )


# --------------------------------------------------------------------------- #
# The release tables and the standard error
# --------------------------------------------------------------------------- #
def _tables(
    fit: list[list[float]], obs: list[list[float]], naive: list[list[float]]
) -> dict[str, pd.DataFrame]:
    genes = {"locusId": [11, 12, 13], "sysName": ["b0001", "b0002", "b0003"]}
    headers = ["set1IT003 D-Glucose (C)", "set1IT004 D-Glucose (C)"]
    f, o, n = np.array(fit), np.array(obs), np.array(naive)
    t = f / np.sqrt(0.1**2 + np.maximum(o, n) ** 2)

    def frame(values: np.ndarray[Any, Any], extra: dict[str, Any]) -> pd.DataFrame:
        return pd.DataFrame({**genes, **extra, **dict(zip(headers, values.T))})

    return {
        "fitness": frame(f, {"desc": ["x", "y", "z"], "comb": ["a", "b", "c"]}),
        "se_obs": frame(o, {"desc": ["x", "y", "z"]}),
        "se_naive": frame(n, {"desc": ["x", "y", "z"]}),
        "t": frame(t, {"desc": ["x", "y", "z"]}),
    }


FIT = [[-1.0, -2.0], [0.5, 0.25], [3.0, -0.125]]
OBS = [[0.2, 0.4], [0.1, 0.3], [0.5, 0.6]]
NAIVE = [[0.3, 0.1], [0.1, 0.2], [0.7, 0.05]]


def test_the_standard_error_is_the_larger_estimate_and_reproduces_t() -> None:
    tables = _tables(FIT, OBS, NAIVE)
    release = p.align_release(**tables, samples=["set1IT004", "set1IT003"], sigma=0.1)
    assert release.samples == ("set1IT004", "set1IT003")
    assert release.b_numbers == ("b0001", "b0002", "b0003")
    assert release.locus_ids == (11, 12, 13)
    np.testing.assert_array_equal(
        release.fitness, [[-2.0, -1.0], [0.25, 0.5], [-0.125, 3.0]]
    )
    np.testing.assert_array_equal(
        release.standard_error, [[0.4, 0.3], [0.3, 0.1], [0.6, 0.7]]
    )
    assert release.n_values == 6
    assert release.n_naive_larger == 2
    assert release.t_max_abs_deviation < 1e-12


def test_a_t_that_the_larger_estimate_does_not_reproduce_is_refused() -> None:
    tables = _tables(FIT, OBS, NAIVE)
    tables["t"].iloc[1, 3] = 99.0
    with pytest.raises(ValueError, match="the released t is not fit / sqrt"):
        p.align_release(**tables, samples=["set1IT003", "set1IT004"], sigma=0.1)
    with pytest.raises(ValueError, match="the released t is not fit / sqrt"):
        p.align_release(
            **_tables(FIT, OBS, NAIVE), samples=["set1IT003", "set1IT004"], sigma=0.2
        )


def test_misaligned_or_incomplete_tables_are_refused() -> None:
    samples = ["set1IT003", "set1IT004"]
    swapped = _tables(FIT, OBS, NAIVE)
    swapped["se_obs"].loc[0, "sysName"] = "b0009"
    with pytest.raises(ValueError, match="se_obs: gene rows differ"):
        p.align_release(**swapped, samples=samples, sigma=0.1)
    with pytest.raises(ValueError, match=r"only in fitness \['set1IT004'\]"):
        p.align_release(**_tables(FIT, OBS, NAIVE), samples=["set1IT003"], sigma=0.1)
    renamed = _tables(FIT, OBS, NAIVE)
    renamed["t"] = renamed["t"].rename(
        columns={"set1IT004 D-Glucose (C)": "set1IT004 Glucose"}
    )
    with pytest.raises(ValueError, match=r"t: header text differs for \['set1IT004'\]"):
        p.align_release(**renamed, samples=samples, sigma=0.1)
    holed = _tables(FIT, OBS, NAIVE)
    holed["fitness"].iloc[2, 4] = math.nan
    with pytest.raises(ValueError, match="fitness: 1 missing values"):
        p.align_release(**holed, samples=samples, sigma=0.1)
    zero = _tables(FIT, OBS, NAIVE)
    zero["se_naive"].iloc[0, 3] = 0.0
    with pytest.raises(ValueError, match="not positive"):
        p.align_release(**zero, samples=samples, sigma=0.1)


def test_a_sample_with_two_columns_is_refused() -> None:
    frame = pd.DataFrame({"locusId": [1], "set1IT003 a": [0.0], "set1IT003 b": [1.0]})
    with pytest.raises(ValueError, match="fitness: sample set1IT003 has two columns"):
        p.sample_columns(frame, "fitness")


# --------------------------------------------------------------------------- #
# The ECK route
# --------------------------------------------------------------------------- #
CROSSWALK = EckCrosswalk(
    shared=("ECK0001", "ECK0002", "ECK0018", "ECK0261"),
    pairs=(
        EckPair(eck="ECK0001", mg1655="b0001", bw25113="BW25113_0001"),
        EckPair(eck="ECK0002", mg1655="b0002", bw25113="BW25113_0002"),
        EckPair(eck="ECK0018", mg1655="b0018", bw25113="BW25113_4412"),
    ),
    mg1655_only=("ECK1159",),
    bw25113_only=(),
)
MG1655_ECKS = {
    "b0001": ("ECK0001",),
    "b0002": ("ECK0002",),
    "b0018": ("ECK0018",),
    "b1172": ("ECK1159",),
    "b1370": ("ECK0261",),
    "b9001": (),
}


def test_the_eck_route_places_one_to_one_pairs_and_names_every_miss() -> None:
    placed, unplaced = p.eck_route(
        ["b0001", "b0018", "b1172", "b1370", "b0500", "b9001", "b0002"],
        CROSSWALK,
        MG1655_ECKS,
    )
    assert placed == {
        "b0001": ("ECK0001", "BW25113_0001"),
        "b0018": ("ECK0018", "BW25113_4412"),
        "b0002": ("ECK0002", "BW25113_0002"),
    }
    assert [(u.b_number, u.reason, u.eck) for u in unplaced] == [
        ("b1172", "eck_absent_from_bw25113", ("ECK1159",)),
        ("b1370", "eck_not_one_to_one", ("ECK0261",)),
        ("b0500", "b_number_not_in_mg1655_annotation", ()),
        ("b9001", "mg1655_locus_has_no_eck_synonym", ()),
    ]


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
def _mapping(b: str, tag: str, name: str) -> p.GeneMapping:
    return p.GeneMapping(
        b_number=b,
        eck=f"ECK{b[1:]}",
        locus_tag=tag,
        perturbed_gene_name=name,
        numerics_agree=b[1:] == tag.removeprefix("BW25113_"),
    )


def test_the_genotype_is_a_derived_gene_level_tn5_insertion() -> None:
    genotype = p.build_genotype(_mapping("b0002", "BW25113_0002", "thrA"))
    (perturbation,) = genotype.perturbations
    assert isinstance(perturbation, TransposonInsertionPerturbation)
    assert perturbation.systematic_gene_name == "BW25113_0002"
    assert perturbation.perturbed_gene_name == "thrA"
    assert perturbation.gene_namespace == "ecoli_k12_bw25113_locus_tag"
    assert (perturbation.transposon, perturbation.library_pool) == ("Tn5", "KEIO_ML9")
    assert (perturbation.barcode, perturbation.insertion_position) == (None, None)
    assert "Locus tag DERIVED" in perturbation.description
    assert "one-to-one ECK pair" in perturbation.description


def test_the_phenotype_stores_the_standard_error_and_gaps_the_strain_count() -> None:
    phenotype = p.build_phenotype(-1.5, 0.3, "Keio:set1IT003")
    assert phenotype.measurement_type is MeasurementType.log2_ratio
    assert phenotype.assay_type is AssayType.pooled_competitive_growth_barcode
    assert phenotype.environment_response == -1.5
    assert phenotype.environment_response_uncertainty == 0.3
    assert phenotype.environment_response_uncertainty_type is (
        UncertaintyType.standard_error
    )
    assert phenotype.environment_response_se == 0.3
    assert (phenotype.n_samples, phenotype.sample_unit) == (None, None)
    assert [g.field for g in phenotype.provenance_gaps] == ["n_samples", "sample_unit"]
    assert phenotype.screen_id == "Keio:set1IT003"
    assert phenotype.units == p.UNITS


def test_records_run_sample_by_sample_with_each_samples_publication() -> None:
    specs = _specs()
    kept = [specs["set1IT003"], specs["set2IT026"]]
    release = p.align_release(
        **_tables(FIT, OBS, NAIVE), samples=["set1IT003", "set1IT004"], sigma=0.1
    )
    # set2IT026 is read from the second release column in this synthetic release.
    columns = {"set1IT003": 0, "set2IT026": 1}
    reference_genome = AssemblyReferenceGenome(
        species="Escherichia coli",
        strain="BW25113",
        assembly_set="ecoli_K12_BW25113_ASM75055v1",
        assembly_accession="GCA_000750555.1",
    )
    samples = [
        p.SampleRecords(
            spec=spec,
            column=columns[spec.name],
            environment=p.build_environment(spec, SOLVENTS),
            reference=p.build_reference(
                "D",
                p.build_environment(spec, SOLVENTS),
                reference_genome,
                spec.screen_id,
            ),
            publication=p.PUBLICATIONS[spec.source],
        )
        for spec in kept
    ]
    mapping = {
        "b0001": _mapping("b0001", "BW25113_0001", "thrL"),
        "b0003": _mapping("b0003", "BW25113_0003", "thrB"),
    }
    records = list(p.iter_records("D", samples, release, mapping))
    seen = []
    for experiment, _, publication in records:
        assert isinstance(experiment.genotype, Genotype)
        seen.append(
            (
                experiment.genotype.perturbations[0].systematic_gene_name,
                experiment.phenotype.environment_response,
                experiment.phenotype.environment_response_se,
                experiment.phenotype.screen_id,
                publication.doi,
            )
        )
    assert seen == [
        ("BW25113_0001", -1.0, 0.3, "Keio:set1IT003", "10.1128/mBio.00306-15"),
        ("BW25113_0003", 3.0, 0.7, "Keio:set1IT003", "10.1128/mBio.00306-15"),
        ("BW25113_0001", -2.0, 0.4, "Keio:set2IT026", "10.1038/s41586-018-0124-0"),
        ("BW25113_0003", -0.125, 0.6, "Keio:set2IT026", "10.1038/s41586-018-0124-0"),
    ]
    reference = records[0][1]
    assert reference.phenotype_reference.environment_response == 0.0
    assert reference.phenotype_reference.screen_id == "Keio:set1IT003"
    assert reference.environment_reference == records[0][0].environment
    assert reference.genome_reference.assembly_set == "ecoli_K12_BW25113_ASM75055v1"


def test_the_inventory_counts_samples_genes_and_records_by_rule() -> None:
    specs = tuple(_specs().values())
    identifiers = p.IdentifierReport(
        n_source_genes=10,
        n_mapped=8,
        mapped_fraction=0.8,
        min_fraction=0.99,
        numeric_disagreements=(),
        unmapped=(
            p.UnmappedGene(
                b_number="b0500", reason="b_number_not_in_mg1655_annotation", eck=()
            ),
            p.UnmappedGene(
                b_number="b1172", reason="eck_absent_from_bw25113", eck=("ECK1159",)
            ),
        ),
        unmapped_by_reason={},
        reconcile_status_histogram={},
        reconcile_layer_histogram={},
        n_symbol_names=8,
        n_tag_names=0,
    )
    counts = p.inventory(specs, identifiers, SOLVENTS)
    assert counts.table_s5_samples == 9
    assert counts.samples_by_source == {
        p.SampleSource.wetmore2015.value: 4,
        p.SampleSource.price2018.value: 5,
    }
    assert counts.kept_samples_by_source == {
        p.SampleSource.wetmore2015.value: 2,
        p.SampleSource.price2018.value: 4,
    }
    assert (counts.source_records, counts.kept_records) == (90, 48)
    assert [(d.rule, d.items, d.n_records) for d in counts.drops] == [
        (p.DROP_DISREGARDED, ("set1IT007",), 10),
        (p.DROP_MEDIUM_NOT_IN_LIBRARY, ("set1IT067",), 10),
        (p.DROP_MOTILITY, ("set6IT068",), 10),
        (p.DROP_NO_ECK_PAIR, ("b0500", "b1172"), 12),
    ]
    assert counts.source_records == counts.kept_records + sum(
        d.n_records for d in counts.drops
    )
    assert counts.unidentified_compounds == {"Cephalothin sodium salt": 1}
    # two kept stress samples, both on a water stock (Table S4), so both counts are the
    # same here; on the real sheet they differ (45 / 7 / 3 samples, 29 / 5 / 1 compounds)
    assert counts.stress_solvent_by_sample == {"water": 2}
    assert counts.stress_solvent_by_compound == {"water": 2}


def _resolution(
    name: str, status: GeneNameStatus, systematic: str
) -> GeneNameResolution:
    return GeneNameResolution(
        input_name=name, status=status, systematic_name=systematic
    )


def test_the_supplementary_row_accepts_genes_and_pseudogenes_that_name_themselves() -> (
    None
):
    answers = {
        "BW25113_0001": _resolution(
            "BW25113_0001", GeneNameStatus.CURRENT, "BW25113_0001"
        ),
        "BW25113_0123": _resolution(
            "BW25113_0123", GeneNameStatus.NON_GENE_FEATURE, "BW25113_0123"
        ),
        "BW25113_0999": _resolution(
            "BW25113_0999", GeneNameStatus.RENAMED, "BW25113_0998"
        ),
    }
    passing = p.stored_tags_are_loci(
        {"BW25113_0001", "BW25113_0123"}, answers.__getitem__
    )
    assert passing.level is Level.L1
    assert passing.passed
    assert passing.details["statuses"] == {"current": 1, "non_gene_feature": 1}
    failing = p.stored_tags_are_loci(set(answers), answers.__getitem__)
    assert not failing.passed
    assert failing.details["not_a_locus"] == ["BW25113_0999"]


# --------------------------------------------------------------------------- #
# Raw mirror and the xlsx rendering
# --------------------------------------------------------------------------- #
def _fake_raw(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    source = tmp_path / "staging"
    files = {}
    for i, rel in enumerate(("data/a.tab", "data/sub/b.tab")):
        path = source / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        payload = f"locusId\tset1IT00{i}\n1\t0.5\n".encode()
        path.write_bytes(payload)
        files[rel] = p.RawFile(
            relpath=rel,
            url=f"https://example.org/{rel}",
            sha256=hashlib.sha256(payload).hexdigest(),
            purpose="test",
        )
    monkeypatch.setattr(p, "RAW_FILES", files)
    return source


def test_deposit_is_idempotent_and_refuses_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _fake_raw(tmp_path, monkeypatch)
    data_root = str(tmp_path / "root")
    root = p.deposit_raw_mirror(source_dir=source, data_root=data_root)
    assert root == Path(data_root) / p.RAW_DIR_REL
    manifest = Manifest.model_validate_json((root / "manifest.json").read_text())
    assert [f.path for f in manifest.files] == ["data/a.tab", "data/sub/b.tab"]
    assert {f.retrieval.retriever for f in manifest.files if f.retrieval} == {
        "torchcell.literature.retrieve.direct_url"
    }
    assert "wetmoreRapidQuantificationMutant2015 mirror" in manifest.si_expected[0]
    first = (root / "manifest.json").read_bytes()
    p.deposit_raw_mirror(source_dir=source, data_root=data_root)
    assert (root / "manifest.json").read_bytes() == first
    (root / "data/a.tab").write_bytes(b"changed")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        p.deposit_raw_mirror(source_dir=source, data_root=data_root)
    (source / "data/sub/b.tab").write_bytes(b"drifted")
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        p.deposit_raw_mirror(source_dir=source, data_root=data_root)


def test_a_manifest_recording_other_files_is_never_overwritten(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _fake_raw(tmp_path, monkeypatch)
    data_root = str(tmp_path / "root")
    root = p.deposit_raw_mirror(source_dir=source, data_root=data_root)
    stale = root / "manifest.json"
    payload = json.loads(stale.read_text())
    payload["files"] = payload["files"][:1]
    stale.write_text(json.dumps(payload))
    with pytest.raises(RuntimeError, match="records other files or retrievals"):
        p.deposit_raw_mirror(source_dir=source, data_root=data_root)
    assert json.loads(stale.read_text())["files"] == payload["files"]


def test_xlsx_text_joins_cells_and_rows(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Strain", None, "Transposon"])
    sheet.append([None, None, None])
    sheet.append(["Escherichia coli BW25113", "KEIO_ML9", 12])
    second = workbook.create_sheet("other")
    second.append(["x"])
    path = tmp_path / "t.xlsx"
    workbook.save(path)
    assert p.xlsx_text(str(path)) == (
        "Strain | Transposon / Escherichia coli BW25113 | KEIO_ML9 | 12\nx"
    )


# --------------------------------------------------------------------------- #
# Mirror paths and the recorded retrieval
# --------------------------------------------------------------------------- #
def test_source_path_checks_each_mirror_through_its_own_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _fake_raw(tmp_path, monkeypatch)
    data_root = str(tmp_path / "root")
    p.deposit_raw_mirror(source_dir=source, data_root=data_root)
    pin = p.RAW_FILES["data/a.tab"].sha256
    calls: list[tuple[str, ...]] = []

    def wetmore_raw(relpath: str, root: str) -> Path:
        calls.append(("wetmore", relpath, root))
        return Path("/w") / relpath

    def library(key: str, relpath: str, pin: str, root: str) -> Path:
        calls.append(("library", key, relpath, pin, root))
        return Path("/l") / relpath

    monkeypatch.setattr(w, "raw_path", wetmore_raw)
    monkeypatch.setattr(w, "library_path", library)
    monkeypatch.setattr(
        p,
        "RAW_FILE_NAMES",
        {
            "a.tab": ("price", "data/a.tab", pin),
            "w.tab": ("wetmore", "data/bigfit/w.tab", "0" * 64),
            "s.xlsx": ("library", "si/s.xlsx", "1" * 64),
            "stale.tab": ("price", "data/a.tab", "2" * 64),
        },
    )
    assert p.source_path("a.tab", data_root) == (
        Path(data_root) / p.RAW_DIR_REL / "data/a.tab"
    )
    assert p.source_path("w.tab", data_root) == Path("/w/data/bigfit/w.tab")
    assert p.source_path("s.xlsx", data_root) == Path("/l/si/s.xlsx")
    assert calls == [
        ("wetmore", "data/bigfit/w.tab", data_root),
        ("library", p.CITATION_KEY, "si/s.xlsx", "1" * 64, data_root),
    ]
    with pytest.raises(ManifestPinMismatchError):
        p.source_path("stale.tab", data_root)
    assert p.manifest_sha256(p.load_manifest(data_root), "data/a.tab") == pin
    with pytest.raises(KeyError, match="data/z.tab is not in the"):
        p.manifest_sha256(p.load_manifest(data_root), "data/z.tab")


def test_the_recorded_retrieval_writes_only_pinned_bytes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _fake_raw(tmp_path, monkeypatch)
    served = {raw.url: (source / rel).read_bytes() for rel, raw in p.RAW_FILES.items()}
    monkeypatch.setattr(p, "direct_url", served.__getitem__)
    dest = p.retrieve_raw_files(tmp_path / "fetched")
    assert sorted(str(f.relative_to(dest)) for f in dest.rglob("*.tab")) == [
        "data/a.tab",
        "data/sub/b.tab",
    ]
    served[p.RAW_FILES["data/a.tab"].url] = b"upstream changed"
    with pytest.raises(RawSha256MismatchError):
        p.retrieve_raw_files(tmp_path / "again")


def test_read_release_reads_the_four_linked_tables(tmp_path: Path) -> None:
    for label, name in (
        ("fitness", "fit_logratios_good.tab"),
        ("se_obs", "fit_standard_error_obs.tab"),
        ("se_naive", "fit_standard_error_naive.tab"),
        ("t", "fit_t.tab"),
    ):
        _tables(FIT, OBS, NAIVE)[label].to_csv(tmp_path / name, sep="\t", index=False)
    release = p.read_release(str(tmp_path), ["set1IT003", "set1IT004"])
    assert release.b_numbers == ("b0001", "b0002", "b0003")
    assert release.n_naive_larger == 2


# --------------------------------------------------------------------------- #
# The synthetic K-12 pair: the ECK route, the build, the verifier, the report
# --------------------------------------------------------------------------- #
@pytest.fixture
def k12(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[EcoliK12MG1655Genome, EcoliK12BW25113Genome]:
    """Both synthetic K-12 genomes; the network refuses."""
    files = write_assembly(tmp_path / "tier", MG1655_ASSEMBLY, MG1655_LOCI, MG1655_GAF)
    files |= write_assembly(tmp_path / "tier", BW25113_ASSEMBLY, BW25113_LOCI)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, files)
    return (
        EcoliK12MG1655Genome(genome_root=str(tmp_path / "mg1655"), overwrite=True),
        EcoliK12BW25113Genome(genome_root=str(tmp_path / "bw25113"), overwrite=True),
    )


#: The synthetic release's genes: four map (b0003 crosses numbers, b0004 is a
#: pseudogene), b0005 is not one-to-one, b0006 is absent from BW25113, b0099 unknown.
SYNTHETIC_GENES = ["b0001", "b0002", "b0003", "b0004", "b0005", "b0006", "b0099"]


def test_map_genes_on_the_synthetic_k12_pair(
    k12: tuple[EcoliK12MG1655Genome, EcoliK12BW25113Genome],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    mg1655, bw25113 = k12
    with pytest.raises(LocusTagResolutionError, match="places 4 of 7 genes"):
        p.map_genes(SYNTHETIC_GENES, mg1655, bw25113, label="synthetic")
    monkeypatch.setattr(p, "MIN_ECK_ROUTE_FRACTION", 0.5)
    mapping, report = p.map_genes(SYNTHETIC_GENES, mg1655, bw25113, label="synthetic")
    assert {b: (m.locus_tag, m.perturbed_gene_name) for b, m in mapping.items()} == {
        "b0001": ("BW25113_0001", "thrL"),
        "b0002": ("BW25113_0002", "thrA"),
        "b0003": ("BW25113_4412", "hokC"),
        "b0004": ("BW25113_0004", "BW25113_0004"),
    }
    assert report.numeric_disagreements == (("b0003", "BW25113_4412", "ECK0003"),)
    assert [(u.b_number, u.reason, u.eck) for u in report.unmapped] == [
        ("b0005", "eck_not_one_to_one", ("ECK0005",)),
        ("b0006", "eck_absent_from_bw25113", ("ECK0006",)),
        ("b0099", "b_number_not_in_mg1655_annotation", ()),
    ]
    assert report.reconcile_layer_histogram["gene synonym"] == 4
    assert (report.n_symbol_names, report.n_tag_names) == (3, 1)


def test_the_mapping_refuses_a_reconciler_that_disagrees_with_the_crosswalk() -> None:
    def reconcile(ecks: pd.Series) -> tuple[pd.Series, LocusTagReconciliation]:
        return (
            pd.Series(["BW25113_0001", "BW25113_9999"]),
            LocusTagReconciliation(
                label="x",
                assembly_set="ecoli_K12_BW25113_ASM75055v1",
                gene_namespace="ecoli_k12_bw25113_locus_tag",
                unique_names=2,
                status_histogram={s: 0 for s in GeneNameStatus},
                layer_histogram={},
                remapped=2,
                kept_on_collision=(),
                retired_kept=(),
                ambiguous_kept={},
                case_insensitive=(),
                outside_namespace=(),
            ),
        )

    with pytest.raises(ValueError, match=r"disagree on \['b0002'\]"):
        p.assemble_mapping(
            ["b0001", "b0002"],
            CROSSWALK,
            MG1655_ECKS,
            reconcile,
            lambda tag: tag,
            label="x",
        )


def test_gene_symbol_is_used_only_when_it_names_the_locus_back() -> None:
    answers = {
        "thrA": _resolution("thrA", GeneNameStatus.RENAMED, "BW25113_0002"),
        "dup": _resolution("dup", GeneNameStatus.RENAMED, "BW25113_0003"),
        "yaaP": _resolution("yaaP", GeneNameStatus.NON_GENE_FEATURE, "BW25113_0004"),
    }
    assert p.gene_symbol("thrA", "BW25113_0002", answers.__getitem__) == "thrA"
    assert p.gene_symbol("dup", "BW25113_0002", answers.__getitem__) == "BW25113_0002"
    assert p.gene_symbol("yaaP", "BW25113_0004", answers.__getitem__) == "BW25113_0004"
    assert p.gene_symbol(None, "BW25113_0009", answers.__getitem__) == "BW25113_0009"


def _release_frames(samples: list[str]) -> dict[str, pd.DataFrame]:
    """A release over SYNTHETIC_GENES x ``samples`` whose t follows the rule exactly."""
    n_genes = len(SYNTHETIC_GENES)
    grid = np.arange(n_genes * len(samples), dtype=np.float64).reshape(n_genes, -1)
    fit = grid / 8.0 - 2.0
    obs = 0.1 + grid / 100.0
    naive = 0.3 - grid / 200.0
    t = fit / np.sqrt(0.1**2 + np.maximum(obs, naive) ** 2)
    genes = {"locusId": list(range(100, 100 + n_genes)), "sysName": SYNTHETIC_GENES}
    headers = [f"{s} condition" for s in samples]

    def frame(values: np.ndarray[Any, Any]) -> pd.DataFrame:
        return pd.DataFrame({**genes, **dict(zip(headers, values.T, strict=True))})

    return {
        "fit_logratios_good.tab": frame(fit),
        "fit_standard_error_obs.tab": frame(obs),
        "fit_standard_error_naive.tab": frame(naive),
        "fit_t.tab": frame(t),
    }


#: set1IT003 (Wetmore, glucose), set2IT026 (Price, stress), set1IT007 (Wetmore,
#: withdrawn), set6IT068 (Price, motility): two kept samples.
SYNTHETIC_TABLE = pd.DataFrame([ROWS[0], ROWS[4], ROWS[1], ROWS[3]])
SYNTHETIC_SAMPLES = ["set1IT003", "set2IT026", "set1IT007", "set6IT068"]


def _pin() -> AssemblyReferenceGenome:
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain="BW25113",
        assembly_set="ecoli_K12_BW25113_ASM75055v1",
        assembly_accession="GCA_000750555.1",
    )


#: Supplementary Table 1 Keio rows for the synthetic workbook: b0001 and b0002 map, b0003
#: maps but names a RefSeq tag the annotation does not carry, b0006 and b0007 have no
#: BW25113 locus (ECK0006 / ECK0007 are absent) and so carry no locus_tag, exactly as the
#: real sheet leaves the deleted araBAD and rhaBAD rows blank.
SYNTHETIC_ESSENTIAL_ROWS = [
    ("b0001", "BW25113_RS00005", "thrL", "thr operon leader peptide", "Arole"),
    ("b0002", "BW25113_RS00010", "thrA", "aspartokinase I", "Bspecific"),
    ("b0003", "BW25113_RS09999", "thrW", "tRNA-Thr", "Cvague"),
    ("b0006", None, "araC", "L-arabinose transcriptional regulator", "Dhypo"),
    ("b0007", None, "rhaD", "rhamnulose-1-phosphate aldolase", "Dhypo"),
]
#: The synthetic ``Solvent`` column; the stress samples of ``ROWS`` name the first two.
SYNTHETIC_SOLVENT_ROWS = [
    ("Cephalothin sodium salt", "water"),
    ("Dimethyl Sulfoxide", "water"),
    ("Chloramphenicol", "Ethanol"),
    ("Vanillin", "Dimethyl Sulfoxide"),
]


def _write_workbook(path: Path) -> Path:
    """A workbook with real Table S1 and Table S4 sheets, each under a preamble.

    Both sheets are read by the header-finding reader, so the preamble rows above the
    header are part of what is under test; Table S1 also carries a non-Keio organism, to
    pin the orgId filter.
    """
    workbook = openpyxl.Workbook()
    s1 = workbook.active
    assert s1 is not None
    s1.title = p.TABLE_S1_SHEET
    s1.append(["This table lists all of the likely-essential protein-coding genes"])
    s1.append(["organism -- which bacterium the gene is from"])
    s1.append([None])
    s1.append(
        [
            "organism",
            "orgId",
            "locusId",
            "sysName",
            "locus_tag",
            "protein_id",
            "uniprotId",
            "scaffoldId",
            "begin",
            "end",
            "strand",
            "name",
            "desc",
            "GC",
            "nReads",
            "normreads",
            "nPosCentral",
            "dens",
            "geneClass",
        ]
    )
    for i, (b, tag, name, desc, gene_class) in enumerate(SYNTHETIC_ESSENTIAL_ROWS):
        s1.append(
            [
                "Escherichia coli BW25113",
                "Keio",
                14000 + i,
                b,
                tag,
                "WP_000000000.1",
                f"sp|P0000{i}|X_ECOLI",
                7023,
                100 + i,
                200 + i,
                "+",
                name,
                desc,
                0.5,
                i,
                0.01 * i,
                0,
                0.0,
                gene_class,
            ]
        )
    s1.append(
        [
            "Shewanella oneidensis MR-1",
            "MR1",
            9001,
            "SO0001",
            "SO_RS00005",
            "WP_111111111.1",
            "sp|Q00001|Y_SHEON",
            7000,
            1,
            2,
            "+",
            "soA",
            "a gene of another organism",
            0.5,
            0,
            0.0,
            0,
            0.0,
            "Arole",
        ]
    )
    s4 = workbook.create_sheet(p.TABLE_S4_SHEET)
    s4.append(["The values reported are the half-maximum inhibitory concentrations"])
    s4.append(["Microbe", "Media used for stress experiments"])
    s4.append(["Escherichia coli BW25113", "LB"])
    s4.append([None])
    s4.append(
        [
            "Compound",
            "CAS",
            "CoreSet_forMutantFitnessAssays",
            "Stock solution",
            "Stock solution units",
            "Solvent",
            "Maximum concentration tested",
            "Minimum concentration tested",
            "Escherichia coli BW25113",
        ]
    )
    for compound, solvent in SYNTHETIC_SOLVENT_ROWS:
        s4.append([compound, "1-00-0", "no", 10, "mg/ml", solvent, 1, 0.001, 0.5])
    workbook.save(path)
    return path


@pytest.fixture
def mirrored(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    k12: tuple[EcoliK12MG1655Genome, EcoliK12BW25113Genome],
) -> Path:
    """A tmp ``DATA_ROOT`` whose mirrors hold the synthetic release; returns it."""
    source = tmp_path / "source"
    source.mkdir()
    for name, frame in _release_frames(SYNTHETIC_SAMPLES).items():
        frame.to_csv(source / name, sep="\t", index=False)
    _write_workbook(source / "si3.xlsx")
    monkeypatch.setattr(
        p,
        "RAW_FILE_NAMES",
        {
            name: (
                "price",
                f"data/{name}",
                hashlib.sha256(path.read_bytes()).hexdigest(),
            )
            for name, path in ((n, source / n) for n in sorted(os.listdir(source)))
        },
    )
    monkeypatch.setattr(p, "source_path", lambda name, data_root=None: source / name)
    monkeypatch.setattr(w, "read_superset_experiments", lambda path: SYNTHETIC_TABLE)
    monkeypatch.setattr(
        w,
        "subsumption_record",
        lambda data_root=None: SimpleNamespace(
            carried=("set1IT003", "set1IT007"), disregarded=("set1IT007",)
        ),
    )
    genomes = dict(zip(("MG1655", "BW25113"), k12, strict=True))
    monkeypatch.setattr(
        p, "bacterial_genome", lambda host, strain, data_root=None: genomes[strain]
    )
    monkeypatch.setattr(p, "assembly_reference", lambda strain: _pin())
    monkeypatch.setattr(p, "MIN_ECK_ROUTE_FRACTION", 0.5)
    monkeypatch.setattr(p, "load_dotenv", lambda: None)
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.delenv("TC_DATA_URL", raising=False)
    return data_root


def _dataset_root(data_root: Path) -> Path:
    return data_root / "data/torchcell/rbtnseq_price2018_ecoli"


def test_the_loader_builds_the_synthetic_release_end_to_end(mirrored: Path) -> None:
    root = _dataset_root(mirrored)
    dataset = p.RbTnseqPrice2018EcoliDataset(root=str(root))
    assert len(dataset) == 8
    assert sorted(dataset.gene_set) == [
        "BW25113_0001",
        "BW25113_0002",
        "BW25113_0004",
        "BW25113_4412",
    ]
    references = dataset.experiment_reference_index
    assert references is not None
    assert len(references) == 2
    first = dataset[0]
    phenotype = first["experiment"]["phenotype"]
    assert phenotype["screen_id"] == "Keio:set1IT003"
    assert phenotype["environment_response"] == -2.0
    assert phenotype["environment_response_se"] == pytest.approx(0.3)
    assert first["experiment"]["genotype"]["perturbations"][0][
        "systematic_gene_name"
    ] == ("BW25113_0001")
    assert first["publication"]["doi"] == "10.1128/mBio.00306-15"
    drops = json.loads((root / "preprocess" / "dropped_records.json").read_text())
    assert (drops["source_records"], drops["kept_records"]) == (28, 8)
    assert [(d["rule"], d["items"], d["n_records"]) for d in drops["drops"]] == [
        (p.DROP_DISREGARDED, ["set1IT007"], 7),
        (p.DROP_MEDIUM_NOT_IN_LIBRARY, [], 0),
        (p.DROP_MOTILITY, ["set6IT068"], 7),
        (p.DROP_NO_ECK_PAIR, ["b0005", "b0006", "b0099"], 6),
    ]
    identifiers = json.loads(
        (root / "preprocess" / "identifier_mapping.json").read_text()
    )
    assert (identifiers["n_source_genes"], identifiers["n_mapped"]) == (7, 4)
    standard_error = json.loads(
        (root / "preprocess" / "standard_error.json").read_text()
    )
    assert standard_error["n_values"] == 28
    assert (root / "preprocess" / "build_manifest.json").is_file()
    assert sorted(os.listdir(root / "raw")) == sorted(p.RAW_FILE_NAMES)


def test_the_build_refuses_a_withdrawal_set_other_than_wetmores(
    mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        w,
        "subsumption_record",
        lambda data_root=None: SimpleNamespace(
            carried=("set1IT003", "set1IT007"), disregarded=()
        ),
    )
    with pytest.raises(ValueError, match=r"withdrawn samples \['set1IT007'\] != "):
        p.RbTnseqPrice2018EcoliDataset(root=str(_dataset_root(mirrored)))


def test_a_genome_of_the_wrong_strain_is_refused(
    k12: tuple[EcoliK12MG1655Genome, EcoliK12BW25113Genome],
) -> None:
    dataset = p.RbTnseqPrice2018EcoliDataset.__new__(p.RbTnseqPrice2018EcoliDataset)
    dataset.ecoli_genome = k12[0]
    with pytest.raises(TypeError, match="expected the BW25113 genome"):
        dataset._bw25113()
    dataset.ecoli_genome = k12[1]
    assert dataset._bw25113() is k12[1]
    with pytest.raises(NotImplementedError, match="builds its records in process"):
        dataset.create_experiment()
    frame = pd.DataFrame({"a": [1]})
    assert dataset.preprocess_raw(frame) is frame


def test_verify_runs_the_family_verifier_with_the_bacterial_universe(
    mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    p.RbTnseqPrice2018EcoliDataset(root=str(_dataset_root(mirrored)))
    monkeypatch.setattr(p, "EXPECTED_RECORDS", 8)
    report = p.verify(str(mirrored))
    results = {r.name: r.passed for r in report.results}
    assert results["count"] is True
    assert results["pair_uniqueness"] is True
    assert results["gene_containment_sgd"] is True
    assert results["compound_identity"] is True
    # BW25113_0004 is a pseudogene the collection deleted. The shared row used to FAIL
    # on it (it required status `current`, which a `/pseudo` locus can never have) and
    # this loader's supplementary row carried the real check; the shared row now accepts
    # a non-gene feature that resolves to ITSELF and counts it, so both pass and the
    # supplementary row is the stricter restatement rather than the only honest one.
    assert results["canonical_gene_names"] is True
    assert results["stored_tags_are_loci_of_the_pinned_assembly"] is True
    pseudogene = next(
        r for r in report.results if r.name == "canonical_gene_names"
    ).details["self_resolving_non_gene_features"]
    assert any("BW25113_0004" in entry for entry in pseudogene)
    written = json.loads(
        (
            _dataset_root(mirrored) / "preprocess" / "verification_report.json"
        ).read_text()
    )
    assert written["dataset_name"] == "RbTnseqPrice2018EcoliDataset"


def test_report_recomputes_the_notes_numbers_from_the_mirrors(mirrored: Path) -> None:
    out = p.report(str(mirrored))
    assert out["inventory"]["kept_samples_by_source"] == {
        p.SampleSource.wetmore2015.value: 1,
        p.SampleSource.price2018.value: 1,
    }
    assert out["inventory"]["kept_records"] == 8
    assert out["identifiers"]["n_mapped"] == 4
    assert out["standard_error"]["n_values"] == 28
    assert out["set_prefixes"] == ["set1", "set2", "set6"]


def test_the_command_line_dispatches_each_command(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(p, "retrieve_raw_files", lambda dest: f"retrieved {dest}")
    monkeypatch.setattr(
        p, "deposit_raw_mirror", lambda source_dir: f"deposited {source_dir}"
    )
    monkeypatch.setattr(p, "report", lambda: {"records": 8})
    monkeypatch.setattr(
        p, "verify", lambda: SimpleNamespace(summary=lambda: "verified")
    )
    monkeypatch.setattr(p, "essentiality_report", lambda: {"genes": 320})
    monkeypatch.setattr(
        p,
        "verify_essentiality",
        lambda: SimpleNamespace(summary=lambda: "essentiality verified"),
    )
    p.main(["retrieve", "--dest", "/d"])
    p.main(["deposit", "--source-dir", "/s"])
    p.main(["report"])
    p.main(["verify"])
    p.main(["essentiality-report"])
    p.main(["essentiality-verify"])
    assert capsys.readouterr().out.split("\n") == [
        "retrieved /d",
        "deposited /s",
        "{",
        '  "records": 8',
        "}",
        "verified",
        "{",
        '  "genes": 320',
        "}",
        "essentiality verified",
        "",
    ]


def test_a_sample_whose_group_needs_a_condition_names_one() -> None:
    specs = _specs()
    carbon = specs["set1IT003"]
    with pytest.raises(ValueError, match="a carbon source sample names no condition"):
        p.build_environment(carbon.model_copy(update={"condition": None}), SOLVENTS)
    with pytest.raises(ValueError, match="'D-Glucose' has no dose"):
        p.build_environment(carbon.model_copy(update={"concentration": None}), SOLVENTS)


def test_an_eck_that_is_neither_paired_absent_nor_shared_is_refused() -> None:
    with pytest.raises(ValueError, match="b0042: ECK"):
        p.eck_route(["b0042"], CROSSWALK, {"b0042": ("ECK0042",)})


def test_a_table_missing_a_sample_is_refused() -> None:
    tables = _tables(FIT, OBS, NAIVE)
    tables["se_naive"] = tables["se_naive"].drop(columns=["set1IT004 D-Glucose (C)"])
    with pytest.raises(ValueError, match=r"se_naive: samples absent \['set1IT004'\]"):
        p.align_release(**tables, samples=["set1IT003", "set1IT004"], sigma=0.1)


# --------------------------------------------------------------------------- #
# Supplementary Table 1: the likely-essential genes
# --------------------------------------------------------------------------- #
def test_the_essentiality_loader_is_registered_against_bw25113() -> None:
    cls = p.GeneEssentialityPrice2018EcoliDataset
    assert dataset_registry["GeneEssentialityPrice2018EcoliDataset"] is cls
    assert declared_reference_strain(cls) == "BW25113"
    params = inspect.signature(cls.__init__).parameters
    assert "ecoli_genome" in params
    assert params["root"].default == "data/torchcell/gene_essentiality_price2018_ecoli"
    shell = cls.__new__(cls)
    assert shell.raw_file_names == ["si3.xlsx", "fit_logratios_good.tab"]
    assert shell.experiment_class is BacterialGeneEssentialityExperiment
    assert shell.reference_class is BacterialGeneEssentialityExperimentReference
    assert (p.ESSENTIAL_SOURCE_GENES, p.ESSENTIAL_EXPECTED_RECORDS) == (324, 320)
    assert p.MIN_ESSENTIAL_ECK_ROUTE_FRACTION < p.MIN_ECK_ROUTE_FRACTION


def test_a_sheet_is_read_under_its_own_header_row(tmp_path: Path) -> None:
    workbook = _write_workbook(tmp_path / "si3.xlsx")
    table = p.read_below_header(workbook, p.TABLE_S4_SHEET)
    assert list(table["Compound"]) == [c for c, _ in SYNTHETIC_SOLVENT_ROWS]
    assert "Microbe" not in list(table.columns)


def test_a_sheet_without_exactly_one_header_row_is_refused(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = p.TABLE_S4_SHEET
    sheet.append(["Compound", "Solvent"])
    sheet.append(["Compound", "Solvent"])
    path = tmp_path / "two.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="expected one 'Compound' header row, found 2"):
        p.read_below_header(path, p.TABLE_S4_SHEET)


def test_read_solvents_keys_table_s4_by_lowercased_compound(tmp_path: Path) -> None:
    solvents = p.read_solvents(_write_workbook(tmp_path / "si3.xlsx"))
    assert solvents == {
        "cephalothin sodium salt": "water",
        "dimethyl sulfoxide": "water",
        "chloramphenicol": "Ethanol",
        "vanillin": "Dimethyl Sulfoxide",
    }


def test_read_solvents_refuses_an_empty_or_repeated_compound(tmp_path: Path) -> None:
    def sheet(rows: list[list[Any]]) -> Path:
        book = openpyxl.Workbook()
        active = book.active
        assert active is not None
        active.title = p.TABLE_S4_SHEET
        active.append(["preamble"])
        active.append(["Compound", "Solvent"])
        for row in rows:
            active.append(row)
        path = tmp_path / f"{len(rows)}-{rows[0][1]}.xlsx"
        book.save(path)
        return path

    with pytest.raises(ValueError, match=r"'Vanillin' has no Solvent"):
        p.read_solvents(sheet([["Vanillin", None]]))
    with pytest.raises(ValueError, match="two rows name 'vanillin'"):
        p.read_solvents(sheet([["Vanillin", "water"], ["vanillin", "water"]]))


def test_read_essential_genes_keeps_the_keio_rows_and_their_coverage(
    tmp_path: Path,
) -> None:
    genes = p.read_essential_genes(_write_workbook(tmp_path / "si3.xlsx"))
    assert [g.b_number for g in genes] == [b for b, *_ in SYNTHETIC_ESSENTIAL_ROWS]
    assert [g.refseq_locus_tag for g in genes] == [
        t for _, t, *_ in (SYNTHETIC_ESSENTIAL_ROWS)
    ]
    assert [g.row for g in genes] == [1, 2, 3, 4, 5]
    first = genes[0]
    assert (first.name, first.gene_class, first.locus_id) == ("thrL", "Arole", "14000")
    assert (first.n_reads, first.n_pos_central, first.dens) == (0, 0, 0.0)
    assert genes[1].normreads == pytest.approx(0.01)


def test_a_table_s1_row_that_is_not_a_b_number_is_refused(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = p.TABLE_S1_SHEET
    sheet.append(["preamble"])
    header = [
        "organism",
        "orgId",
        "locusId",
        "sysName",
        "locus_tag",
        "name",
        "desc",
        "GC",
        "nReads",
        "normreads",
        "nPosCentral",
        "dens",
        "geneClass",
    ]
    sheet.append(header)
    sheet.append(
        [
            "E. coli",
            "Keio",
            1,
            "BW25113_0001",
            None,
            "thrL",
            "d",
            0.5,
            0,
            0.0,
            0,
            0.0,
            "Arole",
        ]
    )
    path = tmp_path / "bad.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="'BW25113_0001' is no b-number"):
        p.read_essential_genes(path)
    empty = openpyxl.Workbook()
    other = empty.active
    assert other is not None
    other.title = p.TABLE_S1_SHEET
    other.append(["preamble"])
    other.append(header)
    other.append(
        [
            "S. oneidensis",
            "MR1",
            1,
            "b0001",
            None,
            "x",
            "d",
            0.5,
            0,
            0.0,
            0,
            0.0,
            "Arole",
        ]
    )
    no_keio = tmp_path / "nokeio.xlsx"
    empty.save(no_keio)
    with pytest.raises(ValueError, match="no Keio rows"):
        p.read_essential_genes(no_keio)


def test_the_selection_medium_is_lb_lennox_agar_with_kanamycin() -> None:
    medium = p.selection_medium()
    assert medium.state == "solid"
    assert medium.base_medium == "LB"
    assert medium.is_synthetic is False
    added = {c.compound.name: c for c in medium.components[len(LB_LENNOX.components) :]}
    assert sorted(added) == ["agar", "kanamycin"]
    assert added["agar"].role is MediaComponentRole.gelling_agent
    # LB_AGAR's 2% is another paper's bench value, so no agar amount is asserted
    assert added["agar"].concentration is None
    assert added["kanamycin"].role is MediaComponentRole.selection_agent
    assert added["kanamycin"].concentration is not None
    assert (
        added["kanamycin"].concentration.value,
        added["kanamycin"].concentration.unit,
    ) == (50.0, ConcentrationUnit.ug_per_ml)
    assert medium.components[: len(LB_LENNOX.components)] == LB_LENNOX.components


def test_the_essentiality_environment_is_the_library_selection_condition() -> None:
    environment = p.essentiality_environment()
    assert environment.media == p.selection_medium()
    assert environment.temperature is not None
    assert environment.temperature.value == 37.0
    assert environment.perturbations == []
    assert environment.aerobicity == "aerobic"
    assert [g.field for g in environment.provenance_gaps] == [
        "duration_hours",
        "duration_generations",
    ]
    assert environment != p.build_environment(_specs()["set2IT045"], SOLVENTS)


def _essential_identifiers(
    mapped: int, unmapped: tuple[str, ...]
) -> p.IdentifierReport:
    return p.IdentifierReport(
        n_source_genes=mapped + len(unmapped),
        n_mapped=mapped,
        mapped_fraction=mapped / (mapped + len(unmapped)),
        min_fraction=0.5,
        numeric_disagreements=(),
        unmapped=tuple(
            p.UnmappedGene(b_number=b, reason="eck_absent_from_bw25113", eck=())
            for b in unmapped
        ),
        unmapped_by_reason={"eck_absent_from_bw25113": len(unmapped)},
        reconcile_status_histogram={},
        reconcile_layer_histogram={},
        n_symbol_names=mapped,
        n_tag_names=0,
    )


def test_the_essentiality_inventory_counts_rows_and_proves_disjointness(
    tmp_path: Path,
) -> None:
    genes = p.read_essential_genes(_write_workbook(tmp_path / "si3.xlsx"))
    counts = p.essentiality_inventory(
        genes, _essential_identifiers(3, ("b0006", "b0007")), ["b0004", "b0005"]
    )
    assert (counts.source_genes, counts.kept_genes) == (5, 3)
    assert counts.dropped_genes == ("b0006", "b0007")
    assert counts.dropped_gene_names == ("araC", "rhaD")
    # the drop reason is tied to the pinned BW25113 genotype, not asserted from memory
    assert counts.dropped_genes_explained_by_background == {
        "araC": "(araBAD)567",
        "rhaD": "(rhaBAD)568",
    }
    assert counts.drop_rule == p.DROP_NO_ECK_PAIR
    assert counts.rows_without_a_refseq_locus_tag == ("b0006", "b0007")
    assert counts.shared_with_fitness == ()
    assert counts.fitness_genes == 2
    assert counts.gene_class_histogram == {
        "Arole": 1,
        "Bspecific": 1,
        "Cvague": 1,
        "Dhypo": 2,
    }
    # the three quantities the one boolean cannot hold travel with the records
    assert set(counts.unencodable_quantities) == {
        "the label is 'nearly essential' too",
        "the condition is the library isolation, not an assay",
        "the list's own false-discovery rate",
        "the naive rate against the PEC and Keio list",
        "the call rule and its threshold",
        "the coverage evidence behind each call",
    }
    assert (
        "nearly essential"
        in counts.unencodable_quantities["the label is 'nearly essential' too"]
    )
    assert (
        r"between $6 \%$ and $16 \%$"
        in counts.unencodable_quantities["the list's own false-discovery rate"]
    )


def test_a_gene_in_both_releases_or_a_lost_mapping_is_refused(tmp_path: Path) -> None:
    genes = p.read_essential_genes(_write_workbook(tmp_path / "si3.xlsx"))
    with pytest.raises(ValueError, match="both likely-essential and valued"):
        p.essentiality_inventory(
            genes, _essential_identifiers(3, ("b0006", "b0007")), ["b0001", "b0004"]
        )
    with pytest.raises(ValueError, match="are not the rows Table S1 leaves without"):
        p.essentiality_inventory(
            genes, _essential_identifiers(4, ("b0007",)), ["b0004"]
        )


def test_an_essentiality_record_carries_the_label_caveat() -> None:
    mapping = _mapping("b0001", "BW25113_0001", "thrL")
    environment = p.essentiality_environment()
    experiment = p.build_essentiality_experiment("D", mapping, environment)
    assert experiment.phenotype.is_essential is True
    assert experiment.phenotype.label_name == "is_essential"
    assert isinstance(experiment.genotype, Genotype)
    (perturbation,) = experiment.genotype.perturbations
    assert isinstance(perturbation, TransposonInsertionPerturbation)
    assert perturbation.systematic_gene_name == "BW25113_0001"
    assert perturbation.perturbed_gene_name == "thrL"
    assert (perturbation.transposon, perturbation.library_pool) == ("Tn5", "KEIO_ML9")
    assert perturbation.barcode is None
    assert "nearly essential" in perturbation.description
    assert "6% to 16%" in perturbation.description
    assert "DERIVED" in perturbation.description
    reference = p.build_essentiality_reference("D", _pin(), environment)
    assert reference.phenotype_reference.is_essential is False
    assert reference.environment_reference == environment


# --------------------------------------------------------------------------- #
# The essentiality build, its verifier and its report
# --------------------------------------------------------------------------- #
@pytest.fixture
def essentiality_mirrored(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    k12: tuple[EcoliK12MG1655Genome, EcoliK12BW25113Genome],
) -> Path:
    """A tmp ``DATA_ROOT`` whose mirror holds a synthetic Table S1; returns it.

    The fitness table beside it names only ``b0004`` and ``b0005``, so it is disjoint
    from the five essential rows, as the release's own rule requires.
    """
    source = tmp_path / "source"
    source.mkdir()
    _write_workbook(source / "si3.xlsx")
    pd.DataFrame(
        {
            "locusId": [104, 105],
            "sysName": ["b0004", "b0005"],
            "set1IT003 x": [0.0, 1.0],
        }
    ).to_csv(source / "fit_logratios_good.tab", sep="\t", index=False)
    monkeypatch.setattr(
        p,
        "RAW_FILE_NAMES",
        {
            name: (
                "price",
                f"data/{name}",
                hashlib.sha256((source / name).read_bytes()).hexdigest(),
            )
            for name in ("si3.xlsx", "fit_logratios_good.tab")
        },
    )
    monkeypatch.setattr(p, "source_path", lambda name, data_root=None: source / name)
    genomes = dict(zip(("MG1655", "BW25113"), k12, strict=True))
    monkeypatch.setattr(
        p, "bacterial_genome", lambda host, strain, data_root=None: genomes[strain]
    )
    monkeypatch.setattr(p, "assembly_reference", lambda strain: _pin())
    monkeypatch.setattr(p, "load_dotenv", lambda: None)
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.delenv("TC_DATA_URL", raising=False)
    return data_root


def _essentiality_root(data_root: Path) -> Path:
    return data_root / p.ESSENTIAL_DATASET_ROOT_REL


def test_the_essentiality_loader_builds_the_synthetic_table_end_to_end(
    essentiality_mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(p, "MIN_ESSENTIAL_ECK_ROUTE_FRACTION", 0.5)
    root = _essentiality_root(essentiality_mirrored)
    dataset = p.GeneEssentialityPrice2018EcoliDataset(root=str(root))
    assert len(dataset) == 3
    assert sorted(dataset.gene_set) == ["BW25113_0001", "BW25113_0002", "BW25113_4412"]
    references = dataset.experiment_reference_index
    assert references is not None
    assert len(references) == 1
    first = dataset[0]
    assert first["experiment"]["phenotype"]["is_essential"] is True
    assert first["reference"]["phenotype_reference"]["is_essential"] is False
    assert first["publication"]["doi"] == p.PAPER_DOI
    assert first["experiment"]["environment"]["media"]["state"] == "solid"
    drops = json.loads((root / "preprocess" / "dropped_records.json").read_text())
    assert (drops["source_genes"], drops["kept_genes"]) == (5, 3)
    assert drops["dropped_genes"] == ["b0006", "b0007"]
    assert drops["shared_with_fitness"] == []
    label = json.loads((root / "preprocess" / "essentiality_label.json").read_text())
    assert sorted(label) == sorted(p.ESSENTIALITY_LABEL_VALUES)
    assert label["ESSENTIAL_FDR"]["provenance"]["source_uri"] == p.SUPPLEMENTARY_NOTES
    assert "nearly essential" in label["ESSENTIAL_LABEL"]["quote"]
    rows = pd.read_csv(root / "preprocess" / "essential_genes.csv")
    assert list(rows["b_number"]) == [b for b, *_ in SYNTHETIC_ESSENTIAL_ROWS]
    assert list(rows["record"].astype("Int64")) == [0, 1, 2, pd.NA, pd.NA]
    assert list(rows["locus_tag"].fillna("")) == [
        "BW25113_0001",
        "BW25113_0002",
        "BW25113_4412",
        "",
        "",
    ]
    assert (root / "preprocess" / "build_manifest.json").is_file()


def test_the_essentiality_build_refuses_an_overlap_with_the_fitness_genes(
    essentiality_mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(p, "MIN_ESSENTIAL_ECK_ROUTE_FRACTION", 0.5)
    source = p.source_path("fit_logratios_good.tab")
    pd.DataFrame({"locusId": [101], "sysName": ["b0001"], "set1IT003 x": [0.0]}).to_csv(
        source, sep="\t", index=False
    )
    monkeypatch.setitem(
        p.RAW_FILE_NAMES,
        "fit_logratios_good.tab",
        (
            "price",
            "data/fit_logratios_good.tab",
            hashlib.sha256(Path(source).read_bytes()).hexdigest(),
        ),
    )
    with pytest.raises(ValueError, match="both likely-essential and valued"):
        p.GeneEssentialityPrice2018EcoliDataset(
            root=str(_essentiality_root(essentiality_mirrored))
        )


def test_the_essentiality_route_floor_stops_a_wrong_annotation(
    essentiality_mirrored: Path,
) -> None:
    with pytest.raises(LocusTagResolutionError, match="places 3 of 5 genes"):
        p.GeneEssentialityPrice2018EcoliDataset(
            root=str(_essentiality_root(essentiality_mirrored))
        )


def test_an_essentiality_genome_of_the_wrong_strain_is_refused(
    k12: tuple[EcoliK12MG1655Genome, EcoliK12BW25113Genome],
) -> None:
    cls = p.GeneEssentialityPrice2018EcoliDataset
    dataset = cls.__new__(cls)
    dataset.ecoli_genome = k12[0]
    with pytest.raises(TypeError, match="expected the BW25113 genome"):
        dataset._bw25113()
    dataset.ecoli_genome = k12[1]
    assert dataset._bw25113() is k12[1]
    with pytest.raises(NotImplementedError, match="builds its records in process"):
        dataset.create_experiment()
    frame = pd.DataFrame({"a": [1]})
    assert dataset.preprocess_raw(frame) is frame


def test_verify_essentiality_runs_every_row_over_the_built_store(
    essentiality_mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(p, "MIN_ESSENTIAL_ECK_ROUTE_FRACTION", 0.5)
    p.GeneEssentialityPrice2018EcoliDataset(
        root=str(_essentiality_root(essentiality_mirrored))
    )
    monkeypatch.setattr(p, "ESSENTIAL_EXPECTED_RECORDS", 3)
    report = p.verify_essentiality(str(essentiality_mirrored))
    results = {r.name: r.passed for r in report.results}
    assert results["structural"] is True
    assert results["count"] is True
    assert results["calls_match_table_s1"] is True
    assert results["disjoint_from_the_fitness_genes"] is True
    assert results["label_caveat_on_every_record"] is True
    assert results["media_membership"] is True
    # the shared containment row keeps its name and names the universe in its message
    containment = next(r for r in report.results if r.name == "gene_containment_sgd")
    assert containment.passed is True
    assert "BW25113 genes" in containment.message
    assert results["current_genome_genes"] is True
    assert results["stored_tags_are_loci_of_the_pinned_assembly"] is True
    # BW25113_0004 is the fixture's pseudogene and is not in this gene set, so unlike
    # the fitness dataset every row of the essentiality report passes
    assert [r.name for r in report.results if not r.passed] == []
    written = json.loads(
        (
            _essentiality_root(essentiality_mirrored)
            / "preprocess"
            / "verification_report.json"
        ).read_text()
    )
    assert written["dataset_name"] == "GeneEssentialityPrice2018EcoliDataset"


def test_the_essentiality_rows_fail_on_a_record_that_misstates_the_call() -> None:
    mapping = _mapping("b0001", "BW25113_0001", "thrL")
    environment = p.essentiality_environment()
    experiment = p.build_essentiality_experiment("D", mapping, environment)
    reference = p.build_essentiality_reference("D", _pin(), environment)
    record = {
        "experiment": experiment.model_dump(),
        "reference": reference.model_dump(),
    }
    assert p.essentiality_calls_match_table([record, record], []).passed is False
    assert p.essentiality_label_is_qualified([record]).passed is True
    stripped = json.loads(json.dumps(record))
    stripped["experiment"]["genotype"]["perturbations"][0]["description"] = "x"
    assert p.essentiality_label_is_qualified([stripped]).passed is False
    flipped = json.loads(json.dumps(record))
    flipped["experiment"]["phenotype"]["is_essential"] = False
    result = p.essentiality_calls_match_table([flipped], [])
    assert result.passed is False
    assert "is_essential=False" in result.message


def test_the_call_row_names_a_multi_perturbation_or_a_live_reference() -> None:
    mapping = _mapping("b0001", "BW25113_0001", "thrL")
    environment = p.essentiality_environment()
    experiment = p.build_essentiality_experiment("D", mapping, environment)
    reference = p.build_essentiality_reference("D", _pin(), environment)
    record = {
        "experiment": experiment.model_dump(),
        "reference": reference.model_dump(),
    }
    doubled = json.loads(json.dumps(record))
    perturbations = doubled["experiment"]["genotype"]["perturbations"]
    perturbations.append(json.loads(json.dumps(perturbations[0])))
    result = p.essentiality_calls_match_table([doubled], [])
    assert result.passed is False
    assert "more than one perturbation" in result.message
    live = json.loads(json.dumps(record))
    live["reference"]["phenotype_reference"]["is_essential"] = True
    result = p.essentiality_calls_match_table([live], [])
    assert result.passed is False
    assert "expected viable" in result.message


def test_the_refseq_route_reports_a_tag_that_names_another_locus() -> None:
    gene = p.EssentialGene(
        row=1,
        locus_id="1",
        b_number="b0001",
        refseq_locus_tag="BW25113_RS00005",
        name="thrL",
        desc="d",
        gene_class="Arole",
        gc=0.5,
        n_reads=0,
        normreads=0.0,
        n_pos_central=0,
        dens=0.0,
    )
    mapping = {"b0001": _mapping("b0001", "BW25113_0001", "thrL")}
    elsewhere = _resolution("BW25113_RS00005", GeneNameStatus.RENAMED, "BW25113_9999")
    out = p.refseq_route_agreement([gene], mapping, lambda name: elsewhere)
    assert out["n_agree"] == 0
    assert out["disagree"] == [("b0001", "BW25113_9999", "BW25113_0001")]


def test_the_disjointness_row_names_the_genes_that_break_it(tmp_path: Path) -> None:
    genes = p.read_essential_genes(_write_workbook(tmp_path / "si3.xlsx"))
    assert p.essential_genes_are_not_fitness_genes(genes, ["b0004"]).passed is True
    broken = p.essential_genes_are_not_fitness_genes(genes, ["b0001", "b0004"])
    assert broken.passed is False
    assert broken.details["shared"] == ["b0001"]


def test_the_essentiality_report_recomputes_the_notes_numbers(
    essentiality_mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(p, "MIN_ESSENTIAL_ECK_ROUTE_FRACTION", 0.5)
    out = p.essentiality_report(str(essentiality_mirrored))
    assert out["inventory"]["kept_genes"] == 3
    assert out["identifiers"]["n_mapped"] == 3
    # b0001 and b0002 carry a RefSeq tag the annotation knows; b0003's is unknown there
    assert out["refseq_route_agreement"]["n_agree"] == 2
    assert out["refseq_route_agreement"]["unresolved"] == ["b0003"]
    assert out["refseq_route_agreement"]["disagree"] == []
    assert out["table_s1_names_differing_from_the_genome"] == ["b0003 thrW -> hokC"]


# --------------------------------------------------------------------------- #
# Data-gated: the real mirrors
# --------------------------------------------------------------------------- #
def _audit(name: str, sv: Any) -> None:
    raw_keyed = sv.provenance.source_uri.startswith("data/")
    root = Path(DATA_ROOT) / (p.RAW_ROOT_REL if raw_keyed else p.LIBRARY_DIR_REL)
    path = sv.source_path(root)
    if path.suffix == ".xlsx":
        assert hashlib.sha256(path.read_bytes()).hexdigest() == sv.provenance.sha256
        assert sv.quote in p.xlsx_text(str(path)), name
        return
    result = audit_sourced_value(sv, root)
    assert result.passed, f"{name}: {result.message}"


@pytest.mark.data
@needs_mirrors
@pytest.mark.parametrize("name", sorted(p.SOURCED_VALUES))
def test_every_quote_is_verbatim_in_its_pinned_file(name: str) -> None:
    _audit(name, p.SOURCED_VALUES[name])


@pytest.mark.data
@needs_mirrors
@pytest.mark.parametrize("name", sorted(p.DEFERRED_VALUES))
def test_every_deferred_wetmore_quote_is_verbatim(name: str) -> None:
    _audit(name, p.DEFERRED_VALUES[name])


@pytest.fixture(scope="module")
def real_specs() -> tuple[p.SampleSpec, ...]:
    table = w.read_superset_experiments(p.source_path("si3.xlsx", DATA_ROOT))
    carried = frozenset(w.subsumption_record(DATA_ROOT).carried)
    return p.classify_samples(table, carried)


@pytest.mark.data
@needs_mirrors
def test_the_sample_inventory_by_source_and_rule(
    real_specs: tuple[p.SampleSpec, ...],
) -> None:
    by_rule: dict[str | None, list[str]] = {}
    for spec in real_specs:
        by_rule.setdefault(spec.drop_rule, []).append(spec.name)
    assert len(real_specs) == 162
    assert {k: len(v) for k, v in by_rule.items()} == {
        None: 147,
        p.DROP_DISREGARDED: 4,
        p.DROP_MEDIUM_NOT_IN_LIBRARY: 4,
        p.DROP_MOTILITY: 7,
    }
    assert by_rule[p.DROP_DISREGARDED] == [
        "set1IT007",
        "set1IT008",
        "set1IT043",
        "set1IT044",
    ]
    assert by_rule[p.DROP_MEDIUM_NOT_IN_LIBRARY] == [
        "set1IT067",
        "set1IT068",
        "set1IT069",
        "set1IT070",
    ]
    assert sum(s.source is p.SampleSource.wetmore2015 for s in real_specs) == 92
    kept = [s for s in real_specs if s.drop_rule is None]
    assert sum(s.source is p.SampleSource.wetmore2015 for s in kept) == 84
    assert sum(s.source is p.SampleSource.price2018 for s in kept) == 63
    assert sorted({s.name.split("IT")[0] for s in real_specs}) == [
        "set1",
        "set2",
        "set6",
    ]
    assert set(by_rule[p.DROP_DISREGARDED]) == set(
        w.subsumption_record(DATA_ROOT).disregarded
    )


@pytest.mark.data
@needs_mirrors
def test_the_released_t_is_reproduced_by_the_larger_standard_error(
    real_specs: tuple[p.SampleSpec, ...],
) -> None:
    def read(name: str) -> pd.DataFrame:
        return pd.read_csv(p.source_path(name, DATA_ROOT), sep="\t")

    release = p.align_release(
        fitness=read("fit_logratios_good.tab"),
        se_obs=read("fit_standard_error_obs.tab"),
        se_naive=read("fit_standard_error_naive.tab"),
        t=read("fit_t.tab"),
        samples=[s.name for s in real_specs],
        sigma=0.1,
    )
    assert release.fitness.shape == (3789, 162)
    assert release.n_values == 613818
    assert release.n_naive_larger == 219448
    assert release.t_max_abs_deviation < 1e-12


@pytest.mark.data
@needs_mirrors
def test_the_eck_route_on_the_deposited_annotations() -> None:
    fitness = pd.read_csv(
        p.source_path("fit_logratios_good.tab", DATA_ROOT),
        sep="\t",
        usecols=["sysName"],
    )
    mg1655 = bacterial_genome("ecoli", "MG1655", DATA_ROOT)
    bw25113 = bacterial_genome("ecoli", "BW25113", DATA_ROOT)
    assert isinstance(mg1655, EcoliK12MG1655Genome)
    assert isinstance(bw25113, EcoliK12BW25113Genome)
    mapping, report = p.map_genes(
        list(fitness["sysName"]), mg1655, bw25113, label="test"
    )
    assert (report.n_source_genes, report.n_mapped) == (3789, 3768)
    assert report.numeric_disagreements == (("b0018", "BW25113_4412", "ECK0018"),)
    assert report.unmapped_by_reason == {
        "b_number_not_in_mg1655_annotation": 8,
        "eck_absent_from_bw25113": 7,
        "eck_not_one_to_one": 6,
    }
    assert report.reconcile_status_histogram["renamed"] == 3625
    assert report.reconcile_status_histogram["non_gene_feature"] == 143
    assert report.reconcile_layer_histogram["gene synonym"] == 3768
    assert (report.n_symbol_names, report.n_tag_names) == (3625, 143)
    assert mapping["b0002"].locus_tag == "BW25113_0002"
    assert mapping["b0002"].perturbed_gene_name == "thrA"


@pytest.mark.data
@needs_mirrors
def test_the_raw_mirror_holds_the_pinned_retrievals() -> None:
    manifest = p.load_manifest(DATA_ROOT)
    assert {f.path: f.sha256 for f in manifest.files} == {
        rel: raw.sha256 for rel, raw in p.RAW_FILES.items()
    }
    assert {f.retrieval.source_url for f in manifest.files if f.retrieval} == {
        raw.url for raw in p.RAW_FILES.values()
    }
    for name, (_, _, pin) in p.RAW_FILE_NAMES.items():
        path = p.source_path(name, DATA_ROOT)
        assert hashlib.sha256(path.read_bytes()).hexdigest() == pin, name


BUILT_LMDB = osp.join(
    DATA_ROOT, "data/torchcell/rbtnseq_price2018_ecoli/processed/lmdb"
)


@pytest.mark.data
@pytest.mark.skipif(not osp.isdir(BUILT_LMDB), reason="requires the dev LMDB build")
def test_the_dev_build_holds_every_expected_record() -> None:
    env = lmdb.open(BUILT_LMDB, readonly=True, lock=False)
    try:
        assert env.stat()["entries"] == p.EXPECTED_RECORDS
        with env.begin() as txn:
            payload: bytes | None = txn.get(b"0")
        assert payload is not None, "record 0 is absent from the built LMDB"
        first: dict[str, Any] = pickle.loads(payload)
    finally:
        env.close()
    experiment = first["experiment"]
    assert experiment["experiment_type"] == "bacterial_environment_response"
    perturbation = experiment["genotype"]["perturbations"][0]
    assert perturbation["perturbation_type"] == "transposon_insertion"
    assert perturbation["gene_namespace"] == "ecoli_k12_bw25113_locus_tag"
    assert experiment["phenotype"]["screen_id"] == "Keio:set1IT003"


@pytest.mark.data
@needs_mirrors
def test_the_stress_vehicles_of_every_kept_sample(
    real_specs: tuple[p.SampleSpec, ...],
) -> None:
    """Every kept stress sample's Condition_1 is a Table S4 compound (measured)."""
    solvents = p.read_solvents(p.source_path("si3.xlsx", DATA_ROOT))
    assert len(solvents) == 55
    assert sorted(set(solvents.values())) == ["Dimethyl Sulfoxide", "Ethanol", "water"]
    stress = [s for s in real_specs if s.drop_rule is None and s.group == "stress"]
    assert len(stress) == 55
    assert len({s.condition for s in stress}) == 35
    assert all(
        s.condition is not None and s.condition.strip().lower() in solvents
        for s in stress
    )
    by_sample = Counter(p.stress_solvent(s, solvents).name for s in stress)
    assert by_sample == {"water": 45, "Dimethyl Sulfoxide": 7, "Ethanol": 3}
    by_compound = Counter(
        solvents[c.strip().lower()]
        for c in {s.condition for s in stress if s.condition is not None}
    )
    assert by_compound == {"water": 29, "Dimethyl Sulfoxide": 5, "Ethanol": 1}


@pytest.mark.data
@needs_mirrors
def test_table_s1s_keio_rows_and_their_disjointness_from_the_fitness_genes() -> None:
    """324 released rows, 320 records, intersection 0 with the 3,789 fitness genes."""
    genes = p.read_essential_genes(p.source_path("si3.xlsx", DATA_ROOT))
    assert len(genes) == p.ESSENTIAL_SOURCE_GENES
    assert all(g.gene_class in {"Arole", "Bspecific", "Cvague", "Dhypo"} for g in genes)
    fitness = pd.read_csv(
        p.source_path("fit_logratios_good.tab", DATA_ROOT),
        sep="\t",
        usecols=["sysName"],
    )
    b_numbers = [str(v) for v in fitness["sysName"]]
    assert len(b_numbers) == 3789
    assert {g.b_number for g in genes} & set(b_numbers) == set()
    mg1655 = bacterial_genome("ecoli", "MG1655", DATA_ROOT)
    bw25113 = bacterial_genome("ecoli", "BW25113", DATA_ROOT)
    assert isinstance(mg1655, EcoliK12MG1655Genome)
    assert isinstance(bw25113, EcoliK12BW25113Genome)
    mapping, report = p.map_genes(
        [g.b_number for g in genes],
        mg1655,
        bw25113,
        label="test",
        min_fraction=p.MIN_ESSENTIAL_ECK_ROUTE_FRACTION,
    )
    assert (report.n_source_genes, report.n_mapped) == (324, 320)
    assert report.unmapped_by_reason == {"eck_absent_from_bw25113": 4}
    counts = p.essentiality_inventory(genes, report, b_numbers)
    assert counts.kept_genes == p.ESSENTIAL_EXPECTED_RECORDS
    assert sorted(counts.dropped_genes) == sorted(p.DELETED_OPERON_B_NUMBERS)
    assert sorted(counts.dropped_gene_names) == ["araA", "araB", "rhaA", "rhaB"]
    assert counts.gene_class_histogram == {
        "Arole": 200,
        "Bspecific": 98,
        "Cvague": 7,
        "Dhypo": 19,
    }
    # a second, independent identifier route: Table S1's own RefSeq locus_tag
    agreement = p.refseq_route_agreement(genes, mapping, bw25113.resolve_gene_name)
    assert agreement["n_agree"] == 318
    assert agreement["unresolved"] == ["b1457", "b4047"]
    assert agreement["disagree"] == []


ESSENTIAL_LMDB = osp.join(DATA_ROOT, p.ESSENTIAL_DATASET_ROOT_REL, "processed/lmdb")


@pytest.mark.data
@pytest.mark.skipif(
    not osp.isdir(ESSENTIAL_LMDB), reason="requires the essentiality LMDB build"
)
def test_the_essentiality_dev_build_holds_every_expected_record() -> None:
    env = lmdb.open(ESSENTIAL_LMDB, readonly=True, lock=False)
    try:
        assert env.stat()["entries"] == p.ESSENTIAL_EXPECTED_RECORDS
        with env.begin() as txn:
            payload: bytes | None = txn.get(b"0")
        assert payload is not None, "record 0 is absent from the built LMDB"
        first: dict[str, Any] = pickle.loads(payload)
    finally:
        env.close()
    experiment = first["experiment"]
    assert experiment["experiment_type"] == "bacterial_gene_essentiality"
    assert experiment["phenotype"]["is_essential"] is True


def test_a_gene_bw25113_lacks_for_an_unstated_reason_is_refused(tmp_path: Path) -> None:
    """A drop the pinned background genotype cannot account for stops the build.

    The four real drops are araA, araB, rhaA and rhaB, which ``(araBAD)567`` and
    ``(rhaBAD)568`` of ``BW25113_BACKGROUND_LESIONS`` explain. A fifth gene the ECK route
    could not place would mean something else is wrong with the mapping, so it raises
    rather than being dropped quietly.
    """
    rows = [
        *SYNTHETIC_ESSENTIAL_ROWS,
        ("b0005", None, "proB", "glutamate 5-kinase", "Arole"),
    ]
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = p.TABLE_S1_SHEET
    sheet.append(["preamble"])
    sheet.append(
        [
            "organism",
            "orgId",
            "locusId",
            "sysName",
            "locus_tag",
            "name",
            "desc",
            "GC",
            "nReads",
            "normreads",
            "nPosCentral",
            "dens",
            "geneClass",
        ]
    )
    for b, tag, name, desc, gene_class in rows:
        sheet.append(
            ["E. coli", "Keio", 1, b, tag, name, desc, 0.5, 0, 0.0, 0, 0.0, gene_class]
        )
    path = tmp_path / "extra.xlsx"
    book.save(path)
    genes = p.read_essential_genes(path)
    with pytest.raises(ValueError, match=r"\['proB'\] are absent from BW25113"):
        p.essentiality_inventory(
            genes, _essential_identifiers(3, ("b0005", "b0006", "b0007")), ["b0004"]
        )
