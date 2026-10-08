# tests/torchcell/datasets/scerevisiae/test_kemmeren2014_synthetic.py
# [[tests.torchcell.datasets.scerevisiae.test_kemmeren2014_synthetic]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/scerevisiae/test_kemmeren2014_synthetic.py
"""Kemmeren 2014 microarray loader built end to end on in-memory GEOparse objects.

``raw/`` holds the six ``*_family.soft.gz`` names (placeholder bytes: ``process()`` never
opens them, their presence only stops PyG calling ``download()``), Table S1 as
``kemmeren2014_table_s1.xlsx`` (openpyxl), and pickled REAL ``GEOparse.GEOTypes.GSE``
objects built in memory (``GSM``/``GPL`` with a ``description`` column frame) at the
``<accession>.pkl`` paths ``process()`` checks first. The two deletion GSEs both have
one, so ``GEOparse.get_GEO`` (the network branch) is never reached. The four wildtype
series have pickles too (GSE42215 excepted) but ``process()`` no longer opens them
(#484); a test removes them and gets the same records. The genome is a stub with the
three attributes ``resolve_gene_name_comprehensive`` reads.

Channels follow GEO's own metadata, as on the real series: ``label_ch1`` is Cy5 and
``label_ch2`` Cy3 on every array, and ``source_name_ch1`` / ``source_name_ch2`` say
which channel holds the reference pool ("refpool", or "ref1" on GSE42217). The title
suffix is not consulted.

Platform probes: 1 YAL001C, 2 YBR001C, 3 Q0010, 4 NaN (unmapped). Table S1: YPL177C
CUP9 MATa, YHR127W HSN1 MATalpha, TLC1 (-> YNCB0010W) MATa, CMS1 (-> YLR003C) MATa,
YXX001W with mating type ``diploid`` (skipped).

Deletion samples (responsive GSE42527 then non-responsive GSE42526), signals listed
as YAL001C, YBR001C, Q0010:

    GSM1 "[HS1991] cup9-del-a"   refpool Cy5 4, 4, 1   deletion Cy3 2, 8, 0
    GSM2 "cup9-del-b"            deletion Cy5 8, 2, 4  refpool Cy3 4, 4, 1  (dye swap)
    GSM3 "Sample X"              deletion Cy5 3, 6, 1  refpool Cy3 3, 3, 2; the gene is
                                 named only in ``characteristics_ch1``
    GSM4 "wt control"            wildtype, ignored
    GSM5 "cdk8-del-a"            resolves through the shared reconciler to YPL042C,
                                 which has no Table S1 strain: not written
    GSM6 "zzz9-del-a"            retired name: unresolved

CUP9 -> YPL177C (BY4741): the log2 ratio is taken within each array, log2(deletion /
refpool). YAL001C 2/4 -> -1 and 8/4 -> +1, mean 0, sample SD sqrt(2), SE sqrt(2)/sqrt(2)
= 1, variance 2 (stored as sqrt(2)**2 = 2.0000000000000004); YBR001C the mirror image,
same statistics; Q0010 drops the array with the 0 signal and keeps 4/1 -> +2 alone (n =
1, SE and variance NaN). The linear ``expression`` is the mean deletion signal over the
kept arrays, 5, 5, 4, and the reference ``expression`` the mean refpool signal over the
same arrays, 4, 4, 1. HSN1 -> YHR127W (BY4742): 3/3 -> 0, 6/3 -> +1, 1/2 -> -1, all n
= 1, linear 3, 6, 1, refpool 3, 3, 2.

The reference ``n_replicates`` of a gene is the number of arrays whose refpool value
entered its reference ``expression`` mean, the same arrays as the mutant's count (#484):
CUP9 YAL001C 2, YBR001C 2, Q0010 1; HSN1 1 each. The wildtype arrays (MATa W1, W2, W3;
MATalpha W4) once gave counts of 3, 2, 2 and 1, 1, 1 here; they enter no record.
"""

from __future__ import annotations

import json
import math
import pickle
import re
from pathlib import Path
from typing import Any, cast

import GEOparse
import lmdb
import numpy as np
import openpyxl
import pandas as pd
import pytest
import requests
from GEOparse.GEOTypes import GPL, GSE, GSM

from torchcell.datamodels.schema import (
    Environment,
    Genotype,
    KanMxDeletionPerturbation,
    Media,
    MicroarrayExpressionExperiment,
    MicroarrayExpressionExperimentReference,
    MicroarrayExpressionPhenotype,
    Publication,
    ReferenceGenome,
    Temperature,
)
from torchcell.datasets.scerevisiae import kemmeren2014 as m
from torchcell.sequence.genome.scerevisiae import GeneNameStatus, SCerevisiaeGenome
from torchcell.sequence.genome.scerevisiae.s288c import GeneNameResolution

_DATASET = "MicroarrayKemmeren2014Dataset"
_NAN = "NaN"
_PROBES = {"1": "YAL001C", "2": "YBR001C", "3": "Q0010"}


class _StubGenome:
    """The three attributes ``resolve_gene_name_comprehensive`` reads."""

    gene_attribute_table = pd.DataFrame(
        {"ID": ["YPL177C"], "gene": ["CUP9"], "Alias": ["NONE1"]}
    )
    alias_to_systematic: dict[str, list[str]] = {"CDK8": ["YPL042C"]}

    def resolve_gene_name(self, name: str) -> GeneNameResolution:
        if name.upper() == "CDK8":
            return GeneNameResolution(
                input_name=name,
                status=GeneNameStatus.CURRENT,
                systematic_name="YPL042C",
            )
        return GeneNameResolution(
            input_name=name, status=GeneNameStatus.RETIRED, systematic_name=name
        )


def _genome() -> SCerevisiaeGenome:
    return cast(SCerevisiaeGenome, _StubGenome())


def _describe(table: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        {"description": [f"{c} column" for c in table.columns]}, index=table.columns
    )


def _gsm(
    name: str,
    title: str,
    cy5: list[float],
    cy3: list[float],
    refpool_in: str = "Cy5",
    test: str = "wt",
    reference: str = "refpool",
    characteristics: list[str] | None = None,
) -> GSM:
    """A two-channel array with GEO-style channel metadata.

    ``label_ch1`` is Cy5 and ``label_ch2`` Cy3, as on every real array; ``refpool_in``
    says which dye carries the reference (named ``reference``), and ``test`` names the
    other channel. ``characteristics`` are attached to the test channel.
    """
    table = pd.DataFrame(
        {
            "ID_REF": [1, 2, 3, 4],
            "VALUE": [0.0, 0.0, 0.0, 0.0],
            "Signal Norm_Cy5": [*cy5, 9.0],
            "Signal Norm_Cy3": [*cy3, 9.0],
        }
    )
    reference_channel = 1 if refpool_in == "Cy5" else 2
    test_channel = 3 - reference_channel
    metadata: dict[str, list[str]] = {
        "title": [title],
        "geo_accession": [name],
        "label_ch1": ["Cy5"],
        "label_ch2": ["Cy3"],
        f"source_name_ch{reference_channel}": [reference],
        f"source_name_ch{test_channel}": [test],
        f"characteristics_ch{test_channel}": characteristics or [],
    }
    return GSM(name=name, metadata=metadata, table=table, columns=_describe(table))


def _gpl() -> GPL:
    table = pd.DataFrame(
        {"ID": [1, 2, 3, 4], "ORF": ["YAL001C", "YBR001C", "Q0010", np.nan]}
    )
    return GPL(name="GPL11232", metadata={}, table=table, columns=_describe(table))


def _gse(name: str, gsms: list[GSM], with_platform: bool = True) -> GSE:
    return GSE(
        name=name,
        metadata={"title": [name]},
        gpls={"GPL11232": _gpl()} if with_platform else {},
        gsms={gsm.name: gsm for gsm in gsms},
    )


def _write_table_s1(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["orf name", "gene", "mating type"])
    sheet.append(["YPL177C", "CUP9", "MATa"])
    sheet.append(["YHR127W", "HSN1", "MATalpha"])
    sheet.append(["TLC1", "TLC1", "MATa"])
    sheet.append(["CMS1", None, "MATa"])
    sheet.append(["YXX001W", "XXX1", "diploid"])
    workbook.save(path)


def _write_raw(raw: Path) -> None:
    raw.mkdir(parents=True)
    for accession in (
        "GSE42527",
        "GSE42526",
        "GSE42241",
        "GSE42240",
        "GSE42217",
        "GSE42215",
    ):
        name = f"{accession}_family.soft.gz"
        (raw / name).write_bytes(b"placeholder: process() reads the .pkl")
    _write_table_s1(raw / "kemmeren2014_table_s1.xlsx")
    gses = {
        "GSE42527": _gse(
            "GSE42527",
            [
                _gsm(
                    "GSM1",
                    "[HS1991] cup9-del-a",
                    [4.0, 4.0, 1.0],
                    [2.0, 8.0, 0.0],
                    refpool_in="Cy5",
                    test="cup9-del",
                ),
                _gsm(
                    "GSM2",
                    "cup9-del-b",
                    [8.0, 2.0, 4.0],
                    [4.0, 4.0, 1.0],
                    refpool_in="Cy3",
                    test="cup9-del",
                ),
            ],
        ),
        "GSE42526": _gse(
            "GSE42526",
            [
                _gsm(
                    "GSM3",
                    "Sample X",
                    [3.0, 6.0, 1.0],
                    [3.0, 3.0, 2.0],
                    refpool_in="Cy3",
                    test="hsn1-del",
                    characteristics=["strain: BY4742", "genotype/variation: hsn1-del"],
                ),
                _gsm(
                    "GSM4",
                    "wt control",
                    [1.0, 1.0, 1.0],
                    [1.0, 1.0, 1.0],
                    refpool_in="Cy5",
                    test="wt",
                    characteristics=["genotype/variation: wt"],
                ),
                _gsm(
                    "GSM5",
                    "cdk8-del-a",
                    [1.0, 1.0, 1.0],
                    [1.0, 1.0, 1.0],
                    refpool_in="Cy5",
                    test="cdk8-del",
                ),
                _gsm(
                    "GSM6",
                    "zzz9-del-a",
                    [1.0, 1.0, 1.0],
                    [1.0, 1.0, 1.0],
                    refpool_in="Cy5",
                    test="zzz9-del",
                ),
            ],
            with_platform=False,
        ),
        "GSE42241": _gse(
            "GSE42241",
            [
                _gsm(
                    "W1",
                    "wt-matA-1-a",
                    [1.0, 1.0, 1.0],
                    [0.0, 0.0, 0.0],
                    refpool_in="Cy5",
                    test="wt-matA",
                ),
                _gsm(
                    "W2",
                    "wt-matA-1-b",
                    [0.0, 0.0, 0.0],
                    [2.0, 0.0, 2.0],
                    refpool_in="Cy3",
                    test="wt-matA",
                ),
            ],
        ),
        "GSE42240": _gse(
            "GSE42240",
            [
                _gsm(
                    "W3",
                    "wt-matA-THM012-a",
                    [3.0, 3.0, 0.0],
                    [0.0, 0.0, 0.0],
                    refpool_in="Cy5",
                    test="wt-matA",
                )
            ],
        ),
        "GSE42217": _gse(
            "GSE42217",
            [
                _gsm(
                    "W4",
                    "wt-htp07-a",
                    [5.0, 5.0, 5.0],
                    [0.0, 0.0, 0.0],
                    refpool_in="Cy5",
                    test="wt-htp07-a",
                    reference="ref1",
                )
            ],
        ),
    }
    for accession, gse in gses.items():
        with open(raw / f"{accession}.pkl", "wb") as handle:
            pickle.dump(gse, handle)


@pytest.fixture
def dataset(tmp_path: Path) -> m.MicroarrayKemmeren2014Dataset:
    root = tmp_path / "kemmeren"
    _write_raw(root / "raw")
    return m.MicroarrayKemmeren2014Dataset(root=str(root), genome=_genome())


def _nan_safe(value: Any) -> Any:
    """Replace float NaN by a string so two dumps compare exactly (NaN != NaN)."""
    if isinstance(value, dict):
        return {k: _nan_safe(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_nan_safe(v) for v in value]
    if isinstance(value, float) and math.isnan(value):
        return _NAN
    return value


_ENVIRONMENT = Environment(
    media=Media(name="SC", state="liquid", is_synthetic=True),
    temperature=Temperature(value=30),
)
_PUBLICATION = Publication(
    pubmed_id="24766815",
    pubmed_url="https://pubmed.ncbi.nlm.nih.gov/24766815/",
    doi="10.1016/j.cell.2014.02.054",
    doi_url="https://doi.org/10.1016/j.cell.2014.02.054",
)


def _experiment(
    orf: str,
    log2: dict[str, float],
    se: dict[str, float],
    var: dict[str, float],
    n: dict[str, int],
    linear: dict[str, float],
) -> MicroarrayExpressionExperiment:
    return MicroarrayExpressionExperiment(
        dataset_name=_DATASET,
        genotype=Genotype(
            perturbations=[
                KanMxDeletionPerturbation(
                    systematic_gene_name=orf, perturbed_gene_name=orf
                )
            ]
        ),
        environment=_ENVIRONMENT,
        phenotype=MicroarrayExpressionPhenotype(
            expression=linear,
            expression_log2_ratio=log2,
            expression_log2_ratio_se=se,
            expression_log2_ratio_variance=var,
            n_replicates=n,
        ),
    )


def _reference(
    strain: str, refpool: dict[str, float], n: dict[str, int]
) -> MicroarrayExpressionExperimentReference:
    return MicroarrayExpressionExperimentReference(
        dataset_name=_DATASET,
        genome_reference=ReferenceGenome(
            species="Saccharomyces cerevisiae", strain=strain
        ),
        environment_reference=_ENVIRONMENT,
        phenotype_reference=MicroarrayExpressionPhenotype(
            expression=refpool,
            expression_log2_ratio=dict.fromkeys(refpool, 0.0),
            n_replicates=n,
        ),
    )


_NANF = float("nan")
# The variance is stored as sd**2 with sd = np.std(ddof=1) = sqrt(2), which squares back
# to 2.0000000000000004 in floating point, not 2.0.
_SD2 = math.sqrt(2.0)
_CUP9 = _experiment(
    "YPL177C",
    {"Q0010": 2.0, "YAL001C": 0.0, "YBR001C": 0.0},
    {"Q0010": _NANF, "YAL001C": 1.0, "YBR001C": 1.0},
    {"Q0010": _NANF, "YAL001C": _SD2**2, "YBR001C": _SD2**2},
    {"Q0010": 1, "YAL001C": 2, "YBR001C": 2},
    {"Q0010": 4.0, "YAL001C": 5.0, "YBR001C": 5.0},
)
_HSN1 = _experiment(
    "YHR127W",
    {"Q0010": -1.0, "YAL001C": 0.0, "YBR001C": 1.0},
    dict.fromkeys(("Q0010", "YAL001C", "YBR001C"), _NANF),
    dict.fromkeys(("Q0010", "YAL001C", "YBR001C"), _NANF),
    {"Q0010": 1, "YAL001C": 1, "YBR001C": 1},
    {"Q0010": 1.0, "YAL001C": 3.0, "YBR001C": 6.0},
)
_REF_CUP9 = _reference(
    "BY4741",
    {"Q0010": 1.0, "YAL001C": 4.0, "YBR001C": 4.0},
    {"Q0010": 1, "YAL001C": 2, "YBR001C": 2},
)
_REF_HSN1 = _reference(
    "BY4742",
    {"Q0010": 2.0, "YAL001C": 3.0, "YBR001C": 3.0},
    {"Q0010": 1, "YAL001C": 1, "YBR001C": 1},
)


def test_two_deletions_build_with_within_array_log2_statistics(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """Record 0 is CUP9 (two dye-swapped arrays, MATa), record 1 HSN1 (one array,
    MATalpha, named only in ``characteristics_ch1``); CDK8 resolves but has no Table S1
    strain and ZZZ9 does not resolve, so neither is written. The deleted gene is not
    on the platform, so the ratios are the docstring's within-array values.
    """
    assert len(dataset) == 2
    assert _nan_safe(dataset[0]["experiment"]) == _nan_safe(_CUP9.model_dump())
    assert _nan_safe(dataset[1]["experiment"]) == _nan_safe(_HSN1.model_dump())
    assert dataset[0]["reference"] == _REF_CUP9.model_dump()
    assert dataset[1]["reference"] == _REF_HSN1.model_dump()
    assert dataset[0]["publication"] == _PUBLICATION.model_dump()


def test_side_files_gene_set_reference_index_and_sample_table(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """One reference per strain; ``data.csv`` lists all six deletion-GSE samples with
    the resolved ORF (empty for the wildtype and the unresolved title).
    """
    preprocess = Path(dataset.root) / "preprocess"
    assert json.loads((preprocess / "gene_set.json").read_text()) == [
        "YHR127W",
        "YPL177C",
    ]
    index = json.loads((preprocess / "experiment_reference_index.json").read_text())
    assert [entry["member_indices"] for entry in index] == [[0], [1]]
    assert [entry["reference"] for entry in index] == [
        _REF_CUP9.model_dump(),
        _REF_HSN1.model_dump(),
    ]
    samples = pd.read_csv(preprocess / "data.csv", keep_default_na=False)
    assert samples.to_dict(orient="list") == {
        "geo_accession": ["GSM1", "GSM2", "GSM3", "GSM4", "GSM5", "GSM6"],
        "title": [
            "[HS1991] cup9-del-a",
            "cup9-del-b",
            "Sample X",
            "wt control",
            "cdk8-del-a",
            "zzz9-del-a",
        ],
        "systematic_gene_name": ["YPL177C", "YPL177C", "YHR127W", "", "YPL042C", ""],
        "is_deletion": [True, True, True, False, True, True],
        "is_wildtype": [False, False, False, True, False, False],
        "is_responsive_mutant": [True, True, False, False, False, False],
    }
    manifest = json.loads((preprocess / "build_manifest.json").read_text())
    assert manifest["dataset_name"] == "kemmeren"
    assert manifest["loader_class"] == _DATASET
    assert dataset.experiment_class is MicroarrayExpressionExperiment
    assert dataset.reference_class is MicroarrayExpressionExperimentReference
    frame = pd.DataFrame({"a": [1]})
    assert dataset.preprocess_raw(frame) is frame


def test_mating_type_map_reads_table_s1_and_rewrites_common_names(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """TLC1 and CMS1 in the orf column become YNCB0010W and YLR003C; ``diploid`` is
    neither mating type and is skipped.
    """
    assert dataset._load_mating_type_map() == (
        {
            "YPL177C": "BY4741",
            "YHR127W": "BY4742",
            "YNCB0010W": "BY4741",
            "YLR003C": "BY4741",
        },
        {"CUP9": "YPL177C", "HSN1": "YHR127W", "TLC1": "YNCB0010W", "CMS1": "YLR003C"},
    )


def test_a_table_s1_without_an_orf_column_yields_empty_maps(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """Finding: ``_load_mating_type_map`` raises ``ValueError("Required 'orf name'
    column not found ...")`` inside a ``try`` whose ``except Exception`` only logs, so
    a Table S1 missing its required column returns two empty maps instead of failing.
    """
    workbook = openpyxl.Workbook()
    workbook.active.append(["gene", "mating type"])
    workbook.active.append(["CUP9", "MATa"])
    workbook.save(Path(dataset.raw_dir) / "kemmeren2014_table_s1.xlsx")
    assert dataset._load_mating_type_map() == ({}, {})
    (Path(dataset.raw_dir) / "kemmeren2014_table_s1.xlsx").unlink()
    with pytest.raises(FileNotFoundError, match="Supplementary Table S1 not found at"):
        dataset._load_mating_type_map()


def test_channels_come_from_geo_metadata_not_from_the_title() -> None:
    """``_channel_columns`` reads ``source_name_ch*`` and ``label_ch*``: a "-a" title
    with the refpool in Cy5 gives the test in Cy3, a "-b" title of a gene whose own
    name carries ``-a`` (``ycr087c-a-del-b``) with the refpool in Cy3 gives the test in
    Cy5, and "ref1" counts as the reference. Two reference channels, none, and a label
    pair other than Cy5 + Cy3 raise ``ValueError``; a deletion named ``ref2-del`` is
    not mistaken for a reference.
    """
    columns = m.MicroarrayKemmeren2014Dataset._channel_columns
    ones = [1.0, 1.0, 1.0]
    assert columns(_gsm("A", "cup9-del-a", ones, ones, "Cy5", "cup9-del")) == (
        "Signal Norm_Cy3",
        "Signal Norm_Cy5",
    )
    assert columns(
        _gsm("B", "ycr087c-a-del-b", ones, ones, "Cy3", "ycr087c-a-del")
    ) == ("Signal Norm_Cy5", "Signal Norm_Cy3")
    assert columns(_gsm("C", "wt-htp07-a", ones, ones, "Cy5", "wt", "ref1")) == (
        "Signal Norm_Cy3",
        "Signal Norm_Cy5",
    )
    assert columns(_gsm("D", "ref2-del-a", ones, ones, "Cy5", "ref2-del")) == (
        "Signal Norm_Cy3",
        "Signal Norm_Cy5",
    )
    with pytest.raises(
        ValueError,
        match=re.escape(
            "E: expected the reference pool in exactly one channel, source names "
            "['refpool', 'refpool']"
        ),
    ):
        columns(_gsm("E", "x", ones, ones, "Cy5", "refpool"))
    with pytest.raises(ValueError, match="F: expected the reference pool"):
        columns(_gsm("F", "x", ones, ones, "Cy5", "cup9-del", reference="wt"))
    swapped = _gsm("G", "x", ones, ones, "Cy5", "cup9-del")
    swapped.metadata["label_ch2"] = ["Cy5"]
    with pytest.raises(
        ValueError,
        match=re.escape("G: channel labels ['Cy5', 'Cy5'] are not one Cy5 and one Cy3"),
    ):
        columns(swapped)


def test_extract_channels_pairs_the_two_signals_of_each_row() -> None:
    """``ycr087c-a-del-b`` with the refpool in Cy3: the deletion values are the Cy5
    column and the refpool values the Cy3 column, row by row, the unmapped probe 4
    skipped; a missing signal column raises.
    """
    extract = m.MicroarrayKemmeren2014Dataset._extract_channels_from_gsm_static
    gsm = _gsm(
        "G", "ycr087c-a-del-b", [8.0, 2.0, 4.0], [4.0, 4.0, 1.0], "Cy3", "ycr087c-a-del"
    )
    deletion, refpool = extract(gsm, _PROBES)
    assert dict(deletion) == {"Q0010": 4.0, "YAL001C": 8.0, "YBR001C": 2.0}
    assert dict(refpool) == {"Q0010": 1.0, "YAL001C": 4.0, "YBR001C": 4.0}
    gsm.table = gsm.table.drop(columns=["Signal Norm_Cy3"])
    with pytest.raises(
        ValueError, match=re.escape("G: column 'Signal Norm_Cy3' missing")
    ):
        extract(gsm, _PROBES)


def test_replicate_pairs_are_one_pair_per_array_in_array_order() -> None:
    """GSM1 (refpool Cy5 4, 4, 1; deletion Cy3 2, 8, 0) then GSM2 (deletion Cy5 8, 2,
    4; refpool Cy3 4, 4, 1) give per gene the (deletion, refpool) pairs in that order,
    the 0 signal kept here and dropped later by ``create_expression_experiment``.
    """
    collect = m.MicroarrayKemmeren2014Dataset._collect_replicate_pairs_static
    first = _gsm(
        "GSM1", "cup9-del-a", [4.0, 4.0, 1.0], [2.0, 8.0, 0.0], "Cy5", "cup9-del"
    )
    second = _gsm(
        "GSM2", "cup9-del-b", [8.0, 2.0, 4.0], [4.0, 4.0, 1.0], "Cy3", "cup9-del"
    )
    assert dict(collect([first, second], _PROBES)) == {
        "Q0010": [(0.0, 1.0), (4.0, 1.0)],
        "YAL001C": [(2.0, 4.0), (8.0, 4.0)],
        "YBR001C": [(8.0, 4.0), (2.0, 4.0)],
    }
    assert dict(collect([], _PROBES)) == {}


def test_channel_check_counts_arrays_with_the_deleted_gene_depleted() -> None:
    """For YAL001C (probe 1) on two arrays: 2/4 on the first is depleted, 8/4 on the
    second is not, so the fraction is 0.5; a gene with no probe on the platform, or an
    array with a 0 signal at that probe, is not counted (NaN when none is).
    """
    check = m.MicroarrayKemmeren2014Dataset._validate_channel_assignment
    depleted = _gsm("A", "yal001c-del-a", [4.0, 4.0, 1.0], [2.0, 8.0, 0.0], "Cy5", "x")
    raised = _gsm("B", "yal001c-del-b", [8.0, 2.0, 4.0], [4.0, 4.0, 1.0], "Cy3", "x")
    zero = _gsm("C", "yal001c-del-a", [4.0, 4.0, 1.0], [0.0, 8.0, 0.0], "Cy5", "x")
    assert check({"YAL001C": [depleted, raised, zero]}, _PROBES) == 0.5
    assert check({"YAL001C": [depleted]}, _PROBES) == 1.0
    assert math.isnan(check({"YPL177C": [depleted], "YAL001C": [zero]}, _PROBES))


def test_reference_n_replicates_count_the_arrays_in_the_refpool_mean(
    tmp_path: Path,
) -> None:
    """#484: the reference ``n_replicates`` describes the stored reference value, the
    refpool mean over the mutant's own arrays, so it equals the mutant's own count gene
    by gene (CUP9: Q0010 1 after the 0-signal array is dropped, YAL001C 2, YBR001C 2),
    never a wildtype-series count (the fixture's MATa series would give 3, 2, 2, and the
    real series 28 and 400). With the wildtype pickles removed the build is identical,
    because no stored value reads them.
    """
    root = tmp_path / "kemmeren"
    _write_raw(root / "raw")
    for accession in ("GSE42241", "GSE42240", "GSE42217"):
        (root / "raw" / f"{accession}.pkl").unlink()
    dataset = m.MicroarrayKemmeren2014Dataset(root=str(root), genome=_genome())
    assert len(dataset) == 2
    for i in range(2):
        experiment_n = dataset[i]["experiment"]["phenotype"]["n_replicates"]
        reference = dataset[i]["reference"]["phenotype_reference"]
        assert reference["n_replicates"] == experiment_n
        assert set(reference["n_replicates"]) == set(reference["expression"])
    assert dataset[0]["reference"]["phenotype_reference"]["n_replicates"] == {
        "Q0010": 1,
        "YAL001C": 2,
        "YBR001C": 2,
    }
    assert dataset[0]["reference"] == _REF_CUP9.model_dump()
    assert dataset[1]["reference"] == _REF_HSN1.model_dump()


def test_resolution_passes_in_priority_order(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """Special map (LUG1 -> YCR087C-A even without a strain), Table S1 common name,
    the gene-attribute table's ``ID`` and ``Alias`` columns, a direct Table S1 key, the
    shared reconciler, and ``None`` for a retired name.
    """
    strains = {"YPL177C": "BY4741", "YXX009W": "BY4742"}
    resolve = dataset.resolve_gene_name_comprehensive
    dataset.resolved_by_alias = dataset.resolved_by_excel = 0
    dataset.resolved_by_gene_table = dataset.resolved_by_shared_reconciler = 0
    dataset.unresolved_genes = 0
    assert resolve("lug1", {}, strains) == "YCR087C-A"
    assert resolve("cup9", {"CUP9": "YPL177C"}, strains) == "YPL177C"
    assert resolve("YPL177C", {}, strains) == "YPL177C"
    assert resolve("none1", {}, strains) == "YPL177C"
    assert resolve("yxx009w", {}, strains) == "YXX009W"
    assert resolve("cdk8", {}, strains) == "YPL042C"
    assert resolve("zzz9", {}, strains) is None
    assert (
        dataset.resolved_by_alias,
        dataset.resolved_by_excel,
        dataset.resolved_by_gene_table,
        dataset.resolved_by_shared_reconciler,
        dataset.unresolved_genes,
    ) == (1, 2, 2, 1, 1)


def test_create_expression_experiment_from_pairs() -> None:
    """No strain raises; no pair, or only pairs with a non-positive signal, returns the
    ``(None, None, None)`` skip; pairs (2, 4) and (8, 4) give log2 -1 and +1, mean 0,
    SE 1, variance 2.0000000000000004, n 2, linear 5, refpool 4, and the reference
    ``n_replicates`` is the 2 arrays of that refpool mean.
    """
    build = m.MicroarrayKemmeren2014Dataset.create_expression_experiment
    with pytest.raises(
        ValueError,
        match=re.escape("Strain (BY4741 or BY4742) must be specified in sample_info"),
    ):
        build("d", {"systematic_gene_name": "YAL001C"}, {})
    info = {"systematic_gene_name": "YAL001C", "strain": "BY4741"}
    assert build("d", info, {}) == (None, None, None)
    assert build("d", info, {"YAL001C": [(0.0, 1.0), (1.0, 0.0)]}) == (None, None, None)
    experiment, reference, publication = build(
        "d", info, {"YAL001C": [(2.0, 4.0), (8.0, 4.0)]}
    )
    assert experiment.phenotype.model_dump() == {
        "graph_level": "node",
        "label_name": "expression_log2_ratio",
        "label_statistic_name": "expression_log2_ratio_se",
        "expression": {"YAL001C": 5.0},
        "expression_log2_ratio": {"YAL001C": 0.0},
        "expression_log2_ratio_se": {"YAL001C": 1.0},
        "expression_log2_ratio_variance": {"YAL001C": _SD2**2},
        "n_replicates": {"YAL001C": 2},
        "provenance_gaps": [],
    }
    assert reference.phenotype_reference.model_dump() == {
        "graph_level": "node",
        "label_name": "expression_log2_ratio",
        "label_statistic_name": "expression_log2_ratio_se",
        "expression": {"YAL001C": 4.0},
        "expression_log2_ratio": {"YAL001C": 0.0},
        "expression_log2_ratio_se": None,
        "expression_log2_ratio_variance": None,
        "n_replicates": {"YAL001C": 2},
        "provenance_gaps": [],
    }
    assert reference.genome_reference.strain == "BY4741"
    assert publication == _PUBLICATION


def test_parallel_build_writes_the_same_records(tmp_path: Path) -> None:
    """``process_workers=1`` routes through ``_process_parallel`` and the static batch
    helpers (one worker process, batches of one gene); the two stored records, both
    references and the publication equal the sequential build's.
    """
    root = tmp_path / "kemmeren"
    _write_raw(root / "raw")
    dataset = m.MicroarrayKemmeren2014Dataset(
        root=str(root), genome=_genome(), process_workers=1, batch_size=1
    )
    assert len(dataset) == 2
    assert _nan_safe(dataset[0]["experiment"]) == _nan_safe(_CUP9.model_dump())
    assert _nan_safe(dataset[1]["experiment"]) == _nan_safe(_HSN1.model_dump())
    assert dataset[0]["reference"] == _REF_CUP9.model_dump()
    assert dataset[1]["reference"] == _REF_HSN1.model_dump()
    assert dataset[0]["publication"] == _PUBLICATION.model_dump()


# ---------------------------------------------------------------------------
# 2026.10.06 (Phase 21): download, the batch path, resolution and Table S1 branches.
# ---------------------------------------------------------------------------

_TABLE_S1_URL = "https://uofi.box.com/shared/static/9n6ruj58ueup0cebhnek8ijdcy4om0bi"
_ALL_SERIES = ["GSE42527", "GSE42526", "GSE42241", "GSE42240", "GSE42217", "GSE42215"]


class _FakeResponse:
    def __init__(self, content: bytes, status_error: Exception | None = None) -> None:
        self.content = content
        self.status_error = status_error

    def raise_for_status(self) -> None:
        if self.status_error is not None:
            raise self.status_error


class _FakeGet:
    """Stand-in for ``requests.get`` recording ``(url, kwargs)``."""

    def __init__(self, response: _FakeResponse) -> None:
        self.response = response
        self.calls: list[tuple[str, dict[str, Any]]] = []

    def __call__(self, url: str, **kwargs: Any) -> _FakeResponse:
        self.calls.append((url, kwargs))
        return self.response


class _FakeGeo:
    """Stand-in for ``GEOparse.get_GEO``; fails on the accession named in ``fail``."""

    def __init__(self, fail: str | None = None) -> None:
        self.fail = fail
        self.calls: list[dict[str, Any]] = []

    def __call__(self, **kwargs: Any) -> dict[str, str]:
        self.calls.append(kwargs)
        if kwargs["geo"] == self.fail:
            raise OSError("offline")
        return {"accession": kwargs["geo"]}


def _empty_raw(dataset: m.MicroarrayKemmeren2014Dataset) -> Path:
    raw = Path(dataset.raw_dir)
    for child in raw.iterdir():
        child.unlink()
    return raw


_WORKBOOK_BYTES = b"PK" + b"x" * 1200


def test_download_writes_table_s1_and_six_geo_pickles(
    dataset: m.MicroarrayKemmeren2014Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Finding: ``download`` (kemmeren2014.py lines 204-285) fetches Table S1 from a
    personal ``uofi.box.com`` share link while ``_load_mating_type_map`` tells the user
    the source is Cell's ``mmc1.xlsx`` (line 968), and neither the workbook nor the six
    GEO pickles get a sha256 or a retrieval record. Under the stub (which writes no
    SOFT file) ``raw/`` holds exactly the seven files the loader itself writes; a real
    ``get_GEO`` also leaves each ``*_family.soft.gz`` there (the real raw dir holds 13
    files), none of them recorded either. ``requests.get`` gets the Box URL with ``timeout=60``; ``get_GEO`` gets the
    two deletion series then the four wildtype series, each with ``destdir=<raw_dir>``
    and ``silent=False``. Pinned until the workbook is fetched from the journal SI and
    every file is recorded with its sha256.
    """
    raw = _empty_raw(dataset)
    get = _FakeGet(_FakeResponse(_WORKBOOK_BYTES))
    geo = _FakeGeo()
    monkeypatch.setattr(requests, "get", get)
    monkeypatch.setattr(GEOparse, "get_GEO", geo)
    dataset.download()
    assert get.calls == [(_TABLE_S1_URL, {"timeout": 60})]
    assert geo.calls == [
        {"geo": accession, "destdir": str(raw), "silent": False}
        for accession in _ALL_SERIES
    ]
    assert (raw / "kemmeren2014_table_s1.xlsx").read_bytes() == _WORKBOOK_BYTES
    assert sorted(p.name for p in raw.iterdir()) == sorted(
        ["kemmeren2014_table_s1.xlsx"] + [f"{a}.pkl" for a in _ALL_SERIES]
    )
    for accession in _ALL_SERIES:
        with open(raw / f"{accession}.pkl", "rb") as handle:
            assert pickle.load(handle) == {"accession": accession}


def test_download_keeps_an_existing_table_s1(
    dataset: m.MicroarrayKemmeren2014Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A workbook already in ``raw/`` is not fetched again (``requests.get`` is never
    called) and its bytes are unchanged; the GEO series are still fetched.
    """
    raw = Path(dataset.raw_dir)
    before = (raw / "kemmeren2014_table_s1.xlsx").read_bytes()
    get = _FakeGet(_FakeResponse(_WORKBOOK_BYTES))
    geo = _FakeGeo()
    monkeypatch.setattr(requests, "get", get)
    monkeypatch.setattr(GEOparse, "get_GEO", geo)
    dataset.download()
    assert get.calls == []
    assert [call["geo"] for call in geo.calls] == _ALL_SERIES
    assert (raw / "kemmeren2014_table_s1.xlsx").read_bytes() == before


@pytest.mark.parametrize(
    ("response", "error_text"),
    [
        (
            _FakeResponse(b"small"),
            "Downloaded file too small (5 bytes), likely not the Excel file",
        ),
        (
            _FakeResponse(b"x" * 999),
            "Downloaded file too small (999 bytes), likely not the Excel file",
        ),
        (
            _FakeResponse(_WORKBOOK_BYTES, status_error=OSError("403 Forbidden")),
            "403 Forbidden",
        ),
    ],
    ids=["too-small", "one-byte-short", "http-error"],
)
def test_table_s1_download_refusals(
    response: _FakeResponse,
    error_text: str,
    dataset: m.MicroarrayKemmeren2014Dataset,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A body under 1000 bytes (5 or 999), or an HTTP error, raises ``RuntimeError``
    with the URL, the cause and the manual-save path (full match), writes no workbook
    and fetches no GEO series.
    """
    raw = _empty_raw(dataset)
    geo = _FakeGeo()
    monkeypatch.setattr(requests, "get", _FakeGet(response))
    monkeypatch.setattr(GEOparse, "get_GEO", geo)
    table = raw / "kemmeren2014_table_s1.xlsx"
    message = (
        f"Failed to download Table S1 from {_TABLE_S1_URL}\nError: {error_text}\n"
        f"Please check the URL or save manually as: {table}"
    )
    with pytest.raises(RuntimeError, match=f"^{re.escape(message)}$"):
        dataset.download()
    assert list(raw.iterdir()) == []
    assert geo.calls == []


def test_a_table_s1_body_of_exactly_1000_bytes_is_accepted(
    dataset: m.MicroarrayKemmeren2014Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The size check is ``len(content) < 1000``, so a 1000-byte body is written
    verbatim and all six GEO series are then fetched.
    """
    raw = _empty_raw(dataset)
    body = b"y" * 1000
    geo = _FakeGeo()
    monkeypatch.setattr(requests, "get", _FakeGet(_FakeResponse(body)))
    monkeypatch.setattr(GEOparse, "get_GEO", geo)
    dataset.download()
    assert (raw / "kemmeren2014_table_s1.xlsx").read_bytes() == body
    assert [call["geo"] for call in geo.calls] == _ALL_SERIES


@pytest.mark.parametrize(
    ("fail", "written"),
    [
        ("GSE42526", ["GSE42527.pkl"]),
        ("GSE42217", ["GSE42240.pkl", "GSE42241.pkl", "GSE42526.pkl", "GSE42527.pkl"]),
    ],
    ids=["deletion-series", "wildtype-series"],
)
def test_geo_download_refusal_names_the_failing_series(
    fail: str,
    written: list[str],
    dataset: m.MicroarrayKemmeren2014Dataset,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A ``get_GEO`` failure raises ``Failed to download <accession> from GEO`` for
    the series that failed, in the deletion loop and in the wildtype loop alike; the
    series fetched before it stay pickled and the ones after are not attempted.
    """
    raw = Path(dataset.raw_dir)
    for child in raw.glob("*.pkl"):
        child.unlink()
    geo = _FakeGeo(fail=fail)
    monkeypatch.setattr(GEOparse, "get_GEO", geo)
    with pytest.raises(
        RuntimeError, match=f"^{re.escape(f'Failed to download {fail} from GEO')}$"
    ):
        dataset.download()
    assert [call["geo"] for call in geo.calls] == _ALL_SERIES[
        : _ALL_SERIES.index(fail) + 1
    ]
    assert sorted(p.name for p in raw.glob("*.pkl")) == written


def test_process_refetches_a_missing_deletion_pickle(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Without ``GSE42526.pkl``, ``process`` calls ``get_GEO(geo="GSE42526",
    destdir=<raw_dir>, silent=False)`` once (line 315) and builds the same two records
    from the object it returns.
    """
    root = tmp_path / "kemmeren"
    _write_raw(root / "raw")
    with open(root / "raw" / "GSE42526.pkl", "rb") as handle:
        gse = pickle.load(handle)
    (root / "raw" / "GSE42526.pkl").unlink()
    calls: list[dict[str, Any]] = []

    def get_geo(**kwargs: Any) -> Any:
        calls.append(kwargs)
        return gse

    monkeypatch.setattr(GEOparse, "get_GEO", get_geo)
    rebuilt = m.MicroarrayKemmeren2014Dataset(root=str(root), genome=_genome())
    assert calls == [{"geo": "GSE42526", "destdir": str(root / "raw"), "silent": False}]
    assert len(rebuilt) == 2
    assert _nan_safe(rebuilt[0]["experiment"]) == _nan_safe(_CUP9.model_dump())
    assert _nan_safe(rebuilt[1]["experiment"]) == _nan_safe(_HSN1.model_dump())


def _raw_gsms(dataset: m.MicroarrayKemmeren2014Dataset) -> dict[str, GSM]:
    gsms: dict[str, GSM] = {}
    for accession in ("GSE42527", "GSE42526"):
        with open(Path(dataset.raw_dir) / f"{accession}.pkl", "rb") as handle:
            gsms.update(pickle.load(handle).gsms)
    return gsms


def test_process_batch_matches_process_sequential_record_by_record(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """Five genes in order: YPL177C (GSM1 + GSM2) -> ``_CUP9``; YPL042C (GSM5) has no
    Table S1 strain and is skipped; YAL001C with no arrays has no pairs and is
    skipped; YHR127W (GSM3) -> ``_HSN1``; YNCB0010W on one array whose deletion
    channel is all 0 has only non-positive pairs, so ``create_expression_experiment``
    returns the skip triple. The static batch path (lines 716-769, only ever run in
    worker processes) and the sequential path write the same two records in order.
    """
    gsms = _raw_gsms(dataset)
    zero = _gsm("Z", "tlc1-del-a", [1.0, 1.0, 1.0], [0.0, 0.0, 0.0], "Cy5", "tlc1-del")
    groups = {
        "YPL177C": [gsms["GSM1"], gsms["GSM2"]],
        "YPL042C": [gsms["GSM5"]],
        "YAL001C": [],
        "YHR127W": [gsms["GSM3"]],
        "YNCB0010W": [zero],
    }
    strains, _ = dataset._load_mating_type_map()
    batch = m.MicroarrayKemmeren2014Dataset._process_batch(
        list(groups.items()), _PROBES, strains, _DATASET
    )
    records = [pickle.loads(blob) for blob in batch]
    expected = [
        (_CUP9.model_dump(), _REF_CUP9.model_dump()),
        (_HSN1.model_dump(), _REF_HSN1.model_dump()),
    ]
    assert len(records) == 2
    for record, (experiment, reference) in zip(records, expected, strict=True):
        assert _nan_safe(record["experiment"]) == _nan_safe(experiment)
        assert record["reference"] == reference
        assert record["publication"] == _PUBLICATION.model_dump()

    dataset.close_lmdb()
    lmdb_dir = Path(dataset.processed_dir) / "lmdb"
    for child in lmdb_dir.iterdir():
        child.unlink()
    dataset._process_sequential(groups, _PROBES, strains)
    env = lmdb.open(str(lmdb_dir), readonly=True, lock=False)
    with env.begin() as txn:
        sequential = [pickle.loads(value) for _, value in txn.cursor()]
    env.close()
    assert [_nan_safe(r) for r in sequential] == [_nan_safe(r) for r in records]


def test_probe_mapping_branches(dataset: m.MicroarrayKemmeren2014Dataset) -> None:
    """Finding: the three "Clean and validate gene name" arms (kemmeren2014.py lines
    923-937) all store ``gene_name.upper()``, so a control probe ``Empty`` maps to the
    gene ``EMPTY`` and a ``None`` cell to ``NONE``; only a float NaN (``"nan"``) is
    dropped. ``EMPTY`` and ``NONE`` do not occur on the real GPL11232, but the same
    ``.upper()`` path stores ``SNR10``, a non-systematic name, as an expression key in
    all 1484 served Kemmeren records (audit 1, 2026.10.06). ``SPOT`` is accepted as the
    ID column and ``Gene`` as the gene column; a GSE without platforms and a platform
    without an ID or gene column map nothing. Pinned until the arms reject names that
    are not ORFs.
    """

    def platform(table: pd.DataFrame) -> GSE:
        gpl = GPL(name="GPL", metadata={}, table=table, columns=_describe(table))
        return GSE(name="G", metadata={}, gpls={"GPL": gpl}, gsms={})

    spot = pd.DataFrame(
        {
            "SPOT": [1.0, 2.0, 3.0, 4.0, 5.0],
            "Gene": ["yal001c", "q0010", "Empty", None, float("nan")],
        }
    )
    mapping = dataset._extract_probe_to_gene_mapping
    assert mapping(platform(spot)) == {
        "1": "YAL001C",
        "2": "Q0010",
        "3": "EMPTY",
        "4": "NONE",
    }
    assert mapping(GSE(name="G", metadata={}, gpls={}, gsms={})) == {}
    assert mapping(platform(pd.DataFrame({"ID": [1], "Description": ["x"]}))) == {}
    assert mapping(platform(pd.DataFrame({"NAME": [1], "ORF": ["YAL001C"]}))) == {}


def _write_sheet(path: Path, rows: list[list[Any]]) -> None:
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)


def test_mating_type_map_spellings_duplicates_and_blank_cells(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """``MATα`` (upper-cased to Greek ``MATΑ``) and ``mat alpha`` map to BY4742,
    ``mata`` to BY4741, ``MATa/alpha`` (both) is skipped as unknown, a blank mating
    type or a blank orf skips the row. A repeated orf keeps BOTH common names and the
    LAST row's strain (YAL001C: MATa then MATalpha -> BY4742).
    """
    _write_sheet(
        Path(dataset.raw_dir) / "kemmeren2014_table_s1.xlsx",
        [
            ["ORF Name", "Gene", "Mating Type"],
            ["yal001c", "ONE1", "mata"],
            ["YAL001C", "ONE2", "MATα"],
            ["YBR001C", "TWO1", "mat alpha"],
            ["YCR001W", "THR1", "MATa/alpha"],
            ["YDR001C", "FOR1", None],
            [None, "FIV1", "MATa"],
        ],
    )
    assert dataset._load_mating_type_map() == (
        {"YAL001C": "BY4742", "YBR001C": "BY4742"},
        {"ONE1": "YAL001C", "ONE2": "YAL001C", "TWO1": "YBR001C"},
    )


def test_mating_type_map_without_gene_or_mating_column(
    dataset: m.MicroarrayKemmeren2014Dataset,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Without a ``gene`` column the common-name map holds only the TLC1 and CMS1
    entries the loader adds itself.

    Finding: without a mating-type column (kemmeren2014.py lines 1108-1114) the
    function only logs an error and returns two empty maps, and a missing ``orf name``
    column raises ``ValueError`` at line 1007 only to be swallowed by the broad
    ``except Exception`` at line 1116 (logged, then ``({}, {})``). The build does not
    end quietly, though: with no strain map every gene is skipped, the LMDB holds 0
    entries, and ``post_process`` then refuses the empty gene set with ``ValueError
    ("Cannot set an empty gene_set: every record of a dataset with
    has_gene_perturbations=True carries gene perturbations, so an empty set is a
    broken build")`` (experiment_dataset.py line 850), a message that does not name
    Table S1. Pinned until a Table S1 missing a
    required column raises at load time naming the column.
    """
    table = Path(dataset.raw_dir) / "kemmeren2014_table_s1.xlsx"
    _write_sheet(
        table, [["orf name", "mating type"], ["TLC1", "MATa"], ["YAL001C", "MATa"]]
    )
    assert dataset._load_mating_type_map() == (
        {"YNCB0010W": "BY4741", "YAL001C": "BY4741"},
        {"TLC1": "YNCB0010W"},
    )
    _write_sheet(table, [["orf name", "gene"], ["YAL001C", "ONE1"]])
    assert dataset._load_mating_type_map() == ({}, {})
    _write_sheet(table, [["gene", "mating type"], ["CUP9", "MATa"]])
    caplog.clear()
    with caplog.at_level("ERROR", logger=m.log.name):
        assert dataset._load_mating_type_map() == ({}, {})
    assert [r.getMessage() for r in caplog.records if r.name == m.log.name] == [
        "Failed to load mating type map: Required 'orf name' column not found in "
        "Excel file!"
    ]

    root = tmp_path / "no_mating"
    _write_raw(root / "raw")
    _write_sheet(
        root / "raw" / "kemmeren2014_table_s1.xlsx",
        [["orf name", "gene"], ["YPL177C", "CUP9"], ["YHR127W", "HSN1"]],
    )
    with pytest.raises(
        ValueError,
        match=(
            r"^Cannot set an empty gene_set: every record of a dataset with "
            r"has_gene_perturbations=True carries gene perturbations, so an empty "
            r"set is a broken build$"
        ),
    ):
        m.MicroarrayKemmeren2014Dataset(root=str(root), genome=_genome())


class _ResolveGenome:
    """Gene table, alias map and reconciler for every resolution pass."""

    gene_attribute_table = pd.DataFrame(
        {
            "ID": ["YAA001W", "YBB001W", "YCC001W", "YDD001W"],
            "gene": ["GENEA", "GENEB", "GENEC", None],
            "Alias": [None, "ALIASB", None, "ALIASD"],
        }
    )
    alias_to_systematic: dict[str, list[str]] = {
        "ONEHIT": ["YEE001W", "YZZ999W"],
        "TWOHIT": ["YGG001W", "YFF001W"],
        "CycC": ["YHH001W"],
        "CycD": ["YII001W", "YJJ001W"],
    }

    def __init__(self) -> None:
        self.reconciled: list[str] = []

    def resolve_gene_name(self, name: str) -> GeneNameResolution:
        self.reconciled.append(name)
        if name.upper() == "RENAMED1":
            return GeneNameResolution(
                input_name=name,
                status=GeneNameStatus.RENAMED,
                systematic_name="YKK001W",
            )
        return GeneNameResolution(
            input_name=name, status=GeneNameStatus.AMBIGUOUS, systematic_name="YLL001W"
        )


_STRAINS = dict.fromkeys(
    [
        "YAA001W",
        "YBB001W",
        "YDD001W",
        "YEE001W",
        "YFF001W",
        "YGG001W",
        "YHH001W",
        "YII001W",
        "YJJ001W",
    ],
    "BY4741",
)


def test_resolution_covers_every_pass_with_exact_counters(
    dataset: m.MicroarrayKemmeren2014Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Each name and the pass that resolves it:

    - ``tlc1``: special map, YNCB0010W (not a strain key, returned anyway, alias +1).
    - ``yaa001w``: table ``ID`` and a strain key (gene table +1).
    - ``YCC001W``: in the table ``ID`` but no strain, then no other hit; reconciler is
      ``AMBIGUOUS``, so ``None`` (unresolved +1).
    - ``geneb``: table ``gene`` -> YBB001W (gene table +1); ``genec`` -> YCC001W has no
      strain and falls through to the reconciler (``AMBIGUOUS``, unresolved +1).
    - ``aliasd``: table ``Alias`` -> YDD001W (gene table +1).
    - ``onehit``: alias map, one candidate with a strain -> YEE001W (alias +1).
    - ``twohit``: two candidates with strains -> the sorted first, YFF001W (alias +1).
    - ``cycc`` / ``cycd``: no upper-case key; the case-insensitive pass finds ``CycC``
      (one candidate, YHH001W) and ``CycD`` (two, sorted first YII001W), alias +2.
    - ``renamed1``: the reconciler's ``RENAMED`` -> YKK001W (reconciler +1).

    Totals: excel 0, gene table 3, alias 5, reconciler 1, unresolved 2. The reconciler
    is called with the name as given, not upper-cased.
    """
    genome = _ResolveGenome()
    monkeypatch.setattr(dataset, "genome", genome)
    dataset.resolved_by_alias = dataset.resolved_by_excel = 0
    dataset.resolved_by_gene_table = dataset.resolved_by_shared_reconciler = 0
    dataset.unresolved_genes = 0
    resolve = dataset.resolve_gene_name_comprehensive
    names = [
        "tlc1",
        "yaa001w",
        "YCC001W",
        "geneb",
        "genec",
        "aliasd",
        "onehit",
        "twohit",
        "cycc",
        "cycd",
        "renamed1",
    ]
    assert [resolve(name, {}, _STRAINS) for name in names] == [
        "YNCB0010W",
        "YAA001W",
        None,
        "YBB001W",
        None,
        "YDD001W",
        "YEE001W",
        "YFF001W",
        "YHH001W",
        "YII001W",
        "YKK001W",
    ]
    assert (
        dataset.resolved_by_excel,
        dataset.resolved_by_gene_table,
        dataset.resolved_by_alias,
        dataset.resolved_by_shared_reconciler,
        dataset.unresolved_genes,
    ) == (0, 3, 5, 1, 2)
    assert genome.reconciled == ["YCC001W", "genec", "renamed1"]


def test_resolution_with_a_genome_that_has_no_tables(
    dataset: m.MicroarrayKemmeren2014Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    """With a genome object carrying none of the three attributes, only the special
    map, the Excel common-name map and a direct strain key resolve; anything else is
    ``None`` without calling any reconciler.
    """
    monkeypatch.setattr(dataset, "genome", object())
    dataset.resolved_by_alias = dataset.resolved_by_excel = 0
    dataset.resolved_by_gene_table = dataset.resolved_by_shared_reconciler = 0
    dataset.unresolved_genes = 0
    resolve = dataset.resolve_gene_name_comprehensive
    assert resolve("cms1", {}, {}) == "YLR003C"
    assert resolve("cup9", {"CUP9": "YPL177C"}, {}) == "YPL177C"
    assert resolve("yaa001w", {}, {"YAA001W": "BY4741"}) == "YAA001W"
    assert resolve("geneb", {}, {}) is None
    assert (
        dataset.resolved_by_alias,
        dataset.resolved_by_excel,
        dataset.unresolved_genes,
    ) == (1, 2, 1)


def test_already_assigned_is_accepted_and_ignored(
    dataset: m.MicroarrayKemmeren2014Dataset,
) -> None:
    """Finding: ``resolve_gene_name_comprehensive`` takes ``already_assigned`` and
    ``process`` passes the set of ORFs it has grouped so far (lines 344, 379, 409, 439),
    but the function never reads it past defaulting it (lines 1139-1140): an ORF already in the
    set is returned again, so two differently named titles that resolve to one ORF are
    pooled into one record without any check. Pinned until the set is either used or
    removed.
    """
    resolve = dataset.resolve_gene_name_comprehensive
    common = {"CUP9": "YPL177C"}
    assert resolve("cup9", common, {}, {"YPL177C"}) == "YPL177C"
    assert resolve("cup9", common, {}, set()) == "YPL177C"
    assert resolve("cup9", common, {}, None) == "YPL177C"


def test_convert_gene_name_branches(
    dataset: m.MicroarrayKemmeren2014Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Finding: ``convert_gene_name`` (lines 1266-1295) has no caller in ``torchcell``
    (its docstring says it is "Used for probe-to-gene mappings"; the probe map is built
    by ``_extract_probe_to_gene_mapping`` without it). Its contract: the Excel map,
    then the table ``gene`` column, then the ``Alias`` column, else the input
    UNCHANGED (case kept: ``mixedCase`` stays ``mixedCase``). Pinned until it is called
    or removed.
    """
    monkeypatch.setattr(dataset, "genome", _ResolveGenome())
    convert = dataset.convert_gene_name
    assert convert("cup9", {"CUP9": "YPL177C"}) == "YPL177C"
    assert convert("genea", {}) == "YAA001W"
    assert convert("aliasb", {}) == "YBB001W"
    assert convert("mixedCase", {}) == "mixedCase"
    monkeypatch.setattr(dataset, "genome", object())
    assert convert("genea", {}) == "genea"


def test_processing_summary_cannot_report_duplicate_deletions(
    dataset: m.MicroarrayKemmeren2014Dataset, caplog: pytest.LogCaptureFixture
) -> None:
    """Finding: ``_log_processing_summary`` counts the KEYS of a dict (line 1320), so
    every count is 1 and the "duplicate gene deletions" branch (lines 1324-1326) can
    never run, however many arrays a gene has. The logged lines are exact. Pinned
    until it counts arrays per gene or the branch is removed.
    """
    samples = [
        {"is_deletion": True, "is_wildtype": False},
        {"is_deletion": True, "is_wildtype": False},
        {"is_deletion": False, "is_wildtype": True},
    ]
    with caplog.at_level("INFO", logger=m.log.name):
        dataset._log_processing_summary({"YAL001C": ["a", "b"]}, samples)
    assert [r.getMessage() for r in caplog.records if r.name == m.log.name] == [
        "Processed 1 unique gene deletion experiments",
        "Total samples: 3, Deletion samples: 2, Wildtype: 1",
        "Unique gene deletions: 1",
    ]


def test_process_reads_genes_from_characteristics_and_skips_unresolved(
    tmp_path: Path,
) -> None:
    """A build whose Table S1 matches the resolved ORFs exactly: four arrays titled
    ``Sample A..D`` name CUP9 only as ``genotype/variation: [HS1991] cup9-del`` (the
    ``[HS1991]`` prefix is stripped, line 403); ``hsn1-del-a`` resolves through the
    special map; ``Sample Z`` names ``zzz9-del`` in its characteristics, which does
    not resolve, so it is a deletion with no ORF and makes no record.

    CUP9: four identical arrays, deletion 2 and refpool 1 -> log2 1, sample SD 0, so
    SE 0 and variance 0, n 4, linear 2, refpool 1. HSN1: deletion 4, refpool 1 -> log2
    2, n 1, NaN SE and variance.
    """
    raw = tmp_path / "k2" / "raw"
    raw.mkdir(parents=True)
    for accession in _ALL_SERIES:
        (raw / f"{accession}_family.soft.gz").write_bytes(b"placeholder")
    _write_sheet(
        raw / "kemmeren2014_table_s1.xlsx",
        [
            ["orf name", "gene", "mating type"],
            ["YPL177C", "CUP9", "MATa"],
            ["YHR127W", "HSN1", "MATalpha"],
        ],
    )
    cup9 = [
        _gsm(
            f"C{i}",
            f"Sample {letter}",
            [1.0, 1.0, 1.0],
            [2.0, 2.0, 2.0],
            "Cy5",
            "cup9-del",
            characteristics=["genotype/variation: [HS1991] cup9-del"],
        )
        for i, letter in enumerate("ABCD")
    ]
    hsn1 = _gsm("H1", "hsn1-del-a", [1.0, 1.0, 1.0], [4.0, 4.0, 4.0], "Cy5", "hsn1-del")
    unresolved = _gsm(
        "Z1",
        "Sample Z",
        [1.0, 1.0, 1.0],
        [1.0, 1.0, 1.0],
        "Cy5",
        "zzz9-del",
        characteristics=["genotype/variation: zzz9-del"],
    )
    for accession, gsms in (("GSE42527", [*cup9, hsn1]), ("GSE42526", [unresolved])):
        with open(raw / f"{accession}.pkl", "wb") as handle:
            pickle.dump(_gse(accession, gsms), handle)
    built = m.MicroarrayKemmeren2014Dataset(root=str(raw.parent), genome=_genome())
    genes = ("Q0010", "YAL001C", "YBR001C")
    cup9_expected = _experiment(
        "YPL177C",
        dict.fromkeys(genes, 1.0),
        dict.fromkeys(genes, 0.0),
        dict.fromkeys(genes, 0.0),
        dict.fromkeys(genes, 4),
        dict.fromkeys(genes, 2.0),
    )
    hsn1_expected = _experiment(
        "YHR127W",
        dict.fromkeys(genes, 2.0),
        dict.fromkeys(genes, _NANF),
        dict.fromkeys(genes, _NANF),
        dict.fromkeys(genes, 1),
        dict.fromkeys(genes, 4.0),
    )
    assert len(built) == 2
    assert _nan_safe(built[0]["experiment"]) == _nan_safe(cup9_expected.model_dump())
    assert _nan_safe(built[1]["experiment"]) == _nan_safe(hsn1_expected.model_dump())
    assert (
        built[0]["reference"]
        == _reference(
            "BY4741", dict.fromkeys(genes, 1.0), dict.fromkeys(genes, 4)
        ).model_dump()
    )
    samples = pd.read_csv(raw.parent / "preprocess" / "data.csv", keep_default_na=False)
    assert samples[
        ["geo_accession", "systematic_gene_name", "is_deletion", "is_wildtype"]
    ].to_dict(orient="list") == {
        "geo_accession": ["C0", "C1", "C2", "C3", "H1", "Z1"],
        "systematic_gene_name": ["YPL177C"] * 4 + ["YHR127W", ""],
        "is_deletion": [True] * 6,
        "is_wildtype": [False] * 6,
    }


def test_parallel_path_logs_written_records_as_attempted_genes(
    dataset: m.MicroarrayKemmeren2014Dataset, caplog: pytest.LogCaptureFixture
) -> None:
    """Finding: ``_process_batch`` drops a skipped gene instead of returning ``None``
    for it, so in ``_process_parallel`` (lines 693-713) every result is bytes: the
    ``skipped_genes`` branch can never run and "Total gene deletions attempted" logs
    the WRITTEN count. Five genes go in (the five-group set of the batch test), two
    records are written, and the log says 2 attempted with no skip warning, where
    ``_process_sequential`` says 5 attempted and warns of 3 skipped. Reach: the served
    build takes the sequential path (``build_dataset_lmdb`` passes no
    ``process_workers``); only ``experiments/012-sameith-kemmeren/scripts/
    kemmeren_volcano.py`` (``process_workers=10``) reaches this log line, and no record
    differs. Pinned until the batch path reports its skips.
    """
    gsms = _raw_gsms(dataset)
    zero = _gsm("Z", "tlc1-del-a", [1.0, 1.0, 1.0], [0.0, 0.0, 0.0], "Cy5", "tlc1-del")
    groups = {
        "YPL177C": [gsms["GSM1"], gsms["GSM2"]],
        "YPL042C": [gsms["GSM5"]],
        "YAL001C": [],
        "YHR127W": [gsms["GSM3"]],
        "YNCB0010W": [zero],
    }
    strains, _ = dataset._load_mating_type_map()
    dataset.close_lmdb()
    dataset.process_workers = 1
    dataset.batch_size = 2
    lmdb_dir = Path(dataset.processed_dir) / "lmdb"

    def messages(run: Any) -> list[str]:
        for child in lmdb_dir.iterdir():
            child.unlink()
        caplog.clear()
        with caplog.at_level("INFO", logger=m.log.name):
            run(groups, _PROBES, strains)
        return [
            r.getMessage()
            for r in caplog.records
            if r.name == m.log.name
            and r.getMessage().startswith(("Wrote", "Total", "Skipped"))
        ]

    assert messages(dataset._process_parallel) == [
        "Wrote 2 experiments to LMDB",
        "Total gene deletions attempted: 2",
    ]
    assert messages(dataset._process_sequential) == [
        "Wrote 2 experiments to LMDB",
        "Total gene deletions attempted: 5",
        "Skipped 3 genes (could not calculate log2 ratios)",
    ]


def test_extract_channels_skips_a_non_numeric_cell() -> None:
    """A row whose signal cell is not a number is skipped by the ``ValueError`` guard;
    the other rows keep their (test, reference) pair.
    """
    gsm = _gsm("N", "cup9-del-a", [4.0, 4.0, 1.0], [2.0, 8.0, 0.0], "Cy5", "cup9-del")
    gsm.table["Signal Norm_Cy3"] = gsm.table["Signal Norm_Cy3"].astype(object)
    gsm.table.loc[1, "Signal Norm_Cy3"] = "n/a"
    deletion, refpool = (
        m.MicroarrayKemmeren2014Dataset._extract_channels_from_gsm_static(gsm, _PROBES)
    )
    assert dict(deletion) == {"Q0010": 0.0, "YAL001C": 2.0}
    assert dict(refpool) == {"Q0010": 1.0, "YAL001C": 4.0}


# Phase 24: the parallel writer's skipped-gene accounting and the summary log


class _Done:
    def __init__(self, value: Any) -> None:
        self._value = value

    def result(self) -> Any:
        return self._value


class _InlineExecutor:
    """Runs each submitted batch at once in this process (no pickling, no workers)."""

    def __init__(self, max_workers: int) -> None:
        self.max_workers = max_workers

    def __enter__(self) -> _InlineExecutor:
        return self

    def __exit__(self, *exc: object) -> None:
        return None

    def submit(self, fn: Any, *args: Any) -> _Done:
        return _Done(fn(*args))


def test_parallel_writer_packs_written_records_and_logs_the_skipped_indices(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """Four genes in batches of 2; the batch function returns ``None`` for the second
    gene of each batch, so ``all_results`` is ``[b"r0", None, b"r2", None]``.

    Written records are renumbered densely (keys ``0`` and ``1`` hold ``r0`` and ``r2``)
    and the warnings name the 2 skipped genes and their result indices ``[1, 3]``.
    ``_process_parallel`` runs unbound on a namespace carrying the four attributes it
    reads; ``ProcessPoolExecutor`` is an inline executor on the module.
    """
    import logging
    from types import SimpleNamespace

    monkeypatch.setattr(m, "ProcessPoolExecutor", _InlineExecutor)
    calls: list[tuple[list[str], str]] = []

    def batch(items: list[Any], probes: Any, strains: Any, name: str) -> list[Any]:
        calls.append(([gene for gene, _ in items], name))
        return [
            f"r{int(gene[1:])}".encode() if gene in ("g0", "g2") else None
            for gene, _ in items
        ]

    fake = SimpleNamespace(
        batch_size=2,
        process_workers=1,
        _process_batch=batch,
        processed_dir=str(tmp_path),
        name="MicroarrayKemmeren2014Dataset",
    )
    genes: dict[str, list[Any]] = {f"g{i}": [] for i in range(4)}
    with caplog.at_level(logging.INFO, logger=m.__name__):
        m.MicroarrayKemmeren2014Dataset._process_parallel(
            cast(Any, fake), genes, {}, {}
        )
    assert calls == [
        (["g0", "g1"], "MicroarrayKemmeren2014Dataset"),
        (["g2", "g3"], "MicroarrayKemmeren2014Dataset"),
    ]
    env = lmdb.open(str(tmp_path / "lmdb"), readonly=True, lock=False)
    with env.begin() as txn:
        stored = dict(txn.cursor())
    env.close()
    assert stored == {b"0": b"r0", b"1": b"r2"}
    messages = [r.getMessage() for r in caplog.records if r.name == m.__name__]
    assert messages == [
        "Created 2 batches of size 2",
        "Wrote 2 experiments to LMDB",
        "Total gene deletions attempted: 4",
        "Skipped 2 genes (could not calculate log2 ratios)",
        "Indices of skipped genes: [1, 3]",
    ]


def test_processing_summary_on_a_dict_never_reports_duplicates(
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Finding: ``_log_processing_summary`` counts ``Counter(dict.keys())``; dict keys
    are unique, so every count is 1 and the ``Found N duplicate gene deletions`` lines
    (kemmeren2014.py 1324 to 1326) cannot fire on the dict the signature takes. Pinned:
    the exact summary for two genes over three samples (2 deletion, 1 wild type).
    """
    import logging

    samples = [
        {"is_deletion": True, "is_wildtype": False},
        {"is_deletion": True, "is_wildtype": False},
        {"is_deletion": False, "is_wildtype": True},
    ]
    with caplog.at_level(logging.INFO, logger=m.__name__):
        m.MicroarrayKemmeren2014Dataset._log_processing_summary(
            cast(Any, None), {"YAL001C": [], "YBR001C": []}, samples
        )
    messages = [r.getMessage() for r in caplog.records if r.name == m.__name__]
    assert messages == [
        "Processed 2 unique gene deletion experiments",
        "Total samples: 3, Deletion samples: 2, Wildtype: 1",
        "Unique gene deletions: 2",
    ]
