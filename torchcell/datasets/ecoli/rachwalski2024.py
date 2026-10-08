# torchcell/datasets/ecoli/rachwalski2024
# [[torchcell.datasets.ecoli.rachwalski2024]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/rachwalski2024
# Test file: tests/torchcell/datasets/ecoli/test_rachwalski2024.py
r"""Rachwalski 2024 mobile CRISPRi collection crossed into E. coli deletion backgrounds.

Rachwalski, Tu, Madden, French, Hansen and Brown 2024 (Cell Reports Methods 4:100693,
doi:10.1016/j.crmeth.2023.100693; citation_key ``rachwalskiMobileCRISPRiCollection2024``)
built a 384-well arrayed collection of conjugative CRISPRi plasmids (pFD152, Addgene
125546) targeting the essential genes of *E. coli*, conjugated it into BW25113 and into
BW25113 ``lpp``, and separately conjugated three single knockdowns into the whole Keio
deletion collection. :class:`CrispriCrossRachwalski2024Dataset` serves the three released
growth tables as ``BacterialFitnessExperiment`` records.

RECORD = one (strain, medium, inducer dose) normalized colony growth:

- GENOTYPE: a ``BacterialCrisprInterferencePerturbation`` on the
  ``ecoli_k12_bw25113_locus_tag`` namespace whose ``crispr`` construct carries the dCas9
  effector and the 20 nt Table S1 spacer, and/or a ``BacterialDeletionPerturbation``.
  The crossed rows carry BOTH: an essential gene reachable only by knockdown together
  with a non-essential gene reachable only by deletion, which is the combination no
  deletion collection can produce. 32,499 of the 48,900 records are such
  two-perturbation genotypes.
- ENVIRONMENT: solid LB or solid MOPS-glucose minimal at 37 C, plus anhydrotetracycline
  (aTc) at the column's dose as a ``SmallMoleculePerturbation``. The 0 ng/mL columns
  carry NO perturbation, because no aTc was added to those plates.
- PHENOTYPE: ``FitnessPhenotype``, the released normalized colony growth, with
  ``n_samples = 2`` and ``screen_id`` naming the released table.

WHY ``FitnessPhenotype`` AND NOT ``EnvironmentResponsePhenotype``. Every released value
is a strictly positive relative colony growth whose neutral point is 1, not 0 (measured
on the pinned bytes: minimum 0.00613, maximum 4.752, no blanks, no NaN and nothing at or
below zero in any of the 63,552 cells, so the ``fitness`` clamp never fires). For Tables
S2A and S3 the denominator is a real control strain IN THE SAME CONDITION: the mean of
the 17 empty-vector wells of that plate, which is MEASURED to be exactly 1.0000 in each
of the 24 (medium, dose) columns -- so the reference convention ``fitness == 1`` is the
release's own normalization, not a convention imposed here. ``EnvironmentResponsePhenotype``
would require a log transform to put the baseline at 0, which would store a number the
paper never released.

WHY THE INDUCER DOSE IS THE ENVIRONMENT AND THE KNOCKDOWN IS THE GENOTYPE. The strain is
the strain at every dose: the guide and the dCas9 cassette are in the cell on the 0 ng/mL
plate too. What the dose changes is how far the DESIGNED knockdown is realized, so aTc
rides on the environment exactly as it does for the other CRISPRi loaders (Rapp 2026
stores its 200 nM aTc as a ``SmallMoleculePerturbation``). ``ExpressionRangeMultiplier``
is left unset, because no record states a realized repression magnitude.

THE THREE SCREENS ARE THREE SCREENS, AND ``screen_id`` IS WHAT KEEPS THEM APART. Table S2A
and Table S3 both hold the WT (lpp+) CRISPRi collection on MOPS minimal at the same six
doses, and they are two independent runs: MEASURED over the 377 x 6 overlapping cells,
ZERO values agree to 1e-9 and the per-dose Pearson r runs from 0.514 at 0 ng/mL to 0.945
at 500 ng/mL. Storing both without a screen label would put two different numbers on one
(strain, environment) key, so ``screen_id`` is ``Table S2A``, ``Table S3`` or
``Table S4A``.

THE FIVE-DOSE METHODS SENTENCE IS INCOMPLETE, AND THE TABLES ARE RIGHT. The Methods list
"five concentrations of aTc (0, 5, 10, 50, 500 ng/mL)" while Tables S2A and S3 carry six
columns including 100 ng/mL, which Figure 3A independently calls "6 different
concentrations of aTc" and which the Results use by name ("Looking at 100 ng/mL aTc
specifically"). Trusting the prose would drop a sixth of the dose response. The build
reproduces the paper's own count off the 100 ng/mL column as proof that the column is the
one the paper means: "With 100 ng/mL aTc added to the growth media, 272 from the 356
strains showed at least a 50% reduction in growth" -- MEASURED, exactly 272 of the 360
non-empty-vector Table S2A rows have LB 100 ng/mL growth at or below 0.5 (and 356 is the
number of DISTINCT target symbols those 360 rows carry). Likewise the Results say the
Keio screen used "two concentrations of inducer, 50 and 500 ng/mL" while Table S4A and
Figure 4B both carry three (0, 50, 500).

TABLE S2A AND TABLE S3 ARE TABLE S1'S PLATE ORDER, WHICH IS WHAT GIVES EVERY ROW A GUIDE.
Neither growth table carries a plate position or a spacer; both name a row by its target
symbol only, and 17 rows are named ``Empty_Vector``. MEASURED: the 377 rows of each are
the 377 filled wells of Table S1 in plate order, position for position, with exactly ONE
label disagreement -- Table S1's well G15 is ``gpsA`` and both growth tables spell it
``gspA``. Table S1's own annotation for that well ("Lipid biosythesis -
Glycerol-3-phosphate dehydrogenase", operon ``yibN,secB,gpsA,grxC``) and the Results
("enhancers ... essential phospholipids (gpsA, psd, plsBC, and pssA)") both say gpsA, and
``gspA`` is a different, non-essential gene, so this is a transposition typo in the growth
tables and Table S1's symbol is the one stored. The positional join is asserted at build
time: the row counts must match and there must be exactly one disagreement, at that well.

RECORDS DROPPED (rule, counts and items in ``preprocess/dropped_records.json``):

1. ``row_is_an_empty_vector_control`` (17 rows x 2 tables): the empty-vector wells ARE the
   denominator of Tables S2A and S3 -- their mean is 1.0000 by construction -- and their
   genotype carries no gene perturbation, so 17 copies of the reference would be 17
   strains nothing tells apart. They are typed as the ``phenotype_reference`` instead.
2. ``label_is_not_in_the_bw25113_annotation``: ``lapA`` and ``lapB`` among the targets;
   124 Keio labels (JW strain ids the pinned annotation no longer carries, plus small-RNA
   and non-gene labels).
3. ``label_is_a_fragment_of_a_merged_bw25113_locus``: ``ispU`` and ``uppS``, two target
   symbols of ONE BW25113 locus, and 46 Keio labels.
4. ``label_is_ambiguous_in_bw25113``: 10 Keio labels matching more than one locus.
5. ``construct_sits_in_more_than_one_collection_well`` (10 wells x 2 tables = 20 rows):
   MEASURED, the 360 guide wells carry only 355 distinct spacers. ``rbfA``, ``rplK``,
   ``rplX`` and ``ydfO`` each head two wells whose 20 nt spacers are IDENTICAL, so they
   are one construct pinned twice rather than two guides; the fifth repeated spacer is
   shared by ``ispU`` and ``uppS``, which are two names for one gene (rule 3 removes
   those two anyway). Two records of one strain in one condition cannot both be written,
   and keeping one would be arbitrary (the rule Campos 2018 and Shiver 2016 already
   apply).
6. ``label_heads_more_than_one_row`` (951 rows of Table S4A): 458 Keio deletion labels
   head 2 to 5 rows after resolution. Table S4A releases no plate, well or clone key, so
   two rows naming one deletion are two measurements of a strain identity the release
   gives nothing to separate. The Keio plate layout that would separate them is Baba
   2006's table, which the mirror does not hold.

WHAT IS NOT LOADED, AND WHY. Table S2B (the 52 LB-specific rows) and Table S2C (71
rows at 500 ng/mL) are subsets of Table S2A whose values this build reproduces exactly:
50 of the 52 Table S2C rows match Table S2A's LB and MOPS 500 ng/mL cells to 1e-9, and the
2 that do not (``ygeN``, ``yqcG``) carry a Table S2C LB value duplicated between them
while their MOPS values still match, which is a copy-paste in the subset. Table S3's
``Fold Change Growth (dlpp/WT)`` block and Table S4A's ``Empty Vector Normalized Growth``
block are RATIOS of columns this dataset already stores, and they are not the ratio of
the stored averages (MEASURED: Table S4A's lolA_0aTC differs from col5/col1 in the third
decimal on 498 of the first 500 rows), because the paper normalized per replicate and
then averaged. Storing them as ``GeneInteractionPhenotype`` would also mistype them:
``gene_interaction`` is an epsilon or tau, a deviation from an expected product, and a
fold change is not one. Table S4B is the author's hit list (68 suppressors, 9 enhancers),
a call over Table S4A rather than a measurement.

THE DISPERSION IS A TYPED GAP, NEVER A ZERO. Every released value is a mean of two
replicates and no table carries a spread. Figure S1B plots "the standard error of the
data" for a handful of selected strains, so the quantity exists upstream and was simply
not released per record; computing a sample SD from two values this loader never sees
would be our statistic, not theirs.

AND THE REPLICATE TYPE IS CONTRADICTORY IN THE SOURCE, SO ``sample_unit`` IS A GAP. The
Methods say "The average of the two technical replicates was then calculated and is
reported in Table S2" and "CRISPRi Keio screens were conducted in technical duplicate",
while the normalization paragraph of the SAME section says "the 12 Keio collection assay
plates screened in biological duplicate" and three figure captions say biological (Figure
S1A "two biological replicates (R1 and R2)", Figure S1B "the average of biological
replicates", Figure 4C "two biological replicates of the pFD152_lolA Keio collection").
``n_samples = 2`` is sourced either way; the unit is not, and is not guessed.

THE ONE ENVIRONMENT VALUE THE PAPER DOES NOT PIN. The base medium of the CRISPRi-Keio
assay plates is never named: that paragraph says only "cultures are spotted onto a 1536-
colony density solid agar plates containing different concentrations of aTc", opening with
"the workflow is followed as above with minor modifications". Every named step of that
workflow is on LB agar and the paper names no other base for it, so LB is what this loader
records, as a READING rather than a statement; it is recorded here, in the medium's own
provenance note and in ``preprocess/build_accounting.json``. Tables S2A and S3 need no
such reading: the Table S2 legend names "LB and MOPS minimal media", the Results name
"both rich (LB) and minimal (MOPS minimal) microbiological media", and Figures 3A, 3D and
S3 put the lpp screen on MOPS minimal. Which Table S2A block is which medium is MEASURED
rather than read off the column order: 50 of the 52 Table S2C rows match block 1 at
500 ng/mL as its LB column and block 2 as its MOPS-glucose column, to 1e-9.

THE CARBON SOURCE COMES FROM A COLUMN HEADER. The Methods say only that "MOPS minimal
medium was prepared according to the manufacturer's instructions" and never name a carbon
source; Table S2C's own column header does ("Normalized Growth MOPS-glucose with
500 ng/ml aTc (This study)"). Glucose is therefore a component of the MOPS medium with
its AMOUNT unset, which is the manufacturer's and is nowhere stated.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import os.path as osp
import shutil
from collections import Counter
from collections.abc import Callable, Iterator, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, ClassVar

import openpyxl
import pandas as pd
from pydantic import BaseModel, ConfigDict
from tqdm import tqdm

from torchcell.data import (
    ExperimentDataset,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
    write_verified,
)
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.media import LB, MOPS_MINIMAL
from torchcell.datamodels.schema import (
    BacterialCrisprInterferencePerturbation,
    BacterialDeletionPerturbation,
    BacterialFitnessExperiment,
    BacterialFitnessExperimentReference,
    ComponentDefinition,
    Concentration,
    ConcentrationUnit,
    CrisprConstruct,
    DerivedIdentifierMapping,
    DoseBasis,
    Environment,
    Experiment,
    ExperimentReference,
    FitnessPhenotype,
    Genotype,
    Media,
    MediaComponent,
    MediaComponentRole,
    Publication,
    SmallMoleculePerturbation,
    StrainConstruction,
    Temperature,
)
from torchcell.datasets.bacteria_common import (
    STRAIN_GENE_NAMESPACES,
    LocusTagReconciliation,
    assembly_reference,
    bacterial_genome,
    reconcile_locus_tags,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ROLE_RAW_DATA,
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.literature.provenance import run_retriever
from torchcell.literature.retrieve import pmc_cloud_url
from torchcell.sequence.genome.base import GeneNameStatus
from torchcell.sequence.genome.ecoli.k12 import (
    EcoliK12BW25113Genome,
    EcoliK12Genome,
    EcoliK12StrainName,
)
from torchcell.verification.report import Level, LevelResult, Provenance
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
)

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------- #
# Provenance anchors
# --------------------------------------------------------------------------- #
DOI = "10.1016/j.crmeth.2023.100693"
PMID = "38215765"
PMCID = "PMC10832289"
TITLE = (
    "A mobile CRISPRi collection enables genetic interaction studies for the essential "
    "genes of Escherichia coli"
)
CITATION_KEY = "rachwalskiMobileCRISPRiCollection2024"
RAW_DIR_REL = f"torchcell-raw/{CITATION_KEY}"
LIBRARY_DIR_REL = f"torchcell-library/{CITATION_KEY}"
DATASET_ROOT_REL = "data/torchcell/crispri_cross_rachwalski2024"

PAPER_MD = "paper.md"
PAPER_MD_SHA256 = "f96083b3ceb106a4ded3afe95406425d299bb55054405cb36d8d0de7ab268a27"
#: MinerU OCR of mmc1.pdf (Document S1, Figures S1-S4): the anchor of every SI quote.
SI1_MD = "si/si1.md"
SI1_MD_SHA256 = "12ed8e8ce819b45b620d40e72d75175434d1b5762253d1864a2d6d8fe555207e"

#: The PMC Article Datasets bucket prefix of this article (version 1).
PMC_PREFIX = f"{PMCID}.1"
#: When the four workbooks were re-retrieved and re-hashed before deposit.
RETRIEVED_AT = "2026-10-08"

TABLE_S1_FILE = "mmc2.xlsx"
TABLE_S2_FILE = "mmc3.xlsx"
TABLE_S3_FILE = "mmc4.xlsx"
TABLE_S4_FILE = "mmc5.xlsx"

ADDGENE_URL = "https://www.addgene.org/125546/"
ZENODO_DOI = "10.5281/zenodo.10214517"


class RawFile(BaseModel):
    """One file the loader consumes: its pinned bytes and how they were retrieved."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    name: str
    sha256: str
    bytes: int
    description: str

    @property
    def mirror_relpath(self) -> str:
        """Path inside the raw mirror (``data/<name>``)."""
        return f"data/{self.name}"

    @property
    def bucket_key(self) -> str:
        """The PMC Article Datasets bucket key the bytes came from."""
        return f"{PMC_PREFIX}/{self.name}"

    @property
    def source_url(self) -> str:
        """HTTPS URL of the bucket object."""
        return pmc_cloud_url(self.bucket_key)

    @property
    def retrieval(self) -> RetrievalRecord:
        """The re-runnable retrieval of these bytes."""
        return RetrievalRecord(
            method=RetrievalMethod.pmc_cloud,
            source_url=self.source_url,
            retriever="torchcell.literature.retrieve.pmc_cloud_object",
            params={"key": self.bucket_key},
            sha256=self.sha256,
            retrieved_at=RETRIEVED_AT,
        )


RAW_FILES: tuple[RawFile, ...] = (
    RawFile(
        name=TABLE_S1_FILE,
        sha256="e510a403d544da24a42f9caf0c86ce5445f3e5f632fca0541d2f3eb09c48bb29",
        bytes=51482,
        description="Table S1, the arrayed CRISPRi collection: each of the 384 wells "
        "with its 384_ROW / 384_COLUMN position, target symbol, Goodall essentiality "
        "flag, cellular function, operon and 20 nt guide RNA sequence. It is the only "
        "released source of a spacer and of a well, and its target symbols are what "
        "the two growth tables' row labels are joined to",
    ),
    RawFile(
        name=TABLE_S2_FILE,
        sha256="27c6220cfebeeeae62bc577716b28b15ee1819b800c56407328606f34d4ca616",
        bytes=88931,
        description="Table S2, sheet ST2A: average normalized growth of the 377 filled "
        "CRISPRi wells on solid LB and on solid MOPS-glucose minimal at 0, 5, 10, 50, "
        "100 and 500 ng/mL aTc (4,524 values). Sheets ST2B and ST2C are subsets of "
        "ST2A and are read only for the build's block-order check",
    ),
    RawFile(
        name=TABLE_S3_FILE,
        sha256="4eb17dab479967ccc98cd3592988f0c195949468dc5ca9635b07204489c84a73",
        bytes=107883,
        description="Table S3, sheet ST3: average normalized growth of the same 377 "
        "CRISPRi wells in the wild-type (lpp+) and the dlpp background on solid MOPS "
        "minimal at the same six doses (4,524 values), plus a dlpp/WT fold-change block "
        "this loader does not store",
    ),
    RawFile(
        name=TABLE_S4_FILE,
        sha256="636b4f6b73fd834c39508a6a5cb6045c17172f9a1852d71c3dd2aa92a12d93b1",
        bytes=10183121,
        description="Table S4, sheet ST4A: average normalized growth of 4,542 Keio "
        "deletion rows crossed with the empty vector and with the lolA, pssA and mreD "
        "knockdowns at 0, 50 and 500 ng/mL aTc (54,504 values), plus an "
        "empty-vector-normalized block this loader does not store. Sheet ST4B is the "
        "suppressor / enhancer hit list",
    ),
)
#: ``{raw file name: pinned sha256}``, the build-time check of every consumed file.
DATA_SHA256: dict[str, str] = {f.name: f.sha256 for f in RAW_FILES}
RAW_FILES_BY_NAME: dict[str, RawFile] = {f.name: f for f in RAW_FILES}

#: What the article releases that the loader deliberately does not consume.
NOT_MIRRORED = (
    "mmc1.pdf (Document S1, Figures S1-S4): quoted through the literature mirror's OCR "
    f"si/si1.md (sha256 {SI1_MD_SHA256}) for the normalization, replicate and duration "
    "statements; it carries no per-record value",
    "mmc6.pdf (Document S2, article plus supplemental information): a reprint of the "
    "article, byte-for-byte redundant with paper.pdf in the literature mirror",
    f"Zenodo {ZENODO_DOI}: a 1.4 GB and a 364 MB archive of raw plate images plus the "
    "ImageJ / R analysis code. No record here is built from an image, and the released "
    "tables are the normalized values the code produces, so the archives are retrieval "
    "metadata rather than a build input",
    "Table S2 sheets ST2B (52 LB-specific rows) and ST2C (71 rows at 500 ng/mL): both "
    "are subsets of ST2A, and this build reproduces 50 of ST2C's 52 joinable rows from "
    "ST2A to 1e-9; the two that differ (ygeN, yqcG) share one LB value in ST2C while "
    "their MOPS values still match, a copy-paste in the subset",
    "Table S3's 'Fold Change Growth (dlpp/WT)' block and Table S4A's 'Empty Vector "
    "Normalized Growth of each deletion mutant' block: ratios of columns this dataset "
    "already stores. They are not the ratio of the stored averages (Table S4A's "
    "lolA_0aTC differs from col5/col1 in the third decimal on 498 of the first 500 "
    "rows) because the paper normalized per replicate and then averaged, and a fold "
    "change is not the epsilon or tau that GeneInteractionPhenotype names",
    "Table S4B (68 suppressors and 9 enhancers at a 3-SD cutoff): a hit call over Table "
    "S4A, not a measurement",
    f"pFD152 (Addgene 125546, {ADDGENE_URL}): the conjugative CRISPRi backbone. The "
    "effector and the spacer are stored on CrisprConstruct; the full plasmid is what a "
    "future effector_plasmid_ref would pin",
)


def _paper(
    value: Any, quote: str, *, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to a verbatim quote in the sha256-pinned paper OCR mirror."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=PAPER_MD,
            citation_key=CITATION_KEY,
            sha256=PAPER_MD_SHA256,
            method="MinerU OCR of the publisher PDF (torchcell-library mirror)",
            page=page,
        ),
    )


def _si(value: Any, quote: str, *, page: str, note: str | None = None) -> SourcedValue:
    """Bind a value to a verbatim quote in the SI figure-legend OCR (``si/si1.md``)."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=SI1_MD,
            citation_key=CITATION_KEY,
            sha256=SI1_MD_SHA256,
            method="MinerU OCR of mmc1.pdf (torchcell-library mirror)",
            page=page,
        ),
    )


def _table(
    value: Any, quote: str, *, name: str, page: str, note: str | None = None
) -> SourcedValue:
    """Bind a value to a verbatim cell of one of the pinned workbooks."""
    return SourcedValue(
        value=value,
        quote=quote,
        note=note,
        provenance=Provenance(
            source_uri=f"data/{name}",
            citation_key=CITATION_KEY,
            sha256=RAW_FILES_BY_NAME[name].sha256,
            method="verbatim cell of the released workbook (raw mirror)",
            page=page,
        ),
    )


# --------------------------------------------------------------------------- #
# Sourced values (verbatim quotes, pinned sha256)
# --------------------------------------------------------------------------- #
_MODEL = (
    "EXPERIMENTAL MODEL AND STUDY PARTICIPANT DETAILS, 'Growth conditions for E. coli'"
)
_PLASMIDS = (
    "STAR Methods, METHOD DETAILS, 'Plasmids and construction of CRISPRi collection'"
)
_SCREENING = "STAR Methods, METHOD DETAILS, 'High-throughput conjugation and CRISPRi array screening'"
_NORMALIZATION = (
    "STAR Methods, METHOD DETAILS, 'Plate imaging quantification and analysis'"
)
_RESULTS_COLLECTION = (
    "Results, 'Creation of an arrayed mobile CRISPRi library on the pFD152 backbone'"
)
_RESULTS_LPP = "Results, 'Conjugation of the CRISPRi collection into E. coli Dlpp ...'"
_RESULTS_KEIO = (
    "Results, 'Conjugation of pFD152-lolA into the Keio deletion collection ...'"
)
_FIGURE_3 = "Figure 3 legend"
_FIGURE_4 = "Figure 4 legend"

SOURCED_VALUES: dict[str, SourcedValue] = {
    "host_strain": _paper(
        "BW25113",
        "All fitness screens with the CRISPRi collection were performed in E. coli "
        "BW25113 [F D(araD-araB)567 lacZ4787D::rrnB-3 LAM rph-1 D(rhaD-rhaB)568 "
        "hsdR514], and, all genes deletions used in the study were obtained from the "
        "Keio collection4 (kanamycin resistant single-gene deletions in E. coli "
        "BW25113).",
        page=_PLASMIDS,
        note="the assembly every record pins, and the namespace its locus tags are "
        "written in; the Keio deletions share it",
    ),
    "temperature": _paper(
        37.0,
        "then were grown for $1 6 \\mathrm { ~ h ~ }$ at $3 7 ^ { \\circ } \\mathrm { C }$ "
        "and visualized using transmissive scanners.",
        page=_SCREENING,
        note="the assay plates' incubation temperature",
    ),
    "duration_hours": _paper(
        16.0,
        "then were grown for $1 6 \\mathrm { ~ h ~ }$ at $3 7 ^ { \\circ } \\mathrm { C }$ "
        "and visualized using transmissive scanners.",
        page=_SCREENING,
        note="the CRISPRi-collection screens (Tables S2A and S3). The CRISPRi-Keio "
        "paragraph states no duration of its own, so Table S4A's is a typed gap",
    ),
    "effector": _paper(
        "dCas9",
        "sgRNAs targeting different essential genes in E. coli were cloned into the "
        "conjugative CRISPRi plasmid pFD152 using a previously described single step "
        "Golden Gate assembly protocol.44 pFD152 was a gift from David Bikard (Addgene "
        "plasmid # 125546).",
        page=_PLASMIDS,
        note="pFD152 carries a catalytically dead Cas9; the plasmid itself is the "
        "effector_plasmid_ref a future full-plasmid capture would pin",
    ),
    "guide_design": _paper(
        "CRISPRbact",
        "sgRNAs used in this study were designed using the publicly available "
        "CRISPRbact tool (https://gitlab.pasteur.fr/dbikard/crisprbact), wherein, an "
        "sgRNA with the highest predicted on-target activity score and the least "
        "off-target homology was selected for each gene.",
        page=_PLASMIDS,
        note="one guide per gene by design, which is why n_guides is 1 on every "
        "construct and why two wells carrying one spacer are one construct",
    ),
    "lb_medium": _paper(
        "LB",
        "E. coli was routinely cultured in LB medium at $3 7 ^ { \\circ } \\mathrm { C }$ , "
        "supplemented with antibiotics (kanamycin, $5 0 \\mu \\ g / \\mathrm { m L }$ and "
        "spectinomycin, $1 5 0 \\mu \\ g /$ mL) or $1 0 \\mathsf { m } \\mathsf { M }$ "
        "Diaminopimelic acid when required.",
        page=_MODEL,
        note="the formulation is not stated, so the shared unqualified LB object is the "
        "base; the kanamycin and spectinomycin are named for the handling and "
        "exconjugant-selection steps, not for the aTc assay plates, whose selection is "
        "nowhere stated and so is not asserted",
    ),
    "agar_lb": _paper(
        15.0,
        "For experiments on solid medium, media was prepared with $1 5 { \\mathfrak { g } } / \\mathsf { L }$ agar.",
        page=_MODEL,
        note="the paper's general statement for every solid medium in the study",
    ),
    "agar_mops": _paper(
        1.5,
        "MOPS minimal medium was prepared according to the manufacturer's instructions: "
        "components were filter sterilized after preparing liquid growth medium or added "
        "to sterile water and agar ( $1 . 5 \\%$ w/v) for solid growth medium.",
        page=_MODEL,
        note="the same amount as the LB sentence's 15 g/L, in the unit that sentence "
        "uses; the MOPS base recipe is the manufacturer's and is not printed",
    ),
    "mops_glucose": _table(
        "D-glucose",
        "Normalized Growth MOPS-glucose with 500 ng/ml aTc (This study)",
        name=TABLE_S2_FILE,
        page="sheet ST2C, header row 2, column C",
        note="the ONLY place the minimal medium's carbon source is named; the Methods "
        "say only 'prepared according to the manufacturer's instructions', so the "
        "amount is unset",
    ),
    "two_media": _paper(
        ["LB", "MOPS minimal"],
        "We screened the CRISPRi collection for growth inhibition in both rich (LB) and "
        "minimal (MOPS minimal) microbiological media at varying aTc concentrations and "
        "observed a dose-dependent growth defect (Figure 1D; Table S1A)",
        page=_RESULTS_COLLECTION,
        note="this sentence cites the growth table as 'Table S1A' while the workbook's "
        "own Legend sheet calls it 'Supplemental Table 2A' and Table S1 is the "
        "collection layout with no growth value in it; the Results citation is off by "
        "one and the workbook legend is what this loader follows",
    ),
    "lpp_medium": _paper(
        "MOPS minimal",
        "The wildtype $( I P { p } + )$ and Dlpp CRISPRi collections were then screened "
        "at varying aTc concentrations on minimal media, and the growth observed for "
        "each strain in the Dlpp background was normalized to that of the wild-type "
        "$( I P { p } + )$ collection at each concentration of inducer (Figure S3).",
        page=_RESULTS_LPP,
        note="Table S3's medium. Figure 3D says 'grown on MOPS minimal medium' and "
        "Figure S3 'on MOPS-minimal media'",
    ),
    "six_doses": _paper(
        6,
        "(A) The mobile arrayed CRISPRi collection was introduced into E. coli "
        "$\\Delta I p p$ , and the collection was screened at 6 different concentrations "
        "of aTc in minimal media.",
        page=_FIGURE_3,
        note="SIX, against the Methods sentence that lists five and omits 100 ng/mL, "
        "which the Table S2A and S3 headers do carry and the Results use by name",
    ),
    "five_dose_list": _paper(
        [0.0, 5.0, 10.0, 50.0, 500.0],
        "Cells were then pinned onto solid media containing five concentrations of aTc "
        "(0, 5, 10, 50, $5 0 0 ~ \\mathrm { { n g / m L } }$ ) using the Singer ROTOR "
        "$^ +$ HDA $( { \\sim } 1 \\ \\mu \\ L )$",
        page=_SCREENING,
        note="INCOMPLETE: the released headers carry 100 ng/mL as well. Recorded so the "
        "discrepancy is in the record, not resolved in favour of the prose",
    ),
    "keio_doses": _paper(
        [0.0, 50.0, 500.0],
        "(B) The Keio collection harboring pFD152_lolA was screened at 0, 50, and "
        "$5 0 0 \\mathrm { n g / m L }$ of aTc.",
        page=_FIGURE_4,
        note="THREE, against the Results sentence 'The CRISPRi Keio screen was "
        "conducted at two concentrations of inducer, 50 and 500 ng/mL aTc'; Table S4A "
        "carries three blocks of three columns, so the figure legend is the one the "
        "release agrees with",
    ),
    "keio_constructs": _paper(
        ["lolA", "pssA", "mreD"],
        "In addition to conjugating and screening the lolA CRISPRi knockdown with the "
        "Keio collection, we also conjugated and screened CRISPRi knockdowns of pssA "
        "and mreD, which code for phosphatidylserine synthase64 and an integral "
        "innermembrane component of the elongasome,65 respectively, as controls to "
        "identify genetic interactions specific to lolA.",
        page=_RESULTS_KEIO,
        note="the three knockdowns crossed into the Keio collection, each a column "
        "block of Table S4A beside the empty-vector block",
    ),
    "keio_medium_unstated": _paper(
        "LB",
        "Then, cultures are spotted onto a 1536-colony density solid agar plates "
        "containing different concentrations of aTc, such that each Keio mutant is found "
        "in four proximal colonies– each with a different CRISPRi plasmid or vector "
        "control.",
        page=_SCREENING,
        note="A READING, NOT A STATEMENT. This sentence names no base medium; the "
        "paragraph opens 'For conjugation of a specific CRISPRi plasmid into a genomic "
        "collection, the workflow is followed as above with minor modifications' and "
        "every named step of it is on LB agar, so LB is what this loader records for "
        "Table S4A. It is the one environment value the paper does not pin",
    ),
    "n_samples_s2": _paper(
        2,
        "The average of the two technical replicates was then calculated and is "
        "reported in Table S2.",
        page=_NORMALIZATION,
        note="n_samples for Tables S2A and S3; the TYPE is contradicted by the figure "
        "legends, so sample_unit is a typed gap",
    ),
    "n_samples_s4_technical": _paper(
        2,
        "CRISPRi Keio screens were conducted in technical duplicate.",
        page=_SCREENING,
        note="n_samples for Table S4A, called TECHNICAL here",
    ),
    "n_samples_s4_biological": _paper(
        2,
        "In summary, the entire Keio screen used 24 assay plates at every aTc "
        "concentration – the 12 Keio collection assay plates screened in biological "
        "duplicate – wherein mutants with the same CRISPRi plasmid occupy the same "
        "position in each plate.",
        page=_NORMALIZATION,
        note="the SAME methods section calling the same duplicate BIOLOGICAL; the count "
        "agrees and the type does not",
    ),
    "replicates_are_biological": _si(
        "biological_replicate",
        "Growth of each CRISPRi mutant was normalized to the mean growth of the 17 empty "
        "vector controls on the plate, then two biological replicates (R1 and R2), or "
        "the average of the two biological replicates are plotted.",
        page="Figure S1 legend, panel A",
        note="the third reading of the replicate type, and the statement that fixes the "
        "denominator of Tables S2A and S3: the MEAN of the 17 empty-vector wells of that "
        "plate, which is MEASURED to be exactly 1.0000 in each of the 24 columns",
    ),
    "uncertainty_exists_unreleased": _si(
        "standard error",
        "Bars represent the average of biological replicates and the standard error of "
        "the data is shown with error bars.",
        page="Figure S1 legend, panel B",
        note="a dispersion exists upstream, for the handful of strains that panel plots; "
        "no table releases one per record, so the uncertainty is a typed gap",
    ),
    "keio_normalization": _paper(
        "two-pass inter-quartile-mean normalization",
        "First, spatial effects across the 24 assay plates were normalized by dividing "
        "each colony with the inter-quartile mean (IQM) of every colony occupying the "
        "same position on other plates (with the same CRISPRi plasmid). Then, plate to "
        "plate variability was normalized by dividing the growth of each colony, by the "
        "IQM of the growth of every other colony harboring the same CRISPRi plasmid "
        "within each screening plate.",
        page=_NORMALIZATION,
        note="Table S4A's stored columns are IQM-normalized colony growth, so their "
        "neutral point is the screen population rather than a control strain; MEASURED, "
        "the twelve stored columns have medians 0.9991 to 1.0033",
    ),
    "s2_normalization": _paper(
        "mean of the 17 empty-vector wells",
        "For analysis of CRISPRi collection growth in wildtype $( I p p + )$ and "
        "$\\Delta I p p$ backgrounds, the average raw colony density of the 17 strains "
        "harboring the empty vector was calculated, then, this value was used to "
        "normalize the growth of strains harboring different CRISPRi plasmids.",
        page=_NORMALIZATION,
        note="MEASURED per column rather than once per screen: the 17 empty-vector rows "
        "mean exactly 1.0000 in every one of Table S2A's 12 and Table S3's 12 columns",
    ),
    "fifty_percent_count": _paper(
        272,
        "With $1 0 0 ~ \\mathrm { { n g / m L } }$ aTc added to the growth media, 272 "
        "from the 356 strains showed at least a $50 \\%$ reduction in growth (Figure 1C).",
        page=_RESULTS_COLLECTION,
        note="the build's arithmetic oracle: MEASURED, exactly 272 of Table S2A's 360 "
        "non-empty-vector rows have LB 100 ng/mL growth at or below 0.5, and those 360 "
        "rows carry 356 distinct target symbols",
    ),
    "keio_size": _paper(
        4000,
        "we chose to conjugate the strongest suppressor from the Dlpp CRISPRi "
        "conjugation screen, the lolA CRISPRi knockdown, into all ${ \\sim } 4 { , } 0 0 0 $ "
        "strains of the Keio collection.",
        page=_RESULTS_KEIO,
        note="Table S4A holds 4,542 rows over 4,017 distinct deletion labels",
    ),
    "keio_cassette": _paper(
        "kanamycin-resistance cassette",
        "all genes deletions used in the study were obtained from the Keio collection4 "
        "(kanamycin resistant single-gene deletions in E. coli BW25113).",
        page=_PLASMIDS,
        note="the cassette stored on every Keio deletion perturbation",
    ),
    "lpp_deletion": _paper(
        "lpp",
        "First, we conjugated the mobile CRISPRi collection into E. coli Dlpp. Lpp is "
        "the most abundant lipoprotein in E. coli and acts to covalently tether "
        "peptidoglycan to the outer membrane.53–56",
        page=_RESULTS_LPP,
        note="the query strain of the Table S3 screen; a kanamycin-resistant deletion "
        "like the rest ('the query strain (kanamycin resistant E. coli Dlpp) was arrayed "
        "on LB agar with kanamycin'), so it carries the same collection and cassette",
    ),
}

# --------------------------------------------------------------------------- #
# Constants of the release
# --------------------------------------------------------------------------- #
BW25113_STRAIN: EcoliK12StrainName = "BW25113"
BW25113_NAMESPACE = STRAIN_GENE_NAMESPACES["BW25113"]
KEIO_COLLECTION = "Keio collection"
KEIO_CASSETTE = "kanamycin-resistance cassette"
CRISPRI_COLLECTION = "Rachwalski 2024 mobile CRISPRi collection (pFD152)"
EFFECTOR = "dCas9"
EMPTY_VECTOR = "Empty_Vector"
LPP_SYMBOL = "lpp"

TEMPERATURE_C = 37.0
DURATION_HOURS = 16.0

#: Table S1's nine column headers, in order.
S1_COLUMNS: tuple[str, ...] = (
    "384_ROW",
    "384_COLUMN",
    "gRNA",
    "Essential in Goodall et. al. 2018",
    "Cellular Function",
    "Operon (sgRNA targeted gene is bolded)",
    "Guide RNA Sequence",
    "Cloning Primer 1",
    "Cloning Primer 2",
)
#: The six inducer doses of Tables S2A and S3, in ng/mL, in released column order.
SIX_DOSES_NG_PER_ML: tuple[float, ...] = (0.0, 5.0, 10.0, 50.0, 100.0, 500.0)
#: The three inducer doses of Table S4A, in ng/mL, in released column order.
THREE_DOSES_NG_PER_ML: tuple[float, ...] = (0.0, 50.0, 500.0)
#: Table S2A / S3 dose-column headers, in order, for the header assertion.
SIX_DOSE_HEADERS: tuple[str, ...] = (
    "0 ng/ml aTC",
    "5 ng/ml aTC",
    "10 ng/ml aTC",
    "50 ng/ml aTC",
    "100 ng/ml aTC",
    "500 ng/ml aTC",
)
#: Table S4A dose-column headers within one block, in order.
THREE_DOSE_HEADERS: tuple[str, ...] = ("Growth_0atC", "Growth_50aTC", "Growth_500aTC")
#: Table S4A's four column-block headers, in released order: the vector then the three
#: knockdowns. The block header is what names each block's CRISPRi plasmid.
S4A_BLOCKS: tuple[tuple[str | None, str], ...] = (
    (None, "Average Normalized Growth of pFD152 (Empty Vector) Cross into the Keio"),
    ("lolA", "Average Normalized Growth of pFD152:lolA Cross into the Keio"),
    ("pssA", "Average Normalized Growth of pFD152:pssA Cross into the Keio"),
    ("mreD", "Average Normalized Growth of pFD152:mreD Cross into the Keio"),
)

SCREEN_S2A = "Table S2A"
SCREEN_S3 = "Table S3"
SCREEN_S4A = "Table S4A"

MEDIUM_LB = "LB"
MEDIUM_MOPS = "MOPS-glucose"

#: Counts MEASURED on the pinned bytes; the build refuses a release that moved.
S1_ROWS = 384
S1_FILLED_WELLS = 377
S1_EMPTY_VECTOR_WELLS = 17
S2A_ROWS = 377
S3_ROWS = 377
S4A_ROWS = 4542
S4A_DISTINCT_LABELS = 4017
#: The single positional label disagreement between Table S1 and the growth tables.
S1_TYPO_WELL = "G15"
S1_TYPO_SYMBOL = "gpsA"
GROWTH_TABLE_TYPO_SYMBOL = "gspA"
#: Guide rows that are not an empty-vector control.
S2A_GUIDE_ROWS = 360
S2A_DISTINCT_TARGETS = 356
#: The paper's own 50%-reduction count at 100 ng/mL aTc on LB.
FIFTY_PERCENT_THRESHOLD = 0.5
FIFTY_PERCENT_COUNT = 272
#: Fraction of distinct labels that must resolve to one BW25113 locus, per table.
MIN_RESOLVED_TARGETS = 0.98
MIN_RESOLVED_KEIO = 0.95

DROP_EMPTY_VECTOR = "row_is_an_empty_vector_control"
DROP_NOT_IN_ANNOTATION = "label_is_not_in_the_bw25113_annotation"
DROP_MERGED_LOCUS = "label_is_a_fragment_of_a_merged_bw25113_locus"
DROP_AMBIGUOUS = "label_is_ambiguous_in_bw25113"
DROP_REPEATED_CONSTRUCT = "construct_sits_in_more_than_one_collection_well"
DROP_DUPLICATE_LABEL = "label_heads_more_than_one_row"

ROW_RULES: tuple[tuple[str, str], ...] = (
    (
        DROP_EMPTY_VECTOR,
        "the 17 empty-vector wells of Tables S2A and S3 ARE their denominator -- the "
        "mean of their raw colony densities is what every other value is divided by, "
        "and it is 1.0000 in every column -- and their genotype carries no gene "
        "perturbation, so they are the phenotype_reference rather than 17 records "
        "nothing tells apart",
    ),
    (
        DROP_NOT_IN_ANNOTATION,
        "the label is no locus tag, symbol or synonym of the pinned BW25113 assembly "
        "(Keio JW strain ids the annotation predates, small-RNA and non-gene labels), "
        "and a bacterial leaf stores a locus tag of its declared namespace and nothing "
        "else",
    ),
    (
        DROP_MERGED_LOCUS,
        "two or more labels resolve to ONE current BW25113 locus, so no member IS the "
        "locus tag; the whole group goes rather than merging distinct strains",
    ),
    (
        DROP_AMBIGUOUS,
        "the label matches more than one locus of the pinned assembly, so no single "
        "locus tag can be written for it",
    ),
    (
        DROP_REPEATED_CONSTRUCT,
        "the construct's 20 nt spacer sits in more than one well of Table S1, so the "
        "wells hold one strain pinned twice rather than two guides; two records of one "
        "strain in one condition cannot both be written and keeping one would be "
        "arbitrary",
    ),
    (
        DROP_DUPLICATE_LABEL,
        "the deletion label heads more than one row of Table S4A, which releases no "
        "plate, well or clone key, so the rows are measurements of a strain identity "
        "the release gives nothing to separate",
    ),
)


# --------------------------------------------------------------------------- #
# Media and environments
# --------------------------------------------------------------------------- #
def _agar(
    value: float, unit: ConcentrationUnit, source: SourcedValue
) -> MediaComponent:
    """The gelling agent of a solid plate, at the amount its own sentence states."""
    return MediaComponent(
        compound=resolved_compound("agar"),
        role=MediaComponentRole.gelling_agent,
        definition=ComponentDefinition.defined,
        concentration=Concentration(value=value, unit=unit),
        provenance=[source],
    )


RACHWALSKI2024_LB_AGAR = Media(
    name="LB agar, 15 g/L agar, formulation not stated (Rachwalski 2024)",
    state="solid",
    is_synthetic=False,
    base_medium="LB",
    components=[
        *LB.components,
        _agar(15.0, ConcentrationUnit.g_per_l, SOURCED_VALUES["agar_lb"]),
    ],
    provenance=[SOURCED_VALUES["lb_medium"], SOURCED_VALUES["agar_lb"]],
)
"""The rich assay plate: the shared unqualified LB base plus the paper's 15 g/L agar."""

RACHWALSKI2024_MOPS_GLUCOSE_AGAR = Media(
    name="MOPS-glucose minimal agar, 1.5% (w/v) agar, glucose amount not stated "
    "(Rachwalski 2024)",
    state="solid",
    is_synthetic=True,
    base_medium="MOPS_MINIMAL",
    components=[
        *MOPS_MINIMAL.components,
        MediaComponent(
            compound=resolved_compound("D-glucose"),
            role=MediaComponentRole.carbon_source,
            definition=ComponentDefinition.defined,
            concentration=None,
            provenance=[SOURCED_VALUES["mops_glucose"]],
            note="the carbon source is named only by Table S2C's column header; the "
            "Methods defer the recipe to the manufacturer and print no amount",
        ),
        _agar(1.5, ConcentrationUnit.percent_w_v, SOURCED_VALUES["agar_mops"]),
    ],
    provenance=[
        SOURCED_VALUES["agar_mops"],
        SOURCED_VALUES["mops_glucose"],
        SOURCED_VALUES["lpp_medium"],
    ],
)
"""The minimal assay plate: Neidhardt's carbon-free MOPS base plus glucose and agar."""

MEDIA_BY_LABEL: dict[str, Media] = {
    MEDIUM_LB: RACHWALSKI2024_LB_AGAR,
    MEDIUM_MOPS: RACHWALSKI2024_MOPS_GLUCOSE_AGAR,
}

_KEIO_DURATION_NOTE = (
    "the CRISPRi-Keio paragraph states no incubation time for its assay plates; it "
    "opens 'the workflow is followed as above with minor modifications', whose 16 h is "
    "not claimed here for a plate the paper does not time"
)
KEIO_DURATION_GAPS: tuple[ProvenanceGap, ...] = tuple(
    ProvenanceGap(
        field=field,
        reason=ProvenanceGapReason.not_reported_by_primary,
        note=_KEIO_DURATION_NOTE,
    )
    for field in ("duration_hours", "duration_generations")
)


def atc_perturbation(ng_per_ml: float) -> SmallMoleculePerturbation:
    """The inducer at one released dose, stored in ug/mL (the schema's unit)."""
    if ng_per_ml <= 0.0:
        raise ValueError(
            "the 0 ng/mL plates carry no aTc, so they carry no perturbation"
        )
    return SmallMoleculePerturbation(
        compound=resolved_compound("anhydrotetracycline"),
        concentration=Concentration(
            value=ng_per_ml / 1000.0,
            unit=ConcentrationUnit.ug_per_ml,
            basis=DoseBasis.fixed,
        ),
    )


def environment(medium: str, ng_per_ml: float, *, timed: bool) -> Environment:
    """One assay plate: its medium at 37 C, with the dose's aTc when any was added."""
    return Environment(
        media=MEDIA_BY_LABEL[medium],
        temperature=Temperature(value=TEMPERATURE_C),
        perturbations=[] if ng_per_ml <= 0.0 else [atc_perturbation(ng_per_ml)],
        aerobicity="aerobic",
        duration_hours=DURATION_HOURS if timed else None,
        provenance_gaps=[] if timed else list(KEIO_DURATION_GAPS),
    )


# --------------------------------------------------------------------------- #
# Phenotype
# --------------------------------------------------------------------------- #
_GAP_SAMPLE_UNIT = ProvenanceGap(
    field="sample_unit",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="the source contradicts itself: the Methods say 'The average of the two "
    "technical replicates was then calculated and is reported in Table S2' and 'CRISPRi "
    "Keio screens were conducted in technical duplicate', while the same section's "
    "normalization paragraph says 'the 12 Keio collection assay plates screened in "
    "biological duplicate' and Figures S1A, S1B and 4C all say biological. The count is "
    "2 either way; the unit is not stated consistently and is not guessed",
)
_GAP_UNCERTAINTY = ProvenanceGap(
    field="fitness_uncertainty",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="every released value is the mean of the two replicates and no table carries a "
    "spread. Figure S1B plots 'the standard error of the data' for a handful of selected "
    "strains, so the quantity exists upstream and was not released per record; a sample "
    "SD of two values this loader never sees would be our statistic, not theirs",
)
_GAP_UNCERTAINTY_TYPE = ProvenanceGap(
    field="fitness_uncertainty_type",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="no uncertainty is released, so there is no kind to name",
)
_GAP_SE = ProvenanceGap(
    field="fitness_se",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="no uncertainty is released, so no standard error can be derived",
)
_GAP_STD = ProvenanceGap(
    field="fitness_std",
    reason=ProvenanceGapReason.not_reported_by_primary,
    note="no standard deviation is released",
)
PHENOTYPE_GAPS: tuple[ProvenanceGap, ...] = (
    _GAP_SAMPLE_UNIT,
    _GAP_UNCERTAINTY,
    _GAP_UNCERTAINTY_TYPE,
    _GAP_SE,
    _GAP_STD,
)
N_REPLICATES = 2


def phenotype(growth: float, screen_id: str) -> FitnessPhenotype:
    """One released normalized colony growth, labelled with the screen it came from."""
    return FitnessPhenotype(
        fitness=growth,
        n_samples=N_REPLICATES,
        sample_unit=None,
        screen_id=screen_id,
        provenance_gaps=list(PHENOTYPE_GAPS),
    )


def reference_phenotype() -> FitnessPhenotype:
    """The empty-vector control in the same condition: growth over itself, 1.0.

    MEASURED, not assumed: the mean of the 17 empty-vector wells is exactly 1.0000 in
    each of Table S2A's twelve and Table S3's twelve columns, and Table S4A's twelve
    stored columns have medians 0.9991 to 1.0033 after the two-pass IQM normalization.
    """
    return FitnessPhenotype(
        fitness=1.0,
        n_samples=S1_EMPTY_VECTOR_WELLS,
        sample_unit=None,
        provenance_gaps=list(PHENOTYPE_GAPS),
    )


# --------------------------------------------------------------------------- #
# Reading the four workbooks
# --------------------------------------------------------------------------- #
class TableLayoutError(ValueError):
    """A released workbook's header is not the one the loader was written against."""


def _cell(value: Any) -> str | None:
    """A workbook cell as a stripped string, or None when it is blank."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _rows(path: str | Path, sheet: str, *, skip: int) -> list[tuple[Any, ...]]:
    """Every row of ``sheet`` after ``skip`` header rows whose first cell is filled."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = book[sheet]
        return [
            row
            for row in worksheet.iter_rows(min_row=skip + 1, values_only=True)
            if _cell(row[0]) is not None
        ]
    finally:
        book.close()


def _header(path: str | Path, sheet: str, row: int) -> tuple[Any, ...]:
    """One header row of ``sheet`` (1-based)."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = book[sheet]
        for index, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if index == row:
                return tuple(values)
        raise TableLayoutError(f"{sheet} has no row {row}")
    finally:
        book.close()


def _expect(got: Sequence[Any], want: Sequence[str], what: str) -> None:
    """Refuse a sheet whose header cells are not the expected ones."""
    seen = [_cell(value) for value in got[: len(want)]]
    if seen != list(want):
        raise TableLayoutError(f"{what} header {seen} is not {list(want)}")


class CollectionWell(BaseModel):
    """One filled well of the arrayed CRISPRi collection (a Table S1 row)."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    well: str
    symbol: str
    spacer: str | None
    essential: str | None
    function: str | None

    @property
    def is_empty_vector(self) -> bool:
        """True for the 17 wells carrying pFD152 with no guide."""
        return self.symbol == EMPTY_VECTOR


def read_table_s1(path: str | Path) -> list[CollectionWell]:
    """Table S1's filled wells, in plate order, with their spacers.

    The sheet has one header row and 384 position rows; the seven trailing wells of row
    P carry a position and no gRNA, which is what "Empty positions at the bottom right
    of the plate represent control wells lacking bacteria" names.
    """
    _expect(_header(path, "Sheet1", 1), S1_COLUMNS, "Table S1")
    rows = _rows(path, "Sheet1", skip=1)
    if len(rows) != S1_ROWS:
        raise TableLayoutError(f"Table S1 has {len(rows)} rows, not {S1_ROWS}")
    wells = [
        CollectionWell(
            well=f"{_cell(row[0])}{_cell(row[1])}",
            symbol=str(_cell(row[2])),
            spacer=_cell(row[6]),
            essential=_cell(row[3]),
            function=_cell(row[4]),
        )
        for row in rows
        if _cell(row[2]) is not None
    ]
    if len(wells) != S1_FILLED_WELLS:
        raise TableLayoutError(
            f"Table S1 has {len(wells)} filled wells, not {S1_FILLED_WELLS}"
        )
    vectors = sum(1 for well in wells if well.is_empty_vector)
    if vectors != S1_EMPTY_VECTOR_WELLS:
        raise TableLayoutError(
            f"Table S1 has {vectors} empty-vector wells, not {S1_EMPTY_VECTOR_WELLS}"
        )
    spacerless = [
        well.well for well in wells if not well.is_empty_vector and not well.spacer
    ]
    if spacerless:
        raise TableLayoutError(f"Table S1 wells with no guide sequence: {spacerless}")
    return wells


class GrowthBlock(BaseModel):
    """One released column block: what it varies, and its value per table row."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    #: ``MEDIUM_LB`` or ``MEDIUM_MOPS``.
    medium: str
    #: Locus symbol of a deletion the whole block carries, or None.
    background: str | None
    #: Target symbol of the block's CRISPRi plasmid, or None for the empty vector.
    knockdown: str | None
    #: Inducer dose of each column, ng/mL, in order.
    doses: tuple[float, ...]
    #: One tuple of ``len(doses)`` values per table row, in table order.
    values: tuple[tuple[float, ...], ...]


def _floats(
    row: tuple[Any, ...], columns: Sequence[int], where: str
) -> tuple[float, ...]:
    """The named columns of a row as floats; a blank or non-numeric cell stops the build."""
    out: list[float] = []
    for column in columns:
        value = row[column]
        if value is None or isinstance(value, bool):
            raise TableLayoutError(f"{where}: column {column} is blank")
        out.append(float(value))
    return tuple(out)


def read_table_s2a(path: str | Path) -> tuple[list[str], tuple[GrowthBlock, ...]]:
    """Table S2A's row labels and its two medium blocks (LB first, MOPS second).

    The block order is not read off the group header, which says only "Average Growth of
    Replicates" twice; it is MEASURED against sheet ST2C, whose two columns name their
    medium and hold the same 500 ng/mL values.
    """
    _expect(_header(path, "ST2A", 2)[1:7], SIX_DOSE_HEADERS, "Table S2A block 1")
    _expect(_header(path, "ST2A", 2)[7:13], SIX_DOSE_HEADERS, "Table S2A block 2")
    rows = _rows(path, "ST2A", skip=2)
    if len(rows) != S2A_ROWS:
        raise TableLayoutError(f"Table S2A has {len(rows)} rows, not {S2A_ROWS}")
    labels = [str(_cell(row[0])) for row in rows]
    blocks = tuple(
        GrowthBlock(
            medium=medium,
            background=None,
            knockdown=None,
            doses=SIX_DOSES_NG_PER_ML,
            values=tuple(
                _floats(row, range(first, first + 6), f"Table S2A row {index}")
                for index, row in enumerate(rows)
            ),
        )
        for medium, first in ((MEDIUM_LB, 1), (MEDIUM_MOPS, 7))
    )
    return labels, blocks


def check_s2a_block_order(
    path: str | Path, labels: Sequence[str], blocks: Sequence[GrowthBlock]
) -> dict[str, Any]:
    """MEASURE which Table S2A block is LB by matching sheet ST2C's named columns.

    ST2C gives, for 71 genes, the normalized growth "LB with 500 ng/ml aTc" and
    "MOPS-glucose with 500 ng/ml aTc". Every one of them that heads exactly one ST2A row
    must equal that row's 500 ng/mL cell in the block this loader calls that medium.
    """
    _expect(
        _header(path, "ST2C", 2)[:3],
        (
            "Genes",
            "Normalized Growth LB with 500 ng/ml aTc (This study)",
            "Normalized Growth MOPS-glucose with 500 ng/ml aTc (This study)",
        ),
        "Table S2C",
    )
    counts = Counter(labels)
    by_label = {label: index for index, label in enumerate(labels)}
    lb_block = next(block for block in blocks if block.medium == MEDIUM_LB)
    mops_block = next(block for block in blocks if block.medium == MEDIUM_MOPS)
    matched = 0
    mismatched: list[str] = []
    for row in _rows(path, "ST2C", skip=2):
        label = str(_cell(row[0]))
        if counts.get(label) != 1:
            continue
        index = by_label[label]
        lb_ok = abs(lb_block.values[index][-1] - float(row[1])) < 1e-9
        mops_ok = abs(mops_block.values[index][-1] - float(row[2])) < 1e-9
        if lb_ok and mops_ok:
            matched += 1
        else:
            mismatched.append(label)
    if matched < len(mismatched) * 10:
        raise TableLayoutError(
            f"Table S2C matches the LB-then-MOPS block order on only {matched} rows "
            f"({len(mismatched)} disagree): {mismatched[:10]}"
        )
    return {
        "rule": "table_s2c_fixes_the_block_order",
        "matched_rows": matched,
        "mismatched_rows": sorted(mismatched),
        "tolerance": 1e-9,
    }


def read_table_s3(path: str | Path) -> tuple[list[str], tuple[GrowthBlock, ...]]:
    """Table S3's row labels and its two background blocks (WT first, then dlpp)."""
    header = _header(path, "ST3", 2)
    _expect(header[1:7], SIX_DOSE_HEADERS, "Table S3 WT block")
    _expect(header[8:14], SIX_DOSE_HEADERS, "Table S3 dlpp block")
    group = _header(path, "ST3", 1)
    for index, want in (
        (1, "Normalized Growth of WT CRISPRi Co"),
        (8, "Normalized Growth of "),
    ):
        got = _cell(group[index]) or ""
        if not got.startswith(want):
            raise TableLayoutError(f"Table S3 group header {index} is {got!r}")
    rows = _rows(path, "ST3", skip=2)
    if len(rows) != S3_ROWS:
        raise TableLayoutError(f"Table S3 has {len(rows)} rows, not {S3_ROWS}")
    labels = [str(_cell(row[0])) for row in rows]
    blocks = tuple(
        GrowthBlock(
            medium=MEDIUM_MOPS,
            background=background,
            knockdown=None,
            doses=SIX_DOSES_NG_PER_ML,
            values=tuple(
                _floats(row, range(first, first + 6), f"Table S3 row {index}")
                for index, row in enumerate(rows)
            ),
        )
        for background, first in ((None, 1), (LPP_SYMBOL, 8))
    )
    return labels, blocks


def read_table_s4a(path: str | Path) -> tuple[list[str], tuple[GrowthBlock, ...]]:
    """Table S4A's deletion labels and its four CRISPRi-plasmid blocks."""
    group = _header(path, "ST4A", 1)
    header = _header(path, "ST4A", 2)
    if _cell(header[0]) != "Keio_Deletion":
        raise TableLayoutError(f"Table S4A column A is {_cell(header[0])!r}")
    starts = (1, 5, 9, 13)
    for (knockdown, want), first in zip(S4A_BLOCKS, starts, strict=True):
        if _cell(group[first]) != want:
            raise TableLayoutError(
                f"Table S4A block at column {first} is {_cell(group[first])!r}, "
                f"not {want!r}"
            )
        _expect(header[first : first + 3], THREE_DOSE_HEADERS, f"Table S4A {knockdown}")
    rows = _rows(path, "ST4A", skip=2)
    if len(rows) != S4A_ROWS:
        raise TableLayoutError(f"Table S4A has {len(rows)} rows, not {S4A_ROWS}")
    labels = [str(_cell(row[0])) for row in rows]
    distinct = len(set(labels))
    if distinct != S4A_DISTINCT_LABELS:
        raise TableLayoutError(
            f"Table S4A has {distinct} distinct labels, not {S4A_DISTINCT_LABELS}"
        )
    blocks = tuple(
        GrowthBlock(
            medium=MEDIUM_LB,
            background=None,
            knockdown=knockdown,
            doses=THREE_DOSES_NG_PER_ML,
            values=tuple(
                _floats(row, range(first, first + 3), f"Table S4A row {index}")
                for index, row in enumerate(rows)
            ),
        )
        for (knockdown, _), first in zip(S4A_BLOCKS, starts, strict=True)
    )
    return labels, blocks


# --------------------------------------------------------------------------- #
# Joining the growth tables to the collection layout
# --------------------------------------------------------------------------- #
class PlateJoin(BaseModel):
    """The positional join of a growth table's rows onto Table S1's filled wells."""

    model_config = ConfigDict(extra="forbid")

    table: str
    rows: int
    disagreements: list[tuple[str, str, str]]


def join_growth_rows_to_wells(
    labels: Sequence[str], wells: Sequence[CollectionWell], *, table: str
) -> PlateJoin:
    """Assert a growth table's rows ARE Table S1's filled wells, position for position.

    The join is positional because neither growth table releases a well, and it is
    checkable because both tables name the same target: every row's label must equal its
    well's symbol apart from the single known transposition typo, and any second
    disagreement stops the build rather than being silently absorbed.
    """
    if len(labels) != len(wells):
        raise TableLayoutError(
            f"{table} has {len(labels)} rows against Table S1's {len(wells)} wells"
        )
    disagreements = [
        (well.well, well.symbol, label)
        for label, well in zip(labels, wells, strict=True)
        if label != well.symbol
    ]
    expected = [(S1_TYPO_WELL, S1_TYPO_SYMBOL, GROWTH_TABLE_TYPO_SYMBOL)]
    if disagreements != expected:
        raise TableLayoutError(
            f"{table} disagrees with Table S1's plate order at {disagreements}, "
            f"not only at {expected}"
        )
    return PlateJoin(table=table, rows=len(labels), disagreements=disagreements)


def check_fifty_percent_reduction(
    wells: Sequence[CollectionWell], blocks: Sequence[GrowthBlock]
) -> dict[str, Any]:
    """Reproduce the paper's "272 from the 356 strains" at 100 ng/mL aTc on LB."""
    lb_block = next(block for block in blocks if block.medium == MEDIUM_LB)
    dose_index = lb_block.doses.index(100.0)
    guides = [
        (well, values)
        for well, values in zip(wells, lb_block.values, strict=True)
        if not well.is_empty_vector
    ]
    reduced = [
        well.symbol
        for well, values in guides
        if values[dose_index] <= FIFTY_PERCENT_THRESHOLD
    ]
    distinct = len({well.symbol for well, _ in guides})
    if len(guides) != S2A_GUIDE_ROWS or distinct != S2A_DISTINCT_TARGETS:
        raise TableLayoutError(
            f"Table S2A has {len(guides)} guide rows over {distinct} symbols, not "
            f"{S2A_GUIDE_ROWS} over {S2A_DISTINCT_TARGETS}"
        )
    if len(reduced) != FIFTY_PERCENT_COUNT:
        raise TableLayoutError(
            f"{len(reduced)} of {len(guides)} guide rows are at or below "
            f"{FIFTY_PERCENT_THRESHOLD} at 100 ng/mL aTc on LB, not {FIFTY_PERCENT_COUNT}"
        )
    return {
        "rule": "fifty_percent_reduction_count_reproduced",
        "quote": str(SOURCED_VALUES["fifty_percent_count"].quote),
        "guide_rows": len(guides),
        "distinct_symbols": distinct,
        "at_or_below_half": len(reduced),
        "threshold": FIFTY_PERCENT_THRESHOLD,
    }


def check_empty_vector_mean_is_one(
    wells: Sequence[CollectionWell], blocks: Sequence[GrowthBlock], *, table: str
) -> dict[str, Any]:
    """MEASURE that the 17 empty-vector wells mean exactly 1 in every column."""
    worst = 0.0
    columns = 0
    for block in blocks:
        for position in range(len(block.doses)):
            vectors = [
                values[position]
                for well, values in zip(wells, block.values, strict=True)
                if well.is_empty_vector
            ]
            worst = max(worst, abs(sum(vectors) / len(vectors) - 1.0))
            columns += 1
    if worst > 1e-9:
        raise TableLayoutError(
            f"{table}: the empty-vector mean is not 1 in every column "
            f"(worst |mean - 1| = {worst:.3g})"
        )
    return {
        "rule": "empty_vector_mean_is_one_per_column",
        "table": table,
        "columns": columns,
        "empty_vector_wells": S1_EMPTY_VECTOR_WELLS,
        "worst_abs_deviation": worst,
    }


# --------------------------------------------------------------------------- #
# Resolving labels to BW25113 locus tags
# --------------------------------------------------------------------------- #
def canonical_symbol(genome: EcoliK12BW25113Genome, tag: str) -> str:
    """The genome's own gene symbol for ``tag`` when it resolves back to ``tag``.

    ONE spelling per locus, because a locus named two ways splits into two graph nodes
    and fails the shared ``canonical_gene_names`` rule. This release needs it: Table S4A
    names BW25113_1472 by the Keio strain id ``JW1468`` while Table S1 names the same
    locus ``yddL``, and the same holds for BW25113_2858 (``JW5459`` against ``ygeN``).
    The released spelling is not lost -- it is the perturbation's
    ``identifier_mapping.source_identifier``. A locus with no symbol, or whose symbol
    resolves elsewhere, is named by its tag.
    """
    symbol = genome.genbank.loci[tag].symbol
    if symbol is None:
        return tag
    return symbol if genome.resolve_gene_name(symbol).systematic_name == tag else tag


class Resolution(BaseModel):
    """Which labels of one table can be written, and which rule removed the rest."""

    model_config = ConfigDict(extra="forbid")

    locus_by_label: dict[str, str]
    #: The pinned annotation's own symbol for each writable label's locus.
    symbol_by_label: dict[str, str]
    unwritable: dict[str, list[str]]
    reconciliation: LocusTagReconciliation


def resolve_labels(
    labels: Sequence[str],
    genome: EcoliK12BW25113Genome,
    *,
    label: str,
    min_resolved: float,
) -> Resolution:
    """Map each distinct label to a BW25113 locus tag, attributing every failure."""
    distinct = sorted(set(labels))
    stored, reconciliation = reconcile_locus_tags(
        genome, pd.Series(distinct), label=label
    )
    reconciliation.require_resolved(min_resolved)
    unwritable = {
        DROP_NOT_IN_ANNOTATION: sorted(reconciliation.retired_kept),
        DROP_MERGED_LOCUS: sorted(reconciliation.kept_on_collision),
        DROP_AMBIGUOUS: sorted(reconciliation.ambiguous_kept),
    }
    removed = set().union(*unwritable.values())
    locus_by_label = {
        name: str(tag)
        for name, tag in zip(distinct, stored.tolist(), strict=True)
        if name not in removed
    }
    outside = sorted(
        name
        for name, tag in locus_by_label.items()
        if genome.resolve_gene_name(tag).status
        not in (GeneNameStatus.CURRENT, GeneNameStatus.NON_GENE_FEATURE)
    )
    if outside:
        raise RuntimeError(
            f"{label}: {len(outside)} kept tags are not loci of the pinned assembly: "
            f"{outside[:10]}"
        )
    return Resolution(
        locus_by_label=locus_by_label,
        symbol_by_label={
            name: canonical_symbol(genome, tag) for name, tag in locus_by_label.items()
        },
        unwritable=unwritable,
        reconciliation=reconciliation,
    )


# --------------------------------------------------------------------------- #
# Records
# --------------------------------------------------------------------------- #
def knockdown(
    well: CollectionWell, locus_tag: str, symbol: str | None = None
) -> BacterialCrisprInterferencePerturbation:
    """The CRISPRi knockdown one collection well carries, on its BW25113 locus.

    ``symbol`` is the pinned annotation's own spelling for ``locus_tag`` and defaults to
    the released one, which is what Table S1 spells it.
    """
    return BacterialCrisprInterferencePerturbation(
        systematic_gene_name=locus_tag,
        perturbed_gene_name=symbol or well.symbol,
        gene_namespace=BW25113_NAMESPACE,
        identifier_mapping=DerivedIdentifierMapping(
            source_identifier=well.symbol, route="gene_symbol"
        ),
        crispr=CrisprConstruct(
            effector=EFFECTOR, guide_sequence=well.spacer, n_guides=1
        ),
    )


def deletion(
    label: str, locus_tag: str, *, symbol: str | None = None, well: str | None = None
) -> BacterialDeletionPerturbation:
    """One kanamycin-cassette deletion of the Keio collection, against BW25113.

    ``label`` is the released spelling and rides on ``identifier_mapping``; ``symbol`` is
    the pinned annotation's own spelling and defaults to the released one.
    """
    return BacterialDeletionPerturbation(
        systematic_gene_name=locus_tag,
        perturbed_gene_name=symbol or label,
        gene_namespace=BW25113_NAMESPACE,
        identifier_mapping=DerivedIdentifierMapping(
            source_identifier=label, route="gene_symbol"
        ),
        collection=KEIO_COLLECTION,
        cassette=KEIO_CASSETTE,
        construction=StrainConstruction(well=well) if well is not None else None,
    )


def build_reference(dataset_name: str) -> BacterialFitnessExperimentReference:
    """The empty vector on the uninduced rich plate: the normalization baseline."""
    return BacterialFitnessExperimentReference(
        dataset_name=dataset_name,
        genome_reference=assembly_reference(BW25113_STRAIN),
        environment_reference=environment(MEDIUM_LB, 0.0, timed=True),
        phenotype_reference=reference_phenotype(),
    )


class Row(BaseModel):
    """One writable (strain, screen) row: its genotype and its per-dose values."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    genotype: Genotype
    screen_id: str
    medium: str
    doses: tuple[float, ...]
    values: tuple[float, ...]
    timed: bool


def _guide_rows(
    wells: Sequence[CollectionWell],
    blocks: Sequence[GrowthBlock],
    resolution: Resolution,
    repeated_spacers: frozenset[str],
    *,
    screen_id: str,
) -> Iterator[Row]:
    """Writable rows of a CRISPRi-collection table, one per (well, block)."""
    for block in blocks:
        background: list[BacterialDeletionPerturbation] = []
        if block.background is not None:
            background = [
                deletion(
                    block.background,
                    resolution.locus_by_label[block.background],
                    symbol=resolution.symbol_by_label[block.background],
                )
            ]
        for well, values in zip(wells, block.values, strict=True):
            if well.is_empty_vector:
                continue
            locus_tag = resolution.locus_by_label.get(well.symbol)
            if locus_tag is None or well.spacer in repeated_spacers:
                continue
            yield Row(
                genotype=Genotype(
                    perturbations=[
                        knockdown(
                            well, locus_tag, resolution.symbol_by_label[well.symbol]
                        ),
                        *background,
                    ]
                ),
                screen_id=screen_id,
                medium=block.medium,
                doses=block.doses,
                values=values,
                timed=True,
            )


def _keio_rows(
    labels: Sequence[str],
    blocks: Sequence[GrowthBlock],
    resolution: Resolution,
    repeated_labels: frozenset[str],
    knockdowns: Mapping[str, BacterialCrisprInterferencePerturbation],
) -> Iterator[Row]:
    """Writable rows of Table S4A, one per (deletion row, CRISPRi plasmid block)."""
    for block in blocks:
        plasmid = [knockdowns[block.knockdown]] if block.knockdown is not None else []
        for label, values in zip(labels, block.values, strict=True):
            locus_tag = resolution.locus_by_label.get(label)
            if locus_tag is None or label in repeated_labels:
                continue
            yield Row(
                genotype=Genotype(
                    perturbations=[
                        deletion(
                            label, locus_tag, symbol=resolution.symbol_by_label[label]
                        ),
                        *plasmid,
                    ]
                ),
                screen_id=SCREEN_S4A,
                medium=block.medium,
                doses=block.doses,
                values=values,
                timed=False,
            )


# --------------------------------------------------------------------------- #
# Raw mirror (the loader reads the mirror, never a live URL)
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    """``DATA_ROOT`` from the environment (the mirrors and build tree live under it)."""
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/rachwalskiMobileCRISPRiCollection2024``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def library_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-library/rachwalskiMobileCRISPRiCollection2024``."""
    return Path(data_root or _data_root()) / LIBRARY_DIR_REL


def _sha256(path: str | Path) -> str:
    """Streaming sha256 of a file."""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def retrieve_raw_files(dest_dir: str | Path) -> dict[str, Path]:
    """Run each file's recorded retriever and write the verified bytes to ``dest_dir``.

    The recorded ``RetrievalRecord`` is what runs, so this IS the re-runnable retrieval;
    a byte mismatch raises before anything is written.
    """
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}
    for raw in RAW_FILES:
        path = dest / raw.name
        write_verified(run_retriever(raw.retrieval), path, raw.sha256, raw.source_url)
        out[raw.name] = path
    return out


def deposit_raw_mirror(
    *, sources: Mapping[str, str | Path], data_root: str | None = None
) -> Path:
    """Write the raw mirror from already-retrieved files plus its ``manifest.json``.

    ``sources`` maps every name in ``RAW_FILES`` to a local file. Idempotent by sha256: a
    mirror file with the pinned hash is left alone, and one with any other hash raises
    rather than being overwritten.
    """
    missing = sorted(set(DATA_SHA256) - set(sources))
    if missing:
        raise KeyError(f"no source given for {missing}")
    root = raw_mirror_dir(data_root)
    records: list[ArtifactRecord] = []
    for raw in RAW_FILES:
        src = Path(sources[raw.name])
        got = _sha256(src)
        if got != raw.sha256:
            raise RuntimeError(
                f"{src} sha256 mismatch: got {got}, expected {raw.sha256}"
            )
        dest = root / raw.mirror_relpath
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            if _sha256(dest) != raw.sha256:
                raise RuntimeError(f"{dest} exists with a different sha256; refusing")
        else:
            shutil.copy2(src, dest)
        records.append(
            ArtifactRecord(
                path=raw.mirror_relpath,
                role=ROLE_RAW_DATA,
                bytes=dest.stat().st_size,
                sha256=raw.sha256,
                source=raw.source_url,
                original_filename=raw.name,
                retrieval=raw.retrieval,
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=DOI,
        title=TITLE,
        files=records,
        si_data_sources=[raw.source_url for raw in RAW_FILES]
        + [f"https://doi.org/{DOI}"],
        si_expected=list(NOT_MIRRORED),
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def load_manifest(data_root: str | None = None) -> Manifest:
    """Read the raw mirror's ``manifest.json``."""
    path = raw_mirror_dir(data_root) / "manifest.json"
    return Manifest.model_validate_json(path.read_text())


def manifest_sha256(manifest: Manifest, relpath: str) -> str:
    """The recorded sha256 of one mirror file."""
    for record in manifest.files:
        if record.path == relpath:
            return record.sha256
    raise KeyError(f"{relpath} is not in the raw-mirror manifest")


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
#: Records of the full build, MEASURED on the pinned bytes: 348 writable collection
#: wells x 12 columns for Table S2A, the same 348 x 12 for Table S3, and 3,379 writable
#: Table S4A rows x 12 columns.
EXPECTED_WRITABLE_WELLS = 348
EXPECTED_WRITABLE_KEIO_ROWS = 3379
EXPECTED_RECORDS = (
    EXPECTED_WRITABLE_WELLS * 12
    + EXPECTED_WRITABLE_WELLS * 12
    + EXPECTED_WRITABLE_KEIO_ROWS * 12
)
#: Records whose genotype carries BOTH a knockdown and a deletion.
EXPECTED_CROSSED_RECORDS = EXPECTED_WRITABLE_WELLS * 6 + EXPECTED_WRITABLE_KEIO_ROWS * 9


@register_dataset
class CrispriCrossRachwalski2024Dataset(ExperimentDataset):
    """Rachwalski 2024 normalized colony growth of CRISPRi knockdown by deletion crosses."""

    #: The host of every screen and the background of the Keio collection.
    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = BW25113_STRAIN

    def __init__(
        self,
        root: str = DATASET_ROOT_REL,
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return BacterialFitnessExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return BacterialFitnessExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """Every consumed workbook, linked from the raw mirror."""
        return [raw.name for raw in RAW_FILES]

    def download(self) -> None:
        """Link each mirror file into ``raw/`` after checking the manifest and sha256."""
        data_root = _data_root()
        manifest = load_manifest(data_root)
        os.makedirs(self.raw_dir, exist_ok=True)
        for raw in RAW_FILES:
            check_manifest_pin(
                raw.mirror_relpath,
                manifest_sha256(manifest, raw.mirror_relpath),
                raw.sha256,
            )
            src = raw_mirror_dir(data_root) / raw.mirror_relpath
            if not src.exists():
                raise RuntimeError(f"required raw artifact missing from mirror: {src}")
            link_verified(src, osp.join(self.raw_dir, raw.name), raw.sha256)
        log.info(
            "Rachwalski 2024 raw files linked into %s (sha256 verified)", self.raw_dir
        )

    def _raw(self, name: str) -> str:
        return osp.join(self.raw_dir, name)

    def _genome(self) -> EcoliK12BW25113Genome:
        """The BW25113 genome, injected by a build or opened from its cache root."""
        if self.ecoli_genome is None:  # a direct run; the build entry points inject it
            self.ecoli_genome = bacterial_genome("ecoli", self.REFERENCE_STRAIN)
        genome = self.ecoli_genome
        if not isinstance(genome, EcoliK12BW25113Genome):
            raise TypeError(
                f"{self.name} needs the BW25113 genome, got {type(genome).__name__}"
            )
        return genome

    @post_process
    def process(self) -> None:
        """Parse the four workbooks into per-(strain, condition) records; write LMDB."""
        verify_raw_files(self.raw_dir, DATA_SHA256)
        wells = read_table_s1(self._raw(TABLE_S1_FILE))
        s2a_labels, s2a_blocks = read_table_s2a(self._raw(TABLE_S2_FILE))
        s3_labels, s3_blocks = read_table_s3(self._raw(TABLE_S3_FILE))
        s4a_labels, s4a_blocks = read_table_s4a(self._raw(TABLE_S4_FILE))

        joins = [
            join_growth_rows_to_wells(s2a_labels, wells, table=SCREEN_S2A),
            join_growth_rows_to_wells(s3_labels, wells, table=SCREEN_S3),
        ]
        checks = [
            check_s2a_block_order(self._raw(TABLE_S2_FILE), s2a_labels, s2a_blocks),
            check_fifty_percent_reduction(wells, s2a_blocks),
            check_empty_vector_mean_is_one(wells, s2a_blocks, table=SCREEN_S2A),
            check_empty_vector_mean_is_one(wells, s3_blocks, table=SCREEN_S3),
        ]

        genome = self._genome()
        targets = [well.symbol for well in wells if not well.is_empty_vector]
        target_resolution = resolve_labels(
            targets + [LPP_SYMBOL],
            genome,
            label=f"{self.name}:targets",
            min_resolved=MIN_RESOLVED_TARGETS,
        )
        if LPP_SYMBOL not in target_resolution.locus_by_label:
            raise RuntimeError(
                f"{self.name}: lpp does not resolve, so Table S3's query strain cannot "
                "be written"
            )
        keio_resolution = resolve_labels(
            s4a_labels,
            genome,
            label=f"{self.name}:keio",
            min_resolved=MIN_RESOLVED_KEIO,
        )

        spacer_counts = Counter(
            well.spacer for well in wells if not well.is_empty_vector
        )
        repeated_spacers = frozenset(
            spacer
            for spacer, count in spacer_counts.items()
            if count > 1 and spacer is not None
        )
        keio_counts = Counter(
            label for label in s4a_labels if label in keio_resolution.locus_by_label
        )
        repeated_labels = frozenset(
            label for label, count in keio_counts.items() if count > 1
        )

        knockdowns = {
            symbol: knockdown(
                well,
                target_resolution.locus_by_label[well.symbol],
                target_resolution.symbol_by_label[well.symbol],
            )
            for symbol, _ in S4A_BLOCKS
            if symbol is not None
            for well in wells
            if well.symbol == symbol and well.symbol in target_resolution.locus_by_label
        }
        missing = sorted(
            symbol
            for symbol, _ in S4A_BLOCKS
            if symbol is not None and symbol not in knockdowns
        )
        if missing:
            raise RuntimeError(
                f"{self.name}: Table S4A's blocks name {missing}, which Table S1 has no "
                "resolvable well for"
            )

        rows = [
            *_guide_rows(
                wells,
                s2a_blocks,
                target_resolution,
                repeated_spacers,
                screen_id=SCREEN_S2A,
            ),
            *_guide_rows(
                wells,
                s3_blocks,
                target_resolution,
                repeated_spacers,
                screen_id=SCREEN_S3,
            ),
            *_keio_rows(
                s4a_labels, s4a_blocks, keio_resolution, repeated_labels, knockdowns
            ),
        ]

        reference = build_reference(self.name)
        publication = Publication(doi=DOI, doi_url=f"https://doi.org/{DOI}")
        environments = {
            (medium, dose, timed): environment(medium, dose, timed=timed)
            for medium in MEDIA_BY_LABEL
            for dose in set(SIX_DOSES_NG_PER_ML) | set(THREE_DOSES_NG_PER_ML)
            for timed in (True, False)
        }

        os.makedirs(self.preprocess_dir, exist_ok=True)
        os.makedirs(self.processed_dir, exist_ok=True)
        lmdb_env, interned_env = self._open_write_lmdb(
            osp.join(self.processed_dir, "lmdb")
        )
        idx = 0
        crossed = 0
        with lmdb_env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for row in tqdm(rows, desc="rachwalski2024"):
                is_crossed = len(row.genotype.perturbations) > 1
                for dose, growth in zip(row.doses, row.values, strict=True):
                    experiment = BacterialFitnessExperiment(
                        dataset_name=self.name,
                        genotype=row.genotype,
                        environment=environments[(row.medium, dose, row.timed)],
                        phenotype=phenotype(growth, row.screen_id),
                    )
                    txn.put(
                        f"{idx}".encode(),
                        self._intern_record(experiment, reference, publication, itxn),
                    )
                    idx += 1
                    crossed += is_crossed
        lmdb_env.close()
        interned_env.close()

        if idx != EXPECTED_RECORDS or crossed != EXPECTED_CROSSED_RECORDS:
            raise RuntimeError(
                f"wrote {idx} records ({crossed} two-perturbation), expected "
                f"{EXPECTED_RECORDS} ({EXPECTED_CROSSED_RECORDS})"
            )
        self._write_ledgers(
            wells=wells,
            s4a_labels=s4a_labels,
            target_resolution=target_resolution,
            keio_resolution=keio_resolution,
            repeated_spacers=repeated_spacers,
            repeated_labels=repeated_labels,
            joins=joins,
            checks=checks,
            kept_records=idx,
            crossed_records=crossed,
        )
        log.info(
            "Rachwalski2024: wrote %d records (%d two-perturbation) from %d writable "
            "collection wells over %d target loci and %d writable Table S4A rows; "
            "dropped %d target labels and %d Keio labels",
            idx,
            crossed,
            sum(
                1
                for well in wells
                if not well.is_empty_vector
                and well.symbol in target_resolution.locus_by_label
                and well.spacer not in repeated_spacers
            ),
            len(target_resolution.locus_by_label) - 1,
            sum(
                1
                for label in s4a_labels
                if label in keio_resolution.locus_by_label
                and label not in repeated_labels
            ),
            sum(len(names) for names in target_resolution.unwritable.values()),
            sum(len(names) for names in keio_resolution.unwritable.values()),
        )

    def _write_ledgers(
        self,
        *,
        wells: Sequence[CollectionWell],
        s4a_labels: Sequence[str],
        target_resolution: Resolution,
        keio_resolution: Resolution,
        repeated_spacers: frozenset[str],
        repeated_labels: frozenset[str],
        joins: Sequence[PlateJoin],
        checks: Sequence[Mapping[str, Any]],
        kept_records: int,
        crossed_records: int,
    ) -> None:
        """The drop log, the identifier ledger, the build checks and the accounting."""
        out = Path(self.preprocess_dir)
        guide_wells = [well for well in wells if not well.is_empty_vector]
        repeated_wells = [
            well for well in guide_wells if well.spacer in repeated_spacers
        ]
        repeated_symbols = sorted({well.symbol for well in repeated_wells})
        dropped: dict[str, dict[str, Any]] = {
            DROP_EMPTY_VECTOR: {
                "collection_wells": S1_EMPTY_VECTOR_WELLS,
                "rows": S1_EMPTY_VECTOR_WELLS * 2,
                "labels": [EMPTY_VECTOR],
            },
            DROP_NOT_IN_ANNOTATION: {
                "labels": target_resolution.unwritable[DROP_NOT_IN_ANNOTATION],
                "keio_labels": keio_resolution.unwritable[DROP_NOT_IN_ANNOTATION],
            },
            DROP_MERGED_LOCUS: {
                "labels": target_resolution.unwritable[DROP_MERGED_LOCUS],
                "keio_labels": keio_resolution.unwritable[DROP_MERGED_LOCUS],
            },
            DROP_AMBIGUOUS: {
                "labels": target_resolution.unwritable[DROP_AMBIGUOUS],
                "keio_labels": keio_resolution.unwritable[DROP_AMBIGUOUS],
            },
            DROP_REPEATED_CONSTRUCT: {
                "labels": repeated_symbols,
                "wells": [well.well for well in repeated_wells],
                "spacers": sorted(repeated_spacers),
                # the same rows in Table S2A and in Table S3
                "rows": len(repeated_wells) * 2,
            },
            DROP_DUPLICATE_LABEL: {
                "labels": sorted(repeated_labels),
                "rows": sum(1 for label in s4a_labels if label in repeated_labels),
            },
        }
        (out / "dropped_records.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "kept_records": kept_records,
                    "crossed_records": crossed_records,
                    "released_cells": {
                        SCREEN_S2A: S2A_ROWS * 12,
                        SCREEN_S3: S3_ROWS * 12,
                        SCREEN_S4A: S4A_ROWS * 12,
                    },
                    "rules": [
                        {"rule": rule, "description": description, **dropped[rule]}
                        for rule, description in ROW_RULES
                    ],
                },
                indent=2,
            )
        )
        (out / "identifier_reconciliation.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "identifier_route": "gene_symbol",
                    "targets": target_resolution.reconciliation.model_dump(mode="json"),
                    "keio": keio_resolution.reconciliation.model_dump(mode="json"),
                    "collection_wells": len(wells),
                    "guide_wells": len(guide_wells),
                    "distinct_spacers": len({well.spacer for well in guide_wells}),
                },
                indent=2,
            )
        )
        (out / "extraction.json").write_text(
            json.dumps(
                {
                    "raw_sha256": DATA_SHA256,
                    "plate_order_joins": [
                        join.model_dump(mode="json") for join in joins
                    ],
                    "checks": list(checks),
                },
                indent=2,
            )
        )
        (out / "build_accounting.json").write_text(
            json.dumps(
                {
                    "dataset": self.name,
                    "sourced_values": {
                        key: value.model_dump(mode="json")
                        for key, value in SOURCED_VALUES.items()
                    },
                    "not_loaded": list(NOT_MIRRORED),
                    "unpinned_environment_values": [
                        {
                            "field": "media",
                            "screen": SCREEN_S4A,
                            "reading": "LB",
                            "why": str(SOURCED_VALUES["keio_medium_unstated"].note),
                        }
                    ],
                },
                indent=2,
            )
        )

    def preprocess_raw(
        self, df: pd.DataFrame, preprocess: dict[str, Any] | None = None
    ) -> pd.DataFrame:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled by ``_guide_rows`` / ``_keio_rows``."""
        raise NotImplementedError


# --------------------------------------------------------------------------- #
# Verification (L0-L4) of a built LMDB
# --------------------------------------------------------------------------- #
VERIFIER_PROVENANCE = Provenance(
    source_uri=(
        f"$DATA_ROOT/{RAW_DIR_REL}/data/{TABLE_S1_FILE} (Table S1, the collection "
        f"layout and its spacers) + data/{TABLE_S2_FILE} (Table S2A, LB and "
        f"MOPS-glucose) + data/{TABLE_S3_FILE} (Table S3, WT and dlpp) + "
        f"data/{TABLE_S4_FILE} (Table S4A, the CRISPRi-Keio cross); raw mirror with "
        f"manifest.json, pmc_cloud retrieval re-verified {RETRIEVED_AT}"
    ),
    citation_key=CITATION_KEY,
    sha256=RAW_FILES_BY_NAME[TABLE_S2_FILE].sha256,
    method=(
        "One BacterialFitnessExperiment per (strain, medium, aTc dose). fitness = the "
        "released average normalized colony growth, verbatim: for Tables S2A and S3 the "
        "strain's colony density over the mean of the 17 empty-vector wells of that "
        "plate (MEASURED to be exactly 1.0000 in each of the 24 columns), for Table S4A "
        "the two-pass inter-quartile-mean normalized density. n_samples=2 from 'The "
        "average of the two technical replicates was then calculated and is reported in "
        "Table S2' and 'CRISPRi Keio screens were conducted in technical duplicate'; "
        "sample_unit is a typed gap because the same methods section and three figure "
        "legends call the same duplicate biological, and the dispersion is a typed gap "
        "because no table releases one. screen_id names the released table, which is "
        "what keeps Table S2A's and Table S3's two independent MOPS runs of the same "
        "strains distinguishable. Genotype = a BacterialCrisprInterferencePerturbation "
        "(dCas9, the 20 nt Table S1 spacer of the row's well, n_guides=1) and/or a "
        "BacterialDeletionPerturbation of the Keio collection (kanamycin cassette); the "
        "crossed rows carry both. Environment = solid LB (15 g/L agar) or solid "
        "MOPS-glucose minimal (1.5% w/v agar, glucose amount unset) at 37 C, aerobic, "
        "plus the column's aTc as a SmallMoleculePerturbation in ug/mL; the 0 ng/mL "
        "columns carry no perturbation because no aTc was added. duration_hours is 16 h "
        "for Tables S2A and S3 and a typed gap for Table S4A. Reference = the empty "
        "vector on the uninduced rich plate, fitness 1.0. Build checks: both growth "
        "tables ARE Table S1's plate order with exactly one label disagreement (well "
        "G15, gpsA vs gspA); Table S2C fixes the LB-then-MOPS block order to 1e-9; the "
        "paper's own 272-of-356 count at 100 ng/mL aTc on LB is reproduced exactly"
    ),
    page="Tables S1 to S4",
    retrieved=RETRIEVED_AT,
)


def stored_tags_are_loci(
    records: Sequence[Mapping[str, Any]], genome: EcoliK12BW25113Genome
) -> LevelResult:
    """SUPPLEMENTARY L1: every stored tag resolves to itself as a locus of the genome.

    The shared ``canonical_gene_names`` rule requires status ``current``, which a
    pseudogene locus never has (the bacterial resolver returns ``non_gene_feature``,
    naming the same tag). This row accepts a gene or a pseudogene locus that resolves to
    itself, and counts the pseudogenes. It is added beside the shared row, never in its
    place.
    """
    tags = sorted(
        {
            str(perturbation["systematic_gene_name"])
            for record in records
            for perturbation in record["experiment"]["genotype"]["perturbations"]
        }
    )
    pseudogenes = 0
    outside: list[str] = []
    for tag in tags:
        resolved = genome.resolve_gene_name(tag)
        if resolved.systematic_name != tag or resolved.status not in (
            GeneNameStatus.CURRENT,
            GeneNameStatus.NON_GENE_FEATURE,
        ):
            outside.append(tag)
        elif resolved.status is GeneNameStatus.NON_GENE_FEATURE:
            pseudogenes += 1
    return LevelResult(
        level=Level.L1,
        name="stored_tags_are_loci",
        passed=not outside,
        message=(
            f"SUPPLEMENTARY: {len(tags)} stored tags are loci of the pinned assembly "
            f"({pseudogenes} pseudogene loci)"
            if not outside
            else f"{len(outside)} stored tags are not loci: {outside[:10]}"
        ),
        details={
            "n_tags": len(tags),
            "n_pseudogene_loci": pseudogenes,
            "outside": outside[:50],
        },
    )


def crossed_genotypes_carry_both_kinds(
    records: Sequence[Mapping[str, Any]],
) -> LevelResult:
    """SUPPLEMENTARY L1: a two-perturbation record is one knockdown and one deletion.

    The row this dataset exists for. A genotype of two perturbations must be exactly one
    ``bacterial_crispr_interference`` and one ``bacterial_deletion``; two of either would
    be a cross the release never ran.
    """
    kinds = Counter(
        tuple(
            sorted(
                perturbation["perturbation_type"]
                for perturbation in record["experiment"]["genotype"]["perturbations"]
            )
        )
        for record in records
    )
    allowed = {
        ("bacterial_crispr_interference",),
        ("bacterial_deletion",),
        ("bacterial_crispr_interference", "bacterial_deletion"),
    }
    unexpected = {key: count for key, count in kinds.items() if key not in allowed}
    crossed = kinds.get(("bacterial_crispr_interference", "bacterial_deletion"), 0)
    return LevelResult(
        level=Level.L1,
        name="crossed_genotypes_carry_both_kinds",
        passed=not unexpected and crossed == EXPECTED_CROSSED_RECORDS,
        message=(
            f"SUPPLEMENTARY: {crossed} records cross a knockdown with a deletion"
            if not unexpected and crossed == EXPECTED_CROSSED_RECORDS
            else f"unexpected genotype shapes {unexpected} or {crossed} crossed records"
        ),
        details={
            "shapes": {"+".join(key): count for key, count in sorted(kinds.items())},
            "expected_crossed": EXPECTED_CROSSED_RECORDS,
        },
    )


def run_verification(data_root: str | None = None) -> Any:
    """Run the L0-L4 fitness gate over the built dev-tree LMDB and write the report."""
    from torchcell.verification.fitness import fitness_gene_set, verify_fitness_dataset
    from torchcell.verification.runners import load_records

    root = osp.join(data_root or _data_root(), DATASET_ROOT_REL)
    records = load_records(root)
    genome = bacterial_genome("ecoli", BW25113_STRAIN, data_root)
    if not isinstance(genome, EcoliK12BW25113Genome):
        raise TypeError(f"expected the BW25113 genome, got {type(genome).__name__}")
    report = verify_fitness_dataset(
        records,
        dataset_name=CrispriCrossRachwalski2024Dataset.__name__,
        provenance=VERIFIER_PROVENANCE,
        expected_count=EXPECTED_RECORDS,
        resolve_gene_name=genome.resolve_gene_name,
        sgd_genes=set(genome.genbank.loci),
        gene_universe_label="BW25113 GenBank loci",
    )
    report.add(stored_tags_are_loci(records, genome))
    report.add(crossed_genotypes_carry_both_kinds(records))
    out = osp.join(root, "preprocess", "verification_report.json")
    with open(out, "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    log.info(
        "Rachwalski2024 verification: %d records over %d genes -> %s",
        len(records),
        len(fitness_gene_set(records)),
        out,
    )
    return report


def main() -> int:
    """Build the dataset and verify it, for interactive debugging."""
    from dotenv import load_dotenv

    load_dotenv()
    logging.basicConfig(level=logging.INFO)
    dataset = CrispriCrossRachwalski2024Dataset(
        root=osp.join(_data_root(), DATASET_ROOT_REL)
    )
    print(f"len = {len(dataset)}")
    print(dataset[0])
    dataset.close_lmdb()
    report = run_verification()
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
