#!/usr/bin/env python
# experiments/database/scripts/yoshikawa2009_release_checks.py
# [[experiments.database.scripts.yoshikawa2009_release_checks]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/database/scripts/yoshikawa2009_release_checks.py
r"""Measurements on the Yoshikawa 2009 release that the loader, its note and its PR cite.

Reads, read-only, the two deposited copies of Table S1 in the raw mirror
(``$DATA_ROOT/torchcell-raw/yoshikawaComprehensivePhenotypicAnalysis2009/si/``), the
and the paper OCR in the literature mirror:

1. the source question: the publisher ``.xls`` against YeastPhenome's ``.xlsx``, cell by
   cell (sheet names, dimensions, merged ranges, every cell's value and kind), plus each
   file's own authorship metadata (the ``.xls`` OLE summary stream and BIFF WRITEACCESS
   user, the ``.xlsx`` docProps);
2. the 5% ethanol columns: whether each copy carries them, and every place the paper's
   OCR writes a 5 percent (``5 \%``) or names the condition;
3. the calls: Table S1's per-condition call counts against the paper's Table 2, how
   many calls the both-cultures rule reproduces with Table 1's printed (rounded)
   thresholds, and the threshold interval the 5% ethanol calls imply (the largest
   sensitive and the smallest non-sensitive "larger of the two cultures", and the
   matching pair for tolerant);
4. the name columns as the loader stores them (#1045: no genome lookup): Column A (the
   ORF, upper-cased) and Column B (the release's gene name) verbatim, the lower-case
   Column A cells, the repeated Column B names, and that no two Column A cells coincide
   once upper-cased;
5. with ``--check-source``, the yp-data copy re-fetched by its recorded retriever and
   hashed (network).

Run from the repo root (a worktree: ``PYTHONPATH=<worktree>``)::

    python experiments/database/scripts/yoshikawa2009_release_checks.py [--check-source]

Writes ``experiments/database/results/yoshikawa2009_release_checks.json``.
"""

from __future__ import annotations

import argparse
import json
import os
import os.path as osp
import re
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import olefile  # type: ignore[import-untyped]
import openpyxl
import xlrd
from dotenv import load_dotenv

from torchcell.datasets.scerevisiae import yoshikawa2009 as y
from torchcell.literature.provenance import check_source

RESULTS = "experiments/database/results"
XLS_KINDS = {
    xlrd.XL_CELL_EMPTY: "empty",
    xlrd.XL_CELL_TEXT: "text",
    xlrd.XL_CELL_NUMBER: "number",
    xlrd.XL_CELL_BLANK: "empty",
}


def _kind(value: Any) -> str:
    """The kind of an openpyxl cell value, on the xlrd vocabulary."""
    if value is None:
        return "empty"
    if isinstance(value, str):
        return "text"
    if isinstance(value, int | float) and not isinstance(value, bool):
        return "number"
    return type(value).__name__


def compare_copies(xls_path: Path, xlsx_path: Path) -> dict[str, Any]:
    """Every cell of the two copies, compared by value and kind; plus the frame."""
    book = xlrd.open_workbook(str(xls_path), formatting_info=True)
    sheet = book.sheet_by_index(0)
    workbook = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    ws = workbook.worksheets[0]
    n_rows = max(sheet.nrows, ws.max_row)
    n_cols = max(sheet.ncols, ws.max_column)
    differences: list[dict[str, Any]] = []
    kinds: Counter[str] = Counter()
    formulas = 0
    for r in range(n_rows):
        for c in range(n_cols):
            inside = r < sheet.nrows and c < sheet.ncols
            xls_kind = XLS_KINDS[sheet.cell(r, c).ctype] if inside else "absent"
            xls_value = (
                sheet.cell(r, c).value if xls_kind != "empty" and inside else None
            )
            xlsx_value = ws.cell(row=r + 1, column=c + 1).value
            if isinstance(xlsx_value, str) and xlsx_value.startswith("="):
                formulas += 1
            xlsx_kind = _kind(xlsx_value)
            kinds[f"{xls_kind}/{xlsx_kind}"] += 1
            same = xls_kind == xlsx_kind and (
                float(xls_value) == float(xlsx_value)
                if xls_value is not None and isinstance(xlsx_value, int | float)
                else xls_value == xlsx_value
            )
            if not same:
                differences.append(
                    {
                        "row": r + 1,
                        "column": c + 1,
                        "xls": [xls_kind, xls_value],
                        "xlsx": [xlsx_kind, xlsx_value],
                    }
                )
    data_cells = (sheet.nrows - y.FIRST_DATA_ROW) * sheet.ncols
    return {
        "sheets": {"xls": book.sheet_names(), "xlsx": workbook.sheetnames},
        "dimensions": {
            "xls": [sheet.nrows, sheet.ncols],
            "xlsx": [ws.max_row, ws.max_column],
        },
        "merged_ranges": {
            "xls": sorted(
                f"r{r0 + 1}c{c0 + 1}:r{r1}c{c1}"
                for r0, r1, c0, c1 in sheet.merged_cells
            ),
            "xlsx": sorted(str(rng) for rng in ws.merged_cells.ranges),
        },
        "cells_compared": n_rows * n_cols,
        "data_cells": data_cells,
        "cell_kinds_xls_over_xlsx": dict(sorted(kinds.items())),
        "xlsx_formula_cells": formulas,
        "n_differences": len(differences),
        "differences": differences,
    }


def metadata(xls_path: Path, xlsx_path: Path) -> dict[str, Any]:
    """Who wrote each copy and when, as each file records it."""
    ole = olefile.OleFileIO(str(xls_path))
    meta = ole.get_metadata()
    ole.close()

    def text(value: Any) -> Any:
        return value.decode("latin-1") if isinstance(value, bytes) else value

    props = openpyxl.load_workbook(str(xlsx_path)).properties
    return {
        "xls": {
            "ole_author": text(meta.author),
            "ole_last_saved_by": text(meta.last_saved_by),
            "ole_created": str(meta.create_time),
            "ole_last_saved": str(meta.last_saved_time),
            "ole_last_printed": str(meta.last_printed),
            "ole_application": text(meta.creating_application),
            "biff_writeaccess_user": xlrd.open_workbook(str(xls_path)).user_name,
        },
        "xlsx": {
            "creator": props.creator,
            "last_modified_by": props.lastModifiedBy,
            "created": str(props.created),
            "modified": str(props.modified),
            "last_printed": str(props.lastPrinted),
        },
    }


def five_percent(paper: str, xls_path: Path, xlsx_path: Path) -> dict[str, Any]:
    """The 5% ethanol columns in each copy, and every 5 percent the paper writes."""
    xls_header = xlrd.open_workbook(str(xls_path)).sheet_by_index(0).row_values(0)
    ws = openpyxl.load_workbook(str(xlsx_path)).worksheets[0]
    xlsx_header = [cell.value for cell in ws[1]]
    pattern = re.compile(r"(?<![\d.] )(?<![\d.])\b5 ?\\?%")
    hits = [
        paper[max(0, m.start() - 40) : m.end() + 40] for m in pattern.finditer(paper)
    ]
    return {
        "xls_condition_header": [v for v in xls_header if v],
        "xlsx_condition_header": [v for v in xlsx_header if v],
        "paper_five_percent_spans": hits,
        "paper_mentions_5_percent_ethanol": "5% ethanol" in paper
        or "5 \\%$ ethanol" in paper,
        "paper_eight_percent_spans": len(re.findall(r"8 \\%", paper)),
    }


def calls(rows: list[y.ReleaseRow]) -> dict[str, Any]:
    """Call counts per condition, the both-cultures rule, and the 5% implied interval."""
    tables = {spec.label: spec.table1 for spec in y.CONDITIONS}
    result: dict[str, Any] = {
        "counts": {
            label: dict(
                zip(
                    ("sensitive", "no_change", "tolerant"),
                    y.call_counts(rows, label),
                    strict=True,
                )
            )
            for label in y.RELEASE_CONDITIONS
        },
        "table2": y.TABLE1_COUNTS.value,
    }
    rule: dict[str, Any] = {}
    for label, table in tables.items():
        mismatched: Counter[str] = Counter()
        examples: list[list[Any]] = []
        for row in rows:
            cells = row.conditions[label]
            a, b = cells.rates
            low, high = table.sensitive_below, table.tolerant_above
            predicted = 0 if a < low and b < low else 2 if a > high and b > high else 1
            if predicted != cells.call:
                mismatched[f"released {cells.call} / rule {predicted}"] += 1
                if len(examples) < 5:
                    examples.append([row.orf, a, b, cells.call, predicted])
        rule[label] = {
            "thresholds": [table.sensitive_below, table.tolerant_above],
            "n_mismatched": sum(mismatched.values()),
            "mismatched": dict(mismatched),
            "examples": examples,
        }
    result["both_cultures_rule_with_table1_thresholds"] = rule
    label = "5% ethanol"
    upper = [max(row.conditions[label].rates) for row in rows]
    lower = [min(row.conditions[label].rates) for row in rows]
    sensitive = [row.conditions[label].call == 0 for row in rows]
    tolerant = [row.conditions[label].call == 2 for row in rows]
    result["five_percent_implied"] = {
        "sensitive_threshold_above": max(
            u for u, s in zip(upper, sensitive, strict=True) if s
        ),
        "sensitive_threshold_at_most": min(
            u for u, s in zip(upper, sensitive, strict=True) if not s
        ),
        "tolerant_threshold_below": min(
            v for v, t in zip(lower, tolerant, strict=True) if t
        ),
        "tolerant_threshold_at_least": max(
            v for v, t in zip(lower, tolerant, strict=True) if not t
        ),
    }
    return result


def names(rows: list[y.ReleaseRow]) -> dict[str, Any]:
    """The two name columns as stored: every strain under its own released names."""
    stored = Counter(y.stored_orf(r.orf) for r in rows)
    duplicated = sorted(n for n, k in Counter(r.name for r in rows).items() if k > 1)
    return {
        "stored_strains": len(rows),
        "column_a_lower_case": sorted(r.orf for r in rows if r.orf != r.orf.upper()),
        "column_a_coinciding_upper_cased": sorted(
            o for o, k in stored.items() if k > 1
        ),
        "column_b_equals_column_a": sum(r.name.upper() == r.orf.upper() for r in rows),
        "column_b_repeated_names": duplicated,
    }


def main() -> None:
    """Measure the release; write the JSON."""
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--check-source", action="store_true")
    args = parser.parse_args()
    load_dotenv()
    data_root = os.environ["DATA_ROOT"]
    raw = y.raw_mirror_dir(data_root)
    xls_path = raw / y.TABLE_S1.raw_rel
    xlsx_path = raw / y.TABLE_S1_YEASTPHENOME.raw_rel
    paper = (
        Path(data_root) / "torchcell-library" / y.CITATION_KEY / y.PAPER_MD
    ).read_text(encoding="utf-8")
    rows = y.read_release(xls_path)
    results: dict[str, Any] = {
        "files": {
            "xls": {"path": str(xls_path), "sha256": y.TABLE_S1.sha256},
            "xlsx": {"path": str(xlsx_path), "sha256": y.TABLE_S1_YEASTPHENOME.sha256},
        },
        "comparison": compare_copies(xls_path, xlsx_path),
        "metadata": metadata(xls_path, xlsx_path),
        "five_percent_ethanol": five_percent(paper, xls_path, xlsx_path),
        "release_rows": len(rows),
        "calls": calls(rows),
        "names": names(rows),
    }
    if args.check_source:
        record = y.yeastphenome_retrieval()
        results["yeastphenome_source_check"] = check_source(
            record, now=datetime.now(UTC).isoformat()
        ).model_dump()
    os.makedirs(RESULTS, exist_ok=True)
    out = osp.join(RESULTS, "yoshikawa2009_release_checks.json")
    with open(out, "w") as handle:
        json.dump(results, handle, indent=2, default=str)
        handle.write("\n")
    print(
        json.dumps(
            {k: v for k, v in results.items() if k != "comparison"},
            indent=2,
            default=str,
        )
    )
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
