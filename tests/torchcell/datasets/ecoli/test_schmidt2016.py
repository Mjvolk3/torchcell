# tests/torchcell/datasets/ecoli/test_schmidt2016.py
# [[tests.torchcell.datasets.ecoli.test_schmidt2016]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_schmidt2016.py
"""Schmidt 2016 proteome loader (``torchcell.datasets.ecoli.schmidt2016``).

The synthetic tests write a small workbook with the release's exact three-sheet, three
block layout (six protein rows, one per rule the loader applies; the 22 released
conditions) and drive the readers, the protein-row rules, the two released-statistics
identities, the environment and phenotype builders, the raw-mirror deposit and a full
``process()`` build into ``tmp_path``. The genome, the locus-tag reconciliation and the
assembly pin are in-test objects, so nothing reads ``$DATA_ROOT``.

The ``@pytest.mark.data`` tests read the real mirrors and the dev-tree LMDB and pin the
numbers the dendron note states: 2,359 released rows, 2,329 protein keys, 14 records,
the 5 non-host rows, the 7 twice-filed symbols, the 11 unresolved symbols, and that
Table S6's first two CV headers are swapped.
"""

from __future__ import annotations

import json
import math
import os
import os.path as osp
import statistics
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

import torchcell.datasets.ecoli.schmidt2016 as sm
from torchcell.data import file_sha256
from torchcell.datamodels.schema import (
    AssemblyReferenceGenome,
    BacterialAssemblySet,
    BacterialGeneNamespace,
    EnvironmentPhysicalPerturbation,
    SmallMoleculePerturbation,
)
from torchcell.datasets.bacteria_common import LocusTagReconciliation
from torchcell.sequence.genome.base import GeneNameStatus

ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_BW25113_ASM75055v1"
NAMESPACE: BacterialGeneNamespace = "ecoli_k12_bw25113_locus_tag"
#: ``(gene, uniprot, dataset, organism)``, one row per rule the loader applies.
ROWS: tuple[tuple[str, str, int, str], ...] = (
    ("aceA", "P00001", 2, "Escherichia coli (strain K12)"),
    ("aceB", "P00002", 1, "Escherichia coli (strain K12)"),
    ("dupA", "P00003", 2, "Escherichia coli (strain K12)"),
    ("dupA", "P00004", 2, "Escherichia coli (strain K12)"),
    ("foreign", "P00005", 2, "Bacillus subtilis"),
    ("ghost", "P00006", 2, "Escherichia coli (strain K12)"),
)
KEPT_GENES = ("aceA", "aceB")
UNRESOLVED = ("ghost",)
#: The four conditions dataset 1 did not cover, as the release's NA cells say.
DATASET1_MISSING = ("Glycerol + AA", "Xylose", "Mannose", "Fructose")
IDENTIFIER_HEADERS = (
    sm.COL_UNIPROT,
    sm.COL_DESCRIPTION,
    sm.COL_GENE,
    sm.COL_PEPTIDES,
    "Confidence.score",
    sm.COL_MW,
    sm.COL_DATASET,
)
TAIL_HEADERS = (sm.COL_GENE, sm.COL_BNUMBER)


def _copies(gene: str, column: str) -> float | str:
    """A deterministic copies/cell for one (row, condition), ``NA`` where dataset 1 is."""
    row = next(r for r in ROWS if r[0] == gene)
    if row[2] == 1 and column in DATASET1_MISSING:
        return sm.NOT_AVAILABLE
    index = [r[0] for r in ROWS].index(gene)
    offset = [c.s6_column for c in sm.CONDITIONS].index(column)
    return 100.0 * (index + 1) + offset


def _replicates(gene: str, column: str) -> tuple[float, float, float]:
    """Three normalized abundances whose median and CV the workbook then states.

    The spread widens with the condition, so no two conditions share a coefficient of
    variation and the Table S6 header-order check has something to distinguish.
    """
    index = [r[0] for r in ROWS].index(gene)
    offset = [c.s6_column for c in sm.CONDITIONS].index(column)
    base = 1.0e6 * (index + 1) + 1.0e3 * offset
    spread = 0.01 * (offset + 1)
    return (base * (1.0 - spread), base, base * (1.0 + 1.5 * spread))


def write_workbook(path: Path) -> Path:
    """Write a workbook with the release's sheet, block and header layout."""
    book = openpyxl.Workbook()
    _write_s6(book.active)
    _write_s8(book.create_sheet(sm.SHEET_S8))
    _write_s25(book.create_sheet(sm.SHEET_S25))
    book.save(path)
    return path


def _write_s6(sheet: Any) -> None:
    sheet.title = sm.SHEET_S6
    columns = [c.s6_column for c in sm.CONDITIONS]
    width = len(columns)
    sheet.append([sm._Q_TABLE_S6])
    labels: list[Any] = [None] * len(IDENTIFIER_HEADERS)
    for label in (sm.BLOCK_COPIES, sm.BLOCK_MASS, sm.BLOCK_CV):
        labels.append(label)
        labels.extend([None] * (width - 1))
    sheet.append(labels)
    sheet.append([*IDENTIFIER_HEADERS, *columns, *columns, *columns, *TAIL_HEADERS])
    for gene, uniprot, dataset, organism in ROWS:
        copies = [_copies(gene, column) for column in columns]
        mass = [v if isinstance(v, str) else v * 1.0e-3 for v in copies]
        # The released CV block holds LB first and Glucose second, under headers that
        # read Glucose then LB: the swap this loader asserts and does not consume.
        order = ["LB", "Glucose", *[c for c in columns if c not in ("LB", "Glucose")]]
        cv = [
            sm.NOT_AVAILABLE
            if dataset == 1
            else 100.0
            * statistics.stdev(_replicates(gene, column))
            / statistics.fmean(_replicates(gene, column))
            for column in order
        ]
        sheet.append(
            [
                uniprot,
                f"Test protein {gene} OS={organism} GN={gene} PE=1 SV=1",
                gene,
                7,
                123.4,
                50000.0,
                dataset,
                *copies,
                *mass,
                *cv,
                gene,
                None
                if gene == "ghost"
                else f"b{ROWS.index((gene, uniprot, dataset, organism)):04d}",
            ]
        )


def _write_s8(sheet: Any) -> None:
    loaded = [c for c in sm.CONDITIONS if c.drop is None]
    header: list[str] = [sm.COL_UNIPROT]
    for spec in loaded:
        header.extend(sm._s8_column(name) for name in spec.s25_files)
        header.append(f"medianNormInt_{spec.s8_suffix}")
        header.append(f"cv_{spec.s8_suffix}")
    sheet.append([sm._Q_TABLE_S8])
    sheet.append([None] * len(header))
    sheet.append(header)
    dataset2 = [r for r in ROWS if r[2] == 2]
    # The last dataset-2 row is written twice, so the duplicate-accession rule has an
    # accession to drop.
    for gene, uniprot, _, _ in [*dataset2, dataset2[-1]]:
        row: list[Any] = [uniprot]
        for spec in loaded:
            values = _replicates(gene, spec.s6_column)
            row.extend(values)
            row.append(statistics.median(values))
            row.append(100.0 * statistics.stdev(values) / statistics.fmean(values))
        sheet.append(row)


def _write_s25(sheet: Any) -> None:
    sheet.append([sm._Q_TABLE_S25])
    sheet.append([None])
    sheet.append(
        ["File Name", "Growth Condition (in biological triplicates)", "Strain"]
    )
    for spec in sm.CONDITIONS:
        for name in spec.s25_files:
            sheet.append([name, spec.s25_label, sm.SCHMIDT_REFERENCE_STRAIN])
    for name in sm.GLUCOSE_REPRODUCIBILITY_FILES:
        sheet.append([name, "glucose", sm.SCHMIDT_REFERENCE_STRAIN])


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    """The synthetic workbook, written once per test."""
    return write_workbook(tmp_path / sm.SI2)


# --------------------------------------------------------------------------- #
# The declared condition table
# --------------------------------------------------------------------------- #
def test_condition_table_shape() -> None:
    assert len(sm.CONDITIONS) == 22
    loaded = [c for c in sm.CONDITIONS if c.drop is None]
    assert len(loaded) == 15
    assert len([c for c in loaded if not c.is_reference]) == sm.EXPECTED_RECORDS
    assert [c.s6_column for c in sm.CONDITIONS if c.is_reference] == [
        sm.REFERENCE_CONDITION
    ]
    dropped = {c.s6_column: c.drop.rule for c in sm.CONDITIONS if c.drop is not None}
    assert dropped == {
        "Glycerol + AA": "medium_has_no_media_library_entry",
        "Chemostat µ=0.5": "culture_not_batch",
        "Chemostat µ=0.35": "culture_not_batch",
        "Chemostat µ=0.20": "culture_not_batch",
        "Chemostat µ=0.12": "culture_not_batch",
        "Stationary phase 1 day": "growth_phase_not_representable",
        "Stationary phase 3 days": "growth_phase_not_representable",
    }
    assert len({c.s25_files for c in sm.CONDITIONS}) == len(sm.CONDITIONS)


def test_chemostat_conditions_carry_the_lower_glucose_the_methods_states() -> None:
    chemostats = [c for c in sm.CONDITIONS if c.s6_column.startswith("Chemostat")]
    assert {c.carbon_g_per_l for c in chemostats} == {1.0}
    glucose = sm.CONDITIONS_BY_COLUMN["Glucose"]
    assert glucose.carbon_g_per_l == 5.0


# --------------------------------------------------------------------------- #
# Reading the workbook
# --------------------------------------------------------------------------- #
def test_read_table_s6_parses_the_block_layout(workbook: Path) -> None:
    rows, cv_pairs = sm.read_table_s6(str(workbook))
    assert [row.gene for row in rows] == [r[0] for r in ROWS]
    assert [row.uniprot for row in rows] == [r[1] for r in ROWS]
    assert [row.organism for row in rows] == [r[3] for r in ROWS]
    assert [row.is_host for row in rows] == [
        r[3].startswith("Escherichia coli") for r in ROWS
    ]
    assert [row.n_replicates for row in rows] == [3 if r[2] == 2 else 1 for r in ROWS]
    assert rows[0].row_number == 4
    assert rows[0].copies["Glucose"] == pytest.approx(_copies("aceA", "Glucose"))
    assert rows[1].copies["Xylose"] is None
    assert rows[5].bnumber is None
    # One CV pair per dataset-2 row.
    assert len(cv_pairs) == len([r for r in ROWS if r[2] == 2])


def test_read_table_s6_refuses_a_renamed_condition_column(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    sheet = book[sm.SHEET_S6]
    sheet.cell(row=sm.HEADER_ROW, column=len(IDENTIFIER_HEADERS) + 1, value="Glukose")
    book.save(path)
    with pytest.raises(RuntimeError, match="copies/cell headers"):
        sm.read_table_s6(str(path))


def test_read_table_s6_refuses_a_missing_block_label(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    book[sm.SHEET_S6].cell(
        row=sm.HEADER_ROW - 1, column=len(IDENTIFIER_HEADERS) + 1, value="other"
    )
    book.save(path)
    with pytest.raises(RuntimeError, match="no block labelled"):
        sm.read_table_s6(str(path))


def test_read_table_s6_refuses_a_description_with_no_organism(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    book[sm.SHEET_S6].cell(row=4, column=2, value="Test protein with no organism")
    book.save(path)
    with pytest.raises(RuntimeError, match="names no OS= organism"):
        sm.read_table_s6(str(path))


def test_number_refuses_a_cell_that_is_neither_numeric_nor_na() -> None:
    assert sm._number(1.5) == 1.5
    assert sm._number(sm.NOT_AVAILABLE) is None
    with pytest.raises(RuntimeError, match="neither a number"):
        sm._number("below LOQ")


def test_header_index_refuses_a_repeated_name() -> None:
    with pytest.raises(RuntimeError, match="appears twice"):
        sm._header_index(["a", "b", "a"])
    assert sm._header_index(["a", "b", "a"], range(2)) == {"a": 0, "b": 1}


def test_read_table_s8_drops_a_repeated_accession(workbook: Path) -> None:
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, repeated = sm.read_table_s8(str(workbook), loaded)
    assert repeated == ("P00006",)
    assert "P00006" not in cells
    cell = cells["P00001"]["Glucose"]
    assert cell.replicates == pytest.approx(_replicates("aceA", "Glucose"))
    assert cell.median == pytest.approx(cell.median_released)
    assert cell.cv == pytest.approx(cell.cv_released)


def test_sample_map_check_and_its_refusals(workbook: Path) -> None:
    samples = sm.read_table_s25(str(workbook))
    sm.check_sample_map(samples)
    with pytest.raises(RuntimeError, match="lists no sample"):
        sm.check_sample_map([s for s in samples if s[0] != "A14-07036"])
    relabelled = [
        (name, "something else" if name == "A14-07036" else condition, strain)
        for name, condition, strain in samples
    ]
    with pytest.raises(RuntimeError, match="the module declares"):
        sm.check_sample_map(relabelled)
    with pytest.raises(RuntimeError, match="glucose replicate"):
        sm.check_sample_map(
            [s for s in samples if s[0] not in sm.GLUCOSE_REPRODUCIBILITY_FILES]
        )


# --------------------------------------------------------------------------- #
# The released-statistics identities
# --------------------------------------------------------------------------- #
def test_released_statistics_check_passes_and_catches_a_drifted_cv(
    workbook: Path,
) -> None:
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, _ = sm.read_table_s8(str(workbook), loaded)
    summary = sm.check_released_statistics(cells)
    # Four of the five dataset-2 accessions survive: P00006 is filed twice in Table S8.
    assert summary["n_cells"] == (len([r for r in ROWS if r[2] == 2]) - 1) * len(loaded)
    assert summary["worst_median_rtol"] == pytest.approx(0.0)
    assert summary["worst_cv_rtol"] < sm.IDENTITY_RTOL

    broken = cells["P00001"]["Glucose"]
    cells["P00001"]["Glucose"] = broken.model_copy(
        update={"cv_released": broken.cv_released * 2}
    )
    with pytest.raises(RuntimeError, match="released cv"):
        sm.check_released_statistics(cells)


def test_released_statistics_check_catches_a_drifted_median(workbook: Path) -> None:
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, _ = sm.read_table_s8(str(workbook), loaded)
    cell = cells["P00001"]["LB"]
    cells["P00001"]["LB"] = cell.model_copy(
        update={"median_released": cell.median_released * 1.5}
    )
    with pytest.raises(RuntimeError, match="released median"):
        sm.check_released_statistics(cells)


def test_cv_header_swap_is_asserted_and_a_corrected_export_would_fail(
    workbook: Path,
) -> None:
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, _ = sm.read_table_s8(str(workbook), loaded)
    _, cv_pairs = sm.read_table_s6(str(workbook))
    summary = sm.check_cv_header_swap(cv_pairs, cells)
    # Every dataset-2 row but the twice-filed P00006 can be checked.
    assert summary == {
        "n_rows_checked": len([r for r in ROWS if r[2] == 2]) - 1,
        "swapped": True,
    }
    unswapped = [(key, second, first) for key, first, second in cv_pairs]
    with pytest.raises(RuntimeError, match="LB then Glucose"):
        sm.check_cv_header_swap(unswapped, cells)
    with pytest.raises(RuntimeError, match="no row could check"):
        sm.check_cv_header_swap([], cells)


def test_standard_error_scales_the_replicate_spread_into_copies_per_cell() -> None:
    replicates = (90.0, 100.0, 110.0)
    cell = sm.ReplicateCell(
        replicates=replicates, median_released=100.0, cv_released=10.0
    )
    assert cell.median == 100.0
    assert cell.cv == pytest.approx(10.0)
    # copies = 2 * the median, so the SD doubles and the SE divides by sqrt(3).
    expected = 2.0 * statistics.stdev(replicates) / math.sqrt(3)
    assert cell.standard_error(200.0) == pytest.approx(expected)


# --------------------------------------------------------------------------- #
# The protein-row rules
# --------------------------------------------------------------------------- #
class _Locus:
    def __init__(self, symbol: str | None) -> None:
        self.symbol = symbol


class _Annotation:
    def __init__(self, loci: dict[str, _Locus]) -> None:
        self.loci = loci


class FakeGenome:
    """Carries the kept synthetic gene names as its loci, and nothing else."""

    ASSEMBLY_SET = ASSEMBLY_SET

    def __init__(self) -> None:
        """One locus per kept gene name."""
        self.genbank = _Annotation({gene: _Locus(gene) for gene in KEPT_GENES})


def _reconciliation(names: pd.Series, label: str) -> LocusTagReconciliation:
    unresolved = [n for n in names if n in UNRESOLVED]
    return LocusTagReconciliation(
        label=label,
        assembly_set=ASSEMBLY_SET,
        gene_namespace=NAMESPACE,
        unique_names=len(set(names)),
        status_histogram={
            GeneNameStatus.RENAMED: len(set(names)) - len(unresolved),
            GeneNameStatus.CURRENT: 0,
            GeneNameStatus.NON_GENE_FEATURE: 0,
            GeneNameStatus.RETIRED: len(unresolved),
            GeneNameStatus.AMBIGUOUS: 0,
        },
        layer_histogram={"gene symbol": len(set(names)) - len(unresolved)},
        remapped=len(set(names)) - len(unresolved),
        kept_on_collision=(),
        retired_kept=tuple(unresolved),
        ambiguous_kept={},
        case_insensitive=(),
        outside_namespace=tuple(unresolved),
    )


def _identity_reconciliation(
    genome: Any, names: pd.Series, *, label: str
) -> tuple[pd.Series, LocusTagReconciliation]:
    """Store each gene name as itself, and mark ``UNRESOLVED`` outside the namespace."""
    return names, _reconciliation(names, label)


def _reference_genome() -> AssemblyReferenceGenome:
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain=sm.SCHMIDT_REFERENCE_STRAIN,
        ploidy="haploid",
        assembly_set=ASSEMBLY_SET,
        assembly_accession="GCA_000750555.1",
    )


def test_select_proteins_applies_the_three_rules_in_order(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(sm, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(sm, "MIN_RESOLVED_FRACTION", 0.5)
    rows, _ = sm.read_table_s6(str(workbook))
    selection = sm.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    assert [row.gene for row in selection.kept] == list(KEPT_GENES)
    assert selection.locus_tag == {"P00001": "aceA", "P00002": "aceB"}
    assert selection.dropped_rows == len(ROWS) - len(KEPT_GENES)
    by_rule = {rule.rule: rule for rule in selection.rules}
    assert by_rule["source_organism_is_not_escherichia_coli"].n_items == 1
    assert by_rule["source_organism_is_not_escherichia_coli"].items == [
        "foreign (P00005, Bacillus subtilis)"
    ]
    assert by_rule["gene_symbol_filed_on_two_released_rows"].n_items == 2
    assert by_rule["gene_symbol_resolves_to_no_locus_of_the_pinned_assembly"].items == [
        "ghost (P00006)"
    ]


def test_select_proteins_refuses_a_release_below_the_resolution_floor(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(sm, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(sm, "MIN_RESOLVED_FRACTION", 0.99)
    rows, _ = sm.read_table_s6(str(workbook))
    from torchcell.datasets.bacteria_common import LocusTagResolutionError

    with pytest.raises(LocusTagResolutionError):
        sm.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]


# --------------------------------------------------------------------------- #
# Environment and phenotype
# --------------------------------------------------------------------------- #
def test_minimal_medium_condition_carries_its_carbon_source() -> None:
    environment = sm.build_environment(sm.CONDITIONS_BY_COLUMN["Acetate"])
    assert environment.media.base_medium == "M9"
    assert environment.temperature is not None
    assert environment.temperature.value == 37.0
    assert environment.aerobicity == "aerobic"
    (perturbation,) = environment.perturbations
    assert isinstance(perturbation, EnvironmentPhysicalPerturbation)
    assert perturbation.factor == "carbon_source"
    assert perturbation.magnitude is not None
    assert perturbation.magnitude.value == 3.5
    assert perturbation.magnitude.unit == "g/L"
    assert perturbation.agent is not None
    assert perturbation.agent.name == "sodium acetate"


def test_lb_condition_carries_no_perturbation() -> None:
    environment = sm.build_environment(sm.CONDITIONS_BY_COLUMN["LB"])
    assert environment.perturbations == []
    assert environment.media.is_synthetic is False


def test_stress_conditions_use_the_axis_each_stress_belongs_on() -> None:
    osmotic = sm.build_environment(sm.CONDITIONS_BY_COLUMN["Osmotic-stress glucose"])
    kinds = [p.perturbation_type for p in osmotic.perturbations]
    assert kinds == ["environment_physical", "small_molecule"]
    salt = osmotic.perturbations[1]
    assert isinstance(salt, SmallMoleculePerturbation)
    assert salt.compound.name == "sodium chloride"
    assert (salt.concentration.value, salt.concentration.unit) == (50.0, "mM")

    hot = sm.build_environment(sm.CONDITIONS_BY_COLUMN["42°C glucose"])
    assert hot.temperature is not None
    assert hot.temperature.value == 42.0
    assert [p.perturbation_type for p in hot.perturbations] == ["environment_physical"]

    acid = sm.build_environment(sm.CONDITIONS_BY_COLUMN["pH6 glucose"])
    ph = acid.perturbations[1]
    assert isinstance(ph, EnvironmentPhysicalPerturbation)
    assert ph.factor == "pH"
    assert ph.magnitude is not None
    assert (ph.magnitude.value, ph.magnitude.unit) == (6.0, "pH")
    assert ph.agent is not None
    assert ph.agent.name == "hydrochloric acid"


def test_every_loaded_environment_is_distinct_and_a_repeat_is_refused() -> None:
    loaded = [c for c in sm.CONDITIONS if c.drop is None]
    environments = {c.s6_column: sm.build_environment(c) for c in loaded}
    sm.check_environments_distinct(environments)
    assert len(environments) == 15
    environments["copy"] = sm.build_environment(sm.CONDITIONS_BY_COLUMN["Glucose"])
    with pytest.raises(RuntimeError, match="serialize to one environment"):
        sm.check_environments_distinct(environments)


def test_dropped_chemostat_and_stationary_arms_would_not_be_distinguishable() -> None:
    """The two structural drop rules, stated as the collision they avoid."""
    chemostats = {
        c.s6_column: sm.build_environment(c)
        for c in sm.CONDITIONS
        if c.drop is sm.DROP_CULTURE_NOT_BATCH
    }
    assert len({e.model_dump_json() for e in chemostats.values()}) == 1
    stationary_and_glucose = {
        c.s6_column: sm.build_environment(c)
        for c in sm.CONDITIONS
        if c.drop is sm.DROP_GROWTH_PHASE or c.s6_column == "Glucose"
    }
    assert len(stationary_and_glucose) == 3
    assert len({e.model_dump_json() for e in stationary_and_glucose.values()}) == 1


def test_phenotype_skips_na_and_matches_se_to_the_replicate_count(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(sm, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(sm, "MIN_RESOLVED_FRACTION", 0.5)
    rows, _ = sm.read_table_s6(str(workbook))
    selection = sm.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, _ = sm.read_table_s8(str(workbook), loaded)

    glucose = sm.build_phenotype(selection.kept, selection.locus_tag, cells, "Glucose")
    assert set(glucose.protein_abundance) == set(KEPT_GENES)
    assert glucose.protein_abundance["aceA"] == pytest.approx(
        _copies("aceA", "Glucose")
    )
    assert glucose.n_replicates == {"aceA": 3, "aceB": 1}
    assert glucose.protein_abundance_se is not None
    assert not math.isnan(glucose.protein_abundance_se["aceA"])
    assert math.isnan(glucose.protein_abundance_se["aceB"])
    assert glucose.measurement_type == sm.MEASUREMENT_TYPE

    xylose = sm.build_phenotype(selection.kept, selection.locus_tag, cells, "Xylose")
    assert set(xylose.protein_abundance) == {"aceA"}


def test_restrict_keeps_only_the_requested_keys_and_refuses_a_missing_one(
    workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(sm, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(sm, "MIN_RESOLVED_FRACTION", 0.5)
    rows, _ = sm.read_table_s6(str(workbook))
    selection = sm.select_proteins(rows, FakeGenome(), label="test")  # type: ignore[arg-type]
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, _ = sm.read_table_s8(str(workbook), loaded)
    glucose = sm.build_phenotype(selection.kept, selection.locus_tag, cells, "Glucose")
    narrowed = sm.restrict(glucose, ["aceA"])
    assert set(narrowed.protein_abundance) == {"aceA"}
    assert set(narrowed.n_replicates) == {"aceA"}
    with pytest.raises(RuntimeError, match="quantifies none of"):
        sm.restrict(glucose, ["aceA", "absent"])


def test_publication_carries_the_resolved_identifiers() -> None:
    publication = sm.publication()
    assert publication.pubmed_id == "26641532"
    assert publication.doi == "10.1038/nbt.3418"
    assert publication.doi_url == "https://doi.org/10.1038/nbt.3418"


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def test_deposit_is_idempotent_and_refuses_other_bytes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = write_workbook(tmp_path / "source.xlsx")
    sha = file_sha256(source)
    monkeypatch.setattr(sm, "SI2_SHA256", sha)
    data_root = str(tmp_path / "root")
    root = sm.deposit_raw_mirror(source=source, data_root=data_root)
    manifest = sm.load_manifest(data_root)
    assert sm.manifest_sha256(manifest, sm.SI2_MIRROR_RELPATH) == sha
    (record,) = manifest.files
    assert record.retrieval is not None
    assert record.retrieval.method == "pmc_cloud"
    assert record.retrieval.params == {"key": sm.SI2_PMC_KEY}
    assert record.role == "raw_data"
    assert manifest.doi == sm.PAPER_DOI
    assert any("PXD000498" in s for s in manifest.si_data_sources)
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        sm.manifest_sha256(manifest, "data/other.xlsx")

    sm.deposit_raw_mirror(source=source, data_root=data_root)
    assert file_sha256(root / sm.SI2_MIRROR_RELPATH) == sha

    other = tmp_path / "other.xlsx"
    other.write_bytes(b"not the release")
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        sm.deposit_raw_mirror(source=other, data_root=data_root)
    (root / sm.SI2_MIRROR_RELPATH).write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="different sha256"):
        sm.deposit_raw_mirror(source=source, data_root=data_root)


def test_download_links_the_mirror_and_refuses_a_drifted_pin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = write_workbook(tmp_path / "source.xlsx")
    sha = file_sha256(source)
    monkeypatch.setattr(sm, "SI2_SHA256", sha)
    data_root = tmp_path / "root"
    sm.deposit_raw_mirror(source=source, data_root=str(data_root))
    monkeypatch.setenv("DATA_ROOT", str(data_root))

    root = tmp_path / "build"
    (root / "raw").mkdir(parents=True)
    dataset = sm.ProteomeSchmidt2016Dataset.__new__(sm.ProteomeSchmidt2016Dataset)
    monkeypatch.setattr(
        type(dataset), "raw_dir", property(lambda self: str(root / "raw"))
    )
    dataset.download()
    assert file_sha256(root / "raw" / sm.SI2) == sha

    monkeypatch.setattr(sm, "SI2_SHA256", "0" * 64)
    from torchcell.data import ManifestPinMismatchError

    with pytest.raises(ManifestPinMismatchError):
        dataset.download()

    monkeypatch.setattr(sm, "SI2_SHA256", sha)
    (data_root / sm.RAW_DIR_REL / sm.SI2_MIRROR_RELPATH).unlink()
    with pytest.raises(RuntimeError, match="missing from mirror"):
        dataset.download()


def test_mirror_paths_follow_the_citation_key(tmp_path: Path) -> None:
    assert sm.raw_mirror_dir(str(tmp_path)).name == sm.CITATION_KEY
    assert sm.raw_mirror_dir(str(tmp_path)).parent.name == "torchcell-raw"
    assert sm.library_dir(str(tmp_path)).parent.name == "torchcell-library"


# --------------------------------------------------------------------------- #
# End-to-end build on the synthetic workbook
# --------------------------------------------------------------------------- #
def test_process_builds_one_record_per_loaded_condition(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "proteome_schmidt2016"
    raw = root / "raw"
    raw.mkdir(parents=True)
    path = write_workbook(raw / sm.SI2)
    monkeypatch.setattr(sm, "DATA_SHA256", {sm.SI2: file_sha256(path)})
    monkeypatch.setattr(sm, "reconcile_locus_tags", _identity_reconciliation)
    monkeypatch.setattr(sm, "MIN_RESOLVED_FRACTION", 0.5)
    monkeypatch.setattr(sm, "assembly_reference", lambda strain: _reference_genome())

    dataset = sm.ProteomeSchmidt2016Dataset(
        root=str(root),
        ecoli_genome=FakeGenome(),  # type: ignore[arg-type]
    )
    assert len(dataset) == sm.EXPECTED_RECORDS
    assert dataset.gene_set == set(KEPT_GENES)

    items = [dataset.transform_item(dataset[i]) for i in range(len(dataset))]
    assert {len(i["experiment"].genotype.perturbations) for i in items} == {0}
    assert {i["experiment"].phenotype.measurement_type for i in items} == {
        sm.MEASUREMENT_TYPE
    }
    assert {i["publication"].doi for i in items} == {sm.PAPER_DOI}
    assert {i["reference"].genome_reference.assembly_set for i in items} == {
        ASSEMBLY_SET
    }
    # The reference is the glucose arm, key-matched to each record.
    for item in items:
        experiment, reference = item["experiment"], item["reference"]
        assert set(reference.phenotype_reference.protein_abundance) == set(
            experiment.phenotype.protein_abundance
        )
        for key, value in reference.phenotype_reference.protein_abundance.items():
            assert value == pytest.approx(_copies(key, "Glucose"))
    # The three conditions dataset 1 did not cover carry one key fewer.
    by_keys = {len(i["experiment"].phenotype.protein_abundance) for i in items}
    assert by_keys == {1, 2}
    assert (
        len([i for i in items if len(i["experiment"].phenotype.protein_abundance) == 1])
        == 3
    )
    # Every environment is distinct, which is the record identity here.
    assert len({i["experiment"].environment.model_dump_json() for i in items}) == len(
        items
    )
    dataset.close_lmdb()

    preprocess = Path(dataset.preprocess_dir)
    ledger = json.loads((preprocess / "dropped_records.json").read_text())
    assert ledger["source_conditions"] == 22
    assert ledger["kept_records"] == sm.EXPECTED_RECORDS
    assert ledger["dropped_records"] == 7
    assert ledger["source_protein_rows"] == len(ROWS)
    assert ledger["kept_protein_keys"] == len(KEPT_GENES)
    assert ledger["dropped_protein_rows"] == len(ROWS) - len(KEPT_GENES)
    checks = json.loads((preprocess / "released_statistics_check.json").read_text())
    assert checks["table_s6_cv_header_swap"]["swapped"] is True
    assert checks["median_and_cv_from_replicates"]["worst_median_rtol"] == 0.0
    sourced = json.loads((preprocess / "sourced_values.json").read_text())
    assert sourced["reference_strain"]["value"] == "BW25113"
    identifiers = pd.read_csv(preprocess / "protein_identifiers.csv")
    assert list(identifiers["released_gene"]) == list(KEPT_GENES)
    conditions = pd.read_csv(preprocess / "conditions.csv")
    assert len(conditions) == sm.EXPECTED_RECORDS
    assert sm.REFERENCE_CONDITION not in set(conditions["s6_column"])
    assert (preprocess / "build_manifest.json").exists()

    report = sm.verify_build(
        str(root),
        genome=FakeGenome(),  # type: ignore[arg-type]
        expected_count=sm.EXPECTED_RECORDS,
    )
    assert [(r.level, r.name) for r in report.results if not r.passed] == []
    assert (preprocess / "verification_report.json").exists()


def test_drop_log_check_refuses_an_unaccounted_drop() -> None:
    reconciliation = _reconciliation(pd.Series(["aceA"]), "test")
    log_model = sm.DropLog(
        dataset="test",
        source_conditions=22,
        reference_conditions=["Glucose"],
        candidate_records=14,
        kept_records=14,
        dropped_records=7,
        source_protein_rows=10,
        kept_protein_keys=9,
        dropped_protein_rows=1,
        rules=[
            sm.DropRule(
                rule="r", scope="protein_row", description="d", n_items=1, items=["x"]
            )
        ],
        reconciliation=reconciliation,
        notes=[],
    )
    with pytest.raises(RuntimeError, match="dropped conditions"):
        log_model.check()
    fixed = log_model.model_copy(update={"source_conditions": 15})
    fixed.check()
    mismatched = fixed.model_copy(update={"dropped_protein_rows": 2})
    with pytest.raises(RuntimeError, match="dropped rows stated"):
        mismatched.check()
    short = fixed.model_copy(update={"kept_protein_keys": 8})
    with pytest.raises(RuntimeError, match="dropped rows !="):
        short.check()


# --------------------------------------------------------------------------- #
# Real data
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    from dotenv import load_dotenv

    load_dotenv()
    return os.environ["DATA_ROOT"]


def _real_workbook() -> str:
    path = sm.raw_mirror_dir(_data_root()) / sm.SI2_MIRROR_RELPATH
    if not path.exists():
        pytest.skip(f"raw mirror not deposited: {path}")
    return str(path)


@pytest.mark.data
def test_raw_mirror_matches_the_pin() -> None:
    path = _real_workbook()
    assert file_sha256(path) == sm.SI2_SHA256
    manifest = sm.load_manifest(_data_root())
    assert sm.manifest_sha256(manifest, sm.SI2_MIRROR_RELPATH) == sm.SI2_SHA256


@pytest.mark.data
def test_every_quote_is_verbatim_in_its_pinned_mirror() -> None:
    paper = sm.library_dir(_data_root()) / sm.PAPER_MD
    if not paper.exists():
        pytest.skip(f"paper mirror not present: {paper}")
    assert file_sha256(paper) == sm.PAPER_MD_SHA256
    text = paper.read_text(encoding="utf-8")
    missing = [name for name, quote in sm.PAPER_QUOTES.items() if quote not in text]
    assert missing == []

    book = openpyxl.load_workbook(_real_workbook(), read_only=True, data_only=True)
    try:
        renderings = set()
        for row in book["CONTENT_AND_ABBREVIATIONS"].iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                renderings.add(" | ".join(cells))
    finally:
        book.close()
    assert [name for name, q in sm.SI2_QUOTES.items() if q not in renderings] == []


@pytest.mark.data
def test_real_release_shape_and_the_three_protein_rules() -> None:
    rows, cv_pairs = sm.read_table_s6(_real_workbook())
    assert len(rows) == 2359
    assert len(cv_pairs) == 2058
    assert len([r for r in rows if r.dataset == 2]) == 2058
    assert len([r for r in rows if r.dataset == 1]) == 301
    non_host = [r for r in rows if not r.is_host]
    assert sorted(r.gene for r in non_host) == [
        "MettuDRAFT_4149",
        "addB",
        "cas1",
        "cgtA",
        "ygbT",
    ]
    counts: dict[str, int] = {}
    for row in rows:
        if row.is_host:
            counts[row.gene] = counts.get(row.gene, 0) + 1
    assert sorted(g for g, n in counts.items() if n > 1) == [
        "bioD",
        "clpB",
        "glsA",
        "mrcB",
        "nrdA",
        "rmlA",
        "rpmE",
    ]
    assert len([r for r in rows if r.bnumber is None]) == 69


@pytest.mark.data
def test_real_sample_map_and_released_statistics() -> None:
    path = _real_workbook()
    samples = sm.read_table_s25(path)
    assert len(samples) == 81
    sm.check_sample_map(samples)
    loaded = [c.s6_column for c in sm.CONDITIONS if c.drop is None]
    cells, repeated = sm.read_table_s8(path, loaded)
    assert repeated == ("P63284",)
    summary = sm.check_released_statistics(cells)
    assert summary["worst_median_rtol"] == 0.0
    assert summary["worst_cv_rtol"] < sm.IDENTITY_RTOL
    _, cv_pairs = sm.read_table_s6(path)
    assert sm.check_cv_header_swap(cv_pairs, cells)["swapped"] is True


@pytest.mark.data
def test_built_lmdb_numbers() -> None:
    root = osp.join(_data_root(), "data/torchcell/proteome_schmidt2016")
    if not osp.exists(osp.join(root, "processed", "lmdb")):
        pytest.skip(f"dev store not built: {root}")
    from torchcell.verification.runners import load_records

    records = load_records(root)
    assert len(records) == sm.EXPECTED_RECORDS
    keys = {len(r["experiment"]["phenotype"]["protein_abundance"]) for r in records}
    assert keys == {sm.EXPECTED_PROTEIN_KEYS, sm.EXPECTED_PROTEIN_KEYS_DATASET2_ONLY}
    ledger = json.loads(Path(root, "preprocess", "dropped_records.json").read_text())
    assert ledger["source_protein_rows"] == 2359
    assert ledger["kept_protein_keys"] == sm.EXPECTED_PROTEIN_KEYS
    assert ledger["dropped_protein_rows"] == 30
    assert ledger["reconciliation"]["unique_names"] == 2340
    assert ledger["reconciliation"]["remapped"] == sm.EXPECTED_PROTEIN_KEYS
