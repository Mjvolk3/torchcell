# tests/torchcell/datasets/ecoli/test_potel2018.py
# [[tests.torchcell.datasets.ecoli.test_potel2018]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_potel2018.py
"""Tests for the Potel 2018 phosphosite-intensity loader (#921).

Everything outside the last block runs with NO network and NO real ``$DATA_ROOT``:
Supplementary Tables 2 and 5 are synthesized to the pinned workbooks' exact headers, the
raw mirror and its manifest are written under ``tmp_path`` with the module's pins
monkeypatched to the synthetic digests and counts, and the W3110 reference is pinned
directly. The four synthetic sites are one quantified everywhere, a two-gene site under
one accession (EF-Tu) that no sample quantifies, a three-accession site with a mixed
row, and a histidine site with a diagnostic peak. The last block, skipped without the
mirror, pins the numbers measured on the REAL bytes.
"""

from __future__ import annotations

import hashlib
import math
import os
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.ecoli.potel2018 as po
from torchcell.datamodels.schema import (
    ConcentrationUnit,
    CultureGrowthPhase,
    PhosphoResidue,
    PhosphositeIntensityExperiment,
    PhosphositeIntensityExperimentReference,
    PhosphositeQuantification,
    ReferenceGenome,
)
from torchcell.literature.manifest import ArtifactRecord, Manifest
from torchcell.verification.levels import l3_convention
from torchcell.verification.sourced import ProvenanceGapReason

COLUMNS = [spec.column for spec in po.SAMPLES]

#: (acc, gene names, positions, loc prob, PEP, score, delta, diag, residue, window,
#:  total, twelve intensities in release order).
ROWS: tuple[tuple[Any, ...], ...] = (
    (
        "P00350",
        "gnd",
        86,
        0.999615,
        4.27602e-05,
        66.401,
        58.487,
        None,
        "S",
        "LLMVKAGAGTDAAIDSLKPYLDKGDIIIDGG",
        78,
        (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12),
    ),
    (
        "P0CE47",
        "tufA;tufB",
        "313;313",
        1,
        0,
        100,
        50,
        None,
        "S",
        "TIKPHTKFESEVYILSKDEGGRHTPFFKGYR",
        0,
        (0,) * 12,
    ),
    (
        "P09373;P42632;P68066",
        "pflB;tdcE;grcA",
        "734;738;101",
        0.9,
        0.01,
        60.5,
        20.25,
        None,
        "S",
        "AIEHPEKYPNLTIRVSGYAVRFNALTREQQQ;AMENPEKYPQLTIRVSGYAVRFNSLTKEQQQ;"
        "AVKHPEKYPQLTIRVSGYAVRFNSLTPEQQR",
        100,
        (0, 5, 0, 0, 0, 0, 16, 0, 32, 0, 0, 64),
    ),
    (
        "P0A6F5",
        "groL",
        119,
        0.5,
        0.002,
        80,
        40,
        "+",
        "H",
        "LVVAATDGPMPQTREHILLGRQVGVPYIIVF",
        40,
        (0, 0, 0, 0, 0, 0, 0, 0, 0, 8, 0, 0),
    ),
)
QUANTIFIED = {
    column: sum(1 for row in ROWS if row[11][index] > 0)
    for index, column in enumerate(COLUMNS)
}


def write_table_2(path: Path, rows: tuple[tuple[Any, ...], ...]) -> Path:
    """A workbook with Table 2's sheet and header, and these rows."""
    book = openpyxl.Workbook()
    book.active.title = "Description"
    sheet = book.create_sheet(po.TABLE_2_SHEET)
    sheet.append(list(po.TABLE_2_HEADER))
    for row in rows:
        acc, genes, pos, loc, pep, score, delta, diag, aa, window, total, cells = row
        sheet.append(
            [acc, "a protein", genes, pos, loc, pep, score, delta, diag, 1, aa]
            + [window, "PEPTIDE", 2, 0.1, total, *cells]
        )
    book.save(path)
    return path


def write_table_5(path: Path, rows: tuple[tuple[Any, ...], ...]) -> Path:
    """Table 5's glycerol log2 columns for every site that has one position."""
    book = openpyxl.Workbook()
    book.active.title = "Description"
    sheet = book.create_sheet(po.TABLE_5_SHEET)
    glycerol = COLUMNS[6:]
    sheet.append(
        ["Significance", *(f"log2 ({c})" for c in glycerol)]
        + ["Acc. No.", "Amino acid", "Position", "Protein names", "Gene names"]
        + ["Localization prob"]
    )
    for row in rows:
        acc, genes, pos, loc, *_ = row
        if isinstance(pos, str):
            continue
        cells = row[11][6:]
        log2 = [math.log2(c) if c > 0 else 17.0 for c in cells]
        sheet.append([None, *log2, acc, row[8], pos, "a protein", genes, loc])
    book.save(path)
    return path


def _w3110() -> ReferenceGenome:
    return ReferenceGenome(
        species="Escherichia coli",
        strain="W3110",
        assembly_set="ecoli_K12_W3110_ASM1024v1",
        assembly_accession="GCF_000010245.2",
    )


@pytest.fixture
def mirrored(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A tmp ``DATA_ROOT`` holding the synthetic mirror, with the synthetic pins."""
    data_root = tmp_path / "data_root"
    mirror = data_root / po.RAW_DIR_REL
    (mirror / "data").mkdir(parents=True)
    (mirror / "si").mkdir(parents=True)
    table = write_table_2(mirror / po.TABLE_2_REL, ROWS)
    table_5 = write_table_5(mirror / po.TABLE_5_REL, ROWS)
    digest = hashlib.sha256(table.read_bytes()).hexdigest()
    manifest = Manifest(
        citation_key=po.CITATION_KEY,
        files=[
            ArtifactRecord(
                path=po.TABLE_2_REL,
                role="raw_data",
                bytes=table.stat().st_size,
                sha256=digest,
                source="synthetic",
            )
        ],
    )
    (mirror / "manifest.json").write_text(manifest.model_dump_json())
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.delenv("TC_DATA_URL", raising=False)
    monkeypatch.setattr(po, "TABLE_2_SHA256", digest)
    monkeypatch.setattr(
        po, "TABLE_5_SHA256", hashlib.sha256(table_5.read_bytes()).hexdigest()
    )
    monkeypatch.setattr(po, "SOURCE_SITES", 4)
    monkeypatch.setattr(po, "SOURCE_RESIDUES", {"S": 3, "H": 1})
    monkeypatch.setattr(po, "MULTI_GENE_SITES", 2)
    monkeypatch.setattr(po, "MULTI_ACCESSION_SITES", 1)
    monkeypatch.setattr(po, "ZERO_CELLS", 48 - sum(QUANTIFIED.values()))
    monkeypatch.setattr(po, "SITES_ZERO_IN_ALL_SAMPLES", 1)
    monkeypatch.setattr(po, "QUANTIFIED_PER_SAMPLE", QUANTIFIED)
    # gnd's six glycerol cells and groL's one; Table 5 lists no multi-position site.
    monkeypatch.setattr(po, "TABLE_5_AGREEING_CELLS", 6 + 1)
    monkeypatch.setattr(po, "strain_reference", lambda data_root=None: _w3110())
    return data_root


def _by_label(dataset: po.PhosphositeIntensityPotel2018Dataset) -> dict[str, Any]:
    records = [dataset[index] for index in range(len(dataset))]
    return {r["experiment"]["phenotype"]["sample_label"]: r for r in records}


# --------------------------------------------------------------------- reading
def test_table_2_rows_become_sites_exactly_as_released(tmp_path: Path) -> None:
    rows = po.read_table_2(write_table_2(tmp_path / "t2.xlsx", ROWS))
    assert [row.site.key() for row in rows] == [
        "P00350|S|86",
        "P0CE47|S|313;313",
        "P09373;P42632;P68066|S|734;738;101",
        "P0A6F5|H|119",
    ]
    tuf, pfl, his = rows[1].site, rows[2].site, rows[3].site
    assert (tuf.protein_ids, tuf.gene_names, tuf.positions) == (
        ["P0CE47"],
        ["tufA", "tufB"],
        [313, 313],
    )
    assert pfl.gene_names == ["pflB", "tdcE", "grcA"]
    assert len(pfl.sequence_windows) == 3
    assert (his.residue, his.diagnostic_peak) == (PhosphoResidue.histidine, True)
    assert rows[0].site.diagnostic_peak is False
    assert (tuf.localization_probability, tuf.posterior_error_probability) == (1.0, 0.0)
    assert rows[2].intensities[COLUMNS[1]] == 5
    assert rows[2].total_intensity == 100


def test_a_header_that_is_not_the_released_one_is_refused(tmp_path: Path) -> None:
    path = write_table_2(tmp_path / "t2.xlsx", ROWS)
    book = openpyxl.load_workbook(path)
    book[po.TABLE_2_SHEET]["A1"] = "Accession"
    book.save(path)
    with pytest.raises(RuntimeError, match="Table 2 header"):
        po.read_table_2(path)


@pytest.mark.parametrize("cell", ["x", "++", 1])
def test_a_diagnostic_peak_other_than_plus_or_empty_is_refused(cell: Any) -> None:
    with pytest.raises(ValueError, match="neither '\\+' nor empty"):
        po._diagnostic_peak(cell)


@pytest.mark.parametrize("cell", [-1, 1.5, "7", None, True])
def test_an_intensity_cell_that_is_not_a_count_is_refused(cell: Any) -> None:
    with pytest.raises(ValueError, match="is not a non-negative integer"):
        po._count(cell, "Intensity glucose_exp1")


# ----------------------------------------------------------- samples and medium
def test_the_twelve_samples_are_the_release_columns_with_their_legend() -> None:
    assert COLUMNS == list(po.TABLE_2_HEADER[16:])
    spec = po.SAMPLES[9]
    assert (spec.column, spec.carbon, spec.phase, spec.replicate) == (
        "Intensity glycerol_stat1",
        "glycerol",
        CultureGrowthPhase.stationary,
        1,
    )
    assert spec.legend.quote == (
        "26. Intensity glycerol_stat1, summed up eXtracted Ion Current (XIC) of all "
        "isotopic clusters associated with the identified phosphopeptdie sequence in "
        "the 1st replicate of glycerol-fed stationary E. coli cells."
    )
    assert [s.harvest_optical_density for s in po.SAMPLES[::3]] == [0.6, 1.2, 0.5, 1.0]
    assert po.reference_sample(spec).column == "Intensity glycerol_exp1"
    assert po.reference_sample(po.SAMPLES[0]) is po.SAMPLES[0]


def test_the_medium_is_chapter_v_m9_with_the_carbon_source_as_stated() -> None:
    glucose, glycerol = po.medium("glucose"), po.medium("glycerol")
    amounts = [
        (c.compound.name, c.concentration.value, c.concentration.unit)
        for c in glycerol.components
        if c.concentration is not None
    ]
    assert amounts == [
        ("disodium hydrogen phosphate", 6.0, ConcentrationUnit.g_per_l),
        ("potassium dihydrogen phosphate", 3.0, ConcentrationUnit.g_per_l),
        ("sodium chloride", 0.5, ConcentrationUnit.g_per_l),
        ("ammonium chloride", 1.0, ConcentrationUnit.g_per_l),
        ("magnesium sulfate", 1.0, ConcentrationUnit.millimolar),
        ("calcium chloride", 0.1, ConcentrationUnit.millimolar),
        ("glycerol", 0.5, ConcentrationUnit.percent),
    ]
    carbon = glucose.components[-1]
    assert carbon.compound.name == "D-glucose"
    assert carbon.concentration is not None
    assert carbon.concentration.unit is ConcentrationUnit.percent_w_v
    assert (glucose.base_medium, glucose.state, glucose.is_synthetic) == (
        "M9",
        "liquid",
        True,
    )
    env = po.environment("glucose")
    assert env.temperature is not None
    assert (env.temperature.value, env.aerobicity, env.perturbations) == (
        37.0,
        "aerobic",
        [],
    )


def test_a_released_zero_is_not_quantified_and_the_unit_is_a_typed_gap(
    tmp_path: Path,
) -> None:
    rows = po.read_table_2(write_table_2(tmp_path / "t2.xlsx", ROWS))
    phenotype = po.phenotype(po.SAMPLES[1], rows)
    assert phenotype.phosphosite_intensity == {
        "P00350|S|86": 2.0,
        "P09373;P42632;P68066|S|734;738;101": 5.0,
    }
    assert phenotype.quantification["P0CE47|S|313;313"] is (
        PhosphositeQuantification.not_quantified
    )
    assert len(phenotype.sites) == 4
    assert (phenotype.replicate, phenotype.harvest_growth_phase) == (
        2,
        CultureGrowthPhase.exponential,
    )
    assert phenotype.replicate_unit is None
    [gap] = phenotype.provenance_gaps
    assert (gap.field, gap.reason) == (
        "replicate_unit",
        ProvenanceGapReason.deferred_pending_source_review,
    )


def test_the_accounting_counts_the_release_and_refuses_an_off_pin_count(
    tmp_path: Path, mirrored: Path
) -> None:
    rows = po.read_table_2(write_table_2(tmp_path / "t2.xlsx", ROWS))
    accounting = po.build_accounting("potel", rows)
    accounting.check()
    assert accounting.total_intensity_not_stored.model_dump() == {
        "equal_within_1e_6_relative": 1,
        "total_greater": 1,
        "total_smaller": 1,
        "total_zero": 1,
    }
    assert accounting.residues == {"S": 3, "H": 1}
    assert accounting.diagnostic_peak_by_residue == {"H": 1}
    off = accounting.model_copy(update={"zero_cells": accounting.zero_cells - 1})
    with pytest.raises(RuntimeError, match="build counts"):
        off.check()


# ------------------------------------------------------------- build + verify
def test_the_loader_stores_every_site_in_every_sample(
    tmp_path: Path, mirrored: Path
) -> None:
    root = tmp_path / "dataset"
    dataset = po.PhosphositeIntensityPotel2018Dataset(root=str(root))
    assert len(dataset) == po.EXPECTED_RECORDS
    assert dataset.gene_set == set()
    by_label = _by_label(dataset)
    assert sorted(by_label) == sorted(COLUMNS)
    record = by_label["Intensity glycerol_exp3"]
    phenotype = record["experiment"]["phenotype"]
    assert phenotype["phosphosite_intensity"] == {
        "P00350|S|86": 9.0,
        "P09373;P42632;P68066|S|734;738;101": 32.0,
    }
    assert phenotype["harvest_optical_density"] == 0.5
    assert record["experiment"]["genotype"]["perturbations"] == []
    assert record["reference"]["phenotype_reference"]["sample_label"] == (
        "Intensity glycerol_exp3"
    )
    stat = by_label["Intensity glucose_stat2"]["reference"]["phenotype_reference"]
    assert stat["sample_label"] == "Intensity glucose_exp2"
    assert (root / "raw" / po.TABLE_2_FILENAME).is_symlink()
    accounting = po.BuildAccounting.model_validate_json(
        (root / "preprocess" / "build_accounting.json").read_text()
    )
    assert accounting.quantified_per_sample == QUANTIFIED


def test_the_verifier_passes_the_synthetic_build_and_fails_a_changed_cell(
    tmp_path: Path, mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "dataset"
    po.PhosphositeIntensityPotel2018Dataset(root=str(root))
    monkeypatch.setattr(
        po, "_l3_quotes_are_verbatim", lambda data_root: l3_convention("quotes", True)
    )
    report = po.verify_build(str(root), data_root=str(mirrored))
    assert report.passed, report.summary()
    names = [r.name for r in report.results]
    assert names[-6:] == [
        "quantified_sites_per_sample_are_the_pinned_ones",
        "every_stored_cell_and_site_is_the_released_one",
        "sample_fields_follow_the_column_legend",
        "assembly_pin_is_w3110",
        "quotes",
        "table_5_log2_values_are_log2_of_the_stored_intensities",
    ]
    l4 = report.results[-1]
    assert (l4.details["agree"], l4.details["imputed"], l4.details["joined_sites"]) == (
        7,
        5,
        2,
    )
    assert (root / "preprocess" / "verification_report.json").is_file()

    path = mirrored / po.RAW_DIR_REL / po.TABLE_2_REL
    changed = list(ROWS)
    changed[0] = (*ROWS[0][:11], (2, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12))
    write_table_2(path, tuple(changed))
    monkeypatch.setattr(
        po, "TABLE_2_SHA256", hashlib.sha256(path.read_bytes()).hexdigest()
    )
    failed = {r.name for r in po.verify_build(str(root)).results if not r.passed}
    assert failed == {"every_stored_cell_and_site_is_the_released_one"}


def test_the_loader_declares_its_schema_pair_its_host_and_its_raw_file() -> None:
    cls = po.PhosphositeIntensityPotel2018Dataset
    dataset = cls.__new__(cls)
    assert dataset.experiment_class is PhosphositeIntensityExperiment
    assert dataset.reference_class is (PhosphositeIntensityExperimentReference)
    assert dataset.raw_file_names == [po.TABLE_2_FILENAME]
    assert (cls.REFERENCE_STRAIN, cls.has_gene_perturbations) == ("W3110", False)
    assert str(cls.organism) == "ecoli"
    assert po.publication().doi == "10.1038/nmeth.4580"


# --------------------------------------------------------- the real released bytes
def _real_mirror() -> Path | None:
    data_root = os.environ.get("DATA_ROOT")
    if not data_root:
        return None
    path = Path(data_root) / po.RAW_DIR_REL / po.TABLE_2_REL
    return path if path.is_file() else None


requires_mirror = pytest.mark.skipif(
    _real_mirror() is None, reason="the Potel 2018 raw mirror is not on this machine"
)


@requires_mirror
def test_the_released_table_has_the_pinned_counts() -> None:
    path = _real_mirror()
    assert path is not None
    assert hashlib.sha256(path.read_bytes()).hexdigest() == po.TABLE_2_SHA256
    rows = po.read_table_2(path)
    accounting = po.build_accounting("potel", rows)
    accounting.check()
    assert accounting.total_intensity_not_stored.model_dump() == {
        "equal_within_1e_6_relative": 278,
        "total_greater": 1130,
        "total_smaller": 477,
        "total_zero": 244,
    }
    assert accounting.diagnostic_peak_by_residue == {"H": 31, "Y": 82}


@requires_mirror
def test_every_released_quote_is_verbatim_in_the_mirror() -> None:
    result = po._l3_quotes_are_verbatim(None)
    assert result.passed, result.message
