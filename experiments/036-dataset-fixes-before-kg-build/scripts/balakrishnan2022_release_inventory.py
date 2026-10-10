# experiments/036-dataset-fixes-before-kg-build/scripts/balakrishnan2022_release_inventory.py
# [[experiments.036-dataset-fixes-before-kg-build.scripts.balakrishnan2022_release_inventory]]
# https://github.com/Mjvolk3/torchcell/tree/main/experiments/036-dataset-fixes-before-kg-build/scripts/balakrishnan2022_release_inventory

r"""What Balakrishnan 2022 released, what is scriptable, and what a loader could serve.

Balakrishnan, Mori, Segota, Zhang, Aebersold, Ludwig and Hwa 2022, *Science* 378
eabk2066 (doi 10.1126/science.abk2066, PMC9804519, an NIH author manuscript). Row 57 of
the bacterial candidate schedule. Four measurements, each regenerable:

1. **Access.** Every route to the supplement: the PMC Article Datasets bucket listing,
   the PMC ``bin/`` supplement URL, the Europe PMC supplementary-files service, the
   Science supplement, and the GEO series supplement directory. Which answer with bytes.
2. **Table S3** (the one scriptable data file, deposited on GEO as
   ``GSE205717_Processed_data_Table_S3.xlsx``, sha256-pinned below): per sheet the rows,
   the sample columns, the header faults, the column sums, the identifiers and their
   resolution against the MG1655 genome.
3. **Sister series.** The steady-state sample sheet against Mori 2021's Dataset EV3
   (sha256-pinned in the Mori raw mirror): same strains, media and limitation series,
   and whether any sample is the same culture.
4. **Cross-layer r.** Balakrishnan's reference-condition mRNA number fractions against
   Mori 2021's reference-condition protein mass fractions on shared b-numbers, beside
   the within-layer replicate r of each. This says the transcriptome arm is a different
   layer, not a re-served proteome; it is not a duplication test of Table S4, which is
   not retrievable.

Step 1 and the Table S3 read reach the network (the bytes are verified against the pin).
Run from the repo root::

    python experiments/036-dataset-fixes-before-kg-build/scripts/balakrishnan2022_release_inventory.py
"""

import hashlib
import json
import math
import os
import os.path as osp
import re
from collections import Counter
from io import BytesIO
from typing import Any

import httpx
import openpyxl
from dotenv import load_dotenv
from scipy.stats import pearsonr, spearmanr

load_dotenv()

import torchcell.datasets.ecoli.mori2021 as mori  # noqa: E402
from torchcell.datasets.bacteria_common import bacterial_genome  # noqa: E402
from torchcell.literature.retrieve import direct_url  # noqa: E402

RESULTS = osp.join(
    os.environ["EXPERIMENT_ROOT"], "036-dataset-fixes-before-kg-build", "results"
)
OUT = osp.join(RESULTS, "balakrishnan2022_release_inventory.json")

DOI = "10.1126/science.abk2066"
PMCID = "PMC9804519"
NIHMS = "NIHMS1856867"
GEO_SERIES = "GSE205717"
TABLE_S3_URL = (
    "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE205nnn/GSE205717/suppl/"
    "GSE205717_Processed_data_Table_S3.xlsx"
)
TABLE_S3_SHA256 = "9d7df03bbbcdf1f92e93027014e746da5bcf42bf4d6b4be9061dd1803f71cc20"

#: The supplementary objects the author manuscript's XML names (``<media xlink:href>``).
SUPPLEMENT_FILES = (
    f"{NIHMS}-supplement-Supplementary_Information.pdf",
    f"{NIHMS}-supplement-Table_S3.xlsx",
    f"{NIHMS}-supplement-Table_S4.xlsx",
    f"{NIHMS}-supplement-Table_S5.xlsx",
    f"{NIHMS}-supplement-Table_S6.xlsx",
    f"{NIHMS}-supplement-Table_S7.xlsx",
)
ACCESS_PROBES: dict[str, str] = {
    "pmc_cloud_bucket_listing": (
        f"https://pmc-oa-opendata.s3.amazonaws.com/?list-type=2&prefix={PMCID}."
    ),
    "pmc_bin_table_s6": (
        f"https://pmc.ncbi.nlm.nih.gov/articles/instance/{PMCID[3:]}/bin/"
        f"{NIHMS}-supplement-Table_S6.xlsx"
    ),
    "europepmc_supplementary_files": (
        f"https://www.ebi.ac.uk/europepmc/webservices/rest/{PMCID}/supplementaryFiles"
    ),
    "europepmc_bin_table_s6": (
        f"https://europepmc.org/articles/{PMCID}/bin/{NIHMS}-supplement-Table_S6.xlsx"
    ),
    "science_org_supplement_page": f"https://www.science.org/doi/suppl/{DOI}",
    "geo_series_suppl_directory": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE205nnn/GSE205717/suppl/"
    ),
}

SHEET_SS_DESC = "1 - RNAseq-ss (description)"
SHEET_SS = "2 - RNAseq-ss (fractions)"
SHEET_DEG_DESC = "3 - RNAseq-deg (description)"
SHEET_DEG = "4 - RNAseq-deg (fractions)"
IDENTITY = ("gene", "locus", "gene length (nt)")
B_NUMBER = re.compile(r"^b\d{4}$")

#: Reference-condition pairs for the cross-layer r: (Balakrishnan column, Mori column).
#: Both are NCM3722 wild type, 0.2% glucose, the same base medium and nitrogen source.
CROSS_LAYER_PAIRS = {"M9_reference": ("c5", "C2"), "MOPS_reference": ("r0", "A2")}
#: Within-layer replicate pairs, the yardstick the cross-layer r is read against.
REPLICATE_PAIRS = {
    "balakrishnan_mrna_M9": ("c5", "c0_1"),
    "balakrishnan_mrna_MOPS": ("r0", "r0_1"),
    "mori_protein_MOPS": ("A2", "H1"),
}


def probe(client: httpx.Client, url: str) -> dict[str, Any]:
    """Status, content type and the first bytes' kind of one GET."""
    response = client.get(url)
    body = response.content
    head = body[:400].decode("utf-8", errors="replace")
    kind = (
        "xlsx_or_zip"
        if body[:2] == b"PK"
        else "pdf"
        if body[:4] == b"%PDF"
        else "proof_of_work_page"
        if "POW_CHALLENGE" in body.decode("utf-8", "replace")
        else "not_open_access_error"
        if "not open access" in head
        else "html"
        if "<html" in head.lower() or "<!doctype" in head.lower()
        else "xml"
        if head.lstrip().startswith("<?xml")
        else "other"
    )
    keys = re.findall(r"<Key>([^<]+)</Key>", body.decode("utf-8", "replace"))
    hrefs = re.findall(r'href="(GSE[^"]+)"', body.decode("utf-8", "replace"))
    return {
        "url": url,
        "status": response.status_code,
        "content_type": response.headers.get("content-type"),
        "bytes": len(body),
        "kind": kind,
        "bucket_keys": keys,
        "geo_files": hrefs,
    }


def sheet_rows(book: openpyxl.Workbook, name: str) -> list[tuple[Any, ...]]:
    """Every row of one sheet, trailing all-blank rows dropped."""
    rows = list(book[name].iter_rows(values_only=True))
    while rows and all(cell is None for cell in rows[-1]):
        rows.pop()
    return rows


def describe_samples(rows: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    """The description sheet as one dict per sample, blank group cells carried down."""
    header = [str(c).strip() for c in rows[0] if c is not None]
    out: list[dict[str, Any]] = []
    group = None
    for row in rows[1:]:
        record = dict(zip(header, row[: len(header)], strict=True))
        if record["Group"] is not None:
            group = record["Group"]
        record["Group (carried)"] = group
        out.append(record)
    return out


def fraction_block(rows: list[tuple[Any, ...]]) -> dict[str, Any]:
    """Shape, header faults, sums and identifiers of one fractions sheet."""
    header = [str(c).strip() for c in rows[0]]
    if tuple(header[:3]) != IDENTITY:
        raise RuntimeError(f"identity block is {header[:3]}, expected {IDENTITY}")
    samples = header[3:]
    counts = Counter(samples)
    duplicated = {name: n for name, n in counts.items() if n > 1}
    body = rows[1:]
    duplicate_identity: dict[str, Any] = {}
    for name in duplicated:
        idx = [i for i, h in enumerate(header) if h == name]
        same = sum(1 for r in body if len({r[i] for i in idx}) == 1)
        duplicate_identity[name] = {
            "column_indices_1_based": [i + 1 for i in idx],
            "rows_identical_across_the_copies": same,
            "rows": len(body),
        }
    sums: dict[str, float] = {}
    zeros: dict[str, int] = {}
    non_numeric = 0
    for j, name in enumerate(header[3:], start=3):
        values = [r[j] for r in body]
        numeric = [float(v) for v in values if isinstance(v, (int, float))]
        non_numeric += len(values) - len(numeric)
        key = f"{name}@{j + 1}"
        sums[key] = sum(numeric)
        zeros[key] = sum(1 for v in numeric if v == 0.0)
    loci = [r[1] for r in body]
    genes = [r[0] for r in body]
    return {
        "gene_rows": len(body),
        "sample_columns": len(samples),
        "distinct_sample_headers": len(counts),
        "duplicated_headers": duplicated,
        "duplicate_identity": duplicate_identity,
        "non_numeric_cells": non_numeric,
        "column_sum_min": min(sums.values()),
        "column_sum_max": max(sums.values()),
        "zero_cells_per_column_min": min(zeros.values()),
        "zero_cells_per_column_max": max(zeros.values()),
        "distinct_gene_names": len(set(genes)),
        "blank_loci": sum(1 for x in loci if x is None),
        "loci_not_a_b_number": dict(
            Counter(
                str(x) for x in loci if not (isinstance(x, str) and B_NUMBER.match(x))
            )
        ),
        "b_number_loci": sum(
            1 for x in loci if isinstance(x, str) and B_NUMBER.match(x)
        ),
        "distinct_loci": len({x for x in loci if x is not None}),
        "headers": samples,
    }


def column(rows: list[tuple[Any, ...]], name: str) -> dict[str, float]:
    """``{b-number: value}`` of the FIRST column headed ``name`` (rows with a locus)."""
    header = [str(c).strip() for c in rows[0]]
    j = header.index(name)
    return {
        str(r[1]).strip(): float(r[j])
        for r in rows[1:]
        if r[1] is not None and isinstance(r[j], (int, float))
    }


def mori_column(name: str) -> dict[str, float]:
    """``{b-number: protein mass fraction}`` of one Mori 2021 EV9 column."""
    path = osp.join(
        os.environ["DATA_ROOT"], mori.RAW_DIR_REL, mori.MIRROR_RELPATH[mori.EV9]
    )
    return {
        row.gene_locus: row.mass_fraction[name]
        for row in mori.read_mass_fractions(path, [name])
        if row.gene_locus is not None
    }


def log_r(a: dict[str, float], b: dict[str, float]) -> dict[str, Any]:
    """Pearson on log10 and Spearman over keys positive in both."""
    keys = sorted(k for k in set(a) & set(b) if a[k] > 0 and b[k] > 0)
    x = [math.log10(a[k]) for k in keys]
    y = [math.log10(b[k]) for k in keys]
    return {
        "shared_positive_keys": len(keys),
        "pearson_log10": round(float(pearsonr(x, y)[0]), 4),
        "spearman": round(float(spearmanr(x, y)[0]), 4),
    }


def mori_series() -> list[dict[str, Any]]:
    """Mori 2021 Dataset EV3's C-, A- and R-limitation rows (NCM3722 lineage)."""
    path = osp.join(
        os.environ["DATA_ROOT"], mori.RAW_DIR_REL, mori.MIRROR_RELPATH[mori.EV3]
    )
    rows = mori._rows(path, mori.SHEET_EV3)
    header = [str(c).strip() for c in rows[0] if c is not None]
    out = []
    for row in rows[1:]:
        record = dict(zip(header, row[: len(header)], strict=True))
        if record["Strain"] != "EQ353":
            out.append(record)
    return out


def main() -> None:
    """Measure, print a summary and write the JSON beside the other 036 results."""
    with httpx.Client(
        follow_redirects=True, timeout=120.0, headers={"User-Agent": "Mozilla/5.0"}
    ) as client:
        access = {name: probe(client, url) for name, url in ACCESS_PROBES.items()}

    raw = direct_url(TABLE_S3_URL)
    sha = hashlib.sha256(raw).hexdigest()
    if sha != TABLE_S3_SHA256:
        raise RuntimeError(f"Table S3 sha256 is {sha}, pinned {TABLE_S3_SHA256}")
    book = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
    sheets = {name: sheet_rows(book, name) for name in book.sheetnames}
    book.close()

    ss_samples = describe_samples(sheets[SHEET_SS_DESC])
    deg_samples = describe_samples(sheets[SHEET_DEG_DESC])
    ss = fraction_block(sheets[SHEET_SS])
    deg = fraction_block(sheets[SHEET_DEG])
    described = [s["Sample ID"] for s in ss_samples]
    ss_headers = set(ss["headers"])

    genome = bacterial_genome("ecoli", "MG1655")
    loci = {str(r[1]).strip() for r in sheets[SHEET_SS][1:] if r[1] is not None}
    resolved = sorted(loci & set(genome.gene_set))

    mori_rows = mori_series()
    mori_strains = sorted({r["Strain"] for r in mori_rows})
    bala_strains = sorted({s["Strain"] for s in ss_samples})
    mori_keys = {
        (r["Strain"], r["Growth medium"], r["Supplement"], r["Growth rate (1/h)"])
        for r in mori_rows
    }
    same_culture_candidates = [
        s["Sample ID"]
        for s in ss_samples
        if (s["Strain"], s["Growth medium"], s["Supplement"], s["Growth rate (1/h)"])
        in mori_keys
    ]
    ev9_path = osp.join(
        os.environ["DATA_ROOT"], mori.RAW_DIR_REL, mori.MIRROR_RELPATH[mori.EV9]
    )
    ev9_names = [row.gene_name for row in mori.read_mass_fractions(ev9_path, ["C2"])]
    s3_names = [str(r[0]).strip() for r in sheets[SHEET_SS][1:]]
    mori_rows_by_id = {str(r["Sample ID"]): r for r in mori_rows}

    ss_rows = sheets[SHEET_SS]
    cross_layer = {
        name: {
            "balakrishnan_column": b,
            "mori_column": m,
            "balakrishnan_growth_rate": next(
                s["Growth rate (1/h)"] for s in ss_samples if s["Sample ID"] == b
            ),
            "mori_growth_rate": mori_rows_by_id[m]["Growth rate (1/h)"],
            **log_r(column(ss_rows, b), mori_column(m)),
        }
        for name, (b, m) in CROSS_LAYER_PAIRS.items()
    }
    replicate = {
        name: (
            log_r(mori_column(a), mori_column(b))
            if name.startswith("mori")
            else log_r(column(ss_rows, a), column(ss_rows, b))
        )
        for name, (a, b) in REPLICATE_PAIRS.items()
    }

    report: dict[str, Any] = {
        "doi": DOI,
        "pmcid": PMCID,
        "nihms": NIHMS,
        "supplement_files_named_by_the_manuscript": list(SUPPLEMENT_FILES),
        "access": access,
        "table_s3": {
            "url": TABLE_S3_URL,
            "sha256": sha,
            "bytes": len(raw),
            "sheets": {name: len(rows) for name, rows in sheets.items()},
            "steady_state": {
                "described_samples": len(described),
                "groups": dict(Counter(s["Group (carried)"] for s in ss_samples)),
                "strains": dict(Counter(s["Strain"] for s in ss_samples)),
                "media": dict(Counter(s["Growth medium"] for s in ss_samples)),
                "nitrogen_source_cells": dict(
                    Counter(s["Nitrogen source"] for s in ss_samples)
                ),
                "supplements": dict(Counter(str(s["Supplement"]) for s in ss_samples)),
                "growth_rate_min": min(s["Growth rate (1/h)"] for s in ss_samples),
                "growth_rate_max": max(s["Growth rate (1/h)"] for s in ss_samples),
                "described_but_no_column": sorted(set(described) - ss_headers),
                "column_but_not_described": sorted(ss_headers - set(described)),
                "fractions": ss,
            },
            "decay_time_course": {
                "described_samples": len(deg_samples),
                "series": dict(Counter(s["Group (carried)"] for s in deg_samples)),
                "time_points_min": [s["Time after shift (min)"] for s in deg_samples],
                "fractions": deg,
            },
            "mg1655_resolution": {
                "distinct_loci": len(loci),
                "resolved": len(resolved),
                "fraction": round(len(resolved) / len(loci), 4),
                "unresolved": sorted(loci - set(genome.gene_set)),
            },
        },
        "mori2021_sister_series": {
            "mori_ev3_ncm3722_lineage_rows": len(mori_rows),
            "mori_strains": mori_strains,
            "balakrishnan_strains": bala_strains,
            "same_strain_set": mori_strains == bala_strains,
            "mori_media": sorted({str(r["Growth medium"]) for r in mori_rows}),
            "mori_nitrogen_sources": sorted(
                {str(r["Nitrogen source"]) for r in mori_rows}
            ),
            "gene_rows_mori_ev9": len(ev9_names),
            "gene_rows_table_s3": len(s3_names),
            "same_gene_name_list_in_the_same_order": ev9_names == s3_names,
            "same_gene_name_set": set(ev9_names) == set(s3_names),
            "same_strain_medium_supplement_and_growth_rate": same_culture_candidates,
            "mori_drop_rule_for_these_rows": mori.DROP_NO_MEDIA_ENTRY.rule,
            "mori_needed_media_additions": mori.DROP_NO_MEDIA_ENTRY.needed_addition,
        },
        "cross_layer_r_mrna_vs_mori_protein": cross_layer,
        "within_layer_replicate_r": replicate,
    }
    os.makedirs(RESULTS, exist_ok=True)
    with open(OUT, "w") as handle:
        json.dump(report, handle, indent=2, default=str)
    for name, result in access.items():
        print(f"{name}: {result['status']} {result['kind']} {result['bytes']} B")
    print("steady-state samples", len(described), "columns", ss["sample_columns"])
    print("duplicated", ss["duplicated_headers"], ss["duplicate_identity"])
    print("described_but_no_column", sorted(set(described) - ss_headers))
    print("sums", ss["column_sum_min"], ss["column_sum_max"])
    print("resolution", report["table_s3"]["mg1655_resolution"]["resolved"], len(loci))
    print("same-culture candidates", same_culture_candidates)
    print(
        "same gene list as Mori EV9",
        ev9_names == s3_names,
        set(ev9_names) == set(s3_names),
    )
    print("loci not a b-number", ss["loci_not_a_b_number"])
    print(json.dumps(cross_layer, indent=1))
    print(json.dumps(replicate, indent=1))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
