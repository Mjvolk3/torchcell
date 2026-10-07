# tests/torchcell/datasets/pputida/test_lim2025.py
# [[tests.torchcell.datasets.pputida.test_lim2025]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/pputida/test_lim2025.py
"""The Lim 2025 P. putida isoprenol-TALE loaders.

Synthetic tests run everywhere. They write a WordprocessingML ``si1.docx`` and an
``si2.xlsx`` from scratch, deposit them through the module's own
:func:`deposit_raw_mirror`, and build BOTH dataset classes end to end against a
synthetic KT2440 assembly, so the whole path -- deposit, manifest round trip, docx and
workbook extraction, the cross-source assertions, the retention ledger, the record
builders and the LMDB write -- is exercised with no network call and no read of the real
``$DATA_ROOT``.

The ``@pytest.mark.data`` tests read the real mirror: they pin the module's digests
against the recorded manifest, confirm the genotype finding on the real annotation, and
check the counts the module docstring states. They skip unless the mirror and the KT2440
tier cache are present.

Derived expectations for the pinned supplements, every one measured on
``si1.docx`` sha256 ``7bc74885...`` and ``si2.xlsx`` sha256 ``a3cfd601...``:
Supplementary Table 3 holds 16 lineages, four per starting strain; the mutation matrix
holds 159 rows on 159 distinct ``AE015451`` positions over 49 clone columns (3 founders,
46 evolved) as 443 calls; and the two loaded IPL400 proteome arms hold 2,367 and 2,374
rows, of which 2,361 survive into the record.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import os.path as osp
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from pydantic import ValidationError

import torchcell.datasets.pputida.lim2025 as l25
from torchcell.datamodels.media import M9_NREL_LIM2025
from torchcell.datamodels.schema import (
    AlleleEdit,
    AssayType,
    BacterialDeletionPerturbation,
    ConcentrationUnit,
    MeasurementType,
    SampleUnit,
    SequenceVariantPerturbation,
    SmallMoleculePerturbation,
)
from torchcell.literature.manifest import ROLE_RAW_DATA, RetrievalMethod

DATA_ROOT = os.environ.get("DATA_ROOT")


# --------------------------------------------------------------------------- #
# Synthetic: the genotype finding, asserted on the classes themselves
# --------------------------------------------------------------------------- #
def test_the_variant_leaf_refuses_a_pputida_locus_tag() -> None:
    """Measured reason 1: the only variant leaf admits S288C ORF names only."""
    with pytest.raises(ValidationError, match="Invalid systematic gene name format"):
        SequenceVariantPerturbation(
            systematic_gene_name="PP_3415",
            perturbed_gene_name="gnuR",
            strain_id="A12_F53_I1",
        )


def test_the_variant_leaf_has_no_slot_for_what_a_breseq_row_releases() -> None:
    """Measured reason 3: no field holds a replicon, position or base change."""
    fields = set(SequenceVariantPerturbation.model_fields)
    for absent in (
        "chromosome",
        "position",
        "reference_allele",
        "alternate_allele",
        "amino_acid_change",
        "variant_frequency",
        "gene_namespace",
    ):
        assert absent not in fields


def test_the_background_allele_requires_a_functional_boolean() -> None:
    """Measured reason 4: ``functional`` is not optional, so it cannot be gapped."""
    from torchcell.datamodels.schema import BacterialBackgroundAllele

    assert BacterialBackgroundAllele.model_fields["functional"].is_required()


def test_a_background_permits_one_allele_entry_per_locus() -> None:
    """Measured reason 5: a clone with two calls in one gene has no background form."""
    from torchcell.datamodels.schema import (
        BacterialBackgroundAllele,
        BacterialStrainBackground,
    )

    def allele(name: str) -> Any:
        return BacterialBackgroundAllele(
            systematic_gene_name="PP_3415",
            gene_namespace=l25.KT2440_NAMESPACE,
            gene_name="PP_3415",
            allele_name=name,
            edit=AlleleEdit.sequence_variant,
            functional=False,
            provenance=[l25.SOURCED_VALUES["variant_caller"]],
        )

    with pytest.raises(ValidationError, match="one allele entry per locus"):
        BacterialStrainBackground(
            name="A12_F53_I1",
            reference_strain="KT2440",
            assembly_set=l25.KT2440_ASSEMBLY_SET,
            alleles=[allele("P293S"), allele("V46I")],
            provenance=[l25.SOURCED_VALUES["variant_caller"]],
        )


def test_a_genotype_compares_by_its_perturbation_set() -> None:
    """Measured reason 6: a clone written as its parent collapses onto the parent."""
    parent = l25.strain_genotype("IPL400", {})
    clone = l25.strain_genotype("IPL400", {})
    assert parent == clone


# --------------------------------------------------------------------------- #
# Synthetic: the sourced values and the provenance anchors
# --------------------------------------------------------------------------- #
def test_every_sourced_value_is_auditable_and_quotes_a_mirrored_paper() -> None:
    """Each entry names a citation key, a sha256 and a non-empty quote."""
    assert SOURCED_KEYS <= set(l25.SOURCED_VALUES)
    for key, value in l25.SOURCED_VALUES.items():
        assert value.provenance.source_uri == l25.PAPER_MD, key
        assert value.provenance.sha256 in {l25.PAPER_MD_SHA256, l25.THOMPSON_SHA256}
        assert value.quote.strip()
        assert value.note is None or value.note.strip()


SOURCED_KEYS = {
    "reference_strain",
    "ipl300_genotype",
    "ipl400_extra_deletion",
    "n_tale_replicates",
    "pp3024_fold",
    "proteome_n_replicates",
    "pp2675_construction",
    "ipl_catabolism_abolished",
}


def test_the_thompson_deferral_is_a_separate_citation_key() -> None:
    """The construction statements defer to the paper Lim 2025 cites for them."""
    deferred = {
        key
        for key, value in l25.SOURCED_VALUES.items()
        if value.provenance.citation_key == l25.THOMPSON_KEY
    }
    assert deferred == {
        "pp2675_construction",
        "pp3839_construction",
        "pp4064_construction",
    }


def test_the_isoprenol_identity_comes_from_the_table_and_matches_the_pin() -> None:
    """The committed table row resolves to the pinned InChIKey with no gap."""
    compound = l25.isoprenol_compound()
    assert compound.name == l25.ISOPRENOL_LABEL
    assert compound.inchikey == l25.ISOPRENOL_INCHIKEY == "CPJRRXSHAYUTGL-UHFFFAOYSA-N"
    assert compound.provenance_gaps == []
    assert compound.smiles is not None


def test_a_conflicting_compound_row_stops_the_build(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A table row with another InChIKey raises rather than being stored."""
    from torchcell.datamodels.schema import Compound

    monkeypatch.setattr(
        l25,
        "resolved_compound",
        lambda name: Compound(name=name, inchikey="AAAAAAAAAAAAAA-AAAAAAAAAA-A"),
    )
    with pytest.raises(l25.CompoundIdentityConflictError, match="not CPJRRXSHAYUTGL"):
        l25.isoprenol_compound()


def test_a_matching_compound_row_is_accepted(monkeypatch: pytest.MonkeyPatch) -> None:
    """Once the table carries the right key, the resolver's object is used as is."""
    from torchcell.datamodels.schema import Compound

    monkeypatch.setattr(
        l25,
        "resolved_compound",
        lambda name: Compound(name=name, inchikey=l25.ISOPRENOL_INCHIKEY),
    )
    assert l25.isoprenol_compound().inchikey == l25.ISOPRENOL_INCHIKEY


# --------------------------------------------------------------------------- #
# Synthetic: the genotype and environment builders
# --------------------------------------------------------------------------- #
def test_ipl300_carries_the_six_designed_deletions_the_paper_lists() -> None:
    """Six deletions, every one a ``PP_`` tag the paper itself prints."""
    genotype = l25.strain_genotype("IPL300", {"PP_1385": "ttgB"})
    assert genotype.systematic_gene_names == [
        "PP_2675",
        "PP_3839",
        "PP_4064",
        "PP_4065",
        "PP_4066",
        "PP_4067",
    ]
    assert set(genotype.perturbation_types) == {"bacterial_deletion"}


def test_ipl400_is_ipl300_plus_the_ttgb_deletion() -> None:
    """The paper's one extra deletion, with its own locus tag and symbol."""
    genotype = l25.strain_genotype("IPL400", {"PP_1385": "ttgB"})
    assert len(genotype) == 7
    extra = next(
        p for p in genotype.perturbations if p.systematic_gene_name == "PP_1385"
    )
    assert isinstance(extra, BacterialDeletionPerturbation)
    assert extra.perturbed_gene_name == "ttgB"
    assert extra.gene_namespace == l25.KT2440_NAMESPACE


def test_the_wild_type_is_the_empty_genotype() -> None:
    """KT2440 IS the reference assembly, so it carries no perturbation."""
    assert l25.strain_genotype("KT2440", {}).perturbations == []


def test_an_unknown_strain_is_refused() -> None:
    """Only the campaign's own starting strains have a typed genotype."""
    with pytest.raises(ValueError, match="not a starting strain"):
        l25.strain_genotype("A10_F63_I1", {})


def test_the_pp3024_genotype_carries_its_jbei_accession() -> None:
    """The one reverse-engineered strain with a number keeps its strain accession."""
    genotype = l25.pp3024_genotype({"PP_3024": "PP_3024"})
    assert len(genotype) == 1
    perturbation = genotype.perturbations[0]
    assert isinstance(perturbation, BacterialDeletionPerturbation)
    assert perturbation.systematic_gene_name == "PP_3024"
    assert perturbation.construction is not None
    assert perturbation.construction.strain_accession == l25.PP3024_STRAIN_ACCESSION


def test_a_deletion_names_no_cassette_and_no_collection() -> None:
    """Scarless deletions of this study's own strains: both fields are None."""
    perturbation = l25.deletion("PP_3024", "PP_3024")
    assert perturbation.cassette is None
    assert perturbation.collection is None


def test_the_ipl400_background_types_the_truncation_the_axis_cannot_state() -> None:
    """Eight alleles: seven full deletions and ``PP_2676`` as a partial deletion."""
    background = l25.ipl_background("IPL400", {"PP_1385": "ttgB"})
    assert background.name == "IPL400"
    assert background.reference_strain == "KT2440"
    edits = {allele.systematic_gene_name: allele.edit for allele in background.alleles}
    assert edits["PP_2676"] is AlleleEdit.partial_deletion
    assert edits["PP_2675"] is AlleleEdit.full_deletion
    assert len(background.alleles) == 8
    assert all(allele.functional is False for allele in background.alleles)
    assert all(
        [gap.field for gap in allele.provenance_gaps] == ["deleted_span"]
        for allele in background.alleles
    )
    assert background.genotype_statement == l25.SI1_STATEMENTS["ipl400_genotype"]


def test_the_ipl300_background_omits_the_ttgb_allele() -> None:
    """IPL300 is the same lesion set without ``PP_1385``."""
    background = l25.ipl_background("IPL300", {})
    assert len(background.alleles) == 7
    assert "PP_1385" not in {a.systematic_gene_name for a in background.alleles}


def test_a_background_is_refused_for_a_strain_that_has_none() -> None:
    """Only the two stacked-deletion strains have a background."""
    with pytest.raises(ValueError, match="not a stacked-deletion strain"):
        l25.ipl_background("KT2440", {})


def test_the_isoprenol_environment_is_the_library_medium_plus_a_typed_dose() -> None:
    """One ``SmallMoleculePerturbation`` in g/L on Lim 2025's own medium object."""
    environment = l25.isoprenol_environment(6.0, l25.PANEL_DURATION_GAP)
    assert environment.media is M9_NREL_LIM2025
    assert environment.temperature is None
    assert environment.duration_hours is None
    assert len(environment.perturbations) == 1
    edit = environment.perturbations[0]
    assert isinstance(edit, SmallMoleculePerturbation)
    assert edit.concentration is not None
    assert (edit.concentration.value, edit.concentration.unit) == (
        6.0,
        ConcentrationUnit.g_per_l,
    )
    assert {gap.field for gap in environment.provenance_gaps} == {
        "temperature",
        "duration_hours",
    }


def test_the_unstressed_environment_carries_no_perturbation() -> None:
    """The proteome reference arm is the bare medium."""
    environment = l25.unstressed_environment()
    assert environment.perturbations == []
    assert environment.media is M9_NREL_LIM2025


def test_the_response_phenotype_is_a_log2_ratio_by_liquid_od_growth() -> None:
    """The typed measurement and assay axes, with the uncertainty as a gap."""
    phenotype = l25.response_phenotype(
        0.5, n_samples=4, gaps=l25.RESPONSE_UNCERTAINTY_GAPS
    )
    assert phenotype.measurement_type is MeasurementType.log2_ratio
    assert phenotype.assay_type is AssayType.liquid_od_growth
    assert phenotype.environment_response == 0.5
    assert phenotype.n_samples == 4
    assert phenotype.sample_unit is SampleUnit.biological_replicate
    assert phenotype.environment_response_uncertainty is None
    assert phenotype.environment_response_se is None
    assert {gap.field for gap in phenotype.provenance_gaps} == {
        "environment_response_uncertainty",
        "environment_response_se",
    }


def test_a_phenotype_with_no_replicate_count_states_no_sample_unit() -> None:
    """The Fig. 3A record releases no n, so neither field is invented."""
    phenotype = l25.response_phenotype(
        math.log2(1.6), n_samples=None, gaps=l25.PANEL_UNCERTAINTY_GAPS
    )
    assert phenotype.n_samples is None
    assert phenotype.sample_unit is None


def test_the_abundance_phenotype_divides_the_sample_sd_by_root_three() -> None:
    """``se = log2_std / sqrt(3)``, the back-solved replicate count."""
    phenotype = l25.abundance_phenotype(
        {"PP_0005": (21.0, 0.3), "PP_0006": (22.0, 0.0)}
    )
    assert phenotype.protein_abundance == {"PP_0005": 21.0, "PP_0006": 22.0}
    assert phenotype.protein_abundance_se is not None
    assert phenotype.protein_abundance_se["PP_0005"] == pytest.approx(
        0.3 / math.sqrt(3)
    )
    assert phenotype.n_replicates == {"PP_0005": 3, "PP_0006": 3}
    assert phenotype.measurement_type == l25.PROTEOME_MEASUREMENT_TYPE


# --------------------------------------------------------------------------- #
# Synthetic: the retention ledger
# --------------------------------------------------------------------------- #
def _drop_log(**overrides: Any) -> l25.DropLog:
    payload: dict[str, Any] = {
        "dataset": "probe",
        "source_rows": 16,
        "reference_rows": ["KT2440"],
        "candidate_records": 3,
        "kept_records": 2,
        "dropped_records": 1,
        "rules": [
            l25.DropRule(
                rule="r", scope="record", description="d", n_items=1, items=["x"]
            )
        ],
    }
    payload.update(overrides)
    return l25.DropLog(**payload)


def test_a_ledger_whose_rules_account_for_every_drop_passes() -> None:
    """The happy path: the record-scoped rules sum to the dropped count."""
    ledger = _drop_log()
    ledger.check()
    accounted = sum(rule.n_items for rule in ledger.rules if rule.scope == "record")
    assert accounted == ledger.dropped_records == 1
    assert ledger.kept_records + ledger.dropped_records == ledger.candidate_records


def test_a_ledger_refuses_a_drop_no_record_scoped_rule_accounts_for() -> None:
    """A record-scoped rule must name every dropped record."""
    with pytest.raises(RuntimeError, match="account for 0 of 1 dropped records"):
        _drop_log(
            rules=[l25.DropRule(rule="r", scope="arm", description="d", n_items=1)]
        ).check()


def test_a_ledger_refuses_counts_that_do_not_add_up() -> None:
    """Kept plus dropped must equal the candidate count."""
    with pytest.raises(RuntimeError, match="!= 9 candidates"):
        _drop_log(candidate_records=9).check()


def test_the_genotype_gaps_name_the_parents_unwritable_content() -> None:
    """Two stated gaps: the parents' called variants and PP_2676's truncation."""
    gaps = l25.genotype_gaps()
    assert len(gaps) == 2
    assert all(gap.strain == "IPL300 and IPL400" for gap in gaps)
    assert "PP_4986" in gaps[0].content
    assert gaps[0].source_quote == (
        l25.SOURCED_VALUES["preexisting_parent_mutations"].quote
    )
    assert l25.PP2676_LOCUS in gaps[1].content
    assert gaps[1].source_quote == l25.SI1_STATEMENTS["pp2676_truncation"]


def test_the_genotype_loci_are_every_locus_the_module_writes_or_gaps() -> None:
    """Nine tags: six IPL300 deletions, ttgB, PP_3024 and the gapped PP_2676."""
    loci = l25.genotype_loci()
    assert len(loci) == 9
    assert set(loci) == {
        "PP_2675",
        "PP_2676",
        "PP_3024",
        "PP_3839",
        "PP_1385",
        "PP_4064",
        "PP_4065",
        "PP_4066",
        "PP_4067",
    }


# --------------------------------------------------------------------------- #
# Synthetic: the Welch back-solve and the sheet readers
# --------------------------------------------------------------------------- #
def _proteome_row(**overrides: Any) -> l25.ProteomeRow:
    payload: dict[str, Any] = {
        "locus_tag": "PP_0005",
        "locus_tag_as_written": "PP_0005",
        "symbol": "Mnme",
        "accession": "P0A175",
        "description": "tRNA modification GTPase MnmE",
        "test_mean": 21.39,
        "test_sd": 0.048,
        "parent_mean": 21.14,
        "parent_sd": 0.294,
        "t_statistic": 0.0,
        "log2_fold_change": 0.25,
    }
    payload.update(overrides)
    row = l25.ProteomeRow(**payload)
    return row.model_copy(
        update={
            "t_statistic": overrides.get("t_statistic", l25.welch_t(row, 3)),
            "log2_fold_change": overrides.get(
                "log2_fold_change", row.test_mean - row.parent_mean
            ),
        }
    )


def test_the_welch_identity_holds_for_a_consistent_row() -> None:
    """A row built from the identity passes the assertion with a tiny residual."""
    assert l25.assert_sheet_statistics([_proteome_row()], sheet="probe") < 1e-9


def test_the_welch_identity_refuses_a_row_whose_t_implies_another_n() -> None:
    """A t computed at n = 2 is refused at n = 3, which is the back-solve."""
    row = _proteome_row()
    off = row.model_copy(update={"t_statistic": l25.welch_t(row, 2)})
    with pytest.raises(l25.CrossSourceError, match="misses the released t"):
        l25.assert_sheet_statistics([off], sheet="probe")


def test_the_fold_change_identity_refuses_a_released_value_that_is_not_the_difference() -> (
    None
):
    """The released log2 fold change must equal the difference of the two means."""
    row = _proteome_row().model_copy(update={"log2_fold_change": 9.0})
    with pytest.raises(
        l25.CrossSourceError, match="not the difference of the two means"
    ):
        l25.assert_sheet_statistics([row], sheet="probe")


def test_a_zero_denominator_row_is_skipped_rather_than_dividing_by_zero() -> None:
    """Both SDs zero means no t is computable, so only the fold change is checked."""
    row = l25.ProteomeRow(
        locus_tag="PP_0005",
        locus_tag_as_written="PP_0005",
        symbol="Mnme",
        accession="P0A175",
        description="tRNA modification GTPase MnmE",
        test_mean=21.0,
        test_sd=0.0,
        parent_mean=20.0,
        parent_sd=0.0,
        t_statistic=0.0,
        log2_fold_change=1.0,
    )
    assert l25.assert_sheet_statistics([row], sheet="probe") == 0.0


def test_shared_symbol_loci_finds_both_rows_of_a_doubly_filed_symbol() -> None:
    """A symbol under two loci returns BOTH of its locus tags."""
    rows = [
        _proteome_row(locus_tag="PP_0548", symbol="Ubid"),
        _proteome_row(locus_tag="PP_5213", symbol="Ubid"),
        _proteome_row(locus_tag="PP_0005", symbol="Mnme"),
    ]
    assert l25.shared_symbol_loci(rows) == ["PP_0548", "PP_5213"]


def test_the_parent_arm_skips_the_dropped_loci() -> None:
    """The dropped keys never reach an abundance map."""
    rows = [_proteome_row(locus_tag="PP_0005"), _proteome_row(locus_tag="PP_0548")]
    arm = l25.parent_arm(rows, ["PP_0548"])
    assert set(arm) == {"PP_0005"}


# --------------------------------------------------------------------------- #
# Synthetic: the Supplementary Table 3 parser and its cross-source assertions
# --------------------------------------------------------------------------- #
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
#: One lineage per row: (group label or "", ALE label, dose or "", initial, ending,
#: final, passages, generations, CCD). The numbers are the real Supplementary Table 3,
#: so the module's own stated ranges and doses are what the parser is checked against.
_LINEAGES: tuple[tuple[str, ...], ...] = (
    (
        "Isoprenol TALE KT2440",
        "1",
        "4 g/L",
        "0.152",
        "8.5 g/L",
        "0.113",
        "67",
        "353",
        "3.32",
    ),
    ("", "2", "", "0.152", "8.5 g/L", "0.073", "64", "320", "2.68"),
    ("", "3", "", "0.153", "8.5 g/L", "0.095", "63", "325", "2.92"),
    ("", "4", "", "0.157", "8.5 g/L", "0.057", "63", "313", "2.63"),
    (
        "Isoprenol TALE IPL300",
        "5",
        "4 g/L",
        "0.253",
        "8.5 g/L",
        "0.091",
        "74",
        "320",
        "3.17",
    ),
    ("", "6", "", "0.253", "8.5 g/L", "0.095", "70", "341", "2.68"),
    ("", "7", "", "0.252", "8.5 g/L", "0.060", "64", "297", "2.05"),
    ("", "8*", "", "0.252", "7.5 g/L", "0.137", "29", "143", "1.33"),
    (
        "Isoprenol TALE IPL400",
        "9",
        "4 g/L",
        "0.300",
        "8.5 g/L",
        "0.085",
        "67",
        "340",
        "2.96",
    ),
    ("", "10", "", "0.300", "8 g/L", "0.119", "66", "326", "2.80"),
    ("", "11", "", "0.305", "8 g/L", "0.062", "51", "244", "1.91"),
    ("", "12", "", "0.298", "8 g/L", "0.061", "55", "256", "1.94"),
    ("HCHO TALE KT2440", "13", "1 mM", "0.284", "9 mM", "0.197", "75", "387", "3.56"),
    ("", "14", "", "0.247", "9 mM", "0.204", "71", "363", "3.35"),
    ("", "15", "", "0.294", "9 mM", "0.144", "70", "362", "3.41"),
    ("", "16", "", "0.288", "9 mM", "0.157", "68", "346", "3.11"),
)


def _paragraph(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def _cell(*paragraphs: str) -> str:
    body = "".join(_paragraph(text) for text in paragraphs)
    return f"<w:tc>{body}</w:tc>"


def _row(cells: list[str]) -> str:
    return "<w:tr>" + "".join(_cell(text) for text in cells) + "</w:tr>"


def _table_s3_xml(rows: tuple[tuple[str, ...], ...] = _LINEAGES) -> str:
    header = _row(["", *l25._S3_HEADER])
    body = "".join(_row(list(row)) for row in rows)
    return f"<w:tbl>{header}{body}</w:tbl>"


def write_si1_docx(
    path: Path,
    *,
    rows: tuple[tuple[str, ...], ...] = _LINEAGES,
    tables: str | None = None,
    statements: dict[str, str] | None = None,
) -> None:
    """A minimal .docx whose body carries the SI statements and Supplementary Table 3."""
    said = l25.SI1_STATEMENTS if statements is None else statements
    paragraphs = "".join(
        _paragraph(text) for key, text in said.items() if key != "pp3024_accession"
    )
    # The accession is a cell split across two paragraphs, exactly as Word writes it.
    accession = (
        "<w:tbl><w:tr>"
        + _cell("KT2440 ΔPP_3024", f"({l25.PP3024_STRAIN_ACCESSION})")
        + "</w:tr></w:tbl>"
    )
    body = (
        paragraphs + accession + (tables if tables is not None else _table_s3_xml(rows))
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        f'2006/main"><w:body>{body}</w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document)


def test_the_docx_walk_finds_a_statement_split_across_two_paragraphs(
    tmp_path: Path,
) -> None:
    """Word splits a visual cell; the joined form is what carries the statement."""
    path = tmp_path / "si1.docx"
    write_si1_docx(path)
    text = l25.docx_text(path)
    assert l25.SI1_STATEMENTS["pp3024_accession"] in text
    assert all(statement in text for statement in l25.SI1_STATEMENTS.values())


def test_a_docx_with_no_body_is_refused(tmp_path: Path) -> None:
    """A WordprocessingML file without a body cannot be read."""
    path = tmp_path / "empty.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main"/>',
        )
    with pytest.raises(l25.TableExtractionError, match="no WordprocessingML body"):
        l25.docx_body(path)


def test_the_table_parser_reads_sixteen_lineages_with_forward_filled_groups(
    tmp_path: Path,
) -> None:
    """The group label and the starting dose are stated once per group."""
    path = tmp_path / "si1.docx"
    write_si1_docx(path)
    lineages = l25.parse_table_s3(l25.docx_tables(path))
    assert len(lineages) == 16
    assert [row.ale_number for row in lineages] == list(range(1, 17))
    assert lineages[3].starting_strain == "KT2440"
    assert lineages[3].stressor == "isoprenol"
    assert lineages[3].starting_dose == 4.0
    assert lineages[7].ale_label == "8*"
    assert lineages[7].starting_strain == "IPL300"
    assert lineages[12].stressor == "formaldehyde"
    assert lineages[12].starting_dose_unit == "mM"


def test_the_table_digest_is_independent_of_the_row_order(tmp_path: Path) -> None:
    """Sorting by ALE number makes the digest a content hash."""
    path = tmp_path / "si1.docx"
    write_si1_docx(path)
    lineages = l25.parse_table_s3(l25.docx_tables(path))
    assert l25.table_s3_digest(lineages) == l25.table_s3_digest(lineages[::-1])


def test_two_tables_with_the_same_header_are_refused(tmp_path: Path) -> None:
    """The table is located by its header, which must be unique."""
    path = tmp_path / "si1.docx"
    write_si1_docx(path, tables=_table_s3_xml() + _table_s3_xml())
    with pytest.raises(l25.TableExtractionError, match="2 SI tables carry"):
        l25.parse_table_s3(l25.docx_tables(path))


def test_a_row_with_the_wrong_cell_count_is_refused(tmp_path: Path) -> None:
    """A short row would silently shift every column."""
    short = (*_LINEAGES[:15], ("", "16", "", "0.288", "9 mM", "0.157", "68", "346"))
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=short)
    with pytest.raises(l25.TableExtractionError, match="cells, expected 9"):
        l25.parse_table_s3(l25.docx_tables(path))


def test_an_unknown_row_group_is_refused(tmp_path: Path) -> None:
    """A group label this loader does not map would lose the starting strain."""
    rows = (("Ethanol TALE KT2440", *_LINEAGES[0][1:]), *_LINEAGES[1:])
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=rows)
    with pytest.raises(l25.TableExtractionError, match="row group"):
        l25.parse_table_s3(l25.docx_tables(path))


def test_an_unparsed_ale_cell_is_refused(tmp_path: Path) -> None:
    """The ALE number is the lineage's identity, so it must parse."""
    rows = ((_LINEAGES[0][0], "one", *_LINEAGES[0][2:]), *_LINEAGES[1:])
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=rows)
    with pytest.raises(l25.TableExtractionError, match="is not an ALE number"):
        l25.parse_table_s3(l25.docx_tables(path))


def test_a_group_whose_first_row_states_no_dose_is_refused(tmp_path: Path) -> None:
    """Nothing is forward-filled across a group boundary."""
    rows = ((_LINEAGES[0][0], "1", "", *_LINEAGES[0][3:]), *_LINEAGES[1:])
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=rows)
    with pytest.raises(
        l25.TableExtractionError, match="states no starting concentration"
    ):
        l25.parse_table_s3(l25.docx_tables(path))


def test_a_malformed_dose_cell_is_refused(tmp_path: Path) -> None:
    """A dose is a number and one of the two released units."""
    rows = ((_LINEAGES[0][0], "1", "4 percent", *_LINEAGES[0][3:]), *_LINEAGES[1:])
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=rows)
    with pytest.raises(l25.TableExtractionError, match="is not a dose"):
        l25.parse_table_s3(l25.docx_tables(path))


def test_a_missing_lineage_is_refused(tmp_path: Path) -> None:
    """All sixteen lineages, numbered one to sixteen."""
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=_LINEAGES[:15])
    with pytest.raises(l25.TableExtractionError, match="expected 16"):
        l25.parse_table_s3(l25.docx_tables(path))


def test_out_of_order_ale_numbers_are_refused(tmp_path: Path) -> None:
    """A renumbered table would re-key every arm."""
    rows = (
        *_LINEAGES[:15],
        ("", "17", "", "0.288", "9 mM", "0.157", "68", "346", "3.11"),
    )
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=rows)
    with pytest.raises(l25.TableExtractionError, match="ALE numbers are"):
        l25.parse_table_s3(l25.docx_tables(path))


def _lineages(tmp_path: Path, rows: tuple[tuple[str, ...], ...] = _LINEAGES) -> Any:
    path = tmp_path / "si1.docx"
    write_si1_docx(path, rows=rows)
    return l25.parse_table_s3(l25.docx_tables(path))


def test_the_range_checks_reproduce_every_stated_span(tmp_path: Path) -> None:
    """Five recomputations of the Results' own numbers, all from Table 3's columns."""
    checks = l25.range_checks(_lineages(tmp_path))
    assert [check.name for check in checks] == [
        "isoprenol arm generations",
        "isoprenol arm CCD (x1e12)",
        "HCHO arm generations",
        "HCHO arm CCD (x1e12)",
        "HCHO arm dose, mM converted to g/L",
    ]
    assert all(check.stated == check.measured for check in checks)


def test_a_changed_generation_count_fails_the_stated_range(tmp_path: Path) -> None:
    """The cross-source check is what catches a mis-parsed column."""
    rows = ((*_LINEAGES[0][:7], "999", _LINEAGES[0][8]), *_LINEAGES[1:])
    with pytest.raises(l25.CrossSourceError, match="isoprenol arm generations"):
        l25.range_checks(_lineages(tmp_path, rows))


def test_an_isoprenol_dose_in_the_wrong_unit_is_refused(tmp_path: Path) -> None:
    """The isoprenol arm is dosed in g/L; a mM cell means the arms were swapped."""
    rows = ((_LINEAGES[0][0], "1", "4 mM", *_LINEAGES[0][3:]), *_LINEAGES[1:])
    with pytest.raises(l25.CrossSourceError, match="doses are"):
        l25.range_checks(_lineages(tmp_path, rows))


def test_a_starting_dose_off_the_stated_four_grams_is_refused(tmp_path: Path) -> None:
    """Table 3's own column must agree with the dose the Results state."""
    rows = (
        (_LINEAGES[0][0], "1", "5 g/L", *_LINEAGES[0][3:]),
        *_LINEAGES[1:4],
        (_LINEAGES[4][0], "5", "5 g/L", *_LINEAGES[4][3:]),
        *_LINEAGES[5:8],
        (_LINEAGES[8][0], "9", "5 g/L", *_LINEAGES[8][3:]),
        *_LINEAGES[9:],
    )
    with pytest.raises(l25.CrossSourceError, match="isoprenol starting doses"):
        l25.range_checks(_lineages(tmp_path, rows))


def test_the_pp3024_reading_is_corroborated_by_the_pair_mean(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """1.6 below the two-mutant mean 1.7 is what 'most of' the improvement means."""
    monkeypatch.setattr(l25, "PP3024_FOLD", 1.9)
    with pytest.raises(l25.CrossSourceError, match="is not below the two-mutant mean"):
        l25.range_checks(_lineages(tmp_path))


def test_the_arms_aggregate_four_lineages_each_against_the_wild_type(
    tmp_path: Path,
) -> None:
    """Two records and one reference arm, each a mean over four biological replicates."""
    reference, records = l25.tolerance_arms(_lineages(tmp_path))
    assert reference.strain == "KT2440"
    assert reference.log2_ratio_to_wt == 0.0
    assert reference.mean_growth_rate == pytest.approx(0.1535)
    assert [arm.strain for arm in records] == ["IPL300", "IPL400"]
    assert all(arm.n_samples == 4 for arm in records)
    assert records[0].log2_ratio_to_wt == pytest.approx(
        math.log2(0.2525 / 0.1535), rel=1e-12
    )
    assert records[1].log2_ratio_to_wt == pytest.approx(
        math.log2(0.30075 / 0.1535), rel=1e-12
    )
    assert records[0].ale_labels == ("5", "6", "7", "8*")


def test_an_arm_with_the_wrong_lineage_count_is_refused(tmp_path: Path) -> None:
    """Four lineages per starting strain is the declared replicate design."""
    moved = ("Isoprenol TALE IPL300", "4", "4 g/L", *_LINEAGES[3][3:])
    rows = (*_LINEAGES[:3], moved, ("", "5", "", *_LINEAGES[4][3:]), *_LINEAGES[5:])
    lineages = l25.parse_table_s3(l25.docx_tables(_write(tmp_path, rows)))
    with pytest.raises(l25.CrossSourceError, match="not 4 lineages"):
        l25.tolerance_arms(lineages)


def _write(tmp_path: Path, rows: tuple[tuple[str, ...], ...]) -> Path:
    path = tmp_path / "si1_alt.docx"
    write_si1_docx(path, rows=rows)
    return path


# --------------------------------------------------------------------------- #
# Synthetic: the si2.xlsx workbook
# --------------------------------------------------------------------------- #
#: The locus tags the synthetic proteome sheets quantify, plus the nine genotype loci.
PROTEOME_TAGS: tuple[str, ...] = (
    "PP_0005",
    "PP_0006",
    "PP_0008",
    "PP_0548",
    "PP_5213",
    "PP_0010",
)
#: The one tag released under isoprenol but absent from the unstressed arm, written in
#: the title case the real ``G+4IP`` sheet uses for exactly those keys.
UNREFERENCED_TAG = "PP_0002"
#: Symbols; ``Ubid`` is filed under two loci, as it is in the release.
SYMBOLS: dict[str, str] = {
    "PP_0005": "Mnme",
    "PP_0006": "Yidc",
    "PP_0008": "Rnpa",
    "PP_0548": "Ubid",
    "PP_5213": "Ubid",
    "PP_0010": "Dnaa",
    UNREFERENCED_TAG: "Ileu",
}
#: Per-arm (mean, sd) of the synthetic measurement, one pair per tag and condition.
ARMS: dict[str, tuple[float, float, float, float]] = {
    "PP_0005": (21.39, 0.048, 21.14, 0.294),
    "PP_0006": (22.66, 0.057, 22.67, 0.071),
    "PP_0008": (15.15, 1.708, 17.37, 0.293),
    "PP_0548": (19.62, 0.422, 19.21, 0.292),
    "PP_5213": (19.62, 0.422, 19.21, 0.292),
    "PP_0010": (22.72, 0.063, 22.62, 0.154),
    UNREFERENCED_TAG: (18.10, 0.210, 18.40, 0.180),
}
#: The clone columns of the synthetic mutation matrix: three founders, two evolved.
CLONES: tuple[str, ...] = (
    "A1 F0 I1 R1",
    "A5 F0 I1 R1",
    "A9 F0 I1 R1",
    "A10 F63 I1 R1",
    "A12 F53 I1 R1",
)
#: ``(region, common region, position, type, change, gene, detail, {clone: freq})``.
MUTATIONS: tuple[tuple[Any, ...], ...] = (
    (
        "Region1",
        "",
        187468,
        "DEL",
        "Δ1 bp",
        "PP_0164",
        "coding (280/693 nt)",
        {"A10 F63 I1 R1": 1},
    ),
    (
        "Region2",
        "Common region 1",
        537604,
        "SNP",
        "A→G",
        "PP_mr07, rpoB",
        "intergenic (+134/‑19)",
        {"A12 F53 I1 R1": 1},
    ),
    (
        "Region3",
        "",
        3866001,
        "SNP",
        "G→A",
        "PP_3415",
        "P293S (CCA→TCA)",
        {"A12 F53 I1 R1": 1, "A9 F0 I1 R1": 0.9},
    ),
    (
        "Region3",
        "",
        3866742,
        "SNP",
        "C→T",
        "PP_3415",
        "V46I (GTC→ATC)",
        {"A12 F53 I1 R1": 1},
    ),
    (
        "Region4",
        "",
        1570996,
        "DEL",
        "Δ22000 bp",
        "ttgB,ttgA,ttgR,PP_1388",
        "",
        {"A10 F63 I1 R1": 1, "A5 F0 I1 R1": 1},
    ),
)
SAMPLE_KEY_HEADER = (
    "S. No",
    "Sample name",
    "Condition",
    "Raw Data file name",
    "Replicates",
    "Time points",
)


def _welch(
    test_mean: float, test_sd: float, parent_mean: float, parent_sd: float
) -> float:
    return (test_mean - parent_mean) / math.sqrt(
        (test_sd**2 + parent_sd**2) / l25.PROTEOME_N_REPLICATES
    )


def _proteome_sheet(
    book: Any, sheet: str, *, tags: tuple[str, ...], title_case: bool
) -> None:
    test_suffix, parent_suffix = l25.PROTEOME_ARMS[sheet]
    worksheet = book.create_sheet(sheet)
    worksheet.append(
        [
            "Protein",
            "Locus Tag" if title_case else "Locus tag",
            "Protein.Group",
            "Protein.Names",
            "Protein.Description",
            f"log2_mean_{test_suffix}",
            f"log2_mean_{parent_suffix}",
            f"log2_std_{test_suffix}",
            f"log2_std_{parent_suffix}",
            "t-test_stat",
            "p-value",
            "p_adjusted(BH)",
            "log2_Fold_change_A/B",
        ]
    )
    for tag in tags:
        test_mean, test_sd, parent_mean, parent_sd = ARMS[tag]
        written = tag.capitalize() if title_case and tag == UNREFERENCED_TAG else tag
        worksheet.append(
            [
                SYMBOLS[tag],
                written,
                f"ACC_{tag}",
                f"{SYMBOLS[tag].upper()}_PSEPK",
                f"synthetic protein {tag}",
                test_mean,
                parent_mean,
                test_sd,
                parent_sd,
                _welch(test_mean, test_sd, parent_mean, parent_sd),
                0.5,
                0.9,
                test_mean - parent_mean,
            ]
        )


def write_si2_xlsx(
    path: Path,
    *,
    mutations: tuple[tuple[Any, ...], ...] = MUTATIONS,
    clones: tuple[str, ...] = CLONES,
    replicate_token: str = l25.PROTEOME_REPLICATE_TOKEN,
    extra_isoprenol_tag: bool = True,
    break_duplicate_export: bool = False,
) -> None:
    """A workbook with the mutation matrix, the sample key and the four proteome sheets."""
    book = openpyxl.Workbook()
    book.remove(book.active)

    sheet = book.create_sheet(l25.SHEET_MUTATIONS)
    sheet.append(
        [
            "",
            "",
            "Reference Seq",
            "Position",
            "Mutation Type",
            "Sequence Change",
            "Gene",
            "starting?",
            "Details",
            *clones,
        ]
    )
    sheet.append(["", "", "", "", "", "", "", "", "Fit mean", *[0 for _ in clones]])
    sheet.append(["", "", "", "", "", "", "", "", "Sum mutation", *[0 for _ in clones]])
    for region, common, position, kind, change, gene, detail, calls in mutations:
        sheet.append(
            [
                region,
                common,
                l25.KT2440_REPLICON,
                position,
                kind,
                change,
                gene,
                "Y",
                detail,
                *[calls.get(clone) for clone in clones],
            ]
        )

    key = book.create_sheet(l25.SHEET_SAMPLE_KEY)
    key.append([f"Dataset {l25.PRIDE_ACCESSION}"])
    key.append([])
    key.append(list(SAMPLE_KEY_HEADER))
    for index, (name, condition) in enumerate(
        (
            (l25.SAMPLE_KEY_M9G, "M9 0.4% glucose"),
            (l25.SAMPLE_KEY_IPL, "M9 0.4% glucose + 4 g/L Isoprenol"),
        ),
        start=1,
    ):
        key.append([index, name, condition, f"HGL{index}_", replicate_token, None])

    stressed = (
        (*PROTEOME_TAGS, UNREFERENCED_TAG) if extra_isoprenol_tag else PROTEOME_TAGS
    )
    _proteome_sheet(book, l25.SHEET_PROTEOME_M9G, tags=PROTEOME_TAGS, title_case=False)
    _proteome_sheet(book, l25.SHEET_PROTEOME_IPL, tags=stressed, title_case=True)
    _proteome_sheet(
        book, l25.SHEET_PROTEOME_M9G_ALT, tags=PROTEOME_TAGS, title_case=False
    )
    _proteome_sheet(
        book, l25.SHEET_PROTEOME_IPL_ALT, tags=PROTEOME_TAGS, title_case=False
    )
    if break_duplicate_export:
        book[l25.SHEET_PROTEOME_M9G_ALT].cell(row=2, column=7, value=99.0)
    book.save(path)


def test_the_mutation_matrix_reader_counts_every_dimension(tmp_path: Path) -> None:
    """Rows, clones, founders, calls, frequencies and the two unkeyable shapes."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    matrix = l25.read_mutation_matrix(str(path))
    assert matrix.replicon == l25.KT2440_REPLICON
    assert matrix.n_rows == 5
    assert matrix.n_distinct_positions == 5
    assert matrix.n_clone_columns == 5
    assert matrix.n_founder_columns == 3
    assert matrix.n_evolved_columns == 2
    assert matrix.n_calls == 7
    assert matrix.call_frequencies == {"1": 6, "0.9": 1}
    assert matrix.mutation_types == {"DEL": 2, "SNP": 3}
    assert matrix.n_intergenic == 1
    assert matrix.n_multi_locus == 1
    assert matrix.n_single_locus == 3
    assert matrix.founder_call_counts == {
        "A1 F0 I1 R1": 0,
        "A5 F0 I1 R1": 1,
        "A9 F0 I1 R1": 1,
    }


def test_an_unparsed_clone_column_is_refused(tmp_path: Path) -> None:
    """A clone column must follow the ALEdb convention or it cannot be read."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path, clones=("sample one",))
    with pytest.raises(l25.SheetExtractionError, match="unparsed clone columns"):
        l25.read_mutation_matrix(str(path))


def test_coordinates_on_another_replicon_are_refused(tmp_path: Path) -> None:
    """Every call must be on the single replicon of the pinned assembly."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    book = openpyxl.load_workbook(path)
    book[l25.SHEET_MUTATIONS].cell(row=4, column=3, value="CP000000")
    book.save(path)
    with pytest.raises(l25.SheetExtractionError, match="coordinates are on"):
        l25.read_mutation_matrix(str(path))


def test_every_call_is_typed_with_its_own_blocking_reasons(tmp_path: Path) -> None:
    """The per-row ledger the de Siqueira 2025 loader also writes."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    calls = l25.read_variant_calls(str(path))
    assert len(calls) == 5
    always = {
        l25.BLOCK_NO_BACTERIAL_VARIANT_LEAF,
        l25.BLOCK_NO_ALLELE_SEQUENCE,
        l25.BLOCK_NO_CALL_FIELDS,
        l25.BLOCK_FUNCTIONAL_REQUIRED,
        l25.BLOCK_GENOTYPE_COLLAPSE,
    }
    assert all(always <= set(call.blocking_reasons) for call in calls)
    intergenic = [call for call in calls if call.is_intergenic]
    assert len(intergenic) == 1
    assert l25.BLOCK_INTERGENIC in intergenic[0].blocking_reasons
    assert intergenic[0].loci == ["PP_mr07", "rpoB"]
    multi = [call for call in calls if len(call.loci) > 1 and not call.is_intergenic]
    assert len(multi) == 1
    assert l25.BLOCK_MULTI_LOCUS in multi[0].blocking_reasons
    stacked = [
        call
        for call in calls
        if l25.BLOCK_ONE_ALLELE_PER_LOCUS in call.blocking_reasons
    ]
    assert {call.position for call in stacked} == {3866001, 3866742}
    assert all(call.replicon == l25.KT2440_REPLICON for call in calls)


def test_the_sample_key_reader_returns_the_replicate_token(tmp_path: Path) -> None:
    """The replicate count comes from the released sheet, not from a constant."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    key = l25.read_sample_key(str(path))
    assert key[l25.SAMPLE_KEY_M9G] == l25.PROTEOME_REPLICATE_TOKEN
    assert key[l25.SAMPLE_KEY_IPL] == l25.PROTEOME_REPLICATE_TOKEN


def test_a_sample_key_with_no_rows_is_refused(tmp_path: Path) -> None:
    """An empty key would leave the replicate count unsourced."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    book = openpyxl.load_workbook(path)
    sheet = book[l25.SHEET_SAMPLE_KEY]
    book.remove(sheet)
    blank = book.create_sheet(l25.SHEET_SAMPLE_KEY)
    blank.append([f"Dataset {l25.PRIDE_ACCESSION}"])
    book.save(path)
    with pytest.raises(l25.SheetExtractionError, match="parsed no samples"):
        l25.read_sample_key(str(path))


def test_the_proteome_reader_upper_cases_a_title_cased_locus_tag(
    tmp_path: Path,
) -> None:
    """The one normalization, with the verbatim spelling kept on the row."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    rows = l25.read_proteome_sheet(str(path), l25.SHEET_PROTEOME_IPL)
    written = {row.locus_tag: row.locus_tag_as_written for row in rows}
    assert written[UNREFERENCED_TAG] == UNREFERENCED_TAG.capitalize()
    assert all(row.locus_tag.startswith("PP_") for row in rows)


def test_a_key_that_is_not_a_locus_tag_is_refused(tmp_path: Path) -> None:
    """A heterologous or contaminant key has no gene node to carry an abundance."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    book = openpyxl.load_workbook(path)
    book[l25.SHEET_PROTEOME_M9G].cell(row=2, column=2, value="Tryp_pig")
    book.save(path)
    with pytest.raises(l25.SheetExtractionError, match="is not a KT2440 locus tag"):
        l25.read_proteome_sheet(str(path), l25.SHEET_PROTEOME_M9G)


def test_a_repeated_locus_tag_in_one_sheet_is_refused(tmp_path: Path) -> None:
    """One row per locus, or an abundance map would silently lose a row."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    book = openpyxl.load_workbook(path)
    book[l25.SHEET_PROTEOME_M9G].cell(row=3, column=2, value="PP_0005")
    book.save(path)
    with pytest.raises(l25.SheetExtractionError, match="files one locus tag twice"):
        l25.read_proteome_sheet(str(path), l25.SHEET_PROTEOME_M9G)


def test_a_missing_column_is_refused(tmp_path: Path) -> None:
    """Every consumed column is named, so a renamed export stops the build."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    book = openpyxl.load_workbook(path)
    book[l25.SHEET_PROTEOME_M9G].cell(row=1, column=10, value="t_stat")
    book.save(path)
    with pytest.raises(l25.SheetExtractionError, match="missing columns"):
        l25.read_proteome_sheet(str(path), l25.SHEET_PROTEOME_M9G)


def test_a_sheet_the_workbook_does_not_carry_is_refused(tmp_path: Path) -> None:
    """A renamed sheet is a changed release, not something to search for."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    with pytest.raises(l25.SheetExtractionError, match="has no sheet"):
        l25._sheet_rows(str(path), "Not A Sheet")


def test_the_row_filter_keeps_a_row_whose_leading_cells_are_blank(
    tmp_path: Path,
) -> None:
    """The mutation sheet states a region once per region; the rest have blank leads."""
    path = tmp_path / "si2.xlsx"
    write_si2_xlsx(path)
    _, rows = l25._sheet_rows(str(path), l25.SHEET_MUTATIONS)
    positions = [row[3] for row in rows if row[3] is not None]
    assert positions == [position for _, _, position, *_ in MUTATIONS]


# --------------------------------------------------------------------------- #
# Synthetic: the raw mirror and the synthetic assembly
# --------------------------------------------------------------------------- #
ASSEMBLY_REPORT = """# Assembly name:  ASM756v2
# Organism name:  Pseudomonas putida KT2440 (g-proteobacteria)
# Infraspecific name:  strain=KT2440
# Taxid:          160488
# GenBank assembly accession: GCA_000007565.2
# RefSeq assembly accession: GCF_000007565.2
# RefSeq assembly and GenBank assemblies identical: yes
#
## Assembly-Units:
AE015451.2\tassembled-molecule\tna\tChromosome\tAE015451.2\t=\tNC_002947.4
"""
ASSEMBLY_REPORT_MEMBER = "GCA_000007565.2_ASM756v2_assembly_report.txt"
#: Every locus the synthetic annotation carries: the genotype loci and the proteome keys.
LOCUS_SPECS: tuple[tuple[str, str | None], ...] = tuple(
    (tag, SYMBOLS.get(tag, "ttgB" if tag == "PP_1385" else None))
    for tag in (*l25.genotype_loci(), *PROTEOME_TAGS, UNREFERENCED_TAG)
)


def _synthetic_loci() -> list[Any]:
    from tests.torchcell.sequence.genome._bacterial_fixtures import SyntheticLocus

    loci = []
    cursor = 1
    for index, (tag, symbol) in enumerate(LOCUS_SPECS):
        start, end = cursor, cursor + 11
        cursor = end + 3
        loci.append(
            SyntheticLocus(
                tag=tag,
                parts=((start, end),),
                strand="+" if index % 2 == 0 else "-",
                symbol=symbol,
                product=f"synthetic product {tag}",
                protein_id=f"AAN{index:05d}.1",
                protein="MKV",
            )
        )
    return loci


@pytest.fixture
def synthetic_kt2440(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """The real KT2440 class over a synthetic assembly, with the network refused."""
    import tests.torchcell.sequence.genome._bacterial_fixtures as fixtures
    from torchcell.sequence.genome.pputida.kt2440 import (
        KT2440_ASSEMBLY,
        PPutidaKT2440Genome,
    )

    monkeypatch.setattr(fixtures, "SEQUENCE", fixtures.SEQUENCE * 6)
    files = fixtures.write_assembly(
        tmp_path / "tier",
        KT2440_ASSEMBLY,
        _synthetic_loci(),
        [fixtures.gaf_row("ttgB", "ttgB|PP_1385", "GO:0000001")],
    )
    fixtures.forbid_network(monkeypatch)
    fixtures.serve_tier(monkeypatch, files)
    report = tmp_path / "tier" / ASSEMBLY_REPORT_MEMBER
    report.write_text(ASSEMBLY_REPORT)

    def serve_report(assembly_set: str, filename: str, **_: Any) -> str:
        if filename != ASSEMBLY_REPORT_MEMBER:
            raise FileNotFoundError(f"{assembly_set}/{filename} is not in the fixture")
        return str(report)

    import torchcell.datasets.bacteria_common as bacteria_common

    monkeypatch.setattr(bacteria_common, "resolve", serve_report)
    root = tmp_path / "kt2440"
    root.mkdir()
    return PPutidaKT2440Genome(genome_root=str(root), overwrite=False)


def _sha256_bytes(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def synthetic_mirror(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A raw mirror under ``tmp_path`` written by the module's own deposit function.

    Both supplements are written here, their real digests are monkeypatched over the
    module's pins, and the parsed Table 3 digest is repinned, so the deposit, the
    manifest round trip and both loaders' ``download`` run with no network call and no
    read of the real ``$DATA_ROOT``.
    """
    staging = tmp_path / "staging"
    staging.mkdir()
    docx = staging / l25.SI1_DOCX
    xlsx = staging / l25.SI2_XLSX
    write_si1_docx(docx)
    write_si2_xlsx(xlsx)
    _repin(monkeypatch, docx, xlsx)
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    root = l25.deposit_raw_mirror(
        sources={l25.SI1_DOCX: docx, l25.SI2_XLSX: xlsx}, data_root=str(data_root)
    )
    assert root == l25.raw_mirror_dir(str(data_root))
    return data_root


def _repin(monkeypatch: pytest.MonkeyPatch, docx: Path, xlsx: Path) -> None:
    """Point the module's pins at the synthetic bytes and the synthetic Table 3."""
    raw_files = tuple(
        raw.model_copy(
            update={
                "sha256": _sha256_bytes(docx if raw.name == l25.SI1_DOCX else xlsx),
                "retrieval": raw.retrieval.model_copy(
                    update={
                        "sha256": _sha256_bytes(
                            docx if raw.name == l25.SI1_DOCX else xlsx
                        )
                    }
                ),
            }
        )
        for raw in l25.RAW_FILES
    )
    monkeypatch.setattr(l25, "RAW_FILES", raw_files)
    monkeypatch.setattr(l25, "SI1_DOCX_SHA256", raw_files[0].sha256)
    monkeypatch.setattr(l25, "SI2_XLSX_SHA256", raw_files[1].sha256)
    monkeypatch.setattr(l25, "DATA_SHA256", {raw.name: raw.sha256 for raw in raw_files})
    monkeypatch.setattr(
        l25,
        "TABLE_S3_SHA256",
        l25.table_s3_digest(l25.parse_table_s3(l25.docx_tables(docx))),
    )


def test_the_deposit_writes_both_files_and_a_manifest_that_pins_them(
    synthetic_mirror: Path,
) -> None:
    """Two ``raw_data`` records, each with a re-runnable Elsevier retrieval."""
    manifest = l25.load_manifest(str(synthetic_mirror))
    assert manifest.citation_key == l25.CITATION_KEY
    assert manifest.doi == l25.PAPER_DOI
    assert manifest.title == l25.PAPER_TITLE
    assert [record.path for record in manifest.files] == [
        f"data/{l25.SI1_DOCX}",
        f"data/{l25.SI2_XLSX}",
    ]
    for record in manifest.files:
        assert record.role == ROLE_RAW_DATA
        assert record.retrieval is not None
        assert record.retrieval.method is RetrievalMethod.direct_url
        assert record.retrieval.retriever == (
            "torchcell.literature.retrieve.elsevier_mmc"
        )
        assert record.retrieval.params["pii"] == l25.PAPER_PII
        assert record.retrieval.sha256 == record.sha256
    assert manifest.provenance_complete is True


def test_the_manifest_records_the_three_sequencing_accessions(
    synthetic_mirror: Path,
) -> None:
    """The deposits that are recorded rather than downloaded."""
    manifest = l25.load_manifest(str(synthetic_mirror))
    joined = " ".join(manifest.si_data_sources)
    assert l25.SRA_BIOPROJECT in joined
    assert l25.GEO_SERIES in joined
    assert l25.PRIDE_ACCESSION in joined
    expected = " ".join(manifest.si_expected)
    assert l25.ALEDB_PROJECT in expected
    assert "RECORDED, not downloaded" in expected


def test_manifest_sha256_raises_on_a_path_the_manifest_does_not_list(
    synthetic_mirror: Path,
) -> None:
    """A pin can only be read for a file the mirror actually holds."""
    manifest = l25.load_manifest(str(synthetic_mirror))
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        l25.manifest_sha256(manifest, "data/si3.pdf")


def test_the_deposit_is_idempotent_by_sha256(
    synthetic_mirror: Path, tmp_path: Path
) -> None:
    """Re-depositing the same bytes leaves the mirror alone."""
    deposited = l25.raw_mirror_dir(str(synthetic_mirror)) / f"data/{l25.SI1_DOCX}"
    before = deposited.stat().st_mtime_ns
    l25.deposit_raw_mirror(
        sources={
            l25.SI1_DOCX: tmp_path / "staging" / l25.SI1_DOCX,
            l25.SI2_XLSX: tmp_path / "staging" / l25.SI2_XLSX,
        },
        data_root=str(synthetic_mirror),
    )
    assert deposited.stat().st_mtime_ns == before


def test_the_deposit_refuses_a_mirror_file_whose_bytes_differ(
    synthetic_mirror: Path, tmp_path: Path
) -> None:
    """A differing mirror file raises rather than being overwritten."""
    deposited = l25.raw_mirror_dir(str(synthetic_mirror)) / f"data/{l25.SI2_XLSX}"
    deposited.write_bytes(b"not the workbook")
    with pytest.raises(RuntimeError, match="exists with a different sha256; refusing"):
        l25.deposit_raw_mirror(
            sources={
                l25.SI1_DOCX: tmp_path / "staging" / l25.SI1_DOCX,
                l25.SI2_XLSX: tmp_path / "staging" / l25.SI2_XLSX,
            },
            data_root=str(synthetic_mirror),
        )


def test_the_deposit_refuses_a_source_off_its_pin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A source whose bytes are not the pinned ones never reaches the mirror."""
    staging = tmp_path / "off"
    staging.mkdir()
    (staging / l25.SI1_DOCX).write_bytes(b"not a docx")
    (staging / l25.SI2_XLSX).write_bytes(b"not a workbook")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "root"))
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        l25.deposit_raw_mirror(
            sources={
                l25.SI1_DOCX: staging / l25.SI1_DOCX,
                l25.SI2_XLSX: staging / l25.SI2_XLSX,
            }
        )


def test_the_deposit_refuses_a_missing_source(tmp_path: Path) -> None:
    """Every file in ``RAW_FILES`` must be given a source."""
    with pytest.raises(KeyError, match="no source given for"):
        l25.deposit_raw_mirror(
            sources={l25.SI1_DOCX: tmp_path / "nothing"},
            data_root=str(tmp_path / "root"),
        )


def test_the_retriever_writes_the_verified_bytes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``retrieve_raw_files`` runs the RECORDED retrieval, so a pin mismatch raises."""
    payload = {raw.name: f"{raw.name} bytes".encode() for raw in l25.RAW_FILES}
    raw_files = tuple(
        raw.model_copy(update={"sha256": hashlib.sha256(payload[raw.name]).hexdigest()})
        for raw in l25.RAW_FILES
    )
    monkeypatch.setattr(l25, "RAW_FILES", raw_files)
    monkeypatch.setattr(
        l25,
        "run_retriever",
        lambda record: payload[
            Path(record.params["filename"])
            .stem.replace("mmc1", l25.SI1_DOCX)
            .replace("mmc2", l25.SI2_XLSX)
        ],
    )
    out = l25.retrieve_raw_files(tmp_path / "download")
    assert set(out) == {l25.SI1_DOCX, l25.SI2_XLSX}
    assert out[l25.SI1_DOCX].read_bytes() == payload[l25.SI1_DOCX]


# --------------------------------------------------------------------------- #
# Synthetic: both loaders built end to end
# --------------------------------------------------------------------------- #
@pytest.fixture
def built_tolerance(
    synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path
) -> Any:
    """The tolerance loader built from the synthetic mirror and assembly."""
    return l25.IsoprenolToleranceLim2025Dataset(
        root=str(tmp_path / "build" / "tolerance"), pputida_genome=synthetic_kt2440
    )


@pytest.fixture
def built_proteome(
    synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path
) -> Any:
    """The proteome loader built from the synthetic mirror and assembly."""
    return l25.ProteomeLim2025Dataset(
        root=str(tmp_path / "build" / "proteome"), pputida_genome=synthetic_kt2440
    )


def test_the_tolerance_loader_builds_three_records(built_tolerance: Any) -> None:
    """IPL300 and IPL400 at 4 g/L, plus dPP_3024 at 6 g/L."""
    assert len(built_tolerance) == 3


def test_every_tolerance_record_is_a_log2_ratio_against_a_zero_reference(
    built_tolerance: Any,
) -> None:
    """The reference arm is KT2440 at log2(1) = 0 in each record's own environment."""
    for index in range(len(built_tolerance)):
        item = built_tolerance[index]
        assert item["experiment"]["phenotype"]["measurement_type"] == "log2_ratio"
        assert item["experiment"]["phenotype"]["assay_type"] == "liquid_od_growth"
        reference = item["reference"]
        assert reference["phenotype_reference"]["environment_response"] == 0.0
        assert reference["genome_reference"]["background"] is None
        assert reference["genome_reference"]["assembly_set"] == l25.KT2440_ASSEMBLY_SET


def test_the_tolerance_records_carry_the_three_expected_genotypes(
    built_tolerance: Any,
) -> None:
    """Six, seven and one deletion, each on the KT2440 namespace."""
    sizes = sorted(
        len(built_tolerance[index]["experiment"]["genotype"]["perturbations"])
        for index in range(len(built_tolerance))
    )
    assert sizes == [1, 6, 7]
    namespaces = {
        perturbation["gene_namespace"]
        for index in range(len(built_tolerance))
        for perturbation in built_tolerance[index]["experiment"]["genotype"][
            "perturbations"
        ]
    }
    assert namespaces == {l25.KT2440_NAMESPACE}


def test_the_two_tale_records_reproduce_the_table_three_ratios(
    built_tolerance: Any,
) -> None:
    """Each arm's mean over four lineages, divided by the wild type's."""
    values = {
        tuple(
            sorted(
                perturbation["systematic_gene_name"]
                for perturbation in built_tolerance[index]["experiment"]["genotype"][
                    "perturbations"
                ]
            )
        ): built_tolerance[index]["experiment"]["phenotype"]["environment_response"]
        for index in range(len(built_tolerance))
    }
    ipl300 = values[("PP_2675", "PP_3839", "PP_4064", "PP_4065", "PP_4066", "PP_4067")]
    assert ipl300 == pytest.approx(math.log2(0.2525 / 0.1535), rel=1e-12)
    assert values[("PP_3024",)] == pytest.approx(math.log2(1.6), rel=1e-12)


def test_the_tolerance_build_writes_every_ledger(built_tolerance: Any) -> None:
    """The retention ledger, the identifier histogram and the typed call rows."""
    out = Path(built_tolerance.preprocess_dir)
    for name in (
        "dropped_records.json",
        "identifier_reconciliation.json",
        "variant_accounting.json",
        "called_variants.json",
        "genotype_gaps.json",
        "extraction.json",
        "table_s3.csv",
        "tale_lineages.csv",
    ):
        assert (out / name).exists(), name
    drops = json.loads((out / "dropped_records.json").read_text())
    assert drops["kept_records"] == 3
    assert drops["dropped_records"] == 2
    rules = {rule["rule"] for rule in drops["rules"]}
    assert "hcho_arm_has_no_strain_other_than_the_reference" in rules
    assert "final_growth_rate_is_an_evolved_population" in rules
    assert "evolved_clone_genotype_is_not_representable" in rules
    assert "released_only_as_a_figure" in rules
    assert "no_titer_is_released_per_writable_strain" in rules
    calls = json.loads((out / "called_variants.json").read_text())
    assert len(calls) == 5
    assert all(call["blocking_reasons"] for call in calls)
    extraction = json.loads((out / "extraction.json").read_text())
    assert len(extraction["range_checks"]) == 5
    assert extraction["si1_statements"] == l25.SI1_STATEMENTS


def test_the_tolerance_table_csv_carries_the_reference_arm_as_a_recordless_row(
    built_tolerance: Any,
) -> None:
    """The wild-type arm's four lineages stay readable beside the two records."""
    import pandas as pd

    table = pd.read_csv(Path(built_tolerance.preprocess_dir) / "table_s3.csv")
    assert len(table) == 4
    reference = table[table["record"].isna()]
    assert len(reference) == 1
    assert reference.iloc[0]["strain"] == "KT2440"
    assert reference.iloc[0]["log2_ratio_to_wt"] == 0.0
    assert reference.iloc[0]["ale_labels"] == "1;2;3;4"


def test_the_proteome_loader_builds_one_record_over_the_shared_loci(
    built_proteome: Any,
) -> None:
    """One record: the parent under isoprenol, keyed to the reference arm's loci."""
    assert len(built_proteome) == 1
    item = built_proteome[0]
    abundance = item["experiment"]["phenotype"]["protein_abundance"]
    assert set(abundance) == set(PROTEOME_TAGS) - {"PP_0548", "PP_5213"}
    assert set(item["reference"]["phenotype_reference"]["protein_abundance"]) == set(
        abundance
    )
    assert item["experiment"]["phenotype"]["n_replicates"] == dict.fromkeys(
        abundance, 3
    )


def test_the_proteome_record_stores_the_released_log2_value_verbatim(
    built_proteome: Any,
) -> None:
    """Nothing is exponentiated; the SE is the released SD over sqrt(3)."""
    phenotype = built_proteome[0]["experiment"]["phenotype"]
    _, _, parent_mean, parent_sd = ARMS["PP_0005"]
    assert phenotype["protein_abundance"]["PP_0005"] == pytest.approx(parent_mean)
    assert phenotype["protein_abundance_se"]["PP_0005"] == pytest.approx(
        parent_sd / math.sqrt(3)
    )
    assert phenotype["measurement_type"] == l25.PROTEOME_MEASUREMENT_TYPE
    reference = built_proteome[0]["reference"]["phenotype_reference"]
    assert reference["protein_abundance"]["PP_0005"] == pytest.approx(parent_mean)


def test_the_proteome_reference_carries_the_full_ipl400_background(
    built_proteome: Any,
) -> None:
    """Where the comparison's reference strain IS IPL400, the background is typed."""
    background = built_proteome[0]["reference"]["genome_reference"]["background"]
    assert background is not None
    assert background["name"] == "IPL400"
    edits = {
        allele["systematic_gene_name"]: allele["edit"]
        for allele in background["alleles"]
    }
    assert edits["PP_2676"] == AlleleEdit.partial_deletion.value
    assert edits["PP_2675"] == AlleleEdit.full_deletion.value
    assert len(background["alleles"]) == 8


def test_the_proteome_record_is_stressed_and_its_reference_is_not(
    built_proteome: Any,
) -> None:
    """The environment is the edit: isoprenol on the record, nothing on the reference."""
    item = built_proteome[0]
    assert len(item["experiment"]["environment"]["perturbations"]) == 1
    assert item["reference"]["environment_reference"]["perturbations"] == []
    assert item["experiment"]["environment"]["media"]["name"] == M9_NREL_LIM2025.name


def test_the_proteome_build_writes_its_back_solve_and_its_ledger(
    built_proteome: Any,
) -> None:
    """The measured replicate count and uncertainty type are written beside the build."""
    out = Path(built_proteome.preprocess_dir)
    back = json.loads((out / "statistics_back_solve.json").read_text())
    assert back["n_replicates"] == 3
    assert back["uncertainty_type"].startswith("sample_sd")
    assert all(value < 1e-6 for value in back["welch_t_worst_residual"].values())
    assert back["sample_key"][l25.SAMPLE_KEY_M9G] == l25.PROTEOME_REPLICATE_TOKEN
    for pair, summary in back["duplicate_exports"].items():
        assert summary["bit_identical_on_shared"] is True, pair
    drops = json.loads((out / "dropped_records.json").read_text())
    assert drops["kept_records"] == 1
    assert drops["dropped_records"] == 0
    rules = {rule["rule"]: rule for rule in drops["rules"]}
    assert rules["gene_symbol_filed_under_two_paralogous_loci"]["items"] == [
        "PP_0548",
        "PP_5213",
    ]
    assert rules["no_key_matched_reference_abundance"]["items"] == [UNREFERENCED_TAG]
    assert rules["production_medium_has_no_media_library_entry"]["n_items"] == 6
    assert (out / "ipl400_proteome.csv").exists()


def test_a_sample_key_without_three_replicates_stops_the_proteome_build(
    synthetic_kt2440: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The stored ``n_replicates`` must be the one the released key states."""
    data_root = _mirror_with(
        tmp_path, monkeypatch, xlsx_kwargs={"replicate_token": "R1,R2"}
    )
    with pytest.raises(l25.SheetExtractionError, match="not 'R1,R2,R3'"):
        l25.ProteomeLim2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )
    assert data_root.exists()


def test_a_disagreeing_duplicate_export_stops_the_proteome_build(
    synthetic_kt2440: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Four sheets export one measurement; a disagreement is a changed release."""
    _mirror_with(tmp_path, monkeypatch, xlsx_kwargs={"break_duplicate_export": True})
    with pytest.raises(l25.CrossSourceError, match="disagree on IPL400"):
        l25.ProteomeLim2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_missing_si_statement_stops_the_tolerance_build(
    synthetic_kt2440: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A quote that moved refuses the build rather than being silently dropped."""
    statements = dict(l25.SI1_STATEMENTS)
    statements["pp2676_truncation"] = "a sentence this document does not contain"
    _mirror_with(tmp_path, monkeypatch, docx_kwargs={"statements": statements})
    with pytest.raises(l25.TableExtractionError, match="no longer states"):
        l25.IsoprenolToleranceLim2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_repinned_table_digest_stops_the_tolerance_build(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The parsed table is content-hashed, so a changed value refuses the build."""
    monkeypatch.setattr(l25, "TABLE_S3_SHA256", "0" * 64)
    with pytest.raises(l25.TableExtractionError, match="pinned 0000"):
        l25.IsoprenolToleranceLim2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_mirror_file_off_its_pin_stops_the_build(
    synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path
) -> None:
    """``download`` verifies the manifest pin before linking anything."""
    from torchcell.data.experiment_dataset import RawSha256MismatchError

    target = l25.raw_mirror_dir(str(synthetic_mirror)) / f"data/{l25.SI2_XLSX}"
    target.write_bytes(b"not the workbook")
    with pytest.raises(RawSha256MismatchError, match="sha256 mismatch"):
        l25.ProteomeLim2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_mirror_missing_a_file_stops_the_build(
    synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path
) -> None:
    """A mirror the manifest lists but the disk does not hold raises by name."""
    target = l25.raw_mirror_dir(str(synthetic_mirror)) / f"data/{l25.SI1_DOCX}"
    target.unlink()
    with pytest.raises(RuntimeError, match="missing from mirror"):
        l25.IsoprenolToleranceLim2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def _mirror_with(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    *,
    docx_kwargs: dict[str, Any] | None = None,
    xlsx_kwargs: dict[str, Any] | None = None,
) -> Path:
    """A synthetic mirror whose supplements are written with the given overrides."""
    staging = tmp_path / "doctored"
    staging.mkdir()
    docx = staging / l25.SI1_DOCX
    xlsx = staging / l25.SI2_XLSX
    write_si1_docx(docx, **(docx_kwargs or {}))
    write_si2_xlsx(xlsx, **(xlsx_kwargs or {}))
    _repin(monkeypatch, docx, xlsx)
    data_root = tmp_path / "doctored_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    l25.deposit_raw_mirror(
        sources={l25.SI1_DOCX: docx, l25.SI2_XLSX: xlsx}, data_root=str(data_root)
    )
    return data_root


def test_both_loaders_declare_their_schema_classes_and_consumed_files(
    built_tolerance: Any, built_proteome: Any
) -> None:
    """The interface every registered dataset states."""
    from torchcell.datamodels.schema import (
        BacterialEnvironmentResponseExperiment,
        BacterialEnvironmentResponseExperimentReference,
        BacterialProteinAbundanceExperiment,
        BacterialProteinAbundanceExperimentReference,
    )

    assert built_tolerance.experiment_class is BacterialEnvironmentResponseExperiment
    assert built_tolerance.reference_class is (
        BacterialEnvironmentResponseExperimentReference
    )
    assert built_proteome.experiment_class is BacterialProteinAbundanceExperiment
    assert built_proteome.reference_class is (
        BacterialProteinAbundanceExperimentReference
    )
    assert built_tolerance.raw_file_names == [l25.SI1_DOCX, l25.SI2_XLSX]
    assert built_proteome.raw_file_names == [l25.SI1_DOCX, l25.SI2_XLSX]
    for dataset in (built_tolerance, built_proteome):
        assert dataset.preprocess_raw("df") == "df"
        with pytest.raises(NotImplementedError):
            dataset.create_experiment()


def test_both_loaders_are_in_the_dataset_registry() -> None:
    """``@register_dataset`` is what the build CLI resolves a class name through."""
    from torchcell.datasets.dataset_registry import dataset_registry

    assert dataset_registry["IsoprenolToleranceLim2025Dataset"] is (
        l25.IsoprenolToleranceLim2025Dataset
    )
    assert dataset_registry["ProteomeLim2025Dataset"] is l25.ProteomeLim2025Dataset


def test_a_genome_of_another_assembly_set_is_refused(
    synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path
) -> None:
    """A loader pinned to KT2440 refuses a genome of another assembly."""

    class Wrong:
        ASSEMBLY_SET = "ecoli_K12_MG1655_ASM584v2"

    for cls in (l25.IsoprenolToleranceLim2025Dataset, l25.ProteomeLim2025Dataset):
        dataset = cls.__new__(cls)
        dataset.pputida_genome = Wrong()
        with pytest.raises(
            ValueError, match="needs the pputida_KT2440_ASM756v2 genome"
        ):
            dataset._genome()


def test_the_genome_is_opened_from_the_tier_when_none_is_injected(
    synthetic_kt2440: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A direct run opens the pinned assembly's default cache exactly once."""
    calls: list[tuple[Any, ...]] = []

    def fake(*args: Any, **kwargs: Any) -> Any:
        calls.append(args)
        return synthetic_kt2440

    monkeypatch.setattr(l25, "bacterial_genome", fake)
    dataset = l25.ProteomeLim2025Dataset.__new__(l25.ProteomeLim2025Dataset)
    dataset.pputida_genome = None
    assert dataset._genome() is synthetic_kt2440
    assert dataset._genome() is synthetic_kt2440
    assert calls == [("pputida", "KT2440")]


def test_the_standard_names_keep_a_tag_whose_symbol_does_not_round_trip(
    synthetic_kt2440: Any,
) -> None:
    """A stored (tag, symbol) pair always resolves back to the tag."""
    names = l25.standard_names(synthetic_kt2440, ["PP_1385", "PP_3024"])
    assert names["PP_1385"] == "ttgB"
    assert names["PP_3024"] == "PP_3024"


def test_the_genotype_loci_all_resolve_on_the_pinned_assembly(
    synthetic_kt2440: Any,
) -> None:
    """Checklist item 4: a complete resolution, and the histogram to prove it."""
    symbols, report = l25.reconcile_genotype_loci(synthetic_kt2440, label="probe")
    assert report.resolved_fraction == 1.0
    assert report.unique_names == 9
    assert report.gene_namespace == l25.KT2440_NAMESPACE
    assert set(symbols) == set(l25.genotype_loci())


# --------------------------------------------------------------------------- #
# The real mirror (@pytest.mark.data)
# --------------------------------------------------------------------------- #
MIRROR_PRESENT = bool(DATA_ROOT) and osp.isfile(
    osp.join(DATA_ROOT or "", l25.RAW_DIR_REL, "manifest.json")
)


@pytest.mark.data
@pytest.mark.skipif(not MIRROR_PRESENT, reason="the Lim 2025 raw mirror is not mounted")
def test_the_mirror_manifest_records_the_digests_the_module_pins() -> None:
    """The deposited bytes are the ones every sourced value and reader assumes."""
    manifest = l25.load_manifest(DATA_ROOT)
    assert l25.manifest_sha256(manifest, f"data/{l25.SI1_DOCX}") == l25.SI1_DOCX_SHA256
    assert l25.manifest_sha256(manifest, f"data/{l25.SI2_XLSX}") == l25.SI2_XLSX_SHA256


@pytest.mark.data
@pytest.mark.skipif(not MIRROR_PRESENT, reason="the Lim 2025 raw mirror is not mounted")
def test_the_real_table_three_parses_to_the_pinned_digest() -> None:
    """The pinned Supplementary Table 3 digest is the one this extraction produces."""
    docx = osp.join(DATA_ROOT or "", l25.RAW_DIR_REL, "data", l25.SI1_DOCX)
    lineages = l25.parse_table_s3(l25.docx_tables(docx))
    assert len(lineages) == 16
    assert l25.table_s3_digest(lineages) == l25.TABLE_S3_SHA256
    reference, records = l25.tolerance_arms(lineages)
    assert reference.mean_growth_rate == pytest.approx(0.1535)
    assert [arm.mean_growth_rate for arm in records] == [
        pytest.approx(0.2525),
        pytest.approx(0.30075),
    ]


@pytest.mark.data
@pytest.mark.skipif(not MIRROR_PRESENT, reason="the Lim 2025 raw mirror is not mounted")
def test_the_real_mutation_matrix_holds_the_counts_the_docstring_states() -> None:
    """159 rows over 49 clones as 443 calls, partitioned into the three shapes."""
    xlsx = osp.join(DATA_ROOT or "", l25.RAW_DIR_REL, "data", l25.SI2_XLSX)
    matrix = l25.read_mutation_matrix(xlsx)
    assert matrix.n_rows == 159
    assert matrix.n_distinct_positions == 159
    assert matrix.n_clone_columns == 49
    assert matrix.n_founder_columns == 3
    assert matrix.n_evolved_columns == 46
    assert matrix.n_calls == 443
    assert matrix.call_frequencies == {"1": 431, "0.9": 12}
    assert matrix.mutation_types == {"SNP": 100, "DEL": 47, "INS": 11, "SUB": 1}
    assert (matrix.n_single_locus, matrix.n_intergenic, matrix.n_multi_locus) == (
        123,
        17,
        19,
    )
    assert matrix.founder_call_counts == {
        "A1 F0 I1 R1": 0,
        "A5 F0 I1 R1": 6,
        "A9 F0 I1 R1": 8,
    }


@pytest.mark.data
@pytest.mark.skipif(not MIRROR_PRESENT, reason="the Lim 2025 raw mirror is not mounted")
def test_every_real_call_carries_a_blocking_reason() -> None:
    """All 159 rows are typed, and the two extra shapes are counted."""
    xlsx = osp.join(DATA_ROOT or "", l25.RAW_DIR_REL, "data", l25.SI2_XLSX)
    calls = l25.read_variant_calls(xlsx)
    assert len(calls) == 159
    assert all(len(call.blocking_reasons) >= 5 for call in calls)
    assert sum(l25.BLOCK_INTERGENIC in c.blocking_reasons for c in calls) == 17
    assert sum(l25.BLOCK_ONE_ALLELE_PER_LOCUS in c.blocking_reasons for c in calls) == 8


@pytest.mark.data
@pytest.mark.skipif(not MIRROR_PRESENT, reason="the Lim 2025 raw mirror is not mounted")
def test_the_real_proteome_sheets_back_solve_to_three_sample_sd_replicates() -> None:
    """Welch's t at n = 3 reproduces for every kept row; the six paralogs are dropped."""
    xlsx = osp.join(DATA_ROOT or "", l25.RAW_DIR_REL, "data", l25.SI2_XLSX)
    baseline = l25.read_proteome_sheet(xlsx, l25.SHEET_PROTEOME_M9G)
    stressed = l25.read_proteome_sheet(xlsx, l25.SHEET_PROTEOME_IPL)
    assert (len(baseline), len(stressed)) == (2367, 2374)
    shared = set(l25.shared_symbol_loci(baseline)) | set(
        l25.shared_symbol_loci(stressed)
    )
    assert shared == {"PP_0548", "PP_5213", "PP_1086", "PP_4999", "PP_1237", "PP_2639"}
    for rows, sheet in (
        (baseline, l25.SHEET_PROTEOME_M9G),
        (stressed, l25.SHEET_PROTEOME_IPL),
    ):
        kept = [row for row in rows if row.locus_tag not in shared]
        assert l25.assert_sheet_statistics(kept, sheet=sheet) < 1e-9
    arms = (l25.parent_arm(baseline, shared), l25.parent_arm(stressed, shared))
    assert len(set(arms[0]) & set(arms[1])) == 2361
    assert sorted(set(arms[1]) - set(arms[0])) == [
        "PP_0002",
        "PP_0416",
        "PP_0985",
        "PP_2271",
        "PP_3610",
        "PP_3699",
        "PP_5287",
    ]


# --------------------------------------------------------------------------- #
# Synthetic: the module's own CLI and verification runners
# --------------------------------------------------------------------------- #
def test_the_cli_deposit_writes_the_mirror(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``deposit --download-dir`` copies already-retrieved files into the mirror."""
    staging = tmp_path / "cli_staging"
    staging.mkdir()
    docx = staging / l25.SI1_DOCX
    xlsx = staging / l25.SI2_XLSX
    write_si1_docx(docx)
    write_si2_xlsx(xlsx)
    _repin(monkeypatch, docx, xlsx)
    data_root = tmp_path / "cli_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    assert l25.main(["deposit", "--download-dir", str(staging)]) == 0
    assert (l25.raw_mirror_dir(str(data_root)) / "manifest.json").exists()


def test_the_cli_deposit_can_run_the_recorded_retrieval(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``--retrieve`` runs the recorded retriever before depositing."""
    staging = tmp_path / "cli_retrieve"
    source = tmp_path / "source"
    source.mkdir()
    docx = source / l25.SI1_DOCX
    xlsx = source / l25.SI2_XLSX
    write_si1_docx(docx)
    write_si2_xlsx(xlsx)
    _repin(monkeypatch, docx, xlsx)
    payload = {"mmc1.docx": docx.read_bytes(), "mmc2.xlsx": xlsx.read_bytes()}
    monkeypatch.setattr(
        l25, "run_retriever", lambda record: payload[record.params["filename"]]
    )
    data_root = tmp_path / "cli_retrieve_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    assert l25.main(["deposit", "--download-dir", str(staging), "--retrieve"]) == 0
    assert (l25.raw_mirror_dir(str(data_root)) / "manifest.json").exists()


def test_the_cli_build_builds_both_datasets(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """``build`` resolves the genome once and builds both dev-tree stores."""
    monkeypatch.setattr(l25, "bacterial_genome", lambda *a, **k: synthetic_kt2440)
    monkeypatch.setattr(l25, "TOLERANCE_ROOT_REL", "cli/tolerance")
    monkeypatch.setattr(l25, "PROTEOME_ROOT_REL", "cli/proteome")
    assert l25.main(["build"]) == 0
    for relative, expected in (("cli/tolerance", 3), ("cli/proteome", 1)):
        drops = json.loads(
            (
                synthetic_mirror / relative / "preprocess" / "dropped_records.json"
            ).read_text()
        )
        assert drops["kept_records"] == expected


def test_the_count_oracle_is_the_ledger_the_build_wrote(
    synthetic_mirror: Path, built_tolerance: Any
) -> None:
    """L1's expected count comes from the build's own retention ledger."""
    assert l25._expected_count(built_tolerance.root) == 3


def test_both_verification_runners_pass_on_the_synthetic_builds(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """L0 to L4, including the provenance audit, over the synthetic stores."""
    import torchcell.verification.runners as runners

    monkeypatch.setattr(l25, "bacterial_genome", lambda *a, **k: synthetic_kt2440)
    monkeypatch.setattr(l25, "TOLERANCE_ROOT_REL", "verify/tolerance")
    monkeypatch.setattr(l25, "PROTEOME_ROOT_REL", "verify/proteome")
    assert l25.main(["build"]) == 0
    universe = {locus for locus, _ in LOCUS_SPECS}
    monkeypatch.setattr(
        runners, "_genome_for_reference", lambda *a, **k: synthetic_kt2440
    )
    monkeypatch.setattr(runners, "_gene_set_for_reference", lambda *a, **k: universe)
    monkeypatch.setattr(l25, "_audit_sourced_values", lambda report, data_root: None)
    reports = l25.run_verification(str(synthetic_mirror))
    assert len(reports) == 2
    for report in reports:
        assert report.passed, report.summary()
    names = {result.name for result in reports[0].results}
    assert "gene_containment_kt2440_locus_tags" in names
    proteome_names = {result.name for result in reports[1].results}
    assert "gene_containment_kt2440_quantified_loci" in proteome_names


def test_the_cli_verify_returns_zero_when_both_reports_pass(
    synthetic_mirror: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``verify`` is the module's own L0-L4 entry point."""
    from torchcell.verification.report import (
        Level,
        LevelResult,
        Provenance,
        VerificationReport,
    )

    passing = VerificationReport(
        dataset_name="probe", provenance=Provenance(source_uri="probe")
    )
    passing.add(LevelResult(level=Level.L1, name="count", passed=True, message="ok"))
    monkeypatch.setattr(l25, "run_verification", lambda root: (passing, passing))
    assert l25.main(["verify"]) == 0


def test_the_cli_verify_returns_one_when_a_report_fails(
    synthetic_mirror: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A failing level makes the runner exit non-zero."""
    from torchcell.verification.report import (
        Level,
        LevelResult,
        Provenance,
        VerificationReport,
    )

    failing = VerificationReport(
        dataset_name="probe", provenance=Provenance(source_uri="probe")
    )
    failing.add(LevelResult(level=Level.L1, name="count", passed=False, message="no"))
    monkeypatch.setattr(l25, "run_verification", lambda root: (failing,))
    assert l25.main(["verify"]) == 1


def test_a_store_with_two_genome_references_is_refused(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """One reference per store, so one gene universe decides L4."""
    import torchcell.verification.runners as runners

    monkeypatch.setattr(
        runners,
        "load_records",
        lambda root: [
            {"reference": {"genome_reference": {"strain": "KT2440"}}},
            {"reference": {"genome_reference": {"strain": "IPL400"}}},
        ],
    )
    monkeypatch.setenv("DATA_ROOT", "/nowhere")
    with pytest.raises(ValueError, match="2 distinct genome references"):
        l25.run_tolerance_verification()
    with pytest.raises(ValueError, match="2 distinct genome references"):
        l25.run_proteome_verification()
