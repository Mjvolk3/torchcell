# tests/torchcell/datasets/scerevisiae/test_dechiara2022.py
# [[tests.torchcell.datasets.scerevisiae.test_dechiara2022]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/scerevisiae/test_dechiara2022.py
"""De Chiara 2022 loaders built end to end on a synthetic release and genome tier.

The Peter 2018 genome tier, the SGD FASTA and the reference-gene tarball are the
Caudal 2024 synthetic fixtures (``test_caudal2024_synthetic``): isolates AAA, SACE_YAU
and ZZZ, six pangenome columns. The synthetic Supplementary Tables workbook names AAA,
SACE_YAU and QQQ; QQQ has no Peter genome, so it is refused and counted, and ZZZ is in
the panel but not in the release, so nothing is built for it.

Genotypes the fixtures imply (Caudal fixture header):

- AAA: one sequence variant (``Q0010``), the accessory ``2-EC1118_1F14_0012g`` present at
  copy number 2.5, the reference ``YAL005C`` absent.
- SACE_YAU: variants ``YAL001C`` and ``YBR001W``, the accessory ``5-contig_7`` present
  (copy number NA, so 1.0), the reference ``YCR001W`` absent.

Release rows (Supplementary Table 3, growth environments Glucose_2, Malic_Acid, NH4, YPD):

- AAA: KAc asci 22 / 51 % at 24 / 72 h with dyads 0 / 15.69 %; water sporulation 0;
  water viability 71.1 / 56; SDC 114.5 / 61.8 / 33.2 (one viability above 100); CR 95 /
  82.2 / 62.2; doubling time 2.08 h (Glucose_2), 4.0 h (Malic_Acid), 1.35 h (YPD).
- SACE_YAU: KAc asci 0 at 24 h only (so no dyad value); water sporulation 12.5; no water
  viability; SDC 40 / 10 / 5; CR 90 / 50 / 30; doubling time 3.0 h (NH4), 1.5 h (YPD).
- QQQ: every life-cycle value and two growth values, all refused (no Peter genome).
"""

from __future__ import annotations

import hashlib
import io
import json
import pickle
from pathlib import Path
from typing import Any

import lmdb
import openpyxl
import pandas as pd
import pytest

from tests.torchcell.conftest import REAL_VERIFY_RAW_FILES
from tests.torchcell.datasets.scerevisiae import test_caudal2024_synthetic as caudal
from tests.torchcell.sequence._peter_tier import gff
from torchcell.data import RawSha256MismatchError
from torchcell.datamodels.media import PBS, SPORULATION_KAC, YPD_AGAR
from torchcell.datamodels.schema import (
    Concentration,
    ConcentrationUnit,
    EnvironmentPhysicalPerturbation,
    LifeCycleAssay,
    MeasurementType,
    PhysicalFactor,
    PreCultureSource,
    SampleUnit,
)
from torchcell.datasets.scerevisiae import dechiara2022 as m
from torchcell.literature.manifest import Manifest
from torchcell.sequence import natural_isolate as ni
from torchcell.sequence.genome.registry import PETER2018_1011, SGD_S288C_R64
from torchcell.sequence.natural_isolate import SGD_FSA_NAME
from torchcell.verification.report import Provenance
from torchcell.verification.sourced import SourcedValue

LIFE = [
    "Sporulation in water",
    "Asci at 24h",
    "Dyads at 24h",
    "Asci at 72h",
    "Dyads at 72h",
    "CLS in water (day 8)",
    "CLS in water (day 20)",
    "CLS in SDC (day 7)",
    "CLS in SDC (day 21)",
    "CLS in SDC (day 35)",
    "CLS in crSDC (day 7)",
    "CLS in crSDC (day 21)",
    "CLS in crSDC (day 35)",
]
STEMS = ["Glucose_2", "Malic_Acid", "NH4", "YPD"]
NA = None
#: name -> 13 life-cycle values in LIFE order, then {stem: (dt hours, yield cells)}.
ROWS: dict[str, tuple[list[float | None], dict[str, tuple[float, float]]]] = {
    "AAA": (
        [
            0.0,
            22.0,
            0.0,
            51.0,
            15.68627451,
            71.1,
            56.0,
            114.5,
            61.8,
            33.2,
            95.0,
            82.2,
            62.2,
        ],
        {"Glucose_2": (2.08, 1.07e7), "Malic_Acid": (4.0, 3.0e5), "YPD": (1.35, 1.3e7)},
    ),
    "SACE_YAU": (
        [12.5, 0.0, NA, NA, NA, NA, NA, 40.0, 10.0, 5.0, 90.0, 50.0, 30.0],
        {"NH4": (3.0, 2.0e6), "YPD": (1.5, 9.0e6)},
    ),
    "QQQ": (
        [1.0, 10.0, 5.0, 20.0, 5.0, 50.0, 40.0, 60.0, 30.0, 20.0, 80.0, 60.0, 40.0],
        {"Glucose_2": (2.5, 1.0e7), "YPD": (1.4, 1.2e7)},
    ),
}
ST10_LIFE = [m.ST10_LIFE_CYCLE_NAME[c] for c in LIFE]


def _workbook(rows: dict[str, Any] = ROWS, st4_extra: bool = False) -> bytes:
    """The three sheets the loader reads, laid out as the release lays them out."""
    book = openpyxl.Workbook()
    st3 = book.active
    assert st3 is not None
    st3.title = m.ST3_SHEET
    header = [
        "Tree order",
        "Standard name",
        "Clade",
        "Clade_class",
        *LIFE,
        *(f"{s}_doublingtime" for s in STEMS),
        *(f"{s}_yield" for s in STEMS),
        *(f"{s}_doublingtime_hours" for s in STEMS),
        *(f"{s}_yield_cellscount" for s in STEMS),
    ]
    st3.append(["Supplementary Table 3. synthetic"])
    st3.append(["", "", "", "", "Sporulation"])
    st3.append(header)
    for order, (name, (life, growth)) in enumerate(rows.items(), start=1):
        dt = [growth[s][0] if s in growth else "NA" for s in STEMS]
        yields = [growth[s][1] if s in growth else "NA" for s in STEMS]
        st3.append(
            [
                order,
                name,
                "01.Wine_European",
                "dom",
                *("NA" if v is None else v for v in life),
                *(0.0 if v != "NA" else "NA" for v in dt),
                *(0.0 if v != "NA" else "NA" for v in yields),
                *dt,
                *yields,
            ]
        )
    st4 = book.create_sheet(m.ST4_SHEET)
    st4.append(["Supplementary Table 4. synthetic"])
    st4.append([])
    st4.append(["Source", "Concentration", "Element"])
    st4.append(["Glucose", 2, "Carbon"])
    st4.append(["Malic_Acid", None, "Carbon"])
    st4.append(["NH4", None, "Nitrogen"])
    if st4_extra:
        st4.append(["Sucrose", 2, "Carbon"])
    st10 = book.create_sheet(m.ST10_SHEET)
    st10.append(["Supplementary Table 10. synthetic"])
    st10.append([])
    st10.append(["Phenotype", "Total n. of isolates"])
    frame = pd.DataFrame(
        {c: [r[0][i] for r in rows.values()] for i, c in enumerate(LIFE)}
    )
    for column, name in zip(LIFE, ST10_LIFE, strict=True):
        st10.append([name, int(frame[column].notna().sum())])
    for stem in STEMS:
        n = sum(1 for _, growth in rows.values() if stem in growth)
        st10.append([f"{stem}_doublingtime_hours", n])
        st10.append([f"{stem}_yield_cellscount", n])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


WORKBOOK = _workbook()
WORKBOOK_SHA256 = hashlib.sha256(WORKBOOK).hexdigest()
ORF_FASTA = ">YAL001C TFC3\nA\n>YBR001W NTH2\nA\n>Q0010 COX1\nA\n"


def _write_tiers(data_root: Path) -> None:
    raw = caudal._raw_bytes()
    caudal._write_tier(
        data_root,
        SGD_S288C_R64,
        {
            SGD_FSA_NAME: caudal._SGD_FSA.encode(),
            ni.SGD_GFF_NAME: gff([]),
            "orf_coding_all_R64-4-1_20230830.fasta": ORF_FASTA.encode(),
            "rna_coding_R64-4-1_20230830.fasta": b">YNCA0001W tA\nA\n",
        },
    )
    caudal._write_tier(
        data_root, PETER2018_1011, {name: raw[name] for name in ni.GENOTYPE_FILES}
    )


@pytest.fixture
def data_root(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(root))
    monkeypatch.setattr(m, "TABLES_SHA256", WORKBOOK_SHA256)
    _write_tiers(root)
    return root


def _store(data_root: Path, family: m.StoreFamily) -> Path:
    """A store root holding the synthetic raw files, so ``download`` never runs."""
    root = data_root / "data" / "torchcell" / family.value
    raw = root / "raw"
    raw.mkdir(parents=True)
    (raw / m.TABLES_FILENAME).write_bytes(WORKBOOK)
    for name, data in caudal._raw_bytes().items():
        if name in m.RAW_FILES:
            (raw / name).write_bytes(data)
    return root


def _raw_records(root: Path) -> list[dict[str, Any]]:
    env = lmdb.open(str(root / "processed" / "lmdb"), readonly=True, lock=False)
    with env.begin() as txn:
        out = [pickle.loads(v) for _, v in txn.cursor()]
    env.close()
    return out


# --------------------------------------------------------------------------- #
# Pure parts
# --------------------------------------------------------------------------- #
def test_growth_stems_split_into_source_and_concentration() -> None:
    kinds = {"Glucose": m.NutrientKind.carbon, "Malic_Acid": m.NutrientKind.carbon}
    assert m.parse_growth_stem("Glucose_0.5", kinds).model_dump() == {
        "stem": "Glucose_0.5",
        "source": "Glucose",
        "concentration_percent": 0.5,
        "kind": m.NutrientKind.carbon,
    }
    assert m.parse_growth_stem("Malic_Acid", kinds).concentration_percent is None
    assert m.parse_growth_stem("YPD", kinds).kind is m.NutrientKind.rich
    with pytest.raises(ValueError, match="names no Supplementary Table 4 source"):
        m.parse_growth_stem("Lactate_2", kinds)


def test_a_carbon_environment_doses_the_source_on_the_glucose_free_base() -> None:
    env = m.growth_environment(
        m.GrowthEnvironment(
            stem="Glucose_2",
            source="Glucose",
            concentration_percent=2.0,
            kind=m.NutrientKind.carbon,
        )
    )
    assert env.media is m.GROWTH_CARBON_BASE
    assert env.temperature is not None and env.temperature.value == 30.0
    (edit,) = env.perturbations
    assert isinstance(edit, EnvironmentPhysicalPerturbation)
    assert (edit.factor, edit.magnitude, edit.agent and edit.agent.name) == (
        PhysicalFactor.carbon_source,
        Concentration(value=2.0, unit=ConcentrationUnit.percent_w_v),
        "D-glucose",
    )
    assert "D-glucose" not in {c.compound.name for c in m.GROWTH_CARBON_BASE.components}


def test_malic_acid_states_no_dose_so_its_magnitude_is_a_typed_gap() -> None:
    env = m.growth_environment(
        m.GrowthEnvironment(
            stem="Malic_Acid",
            source="Malic_Acid",
            concentration_percent=None,
            kind=m.NutrientKind.carbon,
        )
    )
    (edit,) = env.perturbations
    assert isinstance(edit, EnvironmentPhysicalPerturbation)
    assert edit.magnitude is None
    assert [(g.field, g.reason.value) for g in edit.provenance_gaps] == [
        ("magnitude", "not_reported_by_primary")
    ]
    assert edit.agent is not None and edit.agent.name == "malic acid"


def test_a_nitrogen_environment_doses_30_mg_nitrogen_per_litre() -> None:
    env = m.growth_environment(
        m.GrowthEnvironment(
            stem="NH4",
            source="NH4",
            concentration_percent=None,
            kind=m.NutrientKind.nitrogen,
        )
    )
    assert env.media is m.GROWTH_NITROGEN_BASE
    (edit,) = env.perturbations
    assert isinstance(edit, EnvironmentPhysicalPerturbation)
    assert (edit.factor, edit.magnitude, edit.agent and edit.agent.name) == (
        PhysicalFactor.nitrogen_source,
        Concentration(value=30.0, unit=ConcentrationUnit.mg_nitrogen_per_l),
        "ammonium sulfate",
    )
    names = {c.compound.name for c in m.GROWTH_NITROGEN_BASE.components}
    assert "ammonium sulfate" not in names and "D-glucose" in names


def test_the_ypd_environment_is_the_library_ypd_agar_restated() -> None:
    env = m.growth_environment(
        m.GrowthEnvironment(
            stem="YPD",
            source="YPD",
            concentration_percent=None,
            kind=m.NutrientKind.rich,
        )
    )
    assert env.perturbations == []
    assert env.media.components == YPD_AGAR.components
    assert env.media.provenance[-1] is m.YPD_RECIPE


def test_the_life_cycle_protocols_state_their_culture() -> None:
    kac = m.protocol_environment(m.LifeCycleProtocol.kac_sporulation)
    assert kac.media.components == SPORULATION_KAC.components
    assert kac.temperature is not None and kac.temperature.value == 23.0
    assert kac.culture_format is not None
    assert (
        kac.culture_format.vessel,
        kac.culture_format.working_volume_ul,
        kac.culture_format.shaking_rpm,
        kac.culture_format.inoculum_od600,
    ) == ("250 ml flask", 50_000.0, 250.0, 0.6)
    assert kac.pre_culture is not None
    assert (kac.pre_culture.source, kac.pre_culture.duration_hours) == (
        PreCultureSource.timed_culture,
        48.0,
    )
    water = m.protocol_environment(m.LifeCycleProtocol.water_sporulation)
    assert water.media.components == PBS.components
    assert water.temperature is None
    assert [g.field for g in water.provenance_gaps] == ["temperature"]
    viability = m.protocol_environment(m.LifeCycleProtocol.water_viability)
    assert viability.temperature is not None and viability.temperature.value == 30.0
    cr = m.protocol_environment(m.LifeCycleProtocol.calorie_restricted_viability)
    glucose = next(c for c in cr.media.components if c.compound.name == "D-glucose")
    assert glucose.concentration == Concentration(
        value=0.5, unit=ConcentrationUnit.percent_w_v
    )
    assert cr.pre_culture is not None
    assert cr.pre_culture.source is PreCultureSource.overnight_culture


def test_protocol_phenotypes_carry_the_sampling_the_source_states() -> None:
    kac = m.protocol_phenotype(
        m.LifeCycleProtocol.kac_sporulation, {"sporulation_efficiency@24h": 22.0}
    )
    assert (kac.assay, kac.cells_per_sample, kac.cells_per_sample_is_minimum) == (
        LifeCycleAssay.microscopy_ascus_count,
        100,
        True,
    )
    assert sorted(g.field for g in kac.provenance_gaps) == ["n_samples", "sample_unit"]
    cls = m.protocol_phenotype(
        m.LifeCycleProtocol.sdc_viability, {"stationary_phase_viability@7d": 50.0}
    )
    assert (cls.n_samples, cls.sample_unit, cls.cells_per_sample) == (
        1,
        SampleUnit.biological_replicate,
        10_000,
    )
    assert cls.provenance_gaps == []


def test_life_cycle_columns_cover_the_release_with_canonical_keys() -> None:
    assert [c.column for c in m.LIFE_CYCLE_COLUMNS] == [
        "Asci at 24h",
        "Dyads at 24h",
        "Asci at 72h",
        "Dyads at 72h",
        "Sporulation in water",
        "CLS in water (day 8)",
        "CLS in water (day 20)",
        "CLS in SDC (day 7)",
        "CLS in SDC (day 21)",
        "CLS in SDC (day 35)",
        "CLS in crSDC (day 7)",
        "CLS in crSDC (day 21)",
        "CLS in crSDC (day 35)",
    ]
    assert [c.key for c in m.LIFE_CYCLE_COLUMNS[:5]] == [
        "sporulation_efficiency@24h",
        "dyad_fraction@24h",
        "sporulation_efficiency@72h",
        "dyad_fraction@72h",
        "sporulation_efficiency@8d",
    ]


def test_read_release_parses_and_cross_checks_the_three_tables(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    path.write_bytes(WORKBOOK)
    release = m.read_release(str(path))
    assert [e.stem for e in release.environments] == STEMS
    assert release.st10_counts["Asci_at_24h"] == 3
    assert m.growth_values(release, m.GrowthTrait.doubling_time, ["SACE_YAU"]) == {
        ("SACE_YAU", "NH4"): 3.0,
        ("SACE_YAU", "YPD"): 1.5,
    }
    assert m.life_cycle_values(release, ["SACE_YAU"]) == {
        ("SACE_YAU", m.LifeCycleProtocol.kac_sporulation): {
            "sporulation_efficiency@24h": 0.0
        },
        ("SACE_YAU", m.LifeCycleProtocol.water_sporulation): {
            "sporulation_efficiency@8d": 12.5
        },
        ("SACE_YAU", m.LifeCycleProtocol.sdc_viability): {
            "stationary_phase_viability@7d": 40.0,
            "stationary_phase_viability@21d": 10.0,
            "stationary_phase_viability@35d": 5.0,
        },
        ("SACE_YAU", m.LifeCycleProtocol.calorie_restricted_viability): {
            "stationary_phase_viability@7d": 90.0,
            "stationary_phase_viability@21d": 50.0,
            "stationary_phase_viability@35d": 30.0,
        },
    }


def test_read_release_refuses_a_table_4_mismatch_and_a_dyad_without_asci(
    tmp_path: Path,
) -> None:
    path = tmp_path / "book.xlsx"
    path.write_bytes(_workbook(st4_extra=True))
    with pytest.raises(ValueError, match="growth stems disagree"):
        m.read_release(str(path))
    life, growth = ROWS["SACE_YAU"]
    broken = dict(ROWS)
    broken["SACE_YAU"] = ([life[0], 0.0, 3.0, *life[3:]], growth)
    path.write_bytes(_workbook(broken))
    with pytest.raises(ValueError, match="Dyads at 24h is not present exactly"):
        m.read_release(str(path))


def test_deposit_copies_the_workbook_and_records_its_retrieval(data_root: Path) -> None:
    library = data_root / "torchcell-library" / m.CITATION_KEY / "si"
    library.mkdir(parents=True)
    (library / "si4.xlsx").write_bytes(WORKBOOK)
    out = Path(m.deposit_raw_mirror(data_root=str(data_root)))
    assert (out / m.TABLES_REL).read_bytes() == WORKBOOK
    manifest = Manifest.model_validate_json((out / "manifest.json").read_text())
    (record,) = manifest.files
    assert (record.path, record.sha256, record.original_filename) == (
        m.TABLES_REL,
        WORKBOOK_SHA256,
        m.TABLES_FILENAME,
    )
    assert record.retrieval is not None
    assert (record.retrieval.method.value, record.retrieval.source_url) == (
        "springer_esm",
        m.TABLES_URL,
    )
    (out / m.TABLES_REL).write_bytes(b"drift")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        m.deposit_raw_mirror(data_root=str(data_root))


def test_deposit_refuses_a_library_file_off_its_pin(
    data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    library = data_root / "torchcell-library" / m.CITATION_KEY / "si"
    library.mkdir(parents=True)
    (library / "si4.xlsx").write_bytes(WORKBOOK)
    monkeypatch.setattr(m, "TABLES_SHA256", "0" * 64)
    with pytest.raises(RawSha256MismatchError):
        m.deposit_raw_mirror(data_root=str(data_root))
    assert not (data_root / m.RAW_MIRROR_REL).exists()


# --------------------------------------------------------------------------- #
# End-to-end builds
# --------------------------------------------------------------------------- #
def _genotype_summary(genotype: dict[str, Any]) -> list[tuple[str, str]]:
    return sorted(
        (p["perturbation_type"], p["systematic_gene_name"])
        for p in genotype["perturbations"]
    )


AAA_GENOTYPE = [
    ("natural_gene_absence", "YAL005C"),
    ("natural_gene_presence", "2-EC1118_1F14_0012g"),
    ("sequence_variant", "Q0010"),
]
SACE_YAU_GENOTYPE = [
    ("natural_gene_absence", "YCR001W"),
    ("natural_gene_presence", "5-contig_7"),
    ("sequence_variant", "YAL001C"),
    ("sequence_variant", "YBR001W"),
]


@pytest.fixture
def life_cycle(data_root: Path) -> m.LifeCycleDeChiara2022Dataset:
    root = _store(data_root, m.StoreFamily.life_cycle)
    return m.LifeCycleDeChiara2022Dataset(root=str(root))


def test_life_cycle_build_writes_one_record_per_isolate_protocol(
    life_cycle: m.LifeCycleDeChiara2022Dataset,
) -> None:
    """9 records: AAA has all five protocols, SACE_YAU four (no water viability)."""
    records = [life_cycle[i] for i in range(len(life_cycle))]
    assert len(records) == 9
    by_isolate: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        strain = record["experiment"]["genotype"]["perturbations"][0]["strain_id"]
        by_isolate.setdefault(strain, []).append(record)
    assert sorted(by_isolate) == ["AAA", "SACE_YAU"]
    assert [r["experiment"]["phenotype"]["life_cycle"] for r in by_isolate["AAA"]] == [
        {
            "dyad_fraction@24h": 0.0,
            "dyad_fraction@72h": 15.68627451,
            "sporulation_efficiency@24h": 22.0,
            "sporulation_efficiency@72h": 51.0,
        },
        {"sporulation_efficiency@8d": 0.0},
        {"stationary_phase_viability@20d": 56.0, "stationary_phase_viability@8d": 71.1},
        {
            "stationary_phase_viability@21d": 61.8,
            "stationary_phase_viability@35d": 33.2,
            "stationary_phase_viability@7d": 114.5,
        },
        {
            "stationary_phase_viability@21d": 82.2,
            "stationary_phase_viability@35d": 62.2,
            "stationary_phase_viability@7d": 95.0,
        },
    ]
    assert by_isolate["SACE_YAU"][0]["experiment"]["phenotype"]["life_cycle"] == {
        "sporulation_efficiency@24h": 0.0
    }
    assert _genotype_summary(by_isolate["AAA"][0]["experiment"]["genotype"]) == (
        AAA_GENOTYPE
    )
    assert _genotype_summary(by_isolate["SACE_YAU"][0]["experiment"]["genotype"]) == (
        SACE_YAU_GENOTYPE
    )
    kac_reference = by_isolate["AAA"][0]["reference"]["phenotype_reference"]
    assert kac_reference["life_cycle"] == {
        "dyad_fraction@24h": 0.0,
        "dyad_fraction@72h": 15.68627451,
        "sporulation_efficiency@24h": 11.0,
        "sporulation_efficiency@72h": 51.0,
    }


def test_the_genotype_is_interned_once_per_isolate(
    life_cycle: m.LifeCycleDeChiara2022Dataset,
) -> None:
    raw = _raw_records(Path(life_cycle.root))
    refs = [r["experiment"]["genotype"]["$ref"] for r in raw]
    assert len(set(refs)) == 2
    assert [r["experiment"]["genotype"]["name"] for r in raw[:1]] == ["genotype AAA"]
    assert refs.count(refs[0]) == 5


def test_the_life_cycle_ledger_counts_the_refused_isolate(
    life_cycle: m.LifeCycleDeChiara2022Dataset,
) -> None:
    ledger = json.loads(
        (Path(life_cycle.preprocess_dir) / "release_ledger.json").read_text()
    )
    assert {k: ledger[k] for k in ledger if k != "refused_blocks"} == {
        "dataset": "LifeCycleDeChiara2022Dataset",
        "n_isolates_in_release": 3,
        "n_isolates_without_peter_genome": 1,
        "isolates_without_peter_genome": ["QQQ"],
        "n_isolates_built": 2,
        "n_released_cells": 39,
        "n_missing_cells": 5,
        "n_refused_cells": 13,
        "n_kept_cells": 21,
        "n_records": 9,
    }


def _audited_library(data_root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """One synthetic sourced value whose quote is in a synthetic mirror file."""
    text = b"Supplementary Table 3 carries the values.\n"
    folder = data_root / "torchcell-library" / m.CITATION_KEY
    folder.mkdir(parents=True)
    (folder / "paper.md").write_bytes(text)
    sourced = SourcedValue(
        value=1,
        quote="carries the values",
        provenance=Provenance(
            source_uri="paper.md",
            citation_key=m.CITATION_KEY,
            sha256=hashlib.sha256(text).hexdigest(),
        ),
    )
    monkeypatch.setattr(m, "SOURCED_VALUES", (sourced,))
    # the isolate-genotype audit re-reads Peter's Table S1 from the library mirror
    tables = data_root / "torchcell-library" / ni.CITATION_KEY / ni.LIBRARY_TABLES_REL
    tables.parent.mkdir(parents=True)
    tables.write_bytes(caudal._WORKBOOK)
    monkeypatch.setattr(
        ni, "TABLES_SHA256", hashlib.sha256(caudal._WORKBOOK).hexdigest()
    )


def test_verify_build_passes_the_synthetic_life_cycle_store(
    life_cycle: m.LifeCycleDeChiara2022Dataset,
    data_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _audited_library(data_root, monkeypatch)
    # CI runs the whole suite in one process: release the fixture's LMDB handle before
    # verify_build opens the same store with its own reader.
    life_cycle.close_lmdb()
    report = m.verify_build(
        life_cycle.root, data_root=str(data_root), expected_viability_above_100=1
    )
    results = {r.name: r for r in report.results}
    assert report.passed, report.summary()
    assert results["count"].details == {"observed": 9, "expected": 9}
    assert results["viability_above_100"].details == {"observed": 1, "declared": 1}
    assert results["sequence_variant_genes_are_s288c_genes"].details["n_genes"] == 3
    assert results["isolate_genotype_round_trip"].message == (
        "2 stored genotypes equal their re-derivation from the pinned Peter 2018 "
        "releases"
    )
    wrong = m.verify_build(
        life_cycle.root, data_root=str(data_root), expected_viability_above_100=0
    )
    assert not wrong.passed
    assert [r.name for r in wrong.results if not r.passed] == ["viability_above_100"]


@pytest.fixture
def doubling_time(data_root: Path) -> m.DoublingTimeDeChiara2022Dataset:
    root = _store(data_root, m.StoreFamily.doubling_time)
    return m.DoublingTimeDeChiara2022Dataset(root=str(root))


def test_doubling_time_build_writes_one_record_per_released_cell(
    doubling_time: m.DoublingTimeDeChiara2022Dataset,
    data_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    records = [doubling_time[i] for i in range(len(doubling_time))]
    rows = [
        (
            r["experiment"]["genotype"]["perturbations"][0]["strain_id"],
            r["experiment"]["environment"]["media"]["name"],
            r["experiment"]["phenotype"]["environment_response"],
            r["reference"]["phenotype_reference"]["environment_response"],
        )
        for r in records
    ]
    assert rows == [
        ("AAA", m.GROWTH_CARBON_BASE.name, 2.08, 2.08),
        ("AAA", m.GROWTH_CARBON_BASE.name, 4.0, 4.0),
        ("AAA", YPD_AGAR.name, 1.35, (1.35 + 1.5) / 2),
        ("SACE_YAU", m.GROWTH_NITROGEN_BASE.name, 3.0, 3.0),
        ("SACE_YAU", YPD_AGAR.name, 1.5, (1.35 + 1.5) / 2),
    ]
    phenotype = records[0]["experiment"]["phenotype"]
    assert (
        phenotype["measurement_type"],
        phenotype["n_samples"],
        phenotype["sample_unit"],
        phenotype["units"],
    ) == (
        MeasurementType.growth_rate,
        4,
        SampleUnit.colony,
        m.DoublingTimeDeChiara2022Dataset.units,
    )
    ledger = json.loads(
        (Path(doubling_time.preprocess_dir) / "release_ledger.json").read_text()
    )
    assert (
        ledger["n_records"],
        ledger["n_missing_cells"],
        ledger["n_refused_cells"],
        ledger["n_kept_cells"],
    ) == (5, 5, 2, 5)
    assert ledger["refused_blocks"][0]["n_cells"] == 7
    _audited_library(data_root, monkeypatch)
    doubling_time.close_lmdb()
    report = m.verify_build(doubling_time.root, data_root=str(data_root))
    assert report.passed, report.summary()
    assert {r.name for r in report.results} >= {
        "released_value_fidelity",
        "positive_values",
        "measurement_type_consistent",
        "reference_is_population_mean",
        "media_membership",
    }


def test_cell_yield_build_types_its_values_as_cell_yield(data_root: Path) -> None:
    root = _store(data_root, m.StoreFamily.cell_yield)
    dataset = m.CellYieldDeChiara2022Dataset(root=str(root))
    assert len(dataset) == 5
    phenotype = dataset[0]["experiment"]["phenotype"]
    assert (phenotype["measurement_type"], phenotype["environment_response"]) == (
        MeasurementType.cell_yield,
        1.07e7,
    )


def test_a_workbook_off_its_pin_is_refused_before_any_record(
    data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _store(data_root, m.StoreFamily.life_cycle)
    monkeypatch.setattr(m, "TABLES_SHA256", "f" * 64)
    # tests/torchcell/conftest.py swaps in a presence-only recorder; the refusal needs
    # the real byte check back.
    monkeypatch.setattr(m, "verify_raw_files", REAL_VERIFY_RAW_FILES)
    with pytest.raises(RawSha256MismatchError, match="expected f{64}"):
        m.LifeCycleDeChiara2022Dataset(root=str(root))
    assert not (root / "processed" / "lmdb").exists()
