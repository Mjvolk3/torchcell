# tests/torchcell/datasets/pputida/test_menasalvas2025.py
# [[tests.torchcell.datasets.pputida.test_menasalvas2025]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/pputida/test_menasalvas2025.py
"""Tests for the Menasalvas 2025 biosensor-coupled CRISPRi selection loader.

Everything here runs with NO network and NO ``$DATA_ROOT``: the Supplementary
Information markdown is synthesized to the shape the pinned OCR has, the two deposited
Dryad workbooks are synthesized with openpyxl into a synthetic inner zip inside a
synthetic arrival zip, the raw mirror is deposited under ``tmp_path`` by the module's
own ``deposit_raw_mirror`` with its pins monkeypatched to the synthetic digests, and the
KT2440 annotation is the real genome class over a synthetic assembly. All four families
build and verify hermetically.

A separate ``@pytest.mark.data`` block, skipped without the mirror and the built stores,
pins the numbers measured on the REAL bytes: 58 selection records over 60 distinct
targets from 28/30 rows, a 170-member inner zip holding the two workbooks, 48 Absolute
and 10 Relative metabolite rows, 39,438 proteome rows over 2,187 protein keys, 584
breseq calls, the growth-phase harvest OD600 of 2.4 / 1.2 / 2.7, and PP_2088's
production folds of 31.895 and 18.555 against Supplementary Note 2's 34 and 20.
"""

from __future__ import annotations

import hashlib
import json
import os.path as osp
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.pputida.menasalvas2025 as mv
from torchcell.datamodels.media import M9_NREL_MOPS_MENASALVAS2025
from torchcell.datamodels.schema import (
    BacterialEnvironmentResponseExperiment,
    BacterialEnvironmentResponseExperimentReference,
    BacterialVariantType,
    SmallMoleculePerturbation,
)
from torchcell.datasets.bacteria_common import reconcile_locus_tags

# --------------------------------------------------------------------------- #
# Synthetic Supplementary Information markdown
# --------------------------------------------------------------------------- #
#: Loci Supplementary Note 2 names as measured; every one must key every stored record.
NOTE2_LOCI: tuple[str, ...] = mv.NOTE2_MEASURED_LOCI
#: Deletions the three designed strains carry, beyond those the selection tables name.
DESIGNED_DELETION_LOCI: tuple[str, ...] = ("PP_2664", "PP_2675", "PP_3540", "PP_4622")
#: Host protein keys that exist only to keep the resolved fraction above the loader's
#: 0.95 floor beside the one retired key, and to carry the released-0 and the
#: single-replicate protein.
FILLER_PROTEIN_LOCI: tuple[str, ...] = (
    "PP_0100",
    "PP_0200",
    "PP_0300",
    "PP_0400",
    "PP_0500",
    "PP_0600",
    "PP_0700",
    "PP_0800",
)
#: Locus tags the synthetic assembly annotates, with the symbol it gives each one.
LOCUS_SPECS: tuple[tuple[str, str | None], ...] = (
    ("PP_1815", "pyrF"),
    ("PP_2428", "sotB"),
    ("PP_3400", "alkB"),
    ("PP_4373", "fleQ"),
    ("PP_5003", "phaA"),
    ("PP_5004", "phaB"),
    ("PP_5005", "phaC-II"),
    ("PP_5007", None),
    ("PP_1656", "relA"),
    ("PP_2710", None),
    ("PP_4485", "hisQ"),
    ("PP_2074", None),
    # The deposited arms: the designed strains' deletions, Note 2's measured loci, the
    # merged-key locus, the one locus data S1-5 names by symbol, and the filler keys.
    *((tag, None) for tag in DESIGNED_DELETION_LOCI),
    *((tag, None) for tag in NOTE2_LOCI),
    ("PP_1000", None),
    ("PP_1100", "cadA-I"),
    *((tag, None) for tag in FILLER_PROTEIN_LOCI),
)
#: The ``Gene`` cells of the synthetic round-1 table, one per expected row.
ROUND1_CELLS: tuple[str, ...] = (
    "phaAZC-II / PP_5003-PP_5005",
    "PP_5007",
    "alkB /PP_3400",
    "sotB / PP_2428",
    "fleQ / PP_4373",
)
#: The ``Gene`` cells of the synthetic round-2 table; no tag may repeat round 1.
ROUND2_CELLS: tuple[str, ...] = ("relA PP_1656", "PP_2710", "hisQ |PP_4485", "PP_2074")

ASSEMBLY_REPORT = """# Assembly name:  ASM756v2
# Organism name:  Pseudomonas putida KT2440 (g-proteobacteria)
# Infraspecific name:  strain=KT2440
# Taxid:          160488
# GenBank assembly accession: GCA_000007565.2
# RefSeq assembly accession: GCF_000007565.2
# RefSeq assembly and GenBank assemblies identical: yes
#
## Assembly-Units:
AE015451.2\tassembled-molecule\tna\tChromosome\tAE015451.2\t=\tNC_002947.4
"""
ASSEMBLY_REPORT_MEMBER = "GCA_000007565.2_ASM756v2_assembly_report.txt"


def _table(number: int, title: str, cells: tuple[str, ...]) -> str:
    """One caption line plus the single-line ``<table>`` the OCR emits for it."""
    header = "".join(
        f'<td colspan="1" rowspan="1">{name}</td>' for name in mv.TABLE_HEADER
    )
    rows = [f"<tr>{header}</tr>"]
    for index, cell in enumerate(cells):
        body = "".join(
            f'<td colspan="1" rowspan="1">{value}</td>'
            for value in (cell, f"synthetic function {index}", "enzyme")
        )
        rows.append(f"<tr>{body}</tr>")
    body_html = "".join(rows)
    return f"Supplementary Table {number}. {title}\n\n<table>{body_html}</table>\n"


def synthetic_markdown(
    round1: tuple[str, ...] = ROUND1_CELLS,
    round2: tuple[str, ...] = ROUND2_CELLS,
    *,
    header: tuple[str, ...] | None = None,
) -> str:
    """The two tables, wrapped in enough surrounding prose to exercise the scanner."""
    text = "# Supplementary Materials for\n\nsome prose\n\n"
    text += _table(
        1, "Lower pedF-RBS-pyrF Isoprenol Threshold: Enriched gRNAs.", round1
    )
    text += "\nmore prose\n\n"
    text += _table(
        2,
        "Second Round pedF-RBS-pyrF with Higher Isoprenol Threshold: Enriched gRNAs.",
        round2,
    )
    text += "\n# Supplementary Table 4. Strains Used in This Study.\n\n<table><tr>"
    text += '<td colspan="1" rowspan="1">TEAM-862</td></tr></table>\n'
    if header is not None:
        text = text.replace(
            "".join(
                f'<td colspan="1" rowspan="1">{name}</td>' for name in mv.TABLE_HEADER
            ),
            "".join(f'<td colspan="1" rowspan="1">{name}</td>' for name in header),
            1,
        )
    return text


# --------------------------------------------------------------------------- #
# Synthetic: the Gene-cell parser
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("cell", "tags", "symbol"),
    [
        ("PP_5007", ("PP_5007",), None),
        ("sotB / PP_2428", ("PP_2428",), "sotB"),
        ("cmpX PP_2087", ("PP_2087",), "cmpX"),
        ("hisQ |PP_4485", ("PP_4485",), "hisQ"),
        ("relA PP_1656", ("PP_1656",), "relA"),
        ("alkB /PP_3400", ("PP_3400",), "alkB"),
        ("phaAZC-II / PP_5003-PP_5005", ("PP_5003", "PP_5004", "PP_5005"), "phaAZC-II"),
    ],
)
def test_gene_cells_parse_to_tags_and_the_sources_own_symbol(
    cell: str, tags: tuple[str, ...], symbol: str | None
) -> None:
    """Every spelling the released tables use resolves to its tags and its symbol."""
    assert mv.parse_gene_cell(cell) == (tags, symbol)


def test_an_operon_range_expands_by_integer_enumeration() -> None:
    """``PP_5003-PP_5005`` is three loci, not the two endpoints the regex would find."""
    tags, _ = mv.parse_gene_cell("phaAZC-II / PP_5003-PP_5005")
    assert tags == ("PP_5003", "PP_5004", "PP_5005")
    assert len(tags) == 3


def test_a_gene_cell_with_no_locus_tag_is_refused() -> None:
    """A cell naming only a symbol cannot be keyed to a locus, so the build stops."""
    with pytest.raises(RuntimeError, match="names no PP_ locus tag"):
        mv.parse_gene_cell("phaF")


def test_a_descending_locus_range_is_refused() -> None:
    """``PP_5005-PP_5003`` is not an ascending operon span."""
    with pytest.raises(RuntimeError, match="does not ascend"):
        mv.parse_gene_cell("x / PP_5005-PP_5003")


def test_a_gene_cell_repeating_a_tag_is_refused() -> None:
    """A tag named twice would silently collapse, so it is a refusal instead."""
    with pytest.raises(RuntimeError, match="repeats a locus tag"):
        mv.parse_gene_cell("PP_5003-PP_5005 PP_5004")


# --------------------------------------------------------------------------- #
# Synthetic: the table readers
# --------------------------------------------------------------------------- #
def test_read_table_returns_the_three_released_columns() -> None:
    """The reader keeps ``Gene``, ``Function`` and ``FunctionalCategory`` verbatim."""
    rows = mv.read_table(synthetic_markdown(), 1)
    assert len(rows) == len(ROUND1_CELLS)
    assert rows[0] == (ROUND1_CELLS[0], "synthetic function 0", "enzyme")


def test_read_table_picks_the_table_its_caption_names() -> None:
    """Table 2's rows are Table 2's, although three tables share the document."""
    rows = mv.read_table(synthetic_markdown(), 2)
    assert [row[0] for row in rows] == list(ROUND2_CELLS)


def test_a_changed_table_header_is_refused() -> None:
    """A renamed column means the export moved; parsing on would mis-key the data."""
    markdown = synthetic_markdown(header=("Locus", "Function", "FunctionalCategory"))
    with pytest.raises(RuntimeError, match="header is"):
        mv.read_table(markdown, 1)


def test_a_missing_table_is_refused() -> None:
    """A table the loader needs and the markdown lacks stops the build."""
    with pytest.raises(RuntimeError, match="Supplementary Table 9 is not in"):
        mv.read_table(synthetic_markdown(), 9)


def test_a_ragged_table_row_is_refused() -> None:
    """A row with a dropped cell would shift the columns, so it is a refusal."""
    markdown = synthetic_markdown().replace(
        '<td colspan="1" rowspan="1">enzyme</td></tr>', "</tr>", 1
    )
    with pytest.raises(RuntimeError, match="row has 2 cells"):
        mv.read_table(markdown, 1)


def test_read_selected_targets_checks_the_per_round_row_counts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A round whose row count is not the Methods' stated one is a refusal."""
    path = tmp_path / "si1.md"
    path.write_text(synthetic_markdown(), encoding="utf-8")
    monkeypatch.setattr(mv, "EXPECTED_ROWS", {1: 28, 2: 30})
    with pytest.raises(RuntimeError, match="the Methods state 28"):
        mv.read_selected_targets(str(path))


def test_read_selected_targets_refuses_a_target_shared_by_the_two_rounds(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The paper states the rounds do not overlap, so an overlap is a refusal."""
    path = tmp_path / "si1.md"
    overlapping = ROUND2_CELLS[:-1] + ("sotB / PP_2428",)
    path.write_text(synthetic_markdown(round2=overlapping), encoding="utf-8")
    monkeypatch.setattr(
        mv, "EXPECTED_ROWS", {1: len(ROUND1_CELLS), 2: len(overlapping)}
    )
    with pytest.raises(RuntimeError, match="the two rounds share"):
        mv.read_selected_targets(str(path))


def test_read_selected_targets_tags_each_row_with_its_round(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Round membership is what keeps the two screens apart, so it is stored."""
    path = tmp_path / "si1.md"
    path.write_text(synthetic_markdown(), encoding="utf-8")
    monkeypatch.setattr(
        mv, "EXPECTED_ROWS", {1: len(ROUND1_CELLS), 2: len(ROUND2_CELLS)}
    )
    targets = mv.read_selected_targets(str(path))
    assert len(targets) == len(ROUND1_CELLS) + len(ROUND2_CELLS)
    assert [t.round_number for t in targets] == [1] * len(ROUND1_CELLS) + [2] * len(
        ROUND2_CELLS
    )
    assert targets[0].locus_tags == ("PP_5003", "PP_5004", "PP_5005")


# --------------------------------------------------------------------------- #
# Synthetic: the genotype
# --------------------------------------------------------------------------- #
def test_the_host_background_is_exactly_the_declared_seven_perturbations() -> None:
    """Every record carries the pyrF lesion, the five pathway genes and the reporter."""
    host = mv.host_perturbations()
    assert len(host) == 7
    assert {p.systematic_gene_name for p in host} == set(mv.HOST_BACKGROUND_NAMES)
    kinds = {p.perturbation_type for p in host}
    assert kinds == {"bacterial_deletion", "heterologous_pathway"}


def test_the_host_background_refuses_to_drift_from_its_declared_name_set(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The verifier's ``background_genes`` must stay the set the genotype builds."""
    monkeypatch.setattr(mv, "HOST_BACKGROUND_NAMES", frozenset({"PP_1815"}))
    with pytest.raises(RuntimeError, match="the host background is"):
        mv.host_perturbations()


def test_the_pyrf_deletion_is_the_strain_tables_own_locus() -> None:
    """``PP_1815`` comes from "Pp KT2440 ΔPP_1815/pyrF", not from a symbol lookup."""
    deletion = mv.pyrf_deletion()
    assert deletion.systematic_gene_name == "PP_1815"
    assert deletion.perturbed_gene_name == "pyrF"
    assert deletion.gene_namespace == "pputida_kt2440_locus_tag"
    assert mv.PYRF_DELETION.quote == "Pp KT2440 ΔPP_1815/pyrF"


def test_the_five_pathway_genes_are_integrated_with_an_unreported_origin() -> None:
    """The pathway is chromosomal here, and its source organisms are not released."""
    pathway = mv.pathway_perturbations()
    assert [p.systematic_gene_name for p in pathway] == [
        "mvaS",
        "mvaE",
        "MKmm",
        "PMDHKQ",
        "aphA",
    ]
    assert {p.localization for p in pathway} == {"chromosomal_integration"}
    assert {p.integration_locus for p in pathway} == {
        "PP_5322intergenic",
        "PP_0871intergenic",
    }
    assert {p.promoter_name for p in pathway} == {"Pcv", "Ptrc"}
    assert {p.source_organism for p in pathway} == {mv.SOURCE_ORGANISM_UNREPORTED}
    assert {p.pathway_name for p in pathway} == {"isoprenol pathway"}


def test_a_pathway_token_absent_from_the_quoted_plasmid_is_refused(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A stored identifier must be a substring of the quote that justifies it."""
    monkeypatch.setattr(
        mv,
        "PATHWAY_GENES",
        ({"token": "nudB", "promoter": "Ptrc", "locus": "PP_0871intergenic"},),
    )
    with pytest.raises(RuntimeError, match="is not in the quoted pIY670 description"):
        mv.pathway_perturbations()


def test_the_reporter_is_the_teredinibacter_pyrf_homolog_on_a_plasmid() -> None:
    """The growth-coupled reporter is a heterologous addition, not the native gene."""
    reporter = mv.reporter_perturbation()
    assert reporter.systematic_gene_name == "TERTU_1389"
    assert reporter.source_organism == "Teredinibacter turnerae T7901"
    assert reporter.localization == "episomal_plasmid"
    assert reporter.promoter_name == "PpedF"


def test_a_knockdown_carries_three_designed_guides_and_no_spacer() -> None:
    """The library designs three guides per gene and released no enriched spacer."""
    knockdown = mv.crispri_perturbation("PP_2428", "sotB")
    assert knockdown.perturbation_type == "bacterial_crispr_interference"
    assert knockdown.crispr is not None
    assert knockdown.crispr.effector == "dCpf1/dCas12a"
    assert knockdown.crispr.n_guides == 3
    assert knockdown.crispr.guide_sequence is None
    assert knockdown.expression_direction == "decreased"


# --------------------------------------------------------------------------- #
# Synthetic: the environment and the phenotype
# --------------------------------------------------------------------------- #
def _compound_names(environment: Any) -> set[str]:
    """Names of the compounds one environment's small-molecule edits carry."""
    return {
        edit.compound.name
        for edit in environment.perturbations
        if isinstance(edit, SmallMoleculePerturbation)
    }


def test_the_selection_environment_is_induced_and_the_control_is_not() -> None:
    """Crystal violet is what makes the selection arm isoprenol-producing."""
    induced = mv._environment(induced=True)
    control = mv._environment(induced=False)
    names = _compound_names(induced)
    assert "crystal violet" in names
    control_names = _compound_names(control)
    assert "crystal violet" not in control_names
    assert "kanamycin" in control_names
    assert induced.duration_hours == 24.0
    assert induced.media.name == M9_NREL_MOPS_MENASALVAS2025.name


def test_the_selection_temperature_is_a_typed_absence() -> None:
    """The paper never states this plate's temperature, so none is carried over."""
    induced = mv._environment(induced=True)
    assert induced.temperature is None
    assert induced.gapped_fields() == {"temperature"}
    assert mv.CONJUGATION_TEMPERATURE.value == 30.0


def test_the_environment_refuses_a_medium_that_is_not_an_m9_derivative(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The served medium must still be the M9 the Methods name."""
    swapped = M9_NREL_MOPS_MENASALVAS2025.model_copy(
        update={"base_medium": "LB", "name": "not M9"}
    )
    monkeypatch.setattr(mv, "M9_NREL_MOPS_MENASALVAS2025", swapped)
    with pytest.raises(RuntimeError, match="is not an M9 derivative"):
        mv._environment(induced=True)


def test_the_record_phenotype_is_a_categorical_biosensor_call_not_a_titer() -> None:
    """The released readout is growth under the biosensor, with no number at all."""
    phenotype = mv.selection_phenotype(1, is_reference=False)
    assert phenotype.measurement_type == "categorical"
    assert phenotype.assay_type == "biosensor_readout"
    assert phenotype.category == "enhanced"
    assert phenotype.category_label == "enriched"
    assert phenotype.environment_response is None
    assert phenotype.n_samples == 4
    assert phenotype.sample_unit == "biological_replicate"
    assert phenotype.screen_id == "round-1-lower-isoprenol-threshold"


def test_the_phenotype_declares_the_score_and_the_dispersion_as_typed_absences() -> (
    None
):
    """No per-guide read count and no replicate dispersion were released."""
    phenotype = mv.selection_phenotype(2, is_reference=False)
    assert phenotype.gapped_fields() == {
        "environment_response",
        "environment_response_uncertainty",
        "environment_response_uncertainty_type",
    }
    assert phenotype.environment_response_se is None


def test_the_reference_phenotype_is_the_uninduced_baseline() -> None:
    """The control arm omits the inducer, so no clone can be enriched there."""
    reference = mv.selection_phenotype(1, is_reference=True)
    assert reference.category == "no_change"
    assert reference.category_label == "no crystal violet"


def test_each_round_keeps_its_own_screen_id() -> None:
    """``screen_id`` is what keeps two independently thresholded screens apart."""
    ids = {
        mv.selection_phenotype(n, is_reference=False).screen_id
        for n, _ in mv.ROUND_TABLES
    }
    assert ids == set(mv.SCREEN_IDS.values())
    assert len(ids) == 2


# --------------------------------------------------------------------------- #
# Synthetic: retention bookkeeping
# --------------------------------------------------------------------------- #
def _accounting(**overrides: Any) -> mv.BuildAccounting:
    """A balanced accounting record, overridable field by field."""
    fields: dict[str, Any] = {
        "dataset": "D",
        "source_rows": 58,
        "candidate_records": 58,
        "kept_records": 58,
        "dropped_records": 0,
        "distinct_targets": 60,
    }
    fields.update(overrides)
    return mv.BuildAccounting(**fields)


def test_balanced_accounting_passes_its_own_check() -> None:
    """Nothing is dropped, and the arithmetic says so."""
    accounting = _accounting()
    accounting.check()
    assert accounting.kept_records == accounting.candidate_records == 58
    assert accounting.dropped_records == 0


def test_accounting_refuses_arithmetic_that_does_not_balance() -> None:
    """A record that is neither kept nor dropped is a silent loss."""
    with pytest.raises(RuntimeError, match="!= 58 candidates"):
        _accounting(kept_records=57).check()


def test_accounting_refuses_a_drop_with_no_declared_rule() -> None:
    """A dropped record with no rule naming it is a silent loss, so it refuses."""
    with pytest.raises(
        RuntimeError, match=r"rules total 0, 1 records are missing from the build"
    ):
        _accounting(kept_records=57, dropped_records=1).check()


def test_accounting_accepts_a_drop_a_rule_accounts_for() -> None:
    """A rule whose ``n_records`` totals the drop balances the arithmetic."""
    accounting = _accounting(
        kept_records=57,
        dropped_records=1,
        rules=[
            mv.DropRule(
                rule="synthetic_record_rule",
                scope="record",
                description="one record removed, and the rule says so",
                n_records=1,
                items=["row-57"],
            )
        ],
    )
    accounting.check()
    assert sum(rule.n_records for rule in accounting.rules) == 1
    assert accounting.rules[0].items == ["row-57"]


def test_accounting_refuses_rules_that_do_not_total_the_drop() -> None:
    """Two rules claiming three records against one drop is not accounting."""
    with pytest.raises(
        RuntimeError, match=r"rules total 3, 1 records are missing from the build"
    ):
        _accounting(
            kept_records=57,
            dropped_records=1,
            rules=[
                mv.DropRule(
                    rule="a",
                    scope="record",
                    description="claims two",
                    n_records=2,
                    items=[],
                ),
                mv.DropRule(
                    rule="b",
                    scope="record",
                    description="claims one",
                    n_records=1,
                    items=[],
                ),
            ],
        ).check()


def test_a_key_scope_rule_drops_no_record_and_names_its_keys() -> None:
    """A rule whose scope is a measurement KEY leaves every record in place."""
    rule = mv.DropRule(
        rule="protein_key_merges_two_protein_groups",
        scope="protein_key",
        description="the key names two protein groups",
        n_records=0,
        items=["Aroe", "Asd"],
    )
    accounting = _accounting(rules=[rule])
    accounting.check()
    assert rule.n_records == 0
    assert accounting.dropped_records == 0


# --------------------------------------------------------------------------- #
# Synthetic: the raw mirror
# --------------------------------------------------------------------------- #
def _sha256_bytes(path: Path) -> str:
    """sha256 of a small file read whole."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


# --------------------------------------------------------------------------- #
# Synthetic deposited workbooks: Supplementary Data 1 and 2, in miniature
#
# Every shape the loader checks is reproduced and nothing else: the two merged banner
# cells of data S1-1, its 18-column header and its class footnote, the seven columns of
# data S1-5 with one row per released mutation form, and the seven columns of the
# proteome sheet with six samples at three replicates. The SIZES are the fixture's own,
# so the module's measured row counts are re-pointed by `_pin_synthetic_sizes`.
# --------------------------------------------------------------------------- #
#: The harvest OD600 the growth-phase ``Specific Concentration`` block must recover,
#: taken from the module's own pins so the L4 row has something exact to agree with.
GROWTH_OD600: dict[str, float] = {
    strain: value
    for (strain, phase), value in mv.RECOVERED_HARVEST_OD600.items()
    if phase == "growth"
}
#: The production-phase OD600s, which the module pins no value for (its L4 row only
#: pins which metabolites DISAGREE in that phase).
PRODUCTION_OD600: dict[str, float] = {
    "TEAM-2595": 3.0,
    "TEAM-3174": 1.5,
    "TEAM-3185": 3.3,
}
#: ``Absolute`` rows of the synthetic data S1-1. Every metabolite the module's
#: production-phase disagreement pins name is here, because that row compares the
#: recovered set to the pinned set EXACTLY.
ABSOLUTE_METABOLITES: tuple[str, ...] = (
    "2-Methylcitrate",
    "ADP",
    "ATP",
    "Citrate",
    "Malonate",
    "NAD",
    "NADH",
    "Pyruvate",
    "Methylmalonate",
    "Succinate",
    "Glutamate",
    "4-Aminobutyric acid",
    "Fumarate",
    "Acetate",
)
#: ``Relative`` rows: a different scale, so the build must store none of them.
RELATIVE_METABOLITES: tuple[str, ...] = (
    "Glucose-6-phosphate",
    "Mevalonate",
    "Isoprenol",
)
#: The metabolite whose released cell carries a non-breaking space, so ``_norm`` is
#: exercised on a key the records are built from.
NBSP_METABOLITE = "4-Aminobutyric acid"
NBSP_METABOLITE_CELL = f"4-Aminobutyric{mv.NBSP}acid"
#: The metabolite whose released value is 0 for one strain-phase: a MEASUREMENT, stored.
ZERO_METABOLITE = "Fumarate"
ZERO_KEY: tuple[str, str] = ("TEAM-2595", "growth")
#: The metabolite whose cell is BLANK for one strain-phase: an ABSENT measurement.
BLANK_METABOLITE = "Acetate"
BLANK_KEY: tuple[str, str] = ("TEAM-3174", "production")
#: Per-strain and per-phase factors; one released average is their product times the
#: metabolite's own base, which keeps every value distinct and finite.
_STRAIN_FACTOR: dict[str, float] = {
    "TEAM-2595": 1.0,
    "TEAM-3174": 2.0,
    "TEAM-3185": 3.0,
}
_PHASE_FACTOR: dict[str, float] = {"growth": 1.0, "production": 5.0}


def _metabolite_average(name: str, strain: str, phase: str) -> float | None:
    """One released ``Average Concentration`` cell, or ``None`` where it is blank."""
    if (name, (strain, phase)) == (BLANK_METABOLITE, BLANK_KEY):
        return None
    if (name, (strain, phase)) == (ZERO_METABOLITE, ZERO_KEY):
        return 0.0
    base = 1.0 + float((ABSOLUTE_METABOLITES + RELATIVE_METABOLITES).index(name))
    return base * _STRAIN_FACTOR[strain] * _PHASE_FACTOR[phase]


def _metabolite_specific(name: str, strain: str, phase: str) -> float | None:
    """The released ``Specific Concentration`` cell: the average over the harvest OD600.

    A metabolite the module pins as DISAGREEING in the production phase is divided by
    twice that strain's OD600, which is what makes the L4 row recover the pinned set
    rather than an empty one.
    """
    average = _metabolite_average(name, strain, phase)
    if average is None:
        return None
    od600 = (GROWTH_OD600 if phase == "growth" else PRODUCTION_OD600)[strain]
    if name in mv.SPECIFIC_BLOCK_DISAGREEMENTS[(strain, phase)]:
        return average / (2.0 * od600)
    return average / od600


def _metabolite_rows() -> list[tuple[str, str]]:
    """``(released cell, class)`` of every synthetic data S1-1 row, in released order."""
    rows = [
        (NBSP_METABOLITE_CELL if name == NBSP_METABOLITE else name, "Absolute")
        for name in ABSOLUTE_METABOLITES
    ]
    rows.extend((name, "Relative") for name in RELATIVE_METABOLITES)
    return rows


#: Released ``mutation`` cells, one per form :func:`mv.classify_mutation` accepts, with
#: the variant type and multi-base verdict each one must get.
MUTATION_FORMS: tuple[tuple[str, BacterialVariantType, bool], ...] = (
    (f"A{mv.ARROW_RIGHT}G", BacterialVariantType.snv, False),
    (f"C{mv.ARROW_RIGHT}T", BacterialVariantType.snv, False),
    ("+C", BacterialVariantType.insertion, False),
    ("+CGGG", BacterialVariantType.insertion, True),
    ("Δ1 bp", BacterialVariantType.deletion, False),
    (f"Δ1,227{mv.NBSP}bp", BacterialVariantType.deletion, True),
    (f"(C)6{mv.ARROW_RIGHT}7", BacterialVariantType.insertion, False),
    (f"(G)6{mv.ARROW_RIGHT}5", BacterialVariantType.deletion, False),
    (f"2 bp{mv.ARROW_RIGHT}CT", BacterialVariantType.substitution, True),
    (f"48 bp{mv.ARROW_RIGHT}33 bp", BacterialVariantType.substitution, True),
)
#: The released ``gene`` cell of the one row data S1-5 names by SYMBOL, carrying a
#: non-breaking hyphen, so ``_norm`` is exercised on a name the resolver must reach.
NB_HYPHEN_GENE_CELL = f"cadA{mv.NB_HYPHEN}I {mv.ARROW_RIGHT}"
NB_HYPHEN_GENE = "cadA-I"
NB_HYPHEN_LOCUS = "PP_1100"
#: The in-locus annotation cell the module's docstring names, verbatim in shape.
IN_LOCUS_ANNOTATION = f"A155A (GCG{mv.ARROW_RIGHT}GCA)"
#: ``(sample, evidence, position, mutation, annotation, gene, description)`` of every
#: synthetic data S1-5 row, in released order.
WGS_ROWS: tuple[tuple[str, str, int, str, str, str, str], ...] = (
    (
        "TEAM-2595",
        "RA",
        100,
        MUTATION_FORMS[0][0],
        IN_LOCUS_ANNOTATION,
        f"PP_2088 {mv.ARROW_RIGHT}",
        "RNA polymerase sigma-70 factor",
    ),
    (
        "TEAM-2595",
        "RA",
        200,
        MUTATION_FORMS[1][0],
        "intergenic (-133/-2)",
        f"PP_4401 {mv.ARROW_LEFT} / {mv.ARROW_RIGHT} PP_4403",
        "synthetic intergenic site",
    ),
    (
        "TEAM-2595",
        "RA",
        300,
        MUTATION_FORMS[2][0],
        "coding (539/1299 nt)",
        f"PP_3511 {mv.ARROW_RIGHT}",
        "synthetic coding insertion",
    ),
    (
        "TEAM-2595",
        "RA",
        400,
        MUTATION_FORMS[3][0],
        "coding (12/999 nt)",
        f"PP_3839 {mv.ARROW_LEFT}",
        "synthetic multibase insertion",
    ),
    (
        "TEAM-2595",
        "MC JC",
        500,
        MUTATION_FORMS[4][0],
        "coding (100/900 nt)",
        f"PP_4619 {mv.ARROW_RIGHT}",
        "synthetic one-base deletion",
    ),
    (
        "TEAM-2595",
        "MC JC",
        600,
        MUTATION_FORMS[5][0],
        "coding (1-1227/1227 nt)",
        f"PP_4620 {mv.ARROW_RIGHT}",
        "synthetic span deletion",
    ),
    (
        "TEAM-2595",
        "RA",
        700,
        MUTATION_FORMS[6][0],
        "intergenic (+140/-75)",
        f"PP_4621 {mv.ARROW_RIGHT} / {mv.ARROW_LEFT} PP_5210",
        "synthetic repeat expansion",
    ),
    (
        "TEAM-2595",
        "RA",
        800,
        MUTATION_FORMS[7][0],
        "coding (55/600 nt)",
        NB_HYPHEN_GENE_CELL,
        "synthetic repeat contraction",
    ),
    (
        "TEAM-2595",
        "JC",
        900,
        MUTATION_FORMS[8][0],
        "coding (200/600 nt)",
        f"PP_1816 {mv.ARROW_RIGHT}",
        "synthetic two-base replacement",
    ),
    (
        "TEAM-2595",
        "JC",
        1000,
        MUTATION_FORMS[9][0],
        "noncoding (10/100 nt)",
        f"PP_0100 {mv.ARROW_LEFT}",
        "synthetic length replacement",
    ),
    (
        "TEAM-3175",
        "RA",
        1100,
        f"A{mv.ARROW_RIGHT}T",
        "coding (10/600 nt)",
        f"PP_0200 {mv.ARROW_RIGHT}",
        "refused clone call",
    ),
    (
        "TEAM-3175",
        "RA",
        1200,
        "Δ924 bp",
        "coding (1-924/924 nt)",
        f"PP_2074 {mv.ARROW_RIGHT}",
        "refused clone deletion",
    ),
    (
        "TEAM-3184",
        "RA",
        1300,
        f"G{mv.ARROW_RIGHT}C",
        "coding (20/600 nt)",
        f"PP_0300 {mv.ARROW_RIGHT}",
        "refused clone call",
    ),
    (
        "TEAM-3184",
        "RA",
        1400,
        "Δ924 bp",
        "coding (1-924/924 nt)",
        f"PP_2074 {mv.ARROW_LEFT}",
        "refused clone deletion",
    ),
)
#: Rows per clone the synthetic data S1-5 holds, which the module's measured pins are
#: re-pointed to.
SYNTHETIC_WGS_PER_CLONE: dict[str, int] = dict(
    sorted(Counter(row[0] for row in WGS_ROWS).items())
)
#: The largest released deletion of the synthetic sheet, non-breaking space and all.
SYNTHETIC_LARGEST_DELETION = MUTATION_FORMS[5][0]

#: Host protein keys, as the released ``Protein`` cell -> the locus it must reach. The
#: sheet title-cases a UniProt "tertiary Protein.ID", so every one reaches its locus by
#: a case-insensitive match and one reaches it by GENE SYMBOL instead.
HOST_PROTEIN_KEYS: dict[str, str] = {
    **{f"Pp_{tag.removeprefix('PP_')}": tag for tag in NOTE2_LOCI},
    **{f"Pp_{tag.removeprefix('PP_')}": tag for tag in FILLER_PROTEIN_LOCI},
    "Sotb": "PP_2428",
}
#: The key the sheet files under TWO ``Protein.Group`` accessions, whose abundance
#: therefore stands for two protein groups and reaches no record.
MERGED_PROTEIN_KEY = "Pp_1000"
MERGED_PROTEIN_GROUPS: tuple[str, str] = ("Q90001", "Q90002")
#: The key whose released name is a retired symbol, so no locus of the assembly holds it.
RETIRED_PROTEIN_KEY = "Oldsym"
#: Every host key the synthetic sheet releases, kept or not.
ALL_HOST_PROTEIN_KEYS: tuple[str, ...] = (
    *sorted(HOST_PROTEIN_KEYS),
    MERGED_PROTEIN_KEY,
    RETIRED_PROTEIN_KEY,
)
#: Loci the proteome build keeps: every host key but the merged and the retired one.
KEPT_PROTEOME_LOCI: tuple[str, ...] = tuple(sorted(set(HOST_PROTEIN_KEYS.values())))
#: ``accession -> UniProt entry name`` of the four search-database contaminants, whose
#: organism codes are what makes the loader read them as non-host.
CONTAMINANT_NAMES: dict[str, str] = {
    "P04264": "K2C1_HUMAN",
    "P13645": "K1C10_HUMAN",
    "P35527": "K1C9_HUMAN",
    "P00761": "TRYP_PIG",
}
#: The nine non-host rows, as ``(Protein.Group, Protein.Names, Protein)``.
NON_HOST_ROWS: tuple[tuple[str, str, str], ...] = (
    *(
        (accession, entry, entry.split("_")[0].title())
        for accession, entry in sorted(dict(mv.PATHWAY_ENZYME_ORGANISMS.value).items())
    ),
    *(
        (accession, CONTAMINANT_NAMES[accession], accession)
        for accession in sorted(mv.PROTEOME_CONTAMINANTS.value)
    ),
)
#: The protein ``PP_2088``'s production-phase mean per strain, chosen so the stored fold
#: over the control is EXACTLY Supplementary Note 2's 34x and 20x.
SIGX_PRODUCTION_MEAN: dict[str, float] = {
    "TEAM-2595": 100.0,
    "TEAM-3174": 100.0 * mv.NOTE2_SIGX_FOLD["TEAM-3174"],
    "TEAM-3185": 100.0 * mv.NOTE2_SIGX_FOLD["TEAM-3185"],
}
SIGX_KEY = "Pp_2088"
#: The key measured in exactly ONE replicate of every sample, whose SE is therefore nan.
SINGLE_REPLICATE_KEY = "Pp_0800"
#: The key whose first replicate of one sample is a released 0, averaged in verbatim.
ZEROED_PROTEIN_KEY = "Pp_0700"
ZEROED_PROTEIN_SAMPLE = "2595_growth"
#: The six released samples, as the sheet's own ``Sample`` cells.
PROTEOME_SAMPLES: tuple[str, ...] = tuple(
    f"{strain.removeprefix('TEAM-')}_{phase}"
    for strain in mv.RELEASED_STRAINS
    for phase in mv.PHASES
)
PROTEOME_REPLICATES: tuple[str, ...] = ("R1", "R2", "R3")


def _counts_sum(key: str, sample: str, replicate: str) -> float:
    """One released ``Counts_sum`` cell."""
    strain = f"TEAM-{sample.split('_')[0]}"
    phase = sample.split("_")[1]
    index = ALL_HOST_PROTEIN_KEYS.index(key)
    if key == SIGX_KEY and phase == "production":
        return SIGX_PRODUCTION_MEAN[strain]
    if key == ZEROED_PROTEIN_KEY and sample == ZEROED_PROTEIN_SAMPLE:
        return 0.0 if replicate == "R1" else 300.0
    return (
        (index + 1) * 100.0
        + _STRAIN_FACTOR[strain] * 10.0
        + _PHASE_FACTOR[phase]
        + float(replicate.removeprefix("R"))
    )


def _proteome_rows() -> list[tuple[str, str, str, str, str, str, float]]:
    """Every synthetic proteome row, in the sheet's own column order."""
    rows: list[tuple[str, str, str, str, str, str, float]] = []
    for sample in PROTEOME_SAMPLES:
        for replicate in PROTEOME_REPLICATES:
            for key in ALL_HOST_PROTEIN_KEYS:
                if key == SINGLE_REPLICATE_KEY and replicate != "R1":
                    continue
                if key == MERGED_PROTEIN_KEY:
                    group = MERGED_PROTEIN_GROUPS[
                        PROTEOME_REPLICATES.index(replicate) % 2
                    ]
                else:
                    group = f"Q{ALL_HOST_PROTEIN_KEYS.index(key):05d}"
                rows.append(
                    (
                        group,
                        f"{key.upper()}_{mv.HOST_ORGANISM_CODE}",
                        key,
                        f"synthetic host protein {key}",
                        sample,
                        replicate,
                        _counts_sum(key, sample, replicate),
                    )
                )
            for group, entry, protein in NON_HOST_ROWS:
                rows.append(
                    (
                        group,
                        entry,
                        protein,
                        f"synthetic non-host protein {group}",
                        sample,
                        replicate,
                        50.0 + float(replicate.removeprefix("R")),
                    )
                )
    return rows


def _write_grid(sheet: Any, first_row: int, grid: list[list[Any]]) -> None:
    """Write a rectangular block of cells at an exact 1-based row offset."""
    for offset, values in enumerate(grid):
        for column, value in enumerate(values, start=1):
            if value is not None:
                sheet.cell(row=first_row + offset, column=column, value=value)


def write_data_1(
    path: Path,
    *,
    header: tuple[str, ...] | None = None,
    banner: str | None = None,
    footnote: str | None = None,
    drop_footnote: bool = False,
    second_classless_row: bool = False,
    bad_class: bool = False,
    duplicate_metabolite: bool = False,
    metabolite_sheet: str = mv.SHEET_METABOLITES,
    wgs_sheet: str = mv.SHEET_WGS,
    wgs_header: tuple[str, ...] | None = None,
    wgs_rows: tuple[tuple[str, str, int, str, str, str, str], ...] = WGS_ROWS,
    wgs_header_row: int = mv.WGS_HEADER_ROW,
) -> None:
    """Supplementary Data 1 in miniature: data S1-1 and data S1-5.

    Every keyword re-points one checked shape, so one builder serves the happy path and
    every refusal branch.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = metabolite_sheet
    width = len(mv.METABOLITE_HEADER)
    index = mv._metabolite_column_index()
    average_first = min(pair[0] for pair in index.values())
    specific_first = min(pair[1] for pair in index.values())
    sheet.cell(row=1, column=1, value="Supplementary Data 1, Sheet 1. Metabolites")
    banner_row: list[Any] = [None] * width
    banner_row[average_first] = banner or mv.METABOLITE_AVERAGE_BANNER
    banner_row[specific_first] = mv.METABOLITE_SPECIFIC_BANNER
    _write_grid(sheet, mv.METABOLITE_BANNER_ROW, [banner_row])
    _write_grid(sheet, mv.METABOLITE_HEADER_ROW, [list(header or mv.METABOLITE_HEADER)])
    body: list[list[Any]] = []
    for cell, released_class in _metabolite_rows():
        name = mv._norm(cell)
        row: list[Any] = [None] * width
        row[0] = cell
        row[1] = "Approximate" if bad_class else released_class
        for key, (average_column, specific_column) in index.items():
            row[average_column] = _metabolite_average(name, *key)
            row[specific_column] = _metabolite_specific(name, *key)
        body.append(row)
    if duplicate_metabolite:
        body.append(list(body[0]))
    if not drop_footnote:
        footnote_row: list[Any] = [None] * width
        footnote_row[0] = footnote or mv._Q_METABOLITE_FOOTNOTE
        body.append(footnote_row)
    if second_classless_row:
        extra: list[Any] = [None] * width
        extra[0] = "a second class-free row"
        body.append(extra)
    _write_grid(sheet, mv.METABOLITE_HEADER_ROW + 1, body)

    wgs = book.create_sheet(wgs_sheet)
    wgs.cell(row=1, column=1, value="Sheet 5. Illumina Genome Resequencing")
    _write_grid(wgs, wgs_header_row, [list(wgs_header or mv.WGS_HEADER)])
    _write_grid(wgs, wgs_header_row + 1, [list(row) for row in wgs_rows])
    book.save(path)


def write_data_2(
    path: Path,
    *,
    sheet_name: str = mv.SHEET_PROTEOME,
    header: tuple[str, ...] | None = None,
    header_row: int = mv.PROTEOME_HEADER_ROW,
    rows: list[tuple[str, str, str, str, str, str, float]] | None = None,
    blank_counts: bool = False,
    bad_replicate: bool = False,
) -> None:
    """Supplementary Data 2 in miniature: the one proteomics sheet the loader reads."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.cell(row=1, column=1, value="Sheet 4. Growth/Production phase Samples")
    _write_grid(sheet, header_row, [list(header or mv.PROTEOME_HEADER)])
    body = [list(row) for row in (rows if rows is not None else _proteome_rows())]
    if blank_counts:
        body[0][6] = None
    if bad_replicate:
        body[0][5] = "rep one"
    _write_grid(sheet, header_row + 1, body)
    book.save(path)


#: Junk members the released inner zip ships beside the data, plus a member whose name
#: traverses out of the destination. None of them may ever be written.
JUNK_MEMBERS: tuple[str, ...] = (
    f"__MACOSX/._{mv.INNER_ZIP_ROOT}",
    f"{mv.INNER_ZIP_ROOT}/.DS_Store",
    f"{mv.INNER_ZIP_ROOT}/~$Menasalvas et al Supplementary Data 1.xlsx",
    f"{mv.INNER_ZIP_ROOT}/a.fcs",
    f"{mv.INNER_ZIP_ROOT}/b.fastq",
    f"{mv.INNER_ZIP_ROOT}/c.mp4",
    "../evil.xlsx",
)
SYNTHETIC_DRYAD_README = (
    "# Synthetic Dryad README\n\nSheet 1. Metabolite concentrations: Average value "
    "from 3 biological replicates. GP = growth phase samples. PP = production phase "
    "samples.\n\nSheet 5: Genomic Coordinates in P. putida AE015451.\n"
)


def write_inner_zip(
    path: Path, data_1: Path, data_2: Path, *, members: tuple[str, ...] | None = None
) -> None:
    """The released inner zip: the two workbooks under their member names, plus junk."""
    held = (mv.SI_DATA_1_MEMBER, mv.SI_DATA_2_MEMBER) if members is None else members
    with zipfile.ZipFile(path, "w") as archive:
        if mv.SI_DATA_1_MEMBER in held:
            archive.writestr(mv.SI_DATA_1_MEMBER, data_1.read_bytes())
        if mv.SI_DATA_2_MEMBER in held:
            archive.writestr(mv.SI_DATA_2_MEMBER, data_2.read_bytes())
        for junk in JUNK_MEMBERS:
            archive.writestr(junk, b"junk the deposit ships beside the data")


def write_dryad_deposit(root: Path, staging: Path) -> dict[str, str]:
    """Write the three deposited Dryad files; return ``relpath -> sha256``.

    The owner's browser download is the arrival zip; the recipe unzips it beside
    itself, which is why all three files sit in the mirror and all three are pinned.
    """
    staging.mkdir(parents=True, exist_ok=True)
    data_1 = staging / mv.SI_DATA_1_BASENAME
    data_2 = staging / mv.SI_DATA_2_BASENAME
    write_data_1(data_1)
    write_data_2(data_2)
    directory = root / mv.DRYAD_DIR_REL
    directory.mkdir(parents=True, exist_ok=True)
    inner = directory / mv.DRYAD_INNER_ZIP_FILENAME
    readme = directory / mv.DRYAD_README_FILENAME
    write_inner_zip(inner, data_1, data_2)
    readme.write_text(SYNTHETIC_DRYAD_README, encoding="utf-8")
    arrival = directory / mv.DRYAD_ZIP_FILENAME
    with zipfile.ZipFile(arrival, "w") as archive:
        archive.write(inner, arcname=inner.name)
        archive.write(readme, arcname=readme.name)
    return {
        mv.DRYAD_ZIP_REL: _sha256_bytes(arrival),
        mv.DRYAD_INNER_ZIP_REL: _sha256_bytes(inner),
        mv.DRYAD_README_REL: _sha256_bytes(readme),
        mv.SI_DATA_1_MEMBER: _sha256_bytes(data_1),
        mv.SI_DATA_2_MEMBER: _sha256_bytes(data_2),
    }


#: Which module constant pins which deposited Dryad file.
_DRYAD_PIN_OF_REL: dict[str, str] = {
    mv.DRYAD_ZIP_REL: "DRYAD_ZIP_SHA256",
    mv.DRYAD_INNER_ZIP_REL: "DRYAD_INNER_ZIP_SHA256",
    mv.DRYAD_README_REL: "DRYAD_README_SHA256",
}


def _pin_synthetic_sizes(monkeypatch: pytest.MonkeyPatch) -> None:
    """Re-point every count the module measured on the REAL bytes at the fixture's own.

    The fixture is a miniature of the released SHAPE, not of its size. The released
    numbers are asserted by the ``@pytest.mark.data`` block instead.
    """
    monkeypatch.setattr(
        mv,
        "EXPECTED_METABOLITE_ROWS",
        {
            mv.CONCENTRATION_ABSOLUTE: len(ABSOLUTE_METABOLITES),
            mv.CONCENTRATION_RELATIVE: len(RELATIVE_METABOLITES),
        },
    )
    monkeypatch.setattr(mv, "EXPECTED_WGS_ROWS", len(WGS_ROWS))
    monkeypatch.setattr(
        mv, "EXPECTED_WGS_ROWS_PER_CLONE", dict(SYNTHETIC_WGS_PER_CLONE)
    )


@pytest.fixture
def synthetic_mirror(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A raw mirror under ``tmp_path`` written by the module's own deposit function.

    The SI PDF and its OCR are staged and their digests re-pointed; the three Dryad
    files are a MANUAL deposit, so the fixture puts them in place exactly as the
    owner's browser download does and re-points their pins too. ``dryad_deposits``
    reads those module constants at CALL time, so the re-pointing reaches it, and
    ``deposit_raw_mirror`` then verifies and describes all five files.
    """
    staging = tmp_path / "staging"
    staging.mkdir()
    pdf = staging / mv.SI_PDF_FILENAME
    pdf.write_bytes(b"%PDF-1.7 synthetic supplementary information\n")
    markdown = staging / mv.SI1_MD_FILENAME
    markdown.write_text(synthetic_markdown(), encoding="utf-8")
    monkeypatch.setattr(mv, "SI_PDF_SHA256", _sha256_bytes(pdf))
    monkeypatch.setattr(mv, "SI1_MD_SHA256", _sha256_bytes(markdown))
    monkeypatch.setattr(
        mv, "EXPECTED_ROWS", {1: len(ROUND1_CELLS), 2: len(ROUND2_CELLS)}
    )
    monkeypatch.setattr(mv, "EXPECTED_RECORDS", len(ROUND1_CELLS) + len(ROUND2_CELLS))
    _pin_synthetic_sizes(monkeypatch)
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    digests = write_dryad_deposit(
        data_root / mv.RAW_DIR_REL, tmp_path / "dryad-staging"
    )
    for relpath, pin in _DRYAD_PIN_OF_REL.items():
        monkeypatch.setattr(mv, pin, digests[relpath])
    monkeypatch.setattr(
        mv,
        "EXTRACTED_MEMBERS",
        {
            mv.SI_DATA_1_MEMBER: (mv.SI_DATA_1_BASENAME, digests[mv.SI_DATA_1_MEMBER]),
            mv.SI_DATA_2_MEMBER: (mv.SI_DATA_2_BASENAME, digests[mv.SI_DATA_2_MEMBER]),
        },
    )
    root = mv.deposit_raw_mirror(
        si_pdf_path=pdf, si_md_path=markdown, data_root=str(data_root)
    )
    assert root == mv.raw_mirror_dir(str(data_root))
    return data_root


@pytest.fixture
def deposited_workbooks(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """The two synthetic workbooks on disk, with the module's row counts re-pointed.

    The readers take a path, so this is what every reader and refusal test uses; it
    neither deposits a mirror nor builds a store.
    """
    _pin_synthetic_sizes(monkeypatch)
    directory = tmp_path / "workbooks"
    directory.mkdir()
    data_1 = directory / mv.SI_DATA_1_BASENAME
    data_2 = directory / mv.SI_DATA_2_BASENAME
    write_data_1(data_1)
    write_data_2(data_2)
    return data_1, data_2


def test_the_deposit_records_a_rerunnable_retrieval_and_an_ocr_processing_step(
    synthetic_mirror: Path,
) -> None:
    """The PDF carries its PMC retrieval; the markdown carries the MinerU run."""
    manifest = mv.load_manifest(str(synthetic_mirror))
    by_path = {record.path: record for record in manifest.files}
    pdf = by_path[mv.SI_PDF_REL]
    assert pdf.retrieval is not None
    assert pdf.retrieval.method == "pmc_cloud"
    assert pdf.retrieval.retriever == ("torchcell.literature.retrieve.pmc_cloud_object")
    assert pdf.retrieval.params == {"key": f"{mv.PMC_PREFIX}/{mv.SI_PDF_FILENAME}"}
    assert pdf.original_filename == mv.SI_PDF_FILENAME
    ocr = by_path[mv.SI1_MD_REL]
    assert ocr.retrieval is None
    assert ocr.processing is not None
    assert ocr.processing.tool == "mineru"
    assert ocr.processing.version == mv.MINERU_VERSION
    assert ocr.processing.input_sha256 == [mv.SI_PDF_SHA256]


def test_the_manifest_enumerates_every_deposit_and_states_the_titer_gap(
    synthetic_mirror: Path,
) -> None:
    """Dryad, Zenodo, PRIDE and the BioProject are listed; none holds a titer."""
    manifest = mv.load_manifest(str(synthetic_mirror))
    assert any(mv.DRYAD_DOI in source for source in manifest.si_data_sources)
    assert any(mv.ZENODO_DOI in source for source in manifest.si_data_sources)
    assert any(mv.PRIDE_ACCESSION in source for source in manifest.si_data_sources)
    assert any(mv.BIOPROJECT_ACCESSION in source for source in manifest.si_data_sources)
    assert any(
        "NO PER-STRAIN ISOPRENOL TITER IS RELEASED" in entry
        for entry in manifest.si_expected
    )
    assert any(
        "NO PER-GUIDE ENRICHMENT MATRIX IS RELEASED" in entry
        for entry in manifest.si_expected
    )


def test_manifest_sha256_reads_a_pin_and_refuses_an_unknown_path(
    synthetic_mirror: Path,
) -> None:
    """A path the mirror does not hold is a KeyError, never a silent None."""
    manifest = mv.load_manifest(str(synthetic_mirror))
    assert mv.manifest_sha256(manifest, mv.SI1_MD_REL) == mv.SI1_MD_SHA256
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        mv.manifest_sha256(manifest, "si/absent.md")


def test_redepositing_the_same_bytes_is_idempotent(synthetic_mirror: Path) -> None:
    """A second deposit of matching bytes rewrites the manifest and nothing else."""
    root = mv.raw_mirror_dir(str(synthetic_mirror))
    before = (root / mv.SI1_MD_REL).read_bytes()
    mv.deposit_raw_mirror(
        si_pdf_path=root / mv.SI_PDF_REL,
        si_md_path=root / mv.SI1_MD_REL,
        data_root=str(synthetic_mirror),
    )
    assert (root / mv.SI1_MD_REL).read_bytes() == before


def test_a_mirror_file_with_a_different_digest_is_never_overwritten(
    synthetic_mirror: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Upstream drift is detected, not followed."""
    root = mv.raw_mirror_dir(str(synthetic_mirror))
    (root / mv.SI1_MD_REL).write_text("drifted", encoding="utf-8")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        mv.deposit_raw_mirror(
            si_pdf_path=root / mv.SI_PDF_REL,
            si_md_path=tmp_path / "staging" / mv.SI1_MD_FILENAME,
            data_root=str(synthetic_mirror),
        )


def test_the_isoprenol_inchikey_is_recorded_and_agrees_with_the_table() -> None:
    """No record names a product here, but the recorded key is the table's own.

    The curation step this constant was recorded for has happened (the table's isoprenol
    row, PubChem CID 12988), so ``resolved_compound`` no longer gaps ``inchikey``; the
    constant is kept as the cross-check that the row is the molecule this paper means.
    """
    assert mv.ISOPRENOL_INCHIKEY == "CPJRRXSHAYUTGL-UHFFFAOYSA-N"
    from torchcell.datamodels.compound_identity import resolved_compound

    isoprenol = resolved_compound("isoprenol")
    assert isoprenol.inchikey == mv.ISOPRENOL_INCHIKEY
    assert isoprenol.gapped_fields() == set()


def test_raw_mirror_dir_falls_back_to_the_environment(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Called with no argument, the mirror path comes from ``DATA_ROOT``."""
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    assert mv.raw_mirror_dir() == tmp_path / mv.RAW_DIR_REL


# --------------------------------------------------------------------------- #
# Hermetic end to end: the synthetic assembly, the mirror, the built store
# --------------------------------------------------------------------------- #
def _synthetic_loci() -> list[Any]:
    """One ``SyntheticLocus`` per :data:`LOCUS_SPECS` entry, laid end to end."""
    from tests.torchcell.sequence.genome._bacterial_fixtures import SyntheticLocus

    loci = []
    cursor = 1
    for index, (tag, symbol) in enumerate(LOCUS_SPECS):
        start, end = cursor, cursor + 11
        cursor = end + 3
        loci.append(
            SyntheticLocus(
                tag=tag,
                parts=((start, end),),
                strand="+" if index % 2 == 0 else "-",
                symbol=symbol,
                product=f"synthetic product {tag}",
                protein_id=f"AAN{index:05d}.1",
                protein="MKV",
            )
        )
    return loci


@pytest.fixture
def synthetic_kt2440(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """The real KT2440 class over a synthetic assembly, with the network refused."""
    import tests.torchcell.sequence.genome._bacterial_fixtures as fixtures
    from torchcell.sequence.genome.pputida.kt2440 import (
        KT2440_ASSEMBLY,
        PPutidaKT2440Genome,
    )

    monkeypatch.setattr(fixtures, "SEQUENCE", fixtures.SEQUENCE * 8)
    files = fixtures.write_assembly(
        tmp_path / "tier",
        KT2440_ASSEMBLY,
        _synthetic_loci(),
        [fixtures.gaf_row("pyrF", "pyrF|PP_1815", "GO:0000001")],
    )
    fixtures.forbid_network(monkeypatch)
    fixtures.serve_tier(monkeypatch, files)
    report = tmp_path / "tier" / ASSEMBLY_REPORT_MEMBER
    report.write_text(ASSEMBLY_REPORT)

    def serve_report(assembly_set: str, filename: str, **_: Any) -> str:
        if filename != ASSEMBLY_REPORT_MEMBER:
            raise FileNotFoundError(f"{assembly_set}/{filename} is not in the fixture")
        return str(report)

    import torchcell.datasets.bacteria_common as bacteria_common

    monkeypatch.setattr(bacteria_common, "resolve", serve_report)
    root = tmp_path / "kt2440"
    root.mkdir()
    return PPutidaKT2440Genome(genome_root=str(root), overwrite=False)


@pytest.fixture
def built(synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path) -> Any:
    """The loader built over the synthetic mirror and the synthetic annotation."""
    return mv.IsoprenolSelectionMenasalvas2025Dataset(
        root=str(tmp_path / "build"), pputida_genome=synthetic_kt2440
    )


def test_the_build_writes_one_record_per_selected_target(built: Any) -> None:
    """Record count equals the parsed rows, and the two references are the rounds'."""
    assert len(built) == len(ROUND1_CELLS) + len(ROUND2_CELLS)
    assert built.experiment_class is BacterialEnvironmentResponseExperiment
    assert built.reference_class is BacterialEnvironmentResponseExperimentReference
    assert built.raw_file_names == [mv.SI1_MD_FILENAME]


def test_each_built_record_carries_the_host_background_plus_its_knockdowns(
    built: Any,
) -> None:
    """A record is the constant host plus one guide's targets, and nothing else."""
    item = built[0]
    perturbations = item["experiment"]["genotype"]["perturbations"]
    knockdowns = [
        p
        for p in perturbations
        if p["perturbation_type"] == "bacterial_crispr_interference"
    ]
    others = {
        p["systematic_gene_name"]
        for p in perturbations
        if p["perturbation_type"] != "bacterial_crispr_interference"
    }
    assert others == set(mv.HOST_BACKGROUND_NAMES)
    assert len(knockdowns) == 3
    assert [p["systematic_gene_name"] for p in knockdowns] == [
        "PP_5003",
        "PP_5004",
        "PP_5005",
    ]


def test_the_stored_gene_name_is_the_annotations_not_the_tables(built: Any) -> None:
    """The source writes ``phaAZC-II``; the assembly annotates ``PP_5004`` as phaB."""
    perturbations = built[0]["experiment"]["genotype"]["perturbations"]
    by_tag = {
        p["systematic_gene_name"]: p["perturbed_gene_name"]
        for p in perturbations
        if p["perturbation_type"] == "bacterial_crispr_interference"
    }
    assert by_tag["PP_5004"] == "phaB"
    assert by_tag["PP_5005"] == "phaC-II"


def test_an_unannotated_locus_keeps_its_tag_as_the_common_name(built: Any) -> None:
    """``PP_5007`` has no symbol in the assembly, so the tag is its common name."""
    names: dict[str, str] = {}
    for index in range(len(built)):
        for p in built[index]["experiment"]["genotype"]["perturbations"]:
            if p["perturbation_type"] == "bacterial_crispr_interference":
                names[p["systematic_gene_name"]] = p["perturbed_gene_name"]
    assert names["PP_5007"] == "PP_5007"
    assert names["PP_2428"] == "sotB"


def test_the_build_records_its_own_retention_and_its_unasserted_chassis(
    built: Any, tmp_path: Path
) -> None:
    """The accounting states the arithmetic and the chassis the records do not carry."""
    accounting = json.loads(
        Path(built.preprocess_dir, "build_accounting.json").read_text()
    )
    assert accounting["kept_records"] == len(built)
    assert accounting["dropped_records"] == 0
    assert accounting["distinct_targets"] == 11
    assert accounting["reconciliation"]["unique_names"] == 11
    assert accounting["reconciliation"]["status_histogram"]["current"] == 11
    assert accounting["reconciliation"]["outside_namespace"] == []
    joined = " ".join(accounting["notes"])
    assert "no ProductTiterPhenotype is written" in joined
    assert "PJ23119-PP_1697" in joined


def test_the_build_keeps_the_tables_verbatim_cells_beside_the_record(
    built: Any,
) -> None:
    """``preprocess/selected_targets.csv`` keeps the source's own spelling and function."""
    import pandas as pd

    frame = pd.read_csv(osp.join(built.preprocess_dir, "selected_targets.csv"))
    assert len(frame) == len(built)
    assert set(frame["round"]) == {1, 2}
    row = frame.loc[frame["gene_cell"] == "phaAZC-II / PP_5003-PP_5005"].iloc[0]
    assert row["source_symbol"] == "phaAZC-II"
    assert row["locus_tags"] == "PP_5003;PP_5004;PP_5005"
    assert row["annotation_symbols"] == "phaA;phaB;phaC-II"


def test_a_target_outside_the_pinned_assembly_stops_the_build(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A tag the annotation does not carry is a refusal, never a dropped record."""
    from torchcell.datasets.bacteria_common import LocusTagResolutionError

    root = mv.raw_mirror_dir(str(synthetic_mirror))
    markdown = root / mv.SI1_MD_REL
    markdown.write_text(
        synthetic_markdown(round1=ROUND1_CELLS[:-1] + ("PP_9999",)), encoding="utf-8"
    )
    digest = _sha256_bytes(markdown)
    monkeypatch.setattr(mv, "SI1_MD_SHA256", digest)
    manifest = mv.load_manifest(str(synthetic_mirror))
    for record in manifest.files:
        if record.path == mv.SI1_MD_REL:
            record.sha256 = digest
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    with pytest.raises(LocusTagResolutionError):
        mv.IsoprenolSelectionMenasalvas2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_target_in_another_hosts_namespace_stops_the_build(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A reconciled name outside the KT2440 namespace is a refusal, not a record.

    No ``PP_`` tag can reach this branch through the parser, so the reconciliation is
    stubbed: the guard exists for a future annotation whose resolution crosses hosts.
    """
    real = reconcile_locus_tags

    def crossing(genome: Any, names: Any, **kwargs: Any) -> Any:
        stored, report = real(genome, names, **kwargs)
        return stored, report.model_copy(update={"outside_namespace": ["b0002"]})

    monkeypatch.setattr(mv, "reconcile_locus_tags", crossing)
    with pytest.raises(RuntimeError, match="selected targets outside"):
        mv.IsoprenolSelectionMenasalvas2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_missing_mirror_file_stops_the_build(
    synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path
) -> None:
    """A mirror the manifest describes but does not hold is a refusal."""
    (mv.raw_mirror_dir(str(synthetic_mirror)) / mv.SI1_MD_REL).unlink()
    with pytest.raises(RuntimeError, match="required raw artifact missing"):
        mv.IsoprenolSelectionMenasalvas2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_a_parsed_target_count_that_is_not_the_stated_one_stops_the_build(
    built: Any, synthetic_kt2440: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The record total is checked against the Methods, not taken from the file."""
    monkeypatch.setattr(mv, "EXPECTED_RECORDS", 99)
    with pytest.raises(RuntimeError, match="not the stated 99"):
        mv.IsoprenolSelectionMenasalvas2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


def test_the_genome_is_opened_from_the_tier_when_none_is_injected(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A loader built with no genome asks ``bacterial_genome`` for one."""
    calls: list[tuple[str, str]] = []

    def fake(host: str, strain: str, *args: Any, **kwargs: Any) -> Any:
        calls.append((host, strain))
        return synthetic_kt2440

    monkeypatch.setattr(mv, "bacterial_genome", fake)
    dataset = mv.IsoprenolSelectionMenasalvas2025Dataset(root=str(tmp_path / "tier"))
    assert len(dataset) == len(ROUND1_CELLS) + len(ROUND2_CELLS)
    assert calls == [("pputida", "KT2440")]


def test_the_unimplemented_interfaces_behave_as_the_family_does(built: Any) -> None:
    """``preprocess_raw`` is a pass-through and ``create_experiment`` is not used."""
    assert built.preprocess_raw("x") == "x"
    with pytest.raises(NotImplementedError):
        built.create_experiment()


# --------------------------------------------------------------------------- #
# Hermetic: the verification rows
# --------------------------------------------------------------------------- #
def test_verify_build_passes_on_the_hermetic_store(
    built: Any, synthetic_kt2440: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The L0-L4 gate is green on a synthetic build, and the report is written."""
    monkeypatch.setattr(
        mv, "bacterial_genome", lambda *args, **kwargs: synthetic_kt2440
    )
    report = mv.verify_build(built.root, family="selection")
    assert report.passed, report.summary()
    names = {result.name for result in report.results}
    assert "stored_targets_are_loci_of_the_pinned_assembly" in names
    assert "assembly_pin_is_kt2440_genbank_with_no_asserted_background" in names
    assert "every_genotype_carries_the_stated_host_background" in names
    assert Path(built.preprocess_dir, "verification_report.json").is_file()


def test_the_supplementary_rows_fail_when_the_records_do_not_hold(
    built: Any, synthetic_kt2440: Any
) -> None:
    """Each supplementary rule is reachable in the failing direction."""
    from torchcell.verification.runners import load_records

    records = load_records(built.root)
    assert mv.stored_tags_are_loci(records, synthetic_kt2440).passed
    assert mv.assembly_pin(records).passed
    assert mv.host_background_present(records).passed

    broken = [
        {
            "experiment": {
                "genotype": {
                    "perturbations": [
                        {
                            "systematic_gene_name": "PP_9999",
                            "perturbation_type": "bacterial_crispr_interference",
                        }
                    ]
                }
            },
            "reference": {
                "genome_reference": {
                    "assembly_set": "ecoli_K12_MG1655_ASM584v2",
                    "assembly_accession": "GCA_000005845.2",
                    "background": None,
                }
            },
        }
    ]
    assert not mv.stored_tags_are_loci(broken, synthetic_kt2440).passed
    assert not mv.assembly_pin(broken).passed
    assert not mv.host_background_present(broken).passed


def test_main_builds_verifies_and_prints(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """The module's own runner builds all four families and verifies every one."""
    monkeypatch.setattr(
        mv, "bacterial_genome", lambda *args, **kwargs: synthetic_kt2440
    )
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "data_root"))
    mv.main()
    out = capsys.readouterr().out
    assert f"len = {len(ROUND1_CELLS) + len(ROUND2_CELLS)}" in out
    assert "distinct_targets" in out
    assert "[ok] L0 structural" in out
    assert "[XX]" not in out
    for cls, _rel in mv.FAMILY_BUILDS.values():
        assert f"{cls.__name__}: len = " in out
    assert "IsoprenolSelectionMenasalvas2025Dataset: PASS" in out
    assert "proteome_menasalvas2025: PASS" in out
    assert "metabolite_growth_menasalvas2025: PASS" in out
    assert "metabolite_production_menasalvas2025: PASS" in out
    assert out.count("ProteomeMenasalvas2025Dataset: len = 6") == 1


# --------------------------------------------------------------------------- #
# Data-gated: the numbers measured on the REAL pinned bytes
# --------------------------------------------------------------------------- #
LIBRARY_SI = osp.join(
    "/scratch/projects/torchcell-scratch/torchcell-library",
    mv.CITATION_KEY,
    "si/si1.md",
)
requires_real_si = pytest.mark.skipif(
    not osp.isfile(LIBRARY_SI), reason="requires the Menasalvas 2025 library mirror"
)


@requires_real_si
@pytest.mark.data
def test_the_real_tables_hold_the_stated_28_and_30_targets() -> None:
    """Measured on the pinned OCR: 58 rows over 60 distinct loci, rounds disjoint."""
    targets = mv.read_selected_targets(LIBRARY_SI)
    assert len(targets) == mv.EXPECTED_RECORDS == 58
    assert sum(1 for t in targets if t.round_number == 1) == 28
    assert sum(1 for t in targets if t.round_number == 2) == 30
    tags = [tag for target in targets for tag in target.locus_tags]
    assert len(tags) == 60
    assert len(set(tags)) == 60
    multi = [t for t in targets if len(t.locus_tags) > 1]
    assert [t.gene_cell for t in multi] == ["phaAZC-II / PP_5003-PP_5005"]
