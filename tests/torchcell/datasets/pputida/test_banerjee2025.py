# tests/torchcell/datasets/pputida/test_banerjee2025.py
# [[tests.torchcell.datasets.pputida.test_banerjee2025]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/pputida/test_banerjee2025.py
"""The Banerjee 2025 P. putida p-coumarate growth-coupling proteome loader.

Synthetic tests run everywhere: they exercise the label canonicalization the four
released workbooks force, the released-file reader, all five build oracles (the protein
counts, the two groups released twice, the replicate back-solve with its merged-label
exception, PP_0897's own strain identification, and the two text-reproducing counts), the
chassis background and the three genotype shapes, all three environments and their typed
gaps, the retention arithmetic, the raw-mirror deposit over a synthetic archive, the
loader built end to end under ``tmp_path`` with no network and no ``$DATA_ROOT``, and
every verification rule.

The ``@pytest.mark.data`` tests read the real ``$DATA_ROOT``: they assert that every
module quote is a verbatim substring of the sha256-pinned mirrored bytes, pin the raw
mirror's recorded digests against the module constants, and run L0 to L4 over the built
LMDB.

Derived expectations for the pinned released files, every one measured: the two promoter
workbooks hold 2,494 protein rows each and the two cross-feeding workbooks 2,470; the
union of their labels is 2,763 of which 272 are dropped (270 outside the namespace, 2
merged), leaving 2,491 candidates; the six released samples carry 13,413 stored
abundances in all; PP_0897's released log2 fold change is -4.042 (pJ23109), -3.283
(PP_0415 promoter), -7.928 (the deletion against its parent) and +0.048 (the deletion
between its two media).
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import os.path as osp
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pytest

import torchcell.datasets.pputida.banerjee2025 as ds
from torchcell.data.experiment_dataset import RawSha256MismatchError
from torchcell.datamodels.media import M9_DEFERRED_BANERJEE2025
from torchcell.datamodels.schema import (
    AlleleEdit,
    BacterialProteinAbundanceExperiment,
    BacterialProteinAbundanceExperimentReference,
    ConcentrationUnit,
    EnvironmentPhysicalPerturbation,
    PhysicalFactor,
)
from torchcell.literature.manifest import ROLE_SI_DATA, RetrievalMethod

ASSEMBLY_REPORT = """# Assembly name:  ASM756v2
# Organism name:  Pseudomonas putida KT2440 (g-proteobacteria)
# Infraspecific name:  strain=KT2440
# Taxid:          160488
# GenBank assembly accession: GCA_000007565.2
# RefSeq assembly accession: GCF_000007565.2
# RefSeq assembly and GenBank assemblies identical: yes
#
## Assembly-Units:
AE015451.2\tassembled-molecule\tna\tChromosome\tAE015451.2\t=\tNC_002947.4
"""
ASSEMBLY_REPORT_MEMBER = "GCA_000007565.2_ASM756v2_assembly_report.txt"


def _serve_assembly_report(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Serve the KT2440 assembly report from ``tmp_path``, with no tier on disk."""
    import torchcell.datasets.bacteria_common as bacteria_common

    report = tmp_path / ASSEMBLY_REPORT_MEMBER
    report.write_text(ASSEMBLY_REPORT)

    def serve(assembly_set: str, filename: str, **_: Any) -> str:
        if filename != ASSEMBLY_REPORT_MEMBER:
            raise FileNotFoundError(f"{assembly_set}/{filename} is not in the fixture")
        return str(report)

    monkeypatch.setattr(bacteria_common, "resolve", serve)


@pytest.fixture
def served_assembly_report(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """The assembly report the reference builder reads, with no genomes tier."""
    _serve_assembly_report(tmp_path, monkeypatch)


# --------------------------------------------------------------------------- #
# The label canonicalization the four workbooks force
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("label", "stored"),
    [
        ("PP_0002", "PP_0002"),
        ("Pp_0002", "PP_0002"),
        ("pp_0002", "PP_0002"),
        (" PP_0897 ", "PP_0897"),
        ("Aapj", "Aapj"),
        ("A8926_6430", "A8926_6430"),
        ("Tryp_pig", "Tryp_pig"),
    ],
)
def test_only_a_locus_tag_shaped_label_is_upper_cased(label: str, stored: str) -> None:
    """The case disagreement is confined to locus tags, so only those are folded."""
    assert ds.canonical_label(label) == stored


def test_a_symbol_that_merely_contains_a_tag_shape_is_left_alone() -> None:
    """``fullmatch`` and not ``search``: a longer label is not a locus tag."""
    assert ds.canonical_label("Pp_0002x") == "Pp_0002x"
    assert ds.canonical_label("xPp_0002") == "xPp_0002"


# --------------------------------------------------------------------------- #
# The group table
# --------------------------------------------------------------------------- #
def test_the_group_table_is_six_samples_over_four_strains_and_three_media() -> None:
    """The record set, named once, with each group's two released columns derived."""
    assert len(ds.PROTEOME_GROUPS) == ds.EXPECTED_PROTEOME_RECORDS == 6
    assert {g.key for g in ds.PROTEOME_GROUPS} == {
        "2370",
        "2370_pJ_PP_0897",
        "2370_0415pPP_0897",
        "2370_M9_pCA_alanine_malate",
        "2487_M9_pCA_alanine_malate",
        "2487_M9_alanine_malate",
    }
    assert {g.strain for g in ds.PROTEOME_GROUPS} == {
        ds.STRAIN_PARENT,
        ds.STRAIN_PJ,
        ds.STRAIN_0415,
        ds.STRAIN_DELETION,
    }
    assert {g.condition for g in ds.PROTEOME_GROUPS} == {
        ds.CONDITION_PCA,
        ds.CONDITION_PCA_ALA_MAL,
        ds.CONDITION_ALA_MAL,
    }
    group = next(g for g in ds.PROTEOME_GROUPS if g.key == "2370_pJ_PP_0897")
    assert group.mean_column == "log2_mean_2370_pJ_PP_0897"
    assert group.sd_column == "log2_std_2370_pJ_PP_0897"


def test_each_arm_has_its_own_parental_reference_and_replicate_count() -> None:
    """The two arms are searched separately, so neither borrows the other's baseline."""
    assert ds.ARM_REFERENCE_KEY == {
        ds.ARM_PROMOTER: "2370",
        ds.ARM_CROSSFEED: "2370_M9_pCA_alanine_malate",
    }
    assert ds.ARM_REPLICATES == {ds.ARM_PROMOTER: 4, ds.ARM_CROSSFEED: 3}
    for arm, key in ds.ARM_REFERENCE_KEY.items():
        reference = next(g for g in ds.PROTEOME_GROUPS if g.key == key)
        assert reference.arm == arm
        assert reference.strain == ds.STRAIN_PARENT


def test_every_group_reads_a_workbook_that_releases_its_column() -> None:
    """A group's workbook is one of the two the workbook-group table pairs it with."""
    for group in ds.PROTEOME_GROUPS:
        assert group.key in ds.WORKBOOK_GROUPS[group.workbook]


# --------------------------------------------------------------------------- #
# The replicate back-solve
# --------------------------------------------------------------------------- #
def _row(
    label: str,
    a_mean: float,
    b_mean: float,
    a_sd: float,
    b_sd: float,
    n: int,
    *,
    accession: str = "Q00001",
    p_value: float = 0.01,
) -> ds.ComparisonRow:
    """A synthetic released row whose t statistic is exact for ``n`` replicates."""
    t = (a_mean - b_mean) / math.sqrt((a_sd**2 + b_sd**2) / n)
    return ds.ComparisonRow(
        label=label,
        accession=accession,
        description=f"synthetic {label}",
        a_mean=a_mean,
        b_mean=b_mean,
        a_sd=a_sd,
        b_sd=b_sd,
        t_statistic=t,
        p_value=p_value,
        log2_fold_change=a_mean - b_mean,
    )


def test_the_back_solve_recovers_the_replicate_count_in_closed_form() -> None:
    """``n`` is recoverable from the released mean, SD and t, with no search."""
    row = _row("PP_0001", 22.0, 20.0, 0.4, 0.5, 4)
    assert ds.back_solved_replicates(row) == pytest.approx(4.0)
    row = _row("PP_0001", 22.0, 20.0, 0.4, 0.5, 3)
    assert ds.back_solved_replicates(row) == pytest.approx(3.0)


def test_a_flat_row_carries_no_information_about_the_replicate_count() -> None:
    """A zero difference leaves ``n`` unconstrained, so the row is skipped."""
    flat = ds.ComparisonRow(
        label="PP_0001",
        accession="Q00001",
        description=None,
        a_mean=20.0,
        b_mean=20.0,
        a_sd=0.4,
        b_sd=0.5,
        t_statistic=0.0,
        p_value=1.0,
        log2_fold_change=0.0,
    )
    assert ds.back_solved_replicates(flat) is None
    moved = flat.model_copy(update={"a_mean": 22.0, "t_statistic": 6.25})
    assert ds.back_solved_replicates(moved) == pytest.approx(0.41 * 6.25**2 / 4.0)


def test_a_row_with_no_spread_carries_no_information_either() -> None:
    """A zero SD on both sides divides by zero, so the row is skipped."""
    row = ds.ComparisonRow(
        label="PP_0001",
        accession="Q00001",
        description=None,
        a_mean=22.0,
        b_mean=20.0,
        a_sd=0.0,
        b_sd=0.0,
        t_statistic=0.0,
        p_value=1.0,
        log2_fold_change=2.0,
    )
    assert ds.back_solved_replicates(row) is None


def _synthetic_workbooks(n_promoter: int = 4, n_cross: int = 3) -> dict[str, list[Any]]:
    """One row per workbook, each exact for its arm's replicate count."""
    return {
        ds.WORKBOOK_PJ: [_row("PP_0001", 22.0, 20.0, 0.4, 0.5, n_promoter)],
        ds.WORKBOOK_0415: [_row("PP_0001", 21.0, 20.0, 0.4, 0.5, n_promoter)],
        ds.WORKBOOK_PAM_STRAINS: [_row("PP_0001", 19.0, 20.0, 0.4, 0.5, n_cross)],
        ds.WORKBOOK_PAM_MEDIA: [_row("PP_0001", 19.0, 18.0, 0.4, 0.5, n_cross)],
    }


def test_the_back_solve_passes_when_every_arm_states_its_own_count() -> None:
    """Four for the promoter workbooks and three for the cross-feeding ones."""
    checked = ds.assert_replicates_back_solve(_synthetic_workbooks())
    assert checked == dict.fromkeys(ds.WORKBOOK_GROUPS, 1)


def test_the_back_solve_refuses_a_workbook_on_the_other_arms_count() -> None:
    """A promoter workbook at n = 3 is refused, which is the whole point of the check."""
    workbooks = _synthetic_workbooks(n_promoter=3)
    with pytest.raises(RuntimeError, match="back-solves to"):
        ds.assert_replicates_back_solve(workbooks)


def test_the_merged_labels_are_asserted_at_twice_the_arms_count() -> None:
    """``Pyrc`` and ``Ubid`` pool both protein groups' replicates, so they sit at 2n."""
    workbooks = _synthetic_workbooks()
    for workbook, rows in workbooks.items():
        arm = next(
            g.arm
            for g in ds.PROTEOME_GROUPS
            if g.key == ds.WORKBOOK_GROUPS[workbook][0]
        )
        n = ds.ARM_REPLICATES[arm] * ds.MERGED_LABEL_REPLICATE_FACTOR
        rows.append(_row("Pyrc", 22.0, 20.0, 0.4, 0.5, n, accession="Q00002"))
        rows.append(_row("Pyrc", 22.0, 20.0, 0.4, 0.5, n, accession="Q00003"))
    assert ds.merged_accession_labels(workbooks) == {"Pyrc"}
    checked = ds.assert_replicates_back_solve(workbooks)
    assert checked == dict.fromkeys(ds.WORKBOOK_GROUPS, 3)


def test_a_merged_label_at_the_plain_count_is_refused() -> None:
    """The exception is a measured fact, not a license to skip the check."""
    workbooks = _synthetic_workbooks()
    for rows in workbooks.values():
        rows.append(_row("Pyrc", 22.0, 20.0, 0.4, 0.5, 4, accession="Q00002"))
        rows.append(_row("Pyrc", 22.0, 20.0, 0.4, 0.5, 4, accession="Q00003"))
    with pytest.raises(RuntimeError, match="merged-accession label"):
        ds.assert_replicates_back_solve(workbooks)


def test_the_back_solve_refuses_a_workbook_with_no_informative_row() -> None:
    """A workbook whose every row is flat proves nothing and is not silently passed."""
    flat = ds.ComparisonRow(
        label="PP_0001",
        accession="Q00001",
        description=None,
        a_mean=20.0,
        b_mean=20.0,
        a_sd=0.4,
        b_sd=0.5,
        t_statistic=0.0,
        p_value=1.0,
        log2_fold_change=0.0,
    )
    workbooks = _synthetic_workbooks()
    workbooks[ds.WORKBOOK_PJ] = [flat]
    with pytest.raises(RuntimeError, match="no row could back-solve"):
        ds.assert_replicates_back_solve(workbooks)


# --------------------------------------------------------------------------- #
# PP_0897's own fold changes identify the strains
# --------------------------------------------------------------------------- #
def _pp_0897_workbooks(
    pj: float = -4.0, weak: float = -3.3, deletion: float = -7.9, flat: float = 0.05
) -> dict[str, list[Any]]:
    """One PP_0897 row per workbook at the stated fold changes."""
    folds = {
        ds.WORKBOOK_PJ: pj,
        ds.WORKBOOK_0415: weak,
        ds.WORKBOOK_PAM_STRAINS: deletion,
        ds.WORKBOOK_PAM_MEDIA: flat,
    }
    return {
        workbook: [_row(ds.PP_0897, 20.0 + fold, 20.0, 0.4, 0.5, 4)]
        for workbook, fold in folds.items()
    }


def test_pp_0897_settles_which_group_is_which_strain() -> None:
    """Down in both titrations, further down in the deletion, flat across its media."""
    folds = ds.assert_pp_0897_identifies_the_strains(_pp_0897_workbooks())
    assert folds[ds.WORKBOOK_PJ] == pytest.approx(-4.0)
    assert folds[ds.WORKBOOK_PAM_MEDIA] == pytest.approx(0.05)


def test_a_promoter_group_whose_pp_0897_is_unchanged_is_refused() -> None:
    """A titration that did not titrate means the group is not the variant strain."""
    with pytest.raises(RuntimeError, match="strong reduction"):
        ds.assert_pp_0897_identifies_the_strains(_pp_0897_workbooks(pj=-0.1))


def test_a_deletion_group_no_further_down_than_a_titration_is_refused() -> None:
    """The deletion must be the extreme, or the assignment is the other way round."""
    with pytest.raises(RuntimeError, match="which is a titration"):
        ds.assert_pp_0897_identifies_the_strains(_pp_0897_workbooks(deletion=-3.0))


def test_a_deleted_gene_that_moves_between_two_media_is_refused() -> None:
    """A gene that is absent in both cultures cannot respond to the carbon source."""
    with pytest.raises(RuntimeError, match="which a deleted gene cannot do"):
        ds.assert_pp_0897_identifies_the_strains(_pp_0897_workbooks(flat=2.0))


def test_a_workbook_missing_pp_0897_is_refused() -> None:
    """The locus every assignment is read from is not allowed to be absent."""
    workbooks = _pp_0897_workbooks()
    workbooks[ds.WORKBOOK_PJ] = []
    with pytest.raises(RuntimeError, match="rows for PP_0897"):
        ds.assert_pp_0897_identifies_the_strains(workbooks)


# --------------------------------------------------------------------------- #
# The protein-count and shared-group oracles
# --------------------------------------------------------------------------- #
def test_a_changed_protein_count_stops_the_build(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The released row count is pinned, so a revision is detected, not absorbed."""
    workbooks = _synthetic_workbooks()
    monkeypatch.setattr(ds, "EXPECTED_PROTEIN_ROWS", dict.fromkeys(workbooks, 1))
    ds.assert_protein_counts(workbooks)
    monkeypatch.setattr(ds, "EXPECTED_PROTEIN_ROWS", dict.fromkeys(workbooks, 2))
    with pytest.raises(RuntimeError, match="released protein counts changed"):
        ds.assert_protein_counts(workbooks)


def _shared_workbooks() -> dict[str, list[Any]]:
    """Workbooks whose two twice-released groups agree cell for cell."""
    promoter_control = _row("PP_0001", 22.0, 20.0, 0.4, 0.5, 4)
    cross_subject = _row("PP_0001", 19.0, 20.0, 0.4, 0.5, 3)
    return {
        ds.WORKBOOK_PJ: [promoter_control],
        ds.WORKBOOK_0415: [promoter_control],
        ds.WORKBOOK_PAM_STRAINS: [cross_subject],
        ds.WORKBOOK_PAM_MEDIA: [cross_subject],
    }


def test_the_twice_released_groups_are_one_sample_when_they_agree() -> None:
    """This is what makes the record count six rather than eight."""
    ds.assert_shared_groups_agree(_shared_workbooks())


def test_a_twice_released_group_whose_cells_differ_stops_the_build() -> None:
    """If the pair differs the two comparisons no longer share a control sample."""
    workbooks = _shared_workbooks()
    workbooks[ds.WORKBOOK_0415] = [_row("PP_0001", 22.0, 20.5, 0.4, 0.5, 4)]
    with pytest.raises(RuntimeError, match="cells differ between"):
        ds.assert_shared_groups_agree(workbooks)


def test_a_twice_released_group_with_a_different_protein_set_stops_the_build() -> None:
    """A sample released twice over two protein sets is two samples, not one."""
    workbooks = _shared_workbooks()
    workbooks[ds.WORKBOOK_0415] = [_row("PP_0002", 22.0, 20.0, 0.4, 0.5, 4)]
    with pytest.raises(RuntimeError, match="different protein sets"):
        ds.assert_shared_groups_agree(workbooks)


def test_group_cells_refuses_two_released_rows_that_disagree() -> None:
    """A repeated label with two values would change the stored mean."""
    group = next(g for g in ds.PROTEOME_GROUPS if g.key == "2370_pJ_PP_0897")
    rows = [
        _row("PP_0001", 22.0, 20.0, 0.4, 0.5, 4),
        _row("PP_0001", 23.0, 20.0, 0.4, 0.5, 4),
    ]
    with pytest.raises(RuntimeError, match="two released rows disagree"):
        ds.group_cells(rows, group)


def test_group_cells_reads_the_b_column_for_a_second_position_group() -> None:
    """A group that is the workbook's B side reads the B mean, not the A mean."""
    rows = [_row("PP_0001", 22.0, 20.0, 0.4, 0.5, 4)]
    a_group = next(g for g in ds.PROTEOME_GROUPS if g.key == "2370_pJ_PP_0897")
    b_group = next(
        g
        for g in ds.PROTEOME_GROUPS
        if g.key == "2370" and g.workbook == ds.WORKBOOK_PJ
    )
    assert ds.group_cells(rows, a_group) == {"PP_0001": (22.0, 0.4)}
    assert ds.group_cells(rows, b_group) == {"PP_0001": (20.0, 0.5)}


# --------------------------------------------------------------------------- #
# The chassis background and the three genotype shapes
# --------------------------------------------------------------------------- #
def test_the_chassis_types_the_four_host_deletions_and_quotes_the_cassette() -> None:
    """Four full deletions with their annotation symbols; the cassette stays verbatim."""
    background = ds.chassis_background()
    assert background.name == ds.CHASSIS
    assert background.reference_strain == ds.WT_STRAIN
    assert background.assembly_set == ds.KT2440_ASSEMBLY_SET
    assert {a.systematic_gene_name: a.edit for a in background.alleles} == {
        tag: AlleleEdit.full_deletion for tag, _, _ in ds.CUTSET_LOCI
    }
    assert {a.gene_name for a in background.alleles} == {
        symbol for _, symbol, _ in ds.CUTSET_LOCI
    }
    assert all(a.functional is False for a in background.alleles)
    assert all(a.cassette is None for a in background.alleles)
    assert background.genotype_statement is not None
    assert "PP_5402-intergenic::PBAD-Sc.bpsA,Bc.sfp,Ps.glnA" in (
        background.genotype_statement
    )
    assert background.provenance is not None
    assert len(background.provenance) == 3


def test_the_chassis_does_not_type_the_three_heterologous_genes() -> None:
    """Their source organisms are never stated, so no allele asserts one."""
    background = ds.chassis_background()
    names = {a.systematic_gene_name for a in background.alleles}
    assert not names & {"Sc.bpsA", "Bc.sfp", "Ps.glnA", "bpsA", "sfp", "glnA"}
    assert ds.PATHWAY_NOT_TYPED.value == ("Sc.bpsA", "Bc.sfp", "Ps.glnA")
    assert ds.PATHWAY_NOT_TYPED.note is not None
    assert "never expand" in ds.PATHWAY_NOT_TYPED.note


def test_the_parental_genotype_carries_no_perturbation() -> None:
    """Everything a parental sample has is the background."""
    assert ds.strain_genotype(ds.STRAIN_PARENT).perturbations == []


def test_the_deletion_genotype_is_one_bacterial_deletion_on_pp_0897() -> None:
    genotype = ds.strain_genotype(ds.STRAIN_DELETION)
    (perturbation,) = genotype.perturbations
    assert perturbation.perturbation_type == "bacterial_deletion"
    assert perturbation.systematic_gene_name == ds.PP_0897
    assert perturbation.gene_namespace == ds.KT2440_NAMESPACE
    assert perturbation.cassette is None


@pytest.mark.parametrize(
    ("strain", "promoter"),
    [
        ("D1b_gf PJ23109-PP_0897", "pJ23109"),
        ("D1b_gf Ppp_0415-PP_0897", "PP_0415 promoter"),
    ],
)
def test_a_promoter_variant_is_a_promoter_replacement_not_a_deletion(
    strain: str, promoter: str
) -> None:
    """The gene stays present and unedited; only the regulatory part changed."""
    (perturbation,) = ds.strain_genotype(strain).perturbations
    assert perturbation.perturbation_type == "promoter_replacement"
    assert perturbation.systematic_gene_name == ds.PP_0897
    assert perturbation.state == "present"
    assert perturbation.mechanism_so_id == "SO:1000032"
    assert perturbation.promoter_name == promoter
    assert perturbation.native_promoter == "endogenous PP_0897 promoter"
    assert perturbation.expression_direction == "decreased"
    assert perturbation.is_inducible is False
    assert perturbation.crispr is None


def test_the_promoter_builder_refuses_a_strain_that_has_no_swap() -> None:
    with pytest.raises(RuntimeError, match="not a promoter-variant strain"):
        ds.promoter_perturbation(ds.STRAIN_DELETION)


def test_the_publication_carries_the_doi_and_no_invented_pubmed_id() -> None:
    publication = ds.publication()
    assert publication.doi == ds.DOI
    assert publication.doi_url == f"https://doi.org/{ds.DOI}"
    assert publication.pubmed_id is None


# --------------------------------------------------------------------------- #
# The three environments
# --------------------------------------------------------------------------- #
def _carbon_sources(condition: str) -> list[EnvironmentPhysicalPerturbation]:
    """The carbon-source factors of one condition, narrowed to the physical leaf."""
    environment = ds.proteome_environment(condition)
    carbon = [
        p
        for p in environment.perturbations
        if isinstance(p, EnvironmentPhysicalPerturbation)
    ]
    assert len(carbon) == len(environment.perturbations)
    assert all(p.factor is PhysicalFactor.carbon_source for p in carbon)
    return carbon


def test_the_pca_only_environment_carries_one_sourced_carbon_source() -> None:
    environment = ds.proteome_environment(ds.CONDITION_PCA)
    assert environment.media is M9_DEFERRED_BANERJEE2025
    (carbon,) = _carbon_sources(ds.CONDITION_PCA)
    assert carbon.agent is not None
    assert carbon.agent.name == "p-coumaric acid"
    assert carbon.magnitude is not None
    assert carbon.magnitude.value == 60.0
    assert carbon.magnitude.unit is ConcentrationUnit.millimolar
    assert carbon.provenance_gaps == []


def test_the_supplemented_environment_carries_all_three_sourced_doses() -> None:
    carbon = _carbon_sources(ds.CONDITION_PCA_ALA_MAL)
    doses: dict[str, float | None] = {}
    for perturbation in carbon:
        assert perturbation.agent is not None
        assert perturbation.magnitude is not None
        doses[perturbation.agent.name] = perturbation.magnitude.value
    assert doses == {"p-coumaric acid": 50.0, "D-alanine": 70.0, "L-malate": 70.0}
    assert all(p.provenance_gaps == [] for p in carbon)


def test_the_undescribed_medium_gaps_each_dose_rather_than_borrowing_one() -> None:
    """The Methods never describe this medium, so no magnitude is asserted."""
    carbon = _carbon_sources(ds.CONDITION_ALA_MAL)
    names = set()
    for perturbation in carbon:
        assert perturbation.agent is not None
        names.add(perturbation.agent.name)
        assert perturbation.magnitude is None
        (gap,) = perturbation.provenance_gaps
        assert gap.field == "magnitude"
        assert gap.note == ds.UNDOSED_MEDIUM_NOTE
    assert names == {"D-alanine", "L-malate"}


@pytest.mark.parametrize(
    "condition", [ds.CONDITION_PCA, ds.CONDITION_PCA_ALA_MAL, ds.CONDITION_ALA_MAL]
)
def test_every_environment_gaps_the_temperature_and_the_duration(
    condition: str,
) -> None:
    """The shotgun-proteomics Methods state neither, so neither is filled in."""
    environment = ds.proteome_environment(condition)
    assert environment.temperature is None
    assert environment.duration_hours is None
    assert environment.aerobicity == "aerobic"
    fields = {gap.field for gap in environment.provenance_gaps}
    assert {"temperature", "duration_hours"} <= fields


def test_proteome_environment_refuses_an_unreleased_condition() -> None:
    with pytest.raises(RuntimeError, match="not a released proteomics condition"):
        ds.proteome_environment("M9 glucose")


# --------------------------------------------------------------------------- #
# The phenotype
# --------------------------------------------------------------------------- #
def test_the_phenotype_divides_the_released_sd_by_the_root_of_its_arms_count() -> None:
    cells = {"PP_0001": (22.0, 0.4), "PP_0002": (20.0, 0.9)}
    phenotype = ds.ProteomeBanerjee2025Dataset._phenotype(cells, 4)
    assert phenotype.measurement_type == ds.MEASUREMENT_TYPE
    assert phenotype.protein_abundance == {"PP_0001": 22.0, "PP_0002": 20.0}
    assert phenotype.protein_abundance_se is not None
    assert phenotype.protein_abundance_se["PP_0002"] == pytest.approx(0.9 / 2.0)
    assert set(phenotype.n_replicates.values()) == {4}


def test_the_phenotype_refuses_a_sample_with_no_measured_protein() -> None:
    with pytest.raises(RuntimeError, match="no measured protein"):
        ds.ProteomeBanerjee2025Dataset._phenotype({}, 4)


# --------------------------------------------------------------------------- #
# The retention arithmetic
# --------------------------------------------------------------------------- #
def _accounting(**overrides: Any) -> ds.BuildAccounting:
    from torchcell.datasets.bacteria_common import LocusTagReconciliation

    reconciliation = LocusTagReconciliation(
        label="x",
        assembly_set=ds.KT2440_ASSEMBLY_SET,
        gene_namespace=ds.KT2440_NAMESPACE,
        unique_names=1,
        status_histogram={},
        layer_histogram={},
        remapped=0,
        kept_on_collision=(),
        retired_kept=(),
        ambiguous_kept={},
        case_insensitive=(),
        outside_namespace=(),
    )
    fields: dict[str, Any] = {
        "dataset": "x",
        "source_rows": 10,
        "candidate_records": 6,
        "kept_records": 6,
        "dropped_records": 0,
        "rules": [],
        "reconciliation": reconciliation,
        "notes": [],
    }
    fields.update(overrides)
    return ds.BuildAccounting(**fields)


def test_the_accounting_refuses_an_arithmetic_that_loses_a_record() -> None:
    _accounting().check()
    with pytest.raises(RuntimeError, match="kept"):
        _accounting(kept_records=5).check()


def test_the_accounting_refuses_a_rule_with_a_negative_count() -> None:
    rule = ds.DropRule(
        rule="r", scope="protein_label", description="d", n_records=-1, items=[]
    )
    with pytest.raises(RuntimeError, match="negative record count"):
        _accounting(rules=[rule]).check()


def test_a_protein_label_rule_does_not_count_toward_the_record_arithmetic() -> None:
    """The labels are dropped from inside every record, not as records."""
    rule = ds.DropRule(
        rule="protein_label_merges_two_accessions",
        scope="protein_label",
        description="d",
        n_records=0,
        items=["Pyrc", "Ubid"],
    )
    accounting = _accounting(rules=[rule])
    accounting.check()
    assert accounting.dropped_records == 0
    assert accounting.kept_records == accounting.candidate_records
    assert [r.items for r in accounting.rules] == [["Pyrc", "Ubid"]]


# --------------------------------------------------------------------------- #
# Hermetic end to end: a synthetic KT2440 assembly, a synthetic archive, and the
# loader built under tmp_path. No network, no $DATA_ROOT.
# --------------------------------------------------------------------------- #
#: The loci the synthetic assembly carries: PP_0897, the four cutset loci, and enough
#: more that the resolved fraction clears MIN_RESOLVED_FRACTION.
LOCUS_SPECS: tuple[tuple[str, str | None], ...] = (
    (ds.PP_0897, None),
    *((tag, symbol) for tag, symbol, _ in ds.CUTSET_LOCI),
    ("PP_0100", None),
    ("PP_0200", None),
    ("PP_0300", None),
    ("PP_0400", None),
    ("PP_0500", None),
    ("PP_0600", None),
    ("PP_0700", None),
    ("PP_0800", None),
    ("PP_1000", None),
    ("PP_1100", None),
    ("PP_1200", None),
    ("PP_1300", None),
    ("PP_1500", None),
    ("PP_1600", None),
    ("PP_1700", None),
    ("PP_1800", None),
    ("PP_1900", None),
    ("PP_1400", "ygiQ"),
)
#: The tags the synthetic sheets key on, all of which resolve.
SYNTHETIC_TAGS: tuple[str, ...] = (
    ds.PP_0897,
    "PP_0100",
    "PP_0200",
    "PP_0300",
    "PP_0400",
    "PP_0500",
    "PP_0600",
    "PP_0700",
    "PP_0800",
    "PP_1000",
    "PP_1100",
    "PP_1200",
    "PP_1300",
    "PP_1500",
    "PP_1600",
    "PP_1700",
    "PP_1800",
    "PP_1900",
)
#: A symbol that resolves, a symbol no layer resolves, and a label filed twice.
SYNTHETIC_SYMBOL = "Ygiq"
SYNTHETIC_UNRESOLVED = "Krt1"
SYNTHETIC_MERGED = "Pyrc"
SYNTHETIC_LABELS: tuple[str, ...] = (
    *SYNTHETIC_TAGS,
    SYNTHETIC_SYMBOL,
    SYNTHETIC_UNRESOLVED,
    SYNTHETIC_MERGED,
)
#: The labels that reach a record's abundance map: the tags plus the resolved symbol.
SYNTHETIC_KEPT: tuple[str, ...] = (*SYNTHETIC_TAGS, "PP_1400")
#: Rows the synthetic workbooks hold, one per label plus the merged label's second row.
SYNTHETIC_ROWS = len(SYNTHETIC_LABELS) + 1
#: The locus-tag CASE each workbook writes, which is the real release's own split.
WORKBOOK_CASE: dict[str, bool] = {
    ds.WORKBOOK_PJ: True,
    ds.WORKBOOK_0415: True,
    ds.WORKBOOK_PAM_STRAINS: False,
    ds.WORKBOOK_PAM_MEDIA: True,
}
#: PP_0897's log2 mean PER GROUP, set so that every workbook's fold change has the
#: real release's sign and magnitude while the two twice-released groups stay identical
#: across their workbooks. Writing a fold change per WORKBOOK instead would make the
#: shared column differ, which is exactly what ``assert_shared_groups_agree`` refuses.
GROUP_PP_0897_MEAN: dict[str, float] = {
    "2370": 24.9,
    "2370_pJ_PP_0897": 20.9,
    "2370_0415pPP_0897": 21.6,
    "2370_M9_pCA_alanine_malate": 26.25,
    "2487_M9_pCA_alanine_malate": 18.35,
    "2487_M9_alanine_malate": 18.3,
}
#: The fold change each workbook therefore carries for PP_0897.
WORKBOOK_PP_0897_FOLD: dict[str, float] = {
    workbook: GROUP_PP_0897_MEAN[group_a] - GROUP_PP_0897_MEAN[group_b]
    for workbook, (group_a, group_b) in ds.WORKBOOK_GROUPS.items()
}
FULL_HEADER: tuple[str, ...] = (
    "Protein",
    "Protein.Group",
    "Protein.Names",
    "Protein.Description",
    "",
    "",
    "",
    "",
    "t-test_stat",
    "p-value",
    "p_adjusted(BH)",
    "log2_Fold_change_A/B",
)


def _synthetic_loci() -> list[Any]:
    """One ``SyntheticLocus`` per :data:`LOCUS_SPECS` entry, laid end to end."""
    from tests.torchcell.sequence.genome._bacterial_fixtures import SyntheticLocus

    loci = []
    cursor = 1
    for index, (tag, symbol) in enumerate(LOCUS_SPECS):
        start, end = cursor, cursor + 11
        cursor = end + 3
        loci.append(
            SyntheticLocus(
                tag=tag,
                parts=((start, end),),
                strand="+" if index % 2 == 0 else "-",
                symbol=symbol,
                product=f"synthetic product {tag}",
                protein_id=f"AAN{index:05d}.1",
                protein="MKV",
            )
        )
    return loci


@pytest.fixture
def synthetic_kt2440(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """The real KT2440 class over a synthetic assembly, with the network refused."""
    import tests.torchcell.sequence.genome._bacterial_fixtures as fixtures
    from torchcell.sequence.genome.pputida.kt2440 import (
        KT2440_ASSEMBLY,
        PPutidaKT2440Genome,
    )

    monkeypatch.setattr(fixtures, "SEQUENCE", fixtures.SEQUENCE * 6)
    files = fixtures.write_assembly(
        tmp_path / "tier",
        KT2440_ASSEMBLY,
        _synthetic_loci(),
        [fixtures.gaf_row("ygiQ", f"ygiQ|{LOCUS_SPECS[-1][0]}", "GO:0000001")],
    )
    fixtures.forbid_network(monkeypatch)
    fixtures.serve_tier(monkeypatch, files)
    _serve_assembly_report(tmp_path, monkeypatch)
    root = tmp_path / "kt2440"
    root.mkdir()
    return PPutidaKT2440Genome(genome_root=str(root), overwrite=False)


def _cell(label: str, group: str, index: int) -> tuple[float, float]:
    """One group's ``(log2 mean, log2 SD)`` for one label, stable across workbooks."""
    return (20.0 + index / 10.0 + len(group) / 100.0, 0.1 + index / 100.0)


def _write_workbook(path: Path, workbook: str) -> None:
    """A synthetic ``Full t-test output`` sheet for one released comparison.

    The t statistic is computed exactly, at the arm's own replicate count for every
    label and at twice it for the merged label, so the back-solve and its one asserted
    exception both hold on the synthetic bytes as they do on the pinned ones.
    """
    group_a, group_b = ds.WORKBOOK_GROUPS[workbook]
    arm = next(g.arm for g in ds.PROTEOME_GROUPS if g.key == group_a)
    n = ds.ARM_REPLICATES[arm]
    header = list(FULL_HEADER)
    header[4] = f"log2_mean_{group_a}"
    header[5] = f"log2_mean_{group_b}"
    header[6] = f"log2_std_{group_a}"
    header[7] = f"log2_std_{group_b}"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = ds.FULL_SHEET
    sheet.append(header)
    for index, label in enumerate(SYNTHETIC_LABELS):
        accessions = (
            ("Q90001", "Q90002") if label == SYNTHETIC_MERGED else (f"Q{index:05d}",)
        )
        written = label if not label.startswith("PP_") else label
        if label.startswith("PP_") and not WORKBOOK_CASE[workbook]:
            written = label.capitalize()
        a_mean, a_sd = _cell(label, group_a, index)
        b_mean, b_sd = _cell(label, group_b, index)
        if label == ds.PP_0897:
            a_mean = GROUP_PP_0897_MEAN[group_a]
            b_mean = GROUP_PP_0897_MEAN[group_b]
        replicates = n * (ds.MERGED_LABEL_REPLICATE_FACTOR if accessions[1:] else 1)
        t = (a_mean - b_mean) / math.sqrt((a_sd**2 + b_sd**2) / replicates)
        for accession in accessions:
            sheet.append(
                [
                    written,
                    accession,
                    f"{label.upper()}_PSEPK",
                    f"synthetic {label}",
                    a_mean,
                    b_mean,
                    a_sd,
                    b_sd,
                    t,
                    0.01,
                    0.05,
                    a_mean - b_mean,
                ]
            )
    if workbook in ds.EXPECTED_SIGNIFICANT:
        total = ds.EXPECTED_SIGNIFICANT[workbook]
        up = total // 2
        for name, count in zip(ds.SIGNIFICANCE_SHEETS, (up, total - up), strict=True):
            significance = book.create_sheet(name)
            significance.append(list(header))
            for index in range(count):
                significance.append([f"PP_{index:04d}", *([None] * (len(header) - 1))])
    book.save(path)
    book.close()


def _write_analysis(path: Path) -> None:
    """A synthetic curated sheet: 134 rows of which 13 carry both fold changes."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = ds.ANALYSIS_SHEET
    sheet.append(["Protein", "log2FC pJ23109", "log2FC PP_0415p", "Subsystem"])
    for index in range(ds.EXPECTED_ANALYSIS_ROWS):
        both = index < ds.EXPECTED_COMMON_PROTEINS
        sign = 1.0 if index % 2 == 0 else -1.0
        sheet.append(
            [
                f"PP_{index:04d}",
                sign * (3.0 + index / 100.0) if both or index % 3 else None,
                sign * (3.0 + index / 100.0) if both else None,
                "synthetic",
            ]
        )
    book.save(path)
    book.close()


def _sha256_bytes(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def synthetic_mirror(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A raw mirror under ``tmp_path`` built by the module's own deposit function."""
    staging = tmp_path / "staging"
    staging.mkdir()
    members: dict[str, tuple[str, str]] = {}
    archive_path = tmp_path / ds.ARCHIVE_FILENAME
    with zipfile.ZipFile(archive_path, "w") as archive:
        for relpath, (member, _) in ds.RAW_MEMBERS.items():
            name = osp.basename(relpath)
            built = staging / name
            if name == ds.WORKBOOK_ANALYSIS:
                _write_analysis(built)
            else:
                _write_workbook(built, name)
            archive.write(built, member)
            members[relpath] = (member, _sha256_bytes(built))
    monkeypatch.setattr(ds, "RAW_MEMBERS", members)
    monkeypatch.setattr(ds, "ARCHIVE_SHA256", _sha256_bytes(archive_path))
    monkeypatch.setattr(
        ds, "EXPECTED_PROTEIN_ROWS", dict.fromkeys(ds.WORKBOOK_GROUPS, SYNTHETIC_ROWS)
    )
    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    root = ds.deposit_raw_mirror(archive_path=archive_path, data_root=str(data_root))
    assert root == ds.raw_mirror_dir(str(data_root))
    return data_root


# --- the deposit itself ---------------------------------------------------- #
def test_the_deposit_writes_every_member_and_a_manifest_that_pins_them(
    synthetic_mirror: Path,
) -> None:
    manifest = ds.load_manifest(str(synthetic_mirror))
    assert manifest.citation_key == ds.CITATION_KEY
    assert manifest.doi == ds.DOI
    assert {record.path for record in manifest.files} == set(ds.RAW_MEMBERS)
    for record in manifest.files:
        assert record.role == ROLE_SI_DATA
        assert record.sha256 == ds.RAW_MEMBERS[record.path][1]
        assert record.retrieval is not None
        assert record.retrieval.method is RetrievalMethod.pmc_cloud
        assert record.retrieval.params["container_sha256"] == ds.ARCHIVE_SHA256
        assert record.retrieval.params["member"] == ds.RAW_MEMBERS[record.path][0]
        on_disk = synthetic_mirror / ds.RAW_DIR_REL / record.path
        assert _sha256_bytes(on_disk) == record.sha256


def test_the_manifest_enumerates_what_it_deliberately_does_not_hold(
    synthetic_mirror: Path,
) -> None:
    """The model archive, the SI PDF, PRIDE, the figure-only readouts and BIOLOG."""
    manifest = ds.load_manifest(str(synthetic_mirror))
    blob = " ".join(manifest.si_expected)
    assert ds.MODEL_ARCHIVE_FILENAME in blob
    assert ds.PRIDE_ACCESSION in blob
    assert "TITER_NOT_A_DATASET" in blob
    assert "GROWTH_NOT_A_DATASET" in blob
    assert "BIOLOG_NOT_A_DATASET" in blob
    assert ds.ARCHIVE_URL in manifest.si_data_sources


def test_the_deposit_is_idempotent_by_sha256(
    synthetic_mirror: Path, tmp_path: Path
) -> None:
    before = {
        path: _sha256_bytes(path)
        for path in (synthetic_mirror / ds.RAW_DIR_REL / "si").iterdir()
    }
    ds.deposit_raw_mirror(
        archive_path=tmp_path / ds.ARCHIVE_FILENAME, data_root=str(synthetic_mirror)
    )
    after = {
        path: _sha256_bytes(path)
        for path in (synthetic_mirror / ds.RAW_DIR_REL / "si").iterdir()
    }
    assert after == before


def test_the_deposit_refuses_a_mirror_file_whose_bytes_differ(
    synthetic_mirror: Path, tmp_path: Path
) -> None:
    target = (
        synthetic_mirror / ds.RAW_DIR_REL / "si" / osp.basename(ds.WORKBOOK_ANALYSIS)
    )
    target.write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="exists with a different sha256"):
        ds.deposit_raw_mirror(
            archive_path=tmp_path / ds.ARCHIVE_FILENAME, data_root=str(synthetic_mirror)
        )


def test_the_deposit_refuses_an_archive_off_its_pin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    archive = tmp_path / "wrong.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("Data S2/x", "x")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "dr"))
    with pytest.raises(RawSha256MismatchError, match="sha256 mismatch"):
        ds.deposit_raw_mirror(archive_path=archive, data_root=str(tmp_path / "dr"))


def test_the_deposit_refuses_a_member_off_its_pin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    archive = tmp_path / "members.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        for member, _ in ds.RAW_MEMBERS.values():
            handle.writestr(member, "x")
    monkeypatch.setattr(ds, "ARCHIVE_SHA256", _sha256_bytes(archive))
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "dr"))
    with pytest.raises(RuntimeError, match="sha256 mismatch"):
        ds.deposit_raw_mirror(archive_path=archive, data_root=str(tmp_path / "dr"))


def test_raw_mirror_dir_reads_data_root_from_the_environment(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setenv("DATA_ROOT", str(tmp_path))
    assert ds.raw_mirror_dir() == tmp_path / ds.RAW_DIR_REL
    assert ds._data_root() == str(tmp_path)


def test_manifest_sha256_refuses_a_path_the_manifest_does_not_hold(
    synthetic_mirror: Path,
) -> None:
    manifest = ds.load_manifest(str(synthetic_mirror))
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        ds.manifest_sha256(manifest, "si/absent.xlsx")


# --- the reader ------------------------------------------------------------- #
def test_read_comparison_types_every_released_cell(synthetic_mirror: Path) -> None:
    path = str(
        synthetic_mirror / ds.RAW_DIR_REL / "si" / osp.basename(ds.WORKBOOK_PAM_STRAINS)
    )
    rows = ds.read_comparison(path, *ds.WORKBOOK_GROUPS[ds.WORKBOOK_PAM_STRAINS])
    assert len(rows) == SYNTHETIC_ROWS
    assert {row.label for row in rows} == set(SYNTHETIC_LABELS)
    row = next(r for r in rows if r.label == ds.PP_0897)
    assert row.log2_fold_change == pytest.approx(
        WORKBOOK_PP_0897_FOLD[ds.WORKBOOK_PAM_STRAINS]
    )
    assert row.description == f"synthetic {ds.PP_0897}"


def test_read_comparison_refuses_a_changed_header(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = ds.FULL_SHEET
    sheet.append(["Protein", "Protein.Group"])
    sheet.append(["PP_0001", "Q1"])
    book.save(path)
    book.close()
    with pytest.raises(RuntimeError, match="carries no"):
        ds.read_comparison(str(path), "2370_pJ_PP_0897", "2370")


def test_read_comparison_refuses_an_empty_release(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    header = list(FULL_HEADER)
    header[4] = "log2_mean_2370_pJ_PP_0897"
    header[5] = "log2_mean_2370"
    header[6] = "log2_std_2370_pJ_PP_0897"
    header[7] = "log2_std_2370"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = ds.FULL_SHEET
    sheet.append(header)
    book.save(path)
    book.close()
    with pytest.raises(RuntimeError, match="holds no protein row"):
        ds.read_comparison(str(path), "2370_pJ_PP_0897", "2370")


def test_read_comparison_refuses_a_row_missing_an_abundance(tmp_path: Path) -> None:
    path = tmp_path / "partial.xlsx"
    header = list(FULL_HEADER)
    header[4] = "log2_mean_2370_pJ_PP_0897"
    header[5] = "log2_mean_2370"
    header[6] = "log2_std_2370_pJ_PP_0897"
    header[7] = "log2_std_2370"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = ds.FULL_SHEET
    sheet.append(header)
    sheet.append(["PP_0001", "Q1", "N", "D", 22.0, None, 0.4, 0.5, 1.0, 0.1, 0.2, 2.0])
    book.save(path)
    book.close()
    with pytest.raises(RuntimeError, match="missing a released abundance"):
        ds.read_comparison(str(path), "2370_pJ_PP_0897", "2370")


# --- the two text-reproducing counts --------------------------------------- #
def test_the_significance_sheets_reproduce_the_results_text(
    synthetic_mirror: Path,
) -> None:
    raw = str(synthetic_mirror / ds.RAW_DIR_REL / "si")
    assert ds.assert_significant_counts_match_the_text(raw) == ds.EXPECTED_SIGNIFICANT


def test_a_changed_significance_count_stops_the_build(
    synthetic_mirror: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = str(synthetic_mirror / ds.RAW_DIR_REL / "si")
    monkeypatch.setattr(
        ds, "EXPECTED_SIGNIFICANT", {ds.WORKBOOK_PJ: 1, ds.WORKBOOK_0415: 1}
    )
    with pytest.raises(RuntimeError, match="no longer reproduce the"):
        ds.assert_significant_counts_match_the_text(raw)


def test_the_curated_sheet_reproduces_the_papers_own_arithmetic(
    synthetic_mirror: Path,
) -> None:
    raw = str(synthetic_mirror / ds.RAW_DIR_REL / "si")
    assert ds.assert_curated_overlap(raw) == (
        ds.EXPECTED_ANALYSIS_ROWS,
        ds.EXPECTED_COMMON_PROTEINS,
    )


def test_a_curated_sheet_with_the_wrong_row_count_is_refused(
    synthetic_mirror: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = str(synthetic_mirror / ds.RAW_DIR_REL / "si")
    monkeypatch.setattr(ds, "EXPECTED_ANALYSIS_ROWS", 7)
    with pytest.raises(RuntimeError, match="rows, not the"):
        ds.assert_curated_overlap(raw)


def test_a_curated_sheet_with_the_wrong_overlap_is_refused(
    synthetic_mirror: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = str(synthetic_mirror / ds.RAW_DIR_REL / "si")
    monkeypatch.setattr(ds, "EXPECTED_COMMON_PROTEINS", 7)
    with pytest.raises(RuntimeError, match="carry both arms"):
        ds.assert_curated_overlap(raw)


def test_a_common_protein_that_moves_in_opposite_directions_is_refused(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = tmp_path / "si"
    raw.mkdir()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = ds.ANALYSIS_SHEET
    sheet.append(["Protein", "log2FC pJ23109", "log2FC PP_0415p"])
    sheet.append(["PP_0001", 3.0, -3.0])
    book.save(raw / osp.basename(ds.WORKBOOK_ANALYSIS))
    book.close()
    monkeypatch.setattr(ds, "EXPECTED_ANALYSIS_ROWS", 1)
    monkeypatch.setattr(ds, "EXPECTED_COMMON_PROTEINS", 1)
    with pytest.raises(RuntimeError, match="OPPOSITE directions"):
        ds.assert_curated_overlap(str(raw))


# --- the loader built end to end ------------------------------------------- #
@pytest.fixture
def built(synthetic_mirror: Path, synthetic_kt2440: Any, tmp_path: Path) -> Any:
    """The proteome loader built over the synthetic mirror and annotation."""
    dataset = ds.ProteomeBanerjee2025Dataset(
        root=str(tmp_path / "build"), pputida_genome=synthetic_kt2440
    )
    yield dataset
    dataset.close_lmdb()


def test_the_loader_writes_one_record_per_released_sample(built: Any) -> None:
    assert len(built) == ds.EXPECTED_PROTEOME_RECORDS
    assert built.experiment_class is BacterialProteinAbundanceExperiment
    assert built.reference_class is BacterialProteinAbundanceExperimentReference
    assert built.raw_file_names == [osp.basename(relpath) for relpath in ds.RAW_MEMBERS]


def test_every_record_shares_the_d1bgf_background(built: Any) -> None:
    """The four strains differ by their own PP_0897 edit, never by their background."""
    for index in range(len(built)):
        reference = built[index]["reference"]["genome_reference"]
        assert reference["strain"] == ds.CHASSIS
        assert reference["assembly_set"] == ds.KT2440_ASSEMBLY_SET
        assert len(reference["background"]["alleles"]) == len(ds.CUTSET_LOCI)


def test_the_records_drop_the_unresolvable_and_the_merged_label(built: Any) -> None:
    """The symbol that resolves survives; the retired one and the merged one do not."""
    abundance = built[0]["experiment"]["phenotype"]["protein_abundance"]
    assert set(abundance) == set(SYNTHETIC_KEPT)
    assert SYNTHETIC_UNRESOLVED not in abundance
    assert SYNTHETIC_MERGED not in abundance
    reference = built[0]["reference"]["phenotype_reference"]["protein_abundance"]
    assert set(reference) == set(abundance)


def test_the_two_arms_keep_their_own_replicate_counts(built: Any) -> None:
    """Four for the three promoter records and three for the three cross-feeding ones."""
    counts: dict[int, int] = {}
    for index in range(len(built)):
        phenotype = built[index]["experiment"]["phenotype"]
        n = next(iter(phenotype["n_replicates"].values()))
        counts[n] = counts.get(n, 0) + 1
        assert phenotype["measurement_type"] == ds.MEASUREMENT_TYPE
    assert counts == {4: 3, 3: 3}


def test_the_stored_abundance_is_the_released_cell_with_no_arithmetic(
    built: Any,
) -> None:
    """The mean is copied; only the SD is divided, by the root of the arm's count."""
    index = SYNTHETIC_LABELS.index("PP_0100")
    group = next(g for g in ds.PROTEOME_GROUPS if g.key == "2370_pJ_PP_0897")
    mean, sd = _cell("PP_0100", group.key, index)
    for record_index in range(len(built)):
        record = built[record_index]
        phenotype = record["experiment"]["phenotype"]
        if phenotype["protein_abundance"]["PP_0100"] == pytest.approx(mean):
            assert phenotype["protein_abundance_se"]["PP_0100"] == pytest.approx(
                sd / math.sqrt(4)
            )
            return
    raise AssertionError("no record carries the promoter-variant cell")


def test_the_perturbations_are_two_promoter_swaps_and_two_deletions(built: Any) -> None:
    kinds: dict[str, int] = {}
    for index in range(len(built)):
        for perturbation in built[index]["experiment"]["genotype"]["perturbations"]:
            kinds[perturbation["perturbation_type"]] = (
                kinds.get(perturbation["perturbation_type"], 0) + 1
            )
            assert perturbation["systematic_gene_name"] == ds.PP_0897
    assert kinds == {"promoter_replacement": 2, "bacterial_deletion": 2}


def test_the_build_writes_its_accounting_samples_and_oracles(built: Any) -> None:
    accounting = json.loads(
        Path(built.preprocess_dir, "build_accounting.json").read_text()
    )
    assert accounting["kept_records"] == ds.EXPECTED_PROTEOME_RECORDS
    assert accounting["dropped_records"] == 0
    assert accounting["source_rows"] == SYNTHETIC_ROWS * len(ds.WORKBOOK_GROUPS)
    rules = {rule["rule"]: rule for rule in accounting["rules"]}
    assert rules["protein_label_merges_two_accessions"]["items"] == [SYNTHETIC_MERGED]
    assert (
        SYNTHETIC_UNRESOLVED
        in (rules["protein_label_is_not_a_locus_of_the_pinned_assembly"]["items"])
    )
    oracles = json.loads(Path(built.preprocess_dir, "build_oracles.json").read_text())
    assert oracles["replicates_per_arm"] == {ds.ARM_PROMOTER: 4, ds.ARM_CROSSFEED: 3}
    assert oracles["curated_sheet_rows"] == ds.EXPECTED_ANALYSIS_ROWS
    assert (
        oracles["proteins_changed_in_both_promoter_variants"]
        == ds.EXPECTED_COMMON_PROTEINS
    )
    samples = Path(built.preprocess_dir, "samples.csv").read_text()
    assert samples.count("\n") == ds.EXPECTED_PROTEOME_RECORDS + 1
    for group in ds.PROTEOME_GROUPS:
        assert group.key in samples
    dropped = Path(built.preprocess_dir, "dropped_protein_labels.csv").read_text()
    assert SYNTHETIC_MERGED in dropped
    assert SYNTHETIC_UNRESOLVED in dropped


def test_the_build_refuses_a_release_whose_labels_all_fall_outside(
    synthetic_mirror: Path,
    synthetic_kt2440: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A keying change is a refusal, never a quietly smaller record."""
    monkeypatch.setattr(ds, "MIN_RESOLVED_FRACTION", 1.0)
    with pytest.raises(Exception, match="resolve"):
        ds.ProteomeBanerjee2025Dataset(
            root=str(tmp_path / "refused"), pputida_genome=synthetic_kt2440
        )


# --- the verification rules ------------------------------------------------ #
def _records(built: Any) -> list[dict[str, Any]]:
    return [
        json.loads(json.dumps(built[index], default=str)) for index in range(len(built))
    ]


def test_the_verification_rules_pass_on_the_hermetic_build(
    built: Any, synthetic_kt2440: Any
) -> None:
    records = _records(built)
    loci = set(synthetic_kt2440.genbank.loci)
    for result in (
        ds.group_uniqueness_rule(records),
        ds.replicate_split_rule(records),
        ds.assembly_pin_rule(records),
        ds.gene_containment_rule(records, loci),
        ds.stored_scale_rule(records, built.raw_dir),
    ):
        assert result.passed, (result.name, result.message)


def test_group_uniqueness_refuses_a_duplicated_sample(built: Any) -> None:
    records = _records(built)
    result = ds.group_uniqueness_rule([records[0], records[0]])
    assert result.passed is False
    assert "duplicated" in result.message


def test_replicate_split_refuses_a_count_no_arm_states(built: Any) -> None:
    records = _records(built)
    for key in records[0]["experiment"]["phenotype"]["n_replicates"]:
        records[0]["experiment"]["phenotype"]["n_replicates"][key] = 7
    result = ds.replicate_split_rule(records)
    assert result.passed is False


def test_assembly_pin_refuses_a_record_on_another_background(built: Any) -> None:
    records = _records(built)
    records[0]["reference"]["genome_reference"]["strain"] = "KT2440"
    assert ds.assembly_pin_rule(records).passed is False


def test_gene_containment_refuses_a_stored_key_off_the_assembly(built: Any) -> None:
    records = _records(built)
    assert ds.gene_containment_rule(records, {"PP_9999"}).passed is False


def test_stored_scale_refuses_a_value_that_is_not_a_released_cell(built: Any) -> None:
    records = _records(built)
    key = next(iter(records[0]["experiment"]["phenotype"]["protein_abundance"]))
    records[0]["experiment"]["phenotype"]["protein_abundance"][key] = -999.0
    assert ds.stored_scale_rule(records, built.raw_dir).passed is False


def test_the_provenance_names_the_released_artifact_and_the_two_counts() -> None:
    provenance = ds._provenance()
    assert provenance.citation_key == ds.CITATION_KEY
    assert provenance.sha256 == ds.RAW_MEMBERS[f"si/{ds.WORKBOOK_PJ}"][1]
    assert provenance.method is not None
    assert "n = 4" in provenance.method
    assert "n = 3" in provenance.method


# --------------------------------------------------------------------------- #
# The real mirrored bytes
# --------------------------------------------------------------------------- #
LIBRARY_QUOTE_FIELDS: tuple[str, ...] = (
    "M9_DEFERRED",
    "PROTEOMICS_CULTURE",
    "PAM_COMPOSITION",
    "PCA_MM",
    "HARVEST_STATE",
    "INSTRUMENT",
    "DIANN_DATABASE",
    "GLOBAL_FDR",
    "QUANTIFICATION",
    "PROTEINS_QUANTIFIED",
    "SIGNIFICANT_COUNTS",
    "COMMON_PROTEINS",
    "PROMOTER_PARTS",
    "PROMOTER_DIRECTION",
    "RECOMBINEERING",
    "CUTSET",
    "CHASSIS_GENOTYPE",
    "CHASSIS_SI_ROW",
    "PROMOTER_REPLICATES",
    "PROMOTER_REPLICATES_ERRORBARS",
    "CROSSFEED_REPLICATES",
    "AEROBICITY",
    "DATA_AVAILABILITY",
    "TITER_NOT_A_DATASET",
    "GROWTH_NOT_A_DATASET",
    "BIOLOG_NOT_A_DATASET",
    "PATHWAY_NOT_TYPED",
)


@pytest.mark.data
def test_every_module_quote_is_verbatim_in_the_pinned_mirrored_bytes() -> None:
    """Each sourced value's quote is a substring of the sha256-pinned artifact."""
    data_root = os.environ["DATA_ROOT"]
    library = Path(data_root) / "torchcell-library" / ds.CITATION_KEY
    cache: dict[str, str] = {}
    sourced = [getattr(ds, name) for name in LIBRARY_QUOTE_FIELDS] + list(
        ds.STRAIN_SI_ROWS.values()
    )
    assert len(sourced) == len(LIBRARY_QUOTE_FIELDS) + len(ds.STRAIN_SI_ROWS)
    for value in sourced:
        uri = value.provenance.source_uri
        if uri not in cache:
            path = library / uri
            assert hashlib.sha256(path.read_bytes()).hexdigest() == (
                value.provenance.sha256
            ), uri
            cache[uri] = path.read_text()
        assert value.quote in cache[uri], (uri, value.quote[:60])
    assert set(cache) == {ds.PAPER_MD, ds.SI1_MD}
    assert ds.PAPER_MD_SHA256 not in (ds.SI1_MD_SHA256,)


@pytest.mark.data
def test_the_raw_mirror_records_the_digests_the_module_pins() -> None:
    manifest = ds.load_manifest()
    for relpath, (_, sha256) in ds.RAW_MEMBERS.items():
        assert ds.manifest_sha256(manifest, relpath) == sha256


@pytest.mark.data
def test_the_built_store_passes_l0_to_l4() -> None:
    data_root = os.environ["DATA_ROOT"]
    root = osp.join(data_root, "data/torchcell/proteome_banerjee2025")
    if not osp.exists(osp.join(root, "processed", "lmdb")):
        pytest.skip(f"dev store absent: {root}")
    report = ds.verify_build(root, data_root)
    assert report.passed, report.summary()
    names = {result.name for result in report.results}
    assert {
        "group_uniqueness",
        "replicate_split",
        "assembly_pin",
        "gene_containment_kt2440",
        "stored_scale_is_the_released_log2_mean",
    } <= names
