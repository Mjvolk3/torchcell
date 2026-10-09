# tests/torchcell/datasets/ecoli/test_schmidt2016_s23_growth_rate.py
# [[tests.torchcell.datasets.ecoli.test_schmidt2016_s23_growth_rate]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_schmidt2016_s23_growth_rate.py
"""Schmidt 2016 Table S23 loader
(``torchcell.datasets.ecoli.schmidt2016_s23_growth_rate``).

The synthetic tests write a workbook whose Table S23 sheet carries the released 12-cell
header and all 26 released rows in their released spellings, including the three that
need normalizing (``'Galactose '`` with its padding space, ``'Osmotic-stress glucose3'``
with its footnote digit, and the four lowercase ``'chemostat ...'`` labels that match
``s25_label`` rather than ``s6_column``). They drive the reader, the condition join, the
partition, the environment builder, the phenotype builder and a full ``process()`` into
``tmp_path``; the genome pin is an in-test object, so nothing reads ``$DATA_ROOT``.

The ``@pytest.mark.data`` tests read the real mirror and pin the numbers the dendron note
states: 26 released rows over the three released strain spellings (22 / 2 / 2), 15
records after the two drop rules, the two negative stationary-phase rates, the four
chemostat rows whose ``Stdev`` is exactly 0, and the Glucose reference at 0.58 h^-1.
"""

from __future__ import annotations

import json
import os
import os.path as osp
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

import torchcell.datasets.ecoli.schmidt2016 as sm
import torchcell.datasets.ecoli.schmidt2016_s23_growth_rate as s23
from torchcell.data import file_sha256
from torchcell.datamodels.schema import (
    ABSOLUTE_MEASUREMENT_TYPES,
    AssayType,
    AssemblyReferenceGenome,
    BacterialAssemblySet,
    MeasurementType,
    SampleUnit,
    UncertaintyType,
)
from torchcell.verification.report import Level

ASSEMBLY_SET: BacterialAssemblySet = "ecoli_K12_BW25113_ASM75055v1"
ACCESSION = "GCA_000750555.1"

#: Table S23's 26 released rows as ``(condition, strain, rate, stdev)``, in sheet order
#: and in the released spellings. The rates and dispersions are the released ones, so
#: the synthetic sheet and the real one agree on every stored number.
RELEASED: tuple[tuple[str, str, float, float], ...] = (
    ("LB", "BW25113", 1.9, 0.03),
    ("LB", "MG1665", 1.78, 0.05),
    ("LB", "NCM3722", 2.3, 0.05),
    ("Glycerol + AA", "BW25113", 1.27, 0.01),
    ("Acetate", "BW25113", 0.3, 0.04),
    ("Fumarate", "BW25113", 0.42, 0.02),
    ("Galactose ", "BW25113", 0.26, 0.003),
    ("Glucose", "BW25113", 0.58, 0.01),
    ("Glucose", "MG1665", 0.67, 0.07),
    ("Glucose", "NCM3722", 1.03, 0.06),
    ("Glucosamine", "BW25113", 0.46, 0.02),
    ("Glycerol", "BW25113", 0.47, 0.01),
    ("Pyruvate", "BW25113", 0.4, 0.01),
    ("Succinate", "BW25113", 0.44, 0.003),
    ("Fructose", "BW25113", 0.65, 0.004),
    ("Mannose", "BW25113", 0.47, 0.01),
    ("Xylose", "BW25113", 0.55, 0.01),
    ("Osmotic-stress glucose3", "BW25113", 0.55, 0.01),
    ("42°C glucose", "BW25113", 0.66, 0.02),
    ("pH6 glucose", "BW25113", 0.63, 0.005),
    ("Stationary phase 1 day", "BW25113", -0.01, 0.003),
    ("Stationary phase 3 days", "BW25113", -0.01, 0.003),
    ("chemostat µ=0.12", "BW25113", 0.12, 0.0),
    ("chemostat µ=0.20", "BW25113", 0.2, 0.0),
    ("chemostat µ=0.35", "BW25113", 0.35, 0.0),
    ("chemostat µ=0.5", "BW25113", 0.5, 0.0),
)
#: The 15 released condition labels that become records, in sheet order.
KEPT_CONDITIONS: tuple[str, ...] = (
    "LB",
    "Acetate",
    "Fumarate",
    "Galactose ",
    "Glucose",
    "Glucosamine",
    "Glycerol",
    "Pyruvate",
    "Succinate",
    "Fructose",
    "Mannose",
    "Xylose",
    "Osmotic-stress glucose3",
    "42°C glucose",
    "pH6 glucose",
)


def write_workbook(path: Path, rows: tuple[tuple[Any, ...], ...] = RELEASED) -> Path:
    """Write a workbook whose Table S23 sheet carries ``rows`` under the real header."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = s23.SHEET_S23
    sheet.append([s23._Q_TABLE_S23])
    sheet.append([None] * len(s23.S23_HEADERS))
    sheet.append(list(s23.S23_HEADERS))
    for condition, strain, rate, stdev in rows:
        sheet.append(
            [condition, strain, rate, stdev, 2.5, 1.2, 20.0, 15.0, 0.4, 0.5, 0.6, 1700]
        )
    book.save(path)
    return path


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    """The synthetic workbook, written once per test."""
    return write_workbook(tmp_path / sm.SI2)


def _reference_genome() -> AssemblyReferenceGenome:
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain="BW25113",
        assembly_set=ASSEMBLY_SET,
        assembly_accession=ACCESSION,
    )


# --------------------------------------------------------------------------- #
# The declared module constants
# --------------------------------------------------------------------------- #
def test_declared_shape_is_26_released_rows_and_15_records() -> None:
    assert s23.EXPECTED_SOURCE_ROWS == 26
    assert sum(s23.EXPECTED_STRAIN_ROWS.values()) == s23.EXPECTED_SOURCE_ROWS
    assert s23.EXPECTED_STRAIN_ROWS == {"BW25113": 22, "MG1665": 2, "NCM3722": 2}
    assert s23.EXPECTED_RECORDS == 15
    assert s23.EXPECTED_UNPERTURBED == 0
    assert s23.EXPECTED_MEDIA_ONLY_EDIT == ("LB",)
    assert s23.N_SAMPLES == 3
    assert s23.REFERENCE_CONDITION == "Glucose"


def test_the_readout_is_absolute_and_its_sourcing_quotes_the_pinned_artifacts() -> None:
    assert MeasurementType.growth_rate in ABSOLUTE_MEASUREMENT_TYPES
    sourced = s23.SOURCED_VALUES
    assert sourced["n_samples"].value == 3
    assert "biological triplicates" in sourced["n_samples"].quote
    assert sourced["n_samples"].note is not None
    assert "conservative" in sourced["n_samples"].note
    assert sourced["growth_table"].provenance.sha256 == sm.SI2_SHA256
    assert sourced["growth_rate_estimator"].provenance.sha256 == sm.PAPER_MD_SHA256
    assert sourced["assay_type"].value == AssayType.other.value
    assert [gap.field for gap in s23.RECORD_GAPS] == ["replicate_id"]


# --------------------------------------------------------------------------- #
# The condition join
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("label", "column"),
    [
        ("Glucose", "Glucose"),
        ("Galactose ", "Galactose"),
        ("Osmotic-stress glucose3", "Osmotic-stress glucose"),
        ("chemostat µ=0.12", "Chemostat µ=0.12"),
    ],
)
def test_condition_spec_normalizes_each_released_spelling(
    label: str, column: str
) -> None:
    assert s23._condition_spec(label).s6_column == column


def test_condition_spec_refuses_a_label_naming_no_condition() -> None:
    with pytest.raises(RuntimeError, match="names no schmidt2016 condition"):
        s23._condition_spec("Sorbitol")


def test_every_released_label_joins_a_condition_and_the_26_rows_cover_22() -> None:
    specs = {s23._condition_spec(label).s6_column for label, *_ in RELEASED}
    assert specs == {spec.s6_column for spec in sm.CONDITIONS}
    assert len(specs) == 22


# --------------------------------------------------------------------------- #
# Reading Table S23
# --------------------------------------------------------------------------- #
def test_read_table_s23_types_all_26_rows_with_their_unstored_cells(
    workbook: Path,
) -> None:
    rows = s23.read_table_s23(workbook)
    assert len(rows) == s23.EXPECTED_SOURCE_ROWS
    assert [row.row_number for row in rows] == list(range(4, 30))
    assert [(row.condition, row.strain) for row in rows] == [
        (condition, strain) for condition, strain, _, _ in RELEASED
    ]
    first = rows[0]
    assert (first.growth_rate, first.stdev) == (1.9, 0.03)
    assert set(first.unstored) == {
        "Single cell volume [fl]1",
        "Doubling time (h-1)",
        "Time exp before harvest (h)",
        "# of doublings at exponential growth before harvesting",
        "OD @ harvesting. replicates",
        "unlabelled_9",
        "unlabelled_10",
        "Number of Proteins Identified (FDR 1%)2",
    }
    assert first.unstored["Number of Proteins Identified (FDR 1%)2"] == 1700
    # the two stationary rows release a negative rate; the reader stores it verbatim
    stationary = [row for row in rows if row.condition.startswith("Stationary")]
    assert [row.growth_rate for row in stationary] == [-0.01, -0.01]
    # the four chemostat rows release a Stdev of exactly 0
    chemostat = [row for row in rows if row.condition.startswith("chemostat")]
    assert [row.stdev for row in chemostat] == [0.0, 0.0, 0.0, 0.0]


def test_read_table_s23_refuses_a_renamed_header(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / sm.SI2)
    book = openpyxl.load_workbook(path)
    book[s23.SHEET_S23].cell(row=sm.HEADER_ROW, column=3).value = "Growth rate"
    book.save(path)
    with pytest.raises(RuntimeError, match="the module declares"):
        s23.read_table_s23(path)


def test_read_table_s23_refuses_an_unreleased_strain(tmp_path: Path) -> None:
    rows = (("Glucose", "W3110", 0.58, 0.01), *RELEASED)
    with pytest.raises(RuntimeError, match="names the strain 'W3110'"):
        s23.read_table_s23(write_workbook(tmp_path / sm.SI2, rows))


def test_read_table_s23_refuses_a_changed_row_count(tmp_path: Path) -> None:
    with pytest.raises(RuntimeError, match="holds 25 rows"):
        s23.read_table_s23(write_workbook(tmp_path / sm.SI2, RELEASED[:-1]))


def test_read_table_s23_refuses_a_changed_per_strain_count(tmp_path: Path) -> None:
    rows = tuple(
        (condition, "BW25113" if strain == "NCM3722" else strain, rate, stdev)
        for condition, strain, rate, stdev in RELEASED
    )
    with pytest.raises(RuntimeError, match="per-strain row counts are"):
        s23.read_table_s23(write_workbook(tmp_path / sm.SI2, rows))


# --------------------------------------------------------------------------- #
# The partition, the environments and the phenotype
# --------------------------------------------------------------------------- #
def test_partition_keeps_15_rows_and_names_every_drop_rule(workbook: Path) -> None:
    kept, ledger = s23.partition(s23.read_table_s23(workbook))
    assert [row.condition for row in kept] == list(KEPT_CONDITIONS)
    assert ledger.source_rows == 26
    assert ledger.kept_records == 15
    assert ledger.dropped_rows == 11
    assert ledger.dropped == {
        s23.DROP_STRAIN_RULE: [
            "LB/MG1665",
            "LB/NCM3722",
            "Glucose/MG1665",
            "Glucose/NCM3722",
        ],
        "medium_has_no_media_library_entry": ["Glycerol + AA/BW25113"],
        "growth_phase_not_representable": [
            "Stationary phase 1 day/BW25113",
            "Stationary phase 3 days/BW25113",
        ],
        "culture_not_batch": [
            "chemostat µ=0.12/BW25113",
            "chemostat µ=0.20/BW25113",
            "chemostat µ=0.35/BW25113",
            "chemostat µ=0.5/BW25113",
        ],
    }
    assert all(row.strain == "BW25113" for row in kept)


def test_build_environments_are_distinct_and_only_lb_has_no_perturbation(
    workbook: Path,
) -> None:
    kept, _ = s23.partition(s23.read_table_s23(workbook))
    environments = s23.build_environments(kept)
    assert len(environments) == 15
    assert len({env.model_dump_json() for env in environments.values()}) == 15
    assert tuple(
        label for label, env in environments.items() if not env.perturbations
    ) == ("LB",)
    hot = environments["42°C glucose"].temperature
    assert hot is not None and hot.value == 42.0


def test_phenotype_stores_the_released_rate_and_the_conservative_replicate_count(
    workbook: Path,
) -> None:
    kept, _ = s23.partition(s23.read_table_s23(workbook))
    glucose = s23.reference_row(kept)
    phenotype = s23.phenotype(glucose)
    assert phenotype.measurement_type is MeasurementType.growth_rate
    assert phenotype.assay_type is AssayType.other
    assert phenotype.environment_response == 0.58
    assert phenotype.environment_response_uncertainty == 0.01
    assert phenotype.environment_response_uncertainty_type is UncertaintyType.sample_sd
    assert phenotype.n_samples == 3
    assert phenotype.sample_unit is SampleUnit.biological_replicate
    assert phenotype.units == s23.UNITS
    assert phenotype.replicate_id is None
    assert [gap.field for gap in phenotype.provenance_gaps] == ["replicate_id"]
    # the schema derives the SE from the released dispersion and the resolved n
    assert phenotype.environment_response_se == pytest.approx(0.01 / 3**0.5)


def test_reference_row_refuses_a_kept_set_without_exactly_one_glucose_row(
    workbook: Path,
) -> None:
    kept, _ = s23.partition(s23.read_table_s23(workbook))
    without = [row for row in kept if row.condition != s23.REFERENCE_CONDITION]
    with pytest.raises(RuntimeError, match="0 kept rows carry the reference condition"):
        s23.reference_row(without)


# --------------------------------------------------------------------------- #
# End-to-end build on the synthetic workbook
# --------------------------------------------------------------------------- #
def test_download_links_the_mirror(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = write_workbook(tmp_path / "source.xlsx")
    sha = file_sha256(source)
    monkeypatch.setattr(sm, "SI2_SHA256", sha)
    data_root = tmp_path / "root"
    sm.deposit_raw_mirror(source=source, data_root=str(data_root))
    monkeypatch.setenv("DATA_ROOT", str(data_root))

    root = tmp_path / "build"
    (root / "raw").mkdir(parents=True)
    dataset = s23.GrowthRateS23Schmidt2016Dataset.__new__(
        s23.GrowthRateS23Schmidt2016Dataset
    )
    monkeypatch.setattr(
        type(dataset), "raw_dir", property(lambda self: str(root / "raw"))
    )
    dataset.download()
    assert file_sha256(root / "raw" / sm.SI2) == sha

    (data_root / sm.RAW_DIR_REL / sm.SI2_MIRROR_RELPATH).unlink()
    with pytest.raises(RuntimeError, match="missing from mirror"):
        dataset.download()


@pytest.fixture
def built(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> s23.GrowthRateS23Schmidt2016Dataset:
    """The 15-record store, built from the synthetic workbook into ``tmp_path``."""
    root = tmp_path / "growth_rate_s23_schmidt2016"
    raw = root / "raw"
    raw.mkdir(parents=True)
    path = write_workbook(raw / sm.SI2)
    monkeypatch.setattr(sm, "DATA_SHA256", {sm.SI2: file_sha256(path)})
    monkeypatch.setattr(
        s23, "assembly_reference", lambda strain, **_: _reference_genome()
    )
    return s23.GrowthRateS23Schmidt2016Dataset(root=str(root))


def test_process_builds_15_records_against_the_glucose_reference(
    built: s23.GrowthRateS23Schmidt2016Dataset,
) -> None:
    assert len(built) == s23.EXPECTED_RECORDS
    items = [built.transform_item(built[i]) for i in range(len(built))]
    assert {len(i["experiment"].genotype.perturbations) for i in items} == {0}
    assert {i["publication"].doi for i in items} == {sm.PAPER_DOI}
    assert {i["reference"].phenotype_reference.environment_response for i in items} == {
        0.58
    }
    assert {i["reference"].genome_reference.assembly_accession for i in items} == {
        ACCESSION
    }
    assert sorted(i["experiment"].phenotype.environment_response for i in items) == [
        0.26,
        0.3,
        0.4,
        0.42,
        0.44,
        0.46,
        0.47,
        0.47,
        0.55,
        0.55,
        0.58,
        0.63,
        0.65,
        0.66,
        1.9,
    ]
    # every kept condition has its own environment, so no two records collide
    assert len({i["experiment"].environment.model_dump_json() for i in items}) == 15
    built.close_lmdb()


def test_process_writes_the_three_declared_ledgers(
    built: s23.GrowthRateS23Schmidt2016Dataset,
) -> None:
    preprocess = Path(built.preprocess_dir)
    built.close_lmdb()
    ledger = json.loads((preprocess / "dropped_records.json").read_text())
    assert (ledger["source_rows"], ledger["kept_records"]) == (26, 15)
    assert ledger["dropped_records"] == 11
    assert {rule["rule"] for rule in ledger["rules"]} == {
        s23.DROP_STRAIN_RULE,
        "medium_has_no_media_library_entry",
        "growth_phase_not_representable",
        "culture_not_batch",
    }
    assert any("Table S28" in note for note in ledger["notes"])
    not_stored = json.loads((preprocess / "not_stored.json").read_text())
    assert not_stored["issue"] == 826
    assert len(not_stored["columns"]) == 8
    # all 26 released (condition, strain) pairs are distinct, so none collapses
    assert len(not_stored["values"]) == s23.EXPECTED_SOURCE_ROWS
    sourced = json.loads((preprocess / "sourced_values.json").read_text())
    assert sourced["n_samples"]["value"] == 3
    frame = pd.read_csv(preprocess / "growth_rates.csv")
    assert len(frame) == 15
    assert frame["is_reference"].sum() == 1
    assert set(frame["strain"]) == {"BW25113"}
    assert set(frame["n_samples"]) == {3}
    assert (preprocess / "build_manifest.json").exists()


def test_verify_build_passes_every_level_on_the_absolute_branch(
    built: s23.GrowthRateS23Schmidt2016Dataset, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = built.root
    built.close_lmdb()
    report = s23.verify_build(root)
    assert [(r.level, r.name) for r in report.results if not r.passed] == []
    assert report.passed
    assert (Level.L4, "assembly_pin_resolves") in [
        (r.level, r.name) for r in report.results
    ]
    assert any(
        "absolute rule" in r.message
        for r in report.results
        if r.name == "reference_zero"
    )


def test_l4_assembly_pin_refuses_a_record_pinning_another_assembly(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        s23, "assembly_reference", lambda strain, **_: _reference_genome()
    )
    records = [
        {
            "reference": {
                "genome_reference": {
                    "assembly_set": "ecoli_K12_MG1655_ASM584v2",
                    "assembly_accession": "GCA_000005845.2",
                }
            }
        }
    ]
    result = s23.l4_assembly_pin(records)
    assert result.passed is False
    assert "the deposited assembly report names" in result.message


# --------------------------------------------------------------------------- #
# Data-gated: the real mirror
# --------------------------------------------------------------------------- #
def _real_workbook() -> str:
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT is not set")
    path = osp.join(data_root, sm.RAW_DIR_REL, sm.SI2_MIRROR_RELPATH)
    if not osp.exists(path):
        pytest.skip(f"{path} is not in the raw mirror")
    return path


@pytest.mark.data
def test_real_table_s23_releases_the_26_rows_the_module_declares() -> None:
    rows = s23.read_table_s23(_real_workbook())
    assert [
        (row.condition, row.strain, row.growth_rate, row.stdev) for row in rows
    ] == [
        (condition, strain, rate, stdev) for condition, strain, rate, stdev in RELEASED
    ]


@pytest.mark.data
def test_real_partition_keeps_the_15_conditions_and_drops_11_rows() -> None:
    kept, ledger = s23.partition(s23.read_table_s23(_real_workbook()))
    assert [row.condition for row in kept] == list(KEPT_CONDITIONS)
    assert (ledger.kept_records, ledger.dropped_rows) == (15, 11)
    assert s23.reference_row(kept).growth_rate == 0.58


@pytest.mark.data
def test_real_dev_store_holds_15_records_and_its_declared_ledger() -> None:
    data_root = os.environ.get("DATA_ROOT")
    if data_root is None:
        pytest.skip("DATA_ROOT is not set")
    root = osp.join(data_root, s23.DATASET_ROOT_REL)
    ledger_path = Path(root, "preprocess", "dropped_records.json")
    if not ledger_path.exists():
        pytest.skip(f"{root} is not built")
    ledger = json.loads(ledger_path.read_text())
    assert ledger["kept_records"] == s23.EXPECTED_RECORDS
    assert ledger["dropped_records"] == 11
