# tests/torchcell/adapters/test_hoepfner2014_adapter.py
"""Unit tests for the Hoepfner 2014 HIP-HOP adapter conf and its node projections.

The conf is the whole surface of this adapter (the class only picks which enable-list to
load), so what is worth pinning is that every method it enables exists, that the
gene-keyed genotype/perturbation pair and the environment-perturbation pair ARE enabled
(the first separates this conf from the Bloom 2019 segregant one, the second is what makes
the dosed compound a node carrying its InChIKey), and that the phenotype projection
carries the screen id the records are keyed on.
"""

from __future__ import annotations

import hashlib
import json
import os.path as osp
from pathlib import Path
from typing import Any, cast

import pytest
import yaml

import torchcell.adapters as adapters_package
import torchcell.adapters.hoepfner2014_adapter as _init_module
from tests.torchcell.adapters._adapter_init_harness import (
    AdapterCase,
    Shape,
    assert_construction,
    assert_missing_conf,
)
from torchcell.adapters.cell_adapter import CellAdapter
from torchcell.adapters.hoepfner2014_adapter import EnvChemgenHoepfner2014Adapter
from torchcell.datamodels.identity import (
    environment_perturbation_identity,
    identity_sha256,
)
from torchcell.datamodels.schema import (
    AssayType,
    Concentration,
    ConcentrationUnit,
    DoseBasis,
    EnvironmentResponsePhenotype,
    MeasurementType,
    SampleUnit,
    SmallMoleculePerturbation,
)
from torchcell.datasets.scerevisiae.hoepfner2014 import EnvChemgenHoepfner2014Dataset

CONF = "env_chemgen_hoepfner2014_adapter.yaml"


def _load() -> dict[str, Any]:
    path = osp.join(osp.dirname(adapters_package.__file__), "conf", CONF)
    with open(path) as handle:
        conf: dict[str, Any] = yaml.safe_load(handle)
    return conf


def test_conf_exists_and_parses() -> None:
    conf = _load()["cell_adapter"]
    assert conf["node_methods"] and conf["edge_methods"]


def test_every_enabled_method_exists_in_the_cell_adapter_tables() -> None:
    """A conf naming a method the adapter does not implement is a silent no-op."""
    import torchcell.adapters.cell_adapter as module

    with open(module.__file__) as handle:
        source = handle.read()
    conf = _load()["cell_adapter"]
    for entry in conf["node_methods"] + conf["edge_methods"]:
        assert f'"{entry["method_name"]}"' in source, entry["method_name"]


def test_gene_keyed_genotype_and_environment_perturbation_are_enabled() -> None:
    conf = _load()["cell_adapter"]
    nodes = {entry["method_name"] for entry in conf["node_methods"]}
    edges = {entry["method_name"] for entry in conf["edge_methods"]}
    assert "genotype (chunked)" in nodes and "perturbation (chunked)" in nodes
    assert "perturbation to genotype (chunked)" in edges
    assert "environment perturbation (chunked)" in nodes
    assert "environment perturbation to environment (chunked)" in edges
    assert "environment response phenotype (chunked)" in nodes
    assert "environment response phenotype reference" in nodes
    # A segregant genotype is Bloom's shape, never this dataset's.
    assert "segregant genotype (chunked)" not in nodes


def test_adapter_points_at_its_own_conf_and_dataset() -> None:
    import inspect

    source = inspect.getsource(EnvChemgenHoepfner2014Adapter)
    assert CONF in source
    assert EnvChemgenHoepfner2014Dataset.__name__ in source


def test_environment_perturbation_node_projects_the_compound_identity() -> None:
    """The node id hashes the perturbation, so the cleaned compound is what joins."""
    perturbation = SmallMoleculePerturbation(
        compound=EnvChemgenHoepfner2014Dataset._vehicle(),
        concentration=Concentration(
            value=2.0, unit=ConcentrationUnit.percent_v_v, basis=DoseBasis.fixed
        ),
    )

    class FakeEnvironment:
        perturbations = [perturbation]

    class FakeExperiment:
        environment = FakeEnvironment()

    adapter = CellAdapter.__new__(CellAdapter)
    undecorated = cast(Any, CellAdapter._environment_perturbation_node).__wrapped__
    nodes = undecorated(
        adapter, {"experiment": FakeExperiment()}, "environment perturbation (chunked)"
    )
    assert len(nodes) == 1
    expected = identity_sha256(environment_perturbation_identity(perturbation))
    assert nodes[0].get_id() == expected
    props = nodes[0].get_properties()
    assert props["inchikey"] == "IAZDPXIOMUYVGZ-UHFFFAOYSA-N"
    assert "serialized_data" not in props


def test_environment_response_phenotype_node_carries_the_screen_id() -> None:
    phenotype = EnvironmentResponsePhenotype(
        measurement_type=MeasurementType.sensitivity_score,
        assay_type=AssayType.pooled_competitive_growth_barcode,
        environment_response=-3.25,
        n_samples=2,
        sample_unit=SampleUnit.technical_replicate,
        units="adjusted MADL sensitivity score",
        screen_id="0077",
    )

    class FakeExperiment:
        pass

    experiment = FakeExperiment()
    experiment.phenotype = phenotype  # type: ignore[attr-defined]
    adapter = CellAdapter.__new__(CellAdapter)
    undecorated = cast(
        Any, CellAdapter._environment_response_phenotype_node
    ).__wrapped__
    node = undecorated(
        adapter, {"experiment": experiment}, "environment response phenotype (chunked)"
    )
    props = node.get_properties()
    assert props["screen_id"] == "0077"
    assert props["environment_response"] == -3.25
    assert "serialized_data" not in props


# ---- #506: the adapter over a small strain-resolved scratch build ---------------- #
_ORF_FASTA = (
    ">YAL001C TFC3 SGDID:S000000001\nATGGTA\n>YBR115C LYS2 SGDID:S000000319\nATGACC\n"
)
_HIP = [
    [
        "Systematic Name",
        "Ad. scores for Exp. 3_50_HIP_0077",
        "Ad. scores for Exp. 4016_60000_HIP_0125",
    ],
    ["YAL001C", "-1.5", "0.25"],
    ["YBR115C", "0.5", "-0.75"],
]
_HOP = [
    [
        "Systematic Name",
        "Ad. scores for Exp. 3_50_HOP_0077",
        "Ad. scores for Exp. 4016_60000_HOP_0126",
    ],
    ["YAL001C", "1.0", "2.0"],
]


class _CurrentGenome:
    """Every ORF resolves as a CURRENT gene."""

    def resolve_gene_name(self, name: str) -> Any:
        from torchcell.sequence.genome.scerevisiae.s288c import (
            GeneNameResolution,
            GeneNameStatus,
        )

        return GeneNameResolution(
            input_name=name, status=GeneNameStatus.CURRENT, systematic_name=name
        )


class _NoWandb:
    """The three members of ``wandb`` the cell adapter touches, all inert."""

    class Table:
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            pass

        def add_data(self, *args: Any) -> None:
            pass

    def log(self, payload: dict[str, Any]) -> None:
        pass

    def init(self) -> None:
        pass


def _small_build(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """A real ``process()`` over two tiny matrices, a two-gene genomes tier and a
    Table S1 with amitriptyline (CMB 3, IC30 50 uM) and hydrochloric acid (CMB 4016,
    IC30 60849 uM), with the raw pins repointed at the fixture's bytes.
    """
    import gc

    import openpyxl

    import torchcell.adapters.cell_adapter as cell_adapter_module
    from torchcell.datasets.scerevisiae import hoepfner2014 as m
    from torchcell.literature.manifest import ArtifactRecord
    from torchcell.sequence.genome.registry import GenomeManifest

    data_root = tmp_path / "data_root"
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.setattr(cell_adapter_module, "wandb", _NoWandb())
    tier = data_root / "torchcell-genomes" / "sgd_S288C_R64-4-1_20230830"
    tier.mkdir(parents=True)
    records = []
    for name, text in (
        ("orf_coding_all_R64-4-1_20230830.fasta", _ORF_FASTA),
        ("rna_coding_R64-4-1_20230830.fasta", ""),
    ):
        (tier / name).write_text(text)
        records.append(
            ArtifactRecord(
                path=name,
                role="sequence",
                bytes=len(text.encode()),
                sha256=hashlib.sha256(text.encode()).hexdigest(),
            )
        )
    (tier / "manifest.json").write_text(
        GenomeManifest(
            assembly_set="sgd_S288C_R64-4-1_20230830",
            organism="Saccharomyces cerevisiae",
            strain_or_population="S288C",
            source="SGD",
            release="R64-4-1_20230830",
            files=records,
            provenance_complete=True,
            created_at="2026-10-02T00:00:00+00:00",
        ).model_dump_json()
    )
    root = tmp_path / "env_chemgen_hoepfner2014"
    raw = root / "raw"
    raw.mkdir(parents=True)
    for name, rows in (("HIP_scores.txt", _HIP), ("HOP_scores.txt", _HOP)):
        (raw / name).write_text(
            "\n".join("\t".join(f'"{c}"' for c in row) for row in rows) + "\n"
        )
    book = openpyxl.Workbook()
    known = book.active
    known.title = "Reference Substances known MoA"
    known.append(["CMB ID", "Common Name", "IC30 (uM)"])
    known.append([3, "Amitriptyline", 50])
    known.append([4016, "Hydrochloric Acid", 60849])
    book.create_sheet("Substances novel MoA").append(["CMB ID", "Common Name"])
    structures = book.create_sheet("All Structures")
    structures.append(["CMB ID", "SMILE string"])
    structures.append([3, "CN(C)CCC=C2c1ccccc1CCc3ccccc23"])
    structures.append([4016, "Cl"])
    book.save(raw / "Table_S1.xls")
    monkeypatch.setattr(
        m,
        "_DRYAD_FILES",
        {
            name: {
                **spec,
                "sha256": hashlib.sha256((raw / name).read_bytes()).hexdigest(),
            }
            for name, spec in m._DRYAD_FILES.items()
        },
    )
    dataset = m.EnvChemgenHoepfner2014Dataset(root=str(root), genome=_CurrentGenome())
    gc.unfreeze()
    return dataset


def test_adapter_emits_strain_resolved_nodes_from_a_scratch_build(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """#506: the records are ``strain_environment_response`` experiments; the genome
    node serializes the BY4743 background (HIS3 0 functional copies, every allele a
    pending gap); the perturbation nodes are the heterozygous and barcoded kanMX4
    leaves; the 60 mM HCl column adds a pH environment-perturbation node next to the
    compound node; the phenotype nodes carry the screen id.
    """
    import gc

    from torchcell.datamodels.schema import (
        StrainEnvironmentResponseExperimentReference,
        StrainReferenceGenome,
    )

    monkeypatch.delenv("TC_DATA_URL", raising=False)
    dataset = _small_build(tmp_path, monkeypatch)
    assert len(dataset) == 6
    adapter = EnvChemgenHoepfner2014Adapter(
        dataset=dataset,
        process_workers=1,
        io_workers=1,
        chunk_size=2,
        loader_batch_size=2,
    )
    nodes = list(adapter.get_nodes())
    gc.unfreeze()
    by_label: dict[str, list[Any]] = {}
    for node in nodes:
        by_label.setdefault(node.get_label(), []).append(node)
    assert {
        "experiment",
        "experiment reference",
        "genome",
        "perturbation",
        "environment perturbation",
        "environment response phenotype",
    } <= set(by_label)
    experiments = [
        json.loads(n.get_properties()["serialized_data"])
        for n in by_label["experiment"]
    ]
    assert {e["experiment_type"] for e in experiments} == {
        "strain_environment_response"
    }
    (genome,) = {n.get_id(): n for n in by_label["genome"]}.values()
    assert genome.get_properties()["strain"] == "BY4743"
    background = StrainReferenceGenome.model_validate_json(
        genome.get_properties()["serialized_data"]
    ).background
    assert background.functional_copies("YOR202W") == 0
    assert not background.is_fully_sourced
    references = [
        StrainEnvironmentResponseExperimentReference.model_validate_json(
            n.get_properties()["serialized_data"]
        )
        for n in by_label["experiment reference"]
    ]
    assert len(references) == 4  # (HIP, 0077), (HIP, 0125), (HOP, 0077), (HOP, 0126)
    assert {
        n.get_properties()["perturbation_type"] for n in by_label["perturbation"]
    } == {"heterozygous_deletion", "barcoded_kanmx_deletion"}
    env_perts = {
        (p["perturbation_type"], p["factor"], p["compound_name"])
        for p in (n.get_properties() for n in by_label["environment perturbation"])
    }
    assert ("environment_physical", "pH", "hydrochloric acid") in {
        (t, f, str(c).lower()) for t, f, c in env_perts
    }
    assert {
        n.get_properties()["screen_id"]
        for n in by_label["environment response phenotype"]
    } >= {"0077", "0125", "0126"}


# 2026.10.06, Phase 21: the constructor (exact conf content, wiring, refusal); the
# checks are in tests/torchcell/adapters/_adapter_init_harness.py.
_INIT_CASES = [
    AdapterCase(
        EnvChemgenHoepfner2014Adapter,
        "env_chemgen_hoepfner2014_adapter.yaml",
        Shape("environment response phenotype", env_perturbation=True),
        EnvChemgenHoepfner2014Dataset,
    )
]


@pytest.mark.parametrize("case", _INIT_CASES, ids=lambda c: c.adapter_cls.__name__)
def test_init_serves_the_exact_conf_and_wires_the_base_adapter(
    case: AdapterCase,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """The conf the constructor loads is exactly the one its dataset's shape needs.

    * ``EnvChemgenHoepfner2014Adapter`` loads ``conf/env_chemgen_hoepfner2014_adapter.yaml``: 17 node and 15 edge methods (``environment response phenotype``; gene-keyed genotype; perturbation nodes; environment-perturbation nodes; memory_reduction_factor 1.0 on every chunked method).

    The adapter is checked against the dataset ``dataset_adapter_map`` pairs it with
    (asserted to be the class above). The conf lists its methods in the order the
    adapter runs them, no edge dangles, every chunked entity node is linked, and the
    phenotype method matches that dataset's ``experiment_class``. The adapter keeps the dataset and the worker / chunk sizes
    it was given (3, 2, 500, 50), calls ``wandb.init`` once and logs the method table
    (event number, name, node/edge, factor or NaN for a non-chunked method) then the
    dataset name and the pinned start time; nothing is printed.
    """
    assert_construction(case, monkeypatch, capsys)


@pytest.mark.parametrize("case", _INIT_CASES, ids=lambda c: c.adapter_cls.__name__)
def test_init_refuses_a_missing_conf_before_wandb(
    case: AdapterCase, monkeypatch: pytest.MonkeyPatch
) -> None:
    """With the conf absent the error names ``<adapters dir>/conf/<conf name>``."""
    assert_missing_conf(case, _init_module, monkeypatch)
