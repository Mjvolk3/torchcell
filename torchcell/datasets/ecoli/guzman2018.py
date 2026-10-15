# torchcell/datasets/ecoli/guzman2018
# [[torchcell.datasets.ecoli.guzman2018]]
# https://github.com/Mjvolk3/torchcell/tree/main/torchcell/datasets/ecoli/guzman2018
# Test file: tests/torchcell/datasets/ecoli/test_guzman2018.py
"""Guzmán 2018 (BMC Syst Biol): the 14 RNA-seq libraries from the study's own release.

Guzmán, Sandberg, LaCroix, Nyerges, Papp, Feist, Palsson and Monk 2018, "Reframing gene
essentiality in terms of adaptive flexibility" (BMC Syst Biol 12:143,
doi:10.1186/s12918-018-0653-z, citation key ``guzmanReframingGeneEssentiality2018``),
issue #964. Plan decision 16: the PRECISE-1K compendium re-quantifies these 14
libraries (``p1k_00219``-``p1k_00232``) and serves none of them; this loader stores the
study's OWN release, attributed to it, and the compendium is an L4 cross-check that
never gates.

WHAT IS STORED. One ``RNASeqExpressionExperiment`` per column of GEO
``GSE117303_deseq2-tpm.csv.gz``: two wild-type BW25113 libraries and twelve libraries of
Keio deletion strains (``thrA``, ``serB``, ``ptsI``) that grew after extended culture on
glucose M9, each in biological duplicate. ``expression_tpm`` is the released matrix,
keyed by the release's own gene ids (BW25113 locus tags, and 106 ``_1`` / ``_2`` split
ids the release writes for split features), verbatim. ``expression_count`` is a typed
gap: the release is TPM only, and the paper's counts were never deposited.

GENOTYPE. A deleted strain is its Keio deletion (``gene_symbol`` route, the released
symbol kept on ``identifier_mapping``) plus every breseq call Additional file 4 releases
for the matching DNA sample (``FPS <gene> A<n> F<n> I<n> R1``), each with its released
frequency. The GEO sample record says whether a library is a ``Clone`` (an isolate,
``I1``) or a ``Population`` (``I0``), which is the call mode. A call in a gene is a
``LocusVariantCallPerturbation`` on that gene's locus; an intergenic call is a
``SiteVariantCallPerturbation`` keyed on ``CP009273.1:<position>`` with the
released flanking-gene cell verbatim and no flanking lookup. 12 calls over the 6 strains.

ENVIRONMENT. M9 with 2 g/L D-glucose, aerobic, 37 C; the deletion strains with 50 mg/L
kanamycin A, the wild type without. The media section of the Methods states 25 mg/L
kanamycin for the growth characterizations; the RNA-seq section and the GEO growth
protocol state 50 mg/L for these libraries, and that is what is stored.

REFERENCE. Wild-type BW25113 in M9 glucose without kanamycin, its TPM the mean of the
two released wild-type libraries (the comparison the study's DESeq2 analysis makes),
with ``expression_count`` a typed gap.

NOT STORED: Additional file 4's calls for DNA samples that have no RNA library (the
other strains and flasks), the 2-to-17-million read totals (a range, not per library),
and the DESeq2 log2 fold changes and adjusted p-values of Table 3 (a curated subset;
the released matrix is the per-library measurement).
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import logging
import math
import os
import os.path as osp
import re
import shutil
from collections.abc import Callable, Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, ClassVar, Final

from pydantic import BaseModel, ConfigDict

from torchcell.data import (
    ExperimentDataset,
    Organism,
    check_manifest_pin,
    link_verified,
    post_process,
    verify_raw_files,
)
from torchcell.data.experiment_dataset import verify_sha256
from torchcell.datamodels.compound_identity import resolved_compound
from torchcell.datamodels.schema import (
    BacterialVariantCall,
    BacterialVariantType,
    Concentration,
    ConcentrationUnit,
    DeletionStrainPerturbation,
    DerivedIdentifierMapping,
    Environment,
    Experiment,
    ExperimentReference,
    GenePerturbationType,
    Genotype,
    LocusVariantCallPerturbation,
    Media,
    MediaComponent,
    MediaComponentRole,
    Publication,
    RNASeqExpressionExperiment,
    RNASeqExpressionExperimentReference,
    RNASeqExpressionPhenotype,
    SiteVariantCallPerturbation,
    Temperature,
    VariantCallMode,
    VariantFrequencyBasis,
    VariantSiteKind,
)
from torchcell.datasets.bacteria_common import (
    STRAIN_GENE_NAMESPACES,
    assembly_reference,
    bacterial_genome,
)
from torchcell.datasets.dataset_registry import register_dataset
from torchcell.literature.manifest import (
    ArtifactRecord,
    Manifest,
    RetrievalMethod,
    RetrievalRecord,
)
from torchcell.sequence.genome.ecoli.k12 import (
    EcoliK12BW25113Genome,
    EcoliK12Genome,
    EcoliK12MG1655Genome,
    EcoliK12StrainName,
)
from torchcell.verification.report import (
    Level,
    LevelResult,
    Provenance,
    VerificationReport,
)
from torchcell.verification.sourced import (
    ProvenanceGap,
    ProvenanceGapReason,
    SourcedValue,
    audit_sourced_value,
)

log = logging.getLogger(__name__)

CITATION_KEY: Final = "guzmanReframingGeneEssentiality2018"
PAPER_DOI: Final = "10.1186/s12918-018-0653-z"
PAPER_TITLE: Final = "Reframing gene essentiality in terms of adaptive flexibility"
PMCID: Final = "PMC6296033"
GEO_SERIES: Final = "GSE117303"
RAW_DIR_REL: Final = f"torchcell-raw/{CITATION_KEY}"
DATASET_ROOT_REL: Final = "data/torchcell/rnaseq_guzman2018"
RETRIEVED_AT: Final = "2026-10-12"

REFERENCE_STRAIN: Final[EcoliK12StrainName] = "BW25113"
GENE_NAMESPACE: Final = STRAIN_GENE_NAMESPACES[REFERENCE_STRAIN]
REPLICON: Final = "CP009273.1"
CALLER: Final = "breseq 0.30.0"
MEASUREMENT_TYPE: Final = "rnaseq_tpm"
KEIO_COLLECTION: Final = "Keio collection"
KEIO_CASSETTE: Final = "Kanamycin resistance cassette"

_PMC: Final = f"https://pmc-oa-opendata.s3.amazonaws.com/{PMCID}.1"
_GEO: Final = "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE117nnn/GSE117303"


class RawFile(BaseModel):
    """One deposited file: where it sits in the raw mirror and how it was fetched."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    name: str
    relpath: str
    sha256: str
    method: RetrievalMethod
    source_url: str
    retriever: str
    params: dict[str, str]


PAPER_TEXT: Final = RawFile(
    name=f"{PMCID}.1.txt",
    relpath=f"paper/{PMCID}.1.txt",
    sha256="392b5f3d016ddeb5158fb8bb22b3e61af7d6b65c62fdd2653432a58a7fc46b71",
    method=RetrievalMethod.pmc_cloud,
    source_url=f"{_PMC}/{PMCID}.1.txt",
    retriever="torchcell.literature.retrieve.pmc_cloud_object",
    params={"key": f"{PMCID}.1/{PMCID}.1.txt"},
)
MUTATIONS: Final = RawFile(
    name="12918_2018_653_MOESM4_ESM.xlsx",
    relpath="si/12918_2018_653_MOESM4_ESM.xlsx",
    sha256="97e587ad5586da20a613055698d46ad10fbe2d1d1a8e0eeb95848a4a67e6f57d",
    method=RetrievalMethod.pmc_cloud,
    source_url=f"{_PMC}/12918_2018_653_MOESM4_ESM.xlsx",
    retriever="torchcell.literature.retrieve.pmc_cloud_object",
    params={"key": f"{PMCID}.1/12918_2018_653_MOESM4_ESM.xlsx"},
)
TPM: Final = RawFile(
    name="GSE117303_deseq2-tpm.csv.gz",
    relpath="data/GSE117303_deseq2-tpm.csv.gz",
    sha256="590b94e9b36be11f52a2681903face3ceef5a35e80bdcf0619d04297c3fdf282",
    method=RetrievalMethod.direct_url,
    source_url=f"{_GEO}/suppl/GSE117303_deseq2-tpm.csv.gz",
    retriever="torchcell.literature.retrieve.direct_url",
    params={"url": f"{_GEO}/suppl/GSE117303_deseq2-tpm.csv.gz"},
)
SERIES_MATRIX: Final = RawFile(
    name="GSE117303_series_matrix.txt.gz",
    relpath="data/GSE117303_series_matrix.txt.gz",
    sha256="5a1c437a591442ec83aabb7bb4caf9b66f289f82de6065a9636bf41dd1ad9302",
    method=RetrievalMethod.direct_url,
    source_url=f"{_GEO}/matrix/GSE117303_series_matrix.txt.gz",
    retriever="torchcell.literature.retrieve.direct_url",
    params={"url": f"{_GEO}/matrix/GSE117303_series_matrix.txt.gz"},
)
RAW_FILES: Final = (PAPER_TEXT, MUTATIONS, TPM, SERIES_MATRIX)
PUBLICATION: Final = Publication(doi=PAPER_DOI, doi_url=f"https://doi.org/{PAPER_DOI}")

PAPER: Final = Provenance(
    source_uri=PAPER_TEXT.relpath, citation_key=CITATION_KEY, sha256=PAPER_TEXT.sha256
)


def _paper(value: Any, quote: str, note: str | None = None) -> SourcedValue:
    return SourcedValue(value=value, provenance=PAPER, quote=quote, note=note)


SOURCED_VALUES: Final[dict[str, SourcedValue]] = {
    "keio": _paper(
        KEIO_COLLECTION,
        "All strains utilized in this study were taken from the single-gene deletion "
        "Keio collection [7].",
    ),
    "wild_type": _paper(
        "BW25113 without deletion or cassette",
        "The reference strain utilized in growth screens and wherever ‘wild type’ is "
        "specified in this manuscript is the parent Keio strain without any deletions "
        "or Kanamycin resistance cassette.",
    ),
    "rnaseq_condition": _paper(
        {"glucose_g_per_l": 2.0, "kanamycin_mg_per_l": 50.0},
        "RNA sequencing data were generated under aerobic, exponential growth "
        "conditions on M9 minimal medium plus 2 g/L D-glucose and 50 mg/L Kanamycin A. "
        "The wild type BW25113 strain control was grown in the absence of antibiotics.",
        note="the media paragraph states 25 mg/L for the growth characterizations; this "
        "sentence is the one about these libraries",
    ),
    "temperature": _paper(
        37.0, "Both the pre-cultures and growth tests were grown at 37"
    ),
    "m9": _paper(
        {"CaCl2_mM": 0.1, "MgSO4_mM": 2.0},
        "The M9 minimal media utilized in the long term growth characterizations "
        "consisted of 0.1mM CaCl2, 2mM MgSO4, 1x Trace elements Solution, 1x M9 salts "
        "solution, and either 2g/L glucose",
    ),
    "m9_salts": _paper(
        "10x stock: 68 g/L Na2HPO2, 30 g/L KH2PO2, 5 g/L NaCl, 10 g/L NH4Cl",
        "The 10x M9 salt solution was composed of 68 g/L Na2HPO2, 30 g/L KH2PO2, 5 g/L "
        "NaCl, and 10 g/L NH4Cl, in Milli-Q water.",
        note="'HPO2' is the paper's spelling; the salts are stored as one component "
        "at 1x with the stock composition in its note",
    ),
    "duplicates": _paper(
        2,
        "For each condition examined, biological duplicate samples were sequenced and "
        "analyzed.",
    ),
    "replicon": _paper(
        REPLICON,
        "identify mutations relative to the E. coli BW25113 genome (NCBI accession "
        "CP009273.1)",
    ),
    "caller": _paper(CALLER, "The breseq pipeline [77] version 0.30.0"),
    "call_threshold": _paper(
        0.2,
        "Mutations considered for analysis in this study were present in a population "
        "at a fraction > 0.2.",
    ),
    "rnaseq_alignment": _paper(
        REPLICON,
        "RNA sequencing reads were mapped to the E. coli BW25113 genome (CP009273.1)",
    ),
    "geo": _paper(
        GEO_SERIES,
        "The RNAseq data is available in the Gene Expression Omnibus (GEO) database "
        "under the accession number GSE117303.",
    ),
}

#: The frequency note of Additional file 4's ``Description`` sheet, checked at read time.
FREQUENCY_NOTE: Final = (
    "Note: Mutation frequencies are reported as a fraction of the population "
    "containing the detected mutation (0-1.0)"
)
#: Breseq's ``Mutation Type`` cells in the rows this loader reads, and what each is.
VARIANT_TYPE: Final[dict[str, BacterialVariantType]] = {
    "SNP": BacterialVariantType.snv,
    "MOB": BacterialVariantType.insertion,
    "INS": BacterialVariantType.insertion,
    "DEL": BacterialVariantType.deletion,
}
LIBRARY: Final = re.compile(
    r"^(?P<gene>[A-Za-z]+)_A(?P<ale>\d+)_F(?P<flask>\d+)_I(?P<isolate>\d+)_R(?P<rep>\d)$"
)
WILD_TYPE: Final = re.compile(r"^wt_R(?P<rep>\d)$")
SOURCE_KIND: Final[dict[str, VariantCallMode]] = {
    "Clone": VariantCallMode.clone,
    "Population": VariantCallMode.population,
}

#: Build oracles, measured on the pinned files.
EXPECTED_RECORDS: Final = 14
EXPECTED_GENES: Final = 4329
EXPECTED_CALLS: Final = 12

_TPM_LOOKED_IN: Final = Provenance(
    source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{TPM.relpath}",
    citation_key=CITATION_KEY,
    sha256=TPM.sha256,
)
COUNT_GAP: Final = ProvenanceGap(
    field="expression_count",
    reason=ProvenanceGapReason.not_reported_by_primary,
    looked_in=_TPM_LOOKED_IN,
    note="GEO GSE117303's only processed file is a TPM matrix; the per-gene counts the "
    "Methods describe (summarizeOverlaps) were not deposited",
)


# --------------------------------------------------------------------------- #
# Reading the release
# --------------------------------------------------------------------------- #
class Library(BaseModel):
    """One released library column, parsed from its own label."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    column: str
    gsm: str
    source_name: str
    gene: str | None
    dna_sample: str | None
    call_mode: VariantCallMode | None


class Call(BaseModel):
    """One released Additional file 4 row with its frequency in one DNA sample."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    position: int
    mutation_type: str
    sequence_change: str
    gene: str
    details: str
    frequency_statement: str
    frequency: float


def read_tpm(path: str | Path) -> tuple[list[str], dict[str, dict[str, float]]]:
    """``(columns, {column: {gene id: TPM}})``, verbatim, refusing an empty cell."""
    with gzip.open(path, "rt") as handle:
        rows = list(csv.reader(handle))
    header = rows[0]
    if header[0] != "":
        raise ValueError(f"unexpected TPM header {header[:2]}")
    columns = header[1:]
    out: dict[str, dict[str, float]] = {c: {} for c in columns}
    for row in rows[1:]:
        gene = row[0]
        for column, cell in zip(columns, row[1:], strict=True):
            out[column][gene] = float(cell)
    return columns, out


def read_series_matrix(path: str | Path) -> dict[str, dict[str, str]]:
    """``{!Sample_* field: {library column: cell}}`` for the per-sample rows."""
    fields: dict[str, list[str]] = {}
    with gzip.open(path, "rt") as handle:
        for line in handle:
            if not line.startswith("!Sample_"):
                continue
            key, *cells = line.rstrip("\n").split("\t")
            fields.setdefault(key, [c.strip('"') for c in cells])
    titles = fields["!Sample_title"]
    columns = [title.split(": ", 1)[1] for title in titles]
    return {
        key: dict(zip(columns, cells, strict=True)) for key, cells in fields.items()
    }


def libraries(
    columns: Sequence[str], matrix: Mapping[str, Mapping[str, str]]
) -> list[Library]:
    """Each TPM column with its GEO accession, its strain and its call mode."""
    out: list[Library] = []
    for column in columns:
        source = matrix["!Sample_source_name_ch1"][column]
        gsm = matrix["!Sample_geo_accession"][column]
        if WILD_TYPE.match(column):
            if source != "wild type":
                raise ValueError(f"{column}: source {source!r}")
            out.append(
                Library(
                    column=column,
                    gsm=gsm,
                    source_name=source,
                    gene=None,
                    dna_sample=None,
                    call_mode=None,
                )
            )
            continue
        match = LIBRARY.match(column)
        if match is None:
            raise ValueError(f"unparsed library column {column!r}")
        kind = source.rsplit(", ", 1)[1]
        out.append(
            Library(
                column=column,
                gsm=gsm,
                source_name=source,
                gene=match["gene"],
                dna_sample=f"FPS {match['gene']} A{match['ale']} F{match['flask']} "
                f"I{match['isolate']} R1",
                call_mode=SOURCE_KIND[kind],
            )
        )
    return out


def read_calls(path: str | Path, samples: Sequence[str]) -> dict[str, list[Call]]:
    """Every Additional file 4 call with a released frequency in each named sample."""
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    note = [
        row[0]
        for row in workbook["Description"].iter_rows(values_only=True)
        if row[0] is not None
    ]
    if FREQUENCY_NOTE not in note:
        raise ValueError("Additional file 4 no longer states its frequency basis")
    out: dict[str, list[Call]] = {}
    for sample in samples:
        sheet = workbook[sample.split()[1]]
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0]
        if header[:5] != (
            "Position",
            "Mutation Type",
            "Sequence Change",
            "Gene",
            "Details",
        ):
            raise ValueError(f"sheet {sheet.title}: header {header[:5]}")
        index = header.index(sample)
        out[sample] = [
            Call(
                position=int(row[0]),
                mutation_type=str(row[1]),
                sequence_change=str(row[2]),
                gene=str(row[3]),
                details=str(row[4]),
                frequency_statement=str(row[index]),
                frequency=float(row[index]),
            )
            for row in rows[1:]
            if row[index] is not None
        ]
    workbook.close()
    return out


# --------------------------------------------------------------------------- #
# Records (pure, given the resolved loci)
# --------------------------------------------------------------------------- #
def is_intergenic(call: Call) -> bool:
    """Breseq writes ``intergenic (...)`` in Details for a call between loci."""
    return call.details.startswith("intergenic")


def variant_call(call: Call, mode: VariantCallMode) -> BacterialVariantCall:
    """The released row as a call, positions read off the released cells."""
    kind = VARIANT_TYPE[call.mutation_type]
    end = call.position - 1 if kind is BacterialVariantType.insertion else call.position
    return BacterialVariantCall(
        variant_type=kind,
        type_statement=call.mutation_type,
        reference_sequence=REPLICON,
        position_start=call.position,
        position_end=end,
        sequence_change=call.sequence_change,
        annotation=call.details,
        call_mode=mode,
        frequency_statement=call.frequency_statement,
        frequency=call.frequency,
        frequency_basis=VariantFrequencyBasis.fraction,
        caller=CALLER,
    )


def call_perturbation(
    call: Call, mode: VariantCallMode, loci: Mapping[str, str]
) -> GenePerturbationType:
    """An intergenic call keyed on its site, any other on its released gene's locus."""
    if is_intergenic(call):
        site = f"{REPLICON}:{call.position}"
        return SiteVariantCallPerturbation(
            systematic_gene_name=site,
            perturbed_gene_name=site,
            gene_namespace=GENE_NAMESPACE,
            call=variant_call(call, mode),
            site_kind=VariantSiteKind.intergenic,
            flanking_gene_statement=call.gene,
        )
    return LocusVariantCallPerturbation(
        systematic_gene_name=loci[call.gene],
        perturbed_gene_name=call.gene,
        gene_namespace=GENE_NAMESPACE,
        identifier_mapping=DerivedIdentifierMapping(
            source_identifier=call.gene, route="gene_symbol"
        ),
        call=variant_call(call, mode),
    )


def genotype(
    library: Library, calls: Sequence[Call], loci: Mapping[str, str]
) -> Genotype:
    """Wild type is empty; a deletion strain is its Keio deletion plus its calls."""
    if library.gene is None or library.call_mode is None:
        return Genotype(perturbations=[])
    perturbations: list[GenePerturbationType] = [
        DeletionStrainPerturbation(
            systematic_gene_name=loci[library.gene],
            perturbed_gene_name=library.gene,
            gene_namespace=GENE_NAMESPACE,
            identifier_mapping=DerivedIdentifierMapping(
                source_identifier=library.gene, route="gene_symbol"
            ),
            collection=KEIO_COLLECTION,
            cassette=KEIO_CASSETTE,
        )
    ]
    perturbations.extend(call_perturbation(c, library.call_mode, loci) for c in calls)
    return Genotype(perturbations=perturbations)


def medium(*, kanamycin: bool) -> Media:
    """M9 + 2 g/L D-glucose as the Methods state it; kanamycin for the deletions."""
    m9 = SOURCED_VALUES["m9"]
    condition = SOURCED_VALUES["rnaseq_condition"]
    components = [
        MediaComponent(
            compound=resolved_compound("D-glucose"),
            role=MediaComponentRole.carbon_source,
            concentration=Concentration(value=2.0, unit=ConcentrationUnit.g_per_l),
            provenance=[condition],
        ),
        MediaComponent(
            compound=resolved_compound("calcium chloride"),
            role=MediaComponentRole.bulk_salt,
            concentration=Concentration(value=0.1, unit=ConcentrationUnit.millimolar),
            provenance=[m9],
        ),
        MediaComponent(
            compound=resolved_compound("magnesium sulfate"),
            role=MediaComponentRole.bulk_salt,
            concentration=Concentration(value=2.0, unit=ConcentrationUnit.millimolar),
            provenance=[m9],
        ),
    ]
    if kanamycin:
        components.append(
            MediaComponent(
                compound=resolved_compound("kanamycin"),
                role=MediaComponentRole.selection_agent,
                concentration=Concentration(
                    value=50.0, unit=ConcentrationUnit.ug_per_ml
                ),
                provenance=[condition],
                note="50 mg/L Kanamycin A, stated for the RNA-seq cultures",
            )
        )
    return Media(
        name="M9 minimal medium, 2 g/L D-glucose"
        + (", 50 mg/L kanamycin A" if kanamycin else "")
        + " (Guzman 2018; 1x M9 salts and 1x trace elements, composition per its "
        "Methods)",
        state="liquid",
        is_synthetic=True,
        base_medium="M9",
        components=components,
        provenance=[m9, SOURCED_VALUES["m9_salts"], condition],
    )


def environment(*, kanamycin: bool) -> Environment:
    """Aerobic exponential culture at 37 C."""
    return Environment(
        media=medium(kanamycin=kanamycin),
        temperature=Temperature(value=float(SOURCED_VALUES["temperature"].value)),
        aerobicity="aerobic",
    )


def phenotype(tpm: Mapping[str, float]) -> RNASeqExpressionPhenotype:
    """One released library's TPM, verbatim, with the count gap."""
    return RNASeqExpressionPhenotype(
        expression_tpm=dict(tpm),
        expression_count=None,
        measurement_type=MEASUREMENT_TYPE,
        provenance_gaps=[COUNT_GAP],
    )


def reference_phenotype(
    wild_type: Sequence[Mapping[str, float]],
) -> RNASeqExpressionPhenotype:
    """The mean of the released wild-type libraries, with the count gap."""
    genes = wild_type[0].keys()
    return RNASeqExpressionPhenotype(
        expression_tpm={
            g: math.fsum(lib[g] for lib in wild_type) / len(wild_type) for g in genes
        },
        expression_count=None,
        measurement_type=MEASUREMENT_TYPE,
        provenance_gaps=[COUNT_GAP],
    )


# --------------------------------------------------------------------------- #
# Raw mirror
# --------------------------------------------------------------------------- #
def _data_root() -> str:
    return os.environ["DATA_ROOT"]


def raw_mirror_dir(data_root: str | None = None) -> Path:
    """``$DATA_ROOT/torchcell-raw/guzmanReframingGeneEssentiality2018``."""
    return Path(data_root or _data_root()) / RAW_DIR_REL


def deposit_raw_mirror(
    sources: Mapping[str, str | Path], data_root: str | None = None
) -> Path:
    """Copy the four verified files into the raw mirror and write ``manifest.json``."""
    root = raw_mirror_dir(data_root)
    records: list[ArtifactRecord] = []
    for raw in RAW_FILES:
        src = Path(sources[raw.name])
        verify_sha256(src, raw.sha256)
        dest = root / raw.relpath
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            verify_sha256(dest, raw.sha256)
        else:
            shutil.copy2(src, dest)
        records.append(
            ArtifactRecord(
                path=raw.relpath,
                role="raw_data",
                bytes=dest.stat().st_size,
                sha256=raw.sha256,
                source=raw.source_url,
                original_filename=raw.name,
                retrieval=RetrievalRecord(
                    method=raw.method,
                    source_url=raw.source_url,
                    retriever=raw.retriever,
                    params=raw.params,
                    sha256=raw.sha256,
                    retrieved_at=RETRIEVED_AT,
                ),
            )
        )
    manifest = Manifest(
        citation_key=CITATION_KEY,
        doi=PAPER_DOI,
        title=PAPER_TITLE,
        files=records,
        si_data_sources=[raw.source_url for raw in RAW_FILES],
        provenance_complete=True,
        created_at=datetime.now(UTC).isoformat(),
    )
    (root / "manifest.json").write_text(manifest.model_dump_json(indent=2))
    return root


def manifest_sha256(relpath: str, data_root: str | None = None) -> str:
    """The sha256 the raw mirror's manifest records for one file."""
    manifest = Manifest.model_validate_json(
        (raw_mirror_dir(data_root) / "manifest.json").read_text()
    )
    return next(f.sha256 for f in manifest.files if f.path == relpath)


# --------------------------------------------------------------------------- #
# The dataset
# --------------------------------------------------------------------------- #
@register_dataset
class RnaseqGuzman2018Dataset(ExperimentDataset):
    """Guzmán 2018's 14 released RNA-seq libraries, from its own GEO release."""

    organism = Organism.ecoli

    REFERENCE_STRAIN: ClassVar[EcoliK12StrainName] = REFERENCE_STRAIN

    def __init__(
        self,
        root: str = DATASET_ROOT_REL,
        io_workers: int = 0,
        transform: Callable[..., Any] | None = None,
        pre_transform: Callable[..., Any] | None = None,
        ecoli_genome: EcoliK12Genome | None = None,
        **kwargs: Any,
    ) -> None:
        """Initialize; ``ecoli_genome`` is injected by the build entry points."""
        self.ecoli_genome = ecoli_genome
        super().__init__(root, io_workers, transform, pre_transform, **kwargs)

    @property
    def experiment_class(self) -> type[Experiment]:
        """Experiment schema class produced by this dataset."""
        return RNASeqExpressionExperiment

    @property
    def reference_class(self) -> type[ExperimentReference]:
        """Experiment-reference schema class produced by this dataset."""
        return RNASeqExpressionExperimentReference

    @property
    def raw_file_names(self) -> list[str]:
        """The three release files the records are built from."""
        return [MUTATIONS.name, TPM.name, SERIES_MATRIX.name]

    def download(self) -> None:
        """Link the mirrored files after checking their manifest pins and sha256."""
        data_root = _data_root()
        os.makedirs(self.raw_dir, exist_ok=True)
        for raw in (MUTATIONS, TPM, SERIES_MATRIX):
            check_manifest_pin(
                raw.relpath, manifest_sha256(raw.relpath, data_root), raw.sha256
            )
            link_verified(
                raw_mirror_dir(data_root) / raw.relpath,
                osp.join(self.raw_dir, raw.name),
                raw.sha256,
            )

    def _genome(self) -> EcoliK12Genome:
        if self.ecoli_genome is None:
            self.ecoli_genome = bacterial_genome("ecoli", REFERENCE_STRAIN)
        return self.ecoli_genome

    @post_process
    def process(self) -> None:
        """One record per released library."""
        verify_raw_files(
            self.raw_dir, {r.name: r.sha256 for r in (MUTATIONS, TPM, SERIES_MATRIX)}
        )
        columns, tpm = read_tpm(osp.join(self.raw_dir, TPM.name))
        matrix = read_series_matrix(osp.join(self.raw_dir, SERIES_MATRIX.name))
        libs = libraries(columns, matrix)
        samples = sorted({lib.dna_sample for lib in libs if lib.dna_sample})
        calls = read_calls(osp.join(self.raw_dir, MUTATIONS.name), samples)
        n_calls = sum(len(c) for c in calls.values())
        if len(libs) != EXPECTED_RECORDS or n_calls != EXPECTED_CALLS:
            raise RuntimeError(f"{len(libs)} libraries and {n_calls} calls")
        if {len(v) for v in tpm.values()} != {EXPECTED_GENES}:
            raise RuntimeError("a library does not carry the pinned gene count")
        loci = resolve_symbols(self._genome(), libs, calls)
        wild = [tpm[lib.column] for lib in libs if lib.gene is None]
        reference = RNASeqExpressionExperimentReference(
            dataset_name=self.name,
            genome_reference=assembly_reference(REFERENCE_STRAIN),
            environment_reference=environment(kanamycin=False),
            phenotype_reference=reference_phenotype(wild),
        )
        os.makedirs(self.processed_dir, exist_ok=True)
        env, interned_env = self._open_write_lmdb(osp.join(self.processed_dir, "lmdb"))
        with env.begin(write=True) as txn, interned_env.begin(write=True) as itxn:
            for index, lib in enumerate(libs):
                experiment = RNASeqExpressionExperiment(
                    dataset_name=self.name,
                    genotype=genotype(lib, calls.get(lib.dna_sample or "", []), loci),
                    environment=environment(kanamycin=lib.gene is not None),
                    phenotype=phenotype(tpm[lib.column]),
                )
                txn.put(
                    f"{index}".encode(),
                    self._intern_record(experiment, reference, PUBLICATION, itxn),
                )
        env.close()
        interned_env.close()
        out = Path(self.preprocess_dir)
        out.mkdir(parents=True, exist_ok=True)
        (out / "libraries.json").write_text(
            json.dumps(
                {
                    "libraries": [lib.model_dump(mode="json") for lib in libs],
                    "calls": {k: [c.model_dump() for c in v] for k, v in calls.items()},
                    "loci": loci,
                },
                indent=2,
            )
        )

    def preprocess_raw(self, df: Any, preprocess: dict[str, Any] | None = None) -> Any:
        """Preprocessing is handled inside process() for this dataset."""
        return df

    def create_experiment(self) -> None:
        """Experiment construction is handled inline in process()."""
        raise NotImplementedError("built in process()")


def resolve_symbols(
    genome: EcoliK12Genome, libs: Sequence[Library], calls: Mapping[str, Sequence[Call]]
) -> dict[str, str]:
    """The released symbols of the deleted genes and in-gene calls, on BW25113 loci.

    Keyed through the genome's name layers (route ``gene_symbol``); every symbol must
    reach exactly one locus, and the released symbol stays on each record.
    """
    symbols = {lib.gene for lib in libs if lib.gene}
    symbols |= {c.gene for v in calls.values() for c in v if not is_intergenic(c)}
    out: dict[str, str] = {}
    for symbol in sorted(symbols):
        resolution = genome.resolve_gene_name(symbol)
        tag = resolution.systematic_name
        if resolution.status.value not in {"current", "renamed"} or tag is None:
            raise RuntimeError(f"{symbol!r} does not reach one BW25113 locus")
        out[symbol] = tag
    return out


# --------------------------------------------------------------------------- #
# L0-L4 verification
# --------------------------------------------------------------------------- #
Record = Mapping[str, Any]

VERIFIER_PROVENANCE: Final = Provenance(
    source_uri=f"$DATA_ROOT/{RAW_DIR_REL}/{TPM.relpath}",
    citation_key=CITATION_KEY,
    sha256=TPM.sha256,
    method="one RNASeqExpressionExperiment per column of GEO GSE117303's "
    "deseq2-tpm matrix, TPM verbatim, counts a typed gap; genotypes from Additional "
    "file 4's calls for the matching DNA sample",
    page="GSE117303",
)


def _l2_tpm_equals_release(
    records: Sequence[Record], tpm: Mapping[str, Mapping[str, float]]
) -> LevelResult:
    stored = sorted(
        (tuple(sorted(r["experiment"]["phenotype"]["expression_tpm"].items())))
        for r in records
    )
    released = sorted(tuple(sorted(v.items())) for v in tpm.values())
    return LevelResult(
        level=Level.L2,
        name="tpm_fidelity_to_release",
        passed=stored == released,
        message=f"{len(records)} stored TPM vectors equal the {len(tpm)} released "
        "columns as a multiset"
        if stored == released
        else "a stored vector differs",
    )


def _l3_genotypes(records: Sequence[Record]) -> LevelResult:
    kinds = sorted(len(r["experiment"]["genotype"]["perturbations"]) for r in records)
    n_wild = sum(k == 0 for k in kinds)
    return LevelResult(
        level=Level.L3,
        name="genotype_shape",
        passed=n_wild == 2,
        message=f"{n_wild} wild-type libraries; perturbation counts {kinds}",
    )


def _l4_crosscheck(
    tpm: Mapping[str, Mapping[str, float]], data_root: str
) -> LevelResult:
    """L4 (a finding): log2(x+1) Pearson r against PRECISE-1K's copy of each library."""
    import numpy as np
    import pandas as pd

    from torchcell.datasets.bacteria_common import eck_crosswalk
    from torchcell.datasets.ecoli import lamoureux2023 as lm

    meta = lm.read_metadata(str(lm.raw_mirror_dir(data_root) / lm.METADATA.relpath))
    fps = meta[meta["project"] == "fps"]
    compendium = pd.read_csv(
        lm.raw_mirror_dir(data_root) / lm.LOG_TPM.relpath, index_col=0
    )
    mg1655 = bacterial_genome("ecoli", "MG1655", data_root)
    bw25113 = bacterial_genome("ecoli", "BW25113", data_root)
    if not isinstance(mg1655, EcoliK12MG1655Genome) or not isinstance(
        bw25113, EcoliK12BW25113Genome
    ):
        raise TypeError("the K-12 genomes did not load as their strain classes")
    walk = eck_crosswalk(mg1655, bw25113)
    to_b = {pair.bw25113: pair.mg1655 for pair in walk.pairs}
    r_by_library: dict[str, float] = {}
    for column, values in tpm.items():
        match = LIBRARY.match(column)
        rep = (WILD_TYPE.match(column) or match)["rep"]  # type: ignore[index]
        label = (
            f"fps__bw_glc__{rep}"
            if match is None
            else f"fps__fps_{match['gene']}_ale{match['ale']}__{rep}"
        )
        sample = fps.index[fps["sample_id"] == label]
        if len(sample) != 1 or sample[0] not in compendium.columns:
            continue
        pairs = [
            (math.log2(v + 1.0), float(str(compendium.at[to_b[g], sample[0]])))
            for g, v in values.items()
            if g in to_b and to_b[g] in compendium.index
        ]
        x, y = np.array(pairs).T
        r_by_library[column] = float(np.corrcoef(x, y)[0, 1])
    floor = min(r_by_library.values()) if r_by_library else float("nan")
    return LevelResult(
        level=Level.L4,
        name="crosscheck_precise1k_copy",
        passed=True,
        message=f"{len(r_by_library)} of {len(tpm)} libraries matched to PRECISE-1K; "
        f"min Pearson r of log2(TPM+1) {floor:.4f} (a finding, never a gate)",
        details={"pearson_r": r_by_library},
    )


def verify_build(dataset_root: str, data_root: str | None = None) -> VerificationReport:
    """The RNA-seq L0-L3 gate, TPM fidelity, genotype shape, and the L4 cross-check."""
    from torchcell.verification.rnaseq import verify_rnaseq_dataset
    from torchcell.verification.runners import load_records

    base = data_root or _data_root()
    records = load_records(dataset_root)
    report = verify_rnaseq_dataset(
        records,
        dataset_name=osp.basename(osp.normpath(dataset_root)),
        provenance=VERIFIER_PROVENANCE,
        expected_count=EXPECTED_RECORDS,
        replicate_aware=True,
    )
    _, tpm = read_tpm(raw_mirror_dir(base) / TPM.relpath)
    report.add(_l2_tpm_equals_release(records, tpm))
    report.add(_l3_genotypes(records))
    for value in SOURCED_VALUES.values():
        report.add(audit_sourced_value(value, Path(base) / "torchcell-raw"))
    report.add(_l4_crosscheck(tpm, base))
    out = osp.join(dataset_root, "preprocess", "verification_report.json")
    os.makedirs(osp.dirname(out), exist_ok=True)
    with open(out, "w") as handle:
        handle.write(report.model_dump_json(indent=2))
    return report


def main(argv: list[str] | None = None) -> int:
    """CLI: ``deposit <dir>``, ``build`` or ``verify``."""
    from dotenv import load_dotenv

    parser = argparse.ArgumentParser(
        prog="python -m torchcell.datasets.ecoli.guzman2018"
    )
    parser.add_argument("command", choices=("deposit", "build", "verify"))
    parser.add_argument("source_dir", nargs="?")
    args = parser.parse_args(argv)
    load_dotenv()
    data_root = _data_root()
    root = osp.join(data_root, DATASET_ROOT_REL)
    if args.command == "deposit":
        src = Path(args.source_dir)
        print(deposit_raw_mirror({r.name: src / r.name for r in RAW_FILES}, data_root))
        return 0
    if args.command == "build":
        print(f"len = {len(RnaseqGuzman2018Dataset(root=root))}")
        return 0
    report = verify_build(root, data_root)
    print(report.summary())
    return 0 if report.passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
