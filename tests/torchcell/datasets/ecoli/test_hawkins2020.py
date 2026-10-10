# tests/torchcell/datasets/ecoli/test_hawkins2020.py
# [[tests.torchcell.datasets.ecoli.test_hawkins2020]]
# https://github.com/Mjvolk3/torchcell/tree/main/tests/torchcell/datasets/ecoli/test_hawkins2020.py
"""Hawkins 2020 mismatch-CRISPRi loader (``torchcell.datasets.ecoli.hawkins2020``).

The hermetic tests write a workbook carrying Table S3's own E. coli header and one row
per behaviour the loader has to get right: a fully complementary spacer, a singly
mismatched variant of it, a row with a mean and no SD, a row with no mean at all, and a
row on a gene that reaches no BW25113 locus. The genomes are the real MG1655 and BW25113
classes over the synthetic assemblies of
``tests/torchcell/sequence/genome/_bacterial_fixtures.py``, so nothing reads
``$DATA_ROOT`` and no network call is permitted.

The ``@pytest.mark.data`` tests read the pinned mirror and assert the numbers the module,
the dendron note and the PR state: 36,290 released rows, 24,149 stored, 12,104 rows with
no released fitness, 37 on ``sokE``, 317 genes, 778 negative values, 1,870 records with
no SD, and a predicted sgRNA activity of exactly 1.0 on every one of the 3,110 parent
spacers.
"""

from __future__ import annotations

import os
import os.path as osp
from collections.abc import Mapping
from pathlib import Path
from typing import Any, cast

import openpyxl
import pandas as pd
import pytest

import torchcell.datasets.ecoli.hawkins2020 as h
from tests.torchcell.sequence.genome._bacterial_fixtures import (
    BW25113_LOCI,
    MG1655_GAF,
    MG1655_LOCI,
    forbid_network,
    serve_tier,
    write_assembly,
)
from torchcell.datamodels.schema import (
    BacterialCrisprInterferencePerturbation,
    ConcentrationUnit,
    MeasurementType,
    MediaComponentRole,
    SampleUnit,
    SmallMoleculePerturbation,
    UncertaintyType,
)
from torchcell.sequence.genome.ecoli.k12 import (
    BW25113_ASSEMBLY,
    MG1655_ASSEMBLY,
    EcoliK12BW25113Genome,
    EcoliK12MG1655Genome,
)

#: ``(variant, original, locus_tag, pam, offset, gene, mean, sd, family, predicted)``.
#: ``b0001`` crosses to ``BW25113_0001`` on ECK0001; ``b0005`` carries ECK0005, which the
#: synthetic BW25113 annotation puts on two loci, so it is NOT one-to-one and reaches no
#: locus.
PARENT = "AAAAAAAAAAAAAAAAAAAA"
MISMATCHED = "AAAAAAAAAAAAAAAAAAAC"
NO_SD = "AAAAAAAAAAAAAAAAAACA"
NO_MEAN = "AAAAAAAAAAAAAAAAACAA"
UNRESOLVED = "AAAAAAAAAAAAAAAACAAA"
SYNTHETIC: tuple[tuple[Any, ...], ...] = (
    (PARENT, PARENT, "b0001", "CGG", 7, "thrL", 0.42, 0.03, True, 1.0),
    (MISMATCHED, PARENT, "b0001", "CGG", 7, "thrL", 0.91, 0.07, True, 0.33),
    (NO_SD, PARENT, "b0001", "CGG", 7, "thrL", -0.11, None, False, 0.80),
    (NO_MEAN, PARENT, "b0001", "CGG", 7, "thrL", None, None, True, 0.61),
    (UNRESOLVED, PARENT, "b0005", "TGG", 9, "proB", 0.77, 0.05, True, 0.55),
)
STORED = (PARENT, MISMATCHED, NO_SD)


def _workbook(path: Path, rows: tuple[tuple[Any, ...], ...] = SYNTHETIC) -> Path:
    """A workbook with Table S3's E. coli sheet name, header and the given rows."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = h.SHEET_ECOLI
    sheet.append(list(h.ECOLI_COLUMNS))
    for row in rows:
        sheet.append([*row, None, None])
    book.save(path)
    return path


@pytest.fixture
def workbook(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """The synthetic workbook, with the module's row oracle repointed at its shape."""
    monkeypatch.setattr(h, "EXPECTED_SOURCE_ROWS", len(SYNTHETIC))
    return _workbook(tmp_path / "si4.xlsx")


@pytest.fixture
def mg1655(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12MG1655Genome:
    """The real MG1655 class over the synthetic assembly; network refused."""
    tier = write_assembly(
        tmp_path / "mg-tier", MG1655_ASSEMBLY, MG1655_LOCI, MG1655_GAF
    )
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, tier)
    return EcoliK12MG1655Genome(genome_root=str(tmp_path / "mg1655"), overwrite=True)


@pytest.fixture
def bw25113(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> EcoliK12BW25113Genome:
    """The real BW25113 class over the synthetic assembly; network refused."""
    tier = write_assembly(tmp_path / "bw-tier", BW25113_ASSEMBLY, BW25113_LOCI)
    forbid_network(monkeypatch)
    serve_tier(monkeypatch, tier)
    return EcoliK12BW25113Genome(genome_root=str(tmp_path / "bw25113"), overwrite=True)


# --------------------------------------------------------------------------- #
# The declared constants and what is refused
# --------------------------------------------------------------------------- #
def test_declared_counts_are_the_measured_ones() -> None:
    assert h.EXPECTED_SOURCE_ROWS == 36290
    assert h.EXPECTED_RECORDS == 24149
    assert h.EXPECTED_DROPS == {h.DROP_NO_FITNESS: 12104, h.DROP_NO_LOCUS: 37}
    assert h.EXPECTED_RECORDS + sum(h.EXPECTED_DROPS.values()) == h.EXPECTED_SOURCE_ROWS
    assert (h.EXPECTED_RELEASED_GENES, h.EXPECTED_GENES) == (318, 317)
    assert h.EXPECTED_PARENT_SPACERS == 3110
    assert h.EXPECTED_STORED_PARENTS == 2948
    assert h.UNRESOLVED_BNUMBER == "b4700"


def test_the_predicted_dose_is_sourced_as_a_model_output_and_never_stored() -> None:
    """The row's schema need: the knockdown LEVEL is a prediction, so it has no slot."""
    predicted = h.SOURCED_VALUES["predicted_activity"]
    assert predicted.value == "predicted sgRNA activity"
    assert "model output" in str(predicted.note)
    assert h.MODEL_FIT.value == (0.56, 0.10, 0.08)
    assert "0 . 5 6" in h.MODEL_FIT.quote
    assert "preliminary version" in h.PRELIMINARY_MODEL.quote
    assert h.COL_PREDICTED not in {
        field for field in h.RowVerdict.model_fields if field.endswith("stored")
    }
    # it survives in the ledger column name, which says what it is and is not
    assert "predicted_sgrna_activity_not_stored"


def test_the_readout_is_a_ratio_whose_control_is_one_not_zero() -> None:
    assert h.FITNESS_DEFINITION.value == 1.0
    assert "relative fitness of 1 grow as well as the wild-type" in (
        h.FITNESS_DEFINITION.quote
    )
    assert h.reference_phenotype().environment_response == 1.0
    assert h.reference_phenotype().measurement_type == (
        MeasurementType.relative_growth_rate
    )
    assert "relative fitness $< 0$" in h.NEGATIVE_FITNESS.quote


def test_read_ecoli_sheet_refuses_a_changed_header_or_row_count(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(h, "EXPECTED_SOURCE_ROWS", len(SYNTHETIC))
    path = _workbook(tmp_path / "ok.xlsx")
    assert len(h.read_ecoli_sheet(path)) == len(SYNTHETIC)

    monkeypatch.setattr(h, "EXPECTED_SOURCE_ROWS", len(SYNTHETIC) + 1)
    with pytest.raises(RuntimeError, match="not the pinned"):
        h.read_ecoli_sheet(path)

    book = openpyxl.load_workbook(path)
    book[h.SHEET_ECOLI].cell(row=1, column=1).value = "spacer"
    renamed = tmp_path / "renamed.xlsx"
    book.save(renamed)
    monkeypatch.setattr(h, "EXPECTED_SOURCE_ROWS", len(SYNTHETIC))
    with pytest.raises(RuntimeError, match="header is"):
        h.read_ecoli_sheet(renamed)


# --------------------------------------------------------------------------- #
# The mismatch design, which is what the record carries instead of the dose
# --------------------------------------------------------------------------- #
def test_a_fully_complementary_spacer_carries_no_mismatch() -> None:
    design = h.mismatch_design(PARENT, PARENT)
    assert (design.mismatch_index, design.substitution) == (None, None)
    assert design.description.endswith(
        f"fully complementary spacer {PARENT} (the parent of its series)"
    )


def test_a_single_mismatch_states_its_position_and_substitution() -> None:
    design = h.mismatch_design(MISMATCHED, PARENT)
    assert (design.mismatch_index, design.substitution) == (19, "A to C")
    assert "0-based position 19 from the 5' end (A to C)" in design.description
    assert "titrates knockdown below the parent's" in design.description


def test_mismatch_design_refuses_a_double_mismatch_and_a_short_spacer() -> None:
    with pytest.raises(RuntimeError, match="differs from parent"):
        h.mismatch_design("CCAAAAAAAAAAAAAAAAAA", PARENT)
    with pytest.raises(RuntimeError, match="not a 20-nt ACGT sequence"):
        h.mismatch_design("ACGT", PARENT)
    with pytest.raises(RuntimeError, match="parent spacer"):
        h.mismatch_design(PARENT, "ACGTN")


# --------------------------------------------------------------------------- #
# Released b-number -> BW25113 locus
# --------------------------------------------------------------------------- #
def test_resolve_targets_crosses_every_b_number_on_its_one_to_one_eck_pair(
    mg1655: EcoliK12MG1655Genome, bw25113: EcoliK12BW25113Genome
) -> None:
    report = h.resolve_targets(["b0001", "b0003", "b0005"], mg1655, bw25113)
    assert report.resolved["b0001"].bw25113_locus == "BW25113_0001"
    assert report.resolved["b0001"].eck == "ECK0001"
    assert report.resolved["b0001"].symbol == "thrL"
    assert report.resolved["b0001"].numerics_agree is True
    # ECK0003 sits on BW25113_4412, whose number disagrees with b0003: the crosswalk is
    # the only route, and the record records that it was used
    assert report.resolved["b0003"].bw25113_locus == "BW25113_4412"
    assert report.resolved["b0003"].numerics_agree is False
    # ECK0005 is carried by two BW25113 loci, so it is not one-to-one and nothing is
    # guessed
    assert "b0005" not in report.resolved
    assert "no one-to-one ECK partner" in report.unresolved["b0005"]
    assert set(report.stored) == {"b0001", "b0003"}


def test_resolve_targets_reports_a_retired_b_number_rather_than_guessing(
    mg1655: EcoliK12MG1655Genome, bw25113: EcoliK12BW25113Genome
) -> None:
    report = h.resolve_targets(["b0001", "b9999"], mg1655, bw25113)
    assert "retired in the pinned MG1655 assembly" in report.unresolved["b9999"]
    assert set(report.stored) == {"b0001"}


def test_resolve_targets_refuses_an_identifier_that_is_not_a_b_number(
    mg1655: EcoliK12MG1655Genome, bw25113: EcoliK12BW25113Genome
) -> None:
    with pytest.raises(RuntimeError, match="is not a b-number"):
        h.resolve_targets(["thrL"], mg1655, bw25113)


# --------------------------------------------------------------------------- #
# Retention and the record
# --------------------------------------------------------------------------- #
def _verdicts(
    workbook: Path, mg1655: EcoliK12MG1655Genome, bw25113: EcoliK12BW25113Genome
) -> dict[str, h.RowVerdict]:
    frame = h.read_ecoli_sheet(workbook)
    report = h.resolve_targets([str(t) for t in frame[h.COL_LOCUS]], mg1655, bw25113)
    return {
        verdict.spacer: verdict
        for verdict in (
            h.classify_row(cast(Mapping[str, Any], row), report.stored)
            for row in frame.to_dict(orient="records")
        )
    }


def test_classify_row_drops_a_blank_mean_and_an_unresolvable_gene(
    workbook: Path, mg1655: EcoliK12MG1655Genome, bw25113: EcoliK12BW25113Genome
) -> None:
    verdicts = _verdicts(workbook, mg1655, bw25113)
    assert [spacer for spacer, v in verdicts.items() if not v.drop_reason] == list(
        STORED
    )
    assert verdicts[NO_MEAN].drop_reason == h.DROP_NO_FITNESS
    assert verdicts[UNRESOLVED].drop_reason == h.DROP_NO_LOCUS
    # a dropped row still carries what the release said, so the drop is reversible
    assert verdicts[NO_MEAN].predicted_activity == 0.61
    assert verdicts[UNRESOLVED].fitness == 0.77
    assert verdicts[NO_SD].sd is None
    assert verdicts[NO_SD].family_retained is False
    assert verdicts[PARENT].bw25113_locus == "BW25113_0001"


def test_the_record_carries_the_mismatch_design_and_the_derived_locus(
    workbook: Path, mg1655: EcoliK12MG1655Genome, bw25113: EcoliK12BW25113Genome
) -> None:
    verdicts = _verdicts(workbook, mg1655, bw25113)
    report = h.resolve_targets(["b0001"], mg1655, bw25113)
    genotype = h.build_genotype(
        h.mismatch_design(MISMATCHED, PARENT), report.stored["b0001"], "thrL"
    )
    (leaf,) = genotype.perturbations
    assert isinstance(leaf, BacterialCrisprInterferencePerturbation)
    assert leaf.systematic_gene_name == "BW25113_0001"
    assert leaf.perturbed_gene_name == "thrL"
    assert leaf.gene_namespace == "ecoli_k12_bw25113_locus_tag"
    assert leaf.expression_direction == "decreased"
    assert leaf.identifier_mapping is not None
    assert (
        leaf.identifier_mapping.source_identifier,
        leaf.identifier_mapping.route,
    ) == ("b0001", "eck_crosswalk")
    assert leaf.crispr.guide_sequence == MISMATCHED
    assert leaf.crispr.effector == "dCas9"
    assert leaf.crispr.n_guides == 1
    assert leaf.crispr.library_pool == h.LIBRARY_POOL
    assert f"parent spacer {PARENT}" in leaf.description
    assert verdicts[MISMATCHED].b_number == "b0001"


def test_build_genotype_refuses_a_target_with_no_locus() -> None:
    target = h.TargetResolution(
        b_number="b4700",
        mg1655_locus=None,
        eck=None,
        bw25113_locus=None,
        symbol=None,
        numerics_agree=None,
    )
    with pytest.raises(RuntimeError, match="reached no BW25113 locus"):
        h.build_genotype(h.mismatch_design(PARENT, PARENT), target, "sokE")


# --------------------------------------------------------------------------- #
# Phenotype, reference and environment
# --------------------------------------------------------------------------- #
def test_a_released_sd_is_a_sample_sd_over_four_replicates() -> None:
    phenotype = h.guide_phenotype(0.91, 0.07)
    assert phenotype.environment_response == 0.91
    assert phenotype.environment_response_uncertainty == 0.07
    assert phenotype.environment_response_uncertainty_type == UncertaintyType.sample_sd
    assert phenotype.n_samples == 4
    assert phenotype.sample_unit == SampleUnit.biological_replicate
    assert phenotype.environment_response_se == pytest.approx(0.07 / 2.0)
    assert phenotype.provenance_gaps == []
    assert phenotype.screen_id == h.SCREEN_ID


def test_a_blank_sd_is_three_typed_gaps_and_never_a_zero() -> None:
    phenotype = h.guide_phenotype(-0.11, None)
    assert phenotype.environment_response == -0.11
    assert phenotype.environment_response_uncertainty is None
    assert phenotype.environment_response_se is None
    assert [gap.field for gap in phenotype.provenance_gaps] == [
        "environment_response_se",
        "environment_response_uncertainty",
        "environment_response_uncertainty_type",
    ]
    assert all(
        gap.note is not None and "100-counts-at-t0 floor" in gap.note
        for gap in phenotype.provenance_gaps
    )


def test_the_reference_is_the_non_targeting_one_with_no_borrowed_spread() -> None:
    reference = h.reference_phenotype()
    assert reference.environment_response == 1.0
    assert reference.n_samples is None
    assert [gap.field for gap in reference.provenance_gaps] == [
        "n_samples",
        "sample_unit",
        "environment_response_se",
        "environment_response_uncertainty",
        "environment_response_uncertainty_type",
    ]
    # the released noise floor is recorded and deliberately NOT stored as the spread
    assert h.NOISE_FLOOR.value == 0.0825
    assert "0.08254" in str(h.NOISE_FLOOR.note)


def test_the_environment_is_lb_with_its_selection_and_the_iptg_inducer() -> None:
    environment = h.screen_environment()
    assert environment.temperature is not None
    assert environment.temperature.value == 37.0
    assert environment.duration_generations == 10.0
    assert environment.aerobicity == "aerobic"
    (inducer,) = environment.perturbations
    assert isinstance(inducer, SmallMoleculePerturbation)
    assert inducer.compound.name == "IPTG"
    assert inducer.concentration is not None
    assert (inducer.concentration.value, inducer.concentration.unit) == (
        1.0,
        ConcentrationUnit.millimolar,
    )
    roles = {
        component.compound.name: component.role
        for component in environment.media.components
    }
    assert roles["ampicillin"] == MediaComponentRole.selection_agent
    amounts = {
        component.compound.name: component.concentration
        for component in environment.media.components
    }
    assert amounts["ampicillin"] is not None
    assert amounts["ampicillin"].value == 100.0
    assert amounts["ampicillin"].unit == ConcentrationUnit.ug_per_ml
    # the formulation is unstated, so every LB ingredient carries no amount
    assert all(
        amounts[name] is None
        for name in ("tryptone", "yeast extract", "sodium chloride")
    )
    assert environment.media.base_medium == "LB"


def test_the_host_background_is_bw25113_with_the_tn7att_dcas9_cassette() -> None:
    background = h.host_background()
    assert background.name == "CAG78830"
    assert background.reference_strain == "BW25113"
    assert background.assembly_set == "ecoli_K12_BW25113_ASM75055v1"
    assert background.alleles == []
    assert background.genotype_statement is not None
    assert "Tn7att::PlLac-O1-dcas9(Gent)" in background.genotype_statement
    assert background.provenance is not None
    assert [value.provenance.citation_key for value in background.provenance] == [
        h.CITATION_KEY
    ] * 3


# --------------------------------------------------------------------------- #
# The retention arithmetic
# --------------------------------------------------------------------------- #
def _accounting(**overrides: Any) -> h.BuildAccounting:
    fields: dict[str, Any] = {
        "dataset": "mismatch_crispri_fitness_hawkins2020",
        "released_rows": h.EXPECTED_SOURCE_ROWS,
        "released_genes": 318,
        "released_parent_spacers": 3110,
        "kept_records": h.EXPECTED_RECORDS,
        "dropped_rows": sum(h.EXPECTED_DROPS.values()),
        "dropped_rows_by_reason": dict(h.EXPECTED_DROPS),
        "stored_genes": 317,
        "stored_parent_spacers": 2948,
        "stored_fully_complementary": 2378,
        "stored_single_mismatch": h.EXPECTED_RECORDS - 2378,
        "stored_without_sd": 1870,
        "stored_in_excluded_series": 4863,
        "stored_negative": 778,
        "unresolved_genes": {"b4700": "retired"},
        "reconciliation": _reconciliation(),
    }
    fields.update(overrides)
    return h.BuildAccounting(**fields)


def _reconciliation() -> Any:
    """A minimal reconciliation report: 317 current names and one retired."""
    from torchcell.datasets.bacteria_common import LocusTagReconciliation
    from torchcell.sequence.genome.base import GeneNameStatus

    return LocusTagReconciliation(
        label="hawkins2020 b-numbers",
        assembly_set="ecoli_K12_MG1655_ASM584v2",
        gene_namespace="ecoli_k12_mg1655_bnumber",
        unique_names=318,
        status_histogram={GeneNameStatus.CURRENT: 317, GeneNameStatus.RETIRED: 1},
        layer_histogram={"locus tag": 317, "not found": 1},
        remapped=0,
        kept_on_collision=(),
        retired_kept=(h.UNRESOLVED_BNUMBER,),
        ambiguous_kept={},
        case_insensitive=(),
        outside_namespace=(),
    )


def test_the_accounting_balances_and_refuses_a_moved_count() -> None:
    _accounting().check()
    with pytest.raises(RuntimeError, match="!= 36290 released rows"):
        _accounting(kept_records=h.EXPECTED_RECORDS - 1).check()
    with pytest.raises(RuntimeError, match="per-reason drops total"):
        _accounting(
            dropped_rows_by_reason={h.DROP_NO_FITNESS: 12104, h.DROP_NO_LOCUS: 36}
        ).check()
    with pytest.raises(RuntimeError, match="are not the pinned"):
        _accounting(
            dropped_rows=sum(h.EXPECTED_DROPS.values()),
            dropped_rows_by_reason={h.DROP_NO_FITNESS: 12141, h.DROP_NO_LOCUS: 0},
        ).check()
    with pytest.raises(RuntimeError, match="mismatched !="):
        _accounting(stored_single_mismatch=0).check()


# --------------------------------------------------------------------------- #
# The real release (data-gated)
# --------------------------------------------------------------------------- #
@pytest.mark.data
def test_the_released_workbook_is_the_pinned_bytes_and_every_quote_is_verbatim() -> (
    None
):
    library = h.library_dir(os.environ["DATA_ROOT"]) / h.TABLE_S3_LIBRARY_REL
    if not library.exists():
        pytest.skip("the literature mirror is not mounted")
    assert h._sha256(library) == h.TABLE_S3_SHA256
    paper = h.library_dir(os.environ["DATA_ROOT"]) / h.PAPER_MD
    assert h._sha256(paper) == h.PAPER_MD_SHA256
    text = paper.read_text()
    missing = [
        name for name, value in h.SOURCED_VALUES.items() if value.quote not in text
    ]
    assert missing == []


@pytest.mark.data
def test_the_release_measures_what_the_module_declares() -> None:
    library = h.library_dir(os.environ["DATA_ROOT"]) / h.TABLE_S3_LIBRARY_REL
    if not library.exists():
        pytest.skip("the literature mirror is not mounted")
    frame = h.read_ecoli_sheet(library)
    assert len(frame) == 36290
    assert frame[h.COL_LOCUS].nunique() == 318
    parents = frame[h.COL_VARIANT] == frame[h.COL_PARENT]
    assert int(parents.sum()) == h.EXPECTED_PARENT_SPACERS
    # the predicted column is the model's ACTIVITY: exactly 1.0 on every parent spacer
    assert set(frame.loc[parents, h.COL_PREDICTED]) == {1.0}
    assert float(frame.loc[~parents, h.COL_PREDICTED].min()) < 0.0
    mean = frame[h.COL_MEAN].notna()
    assert int((~mean).sum()) == h.EXPECTED_DROPS[h.DROP_NO_FITNESS]
    sokE = frame[h.COL_LOCUS] == h.UNRESOLVED_BNUMBER
    assert int((sokE & mean).sum()) == h.EXPECTED_DROPS[h.DROP_NO_LOCUS]
    kept = mean & ~sokE
    assert int(kept.sum()) == h.EXPECTED_RECORDS
    assert int((kept & (frame[h.COL_MEAN] < 0)).sum()) == 778
    assert int((kept & frame[h.COL_SD].isna()).sum()) == 1870
    assert int((kept & ~frame[h.COL_FAMILY].astype(bool)).sum()) == 4863


@pytest.mark.data
def test_the_late_window_carries_the_early_windows_dispersion() -> None:
    """The measurement that refuses the 10-to-15-doubling sheet."""
    library = h.library_dir(os.environ["DATA_ROOT"]) / h.TABLE_S3_LIBRARY_REL
    if not library.exists():
        pytest.skip("the literature mirror is not mounted")
    sheets = pd.read_excel(library, sheet_name=None, engine="openpyxl")
    early = sheets[h.SHEET_ECOLI]
    late = sheets[h.SHEET_ECOLI_LATE]
    both = early[h.COL_SD].notna() & late[h.COL_SD].notna()
    assert int(both.sum()) == 22313
    assert early.loc[both, h.COL_SD].equals(late.loc[both, h.COL_SD])
    controls = sheets[h.SHEET_ECOLI_CONTROLS][h.COL_MEAN]
    late_controls = sheets[h.SHEET_ECOLI_LATE_CONTROLS][h.COL_MEAN]
    assert float(controls.std()) == pytest.approx(h.NOISE_FLOOR.value, abs=5e-5)
    assert float(late_controls.std()) > 2.0 * float(controls.std())


@pytest.mark.data
def test_the_built_store_holds_the_declared_records() -> None:
    from torchcell.verification.runners import stream_records

    root = osp.join(
        os.environ["DATA_ROOT"], "data/torchcell/mismatch_crispri_fitness_hawkins2020"
    )
    if not osp.exists(osp.join(root, "preprocess", "build_manifest.json")):
        pytest.skip("the dev store is absent or mid-rebuild")
    summary = h.summarize_store(stream_records(root))
    assert summary.n_records == h.EXPECTED_RECORDS
    assert len(summary.loci) == h.EXPECTED_GENES
    assert summary.n_spacers == h.EXPECTED_RECORDS
    assert summary.malformed_spacers == ()
    assert summary.designs_without_a_parent == ()
    assert summary.derived_from_a_b_number == h.EXPECTED_RECORDS
    assert summary.negative_responses == 778
    assert summary.screens == {h.SCREEN_ID: h.EXPECTED_RECORDS}


# --------------------------------------------------------------------------- #
# The whole build, hermetically: mirror -> raw/ -> LMDB -> L0-L4
# --------------------------------------------------------------------------- #
def _pin(strain: Any, *, background: Any = None, data_root: str | None = None) -> Any:
    """``assembly_reference`` without the genomes tier: the pin from the constants."""
    from torchcell.datamodels.schema import (
        ASSEMBLY_SET_ACCESSIONS,
        BACTERIAL_ASSEMBLY_SETS,
        AssemblyReferenceGenome,
    )

    assembly_set = cast(Any, BACTERIAL_ASSEMBLY_SETS[strain])
    return AssemblyReferenceGenome(
        species="Escherichia coli",
        strain=strain if background is None else background.name,
        assembly_set=assembly_set,
        assembly_accession=ASSEMBLY_SET_ACCESSIONS[assembly_set][0],
        background=background,
    )


@pytest.fixture
def mirrored(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    mg1655: EcoliK12MG1655Genome,
    bw25113: EcoliK12BW25113Genome,
) -> Path:
    """A whole synthetic world: the library mirror, the pins, and the two genomes.

    ``DATA_ROOT`` points at ``tmp_path``, the workbook's own digest and byte count
    replace the released ones, and ``bacterial_genome`` hands back the synthetic
    assemblies, so the build reads nothing outside ``tmp_path``.
    """
    data_root = tmp_path / "root"
    library = data_root / "torchcell-library" / h.CITATION_KEY / "si"
    library.mkdir(parents=True)
    workbook = _workbook(library / "si4.xlsx")
    monkeypatch.setenv("DATA_ROOT", str(data_root))
    monkeypatch.setattr(h, "TABLE_S3_SHA256", h._sha256(workbook))
    monkeypatch.setattr(h, "TABLE_S3_BYTES", workbook.stat().st_size)
    monkeypatch.setattr(h, "EXPECTED_SOURCE_ROWS", len(SYNTHETIC))
    monkeypatch.setattr(h, "EXPECTED_RELEASED_GENES", 2)
    monkeypatch.setattr(h, "EXPECTED_RECORDS", len(STORED))
    monkeypatch.setattr(h, "EXPECTED_GENES", 1)
    monkeypatch.setattr(h, "EXPECTED_STORED_PARENTS", 1)
    monkeypatch.setattr(h, "EXPECTED_NEGATIVE", 1)
    monkeypatch.setattr(h, "EXPECTED_WITHOUT_SD", 1)
    monkeypatch.setattr(h, "EXPECTED_DROPS", {h.DROP_NO_FITNESS: 1, h.DROP_NO_LOCUS: 1})
    genomes = {"MG1655": mg1655, "BW25113": bw25113}
    monkeypatch.setattr(
        h, "bacterial_genome", lambda host, strain, data_root=None: genomes[str(strain)]
    )
    monkeypatch.setattr(h, "assembly_reference", _pin)
    return data_root


def _build(root: Path) -> h.MismatchCrispriFitnessHawkins2020Dataset:
    """Deposit the raw mirror and build the store under ``root``."""
    h.deposit_raw_mirror()
    return h.MismatchCrispriFitnessHawkins2020Dataset(root=str(root / "store"))


def test_the_deposit_records_the_publisher_retrieval_and_refuses_other_bytes(
    mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    mirror = h.deposit_raw_mirror()
    manifest = h.load_manifest()
    (record,) = manifest.files
    assert record.path == h.TABLE_S3_REL
    assert record.sha256 == h.TABLE_S3_SHA256
    assert record.retrieval is not None
    assert record.retrieval.source_url == h.TABLE_S3_URL
    assert record.retrieval.params == {
        "pii": h.TABLE_S3_PII,
        "filename": h.TABLE_S3_PUBLISHER_FILENAME,
    }
    assert manifest.doi == h.DOI
    assert any("10-15 relfit (eco)" in line for line in manifest.si_expected)
    assert h.manifest_sha256(manifest, h.TABLE_S3_REL) == h.TABLE_S3_SHA256
    with pytest.raises(KeyError, match="not in the raw-mirror manifest"):
        h.manifest_sha256(manifest, "data/absent.xlsx")
    # a second deposit is idempotent, and a changed pin refuses the mirrored bytes
    assert h.deposit_raw_mirror() == mirror
    monkeypatch.setattr(h, "TABLE_S3_SHA256", "0" * 64)
    with pytest.raises(RuntimeError, match="is not the pinned"):
        h.deposit_raw_mirror()


def test_the_build_writes_one_record_per_measured_sgrna_and_its_ledger(
    mirrored: Path,
) -> None:
    dataset = _build(mirrored)
    assert len(dataset) == len(STORED)
    accounting = h.BuildAccounting.model_validate_json(
        Path(dataset.preprocess_dir, "build_accounting.json").read_text()
    )
    accounting.check()
    assert accounting.released_rows == len(SYNTHETIC)
    assert accounting.kept_records == len(STORED)
    assert accounting.dropped_rows_by_reason == {
        h.DROP_NO_FITNESS: 1,
        h.DROP_NO_LOCUS: 1,
    }
    assert accounting.stored_fully_complementary == 1
    assert accounting.stored_single_mismatch == 2
    assert accounting.stored_in_excluded_series == 1
    assert accounting.unresolved_genes == {
        "b0005": accounting.unresolved_genes["b0005"]
    }
    assert "no one-to-one ECK partner" in accounting.unresolved_genes["b0005"]
    assert any("mismatch_design" not in note for note in accounting.notes)
    assert any("predicted sgRNA ACTIVITY" in note for note in accounting.notes)

    ledger = pd.read_csv(Path(dataset.preprocess_dir, "guide_retention.csv"))
    assert len(ledger) == len(SYNTHETIC)
    assert "predicted_sgrna_activity_not_stored" in ledger.columns
    dropped = ledger[ledger["drop_reason"].notna()]
    assert set(dropped["spacer"]) == {NO_MEAN, UNRESOLVED}
    # the refused dose survives in the ledger for every row, kept or dropped
    assert (
        float(
            ledger.loc[
                ledger["spacer"] == PARENT, "predicted_sgrna_activity_not_stored"
            ].iloc[0]
        )
        == 1.0
    )


def test_the_build_refuses_a_record_count_that_moved(
    mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(h, "EXPECTED_RECORDS", len(STORED) + 1)
    with pytest.raises(RuntimeError, match="records written, not the pinned"):
        _build(mirrored)


def test_the_build_refuses_a_released_gene_count_that_moved(
    mirrored: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(h, "EXPECTED_RELEASED_GENES", 3)
    with pytest.raises(RuntimeError, match="released genes, not the pinned"):
        _build(mirrored)


def test_every_l0_to_l4_row_passes_on_the_built_store(mirrored: Path) -> None:
    dataset = _build(mirrored)
    report = h.verify_build(dataset.root, expected_count=len(STORED))
    failed = [result.name for result in report.results if not result.passed]
    assert failed == []
    names = {result.name for result in report.results}
    assert {
        "reference_zero",
        "stored_loci_are_bw25113_loci_derived_from_the_released_b_number",
        "guide_spacers_are_twenty_nt_acgt",
        "perturbation_descriptions_state_the_mismatch_design",
        "one_screen_and_the_negative_tail_is_intact",
    } <= names
    row = next(r for r in report.results if r.name == "reference_zero")
    assert row.details["rule"] == "ratio_reference"
    assert Path(dataset.root, "preprocess", "verification_report.json").exists()


def test_the_supplementary_rows_fail_on_a_store_that_lost_its_design(
    mirrored: Path, bw25113: EcoliK12BW25113Genome
) -> None:
    dataset = _build(mirrored)
    from torchcell.verification.runners import stream_records

    records = list(stream_records(dataset.root))
    summary = h.summarize_store(records)
    assert summary.n_records == len(STORED)
    assert summary.loci == ("BW25113_0001",)
    assert summary.derived_from_a_b_number == len(STORED)
    assert summary.negative_responses == 1
    assert h.one_screen_and_the_depleted_tail(summary).passed

    stripped = summary.model_copy(
        update={
            "designs_without_a_parent": (MISMATCHED,),
            "malformed_spacers": ("ACGT",),
        }
    )
    assert not h.designs_name_their_parent(stripped).passed
    assert not h.spacers_are_twenty_nt(stripped).passed
    moved = summary.model_copy(update={"negative_responses": 0})
    assert not h.one_screen_and_the_depleted_tail(moved).passed
    outside = summary.model_copy(update={"loci": ("BW25113_9999",)})
    assert not h.every_locus_is_derived_and_pinned(outside, bw25113).passed
